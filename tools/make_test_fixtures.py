"""產生 tests/parity/fixtures 的兩個快照（4-2a F4／F5）。用法：
    python3 tools/make_test_fixtures.py <LibreOffice 重算後的 xlsx>
輸出：
  tests/parity/fixtures/contract_snapshot_v13.0.json   下游契約名稱（契約文件所列＋全部 IF_／IFW_／IFC_／L1_／GOV_／CHK_／CAP_）逐格值
  tests/parity/fixtures/dead_unmarked_baseline_v10.3.json   死格中「用途」欄為空者（基準名單；之後新增的死格必須有用途標記）
之後每版若有預期變動，必須在同一個 PR 重跑本工具更新快照，並在報告列出差異（舊值、新值、理由）。"""
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import range_boundaries, get_column_letter as L

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "tests" / "parity"))
import deadcells_lib  # noqa: E402

FIX = REPO / "tests" / "parity" / "fixtures"
CONTRACT = REPO / "docs" / "plan" / "Tokenomics_downstream_contract.md"
PURPOSE = "用途"


def name_cells(attr):
    """具名範圍文字 → [(sheet, coordinate)]；非儲存格參照（常數）→ None。"""
    out = []
    for part in attr.lstrip("=").split(","):
        if "!" not in part:
            return None
        sh, rg = part.rsplit("!", 1)
        sh = sh.strip("'").replace("''", "'")
        try:
            c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
        except Exception:
            return None
        if None in (c1, r1, c2, r2):
            return None
        out += [(sh, f"{L(c)}{r}") for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
    return out


def jsonable(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v
    if v is None:
        return None
    return str(v)


def contract_snapshot(path):
    wv = openpyxl.load_workbook(path, data_only=True)
    wf = openpyxl.load_workbook(path)
    names = {n: d.attr_text for n, d in wf.defined_names.items()}
    chosen = sorted(deadcells_lib.contract_names(names, CONTRACT))
    snap, ncell = {}, 0
    for n in chosen:
        cells = name_cells(names[n])
        if not cells:
            continue
        vals = [jsonable(wv[s][a].value) if s in wv.sheetnames else None for s, a in cells]
        snap[n] = {"ref": names[n], "values": vals}
        ncell += len(vals)
    return {"model": Path(path).name, "names": snap, "n_names": len(snap), "n_cells": ncell}


def dead_unmarked(path):
    dead, _, _ = deadcells_lib.dead_formula_cells(path, CONTRACT)
    wb = openpyxl.load_workbook(path)
    pcol = {}
    for ws in wb.worksheets:
        for c in range(3, ws.max_column + 2):
            if ws.cell(3, c).value == PURPOSE:
                pcol[ws.title] = c
    unmarked = []
    for s, a in sorted(dead):
        ws = wb[s]
        col = pcol.get(s)
        if not (col and ws.cell(ws[a].row, col).value not in (None, "")):
            unmarked.append(f"{s}!{a}")
    return {"model": Path(path).name, "n_dead": len(dead), "unmarked": unmarked}


if __name__ == "__main__":
    p = Path(sys.argv[1])
    FIX.mkdir(parents=True, exist_ok=True)
    cs = contract_snapshot(p)
    (FIX / "contract_snapshot_v13.0.json").write_text(json.dumps(cs, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("契約名稱快照：名稱", cs["n_names"], "格", cs["n_cells"])
    du = dead_unmarked(p)
    (FIX / "dead_unmarked_baseline_v10.3.json").write_text(json.dumps(du, ensure_ascii=False, indent=0), encoding="utf-8")
    print("死格", du["n_dead"], "其中用途欄為空", len(du["unmarked"]))
