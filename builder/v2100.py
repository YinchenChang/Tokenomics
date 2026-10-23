# v21.0 (work order docs/workorders/20261023_v21.0_硬體取數與NonNV.md r0, branch claude/v21-0): Interface O (hardware 取數, R1) and P (non-NVIDIA ratios, R5).
# MAJOR version: new contract names only (13): IF_RackPrice, IF_ScaleOutShare, IF_StoragePerGPU (1x15 each) and IF_NonNV_Out_{AMD,TPU,Trn,CS,OAI}, IF_NonNV_Hold_{...} (1x3 each).
#   Each is a reference to an existing cell; no existing named value or formula cell changes value.
#   W1  Interface section O: IF_RackPrice (Spec_Rack C:G rows 11/12/13), IF_ScaleOutShare (Inputs D13:F13 x5), IF_StoragePerGPU (Inputs D14:F14 x5).
#   W2  Interface section P: IF_NonNV_Out_* (NonNV B:D rows 5-9), IF_NonNV_Hold_* (NonNV E:G rows 5-9); NonNV M5:M9 text.
#   W3  (outside the workbook) contract v0.24, CLAUDE.md section 4 (R3).
#   W4  L1 row 93 old-basis "L_total" text -> L1_LTotalFleetVDisc (text only); Decisions SL57-01; README.
# Every write to an Excel-owned cell is guarded (old text / already-applied test).
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

from common import put, section
import v1410
import v2010

VERSION = "20261023_Tokenomics_v21.0"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-23"
V = "v21.0"
WO = "工作單 docs/workorders/20261023_v21.0_硬體取數與NonNV.md（r0）"
DEC_ID = "SL57-01"
DEC_BY = ("Andy 2026-10-23 13:17「先做 Tokenomics V 二十一」；2026-10-23 12:16「你這樣的規劃，我是認同的」（資料權威歸屬）；2026-10-23 12:08 R4／R5；"
          "依據：《情境卡產生規則》R1.6.1 第 13 條；交流/20261023_算力收支-to-AI半導體_同步請求回覆.md（R1、R3）與 交流/20261023_算力收支-to-AI半導體_R4R5與權威歸屬.md（R5）")

O_TITLE = ("O. 硬體取數（v21.0 R1）：供 AI半導體 改接 company-models/markets/chip_bridge 的 In 頁抄錄值；權威在 Tokenomics（R1.6.1 第 13 條）。"
           "欄序同其他 IF_：C–E Hopper、F–H GB200、I–K GB300、L–N VR200、O–Q Rubin Ultra；每世代內 低成本／基準／高成本。每列只是既有格的引用，不改任何既有格")
P_TITLE = ("P. 非 NVIDIA 比例（v21.0 R5；相對 VR200，同層級、同 SLO；全部 Assumed）：每個名稱 3 格（C＝低、D＝基準、E＝高）。"
           "ASIC 每顆規格（每 GW 顆數、單價、HBM）不在 Tokenomics，歸 AI半導體；昇騰、Maia 待證據後於 v21.1 加入")

RACK_LAB = ("機架價格（買方取得價，CSP／hyperscaler 口徑；涵蓋範圍＝Spec_Rack 第 14 列：Hopper＝8 GPU 伺服器＋CPU、NIC、機內儲存（HGX ×4＝32 GPU 為一架）；"
            "GB200／GB300＝運算匣＋NVLink 交換匣＋機架內冷卻；VR200＝運算匣＋交換匣＋冷卻＋約 $1M NAND，是否含部分網路待查；"
            "Rubin Ultra＝72 封裝單架，價格為 VR200 機架價 × 倍數（Assumed，無公開報價）；"
            "標記沿用 Spec_Rack 第 17 列：Analogy／Interested-party／Derived 與 Interested-party／Interested-party／Derived 與 Assumed；"
            "低成本／基準／高成本＝Spec_Rack 第 11／12／13 列）　[IF_RackPrice]")
SCALE_LAB = ("Scale-out 網路 ÷ 機架價格（來源 Inputs 第 13 列，Analogy；來源不分世代，各世代相同；低成本 0.05／基準 0.10／高成本 0.15，"
             "低成本／基準／高成本＝Inputs D13／E13／F13；比例單位＝÷ 機架價）　[IF_ScaleOutShare]")
STOR_LAB = ("儲存與管理設備（來源 Inputs 第 14 列，Assumed；每 GPU；來源不分世代，各世代相同；低成本 300／基準 500／高成本 800，"
            "低成本／基準／高成本＝Inputs D14／E14／F14）　[IF_StoragePerGPU]")

# (stem, NonNV row, candidate display name for the label)
NONNV = [("AMD", 5), ("TPU", 6), ("Trn", 7), ("CS", 8), ("OAI", 9)]
NONNV_FULL = {"AMD": "AMD MI455X（Helios）", "TPU": "Google TPU v7 Ironwood", "Trn": "AWS Trainium3", "CS": "Cerebras WSE-3", "OAI": "OpenAI／Broadcom 客製"}

NEW_NAMES = ["IF_RackPrice", "IF_ScaleOutShare", "IF_StoragePerGPU"] + [f"IF_NonNV_{k}_{s}" for k in ("Out", "Hold") for s, _ in NONNV]

NONNV_M_OLD = "未接線：B{r}:J{r}（MAG、智譜、AI 半導體待用；目前不進任何輸出）"
NONNV_M_NEW = "已接 Interface P 節（v21.0）：IF_NonNV_Out_{s}／IF_NonNV_Hold_{s}；H:J 欄仍未接線"

L93_REPL = [      # (tuple index in the L1 row, old substring, new substring) — row 93 = PriceReqBE; indices follow the sheet columns C (2), I (8), J (9), K (10)
    (2, "÷ L_total；v8.0 L4", "÷ L1_LTotalFleetVDisc（主口徑 L）；v8.0 L4"),
    (8, "× 單價 ÷ L_total。判讀", "× 單價 ÷ L1_LTotalFleetVDisc（主口徑 L）。判讀"),
    (9, "生產折減、L_total", "生產折減、L1_LTotalFleetVDisc"),
    (10, "L_total：Assumed（g、T1、ρ）", "L1_LTotalFleetVDisc：Assumed（g、T1、ρ）"),
]


# ================================================================== W1／W2: Interface sections O and P (called after v2000.interface_n_ext, before dependency tracing)
def interface_op(wb):
    ws = wb["Interface"]
    last = max(c.row for row in ws.iter_rows() for c in row if c.value is not None)
    assert "IF_LTotalFleetVDisc" in str(ws[f"A{last}"].value) or "IF_TokGWPaidMix" in str(ws[f"A{last}"].value), ws[f"A{last}"].value      # v19.0 N is the last section
    start = last + 2
    section(ws, start, O_TITLE, 17)
    made = []

    def name_it(name, ref):
        wb.defined_names[name] = DefinedName(name, attr_text=ref)

    r = start + 1
    # IF_RackPrice: generation g (Spec_Rack column C..G) -> Interface columns 3g..3g+2 = rows 11 (low) / 12 (base) / 13 (high)
    put(ws, f"A{r}", RACK_LAB); put(ws, f"B{r}", "$M/架")
    for g in range(5):
        for j, srow in enumerate((11, 12, 13)):
            put(ws, f"{L(3 + 3 * g + j)}{r}", f"=Spec_Rack!${L(3 + g)}${srow}", fmt="0.0")
    name_it("IF_RackPrice", f"Interface!$C${r}:$Q${r}"); made.append(("IF_RackPrice", r)); r += 1
    for name, lab, unit, row, fmt in (("IF_ScaleOutShare", SCALE_LAB, "比例（÷ 機架價）", 13, "0.00"), ("IF_StoragePerGPU", STOR_LAB, "$/GPU", 14, "#,##0")):
        put(ws, f"A{r}", lab); put(ws, f"B{r}", unit)
        for g in range(5):
            for j, col in enumerate("DEF"):
                put(ws, f"{L(3 + 3 * g + j)}{r}", f"=Inputs!${col}${row}", fmt=fmt)
        name_it(name, f"Interface!$C${r}:$Q${r}"); made.append((name, r)); r += 1
    r += 1
    section(ws, r, P_TITLE, 17)
    r += 1
    for kind, c0, ulab, fmt in (("Out", "B", "產出比（每 GW 產出 ÷ VR200）", "0.00"), ("Hold", "E", "持有成本比（每 GW 持有成本 ÷ VR200）", "0.00")):
        for s, nrow in NONNV:
            nm = f"IF_NonNV_{kind}_{s}"
            put(ws, f"A{r}", f"{NONNV_FULL[s]}：{ulab}（NonNV 第 {nrow} 列 {'B:D' if kind == 'Out' else 'E:G'}；同層級、同 SLO；Assumed；C＝低、D＝基準、E＝高）　[{nm}]")
            put(ws, f"B{r}", "比（÷ VR200）")
            for j in range(3):
                put(ws, f"{L(3 + j)}{r}", f"=NonNV!${L(ord(c0) - 64 + j)}${nrow}", fmt=fmt)
            name_it(nm, f"Interface!$C${r}:$E${r}"); made.append((nm, r)); r += 1
    return dict(start=start, made=made)


# ================================================================== W2: NonNV M5:M9 text (guarded)
def nonnv_text(wb):
    ws = wb["NonNV"]
    n = 0
    for s, r in NONNV:
        new = NONNV_M_NEW.format(s=s)
        cur = ws[f"M{r}"].value
        if cur == new:
            continue
        assert cur == NONNV_M_OLD.format(r=r), (r, cur)
        ws[f"M{r}"].value = new; n += 1
    return n


# ================================================================== W4: L1 row 93 text (hook in gov._rows_l1, like v2010.l1_patch)
def l1_patch(R):
    out = []
    for row in R:
        row = list(row)
        if row[0] == "PriceReqBE":
            for k, old, new in L93_REPL:
                assert old in row[k], (k, row[k])
                row[k] = row[k].replace(old, new)
                assert "L_total" not in row[k].replace("L1_LTotalFleetVDisc", ""), row[k]
        out.append(tuple(row))
    return out


# ================================================================== Decisions
def _decisions():
    return [[DEC_ID, V, "v21.0 Interface 新增 O 節（硬體取數 R1）與 P 節（非 NVIDIA 比例 R5）；下游契約權威歸屬",
             "W1：IF_RackPrice（Spec_Rack C:G 第 11／12／13 列）、IF_ScaleOutShare（Inputs D13:F13）、IF_StoragePerGPU（Inputs D14:F14）；"
             "W2：IF_NonNV_Out_／IF_NonNV_Hold_ × AMD／TPU／Trn／CS／OAI（NonNV 第 5–9 列），NonNV M5:M9 文字改為已接線（H:J 欄仍未接線）；"
             "W3（活頁簿外）：契約 v0.24 新增「權威歸屬」「證據送件」兩節、CLAUDE.md 第 4 節 R3；W4：L1 第 93 列 L_total 舊口徑文字改為 L1_LTotalFleetVDisc。"
             "既有名稱數值不變；昇騰、Maia 不在本版（待證據後 v21.1）。Andy 原話與依據：" + DEC_BY,
             DATE, DEC_BY, "v21.0 已建", "Interface O、P 節；NonNV M5:M9；L1 第 93 列", f"{WO} W1–W5", "否"]]


def apply(wb, append_text):
    log = {"nonnv_text": nonnv_text(wb)}
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == DEC_ID for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


# ================================================================== README
README_VERSION = (VERSION + "（主版：" + WO + "——Interface 新增 O 節「硬體取數」（IF_RackPrice、IF_ScaleOutShare、IF_StoragePerGPU）與 P 節「非 NVIDIA 比例」"
                  "（IF_NonNV_Out_／IF_NonNV_Hold_ × AMD／TPU／Trn／CS／OAI）共 13 個名稱，皆為既有格的引用；NonNV M5:M9 用途文字；L1 第 93 列 L_total 舊口徑文字改主口徑；Decisions SL57-01。"
                  "既有名稱的值全部不變。）以下為 " + v2010.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v2010.README_TITLE
README_TITLE = V + _T[len(v2010.V):]
README_ROW = ("v21.0 硬體取數與 NonNV",
              "Interface O 節「硬體取數」：IF_RackPrice（機架價 $M/架）、IF_ScaleOutShare、IF_StoragePerGPU，供 AI半導體 改接 chip_bridge In 頁抄錄值（權威在 Tokenomics）。"
              "Interface P 節「非 NVIDIA 比例」：IF_NonNV_Out_／IF_NonNV_Hold_ × AMD／TPU／Trn／CS／OAI（相對 VR200，Assumed；低／基準／高）。"
              "ASIC 每顆規格不在 Tokenomics；昇騰、Maia 待 v21.1。細節見工作單 docs/workorders/20261023_v21.0_硬體取數與NonNV.md 與 Decisions SL57-01。")
