"""第 1 段輸入判讀（嚴重度 高／中／低 共 186 格）逐格比對：v5.36（稽核時）對現行活頁簿。
輸出 s1_cells.json：值是否改變、Gov_Map 登錄（類別／SRC／決策 ID／低高）前後。"""
import openpyxl, os, re, json, sys
from openpyxl.utils import range_boundaries, get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import lib
ROOT = lib.ROOT
OLD = os.path.join(ROOT, 'model', 'archive', '20261015_Tokenomics_v5.36.xlsx')
REP = os.path.join(ROOT, 'docs', 'reports')

old_v = openpyxl.load_workbook(OLD, data_only=True)
cur_v = lib.val()


def gov_index(wb):
    ws = wb['Gov_Map']
    idx = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[0]:
            continue
        sh, rng = r[2], r[3]
        if not sh or not rng:
            continue
        try:
            c1, r1, c2, r2 = range_boundaries(str(rng).replace('$', ''))
        except Exception:
            continue
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                idx[(sh, f'{get_column_letter(cc)}{rr}')] = dict(gm=r[0], cat=r[5], src=r[7], dec=r[9], lo=r[10], hi=r[11], rng=r[12], seg=r[15])
    return idx


gi_old = gov_index(old_v)
gi_cur = gov_index(cur_v)

wb = openpyxl.load_workbook(os.path.join(REP, '20261015_稽核1_明細.xlsx'), read_only=True)
rows = list(wb['輸入判讀'].iter_rows(values_only=True))
h = list(rows[1])
ix = {k: i for i, k in enumerate(h)}
out = []
for r in rows[2:]:
    sev = r[ix['嚴重度']]
    if sev not in ('高', '中', '低'):
        continue
    sh, cl = r[ix['頁']], r[ix['格']]
    def gv(wbk, sh, cl):
        try:
            return wbk[sh][cl].value
        except Exception:
            return 'ERR'
    vo, vc = gv(old_v, sh, cl), gv(cur_v, sh, cl)
    go, gc = gi_old.get((sh, cl)), gi_cur.get((sh, cl))
    out.append(dict(gm=r[ix['GM_ID']], sheet=sh, cell=cl, label=r[ix['列標籤']], sev=sev, cls=r[ix['問題類別']],
                    reason=r[ix['理由']], m_path=r[ix['是否在 m 路徑']], old_val=vo, cur_val=vc, gov_old=go, gov_cur=gc))
json.dump(out, open(os.path.join(HERE, 's1_cells.json'), 'w'), ensure_ascii=False, default=str)
print(len(out))
import collections
print(collections.Counter((o['sev'], (o['old_val'] == o['cur_val'])) for o in out))
for o in out:
    if o['sev'] == '高':
        go, gc = o['gov_old'] or {}, o['gov_cur'] or {}
        print(o['sheet'], o['cell'], '| val', str(o['old_val'])[:8], '->', str(o['cur_val'])[:8], '| cat', go.get('cat'), '->', gc.get('cat'), '| src', go.get('src'), '->', gc.get('src'), '| dec', go.get('dec'), '->', gc.get('dec'), '| rng', go.get('lo'), go.get('hi'), '->', gc.get('lo'), gc.get('hi'))
