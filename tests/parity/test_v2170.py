"""v21.7 專屬測試（工作單 docs/workorders/20261025_v21.7_NonNV_AMD與Trn同源價.md r0）。讀 LibreOffice 重算後的值。

  1. 具名範圍的值與封存 v21.6 相同，只有 IF_NonNV_Out_AMD、IF_NonNV_Hold_AMD（各 3 格）、IF_NonNV_Hold_Trn（基準 1 格）與版本字串 RNG_Version／RNG_Current 例外；名稱沒有增減。
  2. W1：NonNV B5:D5 ＝ 0.70／0.80／1.15、E5:G5 ＝ 0.67／0.85／1.12、F7 ＝ 0.61；其餘輸入不變；H:J 公式照舊。
  3. W2：NonNV L5／L7／L8／L9／A2、Interface A368 說明更新，原文逐字保留於「｜原：」之後。
  4. W3：DB_Evidence E410／E411、Decisions SL64-01、README B5 開頭。
  5. GOV 三項計數 0／134／182；RNG_StaleFlag 一致；全部公式格無錯誤值。
  6. 逐格比較 v21.6 → v21.7：數值差異只落在工作單第 0 節清單；文字／公式差異只落在 W1–W3 格與版本字串。
"""
import csv
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import range_boundaries

from conftest import model_path
from v2170_allow import NUM_DIFFS_V2170 as NUM_DIFFS      # 工作單第 0 節清單（數值差異；對 v21.6）

REPO = Path(__file__).resolve().parents[2]
V216 = REPO / "model" / "archive" / "20261025_Tokenomics_v21.6.xlsx"
VERSION_NAMES = {"RNG_Version", "RNG_Current"}
CHANGED_NAMES = VERSION_NAMES | {"IF_NonNV_Out_AMD", "IF_NonNV_Hold_AMD", "IF_NonNV_Hold_Trn"}
TAIL = "持有比為慣例值（Block 1『接近 1』），未推導；找到 TPU v7、Trainium3 機架功率後以每 kW 設備價法重推（Andy 2026-10-23 Q2）。"
A2_OLD = "持有成本比預設接近 1，依 Block 1 結論：每 GW 持有成本幾乎不隨世代改變"
A368_OLD = "（第 379–382 列；持有比依每 kW 設備價推導，既有五列為慣例值）"


def _flat(v):
    if isinstance(v, (list, tuple)):
        return [x for r in v for x in _flat(r)]
    return [v]


def _num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def _eq(a, b):
    if a in (None, "") and b in (None, ""):
        return True
    if _num(a) and _num(b):
        return abs(a - b) <= 1e-9 * max(abs(a), abs(b), 1e-300) + 1e-12
    return a == b


def _arch_value(pw, d):
    sh, rg = d.attr_text.rsplit("!", 1)
    c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
    return [pw[sh.strip("'")].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


@pytest.fixture(scope="module")
def wbf():
    return openpyxl.load_workbook(model_path())


@pytest.fixture(scope="module")
def pw():
    return openpyxl.load_workbook(V216, data_only=True)


def _row(ws, sid):
    rr = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == sid]
    assert len(rr) == 1, (sid, rr)
    return rr[0]


# ---------------------------------------------------------------- 1. 具名範圍值
def test_v2170_names_changed_only_expected(lo, pw, wbf):
    v = lo.view({})
    bad, n = [], 0
    for name, d in wbf.defined_names.items():
        if "!" not in d.attr_text or "," in d.attr_text or name not in pw.defined_names:
            continue
        want = _arch_value(pw, pw.defined_names[name])
        got = _flat(v.get_name(name))
        n += 1
        if not (len(want) == len(got) and all(_eq(a, b) for a, b in zip(want, got))):
            bad.append(name)
    assert n > 1500, n
    assert sorted(bad) == sorted(CHANGED_NAMES), bad
    assert set(wbf.defined_names) == set(pw.defined_names)      # 名稱沒有增減
    assert [round(x, 12) for x in _flat(v.get_name("IF_NonNV_Out_AMD"))] == [0.70, 0.80, 1.15]
    assert [round(x, 12) for x in _flat(v.get_name("IF_NonNV_Hold_AMD"))] == [0.67, 0.85, 1.12]
    assert [round(x, 12) for x in _flat(v.get_name("IF_NonNV_Hold_Trn"))] == [0.45, 0.61, 0.70]
    assert [round(x, 12) for x in _flat(v.get_name("IF_NonNV_Hold_TPU"))] == [0.65, 0.70, 0.90]      # TPU 不變
    assert str(v.get_name("RNG_Version")).endswith("v21.7") and str(v.get_name("RNG_Current")).endswith("v21.7")


# ---------------------------------------------------------------- 2. W1 NonNV
def test_v2170_nonnv_rows(lo, wbf, pw):
    nn = wbf["NonNV"]
    assert nn["A5"].value == "AMD MI455X（Helios）" and nn["A7"].value == "AWS Trainium3"
    assert [nn[f"{c}5"].value for c in "BCD"] == [0.7, 0.8, 1.15] and [nn[f"{c}5"].value for c in "EFG"] == [0.67, 0.85, 1.12]
    assert [nn[f"{c}7"].value for c in "EFG"] == [0.45, 0.61, 0.7] and [nn[f"{c}7"].value for c in "BCD"] == [0.3, 0.5, 0.8]      # Trn3 產出比、持有比低／高端不變
    assert [nn[f"{c}6"].value for c in "BCDEFG"] == [0.6, 0.8, 1.1, 0.65, 0.7, 0.9]      # TPU 不變
    for r in (5, 7):
        assert (nn[f"H{r}"].value, nn[f"I{r}"].value, nn[f"J{r}"].value) == (f"=F{r}/C{r}", f"=E{r}/D{r}", f"=G{r}/B{r}") and nn[f"K{r}"].value == "Assumed"
    v = lo.view({})
    assert abs(v.get("NonNV", "H5") - 0.85 / 0.8) < 1e-12 and abs(v.get("NonNV", "I5") - 0.67 / 1.15) < 1e-12 and abs(v.get("NonNV", "J5") - 1.12 / 0.7) < 1e-12
    assert abs(v.get("NonNV", "H7") - 0.61 / 0.5) < 1e-12 and abs(v.get("NonNV", "I7") - 0.45 / 0.8) < 1e-12 and abs(v.get("NonNV", "J7") - 0.7 / 0.3) < 1e-12
    pn = pw["NonNV"]
    for r in (6, 8, 9, 10, 11):      # 其他列的值與 v21.6 相同
        for c in "BCDEFGHIJ":
            assert _eq(pn[f"{c}{r}"].value, v.get("NonNV", f"{c}{r}")), (c, r)
    p6 = openpyxl.load_workbook(V216)["NonNV"]
    assert nn["L6"].value == p6["L6"].value      # L6 不動
    # L5：原文（慣例值尾句）逐字保留於「｜原：」之後
    l5 = nn["L5"].value
    assert l5.startswith("MI355X 在 V4-Pro 上 26 天內吞吐提升 110 倍（S20 所屬部落格），軟體成熟度為主要不確定｜產出比 v21.7 依每 kW 峰值算力比 1.066")
    for key in ("MI455X MXFP4 40.3 PF、TBP 2,500 W、Helios 245 kW", "效率 f 0.656／0.757／0.994", "轉換係數 t 0.891，Assumed", "每架 $5.25M 下界／$7.3M＝AMD 自述隱含比 0.865 × VR200",
                "$10.5M＝Futurum 同源比值 1.25 × VR200", "每 token 成本比基準約 1.05（PR #174）"):
        assert key in l5, key
    assert l5.endswith("｜原：" + TAIL) and l5.count("｜原：") == 1
    # L7：v21.6 說明保留，後加 v21.7 句，原文尾句保留
    l7 = nn["L7"].value
    assert l7.startswith("OpenAI v0.5 以 2GW Trainium 合約為輸入｜持有比 v21.6 依每 kW 設備價推導") and "｜v21.7：持有比基準 0.55 → 0.61，改用 KGI 同源比值" in l7
    assert "Trn3 每架約 $1.6–1.8M → 0.593–0.631 取中點；PR #174 §10" in l7 and l7.endswith("｜原：" + TAIL) and l7.count("｜原：") == 1
    assert l7.index("v21.6 依每 kW") < l7.index("v21.7：持有比基準") < l7.index("｜原：")
    for r, head in ((8, "高互動性利基；每 GW 吞吐低、單用戶速度高"), (9, "無公開規格；沿用 v0.5「自研／其他 0.6（0.4–1.0）」")):
        l = nn[f"L{r}"].value
        assert l == head + "｜TPU v7、Trainium3（v21.6）、AMD（v21.7）已以每 kW 法重推；本列尚無機架功率與價格資料，維持慣例值。｜原：" + TAIL, r
    a2 = nn["A2"].value
    assert a2.startswith("每 GW 產出比與每 GW 持有成本比皆相對 VR200（同層級、同 SLO）。持有成本比：TPU、Trn3、AMD、昇騰、Maia 依每 kW 設備價推導（v21.1、v21.6、v21.7）；Cerebras、OpenAI 客製仍為慣例值（Block 1『接近 1』）")
    assert a2.endswith("｜原：" + A2_OLD) and a2.count("｜原：") == 1
    a368 = wbf["Interface"]["A368"].value
    assert a368.startswith("P. 非 NVIDIA 比例（v21.0 R5；") and "（第 379–382 列）；持有比依每 kW 設備價推導：昇騰、Maia（v21.1）、TPU、Trn3（v21.6）、AMD（v21.7）；Cerebras、OpenAI 客製為慣例值；產出比：AMD、昇騰、Maia 已推導，TPU、Trn3 待精度口徑判斷" in a368
    assert a368.endswith("｜原：" + A368_OLD) and a368.count("｜原：") == 1


def test_v2170_interface_gov_mirrors(lo):
    v = lo.view({})
    got = {f"{c}{r}": v.get("Interface", f"{c}{r}") for r in (369, 374, 376) for c in "CDE"}
    assert got == {"C369": 0.7, "D369": 0.8, "E369": 1.15, "C374": 0.67, "D374": 0.85, "E374": 1.12, "C376": 0.45, "D376": 0.61, "E376": 0.7}
    gm = {c: v.get("Gov_Map", c) for c in ("K389", "L389", "Q389", "K390", "L390", "Q390", "Q391", "Q392", "Q393", "Q394", "Q402")}
    assert gm == {"K389": 0.7, "L389": 1.15, "Q389": 0.8, "K390": 0.67, "L390": 1.12, "Q390": 0.85, "Q391": 0.7, "Q392": 1.15, "Q393": 0.67, "Q394": 1.12, "Q402": 0.61}


# ---------------------------------------------------------------- 3. W3 記錄與文字
def test_v2170_evidence_decisions_readme(wbf):
    ev = wbf["DB_Evidence"]
    r = _row(ev, "E410")
    assert r == 334 and ev.cell(r - 1, 1).value == "E409" and _row(ev, "E411") == 335
    assert ev.cell(r, 2).value == "2026-10-25" and ev.cell(r, 10).value == "v21.7" and ev.cell(r, 12).value == "已處理"
    assert "docs/evidence/20261025_NonNV_AMD/" in ev.cell(r, 4).value and "PR #174" in ev.cell(r, 4).value
    assert ev.cell(r, 5).value == "Interested-party（AMD 規格與自述）＋Analogy（CLSA、Futurum；SemiAnalysis 只作佐證）"
    t = " ".join(str(ev.cell(r, c).value) for c in range(1, 19))
    for key in ("產出比 0.70／0.80／1.15、持有比 0.67／0.85／1.12（採用，v21.7 W1）", "每 token 成本比基準約 1.05（現值 1.29）", "分析師絕對價與 TK 買方取得價不同口徑",
                "同源比值全區間 1.12／1.23／1.35", "只換產出比（成本比約 1.13）", "兩錨點 f（產出比 0.74）"):
        assert key in t, key
    r2 = 335
    assert ev.cell(r2, 2).value == "2026-10-25" and ev.cell(r2, 10).value == "v21.7"
    t2 = " ".join(str(ev.cell(r2, c).value) for c in range(1, 19))
    for key in ("0.55 → 0.61", "KGI 同源比值", "PR #172", "PR #174", "0.593–0.631 取中點"):
        assert key in t2, key
    d = wbf["Decisions"]
    r = _row(d, "SL64-01")
    assert r == 336 and d.cell(r - 1, 1).value == "SL63-01" and d.cell(r, 2).value == "v21.7" and d.cell(r, 7).value == "v21.7 已建"
    txt = d.cell(r, 4).value + d.cell(r, 6).value
    for key in ("0.5／0.7／1.0 → 0.70／0.80／1.15", "0.8／0.9／1.0 → 0.67／0.85／1.12", "0.55 → 0.61", "IF_NonNV_Out_AMD、IF_NonNV_Hold_AMD（各 3 格）、IF_NonNV_Hold_Trn（基準 1 格）", "PR #174", "判斷題 9、10"):
        assert key in txt, key
    rm = wbf["README"]
    b5 = str(rm["B5"].value)
    assert b5.startswith("20261025_Tokenomics_v21.7（次版，底版 v21.6：") and "契約名稱的值改變的只有 IF_NonNV_Out_AMD、IF_NonNV_Hold_AMD（各 3 格）、IF_NonNV_Hold_Trn（基準 1 格）與版本字串" in b5
    assert "以下為 v21.6（次版，底版 v21.5：" in b5 and str(rm["A1"].value).startswith("Tokenomics v21.7")
    assert (REPO / "docs" / "workorders" / "20261025_v21.7_預期變動.csv").exists()


# ---------------------------------------------------------------- 4. GOV
def test_v2170_gov_counts(lo):
    v = lo.view({})
    assert (v.get_name("GOV_Errors"), v.get_name("GOV_Warnings"), v.get_name("GOV_Info")) == (0, 134, 182)
    assert v.get_name("RNG_StaleFlag") == "一致" and v.get_name("RNG_OutCount") == 0
    errs = [(s, c) for (s, c), x in v.evaluate_all().items() if isinstance(x, str) and x.startswith("#")]
    assert errs == []


# ---------------------------------------------------------------- 5. 逐格比較 v21.6 → v21.7
EXACT_CELLS = {
    "README": {"A1", "B5"}, "Alloc": {"C177"}, "Interface": {"A368"},
    "NonNV": {"B5", "C5", "D5", "E5", "F5", "G5", "F7", "L5", "L7", "L8", "L9", "A2"},
    "Decisions": {f"{c}336" for c in "ABCDEFGHIJ"},
    "DB_Evidence": {f"{c}{r}" for r in (334, 335) for c in "ABCDEFGHIJKLMNOPQR"},
}
TEXT_MIRRORS = ({("Gov_Map", f"N{r}") for r in range(389, 395)} | {("Gov_Map", f"N{r}") for r in range(401, 419)})      # NonNV L5／L7／L8／L9 說明的鏡像（Gov_Map N 欄公式 =NonNV!$L$5 等）


def test_v2170_cell_diffs_only_expected():
    a, b = openpyxl.load_workbook(V216), openpyxl.load_workbook(model_path())
    av, bv = openpyxl.load_workbook(V216, data_only=True), openpyxl.load_workbook(model_path(), data_only=True)
    assert a.sheetnames == b.sheetnames
    fd, nd, td = {}, set(), set()
    for ws in a:
        wn, va, vb = b[ws.title], av[ws.title], bv[ws.title]
        for r in range(1, max(ws.max_row, wn.max_row) + 1):
            for c in range(1, max(ws.max_column, wn.max_column) + 1):
                x, y = ws.cell(r, c).value, wn.cell(r, c).value
                if x != y:
                    fd.setdefault(ws.title, set()).add(ws.cell(r, c).coordinate)
                p, q = va.cell(r, c).value, vb.cell(r, c).value
                if _num(p) and _num(q) and not _eq(p, q):
                    nd.add((ws.title, ws.cell(r, c).coordinate))
                elif not (_num(p) and _num(q)) and p != q:
                    td.add((ws.title, ws.cell(r, c).coordinate))
    assert fd == EXACT_CELLS, {k: sorted(fd.get(k, set()) ^ EXACT_CELLS.get(k, set())) for k in set(fd) | set(EXACT_CELLS) if fd.get(k, set()) != EXACT_CELLS.get(k, set())}
    assert nd == NUM_DIFFS, sorted(nd ^ NUM_DIFFS)      # 數值差異只落在工作單第 0 節清單
    text_expected = {(s, c) for s, cs in EXACT_CELLS.items() for c in cs} | TEXT_MIRRORS | {("Alloc", "C178")}
    assert td - text_expected == set(), sorted(td - text_expected)
    for sh, cells in fd.items():      # 公式文字不變：改的只有常數與文字
        for coord in cells:
            x, y = a[sh][coord].value, b[sh][coord].value
            assert not (isinstance(x, str) and x.startswith("=")) or sh in ("README", "Alloc"), (sh, coord)
    # 預期變動 CSV 列出全部數值差異格
    rows = list(csv.reader(open(REPO / "docs" / "workorders" / "20261025_v21.7_預期變動.csv", encoding="utf-8-sig")))
    listed = {(r[0], r[1]) for r in rows[1:]}
    assert NUM_DIFFS <= listed, sorted(NUM_DIFFS - listed)
