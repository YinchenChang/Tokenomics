"""O4 查證（唯讀）：在 Python 中逐式複製 Calib L 節（第 209–254 列）的公式，
把「達標所需批次口徑 η_d 倍數」算到 VR200 欄（以及 v6.1 的 GB200／GB300 對照），
目標值換成 MLPerf v6.1 offline 每 GPU 實測。所有物理量與規格一律從現行活頁簿讀取（不硬編碼）；
只有 MLPerf 目標值寫在下面的 TARGETS（出處見 01_scan_輸出.txt）。
本檔不寫回 Excel，只輸出候選數字供 Project 判斷。

先用 v6.0 的原目標（SRC_PERF_079 6,752；SRC_PERF_080 8,987.2）重算 GB200／GB300，
必須與 Calib C254＝1.8297、D254＝1.5964 相同，才表示複製正確。
用法：python3 -I 02_vr200_batch_eta.py <model xlsx> > 02_輸出.txt
"""
import sys, math
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
C, S, R = wb['Calib'], wb['Serving'], wb['Spec_Rack']
col = {'Hopper': 3, 'GB200': 4, 'GB300': 5, 'VR200': 6, 'RubinUltra': 7}

A = C['C101'].value; L = C['C102'].value; dm = C['C103'].value; ke = C['C105'].value
KVb = C['C108'].value; attF = C['C109'].value; ISL = C['C110'].value; OSL = C['C111'].value
Pne = C['C113'].value; Pe = C['C114'].value
hbm_res = S['C9'].value; ne_cap = S['C10'].value; e_cap = S['C11'].value
ep_bytes = S['C12'].value; eta_mem = S['C13'].value; ep_eff = S['C14'].value


def need_mult(gen, target, isl=ISL, osl=OSL):
    c = col[gen]
    peak = R.cell(35, c).value; hbm = R.cell(25, c).value; bw = R.cell(26, c).value
    epbw = R.cell(27, c).value; epbase = R.cell(29, c).value
    eb = R.cell(30, c).value; neb = R.cell(31, c).value
    eta70 = C.cell(70, c).value; lat71 = C.cell(71, c).value; etap = C.cell(72, c).value
    mult = C.cell(156, c).value  # CAL_BatchEtaD 現值
    hbm_u = hbm * (1 - hbm_res)
    tp = max(1, math.ceil(Pne * neb / (ne_cap * hbm_u)))
    ep = max(epbase, math.ceil(Pe * eb / (e_cap * hbm_u)))
    w = Pne * neb / tp + Pe * eb / ep
    ctx = isl + osl / 2
    fl = (2 * A * 1e9 + attF * ctx) / 1e9
    fixed = w / bw + L * lat71 / 1000
    eta = eta70 * mult
    t_cmp = fl / (peak * eta) / 1000
    t_kv = KVb * ctx / (bw * eta_mem) / 1e9
    t_ep = L * ke * dm * ep_bytes * (ep - 1) / ep / (epbw * ep_eff) / 1e9
    t_step = max(t_cmp, t_kv, t_ep)
    B = (hbm_u - w) * 1e9 / (KVb * ctx)
    dec = B / ((fixed + B * t_step) / 1000)
    pre = etap * peak * 1e15 / (((2 * A * 1e9 + attF * isl / 2)) )
    out = osl / (isl / pre + osl / dec)
    need_dec = osl / (osl / target - isl / pre)
    need_t = (1000 * B / need_dec - fixed) / B
    floor = max(t_kv, t_ep)
    sol = (fl / peak / 1000 / need_t) / eta70 if need_t > floor else float('nan')
    return dict(gen=gen, target=target, current_mult=mult, model_out=round(out, 1),
                model_over_target=round(out / target, 4), B=round(B, 1), EP=ep,
                need_mult=round(sol, 4), t_cmp_ms=round(t_cmp, 5), floor_ms=round(floor, 5))


print('== 複製檢查（v6.0 原目標；應等於 Calib C254 1.8297、D254 1.5964）')
print(need_mult('GB200', C['C217'].value), 'Calib C254 =', C['C254'].value)
print(need_mult('GB300', C['D217'].value), 'Calib D254 =', C['D254'].value)

TARGETS = [  # MLPerf v6.1 offline, deepseek-r1, 每 GPU（Result ÷ total_accelerators）；見 01_scan_輸出.txt
    ('VR200', 16435.1, 'NVIDIA Vera Rubin 72 GPU，closed，preview'),
    ('VR200', 15522.3, 'Nebius VR200 36 GPU，closed，preview'),
    ('GB300', 9440.8, 'NVIDIA GB300 72 GPU，closed，available'),
    ('GB300', 9582.8, 'Nebius GB300 72 GPU，closed，available'),
    ('GB200', 7274.9, 'NVIDIA GB200 72 GPU，closed，available'),
    ('GB200', 7528.8, 'Azure GB200 72 GPU，closed，available'),
]
print('== v6.1 offline 目標')
for g, t, note in TARGETS:
    print(note, need_mult(g, t))
