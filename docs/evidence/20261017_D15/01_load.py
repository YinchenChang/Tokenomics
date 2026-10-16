"""D15 第 1 步：唯讀載入 v12.0（openpyxl data_only=True）與稽核第 1 段明細，輸出 01_load.pkl。
用法：python3 01_load.py
不寫 model/；只寫本資料夾。"""
import openpyxl, pickle, re, datetime, os
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..', '..'))
WB = os.path.join(ROOT, 'model', '20261016_Tokenomics_v12.0.xlsx')
A1 = os.path.join(ROOT, 'docs', 'reports', '20261015_稽核1_明細.xlsx')

wb = openpyxl.load_workbook(WB, data_only=True)

def rows(ws, start, ncol):
    out = []
    for r in range(start, ws.max_row + 1):
        out.append([ws.cell(r, c).value for c in range(1, ncol + 1)] + [r])
    return out

gm_ws = wb['Gov_Map']
gm_hdr = [gm_ws.cell(4, c).value for c in range(1, gm_ws.max_column + 1)]
gm = [x for x in rows(gm_ws, 5, gm_ws.max_column) if isinstance(x[0], str) and x[0].startswith('GM')]

ev_ws = wb['DB_Evidence']
ev_hdr = [ev_ws.cell(4, c).value for c in range(1, 18)]
ev = [x for x in rows(ev_ws, 5, 17) if isinstance(x[0], str) and re.match(r'E\d+', x[0])]

dec_ws = wb['Decisions']
dec = [x for x in rows(dec_ws, 5, 10) if isinstance(x[0], str) and x[0].strip()]

lb_ws = wb['Load_Bearing']
lb = [x for x in rows(lb_ws, 5, 14) if isinstance(x[0], str) and x[0].startswith('GM')]

src = {}
src_hdr = {}
for s in wb.sheetnames:
    if not s.startswith('SRC_') or s == 'SRC_Index': continue
    ws = wb[s]
    src_hdr[s] = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
    for r in range(5, ws.max_row + 1):
        i = ws.cell(r, 1).value
        if isinstance(i, str) and i.startswith('SRC_'):
            src[i] = dict(sheet=s, row=r, vals=[ws.cell(r, c).value for c in range(1, ws.max_column + 1)])

# 模型頁格值（供讀現值與列標籤）
cells = {}
for s in wb.sheetnames:
    if s in ('Gov_Map', 'SRC_Index'): continue
    ws = wb[s]
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                cells[(s, c.coordinate)] = c.value

# 稽核第 1 段：輸入判讀頁
a1 = openpyxl.load_workbook(A1, read_only=True, data_only=True)
w = a1['輸入判讀']
arr = list(w.iter_rows(values_only=True))
a1_hdr = list(arr[1]); a1_rows = [list(x) for x in arr[2:] if x and x[0]]
w = a1['高嚴重度']
arr = list(w.iter_rows(values_only=True))
a1_hi = [list(x) for x in arr[2:] if x and x[0]]

pickle.dump(dict(gm_hdr=gm_hdr, gm=gm, ev_hdr=ev_hdr, ev=ev, dec=dec, lb=lb, src=src, src_hdr=src_hdr, cells=cells,
                 a1_hdr=a1_hdr, a1_rows=a1_rows, a1_hi=a1_hi), open(os.path.join(HERE, '01_load.pkl'), 'wb'))
print('Gov_Map', len(gm), 'DB_Evidence', len(ev), 'Decisions', len(dec), 'Load_Bearing', len(lb), 'SRC', len(src))
print('A1 hdr', a1_hdr)
print('A1 rows', len(a1_rows), 'hi', len(a1_hi))
print('GM hdr', gm_hdr)
for s, h in src_hdr.items(): print(s, h[:20]); break
