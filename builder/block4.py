# Block 4 (v5.8): Cap_In, Capability, Price_Frontier, Cache_Store, Fleet_1GW, Amortize, Theory_Rev, Sens_Rev
# Decisions (Andy 2026-10-01): K1 OpenAI price = base; K2 (i) price frontier = cheapest model whose capability
# index >= the OpenAI tier model; Chinese vendors in the candidate table; K3 AA Intelligence Index (METR check);
# K4 (c)+(d): capability->price elasticity base 0 + break-even premium; K5 life Luna/Sol 12, Astra 9 months;
# K6 both amortization bases; v5.9: downstream default (c) = top-down total x bottom-up weights, (d) revenue weights; K7 fleet calibrated to OpenAI 2025; K8 API prices only;
# K9 one price snapshot for all generations; K10 cache storage by analogy to public cache terms;
# K11 utilization 60% / derate 1.0 kept; K12 peak & off-peak both, frontier on hour-weighted;
# K13 all Chinese vendors (open-weight flag); K14 international USD prices.
# New formulas reference key quantities through named ranges (B4_ = Block 4 internal / display; IF_ = downstream).
from common import *
from outputs import COLS15
from openpyxl.workbook.defined_name import DefinedName

TIERS = [("Luna", "Luna（低層）"), ("Sol", "Sol（中層）"), ("Astra", "Astra（頂層）")]
TC = "CDE"                      # tier columns on Cap_In / Price_Frontier summary

def nm(wb, n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)

def head15(ws, note_row=None):
    """rows 4–7: generation, cost case, generation index, column index (C..Q)."""
    put(ws, "A4", "世代", F_BOLD); put(ws, "A5", "成本情境", F_BOLD)
    put(ws, "A6", "世代索引", F_BOLD); put(ws, "A7", "欄索引", F_BOLD)
    for i, X in enumerate(COLS15):
        put(ws, f"{X}4", f"=Unit_Cost!{X}4", F_HLINK); put(ws, f"{X}5", f"=Unit_Cost!{X}5", F_HLINK)
        put(ws, f"{X}6", i // 3 + 1, F_CALC, fmt="0"); put(ws, f"{X}7", i + 1, F_CALC, fmt="0")
        ws.column_dimensions[X].width = 12
    ws.column_dimensions["A"].width = 50; ws.column_dimensions["B"].width = 11
    ws.freeze_panes = "C8"

def row15(ws, r, lab, unit, tpl, fmt="#,##0.000", key=False):
    put(ws, f"A{r}", lab, wrap=True); put(ws, f"B{r}", unit)
    for X in COLS15:
        put(ws, f"{X}{r}", tpl.replace("{X}", X), fmt=fmt, fill=FILL_KEY if key else None)

IF = lambda name: f"INDEX({name},1,{{X}}$7)"          # one-row 15-column IF_ range at this column

# ---------------------------------------------------------------- Cap_In
PRICE_ROWS = [  # model, vendor, country, open weights, fresh, cached, out, peak(1/0), AA index, index ver, price date, tag, source
  ("GPT-6 Luna", "OpenAI", "美國", "否", 0.1, 0.01, 0.5, 0, 37, "v4.3.2", "2026-09-25", "Verified", "S50（價格）；S56（指數）"),
  ("GPT-6 Sol", "OpenAI", "美國", "否", 2.0, 0.2, 10.0, 0, 48, "v4.3.2", "2026-09-25", "Verified", "S50；S56"),
  ("GPT-6 Astra", "OpenAI", "美國", "否", 10.0, 1.0, 50.0, 0, 53, "v4.3", "2026-09-25", "Verified", "S50；S55"),
  ("Claude Fable 5.1", "Anthropic", "美國", "否", 10.0, 0.25, 50.0, 0, 53, "v4.3", "2026-09", "Interested-party", "S51（快取讀取 0.25 僅單一二手來源，待以官方頁核對）；S55"),
  ("Claude Opus 5", "Anthropic", "美國", "否", 5.0, 0.5, 25.0, 0, 51, "v4.3", "2026-08", "Interested-party", "S51；S55"),
  ("Claude Opus 5.5", "Anthropic", "美國", "否", 4.0, 0.2, 20.0, 0, None, "待查", "2026-09-22", "Interested-party", "S51；指數未取得（Anthropic 自稱達 Fable 5.1 水準，未經獨立評測前不納入前緣）"),
  ("DeepSeek V4.1-Flash", "DeepSeek", "中國", "是", 0.30, 0.006, 1.20, 1, 39, "v4.3.2", "2026-10-01", "Verified", "S52（官方頁，尖峰價；離峰半價）；S56"),
  ("DeepSeek V4-Pro-0813", "DeepSeek", "中國", "是", 1.32, 0.044, 3.96, 1, 36, "v4.3", "2026-10-01", "Verified", "S52；S55（指數為『DeepSeek V4 Pro』）"),
  ("GLM-5.3", "Z.ai", "中國", "否（預定開放）", 1.40, 0.26, 4.40, 0, 45, "v4.3.2", "2026-10-01", "Verified", "S53（官方頁）；S56"),
  ("GLM-5.3-Flash", "Z.ai", "中國", "是", 0.15, 0.03, 0.50, 0, 42, "v4.3", "2026-10-01", "Verified", "S53；S55"),
  ("Qwen3.8-Max", "Alibaba", "中國", "預定開放", 2.00, 0.25, 6.00, 0, 40, "v4.3", "2026-09", "Interested-party", "S54（國際站；多個二手來源一致，另一來源載 $2.5／$7.5 附五折，待核）；S55"),
  ("Kimi K3", "Moonshot", "中國", "有爭議", 3.00, 0.30, 15.00, 0, None, "待查", "2026-07-16", "Interested-party", "S54；v4.3 指數未取得（報導僅稱與 GLM-5.3 並列開放權重領先，≥43）"),
  ("MiniMax M3", "MiniMax", "中國", "否", 0.30, None, 1.20, 0, None, "待查", "2026-08", "Interested-party", "S54（≤512K 輸入；快取價未取得）；指數未取得"),
]

def cap_in(wb):
    ws = wb.create_sheet("Cap_In")
    title(ws, "Cap_In — Block 4 輸入（藍字＝輸入；Analogy／Assumed 一律附區間；價格為 2026-09／10 快照）",
          "Block 4 命題：每 1 GW 的理論營收（理想上限）＝SLO 產能 × 利用率 × 層級別有效單價；另列中國廠商在內的單價前緣、訓練攤提、快取儲存與 1 GW 參考機隊。決策 K1–K14 見 README")
    for c, w in zip("ABCDEFGHIJKLM", [44, 10, 12, 12, 12, 16, 70, 10, 10, 9, 11, 15, 60]): ws.column_dimensions[c].width = w
    K = {}
    def hdr(r, cols=("值／Luna", "Sol", "Astra", "標記", "說明")):
        put(ws, f"A{r}", "項目", F_BOLD); put(ws, f"B{r}", "單位", F_BOLD)
        for c, h in zip("CDEFG", cols): put(ws, f"{c}{r}", h, F_BOLD)
    r = 4
    section(ws, r, "A. OpenAI 層級牌價（K1 基準；2026-09-25，短上下文 Standard）", 7); r += 1; hdr(r); r += 1
    for key, lab, vals, note in [
        ("pin", "新鮮輸入", (0.1, 2.0, 10.0), "GPT-6 Luna／Sol／Astra"),
        ("pc", "快取輸入", (0.01, 0.2, 1.0), "＝輸入價 10%"),
        ("pout", "輸出（含思考 token）", (0.5, 10.0, 50.0), "思考 token 依輸出價計費（S50、S51、S54：OpenAI、Anthropic、Moonshot 皆同）")]:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", "$/M")
        for c, v in zip(TC, vals): put(ws, f"{c}{r}", v, fmt="#,##0.000")
        put(ws, f"F{r}", "Verified", F_NOTE); put(ws, f"G{r}", note + "（S50）", F_NOTE, wrap=True); K[key] = r; r += 1
    r += 1
    section(ws, r, "B. 計費與工作負載（單一值；D、E 欄＝低、高）", 7); r += 1
    hdr(r, ("基準", "低", "高", "標記", "說明")); r += 1
    singles = [
      ("disc", "有效折扣（Batch／Flex 半價＋企業議價，加權）", "%", 0.2, 0.1, 0.3, "Assumed", "與 OpenAI 模型 v0.5 一致；Batch 半價為 Verified，企業折扣為 Analogy"),
      ("chi", "參考請求快取命中率 χ（輸入中屬快取命中的比例）", "%", 0.55, 0.3, 0.75, "Assumed", "與 OpenAI 模型 v0.5 一致；成本與營收用同一 χ"),
      ("eps", "能力 → 單價彈性 ε（單價 ∝ 有效算力倍數^ε）", "x", 0, 0, 0.5, "Assumed", "K4 (c)：機制保留、基準 0。斜率無法由公開資料分離（價格由廠商策略與成本決定，見 Capability C 節）"),
      ("cmult", "有效訓練算力倍數（相對 Block 3 基準；情境用）", "x", 1, 0.5, 7, "Assumed", "與 Tech_Registry H_CAP 相乘後代入 ε；高值 7＝J8 前沿錨點中值相對 Astra 基準的倍數"),
      ("pkmode", "尖峰／離峰口徑（1＝尖峰價；2＝依時數加權）", "選擇", 2, 1, 2, "Decision", "K12：兩者並列，前緣基準用 2"),
      ("pkhr", "尖峰時數（每週）", "hr", 35, None, None, "Verified", "DeepSeek：週一至五 UTC 01–04、06–10，共 7 hr × 5（S52）"),
      ("offr", "離峰價 ÷ 尖峰價", "x", 0.5, None, None, "Verified", "DeepSeek（S52）"),
    ]
    for key, lab, unit, b, lo, hi, tag, note in singles:
        put(ws, f"A{r}", lab, wrap=True); put(ws, f"B{r}", unit)
        fmt = "0%" if unit == "%" else ("0.00" if unit == "x" else "0")
        put(ws, f"C{r}", b, fmt=fmt)
        if lo is not None: put(ws, f"D{r}", lo, fmt=fmt); put(ws, f"E{r}", hi, fmt=fmt)
        put(ws, f"F{r}", tag, F_NOTE); put(ws, f"G{r}", note, F_NOTE, wrap=True); K[key] = r; r += 1
    r += 1
    section(ws, r, "C. 模型商業壽命（K5；月）", 7); r += 1; hdr(r); r += 1
    for key, lab, vals, tag, note in [
        ("life", "商業壽命 基準", (12, 12, 9), "Assumed", "旗艦更替快：GPT-5.6（2026-07）→ GPT-6（2026-09）"),
        ("lifelo", "商業壽命 低", (6, 6, 6), "Assumed", ""), ("lifehi", "商業壽命 高", (24, 24, 18), "Assumed", "")]:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", "月")
        for c, v in zip(TC, vals): put(ws, f"{c}{r}", v, fmt="0")
        put(ws, f"F{r}", tag, F_NOTE); put(ws, f"G{r}", note, F_NOTE); K[key] = r; r += 1
    r += 1
    section(ws, r, "D. 1 GW 參考機隊配置（K7：以 OpenAI 2025 算力支出校準；D、E 欄＝低、高）", 7); r += 1
    hdr(r, ("基準", "低", "高", "標記", "說明")); r += 1
    for key, lab, b, lo, hi, tag, note in [
        ("serve", "對外服務占機隊", 0.41, 0.30, 0.60, "Interested-party／Derived",
         "OpenAI 2025 推論 $8.4B ÷（推論 $8.4B＋訓練約 $12B）；以支出比代 GW 比（硬體與雲端加價不同，故為 Derived）"),
        ("free", "其中免費服務占服務", 0.46, 0.35, 0.60, "Interested-party", "OpenAI 2025 非付費用戶推論 $3.9B ÷ $8.4B（The Information 2026-02）")]:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", "%"); put(ws, f"C{r}", b, fmt="0%"); put(ws, f"D{r}", lo, fmt="0%")
        put(ws, f"E{r}", hi, fmt="0%"); put(ws, f"F{r}", tag, F_NOTE); put(ws, f"G{r}", note, F_NOTE, wrap=True); K[key] = r; r += 1
    hdr(r); r += 1
    for key, lab, vals, note in [
        ("mixp", "付費服務 token 層級組合", (0.40, 0.45, 0.15), "OpenAI 模型 v0.5 API tokenMix 2026（Assumed）；三者合計須為 1"),
        ("mixf", "免費服務 token 層級組合", (0.90, 0.10, 0.0), "OpenAI 模型 v0.5 免費方案模型組合（Assumed）")]:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", "%")
        for c, v in zip(TC, vals): put(ws, f"{c}{r}", v, fmt="0%")
        put(ws, f"F{r}", "Assumed", F_NOTE); put(ws, f"G{r}", note, F_NOTE, wrap=True); K[key] = r; r += 1
    r += 1
    section(ws, r, "E. 快取儲存（K10：比照公開快取條款；D、E 欄＝低、高）", 7); r += 1
    hdr(r, ("基準", "低", "高", "標記", "說明")); r += 1
    for key, lab, unit, b, lo, hi, fmt, tag, note in [
        ("stier", "儲存層（1＝SSD；2＝DRAM；3＝HBM）", "選擇", 2, 1, 3, "0", "Assumed", "公開條款：預設保留 5 分鐘、可選 1 小時（S51）；長保留通常下放 DRAM／SSD"),
        ("ret", "保留時間", "hr", 0.0833333333333333, 0.0833333333333333, 1, "0.000", "Analogy", "5 分鐘（預設 TTL）至 1 小時（延長 TTL，寫入價 2 倍，S51）"),
        ("hits", "每次寫入在保留期內的平均命中次數", "次", 5, 1, 20, "0", "Assumed", "代理迴圈多輪重用前綴時高；單次問答低"),
        ("dram", "DRAM 取得成本", "$/GB", 10, 5, 20, "0.00", "Assumed", "2026 記憶體漲價期；含伺服器分攤"),
        ("dlife", "DRAM 折舊年限", "年", 4, 3, 6, "0", "Assumed", ""),
        ("ssd", "SSD 取得成本", "$/GB", 0.10, 0.05, 0.30, "0.00", "Assumed", "企業級 NVMe"),
        ("slife", "SSD 折舊年限", "年", 5, 3, 6, "0", "Assumed", "")]:
        put(ws, f"A{r}", lab, wrap=True); put(ws, f"B{r}", unit)
        put(ws, f"C{r}", b, fmt=fmt); put(ws, f"D{r}", lo, fmt=fmt); put(ws, f"E{r}", hi, fmt=fmt)
        put(ws, f"F{r}", tag, F_NOTE); put(ws, f"G{r}", note, F_NOTE, wrap=True); K[key] = r; r += 1
    r += 1
    section(ws, r, "F. 市場價格與能力候選表（K2 (i)、K13、K14：國際站美元牌價；能力＝Artificial Analysis Intelligence Index）", 13); r += 1
    h = ["模型", "廠商", "國別", "開放權重", "新鮮輸入 $/M", "快取輸入 $/M", "輸出 $/M", "尖峰離峰（1＝有）",
         "能力指數", "指數版本", "價格日期", "標記", "來源", "中國廠商（1＝是）"]
    for i, t in enumerate(h): put(ws, f"{L(i+1)}{r}", t, F_BOLD, wrap=True)
    r += 1; K["tab0"] = r
    for row in PRICE_ROWS:
        for i, v in enumerate(row):
            if v is None: continue
            fmt = "#,##0.000" if i in (4, 5, 6) else ("0" if i in (7, 8) else None)
            put(ws, f"{L(i+1)}{r}", v, F_IN if i < 12 else F_NOTE, fmt=fmt, wrap=(i == 12))
        put(ws, f"N{r}", 1 if row[2] == "中國" else 0, F_IN, fmt="0")   # v5.9：中國廠商旗標（CC 第 7 輪）
        r += 1
    K["tab1"] = r - 1
    put(ws, f"A{r}", "註：前 3 列為 OpenAI 層級模型，其能力指數即各層級門檻。快取價空白者以新鮮輸入價計。能力指數空白者不參與前緣（不代表能力不足）。"
                    "Artificial Analysis 一週內改版三次（v4.1→4.3），跨版本分數不可比；本表全部取 v4.3／v4.3.2（S55、S56）", F_NOTE)
    ws.freeze_panes = "C4"
    # ---- named ranges
    nm(wb, "B4_PriceIn", f"Cap_In!$C${K['pin']}:$E${K['pin']}"); nm(wb, "B4_PriceCache", f"Cap_In!$C${K['pc']}:$E${K['pc']}")
    nm(wb, "B4_PriceOut", f"Cap_In!$C${K['pout']}:$E${K['pout']}")
    for k, n in [("disc", "B4_Disc"), ("chi", "B4_CacheHit"), ("eps", "B4_Eps"), ("cmult", "B4_CompMult"), ("pkmode", "B4_PeakMode"),
                 ("pkhr", "B4_PeakHrs"), ("offr", "B4_OffPeak"), ("serve", "B4_Serve"), ("free", "B4_Free"),
                 ("stier", "B4_StoreTier"), ("ret", "B4_Retain"), ("hits", "B4_Hits"), ("dram", "B4_DRAMcost"),
                 ("dlife", "B4_DRAMlife"), ("ssd", "B4_SSDcost"), ("slife", "B4_SSDlife")]:
        nm(wb, n, f"Cap_In!$C${K[k]}")
    nm(wb, "B4_Life", f"Cap_In!$C${K['life']}:$E${K['life']}")
    nm(wb, "B4_MixPaid", f"Cap_In!$C${K['mixp']}:$E${K['mixp']}"); nm(wb, "B4_MixFree", f"Cap_In!$C${K['mixf']}:$E${K['mixf']}")
    nm(wb, "B4_ISL", "Serving!$C$23:$E$23"); nm(wb, "B4_OSL", "Serving!$C$24:$E$24")
    a, b = K["tab0"], K["tab1"]
    for col, n in zip("ABCDEFGHI", ["B4_MktModel", "B4_MktVendor", "B4_MktCountry", "B4_MktOpen", "B4_MktIn",
                                    "B4_MktCache", "B4_MktOut", "B4_MktPeak", "B4_MktIndex"]):
        nm(wb, n, f"Cap_In!${col}${a}:${col}${b}")
    nm(wb, "B4_MktChina", f"Cap_In!$N${a}:$N${b}")
    ws.column_dimensions["N"].width = 10
    return K

# ---------------------------------------------------------------- Price_Frontier
def price_frontier(wb, K, TR):
    ws = wb.create_sheet("Price_Frontier")
    title(ws, "Price_Frontier — 單價前緣（K2 (i)：能力 ≥ OpenAI 層級模型者之中，參考請求混合單價最低者）",
          "混合單價＝（ISL ×〔(1−χ) 新鮮＋χ 快取〕＋ OSL × 輸出）÷（ISL＋OSL），用各層級參考任務（Serving 列 23–24）。前緣以『整筆請求』比較，不逐 token 類型取最低（避免拼出不存在的組合）。牌價為表列價（未折扣），折扣在 Theory_Rev 對所有廠商一致套用")
    for c, w in zip("ABCDEFGHIJKLMNOP", [24, 10, 8, 10, 10, 10, 9, 11, 11, 11, 11, 11, 11, 11, 11, 11]): ws.column_dimensions[c].width = w
    a, b = K["tab0"], K["tab1"]
    r = 4
    section(ws, r, "A. 候選模型的有效牌價與各層級參考請求混合單價", 16); r += 1
    hdr = ["模型", "國別", "能力指數", "時段係數", "新鮮輸入", "快取輸入", "輸出", "混合 Luna", "混合 Sol", "混合 Astra",
           "合格 Luna", "合格 Sol", "合格 Astra", "中國合格 Luna", "中國合格 Sol", "中國合格 Astra"]
    for i, t in enumerate(hdr): put(ws, f"{L(i+1)}{r}", t, F_BOLD, wrap=True)
    ws.row_dimensions[r].height = 30; r += 1
    P = {"t0": r}
    for j in range(a, b + 1):
        put(ws, f"A{r}", f"=Cap_In!A{j}", F_LINK); put(ws, f"B{r}", f"=Cap_In!C{j}", F_LINK)
        put(ws, f"C{r}", f"=IF(ISNUMBER(Cap_In!I{j}),Cap_In!I{j},\"—\")", F_LINK, fmt="0")
        put(ws, f"D{r}", f"=IF(Cap_In!H{j}=1,IF(B4_PeakMode=1,1,(B4_PeakHrs+(168-B4_PeakHrs)*B4_OffPeak)/168),1)", fmt="0.000")
        put(ws, f"E{r}", f"=Cap_In!E{j}*D{r}", F_LINK, fmt="#,##0.000")
        put(ws, f"F{r}", f"=IF(ISNUMBER(Cap_In!F{j}),Cap_In!F{j},Cap_In!E{j})*D{r}", F_LINK, fmt="#,##0.000")
        put(ws, f"G{r}", f"=Cap_In!G{j}*D{r}", F_LINK, fmt="#,##0.000")
        for t in range(3):
            put(ws, f"{'HIJ'[t]}{r}", f"=(INDEX(B4_ISL,1,{t+1})*((1-B4_CacheHit)*E{r}+B4_CacheHit*F{r})+INDEX(B4_OSL,1,{t+1})*G{r})"
                                      f"/(INDEX(B4_ISL,1,{t+1})+INDEX(B4_OSL,1,{t+1}))", fmt="#,##0.000")
            thr = f"$C${P['t0']+t}"   # first three candidate rows = OpenAI tier models
            put(ws, f"{'KLM'[t]}{r}", f"=IF(AND(ISNUMBER(C{r}),ISNUMBER({thr})),IF(C{r}>={thr},{'HIJ'[t]}{r},1E9),1E9)", fmt="#,##0.000")
            put(ws, f"{'NOP'[t]}{r}", f"=IF(Cap_In!N{j}=1,{'KLM'[t]}{r},1E9)", fmt="#,##0.000")
        r += 1
    P["t1"] = r - 1; t0, t1 = P["t0"], P["t1"]
    put(ws, f"A{r}", "合格欄 1E9＝不合格（能力指數低於該層級 OpenAI 模型或未取得）", F_NOTE); r += 2
    section(ws, r, "B. 各層級前緣（欄＝Luna／Sol／Astra）", 16); r += 1
    put(ws, f"A{r}", "項目", F_BOLD); put(ws, f"B{r}", "單位", F_BOLD)
    for t, (tk, tn) in enumerate(TIERS): put(ws, f"{TC[t]}{r}", tn, F_BOLD)
    r += 1
    def srow(key, lab, unit, f, fmt="#,##0.000", fill=None):
        nonlocal r
        put(ws, f"A{r}", lab, wrap=True); put(ws, f"B{r}", unit)
        for t in range(3):
            put(ws, f"{TC[t]}{r}", f.format(t=t + 1, H="HIJ"[t], Q="KLM"[t], N="NOP"[t], C=TC[t]), fmt=fmt, fill=fill)
        P[key] = r; r += 1
    rng = lambda c: f"${c}${t0}:${c}${t1}"
    srow("thr", "層級能力門檻（OpenAI 層級模型的指數）", "指數", "", "0")
    for t in range(3): ws[f"{TC[t]}{P['thr']}"].value = f"=$C${t0+t}"
    srow("oai", "OpenAI 層級模型 混合單價（表列）", "$/M", "=INDEX(" + rng("{H}") + ",{t})")
    srow("fr", "前緣 混合單價（表列）", "$/M", "=MIN(" + rng("{Q}") + ")", fill=FILL_KEY)
    srow("frm", "前緣模型", "", "")
    for t in range(3): ws[f"{TC[t]}{P['frm']}"].value = f"=INDEX({rng('A')},MATCH({TC[t]}{P['fr']},{rng('KLM'[t])},0))"
    for key, lab, col in [("frin", "前緣 新鮮輸入（表列、時段加權）", "E"), ("frc", "前緣 快取輸入", "F"), ("frout", "前緣 輸出（含思考）", "G")]:
        srow(key, lab, "$/M", "")
        for t in range(3): ws[f"{TC[t]}{P[key]}"].value = f"=INDEX({rng(col)},MATCH({TC[t]}{P['fr']},{rng('KLM'[t])},0))"
    srow("prem", "OpenAI 溢價（OpenAI ÷ 前緣；1＝OpenAI 即前緣）", "x", "", "0.00", FILL_KEY)
    for t in range(3): ws[f"{TC[t]}{P['prem']}"].value = f"={TC[t]}{P['oai']}/{TC[t]}{P['fr']}"
    srow("cn", "中國廠商 合格者最低 混合單價", "$/M", "")
    for t in range(3): ws[f"{TC[t]}{P['cn']}"].value = f"=IF(MIN({rng('NOP'[t])})>=1E9,\"無合格\",MIN({rng('NOP'[t])}))"
    srow("cnm", "中國廠商 合格者最低 模型", "", "")
    for t in range(3): ws[f"{TC[t]}{P['cnm']}"].value = f"=IF(ISNUMBER({TC[t]}{P['cn']}),INDEX({rng('A')},MATCH({TC[t]}{P['cn']},{rng('NOP'[t])},0)),\"—\")"
    srow("cnr", "中國合格者最低 ÷ OpenAI 層級模型", "x", "", "0.00")
    for t in range(3): ws[f"{TC[t]}{P['cnr']}"].value = f"=IF(ISNUMBER({TC[t]}{P['cn']}),{TC[t]}{P['cn']}/{TC[t]}{P['oai']},\"—\")"
    srow("cnk", "中國廠商 合格家數（候選表中）", "個", "", "0")
    for t in range(3): ws[f"{TC[t]}{P['cnk']}"].value = f"=COUNTIF({rng('NOP'[t])},\"<1E9\")"
    r += 1
    section(ws, r, "C. 能力 → 單價倍數（K4 (c)：基準 ε＝0，倍數恆為 1）", 16); r += 1
    srow("hcap", "Tech_Registry H_CAP 有效倍數（能力增量，以有效算力表示）", "x", f"={TR['H_CAP']}", "0.00")
    srow("capf", "能力單價倍數＝（H_CAP × 有效訓練算力倍數）^ε", "x", "", "0.000", FILL_KEY)
    for t in range(3): ws[f"{TC[t]}{P['capf']}"].value = f"=EXP(B4_Eps*LN({TC[t]}{P['hcap']}*B4_CompMult))"
    r += 1
    section(ws, r, "D. OpenAI 有效單價（表列 ×（1−折扣）× 能力單價倍數；依 token 類型；下游連結 IF_Price*）", 16); r += 1
    for key, lab, rngname in [("ein", "新鮮輸入 $/M（有效）", "B4_PriceIn"), ("ec", "快取輸入 $/M（有效）", "B4_PriceCache"),
                              ("eth", "思考 $/M（有效；依輸出價計費）", "B4_PriceOut"), ("eout", "可見輸出 $/M（有效）", "B4_PriceOut")]:
        srow(key, lab, "$/M", "", "#,##0.000", FILL_KEY)
        for t in range(3): ws[f"{TC[t]}{P[key]}"].value = f"=INDEX({rngname},1,{t+1})*(1-B4_Disc)*{TC[t]}{P['capf']}"
    srow("eref", "參考請求混合 $/M 總 token（OpenAI 有效）", "$/M", "", "#,##0.000", FILL_KEY)
    for t in range(3):
        c = TC[t]
        ws[f"{c}{P['eref']}"].value = (f"=(INDEX(B4_ISL,1,{t+1})*((1-B4_CacheHit)*{c}{P['ein']}+B4_CacheHit*{c}{P['ec']})+INDEX(B4_OSL,1,{t+1})*{c}{P['eout']})"
                                       f"/(INDEX(B4_ISL,1,{t+1})+INDEX(B4_OSL,1,{t+1}))")
    srow("fref", "參考請求混合 $/M 總 token（前緣有效＝前緣表列 ×（1−折扣））", "$/M", "", "#,##0.000", FILL_KEY)
    for t in range(3): ws[f"{TC[t]}{P['fref']}"].value = f"={TC[t]}{P['fr']}*(1-B4_Disc)"
    for key, n in [("ein", "B4_EffIn"), ("ec", "B4_EffCache"), ("eout", "B4_EffOut"), ("eref", "B4_EffRef"), ("fref", "B4_FrontRef")]:
        nm(wb, n, f"Price_Frontier!$C${P[key]}:$E${P[key]}")
    ws.freeze_panes = "B6"
    return P

# ---------------------------------------------------------------- Capability
def capability(wb, K, P, TRN, TI):
    ws = wb.create_sheet("Capability")
    title(ws, "Capability — 算力 → 能力 → 單價（K3：AA 指數為主、METR 為檢查；K4：只作檢查與反解，不驅動基準營收）",
          "本頁回答三件事：(1) 各層級門檻模型的能力位置；(2) 本模型各層級的訓練算力與能力的對應（檢查用，Tokenomics 架構為代理，不等於 OpenAI 實際模型）；(3) 為何不由斜率驅動單價")
    for c, w in zip("ABCDEFG", [52, 10, 14, 14, 14, 14, 70]): ws.column_dimensions[c].width = w
    C = {}
    r = 4
    section(ws, r, "A. 各層級（VR200 欄；FLOPs 與世代無關）", 7); r += 1
    put(ws, f"A{r}", "項目", F_BOLD); put(ws, f"B{r}", "單位", F_BOLD)
    for t, (tk, tn) in enumerate(TIERS): put(ws, f"{TC[t]}{r}", tn, F_BOLD)
    put(ws, f"G{r}", "說明", F_BOLD); r += 1
    tcol = "LMN"   # Training VR200 base columns: L/M/N = VR200 Luna/Sol/Astra
    rows = [
      ("flop", "最終訓練 FLOPs（預訓練＋後訓練；Training Cfin）", "FLOP", lambda t: f"=Training!{tcol[t]}{TRN['Cfin']}", "0.00E+00", "Block 3"),
      ("lflop", "log10 FLOPs", "", lambda t: f"=LN({TC[t]}{{flop}})/LN(10)", "0.00", ""),
      ("idx", "門檻模型能力指數（AA）", "指數", lambda t: f"=Price_Frontier!{TC[t]}{P['thr']}", "0", "Cap_In F 節前 3 列"),
      ("slope", "相鄰層級：每 10 倍算力的指數增量", "點／10x", lambda t: "=\"—\"" if t == 0 else
           f"=({TC[t]}{{idx}}-{TC[t-1]}{{idx}})/({TC[t]}{{lflop}}-{TC[t-1]}{{lflop}})", "0.0", "Derived；代理架構，僅示意"),
      ("pgap", "相鄰層級：OpenAI 混合單價倍數", "x", lambda t: "=\"—\"" if t == 0 else
           f"=Price_Frontier!{TC[t]}{P['oai']}/Price_Frontier!{TC[t-1]}{P['oai']}", "0.0", "Price_Frontier B 節"),
      ("cgap", "相鄰層級：本模型 decode 成本倍數（VR200 基準）", "x", lambda t: "=\"—\"" if t == 0 else
           f"=INDEX(IF_CostDec_{TIERS[t][0]},1,11)/INDEX(IF_CostDec_{TIERS[t-1][0]},1,11)", "0.0", "價格倍數 ≈ 成本倍數時，層級價差主要反映服務成本"),
    ]
    for key, lab, unit, f, fmt, note in rows:
        put(ws, f"A{r}", lab, wrap=True); put(ws, f"B{r}", unit); C[key] = r
        for t in range(3):
            put(ws, f"{TC[t]}{r}", f(t).replace("{flop}", str(C.get("flop", 0))).replace("{lflop}", str(C.get("lflop", 0)))
                .replace("{idx}", str(C.get("idx", 0))), fmt=fmt)
        put(ws, f"G{r}", note, F_NOTE, wrap=True); r += 1
    r += 1
    section(ws, r, "B. J8 檢查：前沿預訓練錨點 vs Astra", 7); r += 1
    put(ws, f"A{r}", "前沿錨點中值 ÷ Astra 最終訓練 FLOPs"); put(ws, f"B{r}", "x")
    put(ws, f"C{r}", f"=Train_In!$C${TI['an2']}/E{C['flop']}", fmt="0.0", fill=FILL_KEY)
    put(ws, f"G{r}", "若 Astra 對標 GPT-6 Astra，算力可能低估約此倍數（E006 待查）；Cap_In 有效訓練算力倍數高值即取此量級", F_NOTE, wrap=True)
    C["j8"] = r; r += 2
    section(ws, r, "C. 為何 K4 不採 (a)（摘要；數字見 Checks Block 4 與 Price_Frontier）", 7); r += 1
    for txt in [
        "1. 識別：同層級牌價跨廠商差距大（例：Sol 級輸出價 DeepSeek V4-Pro 離峰 $1.98 至 Kimi K3 $15），主要由廠商策略與成本決定，斜率無法與能力分離。",
        "2. 循環：前緣由接近成本定價的廠商決定（Checks：DeepSeek 價 ÷ 本模型 Hopper 同架構成本約 0.8–1.0），前緣斜率≈服務成本斜率，已由 Unit_Cost 計算。",
        "3. 重複：層級牌價已含能力；只有增量（H_CAP）可再進入，故以 ε 乘在增量上，基準 0。",
        "4. 誤差相乘：J8 算力缺口約 7 倍 × ε 0–0.5 → 單價倍數 1.0–2.6，大於利用率區間的影響。",
        "(d) 反解：Amortize 頁『回本所需單價溢價』與觀測層級價差並列，回答『訓練能否回本』而不假設斜率。",
        "METR 時間長度（K3 檢查）：尚未入表，待查。"]:
        put(ws, f"A{r}", txt, F_NOTE); r += 1
    return C

# ---------------------------------------------------------------- Cache_Store
def cache_store(wb):
    ws = wb.create_sheet("Cache_Store")
    title(ws, "Cache_Store — 快取儲存成本（Block 2 未結事項 4；每 M 快取命中 token；加在 IF_CostCache 之上）",
          "儲存 $/M 快取命中＝KV bytes/token × 10⁶ ÷ 10⁹ × 儲存層 $/GB-hr × 保留時間 ÷ 每次寫入命中次數。HBM 以每 GPU 小時持有成本 ÷ HBM 容量計（把整顆 GPU 成本歸給 HBM，為上限）")
    head15(ws)
    S = {}
    r = 9
    section(ws, r, "共用（儲存層 $/GB-hr）", 17); r += 1
    row15(ws, r, "SSD", "$/GB-hr", "=B4_SSDcost/(B4_SSDlife*8760)", "0.000000"); S["ssd"] = r; r += 1
    row15(ws, r, "DRAM", "$/GB-hr", "=B4_DRAMcost/(B4_DRAMlife*8760)", "0.000000"); S["dram"] = r; r += 1
    row15(ws, r, "HBM（上限）", "$/GB-hr", "=Unit_Cost!{X}$8/INDEX(Spec_Rack!$C$25:$G$25,1,{X}$6)", "0.000000"); S["hbm"] = r; r += 1
    row15(ws, r, "採用儲存層", "$/GB-hr", f"=CHOOSE(B4_StoreTier,{{X}}{S['ssd']},{{X}}{S['dram']},{{X}}{S['hbm']})", "0.000000"); S["sel"] = r; r += 2
    for t, (tk, tn) in enumerate(TIERS):
        section(ws, r, tn, 17); r += 1
        row15(ws, r, "KV bytes/token（Arch 採用列）", "B", f"=Arch!${TC[t]}$25", "#,##0"); kv = r; r += 1
        row15(ws, r, f"儲存 $/M 快取命中 token　[IF_CacheStore_{tk}]", "$/M",
              f"={{X}}{kv}*1E6/1E9*{{X}}${S['sel']}*B4_Retain/B4_Hits", "#,##0.00000", key=True); S[tk] = r; r += 1
        row15(ws, r, "儲存 ÷ 快取命中載入成本（IF_CostCache）", "x", f"=IF({IF(f'IF_CostCache_{tk}')}>0,{{X}}{S[tk]}/{IF(f'IF_CostCache_{tk}')},0)", "0.00"); r += 1
        row15(ws, r, "儲存 ÷ OpenAI 快取輸入有效價", "%", f"={{X}}{S[tk]}/INDEX(B4_EffCache,1,{t+1})", "0.00%"); r += 2
    return S

# ---------------------------------------------------------------- Fleet_1GW
def fleet(wb, K):
    ws = wb.create_sheet("Fleet_1GW")
    title(ws, "Fleet_1GW — 1 GW（IT）參考機隊配置（J11、K7）",
          "用途：對外服務（付費／免費）、最終訓練、研發實驗；服務 GW 依層級 token 組合分配：服務 token 總量 T＝服務 GW × 利用率 ÷ Σ（組合ₜ ÷ 每 GW 產出ₜ）。訓練＋研發＝1−服務；其中最終訓練＝÷ 研發倍數")
    head15(ws)
    F = {}
    r = 9
    section(ws, r, "A. 用途配置（GW）", 17); r += 1
    for key, lab, f in [("sv", "對外服務", "=B4_Serve"), ("pd", "  付費服務", "=B4_Serve*(1-B4_Free)"), ("fr", "  免費服務", "=B4_Serve*B4_Free"),
                        ("tr", "訓練＋研發", "=1-B4_Serve"), ("fin", "  最終訓練", "=(1-B4_Serve)/IF_RDMult"),
                        ("rd", "  研發實驗（不含最終訓練）", "=(1-B4_Serve)*(1-1/IF_RDMult)")]:
        row15(ws, r, lab, "GW", f, "0.000", key=(key in ("pd", "tr"))); F[key] = r; r += 1
    r += 1
    section(ws, r, "B. 服務 token（基準利用率；M tok/年）", 17); r += 1
    inv = lambda mix: "+".join(f"INDEX({mix},1,{t+1})/{IF(f'IF_TokGW_{tk}')}" for t, (tk, _) in enumerate(TIERS))
    # v5.9：SLO 不可達保護（CC 第 7 輪第 6 節第 1 項）——組合權重 > 0 的層級若每 GW 產出為 0，該欄機隊無法服務
    ok = lambda mix: "AND(" + ",".join(f"OR(INDEX({mix},1,{t+1})=0,{IF(f'IF_TokGW_{tk}')}>0)" for t, (tk, _) in enumerate(TIERS)) + ")"
    row15(ws, r, "可服務（組合內各層級皆 SLO 可達＝1）", "旗標", f"=IF(AND({ok('B4_MixPaid')},{ok('B4_MixFree')}),1,0)", "0"); F["ok"] = r; r += 1
    row15(ws, r, "付費服務 token 總量", "M tok/年", f"=IF({{X}}{F['ok']}=1,{{X}}{F['pd']}*IF_Util/({inv('B4_MixPaid')}),\"SLO 不可達\")", "#,##0"); F["Tp"] = r; r += 1
    row15(ws, r, "免費服務 token 總量", "M tok/年", f"=IF({{X}}{F['ok']}=1,{{X}}{F['fr']}*IF_Util/({inv('B4_MixFree')}),\"SLO 不可達\")", "#,##0"); F["Tf"] = r; r += 1
    for t, (tk, tn) in enumerate(TIERS):
        row15(ws, r, f"{tn} 付費 token", "M tok/年", f"=IF({{X}}{F['ok']}=1,{{X}}{F['Tp']}*INDEX(B4_MixPaid,1,{t+1}),\"SLO 不可達\")", "#,##0"); F[f"p{t}"] = r; r += 1
        row15(ws, r, f"{tn} 免費 token", "M tok/年", f"=IF({{X}}{F['ok']}=1,{{X}}{F['Tf']}*INDEX(B4_MixFree,1,{t+1}),\"SLO 不可達\")", "#,##0"); F[f"f{t}"] = r; r += 1
        row15(ws, r, f"{tn} 服務 GW", "GW", f"=IF({{X}}{F['ok']}=1,IF({IF(f'IF_TokGW_{tk}')}>0,({{X}}{F[f'p{t}']}+{{X}}{F[f'f{t}']})/({IF(f'IF_TokGW_{tk}')}*IF_Util),0),\"SLO 不可達\")", "0.000"); F[f"g{t}"] = r; r += 1
    row15(ws, r, "檢查：各層級服務 GW 合計 − 對外服務 GW（應為 0）", "GW",
          f"=IF({{X}}{F['ok']}=1,{{X}}{F['g0']}+{{X}}{F['g1']}+{{X}}{F['g2']}-{{X}}{F['sv']},\"SLO 不可達\")", "0.000000"); F["chk"] = r; r += 2
    section(ws, r, "C. 與 Block 3 一致性", 17); r += 1
    row15(ws, r, "家族研發計畫占 1 GW 年（三層級合計）", "%",
          "=" + "+".join(IF(f"IF_ProgGWyr_{tk}") for tk, _ in TIERS), "0.00%"); F["fam"] = r; r += 1
    row15(ws, r, "隱含每 GW 年可容納的家族研發計畫數（訓練＋研發 ÷ 家族計畫）", "個", f"={{X}}{F['tr']}/{{X}}{F['fam']}", "0.0", key=True)
    F["nprog"] = r; r += 1
    put(ws, f"A{r}", "解讀：前沿實驗室每年約 1–3 個旗艦家族；遠高於此代表 Block 3 單一計畫規模偏小（J8）或機隊訓練占比偏高", F_NOTE)
    return F

# ---------------------------------------------------------------- Amortize
def amortize(wb, F, P, S):
    ws = wb.create_sheet("Amortize")
    title(ws, "Amortize — 訓練攤提（K6：自下而上、由上而下並列；下游預設 (c)＝由上而下總額 × 自下而上權重，Andy 2026-10-01）",
          "自下而上＝層級研發計畫成本 ÷（1 GW 參考機隊在商業壽命內服務的該層級 token）；由上而下＝機隊訓練占比 X ÷（1−X）× 服務成本。"
          "前者同時是『回本所需單價溢價』（K4 (d)）。(c)、(d) 只改變層級間分配，總額同由上而下")
    head15(ws)
    A = {}
    NA = '"SLO 不可達"'
    num = lambda *refs: "AND(" + ",".join(f"ISNUMBER({x})" for x in refs) + ")"
    r = 9
    for t, (tk, tn) in enumerate(TIERS):
        section(ws, r, tn, 17); r += 1
        row15(ws, r, "研發計畫成本（IF_ProgCost）", "$M", f"={IF(f'IF_ProgCost_{tk}')}", "#,##0.0"); pc = r; A[f"pc{t}"] = r; r += 1
        p_, f_ = f"Fleet_1GW!{{X}}{F[f'p{t}']}", f"Fleet_1GW!{{X}}{F[f'f{t}']}"
        row15(ws, r, "年服務 token（付費＋免費）", "M tok/年", f"=IF({num(p_, f_)},{p_}+{f_},{NA})", "#,##0"); A[f"yr{t}"] = r; r += 1
        row15(ws, r, "商業壽命內服務 token（付費＋免費）", "M tok",
              f"=IF(ISNUMBER({{X}}{A[f'yr{t}']}),{{X}}{A[f'yr{t}']}*INDEX(B4_Life,1,{t+1})/12,{NA})", "#,##0"); sv = r; r += 1
        row15(ws, r, f"自下而上攤提＝回本所需單價溢價　[IF_AmortBU_{tk}]", "$/M",
              f"=IF(ISNUMBER({{X}}{sv}),IF({{X}}{sv}>0,{{X}}{pc}*1E6/{{X}}{sv},0),{NA})", "#,##0.00000", key=True)
        A[f"bu{t}"] = r; r += 1
        cost = (f"(INDEX(B4_ISL,1,{t+1})*((1-B4_CacheHit)*{IF(f'IF_CostPre_{tk}')}+B4_CacheHit*({IF(f'IF_CostCache_{tk}')}+Cache_Store!{{X}}{S[tk]}))"
                f"+INDEX(B4_OSL,1,{t+1})*{IF(f'IF_CostDec_{tk}')})/(INDEX(B4_ISL,1,{t+1})+INDEX(B4_OSL,1,{t+1}))/IF_Util")
        A[f"cost{t}"] = cost
        row15(ws, r, "服務成本（參考請求混合，含快取儲存，基準利用率）", "$/M",
              f"=IF(ISNUMBER({IF(f'IF_CostDec_{tk}')}),{cost},{NA})", "#,##0.0000"); A[f"sc{t}"] = r; r += 1
        row15(ws, r, f"由上而下攤提　[IF_AmortTD_{tk}]", "$/M",
              f"=IF(ISNUMBER({{X}}{A[f'sc{t}']}),Fleet_1GW!{{X}}{F['tr']}/(1-Fleet_1GW!{{X}}{F['tr']})*{{X}}{A[f'sc{t}']},{NA})", "#,##0.0000", key=True)
        A[f"td{t}"] = r; r += 1
        row15(ws, r, "由上而下 ÷ 自下而上", "x",
              f"=IF({num('{X}'+str(A[f'bu{t}']), '{X}'+str(A[f'td{t}']))},IF({{X}}{A[f'bu{t}']}>0,{{X}}{A[f'td{t}']}/{{X}}{A[f'bu{t}']},0),{NA})", "#,##0.0"); r += 1
        if t > 0:
            row15(ws, r, "回本所需溢價 ÷ 觀測層級價差（本層級 − 下一層級 OpenAI 有效混合單價）", "%",
                  f"=IF(ISNUMBER({{X}}{A[f'bu{t}']}),{{X}}{A[f'bu{t}']}/(INDEX(B4_EffRef,1,{t+1})-INDEX(B4_EffRef,1,{t})),{NA})", "0.000%"); r += 1
        r += 1
    # ---- v5.9 K6：(c)、(d) 與下游預設
    section(ws, r, "K6 下游預設（Andy 2026-10-01）：(c) 由上而下總額 × 自下而上權重（預設）；(d) 由上而下總額 × 營收權重（並列）", 17); r += 1
    allnum = lambda key: num(*[f"{{X}}{A[f'{key}{t}']}" for t in range(3)])
    tdsum = "+".join(f"{{X}}{A[f'td{t}']}*{{X}}{A[f'yr{t}']}" for t in range(3))
    row15(ws, r, "由上而下年化總額（三層級合計）", "$M/年", f"=IF({allnum('td')},({tdsum})/1E6,{NA})", "#,##0.0")
    A["tdtot"] = r; r += 1
    row15(ws, r, "自下而上年化總額（研發計畫成本 × 12 ÷ 商業壽命，合計）", "$M/年",
          "=" + "+".join(f"{{X}}{A[f'pc{t}']}*12/INDEX(B4_Life,1,{t+1})" for t in range(3)), "#,##0.0"); A["butot"] = r; r += 1
    row15(ws, r, "縮放倍數（由上而下 ÷ 自下而上總額）", "x",
          f"=IF(ISNUMBER({{X}}{A['tdtot']}),IF({{X}}{A['butot']}>0,{{X}}{A['tdtot']}/{{X}}{A['butot']},0),{NA})", "0.00", key=True); A["scale"] = r; r += 1
    pay = lambda t: f"Fleet_1GW!{{X}}{F[f'p{t}']}*INDEX(B4_EffRef,1,{t+1})"
    paidrefs = [f"Fleet_1GW!{{X}}{F[f'p{t}']}" for t in range(3)]
    paysum = "+".join(pay(t) for t in range(3))
    row15(ws, r, "付費營收合計（OpenAI 有效單價）", "$M/年", f"=IF({num(*paidrefs)},({paysum})/1E6,{NA})", "#,##0.0")
    A["revtot"] = r; r += 1
    for t, (tk, tn) in enumerate(TIERS):
        row15(ws, r, f"{tn}：(c) 由上而下總額 × 自下而上權重（下游預設）　[IF_AmortDefault_{tk}]", "$/M",
              f"=IF({num('{X}'+str(A[f'bu{t}']), '{X}'+str(A['scale']))},{{X}}{A[f'bu{t}']}*{{X}}{A['scale']},{NA})", "#,##0.00000", key=True)
        A[f"dc{t}"] = r; r += 1
        row15(ws, r, f"{tn}：(d) 由上而下總額 × 營收權重　[IF_AmortRev_{tk}]", "$/M",
              f"=IF({num('{X}'+str(A['revtot']), '{X}'+str(A[f'yr{t}']))},IF(AND({{X}}{A['revtot']}>0,{{X}}{A[f'yr{t}']}>0),"
              f"{{X}}{A['tdtot']}*({pay(t)}/1E6/{{X}}{A['revtot']})*1E6/{{X}}{A[f'yr{t}']},0),{NA})", "#,##0.00000")
        A[f"dd{t}"] = r; r += 1
    row15(ws, r, "(d) 攤提占付費營收（各層級相同）", "%",
          f"=IF(ISNUMBER({{X}}{A['revtot']}),IF({{X}}{A['revtot']}>0,{{X}}{A['tdtot']}/{{X}}{A['revtot']},0),{NA})", "0.0%"); A["dshare"] = r; r += 1
    put(ws, f"A{r}", "註：(c) 保留由上而下的總額（與觀測支出一致），層級間依物理計畫成本分配；縮放倍數與 Fleet_1GW『隱含家族研發計畫數』同源（J8）。"
                     "若 J8 低估集中於 Astra，(c) 對 Astra 仍偏低。(d) 為聯合成本的相對售價法。毛利率另列不含攤提口徑，以便與公司揭露比較", F_NOTE)
    return A

# ---------------------------------------------------------------- Theory_Rev
def theory_rev(wb, F, A, S, WL=None):
    NA = '"SLO 不可達"'
    WL = WL or {"fp": 22, "cp": 23, "dp": 24, "cost0": 33}
    ws = wb.create_sheet("Theory_Rev")
    title(ws, "Theory_Rev — 每 GW 理論營收（理想上限：SLO 產能 × 利用率 × 層級有效單價；不含需求、市占、訂閱方案）",
          "A 節＝單一層級滿載 1 GW；B 節＝1 GW 參考機隊（付費服務 token × 單價；免費服務營收 0）。OpenAI 單價為主線（K1），前緣單價並列（K2 (i)）。成本＝經濟口徑")
    head15(ws)
    T = {}
    r = 9
    for t, (tk, tn) in enumerate(TIERS):
        cs = S[tk]
        section(ws, r, f"A. {tn}：1 GW 只服務本層級", 17); r += 1
        row15(ws, r, "每 GW 產出（基準利用率）", "M tok/年", f"={IF(f'IF_TokGW_{tk}')}*IF_Util", "#,##0"); tok = r; r += 1
        row15(ws, r, f"理論營收 — OpenAI 有效單價　[IF_RevGW_{tk}]", "$B/年", f"={{X}}{tok}*INDEX(B4_EffRef,1,{t+1})/1E9", "#,##0.0", key=True)
        T[f"rev{t}"] = r; r += 1
        row15(ws, r, f"理論營收 — 前緣單價　[IF_RevGWFront_{tk}]", "$B/年", f"={{X}}{tok}*INDEX(B4_FrontRef,1,{t+1})/1E9", "#,##0.0"); T[f"revf{t}"] = r; r += 1
        row15(ws, r, "理論營收 — 100% 利用率（OpenAI）", "$B/年", f"={IF(f'IF_TokGW_{tk}')}*INDEX(B4_EffRef,1,{t+1})/1E9", "#,##0.0"); r += 1
        row15(ws, r, "年持有成本 — 經濟（IF_HoldEcon）", "$B/年", f"={IF('IF_HoldEcon')}", "#,##0.0"); hold = r; r += 1
        row15(ws, r, "理論營收 ÷ 持有成本", "x", f"={{X}}{T[f'rev{t}']}/{{X}}{hold}", "0.0", key=True); T[f"rh{t}"] = r; r += 1
        row15(ws, r, "前緣營收 ÷ 持有成本", "x", f"={{X}}{T[f'revf{t}']}/{{X}}{hold}", "0.0"); r += 1
        row15(ws, r, "服務成本 $/M（參考請求，含快取儲存，基準利用率）", "$/M", f"=Amortize!{{X}}{A[f'sc{t}']}", "#,##0.0000"); sc = r; r += 1
        g = lambda am: f"=IF(AND(ISNUMBER({{X}}{sc}),ISNUMBER(Amortize!{{X}}{am})),{{X}}{sc}+Amortize!{{X}}{am},{NA})"
        row15(ws, r, f"全成本 $/M（服務＋自下而上攤提）　[IF_FullCost_{tk}]", "$/M", g(A[f'bu{t}']), "#,##0.0000", key=True); T[f"fc{t}"] = r; r += 1
        row15(ws, r, "全成本 $/M（服務＋由上而下攤提）", "$/M", g(A[f'td{t}']), "#,##0.0000"); T[f"fct{t}"] = r; r += 1
        row15(ws, r, f"全成本 $/M（服務＋K6 預設 (c) 攤提）　[IF_FullCostDefault_{tk}]", "$/M", g(A[f'dc{t}']), "#,##0.0000", key=True); T[f"fcd{t}"] = r; r += 1
        gm = lambda fc, price: f"=IF(ISNUMBER({{X}}{fc}),1-{{X}}{fc}/INDEX({price},1,{t+1}),{NA})"
        row15(ws, r, "理論毛利率（OpenAI 單價；不含攤提，可與公司揭露毛利率比較）", "%", gm(sc, "B4_EffRef"), "0%"); r += 1
        row15(ws, r, "理論毛利率（OpenAI 單價；自下而上全成本）", "%", gm(T[f'fc{t}'], "B4_EffRef"), "0%"); r += 1
        row15(ws, r, "理論毛利率（OpenAI 單價；由上而下全成本）", "%", gm(T[f'fct{t}'], "B4_EffRef"), "0%"); r += 1
        row15(ws, r, "理論毛利率（OpenAI 單價；K6 預設 (c) 全成本）", "%", gm(T[f'fcd{t}'], "B4_EffRef"), "0%", key=True); r += 1
        row15(ws, r, "理論毛利率（前緣單價；由上而下全成本）", "%", gm(T[f'fct{t}'], "B4_FrontRef"), "0%"); r += 2
    section(ws, r, "B. 1 GW 參考機隊（Fleet_1GW 配置）", 17); r += 1
    ok = f"Fleet_1GW!{{X}}{F['ok']}=1"
    pay = lambda ref: "+".join(f"Fleet_1GW!{{X}}{F[f'p{t}']}*INDEX({ref},1,{t+1})" for t in range(3))
    # v5.9（Andy 決定 (b)）：機隊合計保留，另列各層級貢獻，使混合數字可拆解（CLAUDE.md 1a）
    for t, (tk, tn) in enumerate(TIERS):
        row15(ws, r, f"{tn} 貢獻 — OpenAI 有效單價　[IF_RevGWFleet_{tk}]", "$B/年",
              f"=IF({ok},Fleet_1GW!{{X}}{F[f'p{t}']}*INDEX(B4_EffRef,1,{t+1})/1E9,{NA})", "#,##0.00"); T[f"flt{t}"] = r; r += 1
    row15(ws, r, "付費服務營收合計（層級組合：付費 token 依 B4_MixPaid）— OpenAI 有效單價　[IF_RevGWFleet]", "$B/年",
          f"=IF({ok},({pay('B4_EffRef')})/1E9,{NA})", "#,##0.0", key=True); T["fl"] = r; r += 1
    for t, (tk, tn) in enumerate(TIERS):
        row15(ws, r, f"{tn} 貢獻 — 前緣單價　[IF_RevGWFleetFront_{tk}]", "$B/年",
              f"=IF({ok},Fleet_1GW!{{X}}{F[f'p{t}']}*INDEX(B4_FrontRef,1,{t+1})/1E9,{NA})", "#,##0.00"); T[f"flft{t}"] = r; r += 1
    row15(ws, r, "付費服務營收合計（層級組合）— 前緣單價　[IF_RevGWFleetFront]", "$B/年", f"=IF({ok},({pay('B4_FrontRef')})/1E9,{NA})", "#,##0.0"); T["flf"] = r; r += 1
    row15(ws, r, "整個 GW 年持有成本（服務＋訓練＋研發）", "$B/年", f"={IF('IF_HoldEcon')}", "#,##0.0"); fh = r; r += 1
    row15(ws, r, "機隊營收 ÷ 持有成本", "x", f"=IF(ISNUMBER({{X}}{T['fl']}),{{X}}{T['fl']}/{{X}}{fh},{NA})", "0.00", key=True); T["flr"] = r; r += 1
    row15(ws, r, "服務 token 加權平均有效單價", "$/M",
          f"=IF({ok},({pay('B4_EffRef')})/(Fleet_1GW!{{X}}{F['Tp']}+Fleet_1GW!{{X}}{F['Tf']}),{NA})", "#,##0.000"); r += 2
    section(ws, r, "C. 每任務營收（OpenAI 有效單價；欄＝Workload 任務類型 C–G；與 Workload 每任務成本對照；harness 依 Tech_Registry T12 混合）", 17); r += 1
    put(ws, f"A{r}", "層級 × 任務", F_BOLD)
    for c in "CDEFG": put(ws, f"{c}{r}", f"=Workload!{c}4", F_HLINK)
    r += 1
    T["task0"] = r
    for t, (tk, tn) in enumerate(TIERS):
        put(ws, f"A{r}", f"{tn} 每任務營收"); put(ws, f"B{r}", "$")
        for c in "CDEFG":
            put(ws, f"{c}{r}", f"=(Workload!{c}{WL['fp']}*INDEX(B4_EffIn,1,{t+1})+Workload!{c}{WL['cp']}*INDEX(B4_EffCache,1,{t+1})+Workload!{c}{WL['dp']}*INDEX(B4_EffOut,1,{t+1}))/1E6",
                fmt="0.0000", fill=FILL_KEY)
        r += 1
    for g, (lab, wr) in enumerate([("VR200", WL["cost0"]), ("GB300", WL["cost0"] + 3)]):
        for t, (tk, tn) in enumerate(TIERS):
            put(ws, f"A{r}", f"{lab} × {tk} 每任務營收 ÷ 服務成本（Workload）"); put(ws, f"B{r}", "x")
            for c in "CDEFG":
                put(ws, f"{c}{r}", f"=IF(ISNUMBER(Workload!{c}{wr+t}),IF(Workload!{c}{wr+t}>0,{c}{T['task0']+t}/Workload!{c}{wr+t},0),\"SLO 不可達\")", fmt="0.0")
            r += 1
    put(ws, f"A{r}", "註：Workload 每任務成本未含快取儲存與攤提；任務的快取命中率取 Workload 列 11，與 A 節參考請求的 χ 不同", F_NOTE)
    return T

# ---------------------------------------------------------------- Sens_Rev
def sens_rev(wb, K, T, A):
    ws = wb.create_sheet("Sens_Rev")
    title(ws, "Sens_Rev — Block 4 敏感度（VR200 × Sol、基準成本；單一輸入變動，其餘不變）",
          "營收 ∝ 利用率 ×（1−折扣）× 混合單價（χ）；攤提 ∝ 1 ÷ 商業壽命。本頁以封閉式重算，不改動 Cap_In")
    for c, w in zip("ABCDEFG", [56, 12, 14, 14, 14, 14, 50]): ws.column_dimensions[c].width = w
    for c, h in zip("ABCDEFG", ["變動", "值", "Sol 營收 ÷ 持有成本", "相對基準", "機隊營收 ÷ 持有成本", "相對基準", "說明"]):
        put(ws, f"{c}4", h, F_BOLD, wrap=True)
    col = "M"   # VR200 base
    base_rh = f"Theory_Rev!{col}{T['rh1']}"; base_fl = f"Theory_Rev!{col}{T['flr']}"
    isl, osl = "INDEX(B4_ISL,1,2)", "INDEX(B4_OSL,1,2)"
    def mix(chi): return f"(({isl}*((1-{chi})*INDEX(B4_PriceIn,1,2)+{chi}*INDEX(B4_PriceCache,1,2))+{osl}*INDEX(B4_PriceOut,1,2))/({isl}+{osl}))"
    rows = [
      ("基準", "—", "1", "1", "Cap_In 基準"),
      ("利用率 40%", 0.4, "B{r}/IF_Util", "B{r}/IF_Util", "營收與服務 token 同比；持有成本不變"),
      ("利用率 80%", 0.8, "B{r}/IF_Util", "B{r}/IF_Util", ""),
      ("折扣 10%", 0.1, "(1-B{r})/(1-B4_Disc)", "(1-B{r})/(1-B4_Disc)", ""),
      ("折扣 30%", 0.3, "(1-B{r})/(1-B4_Disc)", "(1-B{r})/(1-B4_Disc)", ""),
      ("快取命中 χ 30%", 0.3, mix("B{r}") + "/" + mix("B4_CacheHit"), "—", "只算 Sol 混合單價；機隊需重算各層級"),
      ("快取命中 χ 75%", 0.75, mix("B{r}") + "/" + mix("B4_CacheHit"), "—", ""),
      ("對外服務占機隊 30%", 0.3, "1", "B{r}/B4_Serve", "單一層級滿載不受影響"),
      ("對外服務占機隊 60%", 0.6, "1", "B{r}/B4_Serve", ""),
      ("免費占服務 35%", 0.35, "1", "(1-B{r})/(1-B4_Free)", ""),
      ("免費占服務 60%", 0.6, "1", "(1-B{r})/(1-B4_Free)", ""),
      ("前緣單價取代 OpenAI（Sol）", "—", "INDEX(B4_FrontRef,1,2)/INDEX(B4_EffRef,1,2)", "—", "K2 (i)：前緣即 OpenAI 時為 1"),
      ("ε＝0.5、有效算力 ×7（J8 缺口全數轉為能力）", 7, "EXP(0.5*LN(B{r}))/Price_Frontier!D{capf}", "EXP(0.5*LN(B{r}))/Price_Frontier!D{capf}", "K4 (a) 若採用的量級；不入基準"),
    ]
    r = 5
    for lab, v, fr, ff, note in rows:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", v, fmt="0%" if isinstance(v, float) else "0")
        fr2 = fr.replace("{r}", str(r)).replace("{capf}", str(K.get("_capf", 0)))
        ff2 = ff.replace("{r}", str(r)).replace("{capf}", str(K.get("_capf", 0)))
        put(ws, f"D{r}", f"={fr2}", fmt="0.00"); put(ws, f"C{r}", f"={base_rh}*D{r}", fmt="0.0", fill=FILL_KEY)
        if ff2 == "—":
            put(ws, f"F{r}", "—"); put(ws, f"E{r}", "—")
        else:
            put(ws, f"F{r}", f"={ff2}", fmt="0.00"); put(ws, f"E{r}", f"={base_fl}*F{r}", fmt="0.00", fill=FILL_KEY)
        put(ws, f"G{r}", note, F_NOTE, wrap=True); r += 1
    r += 1
    put(ws, f"A{r}", "攤提（自下而上，Sol）對商業壽命", F_BOLD); r += 1
    for lab, m in [("壽命 6 個月", 6), ("壽命 24 個月", 24)]:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", m, fmt="0")
        put(ws, f"C{r}", f"=Amortize!{col}{A['bu1']}*INDEX(B4_Life,1,2)/B{r}", fmt="0.00000", fill=FILL_KEY)
        put(ws, f"D{r}", f"=C{r}/Amortize!{col}{A['bu1']}", fmt="0.00"); put(ws, f"G{r}", "$/M；與 Sol 有效混合單價相比仍極小", F_NOTE); r += 1
    put(ws, f"A{r+1}", "未列入：生產折減（Serving!C18）與利用率之外的產能變動需重算 Perf，見 Sens_Perf；本頁不重算物理產能", F_NOTE)

# ---------------------------------------------------------------- Interface D, Checks, Sources
def interface_b4(wb, start, P, S, A, T):
    ws = wb["Interface"]; r = start; names = []
    section(ws, r, "D. Block 4 產出（理想上限；單價為 2026-09／10 快照、各世代共用（K9）；IF_RevGW 等欄＝世代 × 成本情境）", 17); r += 1
    for t, (tk, tn) in enumerate(TIERS):
        put(ws, f"A{r}", tn, F_BOLD); r += 1
        singles = [(f"IF_PriceFresh_{tk}", "新鮮輸入 $/M — OpenAI 有效", "ein"), (f"IF_PriceCached_{tk}", "快取輸入 $/M — OpenAI 有效", "ec"),
                   (f"IF_PriceThink_{tk}", "思考 $/M — OpenAI 有效（依輸出計費）", "eth"), (f"IF_PriceOut_{tk}", "可見輸出 $/M — OpenAI 有效", "eout"),
                   (f"IF_PriceRef_{tk}", "參考請求混合 $/M — OpenAI 有效", "eref"), (f"IF_FrontRef_{tk}", "參考請求混合 $/M — 前緣有效", "fref")]
        for name, lab, key in singles:
            put(ws, f"A{r}", f"{lab}　[{name}]"); put(ws, f"B{r}", "$/M")
            put(ws, f"C{r}", f"=Price_Frontier!{TC[t]}{P[key]}", fmt="#,##0.000", fill=FILL_KEY if key == "eref" else None)
            names.append((name, f"Interface!$C${r}")); r += 1
        put(ws, f"A{r}", f"前緣模型（顯示）　[IF_FrontModel_{tk}]"); put(ws, f"C{r}", f"=Price_Frontier!{TC[t]}{P['frm']}")
        names.append((f"IF_FrontModel_{tk}", f"Interface!$C${r}")); r += 1
        put(ws, f"A{r}", f"商業壽命　[IF_Life_{tk}]"); put(ws, f"B{r}", "月"); put(ws, f"C{r}", f"=INDEX(B4_Life,1,{t+1})", fmt="0")
        names.append((f"IF_Life_{tk}", f"Interface!$C${r}")); r += 1
        rows = [(f"IF_CacheStore_{tk}", "快取儲存 $/M 快取命中 token", "$/M", "#,##0.00000", f"=Cache_Store!{{X}}{S[tk]}"),
                (f"IF_AmortBU_{tk}", "訓練攤提 自下而上＝回本所需溢價", "$/M", "#,##0.00000", f"=Amortize!{{X}}{A[f'bu{t}']}"),
                (f"IF_AmortTD_{tk}", "訓練攤提 由上而下（機隊訓練占比）", "$/M", "#,##0.0000", f"=Amortize!{{X}}{A[f'td{t}']}"),
                (f"IF_FullCost_{tk}", "全成本 $/M（服務＋快取儲存＋自下而上攤提；基準利用率）", "$/M", "#,##0.0000", f"=Theory_Rev!{{X}}{T[f'fc{t}']}"),
                (f"IF_AmortDefault_{tk}", "訓練攤提 K6 預設 (c)：由上而下總額 × 自下而上權重（下游預設）", "$/M", "#,##0.00000", f"=Amortize!{{X}}{A[f'dc{t}']}"),
                (f"IF_AmortRev_{tk}", "訓練攤提 (d)：由上而下總額 × 營收權重", "$/M", "#,##0.00000", f"=Amortize!{{X}}{A[f'dd{t}']}"),
                (f"IF_FullCostDefault_{tk}", "全成本 $/M（服務＋快取儲存＋K6 預設攤提；基準利用率；下游預設）", "$/M", "#,##0.0000", f"=Theory_Rev!{{X}}{T[f'fcd{t}']}"),
                (f"IF_RevGW_{tk}", "每 GW 理論營收 — OpenAI 有效單價（理想上限）", "$B/年", "#,##0.0", f"=Theory_Rev!{{X}}{T[f'rev{t}']}"),
                (f"IF_RevGWFront_{tk}", "每 GW 理論營收 — 前緣單價（理想上限）", "$B/年", "#,##0.0", f"=Theory_Rev!{{X}}{T[f'revf{t}']}")]
        for name, lab, unit, fmt, tpl in rows:
            put(ws, f"A{r}", f"{lab}　[{name}]"); put(ws, f"B{r}", unit)
            for X in COLS15: put(ws, f"{X}{r}", tpl.replace("{X}", X), fmt=fmt, fill=FILL_KEY if ("RevGW_" in name or "Default" in name) else None)
            names.append((name, f"Interface!$C${r}:$Q${r}")); r += 1
    put(ws, f"A{r}", "1 GW 參考機隊（合計為層級組合：付費 token 依 Cap_In 組合；各層級貢獻另列，合計＝三層級貢獻之和）", F_BOLD); r += 1
    fleet_rows = [(f"IF_RevGWFleet_{tk}", f"付費服務營收 {tn} 貢獻 — OpenAI 有效單價", "$B/年", "#,##0.00", f"=Theory_Rev!{{X}}{T[f'flt{t}']}")
                  for t, (tk, tn) in enumerate(TIERS)]
    fleet_rows += [("IF_RevGWFleet", "付費服務營收合計（層級組合）— OpenAI 有效單價（理想上限）", "$B/年", "#,##0.0", f"=Theory_Rev!{{X}}{T['fl']}")]
    fleet_rows += [(f"IF_RevGWFleetFront_{tk}", f"付費服務營收 {tn} 貢獻 — 前緣單價", "$B/年", "#,##0.00", f"=Theory_Rev!{{X}}{T[f'flft{t}']}")
                   for t, (tk, tn) in enumerate(TIERS)]
    fleet_rows += [("IF_RevGWFleetFront", "付費服務營收合計（層級組合）— 前緣單價（理想上限）", "$B/年", "#,##0.0", f"=Theory_Rev!{{X}}{T['flf']}")]
    for name, lab, unit, fmt, tpl in fleet_rows:
        put(ws, f"A{r}", f"{lab}　[{name}]"); put(ws, f"B{r}", unit)
        for X in COLS15: put(ws, f"{X}{r}", tpl.replace("{X}", X), fmt=fmt, fill=FILL_KEY)
        names.append((name, f"Interface!$C${r}:$Q${r}")); r += 1
    for name, lab, ref in [("IF_ServeShare", "對外服務占機隊（第 0 層參考；下游覆寫）", "=B4_Serve"), ("IF_FreeShare", "免費占服務（第 0 層參考；下游覆寫）", "=B4_Free")]:
        put(ws, f"A{r}", f"{lab}　[{name}]"); put(ws, f"B{r}", "%"); put(ws, f"C{r}", ref, fmt="0%")
        names.append((name, f"Interface!$C${r}")); r += 1
    for n, ref in names: nm(wb, n, ref)
    # replace the old placeholder row
    for row in ws.iter_rows(min_row=17, max_row=start):
        if row[0].value == "每 GW 理論營收":
            row[0].value = "每 GW 理論營收"; ws.cell(row=row[0].row, column=3).value = "見 D 節（IF_RevGW_*、IF_RevGWFleet）"
    return names

def checks_b4(wb, P, F, A, T, U):
    ws = wb["Checks"]
    r0 = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    section(ws, r0, "Block 4 檢查（D＝Hopper 基準、G＝GB200 基準、M＝VR200 基準）", 6)
    put(ws, f"A{r0+1}", "項目", F_BOLD); put(ws, f"B{r0+1}", "本模型", F_BOLD); put(ws, f"C{r0+1}", "外部參照", F_BOLD)
    put(ws, f"D{r0+1}", "單位", F_BOLD); put(ws, f"E{r0+1}", "判讀", F_BOLD); put(ws, f"F{r0+1}", "來源", F_BOLD)
    r = r0 + 2
    # constants (blue) for the OpenAI reconciliation
    put(ws, f"H{r0+1}", "對帳常數", F_BOLD)
    consts = [("OpenAI 2025 營收 $B", 13.07), ("2024 年底 GW", 0.6), ("2025 年底 GW", 1.9), ("2025 Hopper 占機隊", 0.6)]
    cr = {}
    for i, (lab, v) in enumerate(consts):
        put(ws, f"H{r0+2+i}", lab, F_NOTE); put(ws, f"I{r0+2+i}", v, F_IN, fmt="0.00"); cr[i] = f"$I${r0+2+i}"
    ucd = lambda t, key: f"Unit_Cost!D{U[(t, key)]}"
    pf = lambda j: f"INDEX(B4_MktOut,{j},1)"
    rows = [
      ("OpenAI 2025 對帳：機隊付費營收（Hopper／GB200 加權、OpenAI 有效單價）",
       f"=IF(AND(ISNUMBER(Theory_Rev!D{T['fl']}),ISNUMBER(Theory_Rev!G{T['fl']})),{cr[3]}*Theory_Rev!D{T['fl']}+(1-{cr[3]})*Theory_Rev!G{T['fl']},\"SLO 不可達\")", f"={cr[0]}/(({cr[1]}+{cr[2]})/2)", "$B/GW 年",
       "外部參照＝2025 營收 ÷ 平均 GW（GW 口徑未明，D1）；同量級即機隊配置與單價可閉合。2025 實際單價高於 2026 快照", "Cap_In D 節；E010"),
      ("DeepSeek V4-Pro 尖峰輸出價 ÷ 本模型 Hopper Sol decode 成本（基準利用率）",
       f"=IF(ISNUMBER({ucd(2,'cdu')}),{pf(8)}/{ucd(2,'cdu')},\"SLO 不可達\")", "≈1", "x", "≈1：中國廠商定價接近 Hopper 級物理成本；前緣由接近成本者決定（K4 理由 2）", "S52；E011"),
      ("DeepSeek V4-Pro 離峰輸出價 ÷ Hopper Sol decode 成本（100%）",
       f"=IF(ISNUMBER({ucd(2,'cd')}),{pf(8)}*B4_OffPeak/{ucd(2,'cd')},\"SLO 不可達\")", "≈1", "x", "", "S52"),
      ("DeepSeek V4.1-Flash 尖峰輸出價 ÷ Hopper Luna decode 成本（基準利用率）",
       f"=IF(ISNUMBER({ucd(1,'cdu')}),{pf(7)}/{ucd(1,'cdu')},\"SLO 不可達\")", "≈1", "x", "V4.1 架構未必同於 V4-Flash；硬體與中國資本、電力成本不同", "S52"),
      ("單價前緣模型：Luna", f"=Price_Frontier!C{P['frm']}", "—", "", "K2 (i)", "Price_Frontier"),
      ("單價前緣模型：Sol", f"=Price_Frontier!D{P['frm']}", "—", "", "", "Price_Frontier"),
      ("單價前緣模型：Astra", f"=Price_Frontier!E{P['frm']}", "—", "", "Claude Opus 5.5（$4／$20）指數未取得；若 ≥ Astra 門檻，前緣將下移", "Price_Frontier"),
      ("中國合格者最低 ÷ OpenAI（Luna）", f"=Price_Frontier!C{P['cnr']}", ">1", "x", ">1：同能力下中國廠商未比 OpenAI 便宜", "Price_Frontier"),
      ("隱含每 GW 年家族研發計畫數（VR200）", f"=Fleet_1GW!M{F['nprog']}", "1–3", "個", "遠高於 1–3：J8 單一計畫規模偏小或訓練占比偏高", "Fleet_1GW"),
      ("Fleet_1GW 服務 GW 閉合（VR200；應為 0）", f"=Fleet_1GW!M{F['chk']}", "0", "GW", "", "Fleet_1GW"),
      ("自下而上攤提 ÷ Sol 有效混合單價（VR200）", f"=Amortize!M{A['bu1']}/INDEX(B4_EffRef,1,2)", "—", "%", "回本所需溢價占單價的比例", "Amortize"),
      ("K6 縮放倍數：由上而下 ÷ 自下而上年化總額（VR200）", f"=Amortize!M{A['scale']}", "1–3 個計畫時約 1", "x", "與『隱含家族研發計畫數』同源（J8）；(c) 以此倍數放大自下而上攤提", "Amortize"),
      ("K6 (c)：Astra 預設攤提 ÷ Astra 有效混合單價（VR200）", f"=Amortize!M{A['dc2']}/INDEX(B4_EffRef,1,3)", "—", "%", "預設口徑下 Astra 單價中訓練攤提所占比例", "Amortize"),
      ("機隊各層級貢獻合計 − 機隊合計（VR200；應為 0）", f"=Theory_Rev!M{T['flt0']}+Theory_Rev!M{T['flt1']}+Theory_Rev!M{T['flt2']}-Theory_Rev!M{T['fl']}", "0", "$B/年", "", "Theory_Rev B 節"),
    ]
    for i, (a, b, c, d, e, f) in enumerate(rows):
        rr = r + i
        put(ws, f"A{rr}", a, wrap=True); put(ws, f"B{rr}", b, fmt="0.00%" if d == "%" else "#,##0.00", fill=FILL_KEY)
        put(ws, f"C{rr}", c, fmt="#,##0.00"); put(ws, f"D{rr}", d); put(ws, f"E{rr}", e, F_NOTE, wrap=True); put(ws, f"F{rr}", f, F_NOTE)

SOURCES_B4 = [
  ("S50", "OpenAI API 定價頁（GPT-6 Astra／Sol／Luna）", "Standard：Astra $10／$1／$50、Sol $2／$0.2／$10、Luna $0.1／$0.01／$0.5（輸入／快取／輸出）；推理 token 依輸出計費",
   "Verified", "2026-09-25", "developers.openai.com/api/docs/pricing（OpenAI 模型 v0.5 已核對）", "已核對"),
  ("S51", "Anthropic API 定價（二手）", "Fable 5／5.1 $10／$50；Opus 5 $5／$25；Opus 5.5 $4／$20（2026-09-22）；快取寫入 5 分鐘 1.25 倍、1 小時 2 倍；思考 token 依輸出計費",
   "Interested-party（Anthropic 定價；經 dev.to、g2、eesel、orcarouter 轉述）", "2026-07／09", "docs.claude.com 定價頁（未直接讀取）", "多個二手來源一致；Fable 5.1 快取讀取 $0.25 僅單一來源"),
  ("S52", "DeepSeek API 官方定價頁", "V4.1-Flash 尖峰 $0.30／$0.006／$1.20；V4-Pro-0813 尖峰 $1.32／$0.044／$3.96；離峰半價；尖峰＝週一至五 UTC 01–04、06–10",
   "Verified", "2026-10-01", "api-docs.deepseek.com/quick_start/pricing", "已核對原文；第三方頁多為過時價"),
  ("S53", "Z.ai 官方定價頁", "GLM-5.3 $1.4／$0.26／$4.4；GLM-5.3-Flash $0.15／$0.03／$0.5；快取儲存限時免費", "Verified", "2026-10-01",
   "docs.z.ai/guides/overview/pricing", "已核對原文"),
  ("S54", "Moonshot、Alibaba、MiniMax 定價（二手）", "Kimi K3 $3／$0.30／$15；Qwen3.8-Max 國際站 $2／$0.25／$6；MiniMax M3 $0.30／$1.20",
   "Interested-party（廠商定價，經彙整站轉述）", "2026-07／09", "benchlm、morphllm、developersdigest、aireiter 等", "官方頁未直接讀取；Qwen 有 $2.5／$7.5 附五折之衝突記載"),
  ("S55", "Artificial Analysis Intelligence Index v4.3（經報導）", "GPT-6 Astra 53、Fable 5.1 53、Opus 5 51、GPT-5.6 Sol 47、GLM-5.3-Flash 42、Qwen3.8 2.4T 40、DeepSeek V4 Pro 36；一週內三次改版",
   "Verified-measured（獨立評測；改版頻繁，跨版本不可比）", "2026-09", "trendingtopics.eu 報導", "已核對報導原文"),
  ("S56", "Artificial Analysis 發布比較頁 v4.3.2", "GPT-6 Sol 48、GPT-6 Luna 37、DeepSeek V4.1-Flash 39、GLM-5.3 45（max 設定）", "Verified-measured", "2026-09",
   "artificialanalysis.ai/models/releases/comparisons/…", "已核對搜尋摘錄"),
  ("S57", "OpenAI 2025 推論與訓練支出（OpenAI 模型 v0.5 已收錄）", "推論 $8.4B（非付費 $3.9B）、訓練約 $12B；營收 $13.07B；GW 0.6（2024 底）→ 約 1.9（2025 底）",
   "Interested-party（OpenAI 投資人資料，經 The Information、FT、CFO 部落格）", "2026", "openai_token_revenue.json", "待查核（報導）"),
]

def sources_b4(wb):
    ws = wb["Sources"]
    r = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 1
    for i, row in enumerate(SOURCES_B4):
        for c, v in zip("ABCDEFG", row): put(ws, f"{c}{r+i}", v, wrap=True)

EVID_B4 = [
  ("E010", "2026-10-01", "OpenAI 2025 機隊：服務約 41%、免費占服務約 46%（支出比）", "S57", "Interested-party", "Cap_In D 節", "—（新增）", "41%／46%", "採納", "v5.8", "以支出比代 GW 比，為 Derived"),
  ("E011", "2026-10-01", "DeepSeek 官方現價：V4.1-Flash $0.30／$1.20、V4-Pro $1.32／$3.96（尖峰），離峰半價", "S52", "Verified", "Cap_In F 節；Checks", "第三方頁 $0.14／$0.28", "官方 $0.30／$1.20", "採納", "v5.8", "第三方彙整頁多為過時價；只採官方頁"),
  ("E012", "2026-10-01", "中國廠商 2026 年調漲價格（DeepSeek Flash、GLM 5.1、Kimi K2.5、MiniMax）", "S52、S54", "Interested-party", "Cap_In F 節", "—", "上漲約 30% 至 4 倍", "部分採納", "v5.8", "說明單價非單向下降；年降幅（K9 (b)）不入基準"),
  ("E013", "2026-10-01", "Artificial Analysis 指數一週改版三次，GPT-6 Astra 由第五升至第一", "S55", "Verified-measured", "Cap_In F 節 能力指數", "—", "v4.3／v4.3.2", "部分採納", "v5.8", "全表固定同一版本；跨版本不比較"),
  ("E014", "2026-10-01", "Claude Opus 5.5 $4／$20，Anthropic 自稱達 Fable 5.1 水準", "S51", "Interested-party", "Price_Frontier Astra 前緣", "—", "未納入前緣", "待查", "—", "待獨立評測指數；若 ≥ 53，Astra 前緣下移至 $4／$20"),
  ("E015", "2026-10-01", "Kimi K3 為開放權重（Artificial Analysis 標示）／閉源（另一來源）", "S54", "Interested-party", "Cap_In F 節 開放權重欄", "—", "有爭議", "待查", "—", ""),
]

def evidence_b4(wb):
    ws = wb["DB_Evidence"]
    have = {ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1)}
    r = max(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value is not None) + 1
    n = 0
    for row in EVID_B4:
        if row[0] in have: continue
        for i, v in enumerate(row): put(ws, f"{L(i+1)}{r}", v, F_IN, wrap=i in (2, 5, 10))
        ws.row_dimensions[r].height = 30; r += 1; n += 1
    return n
