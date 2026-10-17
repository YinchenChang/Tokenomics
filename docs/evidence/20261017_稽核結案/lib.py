"""稽核結案對帳：載入現行活頁簿（唯讀）。"""
import openpyxl, os, functools
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..'))
CUR = open(os.path.join(ROOT, 'model', 'CURRENT')).read().strip()
PATH = os.path.join(ROOT, 'model', CUR)


@functools.lru_cache(None)
def val():
    return openpyxl.load_workbook(PATH, data_only=True)


@functools.lru_cache(None)
def frm():
    return openpyxl.load_workbook(PATH, data_only=False)


def name_value(n):
    wb = val()
    dn = wb.defined_names.get(n)
    if dn is None:
        return None
    out = []
    for sh, ref in dn.destinations:
        ws = wb[sh]
        c = ws[ref.replace('$', '')]
        if isinstance(c, tuple):
            out.append([[x.value for x in r] for r in c])
        else:
            out.append(c.value)
    return out


def cell(sheet, ref, formulas=False):
    wb = frm() if formulas else val()
    return wb[sheet][ref].value
