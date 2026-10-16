# v10.3 (work order docs/workorders/20261016_瘦身4-2.md, segment 4-2a; branch claude/slim-4-2a): hard-coded parameters exposed, duplicate definitions merged, marks and tests.
#   F1  hard-coded parameters (4-0 sheet 4, 101 judgement cells, values unchanged): L1_JudgeTol (Calib K section; 96 L1 P cells + I3 + Checks E5 read it), SR_RU_PxMultLo／Hi (Spec_Rack section; G11／G13 read them),
#       L1 M30／N30 already read SRC_PRC_001 (only the /4 remains: note), physical conventions (3x training, 6ND, ln4, attention 4, ...) get a note "物理慣例：…" and stay in the formulas.
#   F2  duplicate definitions (4-0 sheet 2): D08 Sens_Rev B20／B21 read Cap_In D23／D24; D09 Cap_In H10 2025 column C reads the H-section cells; D10 no change (see report); D11 permanent test only.
#   F3  Tech_Registry K and T columns enter Checks V as V18／V19 (counted in CHK_Verify*; failure does not enter GOV_Errors).
#   F4  marks: column 用途 on Calib J (105 dead cells), Alloc (27 dead cells read only by names nobody reads), 70 inputs nobody reads.
#   F7  Decisions SL42-*.
# Minor version: only CHK_Verify* changes value. Every write to an Excel-owned cell is guarded (old value / old text); the new inputs are created before restore() so a rebuild keeps the Excel value.
import re
import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as L
from common import put, F_IN, F_NOTE, F_BOLD, F_CALC, F_LINK, FILL_KEY, section
import v1000
import v1020

VERSION = "20261016_Tokenomics_v10.3"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-16"
V = "v10.3"
WO = "工作單 docs/workorders/20261016_瘦身4-2.md（4-2a 段）"
DEC_BY = "Andy 2026-10-16 20:54「依建議繼續作下去」；F1–F7 為 chat 端預設，可撤回"
DASH = "—"
PURPOSE = "用途"
PURPOSE_HDR_ROW = 3
JUDGE_TOL_LABEL = "L1 外部對照判讀容差（L1_JudgeTol）"
RU_LO_LABEL = "Rubin Ultra 機架價格倍數：低（SR_RU_PxMultLo）"
RU_HI_LABEL = "Rubin Ultra 機架價格倍數：高（SR_RU_PxMultHi）"
RU_SEC = "Rubin Ultra 機架價格倍數（v10.3 F1；Assumed，無來源；待證據）"

# base retired inputs: Sens_Rev B20／B21 (blue 6／24) are now links to Cap_In D23／D24 (D08)
SNAP_RETIRED = [("Sens_Rev", ("壽命 6 個月", 0), 2), ("Sens_Rev", ("壽命 24 個月", 0), 2)]


# ================================================================== F1: new inputs (written before restore())
def calib_k(wb):
    """Calib K section (below J): L1_JudgeTol = 0.2 (Assumed; range 0.1–0.3)."""
    ws = wb["Calib"]
    r0 = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    section(ws, r0, "K. 判讀參數（v10.3 F1）：不是物理量，只決定 L1 外部對照欄 P 的文字判讀", 10)
    r = r0 + 1
    put(ws, f"A{r}", JUDGE_TOL_LABEL); put(ws, f"B{r}", "比例")
    put(ws, f"C{r}", 0.2, F_IN, fmt="0%")
    put(ws, f"J{r}", "Assumed｜L1 外部對照的判讀容差，非物理量：外部為單一值時，|推算 ÷ 外部 − 1| ≤ 本值判為「±x% 內」，否則「差距 >x%」（L1 欄 P、Checks I3、Checks E5 讀取）。"
        "v10.3 前寫死在 96 格公式裡（0.2），外露後數值不變。區間 0.1–0.3（Gov_Map）；改本格不影響任何 IF_／L1_ 數值，只改 P 欄文字與 Checks I3 的計數", F_NOTE, wrap=True)
    wb.defined_names["L1_JudgeTol"] = DefinedName("L1_JudgeTol", attr_text=f"Calib!$C${r}")
    return {"sec": r0, "tol": r}


def spec_rack_ru(wb, V602):
    """Spec_Rack section below the model-scale rows: Rubin Ultra rack price multipliers (G11 = F12 × lo, G13 = F12 × hi)."""
    ws = wb["Spec_Rack"]
    r0 = V602["rows"][-1] + 2
    section(ws, r0, RU_SEC, 8)
    rows = {}
    for i, (lab, v, name, note) in enumerate([
            (RU_LO_LABEL, 1.5, "SR_RU_PxMultLo", "Assumed（無來源；待證據）｜Rubin Ultra 單架機架價格低端＝VR200 機架價格（F12，售價口徑）× 本值；v10.3 前寫死在 G11 公式（×1.5）。Gov_Map 區間 1.2–2.0"),
            (RU_HI_LABEL, 2.5, "SR_RU_PxMultHi", "Assumed（無來源；待證據）｜Rubin Ultra 單架機架價格高端＝F12 × 本值；v10.3 前寫死在 G13 公式（×2.5）。Gov_Map 區間 2.0–3.0")]):
        r = r0 + 1 + i
        put(ws, f"A{r}", lab); put(ws, f"B{r}", "x"); put(ws, f"C{r}", v, F_IN, fmt="0.00"); put(ws, f"H{r}", note, F_NOTE)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Spec_Rack!$C${r}")
        rows[name] = r
    return {"sec": r0, "rows": rows}


def inputs_update(wb):
    """After restore(): Spec_Rack G11／G13 read the multipliers (old-formula guard); G14 note names them instead of numbers."""
    log = []
    ws = wb["Spec_Rack"]
    for cell, old, new in (("G11", "=F12*1.5", "=F12*SR_RU_PxMultLo"), ("G13", "=F12*2.5", "=F12*SR_RU_PxMultHi")):
        if ws[cell].value == old:
            ws[cell].value = new; log.append(f"Spec_Rack!{cell}: {old} -> {new} (v10.3 F1)")
    old14 = "72 封裝（單架）；價格＝VR200 機架價格（F12，售價口徑）× 1.5／2.0／2.5（公式；倍數 Assumed；J2）"
    new14 = "72 封裝（單架）；價格＝VR200 機架價格（F12，售價口徑）× 低倍數 SR_RU_PxMultLo／2.0／高倍數 SR_RU_PxMultHi（公式；倍數 Assumed，無來源，待證據；J2；v10.3 F1 外露，倍數見本頁「" + RU_SEC + "」）"
    if ws["G14"].value == old14:
        ws["G14"].value = new14; log.append("Spec_Rack!G14: text")
    return log


def gov_map_rows(wb):
    """Gov_Map rows for the three new inputs (appended by gov.gm_append_c when absent)."""
    out = []
    ws = wb["Calib"]
    r = next(rr for rr in range(150, ws.max_row + 1) if ws.cell(rr, 1).value == JUDGE_TOL_LABEL)
    out.append({'scope': '切片四（v10.3 瘦身 4-2a）', 'sheet': 'Calib', 'cell': f'C{r}', 'label': JUDGE_TOL_LABEL, 'check': JUDGE_TOL_LABEL, 'cls': 'Assumed', 'role': '單值', 'src': '', 'rel': '',
                'dec': 'SL42-F1', 'lo': 0.1, 'hi': 0.3, 'rtext': '0.1–0.3（chat 端預設）',
                'reason': ("L1 外部對照欄 P 的判讀容差（±20%），v10.3 前寫死在 96 格公式；外露後數值不變。非物理量，只決定文字判讀，不進任何 IF_／L1_ 數值；"
                           "0.2 是慣例值，無來源；區間 0.1–0.3 為 chat 端預設。已搜尋：無可比對象（判讀門檻不是可量測物理量）"), 'retag': '', 'seg': '低'})
    ws = wb["Spec_Rack"]
    for lab, name, lo, hi, base, tag in ((RU_LO_LABEL, "SR_RU_PxMultLo", 1.2, 2.0, 1.5, "低端"), (RU_HI_LABEL, "SR_RU_PxMultHi", 2.0, 3.0, 2.5, "高端")):
        r = next(rr for rr in range(44, ws.max_row + 1) if ws.cell(rr, 1).value == lab)
        out.append({'scope': '切片四（v10.3 瘦身 4-2a）', 'sheet': 'Spec_Rack', 'cell': f'C{r}', 'label': lab, 'check': lab, 'cls': 'Assumed', 'role': '單值', 'src': '', 'rel': '',
                    'dec': 'SL42-F1', 'lo': lo, 'hi': hi, 'rtext': f'{lo}–{hi}（chat 端預設）',
                    'reason': (f"Rubin Ultra 單架機架價格{tag}倍數（＝VR200 機架價格 F12 × 本值，基準 {base}）；v10.3 前寫死在 Spec_Rack G{'11' if name.endswith('Lo') else '13'} 公式，外露後數值不變。"
                               "無來源，待證據（Rubin Ultra 為推估欄，不進基準比較）；Assumed 依 CLAUDE.md 第 7 節以區間呈現；區間為 chat 端預設。已搜尋：未見 Rubin Ultra 機架定價公開來源"),
                    'retag': '', 'seg': '低'})
    return out


# ================================================================== F2: duplicate definitions (patched right after the sheet is built, before Deps / contract tracing)
def cap_in_patch(wb):
    """D09: Cap_In H10 2025 column C (m_API／消費訂閱／商用席次／整體、機隊付費 token、m_fleet) reads the H-section cells that already hold the same value (4-0 D09).
    Rows are found by label and every old formula is checked, so a layout change stops the build instead of linking the wrong row."""
    ws = wb["Cap_In"]
    lab = lambda r: str(ws.cell(r, 1).value or "")
    find = lambda start: [r for r in range(1, ws.max_row + 1) if lab(r).startswith(start)]
    pairs = [("m_API", "m_API（2025，三分法）", "206"), ("m_消費訂閱", "m_消費訂閱（2025）", "207"), ("m_商用席次", "m_商用席次（2025）", "208"),
             ("m_整體＝", "隱含 m — 合併（模型 token）", "209"), ("機隊付費 token（點估計）", "機隊付費 token（點估計）", "222"),
             ("m_fleet＝m 分子營收", "機隊 m（點估計）", "223")]
    log = []
    i_start = next(r for r in range(150, ws.max_row + 1) if lab(r).startswith("H10-3."))
    for dst_lab, src_lab, _ in pairs:
        dst = [r for r in find(dst_lab) if r > i_start]
        if dst_lab == "機隊付費 token（點估計）":
            dst = [r for r in dst if r > i_start][:1]
        src = [r for r in find(src_lab) if r < i_start]
        assert len(dst) == 1 and len(src) == 1, (dst_lab, dst, src_lab, src)
        d, s = dst[0], src[0]
        old = ws[f"C{d}"].value
        new = f"=C{s}"
        if old != new:
            assert isinstance(old, str) and old.startswith("="), (d, old)
            ws[f"C{d}"].value = new; ws[f"C{d}"].font = F_CALC
            log.append((d, s))
    return log


# ================================================================== F1／F4: notes, marks (end of build, idempotent)
PHYS = {   # (sheet, row): (note column or None = purpose column, text)
    ("Spec_Rack", 27): ("H", "物理慣例：÷8＝Gb/s 換 GB/s（8 bit／byte），不外露"),
    ("Spec_Rack", 34): ("H", "物理慣例：×8＝pJ/bit 換 pJ/byte（8 bit／byte），不外露"),
    ("Arch", 30): ("G", "物理慣例：係數 4＝注意力 QKᵀ 與 AV 各 2 FLOPs／MAC，共 4，不外露"),
    ("Training", 34): (None, "物理慣例：3＝前向 1＋反向 2（訓練 FLOPs＝3 × 前向），不外露；若日後要改為 Tech_Registry 條目再外露"),
    ("Training", 41): (None, "物理慣例：3＝前向 1＋反向 2（訓練 FLOPs＝3 × 前向），不外露"),
    ("Training", 48): (None, "物理慣例：3＝前向 1＋反向 2（訓練 FLOPs＝3 × 前向），不外露"),
    ("Training", 69): (None, "物理慣例：3＝前向 1＋反向 2（訓練 FLOPs＝3 × 前向），不外露"),
    ("Training", 74): (None, "物理慣例：÷3＝訓練 FLOPs 折回前向（前向 1＋反向 2），不外露"),
    ("Training", 76): (None, "物理慣例：÷3＝訓練 FLOPs 折回前向（前向 1＋反向 2），不外露"),
    ("Cap_In", 86): ("G", "數學慣例：^0.5＝幾何中點（低、高的幾何平均），不外露"),
    ("Cap_In", 294): ("G", "物理慣例：6ND＝訓練 FLOPs ≈ 6 × 啟用參數 × token（前向 2 ＋ 反向 4），不外露"),
    ("Cap_In", 295): ("G", "物理慣例：6ND（同上），不外露"),
    ("Cap_In", 296): ("G", "物理慣例：6ND（同上），不外露"),
    ("Har_In", 77): (None, "數學慣例：LN(4)＝logit(0.8) − logit(0.5)（50% 與 80% 成功率的 logit 差），由時間範圍比值推隱含 β，不外露"),
    ("Har_In", 78): (None, "數學慣例：LN(4)（同上），不外露"),
    ("Har_In", 79): (None, "數學慣例：LN(4)（同上），不外露"),
    ("Har_In", 80): (None, "數學慣例：LN(4)（同上），不外露"),
    ("Alloc", 17): ("G", "索引慣例：3*(g−1)+2＝Interface 欄依「世代 × 3 層級」排列，取世代 g 的第 2 層級（基準成本欄）；3＝層級數，不外露"),
    ("Alloc", 43): ("G", "索引慣例：3*(g−1)+2（同第 17 列），不外露"),
    ("Alloc", 46): ("G", "結構慣例：COUNT(…)=3＝三層級（Luna、Sol、Astra）皆有數才合計，不外露"),
    ("Alloc", 49): ("G", "結構慣例：COUNT(…)=3＝三層級皆有數才合計，不外露"),
    ("Fleet_1GW", 33): (None, "結構慣例：COUNT(…)=3＝三層級（Luna、Sol、Astra）皆有數才合計，不外露（4-0 H99 判為層級數）"),
}
L1_NOTES = {30: " ｜v10.3 F1：M30、N30 已直接引用 SRC_PRC_001（CoreWeave 隨需牌價，每實例）；「÷ 4」＝該實例的 GPU 數（物理慣例，不外露）"}

# 4-0 sheet 6 G5: inputs that no formula reads. Row -> text (applies to the listed columns only)
_END = "區間端點（供 Gov_Map 與情境）"
_SPARE = "備用（無公式讀取）"
UNREAD = {
    "Spec_Rack": {39: f"{_SPARE}：G 欄 BF16 dense（Rubin Ultra）", 45: f"{_END}：ARCH_ScaleGen_Lo（C:G）", 46: f"{_END}：ARCH_ScaleGen_Hi（C:G）"},
    "Workload": {40: f"{_SPARE}：C 欄（參照基準 1；E、F 欄連結 SRC）"},
    "Train_In": {46: f"{_END}：前沿預訓練錨點低（C）；J8 對照用", 48: f"{_END}：前沿預訓練錨點高（C）；J8 對照用"},
    "Cap_In": {13: f"{_END}：D、E 欄", 14: f"{_END}：D、E 欄", 15: f"{_END}：D、E 欄", 16: f"{_END}：D、E 欄（尖峰／離峰口徑情境）",
               23: f"{_END}：C、E 欄（D 欄由 Sens_Rev B20 讀取）", 24: f"{_END}：C、E 欄（D 欄由 Sens_Rev B21 讀取）",
               28: f"{_END}：D、E 欄", 29: f"{_END}：D、E 欄", 36: f"{_END}：D、E 欄", 37: f"{_END}：D、E 欄", 38: f"{_END}：D、E 欄", 39: f"{_END}：D、E 欄",
               40: f"{_END}：D、E 欄", 41: f"{_END}：D、E 欄", 42: f"{_END}：D、E 欄", 133: f"{_SPARE}：C 欄（基準）；{_END}：D、E 欄（H10 顯示用廣告營收）"},
    "Alloc_In": {9: f"{_END}：D、E 欄", 10: f"{_END}：D、E 欄", 11: f"{_END}：D、E 欄", 12: f"{_END}：D、E 欄", 13: f"{_END}：D、E 欄",
                 29: f"{_END}：D、E 欄（研發世代組合 2025）", 30: f"{_END}：D、E 欄", 31: f"{_END}：D、E 欄", 32: f"{_END}：D、E 欄"},
    "Alloc": {137: f"{_SPARE}：C:E 欄（層級索引常數 1／2／3，無公式讀取）"},
}
# Calib J (4-0 R-rows 167-204): calibration and verification copy rows
CALIB_J_ROWS = range(167, 205)
CALIB_J_TXT = "校準與驗證複製列（顯示用；無下游）"
# Alloc: dead cells whose only readers are names that no formula reads (F4 item 2; the readers are listed in the report)
ALLOC_NAME_ROWS = {
    24: "顯示用：具名範圍 AL_UnservedPaid 無公式讀取", 25: "顯示用：具名範圍 AL_UnservedFree 無公式讀取", 32: "顯示用：具名範圍 AL_UnservedGen 無公式讀取",
    38: "顯示用：具名範圍 AL_ServeGen 無公式讀取", 47: "顯示用：只被 Alloc 第 142 列（單一訓練世代對照）讀取；第 142 列只被第 143 列讀取，第 143 列無人讀取",
    48: "顯示用：只被 Alloc 第 142 列（單一訓練世代對照）讀取；第 143 列無人讀取", 51: "顯示用：具名範圍 AL_RDGWg 只被第 59、60 列讀取，後兩者的具名範圍 AL_Q1g／AL_Q2g 無公式讀取",
    59: "顯示用：具名範圍 AL_Q1g 無公式讀取", 60: "顯示用：具名範圍 AL_Q2g 無公式讀取", 67: "顯示用：具名範圍 AL_J8Gap 無公式讀取",
    135: "顯示用：具名範圍 AL_RDShareExcl 無公式讀取", 161: "顯示用：具名範圍 AL_Q1_Tok 無公式讀取", 162: "顯示用：具名範圍 AL_Q2_Tok 無公式讀取",
}


def _purpose_col(ws, create=True):
    last = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None: last = max(last, cell.column)
    for c in range(3, last + 1):
        if ws.cell(PURPOSE_HDR_ROW, c).value == PURPOSE: return c
    if not create: return None
    col = last + 1
    put(ws, f"{L(col)}{PURPOSE_HDR_ROW}", PURPOSE, F_BOLD)
    ws.column_dimensions[L(col)].width = 46
    return col


def _set_mark(ws, row, text):
    """Writes the purpose text; an existing different text stays and the new one is appended (the 4-1 marks are kept)."""
    col = _purpose_col(ws)
    cell = ws.cell(row, col)
    if cell.value in (None, ""):
        put(ws, cell.coordinate, text, F_NOTE); return 1
    if text in str(cell.value): return 0
    cell.value = f"{cell.value} ｜{text}"; return 1


def notes_and_marks(wb, append_text):
    log = {"phys": 0, "marks_calib_j": 0, "marks_alloc": 0, "marks_unread": 0, "l1_notes": 0}
    for (sh, row), (col, txt) in PHYS.items():
        ws = wb[sh]
        cell = ws.cell(row, _purpose_col(ws)) if col is None else ws[f"{col}{row}"]
        log["phys"] += append_text(cell, " ｜" + txt)
    ws = wb["L1"]
    for row, txt in L1_NOTES.items():
        log["l1_notes"] += append_text(ws[f"I{row}"], txt)
    ws = wb["Calib"]
    for r in CALIB_J_ROWS: log["marks_calib_j"] += _set_mark(ws, r, CALIB_J_TXT)
    ws = wb["Alloc"]
    for r, txt in ALLOC_NAME_ROWS.items(): log["marks_alloc"] += _set_mark(ws, r, txt)
    for sh, rows in UNREAD.items():
        ws = wb[sh]
        for r, txt in rows.items(): log["marks_unread"] += _set_mark(ws, r, txt)
    return log


# ================================================================== end of build
def apply(wb, append_text):
    log = {}
    log["notes_marks"] = notes_and_marks(wb, append_text)
    ws = wb["Gov_Map"]; n = 0
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 3).value == "Sens_Rev" and ws.cell(r, 4).value == "B20:B21":
            n += append_text(ws.cell(r, 14), " ｜v10.3 4-2a D08：B20:B21 已改為公式（讀 Cap_In D23:D24 商業壽命低／高），本列保留登錄（情境值的範圍說明）；改值請改 Cap_In D23:D24")
        if ws.cell(r, 3).value == "Spec_Rack" and ws.cell(r, 4).value in ("G11", "G13"):
            n += append_text(ws.cell(r, 14), " ｜v10.3 4-2a F1：倍數已外露為 SR_RU_PxMultLo／SR_RU_PxMultHi（Spec_Rack「" + RU_SEC + "」）；本格為公式 =F12 × 倍數")
    log["gov"] = n
    log["decisions"] = decisions_append(wb)
    return log


# ================================================================== Decisions (10 columns; layout of v800／v810／v1000／v1010)
_D = [
    ("F1", "寫死參數外露（明細頁 4；數值不變）",
     "新增具名輸入 L1_JudgeTol＝0.2（Calib K 節；Assumed，區間 0.1–0.3；非物理量）：L1 欄 P 的 96 格判讀公式、Checks I3、Checks E5 改讀它。"
     "新增 SR_RU_PxMultLo＝1.5、SR_RU_PxMultHi＝2.5（Spec_Rack「Rubin Ultra 機架價格倍數」節；Assumed，無來源，待證據；區間 1.2–2.0／2.0–3.0）：Spec_Rack G11、G13 改讀它們。"
     "L1 M30、N30 已直接引用 SRC_PRC_001，只剩 /4（每實例 GPU 數）：加註。物理慣例（訓練 3 倍、6ND、ln4、注意力係數 4、位元組換算、層級數 3、索引 3*(g−1)+2）不外露，只在說明欄加註『物理慣例：…』。"
     "Fleet_1GW M33 的 COUNT(…)=3 判為層級數（4-0 列為『待判斷：未分類』），加註不外露。",
     "Calib K 節；Spec_Rack 新節；L1 欄 P；Checks I3／E5；Gov_Map 新增 3 列", "v10.3 已建"),
    ("F2", "重複定義處理（明細頁 2）",
     "D08：Sens_Rev B20／B21 改讀 Cap_In D23／D24（Sol 欄；兩處皆為 6／24，數值不變），標籤改為公式。D09：Cap_In H10 2025 欄 C206–C209、C222、C223 改直接引用 H 節對應格（C144、C145、C146、C94、C116、C119；數值不變），第 224 列一致性檢查保留。"
     "D10：不執行——Alloc 第 122 列已於 4-1 清空，第 121 列 C 欄是 2026 快照價（Cap_In C90），Cap_In 第 268 列 C 欄是 2025 年度價（CAP_Price25Mix），兩處已不是同一個定義，也沒有同一個值；"
     "D11：Alloc 與 Fleet_1GW 單位不同（每 GW 對每 1 GW 機隊），不合併，新增永久測試（tests/parity/test_slim.py）。D01、D02–D07、D12–D16 維持原決定。",
     "Sens_Rev B20:B21；Cap_In 第 206–209、222–223 列；tests/parity/test_slim.py", "v10.3 已建（D10 不執行，待 Project 判斷）"),
    ("F3", "Tech_Registry K、T 欄接進 Checks（原則 4）",
     "Checks V 新增 V18（K 欄：『主流』條目數＝公開採用實驗室數 ≥ 門檻的條目數；公式被覆寫時失敗）與 V19（T 欄：非『一致』的條目數＝0）。計入 CHK_VerifyTotal／Pass／Fail，不進 GOV_Errors。",
     "Checks V 節；Tech_Registry K、T 欄", "v10.3 已建"),
    ("F4", "標記（原則 2）與死格測試加強",
     "Calib J 節 105 個死格填『校準與驗證複製列（顯示用；無下游）』；Alloc 27 個死格查明讀取者（皆為無人讀取的具名範圍或只讀死格的列）後填用途；70 個無人引用的輸入填『區間端點（供 Gov_Map 與情境）』或『備用（無公式讀取）』，不刪任何輸入。"
     "test_slim 新增：死格若不在基準名單（v10.3 未標記死格）內，必須有用途標記。",
     "各頁『用途』欄；tests/parity/test_slim.py；tests/parity/fixtures/dead_unmarked_baseline_v10.3.json", "v10.3 已建"),
    ("F5", "測試補強（查核 D 第 4 項）",
     "契約名稱快照測試（tests/parity/fixtures/contract_snapshot_v10.3.json，相對容差 1e-9；之後每版預期變動須同 PR 更新快照並於報告列出）；驗證清單 V1–V19 逐項斷言；GOV 三項計數在新基準下斷言。",
     "tests/parity/test_snapshot.py；tests/parity/fixtures", "v10.3 已建"),
    ("F7a", "V13、V14 上限 3 保留", "V13、V14 的通過帶上限 3（v10.2 E3 chat 端預設）保留（Andy 2026-10-16 20:54 依建議）。", "Checks V13、V14", "已決定（Andy 2026-10-16 20:54）"),
    ("F7b", "η 校準 E1 停損；Hopper 與 GB200 大批次低估原因待查證",
     "η 校準 E1 停損，C67／D67 未寫入（Hopper 在 DeepSeek profile 條件下即使 η_d 無限大，產出也只到目標的 0.94，所以無解；GB200 的解使有效算力高於 GB300）。"
     "Hopper 與 GB200 在大批次下模型低估實測的原因待查證：(a) 通訊與計算重疊；(b) η_d 隨批次變化。查證另段進行（docs/reports/20261016_大批次低估查證.md）。",
     "Calib J 節；Decisions ETA-E1H／ETA-E1G", "待查證"),
]
DECISIONS_V1030 = [[f"SL42-{i}", V, f"瘦身 4-2a：{t}", txt, DATE, DEC_BY, st, where, f"{WO} {i}", "否"] for i, t, txt, where, st in _D]


def decisions_append(wb):
    ws = wb["Decisions"]
    have = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
    r = max([rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")] or [4]) + 1
    n = 0
    for row in DECISIONS_V1030:
        if row[0] in have: continue
        for i, v in enumerate(row): put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 5, 7))
        r += 1; n += 1
    return n


# ================================================================== README
README_VERSION = (VERSION + "（次版：瘦身 4-2a，" + WO + "——F1 寫死參數外露（L1_JudgeTol、SR_RU_PxMultLo／Hi；物理慣例加註）；F2 重複定義（Sens_Rev 商業壽命、Cap_In H10 2025 欄改引用）；"
                  "F3 Tech_Registry K／T 進 Checks V（V18／V19）；F4 用途標記；Decisions SL42-*。既有 IF_／IFW_／L1_ 名稱數值不變，只有 CHK_Verify* 變動）。以下為 " + v1020.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1020.README_TITLE
README_TITLE = VERSION.split("_")[-1] + _T[len(v1020.VERSION.split("_")[-1]):]
README_ROW = ("瘦身 4-2a（v10.3）",
              "寫死參數外露（L1 判讀容差、Rubin Ultra 價格倍數）；重複定義改引用（商業壽命、Cap_In 2025 欄）；Tech_Registry K／T 進驗證計數；死格與無人引用輸入填用途。細節見 Decisions SL42-*。")
