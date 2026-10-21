# v18.0 (work order docs/workorders/20261021_v18.0_m收入面化.md r0, branch claude/v18-0): revenue-side m and VR200 revenue per GW as P x Q.
# MAJOR version: new contract names (L1 rows 117 onward, each with _Lo／_Hi); no existing named value or existing formula cell changes value.
#   W1  L1 rows: L1_M_2025／_2026Q1／_2026Q2／_FY2026, L1_M (main basis = 2026 Q2, J1), L1_M_API／_Sub／_Seat, L1_PxEff_2026, L1_PxList_2026.
#   W2  L1 rows: L1_TokGW_PaidCap_VR200, L1_TokGW_PaidPQ_VR200 (Q), L1_RevGW_PQ_VR200 (+_FY, _Life).
#   W3  L1_TokLedger／L1_TokLedger_2026Q2 (new rows); label-only tails on L1 rows 76／78／80–82／102 and Cap_In A119／A209／A223 (old m = token-ledger diagnostic).
#   W4  DB_Evidence E399 (Q2 API token current value: searched, not found); Cap_In H196 tail.
#   W5  Decisions SL52-01..06, README, contract document (outside the workbook).
# Every write to an Excel-owned cell is guarded (already-applied test); no input value changes.
from openpyxl.utils import get_column_letter as L

from common import put, F_CALC
import v1000
import v1410
import v1720

VERSION = "20261021_Tokenomics_v18.0"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-21"
V = "v18.0"
WO = "工作單 docs/workorders/20261021_v18.0_m收入面化.md（r0）"
DASH = "—"
DEC_BY = ("Andy 2026-10-21 定義 m（README『變現率 m 的定義』節）；2026-10-21 17:08「v18.0 可以執行」；判斷題 J1–J4 Andy 未個別回應，"
          "依情境卡_工作方式 v1.2『未回應即依預設執行並登錄（可撤回）』採預設")
NR = "無區間"      # the E／F text of rows without a range; written as a formula (="無區間") so that Checks E11 (L1 D／E／F must be formulas) stays 0
NR_LEDGER = "無區間（兩者區間來源不同，不組合）"
NONE_EXT = ("無外部對照", DASH, DASH)


def _t(s):
    return f'="{s}"'


# ---------------------------------------------------------------- rows
I_M = ("v18.0（Andy 2026-10-21 定義）：m＝付費服務 token 的實收單價 ÷ 有效牌價（表列牌價 ×（1 − d），d＝B4_Disc，可調）；只有收入面，不牽涉免費 token、訓練或成本。"
       "(1) 2026 Q1／Q2 的 API token 都用 SRC_DEM_011（2026-03 時點值 15 B／分鐘）；Q2 當期值已搜尋，未找到（DB_Evidence E399），Q2 的 API token 可能偏低、m 可能偏高。"
       "(2) m 的未來變化（壽命期內 m_t ÷ m_2026）未建模，L1_RevGW_PQ_VR200_Life 假設 m 不變。"
       "(3) 2026 欄無區間：輸入（每日則數、消費份額 0.65）為 Assumed 單值，區間待補。")
I_M25 = ("v18.0：2025 欄＝L1_MImplied（模型 token 分母的合併路線），區間取 L1_MImplied_Lo／_Hi。m 的未來變化未建模（見 L1_M）。"
         "分母 token 為模型需求面（Alloc 付費 token），2025 無實測分母；與 token 帳診斷 L1_MFleet 的關係見 L1_TokLedger。")
B_M_BASIS = ("L1_M＝付費服務 token 的實收單價 ÷ 有效牌價（Andy 2026-10-21）；分子＝當期營收（Cap_In 第 187 列，已排除廣告與其他，J3），"
             "分母＝需求面付費 token（API ＋ ChatGPT 付費）× 當年有效牌價。2025 欄用 2025 年度有效牌價（CAP_Price25Mix 1.359），2026 欄用 2026 快照（1.961）——各年用各自的牌價")
I_PQ = ("v18.0：P × Q 由基本定義計算。Q＝模型的 VR200 每 GW 付費產能 × 對外服務占比 ×（1 − 免費占比）× 實際利用率 × 生產折減（重解比 L1_FleetProdRatio）；P＝有效單價 × m。"
        "(1) Q 是模型產能推算，token 帳未閉合：2026 Q2 模型付費 token 為需求面的 L1_TokLedger_2026Q2 倍（約 6.26）；若誤差全在模型產能側，實際營收約為本列 ÷ 6.26。"
        "(2) 回本比較不在本版（成本範圍＝該 GW 全部算力成本，Andy 14:55，另立一版）。"
        "(3) L1_RevGW_Fleet_VR200（65.96）是理想上限、口徑不同，不得與本列相減比較。"
        "(4) 壽命期年均假設 m 不變、L 用 VR200 自身起點折現版（2026 快照時點約等於 VR200 起點 2026.75）。")
SHARED_DRV_M = "當年營收、廣告排除、API／ChatGPT 付費 token、當年有效單價、d（B4_Disc）"
SHARED_WEAK_M = "每日則數、消費訂閱份額 0.65、廣告：Assumed；API token：SRC_DEM_011 時點值（2026-03）"


def l1_rows_new():
    C = "Cap_In"
    def m_row(key, lab, col, extra_h=""):
        return (key, lab, f"Cap_In!${col}$209；{B_M_BASIS}", f"={C}!${col}$209", _t(NR), _t(NR), "x",
                NR + "（2026 欄的輸入為 Assumed 單值，區間待補）" + extra_h, I_M, SHARED_DRV_M, SHARED_WEAK_M, *NONE_EXT,
                "Cap_In 第 209 列；B4_Disc", f"Cap_In 第 209 列（{col} 欄）",
                "2026 欄無區間；Q2 API token 為 2026-03 時點值；m 的未來變化未建模")
    def sub_row(key, lab, rownum, name25, v25):
        h = f"{NR}（2026 欄的輸入為 Assumed 單值，區間待補）；2025 值：{name25} {v25}（v18.0 建置時；現值見該列）"      # static text: a formula here would be an unread (dead) cell
        return (key, lab, f"Cap_In!$E${rownum}（2026 Q2）；分子＝該路線營收（m 分子營收 × 份額），分母＝該路線 token × 當年有效單價（Cap_In 第 205 列）",
                f"={C}!$E${rownum}", _t(NR), _t(NR), "x", h, I_M, SHARED_DRV_M, SHARED_WEAK_M, *NONE_EXT,
                "Cap_In 第 188–193／196–202 列；B4_Disc", f"Cap_In 第 {rownum} 列（E 欄）", "營收份額與消費訂閱 token 份額為 Assumed 單值，區間待補")
    return [
        ("M_2025", "收入面 m — 2025 年度（＝L1_MImplied；各年用各自的牌價：CAP_Price25Mix 1.359）",
         "L1_MImplied 的別名；" + B_M_BASIS, "=L1_MImplied", "=L1_MImplied_Lo", "=L1_MImplied_Hi", "x",
         "同 L1_MImplied 的區間（CTL_Monetize 的 Gov_Map GM594 區間來源）", I_M25, "L1_MImplied 的輸入", "同 L1_MImplied", *NONE_EXT,
         "L1_MImplied；Cap_In 第 94 列", "L1 本頁（L1_MImplied 的別名）", "2025 無實測分母 token"),
        m_row("M_2026Q1", "收入面 m — 2026 Q1（季營收 × 4）", "D"),
        m_row("M_2026Q2", "收入面 m — 2026 Q2（季營收 × 4）", "E"),
        m_row("M_FY2026", "收入面 m — FY2026（內部預測；營收為 Assumed）", "F"),
        ("M", "收入面 m — 主口徑（J1：2026 Q2 年化；替代 FY2026 待 Andy 提供『市場預估 360 億』出處）",
         "＝L1_M_2026Q2；" + B_M_BASIS, "=L1_M_2026Q2", _t(NR), _t(NR), "x",
         NR + "（2026 欄的輸入為 Assumed 單值，區間待補）", I_M, SHARED_DRV_M, SHARED_WEAK_M, *NONE_EXT,
         "Cap_In 第 209 列 E 欄；B4_Disc", "L1 本頁（＝L1_M_2026Q2）", "2026 欄無區間；Q2 API token 為 2026-03 時點值"),
        sub_row("M_API", "收入面 m — API 路線（2026 Q2）", 206, "L1_MAPI", "0.813"),
        sub_row("M_Sub", "收入面 m — 消費訂閱路線（2026 Q2）", 207, "L1_MSub", "6.146"),
        sub_row("M_Seat", "收入面 m — 商用席次路線（2026 Q2）", 208, "L1_MSeat", "2.634"),
        ("PxEff_2026", "2026 快照付費 token 混合有效單價（層級組合 Luna／Sol／Astra × 有效牌價；= 表列牌價 ×（1 − d））",
         "Cap_In!$E$205；2026 快照（B4_MixPaid 加權的 IF_PriceRef）", f"={C}!$E$205", _t(NR), _t(NR), "$/M", NR, "v18.0：L1_M 的分母單價；2026 年用 2026 快照（1.961），2025 年用 CAP_Price25Mix（1.359）",
         "層級組合 B4_MixPaid、IF_PriceRef_*、B4_Disc、AL_Kappa", "牌價：C1 官方牌價；組合：Assumed", *NONE_EXT, "Cap_In 第 205 列、第 90 列", "Cap_In 第 205 列（E 欄）", "顯示用"),
        ("PxList_2026", "2026 快照付費 token 混合表列牌價（＝有效單價 ÷（1 − B4_Disc））",
         "＝L1_PxEff_2026 ÷（1 − B4_Disc）；與 d 無關（d 變動時有效單價同比例變動）", "=L1_PxEff_2026/(1-B4_Disc)", _t(NR), _t(NR), "$/M", NR,
         "v18.0：有效牌價＝表列牌價 ×（1 − d）；本列為表列牌價端，d 改變時不變（驗收 W6 情境 v1800_disc30）", "層級組合、表列牌價、B4_Disc", "牌價：C1 官方牌價；組合：Assumed", *NONE_EXT,
         "L1_PxEff_2026；B4_Disc", "L1 本頁", "顯示用"),
        ("TokGW_PaidCap_VR200", "VR200 每 GW 付費產能（100% 利用率、全部作付費服務；層級組合 B4_MixPaid 調和平均）",
         "1 ÷ Σ（層級占比 ÷ 該層級 IF_TokGW 的 VR200 基準欄）；IF_TokGW 第 11 欄＝VR200 基準（同 L1_FleetProdRatio 取法）",
         "=1/(INDEX(B4_MixPaid,1,1)/INDEX(IF_TokGW_Luna,1,11)+INDEX(B4_MixPaid,1,2)/INDEX(IF_TokGW_Sol,1,11)+INDEX(B4_MixPaid,1,3)/INDEX(IF_TokGW_Astra,1,11))",
         _t(NR), _t(NR), "M tok/GW/年", NR, "v18.0：理想上限（利用率 100%、全部付費）；Q 的基底。約 2.53 × 10¹⁷ token／GW／年",
         "IF_TokGW_Luna／Sol／Astra、B4_MixPaid", "各層級每 GW 產能：roofline 推導（Derived）；組合：Assumed", *NONE_EXT, "Interface 第 21–23 列；B4_MixPaid", "L1 本頁",
         "100% 利用率不是實際；實際見 L1_TokGW_PaidPQ_VR200"),
        ("TokGW_PaidPQ_VR200", "VR200 每 GW 付費 token 實際產出 Q（M tok／年）",
         "＝PaidCap × B4_Serve ×（1 − B4_Free）× AL_UtilActual × L1_FleetProdRatio；Q 側付費切分用支出口徑（chat 端 16:53）",
         "=L1_TokGW_PaidCap_VR200*B4_Serve*(1-B4_Free)*AL_UtilActual*L1_FleetProdRatio",
         "=L1_TokGW_PaidCap_VR200*B4_Serve*(1-B4_Free)*AL_UtilActual_Lo*L1_FleetProdRatio",
         "=L1_TokGW_PaidCap_VR200*B4_Serve*(1-B4_Free)*AL_UtilActual_Hi*L1_FleetProdRatio", "M tok/年",
         "低／高＝實際利用率 AL_UtilActual_Lo／_Hi（0.30／0.65）；其餘取基準", I_PQ + " 低／高只動利用率（J4）。",
         "對外服務占比 B4_Serve、免費占比 B4_Free、實際利用率 AL_UtilActual、生產折減重解比", "利用率 0.41（J4 沿用）、服務占比、免費占比：Assumed", *NONE_EXT,
         "Cap_In 第 28／29 列；Alloc_In 第 24 列；L1_FleetProdRatio", "L1 本頁", "token 帳未閉合（見 L1_TokLedger）"),
        ("RevGW_PQ_VR200", "VR200 每 GW 營收 P × Q — 第一年（主口徑 m＝L1_M，2026 Q2）",
         "＝Q × L1_PxEff_2026 × L1_M ÷ 1E9；P＝有效單價 × m", "=L1_TokGW_PaidPQ_VR200*L1_PxEff_2026*L1_M/1E9",
         "=L1_TokGW_PaidPQ_VR200_Lo*L1_PxEff_2026*L1_M/1E9", "=L1_TokGW_PaidPQ_VR200_Hi*L1_PxEff_2026*L1_M/1E9", "$B/GW/年",
         "低／高＝Q 的低／高（利用率 0.30／0.65）；m、單價取基準", I_PQ, "Q、有效單價、m", "m：Assumed 輸入的衍生；利用率：Assumed", *NONE_EXT,
         "L1_TokGW_PaidPQ_VR200；L1_PxEff_2026；L1_M", "L1 本頁", "token 帳未閉合；m 無區間"),
        ("RevGW_PQ_VR200_FY", "VR200 每 GW 營收 P × Q — 第一年（m 改用 FY2026：L1_M_FY2026）",
         "同 L1_RevGW_PQ_VR200，L1_M 換 L1_M_FY2026", "=L1_TokGW_PaidPQ_VR200*L1_PxEff_2026*L1_M_FY2026/1E9",
         "=L1_TokGW_PaidPQ_VR200_Lo*L1_PxEff_2026*L1_M_FY2026/1E9", "=L1_TokGW_PaidPQ_VR200_Hi*L1_PxEff_2026*L1_M_FY2026/1E9", "$B/GW/年",
         "低／高＝Q 的低／高；m、單價取基準", I_PQ, "Q、有效單價、m（FY2026）", "FY2026 營收為內部預測 Assumed", *NONE_EXT,
         "L1_TokGW_PaidPQ_VR200；L1_PxEff_2026；L1_M_FY2026", "L1 本頁", "FY2026 營收為內部預測"),
        ("RevGW_PQ_VR200_Life", "VR200 每 GW 營收 P × Q — 壽命期年均（第一年 × L1_LTotalDisc；假設 m 不變）",
         "＝L1_RevGW_PQ_VR200 × L1_LTotalDisc；低／高各乘 L1_LTotalDisc_Lo／_Hi", "=L1_RevGW_PQ_VR200*L1_LTotalDisc",
         "=L1_RevGW_PQ_VR200_Lo*L1_LTotalDisc_Lo", "=L1_RevGW_PQ_VR200_Hi*L1_LTotalDisc_Hi", "$B/GW/年",
         "低＝低端 Q × 低端 L；高＝高端 Q × 高端 L（兩端同向組合，範圍很寬）", I_PQ, "第一年營收、L（L1_LTotalDisc）", "L：Assumed／Derived（價格迴歸與降級）；m 不變：Assumed", *NONE_EXT,
         "L1_RevGW_PQ_VR200；L1_LTotalDisc", "L1 本頁", "m 的未來變化未建模"),
        ("TokLedger", "token 帳診斷 — 2025：模型付費 token ÷ 需求面付費 token（＝L1_M_2025 ÷ L1_MFleet）",
         "＝L1_M_2025 ÷ L1_MFleet；同價格基準時成立（2025 兩邊單價同為 CAP_Price25Mix／AL_PriceFleet25＝1.359）", "=IF(AND(ISNUMBER(L1_M_2025),ISNUMBER(L1_MFleet)),L1_M_2025/L1_MFleet,AL_UnavailTxt)", _t(NR_LEDGER), _t(NR_LEDGER), "x",
         NR_LEDGER, "v18.0：模型付費 token（機隊，成本鏈推出）÷ 需求面付費 token。＞1 表示模型產能推算的 token 多於需求面觀測；舊 L1_MFleet＝m ÷ 本列。不是 m。",
         "L1_MImplied、L1_MFleet", "兩者分母皆為模型推算 token（Assumed）", *NONE_EXT, "L1_M_2025；L1_MFleet", "L1 本頁", "兩路線差距未分解"),
        ("TokLedger_2026Q2", "token 帳診斷 — 2026 Q2：模型付費 token ÷ 需求面付費 token（＝Cap_In E209 ÷ E223）",
         "＝Cap_In!$E$209 ÷ Cap_In!$E$223；同價格基準時成立（2026 兩邊單價同為 1.961）", "=IF(AND(ISNUMBER(Cap_In!$E$209),ISNUMBER(Cap_In!$E$223)),Cap_In!$E$209/Cap_In!$E$223,AL_UnavailTxt)", _t(NR_LEDGER), _t(NR_LEDGER), "x",
         NR_LEDGER, "v18.0：同 L1_TokLedger，2026 Q2。L1_RevGW_PQ_VR200 的 Q 為模型產能，若誤差全在模型產能側，實際營收約為該列 ÷ 本列。",
         "Cap_In 第 209／223 列", "兩者分母皆為模型推算 token（Assumed）", *NONE_EXT, "Cap_In 第 209／223 列", "Cap_In 第 209 列、第 223 列（E 欄）", "兩路線差距未分解"),
    ]


KEYS = ["M_2025", "M_2026Q1", "M_2026Q2", "M_FY2026", "M", "M_API", "M_Sub", "M_Seat", "PxEff_2026", "PxList_2026", "TokGW_PaidCap_VR200", "TokGW_PaidPQ_VR200",
        "RevGW_PQ_VR200", "RevGW_PQ_VR200_FY", "RevGW_PQ_VR200_Life", "TokLedger", "TokLedger_2026Q2"]
NEW_NAMES = [f"L1_{k}{s}" for k in KEYS for s in ("", "_Lo", "_Hi")]

v1000._n(v1000.R_, "m＝營收 ÷（需求面付費 token × 有效單價）：由營收與需求面 token 算出（v18.0 起為收入面 m，Andy 2026-10-21 定義）；沿用殘差性質",
         "M_2025", "M_2026Q1", "M_2026Q2", "M_FY2026", "M", "M_API", "M_Sub", "M_Seat")
v1000._n(v1000.E_, "有效單價／表列牌價（層級組合 × 牌價，Derived）", "PxEff_2026", "PxList_2026")
v1000._n(v1000.E_, "VR200 每 GW 付費產能與實際產出 Q、P × Q 營收：物理推導（roofline 產能 × 占比 × 利用率 × 生產折減 × 單價 × m）", "TokGW_PaidCap_VR200", "TokGW_PaidPQ_VR200",
         "RevGW_PQ_VR200", "RevGW_PQ_VR200_FY", "RevGW_PQ_VR200_Life")
v1000._n(v1000.U_, "模型付費 token ÷ 需求面付費 token：兩路線的差距，未分解（token 帳診斷，不是 m）", "TokLedger", "TokLedger_2026Q2")


# ---------------------------------------------------------------- label tails (guarded appends; run after the L1 sheet is built)
TAIL_DIAG = "｜v18.0 起：token 帳診斷（＝m ×（真實 token ÷ 模型 token）），不是 m；m 見 L1_M（Andy 2026-10-21 定義）"
TAIL_M25 = "｜v18.0 起即收入面 m 的 2025 值（L1_M_2025）"
TAIL_MSUB = "｜v18.0 起即收入面分項 m 的 2025 值"
TAIL_209 = "（v18.0：本列即收入面 m；L1_M_* 讀本列）"
TAIL_H196 = "；v18.0：Q2 當期值已搜尋，未找到（DB_Evidence 新列 E399）"


def _l1_row(wb, name):
    ws = wb["L1"]
    rows = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == name]
    assert len(rows) == 1, (name, rows)
    return rows[0]


def label_tails(wb, append_text):
    n = 0
    ws = wb["L1"]
    for name, tail in (("L1_MFleet", TAIL_DIAG), ("L1_MFleet_Px2026", TAIL_DIAG), ("L1_MImplied", TAIL_M25),
                       ("L1_MAPI", TAIL_MSUB), ("L1_MSub", TAIL_MSUB), ("L1_MSeat", TAIL_MSUB)):
        n += int(bool(append_text(ws[f"B{_l1_row(wb, name)}"], tail)))
    ci = wb["Cap_In"]
    for addr, start, tail in (("A223", "m_fleet＝m 分子營收", TAIL_DIAG), ("A119", "機隊 m（點估計）", TAIL_DIAG),
                              ("A209", "m_整體＝m 分子營收", TAIL_209), ("H196", "SRC_DEM_011 15 B tok/min", TAIL_H196)):
        assert str(ci[addr].value).startswith(start), (addr, ci[addr].value)
        n += int(bool(append_text(ci[addr], tail)))
    return n


def clear_l1_ratio(wb):
    """L1 column O (推算 ÷ 外部) has no external value on the new rows: cleared like the other external-less rows (v10.1 L1_O_ROWS), so they are not dead formula cells."""
    ws = wb["L1"]; n = 0
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value in {f"L1_{k}" for k in KEYS} and ws[f"O{r}"].value is not None:
            ws[f"O{r}"].value = None; n += 1
    return n


# ---------------------------------------------------------------- DB_Evidence E399
EVIDENCE = [("E399", "OpenAI API 每分鐘 token 2026 Q2／Q3 當期值",
             "搜尋 2026-10-21 17:20（chat 端）：OpenAI 官方 2026-03-31〈Accelerating the next phase of AI〉（最新官方值仍為 15 B／分鐘）、DevDay 2026、第三方統計頁",
             "Assumed（Q2 沿用 2026-03 時點值）", "Cap_In 第 196 列（API token：2026 Q1／Q2＝SRC_DEM_011 時點值年化）", "15 B tok／分鐘（SRC_DEM_011，2026-03）",
             "未找到（Q2／Q3 當期值）", "已搜尋，未找到",
             "Q2 沿用 2026-03 時點值；Q2 的 API token 可能偏低、L1_M_2026Q2 等 m 可能偏高。不改任何值", "L1_M_2026Q1／Q2／FY2026、L1_M、L1_M_API 及其依賴列；不改任何既有 IF_／L1_")]


def evidence_append(wb):
    ws = wb["DB_Evidence"]
    have = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
    r = max(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")) + 1
    n = 0
    for (eid, claim, src, mark, param, cur, new, verdict, note, affected) in EVIDENCE:
        if eid in have:
            continue
        vals = [eid, DATE, claim, src, mark, param, cur, new, verdict, V, note + f"｜{WO} W4", "已處理", "SRC_DEM_011", "—", affected, "—", "—",
                f"{verdict}｜v18.0 新增列（非 106 列積壓）"]
        for i, v in enumerate(vals):
            put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 4, 5, 6, 7, 8, 10, 14, 17))
        r += 1; n += 1
    return n


# ---------------------------------------------------------------- Decisions SL52-*
def _decisions():
    D = [
        ("m 的定義與命名：m 只有收入面，新名稱 L1_M／L1_M_*",
         "m＝付費服務 token 的實收單價 ÷ 有效牌價；有效牌價＝表列牌價 ×（1 − d），d＝B4_Disc（預設 0.20，可調）；m 只有收入面。新增 L1_M_2025／_2026Q1／_2026Q2／_FY2026、L1_M（主口徑）、L1_M_API／_Sub／_Seat、L1_PxEff_2026、L1_PxList_2026。"
         "舊 L1_MFleet 等以成本鏈推 token 再反推，＝ m ×（真實 token ÷ 模型 token），改標為『token 帳診斷』，數值不變；新增 L1_TokLedger／L1_TokLedger_2026Q2。",
         "L1 第 117 列起；L1 第 76／78／80–82／102 列標籤；Cap_In A119／A209／A223"),
        ("J1：L1_M 主口徑用 2026 Q2 年化",
         "L1_M＝L1_M_2026Q2（1.343）。替代：FY2026 1.432（待 Andy 提供『市場預估 360 億』出處）。可撤回。", "L1_M"),
        ("J2：ChatGPT token 的付費／免費以付費占比切分",
         "ChatGPT token 的付費／免費仍以付費占比（token 口徑 1 − AL_FreeShare）切分：m 的分母是觀測的付費 token。可撤回。", "Cap_In 第 198／199 列；L1_M_*"),
        ("J3：廣告營收排除",
         "廣告營收排除於 m 的分子（沿用 Cap_In 第 187 列『m 分子營收＝營收 − 廣告』）。可撤回。", "Cap_In 第 187 列；L1_M_*"),
        ("J4：VR200 新 GW 實際利用率沿用 AL_UtilActual 0.41（0.30–0.65）",
         "L1_TokGW_PaidPQ_VR200 的實際利用率取 AL_UtilActual（低／高 AL_UtilActual_Lo／_Hi），不用 IF_Util。可撤回。", "L1_TokGW_PaidPQ_VR200"),
        ("追加：Q 側付費切分用支出口徑；Q2 API token 當期值未找到",
         "Q 側的付費切分統一用支出口徑 B4_Serve ×（1 − B4_Free）＝0.2214（chat 端 16:53）；生產折減用重解比 L1_FleetProdRatio（0.806）。"
         "OpenAI API 每分鐘 token 的 2026 Q2 當期值已搜尋，未找到（DB_Evidence E399），Q2 沿用 SRC_DEM_011 的 2026-03 時點值。可撤回。",
         "L1_TokGW_PaidPQ_VR200；L1_RevGW_PQ_VR200*；Cap_In 第 196 列；DB_Evidence E399"),
    ]
    return [[f"SL52-{i + 1:02d}", V, f"v18.0 m 收入面化：{t}", txt, DATE, DEC_BY, "v18.0 已建", where, f"{WO} W5", "否"] for i, (t, txt, where) in enumerate(D)]


def apply(wb, append_text):
    log = {}
    log["l1_o_cleared"] = clear_l1_ratio(wb)
    log["label_tails"] = label_tails(wb, append_text)
    log["evidence"] = evidence_append(wb)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == "SL52-01" for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


# ================================================================== README
README_VERSION = (VERSION + "（主版：" + WO + "——m 收入面化：新增 L1_M／L1_M_2025／_2026Q1／_2026Q2／_FY2026、L1_M_API／_Sub／_Seat、L1_PxEff_2026／L1_PxList_2026；"
                  "VR200 每 GW 營收以 P × Q 由基本定義計算（L1_TokGW_PaidCap_VR200、L1_TokGW_PaidPQ_VR200、L1_RevGW_PQ_VR200／_FY／_Life）；"
                  "舊 L1_MFleet 等改標 token 帳診斷（新增 L1_TokLedger／L1_TokLedger_2026Q2），數值不變；DB_Evidence E399；Decisions SL52-01…06。既有名稱的值全部不變。）以下為 "
                  + v1720.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1720.README_TITLE
README_TITLE = V + _T[len(v1720.V):]
README_ROW = ("v18.0 m 收入面化",
              "m＝付費服務 token 的實收單價 ÷ 有效牌價（Andy 2026-10-21）；只有收入面。新名稱 L1_M（主口徑：2026 Q2 年化）、分項 L1_M_API／_Sub／_Seat；"
              "VR200 每 GW 營收 L1_RevGW_PQ_VR200 以 P × Q 由基本定義計算（Q＝付費 token 實際產出，P＝有效單價 × m）；舊 L1_MFleet 為 token 帳診斷（L1_TokLedger）。"
              "細節見工作單 docs/workorders/20261021_v18.0_m收入面化.md 與 Decisions SL52-*。")
