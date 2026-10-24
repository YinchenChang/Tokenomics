"""m 閱讀版的核對（v21.2 重做）：確認閱讀版取自 Tokenomics 的數字與指定版本一致。

用法：
    python3 check_reading_m.py <Tokenomics.xlsx> <m閱讀版.xlsx>
m 閱讀版會先以 LibreOffice 重算（暫存，不改原檔）。全部通過時結束碼 0，否則 1。
檢查：(a) 有 1A_需求公式樹；(b) 9_核對對照 登錄的每格（輸入、區間、原值）等於 Tokenomics 現值（相對容差 1e-9），
差異格 ≤ 1e-9；(c) 本簿重算值與 L1_M_2025（＋區間）、_2026Q1、_2026Q2（＋區間）、L1_M、_FY2026、
_API／_Sub／_Seat、L1_PxEff_2026、L1_PxList_2026、CAP_Price25Mix 一致；(d) 全簿無錯誤值。
"""
import math, re, sys, tempfile
from pathlib import Path
import openpyxl
sys.path.insert(0, str(Path(__file__).parent))
from parity_lib import lo_recalc

if len(sys.argv) != 3:
    sys.exit('用法：python3 check_reading_m.py <Tokenomics.xlsx> <m閱讀版.xlsx>')
TK, MW = sys.argv[1:3]
TV = openpyxl.load_workbook(TK, data_only=True)
TF = openpyxl.load_workbook(TK)
tmp = Path(tempfile.mkdtemp(prefix='chk_m_'))
src = tmp / 'in'; src.mkdir()
import shutil
shutil.copy(MW, src / Path(MW).name)
RC = lo_recalc(src / Path(MW).name, tmp / 'out')
W = openpyxl.load_workbook(RC, data_only=True)
fails = []
REF = re.compile(r"([A-Za-z_][A-Za-z0-9_]*)!\$?([A-Z]+)\$?(\d+)")


def same(a, b):
    try:
        return math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=1e-12)
    except (TypeError, ValueError):
        return a == b


def tk(expr):
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", expr):
        sh, rng = list(TF.defined_names[expr].destinations)[0]
        return TV[sh][rng.replace('$', '')].value
    m = REF.fullmatch(expr)
    if m:
        return TV[m[1]][m[2] + m[3]].value
    s = REF.sub(lambda m: repr(float(TV[m[1]][m[2] + m[3]].value)), expr)
    s = re.sub(r"(?<![0-9.A-Za-z_])[A-Za-z_][A-Za-z0-9_]*", lambda m: repr(float(tk(m[0]))), s)
    assert re.fullmatch(r"[0-9eE+\-*/(). ]+", s), expr
    return eval(s)


def named(book, n, tkbook=False):
    d = (TF if tkbook else openpyxl.load_workbook(RC)).defined_names[n]
    sh, rng = list(d.destinations)[0]
    return (TV if tkbook else book)[sh][rng.replace('$', '')].value


# (a) 規格
if '1A_需求公式樹' not in W.sheetnames:
    fails.append('規格：缺 1A_需求公式樹')
# (b) 對照
ws = W['9_核對對照']
n = 0
for r in ws.iter_rows(min_row=5, values_only=True):
    sh, ad, ex, kd = r[:4]
    if not sh:
        continue
    n += 1
    v = W[sh][ad].value
    if kd == '差異':
        if not (isinstance(v, (int, float)) and abs(v) <= 1e-9):
            fails.append(f'差異 {sh}!{ad} = {v}')
    else:
        t = tk(ex)
        if not same(v, t):
            fails.append(f'{kd} {sh}!{ad} 本簿 {v} ≠ Tokenomics {ex} {t}')
# (c) 結果名稱
PAIR = {'OUT_M_2025': 'L1_M_2025', 'OUT_M_2025_Lo': 'L1_M_2025_Lo', 'OUT_M_2025_Hi': 'L1_M_2025_Hi',
        'OUT_M_2026Q1': 'L1_M_2026Q1', 'OUT_M_2026Q2': 'L1_M_2026Q2', 'OUT_M_2026Q2_Lo': 'L1_M_2026Q2_Lo',
        'OUT_M_2026Q2_Hi': 'L1_M_2026Q2_Hi', 'OUT_M': 'L1_M', 'OUT_M_FY2026': 'L1_M_FY2026',
        'OUT_M_API': 'L1_M_API', 'OUT_M_Sub': 'L1_M_Sub', 'OUT_M_Seat': 'L1_M_Seat',
        'OUT_M_API_2025': 'L1_MAPI', 'OUT_M_Sub_2025': 'L1_MSub', 'OUT_M_Seat_2025': 'L1_MSeat',
        'OUT_PxEff_2026': 'L1_PxEff_2026', 'OUT_PxList_2026': 'L1_PxList_2026', 'OUT_Price25Mix': 'CAP_Price25Mix'}
RF = openpyxl.load_workbook(RC)
for o, t in PAIR.items():
    sh, rng = list(RF.defined_names[o].destinations)[0]
    v = W[sh][rng.replace('$', '')].value
    tv = tk(t)
    if not same(v, tv):
        fails.append(f'結果 {o}={v} ≠ {t}={tv}')

# (e) 本簿內部一致性（r1）
def outv(o):
    sh, rng = list(RF.defined_names[o].destinations)[0]
    return W[sh][rng.replace('$', '')].value


def rowby(ws, col, pred):
    for r in ws.iter_rows():
        v = r[col].value
        if isinstance(v, str) and pred(v):
            return r
    return None


m25, m26 = outv('OUT_M_2025'), outv('OUT_M_2026Q2')
ov = W['1_總覽']
r = rowby(ov, 1, lambda v: v.startswith('m 倍數＝'))
if r is None or not same(r[4].value, m26 / m25) or abs(r[4].value - m26 / m25) > 1e-9:
    fails.append('三因子乘積 ≠ m 的 Q2 ÷ 2025')
for lab in ('3_m重算 的 m ÷', '4_分路線 三路線合計'):
    r = rowby(ov, 1, lambda v: v.startswith(lab))
    if r is None or not (isinstance(r[2].value, (int, float)) and abs(r[2].value) <= 1e-9):
        fails.append(f'1_總覽 檢核格 {lab} = {None if r is None else r[2].value}')
s3 = W['3_m重算']
rm, rc = rowby(s3, 0, lambda v: v == 'S13'), rowby(s3, 0, lambda v: v == 'S14')
for i in range(5, 9):
    if abs(rm[i].value - rc[i].value) > 1e-9:
        fails.append('3_m重算 S14 ≠ S13')
r18 = rowby(W['4_分路線'], 0, lambda v: v == 'R18')
for i in range(5, 9):
    if abs(r18[i].value) > 1e-9:
        fails.append('4_分路線 R18 ≠ 0')
rt = rowby(W['1A_需求公式樹'], 1, lambda v: v.startswith('檢查：分子'))
for i, mv in ((2, m25), (3, m26)):
    if abs(rt[i].value - mv) > 1e-9:
        fails.append('1A 分子÷分母 ≠ m')
sv = W['5_敏感度']
if abs(sv['G6'].value) > 1e-9:
    fails.append('5_敏感度 基準重算檢查 ≠ 0')
rows = [(r[11].value, r[1].value, r[10].value) for r in sv.iter_rows(min_row=7) if isinstance(r[11].value, (int, float)) and isinstance(r[0].value, int)]
top = sorted(rows, reverse=True)[:8]
hd = rowby(ov, 1, lambda v: v.startswith('最不確定的輸入'))
for k, (sw, nm, el) in enumerate(top, 1):
    o = ov.cell(hd[1].row + 1 + k, 2).value, ov.cell(hd[1].row + 1 + k, 3).value
    if o[0] != nm or abs(o[1] - sw) > 1e-9:
        fails.append(f'1_總覽 前 8 名第 {k} 項 {o} ≠ 5_敏感度 {nm} {sw}')

# (d) 錯誤值
for wsx in W.worksheets:
    for row in wsx.iter_rows():
        for c in row:
            if isinstance(c.value, str) and re.match(r'^(#(REF|VALUE|DIV/0|NAME|N/A|NUM)|Err:)', c.value):
                fails.append(f'錯誤值 {wsx.title}!{c.coordinate} {c.value}')
print(f'核對對照 {n} 列；結果名稱 {len(PAIR)} 個')
print('通過' if not fails else f'未通過 {len(fails)} 項')
for x in fails[:50]:
    print(' -', x)
sys.exit(1 if fails else 0)
