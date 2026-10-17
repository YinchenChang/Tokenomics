"""v14.1：新舊兩版（皆為 LibreOffice 重算存檔）逐格差異，依工作單 R1–R6 的預期變動分類；範圍外差異應為 0。
用法：python3 -I 03_classify.py <舊版.xlsx> <新版.xlsx>    （從 repo 根目錄執行）"""
import sys, json, collections
import openpyxl
from openpyxl.utils import get_column_letter as GL

old_p, new_p = sys.argv[1], sys.argv[2]
ov, nv = openpyxl.load_workbook(old_p, data_only=True), openpyxl.load_workbook(new_p, data_only=True)
of, nf = openpyxl.load_workbook(old_p), openpyxl.load_workbook(new_p)


def same(a, b):
    a = None if a == "" else a
    b = None if b == "" else b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= 1e-12 or abs(a - b) <= 1e-9 * max(abs(a), abs(b))
    return a == b


data = json.load(open("builder/v1400_data.json", encoding="utf-8"))["delete"]
del_rows = {(e["sheet"], e["row"]) for e in data}
R5 = {f"C{r}" for r in [185] + list(range(187, 194)) + list(range(196, 203)) + [205] + list(range(216, 223))} | {"G175"}
R5_ROWS = {int(x[1:]) for x in R5}


def pcol(ws):
    for c in range(3, ws.max_column + 2):
        if ws.cell(3, c).value == "用途":
            return c
    return None


gm = nv["Gov_Map"]
gm_rows = {(gm.cell(r, 3).value, gm.cell(r, 4).value): r for r in range(5, gm.max_row + 1)}
GM_R3 = {gm_rows[("Har_In", "C20:G20")], gm_rows[("Har_In", "C32:G32")]}
GM_R6 = {gm_rows[("Cap_In", "C204:F204")]}
ck = nv["Checks"]
V_ROWS = {r for r in range(1, ck.max_row + 1) if ck.cell(r, 1).value in ("V23", "V24")}
cat = collections.Counter()
other = []
for ws in nv.worksheets:
    if ws.title not in ov.sheetnames:
        other.append((ws.title, "new sheet")); continue
    s = ws.title
    pc = pcol(ws)
    mr, mc = max(ov[s].max_row, ws.max_row), max(ov[s].max_column, ws.max_column)
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            va, vb = ov[s].cell(r, c).value, ws.cell(r, c).value
            fa, fb = of[s].cell(r, c).value, nf[s].cell(r, c).value
            if same(va, vb) and (same(fa, fb) or not (isinstance(fa, str) and isinstance(fb, str) and (fa.startswith("=") or fb.startswith("=")))):
                continue
            coord = ws.cell(r, c).coordinate
            L_ = GL(c)
            t = None
            if (s, r) in del_rows and c in (1, 2): t = "R4 標籤清空"
            elif s == "Cap_In" and coord in R5: t = "R5 清空"
            elif s == "Cap_In" and pc and c == pc and (r in R5_ROWS or r in (121, 76)): t = "R5/R1/R6 用途標記"
            elif s == "Cap_In" and r == 119 and L_ in "CDEG": t = "R1 公式/值"
            elif s == "Cap_In" and r == 121 and L_ == "G": t = "R1 說明"
            elif s == "L1" and coord in ("E78", "F89", "E91"): t = "R1 依賴鏈值"
            elif s == "Cap_In" and coord in ("E240", "E242"): t = "R1 依賴鏈值"
            elif s == "Har_In" and r in (20, 32) and (3 <= c <= 7 or c == pc): t = "R3 清空/用途"
            elif s == "Gov_Map" and L_ == "N" and r in GM_R3: t = "R3 Gov_Map 註記"
            elif s == "Gov_Map" and L_ == "N" and r in GM_R6: t = "R6 Gov_Map 註記"
            elif s == "Cap_In" and L_ == "H" and r in (206, 207, 208, 209, 223): t = "R6 錨點文字"
            elif (s == "Training" and r == 122 and c == pc) or (s == "Price_Frontier" and r == 27 and c == pc): t = "R6 用途更正"
            elif s == "Checks" and ((r in V_ROWS and L_ in "BFG") or coord in ("F144", "F145", "A143")): t = "R2 恆等式/計數"
            elif s == "Decisions": t = "Decisions 新增/註記"
            elif s == "README": t = "README 版本"
            elif s == "Alloc" and coord in ("C177", "C178"): t = "區間版本字串"
            elif s == "Interface" and L_ in ("R", "V", "W"): t = "Interface 契約欄（先行追蹤；R1 依賴鏈）"
            if t:
                cat[t] += 1
            else:
                other.append((s, coord, str(fa)[:60], str(fb)[:60], str(va)[:40], str(vb)[:40]))
print("== 分類（格數）")
for t, n in cat.most_common():
    print(" ", t, n)
print("== 範圍外差異", len(other))
for x in other[:60]:
    print("  ", x)
