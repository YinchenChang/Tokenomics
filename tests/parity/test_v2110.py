"""v21.1 專屬測試（工作單 docs/workorders/20261023_v21.1_NonNV與免費占比路徑.md r0）。讀 LibreOffice 重算後的值。

  1. 新名稱存在且形狀正確（IF_ 各 1×3；L1 各 D:F 同一列）。
  2. 值等於來源格（≤1e-12）並等於事前值。
  3. L1 情境列 ÷ 主列 ＝（1 − 免費占比）÷（1 − B4_Free）（2026、2030）。
  4. Cap_In 第 347／348 列 D／E 方向（F 高 → 倍數低）。
  5. B4_Free 仍為 0.46；不回流（新名稱只被預期的頁引用）。
  6. NonNV H10／H11 ＝ 6.875／0.9091。
  7. 非數字防護：SRC_DEM_089／090 設為文字時第 147／148 列回傳 AL_UnavailTxt 而非 #VALUE!；f_registry／cache_off 情境無錯誤值。
  8. 既有名稱數值不變（對封存的 v21.0）；名稱只增加；逐格差異只落在預期範圍。
  9. 記錄與文字：Decisions SL58-01、README、NonNV 文字、Interface 標題與 A364、Alloc_In G24、SRC_DEM_006、SRC_DEM_089／090、DB_Evidence E402。
"""
import re
from pathlib import Path

import openpyxl
import pytest
import yaml
from openpyxl.utils import range_boundaries

from conftest import model_path
from test_v2120 import CHANGED as V2120_CHANGED, REVERT as SB_OLD      # v21.2：AL_SpendBasis 改 1.8 的預期變動名稱（44 個）與舊值輸入

REPO = Path(__file__).resolve().parents[2]
V210 = REPO / "model" / "archive" / "20261023_Tokenomics_v21.0.xlsx"
SC = {s["id"]: s["inputs"] for s in yaml.safe_load((Path(__file__).resolve().parent / "scenarios.yaml").read_text(encoding="utf-8"))["scenarios"]}
F_REGISTRY = SC["f_registry_t07_t09_on"]
CACHE_OFF = SC["d9_cache_cap_off"]
TXT26 = {"SRC_DEM_089": "不適用"}
TXT30 = {"SRC_DEM_090": "不適用"}
PREWARM = {"test_v2110_text_guard": [TXT26, TXT30], "test_v2110_stress_scenarios_no_error_values": [F_REGISTRY, CACHE_OFF]}

IF_PRE = {"IF_NonNV_Out_Ascend": (10, "BCD", (0.05, 0.08, 0.15)), "IF_NonNV_Out_Maia": (11, "BCD", (0.25, 0.55, 0.75)),
          "IF_NonNV_Hold_Ascend": (10, "EFG", (0.45, 0.55, 0.70)), "IF_NonNV_Hold_Maia": (11, "EFG", (0.40, 0.50, 0.60))}
L1_PRE = {"L1_FreeShare_2026": (0.34, 0.25, 0.46), "L1_FreeShare_2030": (0.06, 0.03, 0.30),
          "L1_RevGW_PQ_VR200_F2026": (60.51684715552, 44.28, 95.94), "L1_RevGW_PQ_VR200_F2030": (86.19066110029, 63.07, 136.64)}
NEW_L1 = {f"{k}{s}" for k in L1_PRE for s in ("", "_Lo", "_Hi")}
NEW_NAMES = set(IF_PRE) | NEW_L1 | {"SRC_DEM_089", "SRC_DEM_090"}
SKIP = ("RNG_", "IDX_", "GOV_", "IFC_")      # IDX_：SRC_Index 範圍隨新紀錄延伸；IFC_：契約欄範圍延伸；GOV_：治理計數另行檢查
VERSION_NAMES = {"RNG_Version", "RNG_Current"}


def _flat(v):
    if isinstance(v, (list, tuple)):
        return [x for r in v for x in _flat(r)]
    return [v]


def _num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def _eq(a, b):
    if a in (None, "") and b in (None, ""):
        return True
    if _num(a) and _num(b):
        return abs(a - b) <= 1e-9 * max(abs(a), abs(b), 1e-300) + 1e-12
    return a == b


# ---------------------------------------------------------------- 1, 2. 形狀與值
@pytest.mark.parametrize("nm", sorted(IF_PRE))
def test_v2110_if_names(lo, nm):
    v = lo.view({})
    row, cols, pre = IF_PRE[nm]
    x = _flat(v.get_name(nm))
    assert len(x) == 3
    for j in range(3):
        assert x[j] == v.get("NonNV", f"{cols[j]}{row}")
        assert abs(x[j] - pre[j]) <= 1e-12, (nm, j)
    assert x[0] <= x[1] <= x[2]


@pytest.mark.parametrize("nm", sorted(L1_PRE))
def test_v2110_l1_names(lo, nm):
    v = lo.view({})
    base, lo_, hi = _flat(v.get_name(nm))[0], _flat(v.get_name(nm + "_Lo"))[0], _flat(v.get_name(nm + "_Hi"))[0]
    pre = L1_PRE[nm]
    assert abs(base - pre[0]) <= 1e-9 and abs(lo_ - pre[1]) <= 0.01 and abs(hi - pre[2]) <= 0.01, (nm, base, lo_, hi)
    assert lo_ <= base <= hi
    # 同一列的 D／E／F
    sh, ref = v.name_ref(nm)
    sh2, ref2 = v.name_ref(nm + "_Lo")
    sh3, ref3 = v.name_ref(nm + "_Hi")
    assert sh == sh2 == sh3 == "L1" and ref[0] == "D" and ref2[0] == "E" and ref3[0] == "F" and ref[1:] == ref2[1:] == ref3[1:]


def test_v2110_source_cells(lo):
    v = lo.view({})
    cap = lambda c, r: v.get("Cap_In", f"{c}{r}")
    assert v.get_name("L1_FreeShare_2026") == cap("C", 345) == v.get_name("SRC_DEM_089") == 0.34
    assert (v.get_name("L1_FreeShare_2026_Lo"), v.get_name("L1_FreeShare_2026_Hi")) == (cap("D", 345), cap("E", 345)) == (0.25, 0.46)
    assert v.get_name("L1_FreeShare_2030") == cap("C", 346) == v.get_name("SRC_DEM_090") == 0.06
    assert (v.get_name("L1_FreeShare_2030_Lo"), v.get_name("L1_FreeShare_2030_Hi")) == (cap("D", 346), cap("E", 346)) == (0.03, 0.30)
    for nm, r in (("L1_RevGW_PQ_VR200_F2026", 347), ("L1_RevGW_PQ_VR200_F2030", 348)):
        for suf, src in (("", ""), ("_Lo", "_Lo"), ("_Hi", "_Hi")):
            assert abs(v.get_name(nm + suf) - v.get_name("L1_RevGW_PQ_VR200" + src) * cap("C", r)) <= 1e-12 * max(1, abs(v.get_name(nm + suf)))


# ---------------------------------------------------------------- 3, 4, 5, 6
def test_v2110_ratio_identity(lo):
    v = lo.view({})
    free = v.get_name("B4_Free")
    for yr in (2026, 2030):
        want = (1 - v.get_name(f"L1_FreeShare_{yr}")) / (1 - free)
        got = v.get_name(f"L1_RevGW_PQ_VR200_F{yr}") / v.get_name("L1_RevGW_PQ_VR200")
        assert abs(got - want) <= 1e-12, (yr, got, want)
        assert abs(v.get_name(f"L1_RevGW_PQ_VR200_F{yr}_Lo") / v.get_name("L1_RevGW_PQ_VR200_Lo") - want) <= 1e-12
        assert abs(v.get_name(f"L1_RevGW_PQ_VR200_F{yr}_Hi") / v.get_name("L1_RevGW_PQ_VR200_Hi") - want) <= 1e-12


def test_v2110_multiplier_direction(lo):
    v = lo.view({})
    g = lambda c, r: v.get("Cap_In", f"{c}{r}")
    free = v.get_name("B4_Free")
    for r, fr in ((347, 345), (348, 346)):
        c, d, e = g("C", r), g("D", r), g("E", r)
        assert d <= c <= e                                # F 高 → 倍數低：D（低）用 F 高端、E（高）用 F 低端
        assert abs(d - (1 - g("E", fr)) / (1 - free)) <= 1e-12 and abs(e - (1 - g("D", fr)) / (1 - free)) <= 1e-12
        assert abs(c - (1 - g("C", fr)) / (1 - free)) <= 1e-12
    assert abs(g("C", 347) - 1.2222222222) <= 1e-9 and abs(g("C", 348) - 1.7407407407) <= 1e-9


def test_v2110_b4free_unchanged_and_nonnv_ratios(lo):
    v = lo.view({})
    assert v.get_name("B4_Free") == 0.46 and v.get("Cap_In", "C29") == 0.46
    assert abs(v.get("NonNV", "H10") - 6.875) <= 1e-12 and abs(v.get("NonNV", "H11") - 0.9090909090909) <= 1e-9
    assert abs(v.get("NonNV", "I10") - 0.45 / 0.15) <= 1e-12 and abs(v.get("NonNV", "J10") - 0.70 / 0.05) <= 1e-12


# ---------------------------------------------------------------- 7. 非數字防護
def test_v2110_text_guard(lo):
    b = lo.view({})
    for inp, nm, other in ((TXT26, "L1_RevGW_PQ_VR200_F2026", "L1_RevGW_PQ_VR200_F2030"), (TXT30, "L1_RevGW_PQ_VR200_F2030", "L1_RevGW_PQ_VR200_F2026")):
        v = lo.view(inp)
        unavail = v.get_name("AL_UnavailTxt")
        for suf in ("", "_Lo", "_Hi"):
            assert v.get_name(nm + suf) == unavail, (nm, suf, v.get_name(nm + suf))
            assert v.get_name(other + suf) == b.get_name(other + suf)      # 另一年不受影響
        errs = [(s, c) for (s, c), x in v.evaluate_all().items() if isinstance(x, str) and x.startswith("#")]
        assert errs == [], errs[:5]
        # 既有名稱不受影響（不回流）
        assert v.get_name("L1_RevGW_PQ_VR200") == b.get_name("L1_RevGW_PQ_VR200") and v.get_name("B4_Free") == 0.46


def test_v2110_stress_scenarios_no_error_values(lo):
    for label, inp in (("f_registry", F_REGISTRY), ("cache_off", CACHE_OFF)):
        v = lo.view(inp)
        errs = [(s, c) for (s, c), x in v.evaluate_all().items() if isinstance(x, str) and x.startswith("#")]
        assert errs == [], (label, errs[:5])
        for nm in L1_PRE:
            for suf in ("", "_Lo", "_Hi"):
                x = v.get_name(nm + suf)
                assert not (isinstance(x, str) and x.startswith("#")), (label, nm, suf)


def test_v2110_inputs_propagate(lo):
    chg = {"NonNV!C10": 0.1, "NonNV!F11": 0.6, "Cap_In!D345": 0.2, "Cap_In!E346": 0.4}
    b, h = lo.view({}), lo.view(chg)
    assert _flat(h.get_name("IF_NonNV_Out_Ascend"))[1] == pytest.approx(0.1) and _flat(h.get_name("IF_NonNV_Hold_Maia"))[1] == pytest.approx(0.6)
    assert h.get_name("L1_FreeShare_2026_Lo") == pytest.approx(0.2) and h.get_name("L1_FreeShare_2030_Hi") == pytest.approx(0.4)
    assert h.get_name("L1_RevGW_PQ_VR200_F2026") == b.get_name("L1_RevGW_PQ_VR200_F2026")      # 區間端點不進情境列的基準
    assert abs(h.get("Cap_In", "E347") - (1 - 0.2) / (1 - 0.46)) <= 1e-12 and abs(h.get("Cap_In", "D348") - (1 - 0.4) / (1 - 0.46)) <= 1e-12
    errs = [(s, c) for (s, c), x in h.evaluate_all().items() if isinstance(x, str) and x.startswith("#")]
    assert errs == []


PREWARM["test_v2110_inputs_propagate"] = [{"NonNV!C10": 0.1, "NonNV!F11": 0.6, "Cap_In!D345": 0.2, "Cap_In!E346": 0.4}]


# ---------------------------------------------------------------- 8. 靜態：名稱、不回流、既有名稱不變、逐格差異
@pytest.fixture(scope="module")
def wbf():
    return openpyxl.load_workbook(model_path())


def test_v2110_no_reference_outside(wbf):
    pat_if = re.compile(r"\b(" + "|".join(sorted(IF_PRE, key=len, reverse=True)) + r")\b")
    pat_l1 = re.compile(r"\b(" + "|".join(sorted(["L1_FreeShare_2026", "L1_FreeShare_2030", "L1_RevGW_PQ_VR200_F2026", "L1_RevGW_PQ_VR200_F2030"], key=len, reverse=True)) + r")(_Lo|_Hi)?\b")
    bad = []
    for ws in wbf:
        if ws.title in {"Interface", "Load_Bearing", "Gov_Map"}:
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and (pat_if.search(c.value) or pat_l1.search(c.value)):
                    bad.append((ws.title, c.coordinate))
    assert not bad, bad[:5]
    # B4_Free 的公式與值不被改（C29 為藍字常數）
    assert wbf["Cap_In"]["C29"].value == 0.46


def test_v2110_existing_names_unchanged(lo):
    pw = openpyxl.load_workbook(V210, data_only=True)
    v = lo.view({})
    bad, n = [], 0
    for name, d in pw.defined_names.items():
        if name.startswith(SKIP) or name in VERSION_NAMES or name in V2120_CHANGED or "!" not in d.attr_text or "," in d.attr_text:
            continue
        sh, rg = d.attr_text.rsplit("!", 1)
        c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
        want = [pw[sh.strip("'")].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        got = _flat(v.get_name(name))
        n += 1
        if not (len(want) == len(got) and all(_eq(a, b) for a, b in zip(want, got))):
            bad.append(name)
    assert n > 1500
    assert bad == [], bad[:20]


def test_v2110_names_added_only(wbf):
    a = openpyxl.load_workbook(V210)
    na, nb = set(a.defined_names), set(wbf.defined_names)
    assert na <= nb and nb - na == NEW_NAMES, sorted((nb - na) ^ NEW_NAMES)
    changed = {k for k in na if a.defined_names[k].attr_text != wbf.defined_names[k].attr_text}
    assert all(k.startswith(("IDX_", "RNG_", "GOV_", "IFC_")) for k in changed), sorted(changed)[:10]
    for nm in IF_PRE:
        ref = wbf.defined_names[nm].attr_text
        c1, r1, c2, r2 = range_boundaries(ref.split("!")[1].replace("$", ""))
        assert ref.startswith("Interface!") and r1 == r2 and c1 == 3 and c2 == 5 and 379 <= r1 <= 382, (nm, ref)
    assert [int(wbf.defined_names[n].attr_text.split("$")[-1]) for n in ("L1_FreeShare_2026", "L1_FreeShare_2030", "L1_RevGW_PQ_VR200_F2026", "L1_RevGW_PQ_VR200_F2030")] == [145, 146, 147, 148]
    assert [int(wbf.defined_names[n].attr_text.split("$")[-1]) for n in ("IF_NonNV_Out_Ascend", "IF_NonNV_Out_Maia", "IF_NonNV_Hold_Ascend", "IF_NonNV_Hold_Maia")] == [379, 380, 381, 382]


def test_v2110_diffs_only_expected_vs_v210():
    a, b = openpyxl.load_workbook(V210), openpyxl.load_workbook(model_path())
    assert a.sheetnames == b.sheetnames
    d = {}
    for ws in a:
        wn = b[ws.title]
        for r in range(1, max(ws.max_row, wn.max_row) + 1):
            for c in range(1, max(ws.max_column, wn.max_column) + 1):
                x, y = ws.cell(r, c).value, wn.cell(r, c).value
                if x != y:
                    d.setdefault(ws.title, []).append((ws.cell(r, c).coordinate, x, y))
    allowed = {"Alloc", "Alloc_In", "Cap_In", "Checks", "DB_Evidence", "Decisions", "Gov_Map", "Interface", "L1", "Load_Bearing", "NonNV", "README", "SRC_Demand", "SRC_Index"}
    assert set(d) <= allowed, set(d) - allowed
    assert [t[0] for t in d["Alloc"]] == ["C177", "D1020", "D1021", "D1022"]      # 版本字串；v21.2：D1020:D1022 指紋記錄值（test_v2120 看守）
    assert [t[0] for t in d["Alloc_In"]] == ["C23", "D23", "E23", "G23", "G24"]      # v21.2：C23:E23、G23（test_v2120 看守）
    # Cap_In：只有第 344 列以後新增（含用途欄）；既有格不變
    for coord, x, y in d["Cap_In"]:
        assert (x is None and int(re.sub(r"[A-Z]+", "", coord)) >= 344) or coord == "G28", (coord, x, y)      # v21.2：G28（test_v2120 看守）
    # Interface：A364、A368 文字；其餘為第 379 列以後新增；既有列的 R–Y 契約欄不變
    for coord, x, y in d["Interface"]:
        row = int(re.sub(r"[A-Z]+", "", coord))
        assert coord in ("A364", "A368") or (row >= 379 and x is None), (coord, x, y)
    # L1：T4 表頭；其餘為第 145 列以後新增
    for coord, x, y in d["L1"]:
        row = int(re.sub(r"[A-Z]+", "", coord))
        assert coord in ("T4", "C71", "H71") or (row >= 145 and x is None), (coord, x, y)      # v21.2：C71／H71（test_v2120 看守）
    assert sorted(t[0] for t in d["Checks"]) == ["D118", "D119", "D86", "E118"] or sorted(t[0] for t in d["Checks"]) == ["D118", "D119", "E118"], d["Checks"]
    nn = sorted(t[0] for t in d["NonNV"])
    assert nn == sorted(["A1", "L5", "L6", "L7", "L8", "L9"] + [f"{c}{r}" for r in (10, 11) for c in "ABCDEFGHIJKLM"]), nn
    # SRC_Demand：W10 前置文字、兩筆新紀錄（第 93、94 列）、檢查公式 X 欄範圍隨之延伸
    for coord, x, y in d["SRC_Demand"]:
        row = int(re.sub(r"[A-Z]+", "", coord))
        col = re.sub(r"\d+", "", coord)
        assert coord == "W10" or row in (93, 94) or col == "X", (coord, x, y)
    assert {t[0] for t in d["DB_Evidence"]} <= {f"{c}{r}" for c in "ABCDEFGHIJKLMNOPQR" for r in (326, 327)}      # v21.2：E403 在第 327 列
    assert {t[0] for t in d["Decisions"]} <= {f"{c}{r}" for c in "ABCDEFGHIJ" for r in (330, 331)}      # v21.2：SL59-01 在第 331 列
    print({k: len(v) for k, v in d.items()})      # 供報告核對：Gov_Map／SRC_Index／Load_Bearing 的連帶差異格數


# ---------------------------------------------------------------- 9. 記錄與文字
def test_v2110_records_and_text(wbf):
    d = wbf["Decisions"]
    rows = [r for r in range(5, d.max_row + 1) if d.cell(r, 1).value == "SL58-01"]
    assert len(rows) == 1 and d.cell(rows[0], 2).value == "v21.1" and d.cell(rows[0], 7).value == "v21.1 已建"
    txt = d.cell(rows[0], 4).value + d.cell(rows[0], 6).value
    for key in ("當然要放", "J one J two 都同意", "NonNV Q1–Q4", "J1", "J2"):
        assert key in txt, key
    rm = wbf["README"]
    assert str(rm["B5"].value).startswith("20261023_Tokenomics_v21.2（次版，底版 v21.1：") and "v21.1（主版：" in str(rm["B5"].value) and "v21.0（主版：" in str(rm["B5"].value) and str(rm["A1"].value).startswith("Tokenomics v21.2")      # v21.2：版本字串
    rr = [x for x in range(4, rm.max_row + 1) if rm.cell(x, 1).value == "v21.1 NonNV 昇騰／Maia 與免費占比路徑"]
    assert len(rr) == 1 and "L1_FreeShare" in rm.cell(rr[0], 2).value and "IF_NonNV_Out_／IF_NonNV_Hold_ × Ascend／Maia" in rm.cell(rr[0], 2).value
    nv = wbf["NonNV"]
    assert nv["A10"].value == "華為昇騰 910C（CloudMatrix 384）" and nv["A11"].value == "Microsoft Maia 200"
    assert nv["K10"].value == nv["K11"].value == "Assumed"
    for r in (10, 11):
        assert nv[f"H{r}"].value == f"=F{r}/C{r}" and nv[f"I{r}"].value == f"=E{r}/D{r}" and nv[f"J{r}"].value == f"=G{r}/B{r}"
        assert "已接 Interface P 節（v21.1）" in nv[f"M{r}"].value and "H:J 欄未接線" in nv[f"M{r}"].value
    assert "559 kW" in nv["L10"].value and "0.13" in nv["L10"].value and "PR #149" in nv["L10"].value
    assert "0.68" in nv["L11"].value and "1.7 倍" in nv["L11"].value and "$0.5／0.8／1.2M" in nv["L11"].value
    for r in range(5, 10):
        assert nv[f"L{r}"].value.endswith("（Andy 2026-10-23 Q2）。") and "持有比為慣例值" in nv[f"L{r}"].value, r
        assert wbf["NonNV"][f"M{r}"].value.startswith("已接 Interface P 節（v21.0）")      # v21.0 的 M5:M9 不動
    assert "B:G 已接 Interface P 節；H:J 未接線" in nv["A1"].value and "待用" not in nv["A1"].value
    ws = wbf["Interface"]
    for nm in IF_PRE:
        r = int(wbf.defined_names[nm].attr_text.split("$")[-1])
        assert f"[{nm}]" in ws[f"A{r}"].value and nv[f"A{IF_PRE[nm][0]}"].value in ws[f"A{r}"].value, nm
        assert ws[f"R{r}"].value and ws[f"R{r}"].value != "—", nm
    assert "昇騰（910C／CM384）、Maia 200 於 v21.1 加入（第 379–382 列；持有比依每 kW 設備價推導，既有五列為慣例值）" in ws["A368"].value and "待證據後" not in ws["A368"].value
    a364 = ws["A364"].value
    sr = wbf["Spec_Rack"]
    marks = "；".join(f"{g} {sr[f'{c}17'].value}" for g, c in zip(("Hopper", "GB200", "GB300", "VR200", "Rubin Ultra"), "CDEFG"))
    assert "標記（Spec_Rack 第 17 列，依世代）：" + marks in a364 and "標記沿用" not in a364
    assert "基準倍數 2.0（Spec_Rack G12＝F12 × 2）" in a364
    g24 = wbf["Alloc_In"]["G24"].value
    assert g24.startswith("純利用率＝平均 token 吞吐 ÷ 容量；0.41＝原 0.35 ÷ 生產折減 0.85。無直接量測（2026-10-23 查證，PR #151／#152）") and g24.endswith("暫定、待證據。")
    assert "Azure 尖峰／平均換算 0.25–0.59" in g24 and "Aegaeon 長尾專屬實例 0.13–0.34" in g24
    sd = wbf["SRC_Demand"]
    assert sd["A10"].value == "SRC_DEM_006" and sd["W10"].value.startswith("出處待補；已排除 The Information 2026-02〈Why OpenAI, Anthropic Missed Their Own Gross Margin Forecasts〉（2026-10-23 查證 J5）｜外部最近為 2025-09 預測逾 90 億（2026-10-23 9a 查證 J1，docs/evidence/20261023_Q側缺口_A段/）｜—｜v17.2 Q1")      # v21.2：W2 插入句


def test_v2110_sources_and_evidence(wbf):
    sd = wbf["SRC_Demand"]
    ids = {sd.cell(r, 1).value: r for r in range(5, sd.max_row + 1)}
    assert max(k for k in ids if str(k).startswith("SRC_DEM_")) == "SRC_DEM_090"
    for sid, val, yr in (("SRC_DEM_089", 0.34, 2026), ("SRC_DEM_090", 0.06, 2030)):
        r = ids[sid]
        assert sd[f"C{r}"].value == val and sd[f"F{r}"].value == "比例" and str(yr) in sd[f"B{r}"].value
        assert sd[f"K{r}"].value == 3 and sd[f"L{r}"].value == "利害關係方" and sd[f"N{r}"].value == "二手" and sd[f"O{r}"].value == "Active" and sd[f"Q{r}"].value == "E402"
        assert "付費牆" in sd[f"J{r}"].value and "BigGo 2026-02-25" in sd[f"J{r}"].value and "WallStreetCN 2026-02-26" in sd[f"J{r}"].value
        assert "有呈現付費率上升的誘因" in sd[f"M{r}"].value
    ev = wbf["DB_Evidence"]
    rows = [r for r in range(5, ev.max_row + 1) if ev.cell(r, 1).value == "E402"]
    assert len(rows) == 1 and ev.cell(rows[0] - 1, 1).value == "E401"
    assert "SRC_DEM_089" in ev.cell(rows[0], 13).value and ev.cell(rows[0], 16).value == 3 and ev.cell(rows[0], 17).value == "利害關係方"


def test_v2110_cap_in_q_section(wbf):
    ws = wbf["Cap_In"]
    assert ws["A344"].value.startswith("Q. 免費占比時間路徑（v21.1") and ws.max_row == 348
    assert ws["A345"].value.startswith("免費服務占服務 — 2026（投資人資料預估）") and ws["A346"].value.startswith("免費服務占服務 — 2030（投資人資料預估）")
    assert ws["C345"].value == "=SRC_DEM_089" and ws["C346"].value == "=SRC_DEM_090"
    assert (ws["D345"].value, ws["E345"].value, ws["D346"].value, ws["E346"].value) == (0.25, 0.46, 0.03, 0.30)
    assert ws["F345"].value == ws["F346"].value == "Interested-party（區間：Assumed）"
    assert "口徑同 B4_Free" in ws["G345"].value and "不回流" in ws["G345"].value and "『付費』可能含 5–8 美元廣告方案" in ws["G346"].value
    assert "F 高 → 倍數低" in ws["A347"].value and "F 高 → 倍數低" in ws["A348"].value
    for r in (347, 348):
        for c in "CDE":
            f = ws[f"{c}{r}"].value
            assert f.count("ISNUMBER") >= 2 and "AL_UnavailTxt" in f and "B4_Free" in f, (c, r)
    l1 = wbf["L1"]
    assert l1["T147"].value == l1["T148"].value == "情境" and l1["T145"].value == l1["T146"].value == "估計"
    assert l1["S147"].value == "token 帳未閉合；m 可能隨廣告方案下降（未建模）；利害關係方預估"
    for r in (147, 148):
        assert l1[f"H{r}"].value.startswith("情境列（F 隨時間下降）；只把（1 − F）換成") and "利用率 0.30／0.65（同第 129 列）" in l1[f"H{r}"].value
        for c in "DEF":
            f = l1[f"{c}{r}"].value
            assert f.startswith("=IF(AND(ISNUMBER(L1_RevGW_PQ_VR200") and "AL_UnavailTxt" in f, (c, r)
    assert "情境" in l1["T4"].value


def test_v2110_gov_map_registered(wbf):
    gm = wbf["Gov_Map"]
    have = {(gm.cell(r, 3).value, gm.cell(r, 4).value) for r in range(5, gm.max_row + 1)}
    for r in (10, 11):
        for c in "BCDEFG":
            assert ("NonNV", f"{c}{r}") in have, (c, r)
    assert ("Cap_In", "D345:E345") in have and ("Cap_In", "D346:E346") in have
