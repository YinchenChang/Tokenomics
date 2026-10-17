# v15.2 (work order docs/workorders/20261017_v15.2_m證據B段.md r1, branch claude/v15-2): m evidence B stage -- inputs changed by the A-stage verdicts B1-B6.
# MINOR version: input values change, names and definitions do not.
#   B2  Cap_In C280／D280／E280 (CAP_Px25_Astra, _Lo, _Hi) 7.174／2.089／16.33 -> 4.5／3.3／5.9 (volume-weighted instead of days-on-sale weighted); G280 note, A280 label.
#   B3  Calib D67 (GB200 eta_d multiple) 1 -> 1.15; Gov_Map GM342 range 0.5-1.5 -> 0.9-1.45 and note (MLPerf v6.0 same-software pairing).
#   V2  Gov_Map GM342 category Analogy -> Derived (revision 9 correction exception: original wording kept in Gov_Map AL, Decisions SL47-05).
#   B6  Gov_Map GM457 M note (Astra tier-mix range).   B1  SRC_DEM_004 W note (Lo stays 7.5; J1 withdrawn).   B5  Gov_Map GM644 N note (kappa stays 1.0).
#   V3  J2 label SRC_DEM_018 B22 + W22; B4 Calib G9 label; Calib J67 sentence appended.
#   V4  Decisions SL47-01..05;  V5  DB_Evidence E387..E392.
# Every write to an Excel-owned cell is guarded (old value / old text); appended texts keep the original wording (CLAUDE.md section 6, revision 9).
# Gov_Map edits run in gov.py's update chain (gov_update) so that Load_Bearing (statically expanded from Gov_Map F) is built from the corrected category.
from openpyxl.utils import get_column_letter as L

import v1020
import v1410
import v1510

VERSION = "20261017_Tokenomics_v15.2"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-17"
V = "v15.2"
WO = "工作單 docs/workorders/20261017_v15.2_m證據B段.md（r1）"
A_REP = "docs/reports/20261017_m證據A段.md"
DEC_BY = ("Andy 2026-10-17 20:35「两题照建议」、20:46「判断两题皆同一建议，请继续下去研究，开始研究 A 段。」、21:20「照建議繼續下去」（B1–B6 依預設）；"
          "皆為 chat 端建議，可撤回")

B2_NEW = {"C": 4.5, "D": 3.3, "E": 5.9}
B2_OLD = {"C": 7.174, "D": 2.089, "E": 16.33}
A280_OLD = "（參考請求混合單價，表列未折扣；在售天數加權）"
A280_NEW = "（參考請求混合單價，表列未折扣；用量加權）"
G280_NOTE = ("｜v15.2 B2：改用量加權（原為在售天數加權 7.174，區間 2.089–16.33，保留於本說明）。量指數＝Azure 推論季別支出（Q1 2.075、Q2 2.947、Q3 3.648、Q4＝Q3×1.075，SRC_DEM_018 同源），"
             "GPT-5（2025-08-07）後推理占比 ×2（Assumed；Altman 2025-08-10 用戶比例：免費 <1%→7%、Plus 7%→24%）。只量指數 5.88、×3.4 為 3.67、×7 為 2.82；基準 4.5、區間 3.3–5.9。"
             "各模型 token 占比：已搜尋，未找到。來源 " + A_REP + " 第 1 節")

D67_OLD, D67_NEW = 1, 1.15
GM342_K = (0.5, 0.9)
GM342_L = (1.5, 1.45)
GM342_M_OLD = "×0.5–×1.5（Sens_Perf 情境）"
GM342_M_NEW = "×0.9–×1.45（v15.2 B3；原 ×0.5–×1.5 Sens_Perf 情境）"
GM342_N_OLD = "非 Hopper 世代沿用 GB300 擬合的 η_d（Block 2 最大 Analogy）"
GM342_N_ADD = ("｜v15.2 B3：MLPerf v6.0 NVIDIA 同軟體配對（6.0-0078 GB300×72 對 6.0-0075 GB200×72，TensorRT 10.14）Offline 每 GPU 比 R：DSR1 1.331、Llama 3.1 405B 1.262、Llama 2 70B 1.269、GPT-OSS-120B 1.186 "
               "→ 1.5÷R＝1.13–1.27；InferenceX DSV4-Pro vLLM 16／20 tok/s/user 1.25／1.16（第二來源，Interested-party）；頻寬受限點（R 1.04–1.08）與容量受限點不用。來源 " + A_REP + " 第 2 節")
GM342_AL = ("v15.2 更正（修訂九，Decisions SL47-05）前的原值：類別＝Analogy｜理由（N 欄）＝" + GM342_N_OLD + "｜區間＝0.5–1.5（K 0.5、L 1.5）｜區間文字（M 欄）＝" + GM342_M_OLD +
            "｜v15.2 起：類別＝Derived，區間 0.9–1.45（MLPerf v6.0 同軟體配對，Calib D67＝1.15）")
J67_ADD = "；（v15.2 起 D67＝1.15，依 MLPerf v6.0 配對，見 GM342）"

GM457_ADD = "；v15.2 B6：Astra 0.10–0.30（Assumed；Luna 與 Sol 依 0.4：0.45 補足）"
GM644_ADD = ("｜v15.2 B5：維持 1.0。v5.36A 1.035 與 v15.0 重算 0.852 為同一彙總法（需求加權調和平均＝GW 加權算術平均）；差距＝需求權重（API 0.55→0.39，約 −0.10）＋GPU 秒校準（約 −0.08）；"
             "2025 流量證據範圍 0.67–1.05（A 段第 3 節）。結構問題（κ 乘 token、單價固定參考請求）另立一版")

W8_ADD = ("｜v15.2 B1：低端 7.5 維持。J1（低端改 6.0，Andy 2026-10-17 20:46 同意）已撤回：營業成本內算力約 48–52 億 < 付費＋免費推論約 78 億，免費推論大概率不在營業成本，6.0 的前提不成立（"
          + A_REP + " 第 4 節）")
B22_OLD = "OpenAI 2025 推論支出（Azure 帳單，全年推估）"
B22_NEW = "Microsoft 端推論類工作負載（含研究推論；Azure 帳單，全年推估）"
W22_ADD = ("｜v15.2 J2：標籤更正（調和分析：Microsoft 稱推論的 126 億中約 65.5 億未進 OpenAI 營業成本，最可能為研究推論列研發；docs/reports/20261017_推論支出調和分析.md）")
G9_OLD = "驗證（舊軟體）"
G9_NEW = "驗證（舊軟體；容量受限點，不用於 η 驗證）"


def _row_of(ws, rid, col=1, first=5):
    rows = [r for r in range(first, ws.max_row + 1) if ws.cell(r, col).value == rid]
    assert len(rows) == 1, (ws.title, rid, rows)
    return rows[0]


# ================================================================== Gov_Map (called from gov.py's update chain, before Load_Bearing is expanded)
def gov_update(ws, append_text):
    n = 0
    r = v1020._gm_row(ws, "Calib", "D67")
    if ws[f"F{r}"].value == "Analogy":
        assert ws[f"N{r}"].value == GM342_N_OLD and ws[f"K{r}"].value == GM342_K[0] and ws[f"L{r}"].value == GM342_L[0] and ws[f"M{r}"].value == GM342_M_OLD, (ws[f"N{r}"].value, ws[f"K{r}"].value)
        assert ws[f"AL{r}"].value in (None, ""), ws[f"AL{r}"].value
        ws[f"AL{r}"].value = GM342_AL
        ws[f"F{r}"].value = "Derived"
        ws[f"K{r}"].value = GM342_K[1]
        ws[f"L{r}"].value = GM342_L[1]
        ws[f"M{r}"].value = GM342_M_NEW
        n += 6
    n += int(append_text(ws[f"N{r}"], GM342_N_ADD))
    r57 = v1020._gm_row(ws, "Cap_In", "C31:E31")
    n += int(append_text(ws[f"M{r57}"], GM457_ADD))
    r44 = v1020._gm_row(ws, "Alloc_In", "C26")
    n += int(append_text(ws[f"N{r44}"], GM644_ADD))
    return n


# ================================================================== V1 / V3 (late; Excel-owned cells)
def b2(wb, append_text):
    ws = wb["Cap_In"]
    assert str(ws["A280"].value).startswith("2025 年度層級單價：Astra"), ws["A280"].value
    n = 0
    for c in "CDE":
        cur = ws[f"{c}280"].value
        if cur == B2_NEW[c]:
            continue
        assert cur == B2_OLD[c], (c, cur)
        ws[f"{c}280"].value = B2_NEW[c]; n += 1
    a = ws["A280"].value
    if A280_OLD in a:
        ws["A280"].value = a.replace(A280_OLD, A280_NEW); n += 1
    else:
        assert A280_NEW in a, a
    n += int(append_text(ws["G280"], G280_NOTE))
    return n


def b3_b4(wb, append_text):
    ws = wb["Calib"]
    n = 0
    cur = ws["D67"].value
    if cur != D67_NEW:
        assert cur == D67_OLD, cur
        ws["D67"].value = D67_NEW; n += 1
    n += int(append_text(ws["J67"], J67_ADD))
    g9 = ws["G9"].value
    if g9 != G9_NEW:
        assert g9 == G9_OLD, g9
        ws["G9"].value = G9_NEW; n += 1
    return n


def b1_j2(wb, append_text):
    ws = wb["SRC_Demand"]
    r4, r18 = _row_of(ws, "SRC_DEM_004"), _row_of(ws, "SRC_DEM_018")
    assert ws[f"D{r4}"].value == 7.5, ws[f"D{r4}"].value          # B1: Lo stays 7.5
    n = int(append_text(ws[f"W{r4}"], W8_ADD))
    b = ws[f"B{r18}"].value
    if b != B22_NEW:
        assert b == B22_OLD, b
        ws[f"B{r18}"].value = B22_NEW; n += 1
    n += int(append_text(ws[f"W{r18}"], W22_ADD))
    return n


# ================================================================== V5 DB_Evidence E387..E392 (columns A..Q, R repeats the verdict as v13.1 did)
EV_COMMON = "讀取者：chat 端（" + A_REP + "）｜Andy 2026-10-17 20:35「两题照建议」、20:46「判断两题皆同一建议，请继续下去研究，开始研究 A 段。」、21:20「照建議繼續下去」（B1–B6 依預設）（文字，未改字）；chat 端建議，可撤回"
EVIDENCE = [
    ("E387", "Altman：ChatGPT 每日使用推理模型用戶比例 免費 <1%→7%、Plus 7%→24%（GPT-5 發布後）",
     "x.com/sama/status/1954603417252532479（2025-08-10）", "Interested-party（OpenAI 執行長；用戶比例，不是 token 比例）",
     "Cap_In C280:E280（Astra 2025 單價的推理占比倍數）", "7.174（在售天數加權；區間 2.089–16.33）", "免費 <1%→7%、Plus 7%→24%",
     "部分採用（B2 推理占比倍數的依據；用戶比例非 token 比例）", "用戶比例不等於 token 比例，推理占比 ×2 為 Assumed", "CAP_Px25_Astra*；L1_MFleet*", "利害關係方"),
    ("E388", "OpenRouter／a16z《State of AI》：推理模型 token 占比年初可忽略→年底 >50%；prompt 1.5K→6K+、completion 150→400",
     "https://www.a16z.com/state-of-ai（2025-12）", "Analogy（OpenRouter 平台流量，非 OpenAI 自家流量）",
     "Cap_In C280:E280；Cap_In C31:E31（B6）；Alloc_In C26（κ，B5）", "7.174（在售天數加權）；κ 1.0；付費層級組合各層 ±0.15", "推理 token 占比年底 >50%；prompt 6K+、completion 400",
     "部分採用（Analogy；B2、B6、κ 證據）", "平台流量對 OpenAI 流量為類比；κ 數值不改（B5）", "CAP_Px25_Astra*；L1_MFleet*；AL_Kappa（不改）", "—"),
    ("E389", "MLPerf v6.0 GB300／GB200 同軟體配對（TensorRT 10.14）：Offline 每 GPU 比 DSR1 1.331、Llama 3.1 405B 1.262、Llama 2 70B 1.269、GPT-OSS-120B 1.186（四模型 R 1.19–1.33）",
     "github.com/mlcommons/inference_results_v6.0（6.0-0078 對 6.0-0075）", "Interested-party（提交者 NVIDIA；MLCommons 稽核）",
     "Calib D67（GB200 η_d 倍數；Gov_Map GM342）", "1（沿用 GB300）；區間 0.5–1.5", "1.5÷R＝1.13–1.27；D67＝1.15，區間 0.9–1.45",
     "採用（B3）", "同軟體配對；頻寬受限點（R 1.04–1.08）與容量受限點不用", "IF_TokGW* 的 GB200 欄；GB200 每 GW 營收與成本；L1_MFleet*", "利害關係方"),
    ("E390", "InferenceX DSV4-Pro vLLM GB300／GB200：16／20 tok/s/user R 1.20／1.29（1.5÷R＝1.25／1.16）；27 tok/s/user R 2.82 為容量受限",
     "inferencex.semianalysis.com/blog/gb300-nvl72-vs-gb200-nvl72-dsv4-pro-vllm-fp4", "Interested-party（SemiAnalysis InferenceX；第二來源）",
     "Calib D67（B3 第二來源）；Calib G9（B4 容量受限判讀）", "—", "R 1.20／1.29（16／20 tok/s/user；1.5÷R＝1.25／1.16）；2.82（27 tok/s/user，容量受限）",
     "部分採用（第二來源；B4 容量受限判讀）", "容量受限點（R 2.82）不用於 η 驗證", "Calib G9／G20 標籤；不改 IF_／L1_", "利害關係方"),
    ("E391", "OpenAI 2025 查核財報付 Microsoft 分項：營業成本 60.47、研發 105.9、行銷 5.27、管理 0.42（億美元）；Microsoft 端推論 2024 37.7 對 OpenAI 端 18–20（倍數約 2.1）",
     "techflowpost.com/en-US/article/32082；epoch.ai/data-insights/openai-compute-spend", "查核報表（外流）／Epoch 資料洞察",
     "SRC_DEM_004（B1，Lo 7.5 維持）；SRC_DEM_018（J2 標籤）", "SRC_DEM_004 Lo 7.5；SRC_DEM_018 標籤『OpenAI 2025 推論支出（Azure 帳單，全年推估）』", "Microsoft 稱推論的 126 億中約 65.5 億未進 OpenAI 營業成本",
     "採用（J2、B1；不改數值）", "不改任何數值；只更正標籤與備註", "不改任何 IF_／L1_", "中立"),
    ("E392", "營收分成記帳位置：已搜尋，未找到；Microsoft FY26 10-K 稱 OpenAI 相關營收『inclusive of revenue sharing payments』（媒體轉述）",
     "A 段第 4 節（" + A_REP + "）", "二手（媒體轉述 10-K 字句）",
     "SRC_DEM_004／018 的口徑（營收分成是否在營業成本）", "—", "—",
     "待查", "10-K 原文未讀；分成在 OpenAI 帳上記為成本或營收減項，已搜尋未找到", "不改任何 IF_／L1_", "—"),
]


def evidence_append(wb):
    ws = wb["DB_Evidence"]
    have = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
    r = max(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")) + 1
    n = 0
    from common import put, F_CALC
    for (eid, claim, src, mark, param, cur, new, verdict, note, affected, stance) in EVIDENCE:
        if eid in have:
            continue
        status = "待判定" if verdict.startswith("待查") else "已處理"
        vals = [eid, DATE, claim, src, mark, param, cur, new, verdict, V, note + "｜" + EV_COMMON, status, "—", "—", affected, "—", stance,
                f"{verdict}｜v15.2 新增列（非 106 列積壓）"]
        for i, v in enumerate(vals):
            put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 4, 5, 6, 7, 8, 10, 14, 17))
        r += 1; n += 1
    return n


# ================================================================== V4 Decisions SL47-01..05
def _decisions():
    D = [
        ("κ 維持 1.0（B5）",
         "Alloc_In C26（κ）維持 1.0，不改。20:35「两题照建议」之一；21:20 B5 再確認。理由：v5.36A 1.035 與 v15.0 重算 0.852 為同一彙總法（需求加權調和平均＝GW 加權算術平均）；差距＝需求權重（API 0.55→0.39，約 −0.10）"
         "＋GPU 秒校準（約 −0.08）；2025 流量證據範圍 0.67–1.05。結構問題（κ 乘 token、單價固定參考請求）另立一版。數值不改；Gov_Map GM644 N 欄附註。",
         "Andy 2026-10-17 20:35「两题照建议」（κ 維持 1.0）；21:20「照建議繼續下去」（B1–B6 依預設）；chat 端建議，可撤回",
         "Alloc_In C26；Gov_Map GM644 N"),
        ("推論支出主口徑：SRC_DEM_004 84 億為主、SRC_DEM_018 對照",
         "推論支出以 SRC_DEM_004（8.4 $B，區間 7.5–8.8）為主口徑，SRC_DEM_018（Azure 帳單全年推估 12.6 $B）為對照；不改任何數值。依據：docs/reports/20261017_推論支出調和分析.md（Microsoft 端 126 億中約 65.5 億未進 OpenAI 營業成本）。",
         "Andy 2026-10-17 20:35「两题照建议」（推論支出主口徑 84 億）；chat 端建議，可撤回",
         "SRC_DEM_004、SRC_DEM_018（口徑不變）"),
        ("J1 同意後撤回：SRC_DEM_004 低端維持 7.5（B1）",
         "兩個時點同一筆記錄：(1) 20:46 Andy 同意 J1（低端改 6.0）；(2) 21:20 B1 依 A 段第 4 節撤回——營業成本內算力約 48–52 億 < 付費＋免費推論約 78 億，免費推論大概率不在營業成本，6.0 的前提不成立。"
         "SRC_DEM_004 Lo 維持 7.5（C／D／E 數值不變），理由寫在 W 欄。",
         "Andy 2026-10-17 20:46「判断两题皆同一建议，请继续下去研究，开始研究 A 段。」（J1）；21:20「照建議繼續下去」（B1–B6 依預設，B1 撤回 J1）；chat 端建議，可撤回",
         "SRC_Demand 第 8 列（SRC_DEM_004）W 欄"),
        ("J2 標籤：SRC_DEM_018 改為 Microsoft 端推論類工作負載",
         "SRC_DEM_018 的指標名稱由『OpenAI 2025 推論支出（Azure 帳單，全年推估）』改為『Microsoft 端推論類工作負載（含研究推論；Azure 帳單，全年推估）』，數值 12.6／12.3／13 不變。"
         "依據：調和分析（Microsoft 稱推論的 126 億中約 65.5 億未進 OpenAI 營業成本，最可能為研究推論列研發）。原標籤保留在本筆與 W 欄。",
         "Andy 2026-10-17 20:46「判断两题皆同一建议，请继续下去研究，开始研究 A 段。」（J2）；chat 端建議，可撤回",
         "SRC_Demand 第 22 列（SRC_DEM_018）B、W 欄"),
        ("B2、B3（含 Gov_Map GM342 類別更正）、B4、B6",
         "B2：Cap_In C280／D280／E280（CAP_Px25_Astra、_Lo、_Hi）7.174／2.089／16.33 → 4.5／3.3／5.9（用量加權；Derived 不變）。"
         "B3：Calib D67（GB200 η_d 倍數）1 → 1.15；Gov_Map GM342 區間 0.5–1.5 → 0.9–1.45（MLPerf v6.0 同軟體配對）。"
         "GM342 類別 Analogy → Derived（修訂九更正例外）：更正理由＝GB200 倍數不再是無依據的沿用 GB300，而是由 MLPerf v6.0 同軟體配對（GB300×72 對 GB200×72，四模型 Offline 每 GPU 比 R 1.19–1.33）推得 1.5÷R＝1.13–1.27，"
         "第二來源 InferenceX 1.16–1.25；原字樣『Analogy』與原理由『" + GM342_N_OLD + "』逐字保留在 Gov_Map 同列 N 欄與 AL 欄。"
         "B4：Calib G9 標籤改為『驗證（舊軟體；容量受限點，不用於 η 驗證）』。B6：Cap_In C31:E31 的 Astra 層級組合區間 0.10–0.30（Assumed；基準值不變），寫在 Gov_Map GM457 M 欄。"
         "依據：" + A_REP + " 第 1、2 節。",
         "Andy 2026-10-17 21:20「照建議繼續下去」（B1–B6 依預設）；chat 端建議，可撤回",
         "Cap_In C280:E280（A280、G280）；Calib D67、G9（J67）；Gov_Map GM342（F、K、L、M、N、AL）、GM457 M"),
    ]
    return [[f"SL47-{i + 1:02d}", V, f"v15.2 m 證據 B 段：{t}", txt, DATE, who, "v15.2 已建", where, f"{WO} V4-{i + 1}", "否"]
            for i, (t, txt, who, where) in enumerate(D)]


def apply(wb, append_text):
    log = {}
    log["b2"] = b2(wb, append_text)
    log["b3_b4"] = b3_b4(wb, append_text)
    log["b1_j2"] = b1_j2(wb, append_text)
    log["evidence"] = evidence_append(wb)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == "SL47-01" for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


# ================================================================== README
README_VERSION = (VERSION + "（次版：" + WO + "——m 證據 B 段，依 A 段判定改輸入：B2 Cap_In C280:E280（Astra 2025 單價）7.174／2.089／16.33 → 4.5／3.3／5.9（用量加權）；"
                  "B3 Calib D67（GB200 η_d 倍數）1 → 1.15、Gov_Map GM342 區間 0.9–1.45 且類別 Analogy → Derived（修訂九更正）；B6 Astra 層級組合區間註記；B1、B5 維持（只加註）；"
                  "J2、B4 標籤；Decisions SL47-01…05；DB_Evidence E387…E392；區間端點與指紋重產。具名範圍、定義不變。）以下為 "
                  + v1510.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1510.README_TITLE
README_TITLE = V + _T[len(v1510.V):]
README_ROW = ("v15.2 m 證據 B 段",
              "依 A 段判定改輸入：Astra 2025 單價 7.174 → 4.5（區間 3.3–5.9）；GB200 η_d 倍數 1 → 1.15（區間 0.9–1.45，類別 Derived）；Astra 層級組合區間註記；J2、B4 標籤。名稱與定義不變。細節見 Decisions SL47-01…05。")
