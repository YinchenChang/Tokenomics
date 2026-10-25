# v21.5 (work order docs/workorders/20261025_v21.5_計價倍數新舊約與GM088.md r0, branch claude/v21-5): TEXT / GOVERNANCE ONLY.
# MINOR version: no input value, formula or contract name changes; every contract name keeps its value (except RNG_Version／RNG_Current version strings).
#   W1  DB_Evidence E407 (stock-fleet lessor anchors 0.65-1.32) and E408 (2026 new contracts 1.36-2.30), both after E406; PR #170 L5／L6 noted at the tail of E408.
#   W2  Alloc_In G23 (AL_SpendBasis explanation): v21.5 paragraph (stock fleet vs 2026 new contracts; 1.8 read as a mix; value not changed).
#   W3  Gov_Map GM088 (Spec_Rack F11 H／I／N): HELD. Measured with v2130.GOV_MAP_UPD re-added: Gov_Map H92／I92／N92／R92／AF92 and SRC_Index!A980 change, and so does the value of the named
#       range IDX_SrcID (one cell) = the work order's stop condition. builder/gov.py keeps it off (one-line re-enable); changes listed in docs/workorders/20261025_v21.5_預期變動.csv.
#   W4  Decisions SL62-01, README B5 opening paragraph, README row 75.
# Excel-owned cells (DB_Evidence, Decisions) are written only while absent; Alloc_In G23 gets a prefix only once. Idempotent.
from openpyxl.utils import get_column_letter as L

from common import put, F_CALC
import v1410
import v2140

VERSION = "20261025_Tokenomics_v21.5"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-25"
DASH = "—"
V = "v21.5"
WO = "工作單 docs/workorders/20261025_v21.5_計價倍數新舊約與GM088.md（r0）"
DEC_ID = "SL62-01"
EVD = "docs/evidence/20261025_計價倍數_第二出租方/"
DEC_BY = ("Andy 未回應，照工作單預設（2026-10-25 已在對話列為判斷題 5）：PR #170（計價倍數第二出租方錨點，獨立核對 2026-10-25 r1 後可合併）L1（AL_SpendBasis 1.8 不改值）、"
          "L2（Nebius 實耗讀法只作敏感度）、L4（新舊約混合；分母以 GB300 為基準）；v21.4 報告 P1（v2130.GOV_MAP_UPD 補套用）")

# ---------------------------------------------------------------- W1 DB_Evidence
_OLD_VAL = "AL_SpendBasis 1.8（0.85–2.3，Derived；v21.2 J5 殘差法）"
E407 = [
    "E407", DATE,
    "出租方存量機隊每 GW 年營收 ÷ Tokenomics 持有成本 0.65–1.32（CoreWeave 2026 Q1–Q2、Nebius 2025 Q4、WhiteFiber 2026 Q2）",
    "出租方自揭（CoreWeave 2026 Q1–Q2、Nebius 2025 Q4、WhiteFiber 2026 Q2）；第二出租方錨點查證 " + EVD + "（PR #170）",
    "Interested-party（出租方自揭）", "Alloc_In C23（AL_SpendBasis 1.8；0.85–2.3，不改）",
    _OLD_VAL,
    "存量機隊每 GW 年營收 ÷ Tokenomics 持有成本 0.65–1.32（CoreWeave 2026 Q1–Q2、Nebius 2025 Q4、WhiteFiber 2026 Q2）",
    "弱反對 AL_SpendBasis 1.8；1.8 不改（PR #170 L1，照預設；" + V + " W1）", V,
    "存量機隊平均；Nebius active 為 IT 側可供營收容量（20-F 定義，未明設計或實耗；實耗讀法 0.46–0.75 只作敏感度，L2）；"
    "CoreWeave active power 口徑仍未定義（L3，Q3 財報再查）；WhiteFiber 分母依 10-Q 為冰島 5.5 MW-IT，偏下界｜" + WO,
    "已處理", DASH, DASH, "AL_SpendBasis（不改值）", 3, "利害關係方", "弱反對（1.8 不改，存量機隊平均）｜" + V + " 新增列"]
E408 = [
    "E408", DATE,
    "2026 新約每 GW 年營收 ÷ Tokenomics GB300 持有成本 1.36–2.30（Nebius Q2 大單 ACV 2,000–2,500 萬美元／MW、IREN 近期 3 年約 >2,000 萬／洽談 ~2,500 萬美元 per MW (IT)：1.84–2.30；WhiteFiber 10-Q Blackwell 新約 1.36–1.86）",
    "出租方自揭（Nebius、IREN、WhiteFiber 新約）；第二出租方錨點查證 " + EVD + "（PR #170）",
    "Interested-party（出租方自揭）", "Alloc_In C23（AL_SpendBasis 1.8；0.85–2.3，不改）",
    _OLD_VAL,
    "2026 新約每 GW 年營收 ÷ Tokenomics GB300 持有成本 1.36–2.30",
    "支持 AL_SpendBasis 1.8 作為新增容量價；1.8 讀作新舊約混合，不改（PR #170 L1、L4）", V,
    "分母以 GB300 為基準，VR200 為 1.64–2.04（L4）；Nebius 原文 ACV＝annual contract value；IREN 原文只寫 per MW (IT)，年額讀法為 Derived"
    "（2026 Contracted ARR 40 億 ÷ 0.3 GW 扣 Microsoft 後約 2,060 萬／MW·年佐證）；短約 4,000–5,000 萬／MW（至多 6 個月，3.7–4.6）不納入。"
    "另記：GB300 IT 資本支出 2026 公開點多數高於 TK 32.2（L5）；IREN 推算營運費用 2.6–3.3（L6），只記錄、不改｜" + WO,
    "已處理", DASH, DASH, "AL_SpendBasis（不改值）", 3, "利害關係方", "支持（1.8 讀作新舊約混合，不改）｜" + V + " 新增列"]


def evidence_append(wb):
    ws = wb["DB_Evidence"]
    n = 0
    for prev, row in (("E406", E407), ("E407", E408)):
        if any(ws.cell(r, 1).value == row[0] for r in range(5, ws.max_row + 1)):
            continue
        r = max(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")) + 1
        assert ws.cell(r - 1, 1).value == prev, (ws.cell(r - 1, 1).value, prev)
        for i, v in enumerate(row):
            put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 4, 5, 6, 7, 8, 10, 14, 17))
        n += 1
    return n


# ---------------------------------------------------------------- W2 Alloc_In G23 (newest paragraph first, like v21.3)
G23_HEAD = ("v21.5：出租方錨點依合約時點分兩群——存量機隊平均約 0.65–1.32（CoreWeave、Nebius、WhiteFiber），2026 新約約 1.36–2.30（Nebius、IREN 自報新約 1.84–2.30；WhiteFiber 10-Q 新約 1.36–1.86）；"
            "1.8 讀作新舊約混合，值不改（E407、E408；PR #170）。｜")


def text_fixes(wb):
    ai = wb["Alloc_In"]
    cur = str(ai["G23"].value)
    if cur.startswith(G23_HEAD):
        return 0
    assert cur.startswith("v21.3（2026-10-24 Azure 實付價 A 段 J3"), cur[:60]
    ai["G23"].value = G23_HEAD + cur
    return 1


# ---------------------------------------------------------------- W4 Decisions
def _decisions():
    return [[DEC_ID, V, "v21.5 計價倍數新舊約混合登錄（E407／E408）、Gov_Map GM088 補套用試算（W3 保留現狀；只改文字與治理欄）",
             "W1：DB_Evidence E407（出租方存量機隊每 GW 年營收 ÷ Tokenomics 持有成本 0.65–1.32，弱反對 AL_SpendBasis 1.8）、E408（2026 新約 1.36–2.30，支持 1.8 作為新增容量價；"
             "分母 GB300，VR200 為 1.64–2.04；PR #170 L5、L6 只記錄於 E408 說明欄末）；W2：Alloc_In G23 加 v21.5 段（出租方錨點依合約時點分兩群，1.8 讀作新舊約混合，C23:E23 值與公式不變）；"
             "W3（保留現狀，未套用）：builder/gov.py 試加回 v2130.GOV_MAP_UPD（Gov_Map GM088 H／I／N）並以 LibreOffice 重算比對 v21.4：變動格為 Gov_Map H92／I92／N92／R92（7.8 → 7.52）／AF92（15 → 98）與 SRC_Index!A980（SRC_HW_015 → SRC_HW_098），其中 SRC_Index!A980 使具名範圍 IDX_SrcID 的值改變（1 格），觸發工作單 W3 停止條件，故本版不套用、列待 Project 判斷（清單見 docs/workorders/20261025_v21.5_預期變動.csv）；GOV 三項計數仍 0／134／182。"
             "不做：(1) VR200 P×Q 營收的 IF_ 名稱——新增契約名稱即主版升級（CLAUDE.md 第 4 節），下游已可依契約讀 L1_RevGW_PQ_VR200，留待下一次主版；"
             "(2) IF_PxEffMix_Lo／_Hi——該值本身無區間，屬新機制（判斷類），不在本單。AL_SpendBasis 1.8 不改值；契約名稱與值不變。"
             "Andy 原話與依據：" + DEC_BY,
             DATE, DEC_BY, "v21.5 已建",
             "DB_Evidence E407／E408；Alloc_In G23；README B5／第 75 列；Decisions SL62-01（Gov_Map GM088 未動）",
             f"{WO} W1–W4", "否"]]


def apply(wb, append_text):
    log = {"text_fixes": text_fixes(wb)}
    log["evidence"] = evidence_append(wb)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == DEC_ID for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


# ---------------------------------------------------------------- README
_prev = v2140.README_VERSION.split("_Tokenomics_", 1)[1]
README_VERSION = (VERSION + "（次版，底版 v21.4：" + WO + "——只改文字與治理欄；所有契約名稱的值不變（版本字串除外）；GOV 三項計數 0／134／182 不變。"
                  "DB_Evidence E407（出租方存量機隊錨點 0.65–1.32）、E408（2026 新約錨點 1.36–2.30），AL_SpendBasis 1.8 不改值、讀作新舊約混合（Alloc_In G23 加說明）；"
                  "Gov_Map GM088 補套用 v21.3 更新經試算會使具名範圍 IDX_SrcID 改變，依停止條件本版不套用（待 Project 判斷）；區間端點與時效指紋重產；Decisions SL62-01。）以下為 " + _prev)
_T = v2140.README_TITLE
README_TITLE = V + _T[len(v2140.V):]
README_ROW = ("v21.5 計價倍數新舊約混合登錄（Gov_Map GM088 暫緩）",
              "只改文字與治理欄，所有契約名稱的值不變。AL_SpendBasis（1.8）不改值：出租方錨點依合約時點分兩群——存量機隊平均約 0.65–1.32（E407），2026 新約約 1.36–2.30（E408），"
              "1.8 讀作新舊約混合；Alloc_In G23 加說明。Gov_Map GM088（Spec_Rack F11 VR200 機架價低端）判斷欄補套用 v21.3 更新：試算會使具名範圍 IDX_SrcID 的值改變（SRC_Index A980），依工作單停止條件本版不套用，待 Project 判斷。"
              "細節見工作單 docs/workorders/20261025_v21.5_計價倍數新舊約與GM088.md 與 Decisions SL62-01。")
