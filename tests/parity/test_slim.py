"""瘦身 4-1（v10.1）新增的常設檢查：鏡像與複製頁一致、Alloc G 節基準列對 A–D 推導鏈、死格不增加、Gov_Map 衛生。
鏡像與死格檢查讀公式文字（不重算）；Alloc 檢查讀 LibreOffice 重算值。"""
import re
from functools import lru_cache
from pathlib import Path

import pytest

from conftest import model_path
import mirror_lib
import deadcells_lib

REPO = Path(__file__).resolve().parents[2]
CONTRACT = REPO / "docs" / "plan" / "Tokenomics_downstream_contract.md"
# 死格基準（v10.1 建置值；與 4-0 同一定義，另加 Sens_／_Prod 頁為根）。新增死格需先在工作單說明，否則此測試失敗。
DEAD_BASELINE = 2771      # v14.1（R1／R5）：2798 → 2771（−27：R5 清空 24 個死格；R1 使 Cap_In 第 121 列 C:E 轉為活格 −3）； v14.0（4-3B）：3604 → 2798（−806：刪除 73 列的 514 個死格，A 段『升級（隨附上游）』與新名稱使上游轉為活格，其餘依刪除後的上游重算；無標記死格 0）； v13.1（D15 B 段 P2-2）：3603 → 3604（+1＝Alloc P 節『指紋格數』顯示格，無公式讀取，已填『用途』欄）； v12.0（v11.1 修補 H5）：3598 → 3603（+5＝Alloc N 節「區間（情境）」基準列的即時值公式 C–G 欄 × 1 列，顯示用、無人讀取，已填「用途」欄；其餘列為常數不是死格）； v11.0（大批次修正 G3）：3574 → 3598（+24＝Calib L 節驗證點複本列的顯示格 C、D、E 欄 × 8 列〔第 217、218、249–254 列：實測目標、世代、模型 ÷ 實測、有效 PF、非算力項下限、達標所需 decode tok/s、每序列每步時間、η_d 倍數解〕，皆為 V20–V22 的顯示用複本、無人讀取，已填「用途」欄；同 v10.2 E1 對 J 節的處理）； v10.3（4-2a）：3599 → 3574（F3：Tech_Registry K、T 欄接進 Checks V，−25，Tech_Registry 的 K、T 欄不再是死格；下調＝更嚴）；v10.2 E1：+105＝Calib J 節 C、D 欄（校準點複本列，工作單 E1 要求；顯示模型 ÷ 目標與所需倍數，無下游讀取）；v10.1：3494


FIXTURES = Path(__file__).resolve().parent / "fixtures"
PURPOSE = "用途"


@lru_cache(maxsize=1)
def _model():
    return mirror_lib.Model(model_path())


def test_prod_mirror_no_undocumented_difference():
    """8 個 _Prod 頁與核心頁逐格結構相同；不同處必須是頁首明訂者。"""
    assert _model().prod_undocumented() == []


def test_sens_core_no_undocumented_difference():
    """Sens_Perf／Sens_Train／Sens_Har／Sensitivity 與核心頁的差異必須落在黃底格。"""
    assert _model().sens_core_undocumented() == []


@pytest.mark.parametrize("sp", ["Sensitivity", "Sens_Har"])
def test_yellow_cells_match_differing_cells(sp):
    """黃底公式格與「公式和第一欄不同的格」一一對應。"""
    diff, yellow = _model().row_diff_cells(sp)
    assert set(diff) == set(yellow), (sorted(set(diff) - set(yellow))[:10], sorted(set(yellow) - set(diff))[:10])


def test_blue_scenario_cells_have_yellow():
    assert _model().blue_scenario_without_yellow() == []


def test_perf_batch_differs_only_on_whitelist_rows():
    rows = _model().perf_batch_diff_rows()
    assert set(rows) <= mirror_lib.PERF_BATCH_ROWS, sorted(set(rows) - mirror_lib.PERF_BATCH_ROWS)


def test_theory_rev_prod_header_note():
    import openpyxl
    ws = openpyxl.load_workbook(model_path(), read_only=True)["Theory_Rev_Prod"]
    assert "90" in str(ws["A2"].value) and "96" in str(ws["A2"].value)


def test_alloc_g_base_row_equals_chain(lo):
    """Alloc G 節基準列（第 81 列）的研發 GW、Q1、Q2、服務 GW 與 A–D 推導鏈輸出相同（容差 1e-9）。"""
    v = lo.view({})
    row = lambda c: v.get("Alloc", f"{c}81")
    pairs = [("R", "AL_RDGWg"), ("S", "AL_Q1g"), ("T", "AL_Q2g"), ("W", "IF_AllocServeGW")]
    for col, name in pairs:
        a, b = row(col), v.get_name(name)
        assert abs(a - b) <= 1e-9 * max(1.0, abs(b)), (col, name, a, b)


def test_dead_formula_cells_not_increased():
    dead, nform, per_page = deadcells_lib.dead_formula_cells(model_path(), CONTRACT)
    assert len(dead) <= DEAD_BASELINE, (len(dead), per_page.most_common(5))


def test_gov_map_no_double_registration():
    """同一個儲存格不得被兩列（非停用）登錄。"""
    import openpyxl
    from openpyxl.utils import range_boundaries
    ws = openpyxl.load_workbook(model_path(), read_only=True)["Gov_Map"]
    seen, dup = {}, []
    for r in ws.iter_rows(min_row=5, max_col=4, values_only=True):
        gid, _, sheet, rng = r
        if not (isinstance(gid, str) and gid.startswith("GM") and sheet and rng):
            continue
        if str(rng).startswith("（停用）"):
            continue
        for part in str(rng).split(","):
            part = part.strip().replace("$", "")
            try:
                c1, r1, c2, r2 = range_boundaries(part)
            except Exception:
                continue
            if None in (c1, r1, c2, r2) or (c2 - c1 + 1) * (r2 - r1 + 1) > 400:
                continue
            for cc in range(c1, c2 + 1):
                for rr in range(r1, r2 + 1):
                    k = (sheet, cc, rr)
                    if k in seen and seen[k] != gid:
                        dup.append((sheet, cc, rr, seen[k], gid))
                    seen.setdefault(k, gid)
    assert dup == [], dup[:10]


def test_dead_cells_have_purpose_mark():
    """4-2a F4（原則 2）：每個死格都必須有「用途」標記。基準名單＝v10.3 建置時用途欄為空的死格（fixtures/dead_unmarked_baseline_v10.3.json，
    其中 2,3xx 格屬 4-0 的 B 類『視判斷可刪』，待 4-3 判斷批次逐列決定刪除或標記）；基準之外新增的死格若沒有標記，測試失敗。
    Calib J 節 105 格與 Alloc 27 格已標記，不在基準名單內，所以拿掉標記也會失敗。"""
    import json
    import openpyxl
    dead, _, _ = deadcells_lib.dead_formula_cells(model_path(), CONTRACT)
    base = set(json.loads((FIXTURES / "dead_unmarked_baseline_v10.3.json").read_text(encoding="utf-8"))["unmarked"])
    wb = openpyxl.load_workbook(model_path())
    pcol = {}
    for ws in wb.worksheets:
        for c in range(3, ws.max_column + 2):
            if ws.cell(3, c).value == PURPOSE:
                pcol[ws.title] = c
    new_unmarked = []
    for s, a in sorted(dead):
        ws = wb[s]
        col = pcol.get(s)
        marked = bool(col and ws.cell(ws[a].row, col).value not in (None, ""))
        if not marked and f"{s}!{a}" not in base:
            new_unmarked.append(f"{s}!{a}")
    assert new_unmarked == [], (len(new_unmarked), new_unmarked[:20])
    assert all(isinstance(x, str) and "!" in x for x in base)


def test_alloc_fleet_1gw_same_formula_rows_agree(lo):
    """4-2a F2 D11：Alloc（每 GW 口徑，欄 C–F＝服務世代 Hopper／GB200／GB300／VR200）與 Fleet_1GW（1 GW 機隊，欄＝世代 × 3 成本情境）同式的列，換算單位後相等（容差 1e-9）。
    不合併的理由：單位不同（每 GW 對每 1 GW 機隊）；改一邊漏另一邊時，此測試失敗。對應欄＝Fleet_1GW 第 3*(g−1)+2 個成本欄（Alloc 第 17 列的索引）。"""
    v = lo.view({})
    same_rows = [(19, 38), (20, 39), (21, 40), (22, 41), (23, 42), (26, 18)]          # 可達旗標、保留層級權重合計、可服務旗標：兩邊同值
    conv_rows = [(29, 19, 11), (30, 20, 12)]                                         # 每 GW 年產能 × 服務 GW（Fleet_1GW 第 11／12 列）＝服務 token 總量

    def close(a, b):
        if isinstance(a, str) or isinstance(b, str):
            return a == b
        return abs(a - b) <= 1e-9 * max(1.0, abs(b))

    n = 0
    for g, ac in enumerate("CDEF", start=1):
        idx = 3 * (g - 1) + 2
        fc = openpyxl_col(idx + 2)          # Fleet_1GW 欄 C 為第 1 個成本欄
        assert v.get("Alloc", f"{ac}17") == idx
        for ar, fr in same_rows:
            a, b = v.get("Alloc", f"{ac}{ar}"), v.get("Fleet_1GW", f"{fc}{fr}")
            assert close(a, b), (g, ar, fr, a, b)
            n += 1
        for ar, fr, gr in conv_rows:
            a, b, gw = v.get("Alloc", f"{ac}{ar}"), v.get("Fleet_1GW", f"{fc}{fr}"), v.get("Fleet_1GW", f"{fc}{gr}")
            if isinstance(a, str) or isinstance(b, str):
                assert a == b, (g, ar, fr, a, b)
            else:
                assert close(a * gw, b), (g, ar, fr, a, gw, b)
            n += 1
    assert n == 4 * (len(same_rows) + len(conv_rows))


def openpyxl_col(i):
    from openpyxl.utils import get_column_letter
    return get_column_letter(i)


def test_sens_rev_life_reads_cap_in(lo):
    """4-2a F2 D08：Sens_Rev 的商業壽命情境 B20／B21 讀 Cap_In D23／D24（商業壽命低／高，Sol 欄），標籤隨值。"""
    import openpyxl
    wb = openpyxl.load_workbook(model_path(), read_only=True)
    ws = wb["Sens_Rev"]
    assert ws["B20"].value == "=Cap_In!D23" and ws["B21"].value == "=Cap_In!D24"
    v = lo.view({})
    assert v.get("Sens_Rev", "B20") == v.get("Cap_In", "D23") == 6 and v.get("Sens_Rev", "B21") == v.get("Cap_In", "D24") == 24
    assert v.get("Sens_Rev", "A20") == "壽命 6 個月" and v.get("Sens_Rev", "A21") == "壽命 24 個月"


def test_cap_in_h10_2025_column_reads_h_section(lo):
    """4-2a F2 D09：Cap_In H10 2025 欄（C206–C209、C222–C223）直接引用 H 節對應格；第 224 列一致性檢查保留。"""
    import openpyxl
    ws = openpyxl.load_workbook(model_path(), read_only=True)["Cap_In"]
    for r, src in ((206, 144), (207, 145), (208, 146), (209, 94), (223, 119)):
        assert ws[f"C{r}"].value == f"=C{src}", (r, ws[f"C{r}"].value)
    assert ws["C222"].value is None      # v14.1 R5：C222（機隊付費 token 的 2025 年度複本，無人讀取）已清空；其餘 2025 欄連結不變
    assert "一致性" in str(ws["A224"].value) or "自我檢查" in str(ws["A224"].value)
    assert lo.view({}).get("Cap_In", "C224") == "一致"


def test_exposed_parameters_are_inputs_and_registered(lo):
    """4-2a F1：L1_JudgeTol＝0.2、SR_RU_PxMultLo＝1.5、SR_RU_PxMultHi＝2.5 是具名輸入；L1 欄 P 與 Spec_Rack G11／G13 讀它們；Gov_Map 有登錄且標 Assumed。"""
    import openpyxl
    wb = openpyxl.load_workbook(model_path(), read_only=True)
    v = lo.view({})
    assert v.get_name("L1_JudgeTol") == 0.2 and v.get_name("SR_RU_PxMultLo") == 1.5 and v.get_name("SR_RU_PxMultHi") == 2.5
    sr = wb["Spec_Rack"]
    assert sr["G11"].value == "=F12*SR_RU_PxMultLo" and sr["G13"].value == "=F12*SR_RU_PxMultHi"
    l1 = wb["L1"]
    n = sum(1 for r in range(5, 140) if isinstance(l1[f"P{r}"].value, str) and "L1_JudgeTol" in l1[f"P{r}"].value)
    assert n == 96, n
    assert not any(isinstance(l1[f"P{r}"].value, str) and "<=0.2," in l1[f"P{r}"].value for r in range(5, 140))
    gm = wb["Gov_Map"]
    rows = {tuple(r[:2]): r for r in gm.iter_rows(min_row=5, max_col=7, values_only=True) if r[0]}
    reg = [r for r in gm.iter_rows(min_row=5, max_col=7, values_only=True) if r[2] in ("Calib", "Spec_Rack") and r[4] and ("L1_JudgeTol" in str(r[4]) or "SR_RU_PxMult" in str(r[4]))]
    assert len(reg) == 3 and all(r[5] == "Assumed" for r in reg), reg
    del rows
