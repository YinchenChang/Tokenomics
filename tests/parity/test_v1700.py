"""v17.0 回本判讀主口徑（F1／F2）專屬測試（工作單 docs/workorders/20261020_v17.0_F1F2主口徑.md r0）。讀 LibreOffice 重算後的值；不刪除、不放寬任何既有測試。

W1 Cap_In O 節：折現 L_rel、機隊折現 L（L1_LTotalFleetVDisc 三值 ＝ 事前計算 0.75860／0.47298／1.13827，容差 1e-5；Python 獨立重算 1e-9；t0＝0 時等於 L1_LTotalDisc）。
W2 Cap_In 第 242 列改讀 L1_LTotalFleetVDisc 對 L1_LReqFleet_IdNoK（三態各一情境）；L1 P89；v16.0 口徑對照列；L1_LReqFleet_Id 標籤。
W3 Decisions SL50-01…03、README。
不得變：所有既有名稱逐值不變，唯一例外是版本字串（RNG_Version／RNG_Current）。
"""
import re
from math import exp, log
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import range_boundaries

from conftest import model_path

REPO = Path(__file__).resolve().parents[2]
V160 = REPO / "model" / "archive" / "20261018_Tokenomics_v16.0.xlsx"
V170 = REPO / "model" / "archive" / "20261020_Tokenomics_v17.0.xlsx"      # v17.1：現行版已是 v17.1；v16.0→v17.0 的逐格比較改對封存的 v17.0 釘住（v17.0→v17.1 由 test_v1710 釘住）
NEW_NAMES = {"L1_LTotalFleetVDisc", "L1_LTotalFleetVDisc_Lo", "L1_LTotalFleetVDisc_Hi"}
VERSION_NAMES = {"RNG_Version", "RNG_Current"}
V180_NEW = {f"L1_{k}{s}" for k in ["M_2025", "M_2026Q1", "M_2026Q2", "M_FY2026", "M", "M_API", "M_Sub", "M_Seat", "PxEff_2026", "PxList_2026", "TokGW_PaidCap_VR200", "TokGW_PaidPQ_VR200", "RevGW_PQ_VR200", "RevGW_PQ_VR200_FY", "RevGW_PQ_VR200_Life", "TokLedger", "TokLedger_2026Q2"] for s in ("", "_Lo", "_Hi")}      # v18.0：L1 新增 17 列 × 3 名稱（test_v1800 看守）
from test_v1900 import NEW as V190_NEW
from test_v1910 import NEW as V191_NEW      # v19.1：SRC_MOD_106 與 RNG_CHK_Verify* 4 個區間名稱（test_v1910 看守）      # v19.0：Interface N 節新增 25 個 IF_ 名稱（test_v1900 看守）
PRE = {"L1_LTotalFleetVDisc": 0.75860, "L1_LTotalFleetVDisc_Lo": 0.47853,  # v19.1：0.47298 → 0.47853（Arch!C17 改值）
    "L1_LTotalFleetVDisc_Hi": 1.13827}      # 工作單 W1 的 chat 端事前計算（5 位）


def _flat(v):
    if isinstance(v, (list, tuple)):
        return [x for r in v for x in _flat(r)]
    return [v]


def _same(a, b):
    a = None if a == "" else a
    b = None if b == "" else b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= 1e-12 or abs(a - b) <= 1e-9 * max(abs(a), abs(b))
    return a == b


@pytest.fixture(scope="module")
def wbf():
    return openpyxl.load_workbook(model_path())


def _row(ws, rid, col=1):
    rows = [r for r in range(5, ws.max_row + 1) if ws.cell(r, col).value == rid]
    assert len(rows) == 1, (rid, rows)
    return rows[0]


def _lab(ws, prefix):
    rows = [r for r in range(290, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith(prefix)]
    assert len(rows) == 1, (prefix, rows)
    return rows[0]


def _seg(a, x0, x1):
    return (x1 - x0) if abs(a) < 1e-12 else (exp(a * x1) - exp(a * x0)) / a


def _lrel_disc(g, t1, rho, T, t0, gpre, wacc):
    """獨立重算：折現到該世代起點的 L_rel（工作單 W1 公式）。"""
    d, kp, k = log(1 + wacc), log(1 + gpre), log(1 + g)
    past = exp(-kp * t0) * _seg(kp - d, 0, t0) if t0 > 0 else 0.0
    mid = exp(-k * t0) * _seg(k - d, t0, t1)
    late = rho * exp(-k * t0) * _seg(k - d, t1, T)
    return (past + mid + late) / _seg(-d, 0, T)


def test_v1700_new_values_match_precomputed(lo):
    v = lo.view({})
    for n, want in PRE.items():
        assert abs(v.get_name(n) - want) < 1e-5, (n, v.get_name(n), want)


def test_v1700_independent_python_recomputation(wbf, lo):
    """用 Python 獨立重算機隊折現 L，與 LibreOffice 的值相同（1e-9）。"""
    v = lo.view({})
    ci = wbf["Cap_In"]
    g = {c: v.get("Cap_In", f"{c}226") for c in "CDE"}
    t1 = {c: v.get("Cap_In", f"{c}227") for c in "CDE"}
    rho = {c: v.get("Cap_In", f"{c}228") for c in "CDE"}
    T = {c: v.get("Cap_In", f"{c}229") for c in "CDE"}
    r_sv, r_lo, r_hi, r_a, r_gp = (_lab(ci, p) for p in ("世代服務起點 s_v（年", "世代服務起點 s_v 低", "世代服務起點 s_v 高", "校準錨點年 A", "錨點前中階牌價年變化率 g_pre"))
    A = v.get("Cap_In", f"C{r_a}")
    gp = {c: v.get("Cap_In", f"{c}{r_gp}") for c in "CDE"}
    sv = {k: [v.get("Cap_In", f"{c}{r}") for c in "CDEF"] for k, r in (("b", r_sv), ("lo", r_lo), ("hi", r_hi))}
    share = _flat(v.get_name("AL_ShareGen")); rev = _flat(v.get_name("AL_RevGW25"))
    w = [s * r if s > 0 else 0.0 for s, r in zip(share, rev)]
    wacc = {"C": v.get("Inputs", "E28"), "D": v.get("Inputs", "D28"), "E": v.get("Inputs", "F28")}
    comb = {"C": ("b", "C", gp["C"]), "D": ("hi", "D", gp["E"]), "E": ("lo", "E", gp["D"])}      # 同 v16.0：低端組合 s_v 取較晚、g_pre 取 E 欄；高端組合 s_v 取較早、g_pre 取 D 欄
    got = {}
    for col, (k, sc, g_pre) in comb.items():
        ls = [_lrel_disc(g[sc], t1[sc], rho[sc], T[sc], max(0.0, A - s), g_pre, wacc[sc]) for s in sv[k]]
        got[col] = sum(a * b for a, b in zip(w, ls)) / sum(w)
    assert abs(v.get_name("L1_LTotalFleetVDisc") - got["C"]) < 1e-9
    assert abs(v.get_name("L1_LTotalFleetVDisc_Lo") - got["D"]) < 1e-9 and abs(v.get_name("L1_LTotalFleetVDisc_Hi") - got["E"]) < 1e-9


def test_v1700_t0_zero_equals_l1_ltotaldisc(wbf, lo):
    """t0＝0 時（所有世代起點不早於錨點）折現 L_rel 與機隊折現 L 等於 L1_LTotalDisc（三個端點）。"""
    ci = wbf["Cap_In"]
    rows = [_lab(ci, p) for p in ("世代服務起點 s_v（年", "世代服務起點 s_v 低", "世代服務起點 s_v 高")]
    inputs = {f"Cap_In!{c}{r}": 2025.5 for r in rows for c in "CDEF"}
    z = lo.view(inputs)
    for n, ref in (("L1_LTotalFleetVDisc", "L1_LTotalDisc"), ("L1_LTotalFleetVDisc_Lo", "L1_LTotalDisc_Lo"), ("L1_LTotalFleetVDisc_Hi", "L1_LTotalDisc_Hi")):
        assert abs(z.get_name(n) - z.get_name(ref)) < 1e-12, (n, z.get_name(n), z.get_name(ref))
    r_l = _lab(ci, "折現 L_rel（相對錨點、折現到該世代起點；基準")
    assert all(abs(z.get("Cap_In", f"{c}{r_l}") - z.get_name("L1_LTotalDisc")) < 1e-12 for c in "CDEF")
    # 基準：GB300、VR200（t0＝0）的折現 L_rel 等於 L1_LTotalDisc；Hopper、GB200（t0＞0）不等
    v = lo.view({})
    assert abs(v.get("Cap_In", f"E{r_l}") - v.get_name("L1_LTotalDisc")) < 1e-12 and abs(v.get("Cap_In", f"F{r_l}") - v.get_name("L1_LTotalDisc")) < 1e-12
    assert abs(v.get("Cap_In", f"C{r_l}") - v.get_name("L1_LTotalDisc")) > 1e-3


def test_v1700_three_state_row242(wbf, lo):
    v = lo.view({})
    assert v.get("Cap_In", "C242") == "未定（區間重疊；基準 L 0.759 對所需 L 0.807）"
    assert v.get("Cap_In", "D242").startswith("對照：低端欄 L ") and v.get("Cap_In", "E242").startswith("對照：高端欄 L ")
    assert "（全部算力，低端）" in v.get("Cap_In", "D242") and "（全部算力，高端）" in v.get("Cap_In", "E242")
    neg = lo.view({"Cap_In!C226": -0.9, "Cap_In!D226": -0.9, "Cap_In!E226": -0.85})
    assert neg.get("Cap_In", "C242") == "未回本（L 區間全在所需 L 之下）"
    pos = lo.view({"Cap_In!C226": 0.3, "Cap_In!D226": 0.3, "Cap_In!E226": 0.4, "Cap_In!C227": 6, "Cap_In!D227": 6, "Cap_In!E227": 6})
    assert pos.get("Cap_In", "C242") == "回本（L 區間全在所需 L 之上）"
    # L1 P89 同步
    p89 = v.get("L1", "P89")
    assert p89.startswith("恆等式讀法：未定（世代起點折現 L 0.76 對全部算力所需 L 0.81；v15.3 口徑 2025＝第 0 年：L 0.660）（v17.0 口徑對照：只含推論、經計價倍數的所需 L 0.76）（v16.0 口徑：L 0.716 對 0.756）；字面讀法 ")
    assert neg.get("L1", "P89").startswith("恆等式讀法：未回本") and pos.get("L1", "P89").startswith("恆等式讀法：回本")
    assert "（v16.0 口徑：L " in neg.get("L1", "P89") and "（v16.0 口徑：L " in pos.get("L1", "P89")
    # 基準比值 1.003 寫在 G242
    assert "基準比值（L ÷ 所需 L）＝1.003" in wbf["Cap_In"]["G242"].value
    assert abs(v.get_name("L1_LTotalFleetVDisc") / v.get_name("L1_LReqFleet_IdNoK") - 1.003) < 5e-4


def test_v1700_v160_comparison_row(wbf, lo):
    v = lo.view({})
    ci = wbf["Cap_In"]
    r = _lab(ci, "對照：v16.0 口徑（不折現、含 κ）")
    assert v.get("Cap_In", f"C{r}") == "未定（區間重疊；基準 世代起點 L（v16.0 口徑） 0.716 對所需 L 0.756）"
    for p in ("對照：2025＝第 0 年（v15.3 口徑）", "對照：世代起點 L 對 不含 κ 所需 L", "對照：折現版 L 對 不含 κ 所需 L"):      # v16.0 既有三個對照列保留
        assert _lab(ci, p) < r
    assert r > _lab(ci, "機隊折現 L（世代起點＋折現")


def test_v1700_t0_above_t1_unsupported(wbf, lo):
    """t0 > T1 的世代（目前沒有）：折現 L_rel 輸出文字，機隊折現 L 與判讀回報『未支援／非數字』，不出錯誤值（待 Project 判斷）。"""
    ci = wbf["Cap_In"]
    r_sv = _lab(ci, "世代服務起點 s_v（年")
    r_l = _lab(ci, "折現 L_rel（相對錨點、折現到該世代起點；基準")
    s = lo.view({f"Cap_In!C{r_sv}": 2020.0})
    assert s.get("Cap_In", f"C{r_l}") == "t0 > T1（已降級，未支援）"
    assert s.get_name("L1_LTotalFleetVDisc") == "含 t0 > T1 的世代（已降級，未支援）"
    assert s.get("Cap_In", "C242") == "所需 L 或 L 非數字（本情境）"
    assert s.get("L1", "P89") == "推算值非數字（本情境）"
    assert isinstance(s.get_name("L1_LTotalFleetVDisc_Lo"), float)


def test_v1700_l1_row_label_and_names(wbf):
    l1 = wbf["L1"]
    r = _row(l1, "L1_LTotalFleetVDisc")
    assert l1[f"D{r}"].value == f"=Cap_In!$C${_lab(wbf['Cap_In'], '機隊折現 L（世代起點＋折現')}"
    assert wbf.defined_names["L1_LTotalFleetVDisc"].attr_text == f"L1!$D${r}" and wbf.defined_names["L1_LTotalFleetVDisc_Lo"].attr_text == f"L1!$E${r}" and wbf.defined_names["L1_LTotalFleetVDisc_Hi"].attr_text == f"L1!$F${r}"
    assert l1[f"B{_row(l1, 'L1_LReqFleet_Id')}"].value.split("｜v20.0")[0].endswith("｜v17.0 起非主口徑（F1）")
    assert "非主口徑" not in str(l1[f"B{_row(l1, 'L1_LReqFleet_IdNoK')}"].value).split("｜v20.0")[0]      # 其他既有標籤不動（v17.1：IdNoK 標籤改為『v17.0 起主口徑』，仍不含『非主口徑』）


def test_v1700_decisions_readme(wbf):
    d = wbf["Decisions"]
    for i in range(1, 4):
        r = _row(d, f"SL50-0{i}")
        assert d.cell(r, 2).value == "v17.0" and d.cell(r, 7).value == "v17.0 已建" and d.cell(r, 10).value == "否" and "可撤回" in d.cell(r, 6).value
    assert "沒回應的話，下一版照預設改" in d.cell(_row(d, "SL50-01"), 6).value and "2026-10-18 10:44" in d.cell(_row(d, "SL50-02"), 6).value
    assert "L1_LReqFleet_IdNoK" in d.cell(_row(d, "SL50-01"), 4).value and "L1_LTotalFleetVDisc" in d.cell(_row(d, "SL50-02"), 4).value
    assert d.max_row < 1000      # Gov_Map AA 的 COUNTIF 範圍上限（$A$5:$A$1000）
    rd = wbf["README"]
    assert "20261023_Tokenomics_v21.0（主版：" in str(rd["B5"].value) and "v20.1（次版，底版 v20.0：" in str(rd["B5"].value) and "v20.0（主版：" in str(rd["B5"].value) and "v19.1（次版，底版 v19.0：" in str(rd["B5"].value) and "v19.0（主版：" in str(rd["B5"].value) and "v18.0（主版：" in str(rd["B5"].value) and "v17.2（修補版：" in str(rd["B5"].value) and "v17.1（修補版，工程類：" in str(rd["B5"].value) and "v17.0（主版：" in str(rd["B5"].value) and str(rd["A1"].value).startswith("Tokenomics v21.")      # v19.0：版本字串隨版本；v18.0：版本字串隨版本      # v17.2：版本字串隨版本      # v17.1：版本字串隨版本，v17.0 說明保留在後
    rows = [r for r in range(4, rd.max_row + 1) if rd.cell(r, 1).value == "v17.0 回本判讀主口徑"]
    assert len(rows) == 1 and "世代起點、折現的機隊 L" in rd.cell(rows[0], 2).value and "不含 κ" in rd.cell(rows[0], 2).value


def test_v1700_only_new_names_and_nothing_else_changed(lo):
    """與封存的 v16.0 逐名稱比較：新增名稱恰為 3 個；既有名稱值全部相同，唯一例外是版本字串。"""
    pw = openpyxl.load_workbook(V160, data_only=True)
    pf = openpyxl.load_workbook(V160)
    v = lo.view({})
    V172 = __import__("json").load(open(Path(__file__).resolve().parent / "fixtures" / "v172_changed_names.json", encoding="utf-8"))      # v17.2：W1 第二來源 B1 的預期變動（新增 SRC_ 名稱；GOV 計數、IFC_Updated、IDX_ 範圍），test_v1720 看守
    assert set(v.names) - set(pf.defined_names) - set(V172["new"]) - V180_NEW - V190_NEW - V191_NEW - (set(v.names) - set(openpyxl.load_workbook(Path(__file__).resolve().parents[2] / 'model' / 'archive' / '20261022_Tokenomics_v19.1.xlsx').defined_names)) == NEW_NAMES
    assert set(pf.defined_names) - set(v.names) == set()
    changed = set()
    for n, d in pw.defined_names.items():
        if "!" not in d.attr_text or "," in d.attr_text:
            continue
        sh, rg = d.attr_text.rsplit("!", 1)
        c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
        want = [pw[sh.strip("'")].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        got = _flat(v.get_name(n))
        if len(want) != len(got) or not all(_same(a, b) for a, b in zip(want, got)):
            changed.add(n)
    changed -= set(__import__("json").load(open(Path(__file__).resolve().parent / "fixtures" / "v191_changed_names.json", encoding="utf-8"))["changed"]) - VERSION_NAMES      # v19.1：Arch!C17 下游的預期變動名稱
    changed -= set(V172["changed"]) - VERSION_NAMES      # v17.2：同上
    changed -= {"IFC_Conf", "IFC_Use", "IFC_Layer", "IFC_Updated"}      # v19.0：Interface 契約欄範圍隨 N 節新列延伸（既有列的值不變，由 test_v1900 看守）
    assert changed == VERSION_NAMES, sorted(changed ^ VERSION_NAMES)
    for n, want in (("L1_LTotal", 0.659505673291395), ("L1_LReqFleet_Id", 0.756109635897205), ("L1_LTotalFleetV", 0.715894950036491),
                    ("L1_LReqFleet_IdNoK", 0.756109635897205), ("L1_LTotalDisc", 0.698844526610524), ("L1_MFleet", 0.447574227389844), ("IF_AllocServeGW", 0.89809617101894)):
        assert abs(v.get_name(n) - want) < 1e-9, n
    assert (v.get_name("GOV_Errors"), v.get_name("GOV_Warnings"), v.get_name("GOV_Info")) == (0, 135, 182)      # v21.1: 133 → 135 (SRC_DEM_089/090 lack a second source, W1 +2)
    #      # v19.1：133 → 132（SRC_MOD_018 填 C1 免除標記，Z22 缺第二來源 1 → 0）； v17.2：W1 278 → 131（W1 第二來源 B1：A 86、C1 53、Q11 8 離開；預期變動）
    assert v.get_name("RNG_StaleFlag") == "一致" and v.get_name("RNG_OutCount") == 0


def test_v1700_changed_cells_only_expected():
    """與封存的 v16.0 逐格比對（公式文字或輸入值）：只有預期的頁與格不同。"""
    a, b = openpyxl.load_workbook(V160), openpyxl.load_workbook(V170)      # v17.1：右側改為封存的 v17.0
    diffs = {}
    for ws in a:
        wn = b[ws.title]
        coords = {c.coordinate for r in ws.iter_rows() for c in r if c.value is not None} | {c.coordinate for r in wn.iter_rows() for c in r if c.value is not None}
        d = sorted(x for x in coords if ws[x].value != wn[x].value)
        if d:
            diffs[ws.title] = d
    assert set(diffs) <= {"README", "Cap_In", "L1", "Alloc", "Load_Bearing", "Decisions", "Gov_Map", "SRC_Index"}, set(diffs)
    rows = lambda cells: sorted({int(re.sub("[A-Z]+", "", x)) for x in cells})
    assert sorted(diffs["README"]) == ["A1", "A62", "B5", "B62"]
    cap = diffs["Cap_In"]
    assert {x for x in cap if int(re.sub("[A-Z]+", "", x)) <= 318} == {"A242", "C242", "D242", "E242", "G242"}      # 既有格：只有第 242 列
    assert min(r for r in rows(cap) if r > 242) == 320      # 新節接在第 318 列之後
    l1 = diffs["L1"]
    assert sorted(set(re.sub("[0-9]+", "", x) for x in l1 if re.sub("[A-Z]+", "", x) == "107")) == ["B"]      # L1_LReqFleet_Id 只有標籤（B 欄）
    assert rows(l1) == [89, 107, 116] and {x for x in l1 if re.sub("[A-Z]+", "", x) == "89"} == {"P89"}
    assert sorted(diffs["Alloc"]) == ["C177"]      # RNG_Version（常數）
    assert len(diffs["Decisions"]) == 30
    assert {re.match("[A-Z]+", x).group() for x in diffs.get("Load_Bearing", [])} <= {"J"}
