"""列出指定工作表列範圍的公式與值（唯讀）。
用法：python3 dump_range.py <xlsx> <sheet> <起列> <迄列> [最大欄數]
或：python3 dump_range.py <xlsx> --name <具名範圍>  （列出具名範圍位置）"""
import sys
import openpyxl
from openpyxl.utils import get_column_letter

path = sys.argv[1]
wbf = openpyxl.load_workbook(path, data_only=False)
if sys.argv[2] == "--name":
    for nm in sys.argv[3:]:
        dn = wbf.defined_names.get(nm)
        print(nm, "=>", dn.attr_text if dn is not None else None)
    sys.exit()
wbv = openpyxl.load_workbook(path, data_only=True)
sh = sys.argv[2]
r0, r1 = int(sys.argv[3]), int(sys.argv[4])
maxc = int(sys.argv[5]) if len(sys.argv) > 5 else 16
wf, wv = wbf[sh], wbv[sh]
for r in range(r0, r1 + 1):
    out = []
    for c in range(1, maxc + 1):
        f = wf.cell(r, c).value
        if f is None:
            continue
        v = wv.cell(r, c).value
        col = get_column_letter(c)
        if isinstance(f, str) and f.startswith("="):
            vs = f"{v:.4g}" if isinstance(v, float) else v
            out.append(f"{col}: {f[:140]} => {vs}")
        else:
            out.append(f"{col}: {str(f)[:120]}")
    if out:
        print(f"[{r}] " + " | ".join(out))
