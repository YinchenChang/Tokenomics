# v21.6 (work order docs/workorders/20261025_v21.6_NonNV持有比TPU_Trn.md r0, branch claude/v21-6): NonNV holding-cost ratios of TPU v7 (row 6) and Trainium3 (row 7) become derived values.
# MINOR version: two input rows change (6 blue cells, written only while they still hold the v21.5 value, like v2130.inputs_update); no new/removed contract names; contract v0.27 unchanged.
#   W1  NonNV E6:G6 (TPU v7) 0.7／0.8／1.0 -> 0.65／0.70／0.90;  E7:G7 (Trainium3) 0.6／0.7／0.9 -> 0.45／0.55／0.70.  Output ratios B:D are NOT changed (pending PR #172 Q1).
#       NonNV L6／L7 explanation: the conventional-value tail is replaced by the derivation note; the original tail is kept verbatim after "｜原：" (revision 9).
#   W2  DB_Evidence E409 (after E408), Decisions SL63-01, README B5 opening paragraph.
# Contract names whose values change: IF_NonNV_Hold_TPU, IF_NonNV_Hold_Trn (3 cells each) and the version strings. Everything else unchanged.
# Excel-owned cells (DB_Evidence, Decisions) are written only while absent; NonNV L6／L7 are rewritten by the builder on every build, so the text fix is re-applied deterministically. Idempotent.
from openpyxl.utils import get_column_letter as L

from common import put, F_CALC
import v1410
import v2110
import v2150

VERSION = "20261025_Tokenomics_v21.6"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-25"
DASH = "—"
V = "v21.6"
WO = "工作單 docs/workorders/20261025_v21.6_NonNV持有比TPU_Trn.md（r0）"
DEC_ID = "SL63-01"
EVD = "docs/evidence/20261025_NonNV_TPU_Trainium/"
DEC_BY = ("Andy 未回應，照工作單預設（2026-10-25 對話已列為判斷題 Q1、Q7）：PR #172（NonNV B 段，TPU v7、Trainium3 每 kW 重推，獨立核對 r1 後合併）"
          "待判斷 Q7 預設——先只換持有比，產出比待 Q1（精度口徑）決定後再換")

# ---------------------------------------------------------------- W1 inputs (old-value guards)
ROWS = {6: ("Google TPU v7 Ironwood", (0.7, 0.8, 1.0), (0.65, 0.70, 0.90)),
        7: ("AWS Trainium3", (0.6, 0.7, 0.9), (0.45, 0.55, 0.70))}


def inputs_update(wb):
    log = []
    ws = wb["NonNV"]
    for r, (name, old, new) in ROWS.items():
        assert ws[f"A{r}"].value == name, (r, ws[f"A{r}"].value)
        cells = [ws[f"{c}{r}"] for c in "EFG"]
        if all(c.value is not None and type(c.value) is not bool and c.value == o for c, o in zip(cells, old)):
            for c, n in zip(cells, new):
                c.value = n
            log.append(f"NonNV!E{r}:G{r}: {old} -> {new} (v21.6 W1)")
    return log


# ---------------------------------------------------------------- W1 NonNV L6／L7 text
TAIL = v2110.ROWS_5_9_TAIL      # the v21.0 conventional-value tail, kept verbatim after 原：
NEW_NOTE = {
    6: ("持有比 v21.6 依每 kW 設備價推導（PR #172；TPU 每顆約 $21／25／35K、Trn3 每架 $1.0／1.5／2.2M；Assumed）；"
        "產出比仍為慣例值，待精度口徑判斷（PR #172 Q1：各晶片最高原生精度時 TPU 0.25／0.35／0.40、Trn3 0.06／0.14／0.25；FP8 對齊時 TPU 0.48／0.66／0.84、Trn3 0.13／0.28／0.50）"),
}
NEW_NOTE[7] = NEW_NOTE[6]


def text_fixes(wb):
    ws = wb["NonNV"]
    n = 0
    for r in (6, 7):
        cell = ws[f"L{r}"]
        cur = str(cell.value)
        if "v21.6 依每 kW 設備價推導" in cur:
            continue
        assert cur.endswith(TAIL), (r, cur[-80:])
        head = cur[:-len(TAIL)].rstrip("｜")
        cell.value = head + "｜" + NEW_NOTE[r] + "｜原：" + TAIL
        n += 1
    return n


# ---------------------------------------------------------------- W2 DB_Evidence E409
_OLD_VAL = "NonNV 第 6 列 TPU v7 持有比 0.7／0.8／1.0、第 7 列 Trainium3 持有比 0.6／0.7／0.9（慣例值）；產出比 TPU 0.6／0.8／1.1、Trn3 0.3／0.5／0.8（慣例值）"
E409 = [
    "E409", DATE,
    "NonNV B 段：TPU v7、Trainium3 每 kW 重推——持有比 TPU 0.65／0.70／0.90、Trn3 0.45／0.55／0.70（採用，v21.6 W1）；"
    "產出比 TPU 0.25／0.35／0.40、Trn3 0.06／0.14／0.25（各晶片最高原生精度；未採用，待 Q1）；每 token 成本比基準 TPU 約 2.0、Trn3 約 3.9（現值 1.0、1.4）",
    "Google、AWS 自揭規格；KGI、SemiAnalysis 價格（SemiAnalysis 只作佐證）；NonNV B 段查證 " + EVD + "（PR #172）",
    "Interested-party（Google、AWS 自揭規格）＋Analogy（KGI、SemiAnalysis 價格）",
    "NonNV E6:G7（持有比）；Interface 第 375–376 列（IF_NonNV_Hold_TPU／IF_NonNV_Hold_Trn）",
    _OLD_VAL,
    "持有比 TPU 0.65／0.70／0.90、Trn3 0.45／0.55／0.70（採用）；產出比 TPU 0.25／0.35／0.40、Trn3 0.06／0.14／0.25（未採用，待 Q1）",
    "持有比採用（PR #172 Q7 預設，" + V + " W1）；產出比不改，待 Q1（精度口徑）", V,
    "TPU 效率 f 基準 0.856 為 Assumed（InferenceX 單一來源只作佐證）；Trn3 機架倍數 1.35–1.75（基準 1.52）；AWS 1 kW 為 DCD 推論；"
    "TPU 每顆約 $21／25／35K、Trn3 每架 $1.0／1.5／2.2M；產出比的精度口徑（各晶片最高原生精度 vs FP8 對齊）待 Q1｜" + WO,
    "已處理", DASH, DASH, "IF_NonNV_Hold_TPU、IF_NonNV_Hold_Trn（值改變）；其餘契約名稱不變", 3, "利害關係方",
    "採用（持有比；產出比待 Q1）｜" + V + " 新增列"]


def evidence_append(wb):
    ws = wb["DB_Evidence"]
    if any(ws.cell(r, 1).value == E409[0] for r in range(5, ws.max_row + 1)):
        return 0
    r = max(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")) + 1
    assert ws.cell(r - 1, 1).value == "E408", ws.cell(r - 1, 1).value
    for i, v in enumerate(E409):
        put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 4, 5, 6, 7, 8, 10, 14, 17))
    return 1


# ---------------------------------------------------------------- W2 Decisions
def _decisions():
    return [[DEC_ID, V, "v21.6 NonNV 持有比 TPU v7、Trainium3 改為推導值（產出比待 Q1）",
             "W1：NonNV E6:G6（TPU v7 持有比）0.7／0.8／1.0 → 0.65／0.70／0.90；E7:G7（Trainium3 持有比）0.6／0.7／0.9 → 0.45／0.55／0.70（依每 kW 設備價推導，PR #172；Assumed）；"
             "產出比 B:D 不改（TPU 0.6／0.8／1.1、Trn3 0.3／0.5／0.8 慣例值保留；PR #172 Q1：各晶片最高原生精度時 TPU 0.25／0.35／0.40、Trn3 0.06／0.14／0.25，FP8 對齊時 TPU 0.48／0.66／0.84、Trn3 0.13／0.28／0.50，待精度口徑判斷）；"
             "NonNV L6／L7 說明改寫，原文逐字保留於「｜原：」之後（修訂九）。W2：DB_Evidence E409；README B5 加 v21.6 段。"
             "契約名稱的值改變只有 IF_NonNV_Hold_TPU、IF_NonNV_Hold_Trn（各 3 格）與版本字串；契約不升版（v0.27）。Andy 原話與依據：" + DEC_BY,
             DATE, DEC_BY, "v21.6 已建",
             "NonNV E6:G7（及 H6:J7 公式）；NonNV L6／L7；Interface 第 375–376 列；DB_Evidence E409；README B5；Decisions SL63-01",
             f"{WO} W1–W2", "否"]]


def apply(wb, append_text):
    log = {"text_fixes": text_fixes(wb)}
    log["evidence"] = evidence_append(wb)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == DEC_ID for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


# ---------------------------------------------------------------- README
_prev = v2150.README_VERSION.split("_Tokenomics_", 1)[1]
README_VERSION = (VERSION + "（次版，底版 v21.5：" + WO + "——NonNV 持有比 TPU v7 0.7／0.8／1.0 → 0.65／0.70／0.90、Trainium3 0.6／0.7／0.9 → 0.45／0.55／0.70（依每 kW 設備價推導，PR #172；Assumed），產出比不改（待 PR #172 Q1）；"
                  "契約名稱的值改變的只有 IF_NonNV_Hold_TPU、IF_NonNV_Hold_Trn（各 3 格）與版本字串；其餘不變；GOV 三項計數 0／134／182 不變。"
                  "NonNV L6／L7 說明改寫（原文保留於「｜原：」後）；DB_Evidence E409；Decisions SL63-01。）以下為 " + _prev)
README_TITLE = V + v2150.README_TITLE[len(v2150.V):]
