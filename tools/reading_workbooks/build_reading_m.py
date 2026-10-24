"""m 閱讀版建檔程式（v21.2 重做，2026-10-24）。

用法：
    python3 build_reading_m.py <Tokenomics.xlsx> <輸出.xlsx>
輸出檔尚未重算（沒有公式快取值）；請再用 LibreOffice headless 重算存檔
（parity_lib.lo_recalc），然後以 check_reading_m.py 核對。

設計：
- 全部取自 Tokenomics 的數字（輸入、區間、原值）都在建檔時從 Tokenomics 讀取後貼上，
  並登錄在 9_核對對照（本簿格 ↔ Tokenomics 位置），供 check_reading_m.py 逐格核對。
- 計算全部是 Excel 公式，連到 2_輸入（藍字）；改輸入，3、4、5 頁與 1A 一起變。
- 沿用 v17.0 的分頁命名、顏色規則與「白話說明＋公式／Tokenomics 原值／差異」的欄位做法。
"""
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName

if len(sys.argv) != 3:
    sys.exit('用法：python3 build_reading_m.py <Tokenomics.xlsx> <輸出.xlsx>')
TKP, OUTP = sys.argv[1:3]
TKF = openpyxl.load_workbook(TKP)
TKV = openpyxl.load_workbook(TKP, data_only=True)

# ───────────────────────── Tokenomics 讀取 ─────────────────────────
_REF = re.compile(r"([A-Za-z_][A-Za-z0-9_]*)!\$?([A-Z]+)\$?(\d+)")


def tk(expr):
    """Tokenomics 位置（具名範圍、Sheet!Cell，或由 Sheet!Cell 組成的四則運算）→ 數值。"""
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", expr):
        d = TKF.defined_names[expr]
        sh, rng = list(d.destinations)[0]
        v = TKV[sh][rng.replace('$', '')].value
        return v
    if re.fullmatch(_REF.pattern, expr):
        m = _REF.fullmatch(expr)
        return TKV[m[1]][m[2] + m[3]].value
    s = _REF.sub(lambda m: repr(float(TKV[m[1]][m[2] + m[3]].value)), expr)
    s = re.sub(r"(?<![0-9.A-Za-z_])[A-Za-z_][A-Za-z0-9_]*", lambda m: repr(float(tk(m[0]))), s)
    if not re.fullmatch(r"[0-9eE+\-*/(). ]+", s):
        raise ValueError(expr)
    return eval(s)


GOV = {}
for r in TKV['Gov_Map'].iter_rows(min_row=5):
    if r[0].value:
        GOV[(r[2].value, r[3].value)] = r[0].row


def gov_row(sheet, cell):
    from openpyxl.utils import range_boundaries
    for (sh, rng), row in GOV.items():
        if sh != sheet or rng is None:
            continue
        try:
            c1, r1, c2, r2 = range_boundaries(rng)
            c, rr, _, _ = range_boundaries(cell)
        except Exception:
            continue
        if c1 <= c <= c2 and r1 <= rr <= r2:
            return row
    return None


# ───────────────────────── 樣式 ─────────────────────────
FN = 'Arial'
F_BASE = Font(name=FN, size=10)
F_BOLD = Font(name=FN, size=10, bold=True)
F_BLUE = Font(name=FN, size=10, color='0000FF')
F_GREEN = Font(name=FN, size=10, color='008000')
F_GRAY = Font(name=FN, size=9, color='555555')
F_HDR = Font(name=FN, size=10, bold=True, color='FFFFFF')
F_T1 = Font(name=FN, size=13, bold=True)
F_T2 = Font(name=FN, size=9, color='555555')
FILL_HDR = PatternFill('solid', fgColor='305496')
FILL_GRP = PatternFill('solid', fgColor='D9E1F2')
FILL_M = PatternFill('solid', fgColor='E2EFDA')
FILL_Y = PatternFill('solid', fgColor='FFF2CC')
TOP_WRAP = Alignment(vertical='top', wrap_text=True)
TOP = Alignment(vertical='top')

wb = openpyxl.Workbook()
wb.remove(wb.active)
SN = ['0_說明', '1_總覽', '1A_需求公式樹', '2_輸入', '3_m重算', '4_分路線', '5_敏感度', '6_不確定與證據', '9_核對對照']
WS = {n: wb.create_sheet(n) for n in SN}


def q(n):
    return "'" + n + "'"


MAPROWS = []   # (本簿頁, 格, Tokenomics 位置, 類型)


def put(ws, addr, v, font=F_BASE, fill=None, fmt=None, wrap=True, h=None):
    c = ws[addr]
    c.value = v
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(vertical='top', wrap_text=wrap, horizontal=h)
    return c


def paste(ws, addr, expr, kind='原值', fmt=None, fill=None):
    """貼上 Tokenomics 現值（藍字）並登錄核對對照。"""
    v = tk(expr)
    put(ws, addr, v, F_BLUE, fill, fmt, wrap=False)
    MAPROWS.append((ws.title, addr, expr, kind))
    return v


def widths(ws, ws_w):
    for k, v in ws_w.items():
        ws.column_dimensions[k].width = v


def hdr(ws, row, labels, c0=1):
    for i, t in enumerate(labels):
        put(ws, f'{CL(c0 + i)}{row}', t, F_HDR, FILL_HDR)


def grp(ws, row, text, ncol):
    for i in range(1, ncol + 1):
        c = ws[f'{CL(i)}{row}']
        c.fill = FILL_GRP
        c.font = F_BOLD
    put(ws, f'A{row}', text, F_BOLD, FILL_GRP, wrap=False)


def diff(v, o):
    return f'=IF({o}=0,{v}-{o},{v}/{o}-1)'


def defname(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


# ───────────────────────── 2_輸入 ─────────────────────────
# (ID, 單位, 項目, Tokenomics 位置, 標記, 區間低, 區間高, 不確定度, 黃底, 說明)
# 區間低／高：Tokenomics 位置字串，或 None
def GL(sheet, cell):
    r = gov_row(sheet, cell)
    return (f'Gov_Map!K{r}', f'Gov_Map!L{r}') if r else (None, None)


IN_GROUPS = []


def G(title, rows):
    IN_GROUPS.append((title, rows))


G('A. 營收與廣告（OpenAI；m 的分子）', [
    ('IN_Rev', '$B', '2025 全年營收（毛額）', 'SRC_DEM_007', 'Interested-party（SRC_DEM_007）', None, None, '低', 0,
     'Cap_In!C69；SRC_DEM_007，2025 查核財報營收 130.7 億美元（外流、FT 核實）。2025 無廣告，分子＝營收。'),
    ('IN_RevQ1', '$B', '2026 Q1 季營收', 'SRC_DEM_069', 'Interested-party（SRC_DEM_069）', None, None, '中', 0,
     'Cap_In!D185 = 本格 × 4（年化）。The Information 2026-05-22（二手）。'),
    ('IN_RevQ2', '$B', '2026 Q2 季營收', 'SRC_DEM_042', 'Interested-party（SRC_DEM_042）', None, None, '中', 0,
     'Cap_In!E185 = 本格 × 4（年化）。WSJ 2026-08-18（經 CryptoBriefing 轉述）。主口徑 L1_M 用本格。'),
    ('IN_RevFY', '$B', 'FY2026 全年營收（內部預測）', 'Cap_In!F185', 'Assumed（Gov_Map GM624，單值）', None, None, '中', 0,
     '藍字常數 36，內部預測，不登 SRC；Tokenomics 無區間。'),
    ('IN_Ad_Q1', '$B', '2026 Q1 廣告營收（年化；從分子扣除）', 'Cap_In!D186', 'Assumed（Gov_Map GM640）', 'GL:Cap_In:D186', None, '低', 0,
     'Andy 定義：m 只看收費業務；廣告不進分子（Tokenomics J23）。2025 為 0（常數，非輸入）。'),
    ('IN_Ad_Q2', '$B', '2026 Q2 廣告營收（年化；從分子扣除）', 'Cap_In!E186', 'Assumed（Gov_Map GM641）', 'GL:Cap_In:E186', None, '低', 0,
     '依據 SRC_DEM_072（2026-08 廣告年化達 $1B）；Q2 取 0.5。'),
    ('IN_Ad_FY', '$B', 'FY2026 廣告營收（年化；從分子扣除）', 'Cap_In!F186', 'Assumed（Gov_Map GM642）', 'GL:Cap_In:F186', None, '低', 0,
     'FY2026 取 1.0。'),
])
G('B. 營收份額（只用在 4_分路線；不影響整體 m）', [
    ('IN_SubShare', '%', '2025 營收中消費訂閱占比', 'Cap_In!C70', 'Derived（N10；SRC_DEM_039 FT 約 70%）', 'Cap_In!D70', 'Cap_In!E70', '中', 0,
     '低端＝SRC_DEM_041（8）÷ SRC_DEM_007＝0.612；高端 0.82＝A 段 N10 區間上緣（Tokenomics 未追溯到單一 SRC）。'),
    ('IN_SeatShare', '%', '2025 營收中商用席次占比', 'Cap_In!C131', 'Assumed（Gov_Map GM614）', 'GL:Cap_In:C131', None, '中', 0,
     '基準 0.10、區間 0.05–0.15（J18）；API 營收＝1 − 消費訂閱 − 商用席次（餘額，SRC_DEM_071 Sacra 15–20%）。'),
    ('IN_SubShare_Q1', '%', '2026 Q1 營收份額：消費訂閱', 'Cap_In!D188', 'Assumed（Gov_Map GM625，單值）', None, None, '高', 1,
     '2026-03 OpenAI：消費約 60%、企業 >40%（SRC_DEM_023）；Tokenomics 無區間。'),
    ('IN_SubShare_Q2', '%', '2026 Q2 營收份額：消費訂閱', 'Cap_In!E188', 'Assumed（Gov_Map GM625，單值）', None, None, '高', 1,
     'Tokenomics 無區間。'),
    ('IN_SubShare_FY', '%', 'FY2026 營收份額：消費訂閱', 'Cap_In!F188', 'Assumed（Gov_Map GM625，單值）', None, None, '高', 1,
     '企業 2026-07 超過消費、7–9 月翻倍（N05）。Tokenomics 無區間。'),
    ('IN_SeatShare_Q1', '%', '2026 Q1 營收份額：商用席次', 'Cap_In!D189', 'Assumed（Gov_Map GM626，單值）', None, None, '高', 1,
     '企業增量歸席次；API 份額＝餘額（各欄 0.20）。Tokenomics 無區間。'),
    ('IN_SeatShare_Q2', '%', '2026 Q2 營收份額：商用席次', 'Cap_In!E189', 'Assumed（Gov_Map GM626，單值）', None, None, '高', 1,
     'Tokenomics 無區間。'),
    ('IN_SeatShare_FY', '%', 'FY2026 營收份額：商用席次', 'Cap_In!F189', 'Assumed（Gov_Map GM626，單值）', None, None, '高', 1,
     'Tokenomics 無區間。'),
])
G('C. 付費 token（m 的分母之一；免費 token 不進 m，只在算付費占比時扣除）', [
    ('IN_APItpm', '十億 tok/分鐘', 'API 處理量（2025-10 時點）', 'SRC_DEM_010', 'Interested-party（SRC_DEM_010）', None, None, '低', 0,
     'Cap_In 以 Alloc 第 5 列使用：OpenAI 2025-10-06 公告 60 億 token／分鐘。'),
    ('IN_APIratio', 'x', '2025 API 全年平均 ÷ 10 月時點值', 'Alloc_In!C14', 'Assumed（Gov_Map GM584）', 'GL:Alloc_In:C14', None, '中', 0,
     '把 10 月時點值換成全年平均（成長中，所以＜1）。'),
    ('IN_APItpm26', '十億 tok/分鐘', 'API 處理量（2026-03 時點；2026 Q1／Q2 共用）', 'SRC_DEM_011', 'Interested-party（SRC_DEM_011）', None, None, '中', 0,
     'OpenAI 融資公告（2026-03-31）「超過」15 B tok/min；Q2 當期值已搜尋，未找到（DB_Evidence E399）。'),
    ('IN_APImultFY', 'x', 'FY2026 API token ÷ 2026 Q2 時點值（年均倍數）', 'Cap_In!F195', 'Assumed（Gov_Map GM627，單值）', None, None, '中', 0,
     'FY2026 API token＝Q2 時點值年化 × 1.3；Tokenomics 無區間。'),
    ('IN_Prompts', 'B 則/日', '2025 ChatGPT 每日訊息數', 'SRC_DEM_012', 'Interested-party（SRC_DEM_012）', None, None, '低', 0,
     'OpenAI 告知 Axios（2025-07-21）：25 億則／日。'),
    ('IN_Prompts_Q1', 'B 則/日', '2026 Q1 ChatGPT 每日訊息數', 'Cap_In!D197', 'Assumed（Gov_Map GM628，單值）', None, None, '高', 1,
     '由週活比例放大推得；Tokenomics 無區間。'),
    ('IN_Prompts_Q2', 'B 則/日', '2026 Q2 ChatGPT 每日訊息數', 'Cap_In!E197', 'Assumed（Gov_Map GM628，單值）', None, None, '高', 1,
     '2.5 × 1.0 B ÷ 0.7 B 週活 ≈ 3.6；Tokenomics 無區間。主口徑 L1_M 用本格。'),
    ('IN_Prompts_FY', 'B 則/日', 'FY2026 ChatGPT 每日訊息數', 'Cap_In!F197', 'Assumed（Gov_Map GM628，單值）', None, None, '高', 1,
     'Tokenomics 無區間。'),
    ('IN_Tau', 'tok/則', '每則訊息 token 數 τ', 'Alloc_In!C15', 'Assumed（Gov_Map GM586）', 'GL:Alloc_In:C15', None, '高', 1,
     '含輸入、上下文、推理與輸出；無公開數字。2025、2026 共用。'),
    ('IN_FreeShare', '%', 'ChatGPT token 中免費用戶占比', 'Alloc_In!C16', 'Assumed（Gov_Map GM588）', 'GL:Alloc_In:C16', None, '高', 1,
     'token 口徑（不是用戶數口徑，也不是支出口徑）。2025、2026 共用。'),
    ('IN_SubTokShare', '%', '2025 消費訂閱占 ChatGPT 付費 token 份額', 'Cap_In!C132', 'Assumed（Gov_Map GM616）', 'GL:Cap_In:C132', None, '高', 1,
     '只用在 4_分路線；商用席次 token＝餘額。'),
    ('IN_SubTokShare26', '%', '2026 消費訂閱占 ChatGPT 付費 token 份額', 'Cap_In!D200', 'Assumed（Gov_Map GM643，單值）', 'GL:Cap_In:D200', None, '高', 1,
     '2026 三欄共用 0.65（Q2、FY 連結 Q1）；2026 企業席次增加後消費份額應低於 2025。'),
])
G('D. 牌價與有效單價（m 的分母之二）', [
    ('IN_PIn_L', '$/M', 'Luna 新鮮輸入 表列牌價', 'Cap_In!C6', 'Verified（SRC_PRC_003）', None, None, '低', 0, 'GPT-6 Luna；Standard、短上下文，2026-09-25。'),
    ('IN_PIn_S', '$/M', 'Sol 新鮮輸入 表列牌價', 'Cap_In!D6', 'Verified（SRC_PRC_006）', None, None, '低', 0, 'GPT-6 Sol。'),
    ('IN_PIn_A', '$/M', 'Astra 新鮮輸入 表列牌價', 'Cap_In!E6', 'Verified（SRC_PRC_009）', None, None, '低', 0, 'GPT-6 Astra。'),
    ('IN_PCache_L', '$/M', 'Luna 快取輸入 表列牌價', 'Cap_In!C7', 'Verified（SRC_PRC_004）', None, None, '低', 0, '＝輸入價 10%。'),
    ('IN_PCache_S', '$/M', 'Sol 快取輸入 表列牌價', 'Cap_In!D7', 'Verified（SRC_PRC_007）', None, None, '低', 0, ''),
    ('IN_PCache_A', '$/M', 'Astra 快取輸入 表列牌價', 'Cap_In!E7', 'Verified（SRC_PRC_010）', None, None, '低', 0, ''),
    ('IN_POut_L', '$/M', 'Luna 輸出（含思考）表列牌價', 'Cap_In!C8', 'Verified（SRC_PRC_005）', None, None, '低', 0, '思考 token 依輸出價計費。'),
    ('IN_POut_S', '$/M', 'Sol 輸出（含思考）表列牌價', 'Cap_In!D8', 'Verified（SRC_PRC_008）', None, None, '低', 0, ''),
    ('IN_POut_A', '$/M', 'Astra 輸出（含思考）表列牌價', 'Cap_In!E8', 'Verified（SRC_PRC_011）', None, None, '低', 0, ''),
    ('IN_ISL_L', 'tok', 'Luna 參考請求輸入長度 ISL', 'Serving!C23', 'Assumed（Gov_Map GM245）', 'GL:Serving:C23', None, '中', 0, '決定「新鮮輸入／快取輸入／輸出」三種價在混合價裡的權重。'),
    ('IN_ISL_S', 'tok', 'Sol 參考請求輸入長度 ISL', 'Serving!D23', 'Assumed（Gov_Map GM246）', 'GL:Serving:D23', None, '高', 1, ''),
    ('IN_ISL_A', 'tok', 'Astra 參考請求輸入長度 ISL', 'Serving!E23', 'Assumed（Gov_Map GM247）', 'GL:Serving:E23', None, '高', 1, ''),
    ('IN_OSL_L', 'tok', 'Luna 參考請求輸出長度 OSL（含思考）', 'Serving!C24', 'Assumed（Gov_Map GM248）', 'GL:Serving:C24', None, '中', 0, ''),
    ('IN_OSL_S', 'tok', 'Sol 參考請求輸出長度 OSL（含思考）', 'Serving!D24', 'Assumed（Gov_Map GM249）', 'GL:Serving:D24', None, '高', 1, ''),
    ('IN_OSL_A', 'tok', 'Astra 參考請求輸出長度 OSL（含思考）', 'Serving!E24', 'Assumed（Gov_Map GM250）', 'GL:Serving:E24', None, '高', 1, ''),
    ('IN_Chi', '%', '參考請求快取命中率 χ（輸入中屬快取命中的比例）', 'Cap_In!C13', 'Assumed（Gov_Map GM443）', 'GL:Cap_In:C13', None, '高', 1,
     'DeepSeek 2025-03 為 56.3%；2026 agent 流量約 88%。成本與營收用同一 χ。'),
    ('IN_CapMult_L', 'x', 'Luna 能力單價倍數', 'Price_Frontier!C38', 'Decision（ε＝0，Gov_Map GM445）', None, None, '低', 0,
     '＝（H_CAP × 有效訓練算力倍數）^ε；基準 ε＝0，倍數恆為 1。本簿只列原值，不展開。'),
    ('IN_CapMult_S', 'x', 'Sol 能力單價倍數', 'Price_Frontier!D38', 'Decision（ε＝0，Gov_Map GM445）', None, None, '低', 0, '同上。'),
    ('IN_CapMult_A', 'x', 'Astra 能力單價倍數', 'Price_Frontier!E38', 'Decision（ε＝0，Gov_Map GM445）', None, None, '低', 0, '同上。'),
    ('IN_Disc', '%', '有效折扣 d（預設折扣；Batch／Flex 半價＋企業議價，加權）', 'Cap_In!C12', 'Assumed（Gov_Map GM441）', 'GL:Cap_In:C12', None, '高', 1,
     '【可調輸入，不得寫死】有效牌價＝表列牌價 ×（1 − d）。Batch 半價為牌價事實，占比與企業折扣找不到（DB_Evidence E359）。'),
    ('IN_Kappa', 'x', '請求形狀修正 κ', 'Alloc_In!C26', 'Derived（Gov_Map GM644）', 'GL:Alloc_In:C26', None, '高', 1,
     '實際流量每 GPU 秒總 token ÷ 參考請求；乘入混合有效單價（Tokenomics v9.0 M3）。'),
    ('IN_MixP_L', '%', '付費 token 層級組合：Luna', 'Cap_In!C31', 'Assumed（Gov_Map GM457，群組）', None, None, '中', 0,
     'OpenAI 模型 v0.5 tokenMix；三者合計須為 1。'),
    ('IN_MixP_S', '%', '付費 token 層級組合：Sol', 'Cap_In!D31', 'Assumed（Gov_Map GM457，群組）', None, None, '中', 0, ''),
    ('IN_MixP_A', '%', '付費 token 層級組合：Astra', 'Cap_In!E31', 'Assumed（Gov_Map GM457，群組）', 0.10, 0.30, '高', 1,
     '區間 0.10–0.30 只寫在 Gov_Map GM457 理由欄文字（v15.2 B6：Luna 與 Sol 依 0.4:0.45 補足），不是儲存格值；本簿手填。'),
    ('IN_Px_L', '$/M', '2025 年度層級單價：Luna（表列，未折扣）', 'Cap_In!C278', 'Derived（Gov_Map GM665）', 'Cap_In!D278', 'Cap_In!E278', '低', 0,
     '參考請求混合單價，在售天數加權；只用在 2025 欄。'),
    ('IN_Px_S', '$/M', '2025 年度層級單價：Sol（表列，未折扣）', 'Cap_In!C279', 'Derived（Gov_Map GM667）', 'Cap_In!D279', 'Cap_In!E279', '中', 0, ''),
    ('IN_Px_A', '$/M', '2025 年度層級單價：Astra（表列，未折扣）', 'Cap_In!C280', 'Derived（Gov_Map GM669）', 'Cap_In!D280', 'Cap_In!E280', '高', 1,
     '用量加權（v15.2 B2），各模型 token 占比未公開。'),
])
G('E. 本簿試算用（不是 Tokenomics 的值，不影響 3、4、5 頁）', [
    ('IN_DTest', '%', '試算：把 d 改成多少看 m（1A 第 L4 層使用）', None, '本簿試算', None, None, '—', 0,
     '預設 0.30。只用在 1A_需求公式樹 的「d 可調驗證」兩列；改它不會動其他頁。'),
])

IN = {}      # id → 儲存格（2_輸入!D?）
INLO = {}
INHI = {}
ws = WS['2_輸入']
widths(ws, {'A': 17, 'B': 7, 'C': 44, 'D': 13, 'E': 12, 'F': 30, 'G': 11, 'H': 11, 'I': 20, 'J': 70, 'K': 9})
put(ws, 'A1', 'm 推導的全部輸入（藍字＝輸入值，可改；改動後 3、4、5 頁與 1A 會自動重算）', F_T1, wrap=False)
put(ws, 'A2', '值取自 Tokenomics v21.2（20261023_Tokenomics_v21.2.xlsx，master 45ce1ca）。標記與區間取自該簿 Gov_Map 或 Cap_In 的低／高欄；'
    '「—」＝Tokenomics 無此登錄。黃底＝不確定度高且對 m 影響大（見 5_敏感度）。每格與 Tokenomics 的對應登錄在 9_核對對照。', F_T2, wrap=False)
hdr(ws, 4, ['ID', '群組', '項目', '值', '單位', '標記（Gov_Map／SRC）', '區間低', '區間高', 'Tokenomics 位置', '說明與來源', '不確定度'])
r = 5
for title, rows in IN_GROUPS:
    grp(ws, r, title, 11)
    r += 1
    for (iid, unit, label, expr, mark, lo, hi, unc, yel, note) in rows:
        fill = FILL_Y if yel else None
        put(ws, f'A{r}', iid)
        put(ws, f'B{r}', title[0])
        put(ws, f'C{r}', label)
        if expr:
            paste(ws, f'D{r}', expr, '輸入', fill=fill)
            ws[f'D{r}'].alignment = TOP
        else:
            put(ws, f'D{r}', 0.30, F_BLUE, fill, wrap=False)
        put(ws, f'E{r}', unit)
        put(ws, f'F{r}', mark)
        loc = expr if expr else '—'
        if expr and re.fullmatch(r'SRC_\w+', expr):
            loc = f'{expr}（Cap_In 以此為源）'
        put(ws, f'I{r}', loc)
        put(ws, f'J{r}', note)
        put(ws, f'K{r}', unc)
        IN[iid] = f'D{r}'
        defname(iid, f"{q('2_輸入')}!$D${r}")
        # 區間
        for col, spec, store, suffix in (('G', lo, INLO, '_Lo'), ('H', hi, INHI, '_Hi')):
            if isinstance(lo, str) and lo.startswith('GL:'):
                _, shn, cn = lo.split(':')
                glo, ghi = GL(shn, cn)
                spec = glo if col == 'G' else ghi
            if isinstance(spec, str):
                paste(ws, f'{col}{r}', spec, '區間低' if col == 'G' else '區間高')
                ws[f'{col}{r}'].alignment = TOP
                store[iid] = f'{col}{r}'
                defname(iid + suffix, f"{q('2_輸入')}!${col}${r}")
            elif isinstance(spec, (int, float)):
                put(ws, f'{col}{r}', spec, F_BLUE, wrap=False)
                store[iid] = f'{col}{r}'
                defname(iid + suffix, f"{q('2_輸入')}!${col}${r}")
            else:
                put(ws, f'{col}{r}', '—')
        r += 1
ws.freeze_panes = 'D5'
IN_LAST = r - 1
# 註：IN_ISL_* 等的區間來自 Gov_Map K／L 欄（儲存格值），與 Serving 頁無關

# ───────────────────────── 通用：四欄逐步表 ─────────────────────────
PER4 = ['2025', '2026 Q1', '2026 Q2', 'FY2026']


def table_head(ws, row, labels4, lead):
    """欄：A 步驟 B 項目 C 說明 D 單位 E 位置 F–I 值 J–M 原值 N–Q 差異"""
    hdr(ws, row, ['步驟', '項目', '白話說明與公式', '單位', '原活頁簿位置'] +
        [f'值 {x}' for x in labels4] + [f'Tokenomics 原值 {x}' for x in labels4] + [f'差異 {x}' for x in labels4])


def render_rows(ws, items, r0, ncol=17):
    """items：dict(kind='grp'|'step')；step 的公式可為 callable(R)。回傳 R（key→列）。"""
    R = {}
    r = r0
    for it in items:
        it['row'] = r
        if it.get('key'):
            R[it['key']] = r
        r += 1
    for it in items:
        r = it['row']
        if it['kind'] == 'grp':
            grp(ws, r, it['text'], ncol)
            continue
        hl = it.get('hl')
        fill = FILL_M if hl else None
        font = F_BOLD if hl else F_BASE
        put(ws, f'A{r}', it['id'], F_GRAY if not hl else F_BOLD, fill)
        put(ws, f'B{r}', it['name'], font, fill)
        put(ws, f'C{r}', it['desc'], F_GRAY, fill)
        put(ws, f'D{r}', it.get('unit', ''), F_BASE, fill)
        put(ws, f'E{r}', it.get('loc', ''), F_GRAY, fill)
        fs = it['f'](R) if callable(it['f']) else it['f']
        os_ = it.get('o', [None] * 4)
        for i in range(4):
            vc, oc, dc = CL(6 + i), CL(10 + i), CL(14 + i)
            f = fs[i]
            if f is not None:
                is_link = isinstance(f, str) and re.fullmatch(r"=('?[^'!]+'?!)\$?[A-Z]+\$?\d+", f) is not None
                put(ws, f'{vc}{r}', f, (F_GREEN if is_link and not hl else font), fill, it.get('fmt'), wrap=False)
            elif fill:
                ws[f'{vc}{r}'].fill = fill
            o = os_[i]
            if o is not None and f is not None:
                paste(ws, f'{oc}{r}', o, '原值', it.get('fmt'), fill)
                put(ws, f'{dc}{r}', diff(f'{vc}{r}', f'{oc}{r}'), F_BASE, fill, '0.000%', wrap=False)
                MAPROWS.append((ws.title, f'{dc}{r}', '', '差異'))
            elif fill:
                ws[f'{oc}{r}'].fill = fill
                ws[f'{dc}{r}'].fill = fill
    return R


# ───────────────────────── 3_m重算 ─────────────────────────
S3 = '3_m重算'
S4 = '4_分路線'
ws3 = WS[S3]
widths(ws3, {'A': 7, 'B': 34, 'C': 62, 'D': 12, 'E': 26, 'F': 13, 'G': 13, 'H': 13, 'I': 13,
             'J': 13, 'K': 13, 'L': 13, 'M': 13, 'N': 11, 'O': 11, 'P': 11, 'Q': 11})
put(ws3, 'A1', 'm 的逐步重算（全部為公式；改 2_輸入 的藍字即可看到每一步怎麼變）', F_T1, wrap=False)
put(ws3, 'A2', 'm（付費服務 token 的有效牌價變現率）＝ 實收單價 ÷ 有效牌價 ＝ m 分子營收 ÷（需求面付費 token × 混合有效單價）。只有收入面，不經成本鏈。'
    '每一步列出對應的 Tokenomics 格與原值，最後三組欄為原值與差異。四欄＝2025 年度、2026 Q1 年化、2026 Q2 年化（主口徑）、FY2026 預測。', F_T2, wrap=False)
table_head(ws3, 4, PER4, '')
ws3.freeze_panes = 'F5'

N_ = lambda x: f'{x}'  # noqa
I_ = lambda k: f'IN_{k}'  # noqa
FMT_B = '#,##0.000'
FMT_T = '#,##0'
FMT_M = '0.0000'
FMT_P = '0.0000'

items3 = [
    dict(kind='grp', text='第 A 組：分子——m 分子營收（營收 − 廣告）'),
    dict(kind='step', key='rev', id='S1', name='營收（年化）', unit='$B',
         desc='2025＝全年營收；2026 Q1／Q2＝季營收 × 4；FY2026＝內部預測。毛額，不扣 Microsoft 分潤。',
         loc='Cap_In 第 69 列 C；第 185 列 D–F',
         f=['=IN_Rev', '=IN_RevQ1*4', '=IN_RevQ2*4', '=IN_RevFY'], fmt=FMT_B,
         o=['Cap_In!C69', 'Cap_In!D185', 'Cap_In!E185', 'Cap_In!F185']),
    dict(kind='step', key='ad', id='S2', name='廣告營收（年化）', unit='$B',
         desc='Andy 定義：m 單純看收費業務，廣告不進分子。2025 沒有廣告（Tokenomics 的常數 0）。',
         loc='Cap_In 第 186 列 C–F', f=[0, '=IN_Ad_Q1', '=IN_Ad_Q2', '=IN_Ad_FY'], fmt=FMT_B,
         o=['Cap_In!C186', 'Cap_In!D186', 'Cap_In!E186', 'Cap_In!F186']),
    dict(kind='step', key='num', id='S3', name='m 分子營收', unit='$B',
         desc='＝S1 − S2。這是「付費服務實際收到的錢」。', loc='Cap_In 第 187 列 D–F（2025 為第 69 列）',
         f=lambda R: [f'=F{R["rev"]}-F{R["ad"]}', f'=G{R["rev"]}-G{R["ad"]}', f'=H{R["rev"]}-H{R["ad"]}', f'=I{R["rev"]}-I{R["ad"]}'],
         fmt=FMT_B, o=['Cap_In!C69', 'Cap_In!D187', 'Cap_In!E187', 'Cap_In!F187']),
    dict(kind='grp', text='第 B 組：分母之一——需求面付費 token（API ＋ ChatGPT 付費；免費不進 m）'),
    dict(kind='step', key='api', id='S4', name='API token（全年）', unit='M tok/年',
         desc='2025＝每分鐘 token × 525,600 分鐘 × 全年平均比例；2026 Q1／Q2＝2026-03 時點值 × 525,600（不打折）；FY2026＝Q2 × 年均倍數 1.3。',
         loc='Alloc 第 5 列（AL_DAPI）；Cap_In 第 196 列',
         f=lambda R: ['=IN_APItpm*1000000000*525600*IN_APIratio/1000000', '=IN_APItpm26*1000000000*525600/1000000',
                      '=IN_APItpm26*1000000000*525600/1000000', f'=H{R["api"]}*IN_APImultFY'], fmt=FMT_T,
         o=['AL_DAPI', 'Cap_In!D196', 'Cap_In!E196', 'Cap_In!F196']),
    dict(kind='step', key='pr', id='S5', name='ChatGPT 每日訊息數', unit='B 則/日',
         desc='2025 取 OpenAI 自述 25 億；2026 由週活比例放大（Assumed）。', loc='SRC_DEM_012；Cap_In 第 197 列',
         f=['=IN_Prompts', '=IN_Prompts_Q1', '=IN_Prompts_Q2', '=IN_Prompts_FY'], fmt='0.00',
         o=['SRC_DEM_012', 'Cap_In!D197', 'Cap_In!E197', 'Cap_In!F197']),
    dict(kind='step', key='chat', id='S6', name='ChatGPT token 總量（含免費）', unit='M tok/年',
         desc='＝每日訊息數 × 10⁹ × 每則 token τ × 365 ÷ 10⁶。τ 沒有公開數字（Assumed 4,000）。', loc='Alloc 第 6 列；Cap_In 第 198＋199 列',
         f=lambda R: [f'={c}{R["pr"]}*1000000000*IN_Tau*365/1000000' for c in 'FGHI'], fmt=FMT_T,
         o=['AL_DChat', 'Cap_In!D198+Cap_In!D199', 'Cap_In!E198+Cap_In!E199', 'Cap_In!F198+Cap_In!F199']),
    dict(kind='step', key='free', id='S7', name='ChatGPT 免費 token', unit='M tok/年',
         desc='＝S6 × 免費占比。這部分不產生營收，計入管理費用，不進 m（Andy 10-21 14:37）；這裡只是為了從總量扣掉。',
         loc='Alloc 第 7 列；Cap_In 第 199 列', f=lambda R: [f'={c}{R["chat"]}*IN_FreeShare' for c in 'FGHI'], fmt=FMT_T,
         o=['AL_DChatFree', 'Cap_In!D199', 'Cap_In!E199', 'Cap_In!F199']),
    dict(kind='step', key='paid', id='S8', name='ChatGPT 付費 token', unit='M tok/年',
         desc='＝S6 − S7＝S6 ×（1 − 免費占比）。', loc='Alloc 第 8 列；Cap_In 第 198 列',
         f=lambda R: [f'={c}{R["chat"]}-{c}{R["free"]}' for c in 'FGHI'], fmt=FMT_T,
         o=['AL_DChatPaid', 'Cap_In!D198', 'Cap_In!E198', 'Cap_In!F198']),
    dict(kind='step', key='tok', id='S9', name='需求面付費 token 合計', unit='M tok/年',
         desc='＝S4 ＋ S8。m 的分母用的 token 數量。', loc='Cap_In 第 83 列 C；第 196＋198 列',
         f=lambda R: [f'={c}{R["api"]}+{c}{R["paid"]}' for c in 'FGHI'], fmt=FMT_T,
         o=['Cap_In!C83', 'Cap_In!D196+Cap_In!D198', 'Cap_In!E196+Cap_In!E198', 'Cap_In!F196+Cap_In!F198']),
    dict(kind='grp', text='第 C 組：分母之二——混合有效單價（有效牌價；細節見本頁下方「層級單價展開表」）'),
    dict(kind='step', key='px', id='S10', name='付費 token 混合有效單價', unit='$/M tok',
         desc='＝Σ（層級占比 × 層級有效單價）× κ。2025 的層級有效單價＝2025 年度表列單價 ×（1 − d）；'
              '2026 的層級有效單價＝IF_PriceRef（表列牌價 × 請求形狀 ×（1 − d）× 能力倍數）。',
         loc='CAP_Price25Mix（Cap_In C281）；Cap_In 第 205 列', hl=True,
         f=lambda R: [f'=(F{R["t_mix"]}*F{R["t_px25"]}+G{R["t_mix"]}*G{R["t_px25"]}+H{R["t_mix"]}*H{R["t_px25"]})*IN_Kappa'] +
                     [f'=(F{R["t_mix"]}*F{R["t_ref"]}+G{R["t_mix"]}*G{R["t_ref"]}+H{R["t_mix"]}*H{R["t_ref"]})*IN_Kappa'] * 3,
         fmt=FMT_M, o=['CAP_Price25Mix', 'Cap_In!D205', 'Cap_In!E205', 'Cap_In!F205']),
    dict(kind='step', key='pl', id='S11', name='混合表列牌價（未折扣）', unit='$/M tok',
         desc='＝S10 ÷（1 − d）。與 d 無關：d 變動時有效單價同比例變動，表列牌價不動。2025 欄 Tokenomics 沒有對應名稱。',
         loc='L1_PxList_2026（L1 D126）', f=lambda R: [f'={c}{R["px"]}/(1-IN_Disc)' for c in 'FGHI'], fmt=FMT_M,
         o=[None, 'L1_PxList_2026', 'L1_PxList_2026', 'L1_PxList_2026']),
    dict(kind='grp', text='第 D 組：m ＝ 實收單價 ÷ 有效牌價'),
    dict(kind='step', key='act', id='S12', name='實收單價（每百萬付費 token 實收）', unit='$/M tok',
         desc='＝S3 × 10⁹ ÷ S9。付費服務每賣出 100 萬個 token，實際收到多少錢。Tokenomics 沒有這一格，本簿新增。', loc='（本簿新增）',
         f=lambda R: [f'={c}{R["num"]}*1000000000/{c}{R["tok"]}' for c in 'FGHI'], fmt=FMT_M),
    dict(kind='step', key='m', id='S13', name='m（付費 token 的有效牌價變現率）', unit='x', hl=True,
         desc='＝S12 ÷ S10。m＞1：實收高於有效牌價（訂閱定額月費遠超用量的牌價等值）；m＜1：實收低於有效牌價。',
         loc='Cap_In 第 209 列 C–F（m_整體）',
         f=lambda R: [f'={c}{R["act"]}/{c}{R["px"]}' for c in 'FGHI'], fmt=FMT_M,
         o=['Cap_In!C209', 'Cap_In!D209', 'Cap_In!E209', 'Cap_In!F209']),
    dict(kind='step', key='chk', id='S14', name='檢查：另一種寫法的 m', unit='x',
         desc='＝S3 × 10⁹ ÷（S9 × S10）。應與 S13 完全相同（恆等式）。',
         f=lambda R: [f'={c}{R["num"]}*1000000000/({c}{R["tok"]}*{c}{R["px"]})' for c in 'FGHI'], fmt=FMT_M),
    dict(kind='step', key='l1', id='L1', name='對照 L1 名稱 L1_M_2025／_2026Q1／_2026Q2／_FY2026', unit='x',
         desc='Tokenomics 的 L1 頁輸出名稱（下游契約）；與 S13 同值。', loc='L1 D117、D118、D119、D120',
         f=lambda R: [f'={c}{R["m"]}' for c in 'FGHI'], fmt=FMT_M,
         o=['L1_M_2025', 'L1_M_2026Q1', 'L1_M_2026Q2', 'L1_M_FY2026']),
    dict(kind='step', key='lm', id='L2', name='主口徑 L1_M（＝2026 Q2）', unit='x', hl=True,
         desc='Tokenomics 的主口徑 m＝2026 Q2 年化（J1）；FY2026 待 Andy 提供「市場預估 360 億」出處再議。', loc='L1 D121',
         f=lambda R: [None, None, f'=H{R["m"]}', None], fmt=FMT_M, o=[None, None, 'L1_M', None]),
    dict(kind='grp', text='層級單價展開表（S10 的細項；欄 F／G／H＝Luna／Sol／Astra，不是期間）'),
    dict(kind='step', key='t_pin', id='T1', name='新鮮輸入 表列牌價', unit='$/M', desc='輸入：OpenAI 官方牌價（2_輸入 D 組）。', loc='Cap_In 第 6 列 C–E',
         f=['=IN_PIn_L', '=IN_PIn_S', '=IN_PIn_A', None], fmt='0.0000', o=['Cap_In!C6', 'Cap_In!D6', 'Cap_In!E6', None]),
    dict(kind='step', key='t_pc', id='T2', name='快取輸入 表列牌價', unit='$/M', desc='輸入。', loc='Cap_In 第 7 列 C–E',
         f=['=IN_PCache_L', '=IN_PCache_S', '=IN_PCache_A', None], fmt='0.0000', o=['Cap_In!C7', 'Cap_In!D7', 'Cap_In!E7', None]),
    dict(kind='step', key='t_po', id='T3', name='輸出（含思考）表列牌價', unit='$/M', desc='輸入。思考 token 依輸出價計費。', loc='Cap_In 第 8 列 C–E',
         f=['=IN_POut_L', '=IN_POut_S', '=IN_POut_A', None], fmt='0.0000', o=['Cap_In!C8', 'Cap_In!D8', 'Cap_In!E8', None]),
    dict(kind='step', key='t_cap', id='T4', name='能力單價倍數', unit='x',
         desc='＝（H_CAP × 有效訓練算力倍數）^ε。基準 ε＝0，所以恆為 1（本簿只列原值，不展開）。', loc='Price_Frontier 第 38 列 C–E',
         f=['=IN_CapMult_L', '=IN_CapMult_S', '=IN_CapMult_A', None], fmt='0.0000', o=['Price_Frontier!C38', 'Price_Frontier!D38', 'Price_Frontier!E38', None]),
    dict(kind='step', key='t_isl', id='T5', name='參考請求輸入長度 ISL', unit='tok', desc='輸入。', loc='Serving 第 23 列 C–E',
         f=['=IN_ISL_L', '=IN_ISL_S', '=IN_ISL_A', None], fmt=FMT_T, o=['Serving!C23', 'Serving!D23', 'Serving!E23', None]),
    dict(kind='step', key='t_osl', id='T6', name='參考請求輸出長度 OSL（含思考）', unit='tok', desc='輸入。', loc='Serving 第 24 列 C–E',
         f=['=IN_OSL_L', '=IN_OSL_S', '=IN_OSL_A', None], fmt=FMT_T, o=['Serving!C24', 'Serving!D24', 'Serving!E24', None]),
    dict(kind='step', key='t_chi', id='T7', name='快取命中率 χ', unit='%', desc='輸入（三層共用）。', loc='Cap_In C13',
         f=['=IN_Chi', '=IN_Chi', '=IN_Chi', None], fmt='0.00', o=['Cap_In!C13', 'Cap_In!C13', 'Cap_In!C13', None]),
    dict(kind='step', key='t_fin', id='T8', name='新鮮輸入 有效價', unit='$/M', desc='＝T1 ×（1 − d）× T4。', loc='Price_Frontier 第 41 列 C–E',
         f=lambda R: [f'={c}{R["t_pin"]}*(1-IN_Disc)*{c}{R["t_cap"]}' for c in 'FGH'] + [None], fmt='0.0000',
         o=['Price_Frontier!C41', 'Price_Frontier!D41', 'Price_Frontier!E41', None]),
    dict(kind='step', key='t_cin', id='T9', name='快取輸入 有效價', unit='$/M', desc='＝T2 ×（1 − d）× T4。', loc='Price_Frontier 第 42 列 C–E',
         f=lambda R: [f'={c}{R["t_pc"]}*(1-IN_Disc)*{c}{R["t_cap"]}' for c in 'FGH'] + [None], fmt='0.0000',
         o=['Price_Frontier!C42', 'Price_Frontier!D42', 'Price_Frontier!E42', None]),
    dict(kind='step', key='t_out', id='T10', name='輸出 有效價', unit='$/M', desc='＝T3 ×（1 − d）× T4。', loc='Price_Frontier 第 44 列 C–E',
         f=lambda R: [f'={c}{R["t_po"]}*(1-IN_Disc)*{c}{R["t_cap"]}' for c in 'FGH'] + [None], fmt='0.0000',
         o=['Price_Frontier!C44', 'Price_Frontier!D44', 'Price_Frontier!E44', None]),
    dict(kind='step', key='t_ref', id='T11', name='參考請求混合有效價（IF_PriceRef_*）', unit='$/M', hl=True,
         desc='＝[ISL ×（（1 − χ）× T8 ＋ χ × T9）＋ OSL × T10] ÷（ISL ＋ OSL）。每 1 個 token（輸入加輸出）的平均有效價。',
         loc='Price_Frontier 第 45 列；IF_PriceRef_Luna／Sol／Astra',
         f=lambda R: [f'=({c}{R["t_isl"]}*((1-{c}{R["t_chi"]})*{c}{R["t_fin"]}+{c}{R["t_chi"]}*{c}{R["t_cin"]})+{c}{R["t_osl"]}*{c}{R["t_out"]})/({c}{R["t_isl"]}+{c}{R["t_osl"]})'
                      for c in 'FGH'] + [None], fmt='0.0000',
         o=['IF_PriceRef_Luna', 'IF_PriceRef_Sol', 'IF_PriceRef_Astra', None]),
    dict(kind='step', key='t_mix', id='T12', name='付費 token 層級占比', unit='%', desc='輸入（Assumed）。三者合計＝1。', loc='Cap_In 第 31 列 C–E（B4_MixPaid）',
         f=['=IN_MixP_L', '=IN_MixP_S', '=IN_MixP_A', None], fmt='0.00', o=['Cap_In!C31', 'Cap_In!D31', 'Cap_In!E31', None]),
    dict(kind='step', key='t_px25', id='T13', name='2025 層級有效單價', unit='$/M',
         desc='＝2025 年度層級表列單價（輸入）×（1 − d）。2025 不再乘能力倍數（2025 價格對應 2025 年的模型）。S10 的 2025 欄用本列。',
         loc='Cap_In 第 278–280 列 C（表列）',
         f=['=IN_Px_L*(1-IN_Disc)', '=IN_Px_S*(1-IN_Disc)', '=IN_Px_A*(1-IN_Disc)', None], fmt='0.0000'),
    dict(kind='step', key='t_px25l', id='T14', name='　　其中 2025 年度層級單價（表列，未折扣）', unit='$/M', desc='輸入的原值。', loc='Cap_In 第 278–280 列 C',
         f=['=IN_Px_L', '=IN_Px_S', '=IN_Px_A', None], fmt='0.0000', o=['Cap_In!C278', 'Cap_In!C279', 'Cap_In!C280', None]),
]
R3 = render_rows(ws3, items3, 5)
# 表尾：端點表
r_end0 = max(R3.values()) + 2
ws3[f'A{r_end0}'].value = None
hdr(ws3, r_end0, ['步驟', '項目', '白話說明與公式', '單位', '原活頁簿位置',
                  '值 2025 低端', '值 2025 高端', '值 2026 Q2 低端', '值 2026 Q2 高端',
                  'Tokenomics 原值 2025 低端', '2025 高端', '2026 Q2 低端', '2026 Q2 高端',
                  '差異 2025 低端', '2025 高端', '2026 Q2 低端', '2026 Q2 高端'])
items_e = [
    dict(kind='grp', text='區間端點（低端＝token 取高、m 取低；高端＝token 取低、m 取高。營收與單價取基準，價格區間不交——與 Tokenomics 寫法相同）'),
    dict(kind='step', key='e_tau', id='E1', name='每則 token τ', unit='tok', desc='低端取區間上緣（τ 大→付費 token 多→m 小）；高端取下緣。', loc='Alloc_In D15／E15',
         f=['=IN_Tau_Hi', '=IN_Tau_Lo', '=IN_Tau_Hi', '=IN_Tau_Lo'], fmt=FMT_T, o=['Alloc_In!E15', 'Alloc_In!D15', 'Alloc_In!E15', 'Alloc_In!D15']),
    dict(kind='step', key='e_fr', id='E2', name='ChatGPT 免費占比', unit='%', desc='低端取區間下緣（免費少→付費 token 多）；高端取上緣。', loc='Alloc_In D16／E16',
         f=['=IN_FreeShare_Lo', '=IN_FreeShare_Hi', '=IN_FreeShare_Lo', '=IN_FreeShare_Hi'], fmt='0.00',
         o=['Alloc_In!D16', 'Alloc_In!E16', 'Alloc_In!D16', 'Alloc_In!E16']),
    dict(kind='step', key='e_ar', id='E3', name='API 全年平均比例（只 2025）', unit='x', desc='低端取上緣、高端取下緣；2026 欄用時點值，不適用。', loc='Alloc_In D14／E14',
         f=['=IN_APIratio_Hi', '=IN_APIratio_Lo', None, None], fmt='0.00', o=['Alloc_In!E14', 'Alloc_In!D14', None, None]),
    dict(kind='step', key='e_api', id='E4', name='API token', unit='M tok/年', desc='2025 依 E3 重算；2026 Q2 與基準相同。', loc='Cap_In 第 82 列 D／E',
         f=lambda R: [f'=IN_APItpm*1000000000*525600*F{R["e_ar"]}/1000000', f'=IN_APItpm*1000000000*525600*G{R["e_ar"]}/1000000',
                      f'=$H${R3["api"]}', f'=$H${R3["api"]}'], fmt=FMT_T, o=['Cap_In!D82', 'Cap_In!E82', None, None]),
    dict(kind='step', key='e_cp', id='E5', name='ChatGPT 付費 token', unit='M tok/年', desc='＝每日訊息數 × 10⁹ × τ × 365 ÷ 10⁶ ×（1 − 免費占比），用端點的 τ 與免費占比。', loc='Cap_In 第 81 列 D／E',
         f=lambda R: [f'=IN_Prompts*1000000000*F{R["e_tau"]}*365/1000000*(1-F{R["e_fr"]})', f'=IN_Prompts*1000000000*G{R["e_tau"]}*365/1000000*(1-G{R["e_fr"]})',
                      f'=IN_Prompts_Q2*1000000000*H{R["e_tau"]}*365/1000000*(1-H{R["e_fr"]})', f'=IN_Prompts_Q2*1000000000*I{R["e_tau"]}*365/1000000*(1-I{R["e_fr"]})'],
         fmt=FMT_T, o=['Cap_In!D81', 'Cap_In!E81', None, None]),
    dict(kind='step', key='e_tok', id='E6', name='需求面付費 token 合計', unit='M tok/年', desc='＝E4 ＋ E5。', loc='Cap_In 第 83 列 D／E',
         f=lambda R: [f'={c}{R["e_api"]}+{c}{R["e_cp"]}' for c in 'FGHI'], fmt=FMT_T, o=['Cap_In!D83', 'Cap_In!E83', None, None]),
    dict(kind='step', key='e_num', id='E7', name='m 分子營收（取基準）', unit='$B', desc='連結 S3。', loc='',
         f=[f'=$F${R3["num"]}', f'=$F${R3["num"]}', f'=$H${R3["num"]}', f'=$H${R3["num"]}'], fmt=FMT_B),
    dict(kind='step', key='e_px', id='E8', name='混合有效單價（取基準）', unit='$/M tok', desc='連結 S10。', loc='',
         f=[f'=$F${R3["px"]}', f'=$F${R3["px"]}', f'=$H${R3["px"]}', f'=$H${R3["px"]}'], fmt=FMT_M),
    dict(kind='step', key='e_m', id='E9', name='m 端點', unit='x', hl=True, desc='＝E7 × 10⁹ ÷（E6 × E8）。',
         loc='L1_M_2025_Lo／_Hi（L1 E117／F117）；L1_M_2026Q2_Lo／_Hi（L1 E119／F119）',
         f=lambda R: [f'={c}{R["e_num"]}*1000000000/({c}{R["e_tok"]}*{c}{R["e_px"]})' for c in 'FGHI'], fmt=FMT_M,
         o=['L1_M_2025_Lo', 'L1_M_2025_Hi', 'L1_M_2026Q2_Lo', 'L1_M_2026Q2_Hi']),
]
RE = render_rows(ws3, items_e, r_end0 + 1)
R3.update(RE)
R3['px_row'] = R3['px']

# 輸出名稱（供 check_reading_m.py 核對）
defname('OUT_M_2025', f"{q(S3)}!$F${R3['m']}")
defname('OUT_M_2026Q1', f"{q(S3)}!$G${R3['m']}")
defname('OUT_M_2026Q2', f"{q(S3)}!$H${R3['m']}")
defname('OUT_M_FY2026', f"{q(S3)}!$I${R3['m']}")
defname('OUT_M', f"{q(S3)}!$H${R3['lm']}")
defname('OUT_M_2025_Lo', f"{q(S3)}!$F${R3['e_m']}")
defname('OUT_M_2025_Hi', f"{q(S3)}!$G${R3['e_m']}")
defname('OUT_M_2026Q2_Lo', f"{q(S3)}!$H${R3['e_m']}")
defname('OUT_M_2026Q2_Hi', f"{q(S3)}!$I${R3['e_m']}")
defname('OUT_PxEff_2026', f"{q(S3)}!$H${R3['px']}")
defname('OUT_PxList_2026', f"{q(S3)}!$H${R3['pl']}")
defname('OUT_Price25Mix', f"{q(S3)}!$F${R3['px']}")

# ───────────────────────── 4_分路線 ─────────────────────────
ws4 = WS[S4]
widths(ws4, {'A': 7, 'B': 34, 'C': 62, 'D': 12, 'E': 26, 'F': 13, 'G': 13, 'H': 13, 'I': 13,
             'J': 13, 'K': 13, 'L': 13, 'M': 13, 'N': 11, 'O': 11, 'P': 11, 'Q': 11})
put(ws4, 'A1', 'm 按收費模式拆開：API／消費訂閱／商用席次（營收三分 × token 三分）', F_T1, wrap=False)
put(ws4, 'A2', '每條路線的 m＝該路線營收 ÷（該路線 token × 混合有效單價）。三條路線共用同一個有效單價，所以差別全在「每個 token 實收多少」。'
    '主口徑看 2026 Q2 欄；2025 欄取 Tokenomics 的 2025 三分法。營收份額與消費訂閱 token 份額是 Assumed，只影響分路線，不影響整體 m。', F_T2, wrap=False)
table_head(ws4, 4, PER4, '')
ws4.freeze_panes = 'F5'
L3 = lambda col, key: f"{q(S3)}!{col}{R3[key]}"  # noqa
items4 = [
    dict(kind='grp', text='營收三分（$B）'),
    dict(kind='step', key='rev', id='R1', name='m 分子營收', unit='$B', desc='連結 3_m重算 S3。', loc='Cap_In 第 69 列 C；第 187 列 D–F',
         f=[f'={L3(c, "num")}' for c in 'FGHI'], fmt=FMT_B, o=['Cap_In!C69', 'Cap_In!D187', 'Cap_In!E187', 'Cap_In!F187']),
    dict(kind='step', key='sub', id='R2', name='營收份額：消費訂閱', unit='%', desc='輸入。2025 約 70%（FT）；2026 逐期下降（企業成長）。', loc='Cap_In C70；第 188 列 D–F',
         f=['=IN_SubShare', '=IN_SubShare_Q1', '=IN_SubShare_Q2', '=IN_SubShare_FY'], fmt='0.000', o=['Cap_In!C70', 'Cap_In!D188', 'Cap_In!E188', 'Cap_In!F188']),
    dict(kind='step', key='seat', id='R3', name='營收份額：商用席次', unit='%', desc='輸入（Assumed）。', loc='Cap_In C131；第 189 列 D–F',
         f=['=IN_SeatShare', '=IN_SeatShare_Q1', '=IN_SeatShare_Q2', '=IN_SeatShare_FY'], fmt='0.000', o=['Cap_In!C131', 'Cap_In!D189', 'Cap_In!E189', 'Cap_In!F189']),
    dict(kind='step', key='apis', id='R4', name='營收份額：API', unit='%', desc='＝1 − 消費訂閱 − 商用席次（餘額；前兩項的誤差全落在這裡）。', loc='Cap_In C137÷C69；第 190 列 D–F',
         f=lambda R: [f'=1-{c}{R["sub"]}-{c}{R["seat"]}' for c in 'FGHI'], fmt='0.000', o=['Cap_In!C137/Cap_In!C69', 'Cap_In!D190', 'Cap_In!E190', 'Cap_In!F190']),
    dict(kind='step', key='rsub', id='R5', name='消費訂閱營收', unit='$B', desc='＝R1 × R2。', loc='Cap_In 第 135 列 C；第 191 列 D–F',
         f=lambda R: [f'={c}{R["rev"]}*{c}{R["sub"]}' for c in 'FGHI'], fmt=FMT_B, o=['Cap_In!C135', 'Cap_In!D191', 'Cap_In!E191', 'Cap_In!F191']),
    dict(kind='step', key='rseat', id='R6', name='商用席次營收', unit='$B', desc='＝R1 × R3。', loc='Cap_In 第 136 列 C；第 192 列 D–F',
         f=lambda R: [f'={c}{R["rev"]}*{c}{R["seat"]}' for c in 'FGHI'], fmt=FMT_B, o=['Cap_In!C136', 'Cap_In!D192', 'Cap_In!E192', 'Cap_In!F192']),
    dict(kind='step', key='rapi', id='R7', name='API 營收', unit='$B', desc='＝R1 × R4。', loc='Cap_In 第 137 列 C；第 193 列 D–F',
         f=lambda R: [f'={c}{R["rev"]}*{c}{R["apis"]}' for c in 'FGHI'], fmt=FMT_B, o=['Cap_In!C137', 'Cap_In!D193', 'Cap_In!E193', 'Cap_In!F193']),
    dict(kind='grp', text='付費 token 三分（M tok/年）'),
    dict(kind='step', key='tapi', id='R8', name='API token', unit='M tok/年', desc='連結 3_m重算 S4。', loc='Cap_In 第 140 列 C；第 196 列 D–F',
         f=[f'={L3(c, "api")}' for c in 'FGHI'], fmt=FMT_T, o=['Cap_In!C140', 'Cap_In!D196', 'Cap_In!E196', 'Cap_In!F196']),
    dict(kind='step', key='cp', id='R9', name='ChatGPT 付費 token', unit='M tok/年', desc='連結 3_m重算 S8。', loc='Cap_In 第 81 列 C；第 198 列 D–F',
         f=[f'={L3(c, "paid")}' for c in 'FGHI'], fmt=FMT_T, o=['Cap_In!C81', 'Cap_In!D198', 'Cap_In!E198', 'Cap_In!F198']),
    dict(kind='step', key='sts', id='R10', name='消費訂閱占 ChatGPT 付費 token 份額', unit='%', desc='輸入（Assumed）。2026 三欄共用 0.65。', loc='Cap_In C132；第 200 列 D–F',
         f=['=IN_SubTokShare', '=IN_SubTokShare26', '=IN_SubTokShare26', '=IN_SubTokShare26'], fmt='0.00', o=['Cap_In!C132', 'Cap_In!D200', 'Cap_In!E200', 'Cap_In!F200']),
    dict(kind='step', key='tsub', id='R11', name='消費訂閱 token', unit='M tok/年', desc='＝R9 × R10。', loc='Cap_In 第 141 列 C；第 201 列 D–F',
         f=lambda R: [f'={c}{R["cp"]}*{c}{R["sts"]}' for c in 'FGHI'], fmt=FMT_T, o=['Cap_In!C141', 'Cap_In!D201', 'Cap_In!E201', 'Cap_In!F201']),
    dict(kind='step', key='tseat', id='R12', name='商用席次 token', unit='M tok/年', desc='＝R9 ×（1 − R10）。', loc='Cap_In 第 142 列 C；第 202 列 D–F',
         f=lambda R: [f'={c}{R["cp"]}*(1-{c}{R["sts"]})' for c in 'FGHI'], fmt=FMT_T, o=['Cap_In!C142', 'Cap_In!D202', 'Cap_In!E202', 'Cap_In!F202']),
    dict(kind='grp', text='各路線 m'),
    dict(kind='step', key='px', id='R13', name='混合有效單價（三路線共用）', unit='$/M tok', desc='連結 3_m重算 S10。', loc='CAP_Price25Mix；Cap_In 第 205 列',
         f=[f'={L3(c, "px")}' for c in 'FGHI'], fmt=FMT_M, o=['CAP_Price25Mix', 'Cap_In!D205', 'Cap_In!E205', 'Cap_In!F205']),
    dict(kind='step', key='mapi', id='R14', name='m_API', unit='x', hl=True, desc='＝R7 × 10⁹ ÷（R8 × R13）。按量計費；實收低於有效牌價。',
         loc='L1_MAPI（2025）；Cap_In 第 206 列；L1_M_API（Q2）',
         f=lambda R: [f'={c}{R["rapi"]}*1000000000/({c}{R["tapi"]}*{c}{R["px"]})' for c in 'FGHI'], fmt=FMT_M,
         o=['L1_MAPI', 'Cap_In!D206', 'L1_M_API', 'Cap_In!F206']),
    dict(kind='step', key='msub', id='R15', name='m_消費訂閱', unit='x', hl=True,
         desc='＝R5 × 10⁹ ÷（R11 × R13）。定額月費 ÷ 實際用量：遠大於 1，表示訂閱用戶付的錢遠超過用掉 token 的有效牌價等值（或 token 數被低估）。',
         loc='L1_MSub（2025）；Cap_In 第 207 列；L1_M_Sub（Q2）',
         f=lambda R: [f'={c}{R["rsub"]}*1000000000/({c}{R["tsub"]}*{c}{R["px"]})' for c in 'FGHI'], fmt=FMT_M,
         o=['L1_MSub', 'Cap_In!D207', 'L1_M_Sub', 'Cap_In!F207']),
    dict(kind='step', key='mseat', id='R16', name='m_商用席次', unit='x', hl=True, desc='＝R6 × 10⁹ ÷（R12 × R13）。按席次定額收費，同上。',
         loc='L1_MSeat（2025）；Cap_In 第 208 列；L1_M_Seat（Q2）',
         f=lambda R: [f'={c}{R["rseat"]}*1000000000/({c}{R["tseat"]}*{c}{R["px"]})' for c in 'FGHI'], fmt=FMT_M,
         o=['L1_MSeat', 'Cap_In!D208', 'L1_M_Seat', 'Cap_In!F208']),
    dict(kind='step', key='mall', id='R17', name='三路線合計的 m', unit='x',
         desc='＝（R5＋R6＋R7）× 10⁹ ÷（（R8＋R11＋R12）× R13）。三分只是把同一分子、同一分母拆開，必須等於整體 m。', loc='Cap_In 第 147 列 C；第 209 列 D–F',
         f=lambda R: [f'=({c}{R["rsub"]}+{c}{R["rseat"]}+{c}{R["rapi"]})*1000000000/(({c}{R["tapi"]}+{c}{R["tsub"]}+{c}{R["tseat"]})*{c}{R["px"]})' for c in 'FGHI'],
         fmt=FMT_M, o=['Cap_In!C147', 'Cap_In!D209', 'Cap_In!E209', 'Cap_In!F209']),
    dict(kind='step', key='mchk', id='R18', name='檢查：R17 − 3_m重算 S13', unit='x', desc='應為 0。',
         f=lambda R: [f'={c}{R["mall"]}-{q(S3)}!{c}{R3["m"]}' for c in 'FGHI'], fmt='0.0E+00'),
    dict(kind='step', key='ypm', id='R19', name='每百萬 token 實收：API', unit='$/M tok', desc='＝R7 × 10⁹ ÷ R8。', loc='（本簿新增）',
         f=lambda R: [f'={c}{R["rapi"]}*1000000000/{c}{R["tapi"]}' for c in 'FGHI'], fmt=FMT_M),
    dict(kind='step', key='yps', id='R20', name='每百萬 token 實收：消費訂閱', unit='$/M tok', desc='＝R5 × 10⁹ ÷ R11。', loc='（本簿新增）',
         f=lambda R: [f'={c}{R["rsub"]}*1000000000/{c}{R["tsub"]}' for c in 'FGHI'], fmt=FMT_M),
    dict(kind='step', key='ypt', id='R21', name='每百萬 token 實收：商用席次', unit='$/M tok', desc='＝R6 × 10⁹ ÷ R12。', loc='（本簿新增）',
         f=lambda R: [f'={c}{R["rseat"]}*1000000000/{c}{R["tseat"]}' for c in 'FGHI'], fmt=FMT_M),
]
c3 = lambda key, col: f"{q(S3)}!{col}{R3[key]}"  # noqa
R4 = render_rows(ws4, items4, 5)

# 2025 低／高端（Cap_In D／E 欄；F＝低端、G＝高端）
r_lh = max(R4.values()) + 2
hdr(ws4, r_lh, ['步驟', '項目', '白話說明與公式', '單位', '原活頁簿位置', '值 2025 低端', '值 2025 高端', '', '',
                'Tokenomics 原值 2025 低端', '2025 高端', '', '', '差異 2025 低端', '2025 高端', '', ''])
N2 = [None, None]
def lh(key, id_, name, unit, desc, loc, flo, fhi, olo, ohi, fmt):
    return dict(kind='step', key=key, id=id_, name=name, unit=unit, desc=desc, loc=loc,
                f=(lambda R, flo=flo, fhi=fhi: [flo(R) if callable(flo) else flo, fhi(R) if callable(fhi) else fhi, None, None]),
                o=[olo, ohi, None, None], fmt=fmt)
items_lh = [
    dict(kind='grp', text='2025 低／高端（Tokenomics 的 L1_MAPI／MSub／MSeat 區間；低端＝席次與訂閱營收取低、token 取高；API 為餘額，低高不是同向組合）'),
    lh('l_sub', 'H1', '消費訂閱營收份額', '%', '低端取 IN_SubShare_Lo、高端取 IN_SubShare_Hi。', 'Cap_In D70／E70', '=IN_SubShare_Lo', '=IN_SubShare_Hi', 'Cap_In!D70', 'Cap_In!E70', '0.000'),
    lh('l_seat', 'H2', '商用席次營收份額', '%', '低端 IN_SeatShare_Lo、高端 IN_SeatShare_Hi。', 'Cap_In D131／E131', '=IN_SeatShare_Lo', '=IN_SeatShare_Hi', 'Cap_In!D131', 'Cap_In!E131', '0.000'),
    lh('l_rsub', 'H3', '消費訂閱營收', '$B', '＝營收 × H1。', 'Cap_In D135／E135', lambda R: f'=IN_Rev*F{R["l_sub"]}', lambda R: f'=IN_Rev*G{R["l_sub"]}', 'Cap_In!D135', 'Cap_In!E135', FMT_B),
    lh('l_rseat', 'H4', '商用席次營收', '$B', '＝營收 × H2。', 'Cap_In D136／E136', lambda R: f'=IN_Rev*F{R["l_seat"]}', lambda R: f'=IN_Rev*G{R["l_seat"]}', 'Cap_In!D136', 'Cap_In!E136', FMT_B),
    lh('l_rapi', 'H5', 'API 營收（餘額）', '$B', '＝營收 − H3 − H4。', 'Cap_In D137／E137', lambda R: f'=IN_Rev-F{R["l_rsub"]}-F{R["l_rseat"]}', lambda R: f'=IN_Rev-G{R["l_rsub"]}-G{R["l_rseat"]}', 'Cap_In!D137', 'Cap_In!E137', FMT_B),
    lh('l_tapi', 'H6', 'API token', 'M tok/年', '連結 3_m重算 E4（低端取 API 全年比例上緣）。', 'Cap_In D140／E140', f"={c3('e_api', 'F')}", f"={c3('e_api', 'G')}", 'Cap_In!D140', 'Cap_In!E140', FMT_T),
    lh('l_cp', 'H7', 'ChatGPT 付費 token', 'M tok/年', '連結 3_m重算 E5。', 'Cap_In D81／E81', f"={c3('e_cp', 'F')}", f"={c3('e_cp', 'G')}", 'Cap_In!D81', 'Cap_In!E81', FMT_T),
    lh('l_sts', 'H8', '消費訂閱占 ChatGPT 付費 token 份額', '%', '低端取上緣 IN_SubTokShare_Hi（消費 token 多）、高端取下緣。', 'Cap_In D132／E132', '=IN_SubTokShare_Hi', '=IN_SubTokShare_Lo', 'Cap_In!D132', 'Cap_In!E132', '0.00'),
    lh('l_tsub', 'H9', '消費訂閱 token', 'M tok/年', '＝H7 × H8。', 'Cap_In D141／E141', lambda R: f'=F{R["l_cp"]}*F{R["l_sts"]}', lambda R: f'=G{R["l_cp"]}*G{R["l_sts"]}', 'Cap_In!D141', 'Cap_In!E141', FMT_T),
    lh('l_tseat', 'H10', '商用席次 token', 'M tok/年', '＝H7 ×（1 − H8）。', 'Cap_In D142／E142', lambda R: f'=F{R["l_cp"]}*(1-F{R["l_sts"]})', lambda R: f'=G{R["l_cp"]}*(1-G{R["l_sts"]})', 'Cap_In!D142', 'Cap_In!E142', FMT_T),
    lh('l_px', 'H11', '混合有效單價（取基準）', '$/M tok', '連結 3_m重算 S10（價格區間不交）。', 'CAP_Price25Mix', f"={c3('px', 'F')}", f"={c3('px', 'F')}", 'CAP_Price25Mix', 'CAP_Price25Mix', FMT_M),
    lh('l_mapi', 'H12', 'm_API 2025 低／高', 'x', '＝H5 × 10⁹ ÷（H6 × H11）。', 'Cap_In D144／E144', lambda R: f'=F{R["l_rapi"]}*1000000000/(F{R["l_tapi"]}*F{R["l_px"]})', lambda R: f'=G{R["l_rapi"]}*1000000000/(G{R["l_tapi"]}*G{R["l_px"]})', 'Cap_In!D144', 'Cap_In!E144', FMT_M),
    lh('l_msub', 'H13', 'm_消費訂閱 2025 低／高', 'x', '＝H3 × 10⁹ ÷（H9 × H11）。', 'Cap_In D145／E145', lambda R: f'=F{R["l_rsub"]}*1000000000/(F{R["l_tsub"]}*F{R["l_px"]})', lambda R: f'=G{R["l_rsub"]}*1000000000/(G{R["l_tsub"]}*G{R["l_px"]})', 'Cap_In!D145', 'Cap_In!E145', FMT_M),
    lh('l_mseat', 'H14', 'm_商用席次 2025 低／高', 'x', '＝H4 × 10⁹ ÷（H10 × H11）。', 'Cap_In D146／E146', lambda R: f'=F{R["l_rseat"]}*1000000000/(F{R["l_tseat"]}*F{R["l_px"]})', lambda R: f'=G{R["l_rseat"]}*1000000000/(G{R["l_tseat"]}*G{R["l_px"]})', 'Cap_In!D146', 'Cap_In!E146', FMT_M),
]
R4.update(render_rows(ws4, items_lh, r_lh + 1))
defname('OUT_M_API', f"{q(S4)}!$H${R4['mapi']}")
defname('OUT_M_Sub', f"{q(S4)}!$H${R4['msub']}")
defname('OUT_M_Seat', f"{q(S4)}!$H${R4['mseat']}")
defname('OUT_M_API_2025', f"{q(S4)}!$F${R4['mapi']}")
defname('OUT_M_Sub_2025', f"{q(S4)}!$F${R4['msub']}")
defname('OUT_M_Seat_2025', f"{q(S4)}!$F${R4['mseat']}")

exec(open(__file__.replace('build_reading_m.py', 'build_reading_m_part2.py'), encoding='utf-8').read())
