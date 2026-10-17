# v14.2 (work order docs/workorders/20261017_v14.2修補.md r0, branch claude/v14-2): patch after the v14.1 independent check.
# Minor version: no new contract name; the existing contract names L1_MFleet_P25Lo／L1_MFleet_P25Hi change value (low-end column), everything else unchanged.
#   S1  Cap_In rows 285 (L1_MFleet_P25Lo) and 286 (L1_MFleet_P25Hi), every column C／D／E = MAX(that column's point estimate at its own price, that column's own floor):
#       floor = revenue / (that column's fleet-paid-token ceiling (row 118) x the same price (AL_PriceFleet25_Lo／_Hi)).  Same mechanism as v14.1 R1 (row 119 = MAX(point, floor)),
#       but the floor is taken at the row's OWN price (the row 121 floor is at AL_PriceFleet25, the base 2025 price, so it cannot floor a row priced differently).
#       Excel formulas only; no Python-side computation.  L1_MFleet (all three ends), IF_Alloc*, IF_TokGW_*, RNG_* are not on the chain and do not change.
#   S2  Har_In rows 20／32: A／B labels cleared like v14.1 R4 (original text in Decisions SL45-2); H (source & mark) and I (note) kept (CLAUDE.md section 6).
#   S3  engineering: test rename, v14.1 report header SHAs, CHANGELOG v14.1 merge SHA (outside the workbook).
# Every write to an Excel-owned cell is guarded (old value / old text).
from openpyxl.utils import get_column_letter as L

import v1400
import v1410

VERSION = "20261017_Tokenomics_v14.2"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-17"
V = "v14.2"
WO = "工作單 docs/workorders/20261017_v14.2修補.md"
DEC_BY = "Andy 2026-10-17 13:04（回覆『全部依照你的建議』，同意 chat 端 S1、S2 的預設）"

HAR_ROWS = (20, 32)
HAR_A = "供應商專屬（1＝是）"
HAR_B = "旗標"
S1_ROWS = {285: "AL_PriceFleet25_Lo", 286: "AL_PriceFleet25_Hi"}
S1_NOTE = ("；下限截斷：v14.2 起比照 v14.1 R1，各欄不得低於『同欄、同單價』的物理下限"
           "＝營收 ÷（該欄機隊付費 token 上限（第 118 列）× {px}）；見 Decisions SL45-1")
OLD_F = '=IF(ISNUMBER({c}116),SRC_DEM_007*1000000000/({c}116*{px}),AL_UnavailTxt)'
NEW_F = ('=IF(ISNUMBER({c}116),MAX(SRC_DEM_007*1000000000/({c}116*{px}),'
         'IF(ISNUMBER({c}118),SRC_DEM_007*1000000000/({c}118*{px}),0)),AL_UnavailTxt)')


def _norm(s):
    return str(s).replace("1E+9", "1000000000").replace("1e9", "1000000000").replace("1E9", "1000000000")


def _purpose_col(ws):
    return v1400._purpose_col(ws, create=False)


# ================================================================== S1
def s1_p25(wb, append_text):
    ws = wb["Cap_In"]
    n = 0
    for r, px in S1_ROWS.items():
        assert str(ws[f"A{r}"].value).startswith("m（2025 單價取"), (r, ws[f"A{r}"].value)
        for c in "CDE":
            cur = ws[f"{c}{r}"].value
            old, new = OLD_F.format(c=c, px=px), NEW_F.format(c=c, px=px)
            if _norm(cur) == _norm(new):
                continue
            assert _norm(cur) == _norm(old), (r, c, cur)
            ws[f"{c}{r}"].value = new
            n += 1
        assert isinstance(ws[f"G{r}"].value, str), (r, ws[f"G{r}"].value)
        n += append_text(ws[f"G{r}"], S1_NOTE.format(px=px))
        pc = _purpose_col(ws)
        n += append_text(ws[f"{L(pc)}{r}"], "；v14.2 起下限截斷（比照 v14.1 R1，同欄同單價下限；見 Decisions SL45-1）")
    return n


# ================================================================== S2
S2_MARK_TAIL = "；v14.2 S2：A、B 欄標籤已清空（原『{a}｜{b}』記於 Decisions SL45-2）"


def s2_har_labels(wb, append_text):
    ws = wb["Har_In"]
    pc = _purpose_col(ws)
    orig = []
    for r in HAR_ROWS:
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a is None and b is None:
            continue
        assert (a, b) == (HAR_A, HAR_B), (r, a, b)
        orig.append((r, a, b))
        ws.cell(r, 1).value = None; ws.cell(r, 2).value = None
        append_text(ws[f"{L(pc)}{r}"], S2_MARK_TAIL.format(a=a, b=b))
    return orig


# ================================================================== Decisions
def _decisions(s2):
    lab = "；".join(f"Har_In!{r}：{a}｜{b}" for r, a, b in s2) or f"Har_In!{HAR_ROWS[0]}、Har_In!{HAR_ROWS[1]}：{HAR_A}｜{HAR_B}"
    _D = [
        ("1", "S1 L1_MFleet_P25Lo／P25Hi 比照 R1 截在同欄下限",
         "Cap_In 第 285 列（L1_MFleet_P25Lo，2025 單價取低端 AL_PriceFleet25_Lo）與第 286 列（L1_MFleet_P25Hi，單價取高端 AL_PriceFleet25_Hi）的 C／D／E 各欄，"
         "改為 MAX(該欄點估計, 該欄下限)；下限＝SRC_DEM_007 ÷（該欄機隊付費 token 上限（第 118 列）× 該列自己的單價）。"
         "機制與 v14.1 R1（第 119 列＝MAX(點估計, 第 121 列)）相同，全部是 Excel 公式。下限用『該列自己的單價』，因為第 121 列的下限是在基準 2025 單價（AL_PriceFleet25）下算的，"
         "套用到不同單價的列會不一致。v14.1 獨立查核（PR #108）發現：低端欄（D，token 取高）兩列都低於各自的同欄下限（P25Lo 0.3535 對 0.4141；P25Hi 0.1066 對 0.1249）。"
         "L1_MFleet（低端／基準／高端）、IF_Alloc*、IF_TokGW_*、RNG_* 不在依賴鏈上，不變。這兩個名稱仍是顯示用區間，不乘入 L1_MFleet 的 _Lo／_Hi（下游契約第 92 行）。"
         "新舊值見 docs/reports/20261017_v14.2修補.md。",
         "Cap_In 第 285、286 列 C／D／E", "v14.2 已建"),
        ("2", "S2 Har_In 第 20、32 列 A、B 標籤清空",
         "比照 v14.1 R4：供應商專屬旗標（Har_In 第 20、32 列 C:G）已於 v14.1 清空且無任何公式讀取，A、B 欄標籤一併清空。原文（工作表!列：原 A 欄｜原 B 欄）：" + lab +
         "。H 欄（來源與標記：Verified／Assumed）與 I 欄（說明）依 CLAUDE.md 第 6 節保留原文。",
         "Har_In 第 20、32 列 A、B", "v14.2 已建"),
        ("3", "S3 工程類小項",
         "(1) 測試 test_v1400_q3_floor_and_mfleet_lo_values_unchanged 改名為 test_v1400_q3_floor_row121_unchanged_and_mfleet_lo_is_v141_floor（內容不變）。"
         "(2) v14.1 報告表頭補最新提交 SHA c888b0c、合併 SHA ee5ccf7；CHANGELOG 補記 v14.1 合併 SHA。"
         "(3) v14.1 獨立查核報告第 6 節第 2 項：查核者對其餘項目無異議，沒有需補做的文字或標籤項；P25 項目即 S1。",
         "tests/parity/test_v1400.py；docs/reports；CHANGELOG.md", "v14.2 已建（Excel 外）"),
    ]
    return [[f"SL45-{i}", V, f"v14.2 修補：{t}", txt, DATE, DEC_BY, st, where, f"{WO} S{i}", "否"] for i, t, txt, where, st in _D]


def apply(wb, append_text):
    log = {}
    log["s1"] = s1_p25(wb, append_text)
    s2 = s2_har_labels(wb, append_text)
    log["s2_rows"] = len(s2)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == "SL45-1" for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions(s2))
    return log


# ================================================================== README
README_VERSION = (VERSION + "（次版：" + WO + "——S1 L1_MFleet_P25Lo／P25Hi（Cap_In 第 285、286 列）各欄比照 v14.1 R1 截在同欄同單價下限（MAX(點估計, 營收 ÷（第 118 列 × 該列單價)）；"
                  "低端欄 P25Lo 0.3535 → 0.4141、P25Hi 0.1066 → 0.1249，L1_MFleet 三端與其他名稱不變）；S2 Har_In 第 20、32 列 A、B 標籤清空；S3 工程小項；Decisions SL45-*）。以下為 "
                  + v1410.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1410.README_TITLE
README_TITLE = V + _T[len(v1410.V):]
README_ROW = ("v14.2 修補",
              "L1_MFleet_P25Lo／P25Hi 比照 v14.1 R1 截在同欄同單價下限（低端欄 0.3535 → 0.4141、0.1066 → 0.1249）；Har_In 第 20、32 列 A、B 標籤清空。細節見 Decisions SL45-*。")
