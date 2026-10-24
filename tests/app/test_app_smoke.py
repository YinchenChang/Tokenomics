"""網站檢查（2026-10-24 恢復網站，CLAUDE.md 修訂十一）。

1. 每一頁都能渲染、無例外（不寫死列數，Excel 換版不必改測試）。
2. 網站讀到的值＝LibreOffice 對現行活頁簿重新計算的值（全部具名範圍逐格比對）。
   網站讀的是活頁簿存下的值；本項確認 model/ 內的檔案確實是重算後存檔，沒有過期的存值。
"""
import shutil
import sys
from pathlib import Path

import pytest
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests" / "parity"))
HEAD = f"import sys; sys.path.insert(0, {str(ROOT)!r})\n"
PAGES = ("payback", "lookup", "overview", "block2", "block3", "block4", "block5", "alloc", "qa", "evidence", "governance")


@pytest.mark.parametrize("page", PAGES)
def test_page_renders(page):
    at = AppTest.from_string(HEAD + f"from app.views import {page}; {page}.render()", default_timeout=300).run()
    assert not at.exception, [e.message for e in at.exception]


def test_payback_headline_and_lookup_search():
    at = AppTest.from_string(HEAD + "from app.views import payback; payback.render()", default_timeout=300).run()
    labels = [m.label for m in at.metric]
    assert any(l.startswith("L1_LTotalFleetVDisc") for l in labels) and any(l.startswith("L1_LReqFleet_Full") for l in labels)
    at = AppTest.from_string(HEAD + "from app.views import lookup; lookup.render()", default_timeout=300).run()
    at.text_input[0].set_value("L1_RevGW_PQ_VR200").run()
    assert not at.exception
    names = list(at.dataframe[0].value["名稱"])
    assert "L1_RevGW_PQ_VR200" in names and all("L1_RevGW_PQ_VR200" in n for n in names)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="需要 LibreOffice")
def test_site_values_equal_libreoffice_recalc(tmp_path):
    from app.workbook import Workbook
    import openpyxl
    from parity_lib import lo_recalc

    site = Workbook()
    stripped = tmp_path / "in" / site.path.name          # openpyxl 存檔會丟掉所有存值 → LibreOffice 必須從頭計算每個公式
    stripped.parent.mkdir()
    openpyxl.load_workbook(site.path).save(stripped)
    fresh = Workbook(lo_recalc(stripped, tmp_path / "out"))
    bad = []
    for n in site.names:
        a, b = site.get_name(n), fresh.get_name(n)
        flat = lambda v: [x for r in (v if isinstance(v, list) else [v]) for x in (r if isinstance(r, list) else [r])]  # noqa: E731
        for x, y in zip(flat(a), flat(b)):
            if isinstance(x, (int, float)) and isinstance(y, (int, float)) and not isinstance(x, bool):
                if abs(x - y) > 1e-9 * max(1.0, abs(y)):
                    bad.append((n, x, y))
            elif str(x) != str(y):
                bad.append((n, x, y))
    assert not bad, f"{len(bad)} 格不一致，例：{bad[:10]}"
