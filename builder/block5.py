# Block 5 (v5.9): Har_In, Harness, Sens_Har; Interface E; Checks Block 5; Sources S58–S61; DB_Evidence E016–E020
# Decisions (Andy 2026-10-01, "都OK"): L1 harness = per-task parameter set (turns, thinking retention rho, thinking per turn,
# history compaction, cache hit, sub-agents, state retention, vendor flag) replacing the single token multiplier;
# L2 METR-type success rate p = 1/(1+(L/(H50 x horizon multiplier))^beta) with a tier x task override;
# L3 enhanced profiles off in the base (Tech_Registry T12 switch 0) — scenarios only; L4 per-GW theoretical revenue unchanged
# (harness acts on the per-task layer only); L5 cost per success = cost per attempt / p (independent retries, failure detectable);
# L6 non-GPU harness cost excluded (to-do); L7 scenario values only from neutral parties measured under the same conditions.
# New formulas reference key quantities through named ranges (B5_ = Block 5 internal / display; IF_ = downstream).
from common import *
from openpyxl.workbook.defined_name import DefinedName

TIERS = [("Luna", "Luna（低層）"), ("Sol", "Sol（中層）"), ("Astra", "Astra（頂層）")]
TK = "CDEFG"                                   # task columns (Workload C–G)
NA = '"SLO 不可達"'
GENS = [("VR200", "M"), ("GB300", "J")]         # Unit_Cost / Cache_Store base-cost columns (J13: VR200 default, GB300 alongside)
SETS = [("std", "標準 harness"), ("cur", "現行（依 T12 混合）"), ("sel", "選定檔案全採用")]

def nm(wb, n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)

def taskhead(ws, r, lab="項目"):
    put(ws, f"A{r}", lab, F_BOLD); put(ws, f"B{r}", "單位", F_BOLD)
    for c in TK: put(ws, f"{c}{r}", f"=Workload!{c}4", F_HLINK, wrap=True)
    ws.row_dimensions[r].height = 30

def trow(ws, r, lab, unit, f, fmt, key=False, note=None):
    """one row across the 5 task columns; f may use {c} (column) and {k} (1..5)"""
    put(ws, f"A{r}", lab, wrap=True); put(ws, f"B{r}", unit)
    for k, c in enumerate(TK, start=1):
        v = f(c, k) if callable(f) else f.format(c=c, k=k)
        put(ws, f"{c}{r}", v, fmt=fmt, fill=FILL_KEY if key else None)
    if note: put(ws, f"H{r}", note, F_NOTE, wrap=True)

# ---------------------------------------------------------------- Har_In
PROFILES = {  # rows: key, label, unit, fmt, values for 5 tasks, tag, note
 2: ("供應商原生增強（保留推理狀態＋上下文壓縮）", [
   ("T", "輪數倍數", "x", "0.00", [1, 1, 0.7, 0.7, 0.7], "Assumed", "保留推理狀態後重複探索減少；以 ARC-AGI-3『少用約 49% token』校準（Checks Block 5）；區間見 D 節"),
   ("rho", "思考保留 ρ", "%", "0%", [0, 0, 1, 1, 1], "Verified（機制）", "Provider Adapter 在請求之間保留不透明推理狀態（S58）；單輪任務無作用"),
   ("H", "每輪思考倍數", "x", "0.00", [1, 1, 0.8, 0.8, 0.8], "Assumed", "前輪推理可見，每輪重推減少"),
   ("C", "歷史保留比 c（壓縮後）", "x", "0.00", [1, 1, 0.7, 0.7, 0.7], "Assumed", "壓縮較長的對話（S58）；壓縮比例未揭露"),
   ("DChi", "快取命中變動", "百分點", "0%", [0, 0, 0, 0, 0], "Assumed", "加在 Workload 列 11 之上"),
   ("DM", "子代理數增量", "個", "0", [0, 0, 0, 0, 0], "Assumed", "加在 Workload 列 12 之上"),
   ("Ret", "狀態保留時間", "hr", "0.000", [0.0833333333333333, 0.0833333333333333, 0.67, 0.67, 0.67], "Analogy",
    "ARC-AGI-3 每局約 40 分鐘（S58）；單輪任務沿用 5 分鐘 TTL"),
   ("Hz", "時間範圍倍數（本任務）", "x", "0.00", [1, 1, 1.2, 2, 2], "Assumed",
    "L2：harness 乘在層級 50% 時間範圍上；短程 1.2、長程 2（區間 1–4，D 節）。ARC 結果無法換算為此倍數（Checks）"),
   ("Vendor", "供應商專屬（1＝是）", "旗標", "0", [1, 1, 1, 1, 1], "Verified", "不透明推理狀態只能在原廠 API 使用（轉換成本，[I]）"),
 ]),
 3: ("多代理編排（主代理＋並行子代理）", [
   ("T", "輪數倍數", "x", "0.00", [1, 1, 1, 1, 1], "Assumed", ""),
   ("rho", "思考保留 ρ", "%", "0%", [0, 0, 0, 0, 0], "Assumed", ""),
   ("H", "每輪思考倍數", "x", "0.00", [1, 1, 1, 1, 1], "Assumed", ""),
   ("C", "歷史保留比 c（壓縮後）", "x", "0.00", [1, 1, 1, 1, 1], "Assumed", ""),
   ("DChi", "快取命中變動", "百分點", "0%", [0, 0, 0, 0, 0], "Assumed", ""),
   ("DM", "子代理數增量", "個", "0", [0, 0, 3, 0, 3], "Analogy", "Anthropic 多代理：token 約單代理 4 倍（S25；Interested-party）；多代理研究已含 3 個子代理"),
   ("Ret", "狀態保留時間", "hr", "0.000", [0.0833333333333333] * 5, "Assumed", ""),
   ("Hz", "時間範圍倍數（本任務）", "x", "0.00", [1, 1, 1.2, 1.5, 1.5], "Assumed",
    "Anthropic 自報評測分數 +90.2%（內建知識，待查核；非成功率）只作對照（L7）"),
   ("Vendor", "供應商專屬（1＝是）", "旗標", "0", [0] * 5, "Assumed", "編排框架多為開源或開發者自建"),
 ]),
}

def har_in(wb, TR):
    ws = wb.create_sheet("Har_In")
    title(ws, "Har_In — Block 5 輸入：harness 檔案、成功率參數、敏感度區間、證據（藍字＝輸入；Analogy／Assumed 附區間）",
          "命題：harness 不是新的物理量，而是作用在標準任務上的一組參數變換加上任務成功率。標準檔＝Workload 列 5–14；"
          "選定檔案依 Tech_Registry T12 的開關 × 採用比例（w）混合進 Workload。基準 w＝0（L3）。每 GW 理論營收不受影響（L4）")
    for c, w in zip("ABCDEFGHI", [44, 9, 13, 13, 15, 15, 17, 16, 70]): ws.column_dimensions[c].width = w
    H = {}
    t12 = TR["_rows"][0] + 11          # T12 is the 12th registry entry
    r = 4
    section(ws, r, "A. 選擇與混合", 9); r += 1
    put(ws, f"A{r}", "情境檔案（2＝供應商原生增強；3＝多代理編排）"); put(ws, f"B{r}", "選擇"); put(ws, f"C{r}", 2, fmt="0")
    put(ws, f"H{r}", "Decision", F_NOTE); put(ws, f"I{r}", "只在 w > 0 或 Harness 頁『選定檔案全採用』欄起作用", F_NOTE); H["prof"] = r; r += 1
    put(ws, f"A{r}", "選定檔案名稱"); put(ws, f"C{r}", f"=CHOOSE(C{H['prof']}-1,\"{PROFILES[2][0]}\",\"{PROFILES[3][0]}\")"); H["pname"] = r; r += 1
    put(ws, f"A{r}", "混合權重 w＝T12 開關 × 採用比例"); put(ws, f"B{r}", "x")
    put(ws, f"C{r}", f"=Tech_Registry!O{t12}*Tech_Registry!N{t12}", F_LINK, fmt="0.00", fill=FILL_KEY)
    put(ws, f"I{r}", "w＝0：Workload 與標準檔相同，Block 1–4 與 v5.8 一致", F_NOTE); H["w"] = r; r += 2
    nm(wb, "B5_Profile", f"Har_In!$C${H['prof']}"); nm(wb, "B5_ProfileName", f"Har_In!$C${H['pname']}"); nm(wb, "B5_W", f"Har_In!$C${H['w']}")
    section(ws, r, "B. harness 檔案參數（欄＝Workload 任務；標準檔見 Workload 列 5–14，保留時間＝Cap_In 保留時間、時間範圍倍數＝1）", 9); r += 1
    for p in (2, 3):
        pname, rows = PROFILES[p]
        put(ws, f"A{r}", f"檔案 {p}：{pname}", F_BOLD); r += 1
        taskhead(ws, r, "參數"); put(ws, f"H{r}", "標記", F_BOLD); put(ws, f"I{r}", "說明", F_BOLD); r += 1
        for key, lab, unit, fmt, vals, tag, note in rows:
            put(ws, f"A{r}", lab); put(ws, f"B{r}", unit)
            for c, v in zip(TK, vals): put(ws, f"{c}{r}", v, fmt=fmt)
            put(ws, f"H{r}", tag, F_NOTE); put(ws, f"I{r}", note, F_NOTE, wrap=True); H[f"p{p}{key}"] = r; r += 1
        r += 1
    put(ws, f"A{r}", "選定檔案（依 A 節；公式）", F_BOLD); r += 1
    taskhead(ws, r, "參數"); r += 1
    for key, lab, unit, fmt, *_ in PROFILES[2][1]:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", unit)
        for c in TK: put(ws, f"{c}{r}", f"=CHOOSE(B5_Profile-1,{c}{H[f'p2{key}']},{c}{H[f'p3{key}']})", fmt=fmt)
        nm(wb, f"B5_Sel{key}", f"Har_In!$C${r}:$G${r}"); H[f"sel{key}"] = r; r += 1
    r += 1
    section(ws, r, "C. 成功率（L2：METR 型 p＝1 ÷（1＋（任務長度 ÷（層級 50% 時間範圍 × harness 倍數））^β））", 9); r += 1
    put(ws, f"A{r}", "項目", F_BOLD); put(ws, f"B{r}", "單位", F_BOLD)
    for c, h in zip("CDE", [tn for _, tn in TIERS]): put(ws, f"{c}{r}", h, F_BOLD)
    put(ws, f"H{r}", "標記", F_BOLD); put(ws, f"I{r}", "說明", F_BOLD); r += 1
    for key, lab, vals, tag, note in [
        ("h50", "50% 時間範圍 基準", (1.5, 8, 16), "Analogy",
         "METR TH1.1：GPT-5.6 Sol 11.3 小時、Claude Mythos Preview 17.4 小時（最高；套件 16 小時以上量不準）（S59）。GPT-6 各層級未發布；Luna 無對應錨點"),
        ("h50lo", "50% 時間範圍 低", (0.5, 4, 11), "Analogy", ""),
        ("h50hi", "50% 時間範圍 高", (4, 12, 40), "Analogy", "Astra 高值參照 Mythos 級以 ECI 推估 18.8–40 小時（第三方，S59）")]:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", "hr")
        for c, v in zip("CDE", vals): put(ws, f"{c}{r}", v, fmt="0.0")
        put(ws, f"H{r}", tag, F_NOTE); put(ws, f"I{r}", note, F_NOTE, wrap=True); H[key] = r; r += 1
    put(ws, f"A{r}", "斜率 β（C＝基準、D＝低、E＝高）"); put(ws, f"B{r}", "x")
    for c, v in zip("CDE", (0.75, 0.65, 1.0)): put(ws, f"{c}{r}", v, fmt="0.00")
    put(ws, f"H{r}", "Derived", F_NOTE)
    put(ws, f"I{r}", "由 METR 50% ÷ 80% 時間範圍比反推：β＝ln4 ÷ ln（比值）；GPT-5、5.2、5.4 約 0.70–0.75，Gemini 3.1 Pro 約 1.0（見 E 節）", F_NOTE, wrap=True)
    H["beta"] = r; r += 1
    nm(wb, "B5_H50", f"Har_In!$C${H['h50']}:$E${H['h50']}"); nm(wb, "B5_H50Lo", f"Har_In!$C${H['h50lo']}:$E${H['h50lo']}")
    nm(wb, "B5_H50Hi", f"Har_In!$C${H['h50hi']}:$E${H['h50hi']}"); nm(wb, "B5_Beta", f"Har_In!$C${H['beta']}")
    taskhead(ws, r, "任務"); r += 1
    put(ws, f"A{r}", "任務長度（人類完成時間）"); put(ws, f"B{r}", "hr")
    for c, v in zip(TK, [0.02, 0.1, 0.5, 2, 8]): put(ws, f"{c}{r}", v, fmt="0.00")
    put(ws, f"H{r}", "Assumed", F_NOTE); put(ws, f"I{r}", "新增假設（L2）；一般聊天約 1 分鐘、長程 coding 約一個工作日", F_NOTE); H["len"] = r
    nm(wb, "B5_TaskLen", f"Har_In!$C${r}:$G${r}"); r += 1
    put(ws, f"A{r}", "任務期程類別"); put(ws, f"B{r}", "")
    for c, v in zip(TK, ["短程", "短程", "短程", "長程", "長程"]): put(ws, f"{c}{r}", v, F_IN)
    put(ws, f"I{r}", "≥ 1 小時為長程；只作標示，倍數在 B 節逐任務輸入", F_NOTE); r += 1
    for t, (tk, tn) in enumerate(TIERS):
        put(ws, f"A{r}", f"成功率覆寫：{tn}（空白＝用公式）"); put(ws, f"B{r}", "%")
        for c in TK: ws[f"{c}{r}"].number_format = "0%"; ws[f"{c}{r}"].font = F_IN
        put(ws, f"I{r}", "有中立方實測時填入（L7）", F_NOTE); H[f"ovr{t}"] = r
        nm(wb, f"B5_SuccOvr_{tk}", f"Har_In!$C${r}:$G${r}"); r += 1
    r += 1
    section(ws, r, "D. 敏感度區間（Sens_Har；作用於選定檔案的 Coding agent（長程））", 9); r += 1
    put(ws, f"A{r}", "參數", F_BOLD); put(ws, f"C{r}", "低", F_BOLD); put(ws, f"D{r}", "高", F_BOLD); put(ws, f"I{r}", "說明", F_BOLD); r += 1
    for key, lab, lo, hi, fmt, note in [
        ("T", "輪數倍數", 0.5, 1.0, "0.00", ""), ("H", "每輪思考倍數", 0.6, 1.0, "0.00", ""),
        ("rho", "思考保留 ρ", 0.5, 1.0, "0%", ""), ("C", "歷史保留比 c", 0.5, 1.0, "0.00", ""),
        ("Hz", "時間範圍倍數（長程）", 1.0, 4.0, "0.00", "1＝harness 不提高成功率"), ("Ret", "狀態保留時間", 0.0833333333333333, 2, "0.000", "hr")]:
        put(ws, f"A{r}", lab); put(ws, f"C{r}", lo, fmt=fmt); put(ws, f"D{r}", hi, fmt=fmt); put(ws, f"H{r}", "Assumed", F_NOTE)
        put(ws, f"I{r}", note, F_NOTE); H[f"rg{key}"] = r; r += 1
    r += 1
    section(ws, r, "E. 證據點（L7：只有中立方同條件測得者可作情境值；其餘只作對照）", 9); r += 1
    for c, h in zip("ABCDEHI", ["證據", "單位", "成功率", "成本 $", "token 比", "標記", "來源／說明"]): put(ws, f"{c}{r}", h, F_BOLD)
    r += 1
    ev = [
      ("arc0", "ARC-AGI-3：GPT-6 Astra 標準 harness（最高推理）", 0.627, 26098, None, "Verified-measured（ARC Prize 自測，經二手轉述）", "S58"),
      ("arc1", "ARC-AGI-3：GPT-6 Astra Provider Adapter（最高推理，同檔）", 0.986, 17332, None, "Verified-measured（同上）", "S58"),
      ("arc2", "ARC-AGI-3：GPT-6 Astra Provider Adapter（高推理）", 0.999, 18817, 0.51, "Verified-measured（同上）",
       "S58；另一來源 $19,302（衝突）；『少用約 49% token』適用範圍不明，本表記為 token 比 0.51"),
      ("opus0", "ARC-AGI-3：Claude Opus 5 標準 harness", 0.302, None, None, "Verified-measured（同上）", "S58"),
      ("opus1", "ARC-AGI-3：Claude Opus 5 + Nvidia 所建 harness", 1.0, None, None, "Interested-party（Nvidia）",
       "S61；交接原記 Strands，報導為 Nvidia，待核；成本未找到"),
      ("ma", "Anthropic 多代理研究系統（相對單代理）", None, None, 3.75, "Interested-party（Anthropic）",
       "S25：token 約聊天 15 倍、單代理 4 倍 → 3.75；內部評測分數 +90.2%（內建知識，待查核；非成功率）")]
    for key, lab, p, cost, tok, tag, src in ev:
        put(ws, f"A{r}", lab, wrap=True)
        if p is not None: put(ws, f"C{r}", p, fmt="0.0%")
        if cost is not None: put(ws, f"D{r}", cost, fmt="#,##0")
        if tok is not None: put(ws, f"E{r}", tok, fmt="0.00")
        put(ws, f"H{r}", tag, F_NOTE, wrap=True); put(ws, f"I{r}", src, F_NOTE, wrap=True); H[key] = r; r += 1
    for c, h in zip("ABCDEFHI", ["METR TH1.1 50% ÷ 80% 時間範圍", "", "50%（hr）", "80%（hr）", "比值", "隱含 β", "標記", "來源"]): put(ws, f"{c}{r}", h, F_BOLD)
    r += 1
    H["metr0"] = r
    for lab, h50, h80 in [("GPT-5", 3.5667, 0.5333), ("GPT-5.2（high）", 6.5667, 0.9167),        # 3h34／32m、6h34／55m
                          ("GPT-5.4（xhigh）", 5.7, 0.9), ("Gemini 3.1 Pro", 5.8333, 1.5)]:           # 5h42／54m、5h50／1h30
        put(ws, f"A{r}", lab); put(ws, f"C{r}", h50, fmt="0.00"); put(ws, f"D{r}", h80, fmt="0.00")
        put(ws, f"E{r}", f"=C{r}/D{r}", fmt="0.0"); put(ws, f"F{r}", f"=LN(4)/LN(E{r})", fmt="0.00")
        put(ws, f"H{r}", "Verified-measured（METR）", F_NOTE); put(ws, f"I{r}", "S59", F_NOTE); r += 1
    H["metr1"] = r - 1
    ws.freeze_panes = "C4"
    return H

# ---------------------------------------------------------------- Harness
def harness(wb, U, S, WL):
    ws = wb.create_sheet("Harness")
    title(ws, "Harness — 每成功任務的 token、成本與營收（世代 × 層級 × 任務 × harness 檔案；L1–L5）",
          "每次嘗試成本＝新鮮 × prefill 成本＋快取 ×（載入＋儲存 × 保留時間 ÷ Cap_In 保留時間）＋ decode × decode 成本（經濟、基準成本、基準利用率）；"
          "每成功任務＝每次嘗試 ÷ 成功率 p（L5：重試獨立、失敗可偵測）。harness 不改變每 GW 產出與理論營收（L4）。非 GPU 成本未計（L6）")
    for c, w in zip("ABCDEFGH", [52, 10, 14, 14, 16, 16, 18, 60]): ws.column_dimensions[c].width = w
    R = {}
    r = 4
    taskhead(ws, r, "項目"); r += 1
    W = lambda key: f"Workload!{{c}}{WL[key]}"
    section(ws, r, "A. 每次嘗試的 token（標準與選定檔案全採用由本頁計算；現行＝Workload 有效參數）", 8); r += 1
    for sk, sn in (("std", "標準 harness"), ("sel", "選定檔案全採用")):
        put(ws, f"A{r}", sn, F_BOLD); r += 1
        if sk == "std":
            pT, pH, pRho, pChi, pM, pC = W("T"), W("h"), W("rho"), W("chi"), W("m"), W("c")
        else:
            pT = W("T") + "*INDEX(B5_SelT,1,{k})"; pH = W("h") + "*INDEX(B5_SelH,1,{k})"; pRho = "INDEX(B5_SelRho,1,{k})"
            pChi = "MIN(1,MAX(0," + W("chi") + "+INDEX(B5_SelDChi,1,{k})))"; pM = W("m") + "+INDEX(B5_SelDM,1,{k})"; pC = "INDEX(B5_SelC,1,{k})"
        base = r
        rows = [("T", "輪數 T", "輪", "=" + pT, "#,##0.0"), ("h", "每輪思考 h", "tok", "=" + pH, "#,##0"),
                ("rho", "思考保留 ρ", "%", "=" + pRho, "0%"), ("chi", "快取命中 χ", "%", "=" + pChi, "0%"),
                ("m", "子代理數 m", "個", "=" + pM, "0.0"), ("cc", "歷史保留比 c", "x", "=" + pC, "0.00")]
        idx = {}
        for i, (key, lab, unit, f, fmt) in enumerate(rows):
            trow(ws, r, lab, unit, f, fmt); idx[key] = r; r += 1
        g = lambda key: f"{{c}}{idx[key]}"
        inc = r; trow(ws, r, "每輪上下文增量 u＋o＋ρh", "tok", f"={W('u')}+{W('o')}+{g('rho')}*{g('h')}", "#,##0"); r += 1
        hist = f"{g('cc')}*{{c}}{inc}*{g('T')}*({g('T')}-1)/2"
        inp = r; trow(ws, r, "單代理輸入總量", "tok", f"={g('T')}*({W('S')}+{W('u')})+{hist}", "#,##0"); r += 1
        cac = r; trow(ws, r, "其中可快取", "tok", f"={g('chi')}*({g('T')}*{W('S')}+{hist})", "#,##0"); r += 1
        dec = r; trow(ws, r, "單代理 decode", "tok", f"={g('T')}*({g('h')}+{W('o')})", "#,##0"); r += 1
        mul = r; trow(ws, r, "系統倍數（1＋m）× 殘差倍數", "x", f"=(1+{g('m')})*{W('res')}", "0.00"); r += 1
        R[f"{sk}_fp"] = r; trow(ws, r, "任務新鮮 prefill", "tok", f"=({{c}}{inp}-{{c}}{cac})*{{c}}{mul}", "#,##0"); r += 1
        R[f"{sk}_cp"] = r; trow(ws, r, "任務快取 prefill", "tok", f"={{c}}{cac}*{{c}}{mul}", "#,##0"); r += 1
        R[f"{sk}_dp"] = r; trow(ws, r, "任務 decode", "tok", f"={{c}}{dec}*{{c}}{mul}", "#,##0"); r += 1
        R[f"{sk}_tot"] = r; trow(ws, r, "任務總 token", "tok", f"={{c}}{R[sk+'_fp']}+{{c}}{R[sk+'_cp']}+{{c}}{R[sk+'_dp']}", "#,##0", key=True); r += 1
        r += 1
    put(ws, f"A{r}", "現行（Workload）", F_BOLD); r += 1
    for key, wkey, lab in (("fp", "fp", "任務新鮮 prefill"), ("cp", "cp", "任務快取 prefill"), ("dp", "dp", "任務 decode"), ("tot", "tot", "任務總 token")):
        R[f"cur_{key}"] = r; trow(ws, r, lab, "tok", f"=Workload!{{c}}{WL[wkey]}", "#,##0", key=(key == "tot")); r += 1
    R["tokratio"] = r
    trow(ws, r, "選定檔案 ÷ 標準：任務總 token", "x", f"={{c}}{R['sel_tot']}/{{c}}{R['std_tot']}", "0.00", key=True,
         note="ARC-AGI-3 對照：Provider Adapter 少用約 49% token（0.51；適用範圍不明）"); r += 2
    section(ws, r, "B. 狀態保留時間與時間範圍倍數", 8); r += 1
    R["std_ret"] = r; trow(ws, r, "保留時間：標準", "hr", "=B4_Retain", "0.000"); r += 1
    R["cur_ret"] = r; trow(ws, r, "保留時間：現行", "hr", "=B4_Retain+B5_W*(INDEX(B5_SelRet,1,{k})-B4_Retain)", "0.000"); r += 1
    R["sel_ret"] = r; trow(ws, r, "保留時間：選定檔案", "hr", "=INDEX(B5_SelRet,1,{k})", "0.000"); r += 1
    R["std_hz"] = r; trow(ws, r, "時間範圍倍數：標準", "x", "=1", "0.00"); r += 1
    R["cur_hz"] = r; trow(ws, r, "時間範圍倍數：現行", "x", "=1+B5_W*(INDEX(B5_SelHz,1,{k})-1)", "0.00"); r += 1
    R["sel_hz"] = r; trow(ws, r, "時間範圍倍數：選定檔案", "x", "=INDEX(B5_SelHz,1,{k})", "0.00"); r += 2
    section(ws, r, "C. 成功率 p（覆寫欄有值時用覆寫）", 8); r += 1
    for t, (tk, tn) in enumerate(TIERS):
        for sk, sn in SETS:
            R[f"p_{sk}{t}"] = r
            trow(ws, r, f"{tn}｜{sn}", "%",
                 f"=IF(ISNUMBER(INDEX(B5_SuccOvr_{tk},1,{{k}})),INDEX(B5_SuccOvr_{tk},1,{{k}}),"
                 f"1/(1+(INDEX(B5_TaskLen,1,{{k}})/(INDEX(B5_H50,1,{t+1})*{{c}}{R[sk+'_hz']}))^B5_Beta))", "0.0%", key=(sk == "cur")); r += 1
    r += 1
    section(ws, r, "D. 每次嘗試成本（$／任務；經濟、基準成本、基準利用率；含快取儲存）", 8); r += 1
    for gname, col in GENS:
        for t, (tk, tn) in enumerate(TIERS):
            cf, cc, cd = U[(t + 1, "cfu")], U[(t + 1, "ccu")], U[(t + 1, "cdu")]
            for sk, sn in SETS:
                R[f"c_{gname}{t}{sk}"] = r
                trow(ws, r, f"{gname} × {tn}｜{sn}", "$",
                     f"=IF(ISNUMBER(Unit_Cost!{col}{cd}),({{c}}{R[sk+'_fp']}*Unit_Cost!{col}{cf}+{{c}}{R[sk+'_cp']}*(Unit_Cost!{col}{cc}"
                     f"+Cache_Store!{col}{S[tk]}*{{c}}{R[sk+'_ret']}/B4_Retain)+{{c}}{R[sk+'_dp']}*Unit_Cost!{col}{cd})/1E6,{NA})", "$#,##0.0000"); r += 1
    r += 1
    section(ws, r, "E. 每成功任務成本（＝每次嘗試 ÷ p）與 token", 8); r += 1
    for gname, col in GENS:
        for t, (tk, tn) in enumerate(TIERS):
            for sk, sn in SETS:
                cref = f"{{c}}{R[f'c_{gname}{t}{sk}']}"; pref = f"{{c}}{R[f'p_{sk}{t}']}"
                R[f"cs_{gname}{t}{sk}"] = r
                trow(ws, r, f"{gname} × {tn}｜{sn}", "$", f"=IF(AND(ISNUMBER({cref}),{pref}>0),{cref}/{pref},{NA})", "$#,##0.0000",
                     key=(sk == "cur" and gname == "VR200")); r += 1
    for t, (tk, tn) in enumerate(TIERS):
        for sk, sn in SETS:
            R[f"ts_{t}{sk}"] = r
            trow(ws, r, f"每成功任務 token：{tn}｜{sn}", "tok", f"=IF({{c}}{R[f'p_{sk}{t}']}>0,{{c}}{R[sk+'_tot']}/{{c}}{R[f'p_{sk}{t}']},0)", "#,##0"); r += 1
    r += 1
    section(ws, r, "F. 每任務營收（OpenAI 有效單價；營收以每次嘗試計費，每成功任務＝÷ p）", 8); r += 1
    for t, (tk, tn) in enumerate(TIERS):
        for sk, sn in SETS:
            R[f"rv_{t}{sk}"] = r
            trow(ws, r, f"每次嘗試營收：{tn}｜{sn}", "$",
                 f"=({{c}}{R[sk+'_fp']}*INDEX(B4_EffIn,1,{t+1})+{{c}}{R[sk+'_cp']}*INDEX(B4_EffCache,1,{t+1})+{{c}}{R[sk+'_dp']}*INDEX(B4_EffOut,1,{t+1}))/1E6",
                 "$#,##0.0000"); r += 1
    for t, (tk, tn) in enumerate(TIERS):
        for sk, sn in SETS:
            R[f"rs_{t}{sk}"] = r
            trow(ws, r, f"每成功任務營收：{tn}｜{sn}", "$", f"=IF({{c}}{R[f'p_{sk}{t}']}>0,{{c}}{R[f'rv_{t}{sk}']}/{{c}}{R[f'p_{sk}{t}']},0)", "$#,##0.0000"); r += 1
    r += 1
    section(ws, r, "G. R＝每成功任務成本：選定檔案全採用 ÷ 標準（<1＝harness 降低每成功任務成本）", 8); r += 1
    for gname, col in GENS:
        for t, (tk, tn) in enumerate(TIERS):
            a, b = f"{{c}}{R[f'cs_{gname}{t}sel']}", f"{{c}}{R[f'cs_{gname}{t}std']}"
            R[f"R_{gname}{t}"] = r
            trow(ws, r, f"{gname} × {tn}", "x", f"=IF(AND(ISNUMBER({a}),ISNUMBER({b})),{a}/{b},{NA})", "0.00", key=(gname == "VR200")); r += 1
    for t, (tk, tn) in enumerate(TIERS):
        trow(ws, r, f"成功率比（選定 ÷ 標準）：{tn}", "x", f"={{c}}{R[f'p_sel{t}']}/{{c}}{R[f'p_std{t}']}", "0.00"); R[f"pr{t}"] = r; r += 1
    r += 1
    section(ws, r, "H. 成功任務成本前緣（各世代：3 層級 × {標準、選定檔案} 中每成功任務成本最低者；層級替代）", 8); r += 1
    for gname, col in GENS:
        cells = [(f"{TIERS[t][0]}｜標準", R[f"cs_{gname}{t}std"]) for t in range(3)] + [(f"{TIERS[t][0]}｜選定", R[f"cs_{gname}{t}sel"]) for t in range(3)]
        R[f"fr_{gname}"] = r
        trow(ws, r, f"{gname}：最低每成功任務成本", "$", lambda c, k, cells=cells: "=MIN(" + ",".join(f"{c}{rr}" for _, rr in cells) + ")", "$#,##0.0000", key=True); r += 1
        def lab(c, k, cells=cells, fr=r - 1):
            f = '"—"'
            for name, rr in reversed(cells):
                f = f'IF({c}{rr}={c}{fr},"{name}",{f})'
            return "=" + f
        R[f"frl_{gname}"] = r; trow(ws, r, f"{gname}：前緣組合（層級｜檔案）", "", lab, None); r += 1
    put(ws, f"A{r}", "註：標準與現行在 w＝0 時相同。前緣只比較成本，不含延遲；選定檔案『供應商專屬』時，前緣組合綁定該供應商 API（Har_In 旗標）", F_NOTE)
    ws.freeze_panes = "C5"
    for key, n in [("cur_tot", "B5_TaskTok"), ("sel_tot", "B5_TaskTokSel"), ("std_tot", "B5_TaskTokStd")]:
        nm(wb, n, f"Harness!$C${R[key]}:$G${R[key]}")
    return R

# ---------------------------------------------------------------- Sens_Har
def sens_har(wb, U, S, WL, H, R):
    ws = wb.create_sheet("Sens_Har")
    title(ws, "Sens_Har — R 與每成功任務成本的敏感度（VR200、基準成本；Coding agent（長程）；選定檔案全採用 vs 標準）",
          "每欄只改一個參數（區間取 Har_In D、C 節），其餘同基準欄。R＝（選定每次嘗試成本 ÷ p 選定）÷（標準每次嘗試成本 ÷ p 標準）")
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 8
    k = 5; c0 = "G"                                    # Coding agent = task 5 (Workload column G)
    var = [("基準", None, None)]
    for key, lab in [("T", "輪數倍數"), ("H", "每輪思考倍數"), ("rho", "思考保留 ρ"), ("C", "歷史保留比"), ("Hz", "時間範圍倍數"), ("Ret", "保留時間")]:
        var += [(f"{lab} 低", key, f"Har_In!$C${H['rg' + key]}"), (f"{lab} 高", key, f"Har_In!$D${H['rg' + key]}")]
    var += [("β 低", "beta", "Har_In!$D$" + str(H["beta"])), ("β 高", "beta", "Har_In!$E$" + str(H["beta"])),
            ("50% 時間範圍 低", "h50", "lo"), ("50% 時間範圍 高", "h50", "hi")]
    cols = [L(3 + i) for i in range(len(var))]
    for X in cols: ws.column_dimensions[X].width = 11
    put(ws, "A4", "情境", F_BOLD)
    for X, (lab, _, _) in zip(cols, var): put(ws, f"{X}4", lab, F_BOLD, wrap=True)
    ws.row_dimensions[4].height = 42
    W = lambda key: f"Workload!${c0}${WL[key]}"
    basev = {"T": f"INDEX(B5_SelT,1,{k})", "H": f"INDEX(B5_SelH,1,{k})", "rho": f"INDEX(B5_SelRho,1,{k})", "C": f"INDEX(B5_SelC,1,{k})",
             "Hz": f"INDEX(B5_SelHz,1,{k})", "Ret": f"INDEX(B5_SelRet,1,{k})", "beta": "B5_Beta"}
    rows = {}
    r = 5
    for key, lab, fmt in [("T", "輪數倍數", "0.00"), ("H", "每輪思考倍數", "0.00"), ("rho", "思考保留 ρ", "0%"), ("C", "歷史保留比 c", "0.00"),
                          ("Hz", "時間範圍倍數", "0.00"), ("Ret", "保留時間（hr）", "0.000"), ("beta", "β", "0.00")]:
        put(ws, f"A{r}", lab)
        for X, (_, vk, ref) in zip(cols, var):
            put(ws, f"{X}{r}", f"={ref}" if vk == key else f"={basev[key]}", fmt=fmt)
        rows[key] = r; r += 1
    for t, (tk, tn) in enumerate(TIERS):
        put(ws, f"A{r}", f"50% 時間範圍：{tn}（hr）")
        for X, (_, vk, ref) in zip(cols, var):
            nmh = {"lo": "B5_H50Lo", "hi": "B5_H50Hi"}.get(ref, "B5_H50") if vk == "h50" else "B5_H50"
            put(ws, f"{X}{r}", f"=INDEX({nmh},1,{t+1})", fmt="0.0")
        rows[f"h{t}"] = r; r += 1
    r += 1
    def drow(key, lab, f, fmt, key_fill=False):
        nonlocal r
        put(ws, f"A{r}", lab)
        for X in cols: put(ws, f"{X}{r}", f.replace("{X}", X), fmt=fmt, fill=FILL_KEY if key_fill else None)
        rows[key] = r; r += 1
    g = lambda key: f"{{X}}{rows[key]}"
    drow("Tn", "選定：輪數", f"={W('T')}*{g('T')}", "0.0")
    drow("hn", "選定：每輪思考", f"={W('h')}*{g('H')}", "#,##0")
    drow("inc", "選定：每輪增量", f"={W('u')}+{W('o')}+{g('rho')}*{g('hn')}", "#,##0")
    hist = f"{g('C')}*{g('inc')}*{g('Tn')}*({g('Tn')}-1)/2"
    chi = f"MIN(1,MAX(0,{W('chi')}+INDEX(B5_SelDChi,1,{k})))"
    mul = f"(1+{W('m')}+INDEX(B5_SelDM,1,{k}))*{W('res')}"
    drow("inp", "選定：輸入總量（單代理）", f"={g('Tn')}*({W('S')}+{W('u')})+{hist}", "#,##0")
    drow("cac", "選定：可快取", f"={chi}*({g('Tn')}*{W('S')}+{hist})", "#,##0")
    drow("fp", "選定：任務新鮮 prefill", f"=({g('inp')}-{g('cac')})*{mul}", "#,##0")
    drow("cp", "選定：任務快取 prefill", f"={g('cac')}*{mul}", "#,##0")
    drow("dp", "選定：任務 decode", f"={g('Tn')}*({g('hn')}+{W('o')})*{mul}", "#,##0")
    drow("tr", "token 比（選定 ÷ 標準）", f"=({g('fp')}+{g('cp')}+{g('dp')})/Harness!${c0}${R['std_tot']}", "0.00")
    r += 1
    col = "M"
    for t, (tk, tn) in enumerate(TIERS):
        cf, cc, cd = U[(t + 1, "cfu")], U[(t + 1, "ccu")], U[(t + 1, "cdu")]
        drow(f"cs{t}", f"{tn}：選定每次嘗試成本 $",
             f"=IF(ISNUMBER(Unit_Cost!${col}${cd}),({g('fp')}*Unit_Cost!${col}${cf}+{g('cp')}*(Unit_Cost!${col}${cc}+Cache_Store!${col}${S[tk]}*{g('Ret')}/B4_Retain)"
             f"+{g('dp')}*Unit_Cost!${col}${cd})/1E6,{NA})", "$#,##0.0000")
        L_ = f"INDEX(B5_TaskLen,1,{k})"
        drow(f"p0{t}", f"{tn}：p 標準", f"=IF(ISNUMBER(INDEX(B5_SuccOvr_{tk},1,{k})),INDEX(B5_SuccOvr_{tk},1,{k}),1/(1+({L_}/{g('h'+str(t))})^{g('beta')}))", "0.0%")
        drow(f"p1{t}", f"{tn}：p 選定", f"=IF(ISNUMBER(INDEX(B5_SuccOvr_{tk},1,{k})),INDEX(B5_SuccOvr_{tk},1,{k}),1/(1+({L_}/({g('h'+str(t))}*{g('Hz')}))^{g('beta')}))", "0.0%")
        std = f"Harness!${c0}${R[f'c_VR200{t}std']}"
        drow(f"R{t}", f"{tn}：R（每成功任務成本 選定 ÷ 標準）",
             f"=IF(AND(ISNUMBER({g('cs'+str(t))}),ISNUMBER({std})),({g('cs'+str(t))}/{g('p1'+str(t))})/({std}/{g('p0'+str(t))}),{NA})", "0.00", key_fill=(t > 0))
    put(ws, f"A{r+1}", "讀法：R 跨過 1 的欄＝harness 由省錢轉為加成本。Luna 在長程任務 p 很低，R 主要由成功率主導", F_NOTE)
    return rows

# ---------------------------------------------------------------- Interface E, Checks, Sources, Evidence
def interface_b5(wb, start, R):
    ws = wb["Interface"]; r = start; names = []
    section(ws, r, "E. Block 5 產出（每任務層；欄 C–G＝Workload 任務；現行＝依 Tech_Registry T12 混合，基準等於標準 harness；VR200／GB300 基準成本）", 17); r += 1
    put(ws, f"A{r}", "任務　[IF_HdrTask]", F_BOLD)
    for c in TK: put(ws, f"{c}{r}", f"=Workload!{c}4", F_HLINK, wrap=True)
    names.append(("IF_HdrTask", f"Interface!$C${r}:$G${r}")); r += 1
    for name, lab, unit, ref, fmt in [("IF_HarW", "harness 混合權重 w（T12 開關 × 採用比例）", "x", "=B5_W", "0.00"),
                                      ("IF_HarProfile", "選定 harness 檔案（情境）", "", "=B5_ProfileName", None)]:
        put(ws, f"A{r}", f"{lab}　[{name}]"); put(ws, f"B{r}", unit); put(ws, f"C{r}", ref, fmt=fmt)
        names.append((name, f"Interface!$C${r}")); r += 1
    def row5(name, lab, unit, tpl, fmt, key=False):
        nonlocal r
        put(ws, f"A{r}", f"{lab}　[{name}]"); put(ws, f"B{r}", unit)
        for k, c in enumerate(TK, start=1): put(ws, f"{c}{r}", tpl.format(c=c, k=k), fmt=fmt, fill=FILL_KEY if key else None)
        names.append((name, f"Interface!$C${r}:$G${r}")); r += 1
    row5("IF_TaskLen", "任務長度（人類完成時間）", "hr", "=INDEX(B5_TaskLen,1,{k})", "0.00")
    row5("IF_TaskTokFresh", "每次嘗試 新鮮輸入 token（現行）", "tok", f"=Harness!{{c}}{R['cur_fp']}", "#,##0")
    row5("IF_TaskTokCached", "每次嘗試 快取輸入 token（現行）", "tok", f"=Harness!{{c}}{R['cur_cp']}", "#,##0")
    row5("IF_TaskTokDec", "每次嘗試 decode token（思考＋可見；現行）", "tok", f"=Harness!{{c}}{R['cur_dp']}", "#,##0")
    row5("IF_TaskTokSel", "每次嘗試 總 token（選定檔案全採用；情境）", "tok", f"=Harness!{{c}}{R['sel_tot']}", "#,##0")
    row5("IF_HarTokRatio", "總 token：選定檔案 ÷ 標準", "x", f"=Harness!{{c}}{R['tokratio']}", "0.00")
    for t, (tk, tn) in enumerate(TIERS):
        put(ws, f"A{r}", tn, F_BOLD); r += 1
        row5(f"IF_TaskSucc_{tk}", "成功率 p（現行）", "%", f"=Harness!{{c}}{R[f'p_cur{t}']}", "0.0%")
        row5(f"IF_TaskSuccSel_{tk}", "成功率 p（選定檔案全採用；情境）", "%", f"=Harness!{{c}}{R[f'p_sel{t}']}", "0.0%")
        row5(f"IF_CostSuccVR_{tk}", "每成功任務成本 — VR200（現行；經濟、基準成本、基準利用率）", "$", f"=Harness!{{c}}{R[f'cs_VR200{t}cur']}", "$#,##0.0000", key=True)
        row5(f"IF_CostSuccGB_{tk}", "每成功任務成本 — GB300（現行）", "$", f"=Harness!{{c}}{R[f'cs_GB300{t}cur']}", "$#,##0.0000")
        row5(f"IF_RevSucc_{tk}", "每成功任務營收 — OpenAI 有效單價（現行）", "$", f"=Harness!{{c}}{R[f'rs_{t}cur']}", "$#,##0.0000")
        row5(f"IF_HarR_{tk}", "R＝每成功任務成本 選定 ÷ 標準（VR200）", "x", f"=Harness!{{c}}{R[f'R_VR200{t}']}", "0.00")
    for n, ref in names: nm(wb, n, ref)
    return names

def checks_b5(wb, R, H, SH):
    ws = wb["Checks"]
    r0 = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    section(ws, r0, "Block 5 檢查（Coding agent（長程）＝Harness／Workload 欄 G）", 6)
    for c, h in zip("ABCDEF", ["項目", "本模型", "外部參照", "單位", "判讀", "來源"]): put(ws, f"{c}{r0+1}", h, F_BOLD)
    a0, a1, a2 = H["arc0"], H["arc1"], H["arc2"]
    rows = [
      ("混合權重 w（基準應為 0：L3）", "=B5_W", "0", "x", "w＝0 時 Workload＝標準 harness，Block 1–4 與 v5.8 逐格一致", "Tech_Registry T12"),
      ("Workload 總 token − Harness 標準總 token（5 任務絕對差合計）", "=" + "+".join(f"ABS(Workload!{c}35-Harness!{c}{R['std_tot']})" for c in TK), "0（w＝0 時）", "tok",
       "w＞0 時應 > 0", "Workload、Harness"),
      ("ARC 重現（token）：選定檔案 ÷ 標準 總 token", f"=Harness!G{R['tokratio']}", f"=Har_In!E{a2}", "x",
       "外部參照＝『少用約 49% token』；只看方向（ARC 為互動遊戲，不是 Workload 任務）", "S58"),
      ("ARC 重現（R）：Astra × Coding agent（VR200）", f"=Harness!G{R['R_VR2002']}", f"=(Har_In!D{a1}/Har_In!C{a1})/(Har_In!D{a0}/Har_In!C{a0})", "x",
       "外部參照＝同為最高推理的成本比 ÷ 成功率比（≈0.42）；方向一致即可", "S58"),
      ("ARC 成功率換算為時間範圍倍數（β 基準）", f"=((1/Har_In!C{a0}-1)/(1/Har_In!C{a1}-1))^(1/B5_Beta)", "1–4（Har_In 區間）", "x",
       "遠大於區間：ARC 的提升無法用 METR 型曲線表達，故 ARC 只作方向檢查、不校準時間範圍倍數", "S58、S59"),
      ("METR 隱含 β（四個模型平均）", f"=AVERAGE(Har_In!F{H['metr0']}:F{H['metr1']})", "=B5_Beta", "x", "外部參照欄＝本模型採用值", "S59"),
      ("Astra 標準 harness：Coding agent 成功率", f"=Harness!G{R['p_std2']}", "—", "%", "任務長度 8 小時 ÷ 時間範圍 16 小時", "Har_In C 節"),
      ("每成功任務：Coding agent 前緣組合（VR200）", f"=Harness!G{R['frl_VR200']}", "—", "", "層級替代：選定檔案全採用時較低層級能否勝出", "Harness H 節"),
      ("Sens_Har：Astra R 最小～最大（VR200、Coding agent）", f"=MIN(Sens_Har!C{SH['R2']}:Z{SH['R2']})", f"=MAX(Sens_Har!C{SH['R2']}:Z{SH['R2']})", "x",
       "最大值 > 1 即 harness 在區間內可能提高每成功任務成本", "Sens_Har"),
    ]
    for i, (a, b, c, d, e, f) in enumerate(rows):
        rr = r0 + 2 + i
        put(ws, f"A{rr}", a, wrap=True); put(ws, f"B{rr}", b, fmt="0.0%" if d == "%" else "#,##0.00", fill=FILL_KEY)
        put(ws, f"C{rr}", c, fmt="#,##0.00"); put(ws, f"D{rr}", d); put(ws, f"E{rr}", e, F_NOTE, wrap=True); put(ws, f"F{rr}", f, F_NOTE)

SOURCES_B5 = [
  ("S58", "ARC-AGI-3 harness 對照（ARC Prize 自測，經報導）",
   "GPT-6 Astra 標準 harness 最高推理 62.7%、$26,098；OpenAI Provider Adapter 高推理 99.9%、$18,817（另載 $19,302）；同為最高推理 98.6%、$17,332；耗時約快 3.66 倍；少用約 49% token（範圍不明）；adapter 保留不透明推理狀態並壓縮長對話，使用公開 API 功能",
   "Verified-measured（ARC Prize）／Interested-party（OpenAI 發布 99.9%）", "2026-09", "thenextweb、ibl.ai、ecosistemastartup 等轉述", "ARC Prize 原頁待核；金額衝突待核"),
  ("S59", "METR 時間範圍（Time Horizon 1.1）",
   "GPT-5 3h34／32m；GPT-5.2（high）6h34／55m；GPT-5.4（xhigh）5h42／54m；Gemini 3.1 Pro 5h50／1h30；GPT-5.6 Sol 11.3h；Claude Mythos Preview 17.4h（50%／80%）；套件 16 小時以上量不準",
   "Verified-measured（METR，獨立評測）", "2026-01～09", "Wikipedia METR 條目、futuresearch 轉述 metr.org", "metr.org 原頁待核；GPT-6 各層級未發布"),
  ("S60", "Artificial Analysis 每任務成本", "GPT-6 Astra 每任務約 $4.72、Claude Fable 5.1 約 $9.18（其 agent 評測）", "Verified-measured（獨立評測）", "2026-09",
   "sentisense 轉述", "只作對照：AA 的任務組合與 Workload 不同"),
  ("S61", "Claude Opus 5 + Nvidia 所建 harness 通關 ARC-AGI-3", "Opus 5 標準 30.2% → 全部關卡通關", "Interested-party（Nvidia）", "2026-09", "thenextweb 轉述", "交接原記 Strands；歸屬與成本待核"),
]

def sources_b5(wb):
    ws = wb["Sources"]
    r = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 1
    for i, row in enumerate(SOURCES_B5):
        for c, v in zip("ABCDEFG", row): put(ws, f"{c}{r+i}", v, wrap=True)

EVID_B5 = [
  ("E016", "2026-10-01", "ARC-AGI-3：Provider Adapter 同檔成本 0.66 倍、成功率 1.57 倍 → 每成功任務成本約 0.42 倍", "S58", "Verified-measured／Interested-party",
   "Har_In B 節檔案 2；Checks Block 5", "—（新增）", "R≈0.42", "部分採納", "v5.9", "只作方向檢查；ARC 不是 Workload 任務"),
  ("E017", "2026-10-01", "Anthropic 多代理：token 約單代理 4 倍、評測分數 +90.2%", "S25", "Interested-party", "Har_In B 節檔案 3", "—", "R≈2.0", "部分採納", "v5.9",
   "90.2% 非成功率且為內建知識，待查核；只作對照（L7）"),
  ("E018", "2026-10-01", "METR TH1.1：GPT-5.6 Sol 50% 時間範圍 11.3 小時；Mythos Preview 17.4 小時", "S59", "Verified-measured", "Har_In C 節 時間範圍", "—", "Luna 1.5／Sol 8／Astra 16 hr", "部分採納", "v5.9",
   "GPT-6 各層級未發布，以相近模型類比"),
  ("E019", "2026-10-01", "METR 50% ÷ 80% 時間範圍比 3.9–7.2 → β 約 0.70–1.0", "S59", "Verified-measured", "Har_In C 節 β", "—", "0.75（0.65–1.0）", "採納", "v5.9", ""),
  ("E020", "2026-10-01", "Opus 5 搭配 harness 由 30.2% 至全部通關（歸屬 Nvidia 或 Strands 不一）", "S61", "Interested-party", "Har_In E 節", "Strands（交接）", "Nvidia（報導）", "待查", "—", "成本未找到"),
]

def evidence_b5(wb):
    ws = wb["DB_Evidence"]
    have = {ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1)}
    r = max(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value is not None) + 1
    n = 0
    for row in EVID_B5:
        if row[0] in have: continue
        for i, v in enumerate(row): put(ws, f"{L(i+1)}{r}", v, F_IN, wrap=i in (2, 5, 10))
        ws.row_dimensions[r].height = 30; r += 1; n += 1
    return n
