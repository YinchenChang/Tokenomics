"""L 閱讀版的換版核對：確認閱讀版裡取自 Tokenomics 的數字與指定版本一致。

v18.0 起只核對 L 閱讀版（Andy 2026-10-21：m 閱讀版不再使用、日後重做；原 m 閱讀版已移到 docs/plan/_待刪除/）。
用法：
    python3 check_reading.py <Tokenomics.xlsx> <L閱讀版.xlsx>
L 閱讀版需先用 LibreOffice 重算（讀公式結果）。全部通過時結束碼 0，否則 1。
"""
import csv, math, re, sys
from pathlib import Path
import openpyxl

if len(sys.argv) != 3:
    sys.exit('用法：python3 check_reading.py <Tokenomics.xlsx> <L閱讀版.xlsx>')
TK, LW = sys.argv[1:3]
T = openpyxl.load_workbook(TK, data_only=True)
TF = openpyxl.load_workbook(TK)
fails = []


def same(a, b):
    try:
        return math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=1e-12)
    except (TypeError, ValueError):
        return a == b


def cell(ref):
    sh, c = ref.split('!')
    return T[sh.strip("'")][c].value


def named(n):
    d = TF.defined_names.get(n)
    sh, rng = list(d.destinations)[0]
    return T[sh][rng.replace('$', '')].value


def check(label, ok, detail=''):
    if not ok:
        fails.append(f'{label} {detail}')


# 0. 規格：必須有需求公式樹
L = openpyxl.load_workbook(LW, data_only=True)
check('規格', '1_公式樹' in L.sheetnames, 'L 閱讀版缺 1_公式樹')

# 1. L 閱讀版：情境摘要對具名範圍；所需 L
ws = L['3_L逐世代']
for r in range(6, 13):
    d, nm, c = ws.cell(r, 4).value, ws.cell(r, 6).value, ws.cell(r, 3).value
    check('L 摘要原值', same(d, named(nm)), f'{nm} {d} → {named(nm)}')
    check('L 摘要重算', same(c, named(nm)), f'{nm} 本簿 {c} → {named(nm)}')
# 所需 L：v21.2 起主口徑＝L1_LReqFleet_Full（全部算力持有成本）；L1_LReqFleet_IdNoK 為舊口徑對照（L9）
ft = L['1_公式樹']
for cellref, nm in (('C8', 'L1_LReqFleet_Full'), ('C19', 'L1_LReqFleet_Full'), ('C20', 'L1_LReqFleet_Full_Lo'),
                    ('C21', 'L1_LReqFleet_Full_Hi'), ('C17', 'L1_HoldFleet25'), ('C11', 'L1_HoldFleet25'),
                    ('C79', 'L1_LReqFleet_IdNoK'), ('C7', 'L1_LTotalFleetVDisc')):
    check('所需 L／持有成本', same(ft[cellref].value, named(nm)), f'1_公式樹 {cellref} {ft[cellref].value} → {nm} {named(nm)}')
check('所需 L 區間文字', ft['C13'].value == f"{named('L1_LReqFleet_Full_Lo'):.3f}–{named('L1_LReqFleet_Full_Hi'):.3f}", f"{ft['C13'].value}")
check('舊口徑區間文字', ft['C80'].value == f"{named('L1_LReqFleet_IdNoK_Lo'):.3f}–{named('L1_LReqFleet_IdNoK_Hi'):.3f}", f"{ft['C80'].value}")
check('L 區間文字', ft['C12'].value == f"{named('L1_LTotalFleetVDisc_Lo'):.3f}–{named('L1_LTotalFleetVDisc_Hi'):.3f}", f"{ft['C12'].value}")
# 三態判讀：同 Tokenomics Cap_In C242（文字全等）
check('主口徑判讀', ft['C14'].value == T['Cap_In']['C242'].value, f"{ft['C14'].value} ≠ {T['Cap_In']['C242'].value}")
# 2_輸入：全部算力持有成本（基準／低／高）
wi = L['2_輸入']
for col, nm in (('C', 'L1_HoldFleet25'), ('D', 'L1_HoldFleet25_Lo'), ('E', 'L1_HoldFleet25_Hi')):
    check('2_輸入 IN_Hold', same(wi[f'{col}34'].value, named(nm)), f'{col}34 {wi[col+"34"].value} → {nm} {named(nm)}')

# 1b. 2_輸入 第 6–34 列全部藍字輸入（基準／低端組合／高端組合）逐格對 Tokenomics 來源
# 欄序：Tokenomics 欄 → 閱讀版 (C基準, D低端組合, E高端組合)。swap＝低／高欄序對調。
def tk_row(sheet, rng):
    return [c.value for row in T[sheet][rng] for c in row]

def tk3(sheet, rng, order):
    v = tk_row(sheet, rng)
    return [v[i] for i in order]

wi = L['2_輸入']
IN_MAP = {  # 閱讀版列: (來源列表 [C,D,E])
    6: lambda: [tk_row('SRC_Demand', 'C11:C11')[0]] * 3,
    7: lambda: tk3('SRC_Demand', 'C8:E8', (0, 1, 2)),
    8: lambda: tk3('Alloc_In', 'C23:E23', (0, 2, 1)),
    10: lambda: tk3('Cap_In', 'C226:E226', (0, 1, 2)),
    11: lambda: tk3('Cap_In', 'C227:E227', (0, 1, 2)),
    12: lambda: tk3('Cap_In', 'C228:E228', (0, 1, 2)),
    13: lambda: tk3('Cap_In', 'C229:E229', (0, 1, 2)),
    14: lambda: tk3('Cap_In', 'C304:E304', (0, 2, 1)),
    15: lambda: tk3('Inputs', 'D28:F28', (1, 0, 2)),
    16: lambda: [tk_row('Cap_In', 'C303:C303')[0]] * 3,
    31: lambda: tk3('Cap_In', 'C70:E70', (0, 1, 2)),
    32: lambda: tk3('Cap_In', 'C131:E131', (0, 1, 2)),
    34: lambda: [named('L1_HoldFleet25'), named('L1_HoldFleet25_Lo'), named('L1_HoldFleet25_Hi')],
    35: lambda: tk3('Cap_In', 'C331:E331', (0, 1, 2)),   # v21.4 新增 IN_GW（說明句的 GW 端點）
}
for k, r in enumerate((18, 19, 20, 21)):  # 服務起點：Tokenomics 列 300 基準、301 較早、302 較晚；低端組合取較晚
    IN_MAP[r] = (lambda k=k: [tk_row('Cap_In', 'C300:F300')[k], tk_row('Cap_In', 'C302:F302')[k], tk_row('Cap_In', 'C301:F301')[k]])
for k, r in enumerate((22, 23, 24, 25)):
    IN_MAP[r] = (lambda k=k: [tk_row('Alloc', 'C18:F18')[k]] * 3)
for k, r in enumerate((26, 27, 28, 29)):
    IN_MAP[r] = (lambda k=k: [tk_row('Alloc', 'C153:F153')[k]] * 3)
n_in = 0
for r in range(6, 36):   # 第 35 列 IN_GW 為 v21.4 新增
    if wi.cell(r, 1).value is None or wi.cell(r, 3).value is None:
        continue  # 節標題列
    if r not in IN_MAP:
        fails.append(f'2_輸入 第 {r} 列（{wi.cell(r,1).value}）沒有對應的核對來源')
        continue
    exp = IN_MAP[r]()
    for col, e in zip((3, 4, 5), exp):
        n_in += 1
        check('2_輸入', same(wi.cell(r, col).value, e), f'{wi.cell(r,1).value} 第 {r} 列 {"CDE"[col-3]} 欄 {wi.cell(r,col).value} → Tokenomics {e}')
check('2_輸入 格數', n_in == 78, f'核對格數 {n_in}（應為 78：第 6–32 列 72 格＋IN_Hold 3 格＋IN_GW 3 格）')

# 1c. 1_公式樹 G 欄全部「Tokenomics 原值」貼上值
Lo_, Hi_ = named('L1_LTotalFleetVDisc_Lo'), named('L1_LTotalFleetVDisc_Hi')
RLo, RHi = named('L1_LReqFleet_Full_Lo'), named('L1_LReqFleet_Full_Hi')
g322 = tk_row('Cap_In', 'C322:F322')
G_MAP = {
    7: named('L1_LTotalFleetVDisc'), 8: named('L1_LReqFleet_Full'), 11: named('L1_HoldFleet25'),
    12: f'{Lo_:.3f}–{Hi_:.3f}', 13: f'{RLo:.3f}–{RHi:.3f}', 14: T['Cap_In']['C242'].value,
    17: named('L1_HoldFleet25'), 18: tk_row('SRC_Demand', 'C11:C11')[0], 19: named('L1_LReqFleet_Full'),
    20: RLo, 21: RHi,
    30: named('L1_LPrice'), 31: named('L1_LTotal'), 33: named('L1_LTotalDisc'), 35: named('L1_LTotalFleetVDisc'), 37: named('L1_LTotalFleetV'),
    42: g322[0], 45: g322[1],
    70: tk_row('Cap_In', 'C70:C70')[0] + tk_row('Cap_In', 'C131:C131')[0],
    76: tk_row('SRC_Demand', 'C8:C8')[0], 77: named('AL_SpendBasis'), 79: named('L1_LReqFleet_IdNoK'),
    80: f"{named('L1_LReqFleet_IdNoK_Lo'):.3f}–{named('L1_LReqFleet_IdNoK_Hi'):.3f}",
}
n_g = 0
for r in range(1, ft.max_row + 1):
    v = ft.cell(r, 7).value
    if v is None or r == 5:
        continue
    n_g += 1
    if r not in G_MAP:
        fails.append(f'1_公式樹 G{r} 有貼上值但沒有核對來源')
    else:
        check('1_公式樹 G 欄', same(v, G_MAP[r]), f'G{r} {v} → Tokenomics {G_MAP[r]}')
for r in G_MAP:
    check('1_公式樹 G 欄', ft.cell(r, 7).value is not None, f'G{r} 缺值')

# 1d. 3_L逐世代 各世代 L_rel 對 Cap_In 第 322–324 列（S4 基準、S6 低端組合、S7 高端組合；S3 折現無世代起點＝各世代同值）
w3 = L['3_L逐世代']
for first, crow in ((28, 322), (36, 323), (40, 324)):
    exp = tk_row('Cap_In', f'C{crow}:F{crow}')
    for k in range(4):
        check('3_L逐世代 L_rel', same(w3.cell(first + k, 17).value, exp[k]), f'第 {first+k} 列 {w3.cell(first+k,17).value} → Cap_In 第 {crow} 列 {exp[k]}')
check('3_L逐世代 L_rel S3', same(w3.cell(24, 17).value, tk_row('Cap_In', 'E322:E322')[0]), 'S3 折現')

# 2. 「差異」欄與錯誤值
for name, W in (('L', L),):
    F = openpyxl.load_workbook(LW)
    for wsf in F.worksheets:
        wv = W[wsf.title]
        for row in wsf.iter_rows():
            for c in row:
                v = wv[c.coordinate].value
                if isinstance(v, str) and re.match(r'^(#(REF|VALUE|DIV/0|NAME|N/A)|Err:)', v):
                    fails.append(f'{name} 錯誤值 {wsf.title}!{c.coordinate} {v}')
                f = c.value
                if isinstance(f, str) and f.startswith('=') and re.search(r'/[A-Z]+\d+-1', f) and isinstance(v, (int, float)) and abs(v) > 1e-9:
                    fails.append(f'{name} 差異 {wsf.title}!{c.coordinate} {v}')

# 3. 說明句數字檢查（v21.4）：盤點表 text_inventory_L.csv
sys.path.insert(0, str(Path(__file__).resolve().parent))
from text_check import check_text_numbers
n_cur, n_inv = check_text_numbers(openpyxl.load_workbook(LW), L, Path(__file__).resolve().parent / 'text_inventory_L.csv', fails)
print(f'說明句數字檢查：盤點 {n_inv} 格，其中現行值（須為公式）{n_cur} 格')

print('通過' if not fails else f'未通過 {len(fails)} 項')
for x in fails[:50]:
    print(' -', x)
sys.exit(1 if fails else 0)
