"""4-3 B 段 Q1：Sens_*／Sensitivity／*_Prod／Perf_Batch 頁中，標籤與刪除列相同的列，逐項確認不是比對基準（只讀 v13.1）。
用法：python3 docs/evidence/20261017_4-3B/02_sens_labels.py <model.xlsx>
判定：該列在複本頁有公式格，且複本頁公式讀取了『被刪除的核心格』→ 是比對基準（須保留核心列）；
      否則（複本頁同標籤列自成一體，或不讀被刪格）→ 不是比對基準。
同時列出常設測試（tests/parity/mirror_lib）會比對的條件：Sens_Perf／Sens_Train 同列同標籤、*_Prod 同格。"""
import json
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(Path(__file__).parent))
import importlib.util
spec = importlib.util.spec_from_file_location("deps01", Path(__file__).parent / "01_delete_deps.py")
d01 = importlib.util.module_from_spec(spec); spec.loader.exec_module(d01)

COPY_PAGES = ["Sens_Perf", "Sens_Train", "Sensitivity", "Sens_Har", "Sens_Rev", "Perf_Batch"] + [f"{p}_Prod" for p in
              ["Arch", "Serving", "Workload", "Training", "Cap_In", "Theory_Rev", "Amortize", "Harness", "Perf", "Unit_Cost", "Fleet_1GW"]]


def main(model):
    wb = openpyxl.load_workbook(model)
    rows = d01.delete_rows()
    dele = set()
    for sh, r, lab, rng in rows:
        for a in d01.cells_of(rng):
            dele.add((sh, a))
    out = []
    for sh, r, lab, rng in rows:
        key = str(wb[sh].cell(r, 1).value or "").strip()
        for cp in wb.sheetnames:
            if cp == sh or cp not in COPY_PAGES:
                continue
            ws = wb[cp]
            for rr in range(1, ws.max_row + 1):
                v = ws.cell(rr, 1).value
                if v is not None and str(v).strip() == key and key:
                    fcells = [c for c in ws[rr] if isinstance(c.value, str) and c.value.startswith("=")]
                    reads = [c.coordinate for c in fcells if any(f"{sh}!{a}" in c.value.replace("$", "").replace("'", "") for (s2, a) in dele if s2 == sh)]
                    out.append(dict(core=f"{sh}!{r}", label=key, copy=f"{cp}!{rr}", n_formula=len(fcells), reads_deleted=len(reads), same_row=(rr == r)))
    print("標籤相同的複本頁列", len(out))
    for o in out:
        print(o)
    bad = [o for o in out if o["reads_deleted"]]
    print("讀取被刪格的複本頁列（比對基準）", len(bad))


if __name__ == "__main__":
    main(sys.argv[1])
