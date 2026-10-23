# B 段：在 A 段 xlsx 上新增「B 段」頁並更新「證據」「待判斷」「對帳」註記。公式引用 TK 拆解頁。
import openpyxl, sys
from openpyxl.styles import Font, PatternFill, Alignment
src,out=sys.argv[1],sys.argv[2]
wb=openpyxl.load_workbook(src)
BLUE=Font(color='0000FF'); BOLD=Font(bold=True); HDR=PatternFill('solid',fgColor='DDEBF7'); SEC=PatternFill('solid',fgColor='F2F2F2')
T="'TK 拆解'!"
s=wb.create_sheet('B 段')
s['A1']='B 段補充：Oracle 並列、Epoch 12.8 矛盾評估、IREN 原文對帳、每 GPU 小時視角、CoreWeave／Nebius 財報結構'; s['A1'].font=Font(bold=True,size=13)
s['A2']='藍字＝外部輸入（出處見「證據」頁 E19 以後）；黑字＝公式。TK 值一律引用「TK 拆解」頁。'
r=[4]
def sec(t):
    r[0]+=1
    for i in range(1,10): s.cell(r[0],i).fill=SEC
    s.cell(r[0],1,t).font=BOLD; r[0]+=1
def hdr(vals):
    for i,v in enumerate(vals,1):
        c=s.cell(r[0],i,v); c.font=BOLD; c.fill=HDR; c.alignment=Alignment(wrap_text=True,vertical='top')
    r[0]+=1
def put(vals,fmts=None,blue=()):
    for i,v in enumerate(vals,1):
        if v is None: continue
        c=s.cell(r[0],i,v); c.alignment=Alignment(wrap_text=True,vertical='top')
        if i in blue: c.font=BLUE
        if fmts and i in fmts: c.number_format=fmts[i]
    r[0]+=1; return r[0]-1
N='0.00'
# ---- 1 Oracle
sec('1. Oracle–OpenAI：5 GW 與 4.5 GW 並列（$300B／5 年；匿名消息，Interested-party）')
hdr(['版本','總額 $B','年數','GW','GW 口徑假設','每年 $B','每 IT-GW·年 $B','倍數 ÷VR200','倍數 ÷GB300','倍數 ÷組合 9.35','出處'])
rows={}
def ora(label,gw,basis,src_):
    rr=r[0]
    f=f'=F{rr}/D{rr}' if basis=='IT' else f'=F{rr}/(D{rr}/{T}C9)'
    rows[label]=put([label,300,5,gw,basis,f'=B{rr}/C{rr}',f,f'=G{rr}/{T}F66',f'=G{rr}/{T}E66',f'=G{rr}/{T}G66',src_],{6:N,7:N,8:N,9:N,10:N},blue=(2,3,4))
ora('5GW 當 IT（Andy 的 12.0）',5,'IT','The Register 2025-09-11 轉述 WSJ："$300 billion ... five gigawatts ... over the course of five years ... begin in 2027"；repo docs/evidence/20261021_token帳閉合A段/供給側查證.md 第59列')
ora('5GW 當設施，換 IT',5,'設施','同上')
ora('4.5GW 當 IT',4.5,'IT','A 段所用報導（WSJ 經 Techmeme 2025-09-10 摘要："requiring 4.5 gigawatts"）；https://www.techmeme.com/250910/p43')
ora('4.5GW 當設施，換 IT',4.5,'設施','同上')
s.cell(r[0],1,'兩份報導的 GW 數不同（5 對 4.5）且 IT／設施口徑均未明：每 IT-GW·年介於 12.0–16.0，對 VR200 倍數 0.98–1.31。A 段「複算不出 12.0」是因我只用了 4.5GW 的報導；Andy 的 12.0＝5GW 當 IT。').alignment=Alignment(wrap_text=True); r[0]+=1
# ---- 2 Epoch
sec('2. Epoch／OpenAI 2025 雲端支出 ÷ GW（機隊層級對照；保留，不當單價）')
hdr(['項目','數值','單位','','標記／來源'])
a=put(['雲端算力支出 2025（Epoch 引用媒體報導）',16.0,'$B','','Interested-party；二手。Epoch 方法頁另列 16.3'],blue=(2,))
b0=put(['上線 GW 2024 年底',0.6,'GW','','官方（OpenAI 財務長揭露，見供給側查證.md 第27–31列）'],blue=(2,))
b1=put(['上線 GW 2025 年底',1.9,'GW','','同上'],blue=(2,))
lin=put(['全年平均 GW：線性內插',f'=(B{b0}+B{b1})/2','GW','','Derived（=1.25）'],{2:N})
exp_=put(['全年平均 GW：指數內插',f'=(B{b1}-B{b0})/LN(B{b1}/B{b0})','GW','','Derived（=1.13）'],{2:'0.000'})
pl=put(['每 GW·年：線性',f'=B{a}/B{lin}','$B/GW·年','','Derived（=12.8）'],{2:N})
pe=put(['每 GW·年：指數',f'=B{a}/B{exp_}','$B/GW·年','','Derived（=14.2）'],{2:N})
put(['倍數 ÷ 2025 組合 9.35：線性 ／ 指數',f'=B{pl}/{T}G66',f'=B{pe}/{T}G66','','Andy 記錄 1.37／1.52'],{2:N,3:N})
s.cell(r[0],1,'與長約 0.89–0.93 矛盾的可能原因逐項評估（「必須為真」是反推，不是證據）').font=BOLD; r[0]+=1
hdr(['可能原因','方向','量級（公式）','評估','證據狀態','能解決它的輸入'])
# what-must-be-true
w93=put(['若真實倍數為 0.93：分子須為 ($B)',None,f'=0.93*{T}G66*B{lin}','與 16.0 的差即「必須是非 GPU 或其他」的部分','Derived',None],{3:N})
w93b=r[0]-1
s.cell(w93b,2,'分子須下修'); 
put(['  須排除的支出 ($B)／占比',None,f'=B{a}-C{w93b}',f'=C{r[0]}/B{a}','Derived',None],{3:N,4:'0%'})
put(['  同上，用指數 GW 1.13',None,f'=0.93*{T}G66*B{exp_}',f'=1-C{r[0]}/B{a}','Derived（此列 D＝須排除占比）',None],{3:N,4:'0%'})
put(['(a) 非 GPU 雲端服務與研究實驗（儲存、CPU、網路出口、推論以外的雲端帳單）','降低倍數','上段「須排除」約 $5.1–6.2B（占 32–39%）即所需規模','合理但無數字：OpenAI 營業成本 75 億 vs 推論 84 億已顯示口徑不整齊；若這類占 20–35% 即可解釋','搜尋未找到分項','OpenAI 雲端帳單分項（GPU 時數 vs 其他服務）'])
put(['(b) 2025 年初 Hopper 舊約','提高（Hopper 租金 > TK Hopper 持有成本）',f'=B{pl}*1000000000/({T}G35*8760)','隱含每 GPU 小時 ($)；若 2025 Hopper 約定價 $2.0–2.5 即與 12.8 一致；GB300 長約隱含 $2.27，兩者同量級→租金不隨世代上升，TK 持有成本隨世代上升','Derived','OpenAI 各供應商實付 $/GPU-hr（已搜尋，未找到）'])
put(['(c) 雲端商加價（Azure 轉售）','提高','倍數 ×(1+加價)：1.0–1.3（Assumed）','IREN 賣給 Microsoft 是批發價；Microsoft 轉給 OpenAI 若加價 20% → 0.89×1.2＝1.07；與 Microsoft–OpenAI 關係（持股、額度）交織','Assumed','OpenAI 向 Microsoft 的實付價（找不到）'])
put(['(d) GW 口徑','IT 較低→倍數提高；設施較低→倍數降低','若 1.25 為設施：IT＝1.25/PUE',f'=B{a}/(B{lin}/{T}C9)','Derived（本列 D＝每 IT-GW）','OpenAI 財務長 GW 的定義（IT 或設施）'],{4:N})
put(['(e) 非現金額度','降低現金口徑，但「支出」若含額度則提高','無數字','The Register 提醒 Microsoft 端數字可能含額度','二手','OpenAI 現金支付表'])
put(['(f) 時間窗與預付','分子含預留未上線容量→提高','無數字','分母用年底上線 GW 內插，未含「已付費但未上線」容量','Assumed','各供應商合約生效與上線日'])
put(['綜合','',None,'(a) 與 (b)(c) 方向相反且量級都能到 ±20–30%；現有資料無法分離。→ 列為「不可分離」，缺 OpenAI 雲端帳單分項與實付 GPU 小時單價。','',''])
# ---- 3 IREN
sec('3. IREN 原文對帳（sec.gov 8-K 附件 99-1、99-2；chat 端 WebFetch 取得，[Verified]；IREN＝Interested-party）')
hdr(['項目','IREN 值','單位','TK 值','IREN ÷ TK','說明'])
g=put(['GPU 與周邊 capex（含部署服務、伺服器、InfiniBand、線材、軟體授權）',5.8,'$B',None,None,'原文'],blue=(2,))
mw=put(['容量（critical IT load）',0.2,'GW IT',None,None,'原文'],blue=(2,))
perg=put(['  GPU 與周邊 capex 每 IT-GW',f'=B{g}/B{mw}','$B/GW',f'={T}E39',f'=B{r[0]}/D{r[0]}','TK GB300 IT 設備（機架＋網路＋儲存）'],{2:N,4:N,5:N})
body_lo=put(['資料中心本體 低',9,'$M/MW',f'={T}E46*1000/{T}E7',f'=B{r[0]}/D{r[0]}','Tier III 同級；TK 廠房 capex 每 MW＝12.67（土地、建物、電力、機械、光纖保全、預備費）'],{2:'0',4:N,5:N},blue=(2,))
body_hi=put(['資料中心本體 高',11,'$M/MW',f'={T}E46*1000/{T}E7',f'=B{r[0]}/D{r[0]}',''],{2:'0',4:N,5:N},blue=(2,))
cl=put(['100MW 超級叢集架構部署',3,'$M/MW',None,None,'是否已含在 5.8B 的 GPU 周邊內：原文未明（不可分離）'],blue=(2,))
ac=put(['2026 年交付 200MW 的加速費用',2,'$M/MW',None,None,'一次性，非穩態'],blue=(2,))
put(['每 MW 全部 capex（不含 GPU）：低／高','=B%d+B%d+B%d'%(body_lo,cl,ac),'$M/MW','=B%d+B%d+B%d'%(body_hi,cl,ac),None,'低 14、高 16，與原文一致'],{2:'0',4:'0'})
s.cell(r[0],1,'TK 廠房 capex 是否含網路與叢集部署？—不含。DC_Cost 第25–31列只有土地、建物、電力設施、機械散熱、光纖保全、預備費；網路在第21列（機架價 ×10%）、儲存在第23列，都屬 IT 設備；機架價含機架內冷卻。').alignment=Alignment(wrap_text=True); r[0]+=1
hdr(['配置方式（資本支出配置有兩種讀法）','IREN IT 側 $B/GW','IREN 廠房側 $B/GW','TK IT','TK 廠房','IT 比','廠房比','廠房 capex 若改為 IREN 值，TK GB300 持有成本變化 $B/GW·年'])
k=f'(({T}E54+{T}E62+0.005*{T}E46)/{T}E46)'
rA=put(['A. 叢集部署歸 IT 側（29.0＋3）；廠房＝本體 10（中點）',f'=B{perg}+B{cl}',10,f'={T}E39',f'={T}E46','=B%d/D%d'%(r[0],r[0]),'=C%d/E%d'%(r[0],r[0]),f'=(C{r[0]}-E{r[0]})*{k}'],{2:N,3:N,4:N,5:N,6:N,7:N,8:N},blue=(3,))
rB=put(['B. 叢集部署歸廠房側（29.0；10＋3）',f'=B{perg}',13,f'={T}E39',f'={T}E46','=B%d/D%d'%(r[0],r[0]),'=C%d/E%d'%(r[0],r[0]),f'=(C{r[0]}-E{r[0]})*{k}'],{2:N,3:N,4:N,5:N,6:N,7:N,8:N},blue=(3,))
rC=put(['C. 全包（A 段用法）：29.0＋15',f'=B{perg}',15,f'={T}E39',f'={T}E46','=B%d/D%d'%(r[0],r[0]),'=C%d/E%d'%(r[0],r[0]),f'=(C{r[0]}-E{r[0]})*{k}'],{2:N,3:N,4:N,5:N,6:N,7:N,8:N},blue=(3,))
s.cell(r[0],1,'讀法：IT 設備與 IREN 一致（A 配置 0.99、B/C 配置 0.90）；廠房本體 TK 12.67 高於 IREN 9–11（高 15–41%），但加上叢集部署 3 後（12–14）TK 落在區間內。→ A 段「廠房可能偏低」應撤回：以原文的本體 9–11 比，TK 偏高；以含叢集部署比，TK 持平。總 capex 不變（約 44），不影響前述 0.98 倍。').alignment=Alignment(wrap_text=True); r[0]+=1
put(['營運費用下界','=0.15*9.7','$B/GW·年',f'={T}E65',f'=B{r[0]}/D{r[0]}','15%×9.70＝1.455；扣電費、維修、其他直接營運費，不含企業管理費、財務費用、非現金項目 → 是下界（原文有依據）；TK 1.98 含站點營運與財產稅'],{2:'0.000',4:N,5:N})
# ---- 4 per GPU hour
sec('4. 每 GPU 小時視角：長約租金是否不隨世代上升，而 TK 持有成本隨世代上升？（Derived）')
hdr(['世代','TK 每 GW GPU 數','TK 持有成本 $/GPU-hr','外部租金 $/GPU-hr（隱含）','外部 ÷ TK','外部來源'])
def ph(lab,col,ext,src_,fmt='0.00'):
    rr=r[0]; return put([lab,f'={T}{col}35',f'={T}{col}76',ext.replace('#',str(rr)),f'=D{rr}/C{rr}',src_],{2:'#,##0',3:N,4:N,5:N})
ph('Hopper：Silicon Data H100 指數（2026-10）','C','2.53','Interested-party（指數商）')
ph('Hopper：Ornn H100 成交價（2026-10）','C','3.15','Interested-party')
ph('機隊層級（2025 組合，含 Hopper）：Andy 12.8 隱含','G',f'=B{pl}*1000000000/(B#*8760)','Derived：機隊層級，非單一世代')
ph('GB300：IREN 9.70','E',f'=9.7*1000000000/(B#*8760)','Interested-party（8-K）')
ph('GB300：Nscale Narvik（第三方推算）','E','=6200000000/5/52000/8760','Derived（單一第二手）')
ph('VR200：Oracle 12.0（5GW 當 IT）','F',f'=G{rows["5GW 當 IT（Andy 的 12.0）"]}*1000000000/(B#*8760)','Interested-party（匿名消息）')
s.cell(r[0],1,'觀察：外部隱含租金約 $2.3–2.7／GPU-hr，幾乎不隨世代變；TK 持有成本由 Hopper 1.54 升到 VR200 4.78。→ 倍數隨世代由 >1（Hopper、舊世代組合）降到 <1（GB300）。這同時解釋「機隊層級 1.37」與「GB300 長約 0.89」，但依賴「租金不隨世代變」，且資料點少。反例：Oracle→VR200 隱含 $4.7（匿名消息、且假設交付 VR200；若交付 GB300 則隱含約 $2.8，與「租金不隨世代變」一致）。CoreWeave 2026Q2 隱含 $1.5（見第5節）則低於此區間，與假設不合，原因不可分離。').alignment=Alignment(wrap_text=True); r[0]+=1
# ---- 5 CRWV / Nebius
sec('5. CoreWeave 2026Q2（10-Q 部分）與 Nebius 2026Q2（6-K）財報結構（百萬美元）')
hdr(['項目','CoreWeave Q2 2026','Nebius Q2 2026','說明'])
rv=put(['營收',2575,582.3,'10-Q／6-K'],blue=(2,3))
cr=put(['營業成本（cost of revenue）',879,133.6,''],blue=(2,3))
ti=put(['技術與基礎設施',1507,None,'CoreWeave 項目'],blue=(2,))
dp=put(['折舊',1400,259.7,'CoreWeave：是否含在技術與基礎設施內，部分 10-Q 未明（假設含）'],blue=(2,3))
put(['營業成本 ÷ 營收',f'=B{cr}/B{rv}',f'=C{cr}/C{rv}',''],{2:'0%',3:'0%'})
put(['折舊 ÷ 營收',f'=B{dp}/B{rv}',f'=C{dp}/C{rv}',''],{2:'0%',3:'0%'})
put(['營收 − 營業成本 − 技術與基礎設施（SG&A 前）',f'=B{rv}-B{cr}-B{ti}',None,'若折舊含於技術與基礎設施：營業利益率'],{2:'0'})
put(['  ÷ 營收',f'=B{r[0]-1}/B{rv}',None,''],{2:'0.0%'})
put(['TK 對照：廠房側成本占持有成本（廠房年金＋電費＋廠房維護＋站點營運＋廠房財產稅）',f'=({T}G54+{T}G60+{T}G62+{T}G63+0.005*{T}G46)/{T}G66',None,'CoreWeave 營業成本含租賃機房與電力 → 與此對照（34% vs 約 30%），量級吻合；營業成本內容未逐項揭露（Derived）'],{2:'0%'})
put(['TK 對照：IT 帳面折舊 ÷ 持有成本（組合）',f'={T}G39/{T}E11/{T}G66',None,'CoreWeave 折舊占營收 54%；TK 組合 IT 折舊約 47%（以 9.35 為營收）；差距可由新世代較貴、含租賃改良折舊解釋（Derived）'],{2:'0%'})
put(['active power（GW）：Q1 末／Q2 末','=1.0','=1.5','Q2 末 1.5 GW、Q2 淨增「近 500MW」（搜尋摘要，二手）；Q1 末由此反推約 1.0（與「Q1 突破 1GW」一致）；IT 或設施口徑未取得'],{2:'0.0',3:'0.0'})
put(['平均 active power Q2（GW）',f'=(B{r[0]-1}+C{r[0]-1})/2',None,''],{2:'0.00'})
put(['營收年化 ÷ 平均 active power（$B/GW·年）',f'=B{rv}*4/1000/B{r[0]-1}',None,'Q4 2025 為 8.73（A 段）；新容量爬坡拉低比值'],{2:N})
put(['  ÷ TK 組合 9.35',f'=B{r[0]-1}/{T}G66',None,''],{2:N})
put(['  隱含 $/GPU-hr（用 TK 組合每 GW GPU 數）',f'=B{r[0]-2}*1000000000/({T}G35*8760)',None,'明顯低於 H100 指數 2.5–3.2：爬坡落差（電通但未計費）＋舊約低價＋口徑；不可分離'],{2:N})
put(['Nebius 設備耐用年限','4 → 5 年（2026Q1 起）','','折舊率下降約 20%，抬高帳面利潤，不改現金成本；MW 定義未揭露'])
put(['未取得（待 Drive inbox 的 10-Q 全文）','active power 定義、GPU 耐用年限、債務利率','','→ 有了債務利率才能把「出租方折現率 3–8%」落到 CoreWeave 實值'])
for c,w in zip('ABCDEFGHIJK',[58,16,16,22,22,18,14,22,16,16,80]): s.column_dimensions[c].width=w
# ---- 證據
z=wb['證據']
EV=[('E19','Oracle–OpenAI（5GW 版本）','The Register 2025-09-11 轉述 WSJ："$300 billion ... five gigawatts ... over the course of five years ... begin in 2027"；GW 口徑未明。','https://www.theregister.com/2025/09/11/openai_reportedly_on_the_hook/','2025-09-11（chat 端提供）','Interested-party（匿名消息）','已找到（chat 端）','與 E8 的 4.5GW 版本並列於 B 段頁'),
('E20','Epoch 12.8 的算式','$16B ÷ 1.25 GW；1.25＝(0.6+1.9)/2，為 OpenAI 財務長揭露 2024、2025 年底上線 GW 的線性內插；指數內插 1.13。','docs/evidence/20261021_token帳閉合A段/供給側查證.md（第27–31列）','2026-10-21','官方（GW 端點）＋Derived','已找到','1.25 非 Epoch 揭露的 GW，是內插值；A 段「出處未知」更正為此'),
('E21','IREN 8-K 附件 99-1、99-2（原文）','$9.7B／5 年、20% 預付款（抵第 3–5 年）；200MW critical IT load；GPU 與周邊 capex $5.8B（含部署、伺服器、InfiniBand、線材、軟體授權）；每 MW 全部 capex $14–16M（IT 口徑）＝資料中心本體 9–11＋100MW 超級叢集架構部署 3＋加速費用 2；年化營收 $194M×… 即 $1.94B；專案 EBITDA 率 85%（扣電費、維修、其他直接營運費後；扣企業管理費、財務費用、非現金項目前）。','sec.gov/Archives/edgar/data/1878848/000114036125040072/（ex99-1、ex99-2）','2026-10-23（chat 端 WebFetch）','Verified（原文）；IREN＝Interested-party','已找到（原文）','取代 A 段 E2 的二手摘要；A 段把 14–16 當「資料中心基礎設施」是誤讀，見 B 段頁第3節'),
('E22','CoreWeave 2026Q2 10-Q（部分）','季營收 2,575、營業成本 879、技術與基礎設施 1,507、折舊 1,400（百萬美元）；RPO 1,037 億，41% 於 24 個月內認列。','chat 端提供（10-Q 部分）','2026-10-23','Verified（公司申報）','部分','active power 定義、GPU 耐用年限、債務利率待 10-Q 全文（Drive inbox）'),
('E23','Nebius 2026Q2 6-K','營收 582.3、營業成本 133.6、折舊 259.7（百萬美元）；伺服器與網路設備耐用年限 2026Q1 起 4→5 年。MW 定義未揭露。','chat 端提供','2026-10-23','Verified（公司申報）','已找到；MW 定義已搜尋，未找到','—'),
('E24','CoreWeave active power 2026Q1／Q2','Q1：營收約 $2.08B、active power 突破 1GW、contracted >3.5GW；Q2：active power 增近 500MW 至 1.5GW，contracted 約 3.7GW（6 月底）、4.2GW（8 月 11 日），backlog $104.2B。','https://www.businesswire.com/news/home/20260507558197/en/ ；Q2 新聞稿（搜尋摘要）','2026-05／08','Verified（公司揭露，二手摘要）','已找到（摘要）','Q1 末 1.0GW 是反推；口徑未明')]
for i,rw in enumerate(EV):
    for j,x in enumerate(rw,1):
        c=z.cell(z.max_row+1 if j==1 else z.max_row,j,x); c.alignment=Alignment(wrap_text=True,vertical='top')
# 更正 E2、E8 註記
for row in z.iter_rows(min_row=2):
    if row[0].value=='E2': row[7].value=str(row[7].value)+'｜B 段更正：原文顯示 $14–16M/MW 為「全部 capex（IT 口徑）」＝本體 9–11＋叢集部署 3＋加速 2，不是資料中心本體。'
    if row[0].value=='E8': row[7].value=str(row[7].value)+'｜B 段補：另有 5GW 版本（E19），12.0＝300÷5÷5。'
    if row[0].value=='E17': row[6].value='已解決（B 段）：1.25＝(0.6+1.9)/2 內插，見 E20'
# 對帳 註記
q=wb['對帳']
q['H7']='B 段更正：原文 $14–16M/MW 是不含 GPU 的全部 capex（本體 9–11＋叢集部署 3＋加速 2）。C7=15 沿用為「全包」；與 TK 廠房的 like-for-like 比較見「B 段」頁第3節。'
q['H9']='是下界：15%×9.70；原文確認扣電費、維修、其他直接營運費，不含企業管理費、財務費用、非現金項目。'
# 待判斷
p=wb['待判斷']
dec={2:'決定（Andy，2026-10-23）：同世代對比；9.35 只用於機隊層級',3:'決定：不改模型；同列 10% 與出租方 3–8%',4:'決定：先不改；B 段原文已取得：IREN 營運下界 1.455 有依據',5:'決定：先不改；B 段原文：IREN 本體 9–11、含叢集 12–14；TK 12.67 不含叢集部署，結論見 B 段頁第3節（廠房偏低的說法撤回）',6:'決定：保留作機隊層級對照；算式 160÷1.25，1.25 為內插（E20）',7:'決定：已有算式 300÷5÷5＝12.0（5GW）；B 段並列 5GW 與 4.5GW',8:'決定：不作為 2025 年持有成本偏低的證據',9:'決定：讀原文；已取得 IREN、CoreWeave 部分、Nebius；CoreWeave 10-Q 全文待 Drive inbox'}
for i in range(2,p.max_row+1):
    n=p.cell(i,1).value
    mp={1:2,2:3,3:4,4:5,5:6,6:7,7:8,8:9}
    if n in mp:
        p.cell(i,7,dec[mp[n]]); p.cell(i,7).alignment=Alignment(wrap_text=True,vertical='top')
p.cell(1,7,'決定／B 段狀態').font=BOLD; p.cell(1,7).fill=HDR; p.column_dimensions['G'].width=60
nq=[(9,'Epoch 機隊對照的矛盾（1.37 對 0.89–0.93）要不要進一步拆？','維持「不可分離」，並以「每 GPU 小時租金不隨世代變」作為最可能的解釋方向，列 B 段頁第4節','要求 OpenAI 雲端帳單分項','影響機隊層級 9.35 的使用方式','B 段頁第2、4節'),
(10,'CoreWeave 取得 10-Q 全文後，是否用其債務利率與折舊年限回填「出租方折現率」？','是（待 inbox）','維持 3–8% 區間','把 WACC 差異從假設變成實值','B 段頁第5節'),
(11,'IREN 叢集部署 3 $M/MW 是否已在 5.8B 內？','原文未明，A/B 兩配置並列','向 IREN 簡報原頁確認','不影響總 capex；影響 IT／廠房側拆分','B 段頁第3節')]
for rw in nq:
    for j,x in enumerate(rw,1):
        c=p.cell(p.max_row+1 if j==1 else p.max_row,j,x); c.alignment=Alignment(wrap_text=True,vertical='top')
D2={9:'決定（Andy）：維持「不可分離」。補記：OpenAI 揭露支出分項，或揭露對 Azure 的實付價，才能分開。',10:'決定（Andy）：要；只寫進查證檔，不改模型（Tokenomics 維持 10%）。',11:'決定（Andy）：A／B 兩種配置並列即可，不再追；影響 GB300 持有成本 −0.38 至 +0.33（約 ±3%），不改結論。'}
for i in range(2,p.max_row+1):
    if p.cell(i,1).value in D2:
        p.cell(i,7,D2[p.cell(i,1).value]); p.cell(i,7).alignment=Alignment(wrap_text=True,vertical='top')
wb.save(out); print('saved')
