"""4-3 B 段 Q3（只查明，不改數值）：Cap_In 第 121 列（m 下限，上限讀法）與 L1_MFleet_Lo 的定義與落差。
用法：python3 04_mfleet_floor.py model/20261017_Tokenomics_v14.0.xlsx
輸出：Cap_In 第 108–121 列 C／D／E 三欄的公式與值、L1 第 78 列的連結、以及三種處理（截在下限／修正下限定義／改寫標籤）下 L1_MFleet 區間的新值。
讀 LibreOffice 重算後的值；不寫任何檔案。"""
import sys

import openpyxl

model = sys.argv[1]
wf = openpyxl.load_workbook(model)
wv = openpyxl.load_workbook(model, data_only=True)
ci_f, ci_v = wf["Cap_In"], wv["Cap_In"]

print("== Cap_In 第 108–121 列（C 基準、D token 取高、E token 取低）")
for r in range(108, 122):
    a = str(ci_v.cell(r, 1).value)[:70]
    if all(ci_f.cell(r, c).value is None for c in (3, 4, 5)):
        continue
    print(f"[{r}] {a}")
    for c in "CDE":
        print(f"     {c}: {str(ci_f[f'{c}{r}'].value)[:150]}  ->  {ci_v[f'{c}{r}'].value}")

print("\n== L1 第 78 列（L1_MFleet／_Lo／_Hi）")
l1_f, l1_v = wf["L1"], wv["L1"]
for c in "DEF":
    print(f"  L1!{c}78: {l1_f[f'{c}78'].value}  ->  {l1_v[f'{c}78'].value}")

g = lambda a: ci_v[a].value
m_base, m_lo, m_hi = g("C119"), g("D119"), g("E119")
fl_base, fl_lo, fl_hi = g("C121"), g("D121"), g("E121")
print("\n== 比較")
print(f"  L1_MFleet 基準／_Lo／_Hi     : {m_base:.4f} / {m_lo:.4f} / {m_hi:.4f}")
print(f"  Cap_In 121（m 下限）C／D／E   : {fl_base:.4f} / {fl_lo:.4f} / {fl_hi:.4f}")
print(f"  基準 m ÷ 基準下限            : {m_base / fl_base:.4f}（>1 ＝基準在下限之上）")
print(f"  _Lo ÷ 基準下限（0.184 對 0.301）: {m_lo / fl_base:.4f}（<1 ＝ _Lo 低於基準欄的下限）")
print(f"  _Lo ÷ 同欄下限（D119 對 D121）: {m_lo / fl_lo:.4f}（<1 ＝ _Lo 低於其自身欄的下限）")
t116, t118 = [g(f"{c}116") for c in "CDE"], [g(f"{c}118") for c in "CDE"]
print(f"  機隊付費 token（點估計）116 C／D／E: {t116}")
print(f"  機隊付費 token 上限 118 C／D／E    : {t118}")
print(f"  點估計 ÷ 上限                  : {[round(a / b, 4) for a, b in zip(t116, t118)]}（>1 ＝點估計超過同欄上限）")

sg = [g(f"{c}109") for c in "CDE"]
uc = [g(f"{c}114") for c in "CDE"]
pr = g("C115")
print("\n== 落差分解（0.301 → 0.184 的倍數 ＝ 0.3010 ÷ 0.1839）")
print(f"  總倍數 {fl_base / m_lo:.4f} ＝ D116 ÷ C118 ＝ {t116[1] / t118[0]:.4f}")
print(f"  ① 服務 GW（第 109 列）D ÷ C ＝ {sg[1]:.4f} ÷ {sg[0]:.4f} ＝ {sg[1] / sg[0]:.4f}（token 取高情境的服務 GW 較大）")
print(f"  ② 利用率修正：點估計 D114 {uc[1]:.4f} ÷ 上限讀法 C115 {pr:.4f} ＝ {uc[1] / pr:.4f}（點估計用 AL_UtilActual_Hi，上限讀法用固定的付費產能比，兩者不是同一把尺）")
print(f"  ①×② ＝ {sg[1] / sg[0] * uc[1] / pr:.4f}（與總倍數相符）")

print("\n== 三種處理對 L1_MFleet 區間的影響（只算，不寫入）")
print(f"  現況                       : [{m_lo:.4f}, {m_base:.4f}, {m_hi:.4f}]")
print(f"  (a) 截在下限（Lo＝MAX(Lo, 基準下限 {fl_base:.4f})）: [{max(m_lo, fl_base):.4f}, {m_base:.4f}, {m_hi:.4f}]")
print(f"  (b) 修正下限定義（下限改取各欄自身下限，Lo＝MAX(Lo, D121 {fl_lo:.4f})）: [{max(m_lo, fl_lo):.4f}, {m_base:.4f}, {m_hi:.4f}]")
print("  (c) 改寫標籤：數值不變，Cap_In 121 改稱『基準欄的上限讀法 m（不是 L1_MFleet 區間的下界）』")
