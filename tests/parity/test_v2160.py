"""v21.6 專屬測試（工作單 docs/workorders/20261025_v21.6_NonNV持有比TPU_Trn.md r0）。讀 LibreOffice 重算後的值。

  1. 具名範圍（1,821 個）的值與封存 v21.5 相同，只有 IF_NonNV_Hold_TPU、IF_NonNV_Hold_Trn（各 3 格）與版本字串 RNG_Version／RNG_Current 例外；名稱沒有增減。
  2. W1：NonNV E6:G6 ＝ 0.65／0.70／0.90、E7:G7 ＝ 0.45／0.55／0.70；產出比 B:D 不變；H:J 公式照舊；L6／L7 說明改寫、原文逐字保留於「｜原：」之後。
  3. W2：DB_Evidence E409（接 E408）、Decisions SL63-01、README B5 開頭。
  4. GOV 三項計數 0／134／182；RNG_StaleFlag 一致；全部公式格無錯誤值。
  5. 逐格比較 v21.5 → v21.6：數值差異只落在工作單第 0 節清單；文字／公式差異只落在 W1–W2 格與版本字串。
"""
import csv
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import range_boundaries

from conftest import model_path
from v2160_allow import NUM_DIFFS      # 工作單第 0 節清單（數值差異）；舊版逐格比較測試也用它扣除

REPO = Path(__file__).resolve().parents[2]
V215 = REPO / "model" / "archive" / "20261025_Tokenomics_v21.5.xlsx"
VERSION_NAMES = {"RNG_Version", "RNG_Current"}
CHANGED_NAMES = VERSION_NAMES | {"IF_NonNV_Hold_TPU", "IF_NonNV_Hold_Trn"}
TAIL = "持有比為慣例值（Block 1『接近 1』），未推導；找到 TPU v7、Trainium3 機架功率後以每 kW 設備價法重推（Andy 2026-10-23 Q2）。"


def _flat(v):
    if isinstance(v, (list, tuple)):
        return [x for r in v for x in _flat(r)]
    return [v]


def _num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def _eq(a, b):
    if a in (None, "") and b in (None, ""):
        return True
    if _num(a) and _num(b):
        return abs(a - b) <= 1e-9 * max(abs(a), abs(b), 1e-300) + 1e-12
    return a == b


def _arch_value(pw, d):
    sh, rg = d.attr_text.rsplit("!", 1)
    c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
    return [pw[sh.strip("'")].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


@pytest.fixture(scope="module")
def wbf():
    return openpyxl.load_workbook(model_path())


@pytest.fixture(scope="module")
def pw():
    return openpyxl.load_workbook(V215, data_only=True)


def _row(ws, sid):
    rr = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == sid]
    assert len(rr) == 1, (sid, rr)
    return rr[0]


# ---------------------------------------------------------------- 1. 具名範圍值
def test_v2160_names_changed_only_expected(lo, pw, wbf):
    v = lo.view({})
    bad, n = [], 0
    for name, d in wbf.defined_names.items():
        if "!" not in d.attr_text or "," in d.attr_text or name not in pw.defined_names:
            continue
        want = _arch_value(pw, pw.defined_names[name])
        got = _flat(v.get_name(name))
        n += 1
        if not (len(want) == len(got) and all(_eq(a, b) for a, b in zip(want, got))):
            bad.append(name)
    assert n > 1500, n
    assert sorted(bad) == sorted(CHANGED_NAMES), bad
    assert set(wbf.defined_names) == set(pw.defined_names)      # 名稱沒有增減
    assert [round(x, 12) for x in _flat(v.get_name("IF_NonNV_Hold_TPU"))] == [0.65, 0.70, 0.90]
    assert [round(x, 12) for x in _flat(v.get_name("IF_NonNV_Hold_Trn"))] == [0.45, 0.55, 0.70]
    assert str(v.get_name("RNG_Version")).endswith("v21.6") and str(v.get_name("RNG_Current")).endswith("v21.6")


# ---------------------------------------------------------------- 2. W1 NonNV
def test_v2160_nonnv_rows(lo, wbf, pw):
    nn = wbf["NonNV"]
    assert nn["A6"].value == "Google TPU v7 Ironwood" and nn["A7"].value == "AWS Trainium3"
    assert [nn[f"{c}6"].value for c in "EFG"] == [0.65, 0.7, 0.9] and [nn[f"{c}7"].value for c in "EFG"] == [0.45, 0.55, 0.7]
    assert [nn[f"{c}6"].value for c in "BCD"] == [0.6, 0.8, 1.1] and [nn[f"{c}7"].value for c in "BCD"] == [0.3, 0.5, 0.8]      # 產出比不改（待 PR #172 Q1）
    for r in (6, 7):
        assert (nn[f"H{r}"].value, nn[f"I{r}"].value, nn[f"J{r}"].value) == (f"=F{r}/C{r}", f"=E{r}/D{r}", f"=G{r}/B{r}") and nn[f"K{r}"].value == "Assumed"
    v = lo.view({})
    assert abs(v.get("NonNV", "H6") - 0.7 / 0.8) < 1e-12 and abs(v.get("NonNV", "I6") - 0.65 / 1.1) < 1e-12 and abs(v.get("NonNV", "J6") - 0.9 / 0.6) < 1e-12
    assert abs(v.get("NonNV", "H7") - 0.55 / 0.5) < 1e-12 and abs(v.get("NonNV", "I7") - 0.45 / 0.8) < 1e-12 and abs(v.get("NonNV", "J7") - 0.7 / 0.3) < 1e-12
    pn = pw["NonNV"]
    p5 = openpyxl.load_workbook(V215)["NonNV"]
    for r in range(5, 12):      # 其他列（5、8–11）的值與 v21.5 相同
        if r in (6, 7):
            continue
        for c in "BCDEFGHIJ":
            assert _eq(pn[f"{c}{r}"].value, v.get("NonNV", f"{c}{r}")), (c, r)
    for r in (6, 7):
        l = nn[f"L{r}"].value
        assert "持有比 v21.6 依每 kW 設備價推導（PR #172；TPU 每顆約 $21／25／35K、Trn3 每架 $1.0／1.5／2.2M；Assumed）" in l
        assert "產出比仍為慣例值，待精度口徑判斷（PR #172 Q1：各晶片最高原生精度時 TPU 0.25／0.35／0.40、Trn3 0.06／0.14／0.25；FP8 對齊時 TPU 0.48／0.66／0.84、Trn3 0.13／0.28／0.50）" in l
        assert l.endswith("｜原：" + TAIL) and l.count("｜原：") == 1      # 修訂九：原文逐字保留
    assert nn["L6"].value.startswith("自用為主，無公開同口徑實測｜") and nn["L7"].value.startswith("OpenAI v0.5 以 2GW Trainium 合約為輸入｜")
    for r in (5, 8, 9):      # 其他列說明不動
        assert nn[f"L{r}"].value == p5[f"L{r}"].value


def test_v2160_interface_gov_mirrors(lo):
    v = lo.view({})
    got = {f"{c}{r}": v.get("Interface", f"{c}{r}") for r in (375, 376) for c in "CDE"}
    assert got == {"C375": 0.65, "D375": 0.7, "E375": 0.9, "C376": 0.45, "D376": 0.55, "E376": 0.7}
    gm = {c: v.get("Gov_Map", c) for c in ("K396", "L396", "Q396", "Q399", "Q400", "K402", "L402", "Q402", "Q405", "Q406")}
    assert gm == {"K396": 0.65, "L396": 0.9, "Q396": 0.7, "Q399": 0.65, "Q400": 0.9, "K402": 0.45, "L402": 0.7, "Q402": 0.55, "Q405": 0.45, "Q406": 0.7}


# ---------------------------------------------------------------- 3. W2 記錄與文字
def test_v2160_evidence_decisions_readme(wbf):
    ev = wbf["DB_Evidence"]
    r = _row(ev, "E409")
    assert r == 333 and ev.cell(r - 1, 1).value == "E408"
    assert ev.cell(r, 2).value == "2026-10-25" and ev.cell(r, 10).value == "v21.6" and ev.cell(r, 12).value == "已處理"
    assert "docs/evidence/20261025_NonNV_TPU_Trainium/" in ev.cell(r, 4).value and "PR #172" in ev.cell(r, 4).value
    assert ev.cell(r, 5).value == "Interested-party（Google、AWS 自揭規格）＋Analogy（KGI、SemiAnalysis 價格）"
    t = " ".join(str(ev.cell(r, c).value) for c in range(1, 19))
    for key in ("持有比 TPU 0.65／0.70／0.90、Trn3 0.45／0.55／0.70（採用，v21.6 W1）", "產出比 TPU 0.25／0.35／0.40、Trn3 0.06／0.14／0.25（各晶片最高原生精度；未採用，待 Q1）",
                "每 token 成本比基準 TPU 約 2.0、Trn3 約 3.9（現值 1.0、1.4）", "NonNV E6:G7", "Interface 第 375–376 列", "TPU 效率 f 基準 0.856 為 Assumed（InferenceX 單一來源只作佐證）",
                "Trn3 機架倍數 1.35–1.75（基準 1.52）", "AWS 1 kW 為 DCD 推論", "SemiAnalysis 只作佐證"):
        assert key in t, key
    d = wbf["Decisions"]
    r = _row(d, "SL63-01")
    assert r == 335 and d.cell(r - 1, 1).value == "SL62-01" and d.cell(r, 2).value == "v21.6" and d.cell(r, 7).value == "v21.6 已建"
    txt = d.cell(r, 4).value + d.cell(r, 6).value
    for key in ("0.7／0.8／1.0 → 0.65／0.70／0.90", "0.6／0.7／0.9 → 0.45／0.55／0.70", "產出比 B:D 不改", "待精度口徑判斷", "IF_NonNV_Hold_TPU、IF_NonNV_Hold_Trn（各 3 格）", "PR #172", "Q7"):
        assert key in txt, key
    rm = wbf["README"]
    b5 = str(rm["B5"].value)
    assert b5.startswith("20261025_Tokenomics_v21.6（次版，底版 v21.5：") and "契約名稱的值改變的只有 IF_NonNV_Hold_TPU、IF_NonNV_Hold_Trn（各 3 格）與版本字串" in b5
    assert "以下為 v21.5（次版，底版 v21.4：" in b5 and str(rm["A1"].value).startswith("Tokenomics v21.6")
    assert (REPO / "docs" / "workorders" / "20261025_v21.6_預期變動.csv").exists()


# ---------------------------------------------------------------- 4. GOV
def test_v2160_gov_counts(lo):
    v = lo.view({})
    assert (v.get_name("GOV_Errors"), v.get_name("GOV_Warnings"), v.get_name("GOV_Info")) == (0, 134, 182)
    assert v.get_name("RNG_StaleFlag") == "一致" and v.get_name("RNG_OutCount") == 0
    errs = [(s, c) for (s, c), x in v.evaluate_all().items() if isinstance(x, str) and x.startswith("#")]
    assert errs == []


# ---------------------------------------------------------------- 5. 逐格比較 v21.5 → v21.6
EXACT_CELLS = {
    "README": {"A1", "B5"}, "Alloc": {"C177"},
    "NonNV": {"E6", "F6", "G6", "E7", "F7", "G7", "L6", "L7"},
    "Decisions": {f"{c}335" for c in "ABCDEFGHIJ"},
    "DB_Evidence": {f"{c}333" for c in "ABCDEFGHIJKLMNOPQR"},
}
TEXT_MIRRORS = {("Gov_Map", f"N{r}") for r in range(395, 407)}      # NonNV L6／L7 說明的鏡像（Gov_Map N 欄公式 =NonNV!$L$6／$L$7，文字隨 W1 改寫）


def test_v2160_cell_diffs_only_expected():
    a, b = openpyxl.load_workbook(V215), openpyxl.load_workbook(model_path())
    av, bv = openpyxl.load_workbook(V215, data_only=True), openpyxl.load_workbook(model_path(), data_only=True)
    assert a.sheetnames == b.sheetnames
    fd, nd, td = {}, set(), set()
    for ws in a:
        wn, va, vb = b[ws.title], av[ws.title], bv[ws.title]
        for r in range(1, max(ws.max_row, wn.max_row) + 1):
            for c in range(1, max(ws.max_column, wn.max_column) + 1):
                x, y = ws.cell(r, c).value, wn.cell(r, c).value
                if x != y:
                    fd.setdefault(ws.title, set()).add(ws.cell(r, c).coordinate)
                p, q = va.cell(r, c).value, vb.cell(r, c).value
                if _num(p) and _num(q) and not _eq(p, q):
                    nd.add((ws.title, ws.cell(r, c).coordinate))
                elif not (_num(p) and _num(q)) and p != q:
                    td.add((ws.title, ws.cell(r, c).coordinate))
    assert fd == EXACT_CELLS, {k: sorted(fd.get(k, set()) ^ EXACT_CELLS.get(k, set())) for k in set(fd) | set(EXACT_CELLS) if fd.get(k, set()) != EXACT_CELLS.get(k, set())}
    assert nd == NUM_DIFFS, sorted(nd ^ NUM_DIFFS)      # 數值差異只落在工作單第 0 節清單
    text_expected = {(s, c) for s, cs in EXACT_CELLS.items() for c in cs} | TEXT_MIRRORS | {("Alloc", "C178")}
    assert td - text_expected == set(), sorted(td - text_expected)
    for sh, cells in fd.items():      # 公式文字不變：改的只有常數與文字
        for coord in cells:
            x, y = a[sh][coord].value, b[sh][coord].value
            assert not (isinstance(x, str) and x.startswith("=")) or sh in ("README", "Alloc"), (sh, coord)
    # 預期變動 CSV 列出全部數值差異格
    rows = list(csv.reader(open(REPO / "docs" / "workorders" / "20261025_v21.6_預期變動.csv", encoding="utf-8-sig")))
    listed = {(r[0], r[1]) for r in rows[1:]}
    assert NUM_DIFFS <= listed, sorted(NUM_DIFFS - listed)
