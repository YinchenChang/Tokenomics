# v10.1 1-d: Gov_Map registry hygiene (4-0 sheet 6, 11 items). Registry pages only; no input value, no formula of a model page and no name changes.
#   1. the same cell registered twice: Har_In E12:G12 (GM481 narrowed to C12:D12, GM655 stays) and Train_In C34:E34 (GM652-654 stay, GM540 marked 停用)
#   2. Sens_Perf E52:G52 and I54:K54 contain one formula cell (the baseline): marked 含公式格, state column U says so
#   3. labels that do not match the target row (Sens_Train O8:R8, Y8:AB8, Sens_Rev B18): the label now starts with the target row's label
#   4. Alloc I94:L95 (8 blue cells, one-hot service-generation scenarios) get a Gov_Map row of the same category as the other scenario cells (情境值)
# Every edit is guarded by the old text, so a rebuild from v10.1 changes nothing.
import re
from openpyxl.utils import get_column_letter as L

NOTE = "｜v10.1 4-1："
DISABLED = "（停用）"


def _row_of(ws, gid):
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == gid: return r
    return None


def _append(cell, txt):
    v = cell.value
    if isinstance(v, str) and txt in v: return False
    cell.value = (v if v not in (None, "") else "") + txt
    return True


def _prepend(cell, txt):
    v = cell.value
    if isinstance(v, str) and v.startswith(txt): return False
    cell.value = txt + (v if v is not None else "")
    return True


def gov_fixes(wb):
    ws = wb["Gov_Map"]; n = {}
    # 1a Har_In E12:G12 is registered by GM481 (C12:G12) and GM655 (E12:G12): GM481 keeps only C12:D12
    r = _row_of(ws, "GM481")
    if ws.cell(r, 4).value == "C12:G12":
        ws.cell(r, 4).value = "C12:D12"; _append(ws.cell(r, 14), f" {NOTE}E12:G12 與 GM655 重複登錄，本列改為只登錄 C12:D12（4-0 盤點 G1）"); n["GM481"] = 1
    # 1b Train_In C34:E34: GM652-654 register the three cells one by one (with their own intervals), GM540 is switched off
    r = _row_of(ws, "GM540")
    if ws.cell(r, 4).value == "C34:E34":
        ws.cell(r, 4).value = DISABLED + "C34:E34"; _append(ws.cell(r, 14), f" {NOTE}停用——C34、D34、E34 已由 GM652、GM653、GM654 逐格登錄（4-0 盤點 G1）；本列保留作為歷史與區間說明"); n["GM540"] = 1
    # 2 Sens_Perf: one formula cell in the middle of each range
    for gid, rng, mid, src in (("GM424", "E52:G52", "F52", "Tech_Registry!$E$20"), ("GM425", "I54:K54", "J54", "Tech_Registry!$E$21")):
        r = _row_of(ws, gid)
        if ws.cell(r, 3).value == "Sens_Perf" and ws.cell(r, 4).value == rng:
            ws.cell(r, 21).value = "含公式格"
            _prepend(ws.cell(r, 14), f"【含公式格】範圍中間格 {mid} 是公式（基準欄，讀 {src}），其餘兩格是情境輸入。 "); n[gid] = 1
    # 3 labels start with the target row's label
    for gid, old, new in (("GM437", "情境倍數：RL rollout token × 0.3／× 3", "情境倍數（×）：RL rollout token × 0.3／× 3"),
                          ("GM438", "情境倍數：合成資料 token × 0／× 3", "情境倍數（×）：合成資料 token × 0／× 3"),
                          ("GM439", "ε（K4 (a) 量級情境）", "ε（上一列量級情境的指數；K4 (a) 量級情境）")):
        r = _row_of(ws, gid)
        if ws.cell(r, 5).value == old: ws.cell(r, 5).value = new; n[gid] = 1
    return n


def add_alloc_rows(wb):
    """4. Alloc I94:L95 one-hot scenario inputs: a copy of the GM424 row (same category 情境值) with the target, label and note changed."""
    ws = wb["Gov_Map"]
    if _row_of(ws, "GM681") is not None: return 0
    last = max(r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value)
    src = _row_of(ws, "GM424"); new = last + 1
    for c in range(1, ws.max_column + 1):
        s = ws.cell(src, c); d = ws.cell(new, c)
        v = s.value
        if isinstance(v, str) and v.startswith("="): v = re.sub(r"(?<![0-9])%d(?![0-9])" % src, str(new), v)
        d.value = v
        d._style = s._style
    ws.cell(new, 1).value = "GM681"
    ws.cell(new, 8).value = None      # SRC_ID blank: SRC_Index row 1448 reads Gov_Map!H685, and a dash there would change the name IDX_SrcID
    ws.cell(new, 21).value = "—"      # U（建置時狀態）：範圍列一律「—」，與重建時 builder 的寫法一致（不沿用 GM424 的「含公式格」）
    ws.cell(new, 2).value = "切片四（v10.1 瘦身 4-1）"
    ws.cell(new, 3).value = "Alloc"; ws.cell(new, 4).value = "I94:L95"
    ws.cell(new, 5).value = "服務世代組合：全 GB200／全 VR200（G 節敏感度 one-hot 輸入）"
    ws.cell(new, 14).value = ("敏感度情境設定：G 節第 94 列＝全 GB200（J94＝1）、第 95 列＝全 VR200（L95＝1），其餘為 0；黃底藍字。"
                              "v6.2 起存在但未登錄，v10.1 補登（4-0 盤點 G2）；不改任何輸入值")
    return 1


def apply(wb):
    return dict(fixes=gov_fixes(wb), added=add_alloc_rows(wb))
