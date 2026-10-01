"""函數與語意覆蓋：v5.2 用到的函數、運算子與型別行為，逐式對照 LibreOffice。

每個式子放在含中文頁名的小型活頁簿內，引擎與 LibreOffice 重算後比對。
不符者記為「引擎語意缺口」，須在報告列出（不得改寫 Excel、不得在 Python 另寫旁路計算）。
"""
import pytest
import openpyxl

from engine import Engine
from parity_lib import compare, excel_values, lo_recalc

SHEET = "輸入頁"   # 含中文頁名
# 輸入（B 欄）：B1=3.7, B2=-2.5, B3=0.25, B4=文字, B5=空白, B6=10, B7=0, B8="SLO 不可達"
INPUTS = {"F1": "A", "F2": "B", "F3": "A", "F4": "B", "G1": 1.5, "G2": 2, "G3": 0.5, "G4": 1,   # v5.5：Tech_Registry 掛鉤彙總、Training 條件加總
          "B1": 3.7, "B2": -2.5, "B3": 0.25, "B4": "abc", "B6": 10, "B7": 0, "B8": "SLO 不可達"}
FORMULAS = [
    "=2^10", "=B1^2", "=B3^0.5", "=(-8)^(1/3)", "=2^-2", "=B6^B3",            # ^
    '="SLO（"&B4&"）"', '="x"&B8', '=B4&B4',                                  # &（字串）
    "=CEILING(B1,1)", "=CEILING(B6*B1/(2*B3),1)", "=CEILING(10,1)", "=CEILING(0.0001,1)",
    "=FLOOR(B1,1)", "=FLOOR(B6*1000/B1,1)", "=FLOOR(9.999999999,1)", "=FLOOR(4,1)",
    "=ABS(B2)", "=ABS(B7)",
    "=MAX(B1,B2,B6)", "=MIN(B1,B2,B6)", "=MAX(B1:B3)", "=MIN(B1:B3)", "=MAX(B1:B8)", "=MIN(B1:B8)",
    "=ISNUMBER(B1)", "=ISNUMBER(B4)", "=ISNUMBER(B5)", "=ISNUMBER(B1/B6)",
    '=IF(B4="abc",1,0)', '=IF(B4="ABC",1,0)', '=IF(B5="",1,0)', '=IF(B8="SLO 不可達","是","否")',
    '=IF(B1>B6,"高","低")', '=IF(B7<=0,"SLO 不可達",B1)', '=IF(B5="",1,B5)',
    "=AND(B1>0,B6>0)", "=AND(B1>0,B7>0)", "=AND(B1>0,B4=\"abc\",B6=10)",
    "=CHOOSE(2,B1,B2,B6)", "=CHOOSE(B7+1,B1,B2)", '=CHOOSE(3,"a","b","c")',
    "=INDEX(B1:B3,2)", "=INDEX(B1:B3,3)", "=INDEX(B1:B3,1+1)", "=INDEX(B1:B6,3*(2-1)+1)",
    '=TEXT(10.5/B6,"0.0")', '=TEXT(3.57/B1,"0.00")', '=TEXT(B3,"0.0")', '=TEXT(0.05,"0.0")', '=TEXT(2.675,"0.00")',
    '="牌價 ÷ 持有成本＝"&TEXT(10.5/B6,"0.0")&" 倍"',
    '=COUNTIF(B1:B8,"SLO 不可達")', '=COUNTIF(B1:B8,"abc")', '=COUNTIF(B1:B8,"slo 不可達")',
    # v5.5：Tech_Registry 掛鉤彙總（同代碼多條目取乘積）與 Training 的 SUMPRODUCT 條件加總
    '=EXP(SUMPRODUCT((F1:F4="A")*LN(G1:G4)))', '=EXP(SUMPRODUCT((F1:F4="B")*LN(G1:G4)))', '=EXP(SUMPRODUCT((F1:F4="Z")*LN(G1:G4)))',
    '=SUMPRODUCT((F1:F4="A")*G1:G4)', "=SUMPRODUCT((F1:F4=F1)*G1:G4)", "=SUMPRODUCT(G1:G4,G1:G4)",
    "=B1=B1", '=B4="abc"', "=B1<>B6", '=B8<>"x"', "=B5=0", '=B4=B8',
]


@pytest.fixture(scope="module")
def results(tmp_path_factory):
    d = tmp_path_factory.mktemp("fn")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    for k, v in INPUTS.items():
        ws[k] = v
    for i, f in enumerate(FORMULAS, start=1):
        ws[f"D{i}"] = f
    src = d / "fn.xlsx"
    wb.save(src)
    eng = Engine(src)
    got = eng.evaluate_all()
    ref = excel_values(lo_recalc(src, d / "lo"), eng.formula_cells)
    return got, ref


@pytest.mark.parametrize("i,f", list(enumerate(FORMULAS, start=1)), ids=[f for f in FORMULAS])
def test_formula_semantics(i, f, results):
    got, ref = results
    key = (SHEET, f"D{i}")
    res = compare({key: got[key]}, {key: ref[key]})
    assert not res["mismatches"], f"{f}: 引擎 {got[key]!r} ≠ LibreOffice {ref[key]!r}"
