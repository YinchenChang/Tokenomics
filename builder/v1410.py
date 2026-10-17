# v14.1 (work order docs/workorders/20261017_v14.1修補.md r1, branch claude/v14-1): patch after the 4-3B slimming (v14.0).
# Minor version: no new contract name; one existing contract name changes value (L1_MFleet_Lo, plus its dependants), everything else unchanged.
#   R1  m low end on "same-column floor": Cap_In row 119 (機隊 m) of every column = MAX(that column's point estimate, that column's own floor in row 121).
#       Row 121 is computed with the same column's service GW and utilisation convention as row 119, so the low-end column (D, token-high) is floored at 0.215 (not at the base column's 0.301).
#       The mechanism is an Excel formula in Cap_In (patch_specs below edits the v5.33 spec before the sheet is written, so the build-time precedent tracing sees it);
#       no Python-side computation. L1_MFleetFloor is NOT raised to a defined name (nothing downstream needs it; row 121 stays a display/mechanism row).
#   R2  Checks V23／V24 marked "恆等式（非驗證）" like V18 (builder/v1000.py checks_v): out of the verification count (24／16／7／1 -> 22／14／7／1); their must-fail tests stay.
#   R3  Har_In rows 20／32 (vendor-specific flags; no formula reads them since the Har_In 44 deletion in v14.0): values C:G cleared; Gov_Map rows get the note; original text kept.
#   R4  the A／B labels of the 73 rows deleted in v14.0 are cleared; the list (sheet, row, original A／B) is written to Decisions SL44-4.
#   R5  Cap_In 24 cells (C of rows 185, 187–193, 196–202, 205, 216–222 and the text cell G175) cleared after a reader check (every reader of the group is in the group); original formulas in Decisions SL44-5.
#   R6  small items of the v14.0 independent check (PR #104): Gov_Map GM629 note, three 用途 corrections (Cap_In 76, Training 122, Price_Frontier 27), stale "錨點" texts of Cap_In 206–209 and 223.
# Every write to an Excel-owned cell is guarded (old value / old text).
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import PatternFill

from common import put, F_NOTE, F_CALC
import v1400

VERSION = "20261017_Tokenomics_v14.1"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-17"
V = "v14.1"
WO = "工作單 docs/workorders/20261017_v14.1修補.md"
DEC_BY = "Andy 2026-10-17 10:46（回覆『照預設』，同意 chat 端 10:40 的 7 項預設；R1 選 (b)）"
HERE = Path(__file__).resolve().parent
NO_FILL = PatternFill(fill_type=None)
PURPOSE = "用途"
_PATCHED = False


# ================================================================== R1: Cap_In row 119 = MAX(point estimate, same-column floor row 121)
def patch_specs():
    """Call right after v810.patch_specs() and before the Cap_In H7 section is written."""
    global _PATCHED
    if _PATCHED:
        return
    import v533
    for i, it in enumerate(v533.SPEC):
        if it[0] != "m_fleet":
            continue
        it = list(it)
        for k, col in ((3, "C"), (4, "D"), (5, "E")):
            old = it[k]
            m = re.fullmatch(r'=IF\(ISNUMBER\(%s\{tok_fleet\}\),(SRC_DEM_007\*1000000000/\(%s\{tok_fleet\}\*AL_PriceFleet25\)),"SLO 不可達"\)' % (col, col), old)
            assert m, ("m_fleet spec changed", old)
            x = m.group(1)
            it[k] = ('=IF(ISNUMBER(%s{tok_fleet}),MAX(%s,IF(ISNUMBER(%s{m_floor}),%s{m_floor},0)),"SLO 不可達")' % (col, x, col, col))
        v533.SPEC[i] = tuple(it)
        break
    else:
        raise AssertionError("m_fleet spec not found")
    _PATCHED = True


# ================================================================== helpers
def _purpose_col(ws):
    return v1400._purpose_col(ws, create=False)


def _set_text(ws, ref, old, new):
    """Guarded text write: the cell must hold `old` (or already `new`)."""
    cur = ws[ref].value
    if cur == new:
        return 0
    assert cur == old, (ws.title, ref, cur)
    ws[ref].value = new
    return 1


def _gm_rows(wb, sheet, cell):
    ws = wb["Gov_Map"]
    return [r for r in range(5, ws.max_row + 1) if (ws[f"C{r}"].value, ws[f"D{r}"].value) == (sheet, cell)]


# ================================================================== R1 texts
R1_G119_TAIL = "；下限截斷：v14.1 起，低端不得低於同欄物理下限（第 121 列同欄）"
R1_G121_TAIL = "；v14.1 起第 119 列各欄＝MAX(點估計, 本列同欄)，故 L1_MFleet 的 _Lo／基準／_Hi 都不低於各自欄的本列"
R1_MARK_OLD = "顯示用：C121:E121（待定義，見 4-3B 報告）｜物理下限 0.301 與 L1_MFleet_Lo 0.184 的定義不同，三種處理待 Project 判斷；本版不升級為契約名稱"
R1_MARK_NEW = ("公式列：C121:E121（機隊 m 的同欄物理下限；v14.1 正式定義）｜每一欄用該欄自己的服務 GW 與利用率口徑算出『token 上限 → m 下限』；"
               "第 119 列（機隊 m，→ L1_MFleet／_Lo／_Hi）各欄＝MAX(點估計, 本列同欄)；不升級為契約名稱 L1_MFleetFloor（下游不讀此列，只讀截斷後的 L1_MFleet 區間）；"
               "原標記『待定義，見 4-3B 報告』（v14.0）")


def r1_texts(wb, append_text):
    ws = wb["Cap_In"]
    n = 0
    for ref, tail in (("G119", R1_G119_TAIL), ("G121", R1_G121_TAIL)):
        cur = ws[ref].value
        assert isinstance(cur, str) and cur.startswith("="), (ref, cur)
        n += append_text(ws[ref], tail)
    pc = _purpose_col(ws)
    n += _set_text(ws, f"{L(pc)}121", R1_MARK_OLD, R1_MARK_NEW)
    return n


# ================================================================== R3: Har_In 20／32
HAR_ROWS = (20, 32)
R3_MARK = "已清空（v14.1，R3）：供應商專屬旗標（輸入）無任何公式讀取（對應的第 44 列已於 v14.0 刪除）；原值 {vals}，原『來源與標記』欄 H、說明欄 I 原文保留；見 Decisions SL44-3"
R3_GM_TAIL = " ｜v14.1 已移除（R3，修訂九更正註記）：Har_In {rng} 的旗標值已清空（無任何公式讀取）；本列原文與原值保留於此與 Decisions SL44-3"


def r3_har_flags(wb, append_text):
    ws = wb["Har_In"]
    pc = _purpose_col(ws)
    orig = {}
    n = 0
    for r in HAR_ROWS:
        vals = [ws.cell(r, c).value for c in range(3, 8)]
        if all(v is None for v in vals):
            continue                              # already cleared (rebuild from v14.1)
        assert ws.cell(r, 1).value == "供應商專屬（1＝是）", (r, ws.cell(r, 1).value)
        orig[r] = vals
        for c in range(3, 8):
            ws.cell(r, c).value = None; ws.cell(r, c).fill = NO_FILL
        put(ws, f"{L(pc)}{r}", R3_MARK.format(vals="／".join(str(v) for v in vals)), F_NOTE)
        n += 1
    gm = 0
    for r in HAR_ROWS:
        rng = f"C{r}:G{r}"
        rows = _gm_rows(wb, "Har_In", rng)
        assert len(rows) == 1, (rng, rows)
        gm += append_text(wb["Gov_Map"][f"N{rows[0]}"], R3_GM_TAIL.format(rng=rng))
    return dict(rows=n, gov_map=gm, orig=orig)


# ================================================================== R4: A／B labels of the 73 v14.0-deleted rows
def r4_labels(wb):
    data = v1400.DATA["delete"]
    lst = []
    n = 0
    for e in data:
        ws = wb[e["sheet"]]
        a, b = ws.cell(e["row"], 1).value, ws.cell(e["row"], 2).value
        if a is None and b is None:
            continue
        lst.append((e["sheet"], e["row"], a, b))
        ws.cell(e["row"], 1).value = None; ws.cell(e["row"], 2).value = None
        n += 1
    return dict(rows=n, list=lst)


# ================================================================== R5: Cap_In 24 cells
R5_ROWS = [185] + list(range(187, 194)) + list(range(196, 203)) + [205] + list(range(216, 223))
R5_CELLS = [f"C{r}" for r in R5_ROWS] + ["G175"]
R5_MARK_OLD_PREFIX = "顯示用：C{r}（H10 的 2025 年度欄"
R5_MARK_NEW = "已清空 C{r}（v14.1，R5）：H10 的 2025 年度欄，原本只與 D–F 欄並列對照、無任何公式讀取（刪除前已逐格確認）；原公式記於 Decisions SL44-5"
R5_G175_MARK_OLD = "顯示用：G175（說明動態文字：現值取代過時數字）｜C175:E175 為殘差讀數，v14.0 起由 Checks V25 讀取"
R5_G175_MARK_NEW = "已清空 G175（v14.1，R5）：說明動態文字無任何公式讀取；C175:E175 為殘差讀數，v14.0 起由 Checks V25 讀取；原公式記於 Decisions SL44-5"


def r5_cap_in(wb):
    from deps import Deps
    ws = wb["Cap_In"]
    live = [c for c in R5_CELLS if ws[c].value is not None]
    if not live:
        return dict(cleared=0, orig=[], kept=[])
    D = Deps(wb)
    group = {("Cap_In", c) for c in live}
    rev = {}
    for k in D.formulas:
        for p in D.direct(k):
            if p in group:
                rev.setdefault(p, set()).add(k)
    kept = []
    for c in live:
        outside = sorted(k for k in rev.get(("Cap_In", c), ()) if k not in group)
        if outside:
            kept.append((c, outside))
    keep_set = {c for c, _ in kept}
    # a cell whose reader is kept is itself kept (closure); loop until stable
    changed = True
    while changed:
        changed = False
        for c in live:
            if c in keep_set:
                continue
            if any(k[0] == "Cap_In" and k[1] in keep_set for k in rev.get(("Cap_In", c), ())):
                keep_set.add(c); changed = True
    pc = _purpose_col(ws)
    orig = []
    for c in live:
        if c in keep_set:
            continue
        orig.append((c, ws[c].value))
        ws[c].value = None; ws[c].fill = NO_FILL
        r = int(c[1:])
        if c == "G175":
            _set_text(ws, f"{L(pc)}{r}", R5_G175_MARK_OLD, R5_G175_MARK_NEW)
        else:
            cur = ws[f"{L(pc)}{r}"].value
            assert str(cur).startswith(R5_MARK_OLD_PREFIX.format(r=r)), (r, cur)
            ws[f"{L(pc)}{r}"].value = R5_MARK_NEW.format(r=r)
    return dict(cleared=len(orig), orig=orig, kept=[(c, o) for c, o in kept])


# ================================================================== R6
GM629_TAIL = " ｜v14.0 已清空（修訂九更正註記；v14.1 R6 補記）：Cap_In C204:F204 的值已於 v14.0（4-3B Q1）清空，本列原文保留"
ANCHOR_ROWS = (206, 207, 208, 209, 223)
ANCHOR_NEW = "錨點：以本列 C–F 欄現值為準（v5.35 B1 建置時的數字已刪除，避免與現值不一致；原文見 Decisions SL44-6）"
MARK_FIX = [
    ("Cap_In", 76, "顯示用：C76（H2 訂閱人數（SRC_DEM 複本與 2025 期間平均）：每人直覺檢查與 ARPU 的上游）",
     "顯示用：C76（H2 訂閱人數的 SRC_DEM 複本，只供與 2025 期間平均並列對照；目前沒有任何公式讀取；v14.1 R6 更正：原標記稱『每人直覺檢查與 ARPU 的上游』與實際依賴不符）"),
    ("Training", 122, "顯示用：C122:Q122（核心列：Sens_Train 有同列同標籤的列，常設測試以此為逐格比對基準；下游用 IF_TrainCost_*／IF_ProgCost_*）",
     "顯示用：C122:Q122（核心列：Sens_Train 有同列同標籤的列，常設測試以此為逐格比對基準；本列沒有公式讀取，下游 IF_TrainCost_*／IF_ProgCost_* 實際讀第 120 列（研發計畫 GPU 小時，TRN_GPUhProg）；v14.1 R6 更正）"),
    ("Price_Frontier", 27, "顯示用：C27:E27（前緣分項牌價（新鮮輸入／快取輸入／輸出）：B 節前緣表的形狀；下游用 IF_FrontRef_*（混合））",
     "顯示用：C27:E27（前緣分項牌價（新鮮輸入／快取輸入／輸出）：B 節前緣表的形狀；本列沒有公式讀取，下游 IF_FrontRef_* 實際讀第 46 列（B4_FrontRef，前緣混合單價）；v14.1 R6 更正）"),
]


def r6_small(wb, append_text):
    log = {}
    gm = _gm_rows(wb, "Cap_In", "C204:F204")
    assert len(gm) == 1, gm
    log["gm629"] = append_text(wb["Gov_Map"][f"N{gm[0]}"], GM629_TAIL)
    n = 0
    for sh, r, old, new in MARK_FIX:
        ws = wb[sh]
        pc = _purpose_col(ws)
        n += _set_text(ws, f"{L(pc)}{r}", old, new)
    log["marks"] = n
    ws = wb["Cap_In"]
    orig = []
    for r in ANCHOR_ROWS:
        cur = ws[f"H{r}"].value
        if cur == ANCHOR_NEW:
            continue
        assert isinstance(cur, str) and cur.startswith("錨點：") and not cur.startswith("錨點：以本列"), (r, cur)
        orig.append((r, cur))
        ws[f"H{r}"].value = ANCHOR_NEW
    log["anchors"] = len(orig)
    log["anchor_orig"] = orig
    return log


# ================================================================== Decisions
def _decisions(r3, r4, r5, r6):
    lab = "；".join(f"{s}!{r}：{a}｜{b}" for s, r, a, b in r4["list"])
    r5t = "；".join(f"Cap_In!{c}：{f}" for c, f in r5["orig"])
    r6t = "；".join(f"Cap_In!H{r}：{t}" for r, t in r6["anchor_orig"])
    _D = [
        ("1", "R1 機隊 m 低端改用『同欄下限』（下限截斷）",
         "Cap_In 第 119 列（機隊 m，C 基準／D token 取高／E token 取低）各欄改為 MAX(該欄點估計, 該欄第 121 列的物理下限)；第 121 列每一欄用該欄自己的服務 GW 與利用率口徑算出，"
         "所以低端欄（D）的下限是 0.215（不是基準欄的 0.301）。原因（4-3B 報告第 6 節）：D 欄 token 取高的點估計 token（423 億）高於它自己那一欄的上限（361 億，比值 1.17），"
         "兩者差在服務 GW 與利用率修正的口徑不同，造成 L1_MFleet_Lo 0.184 低於物理下限。處理：下限截斷（Andy 選 (b)）。基準 0.407 與高端 1.179 高於各自欄的下限，不受影響。"
         "機制寫在 Excel 公式（Cap_In 第 119 列）；說明欄寫明『下限截斷：v14.1 起，低端不得低於同欄物理下限』。Cap_In 第 121 列用途由『待定義』改為正式定義；不升級為契約名稱 L1_MFleetFloor（下游不讀）。"
         "受影響的名稱：L1_MFleet_Lo、L1_PriceRealFleet_Lo、L1_LReqFleet_Hi（及其依賴格 Cap_In E240／E242、CHK_L1Order、GOV_Warnings 的輸入；新舊值見報告）。",
         "Cap_In 第 119、121 列；L1 第 78 列", "v14.1 已建"),
        ("2", "R2 V23、V24 標『恆等式（非驗證）』",
         "Checks V23（Perf 實際速度 ≥ SLO）與 V24（Cap_In 第 147 列＝L1_MImplied）是結構性恆等式（4-3B 報告第 5 節）：定義上恆成立，無法只靠改輸入使其失敗。比照 V18，F 欄改顯示文字『恆等式（非驗證）』，"
         "移出驗證項目數／通過／失敗；公式被改壞時 F 欄顯示『恆等式不成立』，專屬必須失敗測試保留（改公式格）。驗證計數 24／16／7／1 → 22／14／7／1。",
         "Checks V23、V24；CHK_VerifyTotal／Pass", "v14.1 已建"),
        ("3", "R3 Har_In 第 20、32 列供應商旗標清空",
         "供應商專屬旗標（Har_In 第 20、32 列 C:G）自第 44 列於 v14.0 刪除後已無任何公式讀取，v14.1 清空。原值：" +
         "；".join(f"第 {r} 列 " + "／".join(str(v) for v in vals) for r, vals in sorted(r3["orig"].items())) +
         "。H 欄（來源與標記）與 I 欄（說明）原文保留；Gov_Map 該兩列（列 493、501）加註『v14.1 已移除』，原文保留（修訂九）。需要時再接線。",
         "Har_In 第 20、32 列；Gov_Map 列 493、501（GM489、GM497）", "v14.1 已建"),
        ("4", "R4 v14.0 刪除的 73 列：A、B 欄標籤清空",
         "v14.0（SL43-1）刪除的 73 列，v14.1 一併清空 A、B 欄標籤。可追溯清單（工作表!列：原 A 欄｜原 B 欄）：" + lab,
         "各頁（見清單）", "v14.1 已建"),
        ("5", "R5 Cap_In 24 格清空",
         "Cap_In 第 185、187–193、196–202、205、216–222 列 C 欄（23 格）與第 175 列說明文字格 G175，刪除前以依賴圖逐格確認：這 24 格的讀者全部在這 24 格之內（沒有任何其他格、名稱、Checks 讀取），所以整組清空。"
         "保留（有組外讀者）：" + ("無" if not r5["kept"] else "；".join(f"{c}：{o}" for c, o in r5["kept"])) + "。原公式：" + r5t,
         "Cap_In C 欄上述各列；G175", "v14.1 已建"),
        ("6", "R6 v14.0 獨立查核（PR #104）的小項",
         "(1) Gov_Map GM629（Cap_In C204:F204）加註『v14.0 已清空』，原文保留。(2) 用途標記三筆更正：Cap_In 76（實際無公式讀取）、Training 122（下游實為第 120 列）、Price_Frontier 27（下游實為第 46 列）；原標記文字保留於各列 O／該頁用途欄的更正句。"
         "(3) Cap_In 第 206–209、223 列 H 欄『錨點』舊文字與現值不符，改為『以本列 C–F 欄現值為準』；原文：" + r6t + "。數值全部不變。",
         "Gov_Map GM629；Cap_In 76、206–209、223；Training 122；Price_Frontier 27", "v14.1 已建"),
    ]
    return [[f"SL44-{i}", V, f"v14.1 修補：{t}", txt, DATE, DEC_BY, st, where, f"{WO} R{i}", "否"] for i, t, txt, where, st in _D]


def decisions_append(wb, rows):
    ws = wb["Decisions"]
    have = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
    r = max([rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")] or [4]) + 1
    n = 0
    for row in rows:
        if row[0] in have:
            continue
        for i, v in enumerate(row):
            put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 5, 7))
        r += 1; n += 1
    return n


def old_decision_tail(wb, append_text):
    """SL43-4 (v14.0) says V23–V26 are counted in the verification count: add the v14.1 R2 correction."""
    ws = wb["Decisions"]
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == "SL43-4":
            return append_text(ws[f"D{r}"], " ｜v14.1 R2：V23、V24 已改標『恆等式（非驗證）』，移出驗證計數（見 SL44-2）；V25、V26 仍計入")
    return 0


# ================================================================== apply
def apply(wb, append_text):
    log = {}
    log["r1_texts"] = r1_texts(wb, append_text)
    log["r3"] = r3_har_flags(wb, append_text)
    log["r4"] = r4_labels(wb)
    log["r5"] = r5_cap_in(wb)
    log["r6"] = r6_small(wb, append_text)
    ws = wb["Decisions"]
    have_dec = any(ws.cell(r, 1).value == "SL44-1" for r in range(5, ws.max_row + 1))
    if not have_dec:
        # lists come from the live sheets before clearing; on a rebuild from v14.1 the sheets are rebuilt so they are complete again
        log["decisions"] = decisions_append(wb, _decisions(log["r3"], log["r4"], log["r5"], log["r6"]))
        log["sl43_4"] = old_decision_tail(wb, append_text)
    return log


# ================================================================== README
README_VERSION = (VERSION + "（次版：" + WO + "——R1 機隊 m 低端改用同欄下限（Cap_In 第 119 列各欄＝MAX(點估計, 第 121 列同欄)；L1_MFleet_Lo 0.184 → 0.215，基準與高端不變）；"
                  "R2 Checks V23、V24 標恆等式（驗證計數 24／16／7／1 → 22／14／7／1）；R3 Har_In 第 20、32 列供應商旗標清空；R4 v14.0 刪除列的 A、B 標籤清空；R5 Cap_In 24 格清空；"
                  "R6 v14.0 獨立查核小項（GM629 註記、三筆用途更正、過時錨點文字）；Decisions SL44-*）。以下為 "
                  + v1400.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1400.README_TITLE
README_TITLE = V + _T[len(v1400.V):]
README_ROW = ("v14.1 修補",
              "機隊 m 低端改用同欄下限（L1_MFleet_Lo 0.184 → 0.215）；V23、V24 標恆等式；清空無人讀取的旗標與標籤；用途標記與過時文字更正。細節見 Decisions SL44-*。")
