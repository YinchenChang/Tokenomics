#!/usr/bin/env python3
"""Stage 0 盤點工具（唯讀）：藍字輸入的標記分類、靜態依賴圖、擾動敏感度。

用法
  盤點：  python3 tools/stage0_inventory.py inventory model/<現行>.xlsx --out docs/reports/YYYYMMDD_stage0_inventory.xlsx
  Gate 1：python3 tools/stage0_inventory.py gate1 model/<現行>.xlsx --prev model/archive/<前版>.xlsx --out docs/reports/YYYYMMDD_gate1.md
  影響：  python3 tools/stage0_inventory.py impact model/<現行>.xlsx 'Sheet!C5' ['Sheet!D5' ...]
          （Stage 1 Gate 1 重用：改一筆輸入，列出受影響的 IF_ 名稱、主要輸出與工作表）

規則：
- 不修改輸入的 xlsx（唯讀；LibreOffice 抽驗只動暫存複本）。
- 敏感度一律透過 engine/（pycel）重算；不在此另寫任何公式。
- 本檔不加入 CI。常數 EXPECT_* 是 chat 端 2026-10-02 對 v5.10 的量測值，供核對表使用，不是模型參數。
"""
from __future__ import annotations

import argparse
import collections
import hashlib
import json
import math
import multiprocessing as mp
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formula import Tokenizer
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

# ───────────────────────── 規格常數（來自第 10 輪指令） ─────────────────────────
MARKS = ["Verified", "Interested-party", "Analogy", "Assumed", "Decision", "Derived"]   # 主標記優先序
S_RE = re.compile(r"(?<![A-Za-z0-9])S\d{1,3}(?!\d)")
E_RE = re.compile(r"(?<![A-Za-z0-9])E\d{3}")
OPTION_RE = re.compile(r"[（(][^）)]*\d[^）)]*[／/][^）)]*\d[^）)]*[）)]")      # 「（1／2／3）」類選項說明
OPTION_FLAG_RE = re.compile(r"[（(][^）)]*[01]\s*[＝=]\s*[有無是否][^）)]*[）)]")           # 「（1＝有）」「（0＝無、1＝有）」「（1＝是）」：0／1 旗標（本工具的解讀）

# (ID, 具名範圍, 欄 k, 內容, 單位, 指令表列出的基準值)
OUTPUTS = [
    ("O01", "IF_CapexTotal", 11, "VR200 每 GW 資本支出合計", "$B", 47.5820),
    ("O02", "IF_HoldEcon", 11, "VR200 每 GW 年持有成本（經濟）", "$B/年", 12.0544),
    ("O03", "IF_TokGW_Sol", 11, "VR200 Sol 每 GW 總產出（100%）", "M tok/年", 2.38098e11),
    ("O04", "IF_CostDec_Sol", 11, "VR200 Sol decode $/M（經濟、100%）", "$/M", 0.322576),
    ("O05", "IF_CostDec_Astra", 11, "VR200 Astra decode $/M", "$/M", 1.66066),
    ("O06", "IF_CostDec_Sol", 8, "GB300 Sol decode $/M", "$/M", 0.556984),
    ("O07", "IF_TrainGPUh_Astra", 11, "VR200 Astra 最終訓練 GPU 小時", "GPU-hr", 1.16777e7),
    ("O08", "IF_ProgCost_Astra", 11, "VR200 Astra 研發計畫成本（經濟）", "$M", 440.640),
    ("O09", "IF_PostShareGPUh_Sol", 11, "VR200 Sol 後訓練占比（GPU 小時）", "%", 0.255756),
    ("O10", "IF_FullCostDefault_Sol", 11, "VR200 Sol 全成本 $/M（下游預設）", "$/M", 0.104696),
    ("O11", "IF_RevGW_Sol", 11, "VR200 Sol 每 GW 理論營收", "$B/年", 229.590),
    ("O12", "IF_RevGWFleet", 11, "VR200 1 GW 參考機隊付費服務營收", "$B/年", 49.1065),
    ("O13", "IF_CostSuccVR_Sol", 5, "Coding agent 每成功任務成本（Sol）", "$", 0.0436357),
    ("O14", "IF_FrontSuccVR", 5, "Coding agent 成功任務成本前緣", "$（或文字「無合格」）", 0.0209162),
]
DISCRETE = [("IF_FrontSuccVRName", k) for k in range(1, 6)] + [
    (n, None) for n in ("IF_FrontModel_Luna", "IF_FrontModel_Sol", "IF_FrontModel_Astra")]

EXPECT_ROWS_BY_SHEET = {
    "Inputs": 26, "Spec_Rack": 22, "Arch": 18, "Serving": 16, "Workload": 11, "Calib": 26, "Tech_Registry": 12,
    "Perf": 2, "Sens_Perf": 9, "Unit_Cost": 1, "Train_In": 36, "Perf_Batch": 4, "Training": 2, "Sens_Train": 8,
    "Cap_In": 37, "Sens_Rev": 13, "Har_In": 42, "Energy": 3, "NonNV": 5, "Checks": 11, "DB_Evidence": 20,
}
EXPECT_TOTALS = {"藍字列": 324, "藍字格": 1491, "藍字公式格": 31}
EXPECT_CLASS = {
    "原始數據（候選）": 69, "Analogy": 31, "Assumed": 97, "Decision": 11, "Derived 藍字（異常）": 2,
    "情境選擇": 39, "待分類": 44, "Checks 外部參照": 11, "Evidence 登錄（非輸入）": 20,
}
EXPECT_PRIMARY = {"Verified": 52, "Interested-party": 35, "Analogy": 34, "Assumed": 97, "Decision": 11,
                  "Derived": 2, "無標記": 93}
EXPECT_MULTI = {"多標記列": 27, "多標記：Arch": 13, "多標記：Spec_Rack": 4, "有 S 編號列": 108}

PCT_UNITS = {"%"}
GRADE_HI, GRADE_MID = 0.5, 0.1
REL = 1e-9

# ───────────────────────── 基本工具 ─────────────────────────


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def is_formula(v) -> bool:
    if hasattr(v, "text") and hasattr(v, "ref"):          # ArrayFormula
        return True
    return isinstance(v, str) and v.startswith("=")


def is_blue(c) -> bool:
    col = c.font.color
    return (col is not None and col.type == "rgb" and str(col.rgb).upper().endswith("0000FF")
            and c.value not in (None, ""))


def is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def fmt_num(v) -> str:
    return f"{v:.12g}" if isinstance(v, float) else str(v)


def clean(s):
    return ILLEGAL_CHARACTERS_RE.sub("", s) if isinstance(s, str) else s


# ───────────────────────── 1. 掃描藍字格 ─────────────────────────


def scan_workbook(wb):
    """回傳 (rows, blue_formula_cells)。每個 row 是一個 dict（工作表 × 列）。"""
    rows, bformulas = [], []
    for ws in wb:
        for row in ws.iter_rows():
            blues, texts = [], []
            for c in row:
                v = c.value
                if v in (None, ""):
                    continue
                blue = is_blue(c)
                if is_formula(v):
                    if blue or (c.font.color is not None and c.font.color.type == "rgb"
                                and str(c.font.color.rgb).upper().endswith("0000FF")):
                        bformulas.append((ws.title, c.coordinate, v if isinstance(v, str) else v.text, c.row))
                    continue
                if isinstance(v, str):
                    texts.append(c)
                if blue:
                    blues.append(c)
            if not blues:
                continue
            rows.append(_build_row(ws, row[0].row, blues, texts))
    # 藍字公式格的列標籤
    lab = {(r["sheet"], r["row"]): r["label"] for r in rows}
    out = []
    for s, a, f, rr in bformulas:
        out.append((s, a, f, lab.get((s, rr)) or _row_label(wb[s], rr, None)))
    return rows, out


def _row_label(ws, r, first_blue_col):
    a = ws.cell(row=r, column=1).value
    if isinstance(a, str) and a.strip() and not is_formula(a):
        return a
    if first_blue_col is None:
        first_blue_col = ws.max_column + 1
    for col in range(first_blue_col - 1, 0, -1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.strip() and not is_formula(v):
            return v
    return ""


def _build_row(ws, r, blues, texts):
    label = _row_label(ws, r, blues[0].column)
    b = ws.cell(row=r, column=2).value
    unit = b if isinstance(b, str) and not is_formula(b) else ""
    all_text = [c for c in texts]
    joined = "\n".join(str(c.value) for c in all_text)
    marks = [m for m in MARKS if m in joined]
    primary = marks[0] if marks else ""
    blue_coords = {c.coordinate for c in blues}
    keep_blue_text = ws.title in ("DB_Evidence", "Checks")
    note_cells = [c for c in all_text if c.column > 2 and (c.coordinate not in blue_coords or keep_blue_text)]
    notes = " ｜ ".join(f"{c.coordinate}: {c.value}" for c in note_cells)
    return {
        "sheet": ws.title, "row": r, "label": label, "unit": unit,
        "blue": [(c.coordinate, c.column, c.value) for c in blues],
        "marks": marks, "primary": primary, "multi": len(marks) >= 2,
        "S": sorted(set(S_RE.findall(joined)), key=lambda x: int(x[1:])),
        "E": sorted(set(E_RE.findall(joined))),
        "notes": notes,
    }


def classify(r):
    """初步分類規則 R1–R7（先符合者為準）。回傳 (分類, 代碼)。"""
    if r["sheet"] == "DB_Evidence":
        return "Evidence 登錄（非輸入）", "R1"
    if r["sheet"] == "Checks":
        return "Checks 外部參照", "R2"
    p = r["primary"]
    if p in ("Verified", "Interested-party"):
        return "原始數據（候選）", "R3"
    if p in ("Analogy", "Assumed", "Decision"):
        return p, "R4"
    if p == "Derived":
        return "Derived 藍字（異常）", "R5"
    if r["sheet"].startswith("Sens_") or r["unit"] == "選擇" or "索引" in r["label"]:
        return "情境選擇", "R6"
    return "待分類", "R7"


def column_headers(ws, cells):
    """欄表頭：該欄由上往下第一個非空文字格（須在該藍字格之上）。"""
    out = {}
    for coord, col, _ in cells:
        r0 = ws[coord].row
        h = ""
        for rr in range(1, r0):
            v = ws.cell(row=rr, column=col).value
            if isinstance(v, str) and v.strip() and not is_formula(v):
                h = v
                break
        out[coord] = h
    return out


# ───────────────────────── 2. 靜態依賴圖 ─────────────────────────

_REF = re.compile(r"^(?:(?:'((?:[^']|'')+)'|([^!':]+))!)?(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)$")


def _cells_of(sheet, ref):
    c1, r1, c2, r2 = range_boundaries(ref.replace("$", ""))
    return [(sheet, f"{get_column_letter(c)}{r}") for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


class Graph:
    def __init__(self, wb):
        self.wb = wb
        self.names = {}                      # 名稱 → [(sheet, coord)]
        self.name_ref = {}                   # 名稱 → (sheet, ref 文字)
        for k, v in wb.defined_names.items():
            m = _REF.match(v.attr_text)
            if m:
                sheet = (m.group(1) or m.group(2) or "").replace("''", "'")
                self.names[k] = _cells_of(sheet, m.group(3))
                self.name_ref[k] = (sheet, m.group(3).replace("$", ""))
        self.dependents = collections.defaultdict(set)     # 儲存格 → 引用它的公式格
        self.formula_cells = set()
        self.functions = collections.Counter()
        self.unparsed = []                                  # 解析不了的參照
        self.dynamic = []                                   # OFFSET／INDIRECT 等
        self._build()

    def _build(self):
        for ws in self.wb:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if not is_formula(v):
                        continue
                    text = v if isinstance(v, str) else v.text
                    me = (ws.title, c.coordinate)
                    self.formula_cells.add(me)
                    for t in Tokenizer(text).items:
                        if t.type == "FUNC" and t.subtype == "OPEN":
                            fn = t.value[:-1].upper()
                            self.functions[fn] += 1
                            if fn in ("OFFSET", "INDIRECT", "INDEX_DYN", "HYPERLINK"):
                                self.dynamic.append((ws.title, c.coordinate, fn))
                        elif t.type == "OPERAND" and t.subtype == "RANGE":
                            for ref in self._resolve(t.value, ws.title):
                                self.dependents[ref].add(me)
                            # 解析不了者在 _resolve 內登記

    def _resolve(self, token, cur_sheet):
        m = _REF.match(token)
        if m:
            sheet = (m.group(1) or m.group(2) or cur_sheet).replace("''", "'")
            return _cells_of(sheet, m.group(3))
        if token in self.names:
            return self.names[token]
        self.unparsed.append((cur_sheet, token))
        return []

    def reach(self, seeds):
        """遞移下游（不含種子本身）；回傳公式格集合。"""
        seen, stack = set(), list(seeds)
        while stack:
            x = stack.pop()
            for d in self.dependents.get(x, ()):
                if d not in seen:
                    seen.add(d)
                    stack.append(d)
        return seen


def if_names(graph):
    return [n for n in graph.names if n.startswith("IF_")]


def output_cells(graph):
    out = {}
    for oid, name, k, *_ in OUTPUTS:
        sheet, ref = graph.name_ref[name]
        cells = _cells_of(sheet, ref)
        out[oid] = cells[k - 1]
    return out


def summarize_reach(graph, reach, seeds, outs):
    names = {}
    for n in if_names(graph):
        cnt = sum(1 for c in graph.names[n] if c in reach)
        if cnt:
            names[n] = cnt
    oids = [o for o, c in outs.items() if c in reach]
    sheets = sorted({s for s, _ in reach})
    return {"if_names": names, "outputs": oids, "sheets": sheets, "formula_cells": len(reach)}


# ───────────────────────── 3. 引擎與擾動 ─────────────────────────

_ENG = {}                                   # 每個 worker 一個引擎


def _init_worker(path, probe):
    from engine import Engine
    _ENG["e"] = Engine(path)
    _ENG["probe"] = probe
    _ENG["base"] = _probe(_ENG["e"], probe)
    _ENG["full0"] = _ENG["e"].evaluate_all()                 # 起點全簿快照：每次還原後逐格比對


def _probe(eng, probe):
    addrs = [eng._addr(s, c) for s, c in probe]
    vals = eng._xl.evaluate(addrs)
    return {k: ("" if v is None else v) for k, v in zip(probe, vals)}


def _set(eng, sheet, coord, value):
    """改寫輸入格。沒有任何公式引用的藍字格不在 pycel 的計算圖內（例：Cap_In 低／高欄），
    先取值一次讓它進入計算圖再設值；因無下游，改值不影響任何輸出（靜態圖亦同）。"""
    addr = eng._addr(sheet, coord)
    if addr not in eng._xl.cell_map:
        eng._xl.evaluate(addr)
    eng.set_input(sheet, coord, value)


def _changed(a, b):
    if is_num(a) and is_num(b):
        return abs(a - b) > 1e-12 * max(abs(a), abs(b), 1e-300)
    return a != b


def _worker(job):
    """job＝(inv_id, [(sheet, coord, base)], 方向列表)。回傳每個方向的探測值與改變格。"""
    eng, probe, base = _ENG["e"], _ENG["probe"], _ENG["base"]
    inv, cells, dirs = job
    res = {}
    for d in dirs:
        err = ""
        try:
            for s, c, v in cells:
                _set(eng, s, c, v * (1 + d))
            now = _probe(eng, probe)
        except Exception as ex:                                   # 引擎層級例外也要留痕
            now, err = None, f"{type(ex).__name__}: {ex}"
        finally:
            for s, c, v in cells:
                try:
                    _set(eng, s, c, v)
                except Exception as ex:
                    err += f" 還原失敗 {s}!{c}: {ex}"
        full1 = eng.evaluate_all()                                # 還原後全簿重算，與起點逐格比對
        bad = sum(1 for k, v in _ENG["full0"].items() if _changed(v, full1[k]))
        if now is None:
            res[d] = {"err": err, "restore_bad": bad}
            continue
        changed = [k for k in probe if _changed(base[k], now[k])]
        res[d] = {"vals": now, "changed": changed, "restore_bad": bad}
    return inv, res


# ───────────────────────── 4. 擾動規則 ─────────────────────────


def perturbable(r, cls):
    """回傳 (可擾動格 [(sheet, coord, base)], 原因代碼或 None, 單邊旗標)。"""
    if cls in ("Evidence 登錄（非輸入）", "Checks 外部參照"):
        return [], "X4", False
    nums = [(c, col, v) for c, col, v in r["blue"] if is_num(v)]
    if not nums:
        return [], "X1", False
    selection_row = (cls == "情境選擇" or r["unit"] == "選擇" or "索引" in r["label"]
                     or bool(OPTION_RE.search(r["label"])) or bool(OPTION_FLAG_RE.search(r["label"])))
    keep = []
    for c, col, v in nums:
        if selection_row and float(v).is_integer():          # 整數選擇格（逐格判定；非整數覆寫值仍可擾動）
            continue
        keep.append((r["sheet"], c, v))
    if not keep:
        return [], "X2", False
    nz = [x for x in keep if x[2] != 0]
    if not nz:
        return [], "X3", False
    one_sided = r["unit"] in PCT_UNITS and any(abs(v) * 1.1 > 1 for _, _, v in nz)
    return nz, None, one_sided


def elasticities(y0, yp, ym, one_sided):
    ep = (yp / y0 - 1) / 0.1 if yp is not None else None
    em = (1 - ym / y0) / 0.1 if ym is not None else None
    if ep is not None and em is not None and not one_sided:
        e = (yp - ym) / (y0 * 0.2)
    else:
        e = em if em is not None else ep
    return ep, em, e


def grade(score):
    if score is None or score == 0:
        return "無"
    return "高" if score >= GRADE_HI else ("中" if score >= GRADE_MID else "低")


# ───────────────────────── 5. LibreOffice ─────────────────────────


def lo_recalc(xlsx: Path, outdir: Path) -> Path:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="lo_profile_"))
    try:
        cmd = ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
               "--convert-to", "xlsx", "--outdir", str(outdir), str(xlsx)]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = outdir / xlsx.name
    if not out.exists():
        raise RuntimeError(f"LibreOffice 重算失敗：{r.stdout}\n{r.stderr}")
    return out


def lo_outputs(path, outs):
    wb = openpyxl.load_workbook(path, data_only=True)
    return {o: wb[s][c].value for o, (s, c) in outs.items()}


# ───────────────────────── 6. 主流程 ─────────────────────────


def build_inventory(xlsx: Path, workers: int, do_perturb: bool, lo_n: int, log=print):
    t0 = time.time()
    sha_before = sha256(xlsx)
    wb = openpyxl.load_workbook(xlsx)
    rows, bform = scan_workbook(wb)
    graph = Graph(wb)
    outs = output_cells(graph)
    log(f"掃描完成：{len(rows)} 列；靜態圖公式格 {len(graph.formula_cells)}；函數 {sorted(graph.functions)}")

    # 分類、標頭、編號
    cells_table, inv = [], []
    for i, r in enumerate(rows, 1):
        r["inv"] = f"INV-{i:03d}"
        r["cls"], r["rule"] = classify(r)
        hdr = column_headers(wb[r["sheet"]], r["blue"])
        for coord, col, v in r["blue"]:
            cells_table.append((r["inv"], r["sheet"], coord, v, "數值" if is_num(v) else "文字", hdr[coord]))
        # 靜態可達
        seeds = {(r["sheet"], c) for c, _, _ in r["blue"]}
        reach = graph.reach(seeds)
        r["reach"] = summarize_reach(graph, reach, seeds, outs)
    log(f"依賴圖完成（{time.time() - t0:.1f}s）")

    # 擾動候選
    for r in rows:
        cells, why, one = perturbable(r, r["cls"])
        r["pcells"], r["skip"], r["one_sided"] = cells, why, one
        if why is None and not r["reach"]["outputs"]:
            r["pcells"], r["skip"] = [], ("X5a" if not r["reach"]["if_names"] else "X5b")
    result = {"sha_before": sha_before, "rows": rows, "cells": cells_table, "bformulas": bform, "graph": graph,
              "outs": outs, "wb": wb}
    if do_perturb:
        run_perturbation(result, xlsx, workers, lo_n, log)
    result["sha_after"] = sha256(xlsx)
    result["seconds"] = time.time() - t0
    return result


def run_perturbation(res, xlsx, workers, lo_n, log):
    from engine import Engine
    rows, graph, outs = res["rows"], res["graph"], res["outs"]
    # 探測集合：Interface 全表＋全部 IF_ 名稱格＋O01–O14＋離散事件
    wb = res["wb"]
    probe = set()
    ws = wb["Interface"]
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, ""):
                probe.add(("Interface", c.coordinate))
    for n in if_names(graph):
        probe.update(graph.names[n])
    probe.update(outs.values())
    probe = sorted(probe)
    if_cells = {n: set(graph.names[n]) for n in if_names(graph)}

    t0 = time.time()
    eng = Engine(xlsx)
    base = _probe(eng, probe)
    log(f"引擎建立＋基準探測 {time.time() - t0:.1f}s；探測格 {len(probe)}")
    # 基準 vs 快取值
    cached = openpyxl.load_workbook(xlsx, data_only=True)
    res["O_base"] = {}
    for oid, name, k, desc, unit, tbl in OUTPUTS:
        s, c = outs[oid]
        y = base[(s, c)]
        cv = cached[s][c].value
        res["O_base"][oid] = {"engine": y, "cached": cv, "table": tbl}
    full0 = eng.evaluate_all()                    # 起點全簿快照（還原驗證用）
    del eng

    jobs = []
    for r in rows:
        if not r["pcells"]:
            continue
        dirs = [-0.1] if r["one_sided"] else [0.1, -0.1]
        jobs.append((r["inv"], r["pcells"], dirs))
    log(f"擾動 {len(jobs)} 列（{sum(len(j[2]) for j in jobs)} 次重算），workers={workers}")
    results = {}
    t1 = time.time()
    if workers > 1:
        with mp.get_context("fork").Pool(workers, initializer=_init_worker, initargs=(str(xlsx), probe)) as pool:
            for k, (inv, rr) in enumerate(pool.imap_unordered(_worker, jobs, chunksize=1), 1):
                results[inv] = rr
                if k % 25 == 0:
                    log(f"  {k}/{len(jobs)}  {time.time() - t1:.0f}s")
    else:
        _init_worker(str(xlsx), probe)
        for k, j in enumerate(jobs, 1):
            inv, rr = _worker(j)
            results[inv] = rr
            if k % 25 == 0:
                log(f"  {k}/{len(jobs)}  {time.time() - t1:.0f}s")
    res["perturb_seconds"] = time.time() - t1

    # 還原驗證：新引擎重算全簿，與起點逐格比對（另以同一 worker 引擎的基準於 worker 結束時不再可得）
    eng = Engine(xlsx)
    full1 = eng.evaluate_all()
    res["restore_diff"] = sum(1 for k in full0 if _changed(full0[k], full1[k]))
    res["restore_cells"] = len(full0)
    res["restore_runs"] = sum(len(v) for v in results.values())
    res["restore_bad_runs"] = sum(1 for v in results.values() for x in v.values() if x.get("restore_bad"))
    log(f"還原檢查：每次擾動後全簿 {res['restore_cells']} 格重算，{res['restore_runs']} 次中不一致 {res['restore_bad_runs']} 次；"
        f"全部完成後新引擎基準 vs 起點不一致 {res['restore_diff']} 格")

    # 彙整
    sens, viol, errs = [], [], []
    for r in rows:
        r["max_eps"], r["max_out"], r["iface_changed"], r["flips"], r["observed"] = None, "", 0, [], set()
        r["sens"] = {}
        rr = results.get(r["inv"])
        if rr is None:
            continue
        iface_changed = set()
        for d, x in rr.items():
            if "err" in x:
                errs.append((r["inv"], "引擎例外", x["err"]))
                continue
            for k in x["changed"]:
                if k[0] == "Interface":
                    iface_changed.add(k)
            for n, cs in if_cells.items():
                if any(k in cs for k in x["changed"]):
                    r["observed"].add(n)
        r["iface_changed"] = len(iface_changed)
        for oid, name, k, desc, unit, tbl in OUTPUTS:
            s, c = outs[oid]
            y0 = base[(s, c)]
            yp = rr.get(0.1, {}).get("vals", {}).get((s, c)) if 0.1 in rr else None
            ym = rr.get(-0.1, {}).get("vals", {}).get((s, c)) if -0.1 in rr else None
            note = ""
            bad = [v for v in (yp, ym) if isinstance(v, str)]
            if bad:
                note = "；".join(sorted(set(str(v) for v in bad)))
            yp_n, ym_n = (yp if is_num(yp) else None), (ym if is_num(ym) else None)
            if not is_num(y0) or y0 == 0 or (yp_n is None and ym_n is None):
                ep = em = e = None
            else:
                ep, em, e = elasticities(y0, yp_n, ym_n, r["one_sided"])
            r["sens"][oid] = (y0, yp, ym, ep, em, e, "單邊" if r["one_sided"] else "", note)
            if note:
                errs.append((r["inv"], oid, note))
        # 離散事件：前緣名稱／模型是否改變
        for d, x in rr.items():
            if "changed" in x:
                for n, cs in if_cells.items():
                    if (n == "IF_FrontSuccVRName" or n.startswith("IF_FrontModel_")) and any(k in cs for k in x["changed"]):
                        r["flips"].append(f"{n}（{d:+.0%}）")
        scores = [(abs(v[5]), oid) for oid, v in r["sens"].items()
                  if oid != "O14" and v[5] is not None]
        if scores:
            r["max_eps"], r["max_out"] = max(scores)
        else:
            r["max_eps"] = 0.0
        # 觀測 ⊆ 靜態
        extra = r["observed"] - set(r["reach"]["if_names"])
        if extra:
            viol.append((r["inv"], sorted(extra)))
        r["missed_static"] = len(set(r["reach"]["if_names"]) - r["observed"])
    res["violations"], res["errors"] = viol, errs
    # 排名
    perturbed = [r for r in rows if r["inv"] in results]
    perturbed.sort(key=lambda r: (-(r["max_eps"] or 0), -r["iface_changed"]))
    for i, r in enumerate(perturbed, 1):
        r["rank"] = i
    for r in rows:
        r["grade"] = grade(r["max_eps"]) if r["inv"] in results else "無"
    res["results"] = results
    res["base"] = base
    res["probe"] = probe

    if lo_n:
        lo_spot_check(res, xlsx, perturbed[:lo_n], log)


def lo_spot_check(res, xlsx, top, log):
    outs = res["outs"]
    out = []
    tmp = Path(tempfile.mkdtemp(prefix="s0_lo_"))
    try:
        for r in top:
            d = -0.1 if r["one_sided"] else 0.1
            wb = openpyxl.load_workbook(xlsx)
            for s, c, v in r["pcells"]:
                wb[s][c].value = v * (1 + d)
            f = tmp / f"{r['inv']}.xlsx"
            wb.save(f)
            rec = lo_recalc(f, tmp / "out")
            lo = lo_outputs(rec, outs)
            worst, rows_ = 0.0, {}
            for oid, name, *_ in OUTPUTS:
                eng = res["results"][r["inv"]][d]["vals"][outs[oid]]
                x = lo[oid]
                if is_num(eng) and is_num(x):
                    rel = abs(eng - x) / max(abs(eng), abs(x), 1e-300) if eng != x else 0.0
                else:
                    rel = 0.0 if eng == x else float("inf")
                rows_[oid] = (eng, x, rel)
                worst = max(worst, rel)
            out.append({"inv": r["inv"], "dir": d, "worst": worst, "rows": rows_})
            log(f"  LO 抽驗 {r['inv']} {d:+.0%} 最大相對誤差 {worst:.2e}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    res["lo_check"] = out


# ───────────────────────── 7. 寫 Excel ─────────────────────────

HDR_FILL = PatternFill("solid", fgColor="DDE6F0")


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append([clean(x) for x in r])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h)
        ws.column_dimensions[get_column_letter(i)].width = w or min(max(len(str(h)) * 2 + 2, 10), 40)
    return ws


def counts(res):
    rows = res["rows"]
    by_sheet = collections.Counter(r["sheet"] for r in rows)
    cls = collections.Counter(r["cls"] for r in rows)
    non12 = [r for r in rows if r["rule"] not in ("R1", "R2")]
    prim = collections.Counter((r["primary"] or "無標記") for r in rows)       # 不套 R1、R2：全部 324 列
    multi = [r for r in rows if r["multi"]]
    return {
        "by_sheet": by_sheet, "cls": cls, "prim": prim, "multi": multi, "non12": non12,
        "with_S": sum(1 for r in rows if r["S"]),
        "ncells": len(res["cells"]),
    }


def expectation_table(res):
    c = counts(res)
    rows = res["rows"]
    t = []
    t.append(("總計", "藍字列", EXPECT_TOTALS["藍字列"], len(rows)))
    t.append(("總計", "藍字格", EXPECT_TOTALS["藍字格"], c["ncells"]))
    t.append(("總計", "藍字公式格", EXPECT_TOTALS["藍字公式格"], len(res["bformulas"])))
    for s, n in EXPECT_ROWS_BY_SHEET.items():
        t.append(("工作表列數", s, n, c["by_sheet"].get(s, 0)))
    for k, n in EXPECT_CLASS.items():
        t.append(("初步分類", k, n, c["cls"].get(k, 0)))
    for k, n in EXPECT_PRIMARY.items():
        t.append(("主標記分布", k, n, c["prim"].get(k, 0)))
    mm = collections.Counter(r["sheet"] for r in c["multi"])
    t.append(("多標記", "多標記列", EXPECT_MULTI["多標記列"], len(c["multi"])))
    t.append(("多標記", "Arch", EXPECT_MULTI["多標記：Arch"], mm.get("Arch", 0)))
    t.append(("多標記", "Spec_Rack", EXPECT_MULTI["多標記：Spec_Rack"], mm.get("Spec_Rack", 0)))
    t.append(("多標記", "有 S 編號列", EXPECT_MULTI["有 S 編號列"], c["with_S"]))
    t.append(("動態引用", "OFFSET／INDIRECT 公式數", 0, len(res["graph"].dynamic)))
    return t


def restore_composition(res):
    """本母體（藍字格）與 builder/preserve.py 快照（restore_log 的『藍字輸入』）的差異，依工作表。builder 不在時回傳空。"""
    try:
        sys.path.insert(0, str(REPO / "builder"))
        import preserve
    except Exception:
        return []
    snap = collections.Counter(k[0] for k in preserve.snapshot(res["wb"]))
    mine = collections.Counter()
    for r in res["rows"]:
        mine[r["sheet"]] += len(r["blue"])
    out = []
    for sh in res["wb"].sheetnames:
        if not (mine[sh] or snap[sh]):
            continue
        rebuilt = sh in preserve.REBUILT
        diff = mine[sh] - snap[sh]
        if diff == 0:
            why = ""
        elif not rebuilt:
            why = "builder 不重建此頁（輸入頁或登錄頁）"
        else:
            why = "重建頁內不在快照：A 欄藍字（如 Cap_In 模型名稱列）或 Spec_Rack 第 20 列以前（Block 1 區，由 Excel 擁有）"
        out.append((sh, rebuilt, mine[sh], snap[sh], why))
    return out


def write_workbook(res, out: Path, xlsx: Path):
    rows, graph, outs = res["rows"], res["graph"], res["outs"]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    perturbed = "results" in res

    # README
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 60
    L = lambda *a: ws.append([clean(x) for x in a])
    L(f"Stage 0 盤點（唯讀）— {xlsx.name}")
    L("產生程式", "tools/stage0_inventory.py（輸入 xlsx 路徑，輸出本檔；可重跑，不在 CI）")
    L("活頁簿 SHA-256", res["sha_before"], "", "", "", "前後一致：" + ("是" if res["sha_before"] == res["sha_after"] else "否"))
    L()
    L("目的", "盤點藍字輸入：依既有標記機械分類、建立靜態依賴圖、以擾動求 load-bearing 排序。分類只是『初步』；待分類與爭議列由 chat 端複核，Andy 核准遷移範圍（Gate 0）。")
    L("母體", "藍字格＝字型色 RGB 結尾 0000FF、值非空、非公式；盤點單位＝工作表 × 列。藍字但為公式者另列（藍字公式格）。")
    L()
    L("初步分類規則（依序，先符合者為準）")
    for code, cond, cl in [
        ("R1", "工作表＝DB_Evidence", "Evidence 登錄（非輸入）"), ("R2", "工作表＝Checks", "Checks 外部參照"),
        ("R3", "主標記＝Verified 或 Interested-party", "原始數據（候選）"),
        ("R4", "主標記＝Analogy／Assumed／Decision", "同主標記"), ("R5", "主標記＝Derived", "Derived 藍字（異常）"),
        ("R6", "無標記，且工作表以 Sens_ 開頭、或單位＝「選擇」、或標籤含「索引」", "情境選擇"),
        ("R7", "其餘", "待分類")]:
        L(code, cond, cl)
    L("標記擷取", "標記＝該列所有非公式文字格（含 A、B 欄與藍字文字格）的子字串比對（區分大小寫）：Verified、Interested-party、Analogy、Assumed、Decision、Derived；主標記依序取第一個；兩個以上者標『多標記』。")
    L("擾動方法", "engine（pycel）重算；每列所有數值藍字格同時 ×1.1 與 ×0.9，之後還原；ε₊＝(y₊/y₀−1)/0.1、ε₋＝(1−y₋/y₀)/0.1、ε＝(y₊−y₋)/(y₀×0.2)；單位 % 且 ×1.1>1 者只做 −10%（單邊）。")
    L("整數選擇格（本工具的解讀）", "列屬『情境選擇』、或單位＝選擇、或標籤含『索引』、或標籤含『（1／2／3）』類選項說明，或標籤含『（1＝有）』『（1＝是）』類 0／1 旗標說明時，其中值為整數的藍字格不擾動；同列非整數數值格仍擾動。X2＝該列沒有任何可擾動格。")
    L("不擾動原因代碼", "X1 文字值（列內無數值格）；X2 整數選擇格；X3 基準值為 0；X4 Evidence／Checks 列；X5a 靜態不可達任何 IF_；X5b 可達 IF_ 但未達 O01–O14（彈性記 0）。")
    L("分段（僅供排序閱讀）", "高 ≥ 0.5；中 0.1–0.5；低 < 0.1；無＝0（含未擾動）。分數＝O01–O13 的 |ε| 最大值。門檻由 chat 端日後判斷是否調整。")
    L()
    L("期望值核對表（第 2、3.3 節）")
    L("類別", "項目", "期望", "實得", "差異")
    for cat, item, e, a in expectation_table(res):
        L(cat, item, e, a, a - e)
    L()
    comp = restore_composition(res)
    if comp:
        L("與 builder restore_log 口徑的差異（依工作表）")
        L("工作表", "builder 是否重建", "本母體藍字格", "restore 快照格", "差異", "差異來源")
        for sh, rebuilt, a, b, why in comp:
            L(sh, "是" if rebuilt else "否", a, b, a - b, why)
        L("合計", "", sum(x[2] for x in comp), sum(x[3] for x in comp), sum(x[2] - x[3] for x in comp),
          "restore_log 口徑：builder 重建頁、欄 ≥ B、字型色 FF0000FF、列標籤（A 欄）非公式；本母體另含不重建的頁、A 欄藍字與 0000FF 尾碼的其他色")
    L()
    L("依賴圖")
    L("公式格數", len(graph.formula_cells))
    L("使用函數", "、".join(f"{k}×{v}" for k, v in sorted(graph.functions.items())))
    L("動態引用（OFFSET／INDIRECT）", len(graph.dynamic), "；".join(f"{s}!{c}:{f}" for s, c, f in graph.dynamic[:20]))
    L("無法解析的參照", len(graph.unparsed), "；".join(f"{s}:{t}" for s, t in graph.unparsed[:20]))
    L("區域引用", "一律展開為每一格；INDEX、CHOOSE 因此為保守上界。")
    if perturbed:
        L()
        L("驗證（第 5.4 節）")
        L("1. 觀測 ⊆ 靜態：違反數", len(res["violations"]), "；".join(f"{i}:{v}" for i, v in res["violations"][:20]))
        missed = sum(r.get("missed_static", 0) for r in rows if "missed_static" in r)
        L("   靜態可達但未觀測到改變（名稱 × 列，僅計數）", missed)
        if "lo_check" in res:
            w = max((x["worst"] for x in res["lo_check"]), default=0)
            L("2. LibreOffice 抽驗（分數最高 %d 列）最大相對誤差" % len(res["lo_check"]), w, "通過" if w <= REL else "未通過")
        L("3. 還原：每次擾動後全簿公式格重算（%d 格）" % res["restore_cells"], f"{res['restore_runs']} 次中不一致 {res['restore_bad_runs']} 次")
        L("   全部完成後新引擎基準 vs 起點", res["restore_cells"], f"不一致 {res['restore_diff']} 格")
        L("4. 擾動執行時間（秒）", round(res["perturb_seconds"], 1), f"總執行 {res['seconds']:.0f} 秒")
        bad = max((abs(v["engine"] - v["cached"]) / max(abs(v["cached"]), 1e-300)
                   for v in res["O_base"].values() if is_num(v["engine"]) and is_num(v["cached"])), default=0)
        L("基準 engine vs 活頁簿快取值（O01–O14）最大相對誤差", bad, "通過" if bad <= REL else "未通過")
    ws["A1"].font = Font(bold=True, size=13)

    # 盤點
    hdr = ["INV_ID", "工作表", "列", "標籤", "單位", "藍字格位址", "藍字格數", "基準值", "全部標記", "主標記", "多標記",
           "來源編號（S）", "證據編號（E）", "說明原文", "初步分類", "分類依據", "可達 IF_ 數", "可達主要輸出",
           "最大 |ε|", "分段", "排名", "可達公式格數", "可達工作表", "不擾動原因",
           "chat 複核分類", "SRC 分頁候選", "理由", "Gate 0 爭議"]
    body = []
    for r in rows:
        body.append([
            r["inv"], r["sheet"], r["row"], r["label"], r["unit"], "、".join(c for c, _, _ in r["blue"]), len(r["blue"]),
            "；".join(fmt_num(v) for _, _, v in r["blue"] if is_num(v)), "、".join(r["marks"]), r["primary"],
            "多標記" if r["multi"] else "", "、".join(r["S"]), "、".join(r["E"]), r["notes"], r["cls"], r["rule"],
            len(r["reach"]["if_names"]), "、".join(r["reach"]["outputs"]),
            r.get("max_eps") if perturbed and r["inv"] in res["results"] else None,
            r.get("grade", "") if perturbed else "", r.get("rank") if perturbed else None,
            r["reach"]["formula_cells"], "、".join(r["reach"]["sheets"]), r["skip"] or "", "", "", "", ""])
    _sheet(wb, "盤點", hdr, body, {"標籤": 36, "說明原文": 60, "基準值": 26, "藍字格位址": 20, "可達主要輸出": 30, "可達工作表": 30})

    _sheet(wb, "藍字格", ["INV_ID", "工作表", "格位址", "值", "型別", "欄表頭"],
           [(i, s, c, v if not isinstance(v, bool) else str(v), t, h) for i, s, c, v, t, h in res["cells"]],
           {"值": 40, "欄表頭": 24})
    _sheet(wb, "藍字公式格", ["工作表", "格位址", "公式", "該列標籤"], res["bformulas"], {"公式": 70, "該列標籤": 40})

    dep = []
    for r in rows:
        for n, cnt in sorted(r["reach"]["if_names"].items()):
            obs = ""
            if perturbed and r["inv"] in res["results"]:
                obs = "是" if n in r["observed"] else "否"
            dep.append((r["inv"], n, cnt, obs))
    _sheet(wb, "依賴", ["INV_ID", "IF_ 名稱", "名稱內可達格數", "擾動是否觀測到改變（空白＝未擾動）"], dep, {"IF_ 名稱": 34})

    if perturbed:
        sens = []
        for r in rows:
            if r["inv"] not in res["results"]:
                continue
            for oid, *_ in OUTPUTS:
                y0, yp, ym, ep, em, e, one, note = r["sens"][oid]
                sens.append((r["inv"], oid, y0, yp, ym, ep, em, e, one, note))
        _sheet(wb, "敏感度", ["INV_ID", "輸出 ID", "y₀", "y₊", "y₋", "ε₊", "ε₋", "ε", "單邊", "錯誤／文字結果"], sens)

        rk = []
        for r in sorted((r for r in rows if r["inv"] in res["results"]), key=lambda r: r["rank"]):
            e14 = r["sens"]["O14"][5]
            rk.append((r["rank"], r["inv"], r["sheet"], r["row"], r["label"], r["cls"], r["primary"], r["max_eps"], r["max_out"],
                       r["grade"], r["iface_changed"], "；".join(r["flips"]), len(r["pcells"]), r["one_sided"] and "單邊" or "",
                       abs(e14) if e14 is not None else "（文字或無值）"))
        _sheet(wb, "排序", ["排名", "INV_ID", "工作表", "列", "標籤", "初步分類", "主標記", "最大 |ε|（O01–O13）", "所屬輸出", "分段",
                            "Interface 改變格數", "離散事件（前緣翻轉）", "擾動格數", "單邊", "O14 |ε|（另列）"], rk, {"標籤": 36})

        top = []
        for oid, *_ in OUTPUTS:
            ranked = sorted(((abs(r["sens"][oid][5]), r) for r in rows if r["inv"] in res["results"] and r["sens"][oid][5] is not None),
                            key=lambda x: -x[0])[:5]
            for k, (e, r) in enumerate(ranked, 1):
                top.append((oid, k, r["inv"], r["sheet"], r["row"], r["label"], r["cls"], r["primary"], e))
        _sheet(wb, "各輸出前5", ["輸出 ID", "名次", "INV_ID", "工作表", "列", "標籤", "初步分類", "主標記", "|ε|"], top, {"標籤": 36})

    nz = []
    for r in rows:
        if r["skip"]:
            nz.append((r["inv"], r["sheet"], r["row"], r["label"], r["cls"], r["skip"], "、".join(r["reach"]["outputs"]),
                       len(r["reach"]["if_names"])))
    _sheet(wb, "不擾動", ["INV_ID", "工作表", "列", "標籤", "初步分類", "原因代碼", "靜態可達主要輸出", "可達 IF_ 數"], nz, {"標籤": 36})

    ob = []
    for oid, name, k, desc, unit, tbl in OUTPUTS:
        o = res.get("O_base", {}).get(oid)
        s, c = outs[oid]
        if o:
            dlt = abs(o["engine"] - o["cached"]) / max(abs(o["cached"]), 1e-300) if is_num(o["engine"]) and is_num(o["cached"]) else None
            dtb = abs(o["engine"] - tbl) / max(abs(tbl), 1e-300) if is_num(o["engine"]) else None
            ob.append((oid, name, k, f"{s}!{c}", desc, unit, tbl, o["cached"], o["engine"], dlt, dtb))
        else:
            ob.append((oid, name, k, f"{s}!{c}", desc, unit, tbl, None, None, None, None))
    _sheet(wb, "主要輸出", ["ID", "具名範圍", "欄", "儲存格", "內容", "單位", "指令表基準值（6 位有效）", "活頁簿快取值（LibreOffice）",
                            "engine 基準值", "engine vs 快取 相對誤差", "engine vs 指令表 相對誤差"], ob, {"內容": 36, "具名範圍": 26})

    an = []
    for r in res["rows"]:
        if r["multi"]:
            an.append(("多標記", r["inv"], r["sheet"], r["row"], r["label"], "、".join(r["marks"]), r["cls"]))
    for r in res["rows"]:
        if r["cls"] == "Derived 藍字（異常）":
            an.append(("Derived 藍字", r["inv"], r["sheet"], r["row"], r["label"], "、".join(r["marks"]), r["cls"]))
    for s, a, f, lab in res["bformulas"]:
        an.append(("藍字公式格", "", s, a, lab, f, ""))
    for r in res["rows"]:
        if r["skip"] in ("X5a", "X5b") and r["cls"] in ("原始數據（候選）", "Analogy", "Assumed"):
            an.append(("靜態不可達（原始數據／Analogy／Assumed）", r["inv"], r["sheet"], r["row"], r["label"], r["skip"], r["cls"]))
    for r in res["rows"]:                                      # 孤立輸入：藍字格沒有任何公式引用（改它不會改變任何公式格）
        if r["cls"] in ("Evidence 登錄（非輸入）", "Checks 外部參照"):
            continue
        for c, _, v in r["blue"]:
            if not graph.dependents.get((r["sheet"], c)):
                an.append(("藍字格無任何公式引用", r["inv"], r["sheet"], c, r["label"], str(v), r["cls"]))
    for inv, oid, msg in res.get("errors", []):
        an.append(("擾動錯誤／文字結果", inv, "", "", oid, msg, ""))
    for s, t in graph.unparsed:
        an.append(("無法解析的參照", "", s, "", "", t, ""))
    _sheet(wb, "異常", ["類型", "INV_ID", "工作表", "列／格", "標籤／輸出", "內容", "初步分類"], an, {"內容": 60, "標籤／輸出": 36})
    wb.save(out)


# ───────────────────────── 8. 影響範圍（Stage 1 重用） ─────────────────────────


def impact(xlsx: Path, targets):
    wb = openpyxl.load_workbook(xlsx)
    g = Graph(wb)
    outs = output_cells(g)
    seeds = set()
    for t in targets:
        s, a = t.split("!")
        seeds.update(_cells_of(s.strip("'"), a))
    reach = g.reach(seeds)
    sm = summarize_reach(g, reach, seeds, outs)
    print(f"種子 {len(seeds)} 格；可達公式格 {sm['formula_cells']}；工作表 {', '.join(sm['sheets'])}")
    print("可達主要輸出：", "、".join(sm["outputs"]) or "—")
    print("可達 IF_ 名稱（名稱內可達格數）：")
    for n, c in sorted(sm["if_names"].items()):
        print(f"  {n}  {c}/{len(g.names[n])}")
    return sm


# ───────────────────────── 7. Gate 1 驗收（第 11 輪；唯讀） ─────────────────────────

GOVERNANCE_SHEETS = {"Gov_Map", "Decisions", "SRC_HW", "SRC_DC", "SRC_Model", "SRC_Perf", "L1", "Checks", "Sources",
                     "DB_Evidence", "README"}
UNIT_CONSTS = {0, 1, 2, 3, 4, 8, 10, 12, 24, 60, 100, 168, 365, 1000, 3600, 8760,
               1e3, 1e6, 1e9, 1e12, 1e15, 1e18}               # 單位換算常數（指令第 3.2 節；0 為本工具加入的空值）
POSITION_FUNCS = {"INDEX", "CHOOSE", "MATCH"}                   # 位置索引參數不計（MATCH 的比對型態 0／1／-1）
SRC_PERTURB = ("SRC_HW", "C42", 1.1)                            # 驗收測試：SRC_HW_038 ×1.1


def _cmp_sheet(a, b, sheet, rel=REL):
    """兩份已重算活頁簿的同名工作表逐格比對；回傳 (格數, 差異清單[(coord, x, y, 相對差)])。"""
    wa, wb_ = a[sheet], b[sheet]
    coords = {c.coordinate for r in wa.iter_rows() for c in r if c.value is not None} | \
             {c.coordinate for r in wb_.iter_rows() for c in r if c.value is not None}
    diffs = []
    for co in coords:
        x, y = wa[co].value, wb_[co].value
        if x == y:
            continue
        if is_num(x) and is_num(y):
            d = abs(x - y) / max(abs(x), abs(y), 1e-300)
            if d <= rel:
                continue
            diffs.append((co, x, y, d))
        else:
            diffs.append((co, x, y, None))
    return len(coords), diffs


def scan_constants(wb):
    """G0-5：模型頁公式內的數值常數（排除單位換算常數與 INDEX／CHOOSE／MATCH 的位置參數）。"""
    rows = []
    kept = collections.defaultdict(list)           # 被排除的整數常數（2、3、4…）：依「工作表＋常數＋公式型態」歸併，供人工複核
    for ws in wb:
        if ws.title in GOVERNANCE_SHEETS:
            continue
        for row in ws.iter_rows():
            for c in row:
                if not is_formula(c.value):
                    continue
                text = c.value if isinstance(c.value, str) else c.value.text
                stack, argi = [], []
                for tk in Tokenizer(text).items:
                    if tk.type == "FUNC" and tk.subtype == "OPEN":
                        stack.append(tk.value[:-1].upper()); argi.append(0)
                    elif tk.type == "FUNC" and tk.subtype == "CLOSE":
                        stack.pop(); argi.pop()
                    elif tk.type == "SEP" and tk.subtype == "ARG" and argi:
                        argi[-1] += 1
                    elif tk.type == "OPERAND" and tk.subtype == "NUMBER":
                        v = float(tk.value)
                        if v in UNIT_CONSTS and v not in (0, 1) and v < 1e3 and not (stack and stack[-1] in POSITION_FUNCS and argi[-1] >= 1):
                            kept[(ws.title, tk.value, re.sub(r"\$?[A-Z]{1,3}\$?\d+", "#", text)[:90])].append(c.coordinate)
                        if v in UNIT_CONSTS or v >= 1e90:        # 1e90 以上＝MIN() 的「無窮大」哨兵值（Harness 前緣 9E+99；Price_Frontier 的 1E9 已在單位表），非數據
                            continue
                        if stack and stack[-1] in POSITION_FUNCS and (stack[-1] == "CHOOSE" and argi[-1] == 0 or stack[-1] != "CHOOSE" and argi[-1] >= 1):
                            continue
                        rows.append((ws.title, c.coordinate, text, tk.value))
    scan_constants.excluded = kept
    return rows


def gate1(xlsx: Path, prev: Path, out: Path, workdir: Path, log=print):
    t0 = time.perf_counter()
    res = {}
    wd = workdir; wd.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(xlsx)
    graph = Graph(wb)
    # —— LibreOffice 重算：現行、前版、兩個擾動檔
    def stage(name, src_wb_edit=None):
        f = wd / f"{name}.xlsx"
        w = openpyxl.load_workbook(xlsx)
        if src_wb_edit:
            src_wb_edit(w)
        w.save(f)
        o = lo_recalc(f, wd / "lo")
        return openpyxl.load_workbook(o, data_only=True)
    base = stage("base")
    prevf = wd / "prev.xlsx"; shutil.copy(prev, prevf)
    old = openpyxl.load_workbook(lo_recalc(prevf, wd / "lo_prev"), data_only=True)
    def e_hw(w):
        s, c, k = SRC_PERTURB; w[s][c].value = w[s][c].value * k
    hw = stage("hw038", e_hw)
    gw2 = stage("gw2", lambda w: w["Inputs"].__setitem__("E5", 2))
    log(f"重算完成 {time.perf_counter() - t0:.0f}s")

    # —— 3.1 對 v5.10：Interface 與全部模型頁
    n_if, d_if = _cmp_sheet(old, base, "Interface")
    model_diff = {}
    for s in old.sheetnames:
        if s in GOVERNANCE_SHEETS or s not in base.sheetnames:
            continue
        n, d = _cmp_sheet(old, base, s)
        model_diff[s] = (n, len(d), d[:5])
    res["3.1"] = {"interface_cells": n_if, "interface_diffs": len(d_if), "interface_diff_sample": d_if[:5], "model_pages": model_diff}

    # —— 3.2 原始數據寫死
    gm = wb["Gov_Map"]; hard, linked, cat = [], 0, collections.Counter()
    for r in range(5, gm.max_row + 1):
        if gm.cell(r, 1).value is None:
            continue
        cat[gm.cell(r, 6).value] += 1
        if gm.cell(r, 6).value != "原始數據":
            continue
        sheet, ref = gm.cell(r, 3).value, str(gm.cell(r, 4).value)
        for s_, c_ in _cells_of(sheet, ref):
            v = wb[s_][c_].value
            txt = v if isinstance(v, str) else getattr(v, "text", "")
            if is_formula(v) and "SRC_" in txt:
                linked += 1
            else:
                hard.append((gm.cell(r, 1).value, s_, c_, v))
    consts = scan_constants(wb)
    res["3.2"] = {"gov_categories": dict(cat), "raw_linked": linked, "raw_hardcoded": hard, "constants": consts}

    # —— 3.3 Evidence
    ev_ids = {str(c.value) for c in wb["DB_Evidence"]["A"][4:] if c.value}
    missing, n_active = [], 0
    for s in ("SRC_HW", "SRC_DC", "SRC_Model", "SRC_Perf"):
        ws = wb[s]
        for r in range(5, ws.max_row + 1):
            if ws.cell(r, 1).value is None or ws.cell(r, 15).value != "Active":
                continue
            n_active += 1
            ids = E_RE.findall(str(ws.cell(r, 17).value))
            if not ids or any(i not in ev_ids for i in ids):
                missing.append((ws.cell(r, 1).value, ws.cell(r, 17).value))
    chk = base["Checks"]
    e8 = next(chk.cell(r, 4).value for r in range(1, chk.max_row + 1) if chk.cell(r, 1).value == "E8")
    res["3.3"] = {"active": n_active, "evidence_ids": len(ev_ids), "missing_or_unknown": missing, "E8": e8}

    # —— 3.4 Checks
    res["3.4"] = {k: base["Checks"][wb.defined_names[k].attr_text.split("!")[1].replace("$", "")].value
                  for k in ("GOV_Errors", "GOV_Warnings", "GOV_Info")}

    # —— 3.5 驗收測試
    s_, c_, k_ = SRC_PERTURB
    reach = graph.reach({(s_, c_)})
    outs = output_cells(graph)
    static = summarize_reach(graph, reach, {(s_, c_)}, outs)
    n_if2, d_if2 = _cmp_sheet(base, hw, "Interface")
    obs_names = collections.Counter()
    for n in if_names(graph):
        cells = {f"{c}" for _, c in graph.names[n]}
        obs_names[n] = sum(1 for co, *_ in d_if2 if co in cells)
    obs_names = {n: v for n, v in obs_names.items() if v}
    other, observed_not_static = {}, []
    for s in base.sheetnames:
        n, d = _cmp_sheet(base, hw, s)
        if d:
            other[s] = len(d)
        observed_not_static += [(s, co) for co, *_ in d if (s, co) not in reach and (s, co) != (s_, c_)]
    cols = sorted({re.match(r"[A-Z]+", co).group(0) for co, *_ in d_if2})
    res["3.5"] = {"static_if": static["if_names"], "static_cells": static["formula_cells"], "obs_if_cells": len(d_if2),
                  "obs_names": obs_names, "obs_cols": cols, "other": other, "observed_not_static": observed_not_static,
                  "obs_subset_static": not observed_not_static and set(obs_names) <= set(static["if_names"])}

    # —— 3.6 F14：E5＝2
    names_dev, worst = [], (0.0, None)
    for n in if_names(graph):
        for s2, co in graph.names[n]:
            x, y = base[s2][co].value, gw2[s2][co].value
            if is_num(x) and is_num(y):
                d = abs(x - y) / max(abs(x), abs(y), 1e-300)
                if d > worst[0]:
                    worst = (d, f"{n}（{s2}!{co}）")
            elif x != y:
                names_dev.append((n, co, x, y))
    res["3.6"] = {"max_rel": worst[0], "where": worst[1], "non_numeric_changes": names_dev}
    res["seconds"] = time.perf_counter() - t0
    write_gate1_md(res, out, xlsx, prev)
    return res


def write_gate1_md(res, out, xlsx, prev):
    L = [f"# Gate 1 驗收輸出（`tools/stage0_inventory.py gate1`）", "",
         f"- 現行：`{xlsx.name}`；前版：`{prev.name}`；執行 {res['seconds']:.0f} 秒（含 4 次 LibreOffice 重算）。",
         "- 本檔只列事實；通過與否的判定與意見見同輪 `*_v5.11_sync.md`。", ""]
    r = res["3.1"]
    L += ["## 3.1 v5.11 對 v5.10（LibreOffice 重算值，相對誤差 1e-9）", "",
          f"- Interface：{r['interface_cells']} 格，差異 {r['interface_diffs']} 格。", "",
          "| 模型頁 | 格數 | 差異格 |", "|---|---|---|"]
    L += [f"| {s} | {n} | {d} |" for s, (n, d, _) in r["model_pages"].items()]
    r = res["3.2"]
    L += ["", "## 3.2 原始數據寫死", "", f"- Gov_Map 類別分布：{r['gov_categories']}",
          f"- 類別＝原始數據：已連結 SRC {r['raw_linked']} 格；寫死 {len(r['raw_hardcoded'])} 格 {r['raw_hardcoded'][:10]}", "",
          f"### 公式內數值常數（G0-5；排除單位換算與位置索引；{len(r['constants'])} 格）", "",
          "| 工作表 | 格 | 常數 | 公式 |", "|---|---|---|---|"]
    L += [f"| {s} | {c} | {v} | `{f[:120].replace('|', '¦')}` |" for s, c, f, v in r["constants"]]
    L += ["", "### 被單位表排除的整數常數（2、3、4、8、10、12、24、60、100、168、365；依公式型態歸併，供人工複核是否為情境倍數）", "",
          "| 工作表 | 常數 | 格數 | 範例格 | 公式型態 |", "|---|---|---|---|---|"]
    L += [f"| {s} | {v} | {len(cs)} | {', '.join(dict.fromkeys(cs))[:40]} | `{f.replace('|', '¦')}` |" for (s, v, f), cs in sorted(scan_constants.excluded.items())]
    r = res["3.3"]
    L += ["", "## 3.3 Evidence", "", f"- Active SRC {r['active']} 筆；DB_Evidence 共 {r['evidence_ids']} 個 ID；缺或不存在 {len(r['missing_or_unknown'])}：{r['missing_or_unknown'][:10]}；Checks E8＝{r['E8']}"]
    r = res["3.4"]
    L += ["", "## 3.4 Checks（LibreOffice）", "", f"- {r}"]
    r = res["3.5"]
    L += ["", "## 3.5 SRC_HW_038 ×1.1", "",
          f"- 靜態依賴：可達公式格 {r['static_cells']}；IF_ 名稱 {len(r['static_if'])} 個。",
          f"- 觀測變動：Interface {r['obs_if_cells']} 格；IF_ 名稱 {len(r['obs_names'])} 個；欄 {r['obs_cols']}。",
          f"- 其他頁變動格數：{r['other']}",
          f"- 觀測 ⊆ 靜態：{r['obs_subset_static']}（觀測但不在靜態可達：{r['observed_not_static'][:10]}）", "",
          "觀測變動的 IF_ 名稱：" + "、".join(sorted(r["obs_names"])), "",
          "僅靜態依賴、數值未變的 IF_ 名稱：" + ("、".join(sorted(set(r["static_if"]) - set(r["obs_names"]))) or "—")]
    r = res["3.6"]
    L += ["", "## 3.6 Inputs!E5＝2", "", f"- 所有 IF_ 數值格最大相對差：{r['max_rel']:.3e}，位於 {r['where']}",
          f"- 非數值格的變動：{r['non_numeric_changes'][:10]}"]
    out.write_text("\n".join(L) + "\n", encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("inventory")
    a.add_argument("xlsx", type=Path)
    a.add_argument("--out", type=Path, required=True)
    a.add_argument("--workers", type=int, default=3)
    a.add_argument("--lo-check", type=int, default=10, help="LibreOffice 抽驗列數（0＝略過）")
    a.add_argument("--no-perturb", action="store_true")
    a.add_argument("--summary-json", type=Path)
    b = sub.add_parser("impact")
    b.add_argument("xlsx", type=Path)
    b.add_argument("cells", nargs="+")
    g = sub.add_parser("gate1")
    g.add_argument("xlsx", type=Path)
    g.add_argument("--prev", type=Path, required=True, help="前一版（model/archive/…）")
    g.add_argument("--out", type=Path, required=True)
    g.add_argument("--workdir", type=Path, default=Path(tempfile.gettempdir()) / "gate1")
    args = ap.parse_args()
    if args.cmd == "gate1":
        gate1(args.xlsx, args.prev, args.out, args.workdir)
        return
    if args.cmd == "impact":
        impact(args.xlsx, args.cells)
        return
    res = build_inventory(args.xlsx, args.workers, not args.no_perturb, args.lo_check)
    write_workbook(res, args.out, args.xlsx)
    if args.summary_json:
        c = counts(res)
        summ = {
            "sha_before": res["sha_before"], "sha_after": res["sha_after"], "seconds": res["seconds"],
            "perturb_seconds": res.get("perturb_seconds"), "restore_diff": res.get("restore_diff"), "restore_runs": res.get("restore_runs"), "restore_bad_runs": res.get("restore_bad_runs"),
            "violations": res.get("violations"), "lo_check": [{"inv": x["inv"], "worst": x["worst"]} for x in res.get("lo_check", [])],
            "class": dict(c["cls"]), "primary": dict(c["prim"]), "by_sheet": dict(c["by_sheet"]),
        }
        args.summary_json.write_text(json.dumps(summ, ensure_ascii=False, indent=1, default=str))
    print("完成：", args.out)


if __name__ == "__main__":
    main()
