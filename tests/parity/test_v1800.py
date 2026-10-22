"""v18.0 m 收入面化與 VR200 每 GW 營收 P × Q（工作單 docs/workorders/20261021_v18.0_m收入面化.md r0）專屬測試。讀 LibreOffice 重算後的值。

W6 驗收：
  1. W1、W2、W3-1 各名稱基準／低／高等於工作單事前計算（相對差 ≤ 1e-6；工作單只印 4–5 位小數，故以「印出位數的半個單位」為下限，
     壽命期年均三格另放寬到 2e-5——chat 端事前計算用的是四捨五入到 5 位的 L 值，見報告）；Python 以元件值獨立重算 1e-9。
  2. 恆等式：Σ（分項 m × 分項 token）÷（API token ＋ ChatGPT 付費 token）＝ L1_M（Cap_In 第 196、198、201、202 列；1e-9）。
  3. d 可調：情境 v1800_disc30（B4_Disc＝0.30）：L1_M ×（1 − B4_Disc）與 L1_RevGW_PQ_VR200 與基準相對差 ≤ 1e-9；L1_PxList_2026 不變。
  4. L1_TokLedger × L1_MFleet ＝ L1_M_2025。
另：逐格比較（對封存的 v17.2）只落在預期變動；既有具名範圍的值不變（版本字串除外）；標籤尾註、Decisions、DB_Evidence、README。
"""
import json
import re
from pathlib import Path

import openpyxl
import pytest
import yaml

PREWARM = {"test_v1800_new_values_match_precomputed": [{"SRC_MOD_106": 1024}]}
from openpyxl.utils import range_boundaries

from conftest import model_path

REPO = Path(__file__).resolve().parents[2]
V172 = REPO / "model" / "archive" / "20261020_Tokenomics_v17.2.xlsx"
V180 = REPO / "model" / "archive" / "20261021_Tokenomics_v18.0.xlsx"      # v19.0：現行版已是 v19.0；v17.2→v18.0 的靜態檢查改對封存的 v18.0 釘住（v18.0→v19.0 由 test_v1900 檢查）
SC = yaml.safe_load((Path(__file__).resolve().parent / "scenarios.yaml").read_text(encoding="utf-8"))["scenarios"]
DISC30 = next(s["inputs"] for s in SC if s["id"] == "v1800_disc30")
VERSION_NAMES = {"RNG_Version", "RNG_Current"}

KEYS = ["M_2025", "M_2026Q1", "M_2026Q2", "M_FY2026", "M", "M_API", "M_Sub", "M_Seat", "PxEff_2026", "PxList_2026", "TokGW_PaidCap_VR200", "TokGW_PaidPQ_VR200",
        "RevGW_PQ_VR200", "RevGW_PQ_VR200_FY", "RevGW_PQ_VR200_Life", "TokLedger", "TokLedger_2026Q2"]
NEW_NAMES = {f"L1_{k}{s}" for k in KEYS for s in ("", "_Lo", "_Hi")}
NR = "無區間"
NRL = "無區間（兩者區間來源不同，不組合）"

# 工作單事前計算：名稱 -> (基準, 低, 高)；低／高為文字者以該文字比對。M tok 類與 $B 類為工作單印出的位數。
PRE = {
    "L1_M_2025": (2.51355, 1.08510, 4.25970),
    "L1_M_2026Q1": (1.19234, NR, NR), "L1_M_2026Q2": (1.34320, NR, NR), "L1_M_FY2026": (1.43170, NR, NR), "L1_M": (1.34320, NR, NR),
    "L1_M_API": (0.34028, NR, NR), "L1_M_Sub": (5.39864, NR, NR), "L1_M_Seat": (4.55729, NR, NR),
    "L1_PxEff_2026": (1.96068, NR, NR), "L1_PxList_2026": (2.45084, NR, NR),
    "L1_TokGW_PaidCap_VR200": (253_258_912_515.9, NR, NR),
    "L1_TokGW_PaidPQ_VR200": (18_525_220_456.3, 13_555_039_358.2, 29_369_251_942.9),
    "L1_RevGW_PQ_VR200": (48.7877, 35.6983, 77.3463),
    "L1_RevGW_PQ_VR200_FY": (52.0021, 38.0503, 82.4423),
    "L1_RevGW_PQ_VR200_Life": (34.0951, 15.7101, 84.6284),
    "L1_TokLedger": (5.46420, NRL, NRL), "L1_TokLedger_2026Q2": (6.25811, NRL, NRL),
}
LIFE = {"L1_RevGW_PQ_VR200_Life"}


def _flat(v):
    if isinstance(v, (list, tuple)):
        return [x for r in v for x in _flat(r)]
    return [v]


def _decimals(x):
    s = repr(float(x))
    return len(s.split(".")[1].rstrip("0")) if "." in s else 0


def _tol(want, name):
    """相對容差：1e-6，不小於「印出最後一位的半個單位」；壽命期年均 2e-5。"""
    half = 0.5 * 10 ** (-_decimals(want)) / abs(want)
    return max(1e-6, half) if name not in LIFE else 2e-5


def _close(got, want, tol):
    return isinstance(got, (int, float)) and abs(got - want) <= tol * abs(want)


def test_v1800_new_values_match_precomputed(lo):
    v = lo.view({"SRC_MOD_106": 1024})      # v19.1：Arch!C17 改值（4224 → 2176）後 Luna 相關值變動；事前值在 C17 回 4224 時仍成立（SRC_MOD_106＝1024）
    for n, trio in PRE.items():
        for suf, want in zip(("", "_Lo", "_Hi"), trio):
            got = v.get_name(n + suf)
            if isinstance(want, str):
                assert got == want, (n + suf, got, want)
            else:
                assert _close(got, want, _tol(want, n)), (n + suf, got, want)


def test_v1800_independent_python_recomputation(lo):
    """以元件值（IF_TokGW、B4_*、AL_*、Cap_In 第 187／196／198／205 列、L1_LTotalDisc）獨立重算 Q、P × Q 與 m（1e-9）。"""
    v = lo.view({})
    g = lambda n: _flat(v.get_name(n))
    mix = g("B4_MixPaid")
    cap = 1 / (mix[0] / g("IF_TokGW_Luna")[10] + mix[1] / g("IF_TokGW_Sol")[10] + mix[2] / g("IF_TokGW_Astra")[10])
    serve, free, prod = g("B4_Serve")[0], g("B4_Free")[0], g("L1_FleetProdRatio")[0]
    px = v.get("Cap_In", "E205")
    for suf, util in (("", "AL_UtilActual"), ("_Lo", "AL_UtilActual_Lo"), ("_Hi", "AL_UtilActual_Hi")):
        q = cap * serve * (1 - free) * g(util)[0] * prod
        assert abs(v.get_name("L1_TokGW_PaidPQ_VR200" + suf) / q - 1) < 1e-9, suf
        assert abs(v.get_name("L1_RevGW_PQ_VR200" + suf) / (q * px * v.get("Cap_In", "E209") / 1e9) - 1) < 1e-9, suf
        assert abs(v.get_name("L1_RevGW_PQ_VR200_FY" + suf) / (q * px * v.get("Cap_In", "F209") / 1e9) - 1) < 1e-9, suf
    assert abs(v.get_name("L1_TokGW_PaidCap_VR200") / cap - 1) < 1e-9
    life = {"": "L1_LTotalDisc", "_Lo": "L1_LTotalDisc_Lo", "_Hi": "L1_LTotalDisc_Hi"}
    for suf, ln in life.items():
        assert abs(v.get_name("L1_RevGW_PQ_VR200_Life" + suf) / (v.get_name("L1_RevGW_PQ_VR200" + suf) * v.get_name(ln)) - 1) < 1e-9, suf
    # m：營收 ÷（需求面付費 token × 當年有效單價）
    for col, n in (("D", "L1_M_2026Q1"), ("E", "L1_M_2026Q2"), ("F", "L1_M_FY2026")):
        m = v.get("Cap_In", f"{col}187") * 1e9 / ((v.get("Cap_In", f"{col}196") + v.get("Cap_In", f"{col}198")) * v.get("Cap_In", f"{col}205"))
        assert abs(v.get_name(n) / m - 1) < 1e-9, n
    assert v.get_name("L1_M") == v.get_name("L1_M_2026Q2")                      # J1：主口徑＝2026 Q2
    assert v.get_name("L1_M_2025") == v.get_name("L1_MImplied") and v.get_name("L1_M_2025_Lo") == v.get_name("L1_MImplied_Lo") and v.get_name("L1_M_2025_Hi") == v.get_name("L1_MImplied_Hi")
    assert abs(v.get_name("L1_PxList_2026") * (1 - g("B4_Disc")[0]) / v.get_name("L1_PxEff_2026") - 1) < 1e-12


def test_v1800_identity_sum_of_route_m(lo):
    """恆等式：(m_API × API token ＋ m_消費訂閱 × 消費訂閱 token ＋ m_商用席次 × 商用席次 token) ÷ (API token ＋ ChatGPT 付費 token) ＝ L1_M（2026 Q2 欄）。"""
    v = lo.view({})
    api, paid, sub, seat = (v.get("Cap_In", f"E{r}") for r in (196, 198, 201, 202))
    lhs = (v.get_name("L1_M_API") * api + v.get_name("L1_M_Sub") * sub + v.get_name("L1_M_Seat") * seat) / (api + paid)
    assert abs(lhs / v.get_name("L1_M") - 1) <= 1e-9
    assert abs((sub + seat) / paid - 1) <= 1e-12            # 消費訂閱＋商用席次 token ＝ ChatGPT 付費 token


def test_v1800_disc_adjustable(lo):
    """d 可調（情境 v1800_disc30）：實收單價與 d 無關——L1_M ×（1 − B4_Disc）與 L1_RevGW_PQ_VR200 不變；表列牌價 L1_PxList_2026 不變；有效單價與 m 反向同比例。"""
    b, s = lo.view({}), lo.view(DISC30)
    db, ds = b.get_name("B4_Disc"), s.get_name("B4_Disc")
    assert (db, ds) == (0.2, 0.3)
    rel = lambda n: s.get_name(n) / b.get_name(n) - 1
    assert abs(s.get_name("L1_M") * (1 - ds) / (b.get_name("L1_M") * (1 - db)) - 1) <= 1e-9
    assert abs(rel("L1_RevGW_PQ_VR200")) <= 1e-9
    assert abs(rel("L1_PxList_2026")) <= 1e-9
    assert abs(s.get_name("L1_PxEff_2026") / b.get_name("L1_PxEff_2026") - (1 - ds) / (1 - db)) <= 1e-9
    assert abs(rel("L1_M") - ((1 - db) / (1 - ds) - 1)) <= 1e-9
    for n in ("L1_RevGW_PQ_VR200_FY", "L1_RevGW_PQ_VR200_Life", "L1_TokGW_PaidPQ_VR200"):
        assert abs(rel(n)) <= 1e-9, n
    assert s.get_name("GOV_Errors") == 0


def test_v1800_token_ledger(lo):
    v = lo.view({})
    assert abs(v.get_name("L1_TokLedger") * v.get_name("L1_MFleet") / v.get_name("L1_M_2025") - 1) <= 1e-12
    assert abs(v.get_name("L1_TokLedger_2026Q2") - v.get("Cap_In", "E209") / v.get("Cap_In", "E223")) <= 1e-12
    # 同價格基準（工作單 W3-1）：2025 兩邊單價 CAP_Price25Mix＝AL_PriceFleet25；2026 兩邊 Cap_In 第 205 列＝B4_PriceFleet 第 3 欄
    assert abs(_flat(v.get_name("CAP_Price25Mix"))[0] / _flat(v.get_name("AL_PriceFleet25"))[0] - 1) <= 1e-12
    assert abs(v.get("Cap_In", "E205") / _flat(v.get_name("B4_PriceFleet"))[2] - 1) <= 1e-12
    # 診斷關係：L1_MFleet（2025）＝ L1_M_2025 ÷ 帳比
    assert abs(v.get_name("L1_MFleet") - v.get_name("L1_M_2025") / v.get_name("L1_TokLedger")) <= 1e-12


# ---------------------------------------------------------------- static (formula mode / archive comparison)
@pytest.fixture(scope="module")
def wbf():
    return openpyxl.load_workbook(V180)


def _row(ws, rid, col=1):
    rows = [r for r in range(5, ws.max_row + 1) if ws.cell(r, col).value == rid]
    assert len(rows) == 1, (rid, rows)
    return rows[0]


def test_v1800_l1_rows_formulas_and_tails(wbf):
    l1 = wbf["L1"]
    first = _row(l1, "L1_M_2025")
    assert first == 117 and [l1.cell(first + i, 1).value for i in range(len(KEYS))] == [f"L1_{k}" for k in KEYS]
    for k in KEYS:
        r = _row(l1, f"L1_{k}")
        for c in "DEF":
            assert isinstance(l1[f"{c}{r}"].value, str) and l1[f"{c}{r}"].value.startswith("="), (k, c)      # Checks E11：D／E／F 皆為公式
        assert l1[f"O{r}"].value is None and l1[f"T{r}"].value in ("估計", "殘差", "未解釋")
    r = _row(l1, "L1_M")
    assert l1[f"D{r}"].value == "=L1_M_2026Q2"
    for k, c in (("M_2026Q1", "D"), ("M_2026Q2", "E"), ("M_FY2026", "F")):
        assert l1[f"D{_row(l1, 'L1_' + k)}"].value == f"=Cap_In!${c}$209"
    for k, rr in (("M_API", 206), ("M_Sub", 207), ("M_Seat", 208)):
        assert l1[f"D{_row(l1, 'L1_' + k)}"].value == f"=Cap_In!$E${rr}"
    assert l1[f"D{_row(l1, 'L1_PxEff_2026')}"].value == "=Cap_In!$E$205" and l1[f"D{_row(l1, 'L1_PxList_2026')}"].value == "=L1_PxEff_2026/(1-B4_Disc)"
    assert l1[f"D{_row(l1, 'L1_TokLedger')}"].value == "=IF(AND(ISNUMBER(L1_M_2025),ISNUMBER(L1_MFleet)),L1_M_2025/L1_MFleet,AL_UnavailTxt)" and l1[f"D{_row(l1, 'L1_TokLedger_2026Q2')}"].value == "=IF(AND(ISNUMBER(Cap_In!$E$209),ISNUMBER(Cap_In!$E$223)),Cap_In!$E$209/Cap_In!$E$223,AL_UnavailTxt)"
    assert l1[f"E{_row(l1, 'L1_TokGW_PaidPQ_VR200')}"].value.count("AL_UtilActual_Lo") == 1 and l1[f"F{_row(l1, 'L1_TokGW_PaidPQ_VR200')}"].value.count("AL_UtilActual_Hi") == 1
    life = _row(l1, "L1_RevGW_PQ_VR200_Life")
    assert l1[f"E{life}"].value == "=L1_RevGW_PQ_VR200_Lo*L1_LTotalDisc_Lo" and l1[f"F{life}"].value == "=L1_RevGW_PQ_VR200_Hi*L1_LTotalDisc_Hi"
    # I 欄必寫三點（m 列）與四點（P × Q 列）
    for k in ("M_2026Q1", "M_2026Q2", "M_FY2026", "M", "M_API", "M_Sub", "M_Seat"):
        t = l1[f"I{_row(l1, 'L1_' + k)}"].value
        assert "SRC_DEM_011" in t and "15 B／分鐘" in t and "未找到" in t and "m_t ÷ m_2026" in t and "區間待補" in t, k
    for k in ("TokGW_PaidPQ_VR200", "RevGW_PQ_VR200", "RevGW_PQ_VR200_FY", "RevGW_PQ_VR200_Life"):
        t = l1[f"I{_row(l1, 'L1_' + k)}"].value
        assert "token 帳未閉合" in t and "6.26" in t and "回本比較不在本版" in t and "65.96" in t and "m 不變" in t, k
    # W3 標籤尾註
    for name in ("L1_MFleet", "L1_MFleet_Px2026"):
        assert "v18.0 起：token 帳診斷（＝m ×（真實 token ÷ 模型 token）），不是 m；m 見 L1_M（Andy 2026-10-21 定義）" in l1[f"B{_row(l1, name)}"].value
    assert l1[f"B{_row(l1, 'L1_MImplied')}"].value.endswith("｜v18.0 起即收入面 m 的 2025 值（L1_M_2025）")
    for name in ("L1_MAPI", "L1_MSub", "L1_MSeat"):
        assert l1[f"B{_row(l1, name)}"].value.endswith("｜v18.0 起即收入面分項 m 的 2025 值")
    ci = wbf["Cap_In"]
    for a in ("A119", "A223"):
        assert "v18.0 起：token 帳診斷" in ci[a].value
    assert ci["A209"].value.endswith("（v18.0：本列即收入面 m；L1_M_* 讀本列）")
    assert ci["H196"].value.endswith("；v18.0：Q2 當期值已搜尋，未找到（DB_Evidence 新列 E399）")


def test_v1800_evidence_decisions_readme(wbf):
    ev = wbf["DB_Evidence"]
    r = _row(ev, "E399")
    assert ev.cell(r, 3).value == "OpenAI API 每分鐘 token 2026 Q2／Q3 當期值"
    assert ev.cell(r, 9).value == "已搜尋，未找到" and "2026-10-21 17:20" in ev.cell(r, 4).value and "Accelerating the next phase of AI" in ev.cell(r, 4).value
    assert ev.cell(r, 10).value == "v18.0" and ev.cell(r, 12).value == "已處理"
    d = wbf["Decisions"]
    ids = [f"SL52-0{i}" for i in range(1, 7)]
    for i in ids:
        rr = _row(d, i)
        assert d.cell(rr, 2).value == "v18.0" and d.cell(rr, 7).value == "v18.0 已建" and d.cell(rr, 10).value == "否" and "17:08" in d.cell(rr, 6).value, i
        assert i == "SL52-01" or "可撤回" in d.cell(rr, 4).value, i
    txt = " ".join(str(d.cell(_row(d, i), 3).value) + str(d.cell(_row(d, i), 4).value) for i in ids)
    for k in ("J1", "J2", "J3", "J4", "支出口徑", "E399", "L1_TokLedger"):
        assert k in txt, k
    assert d.max_row < 1000      # Gov_Map AA 的 COUNTIF 範圍上限
    rd = wbf["README"]
    assert str(rd["B5"].value).startswith("20261021_Tokenomics_v18.0（主版：") and "v17.2（修補版：" in str(rd["B5"].value) and str(rd["A1"].value).startswith("Tokenomics v18.0")
    rows = [x for x in range(4, rd.max_row + 1) if rd.cell(x, 1).value == "v18.0 m 收入面化"]
    assert len(rows) == 1 and "L1_M" in rd.cell(rows[0], 2).value and "L1_RevGW_PQ_VR200" in rd.cell(rows[0], 2).value


def _diffs(data_only):
    a, b = openpyxl.load_workbook(V172, data_only=data_only), openpyxl.load_workbook(V180, data_only=data_only)
    assert a.sheetnames == b.sheetnames
    out = []
    for ws in a:
        wn = b[ws.title]
        for r in range(1, max(ws.max_row, wn.max_row) + 1):
            for c in range(1, max(ws.max_column, wn.max_column) + 1):
                x, y = ws.cell(r, c).value, wn.cell(r, c).value
                if x != y and not (isinstance(x, (int, float)) and isinstance(y, (int, float)) and abs(x - y) <= 1e-12 * max(1, abs(x))):
                    out.append((ws.title, ws.cell(r, c).coordinate, x, y))
    return out, a, b


def _allowed(sh, coord):
    col, row = re.match(r"([A-Z]+)(\d+)", coord).groups(); row = int(row)
    if sh == "L1":
        return row >= 117 or (col == "B" and row in (76, 78, 80, 81, 82, 102))
    if sh == "Cap_In":
        return coord in ("A119", "A209", "A223", "H196")
    if sh == "README":
        return coord in ("A1", "B5", "A65", "B65")
    if sh == "Alloc":
        return coord in ("C177", "C178")                 # RNG_Version／RNG_Current（版本字串）
    if sh == "Load_Bearing":
        return col == "J"                                 # 受影響名稱清單（自動追蹤）：只多了新 L1 名稱
    return sh in ("Decisions", "DB_Evidence")


def test_v1800_diffs_only_expected_vs_v172():
    for mode in (False, True):
        d, a, b = _diffs(mode)
        bad = [t for t in d if not _allowed(t[0], t[1])]
        assert bad == [], bad[:10]
        for sh in ("Decisions", "DB_Evidence"):
            rows = {re.sub(r"[A-Z]+", "", t[1]) for t in d if t[0] == sh}
            assert all(t[2] is None for t in d if t[0] == sh) and len(rows) == (6 if sh == "Decisions" else 1)      # 只新增列
        lb = [t for t in d if t[0] == "Load_Bearing"]
        for t in lb:
            assert isinstance(t[2], str) and t[3].startswith(t[2]) and set(re.findall(r"L1_\w+", t[3][len(t[2]):])) <= NEW_NAMES, t      # 只在尾端新增新 L1 名稱
        assert [t for t in d if t[0] == "L1" and int(re.sub(r"[A-Z]+", "", t[1])) < 117 and t[2] is not None and not str(t[3]).startswith(str(t[2]))] == []      # 既有列的舊文字完整保留
    names_a, names_b = set(a.defined_names), set(b.defined_names)
    assert names_a <= names_b and names_b - names_a == NEW_NAMES
    assert not [k for k in names_a if a.defined_names[k].attr_text != b.defined_names[k].attr_text]
    # 數值：既有數字格只允許 L1 新列（≥ 117）
    dv, _, _ = _diffs(True)
    num = [t for t in dv if isinstance(t[2], (int, float)) or isinstance(t[3], (int, float))]
    assert all(t[0] == "L1" and int(re.sub(r"[A-Z]+", "", t[1])) >= 117 for t in num), [t for t in num if not (t[0] == "L1")][:5]


def test_v1800_existing_names_unchanged():
    """與封存的 v17.2 逐名稱比較（對封存的 v18.0 重算值）：既有名稱的值全部不變，唯一例外是版本字串（RNG_Version／RNG_Current）。"""
    pw, pf = openpyxl.load_workbook(V172, data_only=True), openpyxl.load_workbook(V172)
    cw, cf = openpyxl.load_workbook(V180, data_only=True), openpyxl.load_workbook(V180)

    def val(wbk, n):
        sh, rg = wbk.defined_names[n].attr_text.rsplit("!", 1)
        sh = sh.strip("'")
        c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
        v = [wbk[sh].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        return v
    assert set(cf.defined_names) - set(pf.defined_names) == NEW_NAMES and not (set(pf.defined_names) - set(cf.defined_names))
    changed = set()
    for n, d in pw.defined_names.items():
        if "!" not in d.attr_text or "," in d.attr_text:
            continue
        want, got = val(pw, n), val(cw, n)
        same = len(want) == len(got) and all((a == b) or (isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) <= 1e-9 * max(abs(a), abs(b), 1e-300) + 1e-12)
                                             or (a in (None, "") and b in (None, "")) for a, b in zip(want, got))
        if not same:
            changed.add(n)
    assert changed == VERSION_NAMES, sorted(changed ^ VERSION_NAMES)
    assert tuple(val(cw, n)[0] for n in ("GOV_Errors", "GOV_Warnings", "GOV_Info")) == (0, 135, 182)      # v17.2 同；v18.0 不增減任何治理計數
    assert val(cw, "RNG_StaleFlag")[0] == "一致" and val(cw, "RNG_OutCount")[0] == 0 and val(cw, "CHK_L1Order")[0] == 1
