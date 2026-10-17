"""產生（或核對）批次口徑倍數區間端點檔 builder/v1300_ranges.json（4-2b-1 K1／K2）。

用法：
    python3 tools/gen_ranges.py <xlsx> [--version <檔名主幹>] [--out builder/v1300_ranges.json]   寫出端點檔
    python3 tools/gen_ranges.py <xlsx> --check                                                      重新推導並與現有端點檔逐值比對（容差 1e-9）

做法（與 scenarios.yaml 的 batch_gb200_lo、batch_range_lo 相同的兩個端點情境）：
  高端情境：CAL_BatchEtaD[2]（GB200）＝1.59（v15.3 起；v13.0–v15.2 為 1.83）；低端情境：GB200＝3.04 且 Hopper（CAL_BatchEtaD[1]）＝3.74。
  凡契約前綴（IF_／IFW_／IFC_／L1_／CHK_／CAP_）的具名範圍，只要在任一端點情境下與基準不同，就納入；
  RNG_<名稱>_Lo／_Hi 的每一格＝兩個端點情境該格值的較小／較大者（不是情境的「低端／高端」：有些名稱隨倍數增加而上升）。
  另外重算兩個交叉角落（GB200 1.59 × Hopper 3.74、GB200 3.04 × Hopper 1）與 Hopper 單獨 3.74，確認它們都落在端點之內；
  v15.0 T3-2：另以 VR200 倍數的兩個端點情境（CAL_BatchEtaD[4]＝0.70、1.0；基準 0.75）重算；凡在上述兩個批次端點情境不動、但隨 VR200 倍數變動的契約名稱，
  也建 RNG_<名稱>_Lo／_Hi（端點＝vr200_lo／vr200_hi 兩情境每格較小／較大者，JSON 中標 "source": "vr200"）。既有 59 個名稱的端點與其 VR200 欄（點值）不變。
區間端點是情境重算結果，只能以數值寫入 Excel；每次改版（或改動任何上游輸入）後必須重跑本工具，並在同一個 PR 更新 builder/v1300_ranges.json。"""
import argparse
import json
import re
import shutil
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl
from openpyxl.utils import range_boundaries

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "tests" / "parity"))
from parity_lib import lo_recalc, parse_ref, set_inputs  # noqa: E402
import deadcells_lib  # noqa: E402

PREFIX = ("IF_", "IFW_", "IFC_", "L1_", "CHK_", "CAP_")
END_HI = {"CAL_BatchEtaD[2]": 1.59}                          # 研發 GW 區間的高端（IF_AllocRDGW 最高）；v15.3 C2：1.83 → 1.59（Calib C254 在 D67＝1.15 後的 V20 反推點；Gov_Map GM685 K）
END_LO = {"CAL_BatchEtaD[2]": 3.04, "CAL_BatchEtaD[1]": 3.74}      # 低端
CHECK = {
    "corner_gb200_lo_hopper_hi": {"CAL_BatchEtaD[2]": 1.59, "CAL_BatchEtaD[1]": 3.74},
    "corner_gb200_hi_hopper_1": {"CAL_BatchEtaD[2]": 3.04},
    "hopper_hi_only": {"CAL_BatchEtaD[1]": 3.74},
}
VR_LO = {"CAL_BatchEtaD[4]": 0.70}      # v15.0 T3-2：VR200 欄批次倍數區間下端（Gov_Map GM595；MLPerf v6.1 offline Nebius 0.70）
VR_HI = {"CAL_BatchEtaD[4]": 1.0}       # 上端（Andy 2026-10-07 假設，Assumed）
VR = {"vr200_lo": VR_LO, "vr200_hi": VR_HI}
OUT = REPO / "builder" / "v1300_ranges.json"


def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(a - b) <= 1e-12 or abs(a - b) <= 1e-9 * max(abs(a), abs(b))
    return a == b


def run_all(src, scen):
    wbn = openpyxl.load_workbook(src)
    names = {k: v.attr_text for k, v in wbn.defined_names.items()}
    readme = wbn["README"]["B5"].value

    def one(item):
        k, inp = item
        tmp = Path(tempfile.mkdtemp(prefix="rng_"))
        try:
            s = tmp / src.name
            set_inputs(src, s, inp)
            o = lo_recalc(s, tmp / "lo")
            wb = openpyxl.load_workbook(o, data_only=True)
            res = {}
            for n, t in names.items():
                try:
                    sh, ref = parse_ref(t)
                except Exception:
                    continue
                c1, r1, c2, r2 = range_boundaries(ref)
                res[n] = [wb[sh].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
            wb.close()
            return k, res
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    with ThreadPoolExecutor(2) as ex:
        res = dict(ex.map(one, scen.items()))
    return names, res, readme


def graph_upstream(src):
    """依賴圖：具名範圍的上游（含間接）包含 CAL_BatchEtaD 者。"""
    cells, names, maxrc, _ = deadcells_lib.load_formulas(src)
    edges = deadcells_lib.build_graph(cells, names, maxrc)
    import collections
    rev = collections.defaultdict(set)
    for k, vs in edges.items():
        for v in vs:
            rev[v].add(k)
    seen, stack = set(), [("NAME", "CAL_BatchEtaD")]
    while stack:
        n = stack.pop()
        if n in seen:
            continue
        seen.add(n)
        stack.extend(rev.get(n, ()))
    aff = {n for n in names if ("NAME", n) in seen or any(t in seen for t in edges.get(("NAME", n), ()))}
    aff.discard("CAL_BatchEtaD")
    return aff


def derive(src, version=None):
    scen = {"base": {}, "hi": END_HI, "lo": END_LO, **CHECK, **VR}
    names, res, readme = run_all(src, scen)
    if version is None:
        version = str(readme).split("（", 1)[0]
    base = res["base"]
    contract = [n for n in names if n.startswith(PREFIX) and n in base]
    moved = [n for n in contract if any(not same(a, b) for sc in ("hi", "lo") for a, b in zip(base[n], res[sc][n]))]
    out_names = {}
    problems = []
    for n in sorted(moved):
        lo_v, hi_v = [], []
        for i, b in enumerate(base[n]):
            a, c = res["hi"][n][i], res["lo"][n][i]
            if not all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in (a, c, b)):
                if not (same(a, b) and same(c, b)):
                    problems.append(f"{n}[{i}] 非數值但有變動")
                lo_v.append(b); hi_v.append(b); continue
            lo_i, hi_i = min(a, c), max(a, c)
            lo_v.append(lo_i); hi_v.append(hi_i)
            tol = 1e-9 * max(abs(lo_i), abs(hi_i), 1e-300)
            if b < lo_i - tol or b > hi_i + tol:
                problems.append(f"{n}[{i}] 基準 {b} 在端點 [{lo_i},{hi_i}] 之外")
            for k in CHECK:
                v = res[k][n][i]
                if isinstance(v, (int, float)) and (v < lo_i - tol or v > hi_i + tol):
                    problems.append(f"{n}[{i}] 角落 {k}={v} 在端點 [{lo_i},{hi_i}] 之外")
        out_names[n] = {"ref": names[n], "n": len(base[n]), "lo": lo_v, "hi": hi_v,
                        "inverted_cells": [i for i in range(len(base[n])) if isinstance(res["lo"][n][i], (int, float)) and isinstance(res["hi"][n][i], (int, float))
                                           and res["lo"][n][i] > res["hi"][n][i] + 1e-12 * abs(res["hi"][n][i])]}
    # v15.0 T3-2：只經 VR200 欄的名稱（兩個批次端點情境下不動、VR200 倍數改變時才動）也建區間；端點＝vr200_lo／vr200_hi 兩個情境每格的較小／較大者
    vr_moved = sorted(n for n in contract if n not in moved and any(not same(a, b) for sc in ("vr200_lo", "vr200_hi") for a, b in zip(base[n], res[sc][n])))
    for n in vr_moved:
        lo_v, hi_v = [], []
        for i, b in enumerate(base[n]):
            a, c = res["vr200_hi"][n][i], res["vr200_lo"][n][i]
            if not all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in (a, c, b)):
                if not (same(a, b) and same(c, b)):
                    problems.append(f"{n}[{i}] 非數值但有變動（VR200）")
                lo_v.append(b); hi_v.append(b); continue
            lo_i, hi_i = min(a, c), max(a, c)
            lo_v.append(lo_i); hi_v.append(hi_i)
            tol = 1e-9 * max(abs(lo_i), abs(hi_i), 1e-300)
            if b < lo_i - tol or b > hi_i + tol:
                problems.append(f"{n}[{i}] 基準 {b} 在 VR200 端點 [{lo_i},{hi_i}] 之外")
        out_names[n] = {"ref": names[n], "n": len(base[n]), "lo": lo_v, "hi": hi_v, "source": "vr200",
                        "inverted_cells": [i for i in range(len(base[n])) if isinstance(res["vr200_lo"][n][i], (int, float)) and isinstance(res["vr200_hi"][n][i], (int, float))
                                           and res["vr200_lo"][n][i] > res["vr200_hi"][n][i] + 1e-12 * abs(res["vr200_hi"][n][i])]}
    up = sorted(n for n in graph_upstream(src) if n in contract and n not in moved and n not in vr_moved)
    vr_only = []      # v15.0：不再有『只經 VR200 欄而無區間』的名稱（已全部納入 vr200 區間）
    still = [n for n in up if n not in vr_only]
    non_contract_moved = sorted(n for n in names if n in base and not n.startswith(PREFIX) and n != "CAL_BatchEtaD" and not n.startswith("RNG_")
                                and any(not same(a, b) for sc in ("hi", "lo") for a, b in zip(base[n], res[sc][n])))
    return {
        "version": version,
        "generated_from": src.name,
        "scenarios": {"hi": END_HI, "lo": END_LO, "vr200_lo": VR_LO, "vr200_hi": VR_HI},
        "corner_checks": CHECK,
        "names": out_names,
        "graph_only_vr200": vr_only,
        "graph_only_static": still,
        "display_names_moved": non_contract_moved,
        "problems": problems,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--version")
    ap.add_argument("--out", default=str(OUT))
    ap.add_argument("--check", action="store_true")
    a = ap.parse_args()
    src = Path(a.xlsx).resolve()
    d = derive(src, a.version)
    print("納入名稱", len(d["names"]), "格", sum(v["n"] for v in d["names"].values()), "| 問題", len(d["problems"]),
          "| 其中 VR200 區間", sum(1 for v in d["names"].values() if v.get("source") == "vr200"), "| 圖上游含且各情境不動", len(d["graph_only_static"]))
    for p in d["problems"]:
        print("  問題：", p)
    if a.check:
        cur = json.loads(Path(a.out).read_text(encoding="utf-8"))
        bad = []
        if set(cur["names"]) != set(d["names"]):
            bad.append(("名稱集合", sorted(set(cur["names"]) ^ set(d["names"]))))
        for n in set(cur["names"]) & set(d["names"]):
            for k in ("lo", "hi"):
                if not all(same(x, y) for x, y in zip(cur["names"][n][k], d["names"][n][k])) or len(cur["names"][n][k]) != len(d["names"][n][k]):
                    bad.append((n, k))
        for k in ("graph_only_vr200", "graph_only_static"):
            if cur[k] != d[k]:
                bad.append((k, sorted(set(cur[k]) ^ set(d[k]))))
        print("核對：", "全部一致" if not bad else f"不一致 {len(bad)}：{bad[:10]}")
        sys.exit(1 if bad or d["problems"] else 0)
    Path(a.out).write_text(json.dumps(d, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print("已寫入", a.out)
    sys.exit(1 if d["problems"] else 0)


if __name__ == "__main__":
    main()
