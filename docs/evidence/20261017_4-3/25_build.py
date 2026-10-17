"""4-3 A 段：合併 rows.pkl（v13.1 B 級死格逐列）＋24_decisions.DEC＋文件引用，做完整性與依賴一致性檢查，輸出統計 pickle 與明細 xlsx。
用法：python3 25_build.py <dead.pkl> <rows.pkl> <mentions.pkl> <out.xlsx> <stats.pkl>   （工作目錄＝repo 根）"""
import sys, re, pickle, collections, importlib.util
from pathlib import Path
REPO = Path.cwd()
sys.path.insert(0, str(REPO / "tests" / "parity"))
import deadcells_lib as D
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

spec = importlib.util.spec_from_file_location("dec", str(Path(__file__).with_name("24_decisions.py")))
dec = importlib.util.module_from_spec(spec); spec.loader.exec_module(dec)
DEC = dec.DEC
P = pickle.load(open(sys.argv[1], "rb"))
rows = pickle.load(open(sys.argv[2], "rb"))
MEN = pickle.load(open(sys.argv[3], "rb"))
model = REPO / "model" / P["model"]
CONTRACT = REPO / "docs" / "plan" / "Tokenomics_downstream_contract.md"

# ---------- 完整性 ----------
keys = [(x["sheet"], x["row"]) for x in rows]
missing = [k for k in keys if k not in DEC]
extra = [k for k in DEC if k not in set(keys)]
print("rows", len(rows), "missing", missing, "extra", extra)
assert not missing and not extra, (missing, extra)

# ---------- 依賴一致性 ----------
cells, names, maxrc, sheets = D.load_formulas(model)
edges = D.build_graph(cells, names, maxrc)
deadB = set(P["unmarked"]); deadM = set(P["marked"]); deadall = deadB | deadM
row_of = lambda a: int(re.search(r"\d+", a).group())
act_of = lambda k: DEC[(k[0], row_of(k[1]))]["act"] if k in deadB else "M"
conflicts = []
keepers = [k for k in deadall if act_of(k) != "刪除"]
for kc in keepers:
    seen = set(); stack = [kc]
    while stack:
        n = stack.pop()
        for p in edges.get(n, ()):
            if p in seen: continue
            seen.add(p)
            if p[0] in ("NAME", "RANGE"):
                stack.append(p)
            elif p in deadall:
                if p in deadB and act_of(p) == "刪除":
                    conflicts.append((kc, p))
                stack.append(p)
# 去重到列級
cf_rows = collections.Counter(((a[0], row_of(a[1])), (b[0], row_of(b[1]))) for a, b in conflicts)
print("依賴衝突（保留者上游含刪除列）", len(cf_rows))
for k, v in list(cf_rows.items())[:30]: print("  ", k, v)

# 隨附上游檢查：升級／接進 Checks 的格，其上游死格列是否都已被處置為保留／升級／隨附／Checks（不是刪除）→ 已含於上面；另列出 UPA 列是否真的被某個 UP／CHK／KEEP 讀取
need = collections.defaultdict(set)
for kc in keepers:
    seen = set(); stack = [kc]
    while stack:
        n = stack.pop()
        for p in edges.get(n, ()):
            if p in seen: continue
            seen.add(p)
            if p[0] in ("NAME", "RANGE"): stack.append(p)
            elif p in deadall:
                need[(p[0], row_of(p[1]))].add((kc[0], row_of(kc[1])))
                stack.append(p)
unused_upa = [k for k in keys if DEC[k]["act"] == "升級（隨附上游）" and k not in need]
print("UPA 但沒有被任何保留／升級列讀取", unused_upa)

# ---------- 4-0 規則編號（依 v13.1 重判；決策樹同 06_recommend_dead.py）----------
VERIFY = re.compile(r"檢查|核對|判讀|÷ MLPerf|模型 ÷ 實測|上限檢查|一致性")
CMP = re.compile(r"對照|外部|揭露|參照|Requesty|Plus|Pro 方案|方案上限")
HEADER = {"世代", "層級", "成本情境", "項目", "任務", "參數", "實驗室"}
def rule_of(x):
    s, lab, cols = x["sheet"], x["label"], x["cols"]
    whole = not x["row_has_core"]
    if s == "Perf_Batch" and set(cols) <= {"R", "S"}: return "R04"
    if s == "L1": return "R07"
    if s == "Cap_In" and x["row"] in (23, 24): return "R09b"
    if whole and lab in HEADER: return "R12"
    if whole and VERIFY.search(lab): return "R14"
    if whole and CMP.search(lab): return "R15"
    if s == "Cap_In" and whole: return "R16"
    if s == "Calib": return "R20"
    if not whole: return "R21"
    if s in ("Perf", "Perf_Batch") and "速覽" in lab: return "R22c"
    if "無任何格引用" in x["users"]: return "R22a"
    return "R22b"

# ---------- 鏡像頁同列 ----------
MIRROR = {"Perf": "Perf_Prod", "Perf_Batch": "Perf_Batch_Prod", "Training": "Training_Prod", "Unit_Cost": "Unit_Cost_Prod", "Fleet_1GW": "Fleet_1GW_Prod",
          "Amortize": "Amortize_Prod", "Theory_Rev": "Theory_Rev_Prod", "Interface": "Interface_Prod", "Perf_Sens": "Sens_Perf"}
mir_cnt = {}
for x in rows:
    m = MIRROR.get(x["sheet"])
    if not m: continue
    n = sum(1 for (s, a), f in cells.items() if s == m and row_of(a) == x["row"] and isinstance(f, str) and f.startswith("="))
    if n: mir_cnt[(x["sheet"], x["row"])] = (m, n)

# ---------- 統計 ----------
st = dict(model=P["model"], nform=P["nform"], n_dead=len(deadall), n_marked=len(deadM), n_B=len(deadB), n_rows=len(rows))
by_act = collections.OrderedDict()
for x in rows:
    d = DEC[(x["sheet"], x["row"])]
    a = by_act.setdefault(d["act"], dict(rows=0, cells=0, andy=0, andy_cells=0))
    a["rows"] += 1; a["cells"] += x["n"]
    if d["andy"] == "是": a["andy"] += 1; a["andy_cells"] += x["n"]
st["by_act"] = by_act
st["by_page_act"] = {}
for x in rows: st["by_page_act"].setdefault(x["sheet"], collections.Counter())[DEC[(x["sheet"], x["row"])]["act"]] += x["n"]
st["by_rule"] = {}
for x in rows:
    r = rule_of(x); v_ = st["by_rule"].setdefault(r, [0, 0]); v_[0] += 1; v_[1] += x["n"]
st["del_cells"] = by_act.get("刪除", {}).get("cells", 0)
st["mirror_extra_cells"] = sum(n for (k, (m, n)) in mir_cnt.items() if DEC[k]["act"] == "刪除")
st["mirror_del_rows"] = [(k, v) for k, v in mir_cnt.items() if DEC[k]["act"] == "刪除"]
st["conflicts"] = list(cf_rows.items())
st["unused_upa"] = unused_upa
up_cells = sum(x["n"] for x in rows if DEC[(x["sheet"], x["row"])]["act"] in ("升級", "升級（隨附上游）", "接進 Checks"))
st["becomes_core_cells"] = up_cells
st["nform_after"] = st["nform"] - st["del_cells"]
st["nform_after_with_mirror"] = st["nform"] - st["del_cells"] - st["mirror_extra_cells"]
st["dead_after"] = st["n_dead"] - st["del_cells"] - up_cells
pickle.dump(dict(st=st, rows=rows, DEC=DEC, mir=mir_cnt, rule={ (x['sheet'],x['row']): rule_of(x) for x in rows }), open(sys.argv[5], "wb"))
print({k: v for k, v in st.items() if k not in ("by_page_act", "by_rule", "conflicts", "mirror_del_rows", "by_act")})
for a, v in by_act.items(): print(a, v)

# ---------- xlsx ----------
HF = Font(bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="305496")
def sheet(wb, title, header, data, widths, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(header), 8)); ws.row_dimensions[1].height = 60; r0 = 2
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h); c.font = HF; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(data, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v); c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths, 1): ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = f"A{r0}:{L(len(header))}{r0 + len(data)}"
    return ws

wb = openpyxl.Workbook(); wb.remove(wb.active)
# 0 總覽
ov = [["處置", "列數", "格數", "其中需 Andy 判斷（列）", "需 Andy 判斷（格）"]]
for a, v in by_act.items(): ov.append([a, v["rows"], v["cells"], v["andy"], v["andy_cells"]])
ov.append(["合計", len(rows), len(deadB), sum(v["andy"] for v in by_act.values()), sum(v["andy_cells"] for v in by_act.values())])
ws0 = wb.create_sheet("0_總覽")
r = 1
ws0.cell(r, 1, f"4-3 A 段｜v13.1（{P['model']}）B 級死格逐列建議（唯讀；全簿公式格 {P['nform']:,}；死格 {len(deadall):,}，其中已標用途 {len(deadM):,}、無『用途』標記的 B 級 {len(deadB):,}＝fixture 基準名單）").font = Font(bold=True)
r = 3
for row in ov:
    for j, v in enumerate(row, 1):
        c = ws0.cell(r, j, v)
        if r == 3: c.font = HF; c.fill = HFILL
    r += 1
r += 1
ws0.cell(r, 1, "依頁與處置（格數）").font = Font(bold=True); r += 1
acts = list(by_act)
for j, h in enumerate(["頁"] + acts + ["合計"], 1):
    c = ws0.cell(r, j, h); c.font = HF; c.fill = HFILL
r += 1
for s, cnt in sorted(st["by_page_act"].items(), key=lambda kv: -sum(kv[1].values())):
    ws0.cell(r, 1, s)
    for j, a in enumerate(acts, 2): ws0.cell(r, j, cnt.get(a, 0))
    ws0.cell(r, len(acts) + 2, sum(cnt.values())); r += 1
r += 1
ws0.cell(r, 1, "4-0 規則編號（依 v13.1 重判）").font = Font(bold=True); r += 1
for j, h in enumerate(["規則", "列數", "格數"], 1):
    c = ws0.cell(r, j, h); c.font = HF; c.fill = HFILL
r += 1
for k, (a, b) in sorted(st["by_rule"].items()):
    ws0.cell(r, 1, k); ws0.cell(r, 2, a); ws0.cell(r, 3, b); r += 1
r += 1
for t in [f"預估 B 段完成後公式格總數：{st['nform']:,} − 刪除 {st['del_cells']:,} ＝ {st['nform_after']:,}（若連同鏡像頁同列複本格一併清除，再減 {st['mirror_extra_cells']:,} ＝ {st['nform_after_with_mirror']:,}）",
          f"預估 B 段後死格：{st['n_dead']:,} − 刪除 {st['del_cells']:,} − 升級／隨附／接進 Checks 後成為核心 {up_cells:,} ＝ {st['dead_after']:,}（保留顯示用的格仍為死格，但補『用途』標記後不再是『無標記』）"]:
    ws0.cell(r, 1, t); r += 1
ws0.column_dimensions["A"].width = 24
for j in range(2, 12): ws0.column_dimensions[L(j)].width = 14

# 1 逐列
hdr = ["頁", "列", "標籤", "代表值", "格數", "格範圍", "4-0 規則", "處置", "回答的問題（編號）", "建議名稱（升級時）", "理由", "是否需 Andy 判斷", "代表公式", "文件引用（檔名）", "鏡像頁同列公式格", "承重頁（Load_Bearing）"]
data = []
for x in rows:
    d = DEC[(x["sheet"], x["row"])]
    cs = x["cells"]; rng = cs[0] if len(cs) == 1 else f"{cs[0]}:{cs[-1]}"
    v = x["value"]; v = round(v, 6) if isinstance(v, float) else v
    m = mir_cnt.get((x["sheet"], x["row"]))
    data.append([x["sheet"], x["row"], x["label"][:120], str(v)[:40], x["n"], rng, rule_of(x), d["act"], d["q"], d["name"], d["why"], d["andy"], x["formula"][:200],
                 "、".join(MEN.get((x["sheet"], x["row"]), [])[:4]), (f"{m[0]}×{m[1]}" if m else ""), "是" if x["sheet"] in dec.LB_PAGES else ""])
ws1 = sheet(wb, "1_逐列建議", hdr, data, [14, 6, 40, 12, 7, 12, 8, 16, 10, 24, 80, 10, 50, 30, 18, 10],
            note="每列＝v13.1 一個 B 級死格列（同列內各欄合併）。處置四選一（升級／保留：顯示用／接進 Checks／刪除），不確定者標『需 Andy 判斷』；『升級（隨附上游）』＝本身不新增名稱，但因下游列升級而成為核心。問題編號對應工作單『已知問題清單』1–11，— 表示不回答任何一題。")
RX = {(x["sheet"], x["row"]): x for x in rows}
# 2 升級名稱（單位、層級、token 類型標籤＝CLAUDE.md 1a）
UPG = [
 ("Cap_In", 206, "L1_MAPI_Ser", "倍（營收 ÷ 標價 token 成本）", "API（依模型層級混合，與 L1_MAPI 同口徑）", "新鮮輸入＋快取輸入＋輸出（付費 token，混合）", "C:F＝2025／2026 Q1／Q2／FY2026", "1"),
 ("Cap_In", 207, "L1_MSub_Ser", "倍", "消費訂閱方案", "付費 token（混合）", "同上", "1"),
 ("Cap_In", 208, "L1_MSeat_Ser", "倍", "商用席次方案", "付費 token（混合）", "同上；席次占比與 token 份額為 Assumed，下游須帶區間", "1"),
 ("Cap_In", 209, "L1_MImplied_Ser", "倍", "整體（token 加權）", "付費 token（混合）", "同上；FY 欄須等於 L1_MImplied 對應口徑", "1"),
 ("Cap_In", 223, "L1_MFleet_Ser", "倍", "機隊口徑", "付費 token（混合）", "同上", "1"),
 ("Cap_In", 121, "L1_MFleetFloor／_Lo／_Hi", "倍", "機隊口徑（上限讀法）", "付費 token（混合）", "與 L1_MFleet_Lo(0.184) 區間定義不同，升級後衝突會露出，需 Andy 定義", "1"),
 ("Cap_In", 285, "L1_MFleet_P25Lo", "倍", "機隊口徑，2025 單價取低端", "付費 token（混合）", "契約第 4 節：價格區間不交叉進 _Lo／_Hi，故另立名稱", "1"),
 ("Cap_In", 286, "L1_MFleet_P25Hi", "倍", "機隊口徑，2025 單價取高端", "付費 token（混合）", "同上", "1"),
 ("Workload", 38, "IF_TaskTokMult", "倍（總 token ÷ 一般聊天）", "任務別（聊天／搜尋／程式／agent／長 agent），非模型層級", "總 token（新鮮＋快取輸入＋思考＋可見輸出）", "橫列＝5 種任務", "9"),
 ("Training", 109, "IF_TrainFLOP_Luna／_Sol／_Astra", "FLOPs（單一模型最終訓練）", "Luna／Sol／Astra 三世代", "不適用（訓練側，非 token 計價）", "預訓練＋長上下文＋後訓練＋合成資料＋評測", "8"),
 ("Training", 114, "IF_TrainInferShare_Luna／_Sol／_Astra", "比例（0–1）", "三世代", "不適用（GPU 小時構成）", "推論型占最終訓練 GPU 小時", "8"),
 ("Theory_Rev", 24, "IF_GMDefault_Luna", "比例（毛利率）", "Luna；OpenAI 單價；K6 預設 (c) 全成本", "付費 token（混合；依層級另列）", "H 欄以後依層級／token 類型分列", "5"),
 ("Theory_Rev", 42, "IF_GMDefault_Sol", "比例", "Sol", "同上", "同上", "5"),
 ("Theory_Rev", 60, "IF_GMDefault_Astra", "比例", "Astra", "同上", "同上", "5"),
 ("Amortize", 49, "IF_AmortRevShare", "比例（攤提 ÷ 付費營收）", "各層級相同（合計口徑）", "付費營收（混合）", "各層級相同故不分列", "2"),
]
data2 = []
for s, r_, nm, unit, tier, tt, note, q in UPG:
    x = RX[(s, r_)]; v = x["value"]; v = round(v, 6) if isinstance(v, float) else v
    data2.append([nm, s, r_, x["label"][:80], str(v)[:30], x["n"], unit, tier, tt, note, q, "是"])
sheet(wb, "2_升級名稱", ["建議名稱", "來源頁", "來源列", "標籤", "代表值", "格數", "單位", "模型層級／口徑", "token 類型", "備註", "回答的問題", "需 Andy 判斷"], data2,
      [30, 12, 8, 40, 14, 7, 22, 34, 30, 50, 8, 8],
      note="升級＝新名稱進入下游契約（IF_／L1_），須 Andy 同意。層級與 token 類型標籤依 CLAUDE.md 第 1a 節；『付費 token（混合）』者，Segment B 實作時須確認能否依層級分列，不能者在名稱後註明為混合口徑。")
# 3 需 Andy 判斷
data3 = []
for x in rows:
    d = DEC[(x["sheet"], x["row"])]
    if d["andy"] == "是":
        cat = "升級（新名稱進契約）" if d["act"] == "升級" else ("無法決定" if d["act"] == "需 Andy 判斷" else "刪除承重頁的列")
        data3.append([cat, x["sheet"], x["row"], x["label"][:80], x["n"], d["act"], d["q"], d["name"], d["why"]])
data3.sort(key=lambda t: (t[0], t[1], t[2]))
sheet(wb, "3_需Andy判斷", ["分類", "頁", "列", "標籤", "格數", "建議處置", "問題", "建議名稱", "理由"], data3, [22, 12, 6, 40, 7, 14, 8, 28, 90])
# 4 接進 Checks
data4 = [[x["sheet"], x["row"], x["label"][:80], x["n"], x["cells"][0] + ":" + x["cells"][-1], DEC[(x["sheet"], x["row"])]["q"], DEC[(x["sheet"], x["row"])]["why"]]
         for x in rows if DEC[(x["sheet"], x["row"])]["act"] == "接進 Checks"]
sheet(wb, "4_接進Checks", ["頁", "列", "標籤", "格數", "格範圍", "問題", "接線方式與理由"], data4, [12, 6, 40, 7, 14, 8, 100])
# 5 依賴與鏡像
data5 = [[k[0], k[1], v[0], v[1], DEC[k]["act"], "鏡像頁同列有公式格；B 段若刪除須同步清除，並確認鏡像測試（D13／*_Prod）" if DEC[k]["act"] == "刪除" else "保留：不動鏡像頁"] for k, v in sorted(mir_cnt.items())]
sheet(wb, "5_鏡像頁同列", ["頁", "列", "鏡像頁", "鏡像頁同列公式格數", "本列處置", "說明"], data5, [14, 6, 18, 12, 14, 80],
      note=f"依賴檢查：保留／升級列的上游含刪除列＝{len(cf_rows)} 件；『隨附上游』無人讀取＝{len(unused_upa)} 列。")
wb.save(sys.argv[4])
print("saved", sys.argv[4])
