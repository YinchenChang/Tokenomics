"""把 02_audit.py 的結果寫成明細 xlsx。用法：python3 03_write_xlsx.py <audit.pkl> <out.xlsx>"""
import pickle, sys, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L

o = pickle.load(open(sys.argv[1], 'rb'))
wb = Workbook(); wb.remove(wb.active)
HDR = PatternFill('solid', fgColor='FFD9E1F2')

def sheet(name, rule, headers, rows, widths=None):
    ws = wb.create_sheet(name)
    ws.cell(1, 1, rule).alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
    ws.row_dimensions[1].height = max(45, 15 * (rule.count('\n') + 2 + len(rule) // 160))
    for j, h in enumerate(headers, 1):
        c = ws.cell(2, j, h); c.font = Font(bold=True); c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, r in enumerate(rows, 3):
        for j, val in enumerate(r, 1):
            if isinstance(val, (datetime.date, datetime.datetime)): val = val.isoformat()
            if isinstance(val, str) and val.startswith('='): val = "'" + val   # 公式以文字呈現
            ws.cell(i, j, val)
    ws.freeze_panes = 'A3'
    if rows: ws.auto_filter.ref = f'A2:{L(len(headers))}{len(rows) + 2}'
    for j, w in enumerate(widths or [], 1): ws.column_dimensions[L(j)].width = w
    return ws

SEV = {'高': 0, '中': 1, '低': 2}
c1 = sorted(o['c1'], key=lambda z: (SEV[z['sev']], z['item']))
sheet('C1_輸入治理', '規則（C1，以 Gov_Map 為準）：1 藍字輸入是否都登錄、Gov_Map 是否指向實際輸入格；2 SRC_ID 與決策 ID 是否皆為「—」；3 Analogy／Assumed 是否有數值低高（本格本身為低或高端點者豁免）、原始數據或模型頁標記 Verified 是否有 SRC_ID；'
      '4 目前值是否在低高之間；5 SRC 文件日期距 2026-10-15 超過 6 個月、超過 12 個月、無日期（年或月粒度一律取該期間最後一天，偏向少報）；6 被模型公式直接引用 0 次（只計模型頁，不計 Gov_Map、Load_Bearing、SRC_Index 等登錄頁的讀取）。'
      '嚴重度為機械排序用，不代表對錯判斷。',
      ['嚴重度', '檢查項', '工作表', '格', 'GM_ID', '列標籤', '細節'],
      [[z['sev'], z['item'], z['sheet'], z['cell'], z['gm'], z['label'], z['detail']] for z in c1], [7, 42, 13, 10, 12, 40, 70])

c2 = sorted(o['c2'], key=lambda z: (-(z['days'] or 0), z['id']))
sheet('C2_證據積壓', '規則（C2）：DB_Evidence 中，狀態為「待判定」或「待 Andy」、判定為「待查」、或證據日期晚於相關 SRC 的文件日期但判定不是「採納」開頭者，全部列出。積壓天數＝2026-10-15 − 證據日期。'
      'SRC 文件日期為年或月粒度時取該期間最後一天（偏向少報）。「部分採納」「並列」「遷移」「補登」等判定也會被第三條規則列入，是否算積壓需人工判讀。',
      ['Evidence ID', '日期', '積壓天數', '類別', '主張摘要', '對應參數', '狀態', '判定', '相關 SRC_ID', '標記', '來源等級', '列入原因', 'DB_Evidence 列'],
      [[z['id'], z['date'], z['days'], z['cat'].strip('｜'), z['claim'], z['param'], z['status'], z['judge'], z['src'], z['tag'], z['grade'], z['reasons'], z['row']] for z in c2],
      [10, 11, 8, 26, 50, 30, 9, 16, 30, 14, 8, 60, 8])

# 瘦身盤點
cats = ['核心', '敏感度', '鏡像', '治理與來源', '死格']
rows = []
tot = collections.Counter()
for s, cn in o['c3_page'].items():
    rows.append([s] + [cn.get(c, 0) for c in cats] + [cn.get('公式格', 0), cn.get('輸入或常數', 0)])
    for c in cats + ['公式格', '輸入或常數']: tot[c] += cn.get(c, 0)
rows.append(['（合計）'] + [tot[c] for c in cats] + [tot['公式格'], tot['輸入或常數']])
nf = o['nform']; sl = o['slim']
rows.append([]); rows.append(['瘦身估計（占全簿公式格 %d 格）' % nf])
for c in cats: rows.append([f'{c}（公式格）', sl.get(c, 0), f"{sl.get(c, 0) / nf:.1%}"])
sheet('瘦身盤點', '規則（C3）：每格（公式格與輸入／數值常數格；純標籤文字不計）歸入一類。核心＝可由根節點往上游追到、且不在敏感度或鏡像頁；敏感度＝Sens_* 與 Sensitivity；鏡像＝*_Prod；治理與來源＝Gov_Map、Decisions、SRC_*（含 SRC_Index）、DB_Evidence、Checks、Load_Bearing；'
      '死格＝從任何根節點往上游都追不到。根節點＝下游契約所列名稱、全部 IF_／IFW_／IFC_／L1_／GOV_ 具名範圍、Checks 頁全部格。追溯時不穿越 Gov_Map、Decisions、DB_Evidence、Load_Bearing、SRC_Index（否則 Gov_Map「模型值」欄會把每一格輸入都接到 GOV_*）。範圍引用把整個範圍都算上游（保守）。README、Sources_Legacy 不計。',
      ['工作表'] + cats + ['其中公式格', '其中輸入或常數'], rows, [34, 9, 9, 9, 11, 9, 11, 13])

dead = o['dead']
sheet('C3_死格', '規則（C3 死格）：從任何根節點往上游都追不到的公式格或輸入格。「同列是否有核心格」為「是」者多半是同列的對照、顯示或說明欄；「否」表示整列都追不到。「引用者類別」＝直接引用此格的格所屬類別（無任何格引用＝終點格）。'
      '「經治理頁可達」＝若允許穿越 Gov_Map 等登錄頁則可追到（只因被登錄而被讀取）。',
      ['工作表', '格', '列', '列標籤', '種類', '同列是否有核心格', '引用者類別', '經治理頁可達', '公式或值', '目前值'],
      [[z['sheet'], z['cell'], z['row'], z['label'], z['kind'], z['row_has_core'], z['users'], z['reach_via_gov'], z['formula'], str(z['value'])[:60]] for z in dead],
      [13, 7, 6, 40, 12, 18, 18, 9, 60, 16])
sheet('C3_核心推導列', '規則（C3 核心推導列）：含至少一個「核心」格的列，以「頁＋列」為單位，供第 3 段使用。公式格數與輸入格數為該列核心格的計數。',
      ['工作表', '列', '列標籤', '核心公式格數', '核心輸入格數'],
      [[z['sheet'], z['row'], z['label'], z['formulas'], z['inputs']] for z in o['core_rows']], [14, 6, 70, 10, 10])

rows = []
for p, st in o['c4_stats'].items():
    for k, n in st.items(): rows.append(['統計', p, '', '', k, '', '', n, ''])
for z in o['c4']:
    rows.append(['差異', z['page'], z['cell'], z['core_cell'], z['diff'], z['label'], z['documented'], '', z['mirror_f'] + '  ⟷  ' + z['core_f']])
sheet('C4_鏡像一致性', '規則（C4）：*_Prod 與核心頁逐格以相對位置（R1C1）比對；*_Prod 頁名視同核心頁名（頁首說明明訂「讀到鏈上其他格者改讀本組 *_Prod 頁」）；右側 S:AG 自我檢查副本對核心頁 C:Q（平移 16 欄）；Perf_Prod、Perf_Batch_Prod 第 50 列改讀 CTL_ProdDerate 為明訂差異；具名範圍換成其位址後再比對。'
      'Sens_Perf／Sens_Train 依世代索引、層級索引找核心頁同世代同層級欄比對，黃底或藍字情境格不比；列標籤與核心頁同列不同者為本頁專屬列，不比對。Sensitivity、Sens_Har 無逐格對應的核心頁，改比「同列各情境欄」與首欄的結構；Sensitivity 第 64 列起的摘要表、Sens_Rev（封閉式重算）不比對。'
      '「統計」列為各類格數；「差異」列逐格列出。',
      ['類型', '頁', '格', '對照格', '差異／統計項', '列標籤', '是否為明訂差異', '格數', '公式（本頁 ⟷ 對照）'], rows, [6, 15, 9, 9, 50, 40, 26, 7, 100])

c5 = o['c5']
sheet('C5_寫死常數', '規則（C5）：核心公式（C3 歸類為核心的公式格）中，0、1、2、100、1000、1e3、1e6、1e9、1e12、24、365、8760、3600 以外的數字常數，逐格列出。為方便閱讀另依上下文分類（機械分類，非判斷）：'
      '疑似參數或門檻＝最需人工判讀；時間換算＝12、168、525600 等；10 的次方＝1e-9、1e15 等單位換算；哨兵值＝9E+99 等；索引運算＝位於 INDEX／MATCH 內（例如 3*(世代-1)+2）；位置或格式參數＝INDEX 列欄號、TEXT 格式、ROUND 位數等函數參數。',
      ['分類', '工作表', '格', '列', '列標籤', '常數', '所在函數', '公式片段'],
      [[z['ctx'], z['sheet'], z['cell'], z['row'], z['label'], z['const'], z['func'], z['frag']] for z in c5], [22, 13, 7, 6, 40, 14, 12, 80])

rows = [[z['item'], z['sheet'], z['cell'], z['label'], z['unit'], z['other'], z['note']] for z in o['c6']]
rows += [['（資訊）有藍字輸入但沒有「單位」欄的頁', s, '', '', '', '', '這些頁的輸入單位寫在列標籤或說明欄，C6.1 無法機械檢查'] for s in o['no_unit_pages']]
sheet('C6_單位與標籤', '規則（C6）：1 有「單位」欄的頁，藍字數值輸入的單位欄為空白者列出；2a 列標籤去除世代名、括號內文字後相同、但在不同頁的單位不同者列出；2b 列標籤與單位欄的口徑詞互相矛盾（每 GPU 對 每架、$/M 對 $/k、GW 對 MW、年 對 月）者列出。只標可疑，不判斷對錯。',
      ['檢查項', '工作表', '格或列', '列標籤', '單位', '同鍵各頁單位', '備註'], rows, [44, 13, 9, 50, 14, 40, 40])

sheet('C7_第1段輸入清單', '規則（C7-1）：每一格輸入一列（藍字數值、非藍字數值常數、以及 Gov_Map 登錄而非公式的格；藍字文字選項不列）。被引用次數＝被模型頁公式直接引用（穿過具名範圍）的格數；含登錄頁者另列。是否在 m 路徑：見 C7_m路徑。'
      '世代維度：「世代欄頭」＝該頁前 11 列出現 2 個以上世代名；「列內多值」＝同列 C 欄以後有 2 個以上不同數值。來源日期＝Gov_Map SRC_ID 對應 SRC 文件日期（多個取最早）。最後四欄留給第 1 段填寫。',
      ['GM_ID', '頁', '格', '列標籤', '單位', '目前值', '輸入種類', '標記（Gov_Map 類別或頁內標記）', '低', '高', 'SRC_ID', '決策 ID', '來源日期', '被模型公式引用次數', '含登錄頁引用次數', 'C3 分類', '是否在 m 路徑', '世代維度',
       '代表什麼物件', '時點', '應否隨世代或時間變化', '口徑是否與使用它的公式一致'],
      [[z['gm'], z['sheet'], z['cell'], z['label'], z['unit'], z['value'] if not isinstance(z['value'], str) else z['value'][:60], z['kind'], z['tag'], z['lo'], z['hi'], z['src'], z['dec'], z['sdate'], z['refs'], z['refs_all'], z['cls'], z['mpath'], z['gen'], '', '', '', ''] for z in o['c7']],
      [9, 12, 7, 36, 10, 12, 12, 16, 8, 8, 14, 8, 11, 9, 9, 8, 8, 16, 18, 12, 16, 18])

rows = [['根', '', '', t, '', '', ''] for t in o['m_roots']]
rows += [['排除的關鍵詞命中（非 m 語意）', '', '', t, '', '', ''] for t in o['m_excluded']]
rows += [['路徑列', z['sheet'], z['row'], z['label'], z['formulas'], z['inputs'], z['cls']] for z in o['m_rows']]
sheet('C7_m路徑', '規則（C7-2）：以 m 相關名稱為根（CTL_Monetize、L1_M*、L1_Gap*、L1_ScaleServe*、L1_TokPerPromptClose、L1_PriceReal*、L1_LReqFleet、AL_SpendBasis、AL_Kappa，以及名稱或列標籤含「變現／Monetiz／閉合／缺口／落差／κ／Kappa／計價倍數」者；Cap_In 列標籤含上述詞者取該列第一個公式或輸入格）。'
      '往上游追所有格（不穿越 Gov_Map 等登錄頁），以「頁＋列」列出；SRC 頁原始值與 Gov_Map 等登錄頁不列為路徑列。DRV_Close（電力閉合比）與 AL_J8Gap（研發計畫落差）因關鍵詞命中但與 m 無關，排除。',
      ['類型', '工作表', '列', '列標籤或根', '路徑上公式格數', '路徑上輸入格數', 'C3 頁類別'], rows, [12, 14, 6, 80, 10, 10, 10])
wb.save(sys.argv[2])
print('saved', [ws.title for ws in wb.worksheets])
