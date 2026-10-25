"""產生說明句盤點表 CSV（text_inventory_L.csv／text_inventory_m.csv）。
用法：python3 make_text_inventory.py <L|m> <閱讀版.xlsx（未重算亦可）> <*.replaced.json> <輸出.csv>
類別：現行值（已改公式）／已是公式（原本就是公式句）／歷史保留／外部事實保留／原值貼上／預設值保留／待Project判斷（保留）
check_reading*.py 的「說明句數字檢查」讀這份表。
"""
import csv, json, re, sys
import openpyxl
sys.path.insert(0, __file__.rsplit('/', 1)[0] if '/' in __file__ else '.')
from text_scan import scan_static

KIND, WB, RJ, OUT = sys.argv[1:5]

STATIC = {
 'L': {
  ('0_說明', 'B19'): ('歷史保留', '修訂紀錄（v17.0）', '不改'),
  ('0_說明', 'B20'): ('歷史保留', '修訂紀錄（v21.2 r0）', '不改'),
  ('0_說明', 'A21'): ('歷史保留', '修訂紀錄（v21.3 r0）', '不改'),
  ('0_說明', 'B22'): ('歷史保留', '修訂紀錄（v21.4 r0，本次）：格數與列號是本次作業的紀錄', '不改'),
  ('1_公式樹', 'G12'): ('原值貼上', 'Tokenomics L1_LTotalFleetVDisc_Lo／_Hi（check_reading.py 逐格核對）', '貼上值，保留（核對用，不可改公式）'),
  ('1_公式樹', 'G13'): ('原值貼上', 'Tokenomics L1_LReqFleet_Full_Lo／_Hi', '貼上值，保留'),
  ('1_公式樹', 'G14'): ('原值貼上', 'Tokenomics Cap_In!C242', '貼上值，保留'),
  ('1_公式樹', 'G80'): ('原值貼上', 'Tokenomics L1_LReqFleet_IdNoK_Lo／_Hi', '貼上值，保留'),
  ('2_輸入', 'I12'): ('歷史保留', 'ρ 低端版本沿革（0.1704→0.2007→0.2057）與機架價 7.8→7.5', '不改'),
  ('2_輸入', 'I26'): ('歷史保留', '各世代每 GW 營收的 v19.1 前舊值', '不改'),
  ('2_輸入', 'G31'): ('外部事實保留', 'OpenAI 自述約 70%（來源說法，不是模型值）', '不改'),
  ('4_必須為真', 'B19'): ('待Project判斷（保留）', '以 Python 另算的精確值 −5.068%（2026-10-24 chat 端），沒有可引用的格', '不改；換版若 g 的損益兩平改變須手動更新'),
  ('4_必須為真', 'B150'): ('待Project判斷（保留）', '以 Python 另算的精確值 T1 4.9525 年（2026-10-24 chat 端），沒有可引用的格', '不改；同上'),
 },
 'm': {
  ('0_說明', 'B5'): ('歷史保留', 'Andy 2026-10-21 原話（含「折扣20%」引文）', '不改'),
  ('0_說明', 'B30'): ('歷史保留', '修訂紀錄（v17.0）', '不改'),
  ('0_說明', 'B31'): ('歷史保留', '修訂紀錄（v21.2 r0）', '不改'),
  ('0_說明', 'B33'): ('歷史保留', '修訂紀錄（v21.3 r0）', '不改'),
  ('2_輸入', 'F14'): ('外部事實保留', 'SRC_DEM_039 FT 約 70%（來源說法）', '不改'),
  ('2_輸入', 'J16'): ('外部事實保留', 'SRC_DEM_023：2026-03 OpenAI 消費約 60%、企業 >40%', '不改'),
  ('2_輸入', 'J29'): ('待Project判斷（保留）', '推導說明「2.5 × 1.0 B ÷ 0.7 B 週活 ≈ 3.6」：1.0、0.7 是來源數字，本簿沒有對應格；只改 3.6 會與推導式不一致', '不改'),
  ('2_輸入', 'J51'): ('外部事實保留', 'DeepSeek 56.3%、agent 流量約 88%（來源說法）', '不改'),
  ('2_輸入', 'J64'): ('預設值保留', '試算輸入 IN_DTest 的預設值 0.30（設計值，不是模型值）', '不改'),
  ('5_敏感度', 'N28'): ('待Project判斷（保留）', 'ε 的區間 0–0.5 只在 Gov_Map 文字，本簿沒有 ε 的格', '不改'),
  ('6_不確定與證據', 'F11'): ('外部事實保留', 'OpenRouter 等：2026 agent 流量快取約 88%', '不改'),
  ('6_不確定與證據', 'H11'): ('外部事實保留', 'OpenRouter 的 prompt／completion 長度與快取比例', '不改'),
 },
}[KIND]

wb = openpyxl.load_workbook(WB)
rep = json.load(open(RJ))
rows = []
repset = set()
for sh, a, orig, f in rep:
    repset.add((sh, a))
    refs = sorted(set(re.findall(r"IN_\w+|'[^']+'![A-Z]+\d+|(?<![A-Za-z_'])[A-Z]{1,2}\d+(?![\w(])", re.sub(r'"[^"]*"', '', f))))
    rows.append((sh, a, '現行值', orig, '、'.join(refs), '改公式（顯示文字與原文相同）'))
# 原本就是公式的說明句
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('=') and '"&' in v + '&"' and 'TEXT(' in v and (ws.title, c.coordinate) not in repset:
                rows.append((ws.title, c.coordinate, '已是公式', '', '見公式', '原本就是公式句，不改'))
seen = set()
for sh, a, v, tk in scan_static(wb):
    k = (sh, a)
    if k not in STATIC:
        sys.exit(f'未分類的含數字文字格：{sh}!{a} {tk} {v[:60]}')
    cl, src, how = STATIC[k]
    rows.append((sh, a, cl, v, src, how))
    seen.add(k)
miss = set(STATIC) - seen
if miss:
    sys.exit(f'盤點表多列了掃描不到的格：{miss}')
order = {s: i for i, s in enumerate(wb.sheetnames)}
rows.sort(key=lambda r: (order[r[0]], int(re.sub(r'\D', '', r[1])), r[1]))
with open(OUT, 'w', newline='', encoding='utf-8-sig') as fh:
    w = csv.writer(fh)
    w.writerow(['頁', '格', '類別', '原文', '數字對應的來源格或名稱', '處理方式'])
    w.writerows(rows)
from collections import Counter
print(dict(Counter(r[2] for r in rows)))
