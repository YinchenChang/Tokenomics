# v19.0 (work order docs/workorders/20261022_v19.0_Interface取數.md r0, branch claude/v19-0): Interface section N (P x Q 取數) for downstream company models.
# MAJOR version: new Interface contract names only (IF_Disc, IF_MixPaid_*, IF_MixFree_*, IF_UtilActual, IF_ProdDerate, IF_FleetProdRatio_VR200, IF_PxEffMix, IF_PxEffMix25,
#   IF_LTotalDisc, IF_LTotalFleetVDisc, IF_TokGWPaidMix). Each is a reference to, or a combination of, existing cells; no existing named value or existing formula cell changes value.
#   W1  Interface section N (the work order calls it "K"; K／L／M are already taken by v8.0／v10.0／v14.0, so the next free letter is used).
#   W2  Decisions SL53-01, README row, downstream contract document (outside the workbook).
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

from common import put, section
import v1410
import v1800

VERSION = "20261022_Tokenomics_v19.0"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-22"
V = "v19.0"
WO = "工作單 docs/workorders/20261022_v19.0_Interface取數.md（r0）"
DEC_BY = ("Andy 2026-10-22 05:19「Tokenomics 裡面應該只是提供一 G 瓦的……持有成本……算力……三 G 瓦其實應該是擺在 OpenAI 他們自己裡面的，"
          "然後他只是去取用 Tokenomics 裡面一 G 瓦的相關的訊息」；05:23「你這邊做 A 段就好了」")

N_TITLE = ("N. P×Q 取數（v19.0）：下游公司模型做「營收＝P × Q、實際 m×L ÷ 所需 m×L」時，只從本節取每 GW 參數（B4_／AL_／L1_／CTL_ 下游不得直接連結）；"
           "公司層級（多 GW、多世代）計算放在 company-models，Tokenomics 只露出每 GW 參數。本節每列只是既有格的引用或其組合，不改任何既有格")
PCT = "0.0%"

TIERS = (("Luna", 1), ("Sol", 2), ("Astra", 3))

# (name, label, unit, fmt, C, D, E)  D／E None = no range
SCALARS = [
    ("IF_Disc", "折扣 d（有效牌價＝表列牌價 ×（1 − IF_Disc）；基準／低／高）", "%", PCT, "=B4_Disc", "=B4_Disc_Lo", "=B4_Disc_Hi"),
] + [
    (f"IF_MixPaid_{t}", f"付費 token 層級占比 — {t}（B4_MixPaid；付費 token 中該層級的占比；基準）", "%", PCT, f"=INDEX(B4_MixPaid,1,{k})", None, None) for t, k in TIERS
] + [
    (f"IF_MixFree_{t}", f"免費 token 層級占比 — {t}（B4_MixFree；免費 token 中該層級的占比；基準）", "%", PCT, f"=INDEX(B4_MixFree,1,{k})", None, None) for t, k in TIERS
] + [
    ("IF_UtilActual", "實際利用率（營運口徑；實際產出＝產能 × 本列；基準／低／高）", "x", "0.00", "=AL_UtilActual", "=AL_UtilActual_Lo", "=AL_UtilActual_Hi"),
    ("IF_ProdDerate", "生產折減（線性；＝CTL_ProdDerate；其他世代的重解比待補，下游暫用本列；基準）", "x", "0.00", "=CTL_ProdDerate", None, None),
    ("IF_FleetProdRatio_VR200", "VR200 生產折減重解比（＝L1_FleetProdRatio；重解 SLO 後 VR200 每 GW 產出 ÷ 理想產出；僅 VR200；基準）", "x", "0.00000", "=L1_FleetProdRatio", None, None),
    ("IF_PxEffMix", "付費層級組合的有效牌價 — 2026 快照（＝表列 ×（1 − IF_Disc）；Luna／Sol／Astra 依 IF_MixPaid_* 加權；基準）", "$/M", "0.00000", "=Cap_In!$C$90", None, None),
    ("IF_PxEffMix25", "付費層級組合的有效牌價 — 2025 年度（CAP_Price25Mix；各層級 2025 年度加權單價 ×（1 − IF_Disc）；基準／低／高）", "$/M", "0.00000", "=CAP_Price25Mix", "=Cap_In!$D$281", "=Cap_In!$E$281"),
    ("IF_LTotalDisc", "L（新世代自身起點、折現；＝L1_LTotalDisc；壽命期年均產出占第一年的比例；基準／低／高）", "x", "0.00000", "=L1_LTotalDisc", "=L1_LTotalDisc_Lo", "=L1_LTotalDisc_Hi"),
    ("IF_LTotalFleetVDisc", "L（2025 機隊、世代起點、折現；＝L1_LTotalFleetVDisc；基準／低／高）", "x", "0.00000", "=L1_LTotalFleetVDisc", "=L1_LTotalFleetVDisc_Lo", "=L1_LTotalFleetVDisc_Hi"),
]

PAID_MIX = "IF_TokGWPaidMix"
PAID_MIX_LAB = ("付費層級組合下每 GW 年產能（100% 利用、全部付費；層級占比調和平均；實際＝× 服務占比 IF_ServeShare ×（1 − 免費占比 IF_FreeShare）× 實際利用率 × 生產折減；"
                "欄 C–Q＝5 世代 × 3 成本情境，同 IF_TokGW_*；token 類型＝付費 token 混合：新鮮輸入＋快取輸入＋輸出）")


def _paid_mix(k):
    """k: 1-based column position inside C:Q. Non-numeric (or non-positive) denominator parts return AL_UnavailTxt."""
    tg = {t: f"INDEX(IF_TokGW_{t},1,{k})" for t, _ in TIERS}
    ok = ",".join([f"ISNUMBER(IF_MixPaid_{t})" for t, _ in TIERS] + [f"ISNUMBER({tg[t]})" for t, _ in TIERS])
    pos = ",".join(f"{tg[t]}>0" for t, _ in TIERS)
    den = "+".join(f"IF_MixPaid_{t}/{tg[t]}" for t, _ in TIERS)
    return f"=IF(AND({ok}),IF(AND({pos}),1/({den}),AL_UnavailTxt),AL_UnavailTxt)"


def interface_n(wb):
    ws = wb["Interface"]
    start = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    section(ws, start, N_TITLE, 17)
    r = start + 1
    made = []

    def name_it(name, ref):
        wb.defined_names[name] = DefinedName(name, attr_text=ref)

    for name, lab, unit, fmt, c, d, e in SCALARS:
        put(ws, f"A{r}", f"{lab}　[{name}]"); put(ws, f"B{r}", unit)
        put(ws, f"C{r}", c, fmt=fmt)
        name_it(name, f"Interface!$C${r}")
        if d is not None:
            put(ws, f"D{r}", d, fmt=fmt); put(ws, f"E{r}", e, fmt=fmt)
            name_it(name + "_Lo", f"Interface!$D${r}"); name_it(name + "_Hi", f"Interface!$E${r}")
        made.append((name, r)); r += 1
    put(ws, f"A{r}", f"{PAID_MIX_LAB}　[{PAID_MIX}]"); put(ws, f"B{r}", "M tok/GW/年")
    for k in range(3, 18):
        put(ws, f"{L(k)}{r}", _paid_mix(k - 2), fmt="#,##0")
    name_it(PAID_MIX, f"Interface!$C${r}:$Q${r}"); made.append((PAID_MIX, r))
    return dict(start=start, made=made)


NEW_NAMES = []
for _n, _l, _u, _f, _c, _d, _e in SCALARS:
    NEW_NAMES += [_n] + ([_n + "_Lo", _n + "_Hi"] if _d is not None else [])
NEW_NAMES.append(PAID_MIX)


# ---------------------------------------------------------------- Decisions SL53-01
def _decisions():
    D = [("Interface 新增 N 節「P×Q 取數」：公司層級計算放 company-models，Tokenomics 只露出每 GW 參數",
          "新增 IF_Disc／_Lo／_Hi、IF_MixPaid_Luna／Sol／Astra、IF_MixFree_Luna／Sol／Astra、IF_UtilActual／_Lo／_Hi、IF_ProdDerate、IF_FleetProdRatio_VR200、IF_PxEffMix、IF_PxEffMix25／_Lo／_Hi、"
          "IF_LTotalDisc／_Lo／_Hi、IF_LTotalFleetVDisc／_Lo／_Hi、IF_TokGWPaidMix。每一名稱是既有格的引用或組合，既有格數值不變。"
          "理由：下游（company-models 的 OpenAI 模型）做 P × Q 時，B4_／AL_／L1_／CTL_ 依 README 不得直接連結。"
          "Andy 原意：Tokenomics 只提供一 GW 的持有成本、算力等資訊；三 GW 等公司層級計算放在 OpenAI 模型自己裡面，只取用 Tokenomics 一 GW 的相關資訊。"
          "其他世代的重解生產折減比目前只有 VR200（IF_FleetProdRatio_VR200），其他世代下游暫用線性 IF_ProdDerate，是否補齊待 Project 判斷。"
          "工作單稱新節為「K」，但 Interface K／L／M 節已存在（v8.0／v10.0／v14.0），本節取下一個字母 N。",
          "Interface N 節")]
    return [[f"SL53-{i + 1:02d}", V, f"v19.0 Interface 取數：{t}", txt, DATE, DEC_BY, "v19.0 已建", where, f"{WO} W1／W2", "否"] for i, (t, txt, where) in enumerate(D)]


def apply(wb):
    log = {}
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == "SL53-01" for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


# ================================================================== README
README_VERSION = (VERSION + "（主版：" + WO + "——Interface 新增 N 節「P×Q 取數」：IF_Disc、IF_MixPaid_*、IF_MixFree_*、IF_UtilActual、IF_ProdDerate、IF_FleetProdRatio_VR200、"
                  "IF_PxEffMix、IF_PxEffMix25、IF_LTotalDisc、IF_LTotalFleetVDisc（含 _Lo／_Hi）與 IF_TokGWPaidMix；皆為既有格的引用或組合；Decisions SL53-01。既有名稱的值全部不變。）以下為 "
                  + v1800.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1800.README_TITLE
README_TITLE = V + _T[len(v1800.V):]
README_ROW = ("v19.0 Interface 取數",
              "Interface N 節「P×Q 取數」：下游做營收＝P × Q、實際 m×L ÷ 所需 m×L 時，每 GW 參數只從這裡取（IF_Disc、IF_MixPaid_*、IF_MixFree_*、IF_UtilActual、IF_ProdDerate、"
              "IF_FleetProdRatio_VR200、IF_PxEffMix、IF_PxEffMix25、IF_LTotalDisc、IF_LTotalFleetVDisc、IF_TokGWPaidMix）；公司層級計算放 company-models。"
              "細節見工作單 docs/workorders/20261022_v19.0_Interface取數.md 與 Decisions SL53-01。")
