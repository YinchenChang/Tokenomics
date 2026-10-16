"""產生 docs/evidence/20261016_服務條件/證據表.xlsx（證據表、情境計算、MLPerf_v6.0 三頁）。
用法：python3 build_xlsx.py <repo 根目錄>
情境計算頁全部為 Excel 公式；Perf 現值由 v10.0（data_only）讀入，以藍字標示。"""
import sys, csv, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L

R = sys.argv[1].rstrip("/") + "/"
EV = R + "docs/evidence/20261016_服務條件/"
src = openpyxl.load_workbook(R + "model/20261016_Tokenomics_v10.0.xlsx", data_only=True)
perf = src["Perf"]
mfleet = src["L1"]["D78"].value

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "證據表"
hdr = ["題號", "數值", "單位與口徑", "日期", "來源（URL）", "標記", "利害關係方與動機", "是否為獨立的第二來源", "讀取狀態", "來源等級", "備註"]
ws.append(hdr)
AA = "Artificial Analysis（第三方量測商；以 API 探測請求量測，無賣方利益，但與供應商有合作關係）"
OR = "OpenRouter（路由商；數值來自實際流量 P50，無賣方利益）"
OAI = "OpenAI（賣方；宣傳產品速度改善）"
rows = [
 # Q1
 ["Q1", "GPT-6 Astra (high)：OpenAI 52.1 tok/s；首個答案 token 47.5 s；500 token 端到端 57.1 s", "每用戶輸出速度（生成中、首 chunk 之後）；10,000 輸入 token 探測；72 小時中位數", "2026-10（72 小時窗；模型 2026-09 發布）", "https://artificialanalysis.ai/ko/models/gpt-6-astra-high/providers", "Verified", AA, "是（OpenRouter GPT-6 Astra）", "已讀（WebFetch 摘要）", "2", "Astra 層代表模型。英文 non-reasoning 頁讀取失敗"],
 ["Q1", "GPT-6 Astra：OpenAI 32 tps、延遲 4.16 s；Azure 24／9.36 s；Bedrock 52；Flex 40；Fast 68（2 倍價）", "實際流量 P50 throughput 與首 token 延遲；標準路由只用 OpenAI、Azure", "2026-10（3 日窗）", "https://openrouter.ai/models/openai/gpt-6-astra", "Verified", OR, "是（AA）", "已讀（WebFetch 摘要）", "2", "標準端點 24–32；加價端點 52–68"],
 ["Q1", "GPT-6 Sol (max)：OpenAI 102.8 tok/s；Azure 145.8；首個答案 token 113.7 s", "同 AA 口徑；max 推理強度", "2026-10（72 小時窗；模型 2026-09-22 發布）", "https://artificialanalysis.ai/models/gpt-6-sol/providers", "Verified", AA, "是（OpenRouter GPT-6 Sol）", "已讀（WebFetch 摘要）", "2", "頁面標示已被 GPT-6.1 Sol 取代；medium 頁無數據"],
 ["Q1", "GPT-6.1 Sol (medium)：OpenAI 50.7 tok/s；首個答案 token 4.24 s", "同 AA 口徑", "2026-10（72 小時窗；模型 2026-09-29 發布）", "https://artificialanalysis.ai/models/gpt-6-1-sol-medium/providers", "Verified", AA, "是（OpenRouter GPT-6.1 Sol）", "已讀（WebFetch 摘要）", "2", "目前 Sol 層主力"],
 ["Q1", "GPT-6 Sol：OpenAI 63 tps、2.72 s；Azure 22／6.63 s；Azure EU 62；Bedrock US 61；Fast 61；Flex 70", "OpenRouter 實際流量 P50", "2026-10（3 日窗）", "https://openrouter.ai/models/openai/gpt-6-sol", "Verified", OR, "是（AA）", "已讀（WebFetch 摘要）", "2", ""],
 ["Q1", "GPT-6.1 Sol：OpenAI 37 tps、3.36 s；Azure 24／12.05 s；Bedrock 39；Flex 33；Fast 64（2 倍價）；Ultrafast 86（6 倍價）", "OpenRouter 實際流量 P50", "2026-10（3 日窗）", "https://openrouter.ai/models/openai/gpt-6.1-sol", "Verified", OR, "是（AA）", "已讀（WebFetch 摘要）", "2", "標準端點 24–39"],
 ["Q1", "GPT-5.6 Sol (max)：OpenAI 92.8 tok/s；Azure 114.6；Bedrock 87.3", "AA 口徑", "2026-10（72 小時窗；模型 2026-07 發布、已棄用）", "https://artificialanalysis.ai/models/gpt-5-6-sol/providers", "Verified", AA, "是（OpenRouter GPT-5.6 Sol）", "已讀（WebFetch 摘要）", "2", "Serving S26 所引的前一代"],
 ["Q1", "GPT-5.6 Sol：OpenAI 66 tps、4.84 s；Fast 89；Bedrock 75；Azure 28–57；Flex 44", "OpenRouter 實際流量 P50", "2026-10（3 日窗）", "https://openrouter.ai/models/openai/gpt-5.6-sol-20260709", "Verified", OR, "是（AA）", "已讀（WebFetch 摘要）", "2", ""],
 ["Q1", "GPT-6 Luna：OpenAI 116.5 tok/s；Azure 201.2；Flex 45.9", "AA 口徑", "2026-10（72 小時窗；模型 2026-09 發布）", "https://artificialanalysis.ai/models/gpt-6-luna/providers", "Verified", AA, "是（OpenRouter GPT-6 Luna）", "已讀（WebFetch 摘要）", "2", "Luna 層代表模型（價格 0.10／0.50，與 2Q 對應）"],
 ["Q1", "GPT-6 Luna：OpenAI 61 tps、2.59 s；Fast 110；Bedrock 62；Azure 26–47；Flex 37", "OpenRouter 實際流量 P50", "2026-10（3 日窗；2026-09-22 發布）", "https://openrouter.ai/models/openai/gpt-6-luna", "Verified", OR, "是（AA）", "已讀（WebFetch 摘要）", "2", ""],
 ["Q1", "GPT-5.6 Luna (max)：OpenAI 126.7；Bedrock 201.1（另 xhigh 124.1、high 123.8）", "AA 口徑", "2026-10（72 小時窗）", "https://artificialanalysis.ai/models/gpt-5-6-luna/providers", "Verified", AA, "是（OpenRouter GPT-5.6 Luna）", "已讀（WebFetch 摘要）", "2", ""],
 ["Q1", "GPT-5.6 Luna：OpenAI 65 tps、1.83 s；Fast 122；Bedrock 110；Azure 49–104；Flex 61", "OpenRouter 實際流量 P50", "2026-10（3 日窗）", "https://openrouter.ai/models/openai/gpt-5.6-luna", "Verified", OR, "是（AA）", "已讀（WebFetch 摘要）", "2", ""],
 ["Q1", "ChatGPT／Codex 訂閱產品預設速度：GPT-6 Astra 與 GPT-6.1 Sol 由約 30 提高到約 50 tok/s（官方稱「約快 50%」，與 30→50 的 67% 不一致）", "訂閱產品（Codex、Sign in with ChatGPT 夥伴）預設速度；未提 API", "2026-10-05", "https://runtimewire.com/article/openai-gpt-6-models-speed-optimization-tibo-sottiaux", "Interested-party", OAI, "否（PC Watch 為同一則貼文的轉述，非獨立）；OpenRouter GPT-6.1 Sol 37 tps 方向一致", "已讀（WebFetch 摘要）", "3", "原貼文（Tibo Sottiaux）未直接讀取"],
 ["Q1", "同上：TPS 由 30 提高到 50；屬軟體最佳化，未提硬體", "同上", "2026-10-06", "https://pc.watch.impress.co.jp/docs/news/2145903.html", "Interested-party", OAI, "否（同一來源轉述）", "已讀（WebFetch 摘要）", "3", ""],
 ["Q1", "OpenAI 推出約 300 tok/s 的付費高速層（$500 Pro）", "高速層；用量占比未知", "2026-09-29", "https://hermes-ai.net/news/openai-s-gpt-6-astra-hits-300-tokens-per-second-behind-a-500-plan/", "Interested-party", "OpenAI（賣方）；社群新聞站轉述，模型名稱自相矛盾", "否", "已讀（WebFetch 摘要）", "3", "只作背景；不承重"],
 ["Q1", "ChatGPT 網頁版（非 Codex）每用戶速度與首 token 時間的獨立量測：已搜尋，未找到", "—", "—", "WebSearch 2 組關鍵字；AA、OpenRouter（只量 API）；新聞", "—", "—", "—", "已搜尋，未找到", "—", "ChatGPT 端只有 OpenAI 自述（上列 30→50）"],
 # Q2
 ["Q2", "DeepSeek 線上服務：每 H800 節點 decode 平均約 14.8k 輸出 tok/s（＝每 GPU 約 1,850）；平均每用戶輸出 20–22 tok/s；每輸出 token 平均 KV 長度 4,989；decode 單元 18 節點（EP144）", "24 小時平均（2025-02-27 12:00 至 02-28 12:00，UTC+8）；V3＋R1 全部請求", "2025-03-01", "https://github.com/deepseek-ai/open-infra-index/blob/main/202502OpenSourceWeek/day_6_one_more_thing_deepseekV3R1_inference_system_overview.md", "Interested-party", "DeepSeek（自營服務；宣傳成本效率，文中並宣稱理論利潤率 545%）", "部分（Step-3 論文表 8 轉引 1,850，非獨立量測）", "已讀（WebFetch 摘要）", "1", "本段相除：每 decode GPU 並行約 1,850 ÷ 20–22 ＝ 84–92 序列（Derived）。已在 SRC_PERF_041–051"],
 ["Q2", "DeepSeek profile：decode EP128、TP1、提示長度 4K、每 GPU 批次 128 個請求；兩個 micro-batch 重疊", "官方 profile 配置（「與線上部署接近」）", "2025-02", "https://github.com/deepseek-ai/profile-data", "Interested-party", "DeepSeek（同上）", "是（Step-3 論文表 8 引其吞吐 2,324）", "已讀（WebFetch 摘要）", "1", "配合下一列：2,324 ÷ 128 ＝ 每用戶約 18.2 tok/s（Derived）"],
 ["Q2", "Step-3 論文表 8：DSv3-blog 1,850（144 GPU，KV 4,989）；DSv3-profile 2,324（128 GPU，4K）；Step-3 FP8 注意力 4,039（32 GPU，4K）、8K 為 2,643；SLA 皆為 20 tok/s（TPOT 50 ms）；Step-3 不用 MTP；總批次 6,144（約每 GPU 192）", "每 GPU decode 峰值 tok/s（TGS）；Hopper", "2025-07", "https://arxiv.org/pdf/2507.19427", "Interested-party", "StepFun（自家模型與 AFD 系統的宣傳）；DeepSeek 兩點為轉引", "是（對 DeepSeek 兩點為第二出處；對 Step-3 為唯一來源）", "已讀（WebFetch 摘要）", "1", "Step-3：321B／38B 啟用，與 Sol 不同架構，對模型只作 Analogy"],
 ["Q2", "vLLM GLM-5.2-NVFP4（744B／40B 啟用，MTP 3）於 24×B300：4 個 prefill 實例（16 GPU）＋1 個 decode 實例（8 GPU）；SLA 平均 TPOT ≤ 20 ms（約 50 tok/s）、TTFT ≤ 2.5 s；16K 輸入時並行約 300（8K 約 700）；TPOT 由 40 ms 優化到 17–22 ms", "生產 SLA 調校案例；16,384／1,000", "2026-07-23", "https://vllm.ai/blog/2026-07-23-glm-5.2-nvfp4-b300-pd", "Interested-party", "vLLM 專案（開源；宣傳 vLLM 成熟度）", "否（單一來源）", "已讀（WebFetch 摘要）", "1", "本段相除：並行 300 全部算在 8 顆 decode GPU，上限約每 GPU 37 序列（Derived，上限）"],
 ["Q2", "Perplexity：單機（EP8）在每用戶 40–80 tok/s 區間吞吐急降；EP128 在相同輸出速度下吞吐約為單機 5 倍；批次 1 時輸出速度 >100 tok/s", "受控測試；H100／H200；DeepSeek-V3/R1；MTP 2", "未寫（約 2025）", "https://hub-prod.perplexity.ai/hub/blog/lower-latency-and-higher-throughput-with-multi-node-deepseek-deployment", "Interested-party", "Perplexity（自營推論；宣傳自家部署）", "否", "已讀（WebFetch 摘要）", "1", "只給方向；絕對值在圖中"],
 ["Q2", "Huawei CloudMatrix384：DeepSeek-R1 decode 1,943 tok/s/NPU（TPOT < 50 ms）；TPOT 15 ms 時 538；prefill 6,688", "Ascend 910C；INT8", "2025-06", "https://hub.baai.ac.cn/paper/306006c3-51bd-4253-b45e-230d3ed2dc82", "Interested-party", "Huawei（硬體賣方）", "否", "已讀（摘要頁；arXiv 原文讀取被擋）", "3", "互動速度 20→67 tok/s，吞吐降約 3.6 倍；非 NVIDIA，只作 Analogy"],
 ["Q2", "OpenAI 或 Azure 的每 GPU 並行序列數、批次大小：已搜尋，未找到", "—", "—", "WebSearch；OpenAI、Microsoft、AA、OpenRouter 頁面", "—", "—", "—", "已搜尋，未找到", "—", "只能由每用戶速度間接推論"],
 # Q3
 ["Q3", "vLLM wide-EP：DeepSeek 類 MoE 持續吞吐約 2,200 tok/s/H200（先前約 1,500）", "未寫 decode 或總量、批次、ITL、ISL／OSL、GPU 數；CoreWeave H200 叢集", "2025-12-17", "https://vllm.ai/blog/2025-12-17-large-scale-serving", "Interested-party", "vLLM 專案、CoreWeave（硬體出租方）", "否", "已讀（WebFetch 摘要）", "1", "條件不足，只作量級，不作校準或驗證"],
 ["Q3", "vLLM WideEP on GB200：decode 10.1K TPGS、prefill 26.2K TPGS；2K／2K；專家 FP4；decode 1 個 8 GPU 實例、批次填滿 HBM", "每 GPU decode tok/s（是否只算輸出未寫）；無 ITL", "2026-02-03", "https://vllm.ai/blog/dsr1-gb200-part1", "Interested-party", "vLLM 專案（宣傳）", "與 LMSYS GB200 13,386（NVFP4）互相佐證，但非同一數字", "已讀（WebFetch 摘要）", "1", "GB200 大批次候選點（無 SLO）；與 LMSYS 不同來源"],
 ["Q3", "MLPerf v6.0 DeepSeek-R1 server：GB200 NVL72——NVIDIA 72 GPU 336,106（4,668／GPU）；CoreWeave 64 GPU 311,531（4,868／GPU）；Google 72 GPU 156,781（2,178／GPU）", "系統輸出 tok/s ÷ GPU 數（本段相除）；server TPOT p99 80 ms、TTFT 2 s；Closed、available", "2026-04-01", "https://raw.githubusercontent.com/mlcommons/inference_results_v6.0/main/summary_results.json", "Interested-party", "提交者（NVIDIA 賣方；CoreWeave、Google 雲端）", "是（CoreWeave 部落格 4,868；NVIDIA 與 CoreWeave 為不同提交者）", "已讀（原始 JSON 下載，見 MLPerf_v6.0 頁）", "1", "GB200 大批次且有明確 SLO 的點；建議作 V2 新上限"],
 ["Q3", "CoreWeave 部落格：GB200 NVL72（64 GPU）v6.0 server 4,868／GPU、offline 7,323／GPU", "同上", "2026-04-01", "https://coreweave.com/blog/coreweaves-innovation-velocity-drives-mlperf-6-0-leadership", "Interested-party", "CoreWeave（雲端出租方；宣傳）", "是（MLCommons 原始資料）", "已讀（WebFetch 摘要）", "1", ""],
 ["Q4", "MLPerf v6.0 DeepSeek-R1 interactive：GB200 NVL72（NVIDIA，72 GPU）240,318（3,338／GPU）", "interactive TPOT p99 15 ms；Closed、available", "2026-04-01", "https://raw.githubusercontent.com/mlcommons/inference_results_v6.0/main/summary_results.json", "Interested-party", "NVIDIA（賣方）", "是（工作簿 SRC_PERF_025 同值）", "已讀（原始 JSON）", "1", "＝工作簿 V5；server 類別對應值見上上列"],
 ["Q4", "MLPerf v6.0 DeepSeek-R1 server：GB300 NVL72——Nebius 8 GPU 64,510（8,064／GPU）、Nebius 72 GPU 575,580（7,994）、NVIDIA 72 GPU 400,437（5,562）、NVIDIA 288 GPU 1,555,110（5,400）、CoreWeave 8 GPU 44,590（5,574）", "同上", "2026-04-01", "https://raw.githubusercontent.com/mlcommons/inference_results_v6.0/main/summary_results.json", "Interested-party", "提交者", "是（多提交者）", "已讀（原始 JSON）", "1", "工作簿 V3 的 8,064 是 Nebius 8 GPU 單機；72 GPU 機櫃為 5,562–7,994"],
 ["Q4", "MLPerf v6.0 的 Hopper（H100／H200）DeepSeek-R1 提交：原始 JSON 中沒有", "—", "2026-04-01", "https://raw.githubusercontent.com/mlcommons/inference_results_v6.0/main/summary_results.json", "Verified", "—", "—", "已讀（原始 JSON，確認不存在）", "1", "Hopper 上限仍只有 v5.1"],
 ["Q4", "MLPerf v6.1 GB200 NVL72 DeepSeek-R1 結果：已搜尋，未找到", "—", "—", "GitHub mlcommons/inference_results_v6.1（404）、Nebius v6.1 部落格（只有 GB300、VR200）、CoreWeave v6.1 新聞稿、WebSearch", "—", "—", "—", "已搜尋，未找到", "—", ""],
 ["Q3", "SemiAnalysis InferenceX（H200 V4-Pro @75 tok/s/user 4,036 總 tok/s/GPU 等）", "工作簿 SRC_PERF_016、017", "2026", "（工作簿既有）", "Interested-party", "SemiAnalysis（研究銷售）", "否", "未重讀", "—", "依紀律不單獨使用；只作 Hopper 小批次的旁證"],
]
for r in rows:
    ws.append(r)
for i, w in enumerate([6, 60, 40, 18, 50, 14, 36, 30, 20, 8, 44], 1):
    ws.column_dimensions[L(i)].width = w
for c in ws[1]:
    c.font = Font(bold=True)
ws.append([])
ws.append(["說明", "來源等級：1＝原始一手（官方頁、原始資料檔、論文）；2＝獨立第三方量測；3＝二手轉述。標記依 CLAUDE.md 第 7 節。"])

# ---------- 情境計算 ----------
cs = wb.create_sheet("情境計算")
cs.append(["服務條件情境：每用戶速度與 η_d 倍數對每 GW 產出（IF_TokGW）與 L1_MFleet 的粗估（Derived；不是建議值）"])
cs.append(["公式：B ＝ MIN((a×1000 ÷ S − 固定延遲) ÷ (每序列時間 ÷ k), HBM 上限)；decode tok/s/GPU ＝ B × a × 1000 ÷ (固定延遲 ＋ B × 每序列時間 ÷ k)；每請求 GPU 秒 ＝ prefill 部分（不變）＋ OSL ÷ decode；IF_TokGW 比值 ＝ 舊 GPU 秒 ÷ 新 GPU 秒。藍字為 v10.0 Perf 現值。"])
head = ["情境", "世代", "層級", "a（每步接受 token）", "S 基準（tok/s）", "ISL", "OSL", "固定延遲（ms）", "每序列時間（ms）", "HBM 批次上限", "每請求 GPU 秒（基準）", "prefill 占比", "S 新", "k（每序列時間縮小倍數）", "B 新", "decode tok/s/GPU 新", "每請求 GPU 秒 新", "IF_TokGW 新 ÷ 舊", "核對：基準 decode"]
cs.append(head)
cols = ["C", "D", "E", "F", "G", "H", "I", "J", "K"]
gen_short = {"Hopper H100": "Hopper", "GB200 NVL72": "GB200", "GB300 NVL72": "GB300"}
scen = [
 ("S0 基準（核對）", {1: None, 2: None, 3: None}, {"Hopper": 1, "GB200": 1, "GB300": 1}),
 ("A1 觀測中位速度（Luna 90／Sol 55／Astra 40）", {1: 90, 2: 55, 3: 40}, {"Hopper": 1, "GB200": 1, "GB300": 1}),
 ("A2 ChatGPT／Codex 舊預設（Luna 60／Sol 30／Astra 30）", {1: 60, 2: 30, 3: 30}, {"Hopper": 1, "GB200": 1, "GB300": 1}),
 ("A3 A1 速度＋世代換算修正（k：Hopper 3.6、GB200 1.4、GB300 1）", {1: 90, 2: 55, 3: 40}, {"Hopper": 3.6, "GB200": 1.4, "GB300": 1}),
 ("B1 大批次速度 20 tok/s，η_d 不變", {1: 20, 2: 20, 3: 20}, {"Hopper": 1, "GB200": 1, "GB300": 1}),
 ("B2 大批次 20 tok/s，η_d 隨批次（k：Hopper 3.6、Blackwell 1.4）", {1: 20, 2: 20, 3: 20}, {"Hopper": 3.6, "GB200": 1.4, "GB300": 1.4}),
 ("B3 大批次 20 tok/s，η_d 隨批次（k：Hopper 4.3、Blackwell 2.7）", {1: 20, 2: 20, 3: 20}, {"Hopper": 4.3, "GB200": 2.7, "GB300": 2.7}),
]
blue = Font(color="0000FF")
first_rows = {}
for sname, speeds, ks in scen:
    first_rows[sname] = cs.max_row + 1
    for ti, col in enumerate(cols):
        g = gen_short[perf[f"{col}4"].value]
        tier = perf[f"{col}5"].value
        tidx = perf[f"{col}7"].value
        r = cs.max_row + 1
        vals = [sname, g, tier, perf[f"{col}49"].value, perf[f"{col}41"].value, perf[f"{col}43"].value, perf[f"{col}44"].value,
                perf[f"{col}79"].value, perf[f"{col}83"].value, perf[f"{col}86"].value, perf[f"{col}93"].value, perf[f"{col}94"].value]
        snew = speeds[tidx] if speeds[tidx] is not None else f"=E{r}"
        vals += [snew, ks[g],
                 f"=MIN((D{r}*1000/M{r}-H{r})/(I{r}/N{r}),J{r})",
                 f"=O{r}*D{r}*1000/(H{r}+O{r}*I{r}/N{r})",
                 f"=K{r}*L{r}+G{r}/P{r}",
                 f"=K{r}/Q{r}",
                 f"=G{r}/(K{r}*(1-L{r}))"]
        cs.append(vals)
        for cc in range(4, 13):
            cs.cell(r, cc).font = blue
# 彙總
cs.append([])
cs.append(["彙總：機隊加權（層級 token 權重 Luna 0.4／Sol 0.45／Astra 0.15，沿用 C90 付費組合；2025 服務世代 Hopper 0.6／GB200 0.4，Alloc_In C9:C10）。L1_MFleet 新值＝現值 ÷ 機隊 token 比值（假設機隊付費 token 與 IF_TokGW 成正比）"])
cs.append(["情境", "Hopper 加權比值", "GB200 加權比值", "GB300 加權比值", "2025 機隊比值", "L1_MFleet 現值", "L1_MFleet 新值（粗估）"])
w = [0.4, 0.45, 0.15]
for sname, _, _ in scen:
    f0 = first_rows[sname]
    r = cs.max_row + 1
    def gen_expr(off):
        # token 加權的每 token GPU 秒：基準 ÷ 新
        base = "+".join(f"{w[i]}*K{f0+off+i}/(F{f0+off+i}+G{f0+off+i})" for i in range(3))
        new = "+".join(f"{w[i]}*Q{f0+off+i}/(F{f0+off+i}+G{f0+off+i})" for i in range(3))
        return f"=({base})/({new})"
    cs.append([sname, gen_expr(0), gen_expr(3), gen_expr(6), f"=0.6*B{r}+0.4*C{r}" if False else None, mfleet, None])
    # 機隊：tokens ∝ Σ GW_g ÷ c_g；比值 ＝ Σ GW_g/c_g,new ÷ Σ GW_g/c_g,base。用 c_g,base 正規化：
    cbH = "+".join(f"{w[i]}*K{f0+i}/(F{f0+i}+G{f0+i})" for i in range(3))
    cbG = "+".join(f"{w[i]}*K{f0+3+i}/(F{f0+3+i}+G{f0+3+i})" for i in range(3))
    cs.cell(r, 5).value = f"=(0.6/({cbH})*B{r}+0.4/({cbG})*C{r})/(0.6/({cbH})+0.4/({cbG}))"
    cs.cell(r, 6).font = blue
    cs.cell(r, 7).value = f"=F{r}/E{r}"
for i, wd in enumerate([44, 9, 12, 10, 10, 9, 9, 10, 10, 10, 10, 9, 8, 10, 10, 12, 12, 12, 12], 1):
    cs.column_dimensions[L(i)].width = wd

# ---------- MLPerf_v6.0 ----------
ms = wb.create_sheet("MLPerf_v6.0")
ms.append(["MLCommons inference_results_v6.0 summary_results.json 中 deepseek-r1 全部列（2026-10-16 下載）；per_GPU＝Performance_Result ÷ Total Accelerators（本段相除）"])
with open(EV + "mlperf_v6.0_deepseek-r1.csv", newline="") as f:
    for i, row in enumerate(csv.reader(f)):
        if i > 0:
            row = [float(x) if x.replace(".", "", 1).isdigit() else x for x in row]
        ms.append(row)
wb.save(EV + "證據表.xlsx")
print("ok", len(rows))
