#!/usr/bin/env python3
"""v20.1：逐名稱逐格新舊值與倍數（供下游同步通知引用）。
用法：python3 tools/gen_v201_changes_csv.py <舊版 xlsx> <新版 xlsx> <輸出 csv>
只列值有變動的契約前綴名稱（IF_／IFW_／IFC_／L1_／GOV_／CHK_／CAP_），逐格一列；版本字串 RNG_Version／RNG_Current 另列於最後。"""
import csv
import sys

import openpyxl
from openpyxl.utils import get_column_letter as L, range_boundaries

PREFIX = ("IF_", "IFW_", "IFC_", "L1_", "GOV_", "CHK_", "CAP_")


def cells(wb, name):
    sh, rg = wb.defined_names[name].attr_text.rsplit("!", 1)
    c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
    ws = wb[sh.strip("'")]
    return [(f"{L(c)}{r}", ws.cell(r, c).value) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= 1e-9 * max(abs(a), abs(b), 1e-300) + 1e-12
    return a == b


def main(old, new, out):
    a, b = openpyxl.load_workbook(old, data_only=True), openpyxl.load_workbook(new, data_only=True)
    rows = []
    for name in sorted(b.defined_names):
        if not (name.startswith(PREFIX) or name in ("RNG_Version", "RNG_Current")) or "!" not in b.defined_names[name].attr_text or "," in b.defined_names[name].attr_text:
            continue
        if name not in a.defined_names:
            rows.append((name, "（新名稱）", "", "", ""))
            continue
        co, cn = cells(a, name), cells(b, name)
        for (ref, vo), (_, vn) in zip(co, cn):
            if not same(vo, vn):
                mult = vn / vo if isinstance(vo, (int, float)) and isinstance(vn, (int, float)) and vo else ""
                rows.append((name, ref, vo, vn, mult))
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["具名範圍", "儲存格", "舊值", "新值", "倍數（新÷舊）"])
        w.writerows(rows)
    print(len(rows), "rows;", len({r[0] for r in rows}), "names")


if __name__ == "__main__":
    main(*sys.argv[1:4])
