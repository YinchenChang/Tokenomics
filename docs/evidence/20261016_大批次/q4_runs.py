"""大批次低估查證 Q4：在 scratch 複本上量化假設 (a)（通訊重疊）與 (b)（η_d 隨批次上升）。

唯讀：只讀 repo 的 v10.2；所有改動寫在 scratch 複本，以 LibreOffice headless 重算後讀回。
用法：python3 q4_runs.py <v10.2.xlsx> <scratch 目錄> <輸出 json> [情境名,情境名...]

(b) 的做法：η_d 依「每 GPU 每步處理的 token 數 T＝B ×（草稿數＋1）」分段：
  T ≤ T1：k＝1（維持 GB300 小批次擬合值）；T1 < T < T_cal：對數線性上升；T ≥ T_cal：k＝k_cal。
  B 與 η 互相依賴（η 變大 → SLO 下 B 變大），所以先在 Python 以工作簿同一條公式求固定點，
  再把各欄 η_d 以常數寫入 scratch（Perf 第 53 列、Calib 第 140、184 列），LibreOffice 重算後
  檢查工作簿算出的 B* 是否等於固定點（verify 欄）。
"""
import json
import math
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

SRC = Path(sys.argv[1])
WORK = Path(sys.argv[2])
OUT = Path(sys.argv[3])
ONLY = sys.argv[4].split(",") if len(sys.argv) > 4 else None
WORK.mkdir(parents=True, exist_ok=True)

BV = openpyxl.load_workbook(SRC, data_only=True)  # v10.2 快取值（獨立查核已確認與重算一致）


def v(sh, cell):
    return BV[sh][cell].value


SERV_C12 = v("Serving", "C12")  # EP 啟用值 bytes/元素
SERV_C13 = v("Serving", "C13")  # KV 讀取效率
SERV_C14 = v("Serving", "C14")  # EP 通訊效率
KVB = v("Calib", "C108")
LAYERS = v("Calib", "C102")
TOPK = v("Calib", "C105")
DMOD = v("Calib", "C103")

PERF_COLS = {"C": 1, "D": 1, "E": 1, "F": 2, "G": 2, "H": 2}  # Hopper×3 層級、GB200×3 層級
ALL_PERF = "CDEFGHIJKLMNOPQ"
J_COLS = "CDEFG"
H_COLS = "CDEFGH"
SOL_IF = {"Hopper": "D", "GB200": "G", "GB300": "J", "VR200": "M", "RU": "P"}


def solve_B(F, c, kv, ep, a, S, cap, kfun, d, fixedB=None):
    """工作簿同一條式：每步＝F＋B×max(c/k, kv, ep)；SLO 下 B＝(1000a/S−F)/每序列；再受 cap 限制。"""
    if fixedB:
        B = min(fixedB, cap)
        return B, kfun(B * (d + 1))
    B = 1.0 if S else cap
    for _ in range(400):
        k = kfun(B * (d + 1))
        per = max(c / k, kv, ep)
        Bs = (1000 * a / S - F) / per if S else cap
        Bn = max(0.0, min(Bs, cap))
        if abs(Bn - B) < 1e-9:
            break
        B = 0.5 * B + 0.5 * Bn
    return B, kfun(B * (d + 1))


def make_k(T1, Tcal, kcal):
    def k(T):
        if kcal is None or T <= T1:
            return 1.0
        if T >= Tcal:
            return kcal
        return math.exp(math.log(kcal) * math.log(T / T1) / math.log(Tcal / T1))
    return k


def kcal_for(target, B, a, F, c, kv, ep):
    """校準點：給定批次，求讓模型＝目標所需的 k；若所需每序列時間低於非算力項下限，取交叉點並標記。"""
    per = (1000 * B * a / target - F) / B
    floor = max(kv, ep)
    if per <= floor:
        return c / floor, False
    return c / per, True


def j_inputs(col, f_ep=0.0, f_lat=0.0):
    d = v("Calib", f"{col}182")
    F = v("Calib", f"{col}186") - (LAYERS + d) * v("Calib", f"{col}185") / 1000 * f_lat
    S = v("Calib", f"{col}164") or 0
    fixedB = v("Calib", f"{col}165") or None
    return dict(F=F, c=v("Calib", f"{col}187"), kv=v("Calib", f"{col}188"),
                ep=v("Calib", f"{col}189") * (1 - f_ep), a=v("Calib", f"{col}183"), S=S,
                cap=v("Calib", f"{col}191"), d=d, fixedB=fixedB, eta=v("Calib", f"{col}184"),
                gen=v("Calib", f"{col}160"), target=v("Calib", f"{col}167"))


def h_inputs(col, f_ep=0.0, S_override=None):
    d = v("Calib", f"{col}138")
    eta = v("Calib", f"{col}140")
    peak = v("Calib", f"{col}124")
    c = (d + 1) * v("Calib", f"{col}136") / (peak * eta) / 1000
    kv = KVB * v("Calib", f"{col}135") / (v("Calib", f"{col}126") * SERV_C13) / 1e9
    epw = v("Calib", f"{col}133")
    ep = (d + 1) * LAYERS * TOPK * DMOD * SERV_C12 * (epw - 1) / epw / (v("Calib", f"{col}127") * SERV_C14) / 1e9
    cap = (v("Calib", f"{col}131") - v("Calib", f"{col}134")) * 1e9 / (KVB * v("Calib", f"{col}135"))
    S = S_override if S_override is not None else v("Calib", f"{col}118")
    return dict(F=v("Calib", f"{col}144"), c=c, kv=kv, ep=ep * (1 - f_ep), a=v("Calib", f"{col}139"),
                S=S, cap=cap, d=d, fixedB=None, eta=eta, gen=v("Calib", f"{col}117"))


def perf_inputs(col, f_ep=0.0):
    return dict(F=v("Perf", f"{col}79"), c=v("Perf", f"{col}80"), kv=v("Perf", f"{col}81"),
                ep=v("Perf", f"{col}82") * (1 - f_ep), a=v("Perf", f"{col}49"), S=v("Perf", f"{col}41"),
                cap=v("Perf", f"{col}86"), d=v("Perf", f"{col}47"), fixedB=None, eta=v("Perf", f"{col}53"),
                gen=v("Perf", f"{col}6"))


def b_overrides(shape, f_ep=0.0, f_latJ=0.0, univ=False, offline=False):
    """依 shape 求 Hopper、GB200 的 k_cal（用 Calib J 的 C、D 校準點），再對各欄求固定點並回傳 η 覆寫。"""
    T1 = shape["T1"]
    info = {"kcal": {}, "fixedpoint": {}}
    jc = j_inputs("C", f_ep, f_latJ)
    jd = j_inputs("D", f_ep, f_latJ)
    kH, okH = kcal_for(jc["target"], jc["fixedB"], jc["a"], jc["F"], jc["c"], jc["kv"], jc["ep"])
    kG, okG = kcal_for(jd["target"], jd["cap"], jd["a"], jd["F"], jd["c"], jd["kv"], jd["ep"])
    info["kcal"] = {"Hopper": [kH, okH, kH * jc["eta"]], "GB200": [kG, okG, kG * jd["eta"]]}
    kf = {1: make_k(T1, 128 * (jc["d"] + 1), kH), 2: make_k(T1, jd["cap"] * (jd["d"] + 1), kG)}
    if univ:  # 檢查用：GB300 以後也套 GB200 的形狀
        for g in (3, 4, 5):
            kf[g] = kf[2]
    edits = {}
    for col in J_COLS:
        x = j_inputs(col, f_ep, f_latJ)
        B, k = solve_B(x["F"], x["c"], x["kv"], x["ep"], x["a"], x["S"], x["cap"], kf.get(x["gen"], make_k(1, 2, None)), x["d"], x["fixedB"])
        edits[f"Calib!{col}184"] = x["eta"] * k
        info["fixedpoint"][f"J{col}"] = [B, k]
    for col in H_COLS:
        x = h_inputs(col, f_ep, 0.5 if (offline and col in "FH") else None)
        if x["gen"] not in kf:
            continue
        B, k = solve_B(x["F"], x["c"], x["kv"], x["ep"], x["a"], x["S"], x["cap"], kf[x["gen"]], x["d"])
        edits[f"Calib!{col}140"] = x["eta"] * k
        info["fixedpoint"][f"H{col}"] = [B, k]
    for col in PERF_COLS:
        x = perf_inputs(col, f_ep)
        B, k = solve_B(x["F"], x["c"], x["kv"], x["ep"], x["a"], x["S"], x["cap"], kf[x["gen"]], x["d"])
        edits[f"Perf!{col}53"] = x["eta"] * k
        info["fixedpoint"][f"P{col}"] = [B, k]
    return edits, info


def a2J_edits(f):
    return {f"Calib!{c}185": v("Calib", f"{c}185") * (1 - f) for c in J_COLS}


A1 = lambda f: {"Serving!C14": SERV_C14 / (1 - f)}  # EP 通訊項 ×(1−f)（所有 EP 通訊項同時縮小）
OFF = {"Calib!F118": 0.5, "Calib!H118": 0.5}  # MLPerf offline：取消速度下限，批次只受 HBM 限制
TRT = {"Calib!C110": 1000, "Calib!C111": 2000, "Calib!G118": 11.35, "Calib!G125": 141, "Calib!G126": 4.8, "Calib!G133": 8}
SHAPE_EARLY = {"T1": 76}    # 擬合範圍上緣（GB300 擬合點 B 19 × 4 token/步）之後開始上升
SHAPE_LATE = {"T1": 410}    # MLPerf server（GB300 B≈410，模型 0.90）之後才上升


def scen():
    S = {}
    S["S0_基準"] = ({}, None)
    S["OFF_MLPerf離線"] = (OFF, None)
    S["TRT_H200_8GPU"] = (TRT, None)
    S["A1_EP遮50%"] = (A1(0.5), None)
    S["A1_EP遮65%"] = (A1(0.65), None)
    S["A2J_固定延遲遮50%_只校準驗證欄"] = (a2J_edits(0.5), None)
    S["A2all_固定延遲遮50%_含服務"] = ({"Calib!C68": v("Calib", "C68") * 0.5, "Calib!D68": v("Calib", "D68") * 0.5}, None)
    S["B1_早升"] = ({}, dict(shape=SHAPE_EARLY))
    S["B2_晚升"] = ({}, dict(shape=SHAPE_LATE))
    S["B1_早升+A1_50%"] = (A1(0.5), dict(shape=SHAPE_EARLY, f_ep=0.5))
    S["B1_早升+A2J_50%"] = (a2J_edits(0.5), dict(shape=SHAPE_EARLY, f_latJ=0.5))
    S["B1_早升_離線"] = (OFF, dict(shape=SHAPE_EARLY, offline=True))
    S["BU_早升_全世代(檢查)"] = ({}, dict(shape=SHAPE_EARLY, univ=True))
    # 對照：差距若是「世代水準」（不分批次的常數倍數），即 η 校準報告 E1 的假設性倍數
    S["G_const"] = ({"Calib!C67": 4.49, "Calib!D67": 3.04}, None)
    # 建議方向的試算：大批次倍數只放進既有的批次口徑倍數 CAL_BatchEtaD（Calib C156／D156；只作用於 Perf_Batch）
    S["X10_batch"] = ({"Calib!C156": 3.74, "Calib!D156": 3.04}, None)
    return S


def lo_recalc(xlsx, outdir):
    outdir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="lo_profile_"))
    try:
        subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
                        "--convert-to", "xlsx", "--outdir", str(outdir), str(xlsx)],
                       capture_output=True, text=True, timeout=900)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    return outdir / xlsx.name


def read_out(wv):
    r = {"J": {}, "H": {}, "Checks": {}, "Perf": {}, "IF_Sol": {}, "IF_Luna": {}, "IF_Astra": {}}
    for c in J_COLS:
        r["J"][c] = {k: wv["Calib"][f"{c}{row}"].value for k, row in (("B", 193), ("tok", 194), ("ratio", 195), ("eta", 184), ("F", 186))}
    for c in H_COLS:
        r["H"][c] = {k: wv["Calib"][f"{c}{row}"].value for k, row in (("B", 146), ("dec", 147), ("tok", 148), ("ratio", 149), ("eta", 140))}
    for row in range(148, 163):
        r["Checks"][wv["Checks"][f"A{row}"].value] = wv["Checks"][f"E{row}"].value
    r["Checks"]["VF"] = wv["Checks"]["F146"].value
    for c in ALL_PERF:
        r["Perf"][c] = {"B": wv["Perf"][f"{c}87"].value, "dec": wv["Perf"][f"{c}88"].value, "eta": wv["Perf"][f"{c}53"].value}
    for g, c in SOL_IF.items():
        r["IF_Sol"][g] = wv["Interface"][f"{c}31"].value / 1e9
        r["IF_Luna"][g] = wv["Interface"][f"{c}21"].value / 1e9
        r["IF_Astra"][g] = wv["Interface"][f"{c}41"].value / 1e9
    r["L1_MFleet"] = wv["L1"]["D78"].value
    r["IF_AllocRDGW"] = wv["Interface"]["C199"].value
    return r


results = json.loads(OUT.read_text()) if OUT.exists() else {}
for name, (static, bspec) in scen().items():
    if ONLY and name not in ONLY:
        continue
    edits = dict(static)
    info = None
    if bspec:
        e2, info = b_overrides(bspec["shape"], bspec.get("f_ep", 0.0), bspec.get("f_latJ", 0.0),
                               bspec.get("univ", False), bspec.get("offline", False))
        edits.update(e2)
    wb = openpyxl.load_workbook(SRC)
    for key, val in edits.items():
        sh, coord = key.split("!")
        wb[sh][coord].value = val
    tmp = WORK / f"{name}.xlsx"
    wb.save(tmp)
    rec = lo_recalc(tmp, WORK / "recalc")
    wv = openpyxl.load_workbook(rec, data_only=True)
    out = read_out(wv)
    out["edits"] = {k: (round(x, 6) if isinstance(x, float) else x) for k, x in edits.items()}
    out["b_info"] = info
    if info:  # 驗證：工作簿 B* 與 Python 固定點一致
        ver = {}
        for key, (B, k) in info["fixedpoint"].items():
            sec, col = key[0], key[1]
            wbB = out["J"][col]["B"] if sec == "J" else out["H"][col]["B"] if sec == "H" else out["Perf"][col]["B"]
            ver[key] = [B, wbB]
        out["verify_B"] = ver
    results[name] = out
    print(name, "L1_MFleet", out["L1_MFleet"], "J ratio", {c: out["J"][c]["ratio"] for c in J_COLS}, flush=True)
    OUT.write_text(json.dumps(results, ensure_ascii=False, indent=1, default=str))
