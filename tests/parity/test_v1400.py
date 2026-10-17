"""4-3 B 段（v14.0）專屬測試：刪除、18 個名稱、Checks V23–V26 與 D104／D140 的『必須能失敗』測試、既有名稱不變。
讀 LibreOffice 重算後的值；不刪除、不放寬任何既有測試。

V23、V24 是結構性恆等式（實際速度＝MIN(SLO 批次, 記憶體批次) 所對應的速度；Cap_In 第 147 列與 L1_MImplied 由同一組輸入算出），
無法只靠改輸入使其失敗；它們的必須失敗測試改寫一個公式格（與 test_v1200_v19 改寫 Tech_Registry K 欄同一做法），這一點寫入報告『待 Project 判斷』。
V25、V26 與 D104／D140 只改輸入。"""
import json
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import range_boundaries

from conftest import model_path
from parity_lib import lo_recalc

REPO = Path(__file__).resolve().parents[2]
V131_PREV = REPO / "model" / "archive" / "20261017_Tokenomics_v13.1.xlsx"      # v13.1（LibreOffice 重算存檔）：v14.0 的「既有名稱不得變」對照
CONTRACT_PREFIX = ("IF_", "IFW_", "IFC_", "L1_", "CHK_", "CAP_", "RNG_")
Q4_CHANGED = {"CHK_VerifyTotal", "CHK_VerifyPass"}      # 因新增 V23–V26 而變
VERSION_NAMES = {"RNG_Version", "RNG_Current"}
IFC_RANGES = {"IFC_Conf", "IFC_Use", "IFC_Layer", "IFC_Updated"}      # Interface 新增 M 節列，範圍變長
NEW_IF = ["IF_TaskTokMult", "IF_AmortRevShare"] + [f"IF_{p}_{t}" for p in ("TrainFLOP", "TrainInferShare", "GMDefault") for t in ("Luna", "Sol", "Astra")]
NEW_L1 = {"L1_MAPI_Ser": ("Cap_In!$C$206:$F$206", 4), "L1_MSub_Ser": ("Cap_In!$C$207:$F$207", 4), "L1_MSeat_Ser": ("Cap_In!$C$208:$F$208", 4),
          "L1_MImplied_Ser": ("Cap_In!$C$209:$F$209", 4), "L1_MFleet_Ser": ("Cap_In!$C$223:$F$223", 4),
          "L1_MFleet_P25Lo": ("Cap_In!$C$285:$E$285", 3), "L1_MFleet_P25Hi": ("Cap_In!$C$286:$E$286", 3)}
PREWARM = {
    "test_v1400_v25_fails_when_demand_exceeds_ceiling": [{"Alloc_In!C15": 100000}],
    "test_v1400_v26_fails_when_2026_mix_not_100": [{"Cap_In!D214": 0.10}],
    "test_v1400_d104_d140_fail_when_mix_not_100": [{"Alloc_In!C12": 0.2}, {"Alloc_In!C32": 0.2}],
}


@pytest.fixture(scope="module")
def model():
    return model_path()


def _flat(v):
    if isinstance(v, list):
        return [x for r in v for x in _flat(r)]
    return [v]


def _same(a, b):
    a = None if a == "" else a
    b = None if b == "" else b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= 1e-12 or abs(a - b) <= 1e-9 * max(abs(a), abs(b))
    return a == b


def _chk(v, code):
    for r in range(1, 400):
        if v.get("Checks", f"A{r}") == code:
            return r
    raise AssertionError(code)


def _v_row(wb, code):
    ws = wb["Checks"]
    for r in range(1, 400):
        if ws.cell(r, 1).value == code:
            return {"C": ws.cell(r, 3).value, "D": ws.cell(r, 4).value, "E": ws.cell(r, 5).value, "F": ws.cell(r, 6).value}
    raise AssertionError(code)


def _recalc_with_edits(model, tmp_path, edits):
    """複製活頁簿、改寫指定格（含公式格）、以 LibreOffice 重算，回傳唯讀值活頁簿（同 test_parity._recalc_with_edits）。"""
    wb = openpyxl.load_workbook(model)
    for (sh, ref), val in edits.items():
        wb[sh][ref].value = val
    dst = tmp_path / "edit" / model.name
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return openpyxl.load_workbook(lo_recalc(dst, tmp_path / "out"), data_only=True)


# ── 名稱與標籤 ──
def test_v1400_new_names_exist_with_tier_and_token_type(model, lo):
    """18 個新名稱：位置、欄數、值皆為數字；Interface M 節每列的標籤帶『層級』與 token 類型（Luna／Sol／Astra 不得稱『世代』）。"""
    v = lo.view({})
    wb = openpyxl.load_workbook(model)
    ws = wb["Interface"]
    for n in NEW_IF:
        assert n in v.names, n
        sh, ref = v.names[n].split("!")
        assert sh == "Interface"
        row = int(ref.replace("$", "").split(":")[0][1:])
        label = str(ws.cell(row, 1).value)
        assert f"[{n}]" in label, (n, label)
        vals = [x for x in _flat(v.get_name(n)) if isinstance(x, (int, float))]
        assert len(vals) == (5 if n == "IF_TaskTokMult" else 15), (n, len(vals))
        for t in ("Luna", "Sol", "Astra"):
            if n.endswith("_" + t):
                assert f"{t}層級" in label and "世代" not in label.split("；")[0], (n, label)
        if n.startswith("IF_GMDefault"):
            assert "付費 token" in label, label
        if n.startswith(("IF_TrainFLOP", "IF_TrainInferShare")):
            assert "訓練側，無 token 類型" in label, label
        if n in ("IF_TaskTokMult", "IF_AmortRevShare"):
            assert "混合" in label, label
    for n, (ref, k) in NEW_L1.items():
        assert v.names[n] == ref, (n, v.names[n])
        assert len([x for x in _flat(v.get_name(n)) if isinstance(x, (int, float))]) == k, n
    ci = wb["Cap_In"]
    pc = next(c for c in range(3, ci.max_column + 2) if ci.cell(3, c).value == "用途")
    assert "待定義，見 4-3B 報告" in str(ci.cell(121, pc).value)      # Q3：Cap_In 121 保留為顯示用，用途標記


def test_v1400_existing_contract_names_unchanged(lo):
    """v13.1 既有的契約前綴名稱（IF_／IFW_／L1_／CHK_／CAP_／RNG_）逐格值與 v13.1 相同（容差 1e-9；以名稱比較，不比座標）。
    例外只有：CHK_VerifyTotal／Pass（新增 V23–V26）、IFC_*（Interface 新增列，範圍變長）、RNG_Version／Current（版本字串）。"""
    prev = openpyxl.load_workbook(V131_PREV, data_only=True)
    v = lo.view({})
    keep = [n for n in prev.defined_names if n.startswith(CONTRACT_PREFIX) and n not in Q4_CHANGED | VERSION_NAMES | IFC_RANGES]
    assert len(keep) >= 700, len(keep)
    for n in keep:
        sh, rg = prev.defined_names[n].attr_text.rsplit("!", 1)
        c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
        want = [prev[sh.strip("'")].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        got = _flat(v.get_name(n))
        assert len(want) == len(got) and all(_same(a, b) for a, b in zip(want, got)), n
    assert (v.get_name("CHK_VerifyFail"), v.get_name("CHK_VerifyWarn")) == (7, 1)
    assert (v.get_name("CHK_VerifyTotal"), v.get_name("CHK_VerifyPass")) == (24, 16)
    assert v.get_name("RNG_Version") == model_path().name[: -len(".xlsx")] == v.get_name("RNG_Current")


def test_v1400_deleted_cells_empty_and_names_removed(model, lo):
    """Q1：A 段明細處置＝刪除的 72 列＋Har_In 44 的全部數值格為空（514 格）；各列『用途』欄有刪除註記；B5_TaskTok、B5_SelVendor 已移除。"""
    data = json.loads((REPO / "builder" / "v1400_data.json").read_text(encoding="utf-8"))
    assert len(data["delete"]) == 73
    wb = openpyxl.load_workbook(model)
    n = 0
    for e in data["delete"]:
        ws = wb[e["sheet"]]
        c1, r1, c2, r2 = range_boundaries(e["range"])
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                assert ws.cell(r, c).value is None, (e["sheet"], ws.cell(r, c).coordinate)
                n += 1
        pc = next(c for c in range(3, ws.max_column + 2) if ws.cell(3, c).value == "用途")
        assert str(ws.cell(e["row"], pc).value).startswith("已刪除（v14.0"), (e["sheet"], e["row"])
    assert n == 514
    v = lo.view({})
    assert "B5_TaskTok" not in v.names and "B5_SelVendor" not in v.names
    assert "供應商專屬" not in str(wb["Harness"]["A157"].value)


# ── Checks V23–V26（每條都能失敗） ──
def test_v1400_new_checks_pass_on_base(lo):
    v = lo.view({})
    for code in ("V23", "V24", "V25", "V26"):
        r = _chk(v, code)
        assert v.get("Checks", f"G{r}") == "驗證" and v.get("Checks", f"F{r}") == 1, (code, v.get("Checks", f"F{r}"))
    r = _chk(v, "V23")
    assert v.get("Checks", f"D{r}") == 32 and v.get("Checks", f"C{r}") == 0      # 可達欄 32（Perf 15＋Perf_Batch 17）、低於 SLO 的欄數 0
    assert abs(v.get("Checks", f"E{_chk(v, 'V25')}") - 0.245362088692091) < 1e-9
    assert v.get("Checks", f"E{_chk(v, 'V24')}") == 0 and v.get("Checks", f"E{_chk(v, 'V26')}") == 0


@pytest.mark.parametrize("sheet", ["Perf", "Perf_Batch"])
def test_v1400_v23_fails_when_actual_speed_below_slo(model, tmp_path, sheet):
    """V23 必須能失敗：把 C87（採用批次 B*）放大 1000 倍，實際速度遠低於 SLO。無法只靠改輸入（B* 的定義使實際速度恆 ≥ SLO），故改寫公式格。"""
    old = openpyxl.load_workbook(model)[sheet]["C87"].value
    assert old == "=MAX(0,MIN(C85,C86))", old
    wb = _recalc_with_edits(model, tmp_path, {(sheet, "C87"): "=MAX(0,MIN(C85,C86))*1000"})
    r = _v_row(wb, "V23")
    assert r["C"] >= 1 and r["E"] >= 1 and r["F"] == 0, r
    d = wb.defined_names["CHK_VerifyFail"]
    sh, ref = d.attr_text.rsplit("!", 1)
    assert wb[sh.strip("'")][ref.replace("$", "")].value == 8


def test_v1400_v23_not_vacuous_when_some_columns_unreachable(lo):
    """V23 不因『不可達欄』而空轉或誤判：Train_In!E31＝100（Perf_Batch 部分欄 SLO 不可達）時，可達欄數少於 32 但仍 > 0，V23 仍通過。"""
    v = lo.view({"Train_In!E31": 100})
    r = _chk(v, "V23")
    assert 0 < v.get("Checks", f"D{r}") < 32 and v.get("Checks", f"F{r}") == 1, (v.get("Checks", f"D{r}"), v.get("Checks", f"F{r}"))


def test_v1400_v24_fails_when_cap_in_147_differs_from_l1_mimplied(model, tmp_path):
    """V24 必須能失敗：Cap_In!C147 的公式乘 1.01（模擬分子或分母漏項）。無法只靠改輸入（兩條路徑由同一組輸入算出），故改寫公式格。"""
    old = openpyxl.load_workbook(model)["Cap_In"]["C147"].value
    assert old == "=(C135+C136+C137)*1000000000/((C140+C141+C142)*CAP_Price25Mix)", old
    wb = _recalc_with_edits(model, tmp_path, {("Cap_In", "C147"): old + "*1.01"})
    r = _v_row(wb, "V24")
    assert r["E"] > 1e-9 and r["F"] == 0, r
    assert abs(r["E"] - 0.01 * 2.03352783356479) < 1e-6      # 差 1% × 2025 年度 m_整體 2.0335


def test_v1400_v25_fails_when_demand_exceeds_ceiling(lo):
    """V25 必須能失敗：每則提示 token 數（Alloc_In!C15）由 4,000 調到 100,000，需求路線超過上限路線（比值 > 1）。"""
    v = lo.view({"Alloc_In!C15": 100000})
    r = _chk(v, "V25")
    assert v.get("Checks", f"E{r}") > 1 and v.get("Checks", f"F{r}") == 0, (v.get("Checks", f"E{r}"), v.get("Checks", f"F{r}"))
    assert v.get_name("CHK_VerifyFail") >= 8


def test_v1400_v26_fails_when_2026_mix_not_100(lo):
    """V26 必須能失敗：Cap_In!D214（2026 Q1 欄 VR200 組合）由 0.05 改 0.10，該欄合計 105%。"""
    v = lo.view({"Cap_In!D214": 0.10})
    r = _chk(v, "V26")
    assert abs(v.get("Checks", f"E{r}") - 0.05) < 1e-9 and v.get("Checks", f"F{r}") == 0
    assert v.get_name("GOV_Errors") == 0      # 驗證失敗不進 GOV_Errors（與 H1 並列，H1 讀 2025 欄）


def test_v1400_d104_d140_read_alloc_in_and_fail_when_mix_not_100(model, lo):
    """Q4：Checks D104、D140 改讀 Alloc_In C18／C33（單一定義）；服務組合 Alloc_In!C12＝0.2（合計 120%）使 H1 觸發，研發組合 C32＝0.2 使 Z2 觸發。"""
    ws = openpyxl.load_workbook(model)["Checks"]
    assert "Alloc_In!$C$18" in ws["D104"].value and "Alloc_In!$C$33" in ws["D140"].value
    base = lo.view({})
    assert base.get("Checks", f"D{_chk(base, 'H1')}") == 0 and base.get("Checks", f"D{_chk(base, 'Z2')}") == 0 and base.get_name("GOV_Errors") == 0
    a = lo.view({"Alloc_In!C12": 0.2})
    assert a.get("Checks", f"D{_chk(a, 'H1')}") == 1 and a.get("Checks", f"D{_chk(a, 'Z2')}") == 0 and a.get_name("GOV_Errors") >= 1
    b = lo.view({"Alloc_In!C32": 0.2})
    assert b.get("Checks", f"D{_chk(b, 'Z2')}") == 1 and b.get("Checks", f"D{_chk(b, 'H1')}") == 0 and b.get_name("GOV_Errors") >= 1


# ── Q3：只查明，不改數值 ──
def test_v1400_q3_floor_and_mfleet_lo_values_unchanged(lo):
    """Q3：Cap_In 第 121 列（m 下限）與 L1_MFleet 區間的數值本版不變（0.301／0.215／0.637；0.407／0.184／1.179）。"""
    v = lo.view({})
    want = {"C": 0.300967115929835, "D": 0.215398033949784, "E": 0.637342127851415}
    for c, w in want.items():
        assert abs(v.get("Cap_In", f"{c}121") - w) < 1e-12, c
    assert abs(v.get_name("L1_MFleet") - 0.407295434678084) < 1e-12 and abs(v.get_name("L1_MFleet_Lo") - 0.1838665514316) < 1e-12
    assert abs(v.get_name("L1_MFleet_Hi") - 1.17876090506834) < 1e-12
    assert "L1_MFleetFloor" not in v.names      # 不升級
