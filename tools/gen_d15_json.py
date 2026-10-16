"""由 D15 明細（docs/reports/20261017_D15_來源建議_明細.xlsx）產生 builder/v1310_d15.json（D15 B 段 P1 的資料檔）。

用法：python3 tools/gen_d15_json.py            寫出 builder/v1310_d15.json
      python3 tools/gen_d15_json.py --check    重新推導並與現有檔比對（不寫檔）

內容：
  gov      312 列（第 4 頁「無來源無決策_全部」）：Gov_Map ID、頁、格、決策編號（主＋其餘）、建議來源或決策原文、建議標記
  evidence 106 列（第 2 頁「證據積壓」）：ID、建議、理由、是否涉及高敏感度
  decisions 第 5 頁 D15-01…D15-09 原文
決策編號只收「已存在於 Decisions（v13.0）或本版新增（D15-01…D15-12）」的 ID；括號內文字不當作決策編號。
Gov_Map 的 J 欄只放一個決策 ID（AA 檢查以 COUNTIF 比對整格），所以多個編號時，第一個為主、其餘寫入 N 欄。"""
import json
import re
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[1]
XLSX = REPO / "docs" / "reports" / "20261017_D15_來源建議_明細.xlsx"
MODEL = REPO / "model" / "archive" / "20261017_Tokenomics_v13.0.xlsx"      # Decisions IDs as of v13.0 (the base of the D15 B build)
OUT = REPO / "builder" / "v1310_d15.json"
NEW_DEC = [f"D15-{i:02d}" for i in range(1, 13)]
PAREN = re.compile(r"（[^（）]*）")


def known_ids():
    wb = openpyxl.load_workbook(MODEL, read_only=True)
    ws = wb["Decisions"]
    ids = {str(r[0]) for r in ws.iter_rows(min_row=5, max_col=1, values_only=True) if r[0]}
    return ids | set(NEW_DEC)


def parse_decisions(text, known):
    i = text.find("決策：")
    if i < 0:
        return []
    seg = text[i + 3:]
    seg = re.split(r"[；。]", seg, maxsplit=1)[0]
    while PAREN.search(seg):
        seg = PAREN.sub("", seg)
    out = []
    for t in re.findall(r"[A-Za-z0-9][A-Za-z0-9\-]*", seg):
        if t in known and t not in out:
            out.append(t)
    return out


def build():
    known = known_ids()
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    gov = []
    ws = wb["4_無來源無決策_全部"]
    rows = list(ws.iter_rows(values_only=True))[2:]
    for r in rows:
        page, cell, gm, label, cur, in_m, sens, action, src, mark, rng, conf, andy, why, cls = r[:15]
        ids = parse_decisions(str(src), known)
        gov.append({"gm": gm, "page": page, "cell": cell, "label": label, "action": action, "src_text": src, "mark": mark,
                    "ids": ids, "andy": andy, "cls": cls})
    ev = []
    for r in list(wb["2_證據積壓"].iter_rows(values_only=True))[2:]:
        ev.append({"id": r[0], "cur": r[4], "sens": r[8], "rec": r[9], "why": r[10], "andy": r[11]})
    dec = {}
    for r in list(wb["5_新決策建議"].iter_rows(values_only=True))[2:]:
        if r[0]:
            dec[r[0]] = {"item": r[1], "text": r[2], "scope": r[3]}
    return {"source": XLSX.name, "gov": gov, "evidence": ev, "decisions": dec}


def main():
    d = build()
    text = json.dumps(d, ensure_ascii=False, indent=1)
    if "--check" in sys.argv:
        cur = OUT.read_text(encoding="utf-8")
        print("一致" if cur == text else "不一致")
        sys.exit(0 if cur == text else 1)
    OUT.write_text(text, encoding="utf-8")
    noid = [g["gm"] for g in d["gov"] if not g["ids"]]
    print(f"gov {len(d['gov'])} 列（唯一 GM {len({g['gm'] for g in d['gov']})}）；無法解析決策編號 {len(noid)}：{noid}")
    print(f"evidence {len(d['evidence'])} 列；decisions {sorted(d['decisions'])}")


if __name__ == "__main__":
    main()
