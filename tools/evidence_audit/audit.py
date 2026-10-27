"""證據引述回查（API 版，2026-10-27 起）。

目的：逐筆檢查 docs/evidence/ 證據檔中「數字或引述＋來源網址」是否與原文一致，輸出問題清單 Excel。
只產出核對結果，不改任何證據檔或模型；判讀與修正由 chat 端與 Andy 決定。

流程（每個請求檔一次）：
1. 讀 requests/<名>.json：{"scope": ["docs/evidence/20261025_*"], "max_urls": 50}。
2. 抽取（Batch，Sonnet）：每份 .md 抽出「主張、數值、單位、證據檔中的原句、來源網址、標記」。
3. 抓原文：從 GitHub 雲端直接抓每個來源網址（識別用 User-Agent，不帶任何 email）。
   **被網站拒絕（403、需登入、付費牆）一律標「需人工」，不重試、不改用其他管道繞過。**
4. 核對（Batch，Sonnet）：每個網址一則請求，附原文與引用它的全部主張，逐筆判定
   相符／不符／部分相符／原文找不到／口徑不同，附原文片段（15 字內）。
5. 輸出 tools/evidence_audit/out/<名>.xlsx（問題清單在前）與 .md 摘要，commit 回分支。

只用標準函式庫＋openpyxl（PDF 另需 pdftotext）。
"""
import glob
import html
import html.parser
import json
import os
import re
import subprocess
import sys
import time
import urllib.error
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
HERE = os.path.dirname(os.path.abspath(__file__))
API_KEY = os.environ.get("ANTHROPIC_API_KEY", "")
MODEL = os.environ.get("AUDIT_MODEL", "claude-sonnet-5-5")
UA = "Tokenomics evidence audit (github.com/YinchenChang/Tokenomics)"
MAX_SRC_CHARS = 250_000
POLL_MAX_MIN = int(os.environ.get("POLL_MAX_MIN", "100"))
PRICE = {"opus": (2, 10), "sonnet": (1, 5), "haiku": (0.05, 0.25)}   # Batch 價（美元／百萬 token）


# ---------- API ----------
def api(path, body=None, method="GET"):
    req = urllib.request.Request(f"https://api.anthropic.com/v1/{path}",
                                 data=json.dumps(body).encode() if body is not None else None, method=method,
                                 headers={"x-api-key": API_KEY, "anthropic-version": "2023-06-01",
                                          "content-type": "application/json"})
    with urllib.request.urlopen(req, timeout=300) as r:
        return r.read().decode()


def run_batch(items, system, max_tokens=16000):
    """items：{custom_id: 使用者內容} → ({custom_id: 文字}, 費用)。"""
    if not items:
        return {}, 0.0
    reqs = [{"custom_id": cid, "params": {"model": MODEL, "max_tokens": max_tokens, "system": system,
                                         "messages": [{"role": "user", "content": txt}]}}
            for cid, txt in items.items()]
    out, cost = {}, 0.0
    for i in range(0, len(reqs), 500):          # 分批送出，避免單批過大
        b = json.loads(api("messages/batches", {"requests": reqs[i:i + 500]}, "POST"))
        print("batch", b["id"], len(reqs[i:i + 500]))
        t0 = time.time()
        while b["processing_status"] != "ended" and time.time() - t0 < POLL_MAX_MIN * 60:
            time.sleep(30)
            b = json.loads(api(f"messages/batches/{b['id']}"))
        if b["processing_status"] != "ended":
            print("批次逾時，取已完成部分")
        for line in api(f"messages/batches/{b['id']}/results").splitlines():
            r = json.loads(line)
            res = r["result"]
            if res["type"] == "succeeded":
                m = res["message"]
                out[r["custom_id"]] = "".join(c.get("text", "") for c in m["content"])
                u = m.get("usage", {})
                pi, po = next(v for k, v in PRICE.items() if k in MODEL)
                cost += (u.get("input_tokens", 0) * pi + u.get("output_tokens", 0) * po) / 1e6
            else:
                out[r["custom_id"]] = "__error__"
    return out, cost


def parse_json_array(text):
    m = re.search(r"\[.*\]", text or "", re.S)
    try:
        return json.loads(m.group(0)) if m else []
    except json.JSONDecodeError:
        return []


# ---------- 抽取 ----------
SYS_EXTRACT = """你從研究證據檔（Markdown）抽出「可回查的主張」：每一個附有來源網址、且含具體數字或直接引述的句子或表格列。
只輸出 JSON 陣列：[{"claim":主張（一句話，繁體中文）,"value":原檔寫的數字或引述（照抄）,"unit":單位,
"line":證據檔中的原句或表格列（照抄，200 字內）,"url":來源網址（照抄完整網址）,"tag":原檔標記如 [Verified]/[Interested-party]/[Derived]，沒有寫 null}]。
同一網址支撐多個主張時分開列。沒有網址的主張不列。不要推論、不要合併。"""


def extract_claims(files):
    items = {f"f{i}": f"檔案：{os.path.relpath(p, ROOT)}\n\n{open(p, encoding='utf-8').read()[:120000]}"
             for i, p in enumerate(files)}
    res, cost = run_batch(items, SYS_EXTRACT)
    claims = []
    for i, p in enumerate(files):
        for c in parse_json_array(res.get(f"f{i}", "")):
            if isinstance(c, dict) and str(c.get("url", "")).startswith("http"):
                c["file"] = os.path.relpath(p, ROOT)
                c["url"] = c["url"].strip().rstrip(".,;）)」")
                claims.append(c)
    return claims, cost


# ---------- 抓原文 ----------
class _Text(html.parser.HTMLParser):
    def __init__(self):
        super().__init__()
        self.out, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self.skip += 1
        elif tag in ("br", "p", "div", "tr", "li", "h1", "h2", "h3", "h4"):
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript"):
            self.skip -= 1
        elif tag in ("td", "th"):
            self.out.append(" | ")

    def handle_data(self, d):
        if not self.skip:
            self.out.append(d)


def to_text(data, ctype, url):
    if data[:5] == b"%PDF-" or "pdf" in ctype:
        p = subprocess.run(["pdftotext", "-layout", "-", "-"], input=data, capture_output=True)
        return p.stdout.decode("utf-8", "replace")
    s = data.decode("utf-8", "replace")
    if "html" in ctype or s.lstrip()[:200].lower().startswith(("<!doctype", "<html")):
        p = _Text()
        p.feed(s)
        s = html.unescape("".join(p.out))
    return re.sub(r"\n\s*\n+", "\n", re.sub(r"[ \t]+", " ", s))


def fetch(url):
    """回傳 (狀態, 文字)。狀態：ok／需人工（HTTP 碼或原因）。不重試被拒的請求。"""
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            data = r.read(20_000_000)
            t = to_text(data, r.headers.get("Content-Type", ""), url)
            if len(t.strip()) < 300:
                return "需人工（內容過短，可能需 JavaScript 或登入）", t
            return "ok", t[:MAX_SRC_CHARS]
    except urllib.error.HTTPError as e:
        return f"需人工（HTTP {e.code}）", ""
    except Exception as e:
        return f"需人工（{type(e).__name__}）", ""
    finally:
        time.sleep(0.5)


# ---------- 核對 ----------
SYS_VERIFY = """你是獨立核對者。給你一份來源原文，以及研究檔引用這份來源的主張清單。逐筆判定主張是否得到原文支持。
判定只能是：相符／部分相符／不符／原文找不到／口徑不同。
- 數字要逐位比對（單位、年度、季度、年化或全年、含不含某項）；四捨五入差異算相符，但在 note 說明。
- 若原文是記者轉述或推論、而主張寫成公司本人說，判「口徑不同」。
- 標記 [Verified] 但原文只是利害關係人說法，在 note 指出。
只輸出 JSON 陣列：[{"id":主張編號,"verdict":判定,"source_snippet":原文中最相關的片段（15 字以內）,"note":理由（繁體中文，50 字內）}]。"""


def verify(claims, texts):
    by_url = {}
    for i, c in enumerate(claims):
        by_url.setdefault(c["url"], []).append(i)
    items = {}
    for j, (url, idxs) in enumerate(by_url.items()):
        if texts.get(url, ("", ""))[0] != "ok":
            continue
        lst = "\n".join(f"- id={i}｜主張：{claims[i]['claim']}｜數值：{claims[i].get('value')}｜{claims[i].get('unit') or ''}"
                        f"｜標記：{claims[i].get('tag')}｜研究檔原句：{str(claims[i].get('line'))[:200]}" for i in idxs)
        items[f"u{j}"] = f"來源網址：{url}\n\n## 主張清單\n{lst}\n\n## 來源原文\n{texts[url][1]}"
    res, cost = run_batch(items, SYS_VERIFY, max_tokens=8000)
    for txt in res.values():
        for v in parse_json_array(txt):
            try:
                claims[int(v["id"])].update({"verdict": v.get("verdict"), "snippet": v.get("source_snippet"), "vnote": v.get("note")})
            except (KeyError, ValueError, IndexError, TypeError):
                pass
    return cost


# ---------- 輸出 ----------
ORDER = {"不符": 0, "口徑不同": 1, "部分相符": 2, "原文找不到": 3, "需人工": 4, "未判定": 5, "相符": 6}


def write_outputs(name, claims, texts, cost, req):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    for c in claims:
        st = texts.get(c["url"], ("未抓取", ""))[0]
        if st != "ok":
            c["verdict"], c["vnote"] = "需人工", st
        c.setdefault("verdict", "未判定")
    claims.sort(key=lambda c: (ORDER.get(c["verdict"], 5), c["file"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "核對結果"
    head = ["判定", "證據檔", "主張", "數值", "單位", "標記", "證據檔原句", "來源網址", "原文片段", "理由"]
    ws.append(head)
    fills = {"不符": "F8D7DA", "口徑不同": "FCE4D6", "部分相符": "FFF3CD", "原文找不到": "FFF3CD", "需人工": "E2E3E5"}
    for c in claims:
        ws.append([c["verdict"], c["file"], c.get("claim"), str(c.get("value")), c.get("unit"), c.get("tag"),
                   str(c.get("line"))[:300], c["url"], c.get("snippet"), c.get("vnote")])
        if c["verdict"] in fills:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fills[c["verdict"]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "C2"
    for col, w in zip("ABCDEFGHIJ", (10, 34, 40, 14, 8, 12, 50, 40, 24, 40)):
        ws.column_dimensions[col].width = w
    s2 = wb.create_sheet("來源可達性")
    s2.append(["網址", "狀態", "取得字數", "引用筆數"])
    cnt = {}
    for c in claims:
        cnt[c["url"]] = cnt.get(c["url"], 0) + 1
    for url, (st, t) in sorted(texts.items(), key=lambda x: x[1][0]):
        s2.append([url, st, len(t), cnt.get(url, 0)])
    s3 = wb.create_sheet("說明與用量")
    for row in [["範圍", "、".join(req.get("scope", []))], ["模型", MODEL], ["估計費用（美元）", round(cost, 2)],
                ["說明", "本表只列核對結果，不是結論；判定由模型產生，『不符』『口徑不同』須人工對原文確認後才修正證據檔。"
                         "『需人工』表示網站拒絕自動抓取或需登入，依規則不繞過。"]]:
        s3.append(row)
    os.makedirs(f"{HERE}/out", exist_ok=True)
    wb.save(f"{HERE}/out/{name}.xlsx")
    tally = {}
    for c in claims:
        tally[c["verdict"]] = tally.get(c["verdict"], 0) + 1
    ok_urls = sum(1 for st, _ in texts.values() if st == "ok")
    md = [f"# 證據引述回查｜{name}", "", f"- 範圍：{'、'.join(req.get('scope', []))}",
          f"- 主張 {len(claims)} 筆；來源 {len(texts)} 個（可自動取得 {ok_urls}）",
          "- 判定：" + "、".join(f"{k} {v}" for k, v in sorted(tally.items(), key=lambda x: ORDER.get(x[0], 5))),
          f"- 估計費用約 US${cost:.2f}（模型 {MODEL}，Batch 五折）", "",
          "判定由模型產生，只作為人工複核的優先順序；詳見同名 xlsx。"]
    open(f"{HERE}/out/{name}.md", "w", encoding="utf-8").write("\n".join(md) + "\n")
    return tally


def main():
    if not API_KEY:
        print("未設定 ANTHROPIC_API_KEY，結束。")
        return
    for path in sorted(glob.glob(f"{HERE}/requests/*.json")):
        name = os.path.splitext(os.path.basename(path))[0]
        req = json.load(open(path, encoding="utf-8"))
        try:
            files = sorted({f for pat in req.get("scope", []) for f in glob.glob(f"{ROOT}/{pat}/**/*.md", recursive=True)
                            + glob.glob(f"{ROOT}/{pat}", recursive=True) if f.endswith(".md")})
            print(name, "檔案", len(files))
            claims, cost = extract_claims(files)
            urls = list(dict.fromkeys(c["url"] for c in claims))[: req.get("max_urls", 2000)]
            claims = [c for c in claims if c["url"] in urls]
            texts = {u: fetch(u) for u in urls}
            cost += verify(claims, texts)
            print(name, write_outputs(name, claims, texts, cost, req), f"US${cost:.2f}")
            os.makedirs(f"{HERE}/requests/done", exist_ok=True)
            os.replace(path, f"{HERE}/requests/done/{name}.json")
        except Exception as e:
            os.makedirs(f"{HERE}/out", exist_ok=True)
            open(f"{HERE}/out/{name}.md", "w", encoding="utf-8").write(f"# 回查失敗｜{name}\n\n{type(e).__name__}: {e}\n")
            print("失敗", name, e)


if __name__ == "__main__":
    main()
