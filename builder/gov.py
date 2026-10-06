# v5.11 (Stage 1 slice one): Source layer (SRC_HW／SRC_DC／SRC_Model／SRC_Perf), DB_Evidence upgrade, Decisions,
# v5.13 (slice two, package B): SRC_Price／SRC_Cap／SRC_Harness／SRC_Demand (gov_seed2), S30 appended to SRC_Perf,
#   FORMULA_MAP2 links (Cap_In, Har_In, Workload 40), Checks C9:C10／C37／I43:I45 linked to SRC, Gov_Map judgment updates (GOV_MAP_UPD).
# Gov_Map (blue-cell registry), L1 (first batch), governance Checks, F14 (per-GW outputs divided by Inputs!E5).
#
# Ownership (Excel-first, same rule as DB_Evidence since v5.7):
#   - SRC_*, Decisions, DB_Evidence, and the judgment columns of Gov_Map are Excel-owned: created from gov_seed /
#     gov_decisions only when absent, never overwritten afterwards.
#   - Builder-owned (rewritten on every build): FORMULA_MAP links on model pages, helper formulas in SRC_* (X–Z),
#     Gov_Map formula / status columns (Q–AF), the SRC_Index sheet (v5.12), the L1 sheet, the Checks governance section,
#     SRC_/L1_/GOV_/IDX_ names. v5.12 (A) also appends Gov_Map rows for new input cells (GOV_MAP_V512A) when absent.
import re
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from common import put, F_IN, F_CALC, F_LINK, F_BOLD, F_TITLE, F_NOTE, FILL_SEC, FILL_KEY, WRAP, L, title, section
from gov_seed import SRC_RECORDS, PERF_ATTR, EVID_MIG, EVID_UPD, FORMULA_MAP, GOV_MAP
from gov_decisions import DECISIONS
from gov_seed2 import SRC_RECORDS2, PERF_ATTR2, EVID_MIG2, FORMULA_MAP2, GOV_MAP_UPD, GOV_MAP_V513C, GOV_MAP_UPD_E, DEC_UPD
from gov_seed3 import SRC_RECORDS3, EVID_MIG3, DECISIONS_V515, DEC_STATUS_V515, GOV_MAP_V515
import v518                       # v5.18: Stage 2 first write batch (Excel-owned writes with old-value guards; see v518.py)
import v519                       # v5.19: X1 mirrors (Prod sheets) and X2 Gov_Map P promotions; its Gov_Map row and Decisions are registered here

ALL_RECORDS = SRC_RECORDS + SRC_RECORDS2 + SRC_RECORDS3
ALL_PERF_ATTR = {**PERF_ATTR, **PERF_ATTR2}
ALL_FORMULA_MAP = {**FORMULA_MAP, **FORMULA_MAP2}
def _fm(): return v518.formula_map(ALL_FORMULA_MAP)     # v5.18: steps (C44, C69, G2, G4, E3, RU) add or drop links

SRC_SHEETS = ["SRC_HW", "SRC_DC", "SRC_Model", "SRC_Perf", "SRC_Price", "SRC_Cap", "SRC_Harness", "SRC_Demand"]   # v5.13: +4
SRC_LAST = 400                      # record rows 5..SRC_LAST (formula ranges)
SRC_HDR = ["SRC_ID", "指標", "數值", "低", "高", "單位", "口徑", "適用對象", "日期", "出處", "來源等級", "立場", "立場說明",
           "一手／二手", "狀態", "取代者", "Evidence ID", "第二來源 SRC_ID", "審查日", "Andy 原話", "原 S 編號", "模型使用位置", "備註",
           "同指標同口徑 Active 數（公式）", "缺 Evidence（公式）", "利害關係方缺第二來源（公式）"]
PERF_HDR = ["平台", "軟體／日期", "ISL", "OSL", "每用戶速度 tok/s", "MTP（1＝有）", "token 口徑"]
PERF_NAMES = {"AC": "_ISL", "AD": "_OSL", "AE": "_Spd", "AF": "_MTP"}
SRC_TITLE = {"SRC_HW": "晶片與機架的規格、功率、價格", "SRC_DC": "廠房資本支出、電價、折舊慣例、外部成本參照",
             "SRC_Model": "模型架構與訓練揭露", "SRC_Perf": "推論與訓練量測（InferenceX、MLPerf、自揭）及量測條件",
             "SRC_Price": "API 牌價與算力租金（廠商牌價、新雲價目、第三方推算）", "SRC_Cap": "能力評測（Artificial Analysis 指數、METR 時間範圍）",
             "SRC_Harness": "harness 與代理的 token 倍數、成功率、成本（ARC-AGI-3、廠商揭露）", "SRC_Demand": "需求與支出揭露（實驗室推論／訓練支出、營收、算力規模）"}
DASH = "—"


def _nm(wb, n, ref):
    if n in wb.defined_names: del wb.defined_names[n]
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


# ---------------------------------------------------------------- 1. Source sheets (Excel-owned after creation)
def src_sheets(wb):
    made = []
    for sh in SRC_SHEETS:
        if sh in wb.sheetnames: continue
        ws = wb.create_sheet(sh); made.append(sh)
        title(ws, f"{sh} — 第 0 層 Source：{SRC_TITLE[sh]}",
              "只存已取得的原始訊息（非計算、非假設；G1）。藍字＝原始值，由本活頁簿擁有；模型頁以公式連結 SRC_ID 具名範圍。"
              "新訊息先登錄 DB_Evidence，比較後擇優寫入（G2–G6）。X–Z 欄為檢查公式（builder 每次重建）。")
        put(ws, "A3", "來源等級：A＝Andy、1＝一手且已讀原文、2＝二手已讀或一手文件轉載、3＝模型內建知識、只見摘錄或待查核（G0-7）。"
                      "狀態：Active／Alt（口徑不同或多家獨立估計並列，G13）／Superseded。", F_NOTE)
        hdr = SRC_HDR + (PERF_HDR if sh == "SRC_Perf" else [])
        widths = [13, 40, 11, 9, 9, 13, 22, 20, 11, 40, 6, 10, 30, 12, 10, 10, 16, 14, 9, 12, 8, 30, 36, 10, 10, 10] + [16, 22, 7, 7, 9, 8, 14]
        for i, h in enumerate(hdr):
            put(ws, f"{L(i+1)}4", h, F_BOLD, wrap=True); ws.column_dimensions[L(i+1)].width = widths[i]
        r = 5
        for rec in [x for x in ALL_RECORDS if x["sheet"] == sh]:
            _src_row(ws, r, rec); r += 1
        ws.freeze_panes = "C5"
    return made


def _src_row(ws, r, rec):
    vals = [rec["id"], rec["metric"], rec["val"], rec["lo"], rec["hi"], rec["unit"], rec["basis"] or DASH, rec["applies"] or DASH,
            str(rec["date"]), rec["src"] or DASH, rec["grade"], rec["stance"], rec["stance_note"] or DASH, rec["hand"] or DASH,
            rec["status"], DASH, rec["ev"], DASH, DASH, DASH, rec["s"], rec["use"] or DASH, rec["note"] or DASH]
    for i, v in enumerate(vals):
        f = F_IN if i in (2, 3, 4) and isinstance(v, (int, float)) else F_CALC
        put(ws, f"{L(i+1)}{r}", v, f, wrap=i in (1, 9, 12, 21, 22))
    if ws.title == "SRC_Perf" and rec["id"] in ALL_PERF_ATTR:
        a = ALL_PERF_ATTR[rec["id"]]
        for col, k in zip(["AA", "AB", "AC", "AD", "AE", "AF", "AG"], ["platform", "software", "isl", "osl", "spd", "mtp", "tok"]):
            if a[k] is not None or rec["id"] in PERF_ATTR: put(ws, f"{col}{r}", a[k], F_IN if isinstance(a[k], (int, float)) else F_CALC)


def src_append(wb):
    """v5.13: records of SRC_RECORDS2 whose sheet already exists (S30 → SRC_Perf) are appended after its last record, only when
    the ID is absent anywhere on that sheet (Excel-owned afterwards; an ID Andy deleted or renamed is not re-added if its row moved)."""
    added = []
    for rec in SRC_RECORDS2 + SRC_RECORDS3 + v518.src_new_records():
        ws = wb[rec["sheet"]]
        ids = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
        if rec["id"] in ids: continue
        last = max([r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value not in (None, "")] or [4])
        _src_row(ws, last + 1, rec); added.append(rec["id"])
    return added


KEY_COL, KEY_SEP = "AH", "¦"       # v5.12 (A): builder-owned key column (指標¦口徑¦適用對象, Active rows only) used by X

def src_refresh(wb):
    """Helper formulas (X–Z, AH) for every record row, and SRC_ names (value／_Lo／_Hi／Perf attributes)."""
    n_names = 0; index = {}
    for sh in SRC_SHEETS:
        ws = wb[sh]
        put(ws, f"{KEY_COL}4", "同指標鍵（公式；X 欄用，v5.12）", F_BOLD, wrap=True); ws.column_dimensions[KEY_COL].width = 12
        # v5.14: X compares only rows 5..(last record + IDX_HEAD), the same span SRC_Index uses for this sheet (was 5..SRC_LAST).
        # A record beyond the span is already an ERROR (Checks E13, fixed by rebuilding), so the counts are unchanged; the
        # comparison arrays shrink from 8 × 396 to about 700 rows (X columns were ~40% of full-recalc time in v5.13).
        end = _span_end(ws)
        for r in range(5, ws.max_row + 1):
            sid = ws.cell(r, 1).value
            if not (isinstance(sid, str) and sid.startswith("SRC_")): continue
            index[sid] = (sh, r)
            rng = lambda c: f"${c}$5:${c}${end}"
            # v5.12 (A): same count as v5.11 (B, G, H equal and Active) via one key column (AH) and one comparison array,
            # instead of four 396-row arrays per record (the SRC X columns were ~60% of full-recalc time after SRC_Index)
            put(ws, f"{KEY_COL}{r}", f'=IF($O{r}="Active",$B{r}&"{KEY_SEP}"&$G{r}&"{KEY_SEP}"&$H{r},"")')
            put(ws, f"X{r}", f'=IF($O{r}="Active",SUMPRODUCT(({rng(KEY_COL)}=${KEY_COL}{r})*1),0)', fmt="0")
            put(ws, f"Y{r}", f'=IF(AND($O{r}="Active",OR($Q{r}="{DASH}",$Q{r}="")),1,0)', fmt="0")
            put(ws, f"Z{r}", f'=IF(AND($O{r}="Active",$L{r}="利害關係方",OR($R{r}="{DASH}",$R{r}="")),1,0)', fmt="0")
            _nm(wb, sid, f"{sh}!$C${r}"); n_names += 1
            for col, suf in (("D", "_Lo"), ("E", "_Hi")):
                if ws[f"{col}{r}"].value is not None: _nm(wb, sid + suf, f"{sh}!${col}${r}"); n_names += 1
            if sh == "SRC_Perf":
                for col, suf in PERF_NAMES.items():
                    if ws[f"{col}{r}"].value is not None: _nm(wb, sid + suf, f"{sh}!${col}${r}"); n_names += 1
    return n_names, index


# ---------------------------------------------------------------- 2. DB_Evidence upgrade (Excel-owned)
EV_NEW_HDR = ["狀態（已處理／待判定／待 Andy）", "相關 SRC_ID", "取代者", "受影響 IF_／L1_", "來源等級", "立場"]

def evidence_upgrade(wb):
    ws = wb["DB_Evidence"]; added = 0
    if ws["L4"].value is None:
        for i, h in enumerate(EV_NEW_HDR):
            put(ws, f"{L(12+i)}4", h, F_BOLD, wrap=True); ws.column_dimensions[L(12+i)].width = [14, 30, 12, 20, 8, 12][i]
        put(ws, "A3", "v5.11：新增 L–Q 欄（G2 流程的狀態、相關 SRC_ID、取代者、受影響名稱、等級、立場）；E101 起為 Stage 1 遷移紀錄。", F_NOTE)
    have = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}
    for eid, vals in EVID_UPD.items():
        r = have.get(eid)
        if r and ws.cell(r, 12).value is None:
            for i, v in enumerate(vals):
                put(ws, f"{L(12+i)}{r}", v if v != "" else DASH, F_CALC, wrap=i in (1, 5))
    r = max(have.values()) + 1 if have else 5
    for row in EVID_MIG + EVID_MIG2 + EVID_MIG3 + v518.evidence_rows():
        if row[0] in have: continue
        for i, v in enumerate(row):
            put(ws, f"{L(i+1)}{r}", v if v != "" else DASH, F_IN if i < 11 else F_CALC, wrap=i in (2, 10, 12))
        r += 1; added += 1
    return added


# ---------------------------------------------------------------- 3. Decisions (Excel-owned)
DEC_HDR = ["決策 ID", "類別", "項目", "決議", "日期", "Andy 原話／依據", "狀態", "影響範圍（工作表、格、名稱）", "出處", "原話待 Andy 確認"]

def decisions_sheet(wb):
    if "Decisions" in wb.sheetnames: return False
    ws = wb.create_sheet("Decisions")
    title(ws, "Decisions — 決策登錄（G11：交接文件第 2 節移入；之後交接文件只保留決策 ID）",
          "Excel 擁有：builder 只在本頁不存在時建立。模型頁的 Decision 格以決策 ID 引用（見 Gov_Map J 欄）；Checks 檢查 ID 是否存在。"
          "「原話待 Andy 確認」＝是：轉錄自交接文件，Andy 確認原話後改為否。")
    widths = [8, 14, 24, 70, 14, 30, 14, 34, 22, 10]
    for i, h in enumerate(DEC_HDR):
        put(ws, f"{L(i+1)}4", h, F_BOLD, wrap=True); ws.column_dimensions[L(i+1)].width = widths[i]
    for j, row in enumerate(DECISIONS):
        for i, v in enumerate(row):
            v = v.replace("**", "").replace("`", "") if isinstance(v, str) else v
            put(ws, f"{L(i+1)}{5+j}", v, F_CALC, wrap=i in (2, 3, 5, 7))
    ws.freeze_panes = "B5"
    return True


def dec_update(wb):
    """v5.13 E: Andy's review results written to Decisions (Excel-owned), each field only while it still holds the v5.12 value."""
    ws = wb["Decisions"]; col = {"D": 4, "F": 6, "G": 7, "J": 10}; n = 0
    for r in range(5, ws.max_row + 1):
        f = {**DEC_UPD, **DEC_STATUS_V515}.get(ws.cell(r, 1).value)     # v5.15: A1–A8 status "v5.15 已建"
        if not f: continue
        for k, (old, new) in f.items():
            c = ws.cell(r, col[k])
            if c.value == old: c.value = new; n += 1
    return n


def dec_append(wb):
    """v5.15: Decisions A9／A10 are appended only when the ID is absent (Excel-owned afterwards)."""
    ws = wb["Decisions"]; have = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
    r = max([rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")] or [4]) + 1; n = 0
    for row in DECISIONS_V515 + v518.decisions_rows() + v519.DECISIONS_V519:
        if row[0] in have: continue
        for i, v in enumerate(row): put(ws, f"{L(i+1)}{r}", v, F_CALC, wrap=i in (2, 3, 5, 7))
        r += 1; n += 1
    return n


# ---------------------------------------------------------------- 4. Model-page links (builder-owned)
def apply_formula_map(wb):
    n = 0; log = []
    for key, f in _fm().items():
        sh, co = key.split("!")
        c = wb[sh][co]
        if c.value != f:
            log.append(f"{key}: {c.value!r} -> {f}")
            c.value = f; n += 1
        c.font = F_LINK if ("SRC_" in f or "!" in f) else F_CALC
        c.fill = PatternFill(fill_type=None)
    return n, log


# ---------------------------------------------------------------- 5. F14: every "每 GW" output divided by Inputs!E5
def f14(wb):
    n = 0
    def wrap(ws, ref):
        nonlocal n
        v = ws[ref].value
        if isinstance(v, str) and v.startswith("=") and not v.endswith("/CTL_GW"):
            ws[ref].value = "=(" + v[1:] + ")/CTL_GW"; n += 1
    itf = wb["Interface"]
    for r in (6, 7, 8, 9, 10, 11, 12, 13, 15):
        for c in range(3, 18): wrap(itf, f"{L(c)}{r}")
    wrap(wb["Checks"], "B4"); wrap(wb["Checks"], "B5")
    sen = wb["Sensitivity"]
    for c in range(3, 26):
        if sen.cell(7, c).value == "=Inputs!$E$5*1000": sen.cell(7, c).value = "=1000"; n += 1
    sen["A7"].value = "IT 電力（每 GW 基準；F14）"
    # v5.11 (G0-5 scan): scenario values written as formulas (=0.05 etc.) become plain inputs (value unchanged)
    for ref in ("X8", "Y8", "N37"):
        v = sen[ref].value
        if isinstance(v, str) and re.fullmatch(r"=-?\d+(\.\d+)?", v):
            sen[ref].value = float(v[1:]) if "." in v else int(v[1:]); sen[ref].font = F_IN; n += 1
    dc = wb["DC_Cost"]
    dc["A1"].value = "DC_Cost — 設施合計（IT 規模＝Inputs!E5 GW，基準 1）：資本支出與年持有成本；5 世代 × 3 成本情境。每 GW 值見 Interface（F14）"
    for r, t in [(12, "機架數（設施）"), (14, "GPU 數（設施）"), (17, "B. 資本支出（$B；設施）"), (33, "總資本支出（設施）"),
                 (37, "C. 年持有成本 — 會計口徑（$B／年；設施）"), (50, "D. 年持有成本 — 經濟口徑（資本回收年金；設施）")]:
        dc[f"A{r}"].value = t
    wb["Inputs"]["H5"].value = "口徑：IT 關鍵電力。v5.11 起 DC_Cost 為設施合計；Interface、L1 與模型頁的「每 GW」一律已除以本格（F14）"
    return n


# ---------------------------------------------------------------- 6. Gov_Map
GM_HDR = ["GM_ID", "範圍", "工作表", "格", "列標籤", "類別", "區間角色", "SRC_ID（來源或可比對象）", "關係", "決策 ID", "低", "高", "區間文字",
          "理由", "標記變更（v5.11）", "CC 敏感度分段",
          "模型值", "SRC 值", "SRC 狀態", "SRC 等級", "實際狀態（建置時）",
          "E1 原始寫死", "E2 Analogy 缺可比", "E3 缺區間", "E4 缺理由", "E5 Decision 缺 ID", "E6 決策 ID 不存在", "E9 SRC 非 Active",
          "E10 區間順序", "W2 3 級×高段", "E12 SRC 不存在", "SRC_Index 列（MATCH；v5.12）"]
GM_LAST = 700

# ---------------------------------------------------------------- 5a. SRC_Index (v5.12 A; builder-owned, rebuilt every build)
# Why: v5.11 Gov_Map S／T ran 16 MATCH per row over 4 SRC sheets × rows 5..400 (≈6,960 MATCH; full recalc 3.4–3.7 s),
# a cost that grows with Gov_Map rows × SRC sheets. SRC_Index stacks the A (ID), O (status) and K (grade) columns of every
# SRC sheet into one list, so Gov_Map needs one MATCH per row (column AF), shared by S and T.
# A sentinel block mirrors Gov_Map!H, so MATCH always finds the ID (an ID absent from every SRC sheet lands in the sentinel
# block, whose status／grade read "不存在", exactly the v5.11 result) — no #N/A, no new function (no IFERROR).
IDX_HEAD = 50                       # spare slots per SRC sheet after its last record (Checks E13 flags records beyond them)

def _span_end(src):
    """Last row covered for an SRC sheet: last record + IDX_HEAD, capped at SRC_LAST (shared by SRC_Index and the X columns)."""
    last = max([rr for rr in range(5, min(src.max_row, SRC_LAST) + 1) if src.cell(rr, 1).value not in (None, "")] or [4])
    return min(last + IDX_HEAD, SRC_LAST)

def src_index(wb):
    if "SRC_Index" in wb.sheetnames: del wb["SRC_Index"]
    ws = wb.create_sheet("SRC_Index")
    title(ws, "SRC_Index — SRC_* 的 ID、狀態、等級依序堆疊（builder 擁有，每次重建；只供 Gov_Map 查找，不放任何數值）",
          "Gov_Map AF 欄以 1 個 MATCH 找到 SRC_ID 所在列，S（狀態）、T（等級）欄以 INDEX 讀本頁。各 SRC 頁的範圍＝第 5 列到最後一筆紀錄＋"
          f"{IDX_HEAD} 列；紀錄超出範圍時 Checks E13 報錯（重建即可）。最後一段為 Gov_Map H 欄的鏡像（哨兵）：找不到的 SRC_ID 落在這段，狀態與等級為「不存在」。")
    for i, (h, w) in enumerate(zip(["SRC_ID", "工作表", "原列", "狀態", "等級"], [16, 14, 7, 12, 8])):
        put(ws, f"{L(i+1)}4", h, F_BOLD); ws.column_dimensions[L(i+1)].width = w
    r = 5; spans = {}
    for sh in SRC_SHEETS:
        src = wb[sh]
        end = _span_end(src)
        spans[sh] = (r, end)
        for rr in range(5, end + 1):
            put(ws, f"A{r}", f'={sh}!$A{rr}&""'); put(ws, f"B{r}", sh); put(ws, f"C{r}", rr, F_CALC)
            put(ws, f"D{r}", f"={sh}!$O{rr}"); put(ws, f"E{r}", f"={sh}!$K{rr}")
            r += 1
    sent0 = r
    for gr in range(5, GM_LAST + 1):
        put(ws, f"A{r}", f'=Gov_Map!$H{gr}&""'); put(ws, f"B{r}", "（哨兵）", F_NOTE); put(ws, f"C{r}", gr, F_CALC)
        put(ws, f"D{r}", "不存在"); put(ws, f"E{r}", "不存在")
        r += 1
    _nm(wb, "IDX_SrcID", f"SRC_Index!$A$5:$A${r-1}")
    _nm(wb, "IDX_SrcStat", f"SRC_Index!$D$5:$D${r-1}")
    _nm(wb, "IDX_SrcGrade", f"SRC_Index!$E$5:$E${r-1}")
    ws.freeze_panes = "A5"
    return dict(rows=r - 5, src_rows=sent0 - 5, spans={k: v[1] for k, v in spans.items()})

# v5.12 (A): input cells created by moving formula constants out (交接第 8i 節的判定). Each entry: (sheet, column-A label of the
# row, column span, fields). The cell is located by its label, so the row number is never hard-coded here.
GOV_MAP_V512A = [
  ("Arch", "上下文長度（上一列顯示用）", "C", dict(scope="切片一", label="上下文長度（128K 上下文每序列 KV 顯示列）", cls="Assumed", role="單值",
      rtext="結構選擇（無數值區間）", reason="顯示列口徑 128,000 tok（非 131,072），無任何引用；v5.12 自 Arch C31:E31 公式移出（值不變）", seg="無")),
  ("Sens_Train", "情境倍數（×；黃底藍字＝作用中，乘在該情境改動的量上）", "O:R", dict(scope="切片二頁（v5.12 A 包）",
      label="情境倍數：RL rollout token × 0.3／× 3", cls="情境值", role="群組", reason="Sens_Train 情境；v5.12 自 O53:R53 公式移出（值不變）", seg="—")),
  ("Sens_Train", "情境倍數（×；黃底藍字＝作用中，乘在該情境改動的量上）", "Y:AB", dict(scope="切片二頁（v5.12 A 包）",
      label="情境倍數：合成資料 token × 0／× 3", cls="情境值", role="群組",
      reason="Sens_Train 情境；v5.12 自 AA93:AB93 公式與 Y93:Z93 直接寫入的 0 移出（值不變）", seg="—")),
  ("Sens_Rev", "ε（上一列量級情境的指數；能力 ∝ 有效算力^ε）", "B", dict(scope="切片二頁（v5.12 A 包）", label="ε（K4 (a) 量級情境）",
      cls="情境值", role="單值", reason="K4 (a) 若採用的量級情境；v5.12 自 D17、F17 公式移出（值不變）", seg="—")),
  ("Tech_Registry", "門檻", "J", dict(scope="切片二頁（v5.12 A 包）", label="主流判定門檻（公開採用實驗室數）", cls="Decision", role="單值",
      dec="J14", reason="J14：至少兩家實驗室公開採用即為主流；v5.12 自 K5:K16 公式移出（值不變）", seg="—")),
]

def _find_row(ws, label):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == label: return r
    raise KeyError(f"{ws.title}: row label not found: {label}")

def gm_append(wb, ws):
    have = {(ws[f"C{r}"].value, ws[f"D{r}"].value) for r in range(5, ws.max_row + 1)}
    ids = [ws[f"A{r}"].value for r in range(5, ws.max_row + 1) if isinstance(ws[f"A{r}"].value, str) and ws[f"A{r}"].value.startswith("GM")]
    nxt = max(int(x[2:]) for x in ids) + 1 if ids else 1
    r = max([rr for rr in range(5, ws.max_row + 1) if ws[f"C{rr}"].value] or [4]) + 1
    added = 0
    for sh, lab, span, g in GOV_MAP_V512A:
        rr = _find_row(wb[sh], lab)
        c0, c1 = (span.split(":") + [span])[:2]
        cell = f"{c0}{rr}" if c0 == c1 else f"{c0}{rr}:{c1}{rr}"
        if (sh, cell) in have: continue
        vals = [f"GM{nxt:03d}", g["scope"], sh, cell, g["label"], g["cls"], g["role"], g.get("src") or DASH, g.get("rel") or DASH,
                g.get("dec") or DASH, None, None, g.get("rtext") or DASH, g["reason"], DASH, g.get("seg") or DASH]
        for i, v in enumerate(vals):
            put(ws, f"{L(i+1)}{r}", v, F_CALC, wrap=i in (4, 8, 13, 14))
        r += 1; nxt += 1; added += 1
    return added

def gm_append_c(wb, ws):
    """v5.13 C: register the slice-two pages' numeric blue cells (GOV_MAP_V513C) when the (sheet, cell) is not yet in Gov_Map.
    The row's column-A label must still match: a moved or renamed row stops the build instead of registering the wrong cell."""
    have = {(ws[f"C{r}"].value, ws[f"D{r}"].value) for r in range(5, ws.max_row + 1)}
    ids = [ws[f"A{r}"].value for r in range(5, ws.max_row + 1) if isinstance(ws[f"A{r}"].value, str) and ws[f"A{r}"].value.startswith("GM")]
    nxt = max(int(x[2:]) for x in ids) + 1 if ids else 1
    r = max([rr for rr in range(5, ws.max_row + 1) if ws[f"C{rr}"].value] or [4]) + 1
    added = 0
    for g in GOV_MAP_V513C + GOV_MAP_V515 + v519.GOV_MAP_V519:
        if (g["sheet"], g["cell"]) in have: continue
        first = g["cell"].split(":")[0]
        row = int(re.sub(r"[A-Z]+", "", first))
        if wb[g["sheet"]].cell(row, 1).value != g.get("check", g["label"]):     # "check": 模型頁欄 A 的實際標籤（Gov_Map 顯示標籤可不同）
            raise KeyError(f"GOV_MAP_V513C: {g['sheet']}!{g['cell']} row label changed: {wb[g['sheet']].cell(row, 1).value!r} != {g.get('check', g['label'])!r}")
        vals = [f"GM{nxt:03d}", g["scope"], g["sheet"], g["cell"], g["label"], g["cls"], g["role"], g["src"] or DASH, g["rel"] or DASH,
                g["dec"] or DASH, g["lo"], g["hi"], g["rtext"] or DASH, g["reason"] or DASH, g["retag"] or DASH, g["seg"] or DASH]
        for i, v in enumerate(vals):
            font = F_IN if i in (10, 11) and isinstance(v, (int, float)) else (F_LINK if isinstance(v, str) and v.startswith("=") else F_CALC)
            put(ws, f"{L(i+1)}{r}", v, font, wrap=i in (4, 8, 13, 14))
        r += 1; nxt += 1; added += 1
    return added


def gm_update(ws):
    """v5.13: judgment updates to existing Gov_Map rows (Excel-owned A–P). Each field is written only while it still holds the
    v5.12 value, so the update applies once and never overwrites a later Excel edit."""
    loc = {(ws[f"C{r}"].value, ws[f"D{r}"].value): r for r in range(5, ws.max_row + 1) if ws[f"C{r}"].value}
    n = 0
    for sh, cell, fields in GOV_MAP_UPD + GOV_MAP_UPD_E:
        r = loc.get((sh, cell))
        if r is None: continue
        for col, (old, new) in fields.items():
            if ws[f"{col}{r}"].value == old:
                ws[f"{col}{r}"].value = new; n += 1
    return n


def gov_map(wb, src_index):
    ws = wb["Gov_Map"] if "Gov_Map" in wb.sheetnames else None
    if ws is None:
        ws = wb.create_sheet("Gov_Map")
        title(ws, "Gov_Map — 模型頁藍字格登錄（Stage 1 切片一：Inputs、Spec_Rack、Arch、Serving、Workload、Calib、Energy、NonNV 與其敏感度情境格）",
              "A–P 欄為判斷（Excel 擁有；類別、SRC_ID、決策 ID、區間、理由）；Q–AE 欄為 builder 每次重建的公式與建置時狀態。"
              "Checks G 節加總 V–AE 欄。區間角色＝低／高者為另一格的端點，不另檢查區間。")
        for i, h in enumerate(GM_HDR):
            put(ws, f"{L(i+1)}4", h, F_BOLD, wrap=True)
            ws.column_dimensions[L(i+1)].width = [8, 10, 11, 9, 26, 16, 7, 14, 22, 8, 9, 9, 18, 36, 22, 7][i] if i < 16 else 9
        for j, g in enumerate(GOV_MAP):
            r = 5 + j
            vals = [f"GM{j+1:03d}", g["scope"], g["sheet"], g["cell"], g["label"], g["cls"], g["role"], g["src"] or DASH, g["rel"] or DASH,
                    g["dec"] or DASH, g["lo"], g["hi"], g["rtext"] or DASH, g["reason"] or DASH, g["retag"] or DASH, g["seg"] or DASH]
            for i, v in enumerate(vals):
                font = F_IN if i in (10, 11) and isinstance(v, (int, float)) else (F_LINK if isinstance(v, str) and v.startswith("=") else F_CALC)
                put(ws, f"{L(i+1)}{r}", v, font, wrap=i in (4, 8, 13, 14))
        ws.freeze_panes = "E5"
    if ws["AF4"].value is None: put(ws, "AF4", GM_HDR[31], F_BOLD, wrap=True)
    gm_append(wb, ws)
    n_c = gm_append_c(wb, ws)
    n_upd = gm_update(ws) + v518.gov_update(ws) + v519.gov_update(ws, v518._append_text)      # v5.18: judgement columns, P (B method), D5 ranges; v5.19: X2 P promotions
    # ---- builder-owned columns Q..AF
    n = 0; static_raw_hard = 0
    for r in range(5, ws.max_row + 1):
        sh, cell = ws[f"C{r}"].value, ws[f"D{r}"].value
        if not sh or not cell: continue
        n += 1
        single = ":" not in str(cell)
        h = f"$H{r}"
        put(ws, f"Q{r}", f"='{sh}'!{cell}" if single else DASH, F_LINK)
        sid = ws[f"H{r}"].value
        if isinstance(sid, str) and sid in src_index:
            s2, rr = src_index[sid]
            # value field: the one the model cell links to, else the record value, else its low end
            fm = _fm().get(f"{sh}!{cell}", "")
            m = re.search(re.escape(sid) + r"(_Lo|_Hi|_Spd|_ISL|_OSL|_MTP)?\b", fm)
            suf = m.group(1) if m and m.group(1) else ("" if wb[s2][f"C{rr}"].value is not None else "_Lo")
            put(ws, f"R{r}", f"={sid}{suf}", F_LINK)
        else:
            put(ws, f"R{r}", DASH)
        put(ws, f"AF{r}", f'=IF(OR({h}="{DASH}",{h}=""),"{DASH}",MATCH({h},IDX_SrcID,0))')
        put(ws, f"S{r}", f'=IF(ISNUMBER($AF{r}),INDEX(IDX_SrcStat,$AF{r}),"{DASH}")')
        put(ws, f"T{r}", f'=IF(ISNUMBER($AF{r}),INDEX(IDX_SrcGrade,$AF{r}),"{DASH}")')
        if single:
            v = wb[sh][cell].value
            st = ("連結 SRC" if "SRC_" in v else "公式") if isinstance(v, str) and v.startswith("=") else ("藍字（常數）" if v is not None else "空白")
        else:
            st = DASH
        put(ws, f"U{r}", st)
        F = f"$F{r}"; G = f"$G{r}"
        put(ws, f"V{r}", f'=IF(AND({F}="原始數據",$U{r}<>"連結 SRC"),1,0)', fmt="0")
        put(ws, f"W{r}", f'=IF(AND({F}="Analogy",OR({h}="{DASH}",{h}="")),1,0)', fmt="0")
        put(ws, f"X{r}", f'=IF(AND(OR({F}="Analogy",{F}="Assumed"),{G}<>"低",{G}<>"高",NOT(ISNUMBER($K{r})),NOT(ISNUMBER($L{r})),OR($M{r}="{DASH}",$M{r}="")),1,0)', fmt="0")
        put(ws, f"Y{r}", f'=IF(AND(OR({F}="Analogy",{F}="Assumed",{F}="Derived（待改公式）",{F}="Derived（公式）"),OR($N{r}="{DASH}",$N{r}="",$N{r}=0)),1,0)', fmt="0")
        put(ws, f"Z{r}", f'=IF(AND({F}="Decision",OR($J{r}="{DASH}",$J{r}="")),1,0)', fmt="0")
        put(ws, f"AA{r}", f'=IF(OR($J{r}="{DASH}",$J{r}=""),0,IF(COUNTIF(Decisions!$A$5:$A$300,$J{r})=0,1,0))', fmt="0")
        put(ws, f"AB{r}", f'=IF(AND({F}="原始數據",$S{r}<>"Active"),1,0)', fmt="0")
        put(ws, f"AC{r}", f'=IF(AND(ISNUMBER($K{r}),ISNUMBER($L{r}),ISNUMBER($Q{r})),IF(OR($Q{r}<MIN($K{r},$L{r}),$Q{r}>MAX($K{r},$L{r})),1,0),0)', fmt="0")
        put(ws, f"AD{r}", f'=IF(AND(ISNUMBER($T{r}),$P{r}="高"),IF($T{r}=3,1,0),0)', fmt="0")
        put(ws, f"AE{r}", f'=IF($S{r}="不存在",1,0)', fmt="0")      # v5.13: all scopes (slice-two SRC sheets exist)
        if st == "藍字（常數）" and ws[f"F{r}"].value == "原始數據": static_raw_hard += 1
    return n, static_raw_hard, n_upd, n_c


# ---------------------------------------------------------------- 7. L1 (builder-owned; every value a live formula)
L1_HDR = ["L1_ID", "指標或問題", "條件", "基準值", "低", "高", "單位", "區間的定義", "讀法", "主要驅動", "最弱輸入的標記",
          "外部 SRC_ID", "外部值（低）", "外部值（高）", "推算 ÷ 外部", "判讀", "具名範圍", "所在頁與列", "未能回答的部分"]
GENCOL = {"Hopper": 1, "GB200": 4, "GB300": 7, "VR200": 10}     # Interface column index of the low-cost column (base = +1)
COST_RNG = "成本角落情境（低成本／高成本欄）"
UTIL_RNG = "利用率 40–80%（Sens_Rev 情境值；J6 區間）"

def _rows_l1():
    R = []
    for g, c0 in GENCOL.items():
        ext = ("SRC_DC_010", "=SRC_DC_010_Lo", "=SRC_DC_010_Hi") if g == "VR200" else (DASH, None, None)
        R.append((f"CapexGW_{g}", f"每 GW 資本支出（{g}）", "基準功率情境；含廠房", f"=INDEX(IF_CapexTotal,1,{c0+1})",
                  f"=INDEX(IF_CapexTotal,1,{c0})", f"=INDEX(IF_CapexTotal,1,{c0+2})", "$B/GW", COST_RNG,
                  "IT 設備＋廠房；GW＝IT 關鍵電力（D1）", "機架價格、配電設計功率（Block 1 敏感度）", "機架價格：3 級、Interested-party",
                  *ext, "IF_CapexTotal", "Interface 第 11 列", "外部 $50–60B 的 GW 口徑未明（IT 或設施）" if g == "VR200" else DASH))
    R.append(("NvContentGW_VR200", "每 GW NVIDIA 內容（機架，VR200）", "基準功率情境", "=DC_Cost!M19/CTL_GW", "=DC_Cost!L19/CTL_GW",
              "=DC_Cost!N19/CTL_GW", "$B/GW", COST_RNG, "機架取得價 × 每 GW 機架數", "機架價格、配電設計功率", "機架價格：3 級、Interested-party",
              "SRC_DC_011", "=SRC_DC_011", "=SRC_DC_011", "DC_Cost!M19（÷ CTL_GW）", "DC_Cost 第 19 列", "外部值 GW 口徑未明"))
    R.append(("FacCapexMW", "每 MW IT 廠房資本支出", "基準成本情境；與世代無關", "=DC_Cost!M36", "=DC_Cost!L36", "=DC_Cost!N36", "$M/MW",
              COST_RNG, "建物、電力、機械、光纖與預備費（不含 IT 設備）", "電力設施、建物分項（Assumed）", "分項皆 Assumed",
              "SRC_DC_003", "=SRC_DC_003*(1+SRC_DC_004_Lo)", "=SRC_DC_003*(1+SRC_DC_004_Hi)", "DC_Cost!M36", "DC_Cost 第 36 列",
              "外部：JLL 全球平均 × 液冷溢價 7–10%；JLL 的 MW 口徑（IT 或設施）待查"))
    for g, c0 in GENCOL.items():
        R.append((f"HoldEconGW_{g}", f"每 GW 年經濟持有成本（{g}）", "資本回收年金（WACC 10%）＋營運費用", f"=INDEX(IF_HoldEcon,1,{c0+1})",
                  f"=INDEX(IF_HoldEcon,1,{c0})", f"=INDEX(IF_HoldEcon,1,{c0+2})", "$B/GW/年", COST_RNG,
                  "資本回收占 77–81%，電費 4–7%", "IT 折舊年限、機架價格", "機架價格：3 級", DASH, None, None,
                  "IF_HoldEcon", "Interface 第 13 列", DASH))
    for g, c0 in GENCOL.items():
        ext = {"GB200": ("SRC_DC_012", "=SRC_DC_012", "=SRC_DC_012"), "GB300": ("SRC_DC_013", "=SRC_DC_013", "=SRC_DC_013")}.get(g, (DASH, None, None))
        R.append((f"GPUhr_{g}", f"每 GPU 小時持有成本 — 經濟（{g}）", "100% 時數；不含利潤", f"=INDEX(IF_GPUhrEcon,1,{c0+1})",
                  f"=INDEX(IF_GPUhrEcon,1,{c0})", f"=INDEX(IF_GPUhrEcon,1,{c0+2})", "$/GPU-hr", COST_RNG,
                  "每 GW 年持有成本 ÷（GPU 數 × 8,760）", "機架價格、IT 折舊年限、WACC", "機架價格：3 級", *ext, "IF_GPUhrEcon", "Interface 第 14 列",
                  "外部為 SemiAnalysis TCO（假設未公開，須第二來源）；CoreWeave 公開價對照於 v5.12（SRC_Price）" if ext[0] != DASH else DASH))
    for tier in ("Luna", "Sol", "Astra"):
        for g in ("VR200", "GB300"):
            c0 = GENCOL[g]
            R.append((f"CostDec_{tier}_{g}", f"decode（含思考）每 M token 成本 — {tier}（{g}）", "經濟口徑、100% 利用率、SLO 下", f"=INDEX(IF_CostDec_{tier},1,{c0+1})",
                      f"=INDEX(IF_CostDec_{tier},1,{c0})", f"=INDEX(IF_CostDec_{tier},1,{c0+2})", "$/M", COST_RNG,
                      "基準利用率版＝100% 版 ÷ 利用率（IF_Util）", "η_d、每層延遲、SLO、生產折減", "生產折減 1.0：Assumed（K11）；VR200 η_d：Analogy",
                      DASH, None, None, f"IF_CostDec_{tier}", "Interface Block 2 產出", "VR200 無實測，η_d 沿用 GB300" if g == "VR200" else DASH))
    R.append(("CostTotTokIX_GB300", "GB300 每 M 總 token 成本（以 InferenceX 72 tok/s 實測產出計）", "本模型 TCO × 外部實測產出；8K/1K",
              "=DC_Cost!J58/(Calib!C17*3600)*1000000", "=DC_Cost!I58/(Calib!C17*3600)*1000000", "=DC_Cost!K58/(Calib!C17*3600)*1000000", "$/M 總 token",
              COST_RNG, "差異只來自每 GPU 小時 TCO（產出相同）", "每 GPU 小時持有成本", "InferenceX：3 級、Interested-party",
              "SRC_PERF_011", "=SRC_PERF_011", "=SRC_PERF_011", DASH, "Checks 第 19 列（同式）", "SemiAnalysis TCO 假設未公開"))
    for tier, c in (("Luna", "IF_RevGW_Luna"), ("Sol", "IF_RevGW_Sol"), ("Astra", "IF_RevGW_Astra"), ("機隊", "IF_RevGWFleet")):
        key = "RevGW_" + ("Fleet" if tier == "機隊" else tier) + "_VR200"
        lab = f"每 GW 理論營收（理想上限）— {tier}（VR200）" if tier != "機隊" else "1 GW 參考機隊付費營收（理想上限，VR200）"
        # v5.18 (engineering): the ratio is grouped, D*(u/IF_Util), so that u = IF_Util gives F = D exactly; D*u/IF_Util can land 1 ulp off D
        # (scenario e_util_08), which the H3 strict comparison counted for the engine but not for LibreOffice. Same math, values equal to 1e-15.
        R.append((key, lab, "OpenAI 有效單價、基準成本、基準利用率", f"=INDEX({c},1,11)", f"=INDEX({c},1,11)*(Sens_Rev!$B$6/IF_Util)",
                  f"=INDEX({c},1,11)*(Sens_Rev!$B$7/IF_Util)", "$B/GW/年", UTIL_RNG, "單一層級滿載的上限；實際營收（需求、市占）在下游",
                  "利用率、折扣、快取命中 χ（Sens_Rev）", "利用率 60%：Assumed（K11）；生產折減 1.0：Assumed",
                  DASH, None, None, c, "Interface D 節", "需求與市占不在第 0 層（D7）"))
    # ---- v5.13 (D): external comparisons moved from Checks (G9); each row's external columns link SRC
    R.append(("GPUhr_GB200_vsCW", "每 GPU 小時持有成本 — 經濟（GB200）對 CoreWeave 隨需牌價", "100% 時數；不含利潤",
              "=INDEX(IF_GPUhrEcon,1,5)", "=INDEX(IF_GPUhrEcon,1,4)", "=INDEX(IF_GPUhrEcon,1,6)", "$/GPU-hr", COST_RNG,
              "隨需牌價約為持有成本 5 倍：含利潤、閒置與風險溢價；長約價通常較低", "機架價格、IT 折舊年限、WACC", "牌價：3 級（spheron 轉述 CoreWeave 價目）",
              "SRC_PRC_001", "=SRC_PRC_001/4", "=SRC_PRC_001/4", "IF_GPUhrEcon", "Interface 第 14 列；Checks 第 9 列", "長約價未公開"))
    R.append(("GPUhr_GB300_vsBE", "每 GPU 小時持有成本 — 經濟（GB300）對新雲損益兩平租金", "100% 時數；不含利潤",
              "=INDEX(IF_GPUhrEcon,1,8)", "=INDEX(IF_GPUhrEcon,1,7)", "=INDEX(IF_GPUhrEcon,1,9)", "$/GPU-hr", COST_RNG,
              "外部為第三方推算的租金門檻（85% 利用率、9.12% 資金成本、6 年）；本模型為 100% 時數的持有成本", "機架價格、IT 折舊年限、WACC",
              "外部：2 級、機型未明", "SRC_PRC_002", "=SRC_PRC_002", "=SRC_PRC_002", "IF_GPUhrEcon", "Interface 第 14 列；Checks 第 10 列", "外部值機型未明"))
    R.append(("RLshare_Sol_VR200", "RL ÷ 預訓練 GPU 小時（Sol，VR200）", "J9 基準；非同步 RL（T10）", "=Training!$M$78", "=Sens_Train!$O$78",
              "=Sens_Train!$Q$78", "x", "RL rollout token ×0.3–×3（Sens_Train 情境倍數）", "後訓練算力的主要追蹤指標：RL 用掉的 GPU 小時相對預訓練的倍數",
              "rollout token（J9）、rollout 效率、RL trainer MFU", "rollout token：Assumed（J9 校準值）", "SRC_MOD_034", "=SRC_MOD_034/SRC_MOD_040",
              "=SRC_MOD_034/SRC_MOD_040", "TRN_RLRatioH", "Training 第 78 列；Checks 第 34 列", "外部為 DeepSeek R1（2025-01）÷ V3 預訓練，代表較早期的 RL 規模"))
    R.append(("RLshare_Astra_VR200", "RL ÷ 預訓練 GPU 小時（Astra，VR200）", "J9 基準；非同步 RL（T10）", "=Training!$N$78", "=Sens_Train!$P$78",
              "=Sens_Train!$R$78", "x", "RL rollout token ×0.3–×3（Sens_Train 情境倍數）", "同上；Astra 基準校到約 1（RL 與預訓練同量級）",
              "rollout token（J9）、rollout 效率、RL trainer MFU", "rollout token：Assumed（J9 校準值）", "SRC_MOD_036", "=SRC_MOD_036", "=SRC_MOD_036",
              "TRN_RLRatioH", "Training 第 78 列；Checks 第 33 列", "外部為 xAI 宣稱 Grok 4 RL 達「預訓練規模」，口徑（GPU 小時或 FLOPs）不明"))
    R.append(("FinalTrainShare", "最終訓練 ÷ 研發計畫（GPU 小時）", "J10 研發倍數基準 8", "=1/Train_In!$C$13", "=1/Sens_Train!$AE$119",
              "=1/Sens_Train!$AC$119", "%", "研發倍數 4.4–10.4（Sens_Train 情境）", "研發中花在最終訓練的比例；其餘為實驗、消融與失敗嘗試",
              "研發倍數（J10）", "研發倍數：Analogy（由外部區間設定，故本列只驗算一致）", "SRC_DEM_001",
              "=MIN(SRC_DEM_001,SRC_DEM_002,SRC_DEM_003)", "=MAX(SRC_DEM_001,SRC_DEM_002,SRC_DEM_003)", DASH, "Train_In 第 13 列；Checks 第 37 列",
              "外部為支出口徑（Epoch 推估三家），本模型為 GPU 小時口徑"))
    R.append(("RevGWFleet_OAI2025", "OpenAI 2025 對帳：每 GW 機隊付費營收（Hopper／GB200 加權）", ("OpenAI 有效單價（2026 快照）；Hopper 占機隊 57%（Checks 對帳常數，Analogy；v5.18 D4，區間 50–60%）" if v518.step_on("D4") else "OpenAI 有效單價（2026 快照）；Hopper 占機隊 60%（Checks 對帳常數，Assumed）"),
              "=Checks!$B${rev}", "=Checks!$B${rev}", "=Checks!$B${rev}", "$B/GW/年", "無區間（單一對帳值）", "同量級即機隊配置與單價可閉合；2025 實際單價高於 2026 快照",
              "OpenAI 有效單價、機隊配置（K7）", ("Hopper 占機隊：Analogy（Epoch 轉述 NVIDIA 出貨 4/7；全球口徑）" if v518.step_on("D4") else "Hopper 占機隊：Assumed"), "SRC_DEM_007", "=SRC_DEM_007/((SRC_DEM_008+SRC_DEM_009)/2)",
              "=SRC_DEM_007/((SRC_DEM_008+SRC_DEM_009)/2)", DASH, "Checks 第 43 列", "外部 GW 口徑未明（D1）"))
    R.append(("PretrainFLOP_Astra", "Astra 預訓練算力", "J8 基準（啟用參數 × 預訓練 token）", "=Training!$N$31*Training!$N$34*1E21",
              "=Training!$N$31*Training!$N$34*1E21", "=Training!$N$31*Training!$N$34*1E21", "FLOP", "無區間（J8 未結；token 區間見 Gov_Map Train_In E25）",
              "訓練 FLOPs/token × token；與前沿錨點 2e26–2e27 的差距即 J8 缺口", "Astra 啟用參數、預訓練 token", "Astra 架構：Assumed", "SRC_MOD_033",
              "=SRC_MOD_033", "=SRC_MOD_033", DASH, "Training 第 31、34 列；Checks 第 38 列", "外部為 Grok-3 的 Epoch 估計；GPT-6 Astra 實際算力未揭露"))
    from block6 import l1_rows_b6            # v5.15: Answers 1–9 and external comparisons (Block 6)
    R += l1_rows_b6(R, DASH, COST_RNG, UTIL_RNG)
    return R

def l1_sheet(wb):
    global _REV_ROW
    _REV_ROW = _row_by(wb["Checks"], "A", "OpenAI 2025 對帳：機隊付費營收（Hopper／GB200 加權、OpenAI 有效單價）")
    if "L1" in wb.sheetnames: del wb["L1"]
    ws = wb.create_sheet("L1")
    title(ws, "L1 — 第 1 層常用推算值（G9；即時公式、不貼值；附條件、區間與外部對照）",
          "下游取標準推算值時引用 L1_ 名稱；完整構件仍在 Interface（IF_）。外部值一律連結 SRC。判讀：外部為區間時看是否落在區間內；外部為單一值時以 ±20% 判讀。"
          "v5.13 D 起 Checks 的外部比對移入本頁（第 30 列以下）；v5.15 補入 Block 6 的 9 題（L1_Ans1–9）與 3 列外部對照；v5.16 補毛利率兩列與 GPU 小時口徑一列。欄位約定（v5.16）：D＝基準、E＝低、F＝高；無區間時 E＝F＝D，H 欄寫「無區間」；不同口徑或對照量一律另列一列，不放在 E／F。")
    widths = [22, 38, 26, 10, 10, 10, 10, 24, 30, 24, 26, 12, 10, 10, 9, 14, 18, 18, 30]
    for i, h in enumerate(L1_HDR):
        put(ws, f"{L(i+1)}4", h, F_BOLD, wrap=True); ws.column_dimensions[L(i+1)].width = widths[i]
    r = 5; nonformula = 0
    rows_at = {}
    for row in _rows_l1():
        row = tuple(x.replace("{rev}", str(_REV_ROW)) if isinstance(x, str) else x for x in row)
        key, lab, cond, base, lo, hi, unit, rdef, read, drv, weak, sid, elo, ehi, nmref, where, gap = row
        rows_at[key] = r
        vals = [f"L1_{key}", lab, cond, base, lo, hi, unit, rdef, read, drv, weak, sid, elo, ehi, None, None, f"L1_{key}", where, gap]
        for i, v in enumerate(vals):
            if i == 14 or i == 15: continue
            fmt = "#,##0.000" if i in (3, 4, 5, 12, 13) else None
            put(ws, f"{L(i+1)}{r}", v if v is not None else DASH, F_LINK if isinstance(v, str) and v.startswith("=") and i in (12, 13) else None,
                fmt=fmt, wrap=i in (1, 2, 7, 8, 9, 10, 18), fill=FILL_KEY if i == 3 else None)
        for c in "DEF":
            if not str(ws[f"{c}{r}"].value).startswith("="): nonformula += 1
        # v5.14: D (base value) may be text in some scenarios (e.g. L1_RevGWFleet_OAI2025 reads Checks text "SLO 不可達");
        # guard D as well as M/N so O/P never return an error value. Base-case values are unchanged.
        put(ws, f"O{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(M{r}),ISNUMBER(N{r})),D{r}/((M{r}+N{r})/2),"{DASH}")', fmt="0.00")
        put(ws, f"P{r}", f'=IF(ISNUMBER(D{r}),IF(AND(ISNUMBER(M{r}),ISNUMBER(N{r})),IF(M{r}=N{r},IF(ABS(D{r}/M{r}-1)<=0.2,"±20% 內","差距 >20%"),'
                         f'IF(D{r}<M{r},"低於外部區間",IF(D{r}>N{r},"高於外部區間","落在外部區間"))),"無外部對照"),"推算值非數字（本情境）")')
        _nm(wb, f"L1_{key}", f"L1!$D${r}"); _nm(wb, f"L1_{key}_Lo", f"L1!$E${r}"); _nm(wb, f"L1_{key}_Hi", f"L1!$F${r}")
        r += 1
    ws.freeze_panes = "C5"
    return r - 5, nonformula, rows_at


# ---------------------------------------------------------------- 8. Checks governance section (builder-owned)
def checks_gov(wb, l1_rows, l1_nonformula, static_raw_hard, idx_spans):
    ws = wb["Checks"]
    terms = [f'COUNTIF({sh}!$A${e+1}:$A${SRC_LAST},"<>{DASH}")-COUNTIF({sh}!$A${e+1}:$A${SRC_LAST},"")'
             for sh, e in idx_spans.items() if e < SRC_LAST]
    idx_cov = "=" + "+".join(terms) if terms else 0
    r = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    section(ws, r, "G. 治理檢查（Stage 1 切片一，v5.11；v5.12 加 E13；v5.13 SRC 頁 4→8；規劃書第 5 節）：ERROR 合計必須為 0 才可合併", 6); r += 1
    for i, h in enumerate(["編號", "檢查", "等級", "筆數", "範圍與算法"]):
        put(ws, f"{L(i+1)}{r}", h, F_BOLD)
    r += 1
    GM = lambda col: f"Gov_Map!${col}$5:${col}${GM_LAST}"
    srcsum = lambda col, crit: "+".join(f'COUNTIF({s}!${col}$5:${col}${SRC_LAST},"{crit}")' for s in SRC_SHEETS)
    rows = [
      ("E1", "模型頁原始數據寫死（未連結 SRC）", "ERROR", f'=COUNTIF({GM("V")},1)', "Gov_Map 類別＝原始數據、建置時狀態≠連結 SRC（G0-2 保留與切片二另計）"),
      ("E2", "Analogy 缺可比對象 SRC_ID", "ERROR", f'=COUNTIF({GM("W")},1)', "Gov_Map"),
      ("E3", "Analogy／Assumed 缺區間", "ERROR", f'=COUNTIF({GM("X")},1)', "Gov_Map；端點格（區間角色＝低／高）不另檢查"),
      ("E4", "Analogy／Assumed／Derived 缺理由", "ERROR", f'=COUNTIF({GM("Y")},1)', "Gov_Map"),
      ("E5", "Decision 格缺決策 ID", "ERROR", f'=COUNTIF({GM("Z")},1)', "Gov_Map"),
      ("E6", "決策 ID 不在 Decisions 頁", "ERROR", f'=COUNTIF({GM("AA")},1)', "Gov_Map × Decisions"),
      ("E7", "同指標、同口徑、同對象有兩筆以上 Active", "ERROR", "=" + srcsum("X", ">1"), "SRC_* X 欄（逐筆計數，重複者每筆各計 1）"),
      ("E8", "Active SRC 缺 Evidence ID", "ERROR", "=" + srcsum("Y", "1"), "SRC_* Y 欄"),
      ("E9", "模型連結的 SRC 紀錄不是 Active", "ERROR", f'=COUNTIF({GM("AB")},1)', "Gov_Map 類別＝原始數據"),
      ("E10", "基準值不在低／高之間", "ERROR", f'=COUNTIF({GM("AC")},1)', "Gov_Map 數值區間（低、高不分方向）"),
      ("E11", "L1 數值欄不是公式（貼值）", "ERROR", l1_nonformula, "建置時靜態檢查（builder）"),
      ("E12", "引用的 SRC_ID 不存在", "ERROR", f'=COUNTIF({GM("AE")},1)', "Gov_Map"),
      ("E13", "SRC 紀錄超出 SRC_Index 範圍（需重建）", "ERROR", idx_cov, "各 SRC 頁 A 欄在 SRC_Index 範圍之後、第 400 列之前的非空白格（v5.12）"),
      ("W1", "利害關係方 Active 紀錄缺第二來源", "WARN", "=" + srcsum("Z", "1"), "SRC_* Z 欄（SemiAnalysis 規則推廣；Stage 2 補）"),
      ("W2", "3 級紀錄被 CC 高段敏感度參數使用", "WARN", f'=COUNTIF({GM("AD")},1)', "Gov_Map（CC 第 10 輪分段）"),
      ("I1", "DB_Evidence 待判定", "INFO", '=COUNTIF(DB_Evidence!$L$5:$L$500,"待判定")', "DB_Evidence L 欄"),
      ("I2", "DB_Evidence 待 Andy", "INFO", '=COUNTIF(DB_Evidence!$L$5:$L$500,"待 Andy")', "DB_Evidence L 欄"),
      ("I3", "L1 落在外部區間外或差距 >20% 的列", "INFO",
       '=COUNTIF(L1!$P$5:$P$200,"低於外部區間")+COUNTIF(L1!$P$5:$P$200,"高於外部區間")+COUNTIF(L1!$P$5:$P$200,"差距 >20%")', "L1 P 欄"),
      ("I4", "原始數據 G0-2 保留藍字（Stage 2 佇列）", "INFO", f'=COUNTIF({GM("F")},"原始數據（G0-2 保留）")', "Gov_Map"),
      ("I5", "Derived 寫死待改公式", "INFO", f'=COUNTIF({GM("F")},"Derived（待改公式）")', "Gov_Map（Arch 21–23 已於 v5.13 D 改公式；餘為 Train_In C44 等）"),
      ("I6", "結構選擇（無數值區間）", "INFO", f'=COUNTIF({GM("M")},"結構選擇（無數值區間）")', "Gov_Map"),
      ("I7", "原始數據待連結（尚無 SRC 紀錄）", "INFO", f'=COUNTIF({GM("F")},"原始數據（切片二）")', "Gov_Map（v5.13 起：原始數據尚無 SRC 紀錄者，例如 Tech_Registry 採用數、Cap_In 尖峰離峰屬性；Stage 2 補紀錄）"),
      ("I8", "標記變更（v5.11 起；Analogy→Assumed 等）", "INFO", f'=COUNTIF({GM("O")},"<>{DASH}")-COUNTIF({GM("O")},"")', "Gov_Map O 欄；請 Andy 過目"),
    ]
    first = r; err_rows = []; warn_rows = []; info_rows = []
    for code, lab, lvl, f, note in rows:
        put(ws, f"A{r}", code); put(ws, f"B{r}", lab); put(ws, f"C{r}", lvl, F_BOLD if lvl == "ERROR" else None)
        put(ws, f"D{r}", f, F_CALC, fmt="0"); put(ws, f"E{r}", note, F_NOTE)
        {"ERROR": err_rows, "WARN": warn_rows, "INFO": info_rows}[lvl].append(r); r += 1
    put(ws, f"B{r}", "ERROR 合計（CI 讀取 GOV_Errors）", F_BOLD); put(ws, f"D{r}", "=" + "+".join(f"D{x}" for x in err_rows), F_BOLD, fmt="0", fill=FILL_KEY)
    _nm(wb, "GOV_Errors", f"Checks!$D${r}"); r += 1
    put(ws, f"B{r}", "WARN 合計", F_BOLD); put(ws, f"D{r}", "=" + "+".join(f"D{x}" for x in warn_rows), fmt="0"); _nm(wb, "GOV_Warnings", f"Checks!$D${r}"); r += 1
    put(ws, f"B{r}", "INFO 合計", F_BOLD); put(ws, f"D{r}", "=" + "+".join(f"D{x}" for x in info_rows), fmt="0"); _nm(wb, "GOV_Info", f"Checks!$D${r}"); r += 1
    put(ws, f"B{r}", f"建置時靜態檢查：Gov_Map 類別＝原始數據且為藍字常數的格數＝{static_raw_hard}（應為 0；與 E1 一致）", F_NOTE); r += 1
    _nm(wb, "GOV_Table", f"Checks!$A${first}:$D${r-5}")
    return first


# ---------------------------------------------------------------- 9. Block 1 Checks rows 4–8: external refs link to SRC (no literals)
def checks_block1(wb):
    ck = wb["Checks"]
    ck["C4"].value = '=TEXT(SRC_DC_010_Lo,"0")&"–"&TEXT(SRC_DC_010_Hi,"0")'
    ck["E4"].value = '=IF(AND(B4>=SRC_DC_010_Lo,B4<=SRC_DC_010_Hi),"落在區間",IF(B4<SRC_DC_010_Lo,"低於參照","高於參照"))'
    ck["C5"].value = "=SRC_DC_011"
    ck["E5"].value = '=IF(ABS(B5-C5)/C5<=0.2,"±20% 內","差距 >20%")'
    ck["C6"].value = '=TEXT(SRC_DC_003,"0.0")&" × "&TEXT(1+SRC_DC_004_Lo,"0.00")&"–"&TEXT(1+SRC_DC_004_Hi,"0.00")'
    ck["E6"].value = '=IF(AND(B6>=SRC_DC_003*(1+SRC_DC_004_Lo)*0.9,B6<=SRC_DC_003*(1+SRC_DC_004_Hi)*1.3),"合理區間","需檢查")'
    ck["C7"].value = "=SRC_DC_012"; ck["C8"].value = "=SRC_DC_013"
    for ref in ("C4", "C5", "C6", "C7", "C8"): ck[ref].font = F_LINK
    for r, s in ((4, "SRC_DC_010"), (5, "SRC_DC_011"), (6, "SRC_DC_003、SRC_DC_004"), (7, "SRC_DC_012"), (8, "SRC_DC_013")):
        v = str(ck[f"F{r}"].value or "")
        if "SRC_" not in v: ck[f"F{r}"].value = f"{v}［{s}］"
    ck["F19"].value = str(ck["F19"].value or "") + ("［SRC_PERF_011］" if "SRC_" not in str(ck["F19"].value or "") else "")
    ck["C19"].value = "=SRC_PERF_011"; ck["C19"].font = F_LINK


# ---------------------------------------------------------------- 9a. v5.13 (B): Checks external references of slice two link to SRC
def _row_by(ws, col, label):
    for r in range(1, ws.max_row + 1):
        if ws[f"{col}{r}"].value == label: return r
    raise KeyError(f"Checks: {col} label not found: {label}")

def checks_slice2(wb):
    ck = wb["Checks"]; log = []
    def setf(ref, f, tag=None):
        if ck[ref].value != f: log.append(f"Checks!{ref}: {ck[ref].value!r} -> {f}")
        ck[ref].value = f; ck[ref].font = F_LINK
        if tag:
            fr = "F" + ref[1:]
            v = str(ck[fr].value or "")
            if "SRC_" not in v: ck[fr].value = f"{v}［{tag}］"
    setf("C9", "=SRC_PRC_001/4", "SRC_PRC_001；÷4＝每執行個體 GPU 數")           # 4-GPU instance price → $/GPU-hr (unit conversion)
    setf("C10", "=SRC_PRC_002", "SRC_PRC_002")
    r = _row_by(ck, "A", "最終訓練 ÷ 研發計畫")
    lo, hi = "MIN(SRC_DEM_001,SRC_DEM_002,SRC_DEM_003)", "MAX(SRC_DEM_001,SRC_DEM_002,SRC_DEM_003)"
    setf(f"C{r}", f'=TEXT({lo}*100,"0.0")&"%–"&TEXT({hi}*100,"0.0")&"%"', "SRC_DEM_001–003")
    for lab, sid in (("OpenAI 2025 營收 $B", "SRC_DEM_007"), ("2024 年底 GW", "SRC_DEM_008"), ("2025 年底 GW", "SRC_DEM_009")):
        rr = _row_by(ck, "H", lab)
        if ck[f"I{rr}"].value != f"={sid}": log.append(f"Checks!I{rr}: {ck[f'I{rr}'].value!r} -> ={sid}")
        ck[f"I{rr}"].value = f"={sid}"; ck[f"I{rr}"].font = F_LINK
    return log


SAMPLE_OUT = (("樣本外：GB300 SGLang＋MTP 50 tok/s", "SRC_PERF_004"), ("樣本外：GB300 vLLM 無 MTP 27 tok/s", "SRC_PERF_001"),
              ("樣本外：GB200 vLLM 無 MTP 27 tok/s", "SRC_PERF_002"), ("樣本外：GB200／GB300 110 tok/s", "SRC_PERF_014"))
CK_TO_L1 = (("GB200 市場牌價對照", "GPUhr_GB200_vsCW"), ("新雲損益兩平參照", "GPUhr_GB300_vsBE"),
            ("RL ÷ 預訓練 GPU 小時：VR200 Astra", "RLshare_Astra_VR200"), ("RL ÷ 預訓練 GPU 小時：VR200 Sol", "RLshare_Sol_VR200"),
            ("OpenAI 2025 對帳：機隊付費營收（Hopper／GB200 加權、OpenAI 有效單價）", "RevGWFleet_OAI2025"))

def checks_to_l1(wb, l1_at):
    """v5.13 (D): Checks external references read the L1 external columns (one place per comparison, G9); sample-out literals link SRC_PERF."""
    ck = wb["Checks"]; log = []
    def setf(ref, f, note):
        if ck[ref].value != f: log.append(f"Checks!{ref}: {ck[ref].value!r} -> {f}")
        ck[ref].value = f; ck[ref].font = F_LINK
        fr = "F" + ref[1:]; v = str(ck[fr].value or "")
        if note not in v: ck[fr].value = f"{v}［{note}］"
    for lab, sid in SAMPLE_OUT:
        r = _row_by(ck, "A", lab); setf(f"C{r}", f"={sid}", sid); ck[f"C{r}"].number_format = "#,##0"
    for lab, key in CK_TO_L1:
        r = _row_by(ck, "A", lab); setf(f"C{r}", f"=L1!$M${l1_at[key]}", f"→ L1_{key}")
    r = _row_by(ck, "A", "最終訓練 ÷ 研發計畫"); m = l1_at["FinalTrainShare"]
    setf(f"C{r}", f'=TEXT(L1!$M${m}*100,"0.0")&"%–"&TEXT(L1!$N${m}*100,"0.0")&"%"', "→ L1_FinalTrainShare")
    r = _row_by(ck, "A", "前沿錨點中值（5e26）反推 Astra 預訓練 token")
    v = str(ck[f"F{r}"].value or ""); ck[f"F{r}"].value = v if "L1_" in v else f"{v}［Astra 預訓練算力與 SRC_MOD_033 的比較 → L1_PretrainFLOP_Astra］"
    return log


def gov_all(wb):
    made = src_sheets(wb)
    appended = src_append(wb)
    n_src_upd, src_upd_log = v518.src_update(wb)      # v5.18: SRC grades, values, notes (old-value guards)
    n_src_names, idx = src_refresh(wb)
    ev_added = evidence_upgrade(wb)
    dec_made = decisions_sheet(wb)
    dec_made2 = dec_append(wb)
    dec_upd = dec_update(wb)
    n_fm, fm_log = apply_formula_map(wb)
    n_f14 = f14(wb)
    checks_block1(wb)
    ck2_log = checks_slice2(wb)
    n_gm, hard, n_upd, n_c = gov_map(wb, idx)
    SI = src_index(wb)
    n_l1, nonf, l1_at = l1_sheet(wb)
    ck3_log = checks_to_l1(wb, l1_at)
    checks_gov(wb, n_l1, nonf, hard, SI["spans"])
    return dict(src_made=made, src_appended=appended, src_updated_v518=n_src_upd, gm_updated=n_upd, gm_appended_c=n_c, dec_updated=dec_upd, checks_slice2=len(ck2_log), src_names=n_src_names, src_records=len(idx), evidence_added=ev_added, decisions_made=dec_made, decisions_appended=dec_made2,
                src_index_rows=SI["rows"], src_index_src_rows=SI["src_rows"], src_index_spans=SI["spans"],
                formula_map_changed=n_fm, f14_changed=n_f14, gov_rows=n_gm, gov_raw_hardcoded=hard, l1_rows=n_l1, l1_nonformula=nonf,
                fm_log=fm_log + ck2_log + ck3_log + [f"v518 {x}" for x in src_upd_log])
