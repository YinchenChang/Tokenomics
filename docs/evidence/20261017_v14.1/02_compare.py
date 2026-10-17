"""v14.1：新舊兩版（皆為 LibreOffice 重算存檔）逐名稱、逐格比較。
用法：python3 -I 02_compare.py <舊版.xlsx> <新版.xlsx> [輸出 json]    （從 repo 根目錄執行）
輸出：(1) 值有變動的具名範圍（全部名稱，不只契約前綴）舊→新；(2) 逐格差異依工作表與類別（值／公式文字）統計，及前 N 筆明細。"""
import sys, json, collections
import openpyxl
from openpyxl.utils import range_boundaries

old_p, new_p = sys.argv[1], sys.argv[2]
ov, nv = openpyxl.load_workbook(old_p, data_only=True), openpyxl.load_workbook(new_p, data_only=True)
of, nf = openpyxl.load_workbook(old_p), openpyxl.load_workbook(new_p)


def same(a, b):
    a = None if a == "" else a
    b = None if b == "" else b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= 1e-12 or abs(a - b) <= 1e-9 * max(abs(a), abs(b))
    return a == b


def cells(wb, n):
    sh, rg = wb.defined_names[n].attr_text.rsplit("!", 1)
    c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
    ws = wb[sh.strip("'")]
    return [(f"{sh.strip(chr(39))}!{ws.cell(r, c).coordinate}", ws.cell(r, c).value) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


out = {"names_changed": {}, "names_added": [], "names_removed": []}
for n in nv.defined_names:
    if n not in ov.defined_names:
        out["names_added"].append(n); continue
    a, b = cells(ov, n), cells(nv, n)
    if len(a) != len(b):
        out["names_changed"][n] = {"shape": [len(a), len(b)]}; continue
    d = [(ca, va, vb) for (ca, va), (_, vb) in zip(a, b) if not same(va, vb)]
    if d:
        out["names_changed"][n] = {"n": len(d), "cells": d[:12]}
out["names_removed"] = [n for n in ov.defined_names if n not in nv.defined_names]
print("== 值有變動的具名範圍", len(out["names_changed"]))
for n, v in sorted(out["names_changed"].items()):
    print(" ", n, v if "shape" in v else f'{v["n"]} 格', "" if "shape" in v else [(c, x if not isinstance(x, float) else round(x, 6), y if not isinstance(y, float) else round(y, 6)) for c, x, y in v["cells"][:3]])
print("== 新增名稱", out["names_added"], "移除名稱", out["names_removed"])

# cell level
diff = collections.defaultdict(lambda: collections.Counter())
detail = collections.defaultdict(list)
for ws in nv.worksheets:
    if ws.title not in ov.sheetnames:
        diff[ws.title]["new sheet"] += 1; continue
    o_v, n_v, o_f, n_f = ov[ws.title], ws, of[ws.title], nf[ws.title]
    mr, mc = max(o_v.max_row, n_v.max_row), max(o_v.max_column, n_v.max_column)
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            va, vb = o_v.cell(r, c).value, n_v.cell(r, c).value
            fa, fb = o_f.cell(r, c).value, n_f.cell(r, c).value
            k = None
            if not same(fa, fb) and isinstance(fa, str) and isinstance(fb, str) and (fa.startswith("=") or fb.startswith("=")):
                k = "公式文字"
            elif not same(va, vb):
                k = "值"
            if k:
                diff[ws.title][k] += 1
                if len(detail[ws.title]) < 400:
                    detail[ws.title].append((n_v.cell(r, c).coordinate, k, str(va)[:60], str(vb)[:60]))
out["cell_diff"] = {s: dict(c) for s, c in diff.items()}
out["cell_detail"] = detail
print("== 逐格差異（工作表：類別 格數）")
for s, c in diff.items():
    print(" ", s, dict(c))
if len(sys.argv) > 3:
    json.dump(out, open(sys.argv[3], "w"), ensure_ascii=False, indent=1, default=str)
