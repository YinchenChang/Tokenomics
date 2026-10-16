"""兩版活頁簿逐格比較（公式文字與重算後的值）。用法：python3 tools/cmp_models.py <舊.xlsx> <新.xlsx> <輸出.pkl>
兩個檔案都必須是 LibreOffice 重算存檔後的檔（含快取值）。輸出：差異清單與具名範圍值比較。"""
import pickle
import sys
import openpyxl


def load(path):
    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    cells = {}
    for ws in wf.worksheets:
        vs = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                fill = None
                if c.fill is not None and c.fill.fill_type == "solid":
                    fg = c.fill.fgColor
                    fill = str(fg.rgb) if fg.type == "rgb" else f"theme{fg.theme}"
                cells[(ws.title, c.coordinate)] = (c.value, vs[c.coordinate].value, fill)
    names = {n: dn.attr_text for n, dn in wf.defined_names.items()}
    return cells, names, wv


def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= 1e-12 or abs(a - b) <= 1e-9 * max(abs(a), abs(b))
    return a == b


def name_values(names, wv):
    out = {}
    for n, t in names.items():
        if "!" not in t:
            continue
        sh, rng = t.rsplit("!", 1)
        sh = sh.strip("'")
        if sh not in wv.sheetnames:
            out[n] = ("#SHEET_MISSING",)
            continue
        ws = wv[sh]
        try:
            r = ws[rng.replace("$", "")]
        except Exception:
            out[n] = ("#BAD_REF",)
            continue
        if not isinstance(r, tuple):
            out[n] = (r.value,)
        else:
            out[n] = tuple(c.value for row in r for c in (row if isinstance(row, tuple) else (row,)))
    return out


if __name__ == "__main__":
    old, new, outp = sys.argv[1:4]
    co, no, wo = load(old)
    cn, nn, wn = load(new)
    diffs = []
    for k in sorted(set(co) | set(cn)):
        a, b = co.get(k), cn.get(k)
        if a is None or b is None:
            diffs.append((k, a, b, "removed" if b is None else "added"))
            continue
        kind = []
        if a[0] != b[0]:
            kind.append("formula")
        if not same(a[1], b[1]):
            kind.append("value")
        if a[2] != b[2]:
            kind.append("fill")
        if kind:
            diffs.append((k, a, b, "+".join(kind)))
    vo, vn = name_values(no, wo), name_values(nn, wn)
    namediff = []
    for n in sorted(set(vo) | set(vn)):
        if n not in vo or n not in vn:
            namediff.append((n, "added" if n not in vo else "removed"))
        elif len(vo[n]) != len(vn[n]):
            namediff.append((n, "shape", no.get(n), nn.get(n)))
        elif not all(same(x, y) for x, y in zip(vo[n], vn[n])):
            namediff.append((n, "value", no.get(n), nn.get(n)))
    nform = lambda c: sum(1 for v in c.values() if isinstance(v[0], str) and v[0].startswith("="))
    pickle.dump(dict(diffs=diffs, namediff=namediff, nform=(nform(co), nform(cn)), nnames=(len(no), len(nn))), open(outp, "wb"))
    print("diff cells", len(diffs), "| name diffs", len(namediff), "| formula cells", nform(co), "->", nform(cn), "| names", len(no), "->", len(nn))
