# v5.30 (work order docs/workorders/20261008_v5.30.md r0): X15 definition corrections; no input value changes.
#   X15 (a) Interface J four-layer waterfall is a running product: _Prod = _Util x CTL_ProdDerate (token rows: linear; revenue rows keep
#           reading Interface G, which re-solves the SLO), _Life = _Prod x CTL_PriceLife x CTL_Monetize (revenue rows; H section untouched);
#           Front rows get _Prod／_Life formulas instead of "—"; token _Life = _Prod x 1 (unchanged). Self-check: (i) source equality for
#           _Util (revenue: D section) and token _100 (B section); (ii) monotonic _100 >= _Util >= _Prod >= _Life; (iii) G／H／Interface_Prod
#           comparisons removed. Checks K3 sums (i) + (ii).
#   X15 (b) L1_HoldEconMW_* = L1_HoldEconGW_* (no /1000): 1 $B/GW = 1 $M/MW; unit text stays "$M/MW/年"; low／high in step.
#   X15 (c) IFC_Use: every IFW_RevGW*／IFW_RevGWFleet*／IFW_RevGWFront* layer and IFW_TokGW_*_100 carries "；上限，不得作預測"
#           (static flag Y extended; Checks K2 must stay 0).
# Decisions X15a–c are appended only when the ID is absent (Excel-owned afterwards), like v518–v529.
import re
from openpyxl.utils import get_column_letter as L
from common import put, F_BOLD, F_NOTE, section
import v529

VERSION = "20261008_Tokenomics_v5.30"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-08"
V = "v5.30"
DASH = "—"
COLS = range(3, 18)                         # C:Q
ANDY = "Andy 2026-10-08「請直接做」（八家公司模型審視報告 company-models PR #36–#43 共同指出的 Tokenomics 端問題）"
TOL = "0.000000001"                         # 1e-9, relative to max(1, |upper layer|) in the monotonic check

REV, TOK, TAG, LAYERS = v529.REV, v529.TOK, v529.TAG, v529.LAYERS


# ================================================================== X15 (a): Interface J (running product) and self-check rows
def interface_j(wb):
    ws = wb["Interface"]
    start = v529._last_row(ws) + 2
    section(ws, start, "J. 四層瀑布並列（X14 (b) v5.29；X15 (a) v5.30 起逐層累乘）：_100（100% 口徑）→ _Util＝_100 × IF_Util → _Prod＝_Util × CTL_ProdDerate "
                       "→ _Life＝_Prod × CTL_PriceLife × CTL_Monetize，單調遞減；營收 _Prod 讀 G 節（重解 SLO 後的值）、token 產能 _Prod 為線性乘法；"
                       "H 節 _Life（不含折減）與 Interface_Prod（100% 口徑重解）保留原處作並列對照；自我檢查見本節末與 Checks K3", 17)
    r = start + 1; made = []; src_checks = []; groups = []; at = {}
    def base_row(name):
        sh, row = v529._name_row(wb, name); assert sh == "Interface", (name, sh); return row
    def fmt_of(row): return ws.cell(row, 3).number_format
    for base in REV + TOK:
        b = base_row(base); lab = TAG.sub("", ws.cell(b, 1).value); stem = base[3:]
        prod = f"{base}_Prod" if f"{base}_Prod" in wb.defined_names else None
        is_tok = base in TOK
        put(ws, f"A{r}", lab, F_BOLD); r += 1
        rows4 = []
        for suf, ltxt in LAYERS:
            name = f"IFW_{stem}_{suf}"
            note = ""; src = None
            u = at.get(f"IFW_{stem}_Util"); p = at.get(f"IFW_{stem}_Prod")
            if is_tok:
                if suf == "100": f = lambda X: f"={X}{b}"; src = ("Interface", b)
                elif suf == "Util": f = lambda X: f"=IF(ISNUMBER({X}{b}),{X}{b}*IF_Util,{X}{b})"; note = "（B 節 100% 列 × IF_Util；無 D 節對應）"
                elif suf == "Prod":
                    f = lambda X, u=u: f"=IF(ISNUMBER({X}{u}),{X}{u}*CTL_ProdDerate,{X}{u})"
                    note = "（＝Util 列 × CTL_ProdDerate，線性；Interface_Prod 同列為 100% 口徑重解值，保留原處作對照）"
                else: f = lambda X, p=p: f"=IF(ISNUMBER({X}{p}),{X}{p}*1,{X}{p})"; note = "（無 H 節對應：＝Prod 列 × 1；L、m 不作用於產能）"
            else:
                if suf == "100": f = lambda X: f"=IF(ISNUMBER({X}{b}),{X}{b}/IF_Util,{X}{b})"; note = "（D 節列 ÷ IF_Util）"
                elif suf == "Util": f = lambda X: f"={X}{b}"; src = ("Interface", b)
                elif suf == "Prod":
                    if prod:
                        g = base_row(prod); f = lambda X, g=g: f"={X}{g}"; note = f"（＝G 節第 {g} 列：重解 SLO 後的值）"
                    else:
                        f = lambda X, u=u: f"=IF(ISNUMBER({X}{u}),{X}{u}*CTL_ProdDerate,{X}{u})"; note = "（G 節無對應：＝Util 列 × CTL_ProdDerate）"
                else:
                    f = lambda X, p=p: f"=IF(ISNUMBER({X}{p}),{X}{p}*CTL_PriceLife*CTL_Monetize,{X}{p})"
                    note = "（＝Prod 列 × CTL_PriceLife × CTL_Monetize；H 節 _Life 不含折減，並列保留）"
            put(ws, f"A{r}", f"{lab}｜{ltxt}{note}　[{name}]"); put(ws, f"B{r}", ws.cell(b, 2).value)
            for c in COLS:
                X = L(c); put(ws, f"{X}{r}", f(X), fmt=fmt_of(b))
            made.append((name, f"Interface!$C${r}:$Q${r}")); at[name] = r; rows4.append(r)
            if src: src_checks.append((name, r, src))
            r += 1
        groups.append((f"IFW_{stem}", rows4))
    r += 1
    put(ws, f"A{r}", "自我檢查（Checks K3；v5.30 X15 (a)）：(i) IFW_ 列 − 來源列（營收 _Util＝D 節、token _100＝B 節）的絕對差 ≤ 1e-9（兩邊皆為文字時須相同）；"
                     "(ii) 每組單調：_100 ≥ _Util ≥ _Prod ≥ _Life（相鄰兩層皆為數值時才比，容差 1e-9 × max(1, 上層絕對值)；文字列略過）。0＝相符，1＝不符", F_BOLD); r += 1
    f0 = r
    for name, rr, (sh, sr) in src_checks:
        put(ws, f"A{r}", f"{name}：Interface 列 vs {sh} 第 {sr} 列", F_NOTE)
        for c in COLS:
            X = L(c); a, b = f"{X}{rr}", f"{X}{sr}"
            put(ws, f"{X}{r}", f"=IF(AND(ISNUMBER({a}),ISNUMBER({b})),IF(ABS({a}-{b})<=0.000000001,0,1),IF(AND(NOT(ISNUMBER({a})),NOT(ISNUMBER({b}))),IF({a}={b},0,1),1))", fmt="0")
        r += 1
    m0 = r
    for stem, (r100, rutil, rprod, rlife) in groups:
        put(ws, f"A{r}", f"{stem}：Interface 四層單調（第 {r100}–{rlife} 列：_100 ≥ _Util ≥ _Prod ≥ _Life）", F_NOTE)
        for c in COLS:
            X = L(c)
            pair = lambda hi, lo: (f"IF(ISNUMBER({X}{hi}),IF(ISNUMBER({X}{lo}),IF({X}{lo}-{X}{hi}>{TOL}*MAX(1,ABS({X}{hi})),1,0),0),0)")
            put(ws, f"{X}{r}", f"=IF({pair(r100, rutil)}+{pair(rutil, rprod)}+{pair(rprod, rlife)}>0,1,0)", fmt="0")
        r += 1
    f1 = r - 1
    for n, ref in made: v529._nm(wb, n, ref)
    return dict(start=start, made=made, flags=(f0, f1), mono=(m0, f1), src=(f0, m0 - 1), at=at)


# ================================================================== X15 (c): which Interface rows must carry "上限，不得作預測"
_TOK100 = re.compile(r"\[IFW_TokGW_\w+_100\]")


def should_cap(sec, label, layer):
    if sec == "B" and layer == "100%": return 1
    if sec == "D" and ("理論營收" in label or "付費服務營收" in label): return 1
    if sec == "J" and ("[IFW_RevGW" in label or _TOK100.search(label)): return 1      # IFW_RevGW_*, IFW_RevGWFleet*, IFW_RevGWFront* (all layers); IFW_TokGW_*_100
    return 0


# ================================================================== Checks K: K2／K3 wording (formulas built by v529.checks_k from the new flag range)
K2_LABEL = ("R 欄為 A 但 S 欄缺「上限」附註的列數：B 節 100% 產出列、D 節理論營收列，以及 J 節（v5.30 X15 (c)）IFW_RevGW*／IFW_RevGWFleet*／IFW_RevGWFront* 全部四層"
            "與 IFW_TokGW_*_100")
K3_LABEL = ("四層瀑布自我檢查（v5.30 X15 (a)）：(i) 營收 _Util／token _100 不等於來源列（D 節／B 節）的格數＋(ii) 單調違反格數（_100 ≥ _Util ≥ _Prod ≥ _Life；"
            "文字列略過）；原「＝G／H 節、Interface_Prod」比對已移除")


def checks_k(wb, ifc, ifj, lb):
    at = v529.checks_k(wb, ifc, ifj, lb)
    ws = wb["Checks"]
    (s0, s1), (m0, m1) = ifj["src"], ifj["mono"]
    ws[f"B{at['K2']}"].value = K2_LABEL
    ws[f"B{at['K3']}"].value = K3_LABEL
    ws[f"E{at['K3']}"].value = f"Interface 第 {s0}–{s1} 列（來源比對）＋第 {m0}–{m1} 列（單調）× C:Q"
    return at


# ================================================================== X15 (b): L1_HoldEconMW_* (no /1000)
def l1_rows(R):
    out = []
    for row in R:
        key = row[0]
        if key.startswith("HoldEconMW_"):
            gen = key[len("HoldEconMW_"):]
            row = list(row)
            row[1] = f"每 MW 年經濟持有成本（{gen}；數值＝每 GW 的 $B：1 $B/GW＝1 $M/MW；v5.30 X15 (b) 更正）"
            row[3], row[4], row[5] = f"=L1_HoldEconGW_{gen}", f"=L1_HoldEconGW_{gen}_Lo", f"=L1_HoldEconGW_{gen}_Hi"
            row = tuple(row)
        out.append(row)
    return out


# ================================================================== Decisions X15a–c (10 columns; same layout as v529)
_X15 = [
    ("a", "四層瀑布逐層累乘",
     "J 節 IFW_ 四層改為逐層累乘、單調遞減：_Util＝_100 × IF_Util；_Prod＝_Util × CTL_ProdDerate（營收列維持讀 G 節重解值；token 產能列改線性乘法，不再讀 Interface_Prod）；"
     "_Life＝_Prod × CTL_PriceLife × CTL_Monetize（營收列不再讀 H 節；token 產能 _Life＝_Prod）；Front 四列 _Prod／_Life 改公式（取代「—」）。"
     "H 節 _Life 與 Interface_Prod 保留原處作並列對照。Checks K3 改為來源比對（_Util／token _100）＋單調檢查；v5.29 X14 (b) 原意即逐層累乘，實作未落實",
     "Interface J 節（IFW_*_Prod／_Life）、Checks K3"),
    ("b", "L1_HoldEconMW 單位更正",
     "L1_HoldEconMW_*＝L1_HoldEconGW_*（不除以 1000）：1 $B/GW＝1 $M/MW，數值與每 GW 的 $B 相同；單位欄維持「$M/MW/年」；低／高同步。v5.29 值（例：GB300 0.0127）為 $B/MW/年，差 1,000 倍",
     "L1 第 54–57 列（L1_HoldEconMW_Hopper／GB200／GB300／VR200 與 _Lo／_Hi）"),
    ("c", "IFW_ 上限標記",
     "IFC_Use（S 欄）：IFW_RevGW_*、IFW_RevGWFleet*、IFW_RevGWFront* 全部四層與 IFW_TokGW_*_100 附「；上限，不得作預測」，與 D 節、B 節對應列一致；Checks K2 的應附旗標（Y 欄）同步擴充，K2 仍須為 0",
     "Interface J 節 S／Y 欄；Checks K2"),
]
DECISIONS_V530 = [[f"X15{k}", V, f"X15 ({k}) {t}", txt, DATE, ANDY, "v5.30 已建", where, "工作單 v5.30 第 " + {"a": "1", "b": "2", "c": "3"}[k] + " 節", "否"]
                  for k, t, txt, where in _X15]


# ================================================================== README (version string is built from VERSION; v5.29 text is kept after it)
README_VERSION = (VERSION + "（X15 定義更正，不改任何輸入值：(a) Interface J 節四層瀑布改為逐層累乘、單調遞減——IFW_*_Prod＝_Util × CTL_ProdDerate"
                  "（營收列維持讀 G 節重解值；token 產能列改線性乘法）、IFW_*_Life＝_Prod × CTL_PriceLife × CTL_Monetize（營收列）；Front 四列 _Prod／_Life 改公式；"
                  "Checks K3 改為來源比對＋單調檢查；(b) L1_HoldEconMW_* 改為不除以 1000（1 $B/GW＝1 $M/MW；GB300 0.0127 → 12.72 $M/MW/年）；"
                  "(c) IFW_RevGW*／IFW_RevGWFleet*／IFW_RevGWFront* 四層與 IFW_TokGW_*_100 的 Decision Use 附「上限，不得作預測」（Checks K2 擴充）；"
                  "其餘 IF_、L1 與模型頁數值逐格不變；Decisions X15a–c；工作單 docs/workorders/20261008_v5.30.md）。以下為 "
                  + v529.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v529.README_TITLE
README_TITLE = VERSION.split("_")[-1] + _T[len(v529.VERSION.split("_")[-1]):]
README_ROW = ("輸出契約與四層瀑布（v5.29 X14；v5.30 X15 逐層累乘）",
              "Interface R 欄 Confidence：該列公式鏈上（builder 建置時反查，含所有被引用的 Gov_Map 格）的 Gov_Map 格中，取「CC 敏感度分段＝高」者的最弱標記——"
              "含 3 級紀錄或 Analogy／Assumed 缺區間 → C；含非「原始數據且 SRC 等級 1」者 → B；其餘 → A；反查不到任何 Gov_Map 格 → 「—」（Checks K1）。"
              "S 欄 Decision Use 由 R 映射（A 基準輸入／B 情境、相對比較、反轉門檻／C 探索），B 節 100% 產出列、D 節理論營收列，以及 J 節 IFW_RevGW*／IFW_RevGWFleet*／"
              "IFW_RevGWFront* 全部四層與 IFW_TokGW_*_100 另附「上限，不得作預測」（Checks K2；v5.30 X15 (c)）。"
              "T 欄口徑層：100%／IF_Util／CTL_ProdDerate／L×m，依節與標籤判定；U 欄＝R 欄取用的高段 GM 列對應 SRC 審查日的最大值（V 欄為數值暫存，W 欄列出反查到的 GM 列）。"
              "J 節 IFW_（v5.30 X15 (a) 起逐層累乘、單調遞減）：_100＝100% 口徑（營收＝D 節 ÷ IF_Util；token＝B 節）；_Util＝_100 × IF_Util（營收＝D 節）；"
              "_Prod＝_Util × CTL_ProdDerate（營收列讀 G 節重解 SLO 後的值；token 產能列與 Front 列為線性乘法）；_Life＝_Prod × CTL_PriceLife × CTL_Monetize（token 產能 _Life＝_Prod，L、m 不作用於產能）。"
              "H 節 _Life（基準列 × L × m，不含折減）與 Interface_Prod（100% 口徑重解）保留原處作並列對照。自我檢查見 Checks K3（來源比對＋單調檢查）。"
              "Gov_Map AG–AJ、SRC_Index F 為 builder 每次重建的輔助欄。")
