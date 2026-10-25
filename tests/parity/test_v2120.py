"""v21.2 專屬測試（工作單 docs/workorders/20261023_v21.2_計價倍數與說明修補.md r0）。讀 LibreOffice 重算後的值。

  1. Alloc_In C23:E23＝1.8／0.85／2.3，F23 仍為 Derived；AL_SpendBasis* 讀到它們。
  2. 值變動的名稱（44 個，另加兩個版本字串名稱）：基準新值等於事前值，且與封存的 v21.1 不同；變動名稱集合恰為預期。
  3. 不變名稱：全部 IF_／IFW_／IFC_、L1_LReqFleet_Full*、L1_RevGW_PQ_VR200*（含 _FY／_Life／_F2026／_F2030）、L1_M／_Lo／_Hi、L1_LTotal* 與其他所有名稱，值對封存的 v21.1 差 0。
  4. 只有 AL_SpendBasis 驅動：輸入改回 0.85／0.73／1.17，變動名稱回到 v21.1 的值。
  5. W1–W4 文字（G23 原文逐字保留、L1 C71／H71、SRC_DEM_006、Cap_In G28、NonNV L11）；Decisions SL59-01、DB_Evidence E403、README。
  6. W6：區間時效指紋記錄 1.8／0.85／2.3、RNG_StaleFlag＝一致、Alloc C410＝0、GOV 三項計數。
  7. 逐格公式文字差異只落在工作單列明的格。
"""
import re
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import range_boundaries

from conftest import model_path
from test_v2130 import CHANGED as _V2130_CHANGED, LABEL_NAMES as _V2130_LABEL, INDEX_NAMES as _V2130_INDEX, REVERT as R213      # v21.3：機架價低端 7.5／Sol 快取價 0.10 的預期變動（test_v2130 看守）
from v2150_allow import v213_cells      # v21.3：v21.2 → v21.3 的逐格差異集合（規則同 test_v2130，只到工作單 W1–W6 列明的格）

REPO = Path(__file__).resolve().parents[2]
V211 = REPO / "model" / "archive" / "20261023_Tokenomics_v21.1.xlsx"
REVERT = {"Alloc_In!C23": 0.85, "Alloc_In!D23": 0.73, "Alloc_In!E23": 1.17}
REVERT_ALL = {**REVERT, **R213}      # v21.3：另把 Spec_Rack F11 設回 7.8、SRC_PRC_049 設回 0.2，才能單獨看 AL_SpendBasis 的效果
PREWARM = {"test_v2120_only_spend_basis_drives": [REVERT_ALL]}

# 事前值（LibreOffice 對 v21.1 只改 C23:E23 重算；工作單第 0 節）：名稱 -> (v21.1 基準, v21.2 基準)
CHANGED_BASE = {
    "AL_SpendBasis": (0.85, 1.8), "AL_SpendBasis_Lo": (0.73, 0.85), "AL_SpendBasis_Hi": (1.17, 2.3),
    "L1_TokLedger": (5.6159, 2.6520), "L1_TokLedger_2026Q2": (6.3942, 3.0195),
    "L1_MFleet": (0.4476, 0.9478), "L1_MFleet_Lo": (0.2362, 0.2750), "L1_MFleet_Hi": (1.2953, 2.5464),
    "L1_LReqFleet_IdNoK": (0.7561, 0.3571), "L1_LReqFleet": (2.0377, 0.9622), "L1_LReqFleet_Id": (0.7561, 0.3571),
    "L1_LReqFleet_VR200": (0.6923, 0.3269), "L1_GapSpendBasis": (0.85, 1.8), "L1_GapProduct": (1.9007, 4.025),
    "L1_GapUnexplained": (5.6159, 2.652), "L1_PriceRealFleet": (0.6084, 1.2884), "L1_TokPerPromptClose": (44128.7764, 18361.4448),
    "L1_TokCeiling": (29132172730.5955, 13756859345.0035), "L1_MFleet_Px2026": (0.3103, 0.6631),      # v21.3：L1_MFleet_Px2026 v21.2 值 0.6571 → 0.6631（Sol 快取價 0.20 → 0.10：有效牌價 −0.9%；test_v2130 看守）

}
CHANGED = set("""AL_SpendBasis AL_SpendBasis_Hi AL_SpendBasis_Lo L1_GapProduct L1_GapProduct_Hi L1_GapProduct_Lo L1_GapSpendBasis L1_GapSpendBasis_Hi L1_GapSpendBasis_Lo
L1_GapUnexplained L1_GapUnexplained_Hi L1_GapUnexplained_Lo L1_LReqFleet L1_LReqFleet_Hi L1_LReqFleet_Id L1_LReqFleet_IdNoK L1_LReqFleet_IdNoK_Hi L1_LReqFleet_IdNoK_Lo
L1_LReqFleet_Id_Hi L1_LReqFleet_Id_Lo L1_LReqFleet_Lo L1_LReqFleet_VR200 L1_LReqFleet_VR200_Hi L1_LReqFleet_VR200_Lo L1_MFleet L1_MFleet_Hi L1_MFleet_Lo L1_MFleet_P25Hi
L1_MFleet_P25Lo L1_MFleet_Px2026 L1_MFleet_Px2026_Hi L1_MFleet_Px2026_Lo L1_MFleet_Ser L1_PriceRealFleet L1_PriceRealFleet_Hi L1_PriceRealFleet_Lo L1_TokCeiling L1_TokCeiling_Hi
L1_TokCeiling_Lo L1_TokLedger L1_TokLedger_2026Q2 L1_TokPerPromptClose L1_TokPerPromptClose_Hi L1_TokPerPromptClose_Lo""".split())
VERSION_NAMES = {"RNG_Version", "RNG_Current"}
assert len(CHANGED) == 44

# 第 0 節「不變」名稱的前綴／全名（用來另行明確斷言，不只靠「其他全部不變」）
KEEP_PREFIX = ("IF_", "IFW_", "IFC_", "L1_LReqFleet_Full", "L1_RevGW_PQ_VR200", "L1_LTotal")
KEEP_EXACT = {"L1_M", "L1_M_Lo", "L1_M_Hi"}


def _flat(v):
    if isinstance(v, (list, tuple)):
        return [x for r in v for x in _flat(r)]
    return [v]


def _num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def _eq(a, b):
    if a in (None, "") and b in (None, ""):
        return True
    if _num(a) and _num(b):
        return abs(a - b) <= 1e-9 * max(abs(a), abs(b), 1e-300) + 1e-12
    return a == b


def _arch_value(pw, d):
    sh, rg = d.attr_text.rsplit("!", 1)
    c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
    return [pw[sh.strip("'")].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


@pytest.fixture(scope="module")
def wbf():
    return openpyxl.load_workbook(model_path())


@pytest.fixture(scope="module")
def pw():
    return openpyxl.load_workbook(V211, data_only=True)


# ---------------------------------------------------------------- 1. 輸入與標記
def test_v2120_inputs_and_mark(lo, wbf):
    ws = wbf["Alloc_In"]
    assert ws["A23"].value == "推論支出 ÷ 持有成本（計價倍數；含供應商利潤與資本回收）"
    assert (ws["C23"].value, ws["D23"].value, ws["E23"].value) == (1.8, 0.85, 2.3)
    assert ws["F23"].value == "Derived"
    v = lo.view({})
    assert (v.get_name("AL_SpendBasis"), v.get_name("AL_SpendBasis_Lo"), v.get_name("AL_SpendBasis_Hi")) == (1.8, 0.85, 2.3)
    assert wbf.defined_names["AL_SpendBasis"].attr_text == "Alloc_In!$C$23"


# ---------------------------------------------------------------- 2. 變動名稱
@pytest.mark.parametrize("nm", sorted(CHANGED_BASE))
def test_v2120_changed_base_values(lo, nm):
    old, new = CHANGED_BASE[nm]
    v = lo.view({})
    got = _flat(v.get_name(nm))[0]
    assert got == pytest.approx(new, rel=2e-4), nm
    assert abs(got - old) > 1e-6 * max(abs(old), 1e-9), nm


def test_v2120_changed_set_is_exactly_expected(lo, pw, wbf):
    v = lo.view({})
    changed = set()
    for name, d in wbf.defined_names.items():
        if "!" not in d.attr_text or "," in d.attr_text or name in VERSION_NAMES:
            continue
        if name not in pw.defined_names:
            continue
        want = _arch_value(pw, pw.defined_names[name])
        got = _flat(v.get_name(name))
        if not (len(want) == len(got) and all(_eq(a, b) for a, b in zip(want, got))):
            changed.add(name)
    v213 = _V2130_CHANGED | _V2130_LABEL | _V2130_INDEX | {"GOV_Warnings", "GOV_Table"} | {n for n in changed if n.startswith("RNG_")}      # v21.4：GOV_Warnings／GOV_Table（SRC_PRC_049 第二來源「免」，test_v2140 看守）；      # v21.3：W1／W2 的預期變動名稱、改名連帶的文字名稱、SRC_Index 延伸、W6 區間端點（test_v2130 看守）
    assert changed == CHANGED | v213, (sorted(changed - CHANGED - v213), sorted((CHANGED | v213) - changed))
    # 版本字串名稱改變（README B5 的版本主幹）
    assert str(v.get_name("RNG_Version")).endswith("v21.5") and str(v.get_name("RNG_Current")).endswith("v21.5")      # v21.3：版本字串（原 v21.2）


# ---------------------------------------------------------------- 3. 不變名稱
def test_v2120_unchanged_names(lo, pw, wbf):
    v = lo.view({})
    n_keep, bad = 0, []
    for name, d in wbf.defined_names.items():
        if "!" not in d.attr_text or "," in d.attr_text:
            continue
        if not (name.startswith(KEEP_PREFIX) or name in KEEP_EXACT):
            continue
        assert name not in CHANGED, name
        if name in _V2130_CHANGED | _V2130_LABEL | _V2130_INDEX:
            continue      # v21.3：這些名稱改由 test_v2130 看守（機架價低端／Sol 快取價的預期變動）
        assert name in pw.defined_names, name
        want = _arch_value(pw, pw.defined_names[name])
        got = _flat(v.get_name(name))
        n_keep += 1
        if not (len(want) == len(got) and all(_eq(a, b) for a, b in zip(want, got))):
            bad.append(name)
    assert n_keep > 150      # v21.3：扣除 171 個預期變動名稱後的比較數下修（原 > 300）
    assert bad == [], bad[:20]
    # 明確點名（第 0 節）
    for nm in ("L1_LReqFleet_Full", "L1_LReqFleet_Full_Lo", "L1_LReqFleet_Full_Hi", "L1_RevGW_PQ_VR200", "L1_RevGW_PQ_VR200_FY", "L1_RevGW_PQ_VR200_Life",
               "L1_RevGW_PQ_VR200_F2026", "L1_RevGW_PQ_VR200_F2030", "L1_M", "L1_LTotal"):
        assert nm in wbf.defined_names and nm not in CHANGED, nm


# ---------------------------------------------------------------- 4. 只有 AL_SpendBasis 驅動
def test_v2120_only_spend_basis_drives(lo, pw, wbf):
    h = lo.view(REVERT_ALL)      # v21.3：另設回 F11 7.8／SRC_PRC_049 0.2
    for nm in sorted(CHANGED):
        want = _arch_value(pw, pw.defined_names[nm])
        got = _flat(h.get_name(nm))
        assert len(want) == len(got) and all(_eq(a, b) for a, b in zip(want, got)), nm
    errs = [(s, c) for (s, c), x in h.evaluate_all().items() if isinstance(x, str) and x.startswith("#")]
    assert errs == []


# ---------------------------------------------------------------- 5. 文字與記錄
G23_OLD = ("經濟口徑：依據的價格 ÷ 會計成本 1.16 × 會計 ÷ 經濟持有成本 0.73 ≈ 0.85（區間 0.73–1.17）；分母維持經濟口徑持有成本。"
           "小於 1 表示 Azure 計價未涵蓋全部資本成本，這本身是 m 的一項實質命題（決定紀錄 D10）")


def test_v2120_text_w1_to_w4(wbf):
    g = wbf["Alloc_In"]["G23"].value
    assert "｜v21.2（2026-10-23 9a 查證 J5；Andy 21:58「按照預設」、22:02 同意建置）：基準 1.8、區間 0.85–2.3，Derived。" in g and g.startswith("v21.5：") and "｜v21.3（" in g      # v21.3：G23 前面加 v21.3 說明（W3，test_v2130 看守），v21.2 文字改為在其後
    assert "價格因子為殘差，必須約 1.37–2.18 才能重現支出路線（必須為真）" in g and "非實付價直接證據（OpenAI 在 Azure 的實付價找不到）" in g
    assert "高端 2.3≈Microsoft 帳上端" in g and "回本主口徑 L1_LReqFleet_Full、L1_RevGW_PQ_VR200、全部 IF_ 不受影響" in g
    assert g.endswith("｜原說明（v21.1 以前）：" + G23_OLD)
    l1 = wbf["L1"]
    r = int(wbf.defined_names["L1_GapSpendBasis"].attr_text.split("$")[-1])
    assert r == 71
    assert l1[f"C{r}"].value == "Alloc_In 計價倍數（v21.2 起 Derived：基準 1.8、區間 0.85–2.3；X46）｜原：Alloc_In 計價倍數（Assumed）；X46"
    assert l1[f"H{r}"].value == "0.85–2.3（Derived；v21.2 9a J5）｜原：1.0–1.6（Assumed）"
    assert l1[f"K{r}"].value == "本列：Derived（v21.2 9a J5；E403；X46）｜原：本列：Assumed（E262；X46）"
    sd = wbf["SRC_Demand"]
    assert sd["A10"].value == "SRC_DEM_006"
    w = sd["W10"].value
    assert w.startswith("出處待補；已排除 The Information 2026-02〈Why OpenAI, Anthropic Missed Their Own Gross Margin Forecasts〉（2026-10-23 查證 J5）｜"
                        "外部最近為 2025-09 預測逾 90 億（2026-10-23 9a 查證 J1，docs/evidence/20261023_Q側缺口_A段/）｜—｜v17.2 Q1：15 組查詢未找到（R6）")
    assert sd["C10"].value == 12
    ci = wbf["Cap_In"]
    assert ci["A28"].value == "對外服務占機隊"
    assert ci["G28"].value.startswith("支出比代 GW 比，前提是兩邊價格倍數相同（2026-10-23 9a 查證；v21.2）｜OpenAI 2025 推論 $8.4B ÷（推論 $8.4B＋訓練約 $12B）")
    assert ci["C28"].value == pytest.approx(0.41, abs=0.005) or str(ci["C28"].value).startswith("=")      # J2：C28:E28 不改（逐格差異測試看守）
    nn = wbf["NonNV"]
    assert nn["A11"].value == "Microsoft Maia 200"
    l = nn["L11"].value
    assert ("軟體效率因子 0.45–1.0（無第三方實測，主要不確定）。區間 0.25–0.75：高端 0.75 需機架功率 ≤58 kW；低端含軟體成熟度下修（v21.2，v21.1 獨立查核觀察）。機架價格無公開資料") in l
    assert (nn["B11"].value, nn["C11"].value, nn["D11"].value) == (0.25, 0.55, 0.75)


def test_v2120_records(wbf):
    d = wbf["Decisions"]
    rows = [r for r in range(5, d.max_row + 1) if d.cell(r, 1).value == "SL59-01"]
    assert len(rows) == 1 and d.cell(rows[0] - 1, 1).value == "SL58-01"
    assert d.cell(rows[0], 2).value == "v21.2" and d.cell(rows[0], 7).value == "v21.2 已建"
    txt = d.cell(rows[0], 4).value + d.cell(rows[0], 6).value
    for key in ("按照預設", "開個新串來建置下一版", "L1 C71", "H71", "Assumed", "Derived", "原值逐字保留", "SRC_DEM_006", "G28", "L11", "證偽條件"):
        assert key in txt, key
    ev = wbf["DB_Evidence"]
    er = [r for r in range(5, ev.max_row + 1) if ev.cell(r, 1).value == "E403"]
    assert len(er) == 1 and ev.cell(er[0] - 1, 1).value == "E402"
    assert ev.cell(er[0], 5).value == "Derived" and "非實付價直接證據" in ev.cell(er[0], 11).value and "1.8" in ev.cell(er[0], 3).value
    rm = wbf["README"]
    assert str(rm["B5"].value).startswith("20261025_Tokenomics_v21.5（次版，底版 v21.4：") and "v21.4（次版，底版 v21.3：" in str(rm["B5"].value) and "以下為 v21.2（次版，底版 v21.1：" in str(rm["B5"].value) and "v21.1（主版：" in str(rm["B5"].value)      # v21.3：B5 前面加 v21.3 句，v21.2 句保留
    assert str(rm["A1"].value).startswith("Tokenomics v21.5")      # v21.3：版本字串（原 v21.2）
    rr = [x for x in range(4, rm.max_row + 1) if rm.cell(x, 1).value == "v21.2 計價倍數 1.8 與說明修補"]
    assert len(rr) == 1 and "1.8" in rm.cell(rr[0], 2).value


# ---------------------------------------------------------------- 6. W6：指紋、區間、GOV
def test_v2120_fingerprint_ranges_and_gov(lo, wbf):
    v = lo.view({})
    assert v.get_name("RNG_StaleFlag") == "一致"
    assert v.get("Alloc", "C410") == 0
    assert (v.get("Alloc", "D1020"), v.get("Alloc", "D1021"), v.get("Alloc", "D1022")) == (1.8, 0.85, 2.3)
    assert (v.get_name("GOV_Errors"), v.get_name("GOV_Warnings"), v.get_name("GOV_Info")) == (0, 134, 182)      # v21.4：W3 SRC_PRC_049 第二來源「免」，135 → 134（test_v2140 看守）
    # 檔案層級的指紋與端點檔對現行活頁簿一致由 tools/gen_fingerprint.py --check、tools/gen_ranges.py --check 驗證（見報告）


# ---------------------------------------------------------------- 7. 公式文字逐格差異
def test_v2120_diffs_only_expected_vs_v211():
    a, b = openpyxl.load_workbook(V211), openpyxl.load_workbook(model_path())
    assert a.sheetnames == b.sheetnames
    d = {}
    for ws in a:
        wn = b[ws.title]
        for r in range(1, max(ws.max_row, wn.max_row) + 1):
            for c in range(1, max(ws.max_column, wn.max_column) + 1):
                x, y = ws.cell(r, c).value, wn.cell(r, c).value
                if x != y:
                    d.setdefault(ws.title, []).append((ws.cell(r, c).coordinate, x, y))
    got = {k: sorted(t[0] for t in v) for k, v in d.items()}
    want = {
        "Alloc": sorted(["C177", "D1020", "D1021", "D1022"]),
        "Alloc_In": ["C23", "D23", "E23", "G23"],
        "Cap_In": ["G28"],
        "DB_Evidence": sorted(f"{c}327" for c in "ABCDEFGHIJKLMNOPQR"),
        "Decisions": sorted(f"{c}331" for c in "ABCDEFGHIJ"),
        "L1": ["C71", "H71", "K71"],
        "NonNV": ["L11"],
        "README": sorted(["A1", "A72", "B5", "B72"]),
        "SRC_Demand": ["W10"],
    }
    for k, cells in v213_cells().items():      # v21.3：疊上 v21.2 → v21.3 的差異格（規則同 test_v2130）
        want[k] = sorted(set(want.get(k, [])) | cells)
    assert got == want, {k: (sorted(set(got.get(k, [])) ^ set(want.get(k, [])))) for k in set(got) | set(want) if got.get(k) != want.get(k)}
