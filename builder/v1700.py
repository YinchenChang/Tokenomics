# v17.0 (work order docs/workorders/20261020_v17.0_F1F2主口徑.md r0, branch claude/v17-0): payback reading main basis = generation-start + discounted fleet L versus the no-kappa required L (F1, F2 defaults).
# MAJOR version: one new contract name group (L1_LTotalFleetVDisc, _Lo, _Hi); no existing named value changes.
#   W1  Cap_In section O "L 世代起點折現（v17.0）" below section N (nothing above moves): per-generation discounted L_rel (base/low/high combinations, discounted to each generation's own start),
#       fleet discounted L -> L1_LTotalFleetVDisc (+_Lo/_Hi); comparison row "v16.0 口徑 (undiscounted, with kappa)"; L1 row appended after the last row.
#   W2  Cap_In row 242 (A, C, D, E, G): L side reads L1_LTotalFleetVDisc, required-L side reads L1_LReqFleet_IdNoK; L1 P89 reads the new main basis and appends the v16.0 reading;
#       L1 B107 (L1_LReqFleet_Id) label: "v17.0 起非主口徑（F1）".
#   W3  Decisions SL50-01..03, README, contract document (outside the workbook).
# Every write to an Excel-owned cell is guarded (old text / already-applied test); no input value changes.
import re

from openpyxl.utils import get_column_letter as L

from common import put, F_NOTE, F_BOLD, F_CALC, F_LINK, FILL_KEY
import v1000
import v1410
import v1600

VERSION = "20261020_Tokenomics_v17.0"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-20"
V = "v17.0"
WO = "工作單 docs/workorders/20261020_v17.0_F1F2主口徑.md（r0）"
DASH = "—"
DEC_BY = ("chat 端 2026-10-18 10:44 提出判斷題 F1、F2 並寫明「沒回應的話，下一版照預設改」；至 2026-10-20 07:30 Andy 未回應（07:30 回覆「合併」只針對 PR #128、#129）；"
          "依情境卡_工作方式 v1.2 第 3 節『未回應即依預設執行並登錄決策』；可撤回")
DEF_TXT = ("回本判讀＝世代起點、折現的機隊 L（L1_LTotalFleetVDisc）對 不含 κ 的恆等式所需 L（L1_LReqFleet_IdNoK）。"
           "L 的定義不變（某世代硬體在壽命期 T 內，token 單價相對該世代起點單價的平均），改以 WACC 折現（折現到該世代起點，與成本年金同口徑）；所需 L 去掉請求形狀修正 κ。")

GENS = v1600.GENS
ROWS = {}
PURPOSE_MARKS = {}
SEC_TITLE = "O. L 世代起點折現（v17.0；" + WO + "）：" + DEF_TXT

NOLE = "t0 > T1（已降級，未支援）"


# ---------------------------------------------------------------- formulas
def _ln(x): return f"LN(1+{x})"


def _seg(a, x0, x1):
    """(EXP(a*x1)-EXP(a*x0))/a, or x1-x0 when |a| < 1e-12."""
    return f"IF(ABS({a})<0.000000000001,({x1})-({x0}),(EXP(({a})*({x1}))-EXP(({a})*({x0})))/({a}))"


def lrel_disc_formula(t0, c, gp, wacc):
    """Discounted L_rel of one generation (discounted to its own start). t0 = that generation's t0 cell; c = scenario column (C base / D low / E high) of the g, T1, rho, T rows; gp = g_pre cell; wacc = WACC cell."""
    g, t1, rho, T = f"${c}$226", f"${c}$227", f"${c}$228", f"${c}$229"
    d, kp, k = _ln(wacc), _ln(gp), _ln(g)
    past = f"IF({t0}>0,EXP(-({kp})*{t0})*{_seg(f'({kp})-({d})', '0', t0)},0)"
    mid = f"EXP(-({k})*{t0})*{_seg(f'({k})-({d})', t0, t1)}"
    late = f"{rho}*EXP(-({k})*{t0})*{_seg(f'({k})-({d})', t1, T)}"
    den = _seg(f"-({d})", "0", T)
    return f'=IF({t0}>{t1},"{NOLE}",({past}+{mid}+{late})/({den}))'


# ---------------------------------------------------------------- Cap_In O section
LAB = dict(
    lrd="折現 L_rel（相對錨點、折現到該世代起點；基準；各世代，O-1）　WACC＝Inputs E28",
    lrd_lo="折現 L_rel（低端組合：g、T1、ρ 取 D 欄，g_pre 取 E 欄，t0 取低端組合；WACC＝Inputs D28）",
    lrd_hi="折現 L_rel（高端組合：g、T1、ρ 取 E 欄，g_pre 取 D 欄，t0 取高端組合；WACC＝Inputs F28）",
    fleetd="機隊折現 L（世代起點＋折現，錨點 2025）＝Σ w_v × 折現 L_rel_v ÷ Σ w_v　→ L1_LTotalFleetVDisc（D 欄＝低端組合、E 欄＝高端組合）",
    j_v16="對照：v16.0 口徑（不折現、含 κ）：L1_LTotalFleetV 區間 對 所需 L（L1_LReqFleet_Id）區間；三態",
)
NOTE = dict(
    lrd=("d＝LN(1+WACC)，kp＝LN(1+g_pre)，k＝LN(1+g)；seg(a, x0, x1)＝(EXP(a·x1) − EXP(a·x0)) ÷ a（|a| < 1e-12 時＝x1 − x0）；"
         "past＝IF(t0>0, EXP(−kp·t0) × seg(kp − d, 0, t0), 0)；mid＝EXP(−k·t0) × seg(k − d, t0, T1)；late＝ρ × EXP(−k·t0) × seg(k − d, T1, T)；"
         "折現 L_rel＝(past + mid + late) ÷ seg(−d, 0, T)。分母是 T 年的折現年金（折現到該世代起點，與成本年金同口徑）。g、T1、ρ、T 取第 226／227／228／229 列的 C 欄；"
         "t0 ＞ T1 輸出文字『t0 > T1（已降級，未支援）』（目前沒有）。t0＝0 時等於 L1_LTotalDisc（Cap_In 第 314 列）"),
    lrd_lo="低端組合：WACC 取 Inputs D28（0.08）；其餘同基準列（沿用 v16.0 的端點組合）。待 Project 判斷：WACC 與 g 端點同向組合是否合理",
    lrd_hi="高端組合：WACC 取 Inputs F28（0.13）；其餘同基準列（沿用 v16.0 的端點組合）",
    fleetd=("權重 w_v 同 v16.0（Cap_In N 節 2025 營收權重列＝AL_ShareGen × AL_RevGW25）。C 基準；D 低端組合；E 高端組合。chat 端事前計算 0.75860／0.47298／1.13827。"
            "v17.0 起判讀（第 242 列、L1 第 89 列）的 L 側讀本列（F2 依預設：改用折現版）"),
    j_v16="v16.0 主口徑（世代起點 L 對含 κ 的所需 L）保留為對照；文字三態；只作對照",
)


def cap_in_o(wb, n_sec):
    ws = wb["Cap_In"]
    r0 = n_sec["end"] + 2
    assert ws.max_row == n_sec["end"], (ws.max_row, n_sec["end"])
    put(ws, f"A{r0}", SEC_TITLE, F_BOLD, wrap=True)
    r = r0 + 1
    put(ws, f"A{r}", "欄位：世代列 C Hopper／D GB200／E GB300／F VR200；情境列 C 基準／D 低端／E 高端；G 標記、H 說明、O 用途（同 N 節）", F_NOTE, wrap=True)
    for col, g in zip("CDEF", GENS):
        put(ws, f"{col}{r}", g, F_BOLD)
    r += 1
    order = ["lrd", "lrd_lo", "lrd_hi", "fleetd", "j_v16"]
    for k in order:
        ROWS[k] = r; r += 1
    R, N = ROWS, v1600.ROWS

    def row(k, unit, vals, tag, fmt, key=False, purpose=None):
        rr = R[k]
        put(ws, f"A{rr}", LAB[k], wrap=True); put(ws, f"B{rr}", unit)
        for col, v in zip("CDEF", vals):
            if v is None:
                continue
            isf = isinstance(v, str) and v.startswith("=")
            put(ws, f"{col}{rr}", v, (F_LINK if isf and "!" in v else F_CALC), fmt=fmt, fill=FILL_KEY if (key and col == "C") else None)
        put(ws, f"G{rr}", tag, F_NOTE); put(ws, f"H{rr}", NOTE[k], F_NOTE, wrap=True)
        if purpose:
            PURPOSE_MARKS[rr] = purpose

    gp = lambda c: f"${c}${N['gpre']}"
    row("lrd", "x", [lrel_disc_formula(f"{c}{N['t0']}", "C", gp("C"), "Inputs!$E$28") for c in "CDEF"], "Derived", "0.0000")
    row("lrd_lo", "x", [lrel_disc_formula(f"{c}{N['t0_lo']}", "D", gp("E"), "Inputs!$D$28") for c in "CDEF"], "Derived", "0.0000")
    row("lrd_hi", "x", [lrel_disc_formula(f"{c}{N['t0_hi']}", "E", gp("D"), "Inputs!$F$28") for c in "CDEF"], "Derived", "0.0000")
    row("fleetd", "x", [v1600.fleet_formula(R["lrd"], N["w"]), v1600.fleet_formula(R["lrd_lo"], N["w"]), v1600.fleet_formula(R["lrd_hi"], N["w"])], "Derived", "0.0000", key=True)
    row("j_v16", "x", [v1600.tri("L1_LTotalFleetV", "L1_LTotalFleetV_Lo", "L1_LTotalFleetV_Hi", "L1_LReqFleet_Id", "L1_LReqFleet_Id_Lo", "L1_LReqFleet_Id_Hi", "世代起點 L（v16.0 口徑）")],
        "Derived", None, purpose="顯示用：C%d（對照列：v16.0 口徑，不折現、含 κ）" % R["j_v16"])
    return dict(start=r0, end=r - 1, rows=dict(R))


# ---------------------------------------------------------------- Cap_In row 242 (before gov_all)
A242_NEW = ("判讀：折現的世代起點 L 區間 對 不含 κ 的所需 L 區間（恆等式讀法 L1_LReqFleet_IdNoK；三態：未回本／未定／回本；"
            "v17.0 起 L 側讀 L1_LTotalFleetVDisc，F1／F2 依預設）")
G242_ADD = ("｜v17.0 W2（F1、F2 依預設）：本列 L 側改讀 L1_LTotalFleetVDisc（世代起點＋折現的機隊 L，區間 [L1_LTotalFleetVDisc_Lo, L1_LTotalFleetVDisc_Hi]），"
            "所需 L 側改讀 L1_LReqFleet_IdNoK（不含 κ；區間 [_Lo, _Hi]）；三態規則同 v15.3 D1。基準比值（L ÷ 所需 L）＝1.003。"
            "v16.0 口徑（世代起點 L 對含 κ 所需 L）保留為 Cap_In O 節對照列；Decisions SL50-01／02／03")


def inputs_update(wb, append_text):
    log = []
    ws = wb["Cap_In"]
    assert ws["A242"].value in (v1600.A242_NEW, A242_NEW), ws["A242"].value
    n = 0
    for c in "CDE":
        cur = ws[f"{c}242"].value
        assert isinstance(cur, str) and cur.startswith("=IF(AND(ISNUMBER(L1_LTotalFleetV"), cur[:60]
        if "L1_LTotalFleetVDisc" in cur:
            continue
        new = re.sub(r"L1_LTotalFleetV(?!Disc)", "L1_LTotalFleetVDisc", cur)
        new = re.sub(r"L1_LReqFleet_Id(?!NoK)", "L1_LReqFleet_IdNoK", new)
        new = new.replace("所需 L（GW 低端）", "所需 L（不含 κ，低端）").replace("所需 L（GW 高端）", "所需 L（不含 κ，高端）")
        assert new != cur
        ws[f"{c}242"].value = new; n += 1
    ws["A242"].value = A242_NEW
    n += int(append_text(ws["G242"], G242_ADD))
    log.append(f"Cap_In row 242 reads L1_LTotalFleetVDisc / L1_LReqFleet_IdNoK ({n} edits)")
    return log


# ---------------------------------------------------------------- L1 rows (appended after the last existing row)
READ_FD = ("v17.0（F2 依預設）：機隊 L 改為世代起點＋折現。每個世代的 L 以 WACC 折現到該世代起點（與成本年金同口徑），錨點前已服務 t0 年的部分用 g_pre 回推；以 2025 營收（占比 × 每 GW 付費營收）加權。"
           "判讀（Cap_In 第 242 列、L1 第 89 列）的 L 側改讀本列；所需 L 側改讀 L1_LReqFleet_IdNoK（F1）。")
NONE_EXT = ("無外部對照", DASH, DASH)


def l1_rows_new():
    R = ROWS
    def cap(col, key): return f"=Cap_In!${col}${R[key]}"
    return [
        ("LTotalFleetVDisc", "壽命期價格係數 L — 機隊折現 L（世代起點＋折現；錨點 2025；v17.0 判讀主口徑）",
         "各世代從自己的服務起點 s_v 算 L，以 WACC 折現到該世代起點；錨點年 A＝2025.5；以 2025 營收（AL_ShareGen × AL_RevGW25）加權；v17.0",
         cap("C", "fleetd"), cap("D", "fleetd"), cap("E", "fleetd"), "x",
         "低＝低端組合（g −0.20、T1 3、ρ 0.17、s_v 較晚、g_pre −0.09、WACC 8%）；高＝高端組合（g +0.05、T1 5、ρ 0.7、s_v 較早、g_pre −0.30、WACC 13%）", READ_FD,
         "g、T1、ρ、s_v、g_pre、WACC、2025 營收權重", "s_v、g_pre、WACC：Assumed；g、ρ：Derived；T1：Assumed", *NONE_EXT, "AL_ShareGen；AL_RevGW25；Inputs E28／D28／F28", f"Cap_In 第 {R['fleetd']} 列",
         "各世代服務起點無一手資料（以雲端供應日期推定）；WACC 與 g 端點同向組合是否合理待 Project 判斷"),
    ]


v1000._n(v1000.E_, "折現的世代起點機隊 L（v17.0 判讀主口徑）：價格迴歸與降級假設、起點、前期跌價、WACC 為 Assumed", "LTotalFleetVDisc")


# ---------------------------------------------------------------- L1 P89 and B107 label (late: needs the final L1 row numbers)
def _l1_row(wb, name):
    r = int(wb.defined_names[name].attr_text.split("$")[-1])
    assert wb["L1"][f"A{r}"].value == name, (name, r, wb["L1"][f"A{r}"].value)
    return r


def l1_p89(wb):
    ws = wb["L1"]
    fv = _l1_row(wb, "L1_LTotalFleetV")
    fd = _l1_row(wb, "L1_LTotalFleetVDisc")
    nk = _l1_row(wb, "L1_LReqFleet_IdNoK")
    assert ws["A89"].value == "L1_LReqFleet", ws["A89"].value
    cur = ws["P89"].value
    if f"D{fd}" in cur:
        return 0
    tri_old = lambda a, b: (f'IF(AND(ISNUMBER(E{a}),ISNUMBER(F{a}),ISNUMBER(E{b}),ISNUMBER(F{b})),IF(F{a}<E{b},"未回本",IF(E{a}>F{b},"回本","未定")),IF(D{a}>=D{b},"回本","未回本"))')
    reps = [
        (f"ISNUMBER(D89),ISNUMBER(D107),ISNUMBER(D{fv})", f"ISNUMBER(D89),ISNUMBER(D107),ISNUMBER(D{fd}),ISNUMBER(D{nk})", 1),
        (tri_old(fv, 107), tri_old(fd, nk), 2),
        (f"AND(ISNUMBER(D107),ISNUMBER(D{fv}))", f"AND(ISNUMBER(D{nk}),ISNUMBER(D{fd}))", 1),
        (f'"（世代起點 L "&TEXT(D{fv},"0.00")&" 對所需 L "&TEXT(D107,"0.00")&"；v15.3 口徑 2025＝第 0 年：L "&TEXT(L1_LTotal,"0.000")&"）；字面讀法 "',
         f'"（世代起點折現 L "&TEXT(D{fd},"0.00")&" 對不含 κ 所需 L "&TEXT(D{nk},"0.00")&"；v15.3 口徑 2025＝第 0 年：L "&TEXT(L1_LTotal,"0.000")&"）（v16.0 口徑：L "&TEXT(D{fv},"0.000")&" 對 "&TEXT(D107,"0.000")&"）；字面讀法 "', 1),
        (f'"（恆等式讀法：世代起點 L "&TEXT(D{fv},"0.00")&" 對所需 L "&TEXT(D107,"0.00")&"；v15.3 口徑 2025＝第 0 年：L "&TEXT(L1_LTotal,"0.000")&"；字面讀法非數字）"',
         f'"（恆等式讀法：世代起點折現 L "&TEXT(D{fd},"0.00")&" 對不含 κ 所需 L "&TEXT(D{nk},"0.00")&"；v15.3 口徑 2025＝第 0 年：L "&TEXT(L1_LTotal,"0.000")&"）（v16.0 口徑：L "&TEXT(D{fv},"0.000")&" 對 "&TEXT(D107,"0.000")&"）；字面讀法非數字）"', 1),
    ]
    for old, new, cnt in reps:
        assert cur.count(old) == cnt, (old[:80], cur.count(old))
        cur = cur.replace(old, new)
    ws["P89"].value = cur
    return 1


def l1_label(wb, append_text):
    ws = wb["L1"]
    r = [x for x in range(5, ws.max_row + 1) if ws.cell(x, 1).value == "L1_LReqFleet_Id"]
    assert len(r) == 1, r
    return int(append_text(ws[f"B{r[0]}"], "｜v17.0 起非主口徑（F1）"))




# ================================================================== Decisions SL50-01..03
def _decisions():
    D = [
        ("F1：回本判讀的所需 L 拿掉 κ（改讀 L1_LReqFleet_IdNoK）",
         "回本判讀（Cap_In 第 242 列、L1 第 89 列）的所需 L 側由 L1_LReqFleet_Id（含 κ）改讀 L1_LReqFleet_IdNoK（不含 κ；＝SRC_DEM_004 ÷（AL_SpendBasis × SRC_DEM_007））。"
         "L1_LReqFleet_Id 名稱與值不變，標籤加註『v17.0 起非主口徑（F1）』。κ 為 1.0 時兩者基準值相同。本決議取代 SL49-03 中『F1 未回應前不改主口徑』的狀態。",
         DEC_BY, "Cap_In 第 242 列；L1 P89、L1_LReqFleet_Id 標籤"),
        ("F2：回本判讀的 L 改用折現版（新增 L1_LTotalFleetVDisc）",
         "回本判讀的 L 側由 L1_LTotalFleetV（世代起點、不折現）改讀新名稱 L1_LTotalFleetVDisc（＋_Lo／_Hi）：各世代 L 以 WACC 折現到該世代起點（與成本年金同口徑），機隊權重同 v16.0；"
         "WACC 端點低端組合 Inputs D28（0.08）、高端組合 F28（0.13）。新增 Cap_In O 節（含 v16.0 口徑對照列）與 L1 新列。本決議取代 SL49-03 中『F2 未回應前不改主口徑』的狀態。"
         "待 Project 判斷：WACC 端點與 g 端點同向組合是否合理（照 v16.0 慣例）。",
         DEC_BY, "Cap_In O 節；L1_LTotalFleetVDisc；Cap_In 第 242 列；L1 P89"),
        ("回本判讀主口徑切換紀錄（v16.0 → v17.0）",
         "主口徑：世代起點＋折現的機隊 L 對 不含 κ 的恆等式所需 L。基準『未定（區間重疊；基準 L 0.759 對所需 L 0.756）』，基準比值 1.003。"
         "v16.0 口徑（世代起點 L 對含 κ 所需 L；0.716 對 0.756）改為 Cap_In O 節對照列，v16.0 既有三個對照列保留；L1 P89 附『（v16.0 口徑：L 0.716 對 0.756）』。既有名稱的值全部不變。",
         DEC_BY, "Cap_In 第 242 列與對照列；L1 P89；README"),
    ]
    return [[f"SL50-{i + 1:02d}", V, f"v17.0 回本判讀主口徑：{t}", txt, DATE, who, "v17.0 已建", where, f"{WO} W3", "否"]
            for i, (t, txt, who, where) in enumerate(D)]


def purpose_marks(wb):
    ws = wb["Cap_In"]
    import v1400
    pc = v1400._purpose_col(ws, create=False)
    assert pc is not None, "用途 column of Cap_In not found"
    for rr, txt in PURPOSE_MARKS.items():
        put(ws, f"{L(pc)}{rr}", txt, F_NOTE)
    return len(PURPOSE_MARKS)


def clear_l1_ratio(wb):
    ws = wb["L1"]; n = 0
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == "L1_LTotalFleetVDisc" and ws[f"O{r}"].value is not None:
            ws[f"O{r}"].value = None; n += 1
    return n


def apply(wb, append_text):
    log = {}
    log["l1_o_cleared"] = clear_l1_ratio(wb)
    log["purpose"] = purpose_marks(wb)
    log["p89"] = l1_p89(wb)
    log["label107"] = l1_label(wb, append_text)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == "SL50-01" for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


# ================================================================== README
README_VERSION = (VERSION + "（主版：" + WO + "——回本判讀主口徑改為『世代起點＋折現的機隊 L』對『不含 κ 的所需 L』（F1、F2 依預設）：新增 Cap_In O 節與契約名稱 "
                  "L1_LTotalFleetVDisc／_Lo／_Hi；Cap_In 第 242 列與 L1 第 89 列改讀新主口徑；v16.0 口徑（不折現、含 κ）改為對照列；L1_LReqFleet_Id 標籤加註非主口徑；Decisions SL50-01…03。既有名稱的值全部不變。）以下為 "
                  + v1600.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1600.README_TITLE
README_TITLE = V + _T[len(v1600.V):]
README_ROW = ("v17.0 回本判讀主口徑",
              DEF_TXT + "新增 Cap_In O 節與 L1_LTotalFleetVDisc；基準『未定』（L 0.759 對所需 L 0.756，比值 1.003）。v16.0 口徑保留為對照列；L1_LReqFleet_Id（含 κ）名稱與值不變。細節見 Decisions SL50-01…03。")
