# v5.20 (work order docs/workorders/20261006_v5.20.md): X3 eight X2 high-segment rows completed (DB_Evidence E236–E243, GM578 range 0–6,
# C2 list +GM275/GM354), X4 Calib F67 conflict evidence (E244), X5 amortisation-coupling note (Interface A3, README).
# Every write is guarded by the old value / presence, like gov.gm_update, so a rebuild never overwrites a later Excel edit.
from openpyxl.styles import PatternFill
from copy import copy
from common import F_IN

VERSION = "20261006_Tokenomics_v5.20"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-06"
READ = "摘要級，原文未讀"
CHECK = "docs/reports/20261006_v5.19_查核.xlsx「2.3 八列搜尋」"
WHO = "讀取者：CC（v5.19 查核）＋chat 端（2026-10-06 搜尋）；不得標為已讀"
SRC = f"WebSearch 摘要（{CHECK}）；原文網址未取得"

# ------------------------------------------------------------------ Alloc_In E7 (AL_Nrefresh_Hi): 4 -> 6
INPUT_UPD = [("Alloc_In", "E7", 4, 6)]

def inputs_update(wb):
    log = []
    for sh, cell, old, new in INPUT_UPD:
        c = wb[sh][cell]
        if c.value == old and type(c.value) is not bool:
            c.value = new; c.font = copy(F_IN); c.fill = PatternFill(fill_type=None)
            log.append(f"{sh}!{cell}: {old!r} -> {new!r} (v5.20 X3 A3)")
    return log

# ------------------------------------------------------------------ DB_Evidence (17 columns A..Q; same layout as v518.evidence_rows)
def _ev(i, gm, param, claim, now, then, verdict, note, grade="Analogy", stance="中立", affected="—"):
    return [f"E{236 + i}", DATE, claim, SRC, READ, param, then, now, verdict, "v5.20",
            f"{WHO}。工作單 v5.20 {gm}。{note}", "已處理", "—", "—", affected, grade, stance]

EVIDENCE_V520 = [
    _ev(0, "X3 1.1（GM205）", "Arch E8 層數（GM205）",
        "Llama 3.1 405B 126 層；DeepSeek-V3 61 層（開放權重，架構不同）",
        "126 層（稠密）／61 層（MoE）", "100（區間 80–160）",
        "已讀摘要，作為區間依據；區間不改",
        "不得寫入「405B 為 80 層」（誤植，80 層為 70B）。"),
    _ev(1, "X3 1.1（GM203）", "Arch E6 總參數（GM203）",
        "Kimi K3 2.8T（2026-07-16，開放權重）；Llama 4 Behemoth 2T（288B 啟用）；xAI 約 2T（另有 4–6T 估計）；Anthropic、OpenAI 旗艦未揭露",
        "約 2–2.8T（xAI 另有 4–6T 估計，Interested-party）", "基準值（區間 2,000–10,000 B）",
        "已讀摘要，支持區間下端；上端 10T 無可比；區間不改",
        "xAI 數字出自 xAI 宣傳，屬 Interested-party；4–6T 為估計。", stance="混合（xAI 為 Interested-party）"),
    _ev(2, "X3 1.1（GM459）", "Cap_In C36 儲存層代碼（GM459）",
        "公開快取條款保留時間（5 分鐘至 1 小時、延長至 24 小時）與推測的 HBM／DRAM／NVMe 分層一致；分層實作非官方揭露",
        "保留時間 5 分鐘–24 小時；分層為推測", "基準值（區間 D36–E36）",
        "已讀摘要，支持現值；不改",
        "分層實作非官方揭露（報導推測）。"),
    _ev(3, "X3 1.1（GM270）", "Workload G8 每輪思考 token h（GM270）",
        "SWE-bench Verified 每條軌跡推理 token 平均 19.5k（Qwen3-235B-A22B-Thinking；arXiv 2606.24820）；以 Workload G5 輪數 T＝30 換算約 650／輪（chat 端換算）",
        "約 650 token／輪（Derived：19.5k ÷ 30，chat 端換算）", "基準值 600（區間 ×0.5–×2）",
        "已讀摘要，支持基準 600；不改",
        "原文來源 arXiv 2606.24820（摘要）；模型與 harness 與本模型不同（Analogy）。"),
    _ev(4, "X3 1.1（GM290）", "Workload G12 並行子代理數 m（GM290）",
        "Claude Code subagent 同時最多 10、workflow 同時 16（產品上限）；Kimi 300 agent（宣傳數字）；典型使用平均並行數無統計",
        "上限 10–16（Claude Code）；300（Kimi 宣傳）", "基準 0（區間 0 到 max(1, 2×基準)）",
        "上限有可比、典型值無；不改",
        "理由加註：基準 0＝標準 harness 為單代理，屬定義。Kimi 300 為 Interested-party 宣傳數字。"),
    _ev(5, "X3 1.1／1.2（GM578）", "Alloc_In C7 改版計畫數 N_refresh（GM578）",
        "OpenAI GPT-5.x 點版本：5.1（2025-11）、5.2（2025-12）、5.4（2026-03-05）、5.5（2026-04-23）、5.6（約 2026-07）；過去 12 個月約 5 個",
        "約 5 個／12 個月（對外點版本）", "基準 2；區間 0–4（個／年）",
        "只作上界參考；區間上限 4 → 6（Andy 2026-10-06，X3 A3），基準 2 不變",
        "「改版計畫」為相對家族計畫的研發運算當量，不等於對外版本號。",
        affected="AL_Nrefresh、Alloc G 節 N_refresh 高端列；L1_Ans3_Hi（預期不變）"),
    _ev(6, "X3 1.1（GM275）", "Workload G9 每輪可見輸出 o（GM275）",
        "已搜尋（查詢詞：agentic coding average output tokens per turn visible response tokens OpenRouter）；只得「agentic coding 輸出約占總 token 0.5%」單一工作階段樣本（5,300／996,500）",
        "輸出占總 token 約 0.5%（單一樣本）；無每輪分布", "基準值（區間 ×0.5–×2）",
        "無可比對象（已搜尋；Andy 2026-10-06 核准，C2）", "查詢詞照 v5.19 查核檔。"),
    _ev(7, "X3 1.1（GM354）", "Calib F69 η_p 倍數（GM354）",
        "已搜尋（查詢詞：Vera Rubin NVL72 MLPerf Inference v6.1 prefill throughput）；MLPerf v6.1 VR200 只有整體吞吐，無 prefill 拆分",
        "整體吞吐（DeepSeek-R1）；無 prefill 專屬數據", "1（區間 0.5–1.5）",
        "無可比對象（現有資料無法分離；Andy 2026-10-06 核准，C2）",
        "「不可從現有資料分離」是主題的性質，非「不存在」。"),
    _ev(8, "X4 2（GM344／Calib F67）", "Calib F67 η_d 倍數（GM344）",
        "Nebius 公布之 MLPerf Inference v6.1 DeepSeek-R1 數據，第三方換算 VR200 每 GPU 約 GB300 的 1.6–2×；Nebius 量測 VR200 36 GPU、GB300 72 GPU；情境（offline／server／interactive）未讀到",
        "VR200／GB300 每 GPU 約 1.6–2×（第三方換算；Nebius 36／72 GPU）", "1.0（區間 0.5–1.5）；模型 interactive 2.525×，MLPerf v6.1 interactive 2.575×（Calib 第 150／151 列）",
        "待查證；指向 F67 低端；與模型內 MLPerf v6.1 interactive（Calib 第 151 列 2.575×，模型 2.525×）衝突未解",
        "F67 數值與區間不改；GM345（Rubin Ultra G67）不動。原文補查：Andy 手動下載 MLCommons v6.1 結果表後另開工作單。"
        " 36 與 72 GPU 兩個數字為 chat 端補入，不在 v5.19 查核檔內。",
        stance="混合（NVIDIA 宣稱為 Interested-party）", affected="Calib F67；IF_ 各世代 η_d 相關列（數值不改）"),
]

def evidence_rows():
    return EVIDENCE_V520

# ------------------------------------------------------------------ Gov_Map (guarded text appends; GM578 M)
def _ev_id(i): return f"E{236 + i}"
GOV_NOTES = [   # (GM_ID, text appended to N)
    ("GM205", f" ｜v5.20 X3：6.2 齊備（DB_Evidence {_ev_id(0)}；摘要級）"),
    ("GM203", f" ｜v5.20 X3：6.2 齊備（DB_Evidence {_ev_id(1)}；摘要級）"),
    ("GM459", f" ｜v5.20 X3：6.2 齊備（DB_Evidence {_ev_id(2)}；摘要級）"),
    ("GM270", f" ｜v5.20 X3：6.2 齊備（DB_Evidence {_ev_id(3)}；摘要級）"),
    ("GM290", f" ｜v5.20 X3：6.2 齊備（DB_Evidence {_ev_id(4)}；摘要級）；基準 0＝標準 harness 為單代理，屬定義"),
    ("GM578", f" ｜v5.20 X3：6.2 齊備（DB_Evidence {_ev_id(5)}；摘要級）；區間上限 4→6（Andy 2026-10-06，X3）"),
    ("GM275", " ｜v5.20 X3：已搜尋、無可比對象（Andy 2026-10-06 核准，C2）"),
    ("GM354", " ｜v5.20 X3：已搜尋、無可比對象（Andy 2026-10-06 核准，C2）"),
    ("GM344", f" ｜v5.20 X4：低端有一筆未查證第三方換算（DB_Evidence {_ev_id(8)}），原文補查中"),
]
GM578_M = ("—", "0–6（上限依 GPT-5.x 點版本節奏，Andy 2026-10-06，X3）")

def gov_update(ws, append_text):
    rows = {ws[f"A{r}"].value: r for r in range(5, ws.max_row + 1) if ws[f"A{r}"].value}
    n = 0
    for gm, txt in GOV_NOTES:
        r = rows[gm]
        if append_text(ws[f"N{r}"], txt): n += 1
    r = rows["GM578"]
    if ws[f"M{r}"].value == GM578_M[0]:
        ws[f"M{r}"].value = GM578_M[1]; n += 1
    return n

# ------------------------------------------------------------------ Decisions
C2_APPEND = "；v5.20 追加：GM275（Workload G9）、GM354（Calib F69）（Andy 2026-10-06 核准，X3）"
C2_ORIGIN = "；Andy 2026-10-06 核准（X3，經 chat 端轉達）"

DECISIONS_V520 = [
    ['X3', 'v5.20', 'X2 新高段 8 列補齊（6.2）、GM578 區間上限 4 → 6、C2 追加 GM275／GM354',
     'A1 GM205、GM203、GM459、GM270 判齊備；A2 GM290 判齊備（基準 0＝標準 harness 為單代理，屬定義）；A3 GM578（Alloc_In C7）區間上限 4 → 6，基準 2 不變（Alloc_In!E7＝AL_Nrefresh_Hi）；'
     'A4 GM275、GM354 核准為「已搜尋、無可比對象」，併入 Decisions C2 適用列名。證據 DB_Evidence E236–E243（全部 Analogy、摘要級、原文未讀）。',
     '2026-10-06', '「all ok」（回覆 chat 端 A1–A4 建議）', 'v5.20 已建',
     'DB_Evidence E236–E243；Gov_Map GM205／203／459／270／290／578／275／354；Alloc_In!E7；Decisions C2；docs/stage2_source_check_rules.md', '工作單 v5.20 第 0、1 節', '否'],
    ['X4', 'v5.20', 'Calib F67 維持，記錄衝突證據',
     'Calib F67 維持基準 1.0 與區間 0.5–1.5；記錄 Nebius 第三方換算（VR200 每 GPU 約 GB300 的 1.6–2×）與模型內 MLPerf v6.1 interactive（2.575×）的衝突未解（DB_Evidence E244）；'
     '原文補查由 Andy 手動下載 MLCommons v6.1 結果表後另開工作單；GM345（Rubin Ultra G67）不動。',
     '2026-10-06', '「all ok」（回覆 chat 端 B (a)＋(c)）', 'v5.20 已建',
     'DB_Evidence E244；Gov_Map GM344（N 欄）', '工作單 v5.20 第 0、2 節', '否'],
    ['X5', 'v5.20', 'Hopper 攤提耦合維持，加註說明',
     'IF_FullCost_* 含自下而上攤提，攤提分母為層級組合的機隊 token；同一世代任一層級 SLO 不可達時，該世代所有層級的 FullCost 均回傳文字（例：Hopper 在 C18 或 CTL_ProdDerate ≤ 約 0.72）；'
     'IF_RevGW_* 只依自身層級，可能仍為數值。設計不改，只在 Interface A3 與 README 加註。',
     '2026-10-06', '「all ok」（回覆 chat 端 C (a)）', 'v5.20 已建',
     'Interface!A3；README（Interface 攤提耦合）', '工作單 v5.20 第 0、3 節', '否'],
]

def decisions_update(ws, append_text):
    """Append the GM275／GM354 clause to Decisions C2 (D = decision text, F = Andy's words); guarded by presence."""
    n = 0
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == "C2":
            if append_text(ws.cell(r, 4), C2_APPEND): n += 1
            if append_text(ws.cell(r, 6), C2_ORIGIN): n += 1
    return n

# ------------------------------------------------------------------ X5 note (Interface A3, README row)
X5_NOTE = ("IF_FullCost_* 含自下而上攤提，攤提分母為層級組合的機隊 token；同一世代任一層級 SLO 不可達時，該世代所有層級的 FullCost 均回傳文字"
           "（例：Hopper 在 C18 或 CTL_ProdDerate ≤ 約 0.72）。IF_RevGW_* 只依自身層級，可能仍為數值（X5，Andy 2026-10-06 維持此設計）")
README_VERSION = (VERSION + "（X3：X2 新高段 8 列補齊——DB_Evidence E236–E243、Gov_Map GM205／203／459／270／290／578 判 6.2 齊備、GM275／354 併入 Decisions C2；"
                  "Alloc_In E7（N_refresh 區間上限）4 → 6；X4：DB_Evidence E244 記錄 Calib F67 衝突證據（F67 不改）；X5：Interface A3 與 README 加註攤提耦合；"
                  "工作單 docs/workorders/20261006_v5.20.md）。以下為 v5.19（")

README_TITLE = (VERSION.split("_")[-1] + " — Block 1＋2＋3＋4＋5＋6＋第 0 層 Source：機架規格、每 GW 成本、各層級產出與每 token 成本、訓練與研發計畫、"
                "理論營收與單價前緣、harness 與每成功任務成本、研發與服務算力配置（Alloc）、L1 常用推算值、Interface 下游介面")
