#!/usr/bin/env python3
"""產生 builder/gov_seed4.py（v5.18 的種子資料：新 SRC 紀錄、DB_Evidence、Decisions、SRC／Gov_Map 的舊值守衛更新）。

來源：docs/reports/20261005_stage2-1_查核.xlsx（審查用 Excel；工作單 v5.18 第 1、3 節所指的「來源資料以審查用 Excel 為準」）
與 model/archive/20261005_Tokenomics_v5.17.xlsx（讀出各格的 v5.17 舊值，作為守衛：只在格仍為舊值時才寫）。
用法：python3 tools/gen_seed_v518.py   （輸出覆蓋 builder/gov_seed4.py；結果已提交，建置不需執行本檔）
"""
import json, re, sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
REVIEW = ROOT / "docs/reports/20261005_stage2-1_查核.xlsx"
BASE = ROOT / "model/archive/20261005_Tokenomics_v5.17.xlsx"
OUT = ROOT / "builder/gov_seed4.py"
TODAY = "2026-10-05"

rv = openpyxl.load_workbook(REVIEW)
base = openpyxl.load_workbook(BASE)
rec = {}
for r in rv["查核紀錄"].iter_rows(min_row=2, values_only=True):
    if r[0]:
        rec[r[0]] = r

SHEET_OF = {"MOD": "SRC_Model", "HW": "SRC_HW", "PERF": "SRC_Perf", "CAP": "SRC_Cap", "DEM": "SRC_Demand"}
def sheet_of(sid): return SHEET_OF[sid.split("_")[1]]

# ----- 第 1 節：1 級（S1 (a)，值相同、已讀原文）；MOD_015 另依 E3 拆為兩筆
G1_LIST = ("MOD_002 MOD_006 MOD_007 MOD_009 MOD_010 MOD_011 MOD_012 MOD_013 MOD_014 MOD_016 MOD_017 "
           "HW_001 HW_025 HW_028 HW_040 HW_043 HW_047 PERF_036 PERF_037 CAP_008 CAP_011 DEM_013").split()
# 另：HW_014、DEM_011（2→1）、HW_004（升 1，HPE QuickSpecs）、MOD_015／MOD_054（E3）、2 級者（G1 規則與 H1 規則）各自處理
G2_LIST = "HW_003 HW_005 HW_006 HW_010 HW_015".split()       # G1 規則（2 級）
H1_LIST = "PERF_009 PERF_010".split()                        # H1 規則（2 級）

def reader_of(row):
    t = (row[8] or "") + (row[17] or "")
    if "Andy 提供" in t or "Andy 已讀" in t: return "Andy（G2）"
    if "chat 端" in t and "CC 本環境" in t: return "chat 端（CC 本環境另有讀取，見備註）"
    return "CC"

def short(s, n=260):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[:n - 1] + "…"

def first_url(s):
    m = re.search(r"https?://[^\s；（）]+", s or "")
    return m.group(0) if m else short(s, 120)

evid = []                       # DB_Evidence rows
src_upd = {}                    # sid -> {"sheet", "step", "fields": {col: (old,new)|("+",text)}}
nxt = [170]
def eid():
    e = f"E{nxt[0]}"; nxt[0] += 1; return e

def cur_row(sid):
    ws = base[sheet_of(sid)]
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == sid: return ws, r
    raise KeyError(sid)

def val(sid, col):
    ws, r = cur_row(sid); return ws[f"{col}{r}"].value

COLN = {"C": 3, "D": 4, "E": 5, "K": 11, "J": 10, "N": 14, "Q": 17, "R": 18, "S": 19, "W": 23, "M": 13}
def add_upd(sid, fields, step=None):
    d = src_upd.setdefault(sid, {"sheet": sheet_of(sid), "step": None, "fields": {}})
    d["fields"].update(fields)
    if step: d["step"] = step   # 值欄位的更新才需要步驟（歸因用）；等級與備註不屬步驟
    return d

def mk_evidence(sid, grade, rule, hand_new, extra_note=""):
    row = rec[sid]
    e = eid()
    claim = f"{row[1]}（{sid}）：值 {row[3]}"
    ver = "一手（已讀原文）" if grade == 1 else "二手（已讀）"
    note = (f"讀取者：{reader_of(row)}｜讀取日 {row[7]}｜{rule}｜摘錄：{short(row[9], 200)}"
            f"｜換算：{short(row[11], 120)}｜立場：{short(row[14], 120)}{('｜' + extra_note) if extra_note else ''}")
    evid.append([e, TODAY, short(claim, 200), short(first_url(row[6]), 250), ver, short(row[2], 120), short(row[3], 60),
                 short(row[10], 160), f"採納（{rule}）", "v5.18", short(note, 900), "已處理", sid, "—", "IF_*、L1_*（SRC 等級；數值不變）",
                 str(grade), short(row[14].split("：")[0] if row[14] else "—", 20)])
    return e

def grade_fields(sid, new_grade, e, hand_new=None, extra=None):
    ws, r = cur_row(sid)
    f = {"K": (ws[f"K{r}"].value, new_grade), "Q": ("+", f"；{e}"), "S": ("—", TODAY),
         "W": ("+", f" ｜v5.18：等級 {ws[f'K{r}'].value}→{new_grade}（{e}）" + (extra or ""))}
    old_hand = ws[f"N{r}"].value
    if hand_new and old_hand != hand_new: f["N"] = (old_hand, hand_new)
    return f

# 1 級
for k in G1_LIST:
    sid = f"SRC_{k}"
    e = mk_evidence(sid, 1, "S1 (a)；值相同", "一手（已讀）")
    add_upd(sid, grade_fields(sid, 1, e, hand_new="一手（已讀）"))

# HW_040：換算式（÷2）寫入
add_upd("SRC_HW_040", {"W": ("+", " ｜換算式：nvidia.com H100 規格表 FP8 Tensor Core 3,958 teraFLOPS（表末「* With sparsity」，SXM）÷2＝1,979 TFLOPS＝1.979 PF（NVIDIA dense＝含稀疏÷2 的慣例，[Derived]；BF16 1,979÷2＝989.5 對 SRC_HW_019 的 0.989）")})
# HW_047：GB200、GB300、VR200 三欄皆讀到
add_upd("SRC_HW_047", {"W": ("+", " ｜適用範圍：GB200、GB300、VR200 三種機架皆讀到原文，無範圍但書")})
# HW_001：原文直接寫出 10.2 kW
add_upd("SRC_HW_001", {"W": ("+", " ｜原文：DGX H100/H200 使用手冊 Table 3 “10.2 kW max.”，直接寫出，無需反推；手冊標題為 H100/H200 共用，SRC 適用對象為 H100")})
# DEM_011：Andy 讀官方繁中頁（G2）→ 2→1；出處補官方頁
r = rec["SRC_DEM_011"]
e = mk_evidence("SRC_DEM_011", 1, "S1 (a)；Andy 讀官方繁中頁（G2）", "一手（已讀）")
ws, rr = cur_row("SRC_DEM_011")
f = grade_fields("SRC_DEM_011", 1, e, hand_new="一手（已讀）")
f["J"] = ("+", "；OpenAI 官方繁中頁 https://openai.com/zh-Hant/index/accelerating-the-next-phase-ai/（Andy 讀取，2026-03-31：「我們的 API 目前每分鐘可處理超過 150 億個 Token。」）")
add_upd("SRC_DEM_011", f)

# HW_014：出處更正（Schneider 參考設計 MaxP），另列 Supermicro 產品頁（新 SRC_HW_064）；第二來源欄註「未確認獨立」
e1 = eid()
evid.append([e1, TODAY, "VR NVL72 參考設計 MaxP 227／MaxQ 188 kW/rack（SRC_HW_014）", "https://blog.se.com/datacenter/2026/05/08/nvidia-and-schneider-electric-get-in-sync-at-nvidia-gtc-2026-to-deliver-vera-rubin-ai-factories/",
             "一手（已讀原文）", "Spec_Rack F9（F8 低 188）", "227 kW/架（記為 Supermicro DLC-2 設計點）", "MaxP 227；MaxQ 188 kW/rack",
             "採納（S1 (a)；出處更正）", "v5.18",
             "讀取者：chat 端（2026-10-05，Andy 提供）＋CC（1d，curl HTTP 200 自讀）｜逐字：“supports operation at MaxQ 188 kw/rack and MaxP 227 kw/rack”｜"
             "原 SRC 記「Supermicro DLC-2 設計點」，Andy 讀新聞稿確認其全文無 227 kW，實際出處為本部落格（NVIDIA 參考設計 MaxP）｜立場：Interested-party（供電與冷卻設備商）",
             "已處理", "SRC_HW_014", "—", "IF_*、L1_*（等級；數值不變）", "1", "利害關係方"])
e2 = eid()
evid.append([e2, TODAY, "Supermicro Vera Rubin NVL72：DLC-2 液冷 “sized for 227 kW per rack”（SRC_HW_064）", "https://www.supermicro.com/en/accelerators/nvidia/vera-rubin",
             "一手（已讀原文）", "（SRC_HW_014 的第二來源；不直接連結模型格）", "—", "227 kW/架（冷卻 sized-for 容量）", "採納（並列第二來源）", "v5.18",
             "讀取者：CC（1d；curl 為 Akamai「Access Denied」403，改以 WebFetch 讀成，引文與搜尋結果同句）｜口徑＝DLC-2 液冷堆疊每機架設計容量，不是機架實測功率｜"
             "兩來源同為 227 可能同源於 NVIDIA 參考設計，獨立性未確認｜立場：Interested-party（系統廠）",
             "已處理", "SRC_HW_064", "—", "—", "1", "利害關係方"])
add_upd("SRC_HW_014", {"K": (3, 1), "Q": ("+", f"；{e1}"), "S": ("—", TODAY), "N": ("二手", "一手（已讀）"), "R": ("—", "SRC_HW_064（未確認獨立）"),
                       "J": ("+", "；Schneider Electric 部落格 2026-05-08（NVIDIA VR NVL72 參考設計 MaxP 227／MaxQ 188 kW/rack；v5.18 更正出處，原記 Supermicro DLC-2 設計點）"),
                       "W": ("+", f" ｜v5.18：等級 3→1（{e1}）；口徑＝參考設計機架 MaxP（IT 側）；第二來源 SRC_HW_064 未確認獨立")})

# HW_004：升 1 級（HPE QuickSpecs，工作單 G2）；上緣 130→132（步驟 G2）
e = eid()
evid.append([e, TODAY, "GB200 NVL72 機架功率：TDP 132 kW nominal（HPE QuickSpecs V6，2026-09-08）；120 kW（SemiAnalysis）", "HPE QuickSpecs V6（2026-09-08）；https://newsletter.semianalysis.com/p/gb200-hardware-architecture-and-component",
             "一手（已讀原文）", "Spec_Rack D8（低 120）、D9（基準 132）", "120–130 kW/架（基準 130）", "132 kW（TDP nominal）；120 kW", "採納（工作單 v5.18 G2；升 1 級）", "v5.18",
             "讀取者：HPE QuickSpecs＝chat 端（工作單 r2 第 2.2 節 G2 所述；CC 未獨立讀，hpe.com 被出口政策擋）；SemiAnalysis 120／123.6 kW＝CC 於 1d 讀取（二手，2 級）｜"
             "EDPp 192 kW 為電氣設計尖峰口徑，只登錄不採用（見 DB_Evidence「已讀，未採用」列）｜立場：HPE 為系統廠（利害關係方）",
             "已處理", "SRC_HW_004", "—", "IF_*、L1_*（D9 基準 130→132）", "1", "利害關係方"])
add_upd("SRC_HW_004", {"K": (3, 1), "Q": ("+", f"；{e}"), "S": ("—", TODAY), "N": ("二手", "一手（已讀）"),
                       "J": ("+", "；HPE QuickSpecs V6（2026-09-08）TDP 132 kW nominal（chat 端讀取）；SemiAnalysis GB200 hardware architecture（約 120 kW、含轉換損耗總用電 123.6 kW）"),
                       "W": ("+", f" ｜v5.18 G2：基準 130→132（TDP nominal）、低 120 不變；上緣 130→132；等級 3→1（{e}；HPE 頁由 chat 端讀取）")})
add_upd("SRC_HW_004", {"E": (130, 132)}, step="G2")

# HW_005／HW_003／HW_006／HW_010／HW_015：G1 規則（2 級）
G1_NOTES = {
 "SRC_HW_003": " ｜G1 規則：(i) 上限有 SHI 德州政府合約標價 $311,624 一手佐證（chat 端讀取，工作單 r2；CC 的搜尋摘錄見 $311,599.87，差 0.01%，頁面未讀）；(ii) 基準 C12＝1.1 M（275 K×4）；下限 220,000 未讀 [Assumed]。僅一個已讀值，(ii) 的「範圍」為單點",
 "SRC_HW_005": " ｜G1 規則：(i) wing.vc 讀到 132–140 kW（二手已讀）；(ii) 基準 E9 136（中點，Derived）落在 132–140；上緣 142 未讀，v5.18 改為 140（步驟 G3）",
 "SRC_HW_006": " ｜G1 規則：(i) wccftech 讀到 3.1 M（引自 X 貼文 @firstadopter，非 Morgan Stanley；出處更正）；(ii) 基準 D12＝3.1 與其相同；2.8、3.4 端未讀 [Assumed]（3.4 的搜尋摘要實指「72 GPU 模組合計」，口徑不同）",
 "SRC_HW_010": " ｜G1 規則：(i) wing.vc 讀到採購單「just under $5.0M」（不含機架內 CDU；原文 datagravity.dev 被出口政策擋）；(ii) 基準 E12＝5.0 與其相符（值非逐字相同）",
 "SRC_HW_015": " ｜G1 規則：(i) wccftech 讀到 Morgan Stanley Research「estimated BOM $7.8 million」（2026-05-21）；口徑＝成本（BOM），非售價；(ii) 本筆記錄值 7.8 即已讀值；Spec_Rack F12 基準 8.4＝7.8×(1＋代工毛利約 7.5%，區間 5–10%，Assumed)，F11 低 7.8、F13 高 9.1（步驟 G4／r3 I2）",
}
for k in G2_LIST:
    sid = f"SRC_{k}"
    e = mk_evidence(sid, 2, "G1 規則；二手已讀", None, extra_note="區間型紀錄：至少一值讀到原文、基準在已讀範圍")
    add_upd(sid, grade_fields(sid, 2, e, extra=G1_NOTES[sid]))
add_upd("SRC_HW_005", {"E": (142, 140)}, step="G3")
add_upd("SRC_HW_015", {"M": ("+", "；v5.18：口徑＝成本（Morgan Stanley Research 物料成本 BOM），不是售價")})

# PERF_009／010：D1 數值（內插式）＋H1 規則（2 級；W1 待第二來源）
H1_TEXT = {
 "SRC_PERF_009": ("9,384.4→9,342.1", " ｜D1：值改為 InferenceX 快照 0813 原始資料內插至互動性 72：9097.16＋(74.263−72)/(74.263−68.267)×(9746.09−9097.16)＝9,342.1 [Derived]（−0.45%）；頁面無逐字 9,384.4"),
 "SRC_PERF_010": ("3,473.9→3,405.2", " ｜D1：值改為同法內插至互動性 130：3055.6＋(136.29−130)/(136.29−118.29)×(4056.1−3055.6)＝3,405.2 [Derived]（−2.0%）；頁面無逐字 3,473.9"),
}
for k in H1_LIST:
    sid = f"SRC_{k}"
    e = mk_evidence(sid, 2, "H1 規則；上限 2 級", "一手（已讀）",
                    extra_note="H1：(i) 已讀一手原始資料（頁面內嵌資料列與前端程式）、(ii) 已搜尋同口徑第二來源（見下一列）、(iii) 口徑與方法已寫明")
    f = grade_fields(sid, 2, e, hand_new="一手（已讀）", extra=H1_TEXT[sid][1] + "｜W1：待第二來源（H1；SemiAnalysis 不得為唯一依據，MLPerf v6.0 口徑不同，見搜尋紀錄）")
    add_upd(sid, f)
add_upd("SRC_PERF_009", {"C": (9384.4, 9342.1)}, step="D1")
add_upd("SRC_PERF_010", {"C": (3473.9, 3405.2)}, step="D1")
e_h1 = eid()
evid.append([e_h1, TODAY, "InferenceX（SemiAnalysis）GB300 V4-Pro：同口徑第二來源搜尋紀錄（H1 規則 (ii)）", "https://developer.nvidia.com/blog/nvidia-platform-delivers-lowest-token-cost-enabled-by-extreme-co-design/",
             "二手（已讀）", "Calib C17／D17（SRC_PERF_009／010）", "—", "MLPerf Inference v6.0（GB300 NVL72，DeepSeek-R1）Offline 9,821、Server 8,064 tok/s/GPU", "已讀，口徑不同，不作為第二來源（H1 (ii) 已記錄）", "v5.18",
             "查詢詞：「MLPerf Inference v6.0 GB300 NVL72 DeepSeek tok/s/GPU」「InferenceX GB300 V4-Pro 72 tok/s second source」；結果：MLPerf 為 R1 模型（37B 啟用）非 V4-Pro（49B）、ISL／OSL 設定不同、NVIDIA 為賣方且自行挑選配置，只作量級對照；無同口徑獨立第二來源｜"
             "依 H1 規則記 2 級並標 W1「待第二來源」，不擋完成條件｜讀取者：CC（1，2026-10-05）",
             "已處理", "SRC_PERF_009；SRC_PERF_010", "—", "—", "2", "利害關係方"])

# MOD_015：E3 拆為兩筆（Flash 0.512＝21/41；Pro 0.492＝30/61），由 config compress_ratios 計算
e = eid()
evid.append([e, TODAY, "DeepSeek V4 全注意力層比例：由 config.json compress_ratios 結構計算（Flash 21/41＝0.512；Pro 30/61＝0.492）", "https://huggingface.co/deepseek-ai/DeepSeek-V4-Flash/raw/fd53f94/config.json；https://huggingface.co/deepseek-ai/DeepSeek-V4-Pro/raw/main/config.json",
             "一手（已讀原文）", "Arch C16（Luna／Flash）、D16（Sol／Pro）", "0.5（約值，兩層級共用）", "Flash 0.512；Pro 0.492", "採納（工作單 v5.18 E3）", "v5.18",
             "讀取者：CC（2026-10-05）｜Flash：壓縮比 4 的 21 層、128 的 20 層、0 的 2 層（共 43；compress_ratios 長度 44＝43＋1 個 MTP 層，不計）；Pro：4 的 30 層、128 的 31 層（共 61）｜"
             "比例＝壓縮比 4 的層數÷（壓縮比 4＋128 的層數）：Flash 21/41、Pro 30/61 [Derived]｜立場：DeepSeek 自揭（利害關係方）",
             "已處理", "SRC_MOD_015；SRC_MOD_054", "—", "IF_*、L1_*（見歸因 E3）", "1", "利害關係方"])
add_upd("SRC_MOD_015", {"K": (3, 1), "Q": ("+", f"；{e}"), "S": ("—", TODAY), "N": ("一手（未讀）", "一手（已讀）"),
                        "B": ("DeepSeek V4 全注意力層比例（CSA／HCA 約各半）", "DeepSeek V4-Flash 全注意力層比例（壓縮比 4 層／（壓縮比 4＋128 層）＝21/41）"),
                        "G": ("約值", "config.json compress_ratios 前 43 個元素；MTP 層不計"), "H": ("DeepSeek V4", "DeepSeek V4-Flash"),
                        "W": ("+", f" ｜v5.18 E3：0.5→0.512（21/41）；Pro 另列 SRC_MOD_054（{e}）")})
add_upd("SRC_MOD_015", {"C": (0.5, 0.512)}, step="E3")

# 立場欄：HW_040、HW_047、DEM_011 等照報告填寫（已在各自的 mk_evidence 備註；SRC 立場說明補一句）
add_upd("SRC_HW_040", {"M": ("+", "；v5.18 1c：NVIDIA 為賣方（nvidia.com 規格表，含稀疏口徑須 ÷2）")})
add_upd("SRC_HW_047", {"M": ("+", "；v5.18：NVIDIA 為賣方；72 為產品名稱定義，誘因影響極小")})
add_upd("SRC_DEM_011", {"M": ("+", "；v5.18：OpenAI 融資公告，有呈現成長的誘因；Andy 讀官方繁中頁（G2）")})

# ----- F1：VR200 Max-Q 190→188（Schneider 參考設計）；F9 227 不變；F10 230 不變（SRC_HW_012，另列判斷）
add_upd("SRC_HW_011", {"C": (190, 188), "J": ("+", "；Schneider Electric 部落格 2026-05-08（NVIDIA VR NVL72 參考設計 MaxQ 188 kW/rack；v5.18 F1）"),
                       "W": ("+", " ｜v5.18 F1：190→188（MaxQ；見 E 欄證據）；等級仍 3（未列入工作單升級清單，見待判斷）")}, step="F1")

# ----- 舊路線圖：Superseded（保留紀錄）
for sid, why in [("SRC_HW_013", "Rubin Ultra 單架 600 kW（Kyber 舊路線圖）"), ("SRC_HW_045", "Rubin Ultra NVL576 FP8 全架 5,000 PF（Kyber 舊路線圖）"), ("SRC_HW_054", "Rubin Ultra NVL576 144 封裝（Kyber 舊路線圖）")]:
    add_upd(sid, {"O": ("Active", "Superseded"), "P": ("—", "SRC_HW_061／062；Spec_Rack G 欄改單架 72 封裝"),
                  "W": ("+", f" ｜v5.18：Superseded（{why}）；NVIDIA 2026-09 技術部落格改為 MGX NVL 單架 72 GPU，Kyber NVL144 屬 Feynman 世代；保留紀錄")})

# ----- 新 SRC 紀錄
NEWREC = []
def newrec(**k):
    d = dict(s="U-V518", lo=None, hi=None, date="2026-09", status="Active", hand="一手（已讀）", grade=1, stance="利害關係方", ev="", use="", note="")
    d.update(k); NEWREC.append(d)
blog = "https://developer.nvidia.com/blog/nvidia-vera-rubin-pod-seven-chips-five-rack-scale-systems-one-ai-supercomputer"
e_ru = eid()
evid.append([e_ru, TODAY, "Rubin Ultra 欄改單架 72 封裝：NVL576＝8 個 MGX NVL 機架、每架 72 顆 Rubin Ultra GPU、單一 576 GPU NVLink 域", blog, "一手（已讀原文）",
             "Spec_Rack G6、G28（G8–G13、G23、G24、G40 隨之改為 Derived／Assumed）", "144 封裝／架（Kyber）", "72 封裝／架；NVLink 域 576",
             "採納（工作單 v5.18 2.3；E2、F2–F5）", "v5.18",
             "讀取者：CC（1b 起，2026-10-05）｜逐字：“Vera Rubin Ultra NVL576 will combine eight separate MGX NVL racks, each with 72 Rubin Ultra GPUs … single 576-GPU NVLink domain”；“Kyber NVL144 rack … double the NVLink domain per rack to 144 GPUs for the Feynman era”（不納入）｜"
             "「GPU」是封裝或 die 部落格未明載；推斷為封裝（Derived）：Kyber NVL144＝144 封裝＝舊 NVL576（576 die）；Schneider 稱 Rubin Ultra＝Kyber 576 GPU／架、1–1.2 MW，與 NVIDIA 2026-09 衝突，依 Andy 指示列較舊路線圖｜立場：Interested-party",
             "已處理", "SRC_HW_061；SRC_HW_062", "SRC_HW_013；SRC_HW_045；SRC_HW_054", "IF_*（Rubin Ultra 欄）、L1_*", "1", "利害關係方"])
newrec(id="SRC_HW_061", sheet="SRC_HW", metric="Rubin Ultra 單架 GPU 封裝數（MGX NVL 單架）", val=72, unit="封裝/架", basis="每架 72 顆 Rubin Ultra GPU（封裝為推斷 Derived）",
       applies="Rubin Ultra（MGX NVL 單架）", src="NVIDIA 技術部落格 Vera Rubin POD（2026-09）：" + blog, stance_note="NVIDIA：產品規劃揭露（尚未出貨），有行銷誘因",
       use="Spec_Rack!G6｜直接", note="v5.18 新增；取代 SRC_HW_054（Kyber 144）", ev=e_ru)
newrec(id="SRC_HW_062", sheet="SRC_HW", metric="Rubin Ultra NVL576 NVLink 域 GPU 數", val=576, unit="GPU", basis="8 個 MGX NVL 單架互連為單一 NVLink 域（兩層拓撲）",
       applies="Rubin Ultra NVL576", src="NVIDIA 技術部落格 Vera Rubin POD（2026-09）：" + blog, stance_note="NVIDIA：產品規劃揭露（尚未出貨）",
       use="Spec_Rack!G28｜直接", note="v5.18 新增；單架內全互連為 72（Gov_Map G28 低 72）", ev=e_ru)
e_dgx = eid()
evid.append([e_dgx, TODAY, "DGX Vera Rubin NVL72：NVFP4 訓練 2,520 PFLOPS（Dense）、FP8/FP6 訓練 1,260 PFLOPS、NVFP4 推論 3,600 PFLOPS（72 GPU）", "https://www.nvidia.com/en-us/data-center/dgx-vera-rubin-nvl72/",
             "一手（已讀原文）", "Spec_Rack G23、G24、G40（Derived 依據）", "—", "NVFP4 訓練 2,520÷72＝35 PF/GPU（註 Dense）；FP8/FP6 訓練 1,260÷72＝17.5；NVFP4 推論 3,600÷72＝50（註 Sparse）",
             "採納（Derived 依據；工作單 v5.18 2.3）", "v5.18",
             "讀取者：CC（1c，2026-10-05；nvidia.com HTTP 200）｜註腳 “Dense specification”／“Sparse specification”｜Rubin Ultra 為 4 die（Rubin 2 die）：G23＝2×35＝70；G40＝2×17.5＝35.0；G24＝2×50＝100 [Derived]｜立場：NVIDIA 為賣方",
             "已處理", "SRC_HW_063", "—", "IF_*（Rubin Ultra 欄）", "1", "利害關係方"])
newrec(id="SRC_HW_063", sheet="SRC_HW", metric="DGX Vera Rubin NVL72 NVFP4 訓練（整架，72 GPU）", val=2520, unit="PFLOPS/架", basis="Dense specification；72 GPU（35 PF/GPU）",
       applies="VR200 NVL72", date="2026", src="NVIDIA DGX Vera Rubin NVL72 頁：https://www.nvidia.com/en-us/data-center/dgx-vera-rubin-nvl72/", stance_note="NVIDIA 為賣方（規格頁）",
       use="（Spec_Rack G23／G24／G40 的 Derived 依據；不直接連結模型格）", note="v5.18 新增；Rubin Ultra 每封裝值＝2×VR 每 GPU 值（4 die 對 2 die）[Derived]；FP8/FP6 訓練 1,260、NVFP4 推論 3,600 同頁", ev=e_dgx)
newrec(id="SRC_HW_064", sheet="SRC_HW", metric="Supermicro VR NVL72 DLC-2 液冷設計容量", val=227, unit="kW/架", basis="冷卻 sized-for 容量（不是機架實測功率）",
       applies="VR200 NVL72", src="Supermicro Vera Rubin 產品頁：https://www.supermicro.com/en/accelerators/nvidia/vera-rubin", stance_note="Supermicro：系統廠，賣方",
       use="（SRC_HW_014 的第二來源；不直接連結模型格）", note="v5.18 新增；與 SRC_HW_014（Schneider MaxP 227）同為 227、口徑不同，獨立性未確認", ev=e2)
newrec(id="SRC_MOD_054", sheet="SRC_Model", metric="DeepSeek V4-Pro 全注意力層比例（壓縮比 4 層／（壓縮比 4＋128 層）＝30/61）", val=0.492, unit="%", basis="config.json compress_ratios 前 61 個元素",
       applies="DeepSeek V4-Pro", date="2026", src="https://huggingface.co/deepseek-ai/DeepSeek-V4-Pro/raw/main/config.json", stance_note="DeepSeek 自揭",
       use="Arch!D16｜直接", note="v5.18 E3 由 SRC_MOD_015 拆出（Flash 0.512 留在 015）", ev=e)

# ----- 非 SRC 證據（DB_Evidence）：採用為區間依據、已讀未採用、規則與判斷
def ev_row(claim, source, hand, loc, cur, new, verdict, note, status="已處理", sid="—", effect="—", grade="—", stance="—"):
    e = eid()
    evid.append([e, TODAY, short(claim, 200), short(source, 250), hand, short(loc, 120), short(cur, 60), short(new, 160), verdict, "v5.18",
                 short(note, 900), status, sid, "—", effect, grade, stance])
    return e

COP = "https://arxiv.org/abs/2608.00101（Agentic Coding in the Wild: Characterizing GitHub Copilot Traces at Production Scale；HTML 全文 arxiv.org/html/2608.00101v1）"
DSO = "https://raw.githubusercontent.com/deepseek-ai/open-infra-index/main/202502OpenSourceWeek/day_6_one_more_thing_deepseekV3R1_inference_system_overview.md"
E_T = ev_row("Copilot 追蹤：每工作階段使用者輪中位 3／平均 6.1／P90 15；LLM 呼叫中位 15／平均 40.6", COP, "一手（已讀原文）", "Workload 5（輪數 T；Gov_Map E5:G5 區間）", "E:G＝4／4／30",
             "不改（r3 I1）", "已讀，已涵蓋（不改）", "讀取者：CC｜Table 4｜單位對應：Workload T 以每次新輸入（使用者或工具結果）計一輪＝LLM 呼叫數；Copilot『LLM 呼叫平均 40.6』對應 G 欄（Coding agent，區間 15–60，已涵蓋）；『使用者輪中位 3』單位不同，不套用；E、F 非 coding 工作負載｜Analogy｜立場：Microsoft 研究（中立偏利害）",
             effect="—（區間與基準皆不改）", grade="Analogy", stance="中立偏利害")
E_CHI = ev_row("Copilot：KV 快取命中率輪內平均約 90%、跨輪邊界 55%、呼叫層中位約 98%", COP, "一手（已讀原文）", "Workload 11（歷史快取命中率 χ；Gov_Map E11:G11 區間）", "E:G 區間 0.7–0.95",
               "不改（r3 I1）", "已讀，已涵蓋（不改）", "讀取者：CC｜單位對應：T 以呼叫計，相鄰呼叫多屬輪內（命中 90–94%）；跨使用者輪 55% 約占 40 次中 2 次，加權約 0.90；現區間 0.70–0.95 已涵蓋｜Analogy｜立場：Microsoft 研究（中立偏利害）",
               effect="—（區間與基準皆不改）", grade="Analogy", stance="中立偏利害")
E_CHI2 = ev_row("DeepSeek 推論系統概述：24 小時輸入 608 B token 中 342 B 命中（56.3%）；每節點吞吐 73.7k／14.8k tok/s 為 prefill／decode 階段中的值，非 24 小時均值", DSO, "一手（已讀原文）",
                "Workload 11 C:D（一般對話；區間 0.3–0.7 已涵蓋 0.563，不動）；Serving 17／18（未採用）", "C:D 區間 0.3–0.7", "χ C:D 不改（現區間 0.3–0.7 已涵蓋 0.563）；Serving 17／18 只蒐集證據", "已讀，已涵蓋（χ C:D 不改）；Serving 17／18 未採用（S2 (a)）",
                "讀取者：CC｜節點利用（平均÷尖峰節點）＝226.75/278＝0.816 [Derived]｜硬體 H800、模型 V3／R1、2025-02，與第 0 層差異大，僅作量級參考；前次『日均÷73.7k＝0.42／0.58』口徑有誤已更正｜立場：DeepSeek 自揭（利害關係方）",
                effect="—", grade="Analogy", stance="利害關係方")
E_ISL = ev_row("Copilot：每呼叫提示 token 中位 68K（快取 63K）；每使用者輪提示中位 160.2K", COP, "一手（已讀原文）", "Serving 23（參考 ISL；Gov_Map C23:E23 區間）", "×0.5–×2",
               "D23、E23 區間高擴至 68,000；C23 不改（r3 I1：只套中、頂層）", "採用為區間依據（Analogy；D23、E23）", "讀取者：CC｜Figure 10、Table 4｜不用每使用者輪 160K（口徑不同）｜差異：編碼代理的累積上下文，僅代理層有對應；低層（Luna）不套用｜基準不動｜立場：Microsoft 研究",
               effect="區間（基準輸出影響 0）", grade="Analogy", stance="中立偏利害")
E_OSL = ev_row("Copilot：每呼叫輸出 token 中位 247（88% 呼叫 <1,000；不含推理 token）", COP, "一手（已讀原文）", "Serving 24（參考 OSL，含思考）", "1024／2048／4096",
               "不套用（判為不可比）", "已讀，未採用", "讀取者：CC｜不含推理 token，口徑與模型 OSL（含思考）不同；單一產品；僅作下界參考｜OpenRouter State of AI PDF 被出口政策擋（R5），仍缺含推理 token 的分布｜工作單 2.4：Serving 24 不套用（chat 端判讀，Andy 可改）",
               effect="—", grade="—", stance="中立偏利害")
E_MFU = ev_row("MegaScale-MoE（arXiv 2505.11432）：Hopper 上 MFU 27.89–32.48%；原文稱 MFU 隨 GPU 運算能力增加而下降", "https://arxiv.org/abs/2505.11432", "一手（已讀原文）", "Train_In 21（預訓練 MFU 世代倍數）", "Hopper 0.8→1.0（E1）",
               "—", "已讀，未採用", "讀取者：CC｜方向性證據（Blackwell 的 MFU 不高於 Hopper）與模型『Hopper 0.8、GB200 1』相反；僅有 Hopper 世代內量測，無 Blackwell／Rubin 絕對值；E1 依 Andy 決定改 Hopper 1.0（區間 0.8–1.2），此證據不採用，列待判斷",
               effect="—", grade="Analogy", stance="利害關係方")
E_EDP = ev_row("HPE QuickSpecs：GB200 EDPp 約 192 kW、GB300 EDPp 約 155 kW（電氣設計尖峰）", "HPE QuickSpecs V6（2026-09-08；搜尋摘要，hpe.com 被出口政策擋）", "二手（只見摘錄）", "Spec_Rack D8:E10（機架功率）", "—",
               "—", "已讀，未採用（口徑不同）", "電氣設計尖峰口徑（≈1.5×TDP），與模型採用的 TDP／設計功率口徑不同；只登錄於 DB_Evidence（工作單 v5.18 G2）", effect="—", grade="3", stance="利害關係方")
E_SMC = ev_row("SemiAnalysis（搜尋摘錄）：Supermicro SRS-VR-NVL72 4×110 kW 電源架供 ≤220 kW", "SemiAnalysis（搜尋摘錄，未讀）", "只見摘錄", "Spec_Rack F9（VR200 機架功率）", "—",
               "—", "已讀（摘錄），未採用", "僅搜尋摘錄、未讀原文；若屬實可用容量略低於 MaxP 227；Supermicro Datasheet 440 kW（含備援，Verified）、CDU 2.2 MW÷8＝275 kW/架（Derived）並不衝突｜工作單 v5.18 3.5 指定登錄為「已讀，未採用」",
               effect="—", grade="3", stance="利害關係方")
E_FLEET = ev_row("Epoch（轉述 NVIDIA 揭露）：至 2025-10 累計出貨 Hopper 4 M、Blackwell 3 M → Hopper 占 4/7＝0.571", "Epoch AI（1d 重跑讀取；NVIDIA 出貨揭露轉述）", "二手（已讀）", "Checks 對帳常數『2025 Hopper 占機隊』（I46）", "0.6（Assumed）",
                 "0.57；區間 0.5–0.6（D4）", "採納（Analogy）", "讀取者：CC｜口徑：全球累計出貨，非 OpenAI 機隊；OpenAI 自身組成未找到出處；Analogy 區間 0.5–0.6", effect="L1_RevGWFleet_OAI2025、Checks 43", grade="Analogy", stance="中立")
E_192 = ev_row("InferenceX 頁面 tok/s/MW＝tok/s/GPU÷固定常數（GB300 2.12 kW/GPU）：9,506.83÷4,484,352＝2,120.0 W", "https://inferencex.semianalysis.com/compare-per-dollar/deepseek-v4-b300-vs-gb300（頁面資料列與前端程式）", "一手（已讀原文）",
               "Checks 第 21 列（GB300 每 GPU 功率對照 2.12 kW）", "『9,384÷4.43 M 反推』", "改述為『InferenceX 內部常數，非獨立對照』", "採納（工作單 v5.18 3.3）",
               "讀取者：CC｜頁面的 4.43 M tok/s/MW 即由其前端程式固定常數換算，由 9,384 與 4.43 M 反推 2.12 是循環；常數是否含設施頁面未載明；模型 IT 平均 1.52 kW/GPU，2.12÷1.52＝1.39｜立場：SemiAnalysis（利害關係方）",
               effect="Checks 21（不再作外部驗證）", grade="—", stance="利害關係方")
E_KV = ev_row("Arch C34／D34 每層 KV bytes：由 config 結構重算 (a)只計壓縮 KV：Flash 64.37／Pro 64.98；(b)計入 indexer key：80.00／80.72；(c)bf16 KV：128.74／129.97", "https://huggingface.co/deepseek-ai/DeepSeek-V4-Flash/raw/fd53f94/config.json；https://huggingface.co/deepseek-ai/DeepSeek-V4-Pro/raw/main/config.json；官方 inference/model.py（commit fd53f94）",
              "一手（已讀原文）", "Arch C34:D34（Gov_Map 區間）", "65／65（基準；不動）", "基準 65；高 130（D3／E4）", "採用為區間上緣（bf16 KV 涵蓋 indexer key 情境）",
              "讀取者：CC（B5 結構重算，engine 唯讀）｜(a) 與現值 65 相差 −1.0%／−0.03%，影響 <1e-5；官方參考程式以預設 dtype（bf16）配置 kv_cache，模型假設 FP8（SRC_HW_055），若生產為 bf16 則約 128.7／130.0；生產精度未找到出處｜基準不動，只擴大區間",
              effect="—（區間；(c) 影響最大 L1_Ans9 +1.1e-2、IF_FullCost_Sol(VR200) +7.7e-4）", grade="Derived", stance="利害關係方")
E_BOM = ev_row("G4／I2：F12 基準 8.4＝Morgan Stanley BOM 7.8 ×（1＋代工毛利約 7.5%，區間 5–10%）；低 7.8（零毛利 BOM 下限）；高 9.1（不變）", "wccftech（Morgan Stanley Research，2026-05-21）；毛利率為 Assumed", "二手（已讀）", "Spec_Rack F11:F13（VR200 機架價格）", "6.7／7.8／9.1",
               "7.8／8.4／9.1", "採納（工作單 v5.18 r3 G4／I2；Derived）", "代工毛利 5–10% 為 Assumed（無來源）；MS 為成本（BOM）、Bernstein 9.1 M 為售價口徑（摘要，未讀）；r3 I2：BOM 7.8 本身為分析師估計，區間須保留 BOM 不確定，故低取零毛利 BOM 7.8、高維持 9.1（SRC_HW_017，不變）；原 F11 6.7（SRC_HW_016）不再連結，記錄保留",
               sid="SRC_HW_015", effect="IF_FullCost_Sol(VR200)、L1_Ans3／Ans4（基準預期約 +6%）", grade="Derived", stance="中立")
E_EPOCH = ev_row("Epoch AI：最終訓練占 R&D 算力支出 OpenAI 9.6%、MiniMax 22.6%、Z.ai 12.3%（倍數 10.4／4.4／8.1）", "https://epoch.ai/gradient-updates/r-and-d-vs-training-compute", "一手（已讀原文）", "Train_In 13（研發倍數；SRC_DEM_001–003）", "8（4.4–10.4）",
                 "不改", "已讀，未新增（與現區間完全相同）", "讀取者：CC｜無新獨立對象｜立場：Epoch AI（中立）", effect="—", grade="—", stance="中立")
E_SEARCH_LAT = ev_row("每層延遲倍數（Calib 68）、RL rollout token（Train_In 34）：以查詢詞搜尋，找不到原始出處", "搜尋引擎、github raw、arXiv", "找不到原始出處（R6）", "Calib 68；Train_In 34", "Hopper 1.5／其餘 1；2.36 T／9.6 T",
                      "不動（缺口）", "已搜尋，無候選（缺口）", "查詢詞：「DeepSeek inference system overview 226.75 nodes …」「Hopper expert parallel per-layer latency InfiniBand」「RL rollout tokens frontier lab total」「GRPO rollout tokens per step」；已查：github raw 概述、arXiv、搜尋引擎｜"
                      "DeepSeek 概述只給吞吐、無每層延遲；論文為研究規模，無前沿實驗室總 rollout token；R1 RL÷V3 預訓練 GPU 小時比 0.055 僅為間接對照", effect="—", grade="—", stance="—")
E_J = ev_row("J6 基準利用率／J15 生產折減：蒐集證據（S2 (a)）", DSO, "一手（已讀原文）", "Serving 17（基準利用率 0.6）、Serving 18（實測→生產折減 1.0）", "0.6／1.0",
             "不提議數值、不改暫用值（G0-11）", "已讀，未採用（S2 (a)）", "見 E{0}（DeepSeek 概述）；InferenceX 無 H800／H200 同模型點；η 拆解依 Andy『有新數據再說』暫緩".format(E_CHI2[1:]), effect="—", grade="—", stance="—")
E_IX = ev_row("OpenAI 單位 token 需求、harness 殘差倍數、Alloc 家族計畫數 N 與規模倍數 k、Arch E7 Astra 啟用參數、Train_In E25 Astra 預訓練 token：已知無來源之 Assumed 參數", "（各列 Gov_Map 理由欄）", "—", "Workload 13、Alloc_In C6／C8、Arch E7、Train_In E25 等", "—",
              "—", "區間依據＝Gov_Map 理由欄（Claude 提議，Andy 2026-10-03 確認）；無外部來源", "完成檢查 6.2 逐列列出三項；無外部來源者列為『已搜尋：（無）／候選：（無）／區間依據：Andy 確認』，三項不齊處在報告標出", effect="—", grade="—", stance="—")

# ----- W2 殘餘 3 格的候選證據（不寫入等級；S1 (a) 候選，交 chat 端審查；狀態＝待判定）
E_C022 = ev_row("SRC_HW_022 候選升級：GB200（Blackwell）NVFP4 dense 10 PF/GPU（NVIDIA Blackwell Ultra 技術部落格規格表）", "https://developer.nvidia.com/blog/inside-nvidia-blackwell-ultra-the-chip-powering-the-ai-factory-era/", "一手（已讀原文）",
                "Spec_Rack D23、D24", "10 PF/GPU（3 級）", "10 PF（dense）", "候選（未寫入等級）", "讀取者：CC（1b，2026-10-05）｜“NVFP4 dense | sparse performance … 10 | 20 PetaFLOPS 15 | 20 PetaFLOPS”（左欄 Blackwell＝10、右欄 Blackwell Ultra＝15）；值相同、已讀原文，符合 S1 (a)，但不在工作單 1 級清單；寫入後 W2 −1（GM106）",
                status="待判定", sid="SRC_HW_022", grade="1（建議）", stance="利害關係方")
E_C027 = ev_row("SRC_HW_027 候選升級：Rubin NVFP4 推論 50 PF/GPU（NVIDIA Rubin 平台技術部落格 Table 2；DGX VR 頁 3,600÷72＝50）", "https://developer.nvidia.com/blog/inside-the-nvidia-rubin-platform-six-new-chips-one-ai-supercomputer/；https://www.nvidia.com/en-us/data-center/dgx-vera-rubin-nvl72/", "一手（已讀原文）",
                "Spec_Rack F24", "50 PF/GPU（3 級）", "50 PF（註腳 Transformer Engine compute；DGX VR 頁註 Sparse specification）", "候選（未寫入等級）", "讀取者：CC（1b、1c，2026-10-05）｜口徑與 SRC 的『含自適應壓縮』一致；值相同、已讀原文，符合 S1 (a)，但不在工作單 1 級清單；寫入後 W2 −1（GM108）",
                status="待判定", sid="SRC_HW_027", grade="1（建議）", stance="利害關係方")
E_C053 = ev_row("SRC_HW_053 候選升級：HGX H100 8 GPU 板（DGX H100/H200 使用手冊）", "https://docs.nvidia.com/dgx/dgxh100-user-guide/introduction-to-dgxh100.html", "一手（已讀原文）",
                "Spec_Rack C28（Hopper Scale-up 域 GPU 數）", "8（3 級）", "8 x NVIDIA H100 GPUs；4 x 4th generation NVLinks 900 GB/s GPU-to-GPU", "候選（未寫入等級）", "讀取者：CC（2026-10-05，WebFetch）｜“8 x NVIDIA H100 GPUs that provide 640 GB total GPU memory”；“4 x 4th generation NVLinks that provide 900 GB/s GPU-to-GPU bandwidth”（兩句分列，頁面未以單句寫『8 GPU 同一 NVLink 域』，域大小由 NVSwitch 拓撲推得 [Derived]）；寫入後 W2 −1（GM125）",
                status="待判定", sid="SRC_HW_053", grade="1（建議，Derived 一環）", stance="利害關係方")

E_RUP = ev_row("Rubin Ultra 單架價格：以查詢詞搜尋，無一手報價", "搜尋引擎（無頁面可指定）", "找不到原始出處（R6）", "Spec_Rack G11:G13（Rubin Ultra 機架價格）", "23.4／31.2／39（144 封裝）",
               "11.7／15.6／19.5（72 封裝；Assumed）", "已搜尋，無候選（以 VR 每 GPU 價外推）", "查詢詞：「Rubin Ultra rack price」「Rubin Ultra price per GPU」；結果：僅見推估，無公開報價，未入帳｜價格＝VR 每 GPU 價（SRC_HW_015 BOM 7.8÷72＝0.108 M）× 1.5／2.0／2.5 × 72，倍數本身為 Assumed；"
               "與封裝數線性，單架改定義不新增資訊｜讀取者：CC（1c，2026-10-05）", effect="IF_*（Rubin Ultra 欄）", grade="Assumed", stance="—")

# ----- Decisions（Excel 擁有；只在 ID 不存在時附加）
DEC = [
 ["S1", "Stage 2", "來源等級的升級權限", "(a)：CC 讀到原文、附網址與原文摘錄、原文值與 SRC 現值相同（容差 0；換算後相同須列換算式）時列為升級候選；chat 端抽查後即寫入，不逐筆問 Andy。值不同、涉及 A 級或改變標記者逐筆交 Andy",
  TODAY, "「S1, S2:皆(a)」", "生效（v5.18）", "SRC_* 等級欄；DB_Evidence", "工作單 v5.18 第 0 節", "否"],
 ["S2", "Stage 2", "J6 利用率、J15 生產折減的證據蒐集", "(a)：本批蒐集，只作 Andy 參考：不提議數值、不改暫用值（G0-11 不變）；η 拆解仍依 Andy『有新數據再說』暫緩",
  TODAY, "「S1, S2:皆(a)」", "生效（v5.18）", "Serving 17／18；DB_Evidence", "工作單 v5.18 第 0 節", "否"],
 ["H1", "Stage 2", "不得為唯一依據的來源（SemiAnalysis，含 InferenceX）的等級上限",
  "若某來源依規定不得作為唯一依據，且同時符合 (i) 已讀一手原始資料（網址、讀取日期、數值或內插式）、(ii) 已搜尋同口徑第二來源並記錄查詢詞與結果、(iii) 口徑與方法已寫明，則等級依 G0-7 判定但上限 2 級，W1 標『待第二來源』，不擋完成條件；三點不齊者維持 3 級。只適用於明確列名『不得為唯一依據』的來源（目前為 SemiAnalysis，含 InferenceX）",
  TODAY, "「很好」（回覆 chat 端提出的規則草案）", "生效（v5.18）", "SRC_PERF_009、SRC_PERF_010；Checks W1", "工作單 v5.18 第 1 節", "否"],
 ["G15", "Stage 2", "區間型紀錄的等級（工作單 v5.18 稱『G1 規則』）",
  "區間型紀錄（分析師或經銷商估計）同時符合 (i) 至少一個值讀到原文、(ii) 基準落在已讀證據範圍內，即不為 3 級；未讀端點在備註標 Assumed；等級依讀到的最好來源記 1 或 2。編號 G15：Decisions 既有 G1 為另一條規則（原始數據須連結 SRC）",
  TODAY, "「G1–G4: all ok」", "生效（v5.18）", "SRC_HW_003、005、006、010、015", "工作單 v5.18 第 1 節", "否"],
 ["C1", "Stage 2", "Tokenomics 完成條件（讀法 A）", "完成＝工作單 v5.18 第 6 節三項檢查全數通過：W2（P＝高且所連 SRC 為 3 級）＝0；高段 Assumed 參數逐列三項齊備；H1 規則記 2 級者三點齊備",
  TODAY, "「同意把「完成」定為讀法 A。」；「同意」（Assumed 參數完成標準）", "生效（v5.18）", "Checks G 節 W2；報告第 6 節", "工作單 v5.18 第 0、6 節", "否"],
]

# ----- Gov_Map P 欄（B 法「O＋L1」分段）：現行 P 與審查用 Excel 的候選比對；Inputs!E5 依工作單為「無」
gm = base["Gov_Map"]
gm_row = {}
for r in range(5, gm.max_row + 1):
    if gm.cell(r, 1).value: gm_row[gm.cell(r, 1).value] = r
p_upd = []                      # (sheet, cell, old, new)
mism = []
for r in rv["P欄候選（B法）"].iter_rows(min_row=2, values_only=True):
    if not r[0]: continue
    new = "無" if r[0] == "GM001" else r[7]
    old_rv = r[6]
    rr = gm_row[r[0]]
    old = gm.cell(rr, 16).value
    if old != old_rv: mism.append((r[0], old, old_rv))
    if old != new: p_upd.append((gm.cell(rr, 3).value, gm.cell(rr, 4).value, old, new))
assert not mism, mism
flag = {}                       # IF 全欄補充旗標（備註欄 N 附加）：候選頁『補充：IF 全欄分段』與 B 法分段不同者
for r in rv["P欄候選（B法）"].iter_rows(min_row=2, values_only=True):
    if r[0] and r[9] != (("無" if r[0] == "GM001" else r[7])):
        flag[base["Gov_Map"].cell(gm_row[r[0]], 3).value + "|" + base["Gov_Map"].cell(gm_row[r[0]], 4).value] = r[9]

out = {
    "EVID_MIG4": evid, "SRC_UPD": src_upd, "NEWREC": NEWREC, "DEC": DEC, "P_UPD": p_upd, "IF_FLAG": flag,
    "EV": dict(T=E_T, CHI=E_CHI, CHI2=E_CHI2, ISL=E_ISL, OSL=E_OSL, MFU=E_MFU, EDP=E_EDP, SMC=E_SMC, FLEET=E_FLEET, IX21=E_192, KV=E_KV, BOM=E_BOM,
               EPOCH=E_EPOCH, LAT=E_SEARCH_LAT, J=E_J, ASSUMED=E_IX, RUP=E_RUP, C022=E_C022, C027=E_C027, C053=E_C053, H1=e_h1, RU=e_ru, DGX=e_dgx, SCHNEIDER=e1, SUPERMICRO=e2),
}
hdr = ('# v5.18 種子資料（工作單 docs/workorders/20261005_v5.18.md）：由 tools/gen_seed_v518.py 自審查用 Excel（docs/reports/20261005_stage2-1_查核.xlsx）與 v5.17 舊值產生；\n'
       '# 全部寫入皆 Excel 擁有、只在格仍為舊值時才寫（見 builder/v518.py）。手動修改請改產生器後重跑。\n')
body = "import json\n_J = json.loads(r'''" + json.dumps(out, ensure_ascii=False) + "''')\n" + \
       "EVID_MIG4 = _J['EVID_MIG4']\nSRC_UPD = _J['SRC_UPD']\nNEWREC = _J['NEWREC']\nDEC = _J['DEC']\nP_UPD = [tuple(x) for x in _J['P_UPD']]\nIF_FLAG = _J['IF_FLAG']\nEV = _J['EV']\n"
OUT.write_text(hdr + body, encoding="utf-8")
print("evidence rows", len(evid), "src upd", len(src_upd), "new records", len(NEWREC), "P changes", len(p_upd), "IF flags", len(flag))
