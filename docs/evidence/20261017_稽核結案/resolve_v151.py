"""v15.1（工作單 docs/workorders/20261017_v15.1_結案收尾.md U3）：稽核帳本逐單元重判。

與 resolve_v15.py 的差別：v15.0 依「待判斷編號」（O12、O19、O15…）整批翻轉；本檔不看 O 編號，對每個仍為「未處理」的單元（v14.3 帳本的 161 個）各自找證據：
  證據 1  決策逐字涵蓋：現行活頁簿 Decisions 頁中某筆 SL46-* 的「決議」或「影響範圍」文字，逐字列出該格或該列（解析「頁名 + 第 N 列」「頁名 + 儲存格／範圍」）。
  證據 2  該列自己帶決策標註：該格或該列（含同列 A 欄標籤）在 v15.1 的文字內有 SL46-nn，且 v14.2 沒有。
  證據 3  v14.2 → v15.1 該格或該列的公式／值差異（實質差異＝公式或數值改變，或標籤由文字改為公式；純附加說明文字算「註記」）。
  證據 4  修正點在他處（探針）：該單元的發現本身的修正位置不在該列，改以可重算的探針驗證（例如 SRC_DEM_004_Lo 被公式引用的個數、Checks V16 的文字）。
判定：
  有實質差異 → 已修正；否則被「刻意保留」類決策涵蓋 → 刻意保留；否則有決策涵蓋或自帶標註 → 已記為決策；否則維持「未處理」並列出（不編造狀態）。
其餘 470 個單元沿用 v14.3 帳本判定（更早版本的處理，見該列「處理」欄），本檔只對它們做可重算的核對：S1 格值與帳本相同、決策編號仍在、列標籤仍對得上，結果寫在「證據類型」欄。
輸出欄位：status／how／evidence／residual 之外新增 ev_kind（證據類型）、ev_ptr（證據指標）、ext_src（有無外部來源）。
"""
import functools
import os
import re
import collections

import openpyxl
from openpyxl.utils import column_index_from_string as CI, get_column_letter as GL

import lib

FIX, DEC, KEEP, OPEN = '已修正', '已記為決策', '刻意保留', '未處理'
PREV142 = os.path.join(lib.ROOT, 'model', 'archive', '20261017_Tokenomics_v14.2.xlsx')
KEEP_DECS = {'SL46-07', 'SL46-30'}          # 決議本身就是「記為刻意保留」的兩筆（SL46-03 逐句判斷，見 dec_kind）
ROW_UNIT_SEGS = ('稽核2', '稽核3A', '稽核3B')
GOV_SHEETS = {'Gov_Map', 'Decisions', 'DB_Evidence', 'Checks', 'README', 'Load_Bearing', 'SRC_Index'}


@functools.lru_cache(None)
def prev():
    return openpyxl.load_workbook(PREV142)


def cur():
    return lib.frm()


# ---------------------------------------------------------------- 決策文字解析
def _sheet_re():
    names = sorted(cur().sheetnames, key=len, reverse=True)
    return re.compile(r'(?<![A-Za-z0-9_])(' + '|'.join(re.escape(n) for n in names) + r')(?![A-Za-z0-9_])')


ROWS_RE = re.compile(r'\s*第\s*([0-9][0-9、，,／\-–\s]*?)\s*列(?:的\s*([A-Z](?:／[A-Z])*))?')
CELL_RE = re.compile(r'\s*([A-Z]{1,2})(\d{1,4})(?::([A-Z]{1,2})(\d{1,4}))?(?![A-Za-z0-9_])')
PAREN_RE = re.compile(r'（[^（）]*）')
SEP_RE = re.compile(r'\s*[／、,，]')


def _nums(s):
    out = set()
    for part in re.split(r'[、，,／\s]+', s.strip()):
        if not part:
            continue
        m = re.fullmatch(r'(\d+)[\-–](\d+)', part)
        if m:
            out.update(range(int(m.group(1)), int(m.group(2)) + 1))
        elif part.isdigit():
            out.add(int(part))
    return out


def parse_cover(text):
    """回傳 [dict(sheet, rows(set), cols(set|None), cells[(c1,r1,c2,r2)], sent)]：每個「頁名＋後面接的列／儲存格清單」一筆。"""
    out = []
    sre = _sheet_re()
    for m in sre.finditer(text):
        pos = m.end()
        if text[pos:pos + 1] == '!':
            pos += 1
        ent = dict(sheet=m.group(1), rows=set(), cols=None, cells=[])
        while True:
            mr = ROWS_RE.match(text, pos)
            mc = None if mr else CELL_RE.match(text, pos)
            if mr:
                ent['rows'] |= _nums(mr.group(1))
                if mr.group(2):
                    ent['cols'] = (ent['cols'] or set()) | set(mr.group(2).split('／'))
                pos = mr.end()
            elif mc:
                c1, r1, c2, r2 = mc.group(1), int(mc.group(2)), mc.group(3) or mc.group(1), int(mc.group(4) or mc.group(2))
                ent['cells'].append((CI(c1), r1, CI(c2), r2))
                pos = mc.end()
            else:
                break
            mp = PAREN_RE.match(text, pos)
            if mp:
                pos = mp.end()
            ms = SEP_RE.match(text, pos)
            if not ms:
                break
            pos = ms.end()
        if ent['sheet'] in GOV_SHEETS:      # 治理頁（Gov_Map 編號、Decisions 自引、DB_Evidence 編號等）不是「被決策涵蓋的模型格」
            continue
        if ent['rows'] or ent['cells']:
            a = text.rfind('。', 0, m.start()) + 1
            b = text.find('。', m.end())
            ent['sent'] = text[a:(b + 1 if b >= 0 else len(text))]
            out.append(ent)
    return out


@functools.lru_cache(None)
def dec_index():
    ws = cur()['Decisions']
    idx = {}
    for r in range(5, ws.max_row + 1):
        did = ws.cell(r, 1).value
        if not (isinstance(did, str) and did.startswith('SL46-')):
            continue
        text = '｜'.join(str(ws.cell(r, c).value or '') for c in (4, 8))      # 決議 D、影響範圍 H
        idx[did] = dict(row=r, title=str(ws.cell(r, 3).value or ''), text=text, cover=parse_cover(text))
    return idx


def dec_kind(did, sent):
    if did in KEEP_DECS:
        return KEEP
    if did == 'SL46-03':
        return KEEP if ('刻意保留' in sent or '維持' in sent) else DEC
    return DEC


def covers_cell(sheet, col, row):
    """該格被哪些 SL46-* 逐字列名。回傳 [(決策ID, 精度, 句子)]；精度 cell＝列到該格，row＝列到整列。"""
    res = []
    for did, d in dec_index().items():
        for e in d['cover']:
            if e['sheet'] != sheet:
                continue
            ci = CI(col)
            if any(c1 <= ci <= c2 and r1 <= row <= r2 for c1, r1, c2, r2 in e['cells']):
                res.append((did, 'cell', e['sent'])); break
            if row in e['rows'] and (e['cols'] is None or col in e['cols']):
                res.append((did, 'row' if e['cols'] is None else 'cell', e['sent'])); break
    return res


def covers_row(sheet, row):
    res = []
    for did, d in dec_index().items():
        for e in d['cover']:
            if e['sheet'] != sheet:
                continue
            if row in e['rows'] or any(r1 <= row <= r2 for c1, r1, c2, r2 in e['cells']):
                res.append((did, 'row', e['sent'])); break
    return res


# ---------------------------------------------------------------- 差異與標註
def _kind(old, new):
    so, sn = isinstance(old, str), isinstance(new, str)
    of = so and old.startswith('=')
    nf = sn and new.startswith('=')
    if of or nf:
        return 'formula'
    if so and sn and new.startswith(old) and len(new) > len(old):
        return 'annotation'
    if old is None and sn:
        return 'annotation'
    if isinstance(old, (int, float)) and isinstance(new, (int, float)) or old != new:
        return 'value' if not (so and sn) else 'text-edit'
    return None


def row_diff(sheet, row, cols=None):
    """v14.2 → v15.1 該列的差異 [(座標, 種類, 舊, 新)]（種類：formula／value／annotation／text-edit）。"""
    po, pn = prev()[sheet], cur()[sheet]
    out = []
    for c in range(1, max(po.max_column, pn.max_column) + 1):
        if cols and GL(c) not in cols:
            continue
        a, b = po.cell(row, c).value, pn.cell(row, c).value
        if a == b:
            continue
        k = _kind(a, b)
        if k:
            out.append((f'{GL(c)}{row}', k, a, b))
    return out


def label_aligned(sheet, row):
    return prev()[sheet].cell(row, 1).value == cur()[sheet].cell(row, 1).value or \
        str(cur()[sheet].cell(row, 1).value or '').startswith(str(prev()[sheet].cell(row, 1).value or '\u0000'))


def own_marks(sheet, row, cols=None):
    """該列在 v15.1 帶有、v14.2 沒有的 SL46-nn 標註 {決策ID: 座標}。"""
    po, pn = prev()[sheet], cur()[sheet]
    res = {}
    for c in range(1, pn.max_column + 1):
        if cols and GL(c) not in cols and c != 1:
            continue
        t, o = pn.cell(row, c).value, po.cell(row, c).value
        if isinstance(t, str):
            for m in re.finditer(r'SL46-\d\d', t):
                if not (isinstance(o, str) and m.group(0) in o):
                    res.setdefault(m.group(0), f'{GL(c)}{row}')
    return res


# ---------------------------------------------------------------- 探針（修正點在他處）
def probe_f17():
    pat = re.compile(r'SRC_DEM_004_(Lo|Hi)')
    def cnt(wb):
        return sum(1 for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=') and pat.search(c.value) and c.parent.title != 'Gov_Map')
    a, b = cnt(prev()), cnt(cur())
    return (b > 0 and a == 0), f'SRC_DEM_004_Lo／_Hi 被公式引用的格數：v14.2＝{a}、v15.1＝{b}（Alloc Q 節）；Decisions SL46-01'


def probe_f12():
    ws, wo = cur()['Alloc_In'], prev()['Alloc_In']
    n = sum(1 for r in range(35, 44) for c in range(1, 9) if isinstance(ws.cell(r, c).value, str) and str(ws.cell(r, c).value).startswith('='))
    o = sum(1 for r in range(35, 44) for c in range(1, 9) if isinstance(wo.cell(r, c).value, str) and str(wo.cell(r, c).value).startswith('='))
    ck = cur()['Checks']
    kp = [r for r in range(1, ck.max_row + 1) if ck.cell(r, 1).value in ('KP1', 'KP2')]
    return (n >= 5 and o == 0 and len(kp) == 2), f'Alloc_In 第 35–43 列公式格：v14.2＝{o}、v15.1＝{n}（κ 推導區塊）；Checks KP1／KP2 {len(kp)} 列；Decisions SL46-10'


def probe_g11():
    ck, co = cur()['Checks'], prev()['Checks']
    r = next(r for r in range(1, ck.max_row + 1) if ck.cell(r, 1).value == 'V16')
    fv = lib.val()['Checks'].cell(r, 6).value
    ok = ck.cell(r, 7).value == '恆等式（非驗證）' and co.cell(r, 7).value != '恆等式（非驗證）' and ck.cell(r, 6).value != co.cell(r, 6).value
    return ok, f'Checks V16（第 {r} 列）G 欄：v14.2＝{co.cell(r, 7).value!r}→v15.1＝{ck.cell(r, 7).value!r}；F 欄公式已改，重算值＝{fv!r}；Decisions SL46-08'


PROBES = {'F17': probe_f17, 'F12': probe_f12, 'G11': probe_g11}
PROBE_DEC = {'F17': 'SL46-01', 'F12': 'SL46-10', 'G11': 'SL46-08'}


# ---------------------------------------------------------------- 逐單元判定
def unit_loc(u):
    if u['seg'] == '稽核1':
        m = re.fullmatch(r'([A-Z]+)(\d+)', u['cell'])
        return ('cell', u['sheet'], m.group(1), int(m.group(2)))
    if u['seg'] in ROW_UNIT_SEGS:
        return ('row', u['sheet'], None, int(re.search(r'\d+', u['cell']).group()))
    return ('none', None, None, None)


S0_KEYS = {'S0-C1.1b': '122 格', 'S0-C1.1c': '370 格', 'S0-C1.5': '110 項', 'S0-C6': '單位欄空白'}


def judge(u):
    """對一個「未處理」單元找證據。回傳 dict(status, how, evidence, residual, ev_kind, ev_ptr)；找不到就回 None（維持未處理）。"""
    kind, sh, col, row = unit_loc(u)
    di = dec_index()
    if kind == 'none':                                   # 第 0 段機械項：決策文字須逐字列出該項的數字
        key = S0_KEYS.get(u['uid'])
        for did in sorted(KEEP_DECS):
            if key and key in di[did]['text']:
                return dict(status=KEEP, how=f'Decisions {did}（O15 其餘小項記為刻意保留）', evidence=f'{did} 決議逐字列出「{key}」', residual='',
                            ev_kind='決策逐字涵蓋', ev_ptr=f'{did}（含「{key}」）')
        return None
    cov = covers_cell(sh, col, row) if kind == 'cell' else covers_row(sh, row)
    if kind == 'cell' and any(c[1] == 'cell' for c in cov):      # S1 單格：有列到該格的決策就不用「整列」層級的（避免同列別格的決策誤蓋）
        cov = [c for c in cov if c[1] == 'cell']
    marks = {} if cov else own_marks(sh, row)      # 自帶標註（該列帶有、v14.2 沒有的 SL46-nn）；已有決策逐字列名時不再重複列
    diffs = row_diff(sh, row, cols=({col, 'A'} if kind == 'cell' else None))
    subst = [d for d in diffs if d[1] in ('formula', 'value')]
    ptrs, evid = [], []
    for did, prec, sent in cov:
        ptrs.append(f'{did}（{"列名到該格" if prec == "cell" else "列名到整列"}）')
    for did, at in marks.items():
        if did not in {c[0] for c in cov}:
            ptrs.append(f'{did}（{at} 標註）')
    if subst:
        evid.append('v14.2→v15.1 實質差異：' + '；'.join(f'{d[0]} {d[1]}' for d in subst[:4]) + (' …' if len(subst) > 4 else ''))
    elif diffs:
        evid.append('v14.2→v15.1 僅附加說明：' + '、'.join(d[0] for d in diffs[:4]))
    probe = None
    for fid in u['ids'].split(','):
        fid = fid.strip()
        if fid in PROBES and not (cov or subst):
            ok, txt = PROBES[fid]()
            if ok:
                probe = (fid, txt)
    kinds = [dec_kind(c[0], c[2]) for c in cov]
    if probe:
        return dict(status=FIX, how=f'Decisions {PROBE_DEC[probe[0]]}；修正點在他處', evidence=probe[1], residual='',
                    ev_kind='修正點在他處（探針）', ev_ptr=f'{PROBE_DEC[probe[0]]}；探針 {probe[0]}')
    if subst:
        return dict(status=FIX, how='公式或值已改（' + '、'.join(d[0] for d in subst[:3]) + '）' + (f'；{cov[0][0]}' if cov else ''), evidence='；'.join(evid + ([f'{cov[0][0]} 逐字列名'] if cov else [])),
                    residual='', ev_kind='格差異' + ('＋決策' if cov else ''), ev_ptr='；'.join([d[0] for d in subst[:3]] + [c[0] for c in cov]))
    if cov and KEEP in kinds:
        did = next(c[0] for c, k in zip(cov, kinds) if k == KEEP)
        return dict(status=KEEP, how=f'Decisions {did}（刻意保留）', evidence=f'{did} 逐字列名；' + '；'.join(evid), residual='', ev_kind='決策逐字涵蓋', ev_ptr='；'.join(ptrs))
    if cov or marks:
        did = (cov[0][0] if cov else sorted(marks)[0])
        return dict(status=DEC, how=f'Decisions {did}', evidence=('；'.join(['逐字列名：' + '、'.join(c[0] for c in cov)] * bool(cov) + (['該列自帶標註：' + '、'.join(f'{k}@{v}' for k, v in sorted(marks.items()))] if marks else []) + evid)),
                    residual='', ev_kind='決策逐字涵蓋' if cov else '列自帶決策標註', ev_ptr='；'.join(ptrs))
    return None


# ---------------------------------------------------------------- 外部來源欄
@functools.lru_cache(None)
def gov_rows():
    ws = cur()['Gov_Map']
    rows = []
    for r in range(5, ws.max_row + 1):
        gid, sh, cell = ws.cell(r, 1).value, ws.cell(r, 3).value, ws.cell(r, 4).value
        if isinstance(gid, str) and gid.startswith('GM') and sh and cell:
            rows.append((gid, sh, str(cell).replace('$', ''), ws.cell(r, 6).value, ws.cell(r, 8).value, ws.cell(r, 10).value, ws.cell(r, 16).value))
    return rows


def gov_of(sheet, cell):
    m = re.fullmatch(r'([A-Z]+)(\d+)', cell)
    if not m:
        return None
    ci, rr = CI(m.group(1)), int(m.group(2))
    best = None
    for gid, sh, rg, cat, src, dec, seg in gov_rows():
        if sh != sheet:
            continue
        mm = re.fullmatch(r'([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?', rg)
        if not mm:
            continue
        c1, r1 = CI(mm.group(1)), int(mm.group(2))
        c2, r2 = CI(mm.group(3) or mm.group(1)), int(mm.group(4) or mm.group(2))
        if c1 <= ci <= c2 and r1 <= rr <= r2:
            if (c1, r1, c2, r2) == (ci, rr, ci, rr):
                return dict(gm=gid, cat=cat, src=src, dec=dec, seg=seg)
            best = best or dict(gm=gid, cat=cat, src=src, dec=dec, seg=seg)
    return best


def ext_src(u):
    """帳本新增欄「有無外部來源」：輸入格看 Gov_Map（現行簿）；其餘為繼承輸入。"""
    if u['seg'] != '稽核1':
        return '—（列／機械項，繼承輸入）'
    g = gov_of(u['sheet'], u['cell'])
    if not g:
        return '—（Gov_Map 找不到該格）'
    src = g['src']
    nosrc = src in (None, '', '—')
    if g['cat'] == 'Assumed' and nosrc:
        return '無（Assumed，無 SRC）'
    if nosrc:
        return '無 SRC（類別 %s）' % g['cat']
    return '有（%s；類別 %s）' % (src, g['cat'])


# ---------------------------------------------------------------- 沿用單元的可重算核對
def reverify(u, s1_by_key):
    """對沿用的單元做可重算核對，回傳 (證據類型, 指標)。"""
    kind, sh, col, row = unit_loc(u)
    if kind == 'cell':
        g = gov_of(sh, u['cell'])
        s = s1_by_key.get((u['uid'], sh, u['cell']))
        ok = []
        if s is not None:
            v = cur()[sh][u['cell']].value
            ok.append('格值與帳本相同' if (v == s['cur_val'] or (isinstance(v, str) and v.startswith('='))) else '格值與帳本不同')
        if g and g['dec'] not in (None, '—', ''):
            ok.append(f'Gov_Map {g["gm"]} 決策 {g["dec"]}' + ('（存在於 Decisions）' if _dec_exists(g['dec']) else '（Decisions 找不到）'))
        return '沿用前版判定＋核對', '；'.join(ok) or '沿用前版證據文字'
    if kind == 'row':
        lab = str(cur()[sh].cell(row, 1).value or '')
        want = str(u['label'] or '')
        ok = '列標籤對得上' if want[:12] and (want[:12] in lab) else '列標籤對不上（列號可能已移動）'
        return '沿用前版判定＋核對', ok
    return '沿用前版判定', '第 0 段機械項，沿用前版證據文字'


@functools.lru_cache(None)
def _dec_ids():
    ws = cur()['Decisions']
    return {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}


def _dec_exists(d):
    return d in _dec_ids()


# ---------------------------------------------------------------- 入口
def judge_all(ledger, s1_cells):
    s1_by_key = {('S1-%s' % c['gm'], c['sheet'], c['cell']): c for c in s1_cells}
    n_re, n_open, unresolved = 0, 0, []
    for u in ledger:
        u['ext_src'] = ext_src(u)
        if u['status'] == OPEN:
            n_re += 1
            j = judge(u)
            if j:
                u.update(status=j['status'], how=j['how'], evidence=j['evidence'], residual=j['residual'], ev_kind=j['ev_kind'], ev_ptr=j['ev_ptr'])
                u['rejudged'] = '是'
            else:
                n_open += 1
                u['ev_kind'], u['ev_ptr'], u['rejudged'] = '找不到證據', '—', '是'
                unresolved.append(u)
        else:
            u['ev_kind'], u['ev_ptr'] = reverify(u, s1_by_key)
            u['rejudged'] = '否（沿用）'
    return dict(rejudged=n_re, still_open=n_open, unresolved=unresolved)


def topic_status(FINDINGS, ledger, RANK):
    """主題頁（F／G／B／Y）：原本「未處理」的主題，狀態取其單元中最差者（不再由編號直接指定）。"""
    by = collections.defaultdict(list)
    for u in ledger:
        for i in str(u['ids']).split(','):
            by[i.strip()].append(u)
    n = 0
    for k, F in FINDINGS.items():
        if F['status'] == OPEN and by.get(k):
            w = max(by[k], key=lambda x: RANK[x['status']])
            F.update(status=w['status'], how=w['how'], evidence=w['evidence'], residual=w.get('residual', ''))
            n += 1
    return n
