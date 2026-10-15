"""3A 列清單：第 0 段「C3_核心推導列」扣除「C7_m路徑」（頁＋列），只取 3A 工作表。
用法：python3 00_rows.py <稽核0明細.xlsx> <out.json>"""
import openpyxl,sys,json,collections
SH3A=['Interface','Calib','Perf_Batch','Unit_Cost','Amortize','Workload','Perf','Fleet_1GW','Cache_Store','DC_Cost','Price_Frontier','Tech_Registry','Spec_Rack','Energy','Serving']
wb=openpyxl.load_workbook(sys.argv[1],read_only=True)
core=[(r[0],int(r[1]),r[2],r[3],r[4]) for r in wb['C3_核心推導列'].iter_rows(min_row=3,values_only=True) if r[0] and r[1]]
mp={(r[1],int(r[2])) for r in wb['C7_m路徑'].iter_rows(min_row=3,values_only=True) if r[0]=='路徑列'}
rows=[dict(sheet=s,row=r,label=l,nf=nf,ni=ni) for s,r,l,nf,ni in core if s in SH3A and (s,r) not in mp]
c=collections.Counter(x['sheet'] for x in rows)
print(len(rows),dict(c))
json.dump(rows,open(sys.argv[2],'w'),ensure_ascii=False,indent=0)
