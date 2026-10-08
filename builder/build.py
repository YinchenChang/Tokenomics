# Usage: python3 build.py <base.xlsx> <out.xlsx>
#   chat:  python3 build.py /mnt/project/<latest>.xlsx /home/claude/b2/<new>.xlsx   (or base downloaded from the repo)
#   repo:  python3 builder/build.py model/<current>.xlsx <new>.xlsx
import sys, os, openpyxl
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
if len(sys.argv) != 3:
    sys.exit("usage: build.py <base.xlsx> <out.xlsx>")
BASE, out = sys.argv[1], sys.argv[2]
OUTDIR = os.path.dirname(os.path.abspath(out))
from inputs import spec_rack, arch, serving, energy_inputs
from calib import calib
from outputs import perf_sheet, sens_sheet, unit_cost, workload, nonnv
from finish import interface, checks, sources, readme, interface_b3, checks_b3, write_checks_b3
from training import tech_registry, train_in, perf_batch, training_sheet, sens_train
from finish import evidence_sheet
from preserve import snapshot, restore
from v519 import PROD_SHEETS

wb = openpyxl.load_workbook(BASE)
# ---- v5.12 (A): the pre-Source register is frozen as Sources_Legacy (no formula or name refers to it) ----
if "Sources" in wb.sheetnames and "Sources_Legacy" not in wb.sheetnames:
    wb["Sources"].title = "Sources_Legacy"
    wb["Sources_Legacy"]["A1"].value = ("Sources_Legacy — v5.11 以前的來源清單（v5.12 起凍結，只供追溯原 S 編號；"
                                        "新來源一律先登錄 DB_Evidence，擇優寫入 SRC_*）")
# ---- v5.7 Excel-first: snapshot every input cell (blue font) before the rebuild ----
SNAP = snapshot(wb)
# ---- v5.12 (A): inputs that moved to another row keep their Excel value (restore stays unmatched 0) ----
from training import SNAP_MOVES
for _old, _new in SNAP_MOVES:
    if _old in SNAP and _new not in SNAP: SNAP[_new] = SNAP.pop(_old)
# ---- v5.13 (D): input cells retired because the cell is now a formula (Arch KV rows 21–23, Cap_In N column). Their Excel
#      values are logged, not restored; the new formulas reproduce them with the default inputs (verified by recalculation).
from block4 import PRICE_ROWS
SNAP_RETIRED = [("Arch", (lab, 0), c) for lab in ("KV bytes/token — 情境 1", "KV bytes/token — 情境 2", "KV bytes/token — 情境 3")
                for c in (3, 4, 5)] + [("Cap_In", (row[0], 0), 14) for row in PRICE_ROWS]
RETIRED = {k: SNAP.pop(k) for k in SNAP_RETIRED if k in SNAP}
# ---- strip Block 2/3 content to recover the Block 1 base, then rebuild deterministically ----
for n in ["Arch","Serving","Workload","Calib","Perf","Sens_Perf","Unit_Cost","Energy","NonNV",
          "Tech_Registry","Train_In","Perf_Batch","Training","Sens_Train",
          "Cap_In","Capability","Price_Frontier","Cache_Store","Fleet_1GW","Amortize","Theory_Rev","Sens_Rev",
          "Har_In","Harness","Sens_Har","Alloc_In","Alloc",
          *PROD_SHEETS]:
    if n in wb.sheetnames: del wb[n]
def clear(ws, r0, c1=1, c2=30):
    for r in range(r0, ws.max_row + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.value = None; cell.fill = openpyxl.styles.PatternFill(fill_type=None)
clear(wb["Spec_Rack"], 21); 
for r in range(1, 21): wb["Spec_Rack"].cell(row=r, column=8).value = None
clear(wb["Interface"], 17); clear(wb["Checks"], 12); clear(wb["Sources_Legacy"], 23); clear(wb["README"], 4, 1, 2)
for n in list(wb.defined_names.keys()):
    if n not in ("CTL_GW","CTL_PowerCase","IF_CapexFacility","IF_CapexIT","IF_CapexTotal","IF_FacilityGW",
                 "IF_GPUhrEcon","IF_GPUsPerGW","IF_HoldAcct","IF_HoldEcon","IF_PowerCost","IF_RacksPerGW"):
        del wb.defined_names[n]
SP = spec_rack(wb)
AR = arch(wb)
serving(wb)
CAL = calib(wb, SP, AR)
import v523
v523.calib_i(wb)       # v5.23 X10: Calib I section, CAL_BatchEtaD (after the existing content, nothing moves)
energy_inputs(wb)
TR = tech_registry(wb)
TI = train_in(wb)
PR = perf_sheet(wb, SP, AR, CAL, TR)
SR = sens_sheet(wb, SP, AR, CAL, TR)
PB = perf_batch(wb, SP, AR, CAL, TR, TI)
v523.perf_batch_etad(wb)      # v5.23 X10: Perf_Batch row 53 C:Q x INDEX(CAL_BatchEtaD,1,gen) (before v519.x1 mirrors the chain)
TRN = training_sheet(wb, SP, AR, TI, PB, TR)
STR = sens_train(wb, SP, AR, TI, PB, TR)
U = unit_cost(wb, PR)
WL = workload(wb, U)
nonnv(wb)
interface(wb, PR, U)
def last_row(ws):
    return max(c.row for row in ws.iter_rows() for c in row if c.value is not None)
ifr = last_row(wb["Interface"]) + 2
interface_b3(wb, TRN, ifr, TI)
checks(wb, CAL, PR, SR, U)
_r, _rows = checks_b3(wb, TRN, TI, CAL, PB)
write_checks_b3(wb, _r, _rows, SP)
# ---- v5.8: Block 4 ----
from block4 import cap_in, price_frontier, capability, cache_store, fleet, amortize, theory_rev, sens_rev, interface_b4, checks_b4, sources_b4, evidence_b4
K4 = cap_in(wb)
import v522
v522.cap_in_g(wb)      # v5.22 X7: CTL_PriceLife／CTL_Monetize (Cap_In G section, after the existing content)
P4 = price_frontier(wb, K4, TR)
C4 = capability(wb, K4, P4, TRN, TI)
S4 = cache_store(wb)
F4 = fleet(wb, K4)
A4 = amortize(wb, F4, P4, S4)
T4 = theory_rev(wb, F4, A4, S4, WL)
v522.theory_rev_life(wb)      # v5.22 X7: Theory_Rev D section (base row x L x m) and self-check rows
K4["_capf"] = P4["capf"]
sens_rev(wb, K4, T4, A4)
interface_b4(wb, last_row(wb["Interface"]) + 2, P4, S4, A4, T4)
checks_b4(wb, P4, F4, A4, T4, U)
# ---- v5.9: Block 5 (harness); v5.10: M1 (b) reliability floor + Interface E additions (inside block5.py) ----
from block5 import har_in, harness, sens_har, interface_b5, checks_b5, sources_b5, evidence_b5
H5 = har_in(wb, TR)
R5 = harness(wb, U, S4, WL)
SH5 = sens_har(wb, U, S4, WL, H5, R5)
interface_b5(wb, last_row(wb["Interface"]) + 2, R5)
checks_b5(wb, R5, H5, SH5)
# ---- v5.15: Block 6 (Alloc_In inputs, Alloc derivation, Interface F) ----
from block6 import alloc_in, alloc, interface_b6, checks_h
AIN = alloc_in(wb)
import v529
AIN_529 = v529.alloc_in_rows(wb)       # v5.29 X14 (m): two Assumed inputs and the v5.28 constant (appended after the existing content; restored by preserve)
AL6 = alloc(wb, AIN)
interface_b6(wb, last_row(wb["Interface"]) + 2)
sources(wb)
sources_b4(wb)
sources_b5(wb)
readme(wb)
# ---- v5.7: write back Excel-owned inputs, then the evidence register (created only if absent) ----
_m, _c, _d = restore(wb, SNAP, os.path.join(OUTDIR, "restore_log.txt"))
print(f"restore: matched {_m}, Excel kept over code {_c}, unmatched {len(_d)}")
import v518
V518_LOG = v518.inputs_update(wb)       # v5.18: Excel-owned input writes, each only while the cell still holds its v5.17 value
print("v518 inputs:", len(V518_LOG))
import v520
V518_LOG += v520.inputs_update(wb)       # v5.20 X3: Alloc_In!E7 4 -> 6 (old-value guard)
import v521
V518_LOG += v521.text_update(wb)         # v5.21: Alloc_In!G7 note tail (old-text guard)
V518_LOG += v529.inputs_update(wb)       # v5.29 X14 (l): Alloc_In 每則提示 token 數 2000/1000/6000 -> 4000/2000/6400 (old-value guards) and the G note
import v531
V518_LOG += v531.inputs_update(wb)       # v5.31 X16–X19: Spec_Rack E12 4.3 and text, Inputs rows 30／36–38, DC_Cost A46 label (old-value guards)
evidence_sheet(wb)
print("evidence rows added:", evidence_b4(wb), evidence_b5(wb))
order = ["README","Inputs","Spec_Rack","Arch","Serving","Workload","Calib","Tech_Registry","Perf","Sens_Perf","Unit_Cost","DC_Cost",
         "Train_In","Perf_Batch","Training","Sens_Train",
         "Cap_In","Capability","Price_Frontier","Cache_Store","Fleet_1GW","Amortize","Theory_Rev","Sens_Rev",
         "Har_In","Harness","Sens_Har","Interface","Energy","NonNV","Sensitivity","Checks","Sources_Legacy","DB_Evidence"]
wb._sheets = [wb[n] for n in order] + [s for s in wb._sheets if s.title not in order]   # v5.11: keep Excel-owned governance sheets
# ---- v5.3: Block 1 Checks — reference cells instead of literals (CC 第 1 輪第 6 節第 5 項) ----
from common import put, F_IN
ck = wb["Checks"]
# v5.11: C4–C8 now link to SRC (gov.checks_block1); C9–C10 (SRC_Price) remain literals until slice two (v5.12)
for r, v in [(9, 10.5), (10, 3.57)]:
    put(ck, f"C{r}", v, F_IN, fmt="#,##0.00")
ck["E7"].value = '=IF(B7>C7,"高於參照（參照假設未知）","低於參照")'
ck["E8"].value = '=IF(B8>C8,"高於參照（參照假設未知）","低於參照")'
ck["E9"].value = '="牌價 ÷ 持有成本＝"&TEXT(C9/B9,"0.0")&" 倍"'
ck["E10"].value = '="參照 ÷ 本模型＝"&TEXT(C10/B10,"0.00")'
# ---- v5.3: display-only named ranges (DRV_／CAL_／IF_Hdr；CC 第 1 輪第 6 節第 2、3 項) ----
from openpyxl.workbook.defined_name import DefinedName
def nm(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)
nm("IF_HdrGen", "Interface!$C$4:$Q$4"); nm("IF_HdrCost", "Interface!$C$5:$Q$5")
nm("DRV_Gen", "Perf!$C$4:$Q$4"); nm("DRV_Tier", "Perf!$C$5:$Q$5")
for n, k in [("DRV_FlopDec","Fd"),("DRV_FlopPre","Fp"),("DRV_WeightGB","W"),("DRV_TfixMs","tfix"),("DRV_SeqMs","ceff"),
             ("DRV_SeqBind","cbind"),("DRV_Batch","B"),("DRV_Bind","bind"),("DRV_DecTokGPU","D"),("DRV_PreTokGPU","Pp"),
             ("DRV_PreShare","psh"),("DRV_RackTok","rtot"),("DRV_GWTok","gwtot"),("DRV_Close","close")]:
    nm(n, f"Perf!$C${PR[k]}:$Q${PR[k]}")
nm("CAL_EtaD", f"Calib!$C${CAL['etad_fit']}"); nm("CAL_TlayerUs", f"Calib!$C${CAL['tl_fit']}")
nm("CAL_F_Label", f"Calib!$C${CAL['lab']}:$I${CAL['lab']}")
for n, k in [("CAL_F_Gen","gn"),("CAL_F_Eng","eng"),("CAL_F_Speed","s"),("CAL_F_Meas","T"),("CAL_F_Model","vT"),("CAL_F_Basis","basis")]:
    nm(n, f"Calib!$C${CAL[k]}:$I${CAL[k]}")
for n, k in [("CAL_H_Gen","mgn"),("CAL_H_Speed","ms"),("CAL_H_Meas","mT"),("CAL_H_Model","mO"),("CAL_H_Basis","mbasis")]:
    nm(n, f"Calib!$C${CAL[k]}:$F${CAL[k]}"); nm("CAL_F_Platform", f"Calib!$C${CAL['indep']}:$I${CAL['indep']}")
nm("CAL_F_Ratio", f"Calib!$C${CAL['ratio']}:$I${CAL['ratio']}")
nm("CAL_H_Label", f"Calib!$C${CAL['mhdr']}:$F${CAL['mhdr']}"); nm("CAL_H_Platform", f"Calib!$C${CAL['mplat']}:$F${CAL['mplat']}")
nm("CAL_H_Ratio", f"Calib!$C${CAL['mR']}:$F${CAL['mR']}")
# ---- v5.5: Block 3 display-only named ranges (TRN_；網站推導鏈用，下游不得連結) ----
for n, k in [("TRN_FlopTokPre","Fpt"),("TRN_FlopPre","Cp"),("TRN_GPUhPre","Hp"),("TRN_GPUhRL","Hrl"),("TRN_RLMFU","rlmfu"),
             ("TRN_RLRatioH","rlH"),("TRN_RLRatioF","rlF"),("TRN_GPUhFinal","Hfin"),("TRN_PostShareF","psF"),("TRN_PostShareH","psH"),
             ("TRN_InferShare","Ish"),("TRN_GPUhProg","Hprog"),("TRN_GWyrProg","GWp")]:
    nm(n, f"Training!$C${TRN[k]}:$Q${TRN[k]}")
nm("TRN_Gen", "Training!$C$4:$Q$4"); nm("TRN_Tier", "Training!$C$5:$Q$5")
# ---- v5.6: Tech_Registry display-only named ranges (TR_；網站唯讀表用，下游不得連結) ----
r0, r1 = TR["_rows"]; h0, h1 = TR["_hooks"]
for n, c in [("TR_ID","A"),("TR_Tech","B"),("TR_Hook","C"),("TR_Acts","D"),("TR_Lo","E"),("TR_Base","F"),("TR_Hi","G"),
             ("TR_Sel","H"),("TR_Status","I"),("TR_Labs","J"),("TR_Main","K"),("TR_Override","L"),("TR_InBase","M"),
             ("TR_Adopt","N"),("TR_Switch","O"),("TR_Eff","P"),("TR_Tag","Q"),("TR_Source","R"),("TR_Trigger","S"),("TR_Check","T")]:
    nm(n, f"Tech_Registry!${c}${r0}:${c}${r1}")
nm("TR_HookCode", f"Tech_Registry!$A${h0}:$A${h1}"); nm("TR_HookName", f"Tech_Registry!$B${h0}:$B${h1}")
nm("TR_HookVal", f"Tech_Registry!$E${h0}:$E${h1}")
# ---- v5.11: Stage 1 slice one (Source layer, Evidence upgrade, Decisions, Gov_Map, L1, governance Checks, F14); v5.13: slice two ----
from gov import gov_all
GOV = gov_all(wb)
checks_h(wb)        # v5.15: Checks H section (after the G section; counted in GOV_Errors)
import v519
X1 = v519.x1(wb)    # v5.19 X1: *_Prod mirror sheets, Interface G, Checks I1 (after gov_all: reads the final formulas of the chain)
print("x1 chain cells:", X1['cells'], "sheets:", X1['sheets'])
X7 = v522.interface_h(wb)      # v5.22 X7: Interface H section (after G; IF_*_Life names)
X7["check_row"] = v522.checks_y(wb)
print("x7 interface rows from:", X7["start"], "check row:", X7["check_row"])
import v526
X26 = v526.interface_i(wb)     # v5.26: Interface I section (after H; DC_Cost components IF_DeprLifeIT … IF_OpexGW, sum check row)
print("v526 interface I from:", X26["start"], "dc rows:", X26["dc_rows"], "check row:", X26["check_row"])
# ---- v5.29: Gov_Map helper columns, Interface J (waterfall), Interface R–U (contract; static precedent tracing), Load_Bearing, Checks K ----
from deps import Deps, gov_map_cells
N_GMH = v529.gov_map_helpers(wb)
import v530
IFJ = v530.interface_j(wb)            # v5.30 X15 (a): running-product waterfall; self-check = source equality + monotonic (replaces v529.interface_j)
I2 = v531.interface_i2(wb)            # v5.31 X17／X18: Interface I continuation after J (IF_MaintITWarr, IF_MaintITPost, IF_WarrantyYrs); IF_StaffSW label (X19)
print("v531 interface I2 from:", I2["start"], "names:", [n for n, _ in I2["made"]])
DEPS = Deps(wb); GMC = gov_map_cells(wb)                 # built after every formula sheet exists (J rows included; R–U formulas are not precedents)
IFC = v529.interface_contract(wb, DEPS, GMC, v530.should_cap)      # v5.30 X15 (c): "上限" flag extended to IFW_ revenue rows and IFW_TokGW_*_100
LB = v529.load_bearing(wb, DEPS, GMC)
CK_K = v530.checks_k(wb, IFC, IFJ, LB)    # v5.29 Checks K with the v5.30 K2／K3 wording
print("v529 interface J from:", IFJ["start"], "names:", len(IFJ["made"]), "flags:", IFJ["flags"], "| contract rows:", len(IFC["rows"]), "no-dep rows:", IFC["no_dep"],
      "| Load_Bearing rows:", LB["rows"], "| Gov_Map helper rows:", N_GMH, "| Checks K rows:", CK_K)
GOV["snap_retired"] = len(RETIRED)
GOV["v529"] = dict(contract_rows=len(IFC["rows"]), contract_no_dep=IFC["no_dep"], waterfall_names=len(IFJ["made"]), load_bearing_rows=LB["rows"], gov_map_helper_rows=N_GMH)
GOV["fm_log"] = GOV["fm_log"] + [f"v518 input {x}" for x in V518_LOG] + [f"retired input (now formula) {k[0]} [{k[1][0]}] col {k[2]}: Excel value {v!r}" for k, v in RETIRED.items()]
open(os.path.join(OUTDIR, "gov_log.txt"), "w").write("\n".join([f"{k}: {v}" for k, v in GOV.items() if k != "fm_log"] + ["-- formula map changes --"] + GOV["fm_log"]))
print("gov:", {k: v for k, v in GOV.items() if k != "fm_log"})
order = [n for n in ["README","Inputs","Spec_Rack","Arch","Serving","Workload","Calib","Tech_Registry","Perf","Sens_Perf","Unit_Cost","DC_Cost",
         "Train_In","Perf_Batch","Training","Sens_Train",
         "Cap_In","Capability","Price_Frontier","Cache_Store","Fleet_1GW","Amortize","Theory_Rev","Sens_Rev",
         "Har_In","Harness","Sens_Har","Alloc_In","Alloc",*PROD_SHEETS,"Interface","L1","Load_Bearing","Energy","NonNV","Sensitivity","Checks","Gov_Map","Decisions",
         "SRC_HW","SRC_DC","SRC_Model","SRC_Perf","SRC_Price","SRC_Cap","SRC_Harness","SRC_Demand","SRC_Index","Sources_Legacy","DB_Evidence"]]
assert sorted(order) == sorted(ws.title for ws in wb.worksheets), set(ws.title for ws in wb.worksheets) ^ set(order)
wb._sheets = [wb[n] for n in order]
wb.save(out)
import json; json.dump({"WL":WL,"H5":H5,"R5":R5,"SH5":SH5,"K4":K4,"P4":P4,"C4":C4,"S4":S4,"F4":F4,"A4":{k:v for k,v in A4.items() if not k.startswith("cost")},"T4":T4,"PR":PR,"CAL":CAL,"SR":SR,"PB":PB,"TRN":TRN,"STR":STR,"TI":TI,"TR":{k:v for k,v in TR.items() if not k.startswith("_")},"U":{f"{k[0]}_{k[1]}":v for k,v in U.items()},"SP":SP,"AR":AR}, open(os.path.join(OUTDIR, "rows.json"),"w"))
print("saved")
