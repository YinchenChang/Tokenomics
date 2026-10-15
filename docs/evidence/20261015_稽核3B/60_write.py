"""寫明細 docs/reports/20261015_稽核3B_明細.xlsx（逐列檢查、傳遞強度、高嚴重度）。"""
import json,openpyxl
from openpyxl.styles import Font,Alignment,PatternFill
J=json.load(open('docs/evidence/20261015_稽核3B/50_逐列判讀.json'))
S3=json.load(open('docs/evidence/20261015_稽核3B/30_s3.json'))
X=json.load(open('docs/evidence/20261015_稽核3B/31_extra.json'))
wb=openpyxl.Workbook(); wb.remove(wb.active)
def sheet(name,rule,head,rows,widths):
    ws=wb.create_sheet(name); ws.append([rule]); ws['A1'].font=Font(bold=True); ws['A1'].alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[1].height=100
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=min(len(head),12))
    ws.append(head)
    for cc in ws[2]: cc.font=Font(bold=True); cc.fill=PatternFill('solid',fgColor='FFDDEBF7'); cc.alignment=Alignment(wrap_text=True)
    for r in rows: ws.append(list(r))
    ws.auto_filter.ref=f'A2:{openpyxl.utils.get_column_letter(len(head))}{len(rows)+2}'
    ws.freeze_panes='C3'
    for i,w in enumerate(widths,1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=3):
        for cc in row: cc.alignment=Alignment(wrap_text=True,vertical='top')
LS=[f'L{i}' for i in range(1,10)]
rows=[[x['sheet'],x['row'],x['label'],x['nform'],x['npat'],x['nusers'],x['users']]+[x[k] for k in LS]+[x['upY'],x['ndown'],x['down'],x['bids'],x['sev']] for x in J]
sheet('逐列檢查','規則：第 0 段「C3_核心推導列」扣除「C7_m路徑」後屬 3B 工作表的 441 個頁＋列，逐列檢查 L1–L9（工作單第 1 節）。「—」＝沒有發現問題（不等於已證明正確）。同列跨欄公式先轉相對位置（R1C1）比對，數字常數不同者合併；L3 以 11_align.py 機械檢查（Training／Theory_Rev 同欄對位 0 處不符；INDEX 3*(世代−1)+層級 全部對位；Harness 寫死欄 M＝VR200、J＝GB300 全部與列標籤一致；L1 INDEX 常數逐一對照列標籤與 D／E／F＝基準／低／高）。L8＝依賴圖上游含第 1 段 Y1–Y9 輸入格，並列出傳到哪些下游契約名稱（IF_／IFW_／IFC_／L1_）。B 編號對應報告第 3 節；嚴重度：高＝B01–B09；中＝其他 B；資訊＝只有 L8。只標問題，不判定對錯、不提修法。',
 ['頁','列','列標籤','公式格數','同列結構數（數字常數合併）','被引用的頁＋列數','被誰引用（前 6）']+['L1 量綱','L2 口徑','L3 引用對位','L4 重複／漏算','L5 恆等式與閉合','L6 倒推','L7 條件分支','L8 與 Y1–Y9 的連結','L9 物理合理性']+['上游 Y 類','傳到下游契約名稱數','下游契約名稱（前 10）','B 編號','嚴重度'],rows,
 [11,6,36,8,8,8,30]+[30,50,20,40,40,45,45,45,40]+[14,8,40,10,12])
NAMES=list(S3['base'].keys())
HIGHMAP={'H01':'B03','H01b':'B03','H02':'B04','H03':'B14（中；研發倍數）','H04':'B09','H05':'Y5（第 1 段）','H05b':'Y5','H05c':'—','H05d':'—','H05e':'—','H06':'Y5／B01','H07':'—','H07b':'—','H08':'B01（第 1 段 S 類）','H08b':'B01','H08c':'B01','H09':'B08','H09b':'B08','H09c':'—','H09d':'Y（第 1 段 Har_In 48）','H09e':'B08','H10':'B01／Y2'}
rows=[]
for r in S3['res']:
    p=r['pct']
    def f(n):
        lo=p.get('lo',{}).get(n); hi=p.get('hi',{}).get(n)
        if lo is None and hi is None and n not in p.get('lo',{}) and n not in p.get('hi',{}): return '—'
        g=lambda x:'文字' if x is None else f'{x:+.1f}%'
        return f"{g(lo) if 'lo' in p else '—'}／{g(hi) if 'hi' in p else '—'}"
    mx=max([abs(v) for t in p.values() for v in t.values() if v is not None] or [0])
    rows.append([r['id'],HIGHMAP.get(r['id'],''),r['label'],r['endpoints'],round(mx,1)]+[f(n) for n in NAMES])
rows.sort(key=lambda x:-x[4])
ex=X
rows.append(['X1','B01','服務 GW 改用支出路線（Alloc C61）','替代口徑（算術對照）','',f"IF_AllocQ1 {ex['Q1']['token_route']:.3f} → {ex['Q1']['spend_route']:.3f}；隱含 N×k {ex['ImpliedNk']['token_route']:.2f} → {ex['ImpliedNk']['spend_route']:.2f}；研發端 × ScaleRD 時 Q1＝{ex['Q1_scaled']['Q1']:.3f}（支出比 {ex['Q1_scaled']['spend_ratio']:.3f}）"]+['']*(len(NAMES)-1))
rows.append(['X2','B06','回本所需實收單價的世代口徑','替代口徑（算術對照）','',"；".join(f"{k} {v['req_price']:.2f}" for k,v in ex['PriceReqBE'].items() if isinstance(v,dict))+f"；機隊實收 {ex['PriceReqBE']['機隊實收（L1 M93）']:.2f} $/M"]+['']*(len(NAMES)-1))
rows.append(['X3','B06','L_total 以 WACC 10% 現值加權','替代口徑（算術對照）','',f"L_total {ex['L_total']['workbook']:.4f} → {ex['L_total']['pv_wacc10']:.4f}"]+['']*(len(NAMES)-1))
rows.append(['X4','B05','批次推論速度下限 Train_In E31（Astra）','20／40／60／80／100／150','',"；".join(f"{k}：Hopper–RU Astra decode GPU 秒 {['0' if (isinstance(y,(int,float)) and y==0) else (f'{y:.0f}' if isinstance(y,(int,float)) else y) for y in v['PB103_Astra']]}，IF_AllocQ1 {v['Q1']:.3f}，L1_RLshare_Astra {v['L1_RLshare_Astra']:.2f}" for k,v in ex['batchSLO_zero'].items())]+['']*(len(NAMES)-1))
rows.append(['X5','B11','p_min（Har_In C8）','0.3／0.4／0.5／0.6；另 0.7、0.75','',"；".join(f"{k}：{v['name']} {v['front'] if isinstance(v['front'],str) else round(v['front'],4)}" for k,v in ex['pmin'].items())+'；0.7：Astra｜選定 0.0877；0.75：「無合格」（文字）']+['']*(len(NAMES)-1))
sheet('傳遞強度','規則（工作單第 2 節）：對嚴重度「高」各項的輸入，單一擾動到工作簿內已有的端點（Gov_Map 低高、Sens_Train／Sens_Har 情境倍數、列內寫明的區間），其他不變；以第 2 段 20_eval.py（工作簿公式原文逐格求值，未修改；3B 3,376 個公式格對檔內快取值數值 0 不符）重算下游契約名稱（IF_／L1_）。數值為相對基準 %（低端／高端）；「文字」＝該情境下變成文字。X1–X5 為替代口徑的算術對照（只用工作簿公式值；不是工作簿機制）。排序依最大絕對變化。',
 ['編號','對應問題','輸入','端點','最大 |變化|（%）']+[f'{n}' for n in NAMES],rows,[7,12,32,30,10]+[16]*len(NAMES))
hi=[[x['sheet'],x['row'],x['label'],x['bids'],x['L1'],x['L2'],x['L4'],x['L5'],x['L6'],x['L7'],x['L8'],x['L9'],x['down']] for x in J if x['sev']=='高']
sheet('高嚴重度','規則：「逐列檢查」頁嚴重度＝高的列（屬 B01–B09：作用在下游契約名稱或其外部對照判讀的邏輯問題）。只標問題，不判定對錯、不提修法。',
 ['頁','列','列標籤','B 編號','L1','L2','L4','L5','L6','L7','L8','L9','下游契約名稱'],hi,[11,6,36,10,30,50,40,40,45,45,45,40,40])
wb.save('docs/reports/20261015_稽核3B_明細.xlsx')
print('rows',len(J),'hi',len(hi),'s3',len(rows))
