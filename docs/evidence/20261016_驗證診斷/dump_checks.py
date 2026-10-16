"""列出 Checks 頁 V 節各列的公式與值（唯讀）。用法：python3 dump_checks.py <xlsx> [關鍵字]"""
import sys
import openpyxl

path = sys.argv[1]
kw = sys.argv[2] if len(sys.argv) > 2 else "V"
wbf = openpyxl.load_workbook(path, data_only=False)
wbv = openpyxl.load_workbook(path, data_only=True)
ws_f = wbf["Checks"]
ws_v = wbv["Checks"]
for r in range(1, ws_f.max_row + 1):
    a = ws_f.cell(r, 1).value
    if a is None:
        continue
    if str(a).startswith(kw) or (len(sys.argv) > 3 and sys.argv[3] in str(ws_f.cell(r, 2).value)):
        row = []
        for c in range(1, min(ws_f.max_column, 10) + 1):
            f = ws_f.cell(r, c).value
            v = ws_v.cell(r, c).value
            if f is None:
                continue
            col = openpyxl.utils.get_column_letter(c)
            if isinstance(f, str) and f.startswith("="):
                row.append(f"{col}{r}: {f}  => {v}")
            else:
                row.append(f"{col}{r}: {f}")
        print(" | ".join(row))
