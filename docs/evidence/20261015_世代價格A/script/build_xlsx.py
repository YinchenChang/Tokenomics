# -*- coding: utf-8 -*-
import sys, datetime as dt
sys.path.insert(0,'/tmp/claude-0/-home-claude/cbf8dd04-753c-5475-ac86-d6f2cba4305c/scratchpad')
import points as P
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
OUT='/home/claude/tokenomics/docs/reports/20261015_世代價格A_查證.xlsx'
wb=Workbook()
HF=PatternFill('solid',fgColor='1F3864'); HFont=Font(bold=True,color='FFFFFF')
BLUE=Font(color='0000FF'); BOLD=Font(bold=True)
WR=Alignment(wrap_text=True,vertical='top')
def hdr(ws,row,cols,widths=None):
    for i,c in enumerate(cols,1):
        x=ws.cell(row,i,c); x.fill=HF; x.font=HFont; x.alignment=WR
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[L(i)].width=w
def title(ws,t,sub=None):
    ws['A1']=t; ws['A1'].font=Font(bold=True,size=13)
    if sub: ws['A2']=sub; ws['A2'].alignment=Alignment(wrap_text=False)
# ---------- 說明 ----------
ws=wb.active; ws.title='說明'
title(ws,'世代間每 MW 簽約收入的產業基準｜A 段查證明細（不改 Excel 模型）')
lines=['讀取版本：model/20261014_Tokenomics_v5.35.xlsx（master 9dff7f5）。本檔只新增，不改模型。擷取日 2026-10-15。',
'頁籤：價格點_a_同等能力、價格點_b_層級牌價、價格點_c_算力租賃、E338-E341確認、上市間隔、對照計算。',
'標記：Verified／Interested-party／Analogy／Assumed／Derived。模型商與雲端商自家牌價一律標 Interested-party（銷售方），價格事實以官方頁確認。',
'藍字＝人工輸入的價格或日期（附來源）；黑字＝公式。「對照計算」頁每個結果都是公式，可點格看推導。',
'Wayback Machine 被工具擋下（SITE_BLOCKED）：歷史牌價改用各模型當時的官方公告頁；上市後 0–6 個月的 GPU 租賃官方價格頁快照取不到。',
'SemiAnalysis 未單獨引用；同源轉載（例如 spheron 轉述 CoreWeave 牌價）不算獨立來源。']
for i,t in enumerate(lines,3): ws.cell(i,1,t)
ws.column_dimensions['A'].width=150
# ---------- a ----------
ws=wb.create_sheet('價格點_a_同等能力'); title(ws,'(a) 同等能力 token 價格（固定能力門檻下的最低價格與下降速率）')
cols=['序列','模型／門檻','時點','數值','單位','說明','標記','來源','URL','文件日期','擷取日','利害關係說明']
hdr(ws,3,cols,[34,36,16,10,22,60,12,36,50,16,12,40])
r=4
for x in P.A_:
    for j,v in enumerate([x['series'],x['model'],x['date'],x['val'],x['unit'],x['note'],x['marker'],x['src'],x['url'],x['docdate'],P.RET,x['interest']],1):
        c=ws.cell(r,j,v); c.alignment=WR
        if j==4: c.font=BLUE
    r+=1
r+=1; ws.cell(r,1,'(a′) 同一模型壽命內降價（供 Q2 區分「壽命內衰減」與「世代間價格變化」）').font=BOLD; r+=1
hdr(ws,r,['模型','生效日','舊價','新價','幅度','說明','標記','來源','URL','文件日期','擷取日','利害關係說明']); r+=1
for x in P.C_:
    for j,v in enumerate([x['model'],x['date'],x['old'],x['new'],x['unit'],x['note'],x['marker'],x['src'],x['url'],x['docdate'],P.RET,x['interest']],1):
        ws.cell(r,j,v).alignment=WR
    r+=1
ws.cell(r+1,1,'已搜尋，未找到：2025-02 以後、可回溯到原始牌價的逐點「固定能力最低價」序列（Epoch 2025 表止於 2025-02-05；Epoch 2026-09 報告為成本口徑彙總，圖內數值未提供）。已搜尋管道：epoch.ai 兩篇、arXiv 2511.23455。未查：Artificial Analysis。')
# ---------- b ----------
ws=wb.create_sheet('價格點_b_層級牌價'); title(ws,'(b) 層級牌價（上市時每 M token；輸入、輸出分開）')
cols=['業者','層級','模型','上市／生效日','輸入 $/M','輸出 $/M','3:1 混合 $/M','8:1 混合 $/M','標記','來源','URL','擷取日','利害關係說明','備註']
hdr(ws,3,cols,[11,16,36,14,10,10,12,12,13,46,50,12,36,44])
r=4
for x in P.B:
    vals=[x['prov'],x['tier'],x['model'],x['date'],x['i'],x['o'],None,None,x['marker'],x['src'],x['url'],P.RET,x['interest'],x['note']]
    for j,v in enumerate(vals,1):
        c=ws.cell(r,j,v); c.alignment=WR
        if j in (5,6): c.font=BLUE
    ws.cell(r,7,f'=(3*E{r}+F{r})/4'); ws.cell(r,8,f'=(8*E{r}+F{r})/9')
    r+=1
ws.cell(r+1,1,'已搜尋，未找到：OpenAI o1（2024-12）、GPT-4.5（2025-02-27，openai.com 頁無絕對價）的官方牌價頁（Wayback 被擋）；Anthropic Mythos 5／5.1 無公開價；Gemini 1.5 Pro 上市價的官方頁（只找到社群貼文與降幅）。')
# ---------- c ----------
ws=wb.create_sheet('價格點_c_算力租賃'); title(ws,'(c) 算力租賃價格（隨需、預留、長約分開；每 GPU 小時與每 MW 年）')
cols=['GPU／產品','計價層','世代','時點','數值','單位','業者','來源','URL','文件日期','擷取日','標記','利害關係說明','備註']
hdr(ws,3,cols,[28,26,18,22,12,26,18,46,50,16,12,13,36,44])
r=4
for x in P.Cc:
    for j,v in enumerate([x['gpu'],x['layer'],x['gen'],x['date'],x['val'],x['unit'],x['vendor'],x['src'],x['url'],x['docdate'],P.RET,x['marker'],x['interest'],x['note']],1):
        c=ws.cell(r,j,v); c.alignment=WR
        if j==5: c.font=BLUE
    r+=1
ws.cell(r+1,1,'已搜尋，未找到：GB300 NVL72 在 CoreWeave、Nebius、AWS、Together AI、Gcore 的公開標價（皆「聯絡業務」）；VR200 NVL72 的任何公開價格；H100／B200／GB200 上市後 0–6 個月的官方價格頁快照（Wayback 被擋）；GB200 在 CoreWeave 上市價（2025-02）。Oracle 官方頁價格欄擷取為空，GB300 18.00 採二手彙整。')
# ---------- E338-E341 ----------
ws=wb.create_sheet('E338-E341確認'); title(ws,'E338–E341 一手確認與年變化率重算')
hdr(ws,3,['證據','項目','E338–E341 原值','一手來源與讀到的值','文件日期','判定','URL'],[8,34,26,70,14,28,60])
chk=[
('E338','GPT-4','30/60（2023-03）','$0.03/$0.06 每 1K','2023-03-14','一致','https://openai.com/index/gpt-4-research/'),
('E338','GPT-4 Turbo','10/30（2023-11）','$0.01/$0.03 每 1K','2023-11-06','一致','https://openai.com/index/new-models-and-developer-products-announced-at-devday/'),
('E338','GPT-4o','5/15（2024-05）','OpenAI 官方社群公告 $5/$15（本站頁無絕對價）','2024-05-13','一致（來源為官方社群公告）','https://community.openai.com/t/announcing-gpt-4o-in-the-api/744700'),
('E338','GPT-4o-0806','2.5/10','$2.50/$10.00','2024-08-06','一致','https://openai.com/index/introducing-structured-outputs-in-the-api/'),
('E338','GPT-4.1','2/8（2025-04）','$2.00/$8.00','2025-04-14','一致','https://openai.com/index/gpt-4-1/'),
('E338','GPT-5','1.25/10（2025-08）','$1.25/$10.00','2025-08-07','一致','https://openai.com/index/introducing-gpt-5-for-developers/'),
('E338','GPT-5.2','1.75/14（2025-12）','$1.75/$14','2025-12-11','一致','https://openai.com/index/introducing-gpt-5-2/'),
('E338','GPT-5.4','2.5/15（2026-03）','$2.50/$15','2026-03-05','一致','https://openai.com/index/introducing-gpt-5-4/'),
('E338','GPT-5.6 Terra','2/12（2026-07）','上市價 $2.50/$15（2026-07-09）；2026-07-30 降 20% → 2/12','2026-07-09／07-30','不一致（時點）：2/12 是降價後價；上市價 2.5/15','https://openai.com/index/gpt-5-6/'),
('E338','GPT-6 Sol','2/10（2026-09）','$2/$10；changelog 2026-09-22 上市','2026-09-22','一致（SRC_PRC_006–008 記 09-25，差 3 天）','https://openai.com/index/introducing-gpt-6-sol-and-luna/'),
('E338','GPT-6.1 Sol','2/10、快取 0.10（2026-09-29）','changelog：2026-09-29「$2 input, $0.10 cached input, $2.50 cache write, $10 output」','2026-09-29','一致','https://developers.openai.com/api/docs/changelog'),
('E339','Claude 3 Sonnet 至 Sonnet 4.6','3/15 持平','3 Sonnet、3.5、3.7、4、4.5、4.6 皆 $3/$15（Sonnet 4 由 4.5 公告間接確認）','2024-03-04～2026-02-17','一致','https://www.anthropic.com/news/claude-3-family'),
('E339','Claude Sonnet 5','3/15（2026-09 標準價）','上市價 $2/$10（2026-06-30）；原訂 9/1 起標準價 3/15，2026-08-10 更新「2/10 轉為永久，3/15 不再適用」；定價頁 Sonnet 5、5.5 皆 2/10','2026-06-30／08-10','不一致：Sonnet 現價 2/10，不是 3/15','https://www.anthropic.com/news/claude-sonnet-5'),
('E339','Sonnet 5 換 tokenizer','每單位文字實價未定','Sonnet 5 頁：同一輸入約 1.0–1.35× token','2026-06-30','一致','https://www.anthropic.com/news/claude-sonnet-5'),
('E340','Gemini 1.5 Pro','3.5/10.5','輸入 3.50 僅見社群貼文（2024-05-23）；官方頁取不到','2024-05-23','找不到官方一手（輸入：社群貼文；輸出：無）','https://discuss.ai.google.dev/t/pay-as-you-go-plan-of-gemini-pro-1-5-3-5-million-tokens-meaning/3375'),
('E340','Gemini 1.5 Pro 降價','1.25/5（2024-10）','官方：自 2024-10-01 輸入 −64%、輸出 −52%（價格表為圖片）；3.5×0.36＝1.26、10.5×0.48＝5.04，與 1.25/5 相符','2024-09-24','降幅一手確認；絕對值為推算相符','https://developers.googleblog.com/en/updated-production-ready-gemini-models-reduced-15-pro-pricing-increased-rate-limits-and-more/'),
('E340','Gemini 2.5 Pro','1.25/10','ai.google.dev 定價頁 ≤200K：$1.25/$10.00','現行頁（無日期）；GA 2025-06-17','一致','https://ai.google.dev/gemini-api/docs/pricing'),
('E340','Gemini 3 Pro','2/12','$2/$12（≤200K）；2025-11-18 上市','2025-11-18','一致','https://blog.google/technology/developers/gemini-3-developers/'),
('E340','Gemini 3.1 Pro','2/12','定價頁 Preview ≤200K $2.00/$12.00；2026-02-19 上市','2026-02-19','一致','https://ai.google.dev/gemini-api/docs/pricing'),
('E341','迴歸結果','見下表','無法由 E338–E340 復現（原始取點與腳本未留存）；重算見下表','—','不可復現','—')]
for i,row in enumerate(chk,4):
    for j,v in enumerate(row,1): ws.cell(i,j,v).alignment=WR
r=4+len(chk)+2
ws.cell(r,1,'E341 重算：對數線性迴歸，年變化率＝EXP(SLOPE(LN(混合價),年))−1；取點＝上市日；混合價＝w×輸入＋(1−w)×輸出').font=BOLD
ws.cell(r+1,1,'w（8:1，對應參考請求 ISL 16,384／OSL 2,048）'); ws.cell(r+1,5,'=8/9')
ws.cell(r+2,1,'w（3:1）'); ws.cell(r+2,5,'=3/4')
W8=f'$E${r+1}'; W3=f'$E${r+2}'
r+=4
def yr(d): y,m,dd=map(int,d.split('-')); return dt.date(y,m,dd).toordinal()/365.25
SER={
'OpenAI 中階（E338 原序列）':[('2023-03-14',30,60),('2023-11-06',10,30),('2024-05-13',5,15),('2024-08-06',2.5,10),('2025-04-14',2,8),('2025-08-07',1.25,10),('2025-12-11',1.75,14),('2026-03-05',2.5,15),('2026-07-09',2,12),('2026-09-22',2,10),('2026-09-29',2,10)],
'OpenAI 中階（確認序列：Terra 取上市價 2.5/15）':[('2023-03-14',30,60),('2023-11-06',10,30),('2024-05-13',5,15),('2024-08-06',2.5,10),('2025-04-14',2,8),('2025-08-07',1.25,10),('2025-12-11',1.75,14),('2026-03-05',2.5,15),('2026-07-09',2.5,15),('2026-09-22',2,10),('2026-09-29',2,10)],
'Claude Sonnet（確認序列：Sonnet 5、5.5 取 2/10）':[('2024-03-04',3,15),('2024-06-21',3,15),('2025-02-24',3,15),('2025-05-22',3,15),('2025-09-29',3,15),('2026-02-17',3,15),('2026-06-30',2,10),('2026-09-28',2,10)],
'Claude Sonnet（E339 說法：Sonnet 5、5.5 仍 3/15）':[('2024-03-04',3,15),('2024-06-21',3,15),('2025-02-24',3,15),('2025-05-22',3,15),('2025-09-29',3,15),('2026-02-17',3,15),('2026-06-30',3,15),('2026-09-28',3,15)],
'Gemini Pro（確認序列）':[('2024-05-23',3.5,10.5),('2024-10-01',1.25,5),('2025-04-04',1.25,10),('2025-11-18',2,12),('2026-02-19',2,12)],
}
WIN=[('2023–2026.8','2023-01-01','2026-08-31'),('2024.3–2026.8','2024-03-01','2026-08-31'),('2025–2026.8','2025-01-01','2026-08-31'),('2024.3–2026.10','2024-03-01','2026-10-31')]
res_rows=[]
for name,pts in SER.items():
    ws.cell(r,1,name).font=BOLD; r+=1
    for j,c in enumerate(['日期','輸入','輸出','年(小數)','LN(8:1)','LN(3:1)'],1): ws.cell(r,j,c).font=BOLD
    r+=1; start=r
    for d,i_,o in pts:
        ws.cell(r,1,d); ws.cell(r,2,i_).font=BLUE; ws.cell(r,3,o).font=BLUE
        y,m,dd=map(int,d.split('-')); ws.cell(r,4,f'=DATE({y},{m},{dd})/365.25')
        ws.cell(r,5,f'=LN({W8}*B{r}+(1-{W8})*C{r})'); ws.cell(r,6,f'=LN({W3}*B{r}+(1-{W3})*C{r})'); r+=1
    ws.cell(r,1,'窗口').font=BOLD; ws.cell(r,2,'點數').font=BOLD; ws.cell(r,5,'年變化 8:1').font=BOLD; ws.cell(r,6,'年變化 3:1').font=BOLD; r+=1
    for lab,s,e in WIN:
        idx=[k for k,(d,_,_) in enumerate(pts) if s<=d<=e]
        if len(idx)<3: continue
        a,b=start+idx[0],start+idx[-1]
        ws.cell(r,1,lab); ws.cell(r,2,len(idx))
        ws.cell(r,5,f'=EXP(SLOPE(E{a}:E{b},D{a}:D{b}))-1'); ws.cell(r,6,f'=EXP(SLOPE(F{a}:F{b},D{a}:D{b}))-1')
        ws.cell(r,5).number_format='0.0%'; ws.cell(r,6).number_format='0.0%'; r+=1
    r+=1
ws.cell(r,1,'E341 原值：OpenAI 2023–2026.8 −49%／−38%；2024.3–2026.8 −24%／−10%；2025–2026.8 +8%／+12%（參考請求權重／3:1）；Sonnet 持平；Gemini Pro −18%／+11%。')
for c in 'ABCDEF': pass
ws.column_dimensions['D'].width=70
# ---------- 上市間隔 ----------
ws=wb.create_sheet('上市間隔'); title(ws,'世代上市間隔（步驟 4）：發表會日期與「量產出貨／雲端開始供應」日期分開記')
hdr(ws,3,['世代','事件','日期','事件類型','來源','URL','標記','備註'],[14,46,13,18,44,56,14,40])
EV=[
('Hopper H100','GTC 發表','2022-03-22','發表會','（常識；本次未取一手頁）','—','未驗證','未逐筆取一手頁'),
('Hopper H100','NVIDIA 宣布 H100 全面量產','2022-09-20','量產','NVIDIA 新聞稿「Hopper in Full Production」','https://nvidianews.nvidia.com/news/nvidia-hopper-in-full-production','Interested-party','系統預計 2022-10 起出貨'),
('Hopper H100','CoreWeave、Cirrascale 雲端 GA','2023-03-21','雲端開始供應（基準）','NVIDIA 新聞稿「Hopper GPUs Expand Reach」','https://nvidianews.nvidia.com/news/nvidia-hopper-gpus-expand-reach-as-demand-for-ai-grows','Interested-party','同稿：OCI 限量、AWS 與 Azure 預覽'),
('Hopper H100','Lambda H100 隨需上線','2023-05-10','雲端開始供應','Lambda 部落格','https://lambda.ai/blog/lambda-cloud-deploys-nvidia-h100-tensor-core-gpus','Interested-party',''),
('Hopper H100','AWS P5 GA','2023-07-26','雲端開始供應（超大規模雲）','AWS What\'s New','https://aws.amazon.com/about-aws/whats-new/2023/07/amazon-ec2-p5-instances-generative-ai-hpc-generally-available/','Interested-party',''),
('GB200 NVL72','GTC 發表','2024-03-18','發表會','（常識；本次未取一手頁）','—','未驗證',''),
('GB200 NVL72','NVIDIA：Blackwell「in full production」','2024-11-20','量產','NVIDIA 第三季 FY2025 財報新聞稿','https://nvidianews.nvidia.com/news/nvidia-announces-financial-results-for-third-quarter-fiscal-2025','Interested-party','未載出貨起日'),
('GB200 NVL72','CoreWeave GB200 NVL72 GA（首家）','2025-02-04','雲端開始供應（基準）','NVIDIA 部落格','https://blogs.nvidia.com/blog/blackwell-coreweave-gb200-nvl72-instances-cloud/','Interested-party','CoreWeave 文件稱 2025-02-03 實例上線'),
('GB300 NVL72','GTC 發表','2025-03-18','發表會','（常識；本次未取一手頁）','—','未驗證',''),
('GB300 NVL72','CoreWeave 首家部署 GB300 NVL72（Dell 出貨）','2025-07-03','出貨／部署（基準）','CoreWeave 新聞稿','https://www.coreweave.com/news/coreweave-becomes-first-cloud-provider-to-deploy-nvidia-gb300-nvl72-platform','Interested-party',''),
('GB300 NVL72','CoreWeave GB300 實例上線','2025-08-19','雲端開始供應（替代）','CoreWeave 文件標題（搜尋結果標題，未開頁）','https://docs.coreweave.com/release-notes/gb300-nvl72','Interested-party','僅見搜尋結果標題'),
('VR200 NVL72','NVIDIA：Rubin「in full production」','2026-01-05','量產','NVIDIA CES 新聞稿','https://nvidianews.nvidia.com/news/rubin-platform-ai-supercomputer','Interested-party','合作夥伴 2026 下半年供應'),
('VR200 NVL72','CoreWeave 宣布 VR NVL72 供應，首位客戶 Cognition 正式上線','2026-09-30','雲端開始供應（基準）','CoreWeave 新聞稿','https://www.coreweave.com/news/coreweave-delivers-nvidia-vera-rubin-nvl72-performance-at-production-scale-starting-with-cognition','Interested-party','頁面未稱 GA；Cognition 叢集 2026-09 初建置'),
]
for i,row in enumerate(EV,4):
    for j,v in enumerate(row,1):
        c=ws.cell(i,j,v); c.alignment=WR
        if j==3: 
            y,m,dd=map(int,v.split('-')); c.value=dt.date(y,m,dd); c.number_format='yyyy-mm-dd'; c.font=BLUE
def rowof(gen,ev): 
    for i,row in enumerate(EV,4):
        if row[0]==gen and ev in row[1]: return i
rH=rowof('Hopper H100','CoreWeave、Cirrascale'); rHp=rowof('Hopper H100','全面量產')
rG=rowof('GB200 NVL72','CoreWeave GB200'); rGp=rowof('GB200 NVL72','in full production')
rB=rowof('GB300 NVL72','首家部署'); rBi=rowof('GB300 NVL72','實例上線')
rV=rowof('VR200 NVL72','宣布 VR NVL72'); rVp=rowof('VR200 NVL72','in full production')
r=4+len(EV)+2
ws.cell(r,1,'相鄰世代間隔（年）').font=BOLD; r+=1
hdr(ws,r,['世代對','基準（雲端開始供應→雲端開始供應）','低','高','基準天數','低天數','高天數','說明'])
for c in range(1,9): pass
r+=1; GAP0=r
gaps=[('Hopper→GB200',f'C{rG}-C{rH}',f'C{rG}-C{rH}',f'C{rGp}-C{rHp}','低＝雲端對雲端；高＝量產宣布對量產宣布（2022-09-20→2024-11-20）'),
      ('GB200→GB300',f'C{rB}-C{rG}',f'C{rB}-C{rG}',f'C{rBi}-C{rG}','低＝GB200 雲端 GA→GB300 首家部署；高＝→GB300 實例上線（2025-08-19）'),
      ('GB300→VR200',f'C{rV}-C{rBi}',f'C{rV}-C{rBi}',f'C{rV}-C{rB}','低＝GB300 實例上線→VR200 首位客戶上線；高＝GB300 首家部署→VR200（基準取中點）')]
# 基準：Hopper→GB200 雲端對雲端；GB200→GB300 部署；GB300→VR200 取 兩者平均
for k,(nm,b,lo,hi,note) in enumerate(gaps):
    rr=GAP0+k
    ws.cell(rr,1,nm)
    if nm=='GB300→VR200':
        ws.cell(rr,6,f'={lo}'); ws.cell(rr,7,f'={hi}'); ws.cell(rr,5,f'=AVERAGE(F{rr},G{rr})')
    elif nm=='Hopper→GB200':
        ws.cell(rr,5,f'={b}'); ws.cell(rr,6,f'={lo}'); ws.cell(rr,7,f'={hi}')
    else:
        ws.cell(rr,5,f'={b}'); ws.cell(rr,6,f'={lo}'); ws.cell(rr,7,f'={hi}')
    ws.cell(rr,2,f'=E{rr}/365.25'); ws.cell(rr,3,f'=MIN(F{rr},G{rr})/365.25'); ws.cell(rr,4,f'=MAX(F{rr},G{rr})/365.25')
    for c in (2,3,4): ws.cell(rr,c).number_format='0.00'
    ws.cell(rr,8,note).alignment=WR
ws.cell(GAP0+4,1,'發表會對發表會（參考，不作基準）：Hopper→Blackwell 2022-03-22→2024-03-18＝約 2.0 年；Blackwell→Blackwell Ultra 2024-03-18→2025-03-18＝1.0 年；發表會日期未逐筆取一手頁。')
ws.cell(GAP0+5,1,'標記：區間為 Analogy（三對世代各 1 個樣本，無統計意義）；日期本身為 Interested-party（業者新聞稿）。')
# ---------- 對照計算 ----------
ws=wb.create_sheet('對照計算'); title(ws,'步驟 5 實算對照：世代每 MW 收入指數（上市對上市）與回本餘裕','q、持有成本、每 MW GPU 數取自 model/20261014_Tokenomics_v5.35.xlsx（括號為原格）；Δ 取自「上市間隔」頁；藍字為人工輸入價格')
ws.column_dimensions['A'].width=58
for c in 'BCD': ws.column_dimensions[c].width=17
ws.column_dimensions['E'].width=17; ws.column_dimensions['F'].width=17; ws.column_dimensions['G'].width=70
def put(r,c,v,fmt=None,blue=False,bold=False):
    x=ws.cell(r,c,v)
    if fmt: x.number_format=fmt
    if blue: x.font=BLUE
    if bold: x.font=BOLD
    return x
hdr(ws,4,['A. 輸入（Excel v5.35）','Hopper→GB200','GB200→GB300','GB300→VR200','','','來源格'])
rows_in=[
(5,'q 基準（新世代÷舊世代 每 MW 總產出，Sol）',[4.84519865841925,1.45587602238485,1.56528522129431],'L1!D58、D59、D60'),
(6,'q 低（VR200 為 Sens_Perf η_d×0.5）',[4.84519865841925,1.45587602238485,0.915334889924169],'L1!E58、E59、E60'),
(7,'q 高（VR200 為 η_d×1.5）',[4.84519865841925,1.45587602238485,2.0506531956019],'L1!F58、F59、F60'),
(8,'每 MW 年持有成本 舊世代 $M/MW/年',[9.70308728611469,8.82816556596761,10.8864110021101],'L1!D54、D55、D56（L1_HoldEconMW_*）'),
(9,'每 MW 年持有成本 新世代 $M/MW/年',[8.82816556596761,10.8864110021101,12.2254710045841],'L1!D55、D56、D57'),
(11,'每 MW GPU 數 舊世代（顆）',[718.048,501.768,487.008],'Interface!D7、G7、J7（每 GW GPU 數基準欄）÷1000'),
(12,'每 MW GPU 數 新世代（顆）',[501.768,487.008,291.744],'Interface!G7、J7、M7 ÷1000'),
]
for r_,lab,vals,src in rows_in:
    put(r_,1,lab)
    for k,v in enumerate(vals): put(r_,2+k,v,'0.0000',blue=True)
    put(r_,7,src)
put(10,1,'持有成本比（新÷舊）')
for k in range(3): put(10,2+k,f'={L(2+k)}9/{L(2+k)}8','0.0000')
put(10,7,'公式')
put(13,1,'Δ 基準（年；「上市間隔」頁）')
for k in range(3): put(13,2+k,f"='上市間隔'!B{GAP0+k}",'0.00')
put(14,1,'Δ 低（年）')
for k in range(3): put(14,2+k,f"='上市間隔'!C{GAP0+k}",'0.00')
put(15,1,'Δ 高（年）')
for k in range(3): put(15,2+k,f"='上市間隔'!D{GAP0+k}",'0.00')
# B 門檻
hdr(ws,17,['B. 兩個門檻（年價格變化率 g；價格每年變化 g，持續 Δ 年）','Hopper→GB200','GB200→GB300','GB300→VR200','','','公式'])
put(18,1,'g*（收入指數＝1）：Δ 基準、q 基準')
put(19,1,'g*（收入指數＝1）：Δ 低')
put(20,1,'g*（收入指數＝1）：Δ 高')
put(21,1,'g*（回本餘裕＝1）：Δ 基準、q 基準')
put(22,1,'g*（回本餘裕＝1）：Δ 低')
put(23,1,'g*（回本餘裕＝1）：Δ 高')
put(24,1,'g*（收入指數＝1）：q 低、Δ 基準（VR200 區間低端）')
put(25,1,'g*（收入指數＝1）：q 高、Δ 基準')
for k in range(3):
    c=L(2+k)
    put(18,2+k,f'={c}5^(-1/{c}13)-1','0.0%'); put(19,2+k,f'={c}5^(-1/{c}14)-1','0.0%'); put(20,2+k,f'={c}5^(-1/{c}15)-1','0.0%')
    put(21,2+k,f'=({c}10/{c}5)^(1/{c}13)-1','0.0%'); put(22,2+k,f'=({c}10/{c}5)^(1/{c}14)-1','0.0%'); put(23,2+k,f'=({c}10/{c}5)^(1/{c}15)-1','0.0%')
    put(24,2+k,f'={c}6^(-1/{c}13)-1','0.0%'); put(25,2+k,f'={c}7^(-1/{c}13)-1','0.0%')
put(18,7,'q×(1+g)^Δ＝1 → g＝q^(−1/Δ)−1；實際年價格變化率高於此值，世代每 MW 收入指數＞1')
put(21,7,'q×(1+g)^Δ÷成本比＝1 → g＝(成本比÷q)^(1/Δ)−1；高於此值，回本餘裕＞1')
# C 推算
hdr(ws,27,['C. 推算值 q×(1+g)^Δ（Δ 基準）','Hopper→GB200','GB200→GB300','GB300→VR200','','','說明'])
lab=[('g＝−25%',-0.25),('g＝−10%',-0.10),('g＝+5%',0.05)]
for k,(nm,g) in enumerate(lab):
    put(28+k,1,f'收入指數：{nm}'); put(31+k,1,f'回本餘裕（指數÷成本比）：{nm}')
    for j in range(3):
        c=L(2+j)
        put(28+k,2+j,f'={c}$5*(1+({g}))^{c}$13','0.000'); put(31+k,2+j,f'={c}{28+k}/{c}$10','0.000')
put(28,7,'與 Cap_In I-1 的 g 基準／區間（−0.10；−0.25～+0.05）並列')
# D 實證
hdr(ws,35,['D. 實證價格 → 上市對上市因子與指數','Hopper→GB200','GB200→GB300','GB300→VR200','','','價格點與備註'])
W=ws.cell(36,1,'混合權重 w（輸入占比；0.75＝3:1，8/9＝8:1）'); put(36,2,0.75,'0.000',blue=True)
r=37
def block(name,olds,news,note,cells_ok=(True,True,True)):
    """olds/news: list of (in,out) or None per pair"""
    global r
    put(r,1,name,bold=True); put(r,7,note).alignment=WR; r+=1
    put(r,1,'  舊世代時點混合價 $/M'); put(r+1,1,'  新世代時點混合價 $/M'); put(r+2,1,'  價格因子（新÷舊）'); put(r+3,1,'  收入指數＝q×因子'); put(r+4,1,'  年化價格變化率'); put(r+5,1,'  回本餘裕＝指數÷成本比'); put(r+6,1,'  收入指數＝1 門檻：實證在門檻哪一側'); put(r+7,1,'  回本餘裕＝1 門檻：實證在哪一側')
    for k in range(3):
        c=L(2+k)
        if olds[k] is None:
            for dd in range(1,8): put(r+dd-1 if dd>1 else r,2+k,'無法實算') if dd==1 else None
            put(r,2+k,'無法實算'); put(r+1,2+k,'缺資料'); continue
        (oi,oo),(ni,no)=olds[k],news[k]
        put(r,2+k,f'=$B$36*{oi}+(1-$B$36)*{oo}','0.000'); put(r+1,2+k,f'=$B$36*{ni}+(1-$B$36)*{no}','0.000')
        put(r+2,2+k,f'={c}{r+1}/{c}{r}','0.0000'); put(r+3,2+k,f'={c}$5*{c}{r+2}','0.000'); put(r+4,2+k,f'={c}{r+2}^(1/{c}$13)-1','0.0%'); put(r+5,2+k,f'={c}{r+3}/{c}$10','0.000')
        put(r+6,2+k,f'=IF({c}{r+4}>={c}$18,"高於門檻（指數≥1）","低於門檻（指數<1）")'); put(r+7,2+k,f'=IF({c}{r+4}>={c}$21,"高於門檻（餘裕≥1）","低於門檻（餘裕<1）")')
    r+=9
block('b1 OpenAI 中階牌價（GPT-4 → 4o-0806 → 4.1 → 6.1 Sol）',[(30,60),(2.5,10),(2,8)],[(2.5,10),(2,8),(2,10)],'時點：2023-03-21 最新＝GPT-4（當時為旗艦）；2025-02-04＝GPT-4o-0806；2025-07-03＝GPT-4.1；2026-09-30＝GPT-6.1 Sol。第一對的「中階」其實是旗艦；層級口徑不連續。')
block('b2 Anthropic Sonnet 牌價',[None,(3,15),(3,15)],[None,(3,15),(2,10)],'Hopper→GB200：2023-03 時點無 Claude 3 系列（Claude 3 為 2024-03），缺資料。GB200→GB300 兩端皆 3/15；GB300→VR200 為 Sonnet 4（3/15）→Sonnet 5.5（2/10）。')
block('b3 Google Pro 牌價',[None,(1.25,5),(1.25,10)],[None,(1.25,10),(2,12)],'Hopper→GB200：2023-03 時點無 Gemini Pro 牌價資料（1.5 Pro 為 2024-05）。GB200 時點＝1.5 Pro（降價後）；GB300 時點＝2.5 Pro；VR200 時點＝3.1 Pro（Gemini 4 Argon 2026-10-13 晚於 VR200 基準日）。')
# a1/a2 point-based first pair; rate-based others
put(r,1,'a1 同等能力（Epoch MMLU≥86）：Hopper→GB200 用逐點；其餘用速率',bold=True); put(r,7,'Epoch 表為 3:1 混合價，因此此區塊不受 B36 影響。第一對：37.50（2023-03-14）→0.18（2025-02-05）。第二、三對無逐點資料，改用下列速率換算。').alignment=WR; r+=1
a1=r
put(r,1,'  舊世代時點價 $/M（Epoch，3:1）'); put(r,2,37.5,'0.00',blue=True); put(r+1,1,'  新世代時點價 $/M'); put(r+1,2,0.18,'0.00',blue=True)
put(r+2,1,'  價格因子'); put(r+2,2,f'=B{r+1}/B{r}','0.0000'); put(r+3,1,'  收入指數＝q×因子'); put(r+3,2,f'=B5*B{r+2}','0.000'); put(r+4,1,'  年化價格變化率'); put(r+4,2,f'=B{r+2}^(1/B13)-1','0.0%')
put(r+5,1,'  回本餘裕'); put(r+5,2,f'=B{r+3}/B10','0.000'); put(r+6,1,'  收入指數＝1 門檻'); put(r+6,2,f'=IF(B{r+4}>=B18,"高於門檻","低於門檻（指數<1）")'); put(r+7,1,'  回本餘裕＝1 門檻'); put(r+7,2,f'=IF(B{r+4}>=B21,"高於門檻","低於門檻（餘裕<1）")')
r+=9
put(r,1,'a2 同等能力（Epoch GPQA≥33）：Hopper→GB200',bold=True); r+=1
put(r,1,'  舊 37.50 → 新 0.12（Phi 4，2024-12-13；2025-02-04 時點最低價）'); put(r,2,0.12/37.5,'0.0000'); put(r+1,1,'  收入指數＝q×因子'); put(r+1,2,f'=B5*B{r}','0.000'); put(r+2,1,'  年化價格變化率'); put(r+2,2,f'=B{r}^(1/B13)-1','0.0%'); r+=4
put(r,1,'a3 同等能力：速率換算（每年下降倍數 R；因子＝R^(−Δ)）',bold=True); put(r,7,'速率來源：Gundlach 等 5–10 倍／年（取 7.5）；Epoch 2026 平均 13 倍／年；Epoch 2026「SOTA 後兩年」4.7 倍／年；Epoch 2025 GPQA 40 倍／年。注意：此速率含演算法效率與競價，不只硬體。').alignment=WR; r+=1
for R,nm in [(4.7,'4.7 倍／年（Epoch 2026：SOTA 後兩年）'),(7.5,'7.5 倍／年（Gundlach 等 5–10 倍的中點）'),(13,'13 倍／年（Epoch 2026 平均，≈每季 −47%）'),(40,'40 倍／年（Epoch 2025，GPQA／GPT-4-0314 水準）')]:
    put(r,1,f'  R＝{nm}'); put(r,5,R,'0.0',blue=True)
    put(r+1,1,'    價格因子 R^(−Δ)'); put(r+2,1,'    收入指數＝q×因子'); put(r+3,1,'    回本餘裕'); put(r+4,1,'    年化價格變化率')
    for k in range(3):
        c=L(2+k)
        put(r+1,2+k,f'=$E{r}^(-{c}$13)','0.0000'); put(r+2,2+k,f'={c}$5*{c}{r+1}','0.000'); put(r+3,2+k,f'={c}{r+2}/{c}$10','0.000'); put(r+4,2+k,f'=1/$E{r}-1','0.0%')
    r+=6
# c
put(r,1,'c 算力租賃：每 GPU 小時 → 每 MW 年（單位換算：價格 × 每 MW GPU 數 × 8760 小時 ÷ 1e6，利用率 100%、未含折扣）',bold=True); put(r,7,'價格已含硬體速度，因此不再乘 q。換算用 Interface 第 7 列基準欄的每 GW GPU 數 ÷1000。').alignment=WR; r+=1
put(r,1,'  年小時數'); put(r,2,8760,'0',blue=True); HR=f'$B${r}'; r+=1
put(r,1,'  CoreWeave 現價 H100 8 卡 $/GPU-hr'); put(r,2,6.16,'0.00',blue=True); cH=r; r+=1
put(r,1,'  CoreWeave 現價 GB200 NVL72 $/GPU-hr'); put(r,2,10.50,'0.00',blue=True); cG=r; r+=1
put(r,1,'  同業者橫斷面：GB200÷H100 價格比（現價，非上市對上市）'); put(r,2,f'=B{cG}/B{cH}','0.000'); cR=r; r+=1
put(r,1,'  收入指數（每 MW）＝價格比 × 每 MW GPU 數比'); put(r,2,f'=B{cR}*B12/B11','0.000'); cI=r; r+=1
put(r,1,'  回本餘裕＝指數 ÷ 成本比（Hopper→GB200）'); put(r,2,f'=B{cI}/B10','0.000'); r+=1
put(r,1,'  GB200→GB300、GB300→VR200'); put(r,3,'無法實算'); put(r,4,'無法實算'); put(r,7,'缺：CoreWeave／Nebius 的 GB300 公開標價（聯絡業務）；任何 VR200 公開價格；OCI 的 GB200 價格').alignment=WR; r+=2
put(r,1,'  每 MW 年收入上限（100% 利用率）$M/MW/年',bold=True); r+=1
price_pts=[('CoreWeave H100 隨需 6.16',6.16,'B11'),('Silicon Data 新雲 H100 中位 3.33（2025-07～12）',3.33,'B11'),('Silicon Data 市集 H100 中位 1.95（2025-07～12）',1.95,'B11'),('Lambda H100 SXM 8x 隨需 2.99（2024-10）',2.99,'B11'),('CoreWeave GB200 隨需 10.50',10.50,'B12'),('Oracle GB300 隨需 18.00（二手彙整）',18.00,'B12'),]
# GPU per MW cells: Hopper B11, GB200 B12 (=C11 column?), GB300 = D11
mapping={'B11':'B11','B12':'C11','GB300':'D11'}
def gpumw(label):
    return 'B11' if 'H100' in label else ('C11' if 'GB200' in label else 'D11')
costcell={'B11':'B8','C11':'C8','D11':'D8'}
hdr(ws,r,['  價格點','$/GPU-hr','每 MW GPU 數','$M/MW/年（上限）','持有成本 $M/MW/年','上限÷持有成本',''])
r+=1
for lab,v,_ in price_pts:
    g=gpumw(lab)
    put(r,1,'  '+lab); put(r,2,v,'0.00',blue=True); put(r,3,f'={g}','0.0'); put(r,4,f'=B{r}*C{r}*{HR}/1000000','0.00'); put(r,5,f'={costcell[g]}','0.000'); put(r,6,f'=D{r}/E{r}','0.00'); r+=1
put(r,1,'  IREN—Microsoft GB300 長約（賣方揭露，$9.7B÷5 年÷200 MW）'); put(r,4,'=9.7/5/200*1000','0.00',blue=True); put(r,5,'=D8*0+D9*0+C9','0.000'); 
# GB300 hold cost is C9? C9 = new for pair2 = GB300 hold 10.886
put(r,5,'=C9','0.000'); put(r,6,f'=D{r}/E{r}','0.00'); put(r,7,'$M/MW-IT/年；與 Excel 契約第 4 節對帳提醒 9.70 一致；回本餘裕 0.89（GB300 持有成本 10.886）').alignment=WR; r+=2
put(r,1,'註：隨需牌價÷持有成本遠大於 1，長約 0.89 小於 1：隨需牌價是「上限」，簽約率、利用率與折扣未計入，不能直接當作每 MW 簽約收入。').alignment=WR
wb.save(OUT); print('saved',OUT)
