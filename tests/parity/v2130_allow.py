"""v21.3 看守的「v21.2 → v21.3 逐格差異」集合，供舊版逐格差異測試（*_diffs_only_expected_*）放寬用。

只放寬到工作單 docs/workorders/20261024_v21.3_機架價低端與Sol快取價.md 的 W1–W6 列明的格；
規則與 test_v2130.test_v2130_diffs_only_expected_vs_v212 相同（那裡是權威，這裡重複驗證一次，寬於工作單即斷言失敗）。
舊測試的做法：實際差異與預期清單兩邊都扣掉本集合（`strip`），其餘比較不變。
"""
import re
from functools import lru_cache
from pathlib import Path

import openpyxl

from conftest import model_path

REPO = Path(__file__).resolve().parents[2]
V212 = REPO / "model" / "archive" / "20261023_Tokenomics_v21.2.xlsx"
SHEETS = {"README", "Spec_Rack", "Cap_In", "Alloc_In", "Alloc", "L1", "Checks", "Decisions", "DB_Evidence", "SRC_Price", "SRC_HW", "SRC_Index"}
_row = lambda coord: int(re.sub(r"[A-Z]+", "", coord))
_col = lambda coord: re.sub(r"\d+", "", coord)
_num = lambda x: isinstance(x, (int, float)) and not isinstance(x, bool)


@lru_cache(maxsize=None)
def v213_cells() -> dict:
    a, b = openpyxl.load_workbook(V212), openpyxl.load_workbook(model_path())
    assert a.sheetnames == b.sheetnames
    d = {}
    for ws in a:
        wn = b[ws.title]
        for r in range(1, max(ws.max_row, wn.max_row) + 1):
            for c in range(1, max(ws.max_column, wn.max_column) + 1):
                x, y = ws.cell(r, c).value, wn.cell(r, c).value
                if x != y:
                    d.setdefault(ws.title, []).append((ws.cell(r, c).coordinate, x, y))
    assert set(d) <= SHEETS, set(d) - SHEETS
    exact = {
        "README": {"A1", "A73", "B5", "B73"}, "Spec_Rack": {"F11"}, "Cap_In": {"A47", "D7", "F47", "G7"}, "Alloc_In": {"G23"},
        "L1": {"I128", "I129", "I130", "I131"}, "Checks": {"D86"},
        "Decisions": {f"{c}332" for c in "ABCDEFGHIJ"},
        "DB_Evidence": {f"{c}{r}" for r in (328, 329) for c in "ABCDEFGHIJKLMNOPQR"},
        "SRC_Price": {"B10", "H10", "W10", "B12", "H12", "W12", "O11", "P11", "V11", "O53", "Q53", "V53", "W53"},
    }
    out = {}
    for sh, lst in d.items():
        got = {t[0] for t in lst}
        if sh in exact:
            assert got == exact[sh], (sh, sorted(got ^ exact[sh]))
        elif sh == "SRC_HW":      # 新增兩列（第 102、103 列）、X 欄範圍、V20、W21
            assert all(_row(x) in (102, 103) or (_col(x) == "X" and 5 <= _row(x) <= 101) or x in ("V20", "W21") for x in got)
        elif sh == "SRC_Index":      # 只因 SRC_HW 多兩列而重排（A:F 欄，第 102 列起）
            assert all(_col(x) in "ABCDEF" and _row(x) >= 102 for x in got)
        elif sh == "Alloc":      # 版本字串 C177、區間端點常數（數字對數字）、指紋記錄列（第 1230–1240 列）
            assert all(t[0] == "C177" or (_num(t[1]) and _num(t[2])) or 1230 <= _row(t[0]) <= 1240 for t in lst)
        out[sh] = got
    return out


def strip(sheet: str, coords) -> list:
    """從座標清單扣掉 v21.3 的差異格（舊測試的實際清單與預期清單兩邊都用）。"""
    gone = v213_cells().get(sheet, set())
    return sorted(c for c in coords if c not in gone)


@lru_cache(maxsize=None)
def csv_value_cells() -> dict:
    """預期變動 CSV 的 (頁, 儲存格) -> (v21.2 值, v21.3 值)（工作單第 0 節，chat 端實測）。"""
    import csv
    out = {}
    for r in csv.DictReader(open(REPO / "docs" / "workorders" / "20261024_v21.3_預期變動.csv", encoding="utf-8")):
        a, b = r["v21.2"], r["new"]
        try:
            a, b = float(a), float(b)
        except ValueError:
            pass      # 文字格（動態說明文字）照原字串
        out[(r["sheet"], r["cell"])] = (a, b)
    return out
