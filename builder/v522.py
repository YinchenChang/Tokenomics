# v5.22 (work order docs/workorders/20261006_v5.22.md r4): X7 lifetime price factor L and monetisation rate m as side-by-side outputs;
# X8 supplement (one DB_Evidence row for L, no SRC records); X9 Q2 (MLPerf v6.1 second sources filled in); Decisions X7／X8／X9.
#
# X7 adds two control inputs (Cap_In G section, base 1 = existing outputs unchanged) and four Theory_Rev rows (base row × L × m) that
# Interface H links to as IF_RevGW_Luna／Sol／Astra_Life and IF_RevGWFleet_Life. The existing IF_RevGW_*, IF_RevGWFleet, L1 and every
# *_Prod row keep their formulas and values. No arithmetic outside the workbook.
# Every write to an Excel-owned cell is guarded by the old value / presence, like v518 and v521, so a rebuild never overwrites a later Excel edit.
import re
from copy import copy
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
import v518
import v521
from common import put, F_IN, F_NOTE, F_BOLD, section

VERSION = "20261006_Tokenomics_v5.22"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-06"
V = "v5.22"
EID = "E246"
COLS = range(3, 18)                         # C:Q (5 generations x 3 columns)

# ------------------------------------------------------------------ X7: inputs (Cap_In G section; after the existing content, nothing moves)
CAP_START = 61                              # Cap_In ends at row 59 (note) in v5.21; the G section starts after one blank row
ROW_L, ROW_M = CAP_START + 2, CAP_START + 3
LAB_L = "壽命期價格係數 L（X7；CTL_PriceLife）"
LAB_M = "變現率 m（X7；CTL_Monetize）"
DESC_L = ("設備 6 年壽命期內平均實收單價 ÷ 2026-09／10 單價快照（K9）。1＝不跌價（模型現行假設）；只驅動 Theory_Rev D 節與 Interface H 節的 _Life 並列輸出，"
          "既有 IF_RevGW_*、L1 與 _Prod 列不受影響。區間 0.3–1.0 見 Gov_Map；類比證據見 DB_Evidence " + EID + "（H100 租金指數，摘要級）")
DESC_M = ("付費 token 平均實收單價 ÷ OpenAI 有效單價（表列 ×(1−折扣)）。1＝所有付費 token 依 API 有效價計收（模型現行假設）；只驅動 _Life 並列輸出。"
          "區間 0.3–1.0 見 Gov_Map；無實證（缺每位付費用戶年 token 消耗量）")


def cap_in_g(wb):
    """Two blue inputs (value 1) and their named ranges. Asserts the sheet still ends where v5.21 ended, so nothing is silently moved."""
    ws = wb["Cap_In"]
    assert ws.max_row == CAP_START - 2, f"Cap_In ends at row {ws.max_row}; X7 section expects it to start at row {CAP_START}"
    section(ws, CAP_START, "G. 營收實現（X7，v5.22）：壽命期價格係數 L 與變現率 m — 並列輸出用控制參數（基準 1＝既有輸出不變；"
                           "Theory_Rev D 節、Interface H 節；區間見 Gov_Map）", 7)
    r = CAP_START + 1
    put(ws, f"A{r}", "項目", F_BOLD); put(ws, f"B{r}", "單位", F_BOLD); put(ws, f"C{r}", "值", F_BOLD)
    put(ws, f"F{r}", "標記", F_BOLD); put(ws, f"G{r}", "說明", F_BOLD)
    for row, lab, tag, desc in ((ROW_L, LAB_L, "Assumed（Analogy 證據見 DB_Evidence " + EID + "）", DESC_L), (ROW_M, LAB_M, "Assumed", DESC_M)):
        put(ws, f"A{row}", lab, wrap=True); put(ws, f"B{row}", "x"); put(ws, f"C{row}", 1, fmt="0.00")
        put(ws, f"F{row}", tag, F_NOTE); put(ws, f"G{row}", desc, F_NOTE, wrap=True)
    for n, row in (("CTL_PriceLife", ROW_L), ("CTL_Monetize", ROW_M)):
        if n in wb.defined_names: del wb.defined_names[n]
        wb.defined_names[n] = DefinedName(n, attr_text=f"Cap_In!$C${row}")


# ------------------------------------------------------------------ X7: Theory_Rev D section (base row x L x m) and its self-check rows
TAG = re.compile(r"　\[[A-Za-z0-9_]+\]\s*$")
TARGETS = [(11, "IF_RevGW_Luna"), (29, "IF_RevGW_Sol"), (47, "IF_RevGW_Astra"), (67, "IF_RevGWFleet")]   # (Theory_Rev row, Interface name of the baseline)
SUFFIX = "_Life"
TR_START = 89
TR_ROWS = {base: TR_START + 1 + i for i, (_, base) in enumerate(TARGETS)}             # life rows 90–93
TR_FLAGS = {base: TR_START + 7 + i for i, (_, base) in enumerate(TARGETS)}            # self-check rows 96–99


def theory_rev_life(wb):
    ws = wb["Theory_Rev"]
    assert ws.max_row == TR_START - 2, f"Theory_Rev ends at row {ws.max_row}; X7 section expects it to start at row {TR_START}"
    section(ws, TR_START, "D. 營收實現並列輸出（X7，v5.22）：基準列 × 壽命期價格係數 L（CTL_PriceLife）× 變現率 m（CTL_Monetize）；"
                          "L＝m＝1 時等於基準列；基準列為文字（SLO 不可達）時輸出相同文字", 17)
    for trow, base in TARGETS:
        lab = TAG.sub("", ws.cell(trow, 1).value)
        rr = TR_ROWS[base]
        put(ws, f"A{rr}", f"{lab}（× 壽命期價格係數 L × 變現率 m）　[{base}{SUFFIX}]"); put(ws, f"B{rr}", ws.cell(trow, 2).value)
        for c in COLS:
            col = L(c); src = ws.cell(trow, c)
            put(ws, f"{col}{rr}", f"=IF(ISNUMBER({col}{trow}),{col}{trow}*CTL_PriceLife*CTL_Monetize,{col}{trow})", fmt=src.number_format)
    put(ws, f"A{TR_START + 6}", "自我檢查（Checks Y／X7）：Interface 的 _Life 列 − 基準列 × L × m 的絕對差 ≤ 1e-9（兩邊皆為文字時須文字相同）；0＝相符，1＝不符", F_BOLD)
    for trow, base in TARGETS:
        rr = TR_FLAGS[base]
        put(ws, f"A{rr}", f"{base}{SUFFIX}：Interface 列 vs {base} × L × m", F_NOTE)
        for k, c in enumerate(COLS, start=1):
            a, b = f"INDEX({base}{SUFFIX},1,{k})", f"INDEX({base},1,{k})"
            put(ws, f"{L(c)}{rr}",
                f"=IF(AND(ISNUMBER({a}),ISNUMBER({b})),IF(ABS({a}-{b}*CTL_PriceLife*CTL_Monetize)<=0.000000001,0,1),"
                f"IF(AND(NOT(ISNUMBER({a})),NOT(ISNUMBER({b}))),IF({a}={b},0,1),1))", fmt="0")


# ------------------------------------------------------------------ X7: Interface H section, four named rows C:Q
def interface_h(wb):
    ws = wb["Interface"]
    start = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    section(ws, start, "H. 營收實現並列輸出（X7，v5.22）：D 節基準列 × 壽命期價格係數 L（CTL_PriceLife，Cap_In）× 變現率 m（CTL_Monetize，Cap_In）；"
                       "基準 L＝m＝1 時等於 D 節對應列；SLO 不可達時為文字", 17)
    r = start + 1; made = []
    for trow, base in TARGETS:
        brow = int(re.match(r"Interface!\$C\$(\d+):\$Q\$(\d+)", wb.defined_names[base].attr_text).group(1))
        lab = TAG.sub("", ws.cell(brow, 1).value)
        put(ws, f"A{r}", f"{lab}（× L × m）　[{base}{SUFFIX}]"); put(ws, f"B{r}", ws.cell(brow, 2).value)
        for c in COLS:
            bc = ws.cell(brow, c)
            put(ws, f"{L(c)}{r}", f"=Theory_Rev!{L(c)}{TR_ROWS[base]}", fmt=bc.number_format, fill=copy(bc.fill) if bc.fill.fill_type else None)
        made.append((f"{base}{SUFFIX}", f"Interface!$C${r}:$Q${r}")); r += 1
    for n, ref in made:
        if n in wb.defined_names: del wb.defined_names[n]
        wb.defined_names[n] = DefinedName(n, attr_text=ref)
    return dict(start=start, made=made)


# ------------------------------------------------------------------ X7: Checks Y section (ERROR, counted in GOV_Errors)
def checks_y(wb):
    ws = wb["Checks"]
    r = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    section(ws, r, "Y. 營收實現並列輸出檢查（工作單 X7，v5.22）：X7 為 ERROR 計入 GOV_Errors", 6); r += 1
    for i, h in enumerate(["編號", "檢查", "等級", "筆數", "範圍與算法"]): put(ws, f"{L(i+1)}{r}", h, F_BOLD)
    r += 1
    first, last = TR_FLAGS[TARGETS[0][1]], TR_FLAGS[TARGETS[-1][1]]
    put(ws, f"A{r}", "X7"); put(ws, f"B{r}", "營收實現並列輸出：Interface 的 _Life 列不等於 基準列 × L × m 的格數（絕對差 > 1e-9；文字列須相同）"); put(ws, f"C{r}", "ERROR", F_BOLD)
    put(ws, f"D{r}", f"=SUM(Theory_Rev!C{first}:Q{last})", fmt="0")
    put(ws, f"E{r}", f"Theory_Rev 第 {first}–{last} 列 × C:Q（4 × 15＝60 格）", F_NOTE)
    ref = wb.defined_names["GOV_Errors"].attr_text.split("!")[1].replace("$", "")
    cell = ws[ref]
    if f"D{r}" not in str(cell.value): cell.value = f"{cell.value}+D{r}"
    ws[f"B{int(ref[1:])}"].value = "ERROR 合計（CI 讀取 GOV_Errors；含 G、H、X、Y 節）"
    return r


# ------------------------------------------------------------------ Gov_Map rows (two new input cells; same dict format as v519.GOV_MAP_V519)
GOV_MAP_V522 = [
    {'scope': '切片三（v5.22 X7）', 'sheet': 'Cap_In', 'cell': f'C{ROW_L}', 'label': LAB_L, 'check': LAB_L,
     'cls': '情境值', 'role': '單值', 'src': '', 'rel': '', 'dec': 'X7', 'lo': 0.3, 'hi': 1.0,
     'rtext': '區間 0.3–1.0（Andy 2026-10-06，X7；C1）',
     'reason': f'X7：Assumed；基準 1＝不跌價（模型現行假設）；下端 0.3＝H100 以短缺高峰為起點的第 3 年比值；Analogy 證據見 DB_Evidence {EID}（類比區間 0.42–0.70）',
     'retag': '', 'seg': '高'},
    {'scope': '切片三（v5.22 X7）', 'sheet': 'Cap_In', 'cell': f'C{ROW_M}', 'label': LAB_M, 'check': LAB_M,
     'cls': '情境值', 'role': '單值', 'src': '', 'rel': '', 'dec': 'X7', 'lo': 0.3, 'hi': 1.0,
     'rtext': '區間 0.3–1.0（Andy 2026-10-06，X7；C1）',
     'reason': 'X7：Assumed；無實證（缺每位付費用戶年 token 消耗量）；上端 1.0＝模型現行假設', 'retag': '', 'seg': '高'},
]

# ------------------------------------------------------------------ X9 Q2: SRC_Perf second sources (column R) and the note tail (column W)
PAIRS = [("SRC_PERF_052", "SRC_PERF_056"), ("SRC_PERF_053", "SRC_PERF_057"), ("SRC_PERF_054", "SRC_PERF_058"), ("SRC_PERF_055", "SRC_PERF_059")]
NOTE_TAIL = f" ｜第二來源：不同提交者同情境（{V} X9 Q2）；Nebius VR200 為 36 GPU"


def src_second_sources(wb):
    """R is written only while it is still empty or '—'; W gets the tail appended (v518._append_text: no duplicate). Returns (n written, log)."""
    ws = wb["SRC_Perf"]; n = 0; log = []
    rows = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}
    for a, b in PAIRS:
        for me, other in ((a, b), (b, a)):
            r = rows.get(me)
            if r is None: log.append(f"{me}: record not found"); continue
            if ws.cell(r, 18).value in (None, "", "—"):
                ws.cell(r, 18).value = other; n += 1
                v518._append_text(ws.cell(r, 23), NOTE_TAIL)
            else:
                log.append(f"{me}!R: kept (already {ws.cell(r, 18).value!r})")
    return n, log


# ------------------------------------------------------------------ DB_Evidence (17 columns A..Q; same layout as v518／v520／v521)
READ = "摘要級，原文未讀"
URLS = ["gpusmith.com/articles/h100-rental-price-history-trends", "introl.com/blog/gpu-cloud-price-collapse-h100-market-december-2025",
        "spheron.network/blog/h100-price-per-hour-2026", "xenospectrum.com/en/nvidia-gpu-rental-depreciation",
        "tomtunguz.com/b200-gpu-pricing-spot-market-model-releases", "simplefunctions.dev/answer/b200max"]
CLAIM = ("產業 GPU 租金指數（摘要級，原文未讀）：H100 上市約 4.70、短缺高峰超過 8、2024 年初約 2.85、2026 年中隨需中位數約 2.99 $/GPU-hr（不同來源拼接）；"
         "H100 一年期長約 2025-10 1.70 → 2026-03 2.35（SemiAnalysis）；B200 標準化指數 2026-01-01 4.40 → 2026-08-02 5.66（Silicon Data）；"
         "B200 現貨指數 2026-03 初 2.31 → 2026-04 底 4.95（Ornn）；預測市場 2026-05-28 對 B200 年底高於 5.57 約 90%（Kalshi，成交量小）。")
VERDICT = ("H100 第 3 年比值 r：高峰起點 0.3–0.37、上市起點 0.5–0.64；壽命平均 L＝(1＋5r)／6 的類比區間 0.42–0.70（工作單 v5.22 第 0.1 節 C2）；"
           "B200 第一年不降反升（短缺驅動），未滿 24 個月不判定；基準維持 1（X7）。")
EVIDENCE_V522 = [[
    EID, DATE, CLAIM,
    "轉述頁（摘要級，原文未讀）：" + "；".join(URLS),
    READ,
    f"Cap_In C{ROW_L} 壽命期價格係數 L（CTL_PriceLife；Gov_Map X7 新列）",
    "1（基準＝不跌價；Gov_Map 區間 0.3–1.0）",
    "壽命平均 L 類比區間 0.42–0.70（H100 第 3 年比值 r：高峰起點 0.3–0.37、上市起點 0.5–0.64）",
    VERDICT, V,
    f"讀取者：chat 端（工作單 v5.22 第 2 節）；摘要級、原文未讀，不得標為已讀｜不新增任何 SRC 紀錄（X8 補充：Tokenomics 只放與算力相關的資料；租金走勢由 chat 端定期查核、寫入交接檔，不寫入模型）｜"
    "既有 SRC_PRC_001、SRC_PRC_002 不動｜Andy 2026-10-06「你的建議我接受，就這樣做吧」",
    "已處理", "—", "—", "IF_RevGW_Luna／Sol／Astra_Life、IF_RevGWFleet_Life（L＝m＝1 時數值不變）",
    "Analogy", "利害關係方（指數業者、算力販售者、研究機構）"]]


def evidence_rows():
    return EVIDENCE_V522


# ------------------------------------------------------------------ Decisions X7／X8／X9 (10 columns; same layout as v519／v520／v521)
DECISIONS_V522 = [
    ["X7", V, "壽命期價格係數 L、變現率 m 並列輸出（基準 1）",
     f"新增 L（CTL_PriceLife，Cap_In!C{ROW_L}）與 m（CTL_Monetize，Cap_In!C{ROW_M}）兩個控制參數，基準皆為 1（既有輸出不變），區間皆 0.3–1.0（Gov_Map，Assumed；L 另有類比證據 DB_Evidence {EID}，類比區間 0.42–0.70）；"
     "Interface H 節新增 IF_RevGW_Luna／Sol／Astra_Life、IF_RevGWFleet_Life＝對應基準列 × L × m；既有 IF_RevGW_*、IF_RevGWFleet、L1_Ans5 與 _Prod 列不改；Checks X7 自我檢查；不做 _Prod × L × m 組合列",
     DATE, "「同意」（回覆 chat 端建議：L 與 m 參數化且基準維持 1）；區間：「都照建議」（C1、C2）", "v5.22 已建",
     f"Cap_In!C{ROW_L}、C{ROW_M}（CTL_PriceLife、CTL_Monetize）；Theory_Rev D 節；Interface H 節；Checks X7；DB_Evidence {EID}；Gov_Map X7 兩列", "工作單 v5.22 第 0、1 節", "否"],
    ["X8", V, "X8 補充：租金資料不進 Tokenomics",
     f"Tokenomics 只放與算力相關的資料。公司專屬租金（牌價、合約價、出租率）不再新增，既有 SRC_PRC_001 維持；產業租金指數不建 SRC 時間序列、不建出租方層，只在 DB_Evidence 登錄一筆（{EID}）作為 L 區間依據；"
     "Checks「L 觀察值」取消；租金走勢由 chat 端定期查核，寫入交接檔、不寫入模型",
     DATE, "「你的建議我接受，就這樣做吧」", "v5.22 已建", f"DB_Evidence {EID}；SRC_PRC_001／002（不動）", "工作單 v5.22 第 0、2 節", "否"],
    ["X9", V, "v5.21 報告第九節 7 項處理",
     "Q1 Calib F67 不改，Perf_Batch 側 VR200 效率先做唯讀敏感度（v5.22 報告，不改模型），再決定是否另立批次口徑參數；Q2 Nebius 與 NVIDIA 同世代同情境互填第二來源（SRC_PERF_052–059，W1 −8）；"
     "Q3 不補登其他提交者；Q4 inferred 維持「未查得」；Q5 日期欄維持；Q6 取整不處理；Q7 restore_log 維持 Excel 優先",
     DATE, "「都照建議」", "v5.22 已建", "SRC_Perf R、W 欄（SRC_PERF_052–059）；Checks W1；v5.22 報告第 3 節敏感度", "工作單 v5.22 第 0、3、4 節", "否"],
]

# ------------------------------------------------------------------ README (version string is built from VERSION; v5.21 text is kept after it)
README_VERSION = (VERSION + "（X7：新增壽命期價格係數 L（CTL_PriceLife）與變現率 m（CTL_Monetize）兩個控制參數（Cap_In G 節，基準 1、區間 0.3–1.0），Interface H 節並列輸出 "
                  "IF_RevGW_Luna／Sol／Astra_Life 與 IF_RevGWFleet_Life（＝對應基準列 × L × m），Checks X7 自我檢查；X8 補充：DB_Evidence " + EID + "（產業租金指數，L 的類比證據，不建 SRC）；"
                  "X9 Q2：SRC_PERF_052–059 互填第二來源（W1 −8）；Decisions X7／X8／X9；既有 IF_RevGW_*、IF_RevGWFleet、L1 與 _Prod 列數值逐格不變；"
                  "工作單 docs/workorders/20261006_v5.22.md）。以下為 " + v521.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v521.README_TITLE
README_TITLE = VERSION.split("_")[-1] + _T[len(v521.VERSION.split("_")[-1]):]
