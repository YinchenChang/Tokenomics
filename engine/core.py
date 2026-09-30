"""Tokenomics 公式引擎：以 pycel 直接載入 Excel 活頁簿並重算。

規則（CLAUDE.md 第 1、2 節）：
- 不手抄任何公式；所有數值都由活頁簿內的公式計算。
- 輸入與輸出優先以具名範圍存取；測試才使用儲存格位址。
- 遇到引擎不支援的函數時，回報給 Andy，不得在此另寫旁路計算。
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from pycel import ExcelCompiler

REPO_ROOT = Path(__file__).resolve().parent.parent
MODEL_DIR = REPO_ROOT / "model"
MODEL_NAME_RE = re.compile(r"^\d{8}_Tokenomics_v\d+(\.\d+)?\.xlsx$")
_REF_RE = re.compile(
    r"^(?:'(?P<q>(?:[^']|'')+)'|(?P<u>[^!']+))!"
    r"\$?(?P<c1>[A-Z]+)\$?(?P<r1>\d+)(?::\$?(?P<c2>[A-Z]+)\$?(?P<r2>\d+))?$"
)


def current_model_path(model_dir: Path = MODEL_DIR) -> Path:
    """model/ 內唯一一份現行活頁簿（舊版在 model/archive/）。"""
    files = sorted(p for p in model_dir.glob("*.xlsx") if MODEL_NAME_RE.match(p.name))
    if len(files) != 1:
        raise RuntimeError(
            f"model/ 必須恰有一份現行 xlsx（YYYYMMDD_Tokenomics_vN.xlsx），實際為 {[p.name for p in files]}"
        )
    return files[0]


def parse_ref(attr_text: str) -> tuple[str, str]:
    """'Interface!$C$19:$Q$19' → ('Interface', 'C19:Q19')；頁名可含中文或引號。"""
    m = _REF_RE.match(attr_text)
    if not m:
        raise ValueError(f"無法解析範圍：{attr_text!r}")
    sheet = (m["q"] or m["u"]).replace("''", "'")
    ref = f"{m['c1']}{m['r1']}"
    if m["c2"]:
        ref += f":{m['c2']}{m['r2']}"
    return sheet, ref


def read_defined_names_xml(path: Path) -> dict[str, str]:
    """直接解析 workbook.xml 的 definedName，作為獨立於 openpyxl 的對照來源。"""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("xl/workbook.xml"))
    out = {}
    for el in root.iterfind(".//m:definedNames/m:definedName", ns):
        if el.get("localSheetId") is None:  # 只取活頁簿層級名稱
            out[el.get("name")] = (el.text or "").strip()
    return out


def norm(v):
    """統一空值：None 與空字串視為同一種空值；錯誤值 pycel 以 '#XXX!' 字串回傳。"""
    return "" if v is None else v


class Engine:
    """封裝 pycel：載入、設定輸入、依具名範圍取值。"""

    def __init__(self, path: str | Path | None = None):
        self.path = Path(path) if path else current_model_path()
        self._xl = ExcelCompiler(filename=str(self.path), plugins=["engine.excel_semantics"])
        wb = openpyxl.load_workbook(self.path)  # 公式模式：取得公式格清單與具名範圍
        self.sheetnames = list(wb.sheetnames)
        self.names = {k: v.attr_text for k, v in wb.defined_names.items()}
        self._formula_cells = [
            (ws.title, c.coordinate)
            for ws in wb
            for row in ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        self._max_row = {ws.title: ws.max_row for ws in wb}
        self._max_col = {ws.title: ws.max_column for ws in wb}
        self._warm_up()

    def _warm_up(self) -> None:
        """建立全簿計算圖，並強制全部公式格重算。

        pycel 冷啟動時，公式格會直接回傳檔內的快取值而不計算（實測：recalculate() 後
        3,729 格的浮點位元改變）。若不強制重算，parity 的基準情境只是在比對快取值。
        本引擎因此一律在建構時重算，之後所有取值都是引擎自行計算的結果。
        """
        addrs = [self._addr(s, c) for s, c in self._formula_cells]
        self._xl.evaluate(addrs)      # 建圖；同時滿足 set_value 要求該格已在 cell map
        self._xl.recalculate()        # 清除快取值並重算

    # ── 基本存取 ────────────────────────────────────────────────
    @staticmethod
    def _addr(sheet: str, ref: str) -> str:
        return f"'{sheet}'!{ref}"

    def get(self, sheet: str, ref: str):
        """單格回傳純量；多格回傳巢狀 list（列 × 欄）。"""
        v = self._xl.evaluate(self._addr(sheet, ref))
        if isinstance(v, tuple):
            if v and not isinstance(v[0], tuple):  # pycel 對單列範圍回傳一維 tuple
                v = (v,)
            return [[norm(x) for x in row] for row in v]
        return norm(v)

    def set_input(self, sheet: str, coord: str, value) -> None:
        """改寫輸入格；相依格於下次取值時重算。"""
        addr = self._addr(sheet, coord)
        # pycel 限制：依賴格尚未建圖就設值，後建圖時會讀回檔內原值（F35 實測重現）；
        # 建構時已建全簿計算圖，故此處可直接設值。
        self._xl.set_value(addr, value)

    # ── 具名範圍 ────────────────────────────────────────────────
    def name_ref(self, name: str) -> tuple[str, str]:
        return parse_ref(self.names[name])

    def get_name(self, name: str):
        """單格→純量；單列或單欄→一維 list；其他→巢狀 list。"""
        sheet, ref = self.name_ref(name)
        v = self.get(sheet, ref)
        if isinstance(v, list) and (len(v) == 1 or all(len(r) == 1 for r in v)):
            return [x for row in v for x in row]
        return v

    def set_name(self, name: str, value) -> None:
        """只允許單格具名範圍（例：CTL_GW）。"""
        sheet, ref = self.name_ref(name)
        if ":" in ref:
            raise ValueError(f"{name} 不是單格具名範圍")
        self.set_input(sheet, ref, value)

    # ── 全簿 ────────────────────────────────────────────────────
    @property
    def formula_cells(self) -> list[tuple[str, str]]:
        return list(self._formula_cells)

    def evaluate_all(self) -> dict[tuple[str, str], object]:
        """重算並回傳全部公式格的值。"""
        addrs = [self._addr(s, c) for s, c in self._formula_cells]
        vals = self._xl.evaluate(addrs)
        return {k: norm(v) for k, v in zip(self._formula_cells, vals)}

    # ── 版面搜尋（以標籤定位，不寫死位址）─────────────────────
    def column_labels(self, sheet: str, col: str = "A") -> list[tuple[int, str]]:
        """回傳指定欄的 (列號, 文字) 清單。"""
        out = []
        for r in range(1, self._max_row[sheet] + 1):
            v = self.get(sheet, f"{col}{r}")
            if isinstance(v, str) and v:
                out.append((r, v))
        return out

    def row_values(self, sheet: str, row: int, first_col: str = "C") -> list:
        """自 first_col 起，讀到工作表最後一欄。"""
        c1 = column_index_from_string(first_col)
        c2 = self._max_col[sheet]
        if c2 < c1:
            return []
        rng = f"{first_col}{row}:{get_column_letter(c2)}{row}"
        v = self.get(sheet, rng)
        return v[0] if isinstance(v, list) else [v]


__all__ = ["Engine", "current_model_path", "parse_ref", "read_defined_names_xml", "norm"]
