"""第 2 段機械檢查：m 路徑 561 頁＋列。
(a) 同列跨欄公式結構（R1C1 化）是否一致（L3 候選）；(b) IF 文字回傳與下游未守衛使用（L7 候選）；
(c) 每列代表公式、單位、值。用法：python3 10_mech.py <wb.pkl> <graph.pkl> <稽核0明細.xlsx> <out.pkl>"""
import pickle,sys,re,collections,openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L
d=pickle.load(open(sys.argv[1],'rb'));G=pickle.load(open(sys.argv[2],'rb'))
cells=d['cells'];edges=G['edges']
wb=openpyxl.load_workbook(sys.argv[3],read_only=True)
rows=[r for r in wb['C7_m路徑'].iter_rows(min_row=3,values_only=True) if r[0]=='路徑列']
path=[(r[1],int(r[2]),r[3]) for r in rows]
byrow=collections.defaultdict(list)
for (s,a),x in cells.items(): byrow[(s,x['r'])].append((a,x))
def r1c1(f,r0,c0):
    out=[]
    try: toks=Tokenizer(f).items
    except Exception: return f
    for t in toks:
        v=t.value
        if t.type=='OPERAND' and t.subtype=='RANGE':
            def rep(m):
                sh=m.group(1) or ''; cd,cl,rd,rw=m.group(2),m.group(3),m.group(4),m.group(5)
                c=CI(cl); r=int(rw)
                cs=f'C{c}' if cd else f'C[{c-c0}]'; rs=f'R{r}' if rd else f'R[{r-r0}]'
                return sh+rs+cs
            if re.fullmatch(r"((?:'[^']+'|[A-Za-z_0-9]+)!)?\$?[A-Z]{1,3}\$?[0-9]+(:\$?[A-Z]{1,3}\$?[0-9]+)?",v):
                v=re.sub(r"((?:'[^']+'|[A-Za-z_0-9]+)!)?(\$?)([A-Z]{1,3})(\$?)([0-9]+)",rep,v)
        out.append(v)
    return ''.join(out)
TXT=re.compile(r'"([^"]{2,})"')
res=[]
for s,r,lab in path:
    cs=sorted(byrow[(s,r)],key=lambda t:(len(re.match('[A-Z]+',t[0]).group()),t[0]))
    fcs=[(a,x) for a,x in cs if isinstance(x['f'],str) and x['f'].startswith('=') and re.match('[A-Z]+',a).group() not in('A','B')]
    unit=None
    for a,x in cs:
        if re.match('[A-Z]+',a).group()=='B' and isinstance(x['v'],str) and len(x['v'])<20: unit=x['v']
    pats=collections.OrderedDict()
    for a,x in fcs:
        p=r1c1(x['f'],x['r'],x['c'])
        pats.setdefault(p,[]).append(a)
    txt=set()
    for a,x in fcs:
        for m in TXT.finditer(x['f']):
            if m.group(1) not in('0.00','#,##0','0.0','0'): txt.add(m.group(1))
    textvals=[a for a,x in fcs if isinstance(x['v'],str)]
    res.append(dict(sheet=s,row=r,label=lab,unit=unit,nform=len(fcs),ninput=len([1 for a,x in cs if not(isinstance(x['f'],str) and x['f'].startswith('='))])-2,
        patterns=[(p[:300],v) for p,v in pats.items()],texts=sorted(txt),textvals=textvals,
        rep=(fcs[0][0],fcs[0][1]['f'][:400],fcs[0][1]['v']) if fcs else None))
# L7：下游消費文字回傳格
rev=collections.defaultdict(set)
for n,ps in edges.items():
    for p in ps: rev[p].add(n)
can_text=set()
for (s,a),x in cells.items():
    f=x['f']
    if isinstance(f,str) and f.startswith('=') and TXT.search(f) and re.search(r'IF\(',f): can_text.add((s,a))
agg=re.compile(r'\b(SUM|SUMPRODUCT|MIN|MAX|AVERAGE|PRODUCT)\(')
l7=[]
pathset={(s,r) for s,r,_ in path}
for (s,a) in can_text:
    x=cells[(s,a)]
    if (s,x['r']) not in pathset: continue
    users=set();stack=[(s,a)];seen=set()
    while stack:
        cur=stack.pop()
        for u in rev.get(cur,()):
            if u in seen: continue
            seen.add(u)
            if u[0] in('NAME','RANGE'): stack.append(u)
            else: users.add(u)
    for u in users:
        if u[0] in('Gov_Map','Checks','Load_Bearing','SRC_Index','Decisions'): continue
        uf=cells[u]['f']
        guarded=bool(re.search(r'ISNUMBER|ISTEXT|ISERROR|IFERROR',uf))
        aggr=bool(agg.search(uf))
        if not guarded or aggr:
            l7.append(dict(src=f'{s}!{a}',srcv=x['v'],user=f'{u[0]}!{u[1]}',userf=uf[:250],userv=cells[u]['v'],guarded=guarded,aggr=aggr))
pickle.dump(dict(rows=res,l7=l7),open(sys.argv[4],'wb'))
multi=[x for x in res if len(x['patterns'])>1]
print('path rows',len(res),'multi-pattern rows',len(multi),'L7 candidates',len(l7))
for x in multi: print(x['sheet'],x['row'],len(x['patterns']),[v for p,v in x['patterns']][:6])
