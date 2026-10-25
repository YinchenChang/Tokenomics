# 產生「NonNV_AMD.xlsx」審閱檔（唯讀查證，不碰 model/）
# 用法：python3 build_xlsx.py <輸出.xlsx>，再以 LibreOffice headless 轉存寫入計算值
# 方法沿用 B 段（docs/evidence/20261025_NonNV_TPU_Trainium/build_xlsx.py）與 A 段：
#   產出比＝（候選每 kW 峰值算力 ÷ VR200）× 效率因子 f；持有比＝1 − s ×（1 − k）
# 本段新增：f 由 MLPerf v6.1（MI355X）對 TK 的 GB200／GB300／VR200 三個錨點還原
# 推導頁藍字＝輸入，黑字＝公式；欄 D/E/F＝低端／基準／高端
import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, __file__.rsplit("/", 1)[0] if "/" in __file__ else ".")
from evidence_data import EVIDENCE, QUESTIONS  # noqa: E402

out = sys.argv[1]
wb = Workbook()
BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7")
SEC = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")

ws = wb.active
ws.title = "推導"
ws["A1"] = "推導 — AMD Instinct MI455X（Helios，72 顆）相對 VR200 NVL72 的每 GW 產出比與持有成本比（C 段；方法同 A、B 段）"
ws["A1"].font = BOLD
ws["A2"] = ("欄 D／E／F＝低端／基準／高端（產出列：悲觀→樂觀；持有列：成本低→成本高）。藍字＝輸入；黑字＝公式。"
            "標記沿用 Verified／Interested-party／Analogy／Assumed／Derived。VR200 分母取自 model/20261025_Tokenomics_v21.6.xlsx（基準成本欄 Interface!M）。"
            "證據編號 M01–M14 見『證據』頁。")
ws["A2"].alignment = WRAP
heads = ["編號", "物理量", "單位", "低端", "基準", "高端", "來源（證據編號／模型格）", "標記", "算式或說明"]
for i, h in enumerate(heads, 1):
    c = ws.cell(4, i, h)
    c.font = BOLD
    c.fill = HDR

R = {}
row = [5]


def sec(title):
    r = row[0]
    ws.cell(r, 1, title).font = BOLD
    for c in range(1, 10):
        ws.cell(r, c).fill = SEC
    row[0] += 1


def ref(key, col):
    return f"{col}{R[key]}"


def add(key, name, unit, vals, src, tag, note="", fmt=None):
    r = row[0]
    R[key] = r
    ws.cell(r, 1, key)
    ws.cell(r, 2, name)
    ws.cell(r, 3, unit)
    for j, v in enumerate(vals):
        col = "DEF"[j]
        cell = ws[f"{col}{r}"]
        if isinstance(v, str) and v.startswith("="):
            def rep(m):
                return ref(m.group(1), m.group(2) or col)
            cell.value = re.sub(r"\{([A-Za-z0-9_]+)(?:@([DEF]))?\}", rep, v)
        else:
            cell.value = v
            cell.font = BLUE
        if fmt:
            cell.number_format = fmt
    ws.cell(r, 7, src)
    ws.cell(r, 8, tag)
    ws.cell(r, 9, note)
    row[0] += 1


def same(v):
    return [v, v, v]


# ---- A. VR200 分母（與 B 段相同，值改讀 v21.6；數值與 v21.5 相同）
sec("A. VR200 NVL72 分母（模型 v21.6；同層級、同 SLO）")
add("A1", "VR200 每 GPU 計算用峰值（NVFP4 dense；模型服務精度）", "PF", same(35), "Spec_Rack!F35（＝F23；F22 服務精度 NVFP4）", "Interested-party（NVIDIA 規格）", "AMD 自己的 1.15X 比較也用 35（M02）")
add("A2", "VR200 機架配電設計功率（基準）", "kW", same(227), "Spec_Rack!F9", "Interested-party（Schneider 轉述 NVIDIA）", "Max-P 設計功率")
add("A3", "每機架 GPU 封裝數", "顆", same(72), "Spec_Rack!F6", "Verified", "")
add("A4", "VR200 每顆 GPU 所攤機架功率", "kW", ["={A2}/{A3}"] * 3, "Derived", "Derived", "227÷72")
add("A5", "VR200 每 kW 峰值算力（NVFP4）", "PF/kW", ["={A1}/{A4}"] * 3, "Derived", "Derived", "")
add("A6", "VR200 每 GW 機架數", "架", same(4052), "Interface!M6", "Derived（模型）", "")
add("A7", "VR200 每 GW 機架 IT 功率（＝A6×A2）", "MW", ["={A6}*{A2}/1000"] * 3, "Derived", "Derived", "候選每 GW 架數＝此值 ÷ 候選機架功率（與 A、B 段相同）")
add("A8", "VR200 機架價格（基準）", "$M", same(8.4), "Spec_Rack!F12", "Interested-party", "MS BOM 7.8 ×（1＋代工毛利 7.5%）")
add("A9", "VR200 每 GW IT 設備資本支出（基準）", "$B", same(37.586352), "Interface!M9", "Derived（模型）", "")
add("A10", "IT 資本支出 ÷（機架數×機架價）＝網路等加成", "x", ["={A9}/({A6}*{A8}/1000)"] * 3, "Derived", "Derived", "約 1.10；候選沿用同一加成")
add("A11", "VR200 每 GW 廠房資本支出（基準）", "$B", same(12.67), "Interface!M10", "Derived（模型）", "")
add("A12", "VR200 每 GW 年持有成本 — 經濟口徑（基準）", "$B/年", same(12.2254710045841), "Interface!M13", "Derived（模型）", "分母")
add("A13", "VR200 每 GW 年持有成本 — 會計口徑（基準）", "$B/年", same(8.9923042149807), "Interface!M12", "Derived（模型）", "")
add("A14", "VR200 每 GW IT 折舊（基準）", "$B/年", same(6.264392), "Interface!M220", "Derived（模型）", "")
add("A15", "VR200 每 GW IT 維護（基準）", "$B/年", same(0.591045788314028), "Interface!M224", "Derived（模型）", "")
add("A16", "VR200 每 GW 財產稅與保險（基準）", "$B/年", same(0.25128176), "Interface!M227", "Derived（模型）", "")
add("A17", "資金成本（經濟 − 會計）", "$B/年", ["={A12}-{A13}"] * 3, "Derived", "Derived", "")
add("A18", "IT 設備占總資本支出比例", "—", ["={A9}/({A9}+{A11})"] * 3, "Derived", "Derived", "")
add("A19", "持有成本中隨 IT 設備資本支出成比例的部分", "$B/年", ["={A14}+{A15}+({A16}+{A17})*{A18}"] * 3, "Derived", "Derived", "")
add("A20", "上述部分占 VR200 持有成本比例 s", "—", ["={A19}/{A12}"] * 3, "Derived", "Derived", "候選持有比 ＝ 1 − s×(1 − k)")
add("A21", "VR200 每 GPU FP8 dense", "PF", same(17.5), "Spec_Rack!F40", "Interested-party", "E 節『FP8 對齊』用")
add("A22", "VR200 GPU 功率（能量用）", "W", same(1800), "Spec_Rack!F33", "Interested-party", "")
add("A23", "VR200 機架功率 ÷ 晶片功率總和", "x", ["={A2}*1000/({A3}*{A22})"] * 3, "Derived", "Derived", "1.75；Helios 旁證見 C3m")
add("A24", "VR200 HBM 頻寬", "TB/s", same(22), "Spec_Rack!F26", "Interested-party", "")
add("A25", "VR200 ridge point（峰值 ÷ 頻寬）", "FLOP/byte", ["={A1}*1000/{A24}"] * 3, "Spec_Rack!F36", "Derived", "1,591")

# ---- B. 校準（同 A、B 段）＋ MLPerf 還原
sec("B. 還原既有世代：模型實際產出比 ÷ 『每 kW 峰值算力比』＝有效效率因子 f（同 A、B 段）")
for key, nm, pf, kw, g, src in [
    ("H", "Hopper H100", 1.979, 41, 32, "Spec_Rack!C35,C9,C6"),
    ("G2", "GB200 NVL72", 10, 132, 72, "Spec_Rack!D35,D9,D6"),
    ("G3", "GB300 NVL72", 15, 136, 72, "Spec_Rack!E35,E9,E6"),
]:
    add(f"B_{key}_pf", f"{nm} 每 GPU 計算用峰值", "PF", same(pf), src, "Interested-party", "Hopper 為 FP8（無 FP4），其餘 NVFP4")
    add(f"B_{key}_kw", f"{nm} 機架配電功率（基準）", "kW", same(kw), src, "Interested-party／Derived", "")
    add(f"B_{key}_n", f"{nm} 每機架 GPU", "顆", same(g), src, "Verified", "")
    add(f"B_{key}_kwg", f"{nm} 每顆所攤功率", "kW", [f"={{B_{key}_kw}}/{{B_{key}_n}}"] * 3, "Derived", "Derived", "")
    add(f"B_{key}_r", f"{nm} 每 kW 峰值算力 ÷ VR200", "—", [f"=({{B_{key}_pf}}/{{B_{key}_kwg}})/{{A5}}"] * 3, "Derived", "Derived", "")
vr = {"H": (0.107484660941882, 0.10544360110803, 0.0977954294066359),
      "G2": (0.534632077509924, 0.513240130160715, 0.483657996161255),
      "G3": (0.693216239743152, 0.665656345719338, 0.563845155935545)}
for key, nm in [("H", "Hopper"), ("G2", "GB200"), ("G3", "GB300")]:
    add(f"B_{key}_v", f"{nm} VR-eq（Luna／Sol／Astra）", "—", list(vr[key]), "Interface IF_VReq_Luna/Sol/Astra（列 22/32/42）", "Derived（模型）", "欄 D/E/F＝Luna／Sol／Astra（此行不是低/基準/高）")
    add(f"B_{key}_vm", f"{nm} VR-eq 三層平均", "—", [f"=AVERAGE(D{R[f'B_{key}_v']}:F{R[f'B_{key}_v']})"] * 3, "Derived", "Derived", "")
    add(f"B_{key}_f", f"{nm} 有效效率因子 f", "—", [f"={{B_{key}_vm}}/{{B_{key}_r}}"] * 3, "Derived", "Derived", "")
add("B_V_f", "VR200 有效效率因子 f（定義）", "—", same(1), "定義（VR200 產出比＝1、每 kW 峰值比＝1）", "Derived", "")

sec("B2. MLPerf Inference v6.1 每顆吞吐（tok/s/GPU；Closed）→ MI355X 每峰值效率 → 以 TK 錨點還原 f（MI455X 未提交，M13）")
add("M_pk", "MI355X 每 GPU 峰值（MXFP4 dense）", "PF", same(10.066), "AMD MI355X 規格（MXFP4 10,066 TF；與 M01 同一表格格式）", "Interested-party（AMD）", "MLPerf 中 MI355X 以 MXFP4 為主（M05）")
add("M_bw", "MI355X HBM 頻寬", "TB/s", same(8), "AMD MI355X 規格", "Interested-party（AMD）", "")
bench = [
    ("DS", "DeepSeek-R1 Server", 4698, 5928, 8517, 16427),
    ("GS", "gpt-oss-120b Server", 14155, 12774, 16118, None),
    ("GO", "gpt-oss-120b Offline", 15227, 15301, 16635, None),
    ("LS", "Llama2-70b-99.9 Server", 13068, 12648, 15998, None),
]
for k, nm, a, g2, g3, v in bench:
    add(f"M_{k}_a", f"{nm}：MI355X", "tok/s/GPU", same(a), "M05", "第三方實測（MLPerf）", "")
    add(f"M_{k}_g2", f"{nm}：GB200", "tok/s/GPU", same(g2), "M05", "第三方實測（MLPerf）", "")
    add(f"M_{k}_g3", f"{nm}：GB300", "tok/s/GPU", same(g3), "M05", "第三方實測（MLPerf）", "")
    if v:
        add(f"M_{k}_v", f"{nm}：VR200（Preview）", "tok/s/GPU", same(v), "M06", "第三方實測（MLPerf Preview）", "")
    add(f"M_{k}_f2", f"{nm}：f（GB200 錨點）＝ f_GB200 ×（MI355X÷峰值）÷（GB200÷峰值）", "—",
        [f"={ref('B_G2_f','E')}*({{M_{k}_a}}/{ref('M_pk','E')})/({{M_{k}_g2}}/{ref('B_G2_pf','E')})"] * 3, "Derived", "Derived（MLPerf）", "", fmt="0.000")
    add(f"M_{k}_f3", f"{nm}：f（GB300 錨點）", "—",
        [f"={ref('B_G3_f','E')}*({{M_{k}_a}}/{ref('M_pk','E')})/({{M_{k}_g3}}/{ref('B_G3_pf','E')})"] * 3, "Derived", "Derived（MLPerf）", "", fmt="0.000")
    if v:
        add(f"M_{k}_fv", f"{nm}：f（VR200 錨點）", "—",
            [f"={ref('B_V_f','E')}*({{M_{k}_a}}/{ref('M_pk','E')})/({{M_{k}_v}}/{ref('A1','E')})"] * 3, "Derived", "Derived（MLPerf Preview）", "VR200 為 Preview；TK 的 GB300÷VR200 每顆 0.38、MLPerf 0.51，故此錨點偏高（M06）", fmt="0.000")
add("M_DSavg", "DeepSeek-R1 三錨點平均 f（MI355X）", "—",
    [f"=AVERAGE({ref('M_DS_f2','E')},{ref('M_DS_f3','E')},{ref('M_DS_fv','E')})"] * 3, "Derived", "Derived（MLPerf）", "最像 TK 前沿負載（大型 MoE 推理）", fmt="0.000")
add("M_G23avg", "DeepSeek-R1 只用 GB200、GB300 兩錨點平均 f（MI355X）", "—",
    [f"=AVERAGE({ref('M_DS_f2','E')},{ref('M_DS_f3','E')})"] * 3, "Derived", "Derived（MLPerf）", "r1：排除已知偏高的 VR200 Preview 錨點（Q3 替代）", fmt="0.000")
add("M_ALLavg", "全部九個還原值平均 f（MI355X）", "—",
    [f"=AVERAGE({','.join(ref(f'M_{k}_{s}','E') for k,_,_,_,_,v in bench for s in (['f2','f3','fv'] if v else ['f2','f3']))})"] * 3, "Derived", "Derived（MLPerf）", "含較小模型（gpt-oss、Llama2-70b），只作敏感度", fmt="0.000")
add("M_rg", "MI355X ridge point", "FLOP/byte", ["={M_pk}*1000/{M_bw}"] * 3, "Derived", "Derived", "1,258")
add("M_rgG2", "GB200 ridge point", "FLOP/byte", same("=10*1000/8"), "Spec_Rack!D36", "Derived（模型）", "")
add("M_rgG3", "GB300 ridge point", "FLOP/byte", same("=15*1000/8"), "Spec_Rack!E36", "Derived（模型）", "")
add("M_e", "f 對 ridge 的彈性 e（TK GB200→GB300 類比）", "—",
    [f"=-LN({ref('B_G3_f','E')}/{ref('B_G2_f','E')})/LN({ref('M_rgG3','E')}/{ref('M_rgG2','E')})"] * 3, "Derived", "Analogy", "f ∝ ridge^(−e)；旁證：以此 e 由 GB300 推 VR200 得 0.95（實際定義 1.0）", fmt="0.000")

# ---- C. MI455X
sec("C. AMD Instinct MI455X（Helios：18 運算托盤 × 4 ＝ 72 顆；6 交換托盤）")
add("C1", "每顆 MXFP4 dense 峰值", "PF", same(40.265), "M01（AMD 資料表）", "Interested-party（AMD）", "有原生 FP4；與 VR200 NVFP4 同為最高原生精度")
add("C2", "每顆 FP8 dense 峰值", "PF", same(20.133), "M01", "Interested-party（AMD）", "FP4:FP8＝2:1")
add("C2b", "MI455X HBM 頻寬", "TB/s", same(23.3), "M01", "Interested-party（AMD）", "")
add("C3", "每顆最大 TBP", "kW", same(2.5), "M02（AMD『maximum TBP of 2500W』）", "Interested-party（AMD 本人說）", "旁證：SemiAnalysis 傳聞 2,300→2,500 W（M02 備註）")
add("C4", "Helios 機架功率", "kW", [245, 245, 225], "M03（AMD 數字經媒體：依負載 225–245 kW；245 亦為散熱上限）", "Interested-party（區間）", "基準 245＝與 VR200 227（Max-P 設計功率）同口徑的上限；高端產出 225＝輕負載（口徑偏樂觀）")
add("C4n", "每機架晶片數", "顆", same(72), "M01、M04", "Interested-party（AMD）", "")
add("C3m", "旁證：機架 ÷ 晶片 TBP 總和", "x", ["={C4}/({C4n}*{C3})"] * 3, "Derived", "Derived", "1.36（245）；TK GB300 1.35、VR200 1.75")
add("C5", "每顆所攤機架功率", "kW", ["={C4}/{C4n}"] * 3, "Derived", "Derived", "VR200 為 3.153")
add("C6", "(a) 最高原生精度：每 kW 峰值算力（MXFP4）", "PF/kW", ["={C1}/{C5}"] * 3, "Derived", "Derived", "", fmt="0.000")
add("C6r", "(a) ÷ VR200 每 kW 峰值算力（NVFP4）", "—", ["={C6}/{A5}"] * 3, "Derived", "Derived", "", fmt="0.0000")
add("C6r8", "(b) FP8 對齊：每 kW FP8 峰值 ÷ VR200 每 kW FP8 峰值", "—", ["=({C2}/{C5})/({A21}/{A4})"] * 3, "Derived", "Derived", "與 (a) 相同（兩者 FP4:FP8 皆 2:1）", fmt="0.0000")
add("C7rg", "MI455X ridge point", "FLOP/byte", ["={C1}*1000/{C2b}"] * 3, "Derived", "Derived", "1,728（VR200 1,591）")
add("C7t", "MI355X→MI455X 轉換係數 t ＝（MI455X ridge ÷ MI355X ridge）^(−e)", "—", [f"=({ref('C7rg','E')}/{ref('M_rg','E')})^(-{ref('M_e','E')})"] * 3, "Derived", "Analogy", "MI455X 算力增 4 倍、頻寬增 2.9 倍，每峰值效率預期下降", fmt="0.000")
add("C7c", "佐證：AMD 自述 Kimi K2 每顆 +10%／+15% 換算的 f ＝ 比值 × A1 ÷ C1", "—", [f"=1.10*{ref('A1','E')}/{ref('C1','E')}", f"=1.125*{ref('A1','E')}/{ref('C1','E')}", f"=1.15*{ref('A1','E')}/{ref('C1','E')}"], "M07", "Interested-party（AMD）", "不進 f，只作高端佐證", fmt="0.000")
add("C7", "有效效率因子 f（MI455X）", "—", [f"={ref('M_DS_f3','E')}*{ref('C7t','E')}", f"={ref('M_DSavg','E')}*{ref('C7t','E')}", f"={ref('M_DS_fv','E')}"],
    "低＝MLPerf DeepSeek-R1 GB300 錨點 × t；基準＝三錨點平均 × t；高＝VR200 錨點不打折（AMD 自述 C7c 0.96–1.00 佐證）", "Assumed（區間；依據 MLPerf，轉換為 Analogy）", "SemiAnalysis InferenceX（M08）只作佐證，不進本列", fmt="0.000")
add("C8", "(a) MI455X 每 GW 產出比（原始）＝ C6r × C7", "—", ["={C6r}*{C7}"] * 3, "Derived", "Derived（含 Assumed）", "", fmt="0.0000")

# ---- D. 持有比
sec("D. 持有成本比（每 kW 設備價推導，同 A、B 段）")
add("D7a", "Futurum 同源比值 Helios ÷ Rubin（5.0/4.0、5.25/3.75、5.5/3.5）", "x", ["=5/4", "=5.25/3.75", "=5.5/3.5"], "M09", "第三方估計", "同一來源相除，口徑差抵銷")
add("D9", "AMD 自述隱含 Helios ÷ VR200 每架價（每顆 +10%／+12.5%／+15% ÷ 每美元 +30%）", "x", ["=1.10/1.30", "=1.125/1.30", "=1.15/1.30"], "M07", "Interested-party（AMD）", "AMD 假設的 VR200 價未揭露；比值乘 TK $8.4M", fmt="0.000")
add("D1abs", "分析師絕對價（CLSA $5.0M／Futurum 中點 $5.25M／Futurum 上緣 $5.5M）", "$M", [5.0, 5.25, 5.5], "M10、M09", "第三方估計", "r1：與 TK $8.4M（買方取得價，MS BOM×1.075）口徑不同；Futurum 同源 Rubin 只 $3.5–4M（TK 口徑的 0.42–0.48）")
add("D2", "Helios 每架取得價（TK 買方取得價口徑）", "$M", ["={D1abs@E}", f"={ref('D9','E')}*{ref('A8','E')}", f"={ref('D7a','D')}*{ref('A8','E')}"], "低＝Futurum 絕對價 $5.25M（口徑混用、下界）；基準＝AMD 自述隱含比 0.865 × TK $8.4M；高＝Futurum 同源比值 1.25 × TK $8.4M", "Interested-party／第三方估計（區間）", "r1（Project 決定）：r0 為 $5.0／5.25／8.4M（絕對價 ÷ TK 價，口徑混用）")
add("D3", "每 GW 機架數（以基準功率）", "架", [f"={{A7}}*1000/$E${R['C4']}"] * 3, "Derived", "Derived", "VR200 為 4,052")
add("D4", "每 GW IT 設備資本支出", "$B", ["={D3}*{D2}/1000*{A10}"] * 3, "Derived", "Derived", "乘 VR200 同一網路加成")
add("D5", "k ＝ ÷ VR200 每 GW IT 資本支出", "—", ["={D4}/{A9}"] * 3, "Derived", "Derived", "", fmt="0.000")
add("D6", "MI455X 每 GW 持有成本比（原始）", "—", ["=1-{A20}*(1-{D5})"] * 3, "Derived", "Derived（含 Assumed）", "", fmt="0.0000")
add("D7", "替代：Futurum 同源比值全區間 × TK $8.4M → 持有比", "—", [f"=1-{ref('A20','E')}*(1-{ref('D3','E')}*{{D7a}}*{ref('A8','E')}/1000*{ref('A10','E')}/{ref('A9','E')})"] * 3, "Derived", "Derived", "不進建議值（Q4 替代）", fmt="0.0000")
add("D10", "替代：AMD 自述隱含比全區間 × $8.4M → 持有比", "—", [f"=1-{ref('A20','E')}*(1-{ref('D3','E')}*{{D9}}*{ref('A8','E')}/1000*{ref('A10','E')}/{ref('A9','E')})"] * 3, "Derived", "Derived", "不進建議值", fmt="0.0000")
add("D11", "對照：r0 口徑——分析師絕對價 ÷ TK 價 → 持有比", "—", [f"=1-{ref('A20','E')}*(1-{ref('D3','E')}*{{D1abs}}/1000*{ref('A10','E')}/{ref('A9','E')})"] * 3, "Derived", "Derived", "口徑混用，只作下界對照（Q4 替代）", fmt="0.0000")
add("D12", "旁證：Helios ÷ VR200 每架 HBM 容量", "x", same("=31/20.7"), "StorageReview（31 TB vs 20.7 TB）；Spec_Rack!F14、F16（VR200 BOM 以記憶體為最大項，自購記憶體 $6.7M）", "Derived", "反向證據：HBM 多 50%，BOM 推理不支持 Helios 明顯較便宜", fmt="0.00")

# ---- E. FP8 對齊
sec("E. 精度口徑 (b)：FP8 對齊（VR200 也以 FP8 計）")
add("E1", "(b) MI455X 產出比 ＝ C6r8 × C7", "—", ["={C6r8}*{C7}"] * 3, "Derived", "Derived", "f 沿用 C7（MLPerf 還原時 MI355X 已以 MXFP4 服務；兩晶片 FP4:FP8 同比例，比值不變）", fmt="0.0000")
add("E2", "(a) − (b) 差", "—", ["={C8}-{E1}"] * 3, "Derived", "Derived", "≈0", fmt="0.00000")
add("E3", "對照：VR200 改高情境 50 PF（Spec_Rack!F24）時 MI455X 產出比", "—", ["={C8}*35/50"] * 3, "Derived", "Derived", "只作對照（NVIDIA 推論壓縮口徑）", fmt="0.0000")

# ---- F. 敏感度
sec("F. 敏感度（一次動一個參數，其餘取基準；D＝參數低值對應結果，F＝高值）")
add("F1", "f 取 低／基準／高（功率取基準 245）", "產出比", [f"={ref('C6r','E')}*{ref('C7','D')}", f"={ref('C8','E')}", f"={ref('C6r','E')}*{ref('C7','F')}"], "Derived", "Derived", "", fmt="0.0000")
add("F2", "機架功率 245／235／225（f 取基準）", "產出比", [f"={ref('C1','E')}/(245/72)/{ref('A5','E')}*{ref('C7','E')}", f"={ref('C1','E')}/(235/72)/{ref('A5','E')}*{ref('C7','E')}", f"={ref('C1','E')}/(225/72)/{ref('A5','E')}*{ref('C7','E')}"], "M03", "Derived", "", fmt="0.0000")
add("F3", "f 不乘 t（MI455X 保有 MI355X 每峰值效率）：DS 三錨點平均", "產出比", [f"={ref('C6r','E')}*{ref('M_DS_f3','E')}", f"={ref('C6r','E')}*{ref('M_DSavg','E')}", f"={ref('C6r','E')}*{ref('M_DS_fv','E')}"], "Derived", "Derived", "Q3 替代 a", fmt="0.0000")
add("F4", "f＝全部九個還原值平均 × t", "產出比", [f"={ref('C6r','E')}*{ref('M_ALLavg','E')}*{ref('C7t','E')}"] * 3, "Derived", "Derived", "Q3 替代 b", fmt="0.0000")
add("F8", "f＝DS-R1 兩錨點（GB200、GB300）平均 × t", "產出比", [f"={ref('C6r','E')}*{ref('M_G23avg','E')}*{ref('C7t','E')}"] * 3, "Derived", "Derived", "r1 新增；Q3 替代", fmt="0.0000")
add("F5", "低端再扣首代機架級風險 × 0.8", "產出比", [f"={ref('C8','D')}*0.8"] * 3, "Derived", "Assumed", "Q3 替代 c", fmt="0.0000")
add("F6", "每架價格 $5.25／7.27／10.5M（功率取基準）", "持有比", [f"={ref('D6','D')}", f"={ref('D6','E')}", f"={ref('D6','F')}"], "Derived", "Derived", "", fmt="0.0000")
add("F7", "機架功率對持有比：245／235／225 kW（價格取基準）", "持有比", [f"=1-{ref('A20','E')}*(1-({ref('A7','E')}*1000/{p}*{ref('D2','E')}/1000*{ref('A10','E')})/{ref('A9','E')})" for p in (245, 235, 225)], "Derived", "Derived", "功率↓→每 GW 架數↑→資本支出↑", fmt="0.0000")

for col, w in zip("ABCDEFGHI", [10, 58, 10, 12, 12, 12, 50, 28, 60]):
    ws.column_dimensions[col].width = w
for r in range(5, row[0]):
    for c in (4, 5, 6):
        cell = ws.cell(r, c)
        if cell.number_format == "General" and ws.cell(r, 3).value in ("—", "PF/kW", "x", "kW"):
            cell.number_format = "0.0000"
    for c in (7, 8, 9):
        ws.cell(r, c).alignment = WRAP
ws.freeze_panes = "D5"

# ---------------------------------------------------------------- 建議值
sv = wb.create_sheet("建議值", 0)
sv["A1"] = "建議值 — NonNV 第 5 列（AMD MI455X／Helios）以每 kW 方法重推｜C 段唯讀查證，未寫入 model/；與 B 段 TPU／Trn3、A 段 Maia／昇騰並排"
sv["A1"].font = BOLD
sv["A2"] = ("第 4–5 列＝推導頁公式的原始結果（(a) 最高原生精度、(b) FP8 對齊）；第 8 列＝建議輸入值（r1；產出比四捨五入至 0.05；持有比保留兩位）；第 11 列＝TK v21.6 現值。"
            "H＝F÷C（每 token 成本比）、I＝E÷D（最佳角落）、J＝G÷B（最差角落），同 NonNV。")
sv["A2"].alignment = WRAP
cols = ["候選", "產出比 低", "產出比 基準", "產出比 高", "持有比 低", "持有比 基準", "持有比 高", "每 token 成本比 基準", "最佳角落", "最差角落", "標記", "說明"]
for i, h in enumerate(cols, 1):
    c = sv.cell(3, i, h)
    c.font = BOLD
    c.fill = HDR


def ratio_cols(r):
    sv[f"H{r}"] = f"=F{r}/C{r}"
    sv[f"I{r}"] = f"=E{r}/D{r}"
    sv[f"J{r}"] = f"=G{r}/B{r}"


sv["A4"] = "AMD MI455X（(a) 最高原生精度；原始公式結果）"
sv["A5"] = "AMD MI455X（(b) FP8 對齊；原始公式結果）"
for r, o in ((4, "C8"), (5, "E1")):
    for col, src in zip("BCD", "DEF"):
        sv[f"{col}{r}"] = f"=推導!{src}{R[o]}"
    for col, src in zip("EFG", "DEF"):
        sv[f"{col}{r}"] = f"=推導!{src}{R['D6']}"
    ratio_cols(r)
    sv[f"L{r}"] = "公式（連結推導頁）"
sv["A7"] = "建議輸入值（若寫入 NonNV 第 5 列）"
sv["A7"].font = BOLD
sv["A8"] = "AMD MI455X（Helios）"
for j, v in enumerate((0.70, 0.80, 1.15, 0.67, 0.85, 1.12)):
    sv.cell(8, 2 + j, v).font = BLUE
ratio_cols(8)
sv["K8"] = "Assumed"
sv["L8"] = ("每 kW 法：Helios 72 顆、機架 225–245 kW（AMD 數字經媒體；245 亦為散熱上限）→ 每顆 3.40 kW；MXFP4 40.27 PF（FP8 20.13），兩種精度口徑每 kW 峰值皆為 VR200 的 1.07；"
            "效率因子由 MLPerf v6.1 MI355X DeepSeek-R1 對 TK GB200／GB300／VR200 錨點還原（0.74–0.99），再乘算力÷頻寬轉換 0.89（Analogy）。"
            "持有比依 TK 買方取得價口徑：每架 $5.25M（Futurum 絕對價，口徑混用、下界）／$7.3M（AMD 自述隱含比 0.865 × VR200）／$10.5M（Futurum 同源比值 1.25 × VR200）推導（r1）。證據：docs/evidence/20261025_NonNV_AMD。")
sv["L8"].alignment = WRAP
sv["A10"] = "TK v21.6 現值（NonNV 第 5 列；慣例值、未推導）"
sv["A10"].font = BOLD
sv["A11"] = "AMD MI455X（Helios）（現值）"
for j, v in enumerate((0.5, 0.7, 1.0, 0.8, 0.9, 1.0)):
    sv.cell(11, 2 + j, v)
ratio_cols(11)
sv["K11"] = "Assumed"
sv["L11"] = "model/20261025_Tokenomics_v21.6.xlsx NonNV!B5:G5"

sv["A13"] = "與 A、B 段並排（每 kW 峰值比為各自最高原生精度；TPU／Trn3 取 B 段建議值、Maia／昇騰取 TK v21.6 NonNV 現值）"
sv["A13"].font = BOLD
for i, h in enumerate(["候選", "每 kW 峰值 ÷ VR200（基準）", "產出比 基準", "持有比 基準", "每 token 成本比", "產出比 區間", "持有比 區間", "來源"], 1):
    c = sv.cell(14, i, h)
    c.font = BOLD
    c.fill = HDR
side = [
    ("AMD MI455X（本提案）", f"=推導!E{R['C6r']}", "=C8", "=F8", 'TEXT(B8,"0.00")&"–"&TEXT(D8,"0.00")', 'TEXT(E8,"0.00")&"–"&TEXT(G8,"0.00")', "本表"),
    ("Google TPU v7（B 段）", 0.383, 0.35, 0.70, "0.25–0.40", "0.65–0.90", "B 段（PR #172）建議值；持有比已寫入 v21.6（買方取得價，無口徑混用）"),
    ("AWS Trainium3（B 段）", 0.176, 0.14, 0.55, "0.06–0.25", "0.45–0.70", "B 段（PR #172）建議值；持有比已寫入 v21.6。r1 註：KGI 絕對價 ÷ TK 價為口徑混用，同源比值時約 0.59–0.63（README §10）"),
    ("Microsoft Maia 200（A 段）", 0.676, 0.55, 0.50, "0.25–0.75", "0.40–0.60", "A 段推導 D6；NonNV 第 11 列"),
    ("華為昇騰 910C（A 段）", 0.093, 0.08, 0.55, "0.05–0.15", "0.45–0.70", "A 段推導 C6；NonNV 第 10 列"),
]
for k, (nm, pk, ob, hb, ol, hl, src) in enumerate(side):
    r = 15 + k
    sv.cell(r, 1, nm)
    sv.cell(r, 2, pk)
    sv.cell(r, 3, ob)
    sv.cell(r, 4, hb)
    sv.cell(r, 5, f"=D{r}/C{r}")
    sv.cell(r, 6, ("=" + ol) if ol.startswith("TEXT") else ol)
    sv.cell(r, 7, ("=" + hl) if hl.startswith("TEXT") else hl)
    sv.cell(r, 8, src)
    if not isinstance(pk, str):
        for cc in (2, 3, 4):
            sv.cell(r, cc).font = BLUE
for col, w in zip("ABCDEFGHIJKL", [36, 12, 11, 11, 11, 11, 11, 14, 10, 10, 10, 80]):
    sv.column_dimensions[col].width = w

# ---------------------------------------------------------------- 證據
ev = wb.create_sheet("證據")
eh = ["編號", "主張", "數值", "單位", "口徑", "出處網址", "來源等級", "立場", "標記", "取得方式與備註", "用於"]
for i, h in enumerate(eh, 1):
    c = ev.cell(1, i, h)
    c.font = BOLD
    c.fill = HDR
for rr, rowv in enumerate(EVIDENCE, 2):
    for cc, v in enumerate(rowv, 1):
        ev.cell(rr, cc, v).alignment = WRAP
for col, w in zip("ABCDEFGHIJK", [6, 52, 22, 12, 30, 56, 18, 10, 24, 52, 12]):
    ev.column_dimensions[col].width = w
ev.freeze_panes = "C2"

# ---------------------------------------------------------------- 待判斷
qd = wb.create_sheet("待判斷")
for i, h in enumerate(["編號", "題目", "預設（本提案採用）", "替代方案", "影響"], 1):
    c = qd.cell(1, i, h)
    c.font = BOLD
    c.fill = HDR
for rr, rowv in enumerate(QUESTIONS, 2):
    for cc, v in enumerate(rowv, 1):
        qd.cell(rr, cc, v).alignment = WRAP
for col, w in zip("ABCDE", [6, 50, 50, 50, 40]):
    qd.column_dimensions[col].width = w

wb.save(out)
print("saved", out, "rows", row[0])
