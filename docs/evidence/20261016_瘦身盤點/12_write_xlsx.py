"""組裝明細 xlsx 並輸出統計（供報告引用）。
用法：python3 12_write_xlsx.py <工作目錄(pickles)> <out.xlsx> <stats.json>
需要：wb.pkl a.pkl rec1.pkl rows.pkl c4x.pkl hc.pkl gov.pkl lab.pkl；12_curated.py（同目錄）"""
import pickle, sys, json, collections, re, os, datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from importlib import import_module
cur = import_module('12_curated')
W = sys.argv[1].rstrip('/') + '/'
ld = lambda n: pickle.load(open(W + n, 'rb'))
wb_ = ld('wb.pkl'); cells = wb_['cells']; a = ld('a.pkl'); rec = ld('rec1.pkl'); rows = ld('rows.pkl'); c4x = ld('c4x.pkl'); hc = ld('hc.pkl'); gov = ld('gov.pkl'); lab = ld('lab.pkl')
NFORM = a['nform']
STD = ['頁', '列或格', '標籤', '目前值', '類別', '建議處置', '理由', '性質＝工程類／判斷類']
out = Workbook(); out.remove(out.active)
HDR = PatternFill('solid', fgColor='FFD9E1F2')
def sheet(name, rule, headers, rws, widths=None):
    ws = out.create_sheet(name)
    ws.cell(1, 1, rule).alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
    ws.row_dimensions[1].height = max(48, 15 * (len(rule) // 150 + 2))
    for j, h in enumerate(headers, 1):
        c = ws.cell(2, j, h); c.font = Font(bold=True); c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, r in enumerate(rws, 3):
        for j, v in enumerate(r, 1):
            if isinstance(v, (datetime.date, datetime.datetime)): v = v.isoformat()
            if isinstance(v, str) and v.startswith(('=', '+', '-', '@')): v = "'" + v
            if isinstance(v, str) and len(v) > 32000: v = v[:32000]
            ws.cell(i, j, v)
    ws.freeze_panes = 'A3'
    if rws: ws.auto_filter.ref = f'A2:{L(len(headers))}{len(rws) + 2}'
    for j, w in enumerate(widths or [], 1): ws.column_dimensions[L(j)].width = w
    return ws
def sv(v):
    if isinstance(v, float): return round(v, 6)
    return v if not isinstance(v, str) else v[:80]

# ---------- 1 死格 ----------
d1 = []
for z in rec:
    d1.append([z['sheet'], z['cell'], z['label'][:80], sv(z['value']), z['cat'], z['act'], z['why'], z['nat'],
               z['nform'], z['nin'], z['status'], z['rule'], z['users']])
sheet('1_死格_逐列', '規則：第 0 段方法（依賴圖；根＝下游契約名稱＋全部 IF_／IFW_／IFC_／L1_／GOV_／CHK_／CAP_ 具名範圍＋Checks 全部格；不穿越 Gov_Map、Decisions、DB_Evidence、Load_Bearing、SRC_Index）。每列＝一個頁＋列，列出該列所有死格欄；'
      '「建議處置」由規則 R01–R22（欄「規則」）判定，只建議不執行；「性質」：不改任何輸出且不動輸入／治理登錄者＝工程類，其餘＝判斷類（6b 原則 5）。「引用者」＝直接引用此列的格種類（無任何格引用＝葉列；死格＝只被其他死格引用的中間量）。',
      STD + ['死格公式格數', '死格輸入或常數格數', '與 v5.36 比對', '規則', '引用者'], d1,
      [13, 12, 36, 12, 24, 40, 70, 10, 9, 9, 24, 6, 18])
# ---------- 1b 與 v5.36 ----------
d1b = []
for z in rec:
    if z['status'].startswith(('新增', '部分')):
        d1b.append([z['sheet'], z['cell'], z['label'][:80], sv(z['value']), '新增死格（v5.36 不是死格）', z['act'], z['why'], z['nat'], z['status']])
gone = collections.OrderedDict()
for g in rows['gone']:
    gone.setdefault((g['sheet'], g['row'], g['label'], g['status']), []).append(g['cell'])
for (s, r, labx, st), cs in gone.items():
    d1b.append([s, f'{cs[0]}…（共 {len(cs)} 格）' if len(cs) > 1 else cs[0], str(labx)[:80], '', '已不再是死格或已不存在', '—（無需處置）', st, '—', ''])
sheet('1b_死格與v5.36比對', '比對鍵：頁＋列標籤＋欄（v5.36→v10.0 之間列位置可能移動）；v5.36 清單取自 docs/reports/20261015_稽核0_明細.xlsx 的 C3_死格。'
      '上半：v10.0 新增（或部分新增）的死格列；下半：v5.36 的死格在 v10.0 已接上下游（現為核心格）或該格已空白（已刪除或搬移）。',
      STD + ['說明'], d1b, [13, 18, 36, 12, 26, 30, 70, 10, 30])
# ---------- 2 重複定義 ----------
d2 = [[x[1], x[2], x[3], '重複定義 ' + x[0], x[5], x[4] + '｜依據：' + x[6], x[7]] for x in cur.DUPS]
d2 = [[r[0].split('；')[0].split('（')[0][:20], r[0], r[1], r[2], r[3], r[4], r[5], r[6]] for r in d2]
sheet('2_重複定義', '範圍：同一物理量或同一口徑在兩處以上各自計算（不含 _Prod 鏡像與 Sens_* 複本）。「建議處置」：已決定保留（引用決定）／建議合併／待判斷。機械掃描（07_dups_scan.py）：同目標多名稱 0；同標籤多列 255 組，人工歸類後列出下列 16 項；'
      '其中 Calib／Perf／Perf_Batch／Training 的同標籤多列是「世代或情境欄的同一條推導鏈」，已歸入 D12、D13。',
      ['頁', '位置', '標籤', '目前值', '類別', '建議處置', '理由（含引用決定）', '性質＝工程類／判斷類'], d2, [10, 44, 36, 28, 14, 30, 90, 10])
# ---------- 3 鏡像與複本 ----------
d3 = []
for z in a['c4']:
    if z['page'] == 'Theory_Rev_Prod':
        act, why, nat = ('待判斷：在頁首明訂「第 90–96 列為自我檢查列，不鏡像核心頁 _Life 列」，並列入永久測試白名單', '稽核 0 第 5 節第 7 項遺留：鏡像頁第 90–96 列放自我檢查，核心頁同列是 _Life 輸出列（過渡別名＋D12 列）；_Life 列沒有生產折減版本。4-4 刪過渡別名後，核心頁該區會改變，白名單須同步。', '判斷類')
    else:
        act, why, nat = ('補黃底（情境改動格）；永久測試期望：本格公式與同列首欄不同、且為黃底', '本欄是情境欄，改讀低／高輸入或另一輸入格；頁首規則「黃底＝情境改動」，但此格未上黃底（稽核 0 第 5 節第 6 項）。', '工程類')
    d3.append([z['page'], z['cell'], z['label'][:60], '', '未明訂差異' if z['documented'] != '黃底情境格（預期差異）' else '預期差異', act, why + '｜差異：' + z['diff'][:80], nat])
for r in sorted({z['row'] for z in c4x['pb']}):
    z = next(y for y in c4x['pb'] if y['row'] == r)
    d3.append(['Perf_Batch↔Perf', f'第 {r} 列（C:Q 共 {sum(1 for y in c4x["pb"] if y["row"] == r)} 格）', z['label'], '', '明訂差異（頁首聲明）', '列入永久測試白名單', '頁首明訂「與 Perf 同一套公式，只換速度下限與參考任務」；本列差異：Batch＝' + z['pb'][:70] + '｜Perf＝' + z['pf'][:70], '工程類'])
for z in c4x['scen']:
    if z['sheet'] in ('Sens_Perf', 'Sens_Train'): continue
    d3.append([z['sheet'], z['cell'], z['label'], sv(z['value']), '情境輸入未上黃底', '補黃底（情境輸入）', '藍字數值但無黃底；頁首規則「黃底＝情境改動」。', '工程類'])
sp = collections.Counter(z['sheet'] for z in c4x['scen'] if z['sheet'] in ('Sens_Perf', 'Sens_Train'))
for s_, n in sp.items():
    d3.append([s_, f'第 6、7 列（共 {n} 格）', '世代索引、層級索引（藍字）', '', '非情境格（索引列）', '不需黃底；列入永久測試白名單', '藍字索引列，用來把欄對應到世代與層級，不是情境改動。', '工程類'])
for p_, st in a['c4_stats'].items():
    same = sum(v for k, v in st.items() if '結構相同' in k or '直接讀原頁同格' in k or '與同列首欄' in k or k.startswith('常數相同'))
    nm = [(k, v) for k, v in st.items() if '整列未鏡像' in k]
    d3.append([p_, '（統計）', '', same, '結構一致格數', '永久測試：未明訂差異＝0', '；'.join(f'{k}：{v}' for k, v in st.items() if '整列未鏡像' not in k and '直接讀' not in k and '結構相同' not in k)[:300] + (f'｜整列未鏡像（資訊，頁首明訂只鏡像所需列）：{nm[0][1]}' if nm else ''), '工程類'])
sheet('3_鏡像與複本', '重跑稽核 0 的 C4（8 個 _Prod 頁、Sens_Perf、Sens_Train、Sensitivity、Sens_Har），另補 Perf↔Perf_Batch 與「藍字情境格是否上黃底」。本頁即 4-1 永久測試的期望值：'
      '未明訂差異＝0；明訂差異＝白名單（Perf／Perf_Batch 第 50 列 CTL_ProdDerate、Perf_Batch↔Perf 第 41、43、44、53、76、93、94 列、Theory_Rev_Prod 第 90–96 列）；情境改動格一律黃底。'
      '「（統計）」列＝各頁結構一致的格數（供測試核對基準數）。',
      STD, d3, [16, 22, 36, 12, 22, 44, 90, 10])
# ---------- 4 寫死參數 ----------
d4 = [[z['sheet'], z['cell'], z['label'][:60], sv(z['value']), z['grp'] + '（常數 ' + z['const'] + '）', z['act'], z['why'] + '｜公式片段：' + z['frag'][:90], z['nat']] for z in hc]
sheet('4_寫死參數', '重跑稽核 0 的 C5，只列「疑似參數或門檻」217 格（v5.36 為 188 格；+29 來自 v6.2–v10.0 新增的 Alloc、Cap_In、Fleet_1GW、L1 公式）。其餘類別（位置、索引、10 的次方、時間換算、哨兵值）不列。'
      '分組：H1 L1 ±20% 門檻 96 格；H6 訓練 3 倍 90 格；其餘 31 格。',
      STD, d4, [12, 10, 40, 12, 34, 50, 90, 10])
# ---------- 5 過時標籤 ----------
d5 = [[x[1], x[2], x[3], '', f'{x[0]}（{x[6]}）', x[7], f'說明宣稱：{x[4]}｜實際：{x[5]}', x[8]] for x in cur.STALE]
sheet('5_過時標籤', '方法：規則掃描＋人工核對。規則 S1：文字中的具名範圍或 SRC_ID 不存在（掃全簿模型頁文字格，92 處，絕大多數是頁名或泛稱，人工確認真正有問題者 3 處）；'
      '規則 S2：說明欄的「約／≈／現值／＝」數字與同列數值相差 >20%（88 個含此類字樣的文字格全數人工比對，14 處自動命中、其餘人工發現）；規則 S3：「第 N 列」列參照 232 處，人工抽查指向列，未發現指向錯列。'
      '本頁列出的 19 項都已用 wb 讀到的現值核對。嚴重度：高＝基準數字與現值差 >40% 或會誤導判讀；中＝20–40%；低＝命名或次要。',
      STD, d5, [14, 14, 30, 6, 12, 50, 100, 10])
# ---------- 6 治理頁 ----------
d6 = []
for z in gov['g1']:
    d6.append(['Gov_Map', f"{z['sheet']}!{z['cell']}（{z['gm']}）", z['label'][:50], '', z['item'][:28], '核對並修正登錄（標籤或範圍）' if not z['item'].startswith('C1.1h') else '合併重複登錄（保留一筆，另一筆刪除或改範圍）', z['detail'][:200], '工程類'])
for z in gov['g2']:
    if z['item'].startswith('C1.1 '):
        d6.append(['Gov_Map', f"{z['sheet']}!{z['cell']}", z['label'][:50], z['detail'][:20], '藍字數值輸入未登錄 Gov_Map', '補登 Gov_Map（類別＝Decision，決策 ID＝v6.2 起的服務世代組合情境）', 'Alloc I94:L95 是「服務世代組合：全 GB200／全 VR200」情境 one-hot 輸入（8 格藍字數值），v6.x 新增時未登錄；稽核 0 時此類為 0。', '判斷類'])
g3 = collections.Counter(z['sheet'] for z in gov['g3']); g3c = collections.Counter(re.sub(r'^類別 ', '', z['detail'].split('；')[0].split(' ')[1] if '類別' in z['detail'] else '') for z in gov['g3'])
for s_, n in g3.most_common():
    d6.append(['Gov_Map', f'{s_}（{n} 列）', 'SRC_ID 與決策 ID 皆為「—」', n, '無來源也無決策', '依類別補決策 ID 或 SRC（屬 4-2 D15 一併處理）', '總數 311，與 v5.36 的 311 完全相同（逐頁數量相同，Inputs 49、Workload 41、NonNV 30、Arch 27、Spec_Rack 26）：第 2 批新增的 Gov_Map 列都有決策 ID，沒有再增加，也沒有減少。', '判斷類'])
for z in gov['g5']:
    d6.append([z['sheet'], z['cell'], z['label'][:50], z['detail'][:30], '輸入格無模型公式引用', '標「備用／區間端點（僅 Gov_Map）」或接上敏感度頁；Alloc C137:E137 為無人用的索引常數，建議刪除', z['detail'][:120], '判斷類' if z['sheet'] != 'Alloc' else '工程類'])
sheet('6_治理頁', 'Gov_Map 651→（v10.0）列的盤點：G1 指向公式格或標籤不符或同格重複登錄 11 項；G2 藍字數值未登錄 8 格（藍字文字未登錄 122 格，與 v5.36 的 122 相同）；G3 無來源也無決策 311 列（＝v5.36 的 311）；'
      'G4 決策 ID 在 Decisions 頁不存在 0、SRC_ID 在 SRC 頁不存在 0；G5 輸入格無模型公式引用 70 格（v5.36 為 43）。Gov_Map 指向空白格：0；指向已刪除頁：0。',
      STD, d6, [14, 26, 40, 14, 26, 50, 90, 10])

# ---------- 統計 ----------
tier = collections.Counter(); tier_rows = collections.Counter()
for z in rec:
    a_ = z['act']; nat = z['nat']
    if z['rule'] in ('R01', 'R06', 'R13', 'R18', 'R22c'): k = 'A 確定可刪（工程類）'
    elif z['rule'] in ('R22a', 'R22b', 'R16', 'R04', 'R11', 'R09b'): k = 'B 視判斷可刪（判斷類）'
    elif a_.startswith('保留'): k = 'C 保留（標顯示用／未接線／被複本讀取）'
    else: k = 'D 其他待判斷'
    tier[k] += z['nform']; tier_rows[k] += 1
    tier[k + '｜輸入或常數'] += z['nin']
st = dict(nform=NFORM, dead_form=sum(z['nform'] for z in rec), dead_in=sum(z['nin'] for z in rec), rows=len(rec), whole=sum(1 for z in rec if z['cat'].startswith('死格（整列')), partial=sum(1 for z in rec if z['cat'].startswith('死格（同列')),
          tier=dict(tier), tier_rows=dict(tier_rows), by_rule={r: sum(z['nform'] for z in rec if z['rule'] == r) for r in sorted({z['rule'] for z in rec})},
          new_rows=sum(1 for z in rec if z['status'].startswith('新增')), part_new=sum(1 for z in rec if z['status'].startswith('部分')), gone=len(rows['gone']),
          dups=collections.Counter(x[5].split('（')[0] for x in cur.DUPS), stale=len(cur.STALE), stale_hi=sum(1 for x in cur.STALE if x[6] == '高'), hc=len(hc), hc_judg=sum(1 for z in hc if z['nat'] == '判斷類'),
          c4=dict(unspec=sum(1 for z in a['c4'] if z['documented'] != '黃底情境格（預期差異）'), total=len(a['c4'])), nat=dict(collections.Counter(z['nat'] for z in rec)))
json.dump(st, open(sys.argv[3], 'w'), ensure_ascii=False, indent=1, default=str)
print(json.dumps(st, ensure_ascii=False, indent=1, default=str)[:3000])
# 總覽頁放最前
ov = [['死格（公式格）', st['dead_form'], f"占全簿公式格 {NFORM:,} 的 {st['dead_form']/NFORM:.1%}"], ['死格（輸入或常數格）', st['dead_in'], ''], ['死格列數', st['rows'], f"整列無人用 {st['whole']}；同列有核心格 {st['partial']}"],
      ['A 確定可刪公式格（工程類）', tier['A 確定可刪（工程類）'], f"占全簿公式格 {tier['A 確定可刪（工程類）']/NFORM:.2%}"],
      ['B 視判斷可刪公式格（判斷類）', tier['B 視判斷可刪（判斷類）'], f"占全簿公式格 {tier['B 視判斷可刪（判斷類）']/NFORM:.1%}"],
      ['C 保留（標顯示用等）公式格', tier['C 保留（標顯示用／未接線／被複本讀取）'], f"占全簿公式格 {tier['C 保留（標顯示用／未接線／被複本讀取）']/NFORM:.1%}"],
      ['重複定義項', len(cur.DUPS), '已決定保留 10、建議合併 3、待判斷 3（見 2_重複定義）'], ['過時或錯誤標籤', len(cur.STALE), f"其中高嚴重度 {st['stale_hi']}"],
      ['寫死參數（疑似）', st['hc'], f"建議外露 {st['hc_judg']}（判斷類）；其餘 {st['hc']-st['hc_judg']} 只加註"], ['鏡像／複本未明訂差異', st['c4']['unspec'], '＝Sensitivity 19＋Sens_Har 20（未上黃底）＋Theory_Rev_Prod 7（待判讀）；另有未上黃底的藍字情境格 16']]
ws = out.create_sheet('0_總覽', 0)
ws.append(['項目', '數量', '說明'])
for r in ov: ws.append(r)
for c in ws[1]: c.font = Font(bold=True); c.fill = HDR
ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 12; ws.column_dimensions['C'].width = 90
out.save(sys.argv[2])
