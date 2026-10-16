# v11.0 (work order docs/workorders/20261016_大批次修正.md; branch claude/bigbatch-fix; segment "大批次修正" G1–G4).
#   G1  Calib C68 (Hopper per-layer latency multiplier 1.5): traced (Decisions, Gov_Map, Calib J-column notes, builder history v5.2–v10.3, archived workbooks). Verdict "cannot be confirmed as double counting":
#       the 1.5 is an Assumed round number that exists since v5.2 (2026-09-30; per-layer latency then 240.4 µs = fit 160.28 × 1.5, not 241.92 µs); the DeepSeek 241.92 µs figure first appears in the repo on 2026-10-16.
#       C68 is NOT changed; notes record the trace; the hypothetical C68 = 1.0 effect is written to Decisions for Project.
#   G2  CAL_BatchEtaD (Calib row 156): GB200 D156 = 3.04 (Derived; solved on Calib J column D with the batch multiplier included); Hopper C156 stays 1 (no solution at C68 = 1.5: even an infinite η_d only reaches 0.941 of the
#       DeepSeek profile target, outside ±2%). Calib J columns C and D now carry CAL_BatchEtaD in their η_d (calibration copy rows of the batch path); E, F, G (serving-side validation points V15, V7, V8) are unchanged.
#   G3  Calib L section (V20–V22: MLPerf v6.0 offline GB200 6,752 and GB300 8,987, TensorRT-LLM H200 1,436; batch path = η_d × CAL_BatchEtaD, no SLO), SRC_Perf 079–083, SRC_HW 067–068 (H200 specs), Checks V20–V22.
#   G4  Checks V1–V3 (MLPerf server) are soft caps: ratio > 1 is "警示" (not a failure); new name CHK_VerifyWarn with CHK_VerifyPass + CHK_VerifyFail + CHK_VerifyWarn = CHK_VerifyTotal (edits live in v1000.checks_v). V4–V6 (interactive) stay hard caps.
# Major version (new contract name CHK_VerifyWarn). Every write to an Excel-owned cell is guarded (old value / old text).
import re
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as L
from common import put, F_IN, F_NOTE, F_BOLD, F_CALC, F_LINK, FILL_KEY, section
import v1000
import v1020
import v1030

VERSION = "20261016_Tokenomics_v11.0"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-16"
V = "v11.0"
WO = "工作單 docs/workorders/20261016_大批次修正.md"
DEC_BY = "Andy 2026-10-16 21:55「both ok」（發本段、V2 改軟上限）；其餘細節為 chat 端預設，可撤回"
EID = "E380"
DASH = "—"
REPORT = "docs/reports/20261016_大批次低估查證.md"
EVTAB = "docs/evidence/20261016_大批次/證據表.xlsx"
MLP_CSV = "docs/evidence/20261016_服務條件/mlperf_v6.0_deepseek-r1.csv"
BATCH_GB200 = 3.04          # G2: Calib D156 (Derived; J-column D solution 3.0373, rounded to 3 significant digits; ratio 1.0001)
# Gov_Map interval of D156 (set from the implied multipliers of the batch-path points; see Decisions BB-G2)
D156_LO, D156_HI = 1.5, 6.7


# ================================================================== SRC records (SRC_Perf 079–083, SRC_HW 067–068)
_SRC = f"{REPORT} 與 {EVTAB}（CC 查證，2026-10-16；原始出處見證據表）"


def _rec(sheet, i, metric, val, unit, basis, applies, date, src, grade, stance, stance_note, hand, use, note, second="—"):
    r = v1000._rec(i, sheet, metric, val, None, None, unit, basis, applies, date, src, grade, stance, stance_note, hand, use, note)
    r["ev"] = EID; r["s"] = "U-V1100"
    if second != "—": r["second"] = second
    return r


_MLP_URL = "https://raw.githubusercontent.com/mlcommons/inference_results_v6.0/main/summary_results.json"
SRC_NEW = [
    _rec("SRC_Perf", "SRC_PERF_079", "MLPerf v6.0 DeepSeek-R1 offline：NVIDIA GB200 NVL72（72 GPU）每 GPU", 6752.0, "tok/s/GPU（輸出 token）",
         "offline；Closed；available；486,141 ÷ 72；無延遲限制；批次只受記憶體限制；輸入 800／輸出 3,880", "GB200 NVL72（NVIDIA 提交 6.0-0075）", "2026-04", _MLP_URL, 1, "利害關係方",
         "MLCommons 稽核為中立；提交者 NVIDIA 為賣方；同輪 CoreWeave 64 GPU 7,323、Google 72 GPU 4,276（依軟體調校差很多）", "一手（原始資料檔已讀：" + MLP_CSV + "）",
         "Calib L 節 C 欄；Checks V20", f"v11.0 G3（{EID}）：大批次、無 SLO 的驗證點（不用於校準）；第二來源 SRC_PERF_082", second="SRC_PERF_082"),
    _rec("SRC_Perf", "SRC_PERF_080", "MLPerf v6.0 DeepSeek-R1 offline：NVIDIA GB300 NVL72（72 GPU）每 GPU", 8987.2, "tok/s/GPU（輸出 token）",
         "offline；Closed；available；647,076 ÷ 72；無延遲限制", "GB300 NVL72（NVIDIA 提交 6.0-0078）", "2026-04", _MLP_URL, 1, "利害關係方",
         "MLCommons 稽核為中立；提交者 NVIDIA 為賣方；同輪 Nebius 72 GPU 9,360", "一手（原始資料檔已讀：" + MLP_CSV + "）",
         "Calib L 節 D 欄；Checks V21", f"v11.0 G3（{EID}）：大批次、無 SLO 的驗證點；第二來源 SRC_PERF_083", second="SRC_PERF_083"),
    _rec("SRC_Perf", "SRC_PERF_081", "TensorRT-LLM DeepSeek-R1 decode：H200 8 GPU，並行 1,024（每 rank 128）每 GPU", 1436.2, "tok/s/GPU（輸出 token；同一批 GPU 做 prefill）",
         "ISL 1K／OSL 2K；EP8、attention DP；並行 1,024；每用戶 11.35 tok/s；無尾端延遲限制；是否使用 MTP 文件未寫（模型取無 MTP）", "Hopper H200 8 GPU（NVIDIA TensorRT-LLM）", "2025",
         "https://nvidia.github.io/TensorRT-LLM/0.21.0rc1/blogs/Best_perf_practice_on_DeepSeek-R1_in_TensorRT-LLM.html", 1, "利害關係方",
         "NVIDIA 為硬體與軟體供應商，有誘因展示高吞吐；單一來源（同文 B200 為另一硬體）", "一手（WebFetch 讀全文）",
         "Calib L 節 E 欄；Checks V22", f"v11.0 G3（{EID}）：Hopper 大批次驗證點（不用於校準）；同條件 MLPerf v5.1 server 556（SRC_PERF_060）的 2.6 倍；無第二來源（W1 +1）"),
    _rec("SRC_Perf", "SRC_PERF_082", "MLPerf v6.0 DeepSeek-R1 offline：CoreWeave GB200 NVL72（64 GPU）每 GPU", 7322.9, "tok/s/GPU（輸出 token）",
         "offline；Closed；available；468,665 ÷ 64", "GB200 NVL72（CoreWeave 提交 6.0-0015）", "2026-04", _MLP_URL, 1, "利害關係方",
         "CoreWeave 為雲端出租方；MLCommons 稽核為中立", "一手（原始資料檔已讀）", "Checks V20 說明欄（第二來源）", f"v11.0 G3（{EID}）：V20 第二來源（不同提交者、64 GPU）", second="SRC_PERF_079"),
    _rec("SRC_Perf", "SRC_PERF_083", "MLPerf v6.0 DeepSeek-R1 offline：Nebius GB300 NVL72（72 GPU）每 GPU", 9360.2, "tok/s/GPU（輸出 token）",
         "offline；Closed；available；673,936 ÷ 72", "GB300 NVL72（Nebius 提交 6.0-0081）", "2026-04", _MLP_URL, 1, "利害關係方",
         "Nebius 為雲端出租方；MLCommons 稽核為中立", "一手（原始資料檔已讀）", "Checks V21 說明欄（第二來源）", f"v11.0 G3（{EID}）：V21 第二來源（同輪 72 GPU，不同提交者）", second="SRC_PERF_080"),
    _rec("SRC_HW", "SRC_HW_067", "H200 HBM 容量", 141, "GB/GPU", "—", "Hopper H200", "2026-10-16", "https://www.nvidia.com/en-us/data-center/h200/（Quick Specs；WebFetch 2026-10-16 已讀）", 1, "利害關係方",
         "NVIDIA：產品規格", "一手（已讀）", "Calib L 節 E 欄（V22 的 H200 條件）", f"v11.0 G3（{EID}）：TensorRT-LLM H200 點的硬體規格；模型的 Hopper 欄仍是 H100（SRC_HW_020、021）"),
    _rec("SRC_HW", "SRC_HW_068", "H200 HBM 頻寬", 4.8, "TB/s/GPU", "—", "Hopper H200", "2026-10-16", "https://www.nvidia.com/en-us/data-center/h200/（Quick Specs；WebFetch 2026-10-16 已讀）", 1, "利害關係方",
         "NVIDIA：產品規格", "一手（已讀）", "Calib L 節 E 欄（V22 的 H200 條件）", f"v11.0 G3（{EID}）：同上"),
]


def src_new_records():
    return SRC_NEW


_IDS = "；".join(x["id"] for x in SRC_NEW)
EVIDENCE_V1100 = [[
    EID, DATE,
    "大批次低估查證（A 段，PR #84）：MLPerf v6.0 offline（GB200 NVL72 6,752、GB300 NVL72 8,987；同輪 CoreWeave 7,323、Nebius 9,360）；TensorRT-LLM H200 8 GPU 並行 1,024 每 GPU 1,436；H200 規格 141 GB／4.8 TB/s",
    f"{REPORT} 第 3、4 節；{EVTAB}；{MLP_CSV}；NVIDIA H200 規格頁", "Interested-party（NVIDIA、CoreWeave、Nebius；MLCommons 稽核）",
    "Checks V 節（V20–V22）；Calib L 節", "無大批次（offline）驗證點；V1–V3 為硬上限", "新增 V20–V22（通過帶 0.5–2.0）；V1–V3 改軟上限",
    "採納為驗證材料（不用於校準）；V1–V3 軟上限（MLPerf server TPOT 80 ms 是 p99，模型取平均）", V,
    f"讀取者：chat 端（A 段查證）；CC 轉入｜{DEC_BY}｜V22 無第二來源；H200 規格為一手（NVIDIA 頁）但模型 Hopper 欄仍用 H100 規格",
    "已處理", _IDS, DASH, "Checks V 節（CHK_Verify*）；Calib J 節 C、D 欄；IF_AllocRDGW 與研發 GW 下游（CAL_BatchEtaD GB200）", "1", "利害關係方（見各 SRC 紀錄）"]]


def evidence_rows():
    return EVIDENCE_V1100


# ================================================================== G3: Calib L section (below K): batch-path validation points V20–V22
CL_COLS = "CDE"
CL_LABELS = {
    "C": "V20：GB200 NVL72 — MLPerf v6.0 offline（NVIDIA 72 GPU；輸入 800／輸出 3,880；無 SLO；批次填滿 HBM）",
    "D": "V21：GB300 NVL72 — MLPerf v6.0 offline（同上）",
    "E": "V22：Hopper H200 8 GPU — TensorRT-LLM（輸入 1K／輸出 2K；並行 1,024＝每 rank 128；EP8；無 SLO）",
}
CL_NOTE = {
    "g": "世代索引（Spec_Rack 欄）：1 Hopper、2 GB200、3 GB300；黑字常數（測量條件）",
    "isl": "C、D：MLPerf 資料集平均輸入 800／輸出 3,880（Calib C110、C111，SRC_PERF_029／030）；E：TensorRT-LLM 文件 ISL 1K／OSL 2K",
    "ep": "0＝依容量規則 MAX(EP 基準, …)（同 H 節）；E：EP8（文件：attention DP＋EP8）",
    "bat": "0＝批次填滿 HBM（offline 不限速度）；E：並行 1,024 ÷ 8 GPU＝每 rank 128",
    "mtp": "0＝無 MTP（同 H 節 server 欄；offline 與 TensorRT-LLM 的 MTP 使用文件未寫）",
    "spec": "E 欄的 HBM、頻寬用 H200 規格（SRC_HW_067／068；模型 Hopper 欄為 H100 80 GB／3.35 TB/s，放不下 EP8 的權重）；峰值算力 H200 與 H100 相同",
    "ed": "批次口徑：採用 η_d（Calib 第 70 列）× CAL_BatchEtaD（第 156 列；Perf_Batch 第 53 列同一條式）。GB200＝3.04（vLLM 10.1K 校準，Calib J 節 D 欄）；GB300、Hopper＝1（未校準；Hopper 在 C68＝1.5 下無解）",
    "R": "模型 ÷ 實測（含 prefill 分攤；輸出 token 口徑）。V20–V22 的通過帶 0.5–2.0（Checks V20–V22）；不用於校準",
    "mult": "達標所需批次口徑 η_d 倍數（＝CAL_BatchEtaD 的解）：以 decode tok/s 反解目標輸出（扣除 prefill 時間），再解每序列算力項；若所需每序列時間 ≤ 非算力項（KV、EP 通訊）下限，則『無解』。此列供 Gov_Map 區間參考",
}


def calib_l(wb, CAL, SP):
    """Calib J patch (C、D columns carry CAL_BatchEtaD) and the L section. Returns the L row map; CAL['L'] = result."""
    ws = wb["Calib"]
    J = CAL["J"]
    # ---- J: calibration copy rows C、D now read the batch-path η_d (η_d × CAL_BatchEtaD); E、F、G (serving-side validation points) unchanged
    put(ws, f"A{J['hdr'] - 1}", "J. η 校準與條件對齊複製列（v10.2 E1／E3；v11.0 G2：C、D 欄的 η_d 乘 CAL_BatchEtaD）：比照 H 節公式（decode-only），條件改為各點的實際條件；模型 ÷ 目標顯示在每欄下方。"
        "校準用 C、D（批次口徑，解出 CAL_BatchEtaD）；驗證用 E、F、G（服務口徑，不乘 CAL_BatchEtaD；不重用校準點）", F_BOLD)
    put(ws, f"A{J['jed']}", "η_d（該世代；含 C67／D67；C、D 欄再乘 CAL_BatchEtaD＝批次口徑，v11.0 G2）", None)
    for X in "CD":
        put(ws, f"{X}{J['jed']}", f"=INDEX($C${CAL['etad']}:$G${CAL['etad']},{X}{J['jg']})*INDEX(CAL_BatchEtaD,1,{X}{J['jg']})", fmt="0.00%")
    put(ws, f"A{J['jmult']}", "達標所需 η_d 倍數（相對 CAL_EtaD；C、D 欄＝C67／D67 × CAL_BatchEtaD 的解）", None)
    put(ws, f"J{J['jed']}", "v11.0 G2：C、D 欄（校準複本列）乘批次口徑倍數 CAL_BatchEtaD（Hopper 1、GB200 3.04）；E、F、G 欄（服務口徑驗證點 V15、V7、V8）不乘", F_NOTE, wrap=True)
    r = J["jstate"]
    put(ws, f"A{r}", "校準狀態（v11.0 G2；停損規則）", F_BOLD)
    put(ws, f"C{r}", '="Hopper：CAL_BatchEtaD 現值 "&TEXT(INDEX(CAL_BatchEtaD,1,1),"0.00")&"（C67 "&TEXT($C$67,"0.00")&"）；模型 ÷ 目標 "&TEXT(C' + str(J["jR"]) + ',"0.00")&"；"&IF(ISNUMBER(C' + str(J["jmult"]) + '),"所需總倍數 "&TEXT(C' + str(J["jmult"]) + ',"0.00"),C' + str(J["jmult"]) + ')&"（η_d 再大產出上限 "&TEXT(C' + str(J["jmax"]) + ',"#,##0")&"＝目標的 "&TEXT(C' + str(J["jmaxR"]) + ',"0.00")&"）"', F_CALC)
    put(ws, f"D{r}", '="GB200：CAL_BatchEtaD 現值 "&TEXT(INDEX(CAL_BatchEtaD,1,2),"0.00")&"（D67 "&TEXT($D$67,"0.00")&"）；模型 ÷ 目標 "&TEXT(D' + str(J["jR"]) + ',"0.00")&"；"&IF(ISNUMBER(D' + str(J["jmult"]) + '),"所需總倍數 "&TEXT(D' + str(J["jmult"]) + ',"0.00")&"（η_d "&TEXT(D' + str(J["jetan"]) + ',"0.0%")&"）",D' + str(J["jmult"]) + ')', F_CALC)
    put(ws, f"J{r}", "G2：GB200 的解寫入 CAL_BatchEtaD（D156＝3.04，Derived）；Hopper 在 C68＝1.5 下無解（η_d 無限大只達目標 0.94，超出 ±2%），不寫入、列待 Project 判斷（Decisions BB-G2）。服務側 C67／D67 不動", F_NOTE, wrap=True)

    # ---- L section
    r0 = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    section(ws, r0, "L. 大批次驗證點（v11.0 G3）：MLPerf v6.0 offline（GB200、GB300）與 TensorRT-LLM H200；比照 J 節公式，條件為無 SLO 的大批次；η_d 取批次口徑（× CAL_BatchEtaD）並含 prefill 分攤；不用於校準", 10)
    r = r0 + 1
    R = {"sec": r0}
    put(ws, f"A{r}", "點位", F_BOLD)
    for X in CL_COLS: put(ws, f"{X}{r}", CL_LABELS[X], F_BOLD, wrap=True)
    ws.row_dimensions[r].height = 58; R["hdr"] = r; r += 1
    consts = [("lg", "世代索引", "", "0", [2, 3, 1]), ("lisl", "輸入長度（decode 起始上下文）", "tok", "#,##0", ["=$C$110", "=$C$110", 1000]),
              ("losl", "輸出長度", "tok", "#,##0", ["=$C$111", "=$C$111", 2000]), ("lep0", "EP 寬度設定（0＝依容量規則）", "GPU", "0", [0, 0, 8]),
              ("lbat", "固定批次（0＝批次填滿 HBM）", "序列/GPU", "0", [0, 0, 128]), ("lmtp", "推測解碼（1＝有）", "", "0", [0, 0, 0])]
    for key, lab, unit, fmt, vals in consts:
        R[key] = r; put(ws, f"A{r}", lab); put(ws, f"B{r}", unit)
        for X, v in zip(CL_COLS, vals): put(ws, f"{X}{r}", v, F_CALC if not (isinstance(v, str) and v.startswith("=")) else None, fmt=fmt)
        r += 1
    put(ws, f"J{R['lg']}", CL_NOTE["g"], F_NOTE, wrap=True); put(ws, f"J{R['lisl']}", CL_NOTE["isl"], F_NOTE, wrap=True); put(ws, f"J{R['lep0']}", CL_NOTE["ep"], F_NOTE, wrap=True)
    put(ws, f"J{R['lbat']}", CL_NOTE["bat"], F_NOTE, wrap=True); put(ws, f"J{R['lmtp']}", CL_NOTE["mtp"], F_NOTE, wrap=True)
    R["lT"] = r; put(ws, f"A{r}", "實測輸出 tok/s/GPU（目標）"); put(ws, f"B{r}", "tok/s/GPU")
    for X, s in zip(CL_COLS, ["SRC_PERF_079", "SRC_PERF_080", "SRC_PERF_081"]): put(ws, f"{X}{r}", f"={s}", fmt="#,##0", fill=FILL_KEY)
    put(ws, f"J{r}", "SRC_PERF_079（GB200 offline）／080（GB300 offline）／081（TensorRT-LLM H200）；輸出 token 口徑，含同一批 GPU 的 prefill", F_NOTE, wrap=True); r += 1
    gl = lambda row: f"=INDEX(Spec_Rack!$C${row}:$G${row},{{X}}{{lg}})"
    rows = [
        ("lgn", "世代", "", None, "=INDEX(Spec_Rack!$C$4:$G$4,{X}{lg})"),
        ("lP", "計算用峰值", "PF", "#,##0.000", gl(SP["pk"])), ("lH", "HBM", "GB", "#,##0", gl(SP["hbm"])), ("lB", "HBM 頻寬", "TB/s", "#,##0.00", gl(SP["bw"])),
        ("lLk", "EP 頻寬", "TB/s", "#,##0.00", gl(SP["link"])), ("lEb", "EP 基準", "", "#,##0", gl(SP["ep"])),
        ("lbe", "專家 B/param", "", "0.0000", gl(SP["be"])), ("lbn", "非專家 B/param", "", "0.00", gl(SP["bn"])),
        ("lU", "可用 HBM", "GB", "#,##0", "={X}{lH}*(1-Serving!$C$9)"),
        ("lTP", "注意力 TP", "", "0", "=MAX(1,CEILING($C${rne}*{X}{lbn}/(Serving!$C$10*{X}{lU}),1))"),
        ("lEP", "EP 寬度", "", "#,##0", "=IF({X}{lep0}>0,{X}{lep0},MAX({X}{lEb},CEILING($C${rex}*{X}{lbe}/(Serving!$C$11*{X}{lU}),1)))"),
        ("lW", "每 GPU 權重", "GB", "#,##0.0", "=$C${rne}*{X}{lbn}/{X}{lTP}+$C${rex}*{X}{lbe}/{X}{lEP}"),
        ("lcd", "decode 平均上下文", "tok", "#,##0", "={X}{lisl}+{X}{losl}/2"),
        ("lFd", "decode FLOPs/token", "GFLOP", "#,##0.0", "=(2*$C${rA}*1E9+$C${rac}*{X}{lcd})/1E9"),
        ("ln", "草稿數", "", "0", "=Serving!$C$6*{X}{lmtp}"), ("la", "每步接受", "", "0.00", "=IF({X}{lmtp}=1,Serving!$C$8,1)"),
        ("lx", "批次口徑 η_d 倍數（CAL_BatchEtaD）", "x", "0.00", "=INDEX(CAL_BatchEtaD,1,{X}{lg})"),
        ("led", "η_d（該世代 × CAL_BatchEtaD；批次口徑）", "%", "0.00%", "=INDEX($C${etad}:$G${etad},{X}{lg})*{X}{lx}"),
        ("ltl", "每層延遲（該世代；含 C68）", "µs", "#,##0", "=INDEX($C${tl}:$G${tl},{X}{lg})"),
        ("ltf", "每步固定延遲", "ms", "#,##0.00", "={X}{lW}/{X}{lB}+($C${rL}+{X}{ln})*{X}{ltl}/1000"),
        ("lc1", "每序列時間：算力項", "ms", "0.0000", "=({X}{ln}+1)*{X}{lFd}/({X}{lP}*{X}{led})/1000"),
        ("lc2", "每序列時間：KV 讀取項", "ms", "0.0000", "=$C${rkv}*{X}{lcd}/({X}{lB}*Serving!$C$13)/1E9"),
        ("lc3", "每序列時間：EP 通訊項", "ms", "0.0000", "=({X}{ln}+1)*$C${rL}*$C${rk}*$C${rd}*Serving!$C$12*({X}{lEP}-1)/{X}{lEP}/({X}{lLk}*Serving!$C$14)/1E9"),
        ("lce", "每序列每步時間（三項取最大）", "ms", "0.0000", "=MAX({X}{lc1},{X}{lc2},{X}{lc3})"),
        ("lcap", "容量批次上限", "序列", "#,##0.0", "=({X}{lU}-{X}{lW})*1E9/($C${rkv}*{X}{lcd})"),
        ("lBs", "B*（批次）", "序列", "#,##0.0", "=IF({X}{lbat}>0,MIN({X}{lbat},{X}{lcap}),MAX(0,{X}{lcap}))"),
        ("lD", "模型 decode tok/s/GPU", "tok/s", "#,##0", "=IF({X}{lBs}>0,{X}{lBs}*{X}{la}/(({X}{ltf}+{X}{lBs}*{X}{lce})/1000),0)"),
        ("lpe", "η_p（該世代）", "%", "0.00%", "=INDEX($C$72:$G$72,{X}{lg})"),
        ("lpf", "prefill FLOPs/token", "GFLOP", "#,##0.0", "=(2*$C${rA}*1E9+$C${rac}*{X}{lisl}/2)/1E9"),
        ("lpt", "prefill tok/s/GPU", "tok/s", "#,##0", "={X}{lpe}*{X}{lP}*1E15/({X}{lpf}*1E9)"),
        ("lout", "模型輸出 tok/s/GPU（含 prefill 分攤）", "tok/s", "#,##0", "=IF({X}{lD}>0,{X}{losl}/({X}{lisl}/{X}{lpt}+{X}{losl}/{X}{lD}),0)"),
        ("lR", "模型 ÷ 實測", "x", "0.000", "=IF(ISNUMBER({X}{lT}),IF({X}{lT}>0,{X}{lout}/{X}{lT},\"—\"),\"—\")"),
        ("lEff", "採用 η_d × 計算用峰值（有效 PF/GPU）", "PF", "0.000", "={X}{led}*{X}{lP}"),
        ("lfl", "非算力項下限（KV、EP 通訊取大）", "ms", "0.0000", "=MAX({X}{lc2},{X}{lc3})"),
        ("ldn", "達標所需 decode tok/s（扣除 prefill 時間）", "tok/s", "#,##0", '=IF({X}{losl}/{X}{lT}>{X}{lisl}/{X}{lpt},{X}{losl}/({X}{losl}/{X}{lT}-{X}{lisl}/{X}{lpt}),"無解（prefill 已超過目標時間）")'),
        ("lneed", "達標所需每序列每步時間", "ms", "0.0000", '=IF(AND(ISNUMBER({X}{ldn}),{X}{lBs}>0),(1000*{X}{lBs}*{X}{la}/{X}{ldn}-{X}{ltf})/{X}{lBs},"—")'),
        ("lmult", "達標所需批次口徑 η_d 倍數（CAL_BatchEtaD 的解）", "x", "0.00",
         '=IF(ISNUMBER({X}{lneed}),IF({X}{lneed}<={X}{lfl},"無解（非算力項下限）",(({X}{ln}+1)*{X}{lFd}/{X}{lP}/1000/{X}{lneed})/INDEX($C${etad}:$G${etad},{X}{lg})),"—")'),
    ]
    consts_map = dict(CAL); consts_map.update(R)
    for key, lab, unit, fmt, tpl in rows:
        R[key] = r; consts_map[key] = r; r += 1
    for key, lab, unit, fmt, tpl in rows:
        rr = R[key]; put(ws, f"A{rr}", lab); put(ws, f"B{rr}", unit)
        for X in CL_COLS:
            m = dict(consts_map); m["X"] = X
            f = tpl.format(**m)
            if X == "E" and key == "lH": f = "=SRC_HW_067"
            if X == "E" and key == "lB": f = "=SRC_HW_068"
            put(ws, f"{X}{rr}", f, fmt=fmt, fill=FILL_KEY if key in ("lR", "lmult") else None)
    put(ws, f"J{R['lH']}", CL_NOTE["spec"], F_NOTE, wrap=True); put(ws, f"J{R['led']}", CL_NOTE["ed"], F_NOTE, wrap=True)
    put(ws, f"J{R['lR']}", CL_NOTE["R"], F_NOTE, wrap=True); put(ws, f"J{R['lmult']}", CL_NOTE["mult"], F_NOTE, wrap=True)
    R["end"] = r - 1
    CAL["L"] = R
    return R


CALIB_L_TXT = "大批次驗證點複製列（V20–V22；顯示用）"


# ================================================================== G2: input write and notes (after restore())
def inputs_update(wb):
    """Calib D156 (CAL_BatchEtaD for GB200): 1 -> 3.04 (old-value guard). D156 is a black constant in the base (v523.calib_i); it becomes a blue Derived input; a rebuild restores the Excel value."""
    log = []
    ws = wb["Calib"]
    r = next(rr for rr in range(150, 170) if str(ws.cell(rr, 1).value or "").startswith("批次口徑 η_d 倍數（X10；CAL_BatchEtaD）"))
    assert r == 156, r
    if ws[f"D{r}"].value == 1:
        put(ws, f"D{r}", BATCH_GB200, F_IN, fmt="0.00"); log.append(f"Calib!D{r}: 1 -> {BATCH_GB200} (v11.0 G2)")
    return log


NOTE_156 = ("v11.0 G2：GB200 D156＝3.04（Derived 校準；Calib J 節 D 欄，目標 vLLM GB200 10.1K tok/s/GPU，批次填滿 HBM，解 3.037、模型 ÷ 目標 1.000）。Hopper C156 維持 1："
            "在 C68＝1.5 下無解（Calib J 節 C 欄，η_d 再大產出上限只達目標 2,324 的 0.94，超出 ±2%；C156＝3.74 即已飽和）；若 C68＝1.0 則需 4.32（Decisions BB-G1、BB-G2，待 Project 判斷）。"
            "GB300（E）、Rubin Ultra（G）維持 1；VR200（F）0.75 不動。只作用在 Perf_Batch（批次推論與訓練 rollout），不影響服務端")
NOTE_C68 = ("v11.0 G1 追查（Decisions BB-G1）：1.5 自 v5.2（2026-09-30）建檔即存在，標記 Assumed（Hopper EP 跨節點走 IB，區間 1–3），v5.2–v5.12 每層延遲＝240.4 µs（擬合 160.28 × 1.5），"
            "不是 DeepSeek 論文的 241.92 µs；該數字於 2026-10-16 才進 repo。v10.2 擬合更新為 161.19 µs 後乘積變 241.78 µs，與 241.92 µs 相差 0.06%＝巧合。物理上固定延遲（每層 all-to-all 延遲、同步）與 EP 通訊項（頻寬×資料量）也不是同一量。"
            "判定：無法確認重複計算，C68 不改。假設 C68＝1.0 的試算見 Decisions BB-G1")


def notes_update(wb, append_text):
    log = {"j156": 0, "j68": 0, "gm68": 0}
    ws = wb["Calib"]
    log["j156"] += append_text(ws["J156"], " ｜" + NOTE_156)
    log["j68"] += append_text(ws["J68"], " ｜" + NOTE_C68)
    g = wb["Gov_Map"]
    for r in range(5, g.max_row + 1):
        if g.cell(r, 3).value == "Calib" and g.cell(r, 4).value == "C68":
            log["gm68"] += append_text(g.cell(r, 14), " ｜v11.0 G1 追查：無法確認重複計算（1.5 為 v5.2 起的 Assumed 整數，非反推自 DeepSeek 241.92 µs），C68 不改；見 Decisions BB-G1")
    return log


# ================================================================== Gov_Map row for the new calibrated input (appended by gov.gm_append_c when absent)
def gov_map_rows(wb):
    ws = wb["Calib"]
    r = 156
    assert str(ws.cell(r, 1).value or "").startswith("批次口徑 η_d 倍數（X10；CAL_BatchEtaD）"), r
    return [{'scope': '切片四（v11.0 大批次修正）', 'sheet': 'Calib', 'cell': f'D{r}', 'label': '批次口徑 η_d 倍數：GB200（CAL_BatchEtaD）', 'check': ws.cell(r, 1).value, 'cls': 'Derived', 'role': '單值',
             'src': 'SRC_PERF_074', 'rel': '校準目標（vLLM GB200 10.1K，Calib J 節 D 欄；解 3.037）', 'dec': 'BB-G2', 'lo': D156_LO, 'hi': D156_HI,
             'rtext': f'{D156_LO}–{D156_HI}（各批次口徑點反推值：V20 MLPerf offline、V8 LMSYS、校準點 3.04；見 Decisions BB-G2）',
             'reason': ("GB200 批次口徑 η_d 倍數（只作用在 Perf_Batch：批次推論與訓練 rollout）。校準：Calib J 節 D 欄（目標 vLLM GB200 10.1K tok/s/GPU，批次填滿 HBM，無 SLO）求解 3.037，寫入 3.04（模型 ÷ 目標 1.000）；"
                        "標記 Derived（校準，單一目標點，利害關係方來源）。區間涵蓋各驗證點反推的值（V20 MLPerf offline GB200 與 V8 LMSYS NVFP4 的所需倍數，見 Calib L 節與 J 節）。"
                        "有效算力（η_d × 10 PF）約為 GB300 擬合值的 2 倍：GB300 的批次口徑倍數尚未校準（＝1），V21 顯示 GB300 offline 模型 ÷ 實測；待 Project 判斷是否校準 GB300"),
             'retag': '', 'seg': '高'}]


# ================================================================== end of build
def apply(wb, append_text):
    log = {"notes": notes_update(wb, append_text)}
    ws = wb["Calib"]
    sec = next(r for r in range(150, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("L. 大批次驗證點（v11.0 G3）"))
    n = 0
    for r in range(sec, ws.max_row + 1):
        n += v1030._set_mark(ws, r, CALIB_L_TXT)
    log["marks_calib_l"] = n
    sp = wb["SRC_Perf"]; k2 = 0
    sec = {x["id"]: x["second"] for x in SRC_NEW if x.get("second")}
    for rr in range(5, sp.max_row + 1):
        i = sp.cell(rr, 1).value
        if i in sec and sp.cell(rr, 18).value in (None, DASH, ""):
            sp.cell(rr, 18).value = sec[i]; k2 += 1       # column R = 第二來源 SRC_ID (Excel-owned; written only while it is still the dash)
    log["second_sources"] = k2
    log["decisions"] = decisions_append(wb)
    return log


# ================================================================== Decisions (10 columns; layout of v1020／v1030)
_D = [
    ("G1", "C68（Hopper 每層延遲倍數 1.5）是否重複計算 EP 通訊：無法確認，不改",
     "追查：(1) Calib J68、Gov_Map GM346、Decisions D204／D171／E1H 與 J4（C187）：C68 標『Assumed／倒推（v5.5 建檔時與 C67 手算配對）1.5，區間 1–3』，說明為『Hopper EP 跨節點走 IB』，無來源，Excel 內無公式關聯；"
     "(2) 封存工作簿 v5.2、v5.4、v5.5、v5.12：C68 在 v5.2（2026-09-30，Calib 結構建立時）即為 1.5，每層延遲（採用）＝240.42 µs（GB300 擬合 160.28 × 1.5），不是 241.92 µs；v5.5 只把 C67 由 3 減半為 1.5（峰值更正 989→1,979 TF），C68 未動；"
     "(3) DeepSeek 論文每層 241.92 µs（雙 micro-batch 2×120.96）在 repo 內首見於 2026-10-16 的大批次查證（證據表第 5 列）；v10.2 擬合更新（161.19 µs）後乘積 241.78 µs，與 241.92 µs 相差 0.06%，是巧合；"
     "(4) 物理量不同：C68 作用在每層固定延遲（all-to-all 延遲、kernel 啟動、同步；與批次無關），EP 通訊項是資料量 ÷ 頻寬（隨批次線性）。"
     "判定：1.5 是先於該數字存在的整數假設，不是用 DeepSeek 數字反推或對齊；無證據顯示重複計算。依工作單『無法確認』列：C68 不改，列待 Project 判斷。"
     "假設 C68＝1.0 的試算（scratch，未寫入）：Hopper 服務側 IF_TokGW_Sol 3.363e10 → 3.738e10（+11.1%）；L1_MFleet 0.4073 → 0.3886；IF_AllocRDGW 0.2402 → 0.2359；IF_AllocServeGW 不變；"
     "且 Hopper 的 CAL_BatchEtaD 會有解（4.32；η_d 15.5%）。",
     "Calib C68、J68；Gov_Map GM346；本決策；未改任何數值", "已查，C68 不改；待 Project 判斷"),
    ("G2", "校準 CAL_BatchEtaD：GB200 寫入 3.04；Hopper 無解，不寫入",
     "目標：GB200 vLLM 10.1K（Calib J 節 D 欄條件：2K／2K、批次填滿 HBM、無 SLO、無 MTP）。Calib J 節 C、D 欄的 η_d 改為（該世代 η_d）× CAL_BatchEtaD（批次口徑；E、F、G 欄服務口徑驗證點不變）。"
     "GB200 解 3.0373（η_d 10.9%），寫入 D156＝3.04，模型 ÷ 目標 1.000（容差 ±2%）；標記 Derived；區間 1.5–6.7（涵蓋各批次口徑點反推值：V20 MLPerf offline、V8 LMSYS，見 Calib L 節、J 節）。"
     "Hopper：目標 DeepSeek profile 2,324（批次 128、4K）在 C68＝1.5 下無解——EP 通訊項（0.2975 ms／序列）使產出上限只達目標的 0.941（η_d 無限大；C156＝3.74 即飽和），超出 ±2%，依工作單『無解停下』不寫入，C156 維持 1。"
     "若 C68＝1.0（G1 的另一種判定）則 Hopper 需 4.32。檢查：GB200 η_d × 峰值＝1.086 PF，不超過硬體峰值（η_d 10.9% < 100%）；但約為 GB300 擬合值（0.536 PF）的 2 倍，GB300 批次口徑倍數未校準（＝1）。"
     "影響（只套 GB200）：IF_AllocRDGW 0.2402 → 0.1973；IF_AllocQ1 0.2110 → 0.1801；L1_ScaleRD 5.318 → 6.475；IF_AllocServeGW、L1_MFleet、L1_LReqFleet_Id 不變。"
     "若再加 Hopper C156＝3.74（假設）：IF_AllocRDGW 0.1626。",
     "Calib D156、J 節（C、D 欄 η_d 與狀態列）、Gov_Map 新列；Perf_Batch 第 53 列 → Training → IF_AllocRDGW 與下游", "v11.0 已建（GB200）；Hopper 待 Project 判斷"),
    ("G3", "新增大批次驗證點 V20–V22（Calib L 節；SRC_PERF_079–083、SRC_HW_067–068）",
     "V20：MLPerf v6.0 offline GB200 NVL72 6,752；V21：offline GB300 NVL72 8,987；V22：TensorRT-LLM H200 8 GPU 1,436（ISL 1K／OSL 2K、並行 1,024＝每 rank 128、EP8；H200 規格 141 GB／4.8 TB/s）。"
     "條件：無 SLO、大批次；模型以 η_d × CAL_BatchEtaD（批次口徑）、含 prefill 分攤（輸出 token 口徑）；通過帶 0.5–2.0；標記 Interested-party（NVIDIA 為賣方；MLCommons 稽核；V20、V21 有第二來源，V22 無）；不用於校準。",
     "Calib L 節；Checks V20–V22；SRC_Perf／SRC_HW 新列", "v11.0 已建"),
    ("G4", "MLPerf server（V1–V3）改為軟上限；新增 CHK_VerifyWarn（主版）",
     "V1、V2、V3：比值 ≤ 1 通過；> 1 為『警示』（MLPerf server 的 TPOT 80 ms 是 p99 口徑，模型取平均；同條件 TensorRT-LLM H200 為 MLPerf v5.1 的 2.6 倍），不計失敗。"
     "新增具名範圍 CHK_VerifyWarn（Checks V 節），恆等式 CHK_VerifyPass＋CHK_VerifyFail＋CHK_VerifyWarn＝CHK_VerifyTotal（CHK_VerifyTotal 含警示項）。interactive 類（V4–V6）維持硬上限（chat 端預設：interactive 為低延遲口徑，與模型服務端口徑較接近）。"
     "新增契約名稱，故為主版 v11.0。",
     "Checks V 節（F 欄公式、V 節摘要列）；CHK_VerifyWarn；下游契約文件", "v11.0 已建"),
]
DECISIONS_V1100 = [[f"BB-{i}", V, f"大批次修正：{t}", txt, DATE, DEC_BY, st, where, f"{WO} {i}", "否"] for i, t, txt, where, st in _D]


def decisions_append(wb):
    ws = wb["Decisions"]
    have = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
    r = max([rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")] or [4]) + 1
    n = 0
    for row in DECISIONS_V1100:
        if row[0] in have: continue
        for i, v in enumerate(row): put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 5, 7))
        r += 1; n += 1
    return n


# ================================================================== README
README_VERSION = (VERSION + "（主版：大批次修正，" + WO + "——G1 C68 追查（無法確認重複計算，不改）；G2 CAL_BatchEtaD GB200＝3.04（Hopper 無解不寫入）、Calib J 節 C／D 欄改批次口徑；G3 新增 V20–V22（MLPerf offline、TensorRT-LLM H200；Calib L 節）；"
                  "G4 V1–V3 改軟上限、新增 CHK_VerifyWarn；Decisions BB-*）。以下為 " + v1030.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1030.README_TITLE
README_TITLE = VERSION.split("_")[-1] + _T[len(v1030.VERSION.split("_")[-1]):]
README_ROW = ("大批次修正（v11.0）",
              "C68 追查（不改）；GB200 批次口徑 η_d 倍數 3.04（Hopper 無解未寫入）；新增驗證 V20–V22；MLPerf server（V1–V3）改軟上限，新增 CHK_VerifyWarn。細節見 Decisions BB-*。")
