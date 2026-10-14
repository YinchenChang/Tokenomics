#!/usr/bin/env python3
"""匯出（規劃書第 6 節「CC 可以隨時匯出」）：SRC_*（A–W 欄）與 Interface → data/export/*.csv；並檢查 GOV_Errors。

用法：python3 tools/export_csv.py [--out data/export] [--check-only]
- 值一律由 LibreOffice 重算 model/CURRENT 指向的活頁簿後讀取（修訂八-B：不再使用公式引擎）；不另寫任何計算。
- GOV_Errors ≠ 0 時以非零狀態結束（CI 用；CLAUDE.md 第 3 節、第 11 輪指令第 5 節）。
- 匯出檔只供外部讀取，不寫回 Excel（CLAUDE.md 第 6 節）。
"""
import argparse
import csv
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "tests" / "parity"))
from parity_lib import lo_recalc, parse_ref, set_inputs  # noqa: E402

SRC_SHEETS = ("SRC_HW", "SRC_DC", "SRC_Model", "SRC_Perf", "SRC_Price", "SRC_Cap", "SRC_Harness", "SRC_Demand")


class Recalculated:
    """model/CURRENT 指向的活頁簿經 LibreOffice 重算後的值（唯讀）。介面只含本工具用到的 get／get_name。"""

    def __init__(self):
        self.path = REPO / "model" / (REPO / "model" / "CURRENT").read_text(encoding="utf-8").strip()
        tmp = Path(tempfile.mkdtemp(prefix="export_csv_"))
        try:
            src = tmp / self.path.name
            set_inputs(self.path, src, {})
            out = lo_recalc(src, tmp / "lo")
            wb = openpyxl.load_workbook(out, data_only=True)
            self.names = {k: v.attr_text for k, v in wb.defined_names.items()}
            self._sheets = {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None} for ws in wb}
            wb.close()
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def get(self, sheet, ref):
        d = self._sheets.get(sheet, {})
        ref = ref.replace("$", "")
        if ":" not in ref:
            return "" if d.get(ref) is None else d[ref]
        c1, r1, c2, r2 = range_boundaries(ref)
        return [["" if d.get(f"{get_column_letter(c)}{r}") is None else d[f"{get_column_letter(c)}{r}"] for c in range(c1, c2 + 1)]
                for r in range(r1, r2 + 1)]

    def get_name(self, name):
        v = self.get(*parse_ref(self.names[name]))
        if isinstance(v, list) and (len(v) == 1 or all(len(r) == 1 for r in v)):
            return [x for row in v for x in row]
        return v


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=REPO / "data" / "export")
    ap.add_argument("--check-only", action="store_true")
    a = ap.parse_args()
    eng = Recalculated()
    errors = eng.get_name("GOV_Errors")
    print(f"{eng.path.name}: GOV_Errors={errors} GOV_Warnings={eng.get_name('GOV_Warnings')} GOV_Info={eng.get_name('GOV_Info')}")
    if not a.check_only:
        a.out.mkdir(parents=True, exist_ok=True)
        for s in SRC_SHEETS + ("Interface",):
            last = "W" if s.startswith("SRC_") else "R"
            first = 4 if s.startswith("SRC_") else 1
            grid = eng.get(s, f"A{first}:{last}{600 if s.startswith('SRC_') else 300}")
            rows = [r for r in grid if any(x != "" for x in r)]
            with open(a.out / f"{s}.csv", "w", newline="", encoding="utf-8-sig") as f:
                csv.writer(f).writerows(rows)
            print(f"  {s}.csv：{len(rows)} 列")
    if errors != 0:
        print("GOV_Errors 不為 0", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
