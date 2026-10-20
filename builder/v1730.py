# v17.3 (work order docs/workorders/20261020_W1第二來源_B2.md, branch claude/w1-b2): W1 second sources, B2 segment (DeepSeek V4-Flash CSA top-k 1024 -> 512).
# MINOR version (CLAUDE.md 4: SRC Active value and formula change).  Andy 2026-10-20 17:29「照預設」(review workbook Q2: Flash column to 512).
#   1  New SRC_MOD_106 (Active): DeepSeek V4-Flash CSA top-k = 512 blocks (HF V4-Flash config.json index_topk; 2nd source vLLM blog c4a default 512);
#      grade 2, C1 spec class, R = "免（C1 官方規格；Andy 2026-10-20 17:17）".
#   2  SRC_MOD_018 (1024): 適用對象 DeepSeek V4 -> DeepSeek V4-Pro (revision 9: original kept in W, Decisions SL52-01), R = same C1 exemption, V (使用位置) Arch!D17 only.
#   3  Arch!C17 = SRC_MOD_106*SRC_MOD_019+SRC_MOD_017 (formula map; D17 unchanged).  4224 -> 2176.   Arch!G17 label names both values (label text only).
#   4  Gov_Map GM201 (Arch!C17) H = SRC_MOD_106 (original SRC_MOD_018 kept in N, revision 9).
#   5  DB_Evidence E399 (the new Active record needs an Evidence ID, Checks E8); Decisions SL52-01.
# Every write to an existing (Excel-owned) cell is guarded by its old value; a re-run is a no-op (idempotent).
from openpyxl.utils import get_column_letter as L

import v1720
import v1410

V = "v17.3"
VERSION = "20261021_Tokenomics_v17.3"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-21"
WO = "工作單 docs/workorders/20261020_W1第二來源_B2.md"
EVD = "docs/evidence/20261020_W1第二來源A段/B0/Q2Q5Q6.md"
DASH = "—"
C1_TEXT = "免（C1 官方規格；Andy 2026-10-20 17:17）"
NEW_ID = "SRC_MOD_106"
OLD_ID = "SRC_MOD_018"
DEC_BY = "Andy 2026-10-20 17:29「照預設」（審閱檔 Q2：Flash 欄改 512）；17:17 C 類規則（官方規格免第二來源）"
TAG = f"｜{V} 更正（修訂九；Decisions SL52-01）：原"

FM = {"Arch!C17": f"={NEW_ID}*SRC_MOD_019+SRC_MOD_017"}


def formula_map(fm):
    fm = dict(fm); fm.update(FM); return fm


def src_new_records(wb):
    rec = v1720._rec(
        NEW_ID, "SRC_Model", "DeepSeek V4-Flash CSA top-k 壓縮塊數", 512, None, None, "塊", "config.json index_topk；MTP 層不計", "DeepSeek V4-Flash",
        "2026-10-20（讀取日）",
        "Hugging Face deepseek-ai/DeepSeek-V4-Flash config.json（index_topk）｜https://huggingface.co/deepseek-ai/DeepSeek-V4-Flash/raw/main/config.json；"
        "第二來源：vLLM 部落格（c4a 預設 512）", 2, "利害關係方",
        "DeepSeek 官方 config.json（C1 官方規格類）；模型作者自揭規格，誤報誘因低", "一手（WebFetch 讀取，非逐字原文）", "Arch!C17｜換算",
        f"{V} W1 B2：Arch 頁 C 欄＝Luna＝DeepSeek V4-Flash 類，原與 Pro 共用 SRC_MOD_018＝1024；Flash 的 config 為 512（原行 `\"index_topk\": 512,`，兩次獨立讀取一致），"
        f"Pro 為 1024。512 × 4 ＋ 128 ＝ 2,176（原 4,224）。壓縮塊長 4（SRC_MOD_019）與滑動窗 128（SRC_MOD_017）對 Flash 同樣適用："
        f"Flash config.json compress_ratios 含 4 與 128 交替、sliding_window 128（壓縮塊長 4 未在 config 單獨驗證，只見 compress_ratios 含 4）。"
        f"WebFetch 經小模型摘要，非逐字原文；curl 直連 403。出處與讀取紀錄：{EVD}")
    rec.update(status="Active", s="U-DS", ev="E399")      # Active (feeds Arch!C17); 原 S 編號 column continues the DeepSeek group
    return [rec]


# ---------------------------------------------------------------- Gov_Map GM201 (Arch!C17)
N205_ADD = (f"｜{V} 修訂九（Decisions SL52-01）：H 欄原為 {OLD_ID}（DeepSeek V4 共用 top-k 1024）；Arch!C17（Luna＝V4-Flash）改連 {NEW_ID}（512），"
            f"{OLD_ID} 改為只適用 V4-Pro（GM202＝Arch!D17 仍連 {OLD_ID}）")


def gov_update(ws, append_text):
    n = 0
    for r in range(5, ws.max_row + 1):
        if ws[f"C{r}"].value == "Arch" and ws[f"D{r}"].value == "C17":
            if ws[f"H{r}"].value == OLD_ID:
                ws[f"H{r}"].value = NEW_ID; n += 1
                append_text(ws[f"N{r}"], N205_ADD)
            else:
                assert ws[f"H{r}"].value == NEW_ID, ws[f"H{r}"].value
    return n


# ---------------------------------------------------------------- DB_Evidence E399
EVIDENCE = ("E399", "DeepSeek V4-Flash 與 V4-Pro 的 CSA top-k 不同：Flash index_topk 512、Pro 1024（原模型把 Luna＝Flash 欄與 Pro 共用 1024）",
            "Hugging Face deepseek-ai/DeepSeek-V4-Flash 與 DeepSeek-V4-Pro 的 config.json（WebFetch 讀取，2026-10-20）；vLLM 部落格 c4a 預設 512；" + EVD,
            "Verified（官方 config.json，C1 官方規格）", "Arch!C17（Luna＝V4-Flash 類）：全注意力層跨度上限", "4,224（1,024 × 4 ＋ 128）", "2,176（512 × 4 ＋ 128）",
            "採用（W1 B2；Andy 2026-10-20 17:29「照預設」）",
            "Pro 的 1024 只讀一次（摘要答覆）；Flash 的 512 兩次獨立讀取一致；config 無名為 CSA 的欄位，以 index_topk 對應；"
            "新紀錄 SRC_MOD_106；SRC_MOD_018 改為只適用 V4-Pro｜讀取者：W1 A 段 B0 查證",
            "Arch!C17 及其下游（Perf、Sens_Perf、Perf_Batch、Training、Sens_Train 的 Luna 欄；Gov_Map Q205）；影響的 IF_／L1_ 見報告")


def evidence_append(wb):
    ws = wb["DB_Evidence"]
    if any(ws.cell(r, 1).value == EVIDENCE[0] for r in range(5, ws.max_row + 1)):
        return 0
    from common import put, F_CALC
    r = max(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")) + 1
    eid, claim, src, mark, param, cur, new, verdict, note, affected = EVIDENCE
    vals = [eid, DATE, claim, src, mark, param, cur, new, verdict, V, note, "已處理", NEW_ID, "—", affected, 2, "利害關係方", f"{verdict}｜{V} 新增列"]
    for i, v in enumerate(vals):
        put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 4, 5, 6, 7, 8, 10, 14, 17))
    return 1


# ---------------------------------------------------------------- apply
def _rows(wb):
    ws = wb["SRC_Model"]
    return {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, str)}, ws


G17_OLD = "V4 CSA：top-k 1,024 壓縮塊 × 4 ＋ 128 窗口"
G17_NEW = f"V4 CSA：top-k 壓縮塊 × 4 ＋ 128 窗口；C 欄（Flash）top-k 512（{NEW_ID}）、D 欄（Pro）top-k 1,024（{OLD_ID}）"


def apply(wb, append_text):
    log = {}
    rows, ws = _rows(wb)
    r18, r106 = rows[OLD_ID], rows[NEW_ID]
    assert ws[f"C{r106}"].value == 512 and ws[f"C{r18}"].value == 1024
    # new record: C1 exemption (guarded)
    log["r106"] = v1720._set(ws[f"R{r106}"], DASH, C1_TEXT)
    # SRC_MOD_018 -> V4-Pro (revision 9: originals verbatim in W)
    n = 0
    if ws[f"H{r18}"].value != "DeepSeek V4-Pro":
        assert ws[f"H{r18}"].value == "DeepSeek V4" and ws[f"V{r18}"].value == "Arch!C17｜換算", (ws[f"H{r18}"].value, ws[f"V{r18}"].value)
        append_text(ws[f"W{r18}"], f"{TAG} H 欄（適用對象）＝『DeepSeek V4』｜原 V 欄（模型使用位置）＝『Arch!C17｜換算』（C17 改連 {NEW_ID}，見 Decisions SL52-01；依據 {EVD}）")
        ws[f"H{r18}"].value = "DeepSeek V4-Pro"; ws[f"V{r18}"].value = "Arch!D17｜換算"; n += 1
    n += v1720._set(ws[f"R{r18}"], DASH, C1_TEXT)
    log["r018"] = n
    # Arch!G17 label (text only)
    a = wb["Arch"]
    if a["G17"].value != G17_NEW:
        assert a["G17"].value == G17_OLD, a["G17"].value
        a["G17"].value = G17_NEW; log["g17"] = 1
    log["evidence"] = evidence_append(wb)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == "SL52-01" for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


def _decisions():
    return [["SL52-01", V, "W1 第二來源 B2：DeepSeek V4-Flash CSA top-k 1024 → 512（Arch!C17）",
             f"Arch 頁 C 欄＝Luna＝DeepSeek V4-Flash 類，D 欄＝Pro 類；兩欄原共用 {OLD_ID}（top-k 1024）。官方 config.json：Flash index_topk 512、Pro 1024。"
             f"新增 {NEW_ID}（V4-Flash，512 塊，Active，等級 2，C1 官方規格，R 欄填『{C1_TEXT}』）；Arch!C17 公式改為 {NEW_ID}×SRC_MOD_019+SRC_MOD_017"
             f"（4,224 → 2,176）；D17 不變。{OLD_ID} 適用對象由『DeepSeek V4』改『DeepSeek V4-Pro』（修訂九：原值逐字保留在該列 W 欄）、R 欄填同樣的 C1 免除標記；"
             f"Gov_Map GM201（Arch!C17）H 欄改連 {NEW_ID}（原值留在 N 欄）。SRC_MOD_019（壓縮塊長 4）與 SRC_MOD_017（滑動窗 128）對 Flash 同樣適用（Flash config compress_ratios 含 4 與 128、sliding_window 128）。"
             f"另新增 DB_Evidence E399（Active 紀錄需 Evidence ID，Checks E8）；Arch!G17 說明文字同步。",
             DATE, DEC_BY, "v17.3 已建", f"Arch!C17 及下游（Perf／Sens_Perf／Perf_Batch／Training／Sens_Train 的 Luna 欄、Gov_Map Q205）；SRC_Model {OLD_ID}／{NEW_ID}；Gov_Map GM201；DB_Evidence E399",
             f"{WO}；{EVD}", "否"]]


# ================================================================== README
README_VERSION = (VERSION + "（次版：" + WO + "——W1 第二來源 B2 段：DeepSeek V4-Flash CSA top-k 1024 → 512；新增 SRC_MOD_106（Flash 512），SRC_MOD_018 改只適用 V4-Pro，"
                  "Arch!C17 改連 SRC_MOD_106（4,224 → 2,176）；Gov_Map GM201、DB_Evidence E399、Decisions SL52-01）。以下為 " + v1720.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1720.README_TITLE
README_TITLE = V + _T[len(v1720.V):]
README_ROW = ("v17.3 W1 第二來源 B2",
              "Arch!C17（Luna＝DeepSeek V4-Flash 類）的 CSA top-k 由與 Pro 共用的 1,024 改為 Flash 自己的 512（SRC_MOD_106），全注意力層跨度上限 4,224 → 2,176；"
              "SRC_MOD_018（1,024）只適用 V4-Pro。細節見工作單 docs/workorders/20261020_W1第二來源_B2.md 與 Decisions SL52-01。")
