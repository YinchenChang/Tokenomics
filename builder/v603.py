# v6.3 (work order docs/workorders/20261016_稽核修正2.md r1, section 2b = segment 2-1b): scale-basis repairs after 2-1 and 2Q (minor version).
#   K1 the mixed unit price paired with the Alloc generation-mixed (fleet) tokens uses the same reachable tiers as the capacity side: tier weights are renormalised over the
#      reachable tiers per generation (Alloc section I: AL_PriceGen, AL_PriceFleet), then weighted across generations by paid-token share (share x AL_CapPaid) so that
#      revenue / (fleet tokens x price) uses ONE set of tiers. Cap_In H7 (m_fleet, m floor, "must be true" rows), H10 m_fleet (Cap_In K section: B4_PriceFleet), L1_PriceRealFleet and
#      L1_PriceReqBE read it. Demand-route prices (model tokens AL_DPaid, row 90) keep the full tier mix B4_MixPaid.
#   K2 the generation mix is an observed fleet fact: a generation with share > 0 whose tiers are all unreachable is not dropped and its share is not transferred; AL_BlendPaid／Free,
#      AL_ServeGW, AL_Q1／Q2 and the closure names answer the text flag "整代不可達" (block6.py, v600.py, v534._CAP). AL_UnservedGen stays.
#   K3 Perf／Perf_Batch／Sens_Perf row "decode GPU-s per M" (and the _Prod mirrors) answers "TTFT 不可達" when row 77 (TTFT) is "否" (perf.py).
#   K4 investigation only (report); nothing is changed for it.
#   J1 r1 (2Q): Spec_Rack ARCH_ScaleGen_Lo／_Hi blue cells and scenario d1_model_scale (0.8／0.55／1／1.4) follow the 2Q report Q2; guarded writes (cell still holds its v6.2 value).
from openpyxl.workbook.defined_name import DefinedName
from common import put, F_IN, F_NOTE, F_BOLD, F_CALC, FILL_KEY, section
import v533
import v534
import v536
import v602

VERSION = "20261016_Tokenomics_v6.3"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-16"
V = "v6.3"
WO = "工作單 docs/workorders/20261016_稽核修正2.md r1 第 2b 節"
ANDY = "chat 端預設（K1–K3 依 D8 『不可達＝不可用』與 D9 『價格與 token 同口徑』；J1 依 2Q 報告 Q2），可撤回；docs/reports/20261016_稽核修正2Q_查證.md"
SLO = "SLO 不可達"
GENX = "整代不可達"
TIERS = [("Luna", 1), ("Sol", 2), ("Astra", 3)]
_PATCHED = False


# ================================================================== spec patches (before the specs are written; call right after v601.patch_specs())
def patch_specs():
    global _PATCHED
    if _PATCHED: return
    # ---- K1: Cap_In H7 rows that pair the price with fleet tokens read AL_PriceFleet (Alloc section I) ----
    tag = "機隊口徑混合有效單價（AL_PriceFleet：可達層級重新正規化；v6.3 K1）"
    sp = v533.SPEC
    for key in ("m_fleet", "m_fleet_rs", "m_floor", "need_1", "need_06", "need_03"):
        i = next(i for i, it in enumerate(sp) if it[0] == key)
        it = list(sp[i])
        for k in range(3, 6):
            if isinstance(it[k], str):
                for X in "CDE": it[k] = it[k].replace(f"{X}{{P_price}}", "AL_PriceFleet")
        if key.startswith("need"):       # AL_PriceFleet may be the text 整代不可達 (K2): text passes through, never #VALUE!
            it[3] = f'=IF(ISNUMBER(AL_PriceFleet),{it[3][1:]},AL_PriceFleet)'
            it[4] = f'=IF(ISNUMBER(C{{{key}}}),C{{{key}}}/AL_DPaid,C{{{key}}})'
            it[5] = it[5].replace("ISNUMBER(C{tok_ceil})", f"AND(ISNUMBER(C{{tok_ceil}}),ISNUMBER(C{{{key}}}))", 1)
        it[1] = it[1].replace("混合有效單價", tag).replace("單價", "機隊口徑單價（AL_PriceFleet）") if key.startswith("need") else it[1].replace("混合有效單價", tag)
        assert "P_price" not in "".join(str(x) for x in it[3:6]), key
        sp[i] = tuple(it)
    # ---- K1: Cap_In H10 m_fleet reads the Cap_In K-section row (named B4_PriceFleet, 1x4 C:F) ----
    sp = v534.SPEC_H10
    j = next(i for i, it in enumerate(sp) if it[0] == "q_m_fleet")
    it = list(sp[j])
    for k, X in zip(range(3, 7), "CDEF"):
        assert f"{X}{{q_price}}" in it[k], it[k]
        it[k] = it[k].replace(f"{X}{{q_price}}", f"INDEX(B4_PriceFleet,1,{'CDEF'.index(X) + 1})")
    it[1] = "m_fleet＝m 分子營收 ÷（機隊付費 token × 機隊口徑單價 B4_PriceFleet；Cap_In K 節；v6.3 K1）"
    sp[j] = tuple(it)
    _PATCHED = True


# ================================================================== K1: Alloc section I (appended after H; nothing above moves)
def alloc_fleet_price(wb, R):
    ws = wb["Alloc"]
    r = ws.max_row + 2
    section(ws, r, "I. 機隊口徑混合有效單價（v6.3 K1；欄 C–F＝服務世代 Hopper／GB200／GB300／VR200）：與機隊付費 token 配對的單價，層級權重只在可達層級間重新正規化（同 B 節），"
                   "世代間以付費 token 占比（世代占比 × 各世代付費產能）加權", 7); r += 1
    price = {"Luna": "IF_PriceRef_Luna", "Sol": "IF_PriceRef_Sol", "Astra": "IF_PriceRef_Astra"}
    flags = lambda c: [f"{c}{R[k]}" for k in ("rl", "rs", "ra")]
    put(ws, f"A{r}", "各世代付費 token 混合有效單價＝Σ（可達層級的付費組合權重 × 該層級參考單價）÷ 可達層級權重合計（全不可達的世代為文字）", wrap=True); put(ws, f"B{r}", "$/M")
    for c in "CDEF":
        num = "+".join(f"IF({f}=1,INDEX(B4_MixPaid,1,{t})*{price[tn]},0)" for t, (f, (tn, _)) in enumerate(zip(flags(c), TIERS), 1))
        put(ws, f"{c}{r}", f'=IF({c}{R["ok"]}=1,({num})/{c}{R["wp"]},"{SLO}")', fmt="0.0000")
    put(ws, f"G{r}", "全層級可達時等於 Cap_In 第 90 列（B4_MixPaid × IF_PriceRef_*）；基準下 Hopper 欄排除 Astra（TTFT 不可達），單價較低", F_NOTE)
    wb.defined_names["AL_PriceGen"] = DefinedName("AL_PriceGen", attr_text=f"Alloc!$C${r}:$F${r}")
    pg = r; r += 1
    ss = [f"{c}${R['share']}" for c in "CDEF"]; ok = [f"{c}${R['ok']}" for c in "CDEF"]
    cp = [f"{c}${R['capp']}" for c in "CDEF"]; pgc = [f"{c}${pg}" for c in "CDEF"]
    num = "+".join(f"IF(AND({s}>0,{o}=1),{s}*{k}*{p},0)" for s, o, k, p in zip(ss, ok, cp, pgc))
    put(ws, f"A{r}", "機隊口徑混合有效單價＝Σ（世代占比 × 各世代付費產能 × 各世代單價）÷ Σ（世代占比 × 各世代付費產能）＝ 機隊付費 token 的平均單價", wrap=True); put(ws, f"B{r}", "$/M")
    put(ws, f"C{r}", f'=IF(AL_BlendPaid="{GENX}","{GENX}",IF(AND(ISNUMBER(AL_BlendPaid),AL_GenShareOK>0),IF(AL_BlendPaid>0,({num})/AL_GenShareOK/AL_BlendPaid,"{SLO}"),"{SLO}"))',
        fmt="0.0000", fill=FILL_KEY)
    put(ws, f"G{r}", "分子分母同為可達層級、可服務世代；與 Cap_In H7 的機隊付費 token（AL_BlendPaid 加權）同口徑。全層級可達時等於 Cap_In 第 90 列", F_NOTE)
    wb.defined_names["AL_PriceFleet"] = DefinedName("AL_PriceFleet", attr_text=f"Alloc!$C${r}")
    pf = r; r += 1
    put(ws, f"A{r}", "對照：完整層級組合混合有效單價（Cap_In 第 90 列；需求路線的模型 token 用此價）", wrap=True); put(ws, f"B{r}", "$/M")
    put(ws, f"C{r}", "=Cap_In!$C$90", fmt="0.0000")
    pc = r; r += 1
    put(ws, f"A{r}", "機隊口徑單價 ÷ 完整層級組合單價（<1：被排除的層級單價較高）", wrap=True); put(ws, f"B{r}", "x")
    put(ws, f"C{r}", f'=IF(ISNUMBER(C{pf}),C{pf}/C{pc},C{pf})', fmt="0.0000")
    R["pg"], R["pf"] = pg, pf
    return r


# ================================================================== K1: Cap_In K section (below the reading row; no row above moves)
def cap_in_k(wb, h36):
    ws = wb["Cap_In"]
    r0 = h36["end"] + 2
    assert ws.max_row == h36["end"], (ws.max_row, h36["end"])
    put(ws, f"A{r0}", "K. 機隊口徑單價（v6.3 K1）：H10 m_fleet 的單價。欄位同 H10：C＝2025 年度（Alloc 的 AL_PriceFleet）、D／E／F＝2026 各欄以該欄的世代組合（H10-4 第 {q_mix_h}–{q_mix_vr} 列）重算；"
                          "公式同 Alloc I 節：Σ（世代占比 × 各世代付費產能 × 各世代單價）÷ Σ（世代占比 × 各世代付費產能），只含可服務世代；H7 讀 AL_PriceFleet"
        .format(q_mix_h=v534.ROWS["q_mix_h"], q_mix_vr=v534.ROWS["q_mix_vr"]), F_BOLD)
    r = r0 + 1
    mixrows = [v534.ROWS[k] for k in ("q_mix_h", "q_mix_gb200", "q_mix_gb300", "q_mix_vr")]
    put(ws, f"A{r}", "機隊口徑混合有效單價（付費 token；可達層級重新正規化）　→ H10 m_fleet", wrap=True); put(ws, f"B{r}", "$/M")
    put(ws, f"C{r}", "=AL_PriceFleet", F_CALC, fmt="0.0000", fill=FILL_KEY)
    for col in "DEF":
        s = [f"{col}{m}" for m in mixrows]
        cp = [f"INDEX(AL_CapPaid,1,{i})" for i in (1, 2, 3, 4)]; pg = [f"INDEX(AL_PriceGen,1,{i})" for i in (1, 2, 3, 4)]
        un = "OR(" + ",".join(f"AND({a}>0,NOT(ISNUMBER({b})))" for a, b in zip(s, cp)) + ")"
        num = "+".join(f"IF(AND({a}>0,ISNUMBER({b})),{a}*{b}*{p},0)" for a, b, p in zip(s, cp, pg))
        den = "+".join(f"IF(AND({a}>0,ISNUMBER({b})),{a}*{b},0)" for a, b in zip(s, cp))
        put(ws, f"{col}{r}", f'=IF({un},"{GENX}",IF(({den})>0,({num})/({den}),"{SLO}"))', F_CALC, fmt="0.0000")
    put(ws, f"G{r}", "Derived｜2026 欄的世代單價與付費產能為基準情境（AL_PriceGen、AL_CapPaid），只有世代組合不同", F_NOTE)
    wb.defined_names["B4_PriceFleet"] = DefinedName("B4_PriceFleet", attr_text=f"Cap_In!$C${r}:$F${r}")
    kr = r; r += 1
    put(ws, f"A{r}", "對照：完整層級組合混合有效單價（第 90 列；需求路線的模型 token 用此價）", wrap=True); put(ws, f"B{r}", "$/M")
    for col in "CDEF": put(ws, f"{col}{r}", "=$C$90", F_CALC, fmt="0.0000")
    kc = r; r += 1
    put(ws, f"A{r}", "機隊口徑單價 ÷ 完整層級組合單價", wrap=True); put(ws, f"B{r}", "x")
    for col in "CDEF": put(ws, f"{col}{r}", f"=IF(ISNUMBER({col}{kr}),{col}{kr}/{col}{kc},{col}{kr})", F_CALC, fmt="0.0000")
    return dict(start=r0, end=r, price_row=kr)


# ================================================================== J1 r1: guarded writes (Excel-owned cells; each only while the cell still holds its v6.2 value)
OLD_LOHI = {"lo": [0.1, 0.3, 1, 1, 1], "hi": [0.5, 1.0, 1, 4.4, 1]}
NEW_LOHI = {"lo": [0.5, 0.3, 1, 1, 1], "hi": [1.3, 1.0, 1, 2.5, 1]}
NOTE_LO_OLD_HEAD = "Assumed｜區間低端（相對 2026 Arch 快照的啟用參數倍數）：Hopper 數百億 ÷ 180B ≈ 0.1；GB200 0.3；GB300 與 VR200 為 1（J1）"
NOTE_LO_NEW = ("Assumed（Hopper、GB200 為 Analogy）｜區間低端（相對 2026 Arch 快照的啟用參數倍數）：Hopper 0.5、GB200 0.3、GB300 1、VR200 1、Rubin Ultra 固定 1。"
               "r1（2Q 報告 Q2）：Hopper 代表模型 GPT-4 級啟用參數約 220–280B，大於快照 Astra 的 180B，r0 的 0.1 方向錯誤，改 0.5")
NOTE_HI_OLD_HEAD = "Assumed｜區間高端：Hopper 0.5；GB200 1.0；GB300 1；VR200 400–800B ÷ 180B ≈ 4.4（J1）；R-b 讀法（快照對應 GB200）由 Project 另發 r1"
NOTE_HI_NEW = ("Assumed（Hopper、GB200 為 Analogy）｜區間高端：Hopper 1.3、GB200 1.0、GB300 1、VR200 2.5、Rubin Ultra 固定 1。"
               "r1（2Q 報告 Q2）：Hopper 1.3、VR200 2.5（r0 的 4.4 無證據）；情境 d1_model_scale 取 Hopper 0.8／GB200 0.55／GB300 1／VR200 1.4。暫不分層級（留待 Andy 決定）")


def inputs_update(wb):
    ws = wb["Spec_Rack"]; log = []
    r0 = next(r for r in range(20, ws.max_row + 1) if ws.cell(r, 1).value == v602.SEC_TITLE)
    for off, key, nm in ((2, "lo", "ARCH_ScaleGen_Lo"), (3, "hi", "ARCH_ScaleGen_Hi")):
        r = r0 + off
        assert ws.cell(r, 1).value.endswith(f"（{nm}）"), ws.cell(r, 1).value
        for col, old, new in zip("CDEFG", OLD_LOHI[key], NEW_LOHI[key]):
            c = ws[f"{col}{r}"]
            if c.value is not None and type(c.value) is not bool and c.value == old and old != new:
                c.value = new; log.append(f"Spec_Rack!{col}{r}: {old} -> {new} ({nm}, v6.3 J1 r1)")
        h = ws[f"H{r}"]
        old_h, new_h = (NOTE_LO_OLD_HEAD, NOTE_LO_NEW) if key == "lo" else (NOTE_HI_OLD_HEAD, NOTE_HI_NEW)
        if isinstance(h.value, str) and h.value == old_h:
            h.value = new_h; log.append(f"Spec_Rack!H{r}: note (v6.3 J1 r1)")
    return log


def gov_update(ws, append_text):
    n = 0
    for cell, col, old, new in (
            ("C45:G45", "M", "低端：Hopper 0.1、GB200 0.3、GB300 1、VR200 1", "低端：Hopper 0.5、GB200 0.3、GB300 1、VR200 1（r1，2Q Q2）"),
            ("C46:G46", "M", "高端：Hopper 0.5、GB200 1.0、GB300 1、VR200 4.4", "高端：Hopper 1.3、GB200 1.0、GB300 1、VR200 2.5（r1，2Q Q2）")):
        r = v536._gm_row(ws, "Spec_Rack", cell)
        if ws[f"{col}{r}"].value == old:
            ws[f"{col}{r}"].value = new; n += 1
        n += append_text(ws[f"N{r}"], " ｜v6.3 J1 r1（AF2-D1c，2Q 報告 Q2）：Hopper 0.8（0.5–1.3）、GB200 0.55（0.3–1.0）為 Analogy；VR200 1.4（1.0–2.5）為 Assumed；GB300＝1 為定義；Rubin Ultra 未設定，固定 1")
    r = v536._gm_row(ws, "Spec_Rack", "C44:G44")
    n += append_text(ws[f"N{r}"], " ｜v6.3 J1 r1：情境 d1_model_scale 取 2Q 值 0.8／0.55／1／1.4（基準列仍全 1）")
    return n


def decisions_update(wb):
    ws = wb["Decisions"]; log = []
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == "AF2-D1c":
            add = "｜r1（v6.3，2Q 報告 Q2）：依 2Q 修訂。情境 d1_model_scale＝Hopper 0.8／GB200 0.55／GB300 1／VR200 1.4（Rubin Ultra 固定 1）；區間 Lo＝0.5／0.3／1／1、Hi＝1.3／1.0／1／2.5；r0 的 0.25／0.5／1／2.5 方向與證據相反（GPT-4 啟用參數大於 Astra 180B）。暫不分層級。"
            if isinstance(ws.cell(r, 4).value, str) and add not in ws.cell(r, 4).value:
                ws.cell(r, 4).value = ws.cell(r, 4).value + add; log.append("Decisions AF2-D1c: r1 note (v6.3)")
    return log


# ================================================================== Decisions (10 columns)
_D = [
    ("K1", "機隊口徑單價與可達層級同口徑（價格端）", "與 Alloc 世代混合 token（機隊付費 token）配對的混合單價，層級權重同樣只在可達層級間重新正規化：Alloc 新增 I 節（AL_PriceGen 各世代單價、AL_PriceFleet 機隊單價＝Σ 世代占比 × 付費產能 × 世代單價 ÷ Σ 世代占比 × 付費產能），"
     "Cap_In H7（m_fleet、對照 m、m 下限、『必須為真』三列）讀 AL_PriceFleet，H10 m_fleet 讀 Cap_In K 節 B4_PriceFleet（2026 各欄以該欄世代組合），L1_PriceRealFleet、L1_PriceReqBE 讀 AL_PriceFleet。"
     "需求路線（模型 token AL_DPaid）仍用完整層級組合（Cap_In 第 90 列）。理由：D9 單價與 token 同口徑；同一個 m 的分子分母不可用兩組層級。世代間以付費 token 占比加權為本段的實作選擇（工作單只寫『比照 2-1-c』），見報告待 Project 判斷。",
     "Alloc I 節；Cap_In H7 第 119–125 列、H10 m_fleet、K 節；L1_MFleet、L1_LReqFleet、L1_PriceRealFleet、L1_PriceReqBE、L1_Gap*", "第 2b 節 K1"),
    ("K2", "整代不可達時不轉移世代權重", "撤回 v6.2 2-1-c 第二點：世代組合是觀察到的機隊事實，任一世代占比 > 0 且整代（三個層級）不可達時，不排除該世代、不把占比轉給其餘世代；AL_BlendPaid／Free、AL_ServeGW、AL_Q1／Q2、AL_ImpliedRDGW 等輸出文字旗標『整代不可達』，"
     "不出現數值或負值（Alloc G 表、H 節、Cap_In H10 世代混合同步）。層級層次的排除（K1、2-1-c 第一點）不變。AL_UnservedGen 保留顯示占比。基準不變（Hopper 的 Luna、Sol 可達）。",
     "Alloc B、G、H 節；Cap_In H10 世代混合列；AL_ServeGW、AL_Q1／Q2；L1_ScaleServe、L1_MFleet、L1_TokPerPromptClose、L1_Gap*（Tech_Registry T07／T09 情境）", "第 2b 節 K2"),
    ("K3", "TTFT 不可達的 decode 成本列也輸出文字", "Perf／Perf_Batch／Sens_Perf『decode（每 M GPU 秒）』列（原第 103 列）與 _Prod 鏡像：該欄 TTFT 符合上限為『否』時輸出文字旗標『TTFT 不可達』（先判 decode 是否可達→『SLO 不可達』，再判 TTFT）；"
     "讀該列的 decode 成本名稱隨之變文字（D8）。基準下只有 Hopper × Astra。", "Perf／Perf_Batch／Sens_Perf 及 _Prod 的 decode GPU 秒列；Unit_Cost、Interface 的 decode 成本名稱", "第 2b 節 K3"),
]
DECISIONS_V603 = [[f"AF2b-{i}", V, f"稽核修正第 2 批 2-1b：{t}", txt, DATE, ANDY, "v6.3 已建", where, f"{WO} {sec}", "否"] for i, t, txt, where, sec in _D]

# ================================================================== README
README_VERSION = (VERSION + "（次版：稽核修正第 2 批 2-1b 段，" + WO + "——K1 機隊口徑單價與可達層級同口徑（Alloc I 節 AL_PriceGen／AL_PriceFleet、Cap_In K 節 B4_PriceFleet）；"
                  "K2 整代不可達時不轉移世代權重（文字旗標『整代不可達』）；K3 TTFT 不可達的 decode 成本列輸出文字；J1 r1 依 2Q 報告修訂 ARCH_ScaleGen_Lo／_Hi 與情境 d1_model_scale；Decisions AF2b-K1／K2／K3）。以下為 "
                  + v602.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v602.README_TITLE
README_TITLE = VERSION.split("_")[-1] + _T[len(v602.VERSION.split("_")[-1]):]
README_ROW = ("機隊口徑單價、整代不可達、TTFT 成本（v6.3）",
              "稽核修正 2-1b：機隊 m 的單價與機隊 token 同用可達層級（Alloc I 節、Cap_In K 節）；整代不可達時輸出『整代不可達』而不轉移世代權重；TTFT 不可達的 decode 成本列輸出文字；"
              "模型規模倍數區間與情境依 2Q 報告修訂。標記與理由見 Gov_Map 與 Decisions AF2b-*。")
