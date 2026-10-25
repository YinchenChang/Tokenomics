# -*- coding: utf-8 -*-
"""產生審閱檔 第二出租方錨點.xlsx（換算全用公式；值只放在「輸入」頁）。
用法：python3 build_xlsx.py <輸出路徑>
TK 數值取自 model/20261024_Tokenomics_v21.4.xlsx（openpyxl data_only），在此寫成常數並於「輸入」頁註明來源格。
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

out = sys.argv[1]
wb = Workbook()
inp = wb.active
inp.title = '輸入'
hdr = ['鍵', '項目', '值', '單位', '標記', '來源']
inp.append(hdr)
REF = {}

def I(key, label, val, unit, tag, src):
    inp.append([key, label, val, unit, tag, src])
    REF[key] = "輸入!$C$%d" % inp.max_row

V = 'Verified'
IP = 'Interested-party'
TKsrc = 'TK v21.4（model/20261024_Tokenomics_v21.4.xlsx）'
# ---- TK ----
I('tk_hop', 'TK 持有成本 Hopper（AL_HoldGen／L1_HoldEconGW_Hopper）', 9.70308728611469, '$B/GW·年', 'Derived', TKsrc + ' Alloc!C39')
I('tk_gb200', 'TK 持有成本 GB200', 8.82816556596761, '$B/GW·年', 'Derived', TKsrc + ' Alloc!D39')
I('tk_gb300', 'TK 持有成本 GB300', 10.8864110021101, '$B/GW·年', 'Derived', TKsrc + ' Alloc!E39')
I('tk_vr200', 'TK 持有成本 VR200', 12.2254710045841, '$B/GW·年', 'Derived', TKsrc + ' Alloc!F39')
I('tk_gpu_hop', 'TK 每 IT GW GPU 數 Hopper', 718048, '顆/GW', 'Derived', TKsrc + ' Interface!D7')
I('tk_gpu_gb200', 'TK 每 IT GW GPU 數 GB200', 501768, '顆/GW', 'Derived', TKsrc + ' Interface!G7')
I('tk_gpu_gb300', 'TK 每 IT GW GPU 數 GB300', 487008, '顆/GW', 'Derived', TKsrc + ' Interface!J7')
I('tk_gpu_vr200', 'TK 每 IT GW GPU 數 VR200', 291744, '顆/GW', 'Derived', TKsrc + ' Interface!M7')
I('tk_capexit_lo', 'TK IF_CapexIT GB300 低', 28.5549024, '$B/GW', 'Derived', TKsrc + ' Interface!I9')
I('tk_capexit', 'TK IF_CapexIT GB300 基準', 32.237224, '$B/GW', 'Derived', TKsrc + ' Interface!J9')
I('tk_capexit_hi', 'TK IF_CapexIT GB300 高', 39.2826064, '$B/GW', 'Derived', TKsrc + ' Interface!K9')
I('tk_capexfac', 'TK IF_CapexFacility 基準', 12.67, '$B/GW', 'Derived', TKsrc + ' Interface!J10')
I('tk_opex_lo', 'TK IF_OpexGW GB300 低（含電費、不含折舊）', 1.20671911367581, '$B/GW·年', 'Derived', TKsrc + ' Interface!I228')
I('tk_opex', 'TK IF_OpexGW GB300 基準', 1.98263481314457, '$B/GW·年', 'Derived', TKsrc + ' Interface!J228')
I('tk_opex_hi', 'TK IF_OpexGW GB300 高', 3.00013873514752, '$B/GW·年', 'Derived', TKsrc + ' Interface!K228')
I('tk_racks_gb300', 'TK 每 IT GW 機架數 GB300（IF_RacksPerGW）', 6764, '架/GW', 'Derived', TKsrc + ' Interface!J6')
I('tk_rack_lo', 'TK GB300 機架價 低（IF_RackPrice）', 4.0, '$M/架', 'Derived', TKsrc + ' Interface!I364')
I('tk_rack', 'TK GB300 機架價 基準', 4.3, '$M/架', 'Derived', TKsrc + ' Interface!J364')
I('tk_rack_hi', 'TK GB300 機架價 高', 5.0, '$M/架', 'Derived', TKsrc + ' Interface!K364')
I('tk_draw_lo', 'TK 平均用電 ÷ 設計功率 低（IF_AvgDraw）', 0.7, 'x', 'Derived', TKsrc + ' Interface!I222')
I('tk_draw_hi', 'TK 平均用電 ÷ 設計功率 高', 0.9, 'x', 'Derived', TKsrc + ' Interface!K222')
I('pue_lo', 'PUE 低（若 MW 為設施電力）', 1.2, 'x', 'Assumed', '沿用 CoreWeave 錨點')
I('pue_hi', 'PUE 高', 1.3, 'x', 'Assumed', 'TK 設施係數區間 1.1／1.2／1.3')
I('sb', 'AL_SpendBasis 基準', 1.8, 'x', 'Assumed', TKsrc + ' Alloc_In!C23（區間 0.85–2.3）')
# ---- CoreWeave ----
I('cw_q2rev', 'CoreWeave Q2 2026 營收', 2.575, '$B/季', V + '（10-Q）', 'CoreWeave 10-Q 2026-06-30；8-K 2026-08-11 新聞稿同')
I('cw_h1rev', 'CoreWeave H1 2026 營收', 4.653, '$B', V + '（10-Q）', '同上')
I('cw_q2end', 'CoreWeave 6/30 active power', 1.5, 'GW', IP + '（8-K；口徑未定義）', '8-K 2026-08-11「reach 1.5 GW」')
I('cw_q2add', 'CoreWeave Q2 新增 active power', 0.5, 'GW', IP + '（8-K）', '8-K 2026-08-11「nearly 500 MWs」')
I('cw_junadd', '其中 6 月新增', 0.3, 'GW', IP + '（法說）', 'Q2 法說（webull 轉載逐字稿，二手）「in June, we brought on more than 300 megawatts」')
I('cw_ye25', 'CoreWeave 2025-12-31 active power', 0.85, 'GW', IP + '（10-K 業務章節）', '10-K FY2025「over 850 MW of active power」（下限）')
I('cw_q1end', 'CoreWeave 2026-03-31 active power', 1.0, 'GW', IP + '（8-K）', '8-K 2026-05-07「Surpassed 1 GW of active power」（下限）')
I('cw_price', '2026-07 全 SKU 漲價', 1.25, 'x', IP + '（法說）', 'Q2 法說「approximately 25% increase across SKUs」')
I('cw_q2cor', 'CoreWeave Q2 2026 cost of revenue（不含折舊；含機房租金與電費）', 0.879, '$B/季', V + '（10-Q）', '8-K 2026-08-11 新聞稿損益表 879（$M）')
# ---- Nebius ----
I('nb_arr_ye', 'Nebius ARR 2025-12-31（AI cloud 12 月營收 × 12）', 1.25, '$B/年', IP, '6-K 2026-02-12 附件 99.2（Q4 2025 股東信）')
I('nb_arr_q3', 'Nebius ARR 2025-09-30', 0.551, '$B/年', IP, '同上「more than doubled from $551 million reported as of the end of September」')
I('nb_act_ye', 'Nebius 2025-12-31 active power（約數）', 0.170, 'GW', IP, '同上「~170 MW of active power」；20-F 同')
I('nb_q3cap', 'Nebius 2025-09-30 active power 上界（Q3 信：年底目標 100 MW）', 0.100, 'GW', IP + '／Assumed（上界）', '6-K 2025-11-11 附件 99.2「on track for … 100 MW of active power by the end of 2025」')
I('nb_q4rev', 'Nebius Q4 2025 core AI cloud 營收', 0.2142, '$B/季', IP, 'Q4 2025 股東信「Core AI cloud revenue was $214.2 million」（20-F 年度數 Verified，季度數未查核）')
I('nb_base26', 'Nebius「2026 base」ACV per MW', 12, '$M/MW·年', IP, '6-K 2026-08-12 附件 99.2 第 2 頁圖「ACV per MW is stepping up … 2026 base $12M」；註 (1) ACV＝Annual contract value；(2) 營收認列口徑、不含預付')
I('nb_q2lo', 'Nebius Q2 2026 四筆大單 yield 低', 20, '$M/MW·年', IP, '同上「a yield of $20-25 million per megawatt」')
I('nb_q2hi', 'Nebius Q2 2026 四筆大單 yield 高', 25, '$M/MW·年', IP, '同上')
I('nb_stlo', 'Nebius Q3 短期容量 價格區間 低', 40, '$M/MW·年', IP, '同上「price opportunity in the $40-50 million per MW range」（已簽第一筆）')
I('nb_sthi', 'Nebius Q3 短期容量 價格區間 高', 50, '$M/MW·年', IP, '同上')
I('nb_payback', 'Nebius Q2 新約回本期', 1 + 10 / 12, '年', IP, '同上「payback period for the associated capex and related operating costs for Q2 deals is 1 year and 10 months」')
# ---- IREN ----
I('ir_tcv', 'IREN–Microsoft 合約總值', 9.7, '$B', IP, 'IREN 新聞稿 2025-11-03「approximately $9.7 billion」（含 20% 預付）')
I('ir_yrs', 'IREN–Microsoft 合約年數', 5, '年', IP, '同上「five-year term」')
I('ir_mw', 'IREN Childress Horizon 1–4 關鍵 IT', 200, 'MW-IT', IP + '／Assumed（全數對應本合約）', '同上「200MW of critical IT load (Horizon 1–4)」；對應關係原文未明寫（TK Decisions D139 同註）')
I('ir_newlo', 'IREN 近期 3 年約 每 MW(IT) 營收（下限）', 20, '$M/MW-IT·年', IP, 'IREN FY26 結果新聞稿 2026-08-27「Recent 3-year contracts >$20m revenue per MW (IT)」')
I('ir_newhi', 'IREN 洽談中 每 MW(IT)', 25, '$M/MW-IT·年', IP, '同上「Active discussions at ~$25m per MW (IT)」')
I('ir_payback', 'IREN 近期合約回本期（GPU 與附屬設備）', 2, '年', IP, '同上「~2 year payback」；註 3：GPU 與附屬設備資本支出 ÷（合約營收 − 直接成本）')
I('ir_q4rev', 'IREN Q4 FY26 AI Cloud 營收', 70.5, '$M/季', IP, 'IREN FY26 結果新聞稿 2026-08-27（季度數未查核）')
I('ir_q4cor', 'IREN Q4 FY26 AI Cloud cost of revenue', 9.2, '$M/季', IP, '同上')
I('ir_dell', 'IREN 向 Dell 採購 GB300 設備', 5.8, '$B', IP, 'IREN 新聞稿 2025-11-03（含伺服器、InfiniBand、線材、軟體授權、部署）；TK SRC_DC_014')
# ---- WhiteFiber ----
I('wf_cloud', 'WhiteFiber Q2 2026 雲端服務營收', 23.806, '$M/季', V + '（10-Q）', 'WhiteFiber 10-Q 2026-06-30 損益表')
I('wf_term', '其中 Initial Customer 終止費（一次性）', 12.3, '$M', V + '（10-Q）', '同上「fixed termination fee of $12.3 million that was recognized as revenue」')
I('wf_ice', 'WhiteFiber 冰島已簽約 IT 負載', 5.5, 'MW-IT', IP + '（10-Q MD&A）', '同上「executed contracts for 5.5 MW IT load」；亞特蘭大 2026-02 起另租容量，MW 未揭露')
contracts = [
    ('wfc1', 'Baseten 1,392 B300（36 個月，2026-11 起）', 165.2, 3, 1392, 'gb300', 'GB300'),
    ('wfc2', '冰島 576 B300（5 年，2026-12 起；另有營收分成）', 87.5, 5, 576, 'gb300', 'GB300'),
    ('wfc3', '新客戶 128 B300（36 個月，2026-07 簽）', 16.0, 3, 128, 'gb300', 'GB300'),
    ('wfc4', 'Prime Intellect 576 VR200（36 個月，2027 Q2 起）', 108.2, 3, 576, 'vr200', 'VR200'),
    ('wfc5', 'Prime Intellect 216 GB200（12 個月，2026-07 起）', 6.8, 1, 216, 'gb200', 'GB200'),
    ('wfc6', '既有客戶 384 B200（24 個月，2026-01 起）', 18.1, 2, 384, 'gb200', 'GB200'),
    ('wfc7', '既有客戶 256 台 H100 伺服器（24 個月，2026 Q2 起；以 8 顆／台計 2,048 顆）', 50.2, 2, 2048, 'hop', 'Hopper'),
]
for k, lab, tcv, yrs, g, gen, gname in contracts:
    I(k + '_tcv', 'WhiteFiber 合約 ' + lab + '：總值', tcv, '$M', V + '（10-Q 揭露）', 'WhiteFiber 10-Q 2026-06-30 Recent Developments')
    I(k + '_yrs', '　同上：初始年數', yrs, '年', V + '（10-Q 揭露）', '同上')
    I(k + '_gpu', '　同上：GPU 數', g, '顆', (V + '（10-Q 揭露）') if k != 'wfc7' else 'Assumed（8 顆／台）', '同上' + ('；B300 以 GB300 密度、B200 以 GB200 密度近似（Analogy）' if gname in ('GB300', 'GB200') else ''))
# ---- GB300 單價點 ----
I('dg_rack', 'Data Gravity：GB300 NVL72 採購單價（不含 CDU）', 5.0, '$M/架', 'Analogy（第三方，採購單）', 'wing.vc 轉載 Data Gravity 2026-08「just under $5.0M」；TK SRC_HW_010')
I('dg_dep', 'Data Gravity：GB300 NVL72 部署後合計', 5.68, '$M/架', 'Analogy', '同上「closer to $5.7 million」（不含 DC 建造、叢集核心交換、儲存）')
I('th_lo', 'Tom\'s Hardware：GB300 NVL72 每架（媒體估計）低', 6.0, '$M/架', 'Analogy（媒體）', '同上轉述「$6–6.5M」；TK SRC_HW_007（Alt）')
I('th_hi', 'Tom\'s Hardware：每架 高', 6.5, '$M/架', 'Analogy（媒體）', '同上')
I('dg_oplo', 'Data Gravity：每架年營運成本 低（含機房租金約 31.8 萬）', 0.240, '$M/架·年', 'Analogy', '同上「$240–410K」')
I('dg_ophi', 'Data Gravity：每架年營運成本 高', 0.410, '$M/架·年', 'Analogy', '同上')

for c, w in zip('ABCDEF', [16, 62, 14, 14, 30, 90]):
    inp.column_dimensions[c].width = w
for cell in inp[1]:
    cell.font = Font(bold=True)

R = lambda k: REF[k]

def sheet(name, widths=(58, 16, 16, 16, 60)):
    ws = wb.create_sheet(name)
    for c, w in zip('ABCDE', widths):
        ws.column_dimensions[c].width = w
    return ws

def put(ws, row, vals, bold=False):
    for i, v in enumerate(vals):
        c = ws.cell(row=row, column=i + 1, value=v)
        if bold:
            c.font = Font(bold=True)

# ================= CoreWeave =================
cw = sheet('CoreWeave')
put(cw, 1, ['CoreWeave（承接 PR #163 錨點；同算法重算 Q2，並加 Q1 2026 點）'], True)
put(cw, 3, ['A. 季平均 active GW', '值', '', '', '說明'], True)
put(cw, 4, ['3/31 active GW（＝6/30 − Q2 新增）', '=%s-%s' % (R('cw_q2end'), R('cw_q2add')), '', '', 'GW；與 8-K「surpassed 1 GW」一致'])
put(cw, 5, ['Q2 平均（線性）', '=(B4+%s)/2' % R('cw_q2end'), '', '', 'GW'])
put(cw, 6, ['Q2 平均（6 月集中）', '=B4+(%s-%s)*(2/3)+%s*(1/6)' % (R('cw_q2add'), R('cw_junadd'), R('cw_junadd')), '', '', 'GW；4–5 月新增計 2/3 季、6 月新增計 1/6 季'])
put(cw, 7, ['Q1 營收（H1 − Q2）', '=%s-%s' % (R('cw_h1rev'), R('cw_q2rev')), '', '', '$B'])
put(cw, 8, ['Q1 平均（線性；12/31 0.85 → 3/31 1.0）', '=(%s+%s)/2' % (R('cw_ye25'), R('cw_q1end')), '', '', 'GW；兩端皆為「逾」，平均偏低 → 每 GW 營收偏高'])
put(cw, 10, ['B. 每 active GW 年營收（$B/GW·年）', 'Q2 線性', 'Q2 6 月集中', 'Q1 線性', ''], True)
put(cw, 11, ['IT 口徑', '=%s*4/B5' % R('cw_q2rev'), '=%s*4/B6' % R('cw_q2rev'), '=B7*4/B8', ''])
put(cw, 12, ['設施口徑 PUE 1.2（IT＝設施 ÷ PUE）', '=B11*%s' % R('pue_lo'), '=C11*%s' % R('pue_lo'), '=D11*%s' % R('pue_lo'), ''])
put(cw, 13, ['設施口徑 PUE 1.3', '=B11*%s' % R('pue_hi'), '=C11*%s' % R('pue_hi'), '=D11*%s' % R('pue_hi'), ''])
put(cw, 15, ['C. ÷ TK 持有成本', 'Hopper', 'GB200', 'GB300', ''], True)
put(cw, 16, ['Q2 IT 線性', '=B11/%s' % R('tk_hop'), '=B11/%s' % R('tk_gb200'), '=B11/%s' % R('tk_gb300'), ''])
put(cw, 17, ['Q2 IT 6 月集中', '=C11/%s' % R('tk_hop'), '=C11/%s' % R('tk_gb200'), '=C11/%s' % R('tk_gb300'), ''])
put(cw, 18, ['Q2 設施 PUE 1.2 線性', '=B12/%s' % R('tk_hop'), '=B12/%s' % R('tk_gb200'), '=B12/%s' % R('tk_gb300'), ''])
put(cw, 19, ['Q2 設施 PUE 1.3 6 月集中', '=C13/%s' % R('tk_hop'), '=C13/%s' % R('tk_gb200'), '=C13/%s' % R('tk_gb300'), ''])
put(cw, 20, ['Q1 IT 線性', '=D11/%s' % R('tk_hop'), '=D11/%s' % R('tk_gb200'), '=D11/%s' % R('tk_gb300'), ''])
put(cw, 21, ['Q1 設施 PUE 1.3', '=D13/%s' % R('tk_hop'), '=D13/%s' % R('tk_gb200'), '=D13/%s' % R('tk_gb300'), ''])
put(cw, 23, ['Q2 區間（GB200／GB300 欄，同 PR #163）：最小', '=MIN(C16:D19)', '', '', 'PR #163 為 0.76'])
put(cw, 24, ['Q2 區間：最大', '=MAX(C16:D19)', '', '', 'PR #163 為 1.28'])
put(cw, 25, ['Q2 區間 × 7 月漲價：最小', '=B23*%s' % R('cw_price'), '', '', ''])
put(cw, 26, ['Q2 區間 × 7 月漲價：最大', '=B24*%s' % R('cw_price'), '', '', ''])
put(cw, 27, ['Q1 區間（Hopper／GB200／GB300，IT 與 PUE 1.3）：最小', '=MIN(B20:D21)', '', '', ''])
put(cw, 28, ['Q1 區間：最大', '=MAX(B20:D21)', '', '', ''])

# ================= Nebius =================
nb = sheet('Nebius')
put(nb, 1, ['Nebius（NBIS）：active power＝「capacity being consumed by IT equipment and available for revenue generation」（20-F）→ IT 側'], True)
put(nb, 3, ['A. 機隊平均：2025 Q4（唯一有 active MW 的時點）', '值', '', '', '說明'], True)
put(nb, 4, ['9/30 active GW 估計（依 ARR 比例）', '=%s*%s/%s' % (R('nb_act_ye'), R('nb_arr_q3'), R('nb_arr_ye')), '', '', 'Assumed：每 GW 營收 9 月與 12 月相同'])
put(nb, 5, ['9/30 active GW 上界', '=%s' % R('nb_q3cap'), '', '', 'Q3 信稱年底目標 100 MW → 9 月底應低於此'])
put(nb, 6, ['Q4 平均（線性，ARR 比例起點）', '=(B4+%s)/2' % R('nb_act_ye'), '', '', 'GW'])
put(nb, 7, ['Q4 平均（線性，上界起點）', '=(B5+%s)/2' % R('nb_act_ye'), '', '', 'GW；平均偏大 → 每 GW 營收偏低（保守端）'])
put(nb, 9, ['每 active GW 年營收（$B/GW·年）', '值', '', '', ''], True)
put(nb, 10, ['(1) ARR ÷ 年底 active', '=%s/%s' % (R('nb_arr_ye'), R('nb_act_ye')), '', '', '12 月營收 × 12 ÷ 年底 MW；12 月平均 MW 低於年底 → 此值略偏低'])
put(nb, 11, ['(2) Q4 營收 × 4 ÷ Q4 平均（ARR 比例起點）', '=%s*4/B6' % R('nb_q4rev'), '', '', 'core AI cloud 營收'])
put(nb, 12, ['(3) Q4 營收 × 4 ÷ Q4 平均（上界起點）', '=%s*4/B7' % R('nb_q4rev'), '', '', '保守端'])
put(nb, 14, ['÷ TK 持有成本（機隊：Hopper 為主，B200／GB200 次之，B300／GB300 剛開始）', 'Hopper', 'GB200', 'GB300', ''], True)
for i, r in enumerate((10, 11, 12)):
    put(nb, 15 + i, ['(%d)' % (i + 1), '=B%d/%s' % (r, R('tk_hop')), '=B%d/%s' % (r, R('tk_gb200')), '=B%d/%s' % (r, R('tk_gb300')), ''])
put(nb, 18, ['機隊區間（Hopper／GB200 欄）：最小', '=MIN(B15:C17)', '', '', ''])
put(nb, 19, ['機隊區間：最大', '=MAX(B15:C17)', '', '', ''])
put(nb, 20, ['敏感度：若 active 為實際耗電（設計功率＝active ÷ 平均用電比 0.7–0.9）→ 最小', '=B18*%s' % R('tk_draw_lo'), '', '', '本段判斷不採（「capacity … available for revenue generation」讀作可用容量）'])
put(nb, 21, ['同上 → 最大', '=B19*%s' % R('tk_draw_hi'), '', '', ''])
put(nb, 23, ['B. 新約與存量合約單價（ACV per MW；MW 口徑未定義，以 IT 計）', '單價', '÷GB200', '÷GB300', '÷VR200'], True)
put(nb, 24, ['「2026 base」ACV per MW', '=%s' % R('nb_base26'), '=B24/%s' % R('tk_gb200'), '=B24/%s' % R('tk_gb300'), '=B24/%s' % R('tk_vr200')])
put(nb, 25, ['Q2 四筆大單 低', '=%s' % R('nb_q2lo'), '=B25/%s' % R('tk_gb200'), '=B25/%s' % R('tk_gb300'), '=B25/%s' % R('tk_vr200')])
put(nb, 26, ['Q2 四筆大單 高', '=%s' % R('nb_q2hi'), '=B26/%s' % R('tk_gb200'), '=B26/%s' % R('tk_gb300'), '=B26/%s' % R('tk_vr200')])
put(nb, 27, ['Q3 短期（≤6 個月）低', '=%s' % R('nb_stlo'), '=B27/%s' % R('tk_gb200'), '=B27/%s' % R('tk_gb300'), '=B27/%s' % R('tk_vr200')])
put(nb, 28, ['Q3 短期 高', '=%s' % R('nb_sthi'), '=B28/%s' % R('tk_gb200'), '=B28/%s' % R('tk_gb300'), '=B28/%s' % R('tk_vr200')])
put(nb, 29, ['若 MW 為設施電力：Q2 大單 ÷GB300 × PUE 1.2（低）', '', '', '=D25*%s' % R('pue_lo'), ''])
put(nb, 30, ['同上 × PUE 1.3（高）', '', '', '=D26*%s' % R('pue_hi'), ''])
put(nb, 32, ['C. 回本期檢查（年化讀法是否成立）', '值', '', '', '說明'], True)
put(nb, 33, ['隱含資本支出（Q2 大單低，扣 TK opex 基準）', '=%s*(B25-%s)' % (R('nb_payback'), R('tk_opex')), '', '', '$B/GW；回本期 × (ACV − 營運費用)'])
put(nb, 34, ['隱含資本支出（Q2 大單高）', '=%s*(B26-%s)' % (R('nb_payback'), R('tk_opex')), '', '', ''])
put(nb, 35, ['對照 TK GB300 IT＋廠房資本支出（基準）', '=%s+%s' % (R('tk_capexit'), R('tk_capexfac')), '', '', '若 ACV 是總額而非年額，隱含資本支出只有約 1/3，與任何口徑都不符 → 年化讀法成立'])
put(nb, 36, ['隱含 ÷ TK（低／高）', '=B33/B35', '=B34/B35', '', ''])

# ================= IREN =================
ir = sheet('IREN')
put(ir, 1, ['IREN（AI cloud 部分）：MW 標明「(IT)」；機隊 MW 未揭露 → 只做合約層'], True)
put(ir, 3, ['A. Microsoft 合約（2025-11，GB300，5 年）', '值', '÷GB300', '', '說明'], True)
put(ir, 4, ['年營收', '=%s/%s' % (R('ir_tcv'), R('ir_yrs')), '', '', '$B/年（含 20% 預付，按直線平均）'])
put(ir, 5, ['每 IT GW 年營收', '=B4/(%s/1000)' % R('ir_mw'), '=B5/%s' % R('tk_gb300'), '', '$B/GW·年；TK DB_Evidence E297 已記 9.70'])
put(ir, 7, ['B. 近期合約（2026，3–5 年）', '單價', '÷GB300', '÷VR200', ''], True)
put(ir, 8, ['近期 3 年約（下限）', '=%s' % R('ir_newlo'), '=B8/%s' % R('tk_gb300'), '=B8/%s' % R('tk_vr200'), '「>$20m」；GPU 為「ordered quite some time ago」→ 以 GB300 為主'])
put(ir, 9, ['洽談中', '=%s' % R('ir_newhi'), '=B9/%s' % R('tk_gb300'), '=B9/%s' % R('tk_vr200'), ''])
put(ir, 11, ['C. 直接成本率與營運費用', '值', '', '', ''], True)
put(ir, 12, ['Q4 FY26 AI Cloud 直接成本 ÷ 營收', '=%s/%s' % (R('ir_q4cor'), R('ir_q4rev')), '', '', '不含折舊；IREN 自有機房，主要為電費與站點營運'])
put(ir, 13, ['隱含營運費用：單價 20 × 直接成本率', '=B12*B8', '', '', '$B/GW·年；對照 TK IF_OpexGW GB300 1.98（1.21–3.00）'])
put(ir, 14, ['隱含營運費用：單價 25 × 直接成本率', '=B12*B9', '', '', ''])
put(ir, 16, ['D. 回本期隱含 GPU 資本支出（註 3 定義）', '值', '÷TK IT', '', ''], True)
put(ir, 17, ['回本 × (20 − 隱含營運費用)', '=%s*(B8-B13)' % R('ir_payback'), '=B17/%s' % R('tk_capexit'), '', '$B/GW；GPU 與附屬設備'])
put(ir, 18, ['回本 × (25 − 隱含營運費用)', '=%s*(B9-B14)' % R('ir_payback'), '=B18/%s' % R('tk_capexit'), '', ''])
put(ir, 19, ['對照：Dell 採購 ÷ 200 MW-IT（2025-11）', '=%s/(%s/1000)' % (R('ir_dell'), R('ir_mw')), '=B19/%s' % R('tk_capexit'), '', 'TK SRC_DC_014＝29.0'])

# ================= WhiteFiber =================
wf = sheet('WhiteFiber', widths=(66, 14, 14, 14, 14, 14, 14))
put(wf, 1, ['WhiteFiber（WYFI）：雲端分部；MW 以 IT 負載揭露（冰島 5.5 MW-IT）'], True)
put(wf, 3, ['A. 機隊平均：2026 Q2', '值', '', '', '說明'], True)
put(wf, 4, ['經常性雲端營收（扣終止費）', '=%s-%s' % (R('wf_cloud'), R('wf_term')), '', '', '$M/季'])
put(wf, 5, ['每 IT GW 年營收（÷ 冰島 5.5 MW）', '=B4*4/%s' % R('wf_ice'), '', '', '$M/MW·年＝$B/GW·年；亞特蘭大等未計入 MW → 偏上界'])
put(wf, 6, ['÷ TK 持有成本 Hopper', '=B5/%s' % R('tk_hop'), '', '', '機隊 H100／H200／B200 為主'])
put(wf, 7, ['÷ TK 持有成本 GB200', '=B5/%s' % R('tk_gb200'), '', '', ''])
put(wf, 9, ['B. 合約層（GPU 數 ÷ TK 每 GW GPU 數換算 IT GW；Analogy）', '年營收 $M', 'IT MW', '$B/GW·年', 'TK 持有成本', '比值', '世代'], True)
gpuref = {'hop': 'tk_gpu_hop', 'gb200': 'tk_gpu_gb200', 'gb300': 'tk_gpu_gb300', 'vr200': 'tk_gpu_vr200'}
holdref = {'hop': 'tk_hop', 'gb200': 'tk_gb200', 'gb300': 'tk_gb300', 'vr200': 'tk_vr200'}
r = 10
for k, lab, tcv, yrs, g, gen, gname in contracts:
    put(wf, r, [lab, '=%s/%s' % (R(k + '_tcv'), R(k + '_yrs')), '=%s/%s*1000' % (R(k + '_gpu'), R(gpuref[gen])), '=B%d/C%d' % (r, r), '=%s' % R(holdref[gen]), '=D%d/E%d' % (r, r), gname])
    r += 1
put(wf, r + 1, ['Blackwell 新約（B300／GB300／GB200、2026 下半年起）比值：最小', '=MIN(F10:F12,F14)', '', '', '', '', ''])
put(wf, r + 2, ['同上：最大', '=MAX(F10:F12,F14)', '', '', '', '', ''])
put(wf, r + 3, ['存量世代續約（B200 2026-01、H100 2026 Q2）：最小', '=MIN(F15:F16)', '', '', '', '', ''])
put(wf, r + 4, ['同上：最大', '=MAX(F15:F16)', '', '', '', '', ''])
put(wf, r + 5, ['VR200（2027 Q2 起）', '=F13', '', '', '', '', ''])
WF_ROWS = {'bw_min': r + 1, 'bw_max': r + 2, 'old_min': r + 3, 'old_max': r + 4, 'vr': r + 5}

# ================= 單價與營運費用 =================
gp = sheet('GB300單價與營運費用')
put(gp, 1, ['GB300 NVL72 IT 資本支出與營運費用：公開點對照 TK（只列，不改 TK）'], True)
put(gp, 3, ['A. 每 IT GW 的 IT 資本支出（$B/GW）', '值', '÷TK 基準', '', '口徑'], True)
put(gp, 4, ['TK IF_CapexIT GB300 低／基準／高', '=%s' % R('tk_capexit'), '=B4/%s' % R('tk_capexit'), '', '=TEXT(%s,"0.0")&"／"&TEXT(%s,"0.0")&"／"&TEXT(%s,"0.0")&"；機架＋scale-out 網路＋儲存與管理"' % (R('tk_capexit_lo'), R('tk_capexit'), R('tk_capexit_hi'))])
put(gp, 5, ['TK 機架價 4.3 × 6,764 架（只機架）', '=%s*%s/1000' % (R('tk_rack'), R('tk_racks_gb300')), '=B5/%s' % R('tk_capexit'), '', 'IF_RackPrice 基準；低 4.0、高 5.0'])
put(gp, 6, ['IREN Dell 採購 ÷ 200 MW-IT（2025-11）', '=IREN!B19', '=B6/%s' % R('tk_capexit'), '', '含伺服器、IB、線材、軟體、部署；TK 已有（SRC_DC_014）'])
put(gp, 7, ['Data Gravity 採購單 5.0 × 6,764 架（只機架，2026-08）', '=%s*%s/1000' % (R('dg_rack'), R('tk_racks_gb300')), '=B7/%s' % R('tk_capexit'), '', '不含 CDU、網路、儲存；TK 高情境機架價即此值'])
put(gp, 8, ['Data Gravity 部署後 5.68 × 6,764 架', '=%s*%s/1000' % (R('dg_dep'), R('tk_racks_gb300')), '=B8/%s' % R('tk_capexit'), '', '含部署項目；不含叢集核心交換與儲存'])
put(gp, 9, ['媒體 6.0 × 6,764 架', '=%s*%s/1000' % (R('th_lo'), R('tk_racks_gb300')), '=B9/%s' % R('tk_capexit'), '', 'TK 列 Alt'])
put(gp, 10, ['媒體 6.5 × 6,764 架', '=%s*%s/1000' % (R('th_hi'), R('tk_racks_gb300')), '=B10/%s' % R('tk_capexit'), '', ''])
put(gp, 11, ['IREN 回本期隱含（單價 20）', '=IREN!B17', '=B11/%s' % R('tk_capexit'), '', 'Derived；GPU 與附屬設備'])
put(gp, 12, ['IREN 回本期隱含（單價 25）', '=IREN!B18', '=B12/%s' % R('tk_capexit'), '', ''])
put(gp, 13, ['Nebius 回本期隱含（單價 20／25）÷ TK IT＋廠房', '=Nebius!B33', '=Nebius!B36', '=Nebius!C36', '「associated capex」範圍未定義；C、D 欄為 ÷ TK IT＋廠房 44.9'])
put(gp, 15, ['B. 每 IT GW 年營運費用（不含折舊，$B/GW·年）', '值', '÷TK 基準', '', '口徑'], True)
put(gp, 16, ['TK IF_OpexGW GB300 基準（低 1.21／高 3.00）', '=%s' % R('tk_opex'), '=B16/%s' % R('tk_opex'), '', '含電費、維護、人員、稅保；不含廠房資本成本'])
put(gp, 17, ['IREN 直接成本率 × 單價 20', '=IREN!B13', '=B17/%s' % R('tk_opex'), '', 'Derived；直接成本率 13% 來自 Q4 FY26 混合機隊'])
put(gp, 18, ['IREN 直接成本率 × 單價 25', '=IREN!B14', '=B18/%s' % R('tk_opex'), '', ''])
put(gp, 19, ['CoreWeave Q2 cost of revenue × 4 ÷ Q2 平均 active GW（線性）', '=%s*4/CoreWeave!B5' % R('cw_q2cor'), '=B19/%s' % R('tk_opex'), '', '含機房租金（託管）→ 應對 TK opex＋廠房資本成本，不可直接比'])
put(gp, 20, ['Data Gravity 每架年營運 0.24 × 6,764 架', '=%s*%s/1000' % (R('dg_oplo'), R('tk_racks_gb300')), '=B20/%s' % R('tk_opex'), '', '範圍口徑未明（文中託管租金約 0.318／架落在區間內）；只作量級'])
put(gp, 21, ['Data Gravity 每架年營運 0.41 × 6,764 架', '=%s*%s/1000' % (R('dg_ophi'), R('tk_racks_gb300')), '=B21/%s' % R('tk_opex'), '', ''])

# ================= 彙總 =================
sm = wb.create_sheet('彙總', 0)
for c, w in zip('ABCDEFGH', [14, 30, 16, 38, 12, 12, 30, 40]):
    sm.column_dimensions[c].width = w
put(sm, 1, ['出租方計價倍數錨點彙總：每 GW 年營收 ÷ TK 世代持有成本（AL_SpendBasis 基準 1.8，區間 0.85–2.3）'], True)
put(sm, 3, ['出租方', '項目', '類型', '電力口徑', '比值 低', '比值 高', '標記', '說明'], True)
rows = [
    ('CoreWeave', 'Q2 2026 每 active GW（PR #163）', '機隊平均', 'active power，口徑未定義（IT／設施兩算）', '=CoreWeave!B23', '=CoreWeave!B24', '營收 Verified；MW Interested-party', '÷GB200／GB300'),
    ('CoreWeave', 'Q2 2026 × 7 月漲價 1.25', '機隊平均（漲價後）', '同上', '=CoreWeave!B25', '=CoreWeave!B26', '漲價為法說', ''),
    ('CoreWeave', 'Q1 2026 每 active GW', '機隊平均', '同上', '=CoreWeave!B27', '=CoreWeave!B28', '營收 Verified（H1−Q2）', '÷Hopper／GB200／GB300；IT 與 PUE 1.3'),
    ('Nebius', 'Q4 2025 每 active GW', '機隊平均', 'active＝IT 設備耗用容量（公司定義）', '=Nebius!B18', '=Nebius!B19', 'Interested-party；9 月底 MW 為 Assumed', '÷Hopper／GB200'),
    ('Nebius', '「2026 base」ACV per MW', '存量合約平均', 'MW 未定義（以 IT 計）', '=Nebius!D24', '=Nebius!C24', 'Interested-party', '÷GB300 至 ÷GB200'),
    ('Nebius', 'Q2 2026 四筆大單 $20–25M/MW', '新約邊際', '同上', '=Nebius!D25', '=Nebius!D26', 'Interested-party', '÷GB300；÷VR200 為 1.64–2.04'),
    ('Nebius', 'Q3 短期 $40–50M/MW', '短約邊際', '同上', '=Nebius!D27', '=Nebius!D28', 'Interested-party', '÷GB300；≤6 個月'),
    ('IREN', 'Microsoft 合約（2025-11）', '長約（簽約時）', 'MW-IT（公司標明）', '=IREN!C5', '=IREN!C5', 'Interested-party；MW 對應 Assumed', '÷GB300'),
    ('IREN', '近期 3 年約 >$20M／洽談 ~$25M per MW(IT)', '新約邊際', 'MW-IT（公司標明）', '=IREN!C8', '=IREN!C9', 'Interested-party', '÷GB300；÷VR200 為 1.64–2.04'),
    ('WhiteFiber', 'Q2 2026 經常性雲端 ÷ 冰島 5.5 MW-IT', '機隊平均（偏上界）', 'IT（公司揭露 IT 負載）', '=MIN(WhiteFiber!B6:B7)', '=MAX(WhiteFiber!B6:B7)', '營收 Verified（10-Q）；MW Interested-party', '÷GB200／Hopper'),
    ('WhiteFiber', 'Blackwell 新約（B300／GB200）', '新約邊際', 'GPU 數 × TK 密度（Analogy）', '=WhiteFiber!B%d' % WF_ROWS['bw_min'], '=WhiteFiber!B%d' % WF_ROWS['bw_max'], '合約 Verified（10-Q）；換算 Analogy', '同世代 TK 持有成本'),
    ('WhiteFiber', '存量世代續約（B200、H100）', '存量續約', '同上', '=WhiteFiber!B%d' % WF_ROWS['old_min'], '=WhiteFiber!B%d' % WF_ROWS['old_max'], '同上', ''),
    ('WhiteFiber', 'VR200 576 顆（2027 Q2 起）', '新約邊際', '同上', '=WhiteFiber!B%d' % WF_ROWS['vr'], '=WhiteFiber!B%d' % WF_ROWS['vr'], '同上', ''),
]
for i, row in enumerate(rows):
    put(sm, 4 + i, list(row))
n = 4 + len(rows)
put(sm, n + 1, ['組別', '', '', '', '低', '高', '', ''], True)
put(sm, n + 2, ['機隊平均（不含漲價列、不含 WhiteFiber 上界）', '', '', '', '=MIN(E4,E6,E7)', '=MAX(F4,F6,F7)', 'Derived', 'CoreWeave Q1／Q2、Nebius Q4'])
put(sm, n + 3, ['新約邊際（Nebius、IREN、WhiteFiber Blackwell；不含短約）', '', '', '', '=MIN(E9,E12,E14)', '=MAX(F9,F12,F14)', 'Derived', ''])
put(sm, n + 4, ['AL_SpendBasis 基準', '', '', '', '=%s' % R('sb'), '=%s' % R('sb'), 'Assumed', '區間 0.85–2.3'])

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, float) or (isinstance(c.value, str) and c.value.startswith('=')):
                if ws.title != '輸入':
                    c.number_format = '0.000'
            c.alignment = Alignment(wrap_text=True, vertical='top')
wb.save(out)
print('saved', out)
