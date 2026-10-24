# v21.3 (work order docs/workorders/20261024_v21.3_機架價低端與Sol快取價.md r0, branch claude/v21-3): VR200 rack price low end 7.8 -> 7.5 (Spec_Rack F11),
# GPT-6.1 Sol cached input 0.20 -> 0.10 (Cap_In D7／F47 read SRC_PRC_049), Azure paid-price notes (J3／J4), L1 token-ledger text fix.
# MINOR version: no contract name is added or removed, definitions unchanged; values of 175 contract names change (work order section 0 and 預期變動.csv).
#   W1  SRC_HW_098 (Bernstein 7.52, Active, linked to Spec_Rack!F11) and SRC_HW_099 (Morgan Stanley per-GW capex, Alt) as new records (src_new_records, appended by gov.src_append);
#       Spec_Rack F11 7.8 -> 7.5 (blue input, written only while it still holds 7.8); SRC_HW_017 W prefix; SRC_HW_016 V (link column) text.
#   W2  Cap_In D7／F47 =SRC_PRC_007 -> =SRC_PRC_049 (formula_map hook in gov._fm: the links are builder-owned and rewritten on every build; asserts the old formula);
#       SRC_PRC_049 Alt -> Active, SRC_PRC_007 Active -> Superseded; GPT-6 Sol -> GPT-6.1 Sol in the current labels (Cap_In A47, SRC_Price B10／B12／H10／H12); Cap_In G7 text.
#   W3  Alloc_In G23 prefix (Azure paid price A segment J3).   W4  DB_Evidence E404 and E405.
#   W5  L1 I128:I131 text: token-ledger multiple "about 6.26" -> "about 3.0" (hook in gov._rows_l1 like v2120.l1_patch; the text lives in column I, not C／H).
#   W6  Decisions SL60-01, README; (outside this file) tools/gen_fingerprint.py and tools/gen_ranges.py re-run, downstream contract v0.27.
# Every write to an Excel-owned cell is guarded (old value / old text / already-applied test), so the module is idempotent.
from openpyxl.utils import get_column_letter as L

from common import put, F_CALC
import v1410
import v1720
import v2120

VERSION = "20261024_Tokenomics_v21.3"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-24"
DASH = "—"
V = "v21.3"
WO = "工作單 docs/workorders/20261024_v21.3_機架價低端與Sol快取價.md（r0）"
DEC_ID = "SL60-01"
EID4, EID5 = "E404", "E405"
EVD_AZ = "docs/evidence/20261024_Azure實付價_A段/"
DEC_BY = ("Andy 2026-10-24 決定 #1（VR200 機架價區間）、#6（Sol 快取價）（算力收支回執《Tokenomics 兩項 Andy 決定》，依 Career & Wealth 證據送件 #1、#2、#6）；"
          "Andy 2026-10-24 14:50「同意用預設建議，請繼續」（Azure 實付價 A 段 J3、J4，PR #159）")

RACK_OLD, RACK_NEW = 7.8, 7.5
# build.py: block4 re-writes the OLD Cap_In A47 label on every build and text_fixes renames it afterwards; when the base already carries the new label,
# the Excel values of that row follow the label back to the builder's key (same pattern as v1520.SNAP_MOVES; keeps restore unmatched 0)
SNAP_MOVES = [(("Cap_In", ("GPT-6.1 Sol", 0), c), ("Cap_In", ("GPT-6 Sol", 0), c)) for c in range(2, 40)]
SRC_OLD, SRC_NEW = "=SRC_PRC_007", "=SRC_PRC_049"

# ---------------------------------------------------------------- W1 text
SRC017_HEAD = "v21.3：原出處 Bernstein 已於 2026-10 下修至 7.52（SRC_HW_098），高端 9.1 保留作上緣（Andy 2026-10-24 決定）｜"
SRC016_V_OLD = "Spec_Rack!F11｜直接"
SRC016_V_NEW = "—（v21.3 起 F11 來源為 SRC_HW_098；原：Spec_Rack!F11｜直接）"
BERNSTEIN_URL = "https://www.remio.ai/post/bernstein-ai-data-center-costs-hit-39-5-billion-per-gw-but-depreciation-is-the-b"
MS_URL = "https://247wallst.com/investing/2026/07/15/morgan-stanley-drops-a-50-billion-bombshell-can-big-tech-still-afford-to-build-the-ai-factories-of-the-future/"

# ---------------------------------------------------------------- W2 text
PRC049_HEAD = "v21.3：Andy 2026-10-24 決定採官方頁短上下文 0.10 為主值（C&W 證據送件 #6；官方頁 2026-10-24 再讀仍為 0.10）｜"
PRC049_V_OLD = "—（未連結模型格；主值 SRC_PRC_007 不改）"
PRC049_V_NEW = "Cap_In!D7｜直接；Cap_In!F47｜直接"
PRC007_V_OLD = "Cap_In!D7｜直接；Cap_In!F47｜直接"
PRC007_V_NEW = "—（v21.3 起由 SRC_PRC_049 取代；原：Cap_In!D7｜直接；Cap_In!F47｜直接）"
RENAME_NOTE = "v21.3 改名（原：GPT-6 Sol；值不變）｜"
G7_OLD = "＝輸入價 10%（S50）"
G7_NEW = "v21.3：Sol 快取＝0.10（官方頁短上下文，輸入價 5%；SRC_PRC_049）｜原：＝輸入價 10%（S50）"

# ---------------------------------------------------------------- W3 text
G23_HEAD = ("v21.3（2026-10-24 Azure 實付價 A 段 J3，PR #159；Andy 14:50「同意用預設建議」）：獨立價格錨點約 4 組，重心 1.0–1.5；"
            "唯一直接報價為 2024 A100 $1.30／小時（原文 per server，採每顆為推論，單位未證實）約 1.31 倍；"
            "1.8 為跨供應商混合殘差，未被推翻，維持基準 1.8、區間 0.85–2.3。｜")

# ---------------------------------------------------------------- W5 text
TL_REPL = (("（約 6.26）", "（約 3.0；v21.1 以前約 6.3）"), ("實際營收約為本列 ÷ 6.26", "實際營收約為本列 ÷ 3.0"))


def _find(ws, col, sid):
    rr = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == sid]
    assert len(rr) == 1, (sid, rr)
    return rr[0]


# ================================================================== W2: Cap_In D7／F47 links (builder-owned: gov.FORMULA_MAP rewrites them on every build, so the source is changed here)
FM = {"Cap_In!D7": SRC_NEW, "Cap_In!F47": SRC_NEW}


def formula_map(fm):
    fm = dict(fm)
    for k in FM:
        assert fm.get(k) == SRC_OLD, (k, fm.get(k))
    fm.update(FM)
    return fm


# ================================================================== W1: guarded input write (after restore, before gov_all)
def inputs_update(wb):
    log = []
    ws = wb["Spec_Rack"]
    assert ws["A11"].value == "機架價格 低（$M）", ws["A11"].value
    c = ws["F11"]
    if c.value is not None and type(c.value) is not bool and c.value == RACK_OLD:
        c.value = RACK_NEW; log.append(f"Spec_Rack!F11: {RACK_OLD} -> {RACK_NEW} (v21.3 W1)")
    return log


# ================================================================== W1: new SRC records (gov.src_append: only while the ID is absent)
def src_new_records(wb):
    r98 = v1720._rec(
        "SRC_HW_098", "SRC_HW", "VR200 NVL72 機架價格（Bernstein 2026-10 下修）", 7.52, None, None, "$M/架",
        "賣方研究估計（每 GW 全額資本支出約 39.5 $B，各設計 34.6–39.5）", "VR200 NVL72", "2026-10-10",
        BERNSTEIN_URL + "（2026-10-10，二手轉述）", 2, "中立", "賣方研究（Bernstein）", "二手", "Spec_Rack!F11｜取整（7.52 → 7.5）",
        "每 GW 金額只作對照、不驅動——GW 口徑（IT 或設施）未寫明，Tokenomics 以每 GW IT 由下而上計算；與 SRC_HW_017（Bernstein 6 月 9.09）同機構，下修版"
        f"（{V} W1；C&W 證據送件 #1；Andy 2026-10-24 決定；{WO}）")
    r98.update(status="Active", s="U-V2130", ev=EID5)
    r99 = v1720._rec(
        "SRC_HW_099", "SRC_HW", "每 GW 全額資本支出（Morgan Stanley：Rubin／GB300／GB200）", None, 35, 49, "$B/GW",
        "賣方研究估計；Rubin 約 49、GB300 約 39、GB200 約 35 $B／GW；GW 口徑未寫明", "VR200／GB300／GB200", "2026-07-15",
        MS_URL + "（二手）", 2, "中立", "賣方研究（Morgan Stanley）", "二手", "—（未連結模型格；只作對照）",
        "與 v21.2 L1_CapexGW_VR200／GB300／GB200 50.26／44.91／36.69 同量級；GW 口徑未寫明，只作對照（低＝GB200 35、高＝Rubin 49；GB300 約 39 見口徑欄）"
        f"（{V} W1；C&W 證據送件 #2；{WO}）")
    r99.update(status="Alt", s="U-V2130", ev=EID5)
    return [r98, r99]


# ================================================================== W1-4, W1-5, W2, W3: text and status (guarded)
def _set(ws, ref, old, new):
    cur = ws[ref].value
    if cur == new:
        return 0
    assert cur == old, (ws.title, ref, cur)
    ws[ref].value = new
    return 1


def _prefix(ws, ref, head, old_start):
    cur = str(ws[ref].value)
    if cur.startswith(head):
        return 0
    assert cur.startswith(old_start), (ws.title, ref, cur[:80])
    ws[ref].value = head + cur
    return 1


def text_fixes(wb):
    n = 0
    hw = wb["SRC_HW"]
    r17 = _find(hw, "A", "SRC_HW_017")
    n += _prefix(hw, f"W{r17}", SRC017_HEAD, "模型取 9.1；與 MS 同口徑")
    r16 = _find(hw, "A", "SRC_HW_016")
    n += _set(hw, f"V{r16}", SRC016_V_OLD, SRC016_V_NEW)
    assert hw[f"O{r16}"].value == "Alt"
    assert hw[f"V{_find(hw, 'A', 'SRC_HW_015')}"].value == "Spec_Rack!F12｜直接"          # W1-5: SRC_HW_015 link stays

    sp = wb["SRC_Price"]
    r49, r50, r07 = _find(sp, "A", "SRC_PRC_049"), _find(sp, "A", "SRC_PRC_050"), _find(sp, "A", "SRC_PRC_007")
    assert sp[f"V{r49}"].value in (PRC049_V_OLD, PRC049_V_NEW), sp[f"V{r49}"].value
    n += _set(sp, f"O{r49}", "Alt", "Active")
    n += _set(sp, f"V{r49}", PRC049_V_OLD, PRC049_V_NEW)
    n += _set(sp, f"Q{r49}", DASH, EID5)                  # Active needs an Evidence ID (Checks E8 ERROR otherwise): the record that registers the adoption is E405
    n += _prefix(sp, f"W{r49}", PRC049_HEAD, "v17.2 W1 B1 Q3：官方頁 2026-10-20 已改列 gpt-6.1-sol")
    n += _set(sp, f"O{r07}", "Active", "Superseded")
    n += _set(sp, f"P{r07}", DASH, "SRC_PRC_049")
    n += _set(sp, f"V{r07}", PRC007_V_OLD, PRC007_V_NEW)
    assert sp[f"O{r50}"].value == "Alt"                                                  # W2-4: SRC_PRC_050 stays Alt, untouched
    for sid, lab_old, lab_new in (("SRC_PRC_006", "GPT-6 Sol 新鮮輸入", "GPT-6.1 Sol 新鮮輸入"), ("SRC_PRC_008", "GPT-6 Sol 輸出（含思考）", "GPT-6.1 Sol 輸出（含思考）")):
        r = _find(sp, "A", sid)
        n += _set(sp, f"B{r}", lab_old, lab_new)
        n += _set(sp, f"H{r}", "GPT-6 Sol", "GPT-6.1 Sol")
        w = str(sp[f"W{r}"].value)
        if not w.startswith(RENAME_NOTE):
            assert w.startswith("—｜v17.2 Q3：官方頁 2026-10-20 已改列 gpt-6.1-sol"), w[:60]
            sp[f"W{r}"].value = RENAME_NOTE + w; n += 1
    assert sp[f"B{r07}"].value == "GPT-6 Sol 快取輸入" and sp[f"H{r07}"].value == "GPT-6 Sol"       # SRC_PRC_007 keeps the old label (Superseded)

    ci = wb["Cap_In"]
    assert ci["A7"].value == "快取輸入" and ci["D5"].value == "Sol"
    n += _set(ci, "A47", "GPT-6 Sol", "GPT-6.1 Sol")
    n += _set(ci, "G7", G7_OLD, G7_NEW)

    ai = wb["Alloc_In"]
    assert ai["A23"].value == "推論支出 ÷ 持有成本（計價倍數；含供應商利潤與資本回收）", ai["A23"].value
    n += _prefix(ai, "G23", G23_HEAD, "v21.2（2026-10-23 9a 查證 J5")
    assert ai["F23"].value == "Derived" and (ai["C23"].value, ai["D23"].value, ai["E23"].value) == (1.8, 0.85, 2.3)    # W3: C23:E23／F23 untouched
    return n


# ================================================================== W5: L1 I128:I131 (hook in gov._rows_l1, like v2120.l1_patch)
def l1_patch(R):
    out, n = [], 0
    for row in R:
        row = list(row)
        hit = False
        for k, f in enumerate(row):
            if isinstance(f, str) and "6.26" in f:
                for old, new in TL_REPL:
                    assert old in f, (row[0], k, old)
                    f = f.replace(old, new)
                assert "6.26" not in f, (row[0], k)
                row[k] = f; hit = True
        n += hit
        out.append(tuple(row))
    assert n == 4, n
    return out


# ================================================================== W4: DB_Evidence E404, E405
EVIDENCE = [
    [EID4, DATE, "Azure 對 OpenAI 實付價：2024 A100 $1.30／小時（DCD 轉述 The Information）",
     "DCD（Data Center Dynamics）轉述 The Information；Azure 實付價 A 段查證 " + EVD_AZ + "（PR #159）",
     "Interested-party（二手轉述）", "Alloc_In 第 23 列 C:E（AL_SpendBasis 1.8；0.85–2.3），G23 說明",
     "AL_SpendBasis 1.8（0.85–2.3，Derived；v21.2 J5 殘差法）", "2024 A100 $1.30／小時，約 1.31 倍；換成 2025 口徑約 2.2（1.9–2.6）",
     "採用為 Interested-party 筆數，標弱反對 AL_SpendBasis 1.8；不改值（V21.3 W3、W4；Andy 2026-10-24 14:50「同意用預設建議」）", V,
     "年度 2024、世代 A100、單位 per server 或 per GPU 均未證實（原文 per server，採每顆為推論）；獨立價格錨點約 4 組，重心 1.0–1.5；"
     "1.8 為跨供應商混合殘差，未被推翻，維持基準 1.8、區間 0.85–2.3；換成 2025 口徑約 2.2（1.9–2.6）｜" + WO, "已處理",
     DASH, DASH, "AL_SpendBasis（不改值）；Alloc_In G23 說明", 3, "利害關係方", "採用為 Interested-party 筆數（弱反對 1.8，不改值）｜v21.3 新增列"],
    [EID5, DATE, "C&W 證據送件 #1–#3、#6：VR200 機架價 Bernstein 7.52、每 GW 資本支出對照（Bernstein 39.5、MS 49／39／35）、GPT-6.1 Sol 快取價 0.10",
     "Career & Wealth 證據送件 20261024（GW 資本支出與 OpenAI 機隊配置）；算力收支回執；SRC_HW_098／099；SRC_PRC_049；" + BERNSTEIN_URL + "；" + MS_URL,
     "Interested-party／賣方研究（二手）", "Spec_Rack!F11；Cap_In D7／F47；SRC_HW_098／099、SRC_PRC_049／007",
     "Spec_Rack F11 7.8；Cap_In D7／F47＝SRC_PRC_007（0.20）", "F11 7.5（Bernstein 7.52 取整）；Cap_In D7／F47＝SRC_PRC_049（0.10）",
     "採用（機架價低端 7.5；Sol 快取價 0.10）；每 GW 資本支出只作對照（V21.3 W1、W2；Andy 2026-10-24 決定 #1、#6）", V,
     "每 GW 資本支出：Bernstein 39.5、MS 49／39／35 只作對照，因 GW 口徑（IT 或設施）未明（SRC_HW_098／099），Tokenomics 以每 GW IT 由下而上計算；"
     "送件的「現值 47.6」不適用：v21.2 為 50.26，機架價 7.8 時重算 47.58，推測 47.6 為舊版值；"
     "鴻海董事長 47 的利害關係註記（伺服器供應商）照登；Sol 快取價採用官方頁短上下文 0.10（W2；官方頁 2026-10-24 再讀仍為 0.10）；"
     "L 觀測點不登錄（快取價調整不視為一次 L 觀測；預設，可撤回）｜" + WO, "已處理",
     "SRC_HW_098／SRC_HW_099／SRC_PRC_049", DASH,
     "L1_M／L1_PxEff_2026／IF_PxEffMix／AL_PriceFleet／IF_PriceCached_Sol／IF_PriceRef_Sol／IF_RevGWFleet／L1_HoldEconGW_VR200_Lo／L1_CapexGW_VR200_Lo／IF_RackPrice 等 175 個名稱（見工作單預期變動.csv）",
     2, "中立", "採用（機架價低端、Sol 快取價）；每 GW 對照不驅動｜v21.3 新增列"],
]


def evidence_append(wb):
    ws = wb["DB_Evidence"]
    n = 0
    for ev in EVIDENCE:
        if any(ws.cell(r, 1).value == ev[0] for r in range(5, ws.max_row + 1)):
            continue
        r = max(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")) + 1
        assert ws.cell(r - 1, 1).value == ("E403" if ev[0] == EID4 else EID4), ws.cell(r - 1, 1).value
        for i, v in enumerate(ev):
            put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 4, 5, 6, 7, 8, 10, 14, 17))
        n += 1
    return n


def _decisions():
    return [[DEC_ID, V, "v21.3 VR200 機架價低端 7.5、GPT-6.1 Sol 快取價 0.10、Azure 實付價加註、L1 token 帳說明修補",
             "W1：Spec_Rack F11 7.8 → 7.5（藍字輸入；來源 SRC_HW_098 Bernstein 7.52 取整）；F13 9.1 不改，SRC_HW_017 備註加「原出處已下修至 7.52、高端 9.1 保留作上緣」；"
             "新增 SRC_HW_098（Active）、SRC_HW_099（Morgan Stanley 每 GW 資本支出對照，Alt）；SRC_HW_016 連結欄改為『—（v21.3 起 F11 來源為 SRC_HW_098；原：Spec_Rack!F11｜直接）』。"
             "只動 VR200 低成本情境（基準 12.23 與高端不變）；每 GW 金額只作對照，因 GW 口徑未明。"
             "W2：Cap_In D7／F47 =SRC_PRC_007 → =SRC_PRC_049（Sol 快取價 0.20 → 0.10）；SRC_PRC_049 Alt → Active、SRC_PRC_007 Active → Superseded（取代者 SRC_PRC_049）；"
             "SRC_PRC_050（長上下文 0.20）維持 Alt；現行標籤 GPT-6 Sol → GPT-6.1 Sol（Cap_In A47、SRC_Price B／H 欄第 10、12 列）；Cap_In G7 加註。"
             "有效牌價混合 1.961 → 1.943（−0.9%）；m 同比例 +0.9%；P×Q 營收、L 不變；L 觀測點不登錄（預設，可撤回）。"
             "W3：Alloc_In G23 加註 Azure 實付價 A 段 J3（C23:E23 不改，維持 1.8／0.85／2.3）；W4：DB_Evidence E404、E405。"
             "W5：L1 I128:I131 token 帳倍數『約 6.26』→『約 3.0（v21.1 以前約 6.3）』。"
             "修訂九（原值保留位置）：Cap_In G7 的『｜原：＝輸入價 10%（S50）』；SRC_HW_016 與 SRC_PRC_007 連結欄括號內『原：…』；SRC_Price 第 10、12 列備註前綴『原：GPT-6 Sol』；"
             "SRC_PRC_049 備註原文逐字保留於新前綴之後；Alloc_In G23、SRC_HW_017 備註原文逐字保留於新前綴之後；L1 I128:I131 保留『v21.1 以前約 6.3』。"
             "Andy 原話與依據：" + DEC_BY,
             DATE, DEC_BY, "v21.3 已建",
             "Spec_Rack F11；Cap_In D7／F47／A47／G7；Alloc_In G23；SRC_HW_016／017／098／099；SRC_PRC_007／049；SRC_Price B／H／W 第 10、12 列；L1 I128:I131；DB_Evidence E404／E405",
             f"{WO} W1–W6", "否"]]


def apply(wb, append_text):
    log = {"text_fixes": text_fixes(wb)}
    log["evidence"] = evidence_append(wb)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == DEC_ID for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


# ================================================================== README
README_VERSION = (VERSION + "（次版，底版 v21.2：" + WO + "——Spec_Rack F11 VR200 機架價低端 7.8 → 7.5（SRC_HW_098 Bernstein 7.52；SRC_HW_099 MS 每 GW 對照）；"
                  "Cap_In D7／F47 改讀 SRC_PRC_049（GPT-6.1 Sol 快取價 0.20 → 0.10；SRC_PRC_007 Superseded）；Alloc_In G23 加註 Azure 實付價（J3）；"
                  "DB_Evidence E404／E405；L1 I128:I131 token 帳說明改約 3.0；區間端點與時效指紋重產；Decisions SL60-01。"
                  "無新增或刪除契約名稱；L1_M 1.3432 → 1.3554、IF_PxEffMix 1.9607 → 1.9431、IF_PriceCached_Sol 0.16 → 0.08；"
                  "L1_RevGW_PQ_VR200*、L1_LReqFleet_Full*、L1_LTotal*、L1_TokLedger* 不變。）以下為 "
                  + v2120.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v2120.README_TITLE
README_TITLE = V + _T[len(v2120.V):]
README_ROW = ("v21.3 機架價低端 7.5 與 Sol 快取價 0.10",
              "VR200 機架價低端 7.8 → 7.5（Bernstein 2026-10 下修至 7.52；只動低成本情境，基準 12.23 與高端 9.1 不變；持有成本低端 9.51 → 9.22、資本支出低端 43.49 → 42.21）。"
              "GPT-6.1 Sol 快取價 0.20 → 0.10（官方頁短上下文；有效牌價混合 1.961 → 1.943，m 同比例 +0.9%，P×Q 營收與 L 不變）。"
              "每 GW 資本支出（Bernstein 39.5、MS 49）只作對照，因 GW 口徑未明。細節見工作單 docs/workorders/20261024_v21.3_機架價低端與Sol快取價.md 與 Decisions SL60-01。")
