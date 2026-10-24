"""網站的資料來源：直接讀取現行活頁簿中 LibreOffice 重算後存下的值（唯讀）。

2026-10-24 恢復網站（CLAUDE.md 修訂十一）：不恢復公式引擎（pycel）。
- model/ 內的活頁簿每一版都「經 LibreOffice 重算存檔」（CLAUDE.md 第 4 節），檔內已存有每個公式格的計算結果；
  網站只讀這些值，因此網站顯示的數字與 Excel（LibreOffice 重算）逐格相同，不需要另跑 parity。
- 不改任何輸入、不做任何計算；網站不能做「改輸入、即時重算」的情境試算（那需要恢復引擎，屬 Project 判斷）。
- 介面與舊 Engine 相同（path、names、sheetnames、get、get_name、name_ref），舊頁面不必改寫。
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from openpyxl.utils import range_boundaries

REPO_ROOT = Path(__file__).resolve().parent.parent
MODEL_DIR = REPO_ROOT / "model"
_REF_RE = re.compile(
    r"^(?:'(?P<q>(?:[^']|'')+)'|(?P<u>[^!']+))!"
    r"\$?(?P<c1>[A-Z]+)\$?(?P<r1>\d+)(?::\$?(?P<c2>[A-Z]+)\$?(?P<r2>\d+))?$"
)


def current_model_path(model_dir: Path = MODEL_DIR) -> Path:
    """model/CURRENT 記錄的現行活頁簿。"""
    name = (model_dir / "CURRENT").read_text(encoding="utf-8").strip()
    path = model_dir / name
    if not path.exists():
        raise RuntimeError(f"model/CURRENT 指向 {name}，但 model/ 內找不到這個檔")
    return path


def parse_ref(attr_text: str) -> tuple[str, str]:
    """'Interface!$C$19:$Q$19' → ('Interface', 'C19:Q19')；頁名可含中文或引號。"""
    m = _REF_RE.match(attr_text)
    if not m:
        raise ValueError(f"無法解析範圍：{attr_text!r}")
    sheet = (m["q"] or m["u"]).replace("''", "'")
    ref = f"{m['c1']}{m['r1']}"
    if m["c2"]:
        ref += f":{m['c2']}{m['r2']}"
    return sheet, ref


def norm(v):
    """空值統一為空字串（與舊引擎相同）。"""
    return "" if v is None else v


class Workbook:
    """唯讀：依具名範圍或位址取 LibreOffice 存下的值。"""

    def __init__(self, path: str | Path | None = None):
        self.path = Path(path) if path else current_model_path()
        wb = openpyxl.load_workbook(self.path, data_only=True)
        self._wb = wb
        self.sheetnames = list(wb.sheetnames)
        self.names = {}
        for k, v in wb.defined_names.items():
            try:
                parse_ref(v.attr_text)
            except ValueError:          # 常數或多區段名稱：網站不使用
                continue
            self.names[k] = v.attr_text

    def get(self, sheet: str, ref: str):
        ws = self._wb[sheet]
        if ":" not in ref:
            return norm(ws[ref].value)
        c1, r1, c2, r2 = range_boundaries(ref)
        return [[norm(ws.cell(r, c).value) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]

    def name_ref(self, name: str) -> tuple[str, str]:
        return parse_ref(self.names[name])

    def get_name(self, name: str):
        sheet, ref = self.name_ref(name)
        v = self.get(sheet, ref)
        if isinstance(v, list) and (len(v) == 1 or all(len(r) == 1 for r in v)):
            return [x for row in v for x in row]
        return v

    def row_values(self, sheet: str, row: int, max_col: int | None = None) -> list:
        ws = self._wb[sheet]
        return [norm(c.value) for c in ws[row][: max_col or ws.max_column]]

    def max_row(self, sheet: str) -> int:
        return self._wb[sheet].max_row
