# 產生 Azure實付價_A段.xlsx：換算一律用公式；輸入值（藍字）在「參數」「換算」兩頁
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

D = os.path.dirname(os.path.abspath(__file__))
BLUE = Font(color="0000FF"); BOLD = Font(bold=True)
HF = PatternFill("solid", fgColor="DDEBF7")
wb = Workbook()

# ---------------- 參數 ----------------
wp = wb.active; wp.title = "參數"
wp.append(["名稱","基準","低","高","單位","標記","來源／說明"])
P = {}
def prm(key, label, base, lo, hi, unit, mark, note):
    wp.append([label, base, lo, hi, unit, mark, note])
    r = wp.max_row
    for c in (2,3,4):
        if wp.cell(r,c).value is not None and not str(wp.cell(r,c).value).startswith("="):
            wp.cell(r,c).font = BLUE
    P[key] = (f"參數!$B${r}", f"參數!$C${r}", f"參數!$D${r}")
prm("hold","AL_FleetHold（TK 機隊持有成本，每 IT-GW·年）",9.35311859805586,6.71937130244442,15.7256917850111,"$B/GW·年","Derived（TK v21.2 Alloc!C157/D157/E157）","IT 關鍵電力口徑；Hopper 0.6／GB200 0.4")
prm("spendbasis","AL_SpendBasis（計價倍數）",1.8,0.85,2.3,"x","Derived","TK v21.2 Alloc_In!C23；本段檢驗對象")
prm("gpu_h","IF_GPUsPerGW Hopper",718048,None,None,"顆/GW","Derived（TK）","沿用 9a 筆記")
prm("gpu_b","IF_GPUsPerGW GB200",501768,None,None,"顆/GW","Derived（TK）","沿用 9a 筆記")
prm("w_h","2025 機隊 Hopper 權重",0.6,None,None,"","Derived（TK）","GB200 權重＝1−此值")
prm("hrs","每年小時數",8760,None,None,"小時","常數","")
prm("gw25","AL_FleetGW25（OpenAI 2025 平均上線 GW）",1.12780698259266,1.0,1.25,"GW","Derived（TK Cap_In!C331）","9a：0.6→1.9 GW 端點內插")
prm("spend25","OpenAI 2025 全部算力支出（84＋120）",20.4,16.1,21.7,"$B","Interested-party（外流查核帳；chat 端 9a）","204 億；161–217 億為 chat 端查核帳框出區間")
prm("s","s：204 億中屬當年上線算力費用的比例",0.9,0.79,1.0,"","Derived（chat 端 9a）","基準取 0.9 為本段簡化，非 9a 數字")
prm("pue","設施電力／IT 電力（PUE，設施口徑 GW 換 IT 用）",1.2,1.0,1.3,"x","Assumed","TK 慣例 GW＝IT；外部 GW 口徑未明時才用；Epoch 對 PUE 區間 1.1–1.7")
# 衍生
wp.append([]); wp.append(["衍生（公式）"]); wp.cell(wp.max_row,1).font = BOLD
def drv(key,label,formula,unit,note):
    wp.append([label, formula, None, None, unit, "Derived", note]); r = wp.max_row
    P[key] = (f"參數!$B${r}",)
drv("gpu_mix","混合 GPU 數／GW＝Hopper 權重×顆數＋GB200 權重×顆數",f"={P['w_h'][0]}*{P['gpu_h'][0]}+(1-{P['w_h'][0]})*{P['gpu_b'][0]}","顆/GW","約 63.2 萬顆")
drv("hold_gpuhr","TK 持有成本每 GPU 小時",f"={P['hold'][0]}*1E9/({P['gpu_mix'][0]}*{P['hrs'][0]})","$/GPU·小時","約 1.69")
drv("price18","1.8 倍對應的實付價（每 GW·年）",f"={P['spendbasis'][0]}*{P['hold'][0]}","$B/GW·年","約 16.8")
drv("price18_gpuhr","1.8 倍對應的實付價（每 GPU 小時）",f"={P['spendbasis'][0]}*{P['hold_gpuhr'][0]}","$/GPU·小時","約 3.04")
for c,w in zip("ABCDEFG",(52,14,10,10,12,34,70)): wp.column_dimensions[c].width = w

# ---------------- 換算 ----------------
wc = wb.create_sheet("換算")
wc.append(["區塊／項目","數值","單位","標記","說明（藍字＝輸入；黑字＝公式）"])
for c in wc[1]: c.font = BOLD; c.fill = HF
C = {}
def blk(title):
    wc.append([]); wc.append([title]); wc.cell(wc.max_row,1).font = BOLD
def row(key,label,val,unit,mark="",note=""):
    wc.append([label,val,unit,mark,note]); r = wc.max_row
    if not (isinstance(val,str) and val.startswith("=")): wc.cell(r,2).font = BLUE
    C[key] = f"換算!$B${r}"
hold,holdlo,holdhi = P["hold"]; hrs = P["hrs"][0]
def per_gw_ratio(k, spend_key, gw_key, note=""):
    row(k+"_pergw", "每 GW·年", f"={C[spend_key]}/{C[gw_key]}", "$B/GW·年","Derived",note)
    row(k+"_ratio", "÷ AL_FleetHold（倍數）", f"={C[k+'_pergw']}/{hold}", "x","Derived")

blk("E01 The Information 2024（DCD 轉述）：A100 每小時 1.30 美元、等效 35 萬顆 A100")
row("e01_p","Microsoft 向 OpenAI 收費（每 A100 小時）",1.30,"$/A100·小時","Interested-party","DCD：「about $1.30 per A100 server per hour」；以下用『每顆』讀法：The Information 原句寫 chips；若 35 萬是 8 顆伺服器，則 280 萬顆 A100 約 2.6 GW，遠超 OpenAI 2024 年底總量 0.6 GW。金額 35 萬×1.30×8760＝39.9 億兩種讀法都吻合")
row("e01_n","等效 A100 顆數",350000,"顆","Interested-party","The Information 2024-07：cluster 'has the equivalent of 350,000 Nvidia A100 chips'")
row("e01_spend","隱含年支出",f"={C['e01_p']}*{C['e01_n']}*{hrs}/1E9","$B/年","Derived","應≈ 報導『約 40 億』")
row("e01_zit","對照：Zitron 內部文件 2024 全年 Azure 推論（4 季加總）",3.767,"$B","Interested-party","wheresyoured.at；與上行同量級（相差約 6%）")
row("e01_kw","每顆 A100 的 IT 電力（含伺服器內網路等）基準",0.93,"kW/顆","Assumed","DGX A100 約 6.5kW／8 顆＝0.81；×Epoch IT 開銷係數 1.14＝0.93；區間 0.8–1.1")
row("e01_kwlo","kW 低",0.8,"kW/顆","Assumed","")
row("e01_kwhi","kW 高",1.1,"kW/顆","Assumed","")
row("e01_gw","隱含 IT-GW（基準）",f"={C['e01_n']}*{C['e01_kw']}/1E6","GW","Derived","")
row("e01_gwlo","隱含 IT-GW（kW 低）",f"={C['e01_n']}*{C['e01_kwlo']}/1E6","GW","Derived","")
row("e01_gwhi","隱含 IT-GW（kW 高）",f"={C['e01_n']}*{C['e01_kwhi']}/1E6","GW","Derived","")
per_gw_ratio("e01","e01_spend","e01_gw","2024 年、A100 世代；比較基準仍用 TK 2025 機隊持有成本（假設每 IT-kW 的持有成本跨世代近似不變：伺服器資本 /kW 約 25–31 千美元，Assumed）")
row("e01_ratio_lo","倍數低（kW 高）",f"={C['e01_spend']}/{C['e01_gwhi']}/{hold}","x","Derived")
row("e01_ratio_hi","倍數高（kW 低）",f"={C['e01_spend']}/{C['e01_gwlo']}/{hold}","x","Derived")

row("e01_gpuhr","換成 2025 機隊口徑每 GPU 小時（倍數×TK 每 GPU 小時持有成本，基準）",f"={C['e01_ratio']}*{P['hold_gpuhr'][0]}","$/GPU·小時","Derived","對照證偽條件 (b) 的 2 美元門檻")
row("e01_gpuhr_lo","同上（低）",f"={C['e01_ratio_lo']}*{P['hold_gpuhr'][0]}","$/GPU·小時","Derived","")
row("e01_gpuhr_hi","同上（高）",f"={C['e01_ratio_hi']}*{P['hold_gpuhr'][0]}","$/GPU·小時","Derived","橫跨 2 美元 → (b) 不能判為觸發，也不能判為排除")

blk("E02 同一個 1.30：換成『占牌價比例』再套到 2025 機隊（與 E01 同源，算一筆）")
row("e02_a100list","Azure ND96amsr A100 v4 隨需牌價（8 顆）",32.77,"$/小時","Verified（牌價，三方追蹤站 sparecores）","")
row("e02_a100gpu","每顆 A100 牌價",f"={C['e02_a100list']}/8","$/A100·小時","Derived","")
row("e02_frac","1.30 占牌價比例",f"={C['e01_p']}/{C['e02_a100gpu']}","","Derived","約 0.32（折扣約 68%）")
row("e02_h100","H100 v5 隨需牌價每顆（98.32÷8）",f"=98.32/8","$/H100·小時","Verified（牌價，azurespeed）","")
row("e02_gb200","GB200 v6 隨需牌價每顆（108.16÷4）",f"=108.16/4","$/GB200·小時","Verified（牌價，azurespeed；Microsoft Learn：每 VM 4 顆）","")
row("e02_mix","2025 機隊權重牌價（0.6／0.4）",f"={P['w_h'][0]}*{C['e02_h100']}+(1-{P['w_h'][0]})*{C['e02_gb200']}","$/GPU·小時","Derived","")
row("e02_applied","套用 0.32 後的每 GPU 小時",f"={C['e02_frac']}*{C['e02_mix']}","$/GPU·小時","Derived","把 2024 A100 的折扣比例硬套到 2025 世代，僅作上界示意")
row("e02_ratio","÷ TK 每 GPU 小時持有成本",f"={C['e02_applied']}/{P['hold_gpuhr'][0]}","x","Derived","牌價每 kW 隨世代大漲（見 E10），折扣比例不可跨世代套用；本列為上界，不作支持證據")

blk("E03 Microsoft FY26 10-K：來自 OpenAI 的營收 241 億（含營收分成）")
row("e03_rev","Microsoft 對 OpenAI 營收（FY26＝2025-07 至 2026-06，含營收分成）",24.1,"$B","Verified","10-K：'revenue from commercial arrangements with OpenAI, inclusive of revenue-sharing payments, of $24.1 billion'")
row("e03_rs","扣除營收分成（OpenAI 付 Microsoft 的 20% 分成；金額未披露）",3.0,"$B","Assumed","20% 營收分成的粗估，區間 2.0–4.0；OpenAI FY26 營收未知，無來源")
row("e03_cap","可歸屬算力使用費",f"={C['e03_rev']}-{C['e03_rs']}","$B","Derived","")
row("e03_gw1","假設 Azure 承載 OpenAI 的 FY26 平均 IT-GW（低）",1.0,"GW","Assumed","無來源；OpenAI 2025 年底總量 1.9 GW")
row("e03_gw2","同上（中）",1.75,"GW","Assumed","")
row("e03_gw3","同上（高）",2.5,"GW","Assumed","")
row("e03_r1","倍數（GW 低）",f"={C['e03_cap']}/{C['e03_gw1']}/{hold}","x","Derived","")
row("e03_r2","倍數（GW 中）",f"={C['e03_cap']}/{C['e03_gw2']}/{hold}","x","Derived","")
row("e03_r3","倍數（GW 高）",f"={C['e03_cap']}/{C['e03_gw3']}/{hold}","x","Derived","分母 Azure 承載 GW 無來源 → 不可分離，僅作量級")

blk("E04 2025-10 重組：OpenAI 向 Azure 增購 2,500 億美元（無單價、無期間）")
row("e04_amt","增購承諾",250,"$B","Verified（Microsoft 官方部落格）","'contracted to purchase an incremental $250B of Azure services'；未載單價與期間")
row("e04_gwy1","若單價＝1.0×持有成本，可買",f"={C['e04_amt']}/{hold}","GW·年","Derived","")
row("e04_gwy2","若單價＝1.8×持有成本，可買",f"={C['e04_amt']}/{P['price18'][0]}","GW·年","Derived","期間未知 → 無法反推單價")

blk("E05 Oracle–OpenAI：3,000 億美元、5 年、4.5 GW（Next Platform：'culminates with 4.5 GW'）")
row("e05_tcv","合約總額",300,"$B","Interested-party（媒體；Oracle 申報未點名客戶）","")
row("e05_yrs","年數",5,"年","Interested-party","")
row("e05_gw","終局容量",4.5,"GW","Interested-party","口徑（IT／設施）未明；起點 2027 年中，非 2025 價格")
row("e05_base","每 GW·年（假設 4.5 GW 全期都在、且為 IT 口徑）",f"={C['e05_tcv']}/{C['e05_yrs']}/{C['e05_gw']}","$B/GW·年","Derived","下限情境")
row("e05_ramp","全期平均容量÷終局容量（爬坡）基準",0.8,"","Assumed","區間 0.6–1.0；'culminates' 表示終局才達 4.5 GW")
row("e05_ramplo","爬坡低（平均容量比例小）",0.6,"","Assumed","")
row("e05_pue","設施→IT 換算（PUE）基準",f"={P['pue'][0]}","x","Assumed","若 4.5 GW 為設施口徑")
row("e05_puehi","PUE 高",f"={P['pue'][2]}","x","Assumed","")
row("e05_r_lo","倍數低（IT 口徑、無爬坡）",f"={C['e05_base']}/{hold}","x","Derived","")
row("e05_r_mid","倍數基準（PUE 基準、爬坡 0.8）",f"={C['e05_base']}*{C['e05_pue']}/{C['e05_ramp']}/{hold}","x","Derived","")
row("e05_r_hi","倍數高（PUE 高、爬坡 0.6）",f"={C['e05_base']}*{C['e05_puehi']}/{C['e05_ramplo']}/{hold}","x","Derived","")

blk("E06 供應商毛利率錨點：價格÷供應商成本＝1÷(1−毛利率)")
row("e06_m1","Oracle AI 雲 2025-06～08 季平均毛利率（The Information 內部文件，DCD 轉述）",0.16,"","Interested-party","季間 10–20%；該季 AI 伺服器出租 9 億、毛利 1.25 億（約 14%）")
row("e06_m2","Oracle 預期 AI 基礎設施毛利率上限（itiger 轉述）",0.35,"","Interested-party","'gross margin of up to 35%'；是否全合約期未明")
row("e06_m3","Oracle 目標毛利率（Next Platform 引述）",0.40,"","Interested-party","'try to get 40 percent gross margins'")
row("e06_k1","價格÷成本（16%）",f"=1/(1-{C['e06_m1']})","x","Derived","")
row("e06_k2","價格÷成本（35%）",f"=1/(1-{C['e06_m2']})","x","Derived","")
row("e06_k3","價格÷成本（40%）",f"=1/(1-{C['e06_m3']})","x","Derived","『成本』是供應商會計成本，與 TK 經濟持有成本是否同口徑未知")

blk("E07 CoreWeave（公司申報）：營收÷活躍電力")
row("e07_q3r","2025Q3 營收",1.365,"$B","Verified","CoreWeave 2025Q3 新聞稿（9a W3-03）")
row("e07_q3a","Q2 末活躍電力",470,"MW","Verified（由 Q3 稿『+120 MW』推得）","")
row("e07_q3b","Q3 末活躍電力",590,"MW","Verified","")
row("e07_q4r","2025Q4 營收",1.572,"$B","Verified","CoreWeave 2025Q4 新聞稿（businesswire 2026-02-26）；全年 5.131")
row("e07_q4b","Q4 末活躍電力",850,"MW","Verified","'more than 850 MW'（取下限）")
row("e07_q3","Q3：年化營收÷平均活躍電力",f"={C['e07_q3r']}*4/(({C['e07_q3a']}+{C['e07_q3b']})/2/1000)","$B/GW·年","Derived","全客戶混合；活躍電力口徑（IT／設施）公司未明")
row("e07_q4","Q4：年化營收÷平均活躍電力",f"={C['e07_q4r']}*4/(({C['e07_q3b']}+{C['e07_q4b']})/2/1000)","$B/GW·年","Derived","Q4 大量新增電力尚在爬坡，分母偏大")
row("e07_r3","Q3 倍數",f"={C['e07_q3']}/{hold}","x","Derived","")
row("e07_r4","Q4 倍數",f"={C['e07_q4']}/{hold}","x","Derived","")
row("e07_r34","Q3＋Q4 合併倍數",f"=({C['e07_q3r']}+{C['e07_q4r']})*2/((({C['e07_q3a']}+{C['e07_q3b']})/2+({C['e07_q3b']}+{C['e07_q4b']})/2)/2/1000)/{hold}","x","Derived","")

blk("E08 IREN–Microsoft（公司公告）：97 億美元、5 年、約 200 MW IT")
row("e08_tcv","合約總額",9.7,"$B","Verified（公司公告；techinformed 轉述）","20% 預付款，第 24 個月後抵用；GPU 為 GB300；Microsoft 是買方")
row("e08_yrs","年數",5,"年","Verified","")
row("e08_gw","IT 負載",0.2,"GW","Verified（公告用語 'about 200 MW of IT load'）","")
row("e08_pergw","每 GW·年",f"={C['e08_tcv']}/{C['e08_yrs']}/{C['e08_gw']}","$B/GW·年","Derived","分母是額定 IT 容量、分子含爬坡期 → 實際單價 ≥ 此值")
row("e08_ratio","倍數",f"={C['e08_pergw']}/{hold}","x","Derived","")

blk("E10 Azure 公開牌價（上界錨點）")
row("e10_h_od","ND H100 v5 隨需（每 GPU 小時）",f"=98.32/8","$/小時","Verified（牌價，azurespeed／sparecores；東美）","8 顆")
row("e10_h_1y","ND H100 v5 一年預留（9a 沿用 Spheron）",7.93,"$/小時","Analogy（二手）","")
row("e10_h_3y","ND H100 v5 三年預留（9a 沿用 Spheron）",5.47,"$/小時","Analogy（二手）","")
row("e10_h_sp","ND H100 v5 Spot（azurespeed／sparecores 18.1695÷8）",f"=18.1695/8","$/小時","Verified（牌價，可被收回）","非生產用途；公開最低價")
row("e10_g_od","ND GB200 v6 隨需（108.16÷4）",f"=108.16/4","$/小時","Verified（牌價）","每 VM 4 顆")
row("e10_g_1y","ND GB200 v6 一年預留（69.222÷4）",f"=69.222/4","$/小時","Verified（牌價，devzero）","三年預留未找到")
gh = P['gpu_h'][0]; gb = P['gpu_b'][0]
def lst(key,label,pk,gpuref):
    row(key+"_pergw",label+"：每 GW·年",f"={C[pk]}*{gpuref}*{hrs}/1E9","$B/GW·年","Derived","")
    row(key+"_ratio",label+"：÷ 持有成本",f"={C[key+'_pergw']}/{hold}","x","Derived","")
    row(key+"_disc",label+"：1.8 倍相當於牌價折扣",f"=1-{P['price18'][0]}/{C[key+'_pergw']}","","Derived","")
lst("e10a","H100 隨需","e10_h_od",gh)
lst("e10b","H100 三年預留","e10_h_3y",gh)
lst("e10c","H100 Spot","e10_h_sp",gh)
lst("e10d","GB200 隨需","e10_g_od",gb)
lst("e10e","GB200 一年預留","e10_g_1y",gb)

blk("E13 OpenAI–AWS：380 億美元、7 年、『數十萬顆 GB200／GB300』（無 MW）")
row("e13_tcv","合約總額",38,"$B","Interested-party（公告）","")
row("e13_yrs","年數",7,"年","Interested-party","")
row("e13_n1","GPU 顆數（低）",200000,"顆","Assumed","公告用語 'hundreds of thousands'；無確切數")
row("e13_n2","GPU 顆數（中）",400000,"顆","Assumed","")
row("e13_n3","GPU 顆數（高）",600000,"顆","Assumed","")
for i in (1,2,3):
    row(f"e13_r{i}",f"倍數（顆數情境 {i}，GB200 每 GW 顆數）",f"={C['e13_tcv']}/{C['e13_yrs']}/({C[f'e13_n{i}']}/{gb})/{hold}","x","Derived","分子平均到 7 年、分母取終局顆數 → 偏低；容量尚在 2026 年底前爬坡")

blk("E15 GW 端點與 GPU 顆數對照（Q3）")
row("e15_gw24","OpenAI 2024 年底",0.6,"GW","Interested-party（OpenAI）","9a 已有")
row("e15_gw25","OpenAI 2025 年底",1.9,"GW","Interested-party（OpenAI）","9a 已有；各報導同源（OpenAI 2026-01）")
row("e15_lin","線性內插平均",f"=({C['e15_gw24']}+{C['e15_gw25']})/2","GW","Derived","")
row("e15_ep24","Epoch 2024 年底 H100e（功率模型）",384000,"H100e","Analogy（獨立研究機構）","90% CI 276k–485k")
row("e15_ep25","Epoch 2025 年底 H100e（功率模型）",1740000,"H100e","Analogy","90% CI 1.25–2.19M")
row("e15_epr","H100e 平均÷年底（線性）",f"=(({C['e15_ep24']}+{C['e15_ep25']})/2)/{C['e15_ep25']}","","Derived","若平均 GW 與 H100e 同比例：平均 GW＝年底 GW×此值")
row("e15_ep_avg","對應平均 GW",f"={C['e15_gw25']}*{C['e15_epr']}","GW","Derived","")
row("e15_gpu","Altman 2025-07-20：2025 年底前上線『well over 1 million GPUs』",1000000,"顆","Interested-party（Altman 於 X）","『well over』→ 下限")
row("e15_gpugw","隱含每 GW 顆數下限（1.0M÷1.9 GW）",f"={C['e15_gpu']}/{C['e15_gw25']}","顆/GW","Derived","TK 混合值約 63 萬 → 一致")
row("e15_f","前重情境：新增容量中先到位的比例（落在 t＝0.25）",0.8,"","Assumed","極端前重；其餘落在 t＝0.75")
row("e15_t","新增容量的平均到位時點（年）",f"={C['e15_f']}*0.25+(1-{C['e15_f']})*0.75","年","Derived","")
row("e15_max","前重情境平均上線 GW",f"={C['e15_gw24']}+({C['e15_gw25']}-{C['e15_gw24']})*(1-{C['e15_t']})","GW","Derived","即使 80% 新增在上半年初完成，平均也僅約 1.45 GW；要達 1.8 須約 4 週內到位 → 證偽條件 (a) 不可能由端點資料成立")

blk("T 三角約束：2025 全部算力支出＝平均 GW × 計價倍數 × AL_FleetHold × …（9a 恆等式）")
row("t_const","支出÷持有成本",f"={P['spend25'][0]}/{hold}","GW","Derived","平均 GW＝此值×s÷倍數")
for m in (1.0,1.3,1.5,1.8,2.3):
    k = str(m).replace(".","")
    row(f"t_m{k}",f"倍數 {m}：所需平均 GW（s＝0.79）",f"={C['t_const']}*{P['s'][1]}/{m}","GW","Derived","")
    row(f"t_m{k}h",f"倍數 {m}：所需平均 GW（s＝1.0）",f"={C['t_const']}*{P['s'][2]}/{m}","GW","Derived","")
row("t_minm_s1","平均 GW ≤ 前重上限時，s＝1.0 所需最低倍數",f"={C['t_const']}*{P['s'][2]}/{C['e15_max']}","x","Derived","")
row("t_minm_s79","同上，s＝0.79",f"={C['t_const']}*{P['s'][1]}/{C['e15_max']}","x","Derived","")
row("t_minm_hi","上限 217 億（s＝1）：所需最低倍數（前重上限 GW）",f"={P['spend25'][2]}/{hold}/{C['e15_max']}","x","Derived","s＝0.79 即 161÷204，與支出下限共線，非獨立假設；全部供應商總支出÷總 GW → 為跨供應商混合倍數，非 Azure 單價")
row("t_m_lin","線性內插 GW（1.25）：s＝1.0 所需倍數",f"={C['t_const']}*{P['s'][2]}/{C['e15_lin']}","x","Derived","")
for c,w in zip("ABCDE",(78,14,14,34,100)): wc.column_dimensions[c].width = w

# ---------------- 證據表 ----------------
we = wb.create_sheet("證據表")
H = ["編號","同源群組","項目","來源","日期（出處／讀取）","原文摘錄（引號內逐字；其餘為摘述）","原始數值與單位","口徑","換算步驟（公式在『換算』頁）","每 GW·年（$B）","÷9.353 倍數","倍數低","倍數高","標記","對 1.8 的判斷"]
we.append(H)
for c in we[1]: c.font = BOLD; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="top")
def ev(*a): we.append(list(a))
def f(k): return f"={C[k]}"
ev("E01","S1 The Information 2024","Microsoft 向 OpenAI 收的 A100 單價","DCD 轉述 The Information（2024-07）https://www.datacenterdynamics.com/en/news/openai-training-and-inference-costs-could-reach-7bn-for-2024-ai-startup-set-to-lose-5bn-report/","2024-07／讀取 2026-10-24","'Microsoft has charged OpenAI about $1.30 per A100 server per hour, way below normal rates.'；'350,000 Nvidia A100 chips'","1.30 美元／A100·小時；35 萬顆；約 40 億／年","2024 年、A100、推論叢集；IT-GW 由 kW/顆 Assumed 換算","35 萬×1.30×8760＝年支出；除以 35 萬×0.93kW＝GW",f("e01_pergw"),f("e01_ratio"),f("e01_ratio_lo"),f("e01_ratio_hi"),"Interested-party（報導洩漏；Microsoft 稱數字不盡正確的是另案）；換算含 Assumed","弱反對：2024／A100 約 1.3 倍（1.1–1.5），低於 1.8 但在 0.85–2.3 內；世代與年度不同。單位未證實：DCD 原文為 $1.30 per A100 server per hour、350,000 servers；Willison 轉引 The Information 為 350,000 A100 chips；本段採『每顆』是依容量推論（8 顆／台會得 2.6 GW＞0.6 GW），原文單位未證實。若真為每台，實付約低至 0.16 倍，更傾向觸發證偽 (b)")
ev("E02","S1 The Information 2024","同一個 1.30 換成占牌價比例再套 2025 機隊","同上＋Azure 牌價頁（sparecores、azurespeed）","2024／2026","同 E01；牌價 A100 32.77／8 顆、H100 98.32／8 顆、GB200 108.16／4 顆","1.30÷4.10＝0.32 → ×2025 權重牌價","把折扣比例跨世代硬套","見換算頁 E02（套用結果約 5.8 美元／GPU·小時）",None,f("e02_ratio"),None,None,"Analogy：2024 A100 折扣 → 2025 世代；差在牌價每 kW 隨世代暴漲","無關（上界示意，不採用）：套用結果約 3.4 倍，高於 2.3 上限，顯示折扣不可跨世代套用")
ev("E03","S2 Microsoft 申報","Microsoft FY26 來自 OpenAI 的營收","Microsoft 10-K FY2026 Note 1（sec.gov）https://www.sec.gov/Archives/edgar/data/789019/000119312526323660/msft-20260630.htm","2026-07／讀取 2026-10-24","'For fiscal year 2026, we recorded revenue from commercial arrangements with OpenAI, inclusive of revenue-sharing payments, of $24.1 billion, and accounts receivable from OpenAI as of June 30, 2026 was $6.0 billion.'","241 億；應收 60 億","FY26＝2025-07～2026-06；含 20% 營收分成與其他安排；未分 Azure／分成","需 Azure 承載 GW 才能換算（無來源）→ 以 1.0／1.75／2.5 GW 示範",f("e03_cap")+"/"+C["e03_gw2"],f("e03_r2"),f("e03_r3"),f("e03_r1"),"Verified（金額）；GW 為 Assumed","無關（不可分離）：只有金額，無同期 Azure 分配容量；示範倍數 0.9–2.3")
ev("E04","S3 重組條款","2025-10 重組：增購 Azure 2,500 億","Microsoft 官方部落格 https://blogs.microsoft.com/blog/2025/10/28/the-next-chapter-of-the-microsoft-openai-partnership/","2025-10-28／讀取 2026-10-24","'OpenAI has contracted to purchase an incremental $250B of Azure services'；Microsoft 不再有 compute 優先承購權","2,500 億美元","未載單價、期間、世代","無法換算：僅示範可買幾個 GW·年",None,None,None,None,"Verified（金額）","無關（不可分離）：單價與期間皆未揭露")
ev("E05","S4 Oracle 合約","Oracle–OpenAI","Next Platform 2026-02-02 https://www.nextplatform.com/2026/02/02/oracles-financing-primes-the-openai-pump/；金額原見 WSJ／9a W3-01","2025-09 簽；2026-02／讀取 2026-10-24","'a five year deal that culminates with 4.5 gigawatts of datacenter facilities'","3,000 億美元／5 年／4.5 GW","2027 起、新世代；GW 口徑未明（facility 字樣）","300÷5÷4.5＝13.3；再依 PUE、爬坡調整",f("e05_base"),f("e05_r_mid"),f("e05_r_lo"),f("e05_r_hi"),"Interested-party（Oracle／OpenAI 轉述）+ Analogy（非 Azure、非 2025）","弱支持（2027 起未來價、非 Azure、facility 與爬坡為 Assumed）：最低 1.4，基準 2.1，高 3.1")
ev("E06","S5 Oracle 毛利率","供應商加價的上限錨點","DCD／The Information https://www.datacenterdynamics.com/en/news/internal-documents-show-oracles-average-profit-margin-for-ai-cloud-deals-was-16-report/；itiger；Next Platform","2025-10／讀取 2026-10-24","'average gross profit margin of just 16 percent'（季間 10–20%）；'gross margin of up to 35%'；'try to get 40 percent gross margins'","16%／35%／40%","價格÷供應商會計成本","1÷(1−毛利率)",None,f("e06_k1"),f("e06_k1"),f("e06_k3"),"Interested-party：Oracle（內部文件外流／管理層目標）","單列對照（不併入 1.0–1.5 群）：價格÷供應商會計 COGS，與『÷TK 持有成本』口徑不同；純毛利只對應 1.19–1.67")
ev("E07","S6 CoreWeave 申報","CoreWeave 營收÷活躍電力","CoreWeave 2025Q4 新聞稿 https://www.businesswire.com/news/home/20260226281661/en/CoreWeave-Reports-Strong-Fourth-Quarter-and-Fiscal-Year-2025-Results；Q3 稿（9a）","2026-02-26／讀取 2026-10-24","Q4 營收 1,572 百萬；'Added approximately 260 MW of active power capacity, bringing the total to more than 850 MW'","Q3 1.365B／530MW；Q4 1.572B／720MW","全客戶、Hopper 為主、含網路儲存；活躍電力口徑未明","年化營收÷平均活躍電力",f("e07_q4"),f("e07_r34"),f("e07_r4"),f("e07_r3"),"Verified（原數）／Derived（比值）；Analogy（非 OpenAI、非 Azure）；active power 公司未定義，Q4 用『more than 850 MW』使倍數略高","反對：neocloud 已實現單價約 0.9–1.1 倍；Azure 對 OpenAI 若遠高於此需另有理由")
ev("E08","S7 IREN 公告","Microsoft 向 IREN 租 GB300","techinformed 轉述 IREN 公告 https://techinformed.com/microsoft-signs-9-7-billion-iren-deal-for-nvidia-gb300-capacity/","2025-11／讀取 2026-10-24","'about $9.7 billion'；五年；'about 200 MW of IT load'；20% 預付","97 億／5 年／200 MW IT","IT 口徑、GB300、Microsoft 是買方","9.7÷5÷0.2",f("e08_pergw"),f("e08_ratio"),f("e08_ratio"),None,"Derived（金額為 Verified）+ Analogy（買方價，非賣給 OpenAI）；未處理 200 MW 分批爬坡，倍數偏低","反對（弱）：Microsoft 向 neocloud 租的成本約 1.0 倍；Azure 轉售給 OpenAI 要加到 1.8 倍需加價約 73%")
ev("E10","S8 Azure 牌價","Azure ND H100 v5／ND GB200 v6 公開牌價","azurespeed https://www.azurespeed.com/AzureVmPricing/Standard_ND96isr_H100_v5 、…/Standard_ND128isr_NDR_GB200_v6；sparecores；devzero；Microsoft Learn（GB200 每 VM 4 顆）","2026／讀取 2026-10-24","H100 v5 隨需 98.32 美元／小時（8 顆）；GB200 v6 隨需 108.16（4 顆）、一年預留 69.222；H100 Spot 18.1695","H100 12.29／顆；GB200 27.04／顆","公開牌價（非 OpenAI 實付）；東美最低區","價×每 GW 顆數×8760（基準欄＝H100 三年預留；低＝H100 Spot；高＝GB200 隨需）",f("e10b_pergw"),f("e10b_ratio"),f("e10c_ratio"),f("e10d_ratio"),"Verified（牌價）；非實付價","無關（僅上界）：牌價非實付；3.7–12.7 倍（Spot 1.5 倍）；1.8 倍相當於牌價折扣 51–86%")
ev("E13","S9 AWS 公告","OpenAI–AWS 380 億／7 年","DCD https://www.datacenterdynamics.com/en/news/openai-signs-38bn-multi-year-agreement-with-aws-for-access-to-nvidia-gb200s-and-gb300s/","2025-11／讀取 2026-10-24","'$38 billion seven-year agreement'；'hundreds of thousands' of GB200／GB300","380 億／7 年；數十萬顆","無 MW；顆數 Assumed","38÷7÷(顆數÷501,768)",None,f("e13_r2"),f("e13_r3"),f("e13_r1"),"Interested-party（公告）+ Assumed 顆數","無關（顆數 Assumed，倍數 0.5–1.5 太寬）")
ev("E15","S10 GW 與顆數","OpenAI 2025 平均上線 GW 的新佐證","Altman X 2025-07-20（Tom's Hardware）；Epoch AI 方法頁 https://epoch.ai/data/ai-chip-users-documentation/methodology","2025-07／2026","'we will cross well over 1 million GPUs brought online by the end of this year!'；Epoch 年底 H100e 2024 38.4 萬、2025 174 萬","見換算頁","年底時點值；Epoch 認為 GW 為 IT 口徑（約 90% 機率）","前重情境上限＝0.6＋1.3×(1−0.35)",None,None,None,None,"Interested-party（OpenAI／Altman）+ Analogy（Epoch）","證偽條件 (a) 未觸發：平均 GW 上限約 1.45，遠低於 1.8")
for r in range(2, we.max_row+1):
    for c in range(1, len(H)+1):
        we.cell(r,c).alignment = Alignment(wrap_text=True, vertical="top")
    for c in (10,11,12,13): we.cell(r,c).number_format = "0.00"
for i,w in enumerate((6,16,26,46,18,52,24,30,30,12,12,10,10,34,50),1): we.column_dimensions[get_column_letter(i)].width = w

# ---------------- 未找到 ----------------
wn = wb.create_sheet("未找到與不可分離")
wn.append(["編號","類別","項目","查詢／處理","結果"])
for c in wn[1]: c.font = BOLD; c.fill = HF
NF = [
("N1","已搜尋，未找到","OpenAI 2025 在 Azure 的實付每 GPU 小時","OpenAI price per GPU hour Azure discount；OpenAI pays Microsoft hourly；Zitron 文件頁（wheresyoured.at）全文","無任何 2025 世代單價；Zitron 文件只有季別金額，無容量"),
("N2","已搜尋，未找到","Microsoft 對 OpenAI 的關聯方定價（成本加成／最惠條款）","cost plus／at cost／most favored；Musk v. Altman 公開附件（courtlistener 433688 號 391.62，為 2016 年簡報）","附件為 2016 年 AWS 時期簡報，無 Azure 單價；重組條款頁只有金額"),
("N3","已搜尋，未找到","OpenAI 2025 Azure 帳單金額＋同期分配容量的配對","Microsoft 10-K／10-Q；法說會逐字稿（FY26Q2：RPO 6,250 億，OpenAI 約 45%）","只有金額與 RPO 占比，無容量；10-K 只披露營收 241 億"),
("N4","已搜尋，未找到","2025 年 Azure 承載 OpenAI 的 GW","Azure OpenAI GW；Fairwater OpenAI GW","無；Epoch 稱 Microsoft 自家託管推論最高約占 OpenAI 第一方機隊 25%（方向相反，非 Azure 出租量）"),
("N5","已搜尋，未找到","OpenAI 2025 季度或月度上線 GW","quarterly GW；mid-2025 GW","只有年底端點 0.2／0.6／1.9 GW"),
("N6","已搜尋，未找到","OpenAI–Google Cloud 合約金額與容量","—","9a 亦無；本段無新增"),
("N7","已搜尋，未找到","Nebius–Microsoft 的 MW／顆數","Nebius Microsoft Vineland capacity MW","金額 174 億（可到 194 億）／5 年，容量未揭露 → 無法換算"),
("N8","已搜尋，未找到","GB200 v6 三年預留牌價；H100 v5 官方價格頁逐字","azure.microsoft.com 價格頁（JS 頁）；learn.microsoft.com 價格頁","官方價格頁與 vantage.sh 被工作環境阻擋；改用三方追蹤站（azurespeed、sparecores、devzero）"),
("D1","不存在／不可分離","重組條款 2,500 億","Microsoft 官方部落格","金額有，單價與期間無 → 不可換算"),
("D2","不存在／不可分離","Microsoft FY26 對 OpenAI 營收 241 億","10-K","含營收分成與非 Azure 安排；無同期容量 → 不可分離"),
("D3","不存在／不可分離","Anthropic–Azure 300 億、up to 1 GW","pureai 轉述公告","期間未揭露（'up to one gigawatt'）→ 不可換算；且非 OpenAI"),
("D4","不存在／不可分離","Anthropic–Google TPU 數百億、>1 GW","DCD","TPU 非 GPU；金額只有 'tens of billions' → 範圍過寬"),
("D5","不存在／不可分離","CoreWeave–OpenAI 三筆合約 224 億","9a W3-02","只有金額上限；容量未揭露（CoreWeave 全客戶比值見 E07）"),
("D6","不存在／不可分離","SemiAnalysis／InferenceX","—","依任務規定不可單獨作依據；本段未使用"),
]
for r in NF: wn.append(list(r))
for i,w in enumerate((6,18,46,60,70),1): wn.column_dimensions[get_column_letter(i)].width = w
for row_ in wn.iter_rows(min_row=2):
    for c in row_: c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(os.path.join(D,"Azure實付價_A段.xlsx"))
print("saved")
