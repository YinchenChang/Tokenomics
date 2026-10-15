# v7.0 (work order docs/workorders/20261016_稽核修正2.md r0, section 3 = segment 2-2): generation-price section B (docs/workorders/20261015_世代價格B.md r0 B-1..B-6, revised by section 3)
# plus the small leftovers of batch 1 and 3B (major version: downstream contract section 4 rewritten, new L1 names).
#   2-2-a B-1 Cap_In section L: generation launch gaps CAP_GenGap_GB200／GB300／VR200 (+ _Lo／_Hi), Analogy single sample.
#         B-2 L1_GenPxThr_Rev_<gen>／L1_GenPxThr_Pay_<gen> (annual price change that keeps new-generation revenue (payback margin) per MW equal to the previous generation), no price inside.
#         B-3 contract section 4 (docs/plan/Tokenomics_downstream_contract.md, rewritten outside the workbook).  B-4 DB_Evidence E366–E373 + annotations on E338／E339／E341.
#         B-5 Decisions.  B-6 no existing value changes.
#   2-2-b Interface A231 title text (only the re-solved SLO is stated); L1_FleetMargin and Checks K4 read the re-solved fleet-revenue ratio (new L1_FleetProdRatio) instead of the linear
#         CTL_ProdDerate; L1_FleetBreakeven stays on the _Util layer (label says so); L1_GapUtil = IF_Util ÷ (AL_UtilActual × CTL_ProdDerate) (same basis as Cap_In rows 114／255);
#         Serving C18 fixed at 1.0 (label + note; scenario b_prod_derate drives CTL_ProdDerate); Alloc sensitivity rows "全 GB200"／"全 VR200" are pure-generation constants (block6.py);
#         Calib C68 marker = back-solved (text only); GOV K3 investigation (report only); Decisions Q3 = R1 and J1–J5.
# Every write to an Excel-owned cell is guarded (text not yet present ／ cell still holds its v6.3 value) so a rebuild never overwrites a later Excel edit.
from openpyxl.workbook.defined_name import DefinedName
from common import put, F_IN, F_NOTE, F_BOLD, F_CALC, FILL_KEY
import v602
import v603

VERSION = "20261016_Tokenomics_v7.0"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-16"
V = "v7.0"
WO = "工作單 docs/workorders/20261016_稽核修正2.md r0 第 3 節"
WOB = "docs/workorders/20261015_世代價格B.md r0"
A_REPORT = "docs/reports/20261015_世代價格A_查證.md"
DASH = "—"
UN = "AL_UnavailTxt"
ANDY_GEN = "Andy 2026-10-15 15:22「照預設」（採世代價格 A 段報告第 8 節的 chat 端建議預設）"
ANDY_Q3 = "Andy 2026-10-16 03:10 採 R1"
CHAT_DEF = "chat 端預設，可撤回"

GENS = [("GB200", "Hopper", "Hopper→GB200", 1.88, 1.88, 2.17,
         "低＝雲端開始供應對雲端開始供應：H100（CoreWeave、Cirrascale GA 2023-03-21）→ GB200（CoreWeave GA 2025-02-04）＝1.88 年；高＝量產宣布對量產宣布（2022-09-20 → 2024-11-20）＝2.17 年；基準取低端。單一樣本"),
        ("GB300", "GB200", "GB200→GB300", 0.41, 0.41, 0.54,
         "低＝首家部署對首家部署（GB200 2025-02-04 → GB300 CoreWeave 2025-07-03）＝0.41 年；高＝GB300 實例上線 2025-08-19＝0.54 年；基準取低端。GB200→GB300 是約半年的小改款，Δ 很短，年化價格變化率對 Δ 極敏感"),
        ("VR200", "GB300", "GB300→VR200", 1.18, 1.11, 1.24,
         "基準＝GB300 兩種日期（2025-07-03 首家部署、2025-08-19 實例上線）到 VR200 雲端供應（CoreWeave 2026-09-30）的平均 1.18 年；低＝自 2025-08-19 起 1.11 年；高＝自 2025-07-03 起 1.24 年。VR200 頁面未稱 GA")]
SEC_TITLE = ("L. 世代上市間隔 Δ（世代價格 B 段 B-1，v7.0；Analogy，單一樣本）：L1_GenPxThr_Rev_*／_Pay_* 的分母。欄位：C 基準、D 低、E 高；"
             "來源 docs/reports/20261015_世代價格A_查證.md 第 5 節（基準＝雲端開始供應，兩種日期都記）")
SNAP_MOVES = []      # Serving C18 label is kept (Gov_Map label guards and perf.py use it); J3 is expressed in the F18 note and the scenario input


# ================================================================== B-1: Cap_In section L (below the K section; nothing above moves)
def cap_in_l(wb, k_sec):
    ws = wb["Cap_In"]
    r0 = k_sec["end"] + 2
    assert ws.max_row == k_sec["end"], (ws.max_row, k_sec["end"])
    put(ws, f"A{r0}", SEC_TITLE, F_BOLD)
    rows = {}
    r = r0 + 1
    for g, prev, pair, base, lo, hi, note in GENS:
        put(ws, f"A{r}", f"世代上市間隔 Δ：{pair}（CAP_GenGap_{g}；年）", wrap=True); put(ws, f"B{r}", "年")
        for col, v in zip("CDE", (base, lo, hi)):
            put(ws, f"{col}{r}", v, F_IN, fmt="0.00", fill=FILL_KEY if col == "C" else None)
        put(ws, f"F{r}", "Analogy（單一樣本）")
        put(ws, f"G{r}", f"{note}。來源：{A_REPORT} 第 5 節；用於 L1_GenPxThr_Rev_{g}／_Pay_{g}", F_NOTE)
        wb.defined_names[f"CAP_GenGap_{g}"] = DefinedName(f"CAP_GenGap_{g}", attr_text=f"Cap_In!$C${r}")
        wb.defined_names[f"CAP_GenGap_{g}_Lo"] = DefinedName(f"CAP_GenGap_{g}_Lo", attr_text=f"Cap_In!$D${r}")
        wb.defined_names[f"CAP_GenGap_{g}_Hi"] = DefinedName(f"CAP_GenGap_{g}_Hi", attr_text=f"Cap_In!$E${r}")
        rows[g] = r; r += 1
    return dict(start=r0, end=r - 1, rows=rows)


# ================================================================== B-2 and 2-2-b: L1 rows (appended last) and patches of existing rows
READ_FIXED = ("固定模型口徑：q 為同一模型、換新硬體的產出比（Arch 三層級為單一時點快照，跨世代不變）。實際各世代上市時，層級代表模型可能變大（見 `ARCH_ScaleGen`）。"
              "基準欄為固定模型口徑；模型隨世代變大的門檻見情境 `d1_model_scale`。若要與層級牌價配對，應讀情境值。"
              "讀法：實證的年價格變化高於本門檻，新世代每 MW 收入（或回本餘裕）高於上一世代。")
LABEL_TAG = "｜固定模型口徑（模型變大見 `ARCH_ScaleGen`／情境 `d1_model_scale`）"


def _thr_expr(kind, g, prev, q, d):
    if kind == "Rev":
        return f"{q}^(-1/{d})-1"
    return f"(L1_HoldEconMW_{g}/L1_HoldEconMW_{prev}/{q})^(1/{d})-1"


def _thr(kind, g, prev, qs, ds):
    """qs／ds: three-name tuples (base, lo, hi). Returns (D, E, F) formulas. E／F are the smaller／larger of the two corner values (q low & Δ low, q high & Δ high), so that
    the L1 column convention E ≤ D ≤ F (Checks H3) holds even where q < 1 flips the order (VR200)."""
    guard = lambda q, d: (f"IF(AND(ISNUMBER({q}),ISNUMBER({d})" + (f",ISNUMBER(L1_HoldEconMW_{g}),ISNUMBER(L1_HoldEconMW_{prev})" if kind == "Pay" else "") + f"),"
                          f"IF(AND({q}>0,{d}>0" + (f",L1_HoldEconMW_{prev}>0" if kind == "Pay" else "") + f"),%s,{UN}),{UN})")
    qb, ql, qh = qs; db, dl, dh = ds
    D = "=" + guard(qb, db) % _thr_expr(kind, g, prev, qb, db)
    c1, c2 = _thr_expr(kind, g, prev, ql, dl), _thr_expr(kind, g, prev, qh, dh)
    ok = f"AND(ISNUMBER({ql}),ISNUMBER({qh}),ISNUMBER({dl}),ISNUMBER({dh})" + (f",ISNUMBER(L1_HoldEconMW_{g}),ISNUMBER(L1_HoldEconMW_{prev})" if kind == "Pay" else "") + ")"
    pos = f"AND({ql}>0,{qh}>0,{dl}>0,{dh}>0" + (f",L1_HoldEconMW_{prev}>0" if kind == "Pay" else "") + ")"
    E = f"=IF({ok},IF({pos},MIN({c1},{c2}),{UN}),{UN})"
    F = f"=IF({ok},IF({pos},MAX({c1},{c2}),{UN}),{UN})"
    return D, E, F


def l1_rows_new():
    none = ("無外部對照", DASH, DASH)
    out = []
    out.append(("FleetProdRatio", "機隊營收重解比（VR200）＝IFW_RevGWFleet_Prod ÷ IFW_RevGWFleet_Util",
                "VR200（Interface 第 10／11／12 欄＝低成本／基準／高成本）；L1_FleetMargin 以此取代線性 CTL_ProdDerate；v7.0（2-2-b）",
                "=IF(AND(ISNUMBER(INDEX(IFW_RevGWFleet_Prod,1,11)),ISNUMBER(INDEX(IFW_RevGWFleet_Util,1,11))),IF(INDEX(IFW_RevGWFleet_Util,1,11)>0,INDEX(IFW_RevGWFleet_Prod,1,11)/INDEX(IFW_RevGWFleet_Util,1,11),AL_UnavailTxt),AL_UnavailTxt)",
                "=IF(AND(ISNUMBER(INDEX(IFW_RevGWFleet_Prod,1,10)),ISNUMBER(INDEX(IFW_RevGWFleet_Util,1,10))),IF(INDEX(IFW_RevGWFleet_Util,1,10)>0,INDEX(IFW_RevGWFleet_Prod,1,10)/INDEX(IFW_RevGWFleet_Util,1,10),AL_UnavailTxt),AL_UnavailTxt)",
                "=IF(AND(ISNUMBER(INDEX(IFW_RevGWFleet_Prod,1,12)),ISNUMBER(INDEX(IFW_RevGWFleet_Util,1,12))),IF(INDEX(IFW_RevGWFleet_Util,1,12)>0,INDEX(IFW_RevGWFleet_Prod,1,12)/INDEX(IFW_RevGWFleet_Util,1,12),AL_UnavailTxt),AL_UnavailTxt)",
                "x", "E＝低成本欄、F＝高成本欄（機隊營收不隨成本欄變，三者通常相等）",
                "_Prod 層是 CTL_ProdDerate 下重解 SLO 後的機隊營收（D7），不等於 _Util × CTL_ProdDerate：基準約 0.782（線性為 0.85）。L1_FleetMargin ＝ IF_Util × 本列 × L × m ÷ L1_FleetBreakeven，"
                "即 L × m × 重解後機隊營收 ÷ 持有成本。對照：付費產能比 AL_ProdRatioPaid（2025 世代組合，用於 m 下限）口徑不同",
                "生產折減 CTL_ProdDerate、重解 SLO、層級組合", "CTL_ProdDerate：Assumed（Serving C28）", *none, "IFW_RevGWFleet_Prod／_Util",
                "L1 本頁；Interface J 節", "顯示用；L1_FleetMargin 的中間量"))
    for g, prev, pair, *_ in GENS:
        q = (f"L1_TokMW_Gen_ratio_{g}", f"L1_TokMW_Gen_ratio_{g}_Lo", f"L1_TokMW_Gen_ratio_{g}_Hi")
        d = (f"CAP_GenGap_{g}", f"CAP_GenGap_{g}_Lo", f"CAP_GenGap_{g}_Hi")
        D, E, F = _thr("Rev", g, prev, q, d)
        out.append((f"GenPxThr_Rev_{g}", f"世代每 MW 收入持平所需的年價格變化率（{pair}）{LABEL_TAG}",
                    f"L1_TokMW_Gen_ratio_{g} ^ (−1 ÷ CAP_GenGap_{g}) − 1；不含價格；B-2（世代價格 B 段）", D, E, F, "%/年（小數）",
                    "E／F＝兩個端點（q 低端＋Δ 低端、q 高端＋Δ 高端）中較小／較大者；" + ("GB200、GB300 的 q 無區間，只隨 Δ" if g != "VR200" else "VR200 的 q 低端 0.915 使收入門檻轉為正值（與 D 欄反向），故以 MIN／MAX 排序"),
                    READ_FIXED, f"L1_TokMW_Gen_ratio_{g}、CAP_GenGap_{g}", "Δ：Analogy（單一樣本）；q：固定模型口徑", *none, f"L1_TokMW_Gen_ratio_{g}；CAP_GenGap_{g}",
                    "L1 本頁；Cap_In L 節", "顯示用；Tokenomics 不提供世代間價格基準，下游以自己的價格口徑比較"))
    for g, prev, pair, *_ in GENS:
        q = (f"L1_TokMW_Gen_ratio_{g}", f"L1_TokMW_Gen_ratio_{g}_Lo", f"L1_TokMW_Gen_ratio_{g}_Hi")
        d = (f"CAP_GenGap_{g}", f"CAP_GenGap_{g}_Lo", f"CAP_GenGap_{g}_Hi")
        D, E, F = _thr("Pay", g, prev, q, d)
        out.append((f"GenPxThr_Pay_{g}", f"世代回本餘裕持平所需的年價格變化率（{pair}）{LABEL_TAG}",
                    f"（L1_HoldEconMW_{g} ÷ L1_HoldEconMW_{prev} ÷ L1_TokMW_Gen_ratio_{g}） ^ (1 ÷ CAP_GenGap_{g}) − 1；不含價格；B-2（世代價格 B 段）", D, E, F, "%/年（小數）",
                    "E／F＝兩個端點（q 低端＋Δ 低端、q 高端＋Δ 高端）中較小／較大者；持有成本比取基準（不交叉組合）",
                    READ_FIXED, f"L1_TokMW_Gen_ratio_{g}、L1_HoldEconMW_{g}／_{prev}、CAP_GenGap_{g}", "Δ：Analogy（單一樣本）；持有成本：3 級來源", *none,
                    f"L1_TokMW_Gen_ratio_{g}；L1_HoldEconMW_{g}；L1_HoldEconMW_{prev}；CAP_GenGap_{g}", "L1 本頁；Cap_In L 節", "顯示用；Tokenomics 不提供世代間價格基準"))
    return out


def l1_patch(R):
    out = []
    for row in R:
        row = list(row)
        key = row[0]
        if key == "FleetBreakeven":
            row[1] = "機隊損益兩平乘積（_Util 層）：令機隊營收（IF_RevGWFleet，_Util 層）＝持有成本所需的「利用率 × 折減 × L × m」（VR200）"
            row[8] = row[8] + "。v7.0 起本列維持 _Util 層（未乘生產折減）；L1_FleetMargin 以重解比 L1_FleetProdRatio 代入生產折減"
        elif key == "FleetMargin":
            row[1] = "機隊營收 ÷ 持有成本（四層口徑：IF_Util × 重解比 L1_FleetProdRatio × L × m ÷ 損益兩平乘積；VR200）"
            row[2] = "VR200 基準成本欄；X14 (c)；v7.0（2-2-b）：生產折減讀重解值（D7），不再線性乘 CTL_ProdDerate"
            ok = lambda be, pr: f"IF(AND(ISNUMBER({be}),ISNUMBER({pr})),IF_Util*{pr}*CTL_PriceLife*CTL_Monetize/{be},AL_UnavailTxt)"
            row[3] = "=" + ok("L1_FleetBreakeven", "L1_FleetProdRatio")
            row[4] = "=" + ok("L1_FleetBreakeven_Hi", "L1_FleetProdRatio_Hi")
            row[5] = "=" + ok("L1_FleetBreakeven_Lo", "L1_FleetProdRatio_Lo")
            row[8] = ("含重解折減：＝Interface 第 210 列 IF_RevGWFleet_Prod（CTL_ProdDerate 下重解 SLO 後）÷ 持有成本 × CTL_PriceLife × CTL_Monetize（Checks K4 比對，應為 0；基準約 3.12，v6.3 線性折減為 3.39）。"
                      "對照（不入外部欄，口徑不同：毛利率 對 倍數）：2025 推論毛利隱含值＝1 − SRC_DEM_004 ÷ SRC_DEM_007，見 L1_Ans5_GM 外部欄（r2 第 11 項）")
            row[9] = "利用率、生產折減（重解）、L、m、持有成本、單價快照"
        elif key == "GapUtil":
            row[1] = "落差分解 4：利用率（IF_Util ÷（實際利用率 × 生產折減 CTL_ProdDerate）；與機隊路線 Cap_In 第 114、255 列同口徑）"
            row[2] = "Alloc_In 新輸入（Assumed）；Serving C28 CTL_ProdDerate；X47；v7.0（2-2-b）"
            row[3], row[4], row[5] = "=IF_Util/(AL_UtilActual*CTL_ProdDerate)", "=IF_Util/(AL_UtilActual_Hi*CTL_ProdDerate)", "=IF_Util/(AL_UtilActual_Lo*CTL_ProdDerate)"
            row[7] = "依實際利用率區間（Alloc_In C24:E24）反向，另乘 CTL_ProdDerate"
            row[8] = ("基準 1.7217（0.6 ÷（0.41 × 0.85））；v6.3 為 1.4634（未乘生產折減，與 Cap_In 第 114、255 列不一致）；v5.35 的 1.4571 誤乘折減於分子。"
                      "L1_GapProduct 隨之改變（1.635 → 1.924）")
            row[10] = "IF_Util：Assumed（K11）；實際利用率、生產折減：Assumed"
        elif key == "GapProduct":
            row[8] = row[8] + "；v7.0：L1_GapUtil 改乘生產折減，乘積隨之變動"
        out.append(tuple(row))
    return out


# ================================================================== text patches (guarded) and Interface／Checks patches
SERV_F18_OLD = "C18 與 C28 是否合併留第 2 批"
SERV_F18_NEW = "J3（第 2 批 2-2 段，v7.0；chat 端預設，可撤回）：C18 固定 1.0，不再作為折減入口；生產折減一律由 C28 `CTL_ProdDerate` 表達，情境 b_prod_derate 的輸入改為 CTL_ProdDerate"
CALIB_J68_OLD = "Hopper EP 跨節點走 IB [Assumed 1.5，區間 1–3]"
CALIB_J68_NEW = ("Hopper EP 跨節點走 IB [倒推（v5.5 建檔時與 C67 手算配對；Excel 內無公式關聯）1.5，區間 1–3]；J4（第 2 批 2-2 段，chat 端預設，可撤回）：不重新反解——"
                 "反解需要一個實測目標，用 MLPerf 或 LMSYS 反解會把 D5 的驗證點變成校準點；等第 2-5 段有獨立驗證點後，若模型 ÷ 實測偏離過大再提出")
G226_TAIL = (" ｜v7.0 B-0 Q2（AF2c-GEN）：g 的現有依據（E338–E341）是換代價格比較，壽命內衰減與世代間價格變化應分開；本版不改 g 的數值與區間，"
             "世代間另以門檻 L1_GenPxThr_Rev_*／_Pay_*（Cap_In L 節的 Δ）表達。E338／E339／E341 已更正或註記（E366–E373）")
A231_OLD = "→ _Prod＝_Util × CTL_ProdDerate →"
A231_NEW = "→ _Prod（營收＝CTL_ProdDerate 下重解 SLO 後的值，D7，非 _Util 的線性折減；token 產能為線性乘法）→"
K4_B = ("L1_FleetMargin 與 Interface 第 210 列（IF_RevGWFleet_Prod，CTL_ProdDerate 下重解 SLO 後）÷ 持有成本（IF_HoldEcon）VR200 基準欄 × CTL_PriceLife × CTL_Monetize 之差 > 1e-9"
        "（1＝不等；v7.0 起讀重解值，不再線性乘 CTL_ProdDerate；應為 0）")
K4_D = ('=IF(AND(ISNUMBER(L1_FleetMargin),ISNUMBER(INDEX(IF_RevGWFleet_Prod,1,11)),ISNUMBER(INDEX(IF_HoldEcon,1,11))),'
        'IF(ABS(L1_FleetMargin-INDEX(IF_RevGWFleet_Prod,1,11)/INDEX(IF_HoldEcon,1,11)*CTL_PriceLife*CTL_Monetize)<=0.000000001,0,1),1)')
K4_E = "L1_FleetMargin；IF_RevGWFleet_Prod[11] ÷ IF_HoldEcon[11] × CTL_PriceLife × CTL_Monetize"


def interface_a231(wb, start):
    c = wb["Interface"].cell(start, 1)
    assert A231_OLD in c.value, c.value[:120]
    c.value = c.value.replace(A231_OLD, A231_NEW)
    return [f"Interface!A{start}: title text (v7.0 2-2-b)"]


def checks_k4(wb):
    ws = wb["Checks"]
    rows = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "K4"]
    assert len(rows) == 1, rows
    r = rows[0]
    assert "L1_FleetMargin" in ws.cell(r, 2).value
    ws.cell(r, 2).value = K4_B; ws.cell(r, 4).value = K4_D; ws.cell(r, 5).value = K4_E
    return r


def text_update(wb, append_text):
    log = []
    c = wb["Serving"]["F18"]
    if isinstance(c.value, str) and SERV_F18_OLD in c.value:
        c.value = c.value.replace(SERV_F18_OLD, SERV_F18_NEW); log.append("Serving!F18: note (v7.0 J3)")
    c = wb["Calib"]["J68"]
    if c.value == CALIB_J68_OLD:
        c.value = CALIB_J68_NEW; log.append("Calib!J68: marker 倒推 (v7.0 J4)")
    ci = wb["Cap_In"]
    for r in range(200, ci.max_row + 1):
        if ci.cell(r, 1).value == "中階 token 牌價年變化率 g（I-1；Assumed）":
            if append_text(ci.cell(r, 7), G226_TAIL): log.append(f"Cap_In!G{r}: note (v7.0 Q2)")
    ev = wb["DB_Evidence"]
    notes = {
        "E338": ("｜v7.0 更正（世代價格 A 段 P-a1，E366）：5.6 Terra 上市價為 2.5/15（2026-07-09），2/12 是 2026-07-30 降價 20% 後的價；2025 年以前的「中階」序列實為旗艦"
                 "（GPT-4／4 Turbo／4o 當時為旗艦；層級名稱在 GPT-5.6 與 GPT-6 之間換位）。原主張、來源與標記欄不改"),
        "E339": ("｜v7.0 更正（P-a2，E367）：「Sonnet 持平」不成立：Sonnet 5（2026-06-30）上市即 2/10，2026-08-10 宣布 2/10 永久（原訂 9/1 的 3/15 取消）；Sonnet 5.5（2026-09-28）亦 2/10。"
                 "原主張、來源與標記欄不改"),
        "E341": ("｜v7.0 註記（P-a4，E369）：原數值無法復現（取點日期與腳本未留存）。對數線性、取點＝上市日、混合價 8:1／3:1：以 E338 原序列重算，2023–2026.8 −47.6%／−42.7%、"
                 "2024.3–2026.8 −16.1%／−8.6%、2025–2026.8 +32.0%／+38.2%；以確認序列重算，−45.8%／−40.8%、−11.0%／−3.0%、+53.0%／+60.2%（A 段報告第 3 節）。"
                 "Cap_In I-1 g 是否調整屬 Project 判斷，本版不改"),
    }
    for r in range(5, ev.max_row + 1):
        eid = ev.cell(r, 1).value
        if eid in notes and append_text(ev.cell(r, 11), notes[eid]): log.append(f"DB_Evidence {eid}: note appended (v7.0 B-4)")
    return log


# ================================================================== Gov_Map: new rows for the Cap_In L inputs, updates of existing rows (guarded)
def gov_update(ws, append_text):
    import v536
    n = 0
    r = v536._gm_row(ws, "Calib", "C68"); assert ws[f"A{r}"].value == "GM346", ws[f"A{r}"].value
    if ws[f"F{r}"].value == "Assumed":
        ws[f"F{r}"].value = "倒推"; n += 1
    r = v536._gm_row(ws, "Serving", "C28")
    n += append_text(ws[f"N{r}"], " ｜v7.0 J3（AF2c-J3）：C28 是生產折減的唯一入口；Serving C18 固定 1.0，情境 b_prod_derate 改以 CTL_ProdDerate 驅動")
    for cell in ("C226", "D226:E226"):
        r = v536._gm_row(ws, "Cap_In", cell)
        n += append_text(ws[f"N{r}"], " ｜v7.0 B-0 Q2（AF2c-GEN）：g 的依據是換代價格比較，壽命內衰減與世代間價格變化應分開；世代間以 L1_GenPxThr_* 門檻表達，g 的數值與區間不變")
    return n


def gov_map_rows(wb):
    ws = wb["Cap_In"]
    r0 = next(r for r in range(200, ws.max_row + 1) if ws.cell(r, 1).value == SEC_TITLE)
    out = []
    for i, (g, prev, pair, base, lo, hi, note) in enumerate(GENS):
        r = r0 + 1 + i
        lab = ws.cell(r, 1).value
        reason = (f"世代價格 B 段 B-1（{ANDY_GEN}）：{pair} 世代上市間隔，基準 {base} 年（{lo}–{hi}）；{note}。標記：工作單寫 Analogy（單一樣本）；Gov_Map 類別暫登 Assumed——"
                  f"Gov_Map 的 Analogy 須有可比 SRC_ID（Checks E2），而上市日期尚無 SRC 紀錄（A 段報告第 5 節；列待 Project 判斷）。只進 L1_GenPxThr_Rev_{g}／_Pay_{g}，不進任何 IF_；"
                  f"{WOB} B-1")
        for cell, role, lo_, hi_, rt in ((f"C{r}", "基準", f"=Cap_In!D{r}", f"=Cap_In!E{r}", f"{lo}–{hi} 年（低＝雲端對雲端／首家部署；高＝量產宣布對量產宣布／實例上線）"),
                                         (f"D{r}:E{r}", "低／高", None, None, f"C{r} 的低、高端點")):
            out.append({'scope': '切片三（v7.0 世代價格 B 段）', 'sheet': 'Cap_In', 'cell': cell, 'label': lab, 'check': lab, 'cls': 'Assumed', 'role': role, 'src': '', 'rel': '',
                        'dec': 'AF2c-GEN', 'lo': lo_, 'hi': hi_, 'rtext': rt, 'reason': reason, 'retag': '', 'seg': '中' if role == "基準" else '—'})
    return out


# ================================================================== DB_Evidence E366–E373 (17 columns A..Q; layout of v531–v536)
_PN = ["P-a1", "P-a2", "P-a3", "P-a4", "P-b1", "P-c1", "P-c2", "P-c3"]
EV = {n: f"E{366 + i}" for i, n in enumerate(_PN)}


def _ev(n, claim, src, tag, param, cur, new, verdict, grade, stance, note=""):
    nt = f"{note}｜暫編 {n}｜{ANDY_GEN}；可撤回｜讀取者：chat 端（{WOB} B-4；來源 {A_REPORT} 第 7 節）"
    return [EV[n], DATE, claim, src, tag, param, cur, new, verdict, V, nt, "已處理", DASH, DASH, "Cap_In I-1 g；L1_GenPxThr_*（對照，不入基準）", grade, stance]


EVIDENCE_V700 = [
    _ev("P-a1", "E338 歷史牌價 11 點中 10 點與官方公告頁一致；5.6 Terra 上市價 2.5/15（2026-07-09），2/12 為 2026-07-30 降價後", "openai.com 各公告頁、developers.openai.com changelog",
        "Interested-party（賣方）", "E338 的更正／補註", "E338：5.6 Terra 2/12（2026-07）", "上市 2.5/15（2026-07-09）；2/12 為降價後", "採納為 E338 的更正登錄", "1", "利害關係方",
        "層級名稱在 GPT-5.6 與 GPT-6 之間換位；GPT-4／4 Turbo／4o 當時為旗艦"),
    _ev("P-a2", "Sonnet 5（2026-06-30）上市 2/10，2026-08-10 轉永久；Sonnet 5.5（2026-09-28）2/10；E339「持平」不成立", "anthropic.com/news/claude-sonnet-5、-5-5；platform.claude.com 定價頁",
        "Interested-party", "E339 更正；Cap_In I-1 g 對照", "E339：Sonnet 3/15 持平", "Sonnet 5／5.5 上市即 2/10", "採納為 E339 的更正登錄", "1", "利害關係方"),
    _ev("P-a3", "Gemini 3 Pro 2/12（2025-11-18）、3.1 Pro 2/12（2026-02-19）、2.5 Pro 1.25/10；1.5 Pro 降價幅度官方確認", "blog.google、ai.google.dev、developers.googleblog.com",
        "Interested-party", "E340 補註", DASH, "2.5 Pro 1.25/10；3／3.1 Pro 2/12", "採納", "1（1.5 Pro 絕對值 3.5/10.5：官方頁缺）", "利害關係方"),
    _ev("P-a4", "E341 無法復現；以確認序列重算見 A 段報告第 3 節", A_REPORT, "Derived", "Cap_In I-1 g 依據", "E341：2025 起 +8%／+12%", "重算 +32%～+60%（5 點，斜率對取點敏感）",
        "待 Project 決定是否以重算取代 E341；本版只註記，不改 g", DASH, "中立"),
    _ev("P-b1", "層級牌價 44 點（OpenAI／Anthropic／Google 旗艦與中階）；GPT-5.6→GPT-6 層級名換位", f"{A_REPORT} 附檔 xlsx「價格點_b」", "Interested-party", "SRC_Price 時間序列（K9 年降幅需要）",
        DASH, "44 點", "登錄 DB_Evidence；SRC_Price 不新增 Active 列（本段不設價格基準，Q3＝算力租賃價為口徑、層級牌價作對照）", "1", "利害關係方"),
    _ev("P-c1", "Epoch 同等能力序列（2023-03→2025-02：37.50→0.18 $/M）；Epoch 2026 速率 13 倍／年；演算法效率約 3 倍／年", "epoch.ai；arXiv 2511.23455v2", "Derived", "g_gen／Q2、Q3",
        DASH, "−79%～−97%／年（含演算法效率與競價）", "登錄為對照，不入基準（g_gen 不設）", "2（研究機構整理）", "中立"),
    _ev("P-c2", "GPU 租賃：H100 Lambda 2.40（2023-05）→2.59→2.99（2024-10）→3.99（現價）；B200 Lambda 4.99（2025-08）→6.69；CoreWeave 現價 H100 6.16／B200 8.60／GB200 10.50；AWS P5 2025-06-01 降 44%",
        "各業者定價頁、部落格", "Interested-party", "SRC_Price（隨需／預留分層）", DASH, "35 筆有數值", "登錄；GB200 與 SRC_PRC_001 同源，勿重複計為獨立來源", "1–2", "利害關係方",
        "GB200 只有 CoreWeave 現價 1 點；GB300、VR200 無公開價"),
    _ev("P-c3", "IREN—Microsoft GB300 長約 $9.7B／5 年／200 MW 關鍵 IT 負載＝9.70 $M/MW-IT/年（與契約對帳提醒一致）", "iren.com 2025-11-03 公告", "Interested-party", "契約第 4 節對帳",
        DASH, "9.70 $M/MW-IT/年", "登錄（只記單價）", "1", "利害關係方"),
]


def evidence_rows():
    return EVIDENCE_V700


# ================================================================== Decisions (10 columns; layout of v530–v603)
_D = [
    ("GEN", "世代價格 B 段：門檻表達、情境標籤、g 不拆（Q1／Q2／Q3／Q6／g_gen）", "2026-10-15", ANDY_GEN,
     "Q1：契約第 4 節『中性 q × p ≈ 1』改為情境標籤『持平情境』，不再是基準；Q2：壽命內衰減（現行 g，用於 L_price）與世代間價格變化分開，本段不改 g 的數值與區間，僅在 Decisions 與 Cap_In I-1 說明欄註明 "
     "g 的現有依據（E338–E341）是換代價格比較；Q3：『每 MW 簽約收入』的價格口徑＝算力租賃價（隨需／預留／長約分開），層級牌價作對照，同等能力 token 價格只作下限情境；g_gen 不設，改輸出門檻；"
     "Q6：L1 只新增門檻列 L1_GenPxThr_Rev_*／_Pay_*（GB200／GB300／VR200，基準／低／高），不新增含價格的每 MW 收入指數。Δ 取 Cap_In L 節（CAP_GenGap_*，Analogy 單一樣本）。"
     "E／F 欄為兩個端點（q 低端＋Δ 低端、q 高端＋Δ 高端）中較小／較大者（VR200 的 q 低端 0.915 使收入門檻為正，與基準欄反向，需 MIN／MAX 才符合 Checks H3 的 E ≤ D ≤ F）。"
     "v7.0 依第 2 批 2-1／2-1b 的修訂：基準版本讀 v6.3、說明文字引用 ARCH_ScaleGen、情境 d1_model_scale 下的門檻列於報告。",
     "Cap_In L 節；L1_GenPxThr_*；DB_Evidence E366–E373（E338／E339／E341 註記）；契約第 4 節", "v7.0 已建", f"{WOB} 第 0–2 節；{WO}"),
    ("Q3R1", "Q3＝R1（RL 與輪數倍數的取向）", "2026-10-16", ANDY_Q3,
     "Q3（Andy 2026-10-16 03:10）：採 R1。RL ÷ 預訓練基準 Luna／Sol 0.3（區間 0.1–0.3 上端）、Astra 1.0（0.1–2）；輪數倍數基準 0.7（0.7–1.7 下端）；不改數值。",
     "Train_In RL 比、Har_In 輪數倍數（數值已於 v6.1 3B 寫入，本筆為補登紀錄）", "v7.0 已登錄（數值不變）", f"{WO} 2-2-b"),
    ("J1", "J1 D1(c) 模型規模倍數的情境值（紀錄）", "2026-10-16", CHAT_DEF,
     "chat 端預設，可撤回：基準一律 1.0（截面維持固定模型口徑）；情境值 r0＝Hopper 0.25／GB200 0.5／GB300 1.0／VR200 2.5（區間 0.1–0.5／0.3–1.0／1.0／1.0–4.4），錨點 R-a（2026 快照對應 GB300）或 R-b（對應 GB200）。"
     "已由 2Q 報告 Q2 修訂為 r1（情境 0.8／0.55／1.0／1.4，區間 Lo 0.5／0.3／1／1、Hi 1.3／1.0／1／2.5），見 AF2-D1c r1。",
     "Spec_Rack 模型規模倍數節；情境 d1_model_scale", "v6.3 已建（本筆補登）", "稽核修正2 工作單第 1 節 J1"),
    ("J2", "J2 D5 獨立驗證點與 MLPerf 輪次（紀錄；第 2-5 段執行）", "2026-10-16", CHAT_DEF,
     "chat 端預設，可撤回：驗證點（模型 ÷ 實測，不用於校準）Hopper Sol＝LMSYS H100 decode 約 2,785 tok/s/GPU、GB200 Sol＝LMSYS GB200 decode 13,386（NVFP4），兩點都標『口徑不嚴（無明確 SLO）』，InferenceX 只作旁證；"
     "MLPerf 上限 Hopper、GB200 用 v5.1、GB300 用 v6.0 Closed，模型高於 MLPerf 即判失敗。第 2-5 段才用到。",
     "第 2-5 段（Checks 重整、MLPerf 上限、獨立驗證點）", "未建（2-5 段）", "稽核修正2 工作單第 1 節 J2"),
    ("J3", "J3 Serving C18 與 C28 重複：只留 C28 一個入口", "2026-10-16", CHAT_DEF,
     "chat 端預設，可撤回：C18 固定 1.0，不再作為折減入口（標籤不變，說明欄 F18 註明；生產折減一律由 C28 CTL_ProdDerate 表達）；情境 b_prod_derate 的輸入由 Serving!C18＝0.7 改為 CTL_ProdDerate＝0.7。"
     "理由：同一件事只留一個入口；同時消除 b_prod_derate 情境 Checks K3 由 57 升到 156 的來源（D 節用 0.7、G 節用 0.85 重解）。C18 儲存格仍可輸入（測試以它誘發 SLO 不可達），只是基準與情境不再用它表達折減。",
     "Serving A18／F18；tests/parity/scenarios.yaml b_prod_derate", "v7.0 已建", "稽核修正2 工作單第 1 節 J3"),
    ("J4", "J4 Calib C68（Hopper 每層延遲倍數 1.5）不重新反解", "2026-10-16", CHAT_DEF,
     "chat 端預設，可撤回：不重新反解，標記改為『倒推（v5.5 建檔時與 C67 手算配對；Excel 內無公式關聯）』，區間維持 1–3。理由：反解需要一個實測目標；用 MLPerf 或 LMSYS 反解會把 D5 的驗證點變成校準點，"
     "違反 D5『校準過的輸入不得再用同一目標對照』。等第 2-5 段有獨立驗證點後，若模型 ÷ 實測偏離過大，再提出。",
     "Calib J68；Gov_Map GM346（類別 Assumed → 倒推）", "v7.0 已建（只改標記）", "稽核修正2 工作單第 1 節 J4"),
    ("J5", "J5 Alloc 世代組合敏感度『全 GB200』『全 VR200』改回純世代", "2026-10-16", CHAT_DEF,
     "chat 端預設，可撤回：『全 GB200』列＝GB200 1、其餘 0；『全 VR200』列＝VR200 1、其餘 0；以常數 1／0 寫入，不再讀 AL_Mix*_Lo／_Hi（v6.1 D3 把它們改成 2025 全年平均區間後，這兩列已不是純世代）。"
     "檢查副本（輸入全為基準）讀基準組合，不受影響。",
     "Alloc G 節『服務世代組合：全 GB200』『全 VR200』（v6.3 第 94、95 列；工作單的第 85、86 列為舊列號）", "v7.0 已建", "稽核修正2 工作單第 1 節 J5"),
    ("D7R", "D7 遺留：L1_FleetMargin、Checks K4、L1_GapUtil 改讀重解值／同口徑", "2026-10-16", "稽核判斷會議決定第 1 批遺留（chat 端定案，併入第 2 批）",
     "L1_FleetMargin 與 Checks K4 由線性 × CTL_ProdDerate 改為讀重解值（新增 L1_FleetProdRatio＝IFW_RevGWFleet_Prod ÷ IFW_RevGWFleet_Util，VR200；基準約 0.782，線性 0.85）；"
     "L1_FleetBreakeven 讀 _Util 層，維持不動，標籤寫明『_Util 層』；L1_GapUtil＝IF_Util ÷（AL_UtilActual × CTL_ProdDerate），與 Cap_In 第 114、255 列同口徑，L1_GapProduct 隨之改變；"
     "Interface 第 231 列標題只寫重解 SLO（D7）。重解比取收入比（Prod 營收 ÷ Util 營收）而非付費產能比（AL_ProdRatioPaid 口徑），使 Checks K4 恆為恆等式；兩者差異見報告待 Project 判斷。",
     "L1_FleetMargin；L1_FleetProdRatio（新）；Checks K4；L1_GapUtil／GapProduct；Interface A231", "v7.0 已建", f"{WO} 2-2-b"),
    ("K3", "GOV K3（f_registry 情境 0→6）查明結果：不修，列待 Project 判斷", "2026-10-16", "稽核判斷會議決定第 1 批遺留",
     "Checks K3 在 f_registry_t07_t09_on 情境為 6：GB200 三欄的 IFW_RevGWFleet 與 IFW_RevGWFleetFront 的 _Prod 營收高於 _Util。原因：_Util 層（Fleet_1GW）GB200 的 Astra 可達，_Prod 層（Fleet_1GW_Prod，CTL_ProdDerate 下重解 SLO）Astra 不可達而被排除並重新正規化，"
     "層級籃子不同（_Prod 少了算力最重的 Astra，付費 token 總量 0.82e9 → 3.42e9），因此 _Prod 營收 1.6055 → 3.0409，單調檢查失敗。這是 v6.0 D8-2 的 Fleet_1GW 層級重新正規化造成，不是 2-1-c（Alloc）處理的同一件事，Alloc 的重新正規化不影響 Fleet_1GW_Prod。未修；列待 Project 判斷。",
     "Checks K3；Fleet_1GW／Fleet_1GW_Prod 第 18–29 列；Interface 第 248–251、268–271 列", "查明完成，未修", f"{WO} 2-2-b；docs/reports/20261015_稽核修正1.md 待判斷第 6 項"),
]
DECISIONS_V700 = [[f"AF2c-{i}", V, f"稽核修正第 2 批 2-2：{t}", txt, dt, andy, st, where, src, "否"] for i, t, dt, andy, txt, where, st, src in _D]

# ================================================================== README
README_VERSION = (VERSION + "（主版：稽核修正第 2 批 2-2 段，" + WO + "——世代價格 B 段：Cap_In L 節世代上市間隔 CAP_GenGap_*、L1_GenPxThr_Rev_*／_Pay_* 門檻列（不含價格；契約第 4 節改寫）、DB_Evidence E366–E373；"
                  "遺留：L1_FleetMargin／Checks K4 讀重解值（新增 L1_FleetProdRatio）、L1_GapUtil 乘生產折減、Interface A231 標題、Serving C18 固定 1.0（b_prod_derate 改驅動 CTL_ProdDerate）、"
                  "Alloc 全 GB200／全 VR200 改純世代、Calib C68 標記倒推；Decisions AF2c-*）。以下為 " + v603.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v603.README_TITLE
README_TITLE = VERSION.split("_")[-1] + _T[len(v603.VERSION.split("_")[-1]):]
README_ROW = ("世代價格門檻、遺留項（v7.0）",
              "稽核修正 2-2：世代間每 MW 收入／回本餘裕持平所需的年價格變化率（L1_GenPxThr_*，不含價格；Δ 在 Cap_In L 節）；L1_FleetMargin、Checks K4 改讀重解值；L1_GapUtil 乘生產折減；"
              "Serving C18 固定 1.0；Alloc 全 GB200／全 VR200 改純世代；Calib C68 標記倒推。標記與理由見 Gov_Map 與 Decisions AF2c-*。")
