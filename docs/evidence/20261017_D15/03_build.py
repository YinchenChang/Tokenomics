"""D15 第 3 步：產生明細 xlsx 與證據表（不改 Excel）。
用法：python3 03_build.py   （讀 01_load.pkl、02_stats.pkl；寫 docs/reports/20261017_D15_來源建議_明細.xlsx、本資料夾 證據表.xlsx、03_counts.txt）
建議的內容寫在本檔的對照表（REC49、BACKLOG_OVR、規則函式），是代理的判讀，不是決定。"""
import pickle, os, re, collections
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import range_boundaries, get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..', '..'))
d = pickle.load(open(os.path.join(HERE, '01_load.pkl'), 'rb'))
s = pickle.load(open(os.path.join(HERE, '02_stats.pkl'), 'rb'))
cells = d['cells']
H = d['gm_hdr']; ix = {h: i for i, h in enumerate(H)}
A = d['a1_hdr']; ia = {h: i for i, h in enumerate(A)}
SRC_COL = 'SRC_ID（來源或可比對象）'
def g(r, k): return r[ix[k]]
def blank(x): return x is None or str(x).strip() in ('', '—', '-', '–')
gm_by = {r[0]: r for r in d['gm']}
LB = {r[0] for r in d['lb']}
MP = collections.defaultdict(set)
for r in d['a1_rows']:
    for x in re.findall(r'GM\d+', str(r[0])): MP[x].add(r[ia['是否在 m 路徑']])

def expand(ref):
    ref = str(ref).replace('$', '')
    if ':' not in ref: return [ref]
    c1, r1, c2, r2 = range_boundaries(ref)
    return [f'{L(c)}{r}' for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
def curval(sheet, ref, maxn=6):
    vs = [cells.get((sheet, a)) for a in expand(ref)]
    vs = ['—' if v is None else (round(v, 6) if isinstance(v, float) else v) for v in vs]
    t = '／'.join(str(v) for v in vs[:maxn])
    return t + (f'（共 {len(vs)} 格）' if len(vs) > maxn else '')
def sens(gid):
    r = gm_by.get(gid)
    return f"Load_Bearing：{'是' if gid in LB else '否'}；分段：{g(r, 'CC 敏感度分段') if r else '—'}"
def is_high(gid):
    r = gm_by.get(gid)
    return gid in LB or (r is not None and g(r, 'CC 敏感度分段') == '高')
def mpath(gid):
    v = MP.get(gid)
    if not v: return '未判（v5.36 後新增或稽核未列）'
    return '是' if v == {'是'} else ('部分格是' if '是' in v else '否')

# ---------------------------------------------------------------- 新決策建議
NEWDEC = [
    ('D15-01', '結構索引與情境選擇器', '世代索引、層級索引、Tech_Registry 低／基準／高選擇器與顯示用常數屬模型結構定義，不需要外部來源；Gov_Map 類別維持「情境選擇」，決策 ID 填本號。', 'Calib C10:I10、C117:F117；Perf／Unit_Cost／Perf_Batch／Training／Sens_* 第 6、7 列；Arch C32'),
    ('D15-02', '敏感度情境值', '敏感度頁（Sens_*、Sensitivity、Calib 格點）的情境值只引用對應基準輸入的區間端點；來源與決策隨基準格，不另立來源。', 'Sens_Perf／Sens_Train／Sens_Rev／Sensitivity 情境列；Calib C90:E90'),
    ('D15-03', 'Derived（公式）列', '以公式由其他輸入或 SRC 推得的列，來源隨上游；Gov_Map H 欄寫上游 SRC 或「上游公式」，關係欄寫「Derived」。', 'Spec_Rack C24、C41；Arch C21:E23；Calib C69'),
    ('D15-04', '資料中心造價與營運分項', 'Inputs 第 14–21、31–35 列為 v4 沿用的分項假設；總額以 JLL（SRC_DC_003）校準。Gov_Map H 欄填 SRC_DC_003、關係「總額校準」，分項本身維持 Assumed 區間。', 'Inputs 第 14–21、31–35 列'),
    ('D15-05', '標準任務與 harness 參數組', 'Workload 第 5–12 列與 Har_In B 節參數組為標準任務定義（Assumed）；校準證據依 L7 規則登錄 DB_Evidence，不直接改基準。', 'Workload 5–12、40；Har_In B 節'),
    ('D15-06', '服務配置係數', 'Serving C9–C15 與 Spec_Rack 第 31 列為服務配置與效率係數（Assumed，區間承載）；與 Calib 擬合共同決定產能，不單獨校準。', 'Serving C9:C15；Spec_Rack C31:G31'),
    ('D15-07', '能量係數', 'Energy 頁係數只供能量閉合檢查與 tokens/J 指標，不進成本與營收；Assumed 區間。', 'Energy C5、C7'),
    ('D15-08', '訓練流程參數', 'Train_In 第 8–15、21、27–38 列為 J8／J9 架構下的訓練流程定義與假設（MFU 係數、取樣數、token 量）；可查證者另列。', 'Train_In 8–15、21、27–38'),
    ('D15-09', 'η_p 倍數擴用 C2', '提議把 Decisions C2「已搜尋、無可比對象」的適用列名追加 Calib D69、E69、G69（GM352、GM353、GM355）；C2 明文規定新列須逐列經 Andy 決定。', 'Calib D69、E69、G69'),
]

# ---------------------------------------------------------------- 49 格建議
# 欄：處置, 建議來源或決策, 標記, 區間, 信心, 是否需 Andy, 理由, 結果類（外部來源／既有 SRC／既有決策／新決策）, 外部是否仍找不到
REC49 = {
 'GM004': ('補來源＋補決策；數值不改', 'SRC 新增建議：Google 2025 全機隊年均 PUE 1.09（datacenters.google/efficiency）；Uptime Institute 2025 全球調查平均 1.54、超大規模 1.10–1.15（mgrid 2025-10-01 轉述，原 PDF 未讀）。決策：AF3B-Q6', 'Analogy（超大規模全機隊與通用機房 → 液冷 AI 設施）；Google 為 Interested-party（自報效率，有展示效率的誘因）', '1.1–1.3（不變）', '中', '否', '兩個獨立來源界定區間：超大規模業者 1.09–1.15 接近低端，通用機房 1.54 遠高於高端；液冷 AI 設施落在兩者之間偏低，現區間合理。Hopper 氣冷口徑已由 AF3B-Q6 決定不另立參數（影響 m_fleet ≤ ±1.3%）。', '外部來源', '否'),
 'GM005': None, 'GM006': None,
 'GM028': ('補決策；數值不改；外部來源仍找不到', '決策：AF3B-Q6（Hopper 氣冷口徑）＋D15-04（造價分項，總額以 SRC_DC_003 校準）', 'Assumed', '2.0–3.2（不變）', '低', '否', '已搜尋：datacentres.com 列機械（冷卻）2–4 $M/MW，但未附原始出處；Cushman & Wakefield 2026 造價指南只見「每 MW 上升 21%」，分項在付費牆後。沒有兩個獨立可用來源，維持 Assumed；現值落在無出處區間 2–4 之內，只作參考。', '新決策＋既有決策', '是'),
 'GM029': None, 'GM030': None,
 'GM203': ('補決策；數值不改；頁內標記改為與 Gov_Map 一致（Assumed）', '決策：D5（代表架構）＋J8（Astra 預訓練規模，未結）；可比對象（Analogy，不作來源）：DeepSeek V4-Pro 1.6T、V4-Flash 284B（SRC_MOD_001／005）、Epoch Grok 3 參數 3e12（E249，出處未明）', 'Assumed', '2–10T（不變）', '中', '是', 'Astra 是封閉前沿代理，廠商未公開架構，外部來源找不到；Gov_Map 為 Assumed、Arch F 欄寫「Verified／Assumed」，以哪一邊為準是稽核第 0 段待判斷第 3 項，屬讀法選擇，且本格為 Load_Bearing 高段。', '既有決策', '是'),
 'GM204': ('補決策；數值不改；頁內標記改為與 Gov_Map 一致（Assumed）', '決策：D5＋J8；可比對象：V4-Pro 啟用 49B（SRC_MOD_002）；GPT-4 外洩約 220–280B（稽核修正 2Q 查證；SemiAnalysis＋Hotz）', 'Assumed', '60–250B（不變）', '中', '是', '同 GM203：外部來源找不到；標記讀法待 Andy；Load_Bearing 高段。', '既有決策', '是'),
 'GM205': ('補決策；數值不改；頁內標記改為與 Gov_Map 一致（Assumed）', '決策：D5＋J8＋X3（6.2 齊備）', 'Assumed', '80–160（不變）', '中', '是', '同 GM203；X3 已判 6.2 齊備，但標記讀法仍未定；Load_Bearing 高段。', '既有決策', '是'),
 'GM206': ('補決策；數值不改', '決策：D5＋J8', 'Assumed', '8,192–16,384（不變）', '中', '否', '封閉模型無公開架構；非高敏感度。頁內標記建議隨 GM203 的決定一併改。', '既有決策', '是'),
 'GM207': ('補決策；數值不改', '決策：D5＋J8', 'Assumed', '64–128（不變）', '中', '否', '同 GM206。', '既有決策', '是'),
 'GM208': ('補決策；數值不改', '決策：D5＋J8', 'Assumed', '128–192（不變）', '中', '否', '同 GM206。', '既有決策', '是'),
 'GM209': ('補決策；數值不改', '決策：D5＋J8；可比對象：V4-Pro 384、V4-Flash 256（SRC_MOD_010／011）', 'Assumed', '256–1,024（不變）', '中', '否', '同 GM206。', '既有決策', '是'),
 'GM210': ('補決策；數值不改', '決策：D5＋J8；可比對象：V4 每 token 6 個（SRC_MOD_012）', 'Assumed', '8–32（不變）', '中', '否', '同 GM206。', '既有決策', '是'),
 'GM212': ('補決策；標「定義」', '決策：D5（沿用 V4 結構，SRC_MOD_013＝1）', '定義（結構選擇）', '無數值區間（結構選擇）', '高', '否', '結構選擇，非可量測數字；以 D5 的代表架構為依據。', '既有決策', '否'),
 'GM213': ('補決策；標「定義」', '決策：J1（Astra KV 情境：全層 MLA 為基準）', '定義（結構選擇）', '由 J1 三個 KV 情境承載', '高', '否', 'J1 已把注意力結構做成情境，本格是基準情境的定義。', '既有決策', '否'),
 'GM214': ('補決策；標「定義」', '決策：J1', '定義（結構選擇）', '由 J1 情境承載', '高', '否', '同 GM213。', '既有決策', '否'),
 'GM215': ('補決策；標「定義」', '決策：J1', '定義（結構選擇）', '由 J1 情境承載', '高', '否', '同 GM213。', '既有決策', '否'),
 'GM216': ('補決策；標「定義」', '決策：J1', '定義（結構選擇）', '由 J1 情境承載', '高', '否', '同 GM213。', '既有決策', '否'),
 'GM570': ('補來源（既有 SRC）；類別改 Derived', 'SRC_MOD_009（KV head 512 維）、SRC_MOD_016（HCA 壓縮 128×）、SRC_MOD_018／019（CSA top-k 1,024 塊、塊長 4）、SRC_MOD_015／054（全注意力層比例）；決策：D5', 'Derived（由 V4 模型卡結構推得）', '4–130 B（不變）', '高', '否', '65 B 由 V4 模型卡的結構參數推得（結構重算 64.37／64.98，差 <1.1%）；上游來源都已在 SRC_Model。Luna／Sol 套到 Hopper 世代的時點問題由 D1(c) 模型規模情境處理。', '既有 SRC', '否'),
 'GM245': ('補決策＋補來源（既有 SRC 作對照）；數值不改', '決策：X12（ISL／OSL 基準不改）、AF2e-M3（κ 已套到價格端）；對照：SRC_DEM_014–017（OpenRouter State of AI，2025-12：每請求輸入 >6K、輸出約 400、程式類 3–4 倍；本段 WebFetch 重讀確認）', 'Assumed（OpenRouter 為 Analogy 對照）', '4,096–16,384（不變）', '中', '否', '市場平均 6K 落在 Luna 區間內；X12 已就讀法作出決定，且稽核 Y6「價格端未同步」已由 AF2e-M3 處理。OpenRouter 數字只有一家（a16z 共同作者，同一資料），不足兩個獨立來源，只作對照。', '外部來源', '否'),
 'GM246': None, 'GM247': None,
 'GM248': ('補決策＋補來源（既有 SRC 作對照）；數值不改', '決策：X12（Luna OSL 下限放寬到 400）、AF2e-M3；對照：SRC_DEM_014–017（OpenRouter 輸出約 400，含推理）', 'Assumed（OpenRouter 為 Analogy 對照）', '400–2,048（不變）', '中', '否', 'X12 已比較市場平均 400 與模型輸出占比 11%，決定基準不改。', '外部來源', '否'),
 'GM249': None, 'GM250': None,
 'GM346': ('補決策；數值不改', '決策：BB2-H4（標記更正為 Assumed、區間 1.0–1.5）、BB-G1、AF2c-J4、C2（已搜尋無可比對象）', 'Assumed（BB2-H4；4-2b-1 的 K3 正在改類別欄）', '1.0–1.5（不變）', '高', '否', '外部來源已由 BB-G1 追查，找不到；標記與區間已由 BB2-H4 決定；是否改值仍是 BB-G1 的待判斷，不在本段範圍。', '既有決策', '是'),
 'GM347': ('補決策；標「定義」', '決策：C2（明文列名 GM346–350）', '定義（1＝沿用 GB300 擬合，NVLink 域內無跨節點延遲）', '1–1.5（不變）', '高', '否', 'C2 已列名 Calib C68:G68，只是 Gov_Map J 欄未填。', '既有決策', '否'),
 'GM348': None, 'GM349': None, 'GM350': None,
 'GM352': ('補決策（擴用 C2）；數值不改', '決策：D15-09（提議把 C2 擴用到本格）；同列 F69 已在 C2（X3）', 'Assumed', '0.5–1.5（不變）', '低', '是', 'C2 明文「之後若有新的列要適用，須逐列列名並經 Andy 決定」；本格屬 Load_Bearing。外部來源：v5.20 X3 已搜尋，無新世代 prefill 效率的同口徑實測。', '新決策', '是'),
 'GM353': None,
 'GM354': ('補決策；數值不改', '決策：C2（X3 已追加 GM354）', 'Assumed', '0.5–1.5（不變）', '高', '否', 'C2 已列名本格，只是 Gov_Map J 欄未填。', '既有決策', '是'),
 'GM355': 'GM352',
 'GM441': ('補決策＋補來源（牌價部分）；數值不改', '決策：X40（區間 0.1–0.4）；牌價事實：OpenAI Batch API「50% cost discount」（developers.openai.com／api／docs／guides／batch）、Anthropic Message Batches「50% of the standard API prices」（platform.claude.com 文件），本段 WebFetch 讀取', 'Assumed（Batch 半價為 Verified 牌價，賣方自訂價屬 Interested-party；占比與企業議價為假設）', '0.1–0.4（不變）', '中', '否', '兩家官方文件互為獨立來源，確認「Batch 五折」這一項事實；加權後的有效折扣仍取決於 Batch 占 token 比例與企業議價，兩者 E333、E334 已搜尋未找到。', '外部來源', '否'),
 'GM442': 'GM441',
 'GM443': ('補決策；數值不改', '決策：AF2e-M1（快取命中同時進價格與產能）；對照：SRC_DEM_046–048（Pegatron 99.3%、Requesty Claude Code 91.9%，程式代理）', 'Assumed', '0.3–0.75（不變）', '中', '否', '程式代理的實測命中率遠高於 0.55，但只代表一種用途；參考請求混合聊天與 API，現值與區間由 X21 決定只登錄不改。', '既有決策', '否'),
 'GM590': ('補決策；數值不改', '決策：AF3B-D3（2025 全年平均區間）、X28（2025 基準年世代組合）', 'Derived（3A 推導）', 'Hopper 0.5–0.7／GB200 0.25–0.4／GB300 0–0.1（不變）', '高', '否', '3A 已查證，OpenAI 無直接揭露；Microsoft 2025-10 首個大規模 GB300 叢集已反映在 GB300 0–10% 區間與年底敏感度情境。', '既有決策', '否'),
 'GM585': ('補決策；數值不改', '決策：A9（API 年化方法）；E260 不採納（X 系列）', 'Assumed', 'C14 的低、高端點（不變）', '中', '否', 'A9 明定 API 為 10 月時點值乘全年平均比例；比例本身無來源。', '既有決策', '否'),
 'GM587': ('補決策；區間是否放寬待判斷', '決策：A10（每則提示 token 數）、X29（高端 6,400 → 11,000）', 'Assumed', '2,000–11,000（現值）；閉合值 L1_TokPerPromptClose 見理由', '低', '是', 'Load_Bearing。v12.0 的 L1_TokPerPromptClose 閉合值為 38,186（區間 11,121–89,220；稽核第 1 段 v5.36 時為 14,426），仍是區間上限 11,000 的 3.5 倍；X29 已放寬一次但仍未涵蓋。閉合解由 Alloc_In C23、C24 決定，不是獨立證據（L1 S84），所以不建議直接以它放寬；是否放寬或改讀法屬改數值，交 Andy。', '既有決策', '是'),
 'GM589': ('補決策；數值不改', '決策：X30（免費占比不改，並列支出口徑對照）、AF2e-M2（免費占比一律 token 口徑）', 'Assumed', 'C16 的低、高端點 0.4–0.8（不變）', '中', '否', '口徑混用已由 AF2e-M2 統一為 token 口徑；數值由 X30 決定不改。', '既有決策', '否'),
}
for k, v in list(REC49.items()):
    if v is None:
        base = {'GM005': 'GM004', 'GM006': 'GM004', 'GM029': 'GM028', 'GM030': 'GM028', 'GM246': 'GM245', 'GM247': 'GM245',
                'GM249': 'GM248', 'GM250': 'GM248', 'GM348': 'GM347', 'GM349': 'GM347', 'GM350': 'GM347', 'GM353': 'GM352'}[k]
        REC49[k] = REC49[base]
    elif isinstance(v, str):
        REC49[k] = REC49[v]
# GM249／GM250 的區間文字不同
REC49['GM249'] = REC49['GM248'][:3] + ('1,024–4,096（不變）',) + REC49['GM248'][4:]
REC49['GM250'] = REC49['GM248'][:3] + ('2,048–8,192（不變）',) + REC49['GM248'][4:]
REC49['GM246'] = REC49['GM245'][:3] + ('8,192–68,000（不變）',) + REC49['GM245'][4:]
REC49['GM247'] = REC49['GM245'][:3] + ('16,384–68,000（不變）',) + REC49['GM245'][4:]

# ---------------------------------------------------------------- 其餘列分類規則
def rnum(ref):
    m = re.match(r'[A-Z]+(\d+)', str(ref).replace('$', ''))
    return int(m.group(1)) if m else 0
def classify(r):
    gid = r[0]; sh = g(r, '工作表'); ref = str(g(r, '格')); cat = str(g(r, '類別')); lab = str(g(r, '列標籤')); rn = rnum(ref)
    col = re.match(r'[A-Z]+', ref).group(0) if re.match(r'[A-Z]+', ref) else ''
    rng = str(g(r, '區間文字'))
    # 回傳：分類, 處置, 決策, 標記, 信心, 理由
    if cat == '情境選擇' or '索引' in lab:
        return '物理常數或定義', '補決策 ID', 'D15-01', '定義', '高', '世代／層級索引或情境選擇器，屬模型結構，不需外部來源。'
    if cat == '情境值':
        dec = 'D15-02'
        if sh == 'Har_In': dec = 'AF3B-Q3＋D15-02'
        if sh == 'Alloc': dec = 'AF2c-J5'
        if sh == 'Cap_In' and rn == 15: dec = 'K4'
        extra = '；範圍含 1 個公式格，Gov_Map 範圍另依 4-0 G1 工程修正' if gid in ('GM424', 'GM425') else ''
        return '物理常數或定義', '補決策 ID', dec, '定義（情境設定）', '高', '敏感度情境值只引用基準輸入的區間端點，來源隨基準格' + extra + '。'
    if cat.startswith('Derived'):
        if sh == 'Spec_Rack' and col == 'G':
            return '模型假設', '補決策 ID', 'D2＋D3', 'Derived', '中', 'Rubin Ultra 推估欄由 VR 值等比推得（D2：推估欄不進基準比較）。'
        src = '；H 欄建議填 SRC_HW_019／SRC_HW_040' if gid == 'GM351' else ''
        if sh == 'Arch': return '物理常數或定義', '補決策 ID', 'J1＋D15-03', 'Derived', '高', 'KV bytes 由 Arch「KV 推導輸入」列以公式推得；Astra 三情境為 J1。'
        return '物理常數或定義', '補決策 ID', 'D15-03', 'Derived', '高', '公式推得，來源隨上游' + src + '。'
    # Assumed／倒推
    if sh == 'Inputs':
        if rn == 8: return '模型假設', '補決策 ID', 'D1', 'Assumed', '中', '機架外 IT 占比屬 IT 電力口徑（D1）下的分攤假設；可用營運商揭露查證，但非高敏感。'
        if rn == 10: return '可查證', '補決策 ID；之後以能量閉合檢查校準', 'D3', 'Assumed', '中', 'D3：運轉模式只影響用電與產出；平均用電比例可由逐步能量模型或營運揭露查證。'
        if rn in (14, 16): return '可查證', '補決策 ID', 'D15-04', 'Assumed', '中', 'v4 沿用的造價分項，可由造價指南查證；非高敏感。'
        if 17 <= rn <= 21: return '可查證', '補決策 ID＋H 欄填 SRC_DC_003（總額校準）', 'D15-04', 'Assumed', '中', '說明欄已寫「總額以 JLL 校準」，但 H 欄未填 SRC_DC_003。'
        if 24 <= rn <= 27: return '可查證', '補決策 ID', 'D4', 'Assumed', '中', '設施各分項年限，屬持有成本口徑（D4）；工程慣例可查。'
        if rn == 28: return '物理常數或定義', '補決策 ID', 'D10', 'Assumed', '高', 'D10 明定 WACC 基準 10%、區間 8–13%；本格即區間端點。'
        if 31 <= rn <= 35: return '可查證', '補決策 ID', 'D15-04', 'Assumed', '中', 'v4 沿用的營運成本分項，可由營運商揭露查證；非高敏感。'
    if sh == 'Spec_Rack':
        if col == 'G' and rn in (11, 12, 13): return '模型假設', '補決策 ID', 'C3＋SL42-F1', 'Assumed', '高', 'C3 已定 Rubin Ultra 價格 ＝ VR200 × 1.5／2.0／2.5；倍數已外露為 SR_RU_PxMult*。'
        if col == 'G': return '模型假設', '補決策 ID', 'D2', 'Assumed', '高', 'Rubin Ultra 為推估欄（D2），不進基準比較。'
        if rn in (8, 10): return '可查證', '補決策 ID', 'D3', 'Assumed', '中', '配電設計功率情境（D3）；Hopper／GB200 的低高端可由 DGX／OEM 規格查證。'
        if rn == 31: return '可查證', '補決策 ID', 'D15-06', 'Assumed', '中', '非專家權重以 FP8 存放，可由 V4 模型卡查證（目前無 SRC 紀錄）。'
        if rn == 32: return '模型假設', '補決策 ID', 'D6', 'Assumed', '中', 'prefill 與 decode 分列（D6）下的實例配置。'
        if rn == 41: return '可查證', '補決策 ID', 'J7', 'Assumed', '中', 'NVFP4 訓練峰值可由 NVIDIA 規格查證；J7：NVFP4 只在 T07（未啟用）。'
    if sh == 'Arch':
        if rn in (36, 37): return '模型假設', '補決策 ID', 'J1', 'Assumed', '高', 'Astra KV 情境參數（J1）。'
        if rn == 15: return '模型假設', '補決策 ID', 'D5＋J8', 'Assumed', '中', 'Astra 封閉前沿代理（D5），無公開架構。'
        if rn == 32: return '物理常數或定義', '補決策 ID', 'D15-01', '定義（顯示用）', '高', '顯示列口徑，無任何引用。'
    if sh == 'Serving':
        if rn == 22: return '模型假設', '補決策 ID', 'AF2-TTFT', 'Assumed', '中', 'TTFT 上限只作檢查（第 2 批 TTFT 判定）。'
        return '模型假設', '補決策 ID', 'D15-06', 'Assumed', '中', '服務配置與效率係數，與 Calib 擬合共同決定產能。'
    if sh == 'Workload':
        if rn == 5: return '模型假設', '補決策 ID', 'C4＋D15-05', 'Assumed', '中', 'C4 已定輪數單位與區間不改（C5:D5）；其餘欄同屬標準任務定義。'
        if rn == 8: return '可查證', '補決策 ID；服務端思考占比待查', 'D15-05', 'Assumed', '中', '稽核 Y5：服務端思考占比待查；可由 OpenRouter 推理 token 占比等查證。'
        if rn == 40: return '物理常數或定義', '補決策 ID', 'D15-05＋L7', '定義（比較基準）', '高', '比較基準：一般聊天＝1。'
        return '模型假設', '補決策 ID', 'D15-05', 'Assumed', '中', '標準任務定義；Stage 2 以 Har_In 證據校準（L7）。'
    if sh == 'Har_In':
        if rn == 12: return '模型假設', '補決策 ID', 'AF3B-Q3', 'Assumed', '高', 'Q3 已定輪數倍數區間 0.7–1.7。'
        return '模型假設', '補決策 ID', 'L1＋D15-05', 'Assumed', '中', 'harness 參數組（L1：以 Har_In 表達）。'
    if sh == 'Tech_Registry':
        if col == 'N': return '物理常數或定義', '補決策 ID', 'G0-6＋L3', '定義（採用比例定義域 0–1）', '高', '採用比例開關。'
        return '模型假設', '補決策 ID', 'L3', 'Assumed', '中', '技術倍數；基準是否作用見 Tech_Registry M 欄（L3）。'
    if sh == 'Train_In':
        if rn == 12: return '可查證', '補決策 ID', 'D15-08', 'Assumed', '中', 'GRPO 每題 16 取樣可由 DeepSeekMath／R1 論文查證（尚無 SRC 紀錄）。'
        if rn == 10: return '物理常數或定義', '補決策 ID', 'D15-08', '定義（方法選擇）', '高', 'KL 懲罰有無是方法選擇。'
        return '模型假設', '補決策 ID', 'D15-08', 'Assumed', '中', '訓練流程假設。'
    if sh == 'Calib': return '模型假設', '補決策 ID', 'D15-02', 'Assumed', '中', '校準頁格點。'
    if sh == 'Energy': return '模型假設', '補決策 ID', 'D15-07', 'Assumed', '中', '能量閉合檢查用係數。'
    if sh == 'NonNV': return '模型假設', '補決策 ID', 'D2', 'Assumed', '高', 'D2：非 NVIDIA 世代只放比例列 [Assumed]。'
    if sh == 'Cap_In':
        if rn == 13: return '模型假設', '補決策 ID', 'AF2e-M1', 'Assumed', '中', '同 C13 的區間端點。'
        if rn in (39, 40, 41, 42): return '可查證', '補決策 ID', 'K10', 'Assumed', '中', '快取儲存成本（K10）；DRAM／SSD 價格可由市場報價查證。'
        if rn in (50, 51): return '可查證', '補決策 ID；可由 Anthropic 官方牌價頁查證快取讀取價', 'K2', 'Assumed', '中', '快取讀取價為公開牌價，可查證；本段依工作單不做外部搜尋。'
    if sh == 'Alloc_In':
        if rn == 6: return '模型假設', '補決策 ID', 'A7', 'Assumed', '中', '研發口徑（A7）下的旗艦家族數。'
        if rn == 7: return '模型假設', '補決策 ID', 'A2＋X3', 'Assumed', '高', 'A2 改版計畫計入；X3 已放寬上限到 6。'
        if rn == 8: return '模型假設', '補決策 ID', 'A3＋C2', 'Assumed', '高', 'A3：不以 59% 校準 k；C2 列名 GM580。'
        if rn == 13: return '模型假設', '補決策 ID', 'A6', 'Assumed', '高', 'A6：成長只作情境。'
    if sh in ('Sens_Rev', 'Sensitivity', 'Sens_Perf', 'Sens_Train'):
        return '物理常數或定義', '補決策 ID', 'D15-02', '定義（情境設定）', '高', '敏感度情境設定。'
    if sh == 'Workload' or sh == 'Unit_Cost': pass
    return '模型假設', '補決策 ID', '（待指派）', cat, '低', '規則未涵蓋，需人工指派。'

# ---------------------------------------------------------------- 證據積壓建議
HIGH_SRC = set(); HIGH_REFS = []
for r in d['gm']:
    if is_high(r[0]):
        for x in re.findall(r'SRC_[A-Z]+_\d+', str(g(r, SRC_COL))): HIGH_SRC.add(x)
        HIGH_REFS.append((r[0], g(r, '工作表'), str(g(r, '格')).split(':')[0]))
def ev_high(row):
    rel = re.findall(r'SRC_[A-Z]+_\d+', str(row[12])); param = str(row[5])
    hits = [x for x in rel if x in HIGH_SRC]
    hits += [gid for gid in re.findall(r'GM\d+', param) if is_high(gid)]
    for gid, sh, c in HIGH_REFS:
        if sh in param and re.search(r'(?<![A-Z])' + c + r'(?!\d)', param): hits.append(gid)
    return sorted(set(hits))

src_status = {}
for i, sv in d['src'].items():
    hdr = d['src_hdr'][sv['sheet']]; src_status[i] = sv['vals'][hdr.index('狀態')]
BACKLOG_OVR = {
 'E006': ('並列', '已由 E248／SRC_MOD_055（Epoch GPT-6 Astra 1.0e27）與 X13 取代；E006 的 2e26–2e27 只作 J8 區間對照，狀態改「已處理」。'),
 'E014': ('採納', '價格 $4／$20 已在 SRC_PRC_012–015（Anthropic 官方牌價，賣方自訂價屬 Interested-party 但即為事實）；「達 Fable 5.1 水準」為能力自稱，不採用。建議判定改「部分採納」並結案。'),
 'E015': ('維持待查', '兩來源互相矛盾（AA 標開放權重、另一來源標閉源）；本段未再搜尋。只影響 Cap_In F 節開放權重欄（非高敏感）。'),
 'E020': ('維持待查', '歸屬（Nvidia 或 Strands）不一，且為 Interested-party；不作基準依據。'),
 'E147': ('維持待查', '11 筆 3 級補登（模型內建知識），待 Stage 2 逐筆查核；不改值。'),
 'E204': ('駁回', '已讀，口徑不同，不作第二來源（H1 規則 (ii)）；判定已明確，只是積壓規則把「非採納」都列入。'),
 'E244': ('駁回', '已被 E245 一手結果取代（X6 結案）。'),
 'E245': ('採納', 'MLPerf v6.1 一手結果，X6 已結案；判定文字未以「採納」開頭而被列入。'),
 'E249': ('並列', '參數 3e12 無出處，只作 GM203 區間內參考（X12）。'),
 'E252': ('並列', '兩處 Epoch 數字（4.6e26 與 3.5e26）並列記錄，SRC_MOD_033 不改（X13 b）。'),
 'E254': ('並列', '硬體錨點作比值口徑對照，絕對規模歸 OpenAI 模型校準（X13 c）。'),
 'E260': ('駁回', '判定已為不採納（維持）。'),
 'E269': ('並列', 'J1 (a) 採用後 SRC_HW_007 已改 Alt，媒體報價只作高端對照。'),
 'E278': ('維持待查', 'neocloud 裸機出租是否付 NVIDIA AI Enterprise 授權仍找不到；只影響 Inputs E33（非高敏感）。'),
 'E283': ('並列', '44M 時點不明、二手；預測值不入序列，只作方案組合變化的定性證據。'),
 'E290': ('並列', '20% 分潤已登錄 SRC_DEM_043；$38B 上限仍待一手，營收維持毛額（X22）。'),
 'E301': ('採納', '已由 E314 查證（Altman 2026-07-14，經 Fortune）；判定欄「待查」字樣未更新。'),
 'E314': ('採納', '判定已為「採用（X27）」，只因文字含「E301『待查』」被規則列入。'),
 'E326': ('採納', '與 E289 同一數字，已有 SRC_DEM_042；不重登。'),
 'E329': ('採納', '作為 H10 2026 份額的依據登錄（r0.1），不登 SRC。'),
 'E332': ('採納', '補到 SRC_DEM_060 備註；無 token 量，不新增。'),
 'E335': ('維持待查', '缺口登錄：典型付費用戶 token 用量仍找不到。'),
 'E337': ('採納', '作為 H10 營收份額依據登錄（Andy 15:49）。'),
 'E338': ('維持待查', '等級 3 彙整，待一手確認後才可入 SRC_Price。'),
 'E350': ('並列', '洩漏數字有爭議（Microsoft 稱不準）；SRC_DEM_018 v5.36 起已只作對照列。'),
}
def backlog_rec(b):
    r = b['row']; eid = r[0]; jd = str(r[8] or '')
    if eid in BACKLOG_OVR: return BACKLOG_OVR[eid]
    if jd.startswith('遷移'):
        rel = re.findall(r'SRC_[A-Z]+_\d+', str(r[12]))
        st = collections.Counter(src_status.get(x, '不存在') for x in rel)
        if all(k == 'Active' for k in st):
            return ('採納', f'Stage 1 遷移紀錄（G0-1），不是新資訊；相關 {len(rel)} 筆 SRC 全為 Active。建議判定欄改「採納（遷移）」或積壓規則排除遷移列。')
        return ('並列', f'Stage 1 遷移紀錄；相關 SRC 狀態 {dict(st)}，非 Active 者已有取代或並列紀錄。')
    if jd.startswith('補登'):
        return ('採納', '補登即已建立 SRC 紀錄（G2）；' + ('等級 2、原文待查，建議在 SRC 備註保留「原文待查」。' if '2 級' in jd else '一手原文已讀。'))
    if jd.startswith(('部分採納', '部分採用')):
        return ('並列', '已採用的部分已入 SRC 或區間；未採用的部分保留為對照，不改基準。判定已明確，只因非「採納」開頭被規則列入。')
    if jd.startswith('並列'):
        return ('並列', '判定已為並列（Alt）。')
    if jd.startswith('登錄 DB_Evidence'):
        return ('採納', '作為依據登錄 DB_Evidence，不登 SRC。')
    return ('維持待查', '規則未涵蓋，需人工判讀。')

# ---------------------------------------------------------------- 第二來源
SECOND = [
 ('SRC_HW_067', 'H200 HBM 容量', '141 GB/GPU', 'NVIDIA H200 產品頁（Interested-party：晶片賣方）', '找到', 'Lenovo Press LP1944「NVIDIA H200 141GB GPU」產品指南：141 GB HBM3e（SXM 與 NVL 兩款），更新日 2025-11-07；https://lenovopress.lenovo.com/LP1944', 'Interested-party（Lenovo：伺服器 OEM，有銷售誘因，但與 NVIDIA 為不同主體）', '高', '建議登錄為第二來源（新 SRC 或填 R 欄）；數值相同，不改值', '否'),
 ('SRC_HW_068', 'H200 HBM 頻寬', '4.8 TB/s/GPU', 'NVIDIA H200 產品頁（Interested-party）', '找到', 'Lenovo Press LP1944：4.8 TB/s（兩款同值），更新日 2025-11-07', 'Interested-party（同上）', '高', '建議登錄為第二來源；數值相同，不改值', '否'),
 ('SRC_PERF_081', 'TensorRT-LLM DeepSeek-R1 H200 8 GPU，並行 1,024：每 GPU 1,436.2 tok/s', '1,436.2 tok/s/GPU（ISL 1K／OSL 2K）', 'NVIDIA TensorRT-LLM 0.21.0rc1 文件（Interested-party：軟體與晶片賣方）；本段 WebFetch 重讀：總 11,489.27 tok/s、每 GPU 1,436.16', '只找到不同口徑對照', '(1) MLPerf v5.1 5.1-0375 HPE Cray XD670（8× H200，Open 組，Offline）8,904.83 tok/s＝每 GPU 1,113.1（MLCommons summary_results.json，2026-10-17 下載）；(2) dstack 2025-03-18：H200 8 GPU SGLang Offline 最高 6,311 tok/s（每 GPU 約 789；輸出 800–3,200）', '(1) Interested-party（HPE：伺服器 OEM；MLCommons 審核過的提交，但屬 Open 組，可改模型與資料）；(2) 中立偏工具商（dstack 為開源調度工具，無 GPU 銷售誘因）', '中', '兩者都不是同口徑（並行數、長度、資料集不同），不能消除 W1；建議登錄 (1) 為「近似口徑對照」（模型 ÷ 實測 1,436 ÷ 1,113＝1.29），V22 仍標單一來源', '否'),
 ('SRC_PERF_060', 'MLPerf v5.1 DeepSeek-R1 server：DGX H200 每 GPU 556 tok/s', '556 tok/s/GPU', '轉引 docs/reports/20261015_稽核修正3A_查證.md（未逐一回到原文）', '未找到（且原值無法在 MLCommons 原始資料中重現）', 'MLCommons v5.1 summary_results.json 的 deepseek-r1 共 17 列，H200 只有 HPE Open Offline 1 列（每 GPU 1,113.1）；沒有任何 H200 Server 列，也沒有 556。其他管道：WebSearch 只出 InferenceX（SemiAnalysis，不可單獨使用）', '—', '中', '工作單只列出前三筆，本筆是代理依「最近新增、無第二來源、作驗證點 V1 對照」選出的第 4 筆。建議 3A 原表回查 556 的出處；在查明前 V1 的對照值標「出處待查」', '是'),
]

# ---------------------------------------------------------------- 搜尋紀錄
SEARCH = [
 (1, 'WebSearch', 'Google data center fleet average PUE 2024 environmental report 1.09', 'Inputs D9:F9（GM004–006）', '找到：datacenters.google/efficiency（WebFetch：2025 全年 1.09；2026 Q2 TTM 1.09）'),
 (2, 'WebSearch', 'Uptime Institute Global Data Center Survey 2025 average annual PUE', 'Inputs D9:F9', '找到：Uptime 2025 報告頁（PDF 未讀）；mgrid 2025-10-01 轉述：平均 1.54，超大規模 1.10–1.15，colo／企業 1.58–1.80'),
 (3, 'WebSearch＋WebFetch', 'data center construction cost per MW mechanical cooling share ...（datacentres.com）', 'Inputs D19:F19（GM028–030）', '只找到無出處數字：機械 2–4 $M/MW、總額 11.3 $M/MW（2026 美國均值）'),
 (4, 'WebSearch＋WebFetch', 'Cushman & Wakefield Data Center Development Cost Guide 2025 cost per MW ...（irei.com）', 'Inputs D19:F19', '只見「每 MW 上升 21%」，分項在付費牆後；未找到'),
 (5, 'WebSearch＋WebFetch', 'OpenRouter State of AI average prompt tokens per request ...', 'Serving C23:E24（GM245–250）', '找到：openrouter.ai/assets/State-of-AI.pdf：輸入 >6K、輸出約 400（2025-12，資料至 2025-11）；程式類 >20K 輸入、為一般的 3–4 倍'),
 (6, 'WebSearch', 'OpenAI Batch API 50% discount ... Anthropic Message Batches API 50% discount', 'Cap_In C12:E12（GM441–442）', '結果多為二手彙整，改讀官方頁'),
 (7, 'WebFetch', 'developers.openai.com/api/docs/guides/batch', 'Cap_In C12', '找到：「50% cost discount compared to synchronous APIs」'),
 (8, 'WebFetch', 'platform.claude.com/docs/en/build-with-claude/batch-processing', 'Cap_In C12', '找到：「All usage is charged at 50% of the standard API prices」'),
 (9, 'WebSearch', 'NVIDIA H200 141GB HBM3e 4.8 TB/s specifications AWS P5e instance', 'SRC_HW_067／068', '多為二手規格頁；改查 OEM 文件'),
 (10, 'WebFetch', 'jarvislabs.ai/gpu/nvidia-h200', 'SRC_HW_067', '只列 141 GB，無頻寬；不採用'),
 (11, 'WebSearch＋WebFetch', 'Lenovo Press ThinkSystem NVIDIA H200 141GB product guide ...', 'SRC_HW_067／068', '找到：LP1944：141 GB HBM3e、4.8 TB/s，更新 2025-11-07'),
 (12, 'WebSearch＋WebFetch', 'SGLang DeepSeek-R1 H200 8 GPU decode throughput ...（dstack）', 'SRC_PERF_081', '找到不同口徑：dstack 2025-03-18 H200 Offline 6,311 tok/s（SGLang）、Online 4,176（TRT-LLM）'),
 (13, 'WebFetch', 'nvidia.github.io/TensorRT-LLM/0.21.0rc1/blogs/Best_perf_practice_on_DeepSeek-R1_in_TensorRT-LLM.html', 'SRC_PERF_081', '重讀原值：總 11,489.27、每 GPU 1,436.16 tok/s，ISL 1K／OSL 2K'),
 (14, 'WebSearch', 'MLPerf Inference v5.1 DeepSeek-R1 H200 server tokens per second results', 'SRC_PERF_060／081', '只出 InferenceX（SemiAnalysis，不可單獨使用）與 TRT-LLM 文件'),
 (15, 'WebFetch', 'raw.githubusercontent.com/mlcommons/inference_results_v5.1/main/summary_results.json', 'SRC_PERF_060', '檔案過大，WebFetch 只讀到前段，無 deepseek-r1 列'),
 (16, 'curl（shell，一次）', '同上 JSON 下載', 'SRC_PERF_060／081', '成功；deepseek-r1 共 17 列，H200 只有 HPE Open Offline 1 列（見 04_mlperf51_deepseek列.txt）'),
]

# ================================================================ 寫明細 xlsx
BOLD = Font(bold=True); WRAP = Alignment(wrap_text=True, vertical='top'); HDR = PatternFill('solid', fgColor='FFDDEBF7')
def sheet(wb, title, note, cols, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.cell(1, 1, note).alignment = WRAP
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(cols), 12))
    ws.row_dimensions[1].height = 60
    for j, c in enumerate(cols, 1):
        x = ws.cell(2, j, c); x.font = BOLD; x.fill = HDR; x.alignment = WRAP
    for i, row in enumerate(rows, 3):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v).alignment = WRAP
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[L(j)].width = (widths[j - 1] if widths and j - 1 < len(widths) else 16)
    ws.freeze_panes = 'A3'
    return ws

COLS = ['頁', '格', 'Gov_Map ID', '標籤', '現值', '是否在 m 路徑', '敏感度（Load_Bearing／Gov_Map 分段）', '建議處置', '建議來源或決策', '標記', '區間', '信心', '是否需 Andy 判斷', '理由']
W = [10, 8, 10, 26, 14, 10, 18, 22, 40, 20, 18, 6, 8, 50]

# 1 m 路徑 49 格
rows49 = []; cnt49 = collections.Counter(); andy49 = []
for r in s['c49']:
    gid = re.findall(r'GM\d+', str(r[ia['GM_ID']]))[0]
    sh, cell = r[ia['頁']], r[ia['格']]
    rec = REC49[gid]
    rows49.append([sh, cell, gid, r[ia['列標籤']], curval(sh, cell), '是', sens(gid), rec[0], rec[1], rec[2], rec[3], rec[4], rec[5], rec[6]])
    cnt49[rec[7]] += 1; cnt49['外部仍找不到＝' + rec[8]] += 1
    if rec[5] == '是': andy49.append((sh, cell, gid, r[ia['列標籤']], rec[6]))

# 2 證據積壓
rows_ev = []; cnt_ev = collections.Counter(); andy_ev = []
for b in s['backlog']:
    r = b['row']; rec, why = backlog_rec(b); hi = ev_high(r)
    andy = '否'
    # 是否需 Andy：高敏感度，且建議會改數值或需選讀法（本段建議皆不改數值；以下列為例外）
    if hi and r[0] in ():
        andy = '是'
    rows_ev.append([r[0], str(r[1]), str(r[2])[:300], str(r[5])[:150], str(r[8])[:200], r[11], '；'.join(b['reasons'])[:250],
                    str(r[12])[:150], ('是：' + '、'.join(hi[:6])) if hi else '否', rec, why, andy])
    cnt_ev[rec] += 1
    if andy == '是': andy_ev.append(r[0])

# 3 第二來源
rows2 = [list(x) for x in SECOND]

# 4 全部無來源也無決策列
rows_all = []; cnt_cls = collections.Counter(); cnt_dec = collections.Counter(); andy_all = []
c49_gm = {re.findall(r'GM\d+', str(r[ia['GM_ID']]))[0] for r in s['c49']}
for r in s['g3']:
    gid = r[0]; sh = g(r, '工作表'); ref = g(r, '格')
    lo, hi_ = g(r, '低'), g(r, '高')
    rng = f'{lo}–{hi_}' if not (blank(lo) and blank(hi_)) else str(g(r, '區間文字'))
    if gid in c49_gm:
        rec = REC49[gid]
        cls = '（m 路徑 49 格；見 1_m路徑49格）'
        row = [sh, ref, gid, g(r, '列標籤'), curval(sh, ref), mpath(gid), sens(gid), rec[0], rec[1], rec[2], rec[3], rec[4], rec[5], rec[6], '49 格']
        if rec[5] == '是': andy_all.append(gid)
    else:
        cls, act, dec, tag, conf, why = classify(r)
        andy = '否'
        row = [sh, ref, gid, g(r, '列標籤'), curval(sh, ref), mpath(gid), sens(gid), act, '決策：' + dec, tag, rng, conf, andy, why, cls]
        cnt_dec[dec] += 1
    cnt_cls[cls] += 1
    rows_all.append(row)

wb = openpyxl.Workbook(); wb.remove(wb.active)
# 0 統計
stat = [
 ['項目', '數量', '說明'],
 ['Gov_Map 列（v12.0）', len(d['gm']), ''],
 ['無來源也無決策（SRC_ID 與決策 ID 皆為「—」）', len(s['g3']), 'v5.36 為 311；v12.0 多 1 列＝Alloc I94:L95（GM681，v6.2 起的敏感度 one-hot 情境）'],
 ['　其中 Load_Bearing', sum(1 for r in s['g3'] if r[0] in LB), ''],
 ['　其中 Gov_Map 分段「高」', sum(1 for r in s['g3'] if g(r, 'CC 敏感度分段') == '高'), ''],
 ['m 路徑高嚴重度格（稽核第 1 段，v5.36）', len(s['c49']), 'v12.0 重新比對：49 格仍全部無來源也無決策（已補 0）；分屬 ' + str(len(c49_gm)) + ' 個 Gov_Map 列'],
 ['　找到外部來源（本段新查）', cnt49['外部來源'], 'PUE 3、ISL／OSL 6、有效折扣 3（只確認 Batch 半價一項）'],
 ['　可引用既有 SRC', cnt49['既有 SRC'], 'Arch C34:D34（V4 模型卡結構參數）'],
 ['　建議既有決策 ID（無外部來源）', cnt49['既有決策'], ''],
 ['　建議新決策 ID', cnt49['新決策'] + cnt49['新決策＋既有決策'], 'D15-09（η_p 擴用 C2）3 格；D15-04＋AF3B-Q6（機械與散熱）3 格'],
 ['　外部來源仍找不到（已搜尋或前版已搜尋）', cnt49['外部仍找不到＝是'], 'Astra 架構 8、C68 1、η_p 4、機械與散熱 3、每則提示 token 2（這些格仍有建議的決策 ID）'],
 ['證據積壓（稽核第 0 段 C2 規則，v12.0）', len(s['backlog']), 'v5.36 為 102；新增 4 列（E326 之後的晚於 SRC 列）'],
] + [[f'　建議「{k}」', v, ''] for k, v in cnt_ev.most_common()] + [
 ['缺第二來源（工作單 4 筆）', len(SECOND), '找到 2（H200 規格）、只找到不同口徑 1、未找到 1'],
 ['需 Andy 判斷（全部）', len(andy49) + len(andy_ev) + 1, '49 格中 ' + str(len(andy49)) + ' 格；積壓 ' + str(len(andy_ev)) + ' 列；第二來源 1 筆（SRC_PERF_060）'],
]
sheet(wb, '0_統計', 'D15 來源補齊：查證與建議（4-2b-2A；讀 v12.0，不改 Excel）。本頁為統計；各頁第 1 列為規則。', stat[0], stat[1:], [44, 10, 80])
sheet(wb, '1_m路徑49格', '規則：稽核第 1 段「輸入判讀」頁中，是否在 m 路徑＝是、嚴重度＝高、SRC_ID 與決策 ID 皆為「—」的格（v5.36 共 49 格；v12.0 重新比對仍為 49 格）。現值讀 v12.0（openpyxl data_only）。「是否需 Andy 判斷」＝Load_Bearing 或分段高，且建議會改變數值或需要選擇讀法。', COLS, rows49, W)
sheet(wb, '2_證據積壓', '規則：稽核第 0 段 C2（狀態待判定／判定含待查／證據日期晚於相關 SRC 文件日期且判定不以採納、採用開頭），以 v12.0 DB_Evidence 重算為 106 列。建議四選一：採納、並列、駁回、維持待查。「高敏感度」＝相關 SRC 或對應參數連到 Load_Bearing 或分段高的 Gov_Map 列。',
      ['ID', '日期', '主張（摘要）', '對應參數', '現判定', '狀態', '積壓原因', '相關 SRC', '是否涉及高敏感度格', '建議', '理由', '是否需 Andy 判斷'], rows_ev, [7, 10, 40, 24, 30, 8, 30, 22, 18, 9, 50, 8])
sheet(wb, '3_第二來源', '規則：工作單第 5 項；前 3 筆為工作單列名，第 4 筆（SRC_PERF_060）為代理選出（見備註）。研究紀律：至少兩個獨立來源；SemiAnalysis（含 InferenceX）不可單獨使用。',
      ['SRC_ID', '指標', '現值', '現有來源與立場', '結果', '候選第二來源', '第二來源的標記與利害關係', '信心', '建議', '是否需 Andy 判斷'], rows2, [12, 30, 16, 30, 14, 50, 30, 6, 40, 8])
sheet(wb, '4_無來源無決策_全部', '規則：v12.0 Gov_Map 中 SRC_ID 與決策 ID 皆為「—」的全部列（312 列）。m 路徑 49 格所在的列照第 1 頁建議；其餘只分類（物理常數或定義／可查證／模型假設）並建議決策 ID，不做外部搜尋。「是否在 m 路徑」取自稽核第 1 段（v5.36）。D15-xx 為本段提議的新決策（見 5_新決策建議），尚未存在於 Decisions 頁。',
      COLS + ['分類'], rows_all, W + [14])
sheet(wb, '5_新決策建議', '本段提議、尚未登錄 Decisions 的決策；由 B 段工作單決定是否登錄。', ['提議編號', '項目', '提議內容', '適用範圍'], [list(x) for x in NEWDEC], [10, 20, 80, 40])
sheet(wb, '6_搜尋紀錄', '本段全部外部查詢（WebSearch／WebFetch；shell 只用一次 curl 下載 MLCommons 原始資料）。「已搜尋，未找到」者列出管道。', ['#', '工具', '查詢或網址', '對應格', '結果'], [list(x) for x in SEARCH], [4, 14, 60, 24, 60])
out = os.path.join(ROOT, 'docs', 'reports', '20261017_D15_來源建議_明細.xlsx')
wb.save(out)

# ================================================================ 證據表（外部來源）
EV = [
 ('Inputs D9:F9 PUE', '1.09（2025 全年；2026 Q2 TTM 1.09）', '全機隊年均 PUE（超大規模，含非 AI 機房）', '2025／2026 Q2', 'https://datacenters.google/efficiency', 'Interested-party', 'Google（營運商自報，有展示效率的誘因）', '是（Uptime）', '已讀（WebFetch 摘要）', '1', 'Analogy：全機隊 ≠ 液冷 AI 設施'),
 ('Inputs D9:F9 PUE', '1.54（全球平均，第六年持平）；超大規模 1.10–1.15；colo／企業 1.58–1.80', '年度調查，營運商自填', '2025', 'https://mgrid.org/2025/10/01/uptime-institute-data-center-pue-stagnation-2025-liquid-cooling/（原報告 https://intelligence.uptimeinstitute.com/resource/uptime-institute-global-data-center-survey-2025）', 'Verified（調查）', 'Uptime Institute（中立調查機構）；本段只讀二手轉述', '是（Google）', '二手已讀；原 PDF 未讀', '2', '超大規模與 colo 範圍在轉述中標「待原文確認」'),
 ('Inputs D19:F19 機械與散熱', '2–4 $M/MW（機械）；總額 11.3 $M/MW', '美國 2026 均值', '2026', 'https://www.datacentres.com/guides/data-centre-costs-explained', 'Assumed（來源未附出處）', '網站未說明資料來源', '否', '已讀（WebFetch 摘要）', '3', '不作來源，只作參考'),
 ('Serving C23:E24 ISL／OSL', '每請求輸入 >6K、輸出約 400；程式類 >20K 輸入、為一般的 3–4 倍', 'OpenRouter 全平台流量', '2025-12（資料至 2025-11）', 'https://openrouter.ai/assets/State-of-AI.pdf', 'Verified-measured（單一平台）', 'OpenRouter（路由商，偏 API 與代理流量）＋a16z（投資方）', '否（同一資料集）', '已讀（WebFetch 摘要）', '1', '既有 SRC_DEM_014–017 同源'),
 ('Cap_In C12 有效折扣', 'Batch API 50% 折扣', '牌價條款', '2026-10 讀取', 'https://developers.openai.com/api/docs/guides/batch', 'Verified（賣方牌價）', 'OpenAI（賣方自訂價）', '是（Anthropic）', '已讀', '1', '只確認 Batch 半價；不含占比與企業議價'),
 ('Cap_In C12 有效折扣', 'Message Batches 收費為標準價 50%', '牌價條款', '2026-10 讀取', 'https://platform.claude.com/docs/en/build-with-claude/batch-processing', 'Verified（賣方牌價）', 'Anthropic（賣方自訂價）', '是（OpenAI）', '已讀', '1', '同上'),
 ('SRC_HW_067／068', '141 GB HBM3e；4.8 TB/s', '每 GPU（SXM 與 NVL 同值）', '2025-11-07 更新', 'https://lenovopress.lenovo.com/LP1944', 'Interested-party', 'Lenovo（伺服器 OEM，銷售誘因；與 NVIDIA 為不同主體）', '是（NVIDIA 產品頁）', '已讀（WebFetch 摘要）', '1', ''),
 ('SRC_PERF_081', '總 11,489.27 tok/s；每 GPU 1,436.16', 'H200 8 GPU，並行 1,024，ISL 1K／OSL 2K，TRT-LLM', '2025', 'https://nvidia.github.io/TensorRT-LLM/0.21.0rc1/blogs/Best_perf_practice_on_DeepSeek-R1_in_TensorRT-LLM.html', 'Interested-party', 'NVIDIA（軟體與晶片賣方）', '—（原來源重讀）', '已讀', '1', '原值確認'),
 ('SRC_PERF_081', '8,904.83 tok/s（8 GPU）＝每 GPU 1,113.1', 'MLPerf v5.1 Open 組 Offline，HPE Cray XD670 8× H200', '2025-09', 'https://raw.githubusercontent.com/mlcommons/inference_results_v5.1/main/summary_results.json（5.1-0375）', 'Interested-party', 'HPE（伺服器 OEM；MLCommons 審核，但 Open 組可改模型）', '近似口徑（非同口徑）', '原始資料已讀（curl 下載）', '1', '模型 ÷ 本值＝1.29'),
 ('SRC_PERF_081', 'Offline 最高 6,311 tok/s（SGLang）；Online 最高 4,176（TRT-LLM）', 'H200 8 GPU；輸出 800–3,200', '2025-03-18', 'https://dstack.ai/blog/h200-mi300x-deepskeek-benchmark/', 'Verified-measured', 'dstack（開源調度工具商，無 GPU 銷售誘因）', '近似口徑（非同口徑）', '已讀（WebFetch 摘要）', '2', '並行數未載明'),
 ('SRC_PERF_060', '找不到 556', 'MLPerf v5.1 deepseek-r1 全部 17 列中無 H200 Server', '2025-09', '同上 summary_results.json', '—', '—', '否', '原始資料已讀', '—', '原值出處待查（3A 原表）'),
]
wb2 = openpyxl.Workbook(); wb2.remove(wb2.active)
sheet(wb2, '證據表', 'D15（4-2b-2A）外部查證證據；格式比照 docs/evidence/20261016_服務條件/證據表.xlsx。', ['題號', '數值', '單位與口徑', '日期', '來源（URL）', '標記', '利害關係方與動機', '是否為獨立的第二來源', '讀取狀態', '來源等級', '備註'], [list(x) for x in EV], [18, 30, 24, 14, 50, 16, 30, 14, 14, 6, 30])
wb2.save(os.path.join(HERE, '證據表.xlsx'))

# ================================================================ 計數輸出
with open(os.path.join(HERE, '03_counts.txt'), 'w') as f:
    def P(*a): print(*a); print(*a, file=f)
    P('49 格結果', dict(cnt49))
    P('49 格需 Andy', len(andy49)); [P('  ', x) for x in andy49]
    P('積壓建議', dict(cnt_ev)); P('積壓需 Andy', andy_ev)
    P('積壓涉及高敏感度', sum(1 for x in rows_ev if x[8] != '否'))
    P('全部列分類', dict(cnt_cls))
    P('其餘列決策建議', dict(cnt_dec.most_common()))
    P('待指派', cnt_dec.get('（待指派）', 0))
