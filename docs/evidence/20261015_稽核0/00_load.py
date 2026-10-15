"""載入 v5.36（公式與值各一份），存成 pickle 供後續腳本使用。"""
import openpyxl, pickle, sys, time
from openpyxl.styles.colors import Color
P='model/20261015_Tokenomics_v5.36.xlsx'
t=time.time()
wf=openpyxl.load_workbook(P,data_only=False)
wv=openpyxl.load_workbook(P,data_only=True)
def fcolor(c):
    f=c.font
    if f is None or f.color is None: return None
    col=f.color
    if col.type=='rgb': return str(col.rgb)
    if col.type=='theme': return f'theme{col.theme}'
    if col.type=='indexed': return f'idx{col.indexed}'
    return None
cells={}  # (sheet,coord)->dict
sheets=[]
for ws in wf.worksheets:
    sheets.append((ws.title,ws.max_row,ws.max_column))
    vs=wv[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if c.value is None: continue
            v=vs[c.coordinate].value
            fill=None
            try:
                if c.fill is not None and c.fill.fill_type=='solid': fill=str(c.fill.fgColor.rgb) if c.fill.fgColor.type=='rgb' else f'theme{c.fill.fgColor.theme}'
            except Exception: fill=None
            cells[(ws.title,c.coordinate)]=dict(f=c.value,v=v,color=fcolor(c),fill=fill,r=c.row,c=c.column,nf=c.number_format)
names={}
for n,dn in wf.defined_names.items():
    names[n]=dn.attr_text
for ws in wf.worksheets:
    for n,dn in ws.defined_names.items():
        names[f'{ws.title}!{n}']=dn.attr_text
pickle.dump(dict(cells=cells,sheets=sheets,names=names),open(sys.argv[1],'wb'))
print(len(cells),len(names),time.time()-t)
