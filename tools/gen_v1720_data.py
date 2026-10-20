"""Generate builder/v1720_data.json (v17.2 W1 second sources, B1) from the A-segment evidence folder.
Usage: python3 tools/gen_v1720_data.py <evidence folder docs/evidence/20261020_W1第二來源A段> <base xlsx>
Reads B1_動作清單.csv (the 147 actions) and the 明細 sheet of the review workbook (publisher stance, read channel, B0 note).
New SRC IDs are numbered after the highest existing number per prefix in the base workbook, in CSV order."""
import csv, json, re, sys
from pathlib import Path
import openpyxl

ev, base = Path(sys.argv[1]), sys.argv[2]
OUT = Path(__file__).resolve().parents[1] / "builder" / "v1720_data.json"
SHEETS = ["SRC_HW", "SRC_DC", "SRC_Model", "SRC_Perf", "SRC_Price", "SRC_Cap", "SRC_Harness", "SRC_Demand"]

rows = list(csv.DictReader(open(ev / "B1_動作清單.csv", encoding="utf-8-sig")))
wbr = openpyxl.load_workbook(ev / "20261020_W1第二來源_A段審閱.xlsx", read_only=True)
ws = wbr["明細"]; it = ws.iter_rows(values_only=True); h = next(it)
det = {d["SRC_ID"]: d for d in (dict(zip(h, r)) for r in it)}

wb = openpyxl.load_workbook(base, read_only=True)
mx = {}
for sh in SHEETS:
    for r in wb[sh].iter_rows(min_row=5, max_col=1, values_only=True):
        m = re.match(r"SRC_([A-Z]+)_(\d+)$", str(r[0] or ""))
        if m: mx[m.group(1)] = max(mx.get(m.group(1), 0), int(m.group(2)))

# Stance by publisher (work order 1: 立場依發布者). Mechanical rule on the A-segment's 發布者立場 text: a publisher that sells the
# hardware／service, submitted the result, or whose position was not checked -> 利害關係方; everything else (media, aggregators,
# independent writers, model hosts, open-source projects) -> 中立.  Listed per record in the report.
INTERESTED = ("賣方", "OEM 賣方", "系統廠", "租賃業者", "業者部落格", "提交者", "內容行銷", "供電設備商", "未查其商業立場", "創投平台")
def stance(txt):
    t = str(txt or "")
    if "非 Alibaba 利害關係方" in t: return "中立"
    return "利害關係方" if any(k in t for k in INTERESTED) else "中立"

def num(s):
    t = str(s).replace(",", "").strip()
    try: return float(t) if "." in t else int(t)
    except ValueError: return str(s).strip()

def grade(r):
    return 2 if r.startswith(("R1", "R2", "R3")) else 3     # stage2 G0-7; WebFetch-read records at most 2 (work order 1)

A, C1, Q11 = [], [], []
for x in rows:
    sid = x["SRC_ID"]; d = det[sid]
    if x["類"] == "A":
        pre = re.match(r"SRC_([A-Z]+)_", sid).group(1); mx[pre] += 1
        new = f"SRC_{pre}_{mx[pre]:03d}"
        rr = x["讀取結果"].strip()
        A.append(dict(orig=sid, id=new, sheet=x["工作表"], row=int(x["列"]), val=num(x["第二來源值"]), unit=x["單位與口徑"],
                      src=f'{x["第二來源發布者"]}｜{x["第二來源網址"]}', grade=grade(rr), stance=stance(d["發布者立場"]),
                      stance_note=str(d["發布者立場"] or "—"), hand=("一手" if rr.startswith("R1") else "二手") + "（WebFetch 讀取，非逐字原文）",
                      read=rr, excerpt=x["摘錄"], date=x["讀取日期"], note=x["備註"], channel=str(d["讀取通道(WebFetch/WebSearch 摘要)"] or "—"),
                      b0=str(d["B0複核"] or "—")))
    elif x["類"] == "C1":
        C1.append(dict(id=sid, sheet=x["工作表"], row=int(x["列"])))
    else:
        assert "Q11" in x["動作"], x
        Q11.append(dict(id=sid, sheet=x["工作表"], row=int(x["列"])))
assert (len(A), len(C1), len(Q11)) == (86, 53, 8), (len(A), len(C1), len(Q11))
json.dump(dict(A=A, C1=C1, Q11=Q11, next_ids=mx), open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(OUT, len(A), len(C1), len(Q11), mx)
