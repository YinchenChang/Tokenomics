# 追蹤 L1_MFleet 的前置格（遞迴）
import openpyxl,re,json,sys
from openpyxl.utils import range_boundaries, get_column_letter
f='v153.xlsx'
wf=openpyxl.load_workbook(f); wv=openpyxl.load_workbook(f,data_only=True)
names={}
for n,d in wf.defined_names.items():
    try: names[n.upper()]=list(d.destinations)
    except: pass
ref_re=re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_\.]*))!\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?|(?<![A-Za-z0-9_!\$\.'])\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?(?![A-Za-z0-9_\(])|([A-Za-z_][A-Za-z0-9_\.]*)")
def cells_of(sh,a,b=None):
    if b is None: return [(sh,a)]
    c1,r1,c2,r2=range_boundaries(f"{a}:{b}")
    if (c2-c1+1)*(r2-r1+1)>400: return [(sh,f"{a}:{b}")]  # big range: keep as one node
    return [(sh,f"{get_column_letter(c)}{r}") for r in range(r1,r2+1) for c in range(c1,c2+1)]
FUNCS=set()
def precedents(sh,addr):
    v=wf[sh][addr].value
    if not (isinstance(v,str) and v.startswith('=')): return [],[]
    s=re.sub(r'"[^"]*"','""',v)
    out=[];used=[]
    for m in ref_re.finditer(s):
        if m.group(3):
            out+=cells_of(m.group(1) or m.group(2),m.group(3)+m.group(4),(m.group(5)+m.group(6)) if m.group(5) else None)
        elif m.group(7):
            out+=cells_of(sh,m.group(7)+m.group(8),(m.group(9)+m.group(10)) if m.group(9) else None)
        elif m.group(11):
            tok=m.group(11)
            if tok.upper() in names:
                used.append(tok)
                for dsh,ref in names[tok.upper()]:
                    ref=ref.replace('$','')
                    if ':' in ref: a,b=ref.split(':'); out+=cells_of(dsh,a,b)
                    else: out+=cells_of(dsh,ref)
    return out,used
start=('Cap_In','C119')
seen={};order=[]
stack=[(start,0)]
while stack:
    (sh,addr),d=stack.pop()
    if (sh,addr) in seen: seen[(sh,addr)]=min(seen[(sh,addr)],d); continue
    seen[(sh,addr)]=d; order.append((sh,addr))
    if ':' in addr: continue
    pre,_=precedents(sh,addr)
    for p in pre: stack.append((p,d+1))
print('nodes',len(seen))
from collections import Counter
print(Counter(s for s,a in seen))
ins=[(s,a) for s,a in seen if ':' not in a and not (isinstance(wf[s][a].value,str) and str(wf[s][a].value).startswith('='))]
print('constant leaves',len(ins)); print(Counter(s for s,a in ins))
json.dump({'nodes':[[s,a,seen[(s,a)]] for s,a in order]},open('trace.json','w'),ensure_ascii=False)
