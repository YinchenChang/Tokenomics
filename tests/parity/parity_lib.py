"""Parity 測試工具：LibreOffice headless 重算（基準）與比對。"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl


REL_TOL = 1e-9
ABS_TOL = 1e-12

# ── 從 legacy/engine/core.py 原樣搬入（修訂八-B）：具名範圍與輸入鍵解析 ──
_REF_RE = re.compile(
    r"^(?:'(?P<q>(?:[^']|'')+)'|(?P<u>[^!']+))!"
    r"\$?(?P<c1>[A-Z]+)\$?(?P<r1>\d+)(?::\$?(?P<c2>[A-Z]+)\$?(?P<r2>\d+))?$"
)


def parse_ref(attr_text: str) -> tuple[str, str]:
    """'Interface!$C$19:$Q$19' → ('Interface', 'C19:Q19')；頁名可含中文或引號。"""
    m = _REF_RE.match(attr_text)
    if not m:
        raise ValueError(f"無法解析範圍：{attr_text!r}")
    sheet = (m["q"] or m["u"]).replace("''", "'")
    ref = f"{m['c1']}{m['r1']}"
    if m["c2"]:
        ref += f":{m['c2']}{m['r2']}"
    return sheet, ref


_KEY_RE = re.compile(r"^(?P<name>[A-Za-z_][A-Za-z0-9_.]*)(?:\[(?P<i>\d+)\])?$")


def resolve_key(names: dict[str, str], key: str) -> tuple[str, str]:
    """情境輸入鍵 → (頁名, 儲存格)。鍵可為 'Sheet!A1'（舊情境）、'NAME'（單格具名範圍）
    或 'NAME[k]'（單列或單欄具名範圍的第 k 格，k 從 1 起算）。"""
    if "!" in key:
        sheet, coord = key.split("!")
        return sheet, coord
    m = _KEY_RE.match(key)
    if not m or m["name"] not in names:
        raise KeyError(f"找不到具名範圍：{key!r}")
    sheet, ref = parse_ref(names[m["name"]])
    if ":" not in ref:
        if m["i"]:
            raise ValueError(f"{key}：單格具名範圍不可加索引")
        return sheet, ref
    if not m["i"]:
        raise ValueError(f"{key}：多格具名範圍必須指定索引，例如 {m['name']}[1]")
    c1, c2 = ref.split(":")
    (a, r1), (b, r2) = (re.match(r"([A-Z]+)(\d+)", x).groups() for x in (c1, c2))
    k = int(m["i"]) - 1
    from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
    if r1 == r2:   # 單列：沿欄展開
        return sheet, f"{gl(ci(a) + k)}{r1}"
    if a == b:     # 單欄：沿列展開
        return sheet, f"{a}{int(r1) + k}"
    raise ValueError(f"{key}：只支援單列或單欄具名範圍")


def read_defined_names_xml(path: Path) -> dict[str, str]:
    """直接解析 workbook.xml 的 definedName，作為獨立於 openpyxl 的對照來源。"""
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("xl/workbook.xml"))
    out = {}
    for el in root.iterfind(".//m:definedNames/m:definedName", ns):
        if el.get("localSheetId") is None:  # 只取活頁簿層級名稱
            out[el.get("name")] = (el.text or "").strip()
    return out


# v15.3 C2：Calib D156（CAL_BatchEtaD GB200）由常數改為公式 =SQRT(C254*D201)（兩個批次口徑反推點的幾何平均）。
# 區間端點情境（batch_gb200_lo／hi、batch_range_lo 等）要把它整格改成指定倍數，所以只此一格允許以值覆寫公式；其他公式格仍一律拒絕。
OVERRIDABLE_FORMULA = {("Calib", "D156")}


def set_inputs(src: Path, dst: Path, inputs: dict[str, object]) -> None:
    """複製活頁簿並改寫輸入格。openpyxl 存檔不帶快取值，LibreOffice 載入時必須重算。"""
    wb = openpyxl.load_workbook(src)
    names = {k: v.attr_text for k, v in wb.defined_names.items()}
    for key, v in inputs.items():          # 鍵：'Sheet!A1'、'NAME' 或 'NAME[k]'（具名範圍優先）
        sheet, coord = resolve_key(names, key)
        cell = wb[sheet][coord]
        if isinstance(cell.value, str) and cell.value.startswith("=") and (sheet, coord) not in OVERRIDABLE_FORMULA:
            raise ValueError(f"{key} 是公式格，不是輸入格")
        cell.value = v
    wb.save(dst)


def lo_recalc(xlsx: Path, outdir: Path) -> Path:
    """以 LibreOffice headless 重算並另存 xlsx，回傳輸出路徑。"""
    outdir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="lo_profile_"))
    try:
        cmd = [
            "soffice", f"-env:UserInstallation=file://{profile}",
            "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), str(xlsx),
        ]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = outdir / xlsx.name
    if not out.exists():
        raise RuntimeError(f"LibreOffice 重算失敗：{r.stdout}\n{r.stderr}")
    return out


def excel_values(recalculated: Path, cells: list[tuple[str, str]]) -> dict[tuple[str, str], object]:
    wb = openpyxl.load_workbook(recalculated, data_only=True)
    return {(s, c): ("" if wb[s][c].value is None else wb[s][c].value) for s, c in cells}


def _is_num(x) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def compare(engine_vals: dict, excel_vals: dict) -> dict:
    """比對；數值：相對誤差 ≤1e-9 或絕對誤差 ≤1e-12；其他型別必須完全相等。"""
    mismatches, max_rel, max_rel_all, max_abs, n_num, n_err = [], 0.0, 0.0, 0.0, 0, 0
    for key, e in engine_vals.items():
        x = excel_vals[key]
        if isinstance(e, str) and e.startswith("#") or isinstance(x, str) and x.startswith("#"):
            n_err += 1
        if _is_num(e) and _is_num(x):
            n_num += 1
            d = abs(e - x)
            scale = max(abs(e), abs(x))
            rel = d / scale if scale else 0.0
            max_rel_all = max(max_rel_all, rel)  # 全部數值格（含絕對誤差極小者）
            if d > ABS_TOL:  # 絕對誤差已在容差內者不計入 max_rel_err
                max_rel = max(max_rel, rel)
            max_abs = max(max_abs, d)
            ok = d <= ABS_TOL or rel <= REL_TOL
        else:
            ok = type(e) is type(x) and e == x
        if not ok:
            mismatches.append((key[0], key[1], e, x))
    return {
        "cells": len(engine_vals), "numeric_cells": n_num, "mismatches": mismatches,
        "max_rel_err": max_rel, "max_rel_err_all_cells": max_rel_all, "max_abs_err": max_abs, "error_value_cells": n_err,
    }


def format_mismatches(mm: list, limit: int = 20) -> str:
    lines = [f"不符 {len(mm)} 格；前 {min(limit, len(mm))} 個：", "頁名 | 座標 | 引擎值 | Excel 值"]
    lines += [f"{s} | {c} | {e!r} | {x!r}" for s, c, e, x in mm[:limit]]
    return "\n".join(lines)
