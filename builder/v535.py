# v5.35 (work order docs/workorders/20261014_v5.35.md r0): three fixes to the quarterly m series (H10) and the split of the lifetime price factor
# L into L_price × L_demote (chat-side judgements J22–J24; Andy 2026-10-14 19:38「3. Assume 0.65，其餘照建議」, 15:49「兩件事我都同意。」).
#   B1 Cap_In H10 (J23): (1) the three-way revenue rows use the m numerator (revenue − ads) as their base, so their sum equals the m_整體 numerator;
#      (2) ads annualised per column: 2026 Q1 0 ／ Q2 0.5 ／ FY2026 1.0 (blue, Assumed); (3) new row「2026 消費訂閱占 ChatGPT 付費 token 份額」0.65
#      (D blue; E／F link D; 2025 column links H8 row 132), used by the consumer／seat token rows. The reserved mid-tier price index row (v5.34 row 203)
#      becomes a link to L_price (display only; not multiplied into anything). The H10 spec of v534.py is patched in place before it is written
#      (patch_h10), so every {key} row reference and the Gov_Map registrations resolve to the shifted rows.
#   B2 Cap_In I section「L 拆分」(after H10, before the reading row; C 基準 ／ D L 低端 ／ E L 高端): inputs g ／ T1 ／ ρ ／ T, k＝LN(1+g), L_price,
#      L_total, L_demote, comparison rows (CTL_PriceLife, E246 analogy, Gov_Map GM593 range reading), must-be-true rows (required L for m＝L1_MFleet
#      and m＝L1_MAPI) and their readings.
#   B3 DB_Evidence E338–E343 (N14–N19); SRC_Demand SRC_DEM_072 (N19 ads run-rate $1B).  SRC_Price is not touched (price history is grade 3).
#   B4 L1_LPrice ／ L1_LDemote ／ L1_LTotal ／ L1_LReqFleet (each with _Lo／_Hi; judgement column for L1_LTotal and L1_LReqFleet).
#   B5 Gov_Map rows for the new blue cells; the five v5.34 Gov_Map rows whose Cap_In cells moved down one row (GM629–GM633) are re-pointed (label
#      guard) and GM629 (the index row) is re-classed「顯示用（公式）」; Decisions X42–X44; README rows.
# Every write to an Excel-owned cell is guarded (ID absent ／ text not yet appended ／ cell still at its old reference), like v518–v534.
from copy import copy
from openpyxl.utils import get_column_letter as L
from common import put, F_IN, F_CALC, F_LINK, F_BOLD, F_NOTE, FILL_KEY
import v518
import v533
import v534

VERSION = "20261014_Tokenomics_v5.35"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-14"
V = "v5.35"
DASH = "—"
WO = "工作單 v5.35 r0"
ANDY_1938 = "Andy 2026-10-14 19:38「3. Assume 0.65，其餘照建議」（文字，未改字；對象＝v5.34 報告「待 Project 判斷」第 1–3 項的 chat 端建議）"
ANDY_1549 = "Andy 2026-10-14 15:49「兩件事我都同意。」（文字，未改字；對象＝chat 端 15:4x 提出的 L 拆分）"
ANDY = f"chat 端建議（「1 GW DC 收入：L 與 m」對話串 2026-10-14），{ANDY_1938}；{ANDY_1549}；可撤回"
_N = ["N14", "N15", "N16", "N17", "N18", "N19"]
EV = {n: f"E{338 + i}" for i, n in enumerate(_N)}          # N14–N19 -> E338–E343 (next free number after E337)


def S(n): return f"SRC_DEM_{n:03d}"


# ================================================================== B3: SRC_Demand SRC_DEM_072 (next free number at v5.34: 072)
SRC_NEW = [
    dict(id=S(72), sheet="SRC_Demand", metric="OpenAI 廣告年化營收", val=1.0, lo=None, hi=None, unit="$B", basis="年化（run-rate）；廣告上線約 200 天後；2026-08 時點值",
         applies="OpenAI", date="2026-08-31", src="CNBC 2026-08-31（二手）", grade=2, stance="利害關係方", stance_note="OpenAI 數字經媒體轉述；有呈現成長的誘因", hand="二手",
         status="Active", ev=EV["N19"], s="U-V535", use="Cap_In H10 廣告營收年化 2026 三欄（Q1 0、Q2 0.5、FY2026 1.0；藍字 Assumed，不連結）的依據",
         note=f"v5.35 新增（{EV['N19']}，暫編 N19）：採用：J23 (2) 廣告年化每欄各自輸入的依據；廣告不進任何 m 的分子；{WO} B3"),
]
SRC_LAST_ID = SRC_NEW[-1]["id"]          # SRC_DEM_072


def src_new_records():
    return SRC_NEW


# ================================================================== B1: H10 patch (applied to v534.SPEC_H10 before v534.cap_in_h8_h10 writes it)
LAB_ADS = "廣告營收（2026 起；每欄各自輸入；不進任何 m 的分子；v5.35 J23 (2)；Assumed）"
LAB_CTS26 = "2026 消費訂閱占 ChatGPT 付費 token 份額（J23 (3)；Assumed；2025 連結 H8）"
LAB_IDX_OLD = v534.LAB["q_idx"]          # "中階牌價指數（預留列，v5.35 候選；先填 1，不進任何計算）" — the four blue 1s are retired (now formulas)
LAB_IDX = "L_price（I 節；顯示用，不進 m）"
SNAP_RETIRED = [("Cap_In", (LAB_IDX_OLD, 0), c) for c in (3, 4, 5, 6)]      # build.py: logged, not restored (the cells are formulas from v5.35)
_PATCHED = False


def _idx(spec, key):
    return next(i for i, it in enumerate(spec) if it[0] == key)


def patch_h10():
    """Edit v534.SPEC_H10 in place (idempotent): ads per column (blue), revenue rows on the m numerator, the 0.65 share row, the index row label."""
    global _PATCHED
    if _PATCHED: return
    sp = v534.SPEC_H10
    c3 = v534._h10_cols
    sp[0] = (None, sp[0][1] + "｜v5.35 B1（J23）：三分營收改以 m 分子營收（營收 − 廣告）為基底，三分之和＝m_整體分子；廣告每欄各自輸入（Q1 0、Q2 0.5、FY2026 1.0）；"
             "2026 三欄的消費訂閱占付費 token 份額改為獨立輸入 0.65（0.5–0.8）；第「L_price」列改為連結 I 節（顯示用）")
    sp[_idx(sp, "q_ads")] = ("q_ads", LAB_ADS, "$B", 0, 0, 0.5, 1.0, "Assumed（D–F）",
                             f"2025 為 0（常數，非輸入）；2026 三欄藍字各自輸入（登錄 Gov_Map）：Q1 0（0–0.2）、Q2 0.5（0.3–0.8）、FY2026 1.0（0.5–1.5）；依據 {S(72)}（2026-08 廣告年化達 $1B，{EV['N19']}）；"
                             "H8 第 {ads} 列的 0.5 保留為 H8 的單值，H10 不再連結")
    sp[_idx(sp, "q_rev_m")] = ("q_rev_m", "m 分子營收＝營收 − 廣告", "$B", *c3("={c}{q_rev}-{c}{q_ads}"), "Derived",
                               "m_整體與 m_fleet 的分子；v5.35 J23 (1)：三分營收亦以本列為基底，三分之和＝本列（v5.34 以含廣告營收為基底，三分之和 ≠ m_整體分子）")
    sp[_idx(sp, "q_rev_cons")] = ("q_rev_cons", "消費訂閱營收＝m 分子營收 × 份額", "$B", *c3("={c}{q_rev_m}*{c}{q_sh_cons}"), "Derived", "v5.35：基底改為 m 分子營收（2025 欄廣告 0，數值不變）")
    sp[_idx(sp, "q_rev_seat")] = ("q_rev_seat", "商用席次營收＝m 分子營收 × 份額", "$B", *c3("={c}{q_rev_m}*{c}{q_sh_seat}"), "Derived", "v5.35：基底改為 m 分子營收")
    sp[_idx(sp, "q_rev_api")] = ("q_rev_api", "API 營收＝m 分子營收 × 份額", "$B", *c3("={c}{q_rev_m}*{c}{q_sh_api}"), "Derived", "v5.35：基底改為 m 分子營收；2026 Q2：5.26（v5.34 為 5.36）")
    i = _idx(sp, "q_tok_free") + 1
    if sp[i][0] != "q_cts26":
        sp.insert(i, ("q_cts26", LAB_CTS26, "%", "=$C${cts}", 0.65, "=D{q_cts26}", "=D{q_cts26}", "Derived（C 連結 H8）／Assumed（D；E／F 連結 D）",
                      "藍字一格（D；登錄 Gov_Map）：2026 三欄 0.65（0.5–0.8，J23 (3)）；2025 欄連結 H8 第 {cts} 列 0.75；2026 企業席次增加後消費份額應低於 2025（v5.34 待判斷第 7 項）"))
    sp[_idx(sp, "q_tok_cons")] = ("q_tok_cons", "消費訂閱 token＝ChatGPT 付費 token × 消費份額（第 {q_cts26} 列）", "M tok/年", *c3("={c}{q_tok_chat}*{c}{q_cts26}"), "Derived（輸入 Assumed）",
                                  "v5.35：2026 三欄用 0.65，2025 欄仍為 H8 的 0.75")
    sp[_idx(sp, "q_tok_seat")] = ("q_tok_seat", "商用席次 token＝ChatGPT 付費 token ×（1 − 消費份額）", "M tok/年", *c3("={c}{q_tok_chat}*(1-{c}{q_cts26})"), "Derived（輸入 Assumed）", "")
    sp[_idx(sp, "q_idx")] = ("q_idx", LAB_IDX, "x", 1, "=C{q_idx}", "=C{q_idx}", "=C{q_idx}", "Derived（連結 I 節）",
                             "v5.35 B2：C 欄＝I 節 L_price（cap_in_i 寫入連結）；D–F 連結 C；顯示用，不進任何計算（v5.34 的四個藍字 1 退役）")
    sp[_idx(sp, "q_m_api")] = sp[_idx(sp, "q_m_api")][:-1] + ("錨點：2025 0.564、2026 Q1 0.295、Q2 0.340、FY2026 0.348（v5.35 B1）",)
    sp[_idx(sp, "q_m_cons")] = sp[_idx(sp, "q_m_cons")][:-1] + ("錨點：2025 4.261、2026 Q1 5.744、Q2 5.399、FY2026 5.569（v5.35 B1）",)
    sp[_idx(sp, "q_m_seat")] = sp[_idx(sp, "q_m_seat")][:-1] + ("錨點：2025 1.826、2026 Q1 3.556、Q2 4.557、FY2026 8.044（v5.35 B1）",)
    sp[_idx(sp, "q_m_all")] = sp[_idx(sp, "q_m_all")][:-1] + ("錨點：2025 1.743（＝L1_MImplied）、2026 Q1 1.192、Q2 1.343、FY2026 1.432（分子扣廣告 0／0.5／1.0）",)
    v534.BLUE_H10.update({("q_ads", "D"), ("q_ads", "E"), ("q_ads", "F"), ("q_cts26", "D")})
    v534.BLUE_H10.difference_update({("q_idx", c) for c in "CDEF"})
    v534.LAB["q_idx"] = LAB_IDX
    _PATCHED = True


# ================================================================== B2: Cap_In I section (after H10; the reading row is cleared and rewritten after it)
FMT = {"$B": "#,##0.000", "%": "0.0%", "x": "0.000", "年": "0"}
LAB_I = dict(
    i_g="中階 token 牌價年變化率 g（I-1；Assumed）",
    i_t1="降級前服務中階的年數 T1（年；I-2；Assumed）",
    i_rho="降級後價值比 ρ（I-3；Assumed；D 欄＝IF_PriceRef_Luna ÷ IF_PriceRef_Sol）",
)
_LN = "LN(1+{c}{i_g})"
_LP = f"=IF(ABS({_LN})<0.000000000001,1,((1+{{c}}{{i_g}})^{{c}}{{i_T}}-1)/({{c}}{{i_T}}*{_LN}))"
_LT = (f"=IF(ABS({_LN})<0.000000000001,({{c}}{{i_t1}}+{{c}}{{i_rho}}*({{c}}{{i_T}}-{{c}}{{i_t1}}))/{{c}}{{i_T}},"
       f"(((1+{{c}}{{i_g}})^{{c}}{{i_t1}}-1)+{{c}}{{i_rho}}*((1+{{c}}{{i_g}})^{{c}}{{i_T}}-(1+{{c}}{{i_g}})^{{c}}{{i_t1}}))/({{c}}{{i_T}}*{_LN}))")
_REQ = lambda m: f'=IF(AND(ISNUMBER(L1_FleetBreakeven),ISNUMBER({m})),L1_FleetBreakeven/(IF_Util*CTL_ProdDerate*{m}),"m 非數字（本情境）")'
_JUDGE = lambda k: f'=IF(ISNUMBER({{c}}{{{k}}}),IF({{c}}{{i_ltotal}}>={{c}}{{{k}}},"回本（L_total 口徑）：L_total ≥ 所需 L","未回本（L_total 口徑）：L_total < 所需 L"),"m 非數字（本情境）")'


def _c3(t): return (t.replace("{c}", "C"), t.replace("{c}", "D"), t.replace("{c}", "E"))


# (key, label, unit, C, D, E, tag, note); key None = sub-header. {gm_lo}／{gm_hi} = Gov_Map GM593 low／high cells (CTL_PriceLife range).
SPEC_I = [
    (None, f"I. L 拆分（v5.35；{WO} B2，J22）：L_total＝L_price × L_demote。L_price 由中階 token 牌價年變化率 g 推得（T 年連續平均）；L_demote 由「前 T1 年服務中階、之後降級為低階（價值比 ρ）」推得。"
           "三者為追蹤列與 L1 名稱（L1_LPrice／L1_LDemote／L1_LTotal），CTL_PriceLife 基準 1 不改、Gov_Map 區間 0.3–1.0 不改。欄位：C 基準、D L 低端、E L 高端"),
    ("i_g", LAB_I["i_g"], "%", -0.10, -0.25, 0.05, "Assumed",
     f"藍字三格（登錄 Gov_Map）：基準 −0.10、低 −0.25、高 +0.05。依據：chat 端 2026-10-14 牌價序列迴歸（{EV['N17']}；序列 {EV['N14']}–{EV['N16']}，等級 3 待一手確認）：OpenAI 中階 2023–2026.8 年變化 −49%／−38%"
     "（參考請求權重／3:1 權重）、2024.3–2026.8 −24%／−10%、2025–2026.8 +8%／+12%；Sonnet 持平；Gemini Pro −18%／+11%。基準取近兩年中值附近、高端取近一年持平略升、低端取 2027 後供給釋放與中國廠商競價情境"),
    ("i_t1", LAB_I["i_t1"], "年", 4, 3, 5, "Assumed", "藍字三格（登錄 Gov_Map）：基準 4、區間 3–5；對照 B4_Life 商業壽命 12 月為模型更替，T1 為硬體所服務層級的更替"),
    ("i_rho", LAB_I["i_rho"], "x", 0.5, "=IF_PriceRef_Luna/IF_PriceRef_Sol", 1.0, "Assumed（C、E）／Derived（D）",
     "藍字兩格（C、E；登錄 Gov_Map）：基準 0.5＝降級後保留一半價值；高端 1.0＝不降級；低端 D 欄＝直接以低階單價計（IF_PriceRef_Luna ÷ IF_PriceRef_Sol ≈ 0.05，公式）"),
    ("i_T", "壽命 T（年；I-4）＝IF_DeprLifeIT（Interface I 節 IT 折舊年限，Hopper 基準欄；三欄同值）", "年", *_c3("=INDEX(IF_DeprLifeIT,1,2)"), "Derived（連結 Interface）",
     "IF_DeprLifeIT 為 Interface 第 219 列 C:Q（世代 × 成本情境）；取第 2 欄（Hopper 基準）＝6；各世代基準欄皆為 6（高成本情境為 4，不取）"),
    (None, "I-5. 公式列（k＝LN(1+g)；|k| < 1e-12 時以極限式處理：L_price＝1、L_total＝(T1 ＋ ρ(T − T1)) ÷ T）；錨點 基準 L_price 0.7412、L_total 0.6426、L_demote 0.867"),
    ("i_k", "k＝LN(1 ＋ g)", "x", *_c3("=" + _LN), "Derived", "基準 −0.1054"),
    ("i_lprice", "L_price＝((1+g)^T − 1) ÷ (T × k)　→ L1_LPrice", "x", *_c3(_LP), "Derived", "壽命期內牌價相對第 0 年的連續平均；基準 0.7412、D 0.4762、E 1.1618"),
    ("i_ltotal", "L_total＝[((1+g)^T1 − 1) ＋ ρ × ((1+g)^T − (1+g)^T1)] ÷ (T × k)　→ L1_LTotal", "x", *_c3(_LT), "Derived", "前 T1 年中階牌價、之後 ρ × 中階牌價的連續平均；基準 0.6426、D 0.342、E 1.1618"),
    ("i_ldemote", "L_demote＝L_total ÷ L_price　→ L1_LDemote", "x", *_c3("={c}{i_ltotal}/{c}{i_lprice}"), "Derived", "基準 0.867、D 0.718、E 1.0（不降級）"),
    (None, "I-6. 對照"),
    ("i_ctl", "對照：CTL_PriceLife（Cap_In C63；基準 1，本單不改）", "x", *_c3("=CTL_PriceLife"), "Derived（連結）", "L_total 與 CTL_PriceLife 的差即「模型現行假設（不跌價）」與「牌價序列推得」的差"),
    ("i_e246", "對照：E246 類比區間（H100 租金指數，摘要級；DB_Evidence E246）", "x", "0.42–0.70", DASH, DASH, "Analogy（文字）", "v5.22 X7 的類比證據；L_total 基準 0.6426 落在其內"),
    ("i_judge", "判讀：L_total 對 Gov_Map CTL_PriceLife 區間（GM593 低／高）", "x",
     *_c3('=IF(AND({c}{i_ltotal}>={gm_lo},{c}{i_ltotal}<={gm_hi}),"區間內（"&TEXT({gm_lo},"0.0")&"–"&TEXT({gm_hi},"0.0")&"）","區間外（"&TEXT({gm_lo},"0.0")&"–"&TEXT({gm_hi},"0.0")&"）")'),
     "Derived", "文字；區間 0.3–1.0 由 Gov_Map GM593 K／L 讀取"),
    (None, "I-7. 「必須為真」：回本所需 L＝L1_FleetBreakeven ÷（IF_Util × CTL_ProdDerate × m）；與 H6／H7 的 m「必須為真」並列。D 欄 m 取高端（所需 L 低）、E 欄 m 取低端"),
    ("i_req_fleet", "回本所需 L（m＝L1_MFleet；D 欄 m＝L1_MFleet_Hi、E 欄 L1_MFleet_Lo）　→ L1_LReqFleet", "x", _REQ("L1_MFleet"), _REQ("L1_MFleet_Hi"), _REQ("L1_MFleet_Lo"), "Derived",
     "基準 0.3645（m 0.8098）；L1_FleetBreakeven 0.1505、IF_Util 0.6、CTL_ProdDerate 0.85"),
    ("i_req_api", "回本所需 L（m＝L1_MAPI；D 欄 m＝L1_MAPI_Hi、E 欄 L1_MAPI_Lo）", "x", _REQ("L1_MAPI"), _REQ("L1_MAPI_Hi"), _REQ("L1_MAPI_Lo"), "Derived", "基準 0.5237（m 0.5637）"),
    ("i_judge_fleet", "判讀：L_total ≥ 所需 L（m＝L1_MFleet）；同欄比較", "x", *_c3(_JUDGE("i_req_fleet")), "Derived", "文字；基準 0.6426 ≥ 0.3645 → 回本（L_total 口徑）"),
    ("i_judge_api", "判讀：L_total ≥ 所需 L（m＝L1_MAPI）；同欄比較", "x", *_c3(_JUDGE("i_req_api")), "Derived", "文字；基準 0.6426 ≥ 0.5237 → 回本（L_total 口徑）"),
]
BLUE_I = {("i_g", "C"), ("i_g", "D"), ("i_g", "E"), ("i_t1", "C"), ("i_t1", "D"), ("i_t1", "E"), ("i_rho", "C"), ("i_rho", "E")}
KEY_FILL = {"i_lprice", "i_ltotal", "i_ldemote", "i_req_fleet"}
NOTE_V535 = (" ｜v5.35 補充：(7) L 拆分（I 節）：L_total＝L_price × L_demote；基準 L_price 0.741（g −10%／年、T 6 年）、L_demote 0.867（T1 4 年、ρ 0.5）、L_total 0.643，落在 CTL_PriceLife 區間 0.3–1.0 內、"
             "接近 E246 類比區間 0.42–0.70 的上緣；回本所需 L（m＝L1_MFleet 0.81）0.365、（m＝L1_MAPI 0.56）0.52，L_total 基準高於兩者 → 回本（L_total 口徑）；L 低端（g −25%、T1 3、ρ 0.05）L_total 0.342，"
             "低於 m＝L1_MFleet 的所需 L 0.365 → 低端情境不回本。(8) H10 2026 三欄改以 m 分子營收為基底、廣告 Q1 0／Q2 0.5／FY 1.0、消費份額 0.65：2026 Q2 m_API 0.340、m_消費訂閱 5.40、m_商用席次 4.56、"
             "m_整體 1.343（不變）、m_fleet 0.781（不變）。")
ROWS = {}


def _gm_row(ws, sheet, cell):
    rows = [r for r in range(5, ws.max_row + 1) if (ws[f"C{r}"].value, ws[f"D{r}"].value) == (sheet, cell)]
    assert len(rows) == 1, (sheet, cell, rows)
    return rows[0]


def _write(ws, r, key, lab, unit, vals, tag, note, ctx):
    put(ws, f"A{r}", lab.format(**ctx), wrap=True); put(ws, f"B{r}", unit)
    fmt = FMT.get(unit)
    for col, v in vals:
        if isinstance(v, str) and v.startswith("="):
            v = v.format(**ctx)
            put(ws, f"{col}{r}", v, F_CALC, fmt=fmt, fill=FILL_KEY if key in KEY_FILL and col == "C" else None)
        elif isinstance(v, (int, float)) and type(v) is not bool:
            put(ws, f"{col}{r}", v, F_IN if (key, col) in BLUE_I else F_CALC, fmt=fmt)
        else:
            put(ws, f"{col}{r}", v, F_CALC)
    put(ws, f"F{r}", tag, F_NOTE); put(ws, f"G{r}", note.format(**ctx), F_NOTE, wrap=True)


def cap_in_i(wb, h34):
    """Rows start where v5.34 wrote its reading row (h34['end']); that row is cleared and rewritten after I (no defined name there).
    Also writes the H10 L_price link row (v534 ROWS['q_idx']) once the L_price row number is known."""
    ws = wb["Cap_In"]
    r0 = h34["end"]
    assert str(ws[f"A{r0}"].value).startswith("讀法") and ws.max_row == r0, (r0, ws.max_row)
    ws[f"A{r0}"].value = None
    gm = wb["Gov_Map"]; g = _gm_row(gm, "Cap_In", "C63")
    assert gm[f"A{g}"].value == "GM593" and gm[f"E{g}"].value.startswith("壽命期價格係數 L"), (g, gm[f"A{g}"].value)
    ctx = {"gm_lo": f"Gov_Map!$K${g}", "gm_hi": f"Gov_Map!$L${g}"}
    r = r0
    for item in SPEC_I:
        if item[0] is not None: ROWS[item[0]] = r
        r += 1
    ctx.update(ROWS)
    r = r0
    for item in SPEC_I:
        if item[0] is None:
            put(ws, f"A{r}", item[1].format(**ctx), F_BOLD); r += 1; continue
        key, lab, unit, c, d, e, tag, note = item
        _write(ws, r, key, lab, unit, [("C", c), ("D", d), ("E", e)], tag, note, ctx)
        r += 1
    put(ws, f"A{r}", v533.NOTE_END + v534.NOTE_V534 + NOTE_V535, F_NOTE)
    # H10 L_price link row (patched spec wrote the label and the D–F links; C gets the link to I)
    q = v534.ROWS["q_idx"]
    assert ws[f"A{q}"].value == LAB_IDX, ws[f"A{q}"].value
    put(ws, f"C{q}", f"=$C${ROWS['i_lprice']}", F_CALC, fmt="0.000")
    return dict(start=r0, rows=dict(ROWS), end=r)


# ================================================================== B5: Gov_Map — re-point the v5.34 rows whose Cap_In cells moved (guarded), new rows
INSERT_AT = 200          # v5.34 row number of the first H10 row that moved down by one (the 0.65 row is inserted before it)
H10_END_OLD = 223        # v5.34 last H10 row (self-check)
GM629_TAIL = (f" ｜v5.35 B2：本列改為公式（C 欄連結 I 節 L_price，D–F 連結 C），四個藍字 1 退役；類別改「顯示用（公式）」，不進任何計算；{WO} B5")


def _shift(cell, n):
    import re
    return re.sub(r"([A-Z]+)(\d+)", lambda m: f"{m.group(1)}{int(m.group(2)) + n}", cell)


def gov_map_repoint(wb, ws):
    """Rows (Cap_In, cell) with a v5.34 H10 row number ≥ INSERT_AT are moved one row down when the Cap_In label is now found one row below
    (and no longer at the old row). GM629 (index row) also gets its new label and class. Idempotent: on a v5.35-based rebuild nothing matches."""
    import re
    cap = wb["Cap_In"]; n = 0
    for r in range(5, ws.max_row + 1):
        if ws[f"C{r}"].value != "Cap_In" or not isinstance(ws[f"D{r}"].value, str): continue
        cell = ws[f"D{r}"].value
        m = re.match(r"^[A-Z]+(\d+)(?::[A-Z]+(\d+))?$", cell)
        if not m: continue
        old = int(m.group(1))
        if not (INSERT_AT <= old <= H10_END_OLD): continue
        lab = ws[f"E{r}"].value
        if lab == LAB_IDX_OLD and cap[f"A{old + 1}"].value == LAB_IDX and cap[f"A{old}"].value != LAB_IDX_OLD:
            ws[f"D{r}"].value = _shift(cell, 1); ws[f"E{r}"].value = LAB_IDX; ws[f"F{r}"].value = "顯示用（公式）"; ws[f"G{r}"].value = "單值"
            ws[f"M{r}"].value = "無（公式列；v5.35）"
            v518._append_text(ws[f"N{r}"], GM629_TAIL); n += 1
        elif cap[f"A{old + 1}"].value == lab and cap[f"A{old}"].value != lab:
            ws[f"D{r}"].value = _shift(cell, 1); n += 1
    return n


def gov_map_rows(wb):
    if not ROWS: return []
    ws = wb["Cap_In"]; out = []
    def row(sc, r, lab, cell, role, src, rel, dec, lo, hi, rtext, reason, cls="Assumed"):
        assert ws[f"A{r}"].value == lab, (r, ws[f"A{r}"].value, lab)
        out.append({'scope': sc, 'sheet': 'Cap_In', 'cell': cell.format(r=r), 'label': lab, 'check': lab, 'cls': cls, 'role': role, 'src': src,
                    'rel': rel, 'dec': dec, 'lo': (lo.format(r=r) if isinstance(lo, str) else lo), 'hi': (hi.format(r=r) if isinstance(hi, str) else hi),
                    'rtext': rtext.format(r=r), 'reason': reason, 'retag': '', 'seg': '無'})
    sc = "切片三（v5.35 L 拆分）"
    rg = (f"v5.35 J22：中階 token 牌價年變化率 g Assumed −0.10（−0.25～+0.05）；依據 chat 端牌價序列迴歸（{EV['N17']}；序列 {EV['N14']}–{EV['N16']} 等級 3，待一手確認）；"
          "只進 Cap_In I 節與 L1_LPrice／LTotal／LDemote，不進任何 IF_；CTL_PriceLife 不改")
    row(sc, ROWS["i_g"], LAB_I["i_g"], "C{r}", "基準", "", "", "X42", "=Cap_In!D{r}", "=Cap_In!E{r}", "−0.25～+0.05（D 欄＝L 低端 −0.25、E 欄 +0.05）", rg)
    row(sc, ROWS["i_g"], LAB_I["i_g"], "D{r}:E{r}", "低／高", "", "", "X42", None, None, "C{r} 的低（−0.25）、高（+0.05）端點", rg)
    rt = "v5.35 J22：降級前服務中階的年數 T1 Assumed 4（3–5）；對照 B4_Life 商業壽命 12 月為模型更替，T1 為硬體所服務層級的更替；只進 Cap_In I 節與 L1_LTotal／LDemote"
    row(sc, ROWS["i_t1"], LAB_I["i_t1"], "C{r}", "基準", "", "", "X42", "=Cap_In!D{r}", "=Cap_In!E{r}", "3–5（D 欄＝L 低端 3、E 欄 5）", rt)
    row(sc, ROWS["i_t1"], LAB_I["i_t1"], "D{r}:E{r}", "低／高", "", "", "X42", None, None, "C{r} 的低（3）、高（5）端點", rt)
    rr = "v5.35 J22：降級後價值比 ρ Assumed 0.5（低端＝IF_PriceRef_Luna ÷ IF_PriceRef_Sol ≈ 0.05 為公式、高端 1.0＝不降級）；只進 Cap_In I 節與 L1_LTotal／LDemote"
    row(sc, ROWS["i_rho"], LAB_I["i_rho"], "C{r}", "基準", "", "", "X42", "=Cap_In!D{r}", "=Cap_In!E{r}", "0.05–1.0（D 欄為公式 IF_PriceRef_Luna ÷ IF_PriceRef_Sol、E 欄 1.0）", rr)
    row(sc, ROWS["i_rho"], LAB_I["i_rho"], "E{r}", "高", "", "", "X42", None, None, "C{r} 的高端 1.0（不降級）；低端 D 欄為公式，不登錄", rr)
    # B1: ads per column and the 0.65 share (H10)
    sc1 = "切片三（v5.35 m 序列修正）"
    ra = f"v5.35 J23 (2)：廣告年化每欄各自輸入（Assumed；{S(72)} 2026-08 廣告年化 $1B，{EV['N19']}）；廣告不進任何 m 的分子，只在 H10 自 m 分子營收扣除；顯示用"
    qa = v534.ROWS["q_ads"]
    for col, lo, hi, txt in (("D", 0, 0.2, "單值 0（2026 Q1；0–0.2）"), ("E", 0.3, 0.8, "單值 0.5（2026 Q2；0.3–0.8）"), ("F", 0.5, 1.5, "單值 1.0（FY2026；0.5–1.5）")):
        row(sc1, qa, LAB_ADS, col + "{r}", "單值", S(72), f"對照：2026-08 廣告年化 $1B（{S(72)}）；Q2 年化 0.5 為上線初期、FY2026 1.0 為全年", "X43", lo, hi, txt, ra)
    row(sc1, v534.ROWS["q_cts26"], LAB_CTS26, "D{r}", "單值", "", "", "X43", 0.5, 0.8, "單值 0.65（0.5–0.8）；E／F 連結 D；2025 欄連結 H8 0.75",
        "v5.35 J23 (3)：2026 三欄的消費訂閱占 ChatGPT 付費 token 份額改為獨立 Assumed 輸入 0.65（0.5–0.8），2025 欄維持 H8 0.75；席次用量占比缺證據（SRC_DEM_045 只有 p95／中位數）")
    return out


# ================================================================== B4: L1 rows (appended after every existing L1 row; no L1 name moves)
READ_L = ("L 拆分（J22）：L_price＝壽命期內中階牌價相對第 0 年的連續平均（g 為牌價年變化率、T 為 IT 折舊年限）；L_demote＝前 T1 年服務中階、之後降級為低階（價值比 ρ）的折減；"
          "L_total＝L_price × L_demote，與 CTL_PriceLife（基準 1）對照。三者為追蹤列，不乘進任何 IF_；用途是回本條件 L·m ≥ 門檻的 L 側「必須為真」。")
READ_REQ = ("回本所需 L＝L1_FleetBreakeven ÷（IF_Util × CTL_ProdDerate × m）：給定 m，L 至少要多少才達機隊回本。L_total 基準（0.64）高於 m＝L1_MFleet 的所需 L（0.36）→ 回本（L_total 口徑）；"
            "L 低端（g −25%、T1 3、ρ 0.05）0.34 低於 0.36 → 低端情境不回本。m 另見 Cap_In I-7 的 L1_MAPI 口徑（所需 L 0.52）。")


def l1_rows_new():
    R = ROWS
    def cap(col, key): return f"=Cap_In!${col}${R[key]}"
    none = ("無外部對照（牌價序列迴歸見 DB_Evidence E338–E341；等級 3，待一手確認）", DASH, DASH)
    return [
        ("LPrice", "壽命期價格係數 L — 牌價分項 L_price", "中階 token 牌價年變化率 g（Assumed −0.10；−0.25～+0.05）、壽命 T＝IF_DeprLifeIT（6 年）；L_price＝((1+g)^T − 1) ÷ (T × LN(1+g))；v5.35 J22",
         cap("C", "i_lprice"), cap("D", "i_lprice"), cap("E", "i_lprice"), "x", "低＝g −0.25、高＝g +0.05（Cap_In I 節 D／E 欄）", READ_L,
         "牌價年變化率 g、IT 折舊年限", "g：Assumed（chat 端迴歸；歷史牌價等級 3）", *none, "IF_DeprLifeIT", f"Cap_In 第 {R['i_lprice']} 列",
         "牌價為表列價，非實收（折讓在 m）；L_price 不含層級組合變動"),
        ("LDemote", "壽命期價格係數 L — 降級分項 L_demote", "前 T1 年服務中階（Assumed 4；3–5）、之後降級為低階、價值比 ρ（Assumed 0.5；0.05–1.0）；L_demote＝L_total ÷ L_price；v5.35 J22",
         cap("C", "i_ldemote"), cap("D", "i_ldemote"), cap("E", "i_ldemote"), "x", "低＝T1 3、ρ 0.05（直接以低階單價計）；高＝T1 5、ρ 1.0（不降級）", READ_L,
         "T1、ρ、g、IT 折舊年限", "T1、ρ：Assumed（無直接觀測）", "無外部對照", DASH, DASH, "IF_PriceRef_Luna；IF_PriceRef_Sol；IF_DeprLifeIT", f"Cap_In 第 {R['i_ldemote']} 列",
         "硬體降級時點與降級後單價無公開資料；低端以 Luna ÷ Sol 參考請求單價 0.05 為代理"),
        ("LTotal", "壽命期價格係數 L — L_total＝L_price × L_demote（對照 CTL_PriceLife）", "g、T1、ρ、T 同上；L_total＝[((1+g)^T1 − 1) ＋ ρ((1+g)^T − (1+g)^T1)] ÷ (T × LN(1+g))；v5.35 J22",
         cap("C", "i_ltotal"), cap("D", "i_ltotal"), cap("E", "i_ltotal"), "x", "低＝g −0.25、T1 3、ρ 0.05；高＝g +0.05、T1 5、ρ 1.0（同向組合）", READ_L,
         "g、T1、ρ、IT 折舊年限", "三者皆 Assumed", "CTL_PriceLife（模型假設 1；Gov_Map GM593 區間 0.3–1.0）；E246 類比區間 0.42–0.70", "=CTL_PriceLife", "=CTL_PriceLife",
         "CTL_PriceLife；IF_PriceRef_Luna；IF_PriceRef_Sol；IF_DeprLifeIT", f"Cap_In 第 {R['i_ltotal']} 列", "CTL_PriceLife 基準 1 不改（本列只作追蹤）"),
        ("LReqFleet", "回本所需壽命期價格係數 L（m＝L1_MFleet）", "L1_FleetBreakeven ÷（IF_Util × CTL_ProdDerate × L1_MFleet）；v5.35 B2「必須為真」",
         cap("C", "i_req_fleet"), cap("D", "i_req_fleet"), cap("E", "i_req_fleet"), "x", "低＝m 取 L1_MFleet_Hi（所需 L 低）、高＝m 取 L1_MFleet_Lo", READ_REQ,
         "機隊回本門檻、利用率、生產折減、機隊 m", "L1_MFleet：Assumed 分母（Azure 倍數、實際利用率）", "L_total 基準（L1_LTotal）", "=L1_LTotal", "=L1_LTotal",
         "L1_FleetBreakeven；IF_Util；CTL_ProdDerate；L1_MFleet；L1_LTotal", f"Cap_In 第 {R['i_req_fleet']} 列", "m 口徑不同所需 L 不同（L1_MAPI 口徑 0.52 見 Cap_In I-7）"),
    ]


def l1_judge(wb):
    """L1 P column: L1_LTotal = inside／outside the Gov_Map range of CTL_PriceLife (GM593); L1_LReqFleet = L_total ≥ required L."""
    ws = wb["L1"]; gm = wb["Gov_Map"]; g = _gm_row(gm, "Cap_In", "C63"); out = []
    lo, hi = f"Gov_Map!$K${g}", f"Gov_Map!$L${g}"
    r = int(wb.defined_names["L1_LTotal"].attr_text.split("$")[-1]); assert ws[f"A{r}"].value == "L1_LTotal", (r, ws[f"A{r}"].value)
    rng = f'"落在 "&TEXT({lo},"0.0")&"–"&TEXT({hi},"0.0")&" 區間'
    ws[f"P{r}"].value = f'=IF(ISNUMBER(D{r}),IF(AND(D{r}>={lo},D{r}<={hi}),{rng}內",{rng}外"),"推算值非數字（本情境）")'
    out.append(r)
    r = int(wb.defined_names["L1_LReqFleet"].attr_text.split("$")[-1]); assert ws[f"A{r}"].value == "L1_LReqFleet", (r, ws[f"A{r}"].value)
    ws[f"P{r}"].value = (f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(M{r})),IF(M{r}>=D{r},"回本（L_total 口徑）：L_total "&TEXT(M{r},"0.00")&" ≥ 所需 L "&TEXT(D{r},"0.00"),'
                         f'"未回本（L_total 口徑）：L_total "&TEXT(M{r},"0.00")&" < 所需 L "&TEXT(D{r},"0.00")),"推算值非數字（本情境）")')
    out.append(r)
    return out


# ================================================================== B3: DB_Evidence E338–E343 (17 columns A..Q; layout of v531–v534)
def _ev(n, claim, src, tag, param, cur, new, verdict, note, sids, affected, grade, stance):
    return [EV[n], DATE, claim, src, tag, param, cur, new, verdict, V, f"{note}｜暫編 {n}｜{ANDY}｜讀取者：chat 端（{WO} B3）",
            "已處理", sids, DASH, affected, grade, stance]


L1_NEW_NAMES = "L1_LPrice、L1_LDemote、L1_LTotal、L1_LReqFleet（v5.35 新列）"
EVIDENCE_V535 = [
    _ev("N14", "OpenAI 中階牌價序列（輸入／輸出 $/M）：GPT-4 30/60（2023-03）→ 4 Turbo 10/30（2023-11）→ 4o 5/15（2024-05）→ 4o-0806 2.5/10 → 4.1 2/8（2025-04）→ GPT-5 1.25/10（2025-08）→ 5.2 1.75/14（2025-12）"
        "→ 5.4 2.5/15（2026-03）→ 5.6 Terra 2/12（2026-07）→ GPT-6 Sol 2/10（2026-09）→ 6.1 Sol 2/10、快取 0.10（2026-09-29）", "chat 端依既有知識整理（未逐筆重讀原文）", "等級 3（G0-7）；待一手確認",
        "Cap_In I-1 牌價年變化率 g 的依據（迴歸輸入）", DASH, "序列（見主張）", "登錄 DB_Evidence；SRC_Price 不新增（等級 3）；2026 各點可對 SRC_PRC_006–008 與 BenchLM 2026-10-13 頁面互證", "歷史牌價未逐筆重讀原文",
        "SRC_PRC_006；SRC_PRC_007；SRC_PRC_008", L1_NEW_NAMES, "3", "利害關係方"),
    _ev("N15", "Claude Sonnet 牌價 3/15 自 2024-03 至 Sonnet 5（2026-09 標準價）持平；Sonnet 5 換 tokenizer，每單位文字實價未定", "Finout 2026-07-01 轉述", "等級 3；二手", "Cap_In I-1 g 的對照（Sonnet 持平）", DASH, "持平",
        "登錄 DB_Evidence；SRC_Price 不新增；g 高端（+0.05）取近一年持平略升的依據之一", "tokenizer 變更使每單位文字實價未定", DASH, L1_NEW_NAMES, "3", "利害關係方"),
    _ev("N16", "Gemini Pro 牌價：1.5 Pro 3.5/10.5 → 1.25/5（2024-10）→ 2.5 Pro 1.25/10 → 3／3.1 Pro 2/12", "BenchLM 2026-10-13", "等級 3；二手彙整", "Cap_In I-1 g 的對照（Gemini Pro −18%／+11%）", DASH, "序列（見主張）",
        "登錄 DB_Evidence；SRC_Price 不新增", "二手彙整", DASH, L1_NEW_NAMES, "3", "第三方"),
    _ev("N17", "chat 端迴歸結果（對數線性；參考請求權重／3:1 權重並列）：OpenAI 中階 2023–2026.8 年變化 −49%／−38%、2024.3–2026.8 −24%／−10%、2025–2026.8 +8%／+12%；Sonnet 持平；Gemini Pro −18%／+11%",
        f"{WO} B2 I-1 依據欄（chat 端 2026-10-14 計算）", "Derived（chat 端）", "Cap_In I-1 g 基準 −0.10（−0.25～+0.05）", DASH, "−0.10（−0.25～+0.05）",
        "登錄 DB_Evidence；g 基準取近兩年中值附近、高端取近一年持平略升、低端取 2027 後供給釋放與中國廠商競價情境", "兩種混合權重並列", DASH, L1_NEW_NAMES, DASH, "中立"),
    _ev("N18", "L 換算式與錨點（B2）：L_price＝((1+g)^T−1)/(T·LN(1+g))、L_total＝[((1+g)^T1−1)+ρ((1+g)^T−(1+g)^T1)]/(T·LN(1+g))、L_demote＝L_total ÷ L_price；基準 0.7412／0.6426／0.867；"
        "降級 ρ 低端＝GPT-6 Luna ÷ Sol 參考請求單價 0.05", f"{WO} B2（chat 端 2026-10-14 計算）", "Derived（chat 端）", "Cap_In I 節公式列與 L1_LPrice／LDemote／LTotal", DASH, "見主張",
        "登錄 DB_Evidence；Cap_In I 節驗收錨點來源", "—", DASH, L1_NEW_NAMES, DASH, "中立"),
    _ev("N19", "OpenAI 廣告年化 $1B（上線約 200 天後）", "CNBC 2026-08-31（二手）", "Interested-party／2 級（二手）", "Cap_In H10 廣告營收年化 2026 三欄（Q1 0、Q2 0.5、FY2026 1.0）", "0.5（H8 單值，v5.34）", "1.0 $B（年化，2026-08）",
        f"採用（X43）：新登錄 {S(72)}；Active；J23 (2) 廣告每欄各自輸入的依據", "二手轉述，未取得一手", S(72), DASH, "2", "利害關係方"),
]


def evidence_rows():
    return EVIDENCE_V535


# ================================================================== B5: Decisions X42–X44 (10 columns; layout of v531–v534)
_ANDY_ALL = f"{ANDY_1938}；{ANDY_1549}；chat 端建議（「1 GW DC 收入：L 與 m」對話串），可撤回"
_X44 = ("v5.34 報告「待 Project 判斷」第 4–13 項全部照 CC 所做定案（chat 端裁定）：(1) Cap_In H8 D／E 欄依 m 方向擺放，L1_MAPI／MSeat／MixFactorAPI 的 _Lo／_Hi 取 D／E 的 MIN／MAX；"
        "(2) 新增顯示用具名範圍 B4_Disc_Lo／B4_Disc_Hi（Cap_In D12／E12）；(3) 折讓因子分母寫 (1 − B4_Disc)，不寫死常數；(4) 中階牌價指數預留列本單改為連結 I 節 L_price（顯示用，不進 m）；"
        "(5) Gov_Map 以 Assumed 登錄方案上限四格與 ChatGPT 則數三點（理由欄註明 N06／Derived）；(6) SRC_DEM_071 數值欄取高端 0.20、低欄 0.15；(7) L1_MAPI_Anthropic 外部欄 M／N 放兩口徑（用量計費占營收 0.83、分潤比 0.16）；"
        "(8) W1 Warnings +2（SRC_DEM_069／071 利害關係方 Active 無第二來源）接受；(9) τ* 判讀文字以 AL_TokPerPrompt_Lo／_Hi 動態產生；(10) 全簿重算 2.3 秒不列入引擎選型討論（修訂六）")
_X = [
    ("X42", "J22 L 拆分", "新增 Cap_In I 節：L_total＝L_price × L_demote。L_price 由中階 token 牌價年變化率 g（Assumed −0.10；−0.25～+0.05）與壽命 T＝IF_DeprLifeIT（6 年）推得（連續平均）；"
     "L_demote 由「前 T1 年（Assumed 4；3–5）服務中階、之後降級為低階（價值比 ρ Assumed 0.5；低端＝IF_PriceRef_Luna ÷ IF_PriceRef_Sol 0.05、高端 1.0）」推得。三者皆為追蹤列與 L1 名稱（L1_LPrice／LDemote／LTotal），"
     "CTL_PriceLife 基準 1 不改、Gov_Map 區間 0.3–1.0 不改。L 的回本「必須為真」列（L1_LReqFleet）與 m 並列。基準 L_price 0.7412、L_total 0.6426、L_demote 0.867",
     "Cap_In I 節；L1_LPrice／L1_LDemote／L1_LTotal／L1_LReqFleet；Gov_Map 新列；DB_Evidence E338–E342", "第 1 節、B2、B4"),
    ("X43", "J23 v5.34 待判斷第 1–3 項", "(1) H10 營收三分的分項營收改以「m 分子營收（營收 − 廣告）」為基底，使三分之和＝m_整體分子；(2) 廣告年化每欄各自輸入：2026 Q1 0、Q2 0.5、FY2026 1.0（Assumed；SRC_DEM_072 2026-08 廣告年化 $1B 為依據）；"
     "(3) 2026 三欄的「消費訂閱占 ChatGPT 付費 token 份額」改為獨立 Assumed 輸入 0.65（0.5–0.8），2025 欄維持 H8 0.75。錨點 2026 Q2：m_API 0.340、m_消費訂閱 5.399、m_商用席次 4.557、m_整體 1.343（不變）、m_fleet 0.781（不變）",
     "Cap_In H10 廣告列、三分營收列、新列（0.65）、消費訂閱／商用席次 token 列；Gov_Map 新列；SRC_DEM_072；DB_Evidence E343", "第 1 節、B1、B3"),
    ("X44", "J24 v5.34 待判斷第 4–13 項定案", _X44, "Cap_In H8／H10；L1；Gov_Map GM620–GM633；SRC_DEM_071；本頁", "第 1 節"),
]
DECISIONS_V535 = [[i, V, f"{i} {t}", txt, DATE, _ANDY_ALL, "v5.35 已建", where, f"{WO} {sec}", "否"] for i, t, txt, where, sec in _X]


# ================================================================== README (version string is built from VERSION; v5.34 text is kept after it)
README_VERSION = (VERSION + "（判斷類，工作單 docs/workorders/20261014_v5.35.md r0：m 序列三項修正＋L 拆分——Cap_In H10 三分營收改以 m 分子營收為基底、廣告每欄各自輸入（Q1 0／Q2 0.5／FY2026 1.0）、"
                  "2026 消費訂閱占付費 token 份額 0.65（2026 Q2 m_API 0.340、m_消費訂閱 5.399、m_商用席次 4.557、m_整體 1.343）；新增 I 節「L 拆分」L_total＝L_price × L_demote（基準 0.7412 × 0.867＝0.6426；"
                  "g −0.10、T1 4、ρ 0.5、T＝IF_DeprLifeIT 6）與回本所需 L（m＝L1_MFleet 0.3645）；L1 新增 L1_LPrice、L1_LDemote、L1_LTotal、L1_LReqFleet；SRC_DEM_072、DB_Evidence E338–E343、Decisions X42–X44；"
                  "既有 IF_、L1_ 數值全部不變；CTL_PriceLife／CTL_Monetize 不改）。以下為 " + v534.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v534.README_TITLE
README_TITLE = VERSION.split("_")[-1] + _T[len(v534.VERSION.split("_")[-1]):]
README_ROW = ("L 拆分（v5.35）",
              "Cap_In I 節：L_total＝L_price × L_demote。L_price＝((1+g)^T − 1) ÷ (T × LN(1+g))（g＝中階 token 牌價年變化率 Assumed −0.10、−0.25～+0.05；T＝IF_DeprLifeIT 6 年）＝0.7412；"
              "L_total＝[((1+g)^T1 − 1) ＋ ρ((1+g)^T − (1+g)^T1)] ÷ (T × LN(1+g))（T1＝降級前服務中階年數 Assumed 4、3–5；ρ＝降級後價值比 Assumed 0.5、低端 IF_PriceRef_Luna ÷ IF_PriceRef_Sol 0.05、高端 1.0）＝0.6426；"
              "L_demote＝L_total ÷ L_price＝0.867。對照 CTL_PriceLife（基準 1，不改）與 E246 類比區間 0.42–0.70；回本所需 L＝L1_FleetBreakeven ÷（IF_Util × CTL_ProdDerate × m）：m＝L1_MFleet → 0.3645、m＝L1_MAPI → 0.5237；"
              "L_total 基準 ≥ 所需 L → 回本（L_total 口徑）。L1_LPrice／L1_LDemote／L1_LTotal／L1_LReqFleet 各含 _Lo／_Hi。H10 的「L_price」列為顯示用連結，不乘入單價。"
              "H10 修正（J23）：三分營收以 m 分子營收為基底、廣告每欄各自輸入、2026 消費份額 0.65。")
