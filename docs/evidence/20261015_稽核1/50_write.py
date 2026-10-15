"""產生明細 xlsx 與統計。用法：python3 50_write.py <out.json> <稽核0明細.xlsx> <明細.xlsx> <stats.json>"""
import json, sys, collections, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as L
o=json.load(open(sys.argv[1])); ns={}; exec(open('docs/evidence/20261015_稽核1/22_systemic.py',encoding='utf-8').read(),ns); SYS=ns['SYS']
src=openpyxl.load_workbook(sys.argv[2],read_only=True)
base=list(src['C7_第1段輸入清單'].iter_rows(min_row=2,values_only=True)); hdr=list(base[0]); base=base[1:]
idx={(r[1],r[2]):r for r in base}
RULE=('規則（第 1 段，工作單第 2 節）：前 18 欄沿用第 0 段 C7；後 4 欄＋本段 4 欄為判讀。判讀依據＝列標籤、說明欄、來源欄、Gov_Map 與 Decisions、直接使用處（依賴圖），加上一般知識；只標問題、不判定對錯、不提修法。'
      '類別：T 時點過時或不符；G 應隨世代或時間變化卻是單一值；U 單位或口徑不一致；S 無來源或來源薄弱且敏感度高；E 有未處理證據（第 0 段 C2）；R 數值合理性可疑；D 定義不清；—＝沒有問題。'
      '嚴重度：高＝在 m 路徑上或屬 Load_Bearing／Gov_Map 敏感度分段「高」，且有 T、G、U、S 任一項；中＝有 T／G／U／S 但不屬高敏感範圍，或屬高敏感範圍但只有 E／R／D；低＝其餘有問題者。'
      '同一列的世代或層級格合併判讀（「同列合併判讀涵蓋」欄）。「不確定」與「否（沒問題）」分開；「不明」＝工作簿內找不到時點。')
wb=openpyxl.Workbook(); ws=wb.active; ws.title='輸入判讀'
cols=hdr[:18]+['代表什麼物件','時點','應否隨世代或時間變化','口徑是否與使用它的公式一致','問題類別','嚴重度','理由','需外部查證','同列合併判讀涵蓋','Load_Bearing','Gov_Map 敏感度分段']
def sheet(ws,rule,cols,rows,widths=None):
    ws.append([rule]); ws.append(cols)
    ws.cell(1,1).alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[1].height=90
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=min(len(cols),12))
    for i in range(1,len(cols)+1): ws.cell(2,i).font=Font(bold=True); ws.cell(2,i).fill=PatternFill('solid',fgColor='FFDDEBF7')
    for r in rows: ws.append(r)
    ws.freeze_panes='A3'; ws.auto_filter.ref=f'A2:{L(len(cols))}{ws.max_row}'
    for i,w in enumerate(widths or [],1): ws.column_dimensions[L(i)].width=w
rows=[]
for x in o:
    b=list(idx[(x['sheet'],x['cell'])])[:18]
    rows.append(b+[x['o'],x['t'],x['v'],x['k'],x['c'],x['sev'],x['r'],x['x'],x['covers'],'是' if x['lb'] else '否',x['tier'] or '—'])
sheet(ws,RULE,cols,rows,[8,12,9,28,8,10,10,12,6,6,12,8,10,6,6,8,6,10,30,22,26,30,8,6,40,6,14,6,8])
hi=[r for r in rows if r[23]=='高']
sheet(wb.create_sheet('高嚴重度'),'規則：上一頁嚴重度＝「高」的全部格（在 m 路徑上，或 Load_Bearing／敏感度分段高，且類別含 T、G、U、S）。'+RULE[:0],cols,hi,[8,12,9,28])
# 系統性
cellmap=collections.defaultdict(list)
for x in o: cellmap[(x['sheet'],x['row'])].append(x)
srows=[]
for sid,t,desc,rs in SYS:
    cs=[c for k in rs for c in cellmap.get(k,[])]
    srows.append([sid,t,desc,len(cs),sum(c['m']=='是' for c in cs),sum(c['sev']=='高' for c in cs),'；'.join(f"{s}!{x}" for s,x in {(c['sheet'],c['covers']) for c in cs})])
sheet(wb.create_sheet('系統性問題'),'規則：同一種根因影響多格的情形；格數依「頁＋列」展開到輸入格（同列各世代／層級格都計入）。只描述根因與範圍，不提修法。',
      ['ID','系統性問題','說明','影響格數','其中 m 路徑','其中嚴重度高','相關格'],srows,[5,40,80,8,8,8,80])
# 搜尋紀錄
srch=[[1,'Rubin Ultra NVL576 MGX rack 72 GPUs per rack HBM4e per GPU bandwidth 2026','Spec_Rack G26（HBM 頻寬）','只有二手規格整理，無法確認 72 封裝單架口徑；維持需外部查證'],
      [2,'GB300 NVL72 shipments volume 2025 fourth quarter OpenAI Microsoft first GB300 cluster','Alloc_In C11（2025 GB300 組合 0%）','Microsoft Azure 2025-10 公告為 OpenAI 部署首個大規模 GB300 NVL72 叢集；0% 是否成立需確認（標 R）'],
      [3,'HGX H100 8-GPU server price 2023 $ per server','Spec_Rack C11:C13（Hopper 機架價格）','未找到 2023 部署當時一手價格；維持需外部查證']]
sheet(wb.create_sheet('網路搜尋紀錄'),'規則：工作單第 2 節，只對嚴重度高且取決於可查事實者搜尋，上限 20 次；每次記錄。',['#','查詢','對應格','結果'],srch,[4,60,30,70])
wb.save(sys.argv[3])
# 統計
st=collections.Counter(); grp=collections.Counter()
for x in o:
    zone='m 路徑內' if x['m']=='是' else 'm 路徑外'
    for c in (x['c'] if x['c']!='—' else '—').split(','):
        st[(zone,c,x['sev'])]+=1
    grp[(zone,x['sev'])]+=1
json.dump(dict(cat=[[*k,v] for k,v in st.items()],sev=[[*k,v] for k,v in grp.items()],sys=srows,
               hi=[[r[1],r[2],r[3],r[22],r[24],r[16]] for r in hi]),open(sys.argv[4],'w'),ensure_ascii=False)
print('rows',len(rows),'hi',len(hi)); print(sorted(grp.items()))
