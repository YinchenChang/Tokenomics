# v15.3 (work order docs/workorders/20261018_v15.3_L證據與修補.md r0, branch claude/v15-3): L evidence B stage + v15.2 repairs.
# MINOR version: input values and formulas change, names and definitions do not.
#   W1  E1  Cap_In C226／D226／E226 (g) -0.10／-0.25／+0.05 -> -0.09／-0.20／+0.05; mark Assumed -> Derived (revision 9: original wording kept in Cap_In G, Gov_Map AL).
#       E2  Cap_In D228 (rho low end) formula x INDEX(IF_TokGW_Luna,1,4)/INDEX(IF_TokGW_Sol,1,4) (GB200 column); E228 1 -> 0.7; mark Assumed -> Derived (two readings).
#       E3  Cap_In C227 (T1) unchanged (Gov_Map GM636 N note only).
#   W2  D1  Cap_In row 242 = three-state reading (L_total interval vs required-L interval, identity reading L1_LReqFleet_Id); row 243 relabelled as a literal-reading comparison.
#   W3  C2  Calib D156 (CAL_BatchEtaD GB200) 2.36 -> =SQRT(C254*D201) (= 2.198); Gov_Map GM685 K 1.83 -> 1.59.
#   W4  Gov_Map AA COUNTIF span $A$5:$A$300 -> $A$5:$A$1000 (gov.py); README／CHANGELOG／contract.
#   W5  Decisions SL48-01..08.   W6  DB_Evidence E393..E397.
# Every write to an Excel-owned cell is guarded (old value / old text); appended texts keep the original wording (CLAUDE.md section 6, revision 9).
# The cell-level edits (inputs_update) run BEFORE gov_all so that the contract tracing (Deps) and Load_Bearing see the final formulas;
# Gov_Map edits run in gov.py's update chain; labels on L1, Decisions and DB_Evidence run last (apply).
from openpyxl.utils import get_column_letter as L

import v1020
import v1410
import v1520

VERSION = "20261018_Tokenomics_v15.3"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-18"
V = "v15.3"
WO = "工作單 docs/workorders/20261018_v15.3_L證據與修補.md（r0）"
A_REP = "docs/reports/20261018_L證據A段.md"
EV_DIR = "docs/evidence/20261018_L證據A段"
DEC_BY = ("Andy 2026-10-17 23:55「原文的问题按建议，然后继续做下去。」（C1、C2、D1）；2026-10-18 00:20「请继续」（E1–E4 依預設）；"
          "皆為 chat 端建議，可撤回")

# ---------------------------------------------------------------- W1 (E1, E2, E3)
G_OLD = {"C": -0.1, "D": -0.25, "E": 0.05}
G_NEW = {"C": -0.09, "D": -0.2, "E": 0.05}
RHO_D_OLD = "=IF_PriceRef_Luna/IF_PriceRef_Sol"
RHO_D_NEW = "=IF_PriceRef_Luna/IF_PriceRef_Sol*INDEX(IF_TokGW_Luna,1,4)/INDEX(IF_TokGW_Sol,1,4)"      # column 4 of the 15-cell IF_TokGW_* row = GB200 (low-cost scenario column; the three columns of a generation hold the same output)
RHO_E_OLD, RHO_E_NEW = 1, 0.7
F226_OLD, F226_NEW = "Assumed", "Derived"
F228_OLD, F228_NEW = "Assumed（C、E）／Derived（D）", "Derived（兩種讀法）"

E1_TXT = ("v15.3 E1：2024-05～2026-10 中階牌價（參考請求混合價）月階梯法、三家合併各自截距回歸 −9.1%／年；分廠商 OpenAI −4.4%、Anthropic −19.4%、Google −2.6%；"
          "低端取 Anthropic；高端保留 2024-10 後上漲與已公告漲價；去除快取折扣引入約 0%。窗口敏感。來源 " + A_REP + " 第 1 節")
E2_TXT = ("v15.3 E2：ρ＝同一張 GPU 的價值比＝低階 token 價格比 × 低階模型每 GPU 吞吐倍數。讀法甲（降級服務自家 Luna）＝0.05 × 3.41＝0.17（低端，公式）；"
          "讀法乙（市場機會價值，A100／H100 租金比 0.55～0.67，2026 供給吃緊）；高端 0.7；基準 0.5 偏向讀法乙。"
          "原低端公式（未乘吞吐倍數）與原高端 1.0 保留於本說明")
E3_TXT = ("v15.3 E3：維持。前沿訓練約 2.5～3 年、服務收入年限較長（A100 合約到 2029，利害關係方）；證據未推翻 4（3～5）")

CAP_G226_ADD = ("｜v15.3 E1（修訂九更正，Decisions SL48-04）：基準 −0.10→−0.09、低端 −0.25→−0.20、高端 +0.05 不變；標記由「" + F226_OLD + "」更正為「" + F226_NEW + "」（原字樣保留於此與 Gov_Map AL 欄）。" + E1_TXT)
CAP_G228_ADD = ("｜v15.3 E2（修訂九更正，Decisions SL48-05）：標記由「" + F228_OLD + "」更正為「" + F228_NEW + "」（原字樣保留於此與 Gov_Map AL 欄）；"
                "低端 D 欄公式由『" + RHO_D_OLD + "』改為『" + RHO_D_NEW.lstrip("=") + "』（GB200 欄；≈0.05 × 3.41＝0.17）；高端 E 欄 1.0→0.7；基準 0.5 不變。" + E2_TXT)

# ---------------------------------------------------------------- W2 (D1)
A242_OLD = "判讀：L_total ≥ 所需 L（m＝L1_MFleet）；同欄比較"
A242_NEW = "判讀：L_total 區間 對 所需 L 區間（恆等式讀法 L1_LReqFleet_Id；三態：未回本／未定／回本）"
A243_OLD = "判讀：L_total ≥ 所需 L（m＝L1_MAPI）；同欄比較"
A243_NEW = "對照（字面讀法，口徑不一致，不作判讀；見 L1 第 89 列說明）：L_total ≥ 所需 L（m＝L1_MAPI）；同欄比較"
_NUM = "AND(ISNUMBER(L1_LTotal),ISNUMBER(L1_LTotal_Lo),ISNUMBER(L1_LTotal_Hi),ISNUMBER(L1_LReqFleet_Id),ISNUMBER(L1_LReqFleet_Id_Lo),ISNUMBER(L1_LReqFleet_Id_Hi))"
C242_NEW = ("=IF(" + _NUM + ",IF(L1_LTotal_Hi<L1_LReqFleet_Id_Lo,\"未回本（L 區間全在所需 L 之下）\","
            "IF(L1_LTotal_Lo>L1_LReqFleet_Id_Hi,\"回本（L 區間全在所需 L 之上）\","
            "\"未定（區間重疊；基準 L \"&TEXT(L1_LTotal,\"0.000\")&\" 對所需 L \"&TEXT(L1_LReqFleet_Id,\"0.000\")&\"）\")),\"所需 L 或 L 非數字（本情境）\")")
D242_NEW = ("=IF(" + _NUM + ",\"對照：低端欄 L \"&TEXT(L1_LTotal_Lo,\"0.000\")&\" 對所需 L（GW 低端）\"&TEXT(L1_LReqFleet_Id_Lo,\"0.000\")&\"（\"&IF(L1_LTotal_Lo>=L1_LReqFleet_Id_Lo,\"≥\",\"<\")&\"）\",\"所需 L 或 L 非數字（本情境）\")")
E242_NEW = ("=IF(" + _NUM + ",\"對照：高端欄 L \"&TEXT(L1_LTotal_Hi,\"0.000\")&\" 對所需 L（GW 高端）\"&TEXT(L1_LReqFleet_Id_Hi,\"0.000\")&\"（\"&IF(L1_LTotal_Hi>=L1_LReqFleet_Id_Hi,\"≥\",\"<\")&\"）\",\"所需 L 或 L 非數字（本情境）\")")
C242_OLD_HEAD = "=IF(ISNUMBER(C240),IF(C233>=C240,"
G242_ADD = ("｜v15.3 D1：本列改為三態判讀。C 欄＝L 區間 [L1_LTotal_Lo, L1_LTotal_Hi] 對所需 L 區間 [L1_LReqFleet_Id_Lo, L1_LReqFleet_Id_Hi]（恆等式讀法 L1_LReqFleet_Id）："
            "L 上端 < 所需 L 下端＝未回本；L 下端 > 所需 L 上端＝回本；其餘＝未定（區間重疊）。D／E 欄為同欄端點對照（前綴『對照：』）。"
            "原判讀（字面讀法 m＝L1_MFleet、L_total ≥ L1_LReqFleet 的二態）與其原說明保留於上；字面讀法值仍在第 240 列與 L1 第 89 列")
G243_ADD = "｜v15.3 D1：本列標籤改為『對照（字面讀法，口徑不一致，不作判讀）』；公式不變。判讀只用第 242 列（恆等式讀法、三態）"

# ---------------------------------------------------------------- W3 (C2)
D156_OLD = 2.36
D156_NEW = "=IF(AND(ISNUMBER(C254),ISNUMBER(D201)),SQRT(C254*D201),1)"      # C254 = V20 back-solve (1.5911), D201 = vLLM 10.1K back-solve (3.0373); fallback 1 (uncalibrated, as the Hopper column) when either point has no solution
GM685_K = (1.83, 1.59)
GM685_M_ADD = "｜v15.3 C2：D67 改 1.15 後 V20 反推點（Calib C254）由 1.83 變 1.5911；基準改為兩個反推點的幾何平均 √(1.5911 × 3.0373)＝2.198（Calib D156 公式）；原 2.36（＝√(1.83×3.04)）與原區間 1.83–3.04 保留於本說明"
GM685_N_ADD = ("｜v15.3 C2：D156 由常數 2.36 改為公式 =SQRT(C254*D201)（兩個反推點所在格：V20＝Calib C254、vLLM 10.1K＝Calib D201；v15.2 實測 1.5911、3.0373，幾何平均 2.198）。"
               "區間低端 K 1.83→1.59，高端 L 3.04 不變。原 2.36（＝√(1.83×3.04)）與原區間 1.83–3.04 保留於本說明。"
               "注意：D201 為『相對 CAL_EtaD』的解（含 D67），D67＝1.15 時以 CAL_BatchEtaD 自身口徑應為 3.0373÷1.15＝2.641（待 Project 判斷，見 v15.3 報告）。無解（C254 或 D201 非數字）時退回 1（未校準）")
CAL_J156_ADD = ("｜v15.3 C2（Decisions SL48-02）：D156 由常數 2.36 改為公式 =SQRT(C254*D201)＝2.198（兩個反推點：V20＝C254 1.5911，D67 改 1.15 後由 1.83 變；vLLM 10.1K＝D201 3.0373）。"
                "原 2.36（＝√(1.83×3.04)）與原區間 1.83–3.04 保留於上。無解時退回 1（未校準）")
CAL_H156_ADD = "｜v15.3 C2：GB200（D156）改為兩個反推點的即時幾何平均 √(C254×D201)＝2.198（區間 1.59–3.04）"

# ---------------------------------------------------------------- Gov_Map (E1, E2, E3, C2)
GM_ROWS = {
    ("Cap_In", "C226"): dict(f_old="Assumed", f_new="Derived",
                              m_old="−0.25～+0.05（D 欄＝L 低端 −0.25、E 欄 +0.05）",
                              m_new="−0.20～+0.05（D 欄＝L 低端 −0.20、E 欄 +0.05；v15.3 E1，原 −0.25～+0.05 保留於 AL 欄）", n_add=E1_TXT, dec="SL48-04"),
    ("Cap_In", "D226:E226"): dict(f_old="Assumed", f_new="Derived",
                                   m_old="C226 的低（−0.25）、高（+0.05）端點",
                                   m_new="C226 的低（−0.20）、高（+0.05）端點（v15.3 E1，原低端 −0.25 保留於 AL 欄）", n_add=E1_TXT, dec="SL48-04"),
    ("Cap_In", "C228"): dict(f_old="Assumed", f_new="Derived（兩種讀法）",
                              m_old="0.05–1.0（D 欄為公式 IF_PriceRef_Luna ÷ IF_PriceRef_Sol、E 欄 1.0）",
                              m_new="0.17–0.7（D 欄為公式 IF_PriceRef_Luna ÷ IF_PriceRef_Sol × 低階模型每 GPU 吞吐倍數（GB200 欄）、E 欄 0.7；v15.3 E2，原 0.05–1.0 保留於 AL 欄）", n_add=E2_TXT, dec="SL48-05"),
    ("Cap_In", "E228"): dict(f_old="Assumed", f_new="Derived（兩種讀法）",
                              m_old="C228 的高端 1.0（不降級）；低端 D 欄為公式，不登錄",
                              m_new="C228 的高端 0.7（讀法乙：市場機會價值；v15.3 E2，原 1.0＝不降級保留於 AL 欄）；低端 D 欄為公式，不登錄", n_add=E2_TXT, dec="SL48-05"),
}


def _al_text(dec, f_old, m_old, k_old=None):
    return ("v15.3 更正（修訂九，Decisions " + dec + "）前的原值：類別＝" + f_old + "｜區間文字（M 欄）＝" + m_old + "｜理由（N 欄）原文保留於 N 欄（v15.3 說明附加於其後）")


def gov_update(ws, append_text):
    n = 0
    for (sh, cell), d in GM_ROWS.items():
        r = v1020._gm_row(ws, sh, cell)
        if ws[f"F{r}"].value == d["f_old"]:
            assert ws[f"M{r}"].value == d["m_old"], (cell, ws[f"M{r}"].value)
            assert ws[f"AL{r}"].value in (None, ""), ws[f"AL{r}"].value
            ws[f"AL{r}"].value = _al_text(d["dec"], d["f_old"], d["m_old"])
            ws[f"F{r}"].value = d["f_new"]
            ws[f"M{r}"].value = d["m_new"]
            n += 3
        else:
            assert ws[f"F{r}"].value == d["f_new"], (cell, ws[f"F{r}"].value)
        n += int(append_text(ws[f"N{r}"], "｜" + d["n_add"]))
    r = v1020._gm_row(ws, "Cap_In", "C227")      # E3: no change, note only
    n += int(append_text(ws[f"N{r}"], "｜" + E3_TXT))
    r = v1020._gm_row(ws, "Calib", "D156")       # C2
    if ws[f"K{r}"].value == GM685_K[0]:
        assert ws[f"L{r}"].value == 3.04, ws[f"L{r}"].value
        ws[f"K{r}"].value = GM685_K[1]; n += 1
    else:
        assert ws[f"K{r}"].value == GM685_K[1], ws[f"K{r}"].value
    n += int(append_text(ws[f"M{r}"], GM685_M_ADD))
    n += int(append_text(ws[f"N{r}"], GM685_N_ADD))
    return n


# ---------------------------------------------------------------- cell edits (before gov_all)
def _row_by_label(ws, prefix, lo=200, hi=300):
    rows = [r for r in range(lo, hi) if str(ws.cell(r, 1).value or "").startswith(prefix)]
    assert len(rows) == 1, (prefix, rows)
    return rows[0]


def inputs_update(wb, append_text):
    """W1 inputs/formulas/marks, W2 rows 242-243, W3 Calib D156 (old-value / old-text guards; idempotent on a v15.3 base)."""
    log = []
    from common import put, F_CALC
    ws = wb["Cap_In"]
    rg = _row_by_label(ws, "中階 token 牌價年變化率 g（I-1")
    r3 = _row_by_label(ws, "降級前服務中階的年數 T1")
    rr = _row_by_label(ws, "降級後價值比 ρ（I-3")
    assert (rg, r3, rr) == (226, 227, 228), (rg, r3, rr)
    for c in "CDE":                                                  # E1
        cur = ws[f"{c}{rg}"].value
        if cur == G_NEW[c]:
            continue
        assert cur == G_OLD[c], (c, cur)
        ws[f"{c}{rg}"].value = G_NEW[c]; log.append(f"Cap_In!{c}{rg}: {G_OLD[c]} -> {G_NEW[c]} (v15.3 E1)")
    assert ws[f"C{rr}"].value == 0.5 and [ws[f"{c}{r3}"].value for c in "CDE"] == [4, 3, 5]      # E2 base and E3 unchanged
    cur = ws[f"E{rr}"].value                                          # E2
    if cur != RHO_E_NEW:
        assert cur == RHO_E_OLD, cur
        ws[f"E{rr}"].value = RHO_E_NEW; log.append(f"Cap_In!E{rr}: {RHO_E_OLD} -> {RHO_E_NEW} (v15.3 E2)")
    cur = ws[f"D{rr}"].value
    if cur != RHO_D_NEW:
        assert cur == RHO_D_OLD, cur
        ws[f"D{rr}"].value = RHO_D_NEW; log.append(f"Cap_In!D{rr}: formula -> {RHO_D_NEW} (v15.3 E2)")
    for r, old, new in ((rg, F226_OLD, F226_NEW), (rr, F228_OLD, F228_NEW)):      # marks (revision 9: original kept in G and Gov_Map AL)
        cur = ws[f"F{r}"].value
        if cur != new:
            assert cur == old, (r, cur)
            ws[f"F{r}"].value = new; log.append(f"Cap_In!F{r}: {old} -> {new} (v15.3, revision 9)")
    log.append(f"Cap_In G226/G228 appended: {int(append_text(ws[f'G{rg}'], CAP_G226_ADD)) + int(append_text(ws[f'G{rr}'], CAP_G228_ADD))}")
    # ---- W2: rows 242 / 243
    assert ws["A242"].value in (A242_OLD, A242_NEW) and ws["A243"].value in (A243_OLD, A243_NEW), (ws["A242"].value, ws["A243"].value)
    if ws["A242"].value == A242_OLD:
        assert ws["C242"].value.startswith(C242_OLD_HEAD), ws["C242"].value
    ws["A242"].value = A242_NEW; ws["A243"].value = A243_NEW
    for c, f in (("C", C242_NEW), ("D", D242_NEW), ("E", E242_NEW)):
        put(ws, f"{c}242", f, F_CALC)
    log.append(f"Cap_In G242/G243 appended: {int(append_text(ws['G242'], G242_ADD)) + int(append_text(ws['G243'], G243_ADD))}")
    # ---- W3: Calib D156 (formula; black font = calculated, no longer an input)
    cal = wb["Calib"]
    assert str(cal["A156"].value).startswith("批次口徑 η_d 倍數（X10；CAL_BatchEtaD）"), cal["A156"].value
    assert str(cal["A254"].value).startswith("達標所需批次口徑 η_d 倍數（CAL_BatchEtaD 的解）") and str(cal["A201"].value).startswith("達標所需 η_d 倍數（相對 CAL_EtaD"), (cal["A254"].value, cal["A201"].value)
    cur = cal["D156"].value
    if cur != D156_NEW:
        assert cur == D156_OLD, cur
        put(cal, "D156", D156_NEW, F_CALC, fmt="0.000"); log.append("Calib!D156: 2.36 -> " + D156_NEW + " (v15.3 C2)")
    log.append(f"Calib J156/H156 appended: {int(append_text(cal['J156'], CAL_J156_ADD)) + int(append_text(cal['H156'], CAL_H156_ADD))}")
    return log


# ================================================================== late text edits (L1 labels), Decisions, DB_Evidence
L1_EDITS = [      # (cell, old, new) : engineering-class label text; every old string is quoted in the report
    ("H86", "低＝g −0.25、高＝g +0.05（Cap_In I 節 D／E 欄）", "低＝g −0.20、高＝g +0.05（Cap_In I 節 D／E 欄；v15.3 E1，原低端 −0.25）"),
    ("H87", "低＝T1 3、ρ 0.05（直接以低階單價計）；高＝T1 5、ρ 1.0（不降級）",
     "低＝T1 3、ρ 0.17（低階 token 單價比 0.05 × 低階模型每 GPU 吞吐倍數，GB200 欄；v15.3 E2）；高＝T1 5、ρ 0.7（v15.3 E2；原低端 ρ 0.05、高端 1.0）"),
    ("H88", "低＝g −0.25、T1 3、ρ 0.05；高＝g +0.05、T1 5、ρ 1.0（同向組合）",
     "低＝g −0.20、T1 3、ρ 0.17；高＝g +0.05、T1 5、ρ 0.7（同向組合；v15.3 E1／E2，原低端 g −0.25、ρ 0.05，高端 ρ 1.0）"),
]


def l1_text(wb):
    ws = wb["L1"]
    n = 0
    for ref, old, new in L1_EDITS:
        cur = ws[ref].value
        if cur == new:
            continue
        assert cur == old, (ref, cur)
        ws[ref].value = new; n += 1
    return n


def _decisions():
    D = [
        ("接受 D67 對 Training 與 IF_Train*（GB200 欄）的連帶變動（C1；v15.2 補登）",
         "v15.2 B3（Calib D67 1 → 1.15）同時進入 Training 頁 GB200 欄，使 Training 頁 78 格與 IF_Train* 中 GB200 的 3 格變動約 1.7–9%（工作單 v15.2 曾列為不得變）。"
         "決議：接受此連帶變動，不把 Training 的批次倍數與 D67 脫鉤。本筆為 v15.2 發生之事的補登，v15.3 不再變動這些格（W3 的 D156 另計於 SL48-02）。",
         "Andy 2026-10-17 23:55「原文的问题按建议，然后继续做下去。」（C1）；chat 端建議，可撤回", "Training 頁 GB200 欄；IF_Train*（GB200）；Calib D67"),
        ("批次口徑倍數 CAL_BatchEtaD GB200 改為兩個反推點的幾何平均公式（C2）",
         "Calib D156：2.36（常數，＝√(1.83×3.04)）→ 公式 =SQRT(C254*D201)＝2.198（兩個反推點：V20＝Calib C254 在 D67 改 1.15 後由 1.83 變 1.5911；vLLM 10.1K＝Calib D201 3.0373）。"
         "Gov_Map GM685 區間 1.83–3.04 → 1.59–3.04（K 1.83 → 1.59，L 3.04 不變）；原 2.36 與原區間保留於 Gov_Map M／N 欄與 Calib J156。"
         "區間端點 RNG_* 依 v13.0 規則以 tools/gen_ranges.py 重產（高端情境 GB200 1.83 → 1.59）；Alloc N 節靜態情境表同步重產。"
         "已知疑點（交 Project）：D201 為『相對 CAL_EtaD』的解（含 D67），D67＝1.15 時以 CAL_BatchEtaD 自身口徑為 3.0373 ÷ 1.15＝2.641，幾何平均應為 2.05；本版依工作單照 2.198 建置。",
         "Andy 2026-10-17 23:55「原文的问题按建议，然后继续做下去。」（C2）；chat 端建議，可撤回", "Calib D156；Gov_Map GM685；RNG_*；Alloc N／O／P 節；IF_Alloc*、IF_Train*、L1_*"),
        ("回本判讀改用恆等式讀法、三態顯示（D1）",
         "Cap_In 第 242 列：以 L1_LTotal（基準／低／高）對 L1_LReqFleet_Id（基準／低／高）比較。C 欄三態：L1_LTotal_Hi < L1_LReqFleet_Id_Lo＝未回本（L 區間全在所需 L 之下）；"
         "L1_LTotal_Lo > L1_LReqFleet_Id_Hi＝回本（L 區間全在所需 L 之上）；其餘＝未定（區間重疊，附基準 L 與所需 L）。D／E 欄為同欄端點對照（前綴『對照：』）。"
         "第 243 列（m＝L1_MAPI 的字面讀法）標籤改為『對照（字面讀法，口徑不一致，不作判讀）』，公式保留。L1 P89 已用恆等式讀法的二態點比較，未改（待 Project 判斷）。",
         "Andy 2026-10-17 23:55「原文的问题按建议，然后继续做下去。」（D1）；chat 端建議，可撤回", "Cap_In 第 242／243 列（A、C、D、E、G）"),
        ("g 基準 −0.09、區間 −0.20～+0.05，標記 Assumed → Derived（E1；修訂九更正）",
         "Cap_In C226／D226／E226：−0.10／−0.25／+0.05 → −0.09／−0.20／+0.05。依據：" + A_REP + " 第 1 節——2024-05～2026-10 中階牌價（參考請求混合價）月階梯法、三家合併各自截距回歸 −9.1%／年；"
         "分廠商 OpenAI −4.4%、Anthropic −19.4%、Google −2.6%；低端取 Anthropic；高端保留 2024-10 後上漲與已公告漲價；去除快取折扣引入約 0%；窗口敏感。"
         "更正理由（修訂九）：g 現由牌價序列迴歸推得（Derived），不再是無依據的假設；原字樣『Assumed』逐字保留在 Cap_In G226 與 Gov_Map GM634／GM635 的 AL 欄，原區間文字保留於 AL 欄、原理由保留於 N 欄。"
         "GM635（D226:E226 低／高列）與 GM634 同為 g 的登錄列，一併更正。",
         "Andy 2026-10-18 00:20「请继续」（E1 依預設）；chat 端建議，可撤回", "Cap_In C226:E226（F、G）；Gov_Map GM634、GM635（F、M、N、AL）；L1 H86／H88；L1_LTotal*、L1_LPrice*"),
        ("ρ 低端公式乘吞吐倍數、高端 0.7、基準 0.5 不變，標記 Assumed → Derived（兩種讀法）（E2；修訂九更正）",
         "Cap_In D228：=IF_PriceRef_Luna/IF_PriceRef_Sol（0.050）→ =IF_PriceRef_Luna/IF_PriceRef_Sol*INDEX(IF_TokGW_Luna,1,4)/INDEX(IF_TokGW_Sol,1,4)（GB200 欄，約 0.170）；E228：1.0 → 0.7；C228 基準 0.5 不變。"
         "依據：" + A_REP + " 第 0 節 E2——ρ＝同一張 GPU 的價值比＝低階 token 價格比 × 低階模型每 GPU 吞吐倍數；讀法甲（降級服務自家 Luna）0.05 × 3.41＝0.17；讀法乙（市場機會價值）A100／H100 租金比 0.55～0.67；"
         "基準 0.5 偏向讀法乙，是 Andy 的判斷點（舊 GPU 是否可轉租或轉作他用）。INDEX 第 4 格已確認為 GB200 欄（Interface IF_TokGW_* 為 5 世代 × 3 成本情境，第 4–6 格為 GB200）。"
         "更正理由（修訂九）：低端與高端現由證據推得（Derived，兩種讀法），原字樣『Assumed』（含『Assumed（C、E）／Derived（D）』）逐字保留在 Cap_In G228 與 Gov_Map GM638／GM639 的 AL 欄。",
         "Andy 2026-10-18 00:20「请继续」（E2 依預設）；chat 端建議，可撤回", "Cap_In C228:E228（F、G）；Gov_Map GM638、GM639；L1 H87／H88；L1_LDemote*、L1_LTotal*"),
        ("T1 維持 4（3–5）（E3）",
         "Cap_In C227 不變；Gov_Map GM636 N 欄附加說明：前沿訓練約 2.5～3 年、服務收入年限較長（A100 合約到 2029，利害關係方）；證據未推翻 4（3～5）。",
         "Andy 2026-10-18 00:20「请继续」（E3 依預設）；chat 端建議，可撤回", "Cap_In C227（不變）；Gov_Map GM636 N"),
        ("長約不計入 Tokenomics 的 L；轉知 company-models（E4）",
         "Tokenomics（實驗室視角）的 L 不計長約：長約鎖住承租方（實驗室）的成本與出租方（neocloud）的收入，實驗室的 token 單價照樣下降；長約使出租方合約期 L≈1"
         "（CoreWeave：加權平均約 5 年 take-or-pay、H100 續約約原價 95%，管理層說法）。轉知 company-models 的 neocloud 模型；本簿數值不改。",
         "Andy 2026-10-18 00:20「请继续」（E4 依預設）；chat 端建議，可撤回", "無（只記錄）；DB_Evidence E396"),
        ("Alloc N 節靜態情境表於 v15.2 重產 10 格（工程類；chat 端判定）",
         "v15.2 的 B3 使研發 GW 與 Q1 改變，Alloc N 節（v12.0 靜態情境表，10 格，與 RNG_IF_AllocRDGW／Q1_* 同源）依 v13.0 測試要求同步重產（v15.2 報告待判斷第 4 項）。chat 端判定：歸工程類，接受。"
         "v15.3 因 C2（D156、區間高端 1.59）再次重產，並放寬 Decisions 檢查範圍（Gov_Map AA 欄 COUNTIF $A$5:$A$300 → $A$5:$A$1000）。",
         "chat 端判定（工程類；見工作單 v15.3 W5-4）；可撤回", "Alloc N 節；Gov_Map AA 欄"),
    ]
    return [[f"SL48-{i + 1:02d}", V, f"v15.3 L 證據 B 段＋修補：{t}", txt, DATE, who, "v15.3 已建", where, f"{WO} W5-{i + 1}", "否"]
            for i, (t, txt, who, where) in enumerate(D)]


EV_COMMON = "讀取者：chat 端（" + A_REP + "）｜Andy 2026-10-17 23:55「原文的问题按建议，然后继续做下去。」、2026-10-18 00:20「请继续」（E1–E4 依預設）（文字，未改字）；chat 端建議，可撤回"
EVIDENCE = [
    ("E393", "中階牌價序列與 g：2024-05～2026-10 月階梯法三家合併各自截距回歸 −9.1%／年（分廠商 OpenAI −4.4%、Anthropic −19.4%、Google −2.6%）；2026-10 三家中階同價 $2／$0.10／$10（混合價 $1.96）",
     EV_DIR + "/L1_prices.md", "Verified／Derived（官方牌價；24 處歷史日期為 [Memory] 未經本次擷取驗證，2026 年新模型經摘要讀取）",
     "Cap_In C226:E226（g；Gov_Map GM634／GM635）", "−0.10（−0.25～+0.05），Assumed", "−0.09（−0.20～+0.05），Derived",
     "採用（E1）", "窗口敏感：2024-10 起 +16.8%、去除快取折扣約 0%、全期事件法 −30.8%；只作回歸，chat 端未重跑", "L1_LPrice*；L1_LDemote*；L1_LTotal*；Cap_In 第 238、242 列", "—"),
    ("E394", "H100 市場租金 2023→2026 年化 −9%～−14%（走勢循環）；AWS P5 −44%（2025-06）、P4d −33%",
     EV_DIR + "/L2_hardware.md", "Interested-party（SemiAnalysis、Silicon Data、GetDeploying）／Verified（AWS 官方牌價）",
     "Cap_In C226:E226（g 的交叉檢查；出租方口徑）", "—", "年化 −9%～−14%（租金）；AWS −17.6%／年（H100）",
     "部分採用（交叉檢查，出租方口徑）", "租金是出租方收入，與實驗室 token 單價是兩種口徑；只作交叉檢查，方向與 −9.1% 一致", "不改任何 IF_／L1_", "利害關係方"),
    ("E395", "A100／H100 租金比 0.55～0.67（2026；flopper 1.59／2.9、UOB 1.61／2.65、AWS A100 第 5.6 年牌價／上市牌價 2.74／4.10）",
     EV_DIR + "/L2_hardware.md", "Analogy（現貨租金／AWS 牌價）",
     "Cap_In C228:E228（ρ 讀法乙：市場機會價值；Gov_Map GM638／GM639）", "ρ 0.05（低端）／0.5／1.0（高端），Assumed", "0.17（讀法甲，公式）／0.5／0.7（讀法乙）",
     "部分採用（E2 讀法乙）", "2026 供給吃緊；舊 GPU 若只能降級服務自家最低層則 ρ≈0.17，能轉租則約 0.6", "L1_LDemote*；L1_LTotal*", "—"),
    ("E396", "CoreWeave 加權平均約 5 年 take-or-pay（RPO $60.7B）、H100 續約約原價 95%",
     EV_DIR + "/L2_hardware.md（CoreWeave 10-K FY2025；Q3 2025 電話會議）", "Verified（10-K）／Interested-party（管理層說法）",
     "—（E4：只記錄；Tokenomics 不計長約）", "—", "出租方合約期 L≈1",
     "部分採用（E4，只記錄）", "長約只把價格風險轉給出租方，不提高實驗室的 L；轉知 company-models 的 neocloud 模型", "不改任何 IF_／L1_", "利害關係方"),
    ("E397", "雲端折舊年限：MSFT 2–6、GOOGL 6、ORCL 6、CRWV 6、META 5.5（2025 起）、AMZN 部分伺服器 6→5（2025）",
     EV_DIR + "/L2_hardware.md（各家 10-K）", "Verified（10-K）",
     "IF_DeprLifeIT（6 年；對照，不改）", "6 年", "2–6 年（會計年限）",
     "登錄（對照 IF_DeprLifeIT 6 年，不改）", "會計年限不等於經濟年限；Burry 主張經濟壽命 2～3 年（放空者，Interested-party）", "不改任何 IF_／L1_", "—"),
]


def evidence_append(wb):
    ws = wb["DB_Evidence"]
    have = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
    r = max(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")) + 1
    n = 0
    from common import put, F_CALC
    for (eid, claim, src, mark, param, cur, new, verdict, note, affected, stance) in EVIDENCE:
        if eid in have:
            continue
        vals = [eid, DATE, claim, src, mark, param, cur, new, verdict, V, note + "｜" + EV_COMMON, "已處理", "—", "—", affected, "—", stance,
                f"{verdict}｜v15.3 新增列（非 106 列積壓）"]
        for i, v in enumerate(vals):
            put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 4, 5, 6, 7, 8, 10, 14, 17))
        r += 1; n += 1
    return n


def apply(wb, append_text):
    log = {}
    log["l1_text"] = l1_text(wb)
    log["evidence"] = evidence_append(wb)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == "SL48-01" for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    # final guard: the early cell edits must have survived every later builder step
    ci = wb["Cap_In"]
    assert ci["C242"].value == C242_NEW and ci["D228"].value == RHO_D_NEW and wb["Calib"]["D156"].value == D156_NEW, "v15.3 cell edits were overwritten by a later builder step"
    return log


# ================================================================== README
README_VERSION = (VERSION + "（次版：" + WO + "——L 證據 B 段＋v15.2 修補：E1 Cap_In C226:E226（g）−0.10／−0.25／+0.05 → −0.09／−0.20／+0.05；"
                  "E2 Cap_In D228（ρ 低端）公式乘低階模型每 GPU 吞吐倍數（GB200 欄，約 0.17）、E228 1.0 → 0.7；E1／E2 標記 Assumed → Derived（修訂九更正）；E3 T1 維持；"
                  "D1 Cap_In 第 242 列改三態判讀（L 區間對所需 L 區間，恆等式讀法）、第 243 列改標對照；C2 Calib D156（CAL_BatchEtaD GB200）2.36 → =SQRT(C254*D201)（2.198），Gov_Map GM685 區間 1.59–3.04；"
                  "Gov_Map AA 欄 Decisions 檢查範圍 300 → 1000；Decisions SL48-01…08；DB_Evidence E393…E397；區間端點、指紋、Alloc N 節靜態表重產。具名範圍、定義不變。）以下為 "
                  + v1520.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1520.README_TITLE
README_TITLE = V + _T[len(v1520.V):]
README_ROW = ("v15.3 L 證據 B 段＋修補",
              "g 基準 −0.09（區間 −0.20～+0.05）、ρ 低端乘吞吐倍數 0.17／高端 0.7（皆改 Derived）；回本判讀改三態（L 區間對所需 L 區間）；CAL_BatchEtaD GB200 改兩反推點幾何平均公式（2.198）。名稱與定義不變。細節見 Decisions SL48-01…08。")
