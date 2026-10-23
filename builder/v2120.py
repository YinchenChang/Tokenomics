# v21.2 (work order docs/workorders/20261023_v21.2_計價倍數與說明修補.md r0, branch claude/v21-2): Alloc_In pricing multiple AL_SpendBasis 0.85／0.73／1.17 -> 1.8／0.85／2.3 (J5),
# plus four text fixes. MINOR version: no contract name is added or removed, definitions unchanged; values of the spend-route diagnostic names change.
#   W1  Alloc_In C23:E23 (blue inputs, written only while they still hold the v21.1 values), G23 note (original text kept verbatim after the new text), L1 C71／H71 (Assumed -> Derived, original kept; CLAUDE.md revision 9).
#   W2  SRC_Demand W10 (SRC_DEM_006): insert the 2025-09 external forecast after the v21.1 prefix.
#   W3  Cap_In G28 prefix (B4_Serve: spend ratio stands for GW ratio only if both sides have the same price multiple).
#   W4  NonNV L11 (Maia 200): range 0.25-0.75 sentence.
#   W5  Decisions SL59-01, DB_Evidence E403, README.   W6 (outside this file): tools/gen_fingerprint.py and tools/gen_ranges.py re-run.
# Every write to an Excel-owned cell is guarded (old value / old text / already-applied test), so the module is idempotent.
from openpyxl.utils import get_column_letter as L

from common import put, F_CALC
import v1410
import v2110

VERSION = "20261023_Tokenomics_v21.2"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-23"
DASH = "—"
V = "v21.2"
WO = "工作單 docs/workorders/20261023_v21.2_計價倍數與說明修補.md（r0）"
DEC_ID = "SL59-01"
EID = "E403"
EVD = "docs/evidence/20261023_Q側缺口_A段/"
DEC_BY = ("Andy 2026-10-23 21:58「按照預設」（9a Q 側缺口查證 J1–J5 定案）、22:02「開個新串來建置下一版」（J5 建置同意）；NonNV 說明為 v21.1 獨立查核觀察。"
          "證據：" + EVD + "（PR #155）")

OLD_VALS = (0.85, 0.73, 1.17)
NEW_VALS = (1.8, 0.85, 2.3)

G23_OLD = ("經濟口徑：依據的價格 ÷ 會計成本 1.16 × 會計 ÷ 經濟持有成本 0.73 ≈ 0.85（區間 0.73–1.17）；分母維持經濟口徑持有成本。"
           "小於 1 表示 Azure 計價未涵蓋全部資本成本，這本身是 m 的一項實質命題（決定紀錄 D10）")
G23_NEW_HEAD = ("v21.2（2026-10-23 9a 查證 J5；Andy 21:58「按照預設」、22:02 同意建置）：基準 1.8、區間 0.85–2.3，Derived。"
                "依據 " + EVD + "：OpenAI 2025 查核財報框出全部算力 161–217 億（口徑因子 s 0.79–1.0），GW 因子 0.89–1.11，價格因子為殘差，必須約 1.37–2.18 才能重現支出路線（必須為真）；"
                "Microsoft 帳另一切法 1.38–2.33、支出路線 0.50≈0.46 皆非獨立證據。非實付價直接證據（OpenAI 在 Azure 的實付價找不到）。"
                "低端 0.85＝v21.1 以前基準；高端 2.3≈Microsoft 帳上端。只影響支出路線（Cap_In 第 109、116–124、172–181、218–223、240、255、284–289、315–318、326、343 列與 L1 對應診斷名稱）；"
                "回本主口徑 L1_LReqFleet_Full、L1_RevGW_PQ_VR200、全部 IF_ 不受影響。｜原說明（v21.1 以前）：")
G23_NEW = G23_NEW_HEAD + G23_OLD

C71_OLD = "Alloc_In 計價倍數（Assumed）；X46"
C71_NEW = "Alloc_In 計價倍數（v21.2 起 Derived：基準 1.8、區間 0.85–2.3；X46）｜原：Alloc_In 計價倍數（Assumed）；X46"
H71_OLD = "1.0–1.6（Assumed）"
H71_NEW = "0.85–2.3（Derived；v21.2 9a J5）｜原：1.0–1.6（Assumed）"

DEM006_OLD_PREFIX = "出處待補；已排除 The Information 2026-02〈Why OpenAI, Anthropic Missed Their Own Gross Margin Forecasts〉（2026-10-23 查證 J5）｜"
DEM006_INSERT = "外部最近為 2025-09 預測逾 90 億（2026-10-23 9a 查證 J1，" + EVD + "）｜"

G28_PREFIX = "支出比代 GW 比，前提是兩邊價格倍數相同（2026-10-23 9a 查證；v21.2）｜"

L11_SENT = "產出比＝每 kW 峰值算力比（約 VR200 的 0.68）× 軟體效率因子 0.45–1.0（無第三方實測，主要不確定）。"
L11_INSERT = "區間 0.25–0.75：高端 0.75 需機架功率 ≤58 kW；低端含軟體成熟度下修（v21.2，v21.1 獨立查核觀察）。"


def _spend_row(wb):
    import v529
    ws = wb["Alloc_In"]
    r = v529._row_by_label(ws, v529.LAB_SPEND)
    assert r == 23, r
    return ws, r


# ================================================================== W1: guarded input writes (called after restore, with the other inputs_update hooks)
def inputs_update(wb):
    ws, r = _spend_row(wb)
    log = []
    for col, old, new in zip("CDE", OLD_VALS, NEW_VALS):
        c = ws[f"{col}{r}"]
        if c.value is not None and type(c.value) is not bool and c.value == old:
            c.value = new; log.append(f"Alloc_In!{col}{r}: {old} -> {new} (v21.2 W1 J5)")
    return log


# ================================================================== W1 L1 C71／H71 (hook in gov._rows_l1, like v2010.l1_patch)
def l1_patch(R):
    out, n = [], 0
    for row in R:
        row = list(row)
        if row[0] == "GapSpendBasis":
            assert row[2] == C71_OLD, row[2]
            assert row[7] == H71_OLD, row[7]
            row[2], row[7] = C71_NEW, H71_NEW; n += 1
        out.append(tuple(row))
    assert n == 1, n
    return out


# ================================================================== W1-3, W2, W3, W4: text
def text_fixes(wb):
    n = 0
    ws, r = _spend_row(wb)
    cur = ws[f"G{r}"].value
    if cur != G23_NEW:
        assert cur == G23_OLD, cur
        ws[f"G{r}"].value = G23_NEW; n += 1
    assert ws[f"F{r}"].value == "Derived", ws[f"F{r}"].value      # W1-2: F23 stays Derived

    sd = wb["SRC_Demand"]
    rr = [x for x in range(5, sd.max_row + 1) if sd.cell(x, 1).value == "SRC_DEM_006"]
    assert rr == [10], rr
    w = str(sd["W10"].value)
    if DEM006_INSERT not in w:
        assert w.startswith(DEM006_OLD_PREFIX), w
        sd["W10"].value = DEM006_OLD_PREFIX + DEM006_INSERT + w[len(DEM006_OLD_PREFIX):]; n += 1

    ci = wb["Cap_In"]
    assert ci["A28"].value == "對外服務占機隊", ci["A28"].value
    g = str(ci["G28"].value)
    if not g.startswith(G28_PREFIX):
        assert g.startswith("OpenAI 2025 推論 $8.4B"), g
        ci["G28"].value = G28_PREFIX + g; n += 1

    nn = wb["NonNV"]
    assert str(nn["A11"].value) == "Microsoft Maia 200", nn["A11"].value
    l = str(nn["L11"].value)
    if L11_INSERT not in l:
        assert l.startswith(L11_SENT), l
        nn["L11"].value = L11_SENT + L11_INSERT + l[len(L11_SENT):]; n += 1
    return n


# ================================================================== W5: DB_Evidence E403, Decisions SL59-01
EVIDENCE = [EID, DATE, "計價倍數 AL_SpendBasis：基準 1.8、區間 0.85–2.3（9a Q 側缺口查證 J5；價格因子為殘差，必須約 1.37–2.18 才能重現支出路線）",
            "9a Q 側缺口查證 " + EVD + "（PR #155）：OpenAI 2025 查核財報框出全部算力 161–217 億（口徑因子 s 0.79–1.0）、GW 因子 0.89–1.11；Microsoft 帳另一切法 1.38–2.33；chat 端查證 2026-10-23（本代理依工作單，未外部搜尋）",
            "Derived", "Alloc_In 第 23 列 C:E（AL_SpendBasis／_Lo／_Hi）；L1 第 71 列 C／H 文字",
            "AL_SpendBasis 0.85（0.73–1.17，Derived；v6.1 D10 經濟口徑）", "1.8（0.85–2.3）；殘差必須值約 1.37–2.18", "採用（V21.2 W1；非實付價直接證據）",
            V, "非實付價直接證據：OpenAI 在 Azure 的實付價找不到；殘差法 1.8 是『必須為真』的推導，不是觀測；Microsoft 帳另一切法 1.38–2.33、支出路線 0.50≈0.46 皆非獨立證據。"
            "證偽條件：找到 2025 平均 GW ≥ 1.8，或 Azure 對 OpenAI 單價 ≤ 2 美元／GPU 小時｜" + WO, "已處理",
            DASH, DASH, "AL_SpendBasis（＋_Lo／_Hi）及支出路線診斷名稱（L1_Gap*／L1_LReqFleet*／L1_MFleet*／L1_TokLedger* 等 46 個）；IF_／IFW_／IFC_ 不變", 3, "利害關係方", "採用（Derived）｜v21.2 新增列"]


def evidence_append(wb):
    ws = wb["DB_Evidence"]
    if any(ws.cell(r, 1).value == EID for r in range(5, ws.max_row + 1)):
        return 0
    r = max(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")) + 1
    assert ws.cell(r - 1, 1).value == "E402", ws.cell(r - 1, 1).value
    for i, v in enumerate(EVIDENCE):
        put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 4, 5, 6, 7, 8, 10, 14, 17))
    return 1


def _decisions():
    return [[DEC_ID, V, "v21.2 計價倍數 AL_SpendBasis 改 1.8（0.85–2.3）；四項說明修補",
             "W1：Alloc_In C23／D23／E23 0.85／0.73／1.17 → 1.8／0.85／2.3（Derived；價格因子為殘差，必須約 1.37–2.18 才能重現支出路線；非實付價直接證據），G23 說明改寫、原文逐字保留於其後；"
             "修訂九標記更正：L1 C71 由『Alloc_In 計價倍數（Assumed）；X46』、H71 由『1.0–1.6（Assumed）』更正為 Derived 版本，原值逐字保留在同格『｜原：』之後（更正版本 v21.2，依據 9a J5）。"
             "只影響支出路線與其診斷名稱；回本主口徑 L1_LReqFleet_Full、L1_RevGW_PQ_VR200、全部 IF_／IFW_／IFC_ 不變。"
             "W2：SRC_DEM_006 備註加『外部最近為 2025-09 預測逾 90 億』（J1）；W3：Cap_In G28 加『支出比代 GW 比，前提是兩邊價格倍數相同』（B4_Serve 0.41 不變，J2）；"
             "W4：NonNV L11 加 Maia 區間 0.25–0.75 說明（B11:G11 不變）。證偽條件：找到 2025 平均 GW ≥ 1.8，或 Azure 對 OpenAI 單價 ≤ 2 美元／GPU 小時。Andy 原話與依據：" + DEC_BY,
             DATE, DEC_BY, "v21.2 已建", "Alloc_In 第 23 列；L1 第 71 列；SRC_Demand W10；Cap_In G28；NonNV L11；DB_Evidence E403", f"{WO} W1–W6", "否"]]


def apply(wb, append_text):
    log = {"text_fixes": text_fixes(wb)}
    log["evidence"] = evidence_append(wb)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == DEC_ID for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


# ================================================================== README
README_VERSION = (VERSION + "（次版，底版 v21.1：" + WO + "——Alloc_In 計價倍數 AL_SpendBasis 0.85（0.73–1.17）→ 1.8（0.85–2.3），Derived（9a Q 側缺口查證 J5，價格因子為殘差；非實付價直接證據）；"
                  "只影響支出路線與其診斷名稱（L1_Gap*、L1_LReqFleet*（不含 _Full）、L1_MFleet*、L1_PriceRealFleet*、L1_TokCeiling*、L1_TokLedger* 等 46 個）；"
                  "SRC_DEM_006、Cap_In G28、NonNV L11 說明文字；區間端點與時效指紋重產；Decisions SL59-01；DB_Evidence E403。無新增或刪除契約名稱；全部 IF_／IFW_／IFC_、L1_LReqFleet_Full、L1_RevGW_PQ_VR200 的值不變。）以下為 "
                  + v2110.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v2110.README_TITLE
README_TITLE = V + _T[len(v2110.V):]
README_ROW = ("v21.2 計價倍數 1.8 與說明修補",
              "Alloc_In 計價倍數 AL_SpendBasis 基準 1.8、區間 0.85–2.3（Derived；9a 查證 J5：價格因子為殘差，必須約 1.37–2.18 才能重現支出路線，非實付價直接證據）。"
              "支出路線換算的 2025 服務 GW 由約 1.06 降到約 0.50；token 帳診斷 L1_TokLedger 由約 5.5 降到約 2.7——這是把 2 倍缺口歸給價格因子的結果，不是新的 Q 側證據。"
              "回本主口徑與 IF_ 不受影響。細節見工作單 docs/workorders/20261023_v21.2_計價倍數與說明修補.md 與 Decisions SL59-01。")
