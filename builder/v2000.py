# v20.0 (work order docs/workorders/20261022_v20.0_GW優先校準.md r0, branch claude/v20-0): 2025 calibration GW first, payback on whole-fleet holding cost, m range for 2026, comparison rows.
# MAJOR version: new contract names only; no existing named value changes (judgement text rows excepted; the old "無區間" text of L1_M_2026Q2／L1_M E／F becomes the new numeric range).
#   W1  Cap_In P section "GW 優先校準（v20.0）": AL_FleetGW25 (+_Lo/_Hi) and L1_ServeGW_GW25, L1_TokPaid_GW25, L1_TokLedger_GW25 (each +_Lo/_Hi).
#   W2  L1_HoldFleet25, L1_PayRatio25, L1_RevPerGW25, L1_LReqFleet_Full (each +_Lo/_Hi); Cap_In row 242 and L1 P89 read L1_LReqFleet_Full (v17.0 basis kept as a comparison row).
#   W3  L1_M_2026Q2 E/F range, L1_M E/F; Checks P1 (price basis, WARN); L1 column T nature "觀測比值" for L1_M／L1_M_*.
#   W4  SRC_DEM_088 (Codex tokens per day), L1_DPaid_T4k/T12k/T30k, L1_DPaid_wCodex, AL_ProdDerate_Lo/_Hi (Alloc_In), IF_ProdDerate_Lo/_Hi (Interface N), L1_RevGW_PQ_VR200 I-column tail.
#   W5  Cap_In G232／G233 text, Decisions SL55-01, DB_Evidence E401, README, label tails.
# Every write to an Excel-owned cell is guarded (old text / already-applied test); the new blue inputs are restored from the workbook on rebuild.
import re

from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

from common import put, F_NOTE, F_BOLD, F_CALC, F_LINK, F_IN, FILL_KEY
import v1000
import v1400
import v1410
import v1720
import v1800
import v1910

VERSION = "20261022_Tokenomics_v20.0"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-22"
V = "v20.0"
WO = "工作單 docs/workorders/20261022_v20.0_GW優先校準.md（r0）"
DASH = "—"
U = "AL_UnavailTxt"
NR = "無區間"
NR_LEDGER = "無區間（兩者區間來源不同，不組合）"
DEC_ID = "SL55-01"
EID = "E401"
SRC_ID = "SRC_DEM_088"
DEC_BY = ("Andy 2026-10-22 05:01「總算力支出是從持有的機隊成本去推算出來的」；05:15「計算上一定要先從 tokenomics 開始」「兩百零四億，很明顯是高估了」；"
          "05:17（平均 GW 用 Tokenomics 0.6→1.9 內插）；07:34 公式「Q＝1 GW ×（1 − 訓練占比）×（1 − 免費占比）× 理論每 GW token × 利用率 × 效率」；"
          "16:11（OpenAI 特有資料留 TK 作「產業校準錨點」）；20:57「go」")


def _t(s):
    return f'="{s}"'


def _num(*xs):
    return ",".join(f"ISNUMBER({x})" for x in xs)


def _safe(args, expr):
    return f"=IF(AND({_num(*args)}),{expr},{U})"


def _div(num, den):
    """num / den with numeric and positive-denominator guards (ISNUMBER first, then > 0)."""
    return f"=IF(AND({_num(num, den)}),IF({den}>0,{num}/{den},{U}),{U})"


# ================================================================== Cap_In P section
ROWS = {}
PURPOSE_MARKS = {}
SEC_TITLE = ("P. GW 優先校準與回本全部算力成本口徑（v20.0；" + WO + "）：2025 年校準先由 Tokenomics 的 GW 起算（平均上線 GW → 服務 GW → 付費 token），"
             "再與需求面對照；回本改以全部算力的持有成本為口徑（Andy 2026-10-22）。欄位：C 基準／D 低端／E 高端；G 標記、H 說明、O 用途")

LAB = dict(
    gwlo="2025 平均上線 GW 低端（後段集中；Epoch 方法頁與 Abilene 時程；Derived）　→ AL_FleetGW25_Lo",
    fleetgw="2025 平均上線 GW：C 指數內插 (G1−G0)÷LN(G1÷G0)（SRC_DEM_008→009）、D 後段集中、E 線性 (G0+G1)÷2　→ AL_FleetGW25／_Lo／_Hi",
    servegw="2025 服務 GW＝平均上線 GW × 對外服務占機隊（B4_Serve；D 取 0.30、E 取 0.60）　→ L1_ServeGW_GW25",
    sigma="混合每 GW 付費產能 Σ（AL_MixGen 世代占比 × IF_TokGWPaidMix 基準欄：Hopper 第 2、GB200 第 5、GB300 第 8、VR200 第 11 欄；M tok/GW/年）",
    tokpaid="GW 優先：2025 付費 token＝服務 GW ×（1 − 免費占比）× Σ × 實際利用率 × 生產折減　→ L1_TokPaid_GW25",
    ledger="GW 優先 token 帳＝GW 優先付費 token ÷ 需求面付費 token（AL_DPaid）　→ L1_TokLedger_GW25",
    hold="全部算力年持有成本＝平均上線 GW × 機隊每 GW 年持有成本（AL_FleetHold；$B/年）　→ L1_HoldFleet25",
    payratio="回本比（實際 ÷ 所需）＝2025 營收（SRC_DEM_007）÷ 全部算力年持有成本　→ L1_PayRatio25（D 低＝÷ 高端成本、E 高＝÷ 低端成本）",
    revpergw="每上線 GW 營收＝2025 營收 ÷ 平均上線 GW　→ L1_RevPerGW25（D 低＝÷ 高端 GW、E 高＝÷ 低端 GW）",
    lreq="全部算力回本所需 L＝全部算力年持有成本 ÷ 2025 營收　→ L1_LReqFleet_Full",
    pd="生產折減的效率區間（情境端點；D 低 0.72、E 高 0.95；基準＝CTL_ProdDerate）　→ AL_ProdDerate_Lo／_Hi（v20.0 W4）",
    tokscen="每則提示 token 情境（對照列用；D4000／12000／30000 對應 C／D／E 欄）",
    dscen="付費 D 情境＝AL_DAPI ＋ SRC_DEM_012 × 1e9 × 每則 token × 365 ÷ 1e6 ×（1 − AL_FreeShare）　→ L1_DPaid_T4k／T12k／T30k（對照，不驅動）",
    dcodex="付費 D ＋ Codex＝AL_DPaid ＋ SRC_DEM_088 × 1e12 × 365 ÷ 1e6（假設 Codex token 全為付費）　→ L1_DPaid_wCodex（對照，不驅動）",
    j_v17="對照：v17.0 主口徑（折現世代起點 L 對只含推論、經計價倍數的所需 L）：L1_LTotalFleetVDisc 區間 對 L1_LReqFleet_IdNoK 區間；三態",
)
NOTE = dict(
    pd="Assumed：0.72＝Hopper SLO 不可達門檻附近、0.95＝軟體進步高端。只用於 L1_TokPaid_GW25 的低／高與 IF_ProdDerate_Lo／_Hi，不改 CTL_ProdDerate。藍字輸入，Excel 擁有。位置：工作單寫 Alloc_In 新列，但 Alloc_In 末段在 gov_all 之後才建、Gov_Map 登錄需列標籤已存在，故放 Cap_In P 節（待 Project 判斷）",
    gwlo="D 欄用的低端值。2025 年上線 GW 若集中在年底（後段集中），平均上線 GW 接近年初 0.6 與年底 1.9 之間偏低的位置；1.0 取自 Epoch 方法頁與 Abilene 時程（chat 端查證；Derived）。藍字輸入，Excel 擁有",
    fleetgw=("端點為 OpenAI 官方（SRC_DEM_008 2024 年底 0.6 GW、SRC_DEM_009 2025 年底 1.9 GW；Interested-party，GW 口徑未明），在本模型標為『產業校準錨點』。"
             "基準取指數內插（成長率固定）；高端為線性；低端為後段集中。Andy 2026-10-22 05:17 指示以 0.6→1.9 內插；基準取指數或線性待 Project 判斷（工作單：指數為基準）"),
    servegw="服務占比 B4_Serve（Cap_In 第 28 列）C 基準 0.41、D 0.30、E 0.60；服務 GW 只含對外服務，其餘為訓練與研發",
    sigma="Σ＝Hopper 0.6 × 26,017,336,543 ＋ GB200 0.4 × 128,013,563,495（事前 66,815,827,324）；世代占比取 Alloc_In 2025 服務世代組合（AL_MixGen）；各欄共用同一個 Σ",
    tokpaid="C：免費占比 IF_FreeShare（支出口徑 0.46）、實際利用率 AL_UtilActual、生產折減 CTL_ProdDerate；D 低：服務 0.30、免費 0.60、利用 AL_UtilActual_Lo、效率 AL_ProdDerate_Lo；E 高：服務 0.60、免費 0.35、利用 AL_UtilActual_Hi、效率 AL_ProdDerate_Hi（Andy 07:34 公式；免費用支出口徑）",
    ledger="讀法：GW 優先後供給 ÷ 需求。舊 L1_TokLedger（5.62）為支出路線，保留對照。需求面 AL_DPaid 無區間，Lo／Hi 都除以同一個基準 AL_DPaid",
    hold="機隊每 GW 年持有成本 AL_FleetHold（Alloc 第 157 列基準欄，9.353 $B/GW/年）；低／高只隨 GW 端點變動，不隨持有成本區間變動",
    payratio="＝每上線 GW 營收 ÷ 每 GW 持有成本（Andy 的「實際 ÷ 所需」）；＞ 1 表示營收高於全部算力持有成本",
    revpergw="分母用 2025 平均上線 GW（非年底 GW）",
    lreq="壽命期營收 × L ≥ 全部算力年持有成本；不含計價倍數與 κ，因成本已是持有成本。判讀（Cap_In 第 242 列、L1 第 89 列）的所需 L 側讀本列（v20.0）",
    tokscen="Assumed 單值（3 情境）：4,000 為 Alloc_In 基準；12,000 與 30,000 為對照情境。藍字輸入，Excel 擁有",
    dscen="對照列，不驅動。企業版訊息未計（SRC_DEM_012 只含消費方案，已搜尋未找到企業量）",
    dcodex="對照列，不驅動。OpenAI Codex 每日 token（2025-10，約 21 天 40 兆 → 約 2 兆／日）取自 OpenAI DevDay 2025-10-06（經 Epoch ai_companies_usage_reports.csv）；假設全為付費",
    j_v17="舊主口徑保留為對照（v17.0 基準：L 0.759 對所需 L 0.756 → 未定）；三態文字，只作對照",
)
SCEN_TOK = (4000, 12000, 30000)


def sigma_formula():
    m = lambda k: f"INDEX(AL_MixGen,{k},1)"
    p = lambda col: f"INDEX(IF_TokGWPaidMix,1,{col})"
    ks = ((1, 2), (2, 5), (3, 8), (4, 11))
    allnum = ",".join(f"ISNUMBER({m(k)})" for k, _ in ks)
    bad = ",".join(f"AND({m(k)}>0,NOT(ISNUMBER({p(c)})))" for k, c in ks)
    tot = "+".join(f"IF(ISNUMBER({p(c)}),{m(k)}*{p(c)},0)" for k, c in ks)
    return f"=IF(AND({allnum}),IF(OR({bad}),{U},{tot}),{U})"


def cap_in_p(wb, o_sec):
    ws = wb["Cap_In"]
    r0 = o_sec["end"] + 2
    assert ws.max_row == o_sec["end"], (ws.max_row, o_sec["end"])
    put(ws, f"A{r0}", SEC_TITLE, F_BOLD, wrap=True)
    r = r0 + 1
    order = ["gwlo", "pd", "fleetgw", "servegw", "sigma", "tokpaid", "ledger", "hold", "payratio", "revpergw", "lreq", "tokscen", "dscen", "dcodex", "j_v17"]
    for k in order:
        ROWS[k] = r; r += 1
    R = ROWS

    def row(k, unit, vals, tag, fmt, blue=False, key=False, purpose=None):
        rr = R[k]
        put(ws, f"A{rr}", LAB[k], wrap=True); put(ws, f"B{rr}", unit)
        for col, v in zip("CDE", vals):
            if v is None:
                continue
            isf = isinstance(v, str) and v.startswith("=")
            put(ws, f"{col}{rr}", v, (F_IN if blue else (F_LINK if isf and "!" in v else F_CALC)), fmt=fmt, fill=FILL_KEY if (key and col == "C") else None)
        put(ws, f"G{rr}", tag, F_NOTE); put(ws, f"H{rr}", NOTE[k], F_NOTE, wrap=True)
        if purpose:
            PURPOSE_MARKS[rr] = purpose

    G0, G1 = "SRC_DEM_008", "SRC_DEM_009"
    fg = R["fleetgw"]
    row("gwlo", "GW", [1.0], "Derived", "0.000", blue=True)
    row("pd", "x", [None, 0.72, 0.95], "Assumed", "0.00", blue=True)
    row("fleetgw", "GW", [
        f"=IF(AND({_num(G0, G1)}),IF(AND({G0}>0,{G1}>0),IF({G1}={G0},{G0},(({G1})-({G0}))/LN({G1}/{G0})),{U}),{U})",
        f"=IF(ISNUMBER($C${R['gwlo']}),$C${R['gwlo']},{U})",
        _safe((G0, G1), f"({G0}+{G1})/2")], "Derived", "0.00000", key=True)
    sv = lambda c: f"=IF(AND(ISNUMBER({c}{fg}),ISNUMBER({c}$28)),{c}{fg}*{c}$28,{U})"
    row("servegw", "GW", [sv("C"), sv("D"), sv("E")], "Derived", "0.00000", key=True)
    row("sigma", "M tok/GW/年", [sigma_formula()], "Derived", "#,##0")
    sg, sig = R["servegw"], f"$C${R['sigma']}"
    tp = lambda c, free, util, eff: (f"=IF(AND({_num(f'{c}{sg}', sig, free, util, eff)}),{c}{sg}*(1-{free})*{sig}*{util}*{eff},{U})")
    row("tokpaid", "M tok/年", [tp("C", "IF_FreeShare", "AL_UtilActual", "CTL_ProdDerate"),
                                tp("D", "$E$29", "AL_UtilActual_Lo", "AL_ProdDerate_Lo"),
                                tp("E", "$D$29", "AL_UtilActual_Hi", "AL_ProdDerate_Hi")], "Derived", "#,##0", key=True)
    t = R["tokpaid"]
    row("ledger", "x", [_div(f"{c}{t}", "AL_DPaid") for c in "CDE"], "Derived", "0.00000", key=True)
    row("hold", "$B/年", [_safe((f"{c}{fg}", "AL_FleetHold"), f"{c}{fg}*AL_FleetHold") for c in "CDE"], "Derived", "0.0000", key=True)
    h = R["hold"]
    row("payratio", "x", [_div("SRC_DEM_007", f"C{h}"), _div("SRC_DEM_007", f"E{h}"), _div("SRC_DEM_007", f"D{h}")], "Derived", "0.00000", key=True)
    row("revpergw", "$B/GW/年", [_div("SRC_DEM_007", f"C{fg}"), _div("SRC_DEM_007", f"E{fg}"), _div("SRC_DEM_007", f"D{fg}")], "Derived", "0.0000", key=True)
    row("lreq", "x", [f"=IF(AND({_num('SRC_DEM_007', f'{c}{h}')}),IF(SRC_DEM_007>0,{c}{h}/SRC_DEM_007,{U}),{U})" for c in "CDE"], "Derived", "0.00000", key=True)
    row("tokscen", "tok", list(SCEN_TOK), "Assumed", "#,##0", blue=True)
    ts = R["tokscen"]
    row("dscen", "M tok/年", [_safe(("AL_DAPI", "SRC_DEM_012", "AL_FreeShare", f"{c}{ts}"), f"AL_DAPI+SRC_DEM_012*1E9*{c}{ts}*365/1E6*(1-AL_FreeShare)") for c in "CDE"],
        "Derived", "#,##0")
    row("dcodex", "M tok/年", [_safe(("AL_DPaid", SRC_ID), f"AL_DPaid+{SRC_ID}*1E12*365/1E6")], "Derived", "#,##0")
    import v1600
    row("j_v17", "x", [v1600.tri("L1_LTotalFleetVDisc", "L1_LTotalFleetVDisc_Lo", "L1_LTotalFleetVDisc_Hi", "L1_LReqFleet_IdNoK", "L1_LReqFleet_IdNoK_Lo", "L1_LReqFleet_IdNoK_Hi",
                                  "折現世代起點 L（v17.0 口徑）")], "Derived", None,
        purpose="顯示用：C%d（對照列：v17.0 主口徑，只含推論、經計價倍數）" % R["j_v17"])
    # names (AL_FleetGW25 lives here; the L1 rows read the other cells)
    for suf, col in (("", "C"), ("_Lo", "D"), ("_Hi", "E")):
        wb.defined_names["AL_FleetGW25" + suf] = DefinedName("AL_FleetGW25" + suf, attr_text=f"Cap_In!${col}${fg}")
    wb.defined_names["AL_ProdDerate_Lo"] = DefinedName("AL_ProdDerate_Lo", attr_text=f"Cap_In!$D${R['pd']}")
    wb.defined_names["AL_ProdDerate_Hi"] = DefinedName("AL_ProdDerate_Hi", attr_text=f"Cap_In!$E${R['pd']}")
    return dict(start=r0, end=r - 1, rows=dict(R))


# ---------------------------------------------------------------- Interface N: IF_ProdDerate_Lo／_Hi in D／E of the IF_ProdDerate row
def interface_n_ext(wb):
    ws = wb["Interface"]
    ref = wb.defined_names["IF_ProdDerate"].attr_text
    r = int(ref.split("$")[-1])
    assert "IF_ProdDerate" in str(ws[f"A{r}"].value), ws[f"A{r}"].value
    put(ws, f"D{r}", "=AL_ProdDerate_Lo", fmt="0.00"); put(ws, f"E{r}", "=AL_ProdDerate_Hi", fmt="0.00")
    wb.defined_names["IF_ProdDerate_Lo"] = DefinedName("IF_ProdDerate_Lo", attr_text=f"Interface!$D${r}")
    wb.defined_names["IF_ProdDerate_Hi"] = DefinedName("IF_ProdDerate_Hi", attr_text=f"Interface!$E${r}")
    tail = "（v20.0：D／E 欄＝效率區間 AL_ProdDerate_Lo／_Hi，基準 C 欄不變）"
    if tail not in str(ws[f"A{r}"].value):
        ws[f"A{r}"].value = f"{ws[f'A{r}'].value}{tail}"
    return r


# ================================================================== Gov_Map rows (blue inputs)
def gov_map_rows(wb):
    if not ROWS:
        return []
    out = []
    ws = wb["Cap_In"]
    sc = "切片三（v20.0 GW 優先校準）"
    R = ROWS

    def ent(sheet, row, cell, role, cls, lo, hi, rtext, reason, seg="無"):
        lab = wb[sheet][f"A{row}"].value
        out.append({'scope': sc, 'sheet': sheet, 'cell': cell, 'label': lab, 'check': lab, 'cls': cls, 'role': role, 'src': '', 'rel': '', 'dec': DEC_ID,
                    'lo': lo, 'hi': hi, 'rtext': rtext, 'reason': reason, 'retag': '', 'seg': seg})

    ent("Cap_In", R["gwlo"], f"C{R['gwlo']}", "單值", "Derived", None, None, "單值 1.0 GW（低端；基準 1.128 為指數內插、高端 1.25 為線性）",
        "v20.0 W1（" + WO + "；Decisions SL55-01）：2025 平均上線 GW 的低端（後段集中），Epoch 方法頁與 Abilene 時程（chat 端查證）；端點 SRC_DEM_008／009 為 OpenAI 官方，標產業校準錨點。只進 Cap_In P 節與 AL_FleetGW25_Lo")
    ent("Cap_In", R["tokscen"], f"C{R['tokscen']}:E{R['tokscen']}", "單值", "Assumed", None, None, "4,000／12,000／30,000 tok（C／D／E 欄為三個對照情境）",
        "v20.0 W4（SL55-01）：每則提示 token 對照情境，只用於 L1_DPaid_T4k／T12k／T30k（對照列，不驅動）；4,000 同 Alloc_In 基準")
    ent("Cap_In", R["pd"], f"D{R['pd']}:E{R['pd']}", "低／高", "Assumed", None, None, "0.72（低）～ 0.95（高）；基準 CTL_ProdDerate 不變",
        "v20.0 W4（SL55-01）：生產折減的效率區間。0.72＝Hopper SLO 不可達門檻附近、0.95＝軟體進步高端；只用於 L1_TokPaid_GW25 低／高與 IF_ProdDerate_Lo／_Hi")
    return out


# ================================================================== SRC (SRC_Demand) record
def src_new_records(wb):
    rec = v1720._rec(
        SRC_ID, "SRC_Demand", "OpenAI Codex 每日 token", 2.0, None, None, "T tok/日", "2025-10；約 21 天 40 兆 → 約 2 兆／日", "OpenAI", "2025-10-06",
        "OpenAI DevDay 2025-10-06（經 Epoch ai_companies_usage_reports.csv）", 3, "利害關係方", "OpenAI 自揭使用量（宣傳口徑）；經 Epoch 資料集轉引，未讀原文",
        "二手（Epoch 資料集轉引）", "L1_DPaid_wCodex（對照，不驅動）｜換算",
        f"{V} W4：Codex 每日 token 約 2 兆（40 兆 ÷ 約 21 天），2025-10 時點值；僅作 L1_DPaid_wCodex 對照列（假設全為付費），不進任何驅動格（{WO}）")
    rec.update(status="Active", s="U-V2000", ev=EID)
    return [rec]


# ================================================================== L1 rows
NONE_EXT = ("無外部對照", DASH, DASH)
READ_GW = ("v20.0（Andy 2026-10-22）：2025 校準先由 Tokenomics 的 GW 起算——平均上線 GW（SRC_DEM_008 → 009 內插，產業校準錨點）× 對外服務占比 ×（1 − 免費占比）× 混合每 GW 付費產能 × 實際利用率 × 生產折減。"
           "支出路線（AL_ServeGWSpend、L1_TokLedger 5.62）保留為對照。")
I_REV_PQ_TAIL = "｜殘差未分配：GW 優先校準後 2025 供給 ÷ 需求為 L1_TokLedger_GW25；若殘差在利用率，Q 應下修"


def l1_rows_new():
    R = ROWS
    def cap(col, key): return f"=Cap_In!${col}${R[key]}"
    def tri3(key, lab, cond, unit, rdef, read, drv, weak, gap, fmt_note=""):
        return (key, lab, cond, cap("C", {"ServeGW_GW25": "servegw", "TokPaid_GW25": "tokpaid", "TokLedger_GW25": "ledger", "HoldFleet25": "hold", "PayRatio25": "payratio",
                                          "RevPerGW25": "revpergw", "LReqFleet_Full": "lreq"}[key]),
                cap("D", {"ServeGW_GW25": "servegw", "TokPaid_GW25": "tokpaid", "TokLedger_GW25": "ledger", "HoldFleet25": "hold", "PayRatio25": "payratio",
                          "RevPerGW25": "revpergw", "LReqFleet_Full": "lreq"}[key]),
                cap("E", {"ServeGW_GW25": "servegw", "TokPaid_GW25": "tokpaid", "TokLedger_GW25": "ledger", "HoldFleet25": "hold", "PayRatio25": "payratio",
                          "RevPerGW25": "revpergw", "LReqFleet_Full": "lreq"}[key]),
                unit, rdef, read, drv, weak, *NONE_EXT, "AL_FleetGW25；SRC_DEM_007／008／009", f"Cap_In 第 {R[{'ServeGW_GW25': 'servegw', 'TokPaid_GW25': 'tokpaid', 'TokLedger_GW25': 'ledger', 'HoldFleet25': 'hold', 'PayRatio25': 'payratio', 'RevPerGW25': 'revpergw', 'LReqFleet_Full': 'lreq'}[key]]} 列", gap)
    return [
        tri3("ServeGW_GW25", "2025 服務 GW（GW 優先）＝平均上線 GW × 對外服務占比", "AL_FleetGW25（指數內插 0.6→1.9 GW）× B4_Serve；低＝1.0 GW × 0.30、高＝1.25 GW × 0.60", "GW",
             "低＝後段集中 GW 1.0 × 服務占比 0.30；高＝線性 GW 1.25 × 0.60", READ_GW, "平均上線 GW、服務占比 B4_Serve", "GW 端點：OpenAI 官方（Interested-party）；內插形狀：Derived；服務占比：Interested-party／Derived",
             "平均上線 GW 的內插形狀無一手資料；GW 口徑未明（D1）"),
        tri3("TokPaid_GW25", "2025 付費 token（GW 優先；M tok/年）", "＝L1_ServeGW_GW25 ×（1 − IF_FreeShare）× Σ（AL_MixGen × IF_TokGWPaidMix 基準欄）× AL_UtilActual × CTL_ProdDerate", "M tok/年",
             "低＝服務 0.30、免費 0.60、利用 0.30、效率 AL_ProdDerate_Lo 0.72、GW 1.0；高＝服務 0.60、免費 0.35、利用 0.65、效率 AL_ProdDerate_Hi 0.95、GW 1.25", READ_GW + " 免費切分用支出口徑（Andy 07:34）。",
             "平均上線 GW、服務占比、免費占比、混合每 GW 付費產能、利用率、生產折減", "利用率、效率區間、服務占比：Assumed／Interested-party；產能：roofline 推導（Derived）",
             "token 路線與支出路線的免費切分口徑不同（見 L1_ExtFreeShare）"),
        tri3("TokLedger_GW25", "token 帳（GW 優先）：GW 優先付費 token ÷ 需求面付費 token（AL_DPaid）", "＝L1_TokPaid_GW25 ÷ AL_DPaid；低／高各除以同一個基準 AL_DPaid", "x",
             "低／高＝付費 token 的低／高端 ÷ 基準 AL_DPaid（需求面無區間）", "v20.0：GW 優先後供給 ÷ 需求；舊 L1_TokLedger（5.62）為支出路線，保留對照。＞ 1 表示模型供給多於需求面觀測；＜ 1 則相反。不是 m。",
             "L1_TokPaid_GW25、AL_DPaid", "兩者分母皆為模型推算 token（Assumed）", "兩路線差距未分解；需求面無區間"),
        tri3("HoldFleet25", "2025 全部算力年持有成本（$B/年）＝平均上線 GW × 機隊每 GW 年持有成本", "AL_FleetGW25 × AL_FleetHold（低／高只隨 GW 端點）", "$B/年",
             "低＝GW 1.0 × AL_FleetHold；高＝GW 1.25 × AL_FleetHold（持有成本取基準 9.353 $B/GW/年）", "v20.0：全部算力持有成本（Andy 10-21 14:55 範圍、10-22 05:01 口徑），不再只含推論。",
             "平均上線 GW、AL_FleetHold", "GW：OpenAI 官方端點＋內插；持有成本：Derived", "持有成本區間未組合進低／高"),
        tri3("PayRatio25", "回本比（全部算力）＝2025 營收 ÷ 全部算力年持有成本（實際 ÷ 所需）", "＝SRC_DEM_007 ÷ L1_HoldFleet25；低＝÷ 高端成本、高＝÷ 低端成本", "x",
             "低＝÷ 高端成本（GW 1.25）；高＝÷ 低端成本（GW 1.0）", "v20.0：＞ 1 表示 2025 營收高於全部算力年持有成本；＝ 1 ÷ L1_LReqFleet_Full。", "SRC_DEM_007、L1_HoldFleet25",
             "營收：Interested-party；GW：同上", "營收與 GW 口徑（年度 vs 年底）不同"),
        tri3("RevPerGW25", "每上線 GW 營收（2025；$B/GW/年）＝SRC_DEM_007 ÷ 平均上線 GW", "＝SRC_DEM_007 ÷ AL_FleetGW25；低＝÷ 高端 GW、高＝÷ 低端 GW", "$B/GW/年",
             "低＝÷ GW 1.25；高＝÷ GW 1.0", "v20.0：實際營收的每 GW 口徑（與理論上限 L1_RevGW_Fleet_VR200 口徑不同，不得相減）。", "SRC_DEM_007、AL_FleetGW25", "同上", "GW 口徑未明"),
        tri3("LReqFleet_Full", "回本所需 L — 全部算力持有成本口徑（v20.0 起主口徑）", "＝L1_HoldFleet25 ÷ SRC_DEM_007（不含計價倍數與 κ，因成本已是持有成本）", "x",
             "低＝GW 1.0 的持有成本 ÷ 營收；高＝GW 1.25 的持有成本 ÷ 營收", "v20.0：壽命期營收 × L ≥ 全部算力年持有成本。判讀（Cap_In 第 242 列、L1 第 89 列）的所需 L 側讀本列；v17.0 主口徑 L1_LReqFleet_IdNoK（只含推論、經計價倍數）改為對照。",
             "L1_HoldFleet25、SRC_DEM_007", "同 L1_HoldFleet25", "持有成本區間未組合進低／高"),
        *[(f"DPaid_{nm}", f"付費 D 情境（對照；每則 {tk} token）",
           "＝AL_DAPI ＋ SRC_DEM_012 × 1e9 × 每則 token × 365 ÷ 1e6 ×（1 − AL_FreeShare）；單值情境，不是基準／低／高區間",
           cap(col, "dscen"), _t(NR), _t(NR), "M tok/年", NR,
           "對照列，不驅動。企業版訊息未計（SRC_DEM_012 只含消費方案，已搜尋未找到企業量）。", "AL_DAPI、SRC_DEM_012、每則 token（Cap_In P 節藍字）、AL_FreeShare", "每則 token：Assumed；訊息數：Interested-party",
           *NONE_EXT, "SRC_DEM_012；AL_FreeShare", f"Cap_In 第 {R['dscen']} 列（{col} 欄）", "企業版訊息未計")
          for nm, tk, col in (("T4k", "4,000", "C"), ("T12k", "12,000", "D"), ("T30k", "30,000", "E"))],
        ("DPaid_wCodex", "付費 D ＋ Codex（對照；假設 Codex token 全為付費）", "＝AL_DPaid ＋ SRC_DEM_088 × 1e12 × 365 ÷ 1e6", cap("C", "dcodex"), _t(NR), _t(NR), "M tok/年", NR,
         "對照列，不驅動。Codex 每日約 2 兆 token（OpenAI DevDay 2025-10-06，經 Epoch）。", "AL_DPaid、SRC_DEM_088", "Codex：Interested-party（3 級）", *NONE_EXT, "AL_DPaid；SRC_DEM_088",
         f"Cap_In 第 {R['dcodex']} 列", "Codex 是否全為付費無資料"),
    ]


KEYS = ["ServeGW_GW25", "TokPaid_GW25", "TokLedger_GW25", "HoldFleet25", "PayRatio25", "RevPerGW25", "LReqFleet_Full", "DPaid_T4k", "DPaid_T12k", "DPaid_T30k", "DPaid_wCodex"]
NEW_L1 = [f"L1_{k}{s}" for k in KEYS for s in ("", "_Lo", "_Hi")]
NEW_NAMES = NEW_L1 + [f"AL_FleetGW25{s}" for s in ("", "_Lo", "_Hi")] + ["AL_ProdDerate_Lo", "AL_ProdDerate_Hi", "IF_ProdDerate_Lo", "IF_ProdDerate_Hi", SRC_ID]

v1000._n(v1000.E_, "GW 優先：平均上線 GW → 服務 GW → 付費 token（自下而上；GW 端點為 OpenAI 官方、占比與利用率為 Assumed）", "ServeGW_GW25", "TokPaid_GW25")
v1000._n(v1000.U_, "GW 優先供給 ÷ 需求面：兩路線的差距，未分解（token 帳診斷，不是 m）", "TokLedger_GW25")
v1000._n(v1000.E_, "全部算力持有成本、每上線 GW 營收、回本比：平均上線 GW 乘持有成本或除營收（比值輸出）", "HoldFleet25", "PayRatio25", "RevPerGW25")
v1000._n(v1000.I_, "全部算力年持有成本 ÷ 實收營收：同一口徑的恆等式讀法（＝ 1 ÷ L1_PayRatio25）", "LReqFleet_Full")
v1000._n(v1000.E_, "付費 D 的對照情境（每則 token 與 Codex；不驅動）", "DPaid_T4k", "DPaid_T12k", "DPaid_T30k", "DPaid_wCodex")
OBS = "觀測比值"
v1000._n(OBS, "m＝實收營收 ÷（觀測的付費 token × 有效牌價）：分子分母皆為觀測量（營收、API token、訊息數），不再是殘差（v20.0；Andy 2026-10-22）",
         "M_2025", "M_2026Q1", "M_2026Q2", "M_FY2026", "M", "M_API", "M_Sub", "M_Seat")


def l1_patch(R):
    """M_2026Q2 gets a numeric low/high (token-high / token-low, same recipe as 2025) and L1_M reads it; both rows' "無區間" text is replaced."""
    out = []
    ci = "Cap_In"
    den = lambda tok, free: (f"(({ci}!$E$196)+{ci}!$E$197*1E9*{tok}*365/1E6*(1-{free}))*{ci}!$E$205")
    lo = (f"=IF(AND({_num(f'{ci}!$E$187', f'{ci}!$E$196', f'{ci}!$E$197', f'{ci}!$E$205', 'AL_TokPerPrompt_Hi', 'AL_FreeShare_Lo')}),"
          f"IF({den('AL_TokPerPrompt_Hi', 'AL_FreeShare_Lo')}>0,{ci}!$E$187*1E9/({den('AL_TokPerPrompt_Hi', 'AL_FreeShare_Lo')}),{U}),{U})")
    hi = (f"=IF(AND({_num(f'{ci}!$E$187', f'{ci}!$E$196', f'{ci}!$E$197', f'{ci}!$E$205', 'AL_TokPerPrompt_Lo', 'AL_FreeShare_Hi')}),"
          f"IF({den('AL_TokPerPrompt_Lo', 'AL_FreeShare_Hi')}>0,{ci}!$E$187*1E9/({den('AL_TokPerPrompt_Lo', 'AL_FreeShare_Hi')}),{U}),{U})")
    RD = ("低＝token 取高（每則 AL_TokPerPrompt_Hi 11,000、免費占比 AL_FreeShare_Lo 0.4）；高＝token 取低（每則 AL_TokPerPrompt_Lo 2,000、免費占比 AL_FreeShare_Hi 0.8）；"
          "營收、API token、單價、每日則數不變（v20.0；同 2025 寫法）")
    for row in R:
        row = list(row)
        if row[0] == "M_2026Q2":
            row[4], row[5], row[7] = lo, hi, RD
        elif row[0] == "M":
            row[4], row[5], row[7] = "=L1_M_2026Q2_Lo", "=L1_M_2026Q2_Hi", "＝L1_M_2026Q2 的區間（v20.0）：" + RD
        out.append(tuple(row))
    return out


# ================================================================== apply
TAIL_NOPRIM = "｜v20.0 起非主口徑（只含推論、支出路線）"
TAIL_SPEND = "｜v20.0 起：支出路線為對照，2025 校準主口徑為 GW 優先（L1_ServeGW_GW25）"


def _name_cell(wb, name):
    ref = wb.defined_names[name].attr_text
    sh, cell = ref.split("!")
    return wb[sh.strip("'")], cell.replace("$", "").split(":")[0]


def label_tails(wb, append_text):
    n = 0
    ws = wb["L1"]
    for key in ("LReqFleet_IdNoK", "LReqFleet_Id"):
        r = int(wb.defined_names[f"L1_{key}"].attr_text.split("$")[-1])
        n += int(bool(append_text(ws[f"B{r}"], TAIL_NOPRIM)))
    for name in ("AL_ServeGWSpend", "IF_AllocServeGW"):
        sh, cell = _name_cell(wb, name)
        r = int(re.sub(r"[A-Z]+", "", cell))
        n += int(bool(append_text(sh[f"A{r}"], TAIL_SPEND)))
    # I column of L1_RevGW_PQ_VR200
    r = int(wb.defined_names["L1_RevGW_PQ_VR200"].attr_text.split("$")[-1])
    n += int(bool(append_text(ws[f"I{r}"], I_REV_PQ_TAIL)))
    return n


def judge_update(wb, append_text):
    """W2: Cap_In row 242 and L1 P89 read L1_LReqFleet_Full (v17.0 basis kept as a comparison row in Cap_In P)."""
    log = {}
    ws = wb["Cap_In"]
    n = 0
    for c in "CDE":
        cur = ws[f"{c}242"].value
        assert isinstance(cur, str) and cur.startswith("=IF(AND(ISNUMBER(L1_LTotalFleetVDisc"), cur[:60]
        if "L1_LReqFleet_Full" in cur:
            continue
        new = cur.replace("L1_LReqFleet_IdNoK", "L1_LReqFleet_Full").replace("所需 L（不含 κ，低端）", "所需 L（全部算力，低端）").replace("所需 L（不含 κ，高端）", "所需 L（全部算力，高端）")
        assert new != cur
        ws[f"{c}242"].value = new; n += 1
    A_OLD = "判讀：折現的世代起點 L 區間 對 不含 κ 的所需 L 區間（恆等式讀法 L1_LReqFleet_IdNoK；三態：未回本／未定／回本；v17.0 起 L 側讀 L1_LTotalFleetVDisc，F1／F2 依預設）"
    A_NEW = "判讀：折現的世代起點 L 區間 對 全部算力持有成本的所需 L 區間（L1_LReqFleet_Full；三態：未回本／未定／回本；v20.0 起所需 L 側讀 L1_LReqFleet_Full，L 側讀 L1_LTotalFleetVDisc）"
    if ws["A242"].value != A_NEW:
        assert ws["A242"].value == A_OLD, ws["A242"].value
        ws["A242"].value = A_NEW; n += 1
    old_head = "（v17.0：未定，L 0.759 對所需 L 0.756）"
    new_head = "（v20.0：未定，L 0.759 對所需 L 0.807，比值 0.940；v17.0：未定，L 0.759 對所需 L 0.756）"
    g = str(ws["G242"].value)
    if new_head not in g:
        assert old_head in g, g[:200]
        ws["G242"].value = g.replace(old_head, new_head, 1); n += 1
    n += int(bool(append_text(ws["G242"], "｜v20.0 W2：所需 L 側改讀 L1_LReqFleet_Full（全部算力持有成本 ÷ 2025 營收，區間 [_Lo, _Hi]）；L 側不變（L1_LTotalFleetVDisc）；三態規則同 v15.3 D1。"
                                         "基準比值（L ÷ 所需 L）＝0.940。v17.0 主口徑（L1_LReqFleet_IdNoK，只含推論、經計價倍數）保留為 Cap_In P 節對照列；Decisions SL55-01")))
    log["row242"] = n
    # L1 P89
    l1 = wb["L1"]
    assert l1["A89"].value == "L1_LReqFleet"
    rf = int(wb.defined_names["L1_LReqFleet_Full"].attr_text.split("$")[-1])
    cur = l1["P89"].value
    if f"D{rf}" not in cur:
        old115 = int(wb.defined_names["L1_LReqFleet_IdNoK"].attr_text.split("$")[-1])
        new = re.sub(rf"\b([DEF]){old115}\b", lambda m: f"{m.group(1)}{rf}", cur)
        new = new.replace("不含 κ 所需 L", "全部算力所需 L")
        mark = '"）（v16.0 口徑：L "'
        assert new.count(mark) == 2, new.count(mark)
        new = new.replace(mark, f'"）（v17.0 口徑對照：只含推論、經計價倍數的所需 L "&TEXT(D{old115},"0.00")&"）（v16.0 口徑：L "')
        assert new != cur
        l1["P89"].value = new; log["p89"] = 1
    for a, tail in (("H89", "｜v20.0 起：所需 L 側改讀 L1_LReqFleet_Full（全部算力持有成本 ÷ 2025 營收；主口徑）；L1_LReqFleet_IdNoK 為 v17.0 對照"),):
        append_text(l1[a], tail)
    return log


def text_fixes(wb):
    """W5: Cap_In G232／G233 numbers follow the current C–E values; L1 column T header lists the new class."""
    ws = wb["Cap_In"]
    n = 0
    for r in (232, 233):
        g = str(ws[f"G{r}"].value)
        m = re.search(r"（v17\.0：基準 [0-9.]+、D [0-9.]+、E [0-9.]+）", g)
        if m:
            cur = "（v20.0 建置時：基準 {:.4f}、D {:.4f}、E {:.4f}）"
            vals = {232: (0.7637, 0.5511, 1.1618), 233: (0.6595, 0.4019, 1.0964)}[r]      # LibreOffice values of C232:E233 in v19.1 (rounded); the cells C–E stay authoritative
            ws[f"G{r}"].value = g.replace(m.group(0), cur.format(*vals)); n += 1
    OLD3 = "(3) 2026 欄無區間：輸入（每日則數、消費份額 0.65）為 Assumed 單值，區間待補。"
    NEW3 = "(3) 區間（v20.0 已補）：2026 Q2 與 L1_M 的低／高只隨每則提示 token（AL_TokPerPrompt_Hi／_Lo）與免費占比（AL_FreeShare_Lo／_Hi）變動；每日則數、消費份額 0.65 仍為 Assumed 單值，2026 Q1 與 FY2026 欄仍無區間。"
    for nm in ("L1_M_2026Q2", "L1_M"):
        rr = int(wb.defined_names[nm].attr_text.split("$")[-1])
        cur = wb["L1"][f"I{rr}"].value
        if isinstance(cur, str) and OLD3 in cur:
            wb["L1"][f"I{rr}"].value = cur.replace(OLD3, NEW3); n += 1
    t4 = wb["L1"]["T4"].value
    if t4 and "觀測比值" not in t4:
        wb["L1"]["T4"].value = t4.replace("未解釋）", "未解釋／觀測比值）") if "未解釋）" in t4 else t4 + "（v20.0 增觀測比值）"; n += 1
    return n


# ---------------------------------------------------------------- Checks P1 (WARN): price basis of the token ledger
def checks_p1(wb):
    ws = wb["Checks"]
    r = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    from common import section
    section(ws, r, "P. GW 優先校準檢查（v20.0，工作單 W3）：P1 為 WARN，計入 GOV_Warnings（不阻擋）", 6); r += 1
    for i, h in enumerate(["編號", "檢查", "等級", "筆數", "範圍與算法"]):
        put(ws, f"{L(i + 1)}{r}", h, F_BOLD)
    r += 1
    f = ("=IF(AND(ISNUMBER(CAP_Price25Mix),ISNUMBER(AL_PriceFleet25),ISNUMBER(Cap_In!$E$205),ISNUMBER(INDEX(B4_PriceFleet,1,3))),"
         "IF(AND(AL_PriceFleet25>0,INDEX(B4_PriceFleet,1,3)>0),IF(AND(ABS(CAP_Price25Mix/AL_PriceFleet25-1)<=0.000000001,ABS(Cap_In!$E$205/INDEX(B4_PriceFleet,1,3)-1)<=0.000000001),0,1),1),0)")
    put(ws, f"A{r}", "P1"); put(ws, f"B{r}", "token 帳的價格基準不一致：|CAP_Price25Mix ÷ AL_PriceFleet25 − 1| > 1e-9，或 |Cap_In!E205 ÷ INDEX(B4_PriceFleet,1,3) − 1| > 1e-9：1＝不一致", wrap=True)
    put(ws, f"C{r}", "WARN"); put(ws, f"D{r}", f, F_CALC, fmt="0")
    put(ws, f"E{r}", "CAP_Price25Mix；AL_PriceFleet25；Cap_In!E205；B4_PriceFleet。L1_TokLedger 成立的前提是兩邊價格基準相同（2025 為 CAP_Price25Mix＝AL_PriceFleet25；2026 為 Cap_In E205＝B4_PriceFleet 第 3 欄）", F_NOTE)
    ref = wb.defined_names["GOV_Warnings"].attr_text.split("!")[1].replace("$", "")
    cell = ws[ref]
    if f"D{r}" not in str(cell.value):
        cell.value = f"{cell.value}+D{r}"
    lab = ws[f"B{int(ref[1:])}"]
    if "P1" not in str(lab.value):
        lab.value = (str(lab.value)[:-1] + "、P1）") if str(lab.value).endswith("）") else f"{lab.value}、P1"
    return r


# ---------------------------------------------------------------- DB_Evidence E401, Decisions SL55-01
EVIDENCE = [EID, DATE, "OpenAI Codex 每日 token（2025-10）：約 21 天 40 兆 → 約 2 兆／日",
            "OpenAI DevDay 2025-10-06（經 Epoch ai_companies_usage_reports.csv）；chat 端查證 2026-10-22（本代理依工作單，未外部搜尋）",
            "Interested-party（OpenAI 自揭）", "L1_DPaid_wCodex（對照列，不驅動）", "無", "2.0 T tok/日（SRC_DEM_088）", "採用（對照列；V20.0 W4）",
            V, "OpenAI 自揭使用量；2025-10 單一時點值；全為付費為假設；不進任何驅動格｜" + WO, "已處理", SRC_ID, DASH, "L1_DPaid_wCodex", 3, "利害關係方",
            "採用（對照列）｜v20.0 新增列"]


def evidence_append(wb):
    ws = wb["DB_Evidence"]
    if any(ws.cell(r, 1).value == EID for r in range(5, ws.max_row + 1)):
        return 0
    r = max(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")) + 1
    for i, v in enumerate(EVIDENCE):
        put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 4, 5, 6, 7, 8, 10, 14, 17))
    return 1


def _decisions():
    return [[DEC_ID, V, "v20.0 2025 校準改 GW 優先、回本改全部算力成本、m 區間與對照列",
             "W1：2025 校準先由 Tokenomics 的 GW 起算——AL_FleetGW25（SRC_DEM_008 → 009 指數內插，低＝後段集中 1.0 GW、高＝線性 1.25 GW）→ L1_ServeGW_GW25 → L1_TokPaid_GW25（Andy 07:34 公式）→ L1_TokLedger_GW25。"
             "W2：回本改以全部算力持有成本為口徑：L1_HoldFleet25、L1_PayRatio25、L1_RevPerGW25、L1_LReqFleet_Full；判讀（Cap_In 第 242 列、L1 第 89 列）的所需 L 側改讀 L1_LReqFleet_Full，v17.0 主口徑改為對照。"
             "W3：L1_M_2026Q2 區間；Checks P1（價格基準，WARN）；L1_M／L1_M_* 性質由殘差改觀測比值。W4：對照列 L1_DPaid_T4k／T12k／T30k／L1_DPaid_wCodex（SRC_DEM_088）、效率區間 AL_ProdDerate_Lo／_Hi、IF_ProdDerate_Lo／_Hi。"
             "OpenAI 特有資料留 Tokenomics 作『產業校準錨點』（Andy 16:11）。既有名稱的值全部不變（判讀文字列除外）。Andy 原話：" + DEC_BY,
             DATE, DEC_BY, "v20.0 已建", "Cap_In P 節；L1 第 134 列起；Cap_In 第 242 列；L1 第 89 列；Interface N 節；SRC_Demand SRC_DEM_088；DB_Evidence E401", f"{WO} W1–W5", "否"]]


def apply(wb, append_text):
    log = {}
    log["label_tails"] = label_tails(wb, append_text)
    log["judge"] = judge_update(wb, append_text)
    log["text_fixes"] = text_fixes(wb)
    log["purpose"] = purpose_marks(wb)
    log["l1_o_cleared"] = clear_l1_ratio(wb)
    log["checks_p1"] = checks_p1(wb)
    log["evidence"] = evidence_append(wb)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == DEC_ID for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


def purpose_marks(wb):
    ws = wb["Cap_In"]
    pc = v1400._purpose_col(ws, create=False)
    assert pc is not None, "用途 column of Cap_In not found"
    for rr, txt in PURPOSE_MARKS.items():
        put(ws, f"{L(pc)}{rr}", txt, F_NOTE)
    return len(PURPOSE_MARKS)


def clear_l1_ratio(wb):
    """L1 column O (推算 ÷ 外部) has no external value on the new rows: cleared like the other external-less rows, so they are not dead formula cells."""
    ws = wb["L1"]; n = 0
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value in {f"L1_{k}" for k in KEYS} and ws[f"O{r}"].value is not None:
            ws[f"O{r}"].value = None; n += 1
    return n


# ================================================================== README
README_VERSION = (VERSION + "（主版：" + WO + "——2025 校準改 GW 優先：AL_FleetGW25、L1_ServeGW_GW25、L1_TokPaid_GW25、L1_TokLedger_GW25；回本改全部算力持有成本：L1_HoldFleet25、L1_PayRatio25、L1_RevPerGW25、"
                  "L1_LReqFleet_Full（判讀主口徑，Cap_In 第 242 列與 L1 第 89 列的所需 L 側改讀它）；L1_M_2026Q2 區間；對照列 L1_DPaid_T4k／T12k／T30k／L1_DPaid_wCodex；AL_ProdDerate_Lo／_Hi、IF_ProdDerate_Lo／_Hi；"
                  "SRC_DEM_088；Checks P1；Decisions SL55-01；DB_Evidence E401。既有名稱的值全部不變（判讀文字列除外）。）以下為 "
                  + v1910.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1910.README_TITLE
README_TITLE = V + _T[len(v1910.V):]
README_ROW = ("v20.0 GW 優先校準",
              "2025 校準先由 Tokenomics 的 GW 起算（平均上線 GW → 服務 GW → 付費 token → 與需求面對照，L1_TokLedger_GW25）；OpenAI 特有資料（0.6→1.9 GW 等）留在 Tokenomics 作『產業校準錨點』。"
              "回本主口徑改為全部算力持有成本：L1_LReqFleet_Full（所需 L）；L1_PayRatio25＝實際 ÷ 所需。v17.0 主口徑 L1_LReqFleet_IdNoK（只含推論、支出路線）改為對照。"
              "細節見工作單 docs/workorders/20261022_v20.0_GW優先校準.md 與 Decisions SL55-01。")
