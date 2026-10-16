# v10.1 (work order docs/workorders/20261016_瘦身4.md section 1 = segment 4-1; engineering class, CLAUDE.md section 4a): slimming that must not change any output.
#   1-a  delete cells that nothing reads (Capability sheet, L1 column O without external value, Alloc diagnostic rows, Perf / Perf_Batch quick-look rows 104-105)
#   1-b  mark display-only / not-wired / copy-read rows (column "用途"), yellow fill on scenario cells, NonNV and Theory_Rev_Prod header notes
#   1-c  stale explanatory text (numbers that move with the model are now formulas), wrong names, README name counts
#   1-d  Gov_Map registry hygiene (no input value changes)
# Every step runs at the END of build.py (after Gov_Map / contract / Load_Bearing are final) and is idempotent: a rebuild from v10.1 gives the same cells.
# Nothing here changes a value that an IF_ / IFW_ / IFC_ / L1_ / GOV_ / CHK_ / CAP_ / AL_ name, Interface, Checks or the downstream contract reads.
import re
import openpyxl
from openpyxl.styles import PatternFill
from common import put, F_CALC, F_NOTE, F_BOLD
import v1000
import v1010_text
import v1010_gov

VERSION = "20261016_Tokenomics_v10.1"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-16"
V = "v10.1"
WO = "工作單 docs/workorders/20261016_瘦身4.md 第 1 節（4-1 段）"
DEC_BY = "CC 依工作單執行（工程類）；chat 端預設，可撤回"
NO_FILL = PatternFill(fill_type=None)

# ================================================================== 1-a: deletions (cells are emptied; no row or column is removed, so no reference moves)
# L1 column O (推算 ÷ 外部) rows whose external value is absent (62 formulas; 4-0 rule R06)
L1_O_ROWS = [5, 6, 7, 11, 12, 13, 14, 15, 18, 19, 20, 21, 22, 23, 24, 26, 27, 28, 29, 37, 38, 40, 41, 42, 43, 44, 45, 50, 51, 52, 53, 54, 55, 56, 57, 58, 67, 68, 69,
             71, 72, 73, 86, 87, 90, 91, 92, 94, 95, 96, 97, 98, 99, 100, 101, 102, 103, 104, 105, 106, 107, 111]
# Alloc diagnostic cells nobody reads (4-0 rule R18 minus the 27 cells that a name or another formula still reads: see report "待 Project 判斷")
ALLOC_DELETE = ["C74", "C75", "C76", "C122", "C144", "C145", "C152", "C163", "C164", "C165"]
# Perf / Perf_Batch quick-look rows 104-105 (unit price is defined once, in Unit_Cost rows 16-18; 4-0 rule R22c)
QUICKLOOK = [("Perf", 104, 105), ("Perf_Batch", 104, 105)]
QUICKLOOK_NOTE = "（v10.1 刪除速覽列 $/M：單價的唯一定義在 Unit_Cost 第 16–18 列；4-1 工程類，輸出不變）"


def _clear(ws, ref):
    c = ws[ref]
    c.value = None
    c.fill = NO_FILL


def delete_cells(wb):
    n = {}
    ws = wb["L1"]; k = 0
    for r in L1_O_ROWS:
        if ws[f"O{r}"].value is not None: _clear(ws, f"O{r}"); k += 1
    n["L1!O"] = k
    ws = wb["Alloc"]; k = 0
    for ref in ALLOC_DELETE:
        if ws[ref].value is not None: _clear(ws, ref); k += 1
    n["Alloc"] = k
    for sh, r0, r1 in QUICKLOOK:
        ws = wb[sh]; k = 0
        for r in range(r0, r1 + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                if cell.value is not None and not (r == r0 and c == 1 and cell.value == QUICKLOOK_NOTE):
                    cell.value = None; cell.fill = NO_FILL; k += 1
        put(ws, f"A{r0}", QUICKLOOK_NOTE, F_NOTE)
        n[sh] = k
    if "Capability" in wb.sheetnames:
        n["Capability"] = sum(1 for row in wb["Capability"].iter_rows() for c in row if c.value is not None)
        del wb["Capability"]
    return n


# ================================================================== Decisions (10 columns; layout of v800／v810／v1000)
_CAP_C = ("Capability 頁 C 節（K4 為何不採 (a)）原文摘要，頁面刪除後保留於此：1. 識別：同層級牌價跨廠商差距大（例：Sol 級輸出價 DeepSeek V4-Pro 離峰 $1.98 至 Kimi K3 $15），"
          "主要由廠商策略與成本決定，斜率無法與能力分離。2. 循環：前緣由接近成本定價的廠商決定（Checks：DeepSeek 價 ÷ 本模型 Hopper 同架構成本約 0.8–1.0），前緣斜率≈服務成本斜率，已由 Unit_Cost 計算。"
          "3. 重複：層級牌價已含能力；只有增量（H_CAP）可再進入，故以 ε 乘在增量上，基準 0。4. 誤差相乘：J8 算力缺口約 7 倍 × ε 0–0.5 → 單價倍數 1.0–2.6，大於利用率區間的影響。"
          "(d) 反解：Amortize 頁『回本所需單價溢價』與觀測層級價差並列。METR 時間長度（K3 檢查）尚未入表。J8 檢查（前沿錨點中值 ÷ Astra 最終訓練 FLOPs）可由 Train_In!C47 ÷ Training 頁 Astra 最終訓練 FLOPs 重算。")
_D = [
    ("1a-1", "Capability 頁刪除（已決定，6b）", "刪除 Capability 工作表（55 格，19 個公式，無任何格或名稱引用）；builder 的 block4.capability() 仍會建立，於 build.py 末端由 v1010.delete_cells 刪除。"
     "README 工作表清單與 Cap_In G14 的指向文字同步修改。" + _CAP_C, "Capability 頁；README 工作表清單；Cap_In G14", "v10.1 已建"),
    ("1a-2", "A 級死格刪除（只清空儲存格，不刪列欄，參照不移動）", "L1 O 欄無外部值的 62 格；Alloc 無人讀取的診斷格 10 格（C74–C76、C122、C144–C145、C152、C163–C165）；Perf、Perf_Batch 速覽列（第 104–105 列）64 格，A104 留一行說明。"
     "4-0 盤點 R18 共 37 格，其中 27 格仍有具名範圍（AL_Q1g、AL_Q2g、AL_RDGWg、AL_J8Gap、AL_Q1_Tok、AL_Q2_Tok、AL_RDShareExcl、AL_UnservedPaid／Free／Gen、AL_ServeGen）或另一公式（C142 讀 C47:E48）引用，"
     "刪除會改變既有具名範圍的值，故保留，列待 Project 判斷。", "L1!O；Alloc；Perf／Perf_Batch 第 104–105 列", "v10.1 已建"),
    ("1a-3", "Perf_Batch R–S 欄（FP8 rollout 情境）未刪", "覆核發現 R–S 欄並非未接線：Sens_Train W59:X62（情境『rollout 精度 FP8』，黃底）讀 Perf_Batch R101、S101、R103、S103、R65、S65、R67、S67，這八格依賴 R–S 欄其餘約 120 格。"
     "刪除會使 Sens_Train W–X 兩欄的值改變，屬輸出改變，本段（輸出不變）不執行；改列待 Project 判斷：連 Sens_Train W–X 一併刪、或保留。第 104–105 列的 R–S 兩欄速覽已隨 1a-2 清空。", "Perf_Batch R–S；Sens_Train W–X", "未執行，待 Project 判斷"),
]
DECISIONS_V1010 = [[f"SL41-{i}", V, f"瘦身 4-1：{t}", txt, DATE, DEC_BY, st, where, f"{WO} {i}", "否"] for i, t, txt, where, st in _D]


def decisions_append(wb, extra=()):
    ws = wb["Decisions"]
    have = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
    r = max([rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")] or [4]) + 1
    n = 0
    for row in DECISIONS_V1010 + list(extra):
        if row[0] in have: continue
        for i, v in enumerate(row): put(ws, f"{openpyxl.utils.get_column_letter(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 5, 7))
        r += 1; n += 1
    return n


# ================================================================== README
README_VERSION = (VERSION + "（次版，工程類：瘦身 4-1，" + WO + "——1-a 刪除無人讀取的格（Capability 頁、L1 O 欄、Alloc 診斷格、Perf 速覽列）；1-b 標記顯示用／未接線／複本讀取、補黃底；"
                  "1-c 過時說明文字改引用現值；1-d Gov_Map 登錄整理；1-e 新增永久測試。所有既有具名範圍與契約名稱的值不變）。以下為 " + v1000.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1000.README_TITLE
README_TITLE = VERSION.split("_")[-1] + _T[len(v1000.VERSION.split("_")[-1]):]
README_ROW = ("瘦身 4-1（v10.1）",
              "工程類，輸出不變：刪除 Capability 頁與無人讀取的格；『用途』欄標記顯示用／未接線／複本讀取；情境格黃底；過時說明改引用現值；Gov_Map 登錄整理。細節見 Decisions SL41-*。")


# ================================================================== 1-b: marks (column "用途"), yellow fill, header notes
from v1010_marks import MARKS
from v518 import _append_text

PURPOSE = "用途"
PURPOSE_HDR_ROW = 3          # the header goes in row 3 of the new column (rows 1-2 are title and subtitle on every model page)
PURPOSE_NOTE = "顯示用＝人讀的對照／欄頭／驗證列，不進任何輸出；未接線＝尚未連到任何輸出；複本讀取＝只被 Sens_／_Prod 複本頁讀取（v10.1 4-1 標記，來源：docs/reports/20261016_瘦身盤點4-0_明細.xlsx）"
YEL_FORMULA = PatternFill("solid", fgColor="FFFFF2CC")     # existing convention on the sensitivity pages for scenario-changed formula cells
YEL_INPUT = PatternFill("solid", fgColor="FFFFFF00")       # existing convention for scenario input cells (blue font)
# cells whose formula differs from the first scenario column of the same row (4-0 section 3; the permanent test keeps this list honest)
YELLOW_FORMULA = {
    "Sensitivity": ["F10", "G10", "T14", "U14", "D17", "E17", "H19", "I19", "J25", "K25", "L26", "M26", "O37", "V40", "W40", "R41", "S41", "P50", "Q50"],
    "Sens_Rev": ["B20", "B21"],
    "Sens_Har": ["D5", "E5", "F6", "G6", "H7", "I7", "J8", "K8", "L9", "M9", "N10", "O10", "P11", "Q11", "R12", "S12", "R13", "S13", "R14", "S14"],
}
YELLOW_INPUT = {
    "Sensitivity": ["X8", "Y8", "N37"],
    "Sens_Rev": ["B6", "B7", "B8", "B9", "B10", "B11", "B12", "B13", "B14", "B15", "B17"],      # v10.3 D08: B20／B21 are links to Cap_In D23／D24 now (formula fill below)
}
NONNV_NOTE = " ｜未接線（MAG、智譜、AI 半導體待用；目前不進任何輸出）"
TRP_NOTE = (" ｜第 90–96 列為本頁自我檢查列，不鏡像 Theory_Rev 的 _Life 列（Theory_Rev 同列為 _Life 輸出列）；列入 tests/parity 的鏡像白名單（v10.1）。")


def purpose_marks(wb):
    """Column 用途: one column per page, right of the last used column; the header is found again on a rebuild (idempotent)."""
    n = 0; pages = {}
    for sh, rows in MARKS.items():
        ws = wb[sh]
        col = None
        last = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None: last = max(last, cell.column)
        for c in range(3, last + 1):
            if ws.cell(PURPOSE_HDR_ROW, c).value == PURPOSE:
                col = c; break
        if col is None:
            col = last + 1
            put(ws, f"{openpyxl.utils.get_column_letter(col)}{PURPOSE_HDR_ROW}", PURPOSE, F_BOLD)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 46
        X = openpyxl.utils.get_column_letter(col)
        for r, txt in rows.items():
            put(ws, f"{X}{r}", txt, F_NOTE); n += 1
        pages[sh] = X
    return n, pages


def yellow_fill(wb):
    n = 0
    for sh, refs in YELLOW_FORMULA.items():
        for ref in refs:
            wb[sh][ref].fill = YEL_FORMULA; n += 1
    for sh, refs in YELLOW_INPUT.items():
        for ref in refs:
            wb[sh][ref].fill = YEL_INPUT; n += 1
    return n


def header_notes(wb):
    n = 0
    ws = wb["NonNV"]
    if NONNV_NOTE not in str(ws["A1"].value): ws["A1"].value = str(ws["A1"].value) + NONNV_NOTE; n += 1
    ws = wb["Theory_Rev_Prod"]
    if TRP_NOTE not in str(ws["A2"].value): ws["A2"].value = str(ws["A2"].value) + TRP_NOTE; n += 1
    return n


# ================================================================== entry point (called once from build.py, after everything else and before the sheet-order assertion)
def apply(wb):
    log = {}
    log["deleted"] = delete_cells(wb)
    log["marks"] = purpose_marks(wb)
    log["yellow"] = yellow_fill(wb)
    log["header_notes"] = header_notes(wb)
    done, same, drift = v1010_text.apply_fixes(wb)
    log["text_fixes"] = dict(changed=done, already=same, drift=drift)
    log["interface_labels"] = v1010_text.interface_labels(wb)
    log["harness_labels"] = v1010_text.harness_labels(wb)
    log["src_cites"] = v1010_text.src_cites(wb)
    log["checks_header"] = v1010_text.checks_header(wb)
    log["gov"] = v1010_gov.apply(wb)
    log["decisions"] = decisions_append(wb)
    return log
