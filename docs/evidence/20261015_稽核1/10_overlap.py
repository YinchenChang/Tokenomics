"""第 1 段開工第一步：m 路徑輸入格 × 第 0 段 C1.2（SRC_ID 與決策 ID 皆為「—」）311 列的交叉。
用法：python3 10_overlap.py <稽核0明細.xlsx> <out.csv>"""
import openpyxl, sys, csv, collections
wb=openpyxl.load_workbook(sys.argv[1],read_only=True)
c12={}
for r in wb['C1_輸入治理'].iter_rows(min_row=3,values_only=True):
    if r[1] and r[1].startswith('C1.2 SRC_ID'): c12[r[4]]=r
inp=list(wb['C7_第1段輸入清單'].iter_rows(min_row=3,values_only=True))
m=[r for r in inp if r[16]=='是']
ov=[r for r in m if r[0] in c12]
gms={r[0] for r in ov}
print('C1.2 rows',len(c12),'m-path cells',len(m),'overlap cells',len(ov),'overlap GM rows',len(gms))
print(collections.Counter(r[1] for r in ov)); print(collections.Counter(r[7] for r in ov))
with open(sys.argv[2],'w',newline='') as f:
    w=csv.writer(f); w.writerow(['GM_ID','頁','格','列標籤','單位','目前值','標記','被模型公式引用次數'])
    for r in ov: w.writerow([r[0],r[1],r[2],r[3],r[4],r[5],r[7],r[13]])
