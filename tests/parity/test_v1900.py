"""v19.0 Interface 新增 N 節「P×Q 取數」（工作單 docs/workorders/20261022_v19.0_Interface取數.md r0）專屬測試。讀 LibreOffice 重算後的值。

W3 驗收：
  1. 新名稱基準／低／高等於工作單事前值（相對差 ≤ 1e-6；工作單只印有限位數，故以「印出位數的半個單位」為下限）。
  2. IF_TokGWPaidMix 各欄與由 IF_TokGW_*、B4_MixPaid 獨立重算一致（≤ 1e-12）。
  3. 情境 v1800_disc30 下 IF_Disc＝0.30（Lo／Hi 不變）。
  4. 新列在 Tech_Registry!O11:O13＝1（f_registry_t07_t09_on）與 B4_CacheHit＝0（d9_cache_cap_off）情境下無錯誤值（v18.0 曾因此在全套 parity 失敗）。
  5. 逐格比較（對封存的 v18.0）只落在預期變動；既有具名範圍的值不變（版本字串除外）。
"""
import re
from pathlib import Path

import openpyxl
import pytest
import yaml
from openpyxl.utils import range_boundaries

from conftest import model_path

REPO = Path(__file__).resolve().parents[2]
V180 = REPO / "model" / "archive" / "20261021_Tokenomics_v18.0.xlsx"
V190 = REPO / "model" / "archive" / "20261022_Tokenomics_v19.0.xlsx"      # v19.1：現行版已是 v19.1；v18.0→v19.0 的靜態檢查改對封存的 v19.0 釘住（v19.0→v19.1 由 test_v1910 檢查）
SC = {s["id"]: s["inputs"] for s in yaml.safe_load((Path(__file__).resolve().parent / "scenarios.yaml").read_text(encoding="utf-8"))["scenarios"]}
DISC30 = SC["v1800_disc30"]
F_REGISTRY = SC["f_registry_t07_t09_on"]
CACHE_OFF = SC["d9_cache_cap_off"]
VERSION_NAMES = {"RNG_Version", "RNG_Current"}
IFC_NAMES = {"IFC_Conf", "IFC_Use", "IFC_Layer", "IFC_Updated"}      # Interface 契約欄範圍隨新列變長

NEW_NAMES = (["IF_Disc", "IF_Disc_Lo", "IF_Disc_Hi"] + [f"IF_MixPaid_{t}" for t in ("Luna", "Sol", "Astra")] + [f"IF_MixFree_{t}" for t in ("Luna", "Sol", "Astra")]
             + ["IF_UtilActual", "IF_UtilActual_Lo", "IF_UtilActual_Hi", "IF_ProdDerate", "IF_FleetProdRatio_VR200", "IF_PxEffMix",
                "IF_PxEffMix25", "IF_PxEffMix25_Lo", "IF_PxEffMix25_Hi", "IF_LTotalDisc", "IF_LTotalDisc_Lo", "IF_LTotalDisc_Hi",
                "IF_LTotalFleetVDisc", "IF_LTotalFleetVDisc_Lo", "IF_LTotalFleetVDisc_Hi", "IF_TokGWPaidMix"])
NEW = set(NEW_NAMES)

PRE = {   # 工作單事前值（名稱 -> 純量值）
    "IF_Disc": 0.20, "IF_Disc_Lo": 0.10, "IF_Disc_Hi": 0.40,
    "IF_MixPaid_Luna": 0.40, "IF_MixPaid_Sol": 0.45, "IF_MixPaid_Astra": 0.15,
    "IF_MixFree_Luna": 0.90, "IF_MixFree_Sol": 0.10, "IF_MixFree_Astra": 0,
    "IF_UtilActual": 0.41, "IF_UtilActual_Lo": 0.30, "IF_UtilActual_Hi": 0.65,
    "IF_ProdDerate": 0.85, "IF_FleetProdRatio_VR200": 0.80582, "IF_PxEffMix": 1.96068,
    "IF_PxEffMix25": 1.35936, "IF_PxEffMix25_Lo": 1.01933, "IF_PxEffMix25_Hi": 1.64651,
    "IF_LTotalDisc": 0.69884, "IF_LTotalDisc_Lo": 0.44008, "IF_LTotalDisc_Hi": 1.09414,
    "IF_LTotalFleetVDisc": 0.75860, "IF_LTotalFleetVDisc_Lo": 0.47298, "IF_LTotalFleetVDisc_Hi": 1.13827,
}
PRE_PAID_MIX = [25_658_985_091, 126_261_198_140, 154_004_682_276, 253_258_912_516, 351_710_375_478]      # 基準欄：Hopper／GB200／GB300／VR200／Rubin Ultra（各世代的 3 個成本情境欄）

# 掛在 f_registry／B4_CacheHit＝0 情境上，供 lo fixture 預先重算
PREWARM = {"test_v1900_new_rows_no_error_values_in_stress_scenarios": [F_REGISTRY, CACHE_OFF]}


def _flat(v):
    if isinstance(v, (list, tuple)):
        return [x for r in v for x in _flat(r)]
    return [v]


def _decimals(x):
    s = repr(float(x))
    return len(s.split(".")[1].rstrip("0")) if "." in s else 0


def _tol(want):
    """相對容差：1e-6，不小於「印出最後一位的半個單位」。"""
    if want == 0:
        return 0.0
    return max(1e-6, 0.5 * 10 ** (-_decimals(want)) / abs(want))


# v19.1：test_v1900_new_values_match_precomputed（v19.0 事前值）已停用——Arch!C17 改值後 Luna 相關值變動，v19.1 的新值由 test_v1910 釘住；v19.0 的值留在封存檔 model/archive/20261022_Tokenomics_v19.0.xlsx。


def test_v1900_new_names_equal_existing_cells(lo):
    """每個新名稱只是既有格的引用：與來源名稱逐值相等。"""
    v = lo.view({})
    g = v.get_name
    eq = [("IF_Disc", "B4_Disc"), ("IF_Disc_Lo", "B4_Disc_Lo"), ("IF_Disc_Hi", "B4_Disc_Hi"), ("IF_UtilActual", "AL_UtilActual"), ("IF_UtilActual_Lo", "AL_UtilActual_Lo"),
          ("IF_UtilActual_Hi", "AL_UtilActual_Hi"), ("IF_ProdDerate", "CTL_ProdDerate"), ("IF_FleetProdRatio_VR200", "L1_FleetProdRatio"), ("IF_PxEffMix25", "CAP_Price25Mix"),
          ("IF_LTotalDisc", "L1_LTotalDisc"), ("IF_LTotalDisc_Lo", "L1_LTotalDisc_Lo"), ("IF_LTotalDisc_Hi", "L1_LTotalDisc_Hi"),
          ("IF_LTotalFleetVDisc", "L1_LTotalFleetVDisc"), ("IF_LTotalFleetVDisc_Lo", "L1_LTotalFleetVDisc_Lo"), ("IF_LTotalFleetVDisc_Hi", "L1_LTotalFleetVDisc_Hi")]
    for a, b in eq:
        assert g(a) == _flat(g(b))[0], (a, b)
    for t, k in (("Luna", 0), ("Sol", 1), ("Astra", 2)):
        assert g(f"IF_MixPaid_{t}") == _flat(g("B4_MixPaid"))[k] and g(f"IF_MixFree_{t}") == _flat(g("B4_MixFree"))[k], t
    assert g("IF_PxEffMix") == v.get("Cap_In", "C90")
    assert g("IF_PxEffMix25_Lo") == v.get("Cap_In", "D281") and g("IF_PxEffMix25_Hi") == v.get("Cap_In", "E281")


def test_v1900_paid_mix_independent_recomputation(lo):
    """IF_TokGWPaidMix 各欄 ＝ 1 ÷ Σ（B4_MixPaid ÷ IF_TokGW_*）（1e-12）。VR200 基準欄（第 11 欄）並等於 L1_TokGW_PaidCap_VR200。"""
    v = lo.view({})
    mix = _flat(v.get_name("B4_MixPaid"))
    tg = [_flat(v.get_name(f"IF_TokGW_{t}")) for t in ("Luna", "Sol", "Astra")]
    pm = _flat(v.get_name("IF_TokGWPaidMix"))
    assert len(pm) == 15 and all(len(x) == 15 for x in tg)
    for k in range(15):
        want = 1 / sum(m / x[k] for m, x in zip(mix, tg))
        assert abs(pm[k] / want - 1) <= 1e-12, (k, pm[k], want)
    assert abs(pm[10] / v.get_name("L1_TokGW_PaidCap_VR200") - 1) <= 1e-12


def test_v1900_disc_scenario(lo):
    """情境 v1800_disc30（B4_Disc＝0.30）：IF_Disc＝0.30，Lo／Hi 不變；有效牌價同比例下降。"""
    b, s = lo.view({}), lo.view(DISC30)
    assert (b.get_name("IF_Disc"), s.get_name("IF_Disc")) == (0.2, 0.3)
    assert (s.get_name("IF_Disc_Lo"), s.get_name("IF_Disc_Hi")) == (0.1, 0.4)
    assert abs(s.get_name("IF_PxEffMix") / b.get_name("IF_PxEffMix") - (1 - 0.3) / (1 - 0.2)) <= 1e-9
    assert abs(s.get_name("IF_PxEffMix25") / b.get_name("IF_PxEffMix25") - (1 - 0.3) / (1 - 0.2)) <= 1e-9
    assert _flat(s.get_name("IF_TokGWPaidMix")) == _flat(b.get_name("IF_TokGWPaidMix"))      # 產能與折扣無關
    assert s.get_name("GOV_Errors") == 0


def _section_n_cells(v):
    names = {n: v.names[n] for n in NEW}
    for n in NEW:
        sh, rg = v.name_ref(n)
        c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                yield n, sh, r, c


def test_v1900_new_rows_no_error_values_in_stress_scenarios(lo):
    """f_registry（Tech_Registry!O11:O13＝1）與 B4_CacheHit＝0 情境：新列每一格都不是錯誤值（v18.0 曾因整代不可達的文字相除在全套 parity 失敗）；
    不可達時回傳 AL_UnavailTxt（文字）或數值，不得是 #VALUE!／#DIV/0!。"""
    for label, inputs in (("基準", {}), ("f_registry", F_REGISTRY), ("B4_CacheHit=0", CACHE_OFF)):
        v = lo.view(inputs)
        unavail = v.get_name("AL_UnavailTxt")
        n = 0
        for name, sh, r, c in _section_n_cells(v):
            from openpyxl.utils import get_column_letter as L
            x = v.get(sh, f"{L(c)}{r}")
            n += 1
            assert not (isinstance(x, str) and x.startswith("#")), (label, name, f"{L(c)}{r}", x)
            assert isinstance(x, (int, float)) or x == unavail, (label, name, f"{L(c)}{r}", x)
        assert n >= 24 + 14
        # 整個 Interface 頁與整本活頁簿也不得有錯誤值
        errs = [(s, c) for (s, c), x in v.evaluate_all().items() if isinstance(x, str) and x.startswith("#")]
        assert errs == [], (label, errs[:5])


def test_v1900_paid_mix_guard_returns_unavail_text(lo):
    """防護：任一層級占比或層級產能非數字（整代不可達）時回傳 AL_UnavailTxt。這裡以公式文字檢查，並在 B4_CacheHit＝0 情境確認各欄為數值或該文字。"""
    wb = openpyxl.load_workbook(model_path())
    ws = wb["Interface"]
    r = next(x for x in range(340, ws.max_row + 1) if "[IF_TokGWPaidMix]" in str(ws.cell(x, 1).value))
    for col in range(3, 18):
        f = ws.cell(r, col).value
        assert f.count("ISNUMBER") == 6 and "AL_UnavailTxt" in f and f.startswith("=IF(AND(") and ">0" in f, (col, f)
    v = lo.view(CACHE_OFF)
    for x in _flat(v.get_name("IF_TokGWPaidMix")):
        assert isinstance(x, (int, float)) or x == v.get_name("AL_UnavailTxt"), x


# ---------------------------------------------------------------- static (formula mode / archive comparison)
@pytest.fixture(scope="module")
def wbf():
    return openpyxl.load_workbook(V190)


def _row(ws, rid, col=1):
    rows = [r for r in range(5, ws.max_row + 1) if ws.cell(r, col).value == rid]
    assert len(rows) == 1, (rid, rows)
    return rows[0]


def test_v1900_interface_section_formulas(wbf):
    ws = wbf["Interface"]
    sec = [r for r in range(340, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("N. P×Q 取數")]
    assert len(sec) == 1
    r0 = sec[0]
    tags = [re.search(r"\[(IF_\w+)\]\s*$", str(ws.cell(r, 1).value or "")) for r in range(r0 + 1, ws.max_row + 1)]
    assert [t.group(1) for t in tags if t] == ["IF_Disc", "IF_MixPaid_Luna", "IF_MixPaid_Sol", "IF_MixPaid_Astra", "IF_MixFree_Luna", "IF_MixFree_Sol", "IF_MixFree_Astra",
                                                 "IF_UtilActual", "IF_ProdDerate", "IF_FleetProdRatio_VR200", "IF_PxEffMix", "IF_PxEffMix25", "IF_LTotalDisc", "IF_LTotalFleetVDisc", "IF_TokGWPaidMix"]
    row = {n: next(r for r in range(r0 + 1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").endswith(f"[{n}]")) for n in
           ("IF_Disc", "IF_MixPaid_Luna", "IF_MixPaid_Sol", "IF_MixPaid_Astra", "IF_MixFree_Luna", "IF_MixFree_Sol", "IF_MixFree_Astra", "IF_UtilActual", "IF_ProdDerate",
            "IF_FleetProdRatio_VR200", "IF_PxEffMix", "IF_PxEffMix25", "IF_LTotalDisc", "IF_LTotalFleetVDisc", "IF_TokGWPaidMix")}
    got = lambda n, c: ws[f"{c}{row[n]}"].value
    assert [got("IF_Disc", c) for c in "CDE"] == ["=B4_Disc", "=B4_Disc_Lo", "=B4_Disc_Hi"]
    for k, t in enumerate(("Luna", "Sol", "Astra"), start=1):
        assert got(f"IF_MixPaid_{t}", "C") == f"=INDEX(B4_MixPaid,1,{k})" and got(f"IF_MixFree_{t}", "C") == f"=INDEX(B4_MixFree,1,{k})"
        assert got(f"IF_MixPaid_{t}", "D") is None and got(f"IF_MixFree_{t}", "E") is None
    assert [got("IF_UtilActual", c) for c in "CDE"] == ["=AL_UtilActual", "=AL_UtilActual_Lo", "=AL_UtilActual_Hi"]
    assert got("IF_ProdDerate", "C") == "=CTL_ProdDerate" and got("IF_FleetProdRatio_VR200", "C") == "=L1_FleetProdRatio"
    assert got("IF_PxEffMix", "C") == "=Cap_In!$C$90" and got("IF_PxEffMix", "D") is None
    assert [got("IF_PxEffMix25", c) for c in "CDE"] == ["=CAP_Price25Mix", "=Cap_In!$D$281", "=Cap_In!$E$281"]
    assert [got("IF_LTotalDisc", c) for c in "CDE"] == ["=L1_LTotalDisc", "=L1_LTotalDisc_Lo", "=L1_LTotalDisc_Hi"]
    assert [got("IF_LTotalFleetVDisc", c) for c in "CDE"] == ["=L1_LTotalFleetVDisc", "=L1_LTotalFleetVDisc_Lo", "=L1_LTotalFleetVDisc_Hi"]
    # 名稱指向
    for n in NEW - {"IF_TokGWPaidMix"}:
        txt = wbf.defined_names[n].attr_text
        base = n[:-3] if n.endswith(("_Lo", "_Hi")) and n[:-3] in row else n
        col = {"": "C", "_Lo": "D", "_Hi": "E"}[n[len(base):]]
        assert txt == f"Interface!${col}${row[base]}", (n, txt)
    assert wbf.defined_names["IF_TokGWPaidMix"].attr_text == f"Interface!$C${row['IF_TokGWPaidMix']}:$Q${row['IF_TokGWPaidMix']}"
    # 每個標籤寫出讀法（A 欄含 [名稱]，且說明單位／口徑）
    assert "實際＝" in ws[f"A{row['IF_TokGWPaidMix']}"].value and "IF_ServeShare" in ws[f"A{row['IF_TokGWPaidMix']}"].value
    assert "表列牌價 ×（1 − IF_Disc）" in ws[f"A{row['IF_Disc']}"].value
    assert "表列 ×（1 − IF_Disc）" in ws[f"A{row['IF_PxEffMix']}"].value


def test_v1900_decisions_readme(wbf):
    d = wbf["Decisions"]
    rr = _row(d, "SL53-01")
    assert d.cell(rr, 2).value == "v19.0" and d.cell(rr, 7).value == "v19.0 已建" and d.cell(rr, 10).value == "否"
    assert "05:19" in d.cell(rr, 6).value and "三 G 瓦其實應該是擺在 OpenAI 他們自己裡面" in d.cell(rr, 6).value
    assert "公司層級計算放 company-models" in d.cell(rr, 3).value and "IF_TokGWPaidMix" in d.cell(rr, 4).value and "N" in d.cell(rr, 4).value
    assert d.max_row < 1000      # Gov_Map AA 的 COUNTIF 範圍上限
    rd = wbf["README"]
    assert str(rd["B5"].value).startswith("20261022_Tokenomics_v19.0（主版：") and "v18.0（主版：" in str(rd["B5"].value) and str(rd["A1"].value).startswith("Tokenomics v19.0")
    rows = [x for x in range(4, rd.max_row + 1) if rd.cell(x, 1).value == "v19.0 Interface 取數"]
    assert len(rows) == 1 and "IF_TokGWPaidMix" in rd.cell(rows[0], 2).value and "N 節" in rd.cell(rows[0], 2).value


def _diffs(data_only):
    a, b = openpyxl.load_workbook(V180, data_only=data_only), openpyxl.load_workbook(V190, data_only=data_only)
    assert a.sheetnames == b.sheetnames
    out = []
    for ws in a:
        wn = b[ws.title]
        for r in range(1, max(ws.max_row, wn.max_row) + 1):
            for c in range(1, max(ws.max_column, wn.max_column) + 1):
                x, y = ws.cell(r, c).value, wn.cell(r, c).value
                if x != y and not (isinstance(x, (int, float)) and isinstance(y, (int, float)) and abs(x - y) <= 1e-12 * max(1, abs(x))):
                    out.append((ws.title, ws.cell(r, c).coordinate, x, y))
    return out, a, b


def _allowed(sh, coord, first_new_row):
    col, row = re.match(r"([A-Z]+)(\d+)", coord).groups(); row = int(row)
    if sh == "Interface":
        return row >= first_new_row                       # N 節新列（含契約欄 R–Y）
    if sh == "README":
        return coord in ("A1", "B5", "A66", "B66")
    if sh == "Alloc":
        return coord in ("C177", "C178")                  # RNG_Version／RNG_Current（版本字串）
    if sh == "Checks":
        return coord in ("D118", "E118", "D119", "E119", "D120", "E120", "D121", "E121")      # Interface 契約欄範圍 R6:R344 → R6:R361（Checks K 節）
    return sh == "Decisions"


def test_v1900_diffs_only_expected_vs_v180():
    for mode in (False, True):
        d, a, b = _diffs(mode)
        first_new = min(r for r in range(300, b["Interface"].max_row + 1) if str(b["Interface"].cell(r, 1).value or "").startswith("N. P×Q 取數"))
        assert first_new == 346 and all(c.value is None for row in a["Interface"].iter_rows(min_row=first_new - 1) for c in row)      # v18.0 的 Interface 到第 344 列為止；第 345 列空一列
        bad = [t for t in d if not _allowed(t[0], t[1], first_new)]
        assert bad == [], bad[:10]
        assert all(t[2] is None for t in d if t[0] in ("Decisions",) or (t[0] == "Interface"))      # 只新增，不改既有格
        chk = [t for t in d if t[0] == "Checks"]
        assert all(str(t[2]).replace("344", "361") == str(t[3]) for t in chk), chk      # 只有 Interface 契約欄範圍 R6:R344 → R361
        assert len({re.sub(r"[A-Z]+", "", t[1]) for t in d if t[0] == "Decisions"}) == 1      # Decisions 只新增一列
    names_a, names_b = set(a.defined_names), set(b.defined_names)
    assert names_a <= names_b and names_b - names_a == NEW
    changed = {k for k in names_a if a.defined_names[k].attr_text != b.defined_names[k].attr_text}
    assert changed == IFC_NAMES, changed      # 只有 IFC_* 四個範圍隨新列延伸（$344 → $361）
    # 數值：既有數字格不變（新列在 Interface 列號 ≥ first_new）
    dv, _, _ = _diffs(True)
    num = [t for t in dv if isinstance(t[2], (int, float)) or isinstance(t[3], (int, float))]
    assert all(t[0] == "Interface" for t in num), [t for t in num if t[0] != "Interface"][:5]


# v19.1：test_v1900_existing_names_unchanged（對 v18.0 逐名稱比較現行版）已停用；v18.0→v19.0 的名稱檢查由 test_v1900_diffs_only_expected_vs_v180（靜態，對封存 v19.0）承擔。
