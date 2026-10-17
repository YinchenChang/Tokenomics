# v13.1 (work order docs/workorders/20261017_D15_B段.md, branch claude/d15-b; P1 source completion + P2 findings of check #96).
#   P1  D15 (docs/reports/20261017_D15_來源建議_明細.xlsx) written as recommended where "是否需 Andy 判斷" = 否, plus Andy's five defaults of 2026-10-17 05:36:
#       - Gov_Map J (decision ID) and H (SRC_ID / relation I) for the 312 rows that had neither (data file builder/v1310_d15.json, produced by tools/gen_d15_json.py);
#         J holds exactly ONE decision ID (check AA counts the whole cell), the other IDs and the recommendation text are appended to N.
#       - Decisions D15-01…D15-12 (D15-01…09 from the D15 proposal; D15-10…12 record Arch E6–E8 mark correction, SRC_PERF_060 "待回查", the prompt-token note).
#       - New sources (only added, nothing existing is changed except the empty second-source field R of SRC_HW_067／068): SRC_DC_018／019 (PUE), SRC_PRC_045／046 (Batch half price), SRC_HW_069／070 (Lenovo LP1944).
#       - DB_Evidence: new column R "D15 B 段處置" with the recommended handling of the 106 backlog rows (columns I／L are NOT changed: CLAUDE.md section 6); new rows E381–E386.
#       - Arch F6:F8 "Verified／Assumed" -> "Assumed" (Gov_Map says Assumed): CLAUDE.md section 6 correction exception (revision 9): original wording kept in the same row (G), Decisions D15-10.
#       - SRC_PERF_060 note "待回查", Checks V1 note, L1_GapPrompt note on the unexplained closure value, Gov_Map GM587 note. No input value is changed.
#   P2  (1) VR200 columns of the range block are point values (no evidence for the VR200 batch multiplier): note in Alloc O and the contract.
#       (2) Range freshness fingerprint: Alloc P section lists every numeric constant upstream of the 53 ranged names (builder/v1310_fingerprint.json, tools/gen_fingerprint.py) with its recorded value;
#           RNG_StaleFlag and Checks RG3 (WARN, counted once together with RG2) turn on when any of them differs although the base may still lie inside Lo／Hi. No new defined name (so no contract change).
#       (3) text corrections of the v13.0 report (docs only).  (4) Gov_Map GM346 column O sentence rewritten under revision 9.
# Minor version: no new contract name, no value of an existing IF_／IFW_／L1_／CHK_／RNG_ name changes. Every write to an Excel-owned cell is guarded (old value / old text).
import json
from pathlib import Path

from openpyxl.utils import get_column_letter as L
from common import put, F_IN, F_NOTE, F_BOLD, F_CALC, F_LINK, FILL_KEY, section
import v1000
import v1300

VERSION = "20261017_Tokenomics_v13.1"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-17"
V = "v13.1"
WO = "工作單 docs/workorders/20261017_D15_B段.md"
DEC_BY = "Andy 2026-10-17 05:36（C68 選 (a)；D15 的 5 項判斷照預設）；D15-01…09 為 chat 端 2026-10-17 05:40 定案（工作單），可撤回"
EID = "E381"
DASH = "—"
HERE = Path(__file__).resolve().parent
D15 = json.loads((HERE / "v1310_d15.json").read_text(encoding="utf-8"))
FP_JSON = HERE / "v1310_fingerprint.json"


def load_fp():
    if not FP_JSON.exists():
        return {"version": "", "roots": 0, "leaves": []}
    return json.loads(FP_JSON.read_text(encoding="utf-8"))


FP = load_fp()
REPORT = "docs/reports/20261017_D15_來源建議.md"
EVTAB = "docs/evidence/20261017_D15/證據表.xlsx"


# ================================================================== new source records
def _rec(sheet, i, ev, metric, val, unit, basis, applies, src, grade, stance, stance_note, hand, use, note, second=DASH):
    r = v1000._rec(i, sheet, metric, val, None, None, unit, basis, applies, DATE, src, grade, stance, stance_note, hand, use, note)
    r["ev"] = ev; r["s"] = "U-V1310"
    if second != DASH: r["second"] = second
    return r


SRC_NEW = [
    _rec("SRC_DC", "SRC_DC_018", "E381", "Google 資料中心全機隊年均 PUE（2025 全年）", 1.09, "PUE",
         "全機隊年均（含非 AI 機房），營運商自報；2026 Q2 TTM 亦為 1.09", "Google 資料中心（超大規模全機隊）", "https://datacenters.google/efficiency（WebFetch 摘要 2026-10-17 已讀）", 1, "利害關係方",
         "Google：營運商自報效率，有展示效率的誘因；全機隊≠液冷 AI 設施", "一手（已讀摘要）",
         "Inputs D9:F9（PUE）的可比對象，Gov_Map GM004–006 的 H 欄；不直接連結",
         f"v13.1 D15 B 段新增（E381）：標記 Analogy（全機隊→液冷 AI 設施）；第二來源 SRC_DC_019；PUE 數值不改（1.1／1.2／1.3）", second="SRC_DC_019"),
    _rec("SRC_DC", "SRC_DC_019", "E381", "Uptime Institute 全球資料中心平均 PUE（2025 年度調查）", 1.54, "PUE",
         "年度調查，營運商自填；超大規模 1.10–1.15、colo／企業 1.58–1.80（mgrid 轉述，原文標『待原文確認』）", "全球資料中心（調查樣本）",
         "https://mgrid.org/2025/10/01/uptime-institute-data-center-pue-stagnation-2025-liquid-cooling/（二手轉述）；原報告 https://intelligence.uptimeinstitute.com/resource/uptime-institute-global-data-center-survey-2025（PDF 未讀）",
         2, "中立", "Uptime Institute：中立調查機構；本筆只讀二手轉述", "二手（已讀轉述；原 PDF 未讀）",
         "Inputs D9:F9（PUE）的可比對象（SRC_DC_018 的第二來源）；不直接連結",
         "v13.1 D15 B 段新增（E381）：原 PDF 未讀，超大規模與 colo 的範圍待原文確認；與 SRC_DC_018 互為第二來源", second="SRC_DC_018"),
    _rec("SRC_Price", "SRC_PRC_045", "E382", "OpenAI Batch API 價格 ÷ 標準價", 0.5, "x（占標準價）",
         "Batch API（非同步）；官方文件寫『50% cost discount』", "OpenAI API", "https://developers.openai.com/api/docs/guides/batch（2026-10-17 已讀）", 1, "利害關係方",
         "OpenAI：賣方自訂價", "一手（已讀）", "Cap_In C12（有效折扣）的牌價事實，Gov_Map GM441／442 的 H 欄；只確認 Batch 半價，不含占比與企業議價；不直接連結",
         "v13.1 D15 B 段新增（E382）：第二來源 SRC_PRC_046（Anthropic，不同主體）", second="SRC_PRC_046"),
    _rec("SRC_Price", "SRC_PRC_046", "E382", "Anthropic Message Batches 價格 ÷ 標準價", 0.5, "x（占標準價）",
         "Message Batches；官方文件寫『50% of the standard API prices』", "Anthropic API", "https://platform.claude.com/docs/en/build-with-claude/batch-processing（2026-10-17 已讀）", 1, "利害關係方",
         "Anthropic：賣方自訂價", "一手（已讀）", "Cap_In C12（有效折扣）的牌價事實（SRC_PRC_045 的第二來源）；不直接連結",
         "v13.1 D15 B 段新增（E382）：與 SRC_PRC_045 互為第二來源", second="SRC_PRC_045"),
    _rec("SRC_HW", "SRC_HW_069", "E383", "H200 HBM 容量（第二來源：Lenovo Press LP1944）", 141, "GB/GPU",
         "SXM 與 NVL 兩款同值；產品指南更新日 2025-11-07", "Hopper H200（Lenovo 產品指南）", "https://lenovopress.lenovo.com/LP1944（WebFetch 摘要 2026-10-17 已讀）", 1, "利害關係方",
         "Lenovo：伺服器 OEM，有銷售誘因，但與 NVIDIA 為不同主體", "一手（已讀摘要）", "SRC_HW_067 的第二來源（Calib L 節 E 欄的 H200 條件）；不直接連結",
         "v13.1 D15 B 段新增（E383）：數值與 SRC_HW_067 相同，不改值", second="SRC_HW_067"),
    _rec("SRC_HW", "SRC_HW_070", "E383", "H200 HBM 頻寬（第二來源：Lenovo Press LP1944）", 4.8, "TB/s/GPU",
         "SXM 與 NVL 兩款同值；產品指南更新日 2025-11-07", "Hopper H200（Lenovo 產品指南）", "https://lenovopress.lenovo.com/LP1944（WebFetch 摘要 2026-10-17 已讀）", 1, "利害關係方",
         "Lenovo：伺服器 OEM，有銷售誘因，但與 NVIDIA 為不同主體", "一手（已讀摘要）", "SRC_HW_068 的第二來源（Calib L 節 E 欄的 H200 條件）；不直接連結",
         "v13.1 D15 B 段新增（E383）：數值與 SRC_HW_068 相同，不改值", second="SRC_HW_068"),
]


def src_new_records():
    return SRC_NEW


# ================================================================== DB_Evidence rows E381–E386 (17 columns A..Q) and column R
_IDS = {x["id"] for x in SRC_NEW}
EVIDENCE_V1310 = [
    [ "E381", DATE, "PUE 外部對照：Google 2025 全機隊年均 1.09（2026 Q2 TTM 1.09）；Uptime Institute 2025 全球平均 1.54（超大規模 1.10–1.15、colo／企業 1.58–1.80，mgrid 轉述）",
      "https://datacenters.google/efficiency；mgrid.org 2025-10-01 轉述 Uptime 年度調查（原 PDF 未讀）；" + EVTAB, "Interested-party（Google）／中立調查（Uptime）", "Inputs D9:F9（PUE）", "1.1／1.2／1.3（Assumed）",
      "1.09（Google 全機隊）；1.54（全球平均）", "並列（Analogy 對照，數值不改）", V,
      f"CC 查證 D15 4-2b-2A，B 段寫入｜{DEC_BY}｜全機隊或一般機房≠液冷 AI 設施，故標 Analogy；Uptime 原文未讀", "已處理", "SRC_DC_018、SRC_DC_019", DASH, "不改任何 IF_（PUE 數值不變）", "1", "利害關係方（Google）／中立（Uptime）"],
    [ "E382", DATE, "Batch 半價牌價：OpenAI Batch API『50% cost discount』；Anthropic Message Batches『50% of the standard API prices』",
      "https://developers.openai.com/api/docs/guides/batch；https://platform.claude.com/docs/en/build-with-claude/batch-processing；" + EVTAB, "Verified（賣方牌價）", "Cap_In C12（有效折扣，Gov_Map GM441／442）", "0.2（區間 0.1–0.4，Assumed）",
      "Batch 價＝標準價 × 0.5", "並列（只確認 Batch 半價；占比與企業議價仍找不到，E333、E334）", V,
      f"CC 查證 D15 4-2b-2A，B 段寫入｜{DEC_BY}｜兩家為不同主體，互為第二來源；賣方自訂價屬 Interested-party", "已處理", "SRC_PRC_045、SRC_PRC_046", DASH, "不改任何 IF_（有效折扣不變）", "1", "利害關係方（賣方牌價）"],
    [ "E383", DATE, "H200 規格第二來源：Lenovo Press LP1944 產品指南 141 GB HBM3e、4.8 TB/s（SXM 與 NVL 同值；更新日 2025-11-07）",
      "https://lenovopress.lenovo.com/LP1944；" + EVTAB, "Interested-party", "SRC_HW_067／068（Calib L 節 E 欄的 H200 條件）", "141 GB；4.8 TB/s（NVIDIA 產品頁）", "141 GB；4.8 TB/s", "採納（第二來源，數值相同）", V,
      f"CC 查證 D15 4-2b-2A，B 段寫入｜{DEC_BY}｜Lenovo 為伺服器 OEM，有銷售誘因，但與 NVIDIA 為不同主體", "已處理", "SRC_HW_069、SRC_HW_070", DASH, "不改任何 IF_（規格值不變）", "1", "利害關係方"],
    [ "E384", DATE, "MLCommons v5.1 DeepSeek-R1 原始資料（summary_results.json，共 17 列）找不到 H200 Server 列，也找不到 556（SRC_PERF_060）；H200 只有 HPE Open 組 Offline 一列（每 GPU 1,113.1）",
      "https://raw.githubusercontent.com/mlcommons/inference_results_v5.1/main/summary_results.json；docs/evidence/20261017_D15/04_mlperf51.py；" + EVTAB, "—", "SRC_PERF_060（Checks V1 的對照值）", "556 tok/s/GPU（轉引 3A 查證，未回原文）",
      "原始資料中無此值", "待查（SRC_PERF_060 標『待回查』，V1 對照值未回查）", V,
      f"CC 查證 D15 4-2b-2A，B 段寫入｜{DEC_BY}｜第 4 筆第二來源選 SRC_PERF_060 是代理的選擇，Andy 2026-10-17 照預設；下一步：回查 3A 原表 556 的出處", "待判定", "SRC_PERF_060", DASH, "不改任何 IF_（V1 為驗證項，數值不變）", "—", "—"],
    [ "E385", DATE, "SRC_PERF_081（TensorRT-LLM H200 每 GPU 1,436）的近似口徑對照：MLPerf v5.1 Open 組 Offline（HPE Cray XD670 8× H200）每 GPU 1,113.1；dstack 2025-03 SGLang Offline 最高 6,311 tok/s（8 GPU，每 GPU 約 789）",
      "MLCommons summary_results.json（5.1-0375）；https://dstack.ai/blog/h200-mi300x-deepskeek-benchmark/；" + EVTAB, "Interested-party（HPE）／中立偏工具商（dstack）", "Checks V22（SRC_PERF_081 單一來源）", "1,436.2（單一來源）", "1,113.1；約 789（皆非同口徑）",
      "並列（近似口徑，不登錄 SRC；不能消除 V22 的單一來源警示）", V,
      f"CC 查證 D15 4-2b-2A，B 段寫入｜工作單未要求登錄為 SRC，列待 Project 判斷｜模型 ÷ 1,113 ＝1.29，但並行數、長度、資料集不同", "待判定", "—", DASH, "不改任何 IF_", "—", "—"],
    [ "E386", DATE, "Inputs D19:F19（機械與散熱）外部管道：datacentres.com 寫 2–4 $M/MW（機械）、總額 11.3 $M/MW，未附原始出處；Cushman & Wakefield 2026 指南的分項在付費牆後",
      "https://www.datacentres.com/guides/data-centre-costs-explained；" + EVTAB, "Assumed（來源未附出處）", "Inputs D19:F19", "2.0–3.2 $M/MW（Assumed）", "2–4 $M/MW（無出處）", "不採納（來源未附出處；不作來源，只作參考）", V,
      f"CC 查證 D15 4-2b-2A，B 段寫入｜{DEC_BY}｜外部來源仍找不到；數值不改，決策 AF3B-Q6＋D15-04", "已處理", "—", DASH, "不改任何 IF_", "3", "第三方"],
]


def evidence_rows():
    return EVIDENCE_V1310


EV_R_HDR = "D15 B 段處置（v13.1；建議文字，原判定欄 I 與狀態欄 L 不改）"


def evidence_col_r(wb):
    """Column R of DB_Evidence: the D15 recommendation (採納／並列／駁回／維持待查 + reason) for the 106 backlog rows; guarded (written only while R is empty)."""
    ws = wb["DB_Evidence"]; n = 0
    if ws["R4"].value in (None, ""):
        put(ws, "R4", EV_R_HDR, F_BOLD, wrap=True); ws.column_dimensions["R"].width = 60; n += 1
    row_of = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}
    for e in D15["evidence"]:
        r = row_of.get(e["id"])
        if r is None: continue
        if ws.cell(r, 18).value in (None, ""):
            put(ws, f"R{r}", f"{e['rec']}｜{e['why']}", F_CALC, wrap=True); n += 1
    # the new rows E381–E386 carry their own verdict in column I; R repeats it so that the column can be filtered
    for eid in ("E381", "E382", "E383", "E384", "E385", "E386"):
        r = row_of.get(eid)
        if r is not None and ws.cell(r, 18).value in (None, ""):
            put(ws, f"R{r}", f"{ws.cell(r, 9).value}｜v13.1 新增列（非 106 列積壓）", F_CALC, wrap=True); n += 1
    return n


# ================================================================== Gov_Map: J／H／I／N for the 312 rows (+ GM346 column O sentence, GM587 note)
SRC_DC3 = ("SRC_DC_003", "總額校準（D15-04：JLL 校準總額；分項本身為 Assumed）")
_PUE = ("SRC_DC_018", "可比（Analogy：Google 全機隊年均 PUE 1.09；第二來源 SRC_DC_019 Uptime 全球平均 1.54，超大規模 1.10–1.15；全機隊≠液冷 AI 設施）")
_ISL = ("SRC_DEM_014", "對照（OpenRouter State of AI 2025-12：每請求輸入 >6K、輸出約 400，SRC_DEM_014–017；單一資料集，非來源）")
_BATCH = ("SRC_PRC_045", "牌價事實（Batch 半價：OpenAI SRC_PRC_045、Anthropic SRC_PRC_046；占比與企業議價為假設）")
HMAP = {
    "GM004": _PUE, "GM005": _PUE, "GM006": _PUE,
    **{g: SRC_DC3 for g in ("GM022", "GM023", "GM024", "GM025", "GM026", "GM027", "GM031", "GM032", "GM033", "GM034", "GM035", "GM036")},
    **{f"GM{n}": _ISL for n in range(245, 251)},
    "GM441": _BATCH, "GM442": _BATCH,
    "GM443": ("SRC_DEM_046", "對照（程式代理實測 SRC_DEM_046–048：Pegatron 99.3%、Requesty Claude Code 91.9%，只代表一種用途）"),
    "GM570": ("SRC_MOD_009", "推導自 V4 模型卡結構（SRC_MOD_009 KV head、016 壓縮 128×、018／019 CSA top-k 與塊長、015／054 全注意力層比例）"),
}
SPECIAL_NOTE = {
    "GM203": "｜Andy 2026-10-17 決定：Arch E6 頁內標記以 Gov_Map 為準，改為 Assumed（修訂九：原字樣保留於 Arch G6，見 Decisions D15-10）",
    "GM204": "｜Andy 2026-10-17 決定：Arch E7 頁內標記以 Gov_Map 為準，改為 Assumed（修訂九：原字樣保留於 Arch G7，見 Decisions D15-10）",
    "GM205": "｜Andy 2026-10-17 決定：Arch E8 頁內標記以 Gov_Map 為準，改為 Assumed（修訂九：原字樣保留於 Arch G8，見 Decisions D15-10）",
    "GM352": "｜Andy 2026-10-17 決定：擴用 Decisions C2（已搜尋，無可比對象）至本格（Decisions D15-09）；C2 原列不改",
    "GM353": "｜Andy 2026-10-17 決定：擴用 Decisions C2（已搜尋，無可比對象）至本格（Decisions D15-09）；C2 原列不改",
    "GM355": "｜Andy 2026-10-17 決定：擴用 Decisions C2（已搜尋，無可比對象）至本格（Decisions D15-09）；C2 原列不改",
    "GM587": "｜v13.1 D15（Andy 2026-10-17 照預設，Decisions D15-12）：區間 2,000–11,000 不動；v13.0 閉合值 L1_TokPerPromptClose 38,186 為輸入區間上限 11,000 的 3.5 倍，屬未解釋；閉合值是殘差（由 Alloc_In C23、C24 推出），不作為放寬輸入區間的證據",
}
TOKEN_NOTE = ("｜v13.1 D15（Andy 2026-10-17 照預設，Decisions D15-12）：v13.0 閉合值 L1_TokPerPromptClose 38,186 為輸入區間上限 11,000 的 3.5 倍，屬未解釋；"
              "閉合值是殘差（由 Alloc_In C23、C24 推出），不作為放寬輸入區間的證據；Alloc_In C15 區間不動")
O_OLD = "類別欄原字樣『倒推』保留，見 Decisions BB2-H4"
O_NEW = "原字樣『倒推』保留於 AL 欄；類別依修訂九更正為 Assumed，見 Decisions BB2-H4、BB3-K3"


def _is_blank(v):
    return v in (None, "", DASH)


def gov_update(ws, append_text):
    n = 0
    row_of = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}
    for g in D15["gov"]:
        r = row_of[g["gm"]]
        assert (ws.cell(r, 3).value, ws.cell(r, 4).value) == (g["page"], g["cell"]), (g["gm"], ws.cell(r, 3).value, ws.cell(r, 4).value)
        ids = g["ids"]
        if ids and _is_blank(ws.cell(r, 10).value):
            ws.cell(r, 10).value = ids[0]; n += 1
        h = HMAP.get(g["gm"])
        if h and _is_blank(ws.cell(r, 8).value):
            ws.cell(r, 8).value = h[0]; n += 1
            if _is_blank(ws.cell(r, 9).value):
                ws.cell(r, 9).value = h[1]; n += 1
        # N: everything that J (one ID) and H (one SRC) cannot carry
        src = str(g["src_text"])
        trivial = bool(ids) and src == f"決策：{ids[0]}"
        if not trivial:
            body = src
            if len(ids) > 1:
                body += f"（J 欄放主要決策 {ids[0]}；全部：{'＋'.join(ids)}）"
            n += append_text(ws.cell(r, 14), f" ｜v13.1 D15 補登（原 J／H 欄為『—』）：{body}")
        mark = str(g["mark"] or "")
        cur_f = str(ws.cell(r, 6).value or "")
        if mark.startswith(("定義", "Derived")) and not cur_f.startswith(mark.split("（")[0]) and g["gm"] not in SPECIAL_NOTE:
            n += append_text(ws.cell(r, 14), f" ｜v13.1 D15 建議標記：{mark}（類別欄 F 現為 {cur_f}，本版未改，待 Project 判斷）")
        if g["gm"] in SPECIAL_NOTE:
            n += append_text(ws.cell(r, 14), " " + SPECIAL_NOTE[g["gm"]])
    # P2-4: GM346 (Calib C68) column O: the v13.0 sentence "類別欄原字樣『倒推』保留" no longer matches the cell (revision 9)
    r = row_of["GM346"]
    c = ws.cell(r, 15)
    if isinstance(c.value, str) and O_OLD in c.value:
        c.value = c.value.replace(O_OLD, O_NEW); n += 1
    return n


# ================================================================== Decisions (10 columns; layout of v1300)
def _dtext(i):
    d = D15["decisions"][f"D15-{i:02d}"]
    return d["item"], d["text"], d["scope"]


_EXTRA = {
    4: " ｜本版 Gov_Map H 欄只填於明細列明的 12 格（Inputs 第 17、18、20、21 列的 D／E／F，H＝SRC_DC_003、關係『總額校準』）；第 14–16、19、31–35 列只填決策編號。",
    9: " ｜Andy 2026-10-17 決定照預設（D15 第 4 項）；Decisions C2 原列不改，本決策是擴用紀錄。",
}
_D = []
for _i in range(1, 10):
    _item, _txt, _scope = _dtext(_i)
    _D.append((f"D15-{_i:02d}", f"D15：{_item}", _txt + _EXTRA.get(_i, ""), _scope, "v13.1 已登錄"))
_D += [
    ("D15-10", "D15：Arch E6–E8 頁內標記更正為 Assumed（修訂九首批應用之一）",
     "Arch F6:F8 原字樣『Verified／Assumed』更正為『Assumed』：Gov_Map GM203–205 的類別已是 Assumed，Astra 欄 4T／180B／100 層無公開來源（封閉模型），Luna／Sol 欄的 Verified 來自各自連結的 SRC_MOD（模型卡）。"
     "依 CLAUDE.md 第 6 節更正例外（修訂九）：原字樣『Verified／Assumed』逐字保留在 Arch G6:G8（來源／說明欄）並寫明更正版本 v13.1；本筆為更正紀錄。數值不改。Arch F9（d_model）同樣寫 Verified／Assumed，不在本工作單範圍，未改。",
     "Arch F6:F8、G6:G8；Gov_Map GM203–205 的 N 欄", "v13.1 已建"),
    ("D15-11", "D15：SRC_PERF_060（MLPerf v5.1 DGX H200 Server 556 tok/s/GPU）標『待回查』",
     "MLCommons v5.1 summary_results.json 的 deepseek-r1 共 17 列，H200 只有 HPE Open 組 Offline 一列（每 GPU 1,113.1）；沒有 H200 Server 列，也找不到 556（E384）。"
     "SRC_PERF_060 的備註欄加註『待回查』（狀態欄 O 維持 Active，否則同指標檢查會失效）、Checks V1 說明欄加註『對照值未回查』；數值不改。"
     "第 4 筆第二來源選 SRC_PERF_060 是代理的選擇，Andy 2026-10-17 照預設。下一步：回查 3A 原表 556 的出處。",
     "SRC_PERF_060 W 欄；Checks V1 J 欄；DB_Evidence E384", "v13.1 已建；待回查"),
    ("D15-12", "D15：每則提示 token 數的區間不動，閉合值的差距標為『未解釋』",
     "Alloc_In C15 區間 2,000–11,000 不動。v13.0 閉合值 L1_TokPerPromptClose＝38,186，為區間上限 11,000 的 3.5 倍，屬未解釋；閉合值是殘差（由 Alloc_In C23、C24 推出），不作為放寬輸入區間的證據。"
     "說明寫入 L1_GapPrompt 列的讀法欄與 Gov_Map GM587 的 N 欄。數值不改。",
     "L1 第 70 列 I 欄；Gov_Map GM587 N 欄", "v13.1 已建"),
]
DECISIONS_V1310 = [[i, "D15 來源補齊", t, txt, DATE, DEC_BY, st, where, f"{WO} P1；{REPORT} 明細第 5 頁", "否"] for i, t, txt, where, st in _D]


def decisions_append(wb):
    ws = wb["Decisions"]
    have = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
    r = max([rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")] or [4]) + 1
    n = 0
    for row in DECISIONS_V1310:
        if row[0] in have: continue
        for i, v in enumerate(row): put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 5, 7))
        r += 1; n += 1
    assert r - 1 <= 300, "Decisions beyond Gov_Map AA's COUNTIF span ($A$5:$A$300)"
    return n


# ================================================================== P1 notes in builder-owned and Excel-owned pages
def notes(wb, append_text):
    log = {}
    # Arch F6:F8 (revision 9: original wording kept in G of the same row)
    ws = wb["Arch"]; n = 0
    for r in (6, 7, 8):
        if ws[f"F{r}"].value == "Verified／Assumed":
            ws[f"F{r}"].value = "Assumed"; n += 1
        n += append_text(ws[f"G{r}"], f" ｜v13.1 更正（修訂九，Decisions D15-10）：標記欄原字樣『Verified／Assumed』更正為『Assumed』（以 Gov_Map 為準；Luna／Sol 欄數值連結 SRC_MOD，Astra 欄無公開來源）")
    log["arch_marks"] = n
    # SRC_PERF_060 (Excel-owned): note in 備註 W; 狀態 O stays Active
    sp = wb["SRC_Perf"]; n = 0
    for r in range(5, sp.max_row + 1):
        if sp.cell(r, 1).value == "SRC_PERF_060":
            n += append_text(sp.cell(r, 23), "【狀態：待回查（MLCommons v5.1 原始資料找不到 H200 Server 556；2026-10-17）】原始資料 deepseek-r1 共 17 列，H200 只有 HPE Open 組 Offline 一列（每 GPU 1,113.1）；本欄值轉引 3A 查證，未回原文；見 E384、Decisions D15-11")
    log["perf060"] = n
    # Checks V1 說明欄 J
    ck = wb["Checks"]
    r = next(x for x in range(1, ck.max_row + 1) if ck.cell(x, 1).value == "V1")
    log["checks_v1"] = append_text(ck.cell(r, 10), "｜v13.1 D15：對照值未回查（SRC_PERF_060＝556 待回查：MLCommons v5.1 原始資料找不到 H200 Server 556；2026-10-17；Decisions D15-11）")
    # L1_GapPrompt 讀法欄 I
    l1 = wb["L1"]
    r = next(x for x in range(1, l1.max_row + 1) if l1.cell(x, 1).value == "L1_GapPrompt")
    log["l1_gapprompt"] = append_text(l1.cell(r, 9), TOKEN_NOTE)
    # second-source fields (Excel-owned; written only while R is empty): SRC_HW_067／068 -> Lenovo; the new rows carry theirs through v1100-style apply
    sec = {x["id"]: x["second"] for x in SRC_NEW if x.get("second")}
    k = 0
    for sh in ("SRC_DC", "SRC_Price", "SRC_HW"):
        w = wb[sh]
        for rr in range(5, w.max_row + 1):
            i = w.cell(rr, 1).value
            if i in sec and w.cell(rr, 18).value in (None, DASH, ""):
                w.cell(rr, 18).value = sec[i]; k += 1
            if i in ("SRC_HW_067", "SRC_HW_068") and w.cell(rr, 18).value in (None, DASH, ""):
                w.cell(rr, 18).value = "SRC_HW_069" if i == "SRC_HW_067" else "SRC_HW_070"; k += 1
    log["second_sources"] = k
    return log


# ================================================================== P2-1／P2-2: Alloc O note, Alloc P (fingerprint), RNG_StaleFlag, Checks RG3
VR200_NOTE = ("（c）v13.1 P2-1：VR200 欄無區間：批次倍數無證據，為點值。已納入區間的 36 個名稱（IF_ 的 15 格名稱）各有 3 格屬 VR200 世代，共 108 格；它們的 Lo／Hi 只涵蓋 GB200／Hopper 倍數，"
              "VR200 倍數（CAL_BatchEtaD 第 4 格＝0.75；Gov_Map H156 區間 0.70–1.0）沒有證據，故不設區間。VR200 倍數離開 0.75（例如設為 1.0）時，這 108 格與上面（a）的 12 個名稱會落在區間之外或改變。下游應把 VR200 欄當點值使用。")
FP_TITLE = ("P. 區間時效指紋（v13.1 P2-2）：下列為『53 個有區間的名稱』上游的全部數值常數格（依賴圖，口徑同死格偵測；builder/v1310_fingerprint.json，tools/gen_fingerprint.py 產生）。"
            "D 欄是產生端點時記錄的值（黑字常數，不是公式），C 欄是現值（連結），E 欄＝1 表示不符。任一格不符，RNG_StaleFlag 與 Checks RG3 顯示『區間過期』——"
            "即使基準仍在 Lo／Hi 之內（舊旗標只在基準出界或版本不符時亮起）。上游輸入改動後須重新產生端點檔與指紋檔。")


def _last_row(ws):
    return max(c.row for row in ws.iter_rows() for c in row if c.value is not None)


def alloc_fingerprint(wb):
    ws = wb["Alloc"]
    # P2-1 note directly under the (b) line of section O
    rb = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("（b）依賴圖上有路徑"))
    if ws.cell(rb + 1, 1).value in (None, ""):
        put(ws, f"A{rb + 1}", VR200_NOTE, F_NOTE)
    r0 = _last_row(ws) + 2
    section(ws, r0, FP_TITLE, 18)
    leaves = FP["leaves"]
    r = r0 + 1
    put(ws, f"A{r}", "指紋格數", F_BOLD); r_n = r
    r += 1
    put(ws, f"A{r}", "指紋不符格數（現值≠產生端點時的值）", F_BOLD); r_d = r
    r += 2
    put(ws, f"A{r}", "工作表", F_BOLD); put(ws, f"B{r}", "儲存格", F_BOLD); put(ws, f"C{r}", "現值（連結）", F_BOLD)
    put(ws, f"D{r}", "產生端點時的值（常數）", F_BOLD); put(ws, f"E{r}", "不符（1＝不符）", F_BOLD)
    r += 1
    first = r
    for sh, cell, val in leaves:
        put(ws, f"A{r}", sh); put(ws, f"B{r}", cell)
        put(ws, f"C{r}", f"='{sh}'!{cell}", F_LINK)
        put(ws, f"D{r}", val, F_CALC)
        put(ws, f"E{r}", f"=IF(ABS(C{r}-D{r})<=1E-9*ABS(D{r})+1E-12,0,1)", F_CALC, fmt="0")
        r += 1
    last = r - 1
    put(ws, f"C{r_n}", f"=COUNT(C{first}:C{last})", F_CALC, fmt="0")
    put(ws, f"C{r_d}", f"=SUM(E{first}:E{last})", F_CALC, fmt="0")
    put(ws, f"D{r_n}", "指紋檔的葉格數（builder/v1310_fingerprint.json）；tests/parity 以依賴圖重新推導並要求逐格相等", F_NOTE)
    put(ws, f"D{r_d}", "不為 0＝上游輸入已改動而端點未重新產生；Checks RG3（WARN）", F_NOTE)
    put(ws, f"Y{r_n}", "顯示用：指紋格數（C 欄）無公式讀取；與指紋檔葉格數的對照由 tests/parity 檢查", F_NOTE)      # 死格需有『用途』標記（Alloc 用途欄＝Y）
    # RNG_StaleFlag also requires the fingerprint to match
    ref = wb.defined_names["RNG_StaleFlag"].attr_text.split("!")[1].replace("$", "")
    cell = ws[ref]
    old = str(cell.value)
    if f"C{r_d}" not in old:
        assert old.startswith("=IF(AND(") and ',C' in old, old
        cell.value = old.replace(f'=0),"一致"', f'=0,C{r_d}=0),"一致"', 1)
        assert f"C{r_d}=0" in cell.value, cell.value
    rf = int(ref[1:])
    dcell = ws[f"D{rf}"]
    if "指紋" not in str(dcell.value):
        dcell.value = str(dcell.value) + "；v13.1：另需 P 節指紋全部相符（任一上游輸入與產生端點時不同也顯示『區間過期』）"
    return {"first": first, "last": last, "n": len(leaves), "fp_row": r_d, "count_row": r_n}


def checks_rg3(wb, fp):
    ws = wb["Checks"]
    r2 = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "RG2")
    r3 = r2 + 1
    assert ws.cell(r3, 1).value in (None, "", "RG3"), ws.cell(r3, 1).value
    put(ws, f"A{r3}", "RG3")
    put(ws, f"B{r3}", "批次口徑區間：上游輸入指紋不符（Alloc P 節）：任一上游數值常數格與產生端點時的值不同：1＝是（基準可能仍在 Lo／Hi 之內；計入 GOV_Warnings 時與 RG2 合算最多 1）", wrap=True)
    put(ws, f"C{r3}", "WARN")
    put(ws, f"D{r3}", f"=IF(Alloc!$C${fp['fp_row']}>0,1,0)", F_CALC, fmt="0")
    put(ws, f"E{r3}", "Alloc P 節 C 欄（不符格數）；builder/v1310_fingerprint.json", F_NOTE)
    ref = wb.defined_names["GOV_Warnings"].attr_text.split("!")[1].replace("$", "")
    cell = ws[ref]
    f = str(cell.value)
    if f"MAX(D{r2},D{r3})" not in f:
        assert f.endswith(f"+D{r2}"), f
        cell.value = f[: -len(f"+D{r2}")] + f"+MAX(D{r2},D{r3})"
    lab = ws[f"B{int(ref[1:])}"]
    if "RG3" not in str(lab.value):
        lab.value = str(lab.value).replace("、RG1、RG2）", "、RG1、RG2、RG3）") if "、RG1、RG2）" in str(lab.value) else f"{lab.value}、RG3"
    return r3


# ================================================================== end of build
def apply(wb, append_text):
    log = {}
    log["notes"] = notes(wb, append_text)
    log["evidence_col_r"] = evidence_col_r(wb)
    log["decisions"] = decisions_append(wb)
    fp = alloc_fingerprint(wb)
    log["fingerprint"] = fp
    log["rg3_row"] = checks_rg3(wb, fp)
    return log


# ================================================================== README
README_VERSION = (VERSION + "（次版：" + WO + "——P1 D15 來源補齊（Gov_Map 312 列補 J／H 欄、Decisions D15-01…12、新 SRC：SRC_DC_018／019、SRC_PRC_045／046、SRC_HW_069／070，DB_Evidence R 欄與 E381–E386，"
                  "Arch E6–E8 標記更正，SRC_PERF_060 待回查）；P2 查核 #96 小項（VR200 欄為點值的說明、區間時效指紋 Alloc P 節與 Checks RG3、C68 O 欄文字）；數值全部不變）。以下為 "
                  + v1300.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1300.README_TITLE
README_TITLE = V + _T[len(v1300.V):]
README_ROW = ("D15 來源補齊與區間時效指紋（v13.1）",
              "Gov_Map 無來源無決策列補登決策編號與來源；新增 6 筆 SRC；DB_Evidence 積壓 106 列加建議處置欄；區間時效旗標加上游輸入指紋（Alloc P 節、Checks RG3）。數值不變。細節見 Decisions D15-*。")
