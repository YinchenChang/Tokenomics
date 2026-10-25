# v21.7 (work order docs/workorders/20261025_v21.7_NonNV_AMD與Trn同源價.md r0, branch claude/v21-7): NonNV AMD row 5 derived values, Trainium3 holding ratio on the KGI same-source basis, stale-text fixes.
# MINOR version: input values change (AMD B5:G5 six cells, Trainium3 F7), written only while they still hold the v21.6 value (like v2160.inputs_update); no new/removed contract names; contract v0.27 unchanged.
#   W1  NonNV B5:D5 (AMD output ratio) 0.5／0.7／1.0 -> 0.70／0.80／1.15;  E5:G5 (AMD holding ratio) 0.8／0.9／1.0 -> 0.67／0.85／1.12;  F7 (Trainium3 holding ratio base) 0.55 -> 0.61.
#   W2  Text (revision 9: the original wording is kept verbatim after "｜原："): NonNV L5, L7 (v21.7 sentence added after the v21.6 note), L8, L9, A2; Interface A368.
#   W3  DB_Evidence E410／E411 (after E409), Decisions SL64-01, README B5 opening paragraph.
# Contract names whose values change: IF_NonNV_Out_AMD, IF_NonNV_Hold_AMD (3 cells each), IF_NonNV_Hold_Trn (base cell D376) and the version strings. Everything else unchanged.
# Excel-owned cells (DB_Evidence, Decisions) are written only while absent; the NonNV／Interface text cells are rewritten by the builder on every build, so the text fixes are re-applied deterministically. Idempotent.
from openpyxl.utils import get_column_letter as L

from common import put, F_CALC
import v1410
import v2110
import v2160

VERSION = "20261025_Tokenomics_v21.7"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-25"
DASH = "—"
V = "v21.7"
WO = "工作單 docs/workorders/20261025_v21.7_NonNV_AMD與Trn同源價.md（r0）"
DEC_ID = "SL64-01"
EVD = "docs/evidence/20261025_NonNV_AMD/"
DEC_BY = ("Andy 未回應，照工作單預設（2026-10-25 對話判斷題 9、10）：PR #174（NonNV C 段 AMD，獨立核對 r1 後合併）Q4、Q5 預設——AMD 產出比與持有比都換，持有比用 TK 買方取得價同口徑；"
          "PR #174 獨立核對與 README §10——B 段 Trn3 持有比用 KGI 絕對價 ÷ TK 價屬輕度口徑混用，改用 KGI 同源比值，基準取中點 0.61；v21.6 代理回報的過時說明（Interface A368、NonNV A2、L5、L8、L9）。"
          "TPU、Trn 產出比仍待 Q1，不改")

# ---------------------------------------------------------------- W1 inputs (old-value guards)
AMD_OUT_OLD, AMD_OUT_NEW = (0.5, 0.7, 1.0), (0.70, 0.80, 1.15)
AMD_HOLD_OLD, AMD_HOLD_NEW = (0.8, 0.9, 1.0), (0.67, 0.85, 1.12)
TRN_F7_OLD, TRN_F7_NEW = 0.55, 0.61


def _same(c, o):
    return c.value is not None and type(c.value) is not bool and c.value == o


def inputs_update(wb):
    log = []
    ws = wb["NonNV"]
    assert ws["A5"].value == "AMD MI455X（Helios）", ws["A5"].value
    assert ws["A7"].value == "AWS Trainium3", ws["A7"].value
    for cols, old, new, what in (("BCD", AMD_OUT_OLD, AMD_OUT_NEW, "output"), ("EFG", AMD_HOLD_OLD, AMD_HOLD_NEW, "holding")):
        cells = [ws[f"{c}5"] for c in cols]
        if all(_same(c, o) for c, o in zip(cells, old)):
            for c, n in zip(cells, new):
                c.value = n
            log.append(f"NonNV!{cols[0]}5:{cols[-1]}5 (AMD {what}): {old} -> {new} (v21.7 W1)")
    if _same(ws["F7"], TRN_F7_OLD):
        ws["F7"].value = TRN_F7_NEW
        log.append(f"NonNV!F7 (Trainium3 holding base): {TRN_F7_OLD} -> {TRN_F7_NEW} (v21.7 W1)")
    return log


# ---------------------------------------------------------------- W2 text (original kept verbatim after 原：)
TAIL = v2110.ROWS_5_9_TAIL      # the v21.0 conventional-value tail
L5_NOTE = ("產出比 v21.7 依每 kW 峰值算力比 1.066（MI455X MXFP4 40.3 PF、TBP 2,500 W、Helios 245 kW；高端以 225 kW 計，每 kW 比 1.16；有原生 FP4，兩種精度口徑相同）"
           "× 效率 f 0.656／0.757／0.994（MLPerf v6.1 MI355X DeepSeek-R1 還原 × 轉換係數 t 0.891，Assumed）；"
           "持有比依 TK 買方取得價同口徑（每架 $5.25M 下界／$7.3M＝AMD 自述隱含比 0.865 × VR200／$10.5M＝Futurum 同源比值 1.25 × VR200；Assumed）；"
           "每 token 成本比基準約 1.06（取整後輸入值 0.85／0.80；查證原始值 0.847／0.807＝1.05；PR #174）")
L7_ADD = ("v21.7：持有比基準 0.55 → 0.61，改用 KGI 同源比值（KGI 同報 VR200『超過 700 萬美元』、GB300 $3.9–4.0M，口徑較 TK 低約 9–20%；"
          "Trn3 每架約 $1.6–1.8M → 0.593–0.631 取中點；PR #174 §10）")
L89_NOTE = "TPU v7、Trainium3（v21.6）、AMD（v21.7）已以每 kW 法重推；本列尚無機架功率與價格資料，維持慣例值。"
A2_OLD = "持有成本比預設接近 1，依 Block 1 結論：每 GW 持有成本幾乎不隨世代改變"
A2_NEW = "持有成本比：TPU、Trn3、AMD、昇騰、Maia 依每 kW 設備價推導（v21.1、v21.6、v21.7）；Cerebras、OpenAI 客製仍為慣例值（Block 1『接近 1』）"
A368_OLD = "（第 379–382 列；持有比依每 kW 設備價推導，既有五列為慣例值）"
A368_NEW = ("（第 379–382 列）；持有比已推導：昇騰、Maia、TPU、Trn3、AMD，其餘為慣例值；產出比已推導：AMD、昇騰、Maia（TPU、Trn3 待精度口徑）")


def text_fixes(wb):
    ws = wb["NonNV"]
    n = 0
    # L5, L8, L9: replace the conventional-value tail with the note; original tail kept after 原：
    for r, note in ((5, L5_NOTE), (8, L89_NOTE), (9, L89_NOTE)):
        cell = ws[f"L{r}"]
        cur = str(cell.value)
        if "｜原：" in cur:
            continue
        assert cur.endswith(TAIL), (r, cur[-80:])
        head = cur[:-len(TAIL)].rstrip("｜")
        cell.value = head + "｜" + note + "｜原：" + TAIL
        n += 1
    # L7: v21.7 sentence added after the v21.6 note (before 原：)
    cell = ws["L7"]
    cur = str(cell.value)
    if L7_ADD not in cur:
        assert "持有比 v21.6 依每 kW 設備價推導" in cur and cur.count("｜原：") == 1, cur[-120:]
        i = cur.index("｜原：")
        cell.value = cur[:i] + "｜" + L7_ADD + cur[i:]
        n += 1
    # NonNV A2 subtitle
    a2 = ws["A2"]
    cur = str(a2.value)
    if A2_NEW not in cur:
        assert cur.endswith(A2_OLD), cur
        a2.value = cur[:-len(A2_OLD)] + A2_NEW + "｜原：" + A2_OLD
        n += 1
    # Interface A368
    wi = wb["Interface"]
    t = str(wi["A368"].value)
    if A368_NEW not in t:
        assert t.endswith(A368_OLD) and v2110.P_TITLE_NEW in t, t[-120:]
        wi["A368"].value = t[:-len(A368_OLD)] + A368_NEW + "｜原：" + A368_OLD
        n += 1
    return n


# ---------------------------------------------------------------- W3 DB_Evidence E410／E411
_OLD_AMD = "NonNV 第 5 列 AMD MI455X 產出比 0.5／0.7／1.0、持有比 0.8／0.9／1.0（慣例值）；每 token 成本比基準 1.29"
E410 = [
    "E410", DATE,
    "NonNV C 段 AMD——產出比 0.70／0.80／1.15、持有比 0.67／0.85／1.12（採用，v21.7 W1）；每 token 成本比基準約 1.06（取整後輸入值 0.85／0.80；查證原始值 0.847／0.807＝1.05；現值 1.29）",
    "AMD 規格與自述（MI455X、Helios）、CLSA、Futurum；SemiAnalysis 只作佐證；NonNV C 段查證 " + EVD + "（PR #174）",
    "Interested-party（AMD 規格與自述）＋Analogy（CLSA、Futurum；SemiAnalysis 只作佐證）",
    "NonNV B5:G5；Interface 第 369、374 列（IF_NonNV_Out_AMD／IF_NonNV_Hold_AMD）",
    _OLD_AMD,
    "產出比 0.70／0.80／1.15、持有比 0.67／0.85／1.12（採用）；每 token 成本比基準約 1.06（取整後輸入值 0.85／0.80；查證原始值 0.847／0.807＝1.05）",
    "採用（PR #174 Q4、Q5 預設，" + V + " W1）", V,
    "價格口徑：分析師絕對價與 TK 買方取得價不同口徑，持有比改用同口徑（每架 $5.25M 下界／$7.3M＝AMD 自述隱含比 0.865 × VR200／$10.5M＝Futurum 同源比值 1.25 × VR200）；"
    "替代：同源比值全區間 1.12／1.23／1.35、只換產出比（成本比約 1.13）、兩錨點 f（產出比 0.74）｜" + WO,
    "已處理", DASH, DASH, "IF_NonNV_Out_AMD、IF_NonNV_Hold_AMD（值改變）；其餘契約名稱不變", 3, "利害關係方",
    "採用（產出比與持有比都換）｜" + V + " 新增列"]
E411 = [
    "E411", DATE,
    "B 段 Trn3 持有比口徑修正——0.55 → 0.61（KGI 同源比值）",
    "KGI 同報 VR200 與 GB300 機架價；NonNV B 段查證 docs/evidence/20261025_NonNV_TPU_Trainium/（PR #172）與 " + EVD + " README §10（PR #174）",
    "Analogy（KGI 同源比值）",
    "NonNV F7；Interface D376（IF_NonNV_Hold_Trn 基準）",
    "NonNV 第 7 列 Trainium3 持有比基準 0.55（v21.6，KGI 絕對價 ÷ TK 價，輕度口徑混用）",
    "持有比基準 0.61（低端 0.45、高端 0.70 不變）",
    "採用（PR #174 §10 Project 預設：基準取中點 0.61，" + V + " W1）", V,
    "KGI 同報 VR200『超過 700 萬美元』、GB300 $3.9–4.0M，口徑較 TK 低約 9–20%；Trn3 每架約 $1.6–1.8M → 0.593–0.631 取中點 0.61｜" + WO,
    "已處理", DASH, DASH, "IF_NonNV_Hold_Trn（基準 1 格值改變）", 3, "利害關係方",
    "採用（口徑修正）｜" + V + " 新增列"]


def evidence_append(wb):
    ws = wb["DB_Evidence"]
    n = 0
    for prev, row in (("E409", E410), ("E410", E411)):
        if any(ws.cell(r, 1).value == row[0] for r in range(5, ws.max_row + 1)):
            continue
        r = max(rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")) + 1
        assert ws.cell(r - 1, 1).value == prev, (ws.cell(r - 1, 1).value, prev)
        for i, v in enumerate(row):
            put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 4, 5, 6, 7, 8, 10, 14, 17))
        n += 1
    return n


# ---------------------------------------------------------------- W3 Decisions
def _decisions():
    return [[DEC_ID, V, "v21.7 NonNV AMD 改為推導值、Trainium3 持有比改同源價、過時說明更新",
             "W1：NonNV B5:D5（AMD 產出比）0.5／0.7／1.0 → 0.70／0.80／1.15；E5:G5（AMD 持有比）0.8／0.9／1.0 → 0.67／0.85／1.12（依每 kW 峰值算力比 × 效率 f，持有比依 TK 買方取得價同口徑；PR #174；Assumed）；"
             "F7（Trainium3 持有比基準）0.55 → 0.61（KGI 同源比值，低端 0.45、高端 0.70 不變；PR #174 §10）。"
             "W2：NonNV L5、L7、L8、L9、A2 與 Interface A368 說明更新，原文逐字保留於「｜原：」之後（修訂九）。W3：DB_Evidence E410／E411；README B5 加 v21.7 段。"
             "契約名稱的值改變只有 IF_NonNV_Out_AMD、IF_NonNV_Hold_AMD（各 3 格）、IF_NonNV_Hold_Trn（基準 1 格）與版本字串；契約不升版（v0.27）。TPU、Trn 產出比仍待 Q1，不改。Andy 原話與依據：" + DEC_BY,
             DATE, DEC_BY, "v21.7 已建",
             "NonNV B5:G5、F7；NonNV L5／L7／L8／L9／A2；Interface 第 369、374、376 列與 A368；DB_Evidence E410／E411；README B5；Decisions SL64-01",
             f"{WO} W1–W3", "否"]]


def apply(wb, append_text):
    log = {"text_fixes": text_fixes(wb)}
    log["evidence"] = evidence_append(wb)
    ws = wb["Decisions"]
    if not any(ws.cell(r, 1).value == DEC_ID for r in range(5, ws.max_row + 1)):
        log["decisions"] = v1410.decisions_append(wb, _decisions())
    return log


# ---------------------------------------------------------------- README
_prev = v2160.README_VERSION.split("_Tokenomics_", 1)[1]
README_VERSION = (VERSION + "（次版，底版 v21.6：" + WO + "——NonNV AMD 產出比 0.5／0.7／1.0 → 0.70／0.80／1.15、持有比 0.8／0.9／1.0 → 0.67／0.85／1.12（PR #174；Assumed）；"
                  "Trainium3 持有比基準 0.55 → 0.61（KGI 同源比值）；"
                  "契約名稱的值改變的只有 IF_NonNV_Out_AMD、IF_NonNV_Hold_AMD（各 3 格）、IF_NonNV_Hold_Trn（基準 1 格）與版本字串；其餘不變；GOV 三項計數 0／134／182 不變。"
                  "NonNV L5／L7／L8／L9／A2 與 Interface A368 說明更新（原文保留於「｜原：」後）；DB_Evidence E410／E411；Decisions SL64-01。）以下為 " + _prev)
README_TITLE = V + v2160.README_TITLE[len(v2160.V):]
