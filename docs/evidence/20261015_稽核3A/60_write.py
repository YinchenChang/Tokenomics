"""寫明細 docs/reports/20261015_稽核3A_明細.xlsx（逐列檢查、傳遞強度、高嚴重度、L9 物理量）。"""
import json,openpyxl
from openpyxl.styles import Font,Alignment,PatternFill
from openpyxl.utils import get_column_letter as GL
E='docs/evidence/20261015_稽核3A/'
J=json.load(open(E+'50_逐列判讀.json'));S3=json.load(open(E+'30_s3.json'));A=json.load(open(E+'33_alt.json'));L9=json.load(open(E+'31_l9.json'))
wb=openpyxl.Workbook();wb.remove(wb.active)
def sheet(name,rule,head,rows,widths):
    ws=wb.create_sheet(name);ws.append([rule]);ws['A1'].font=Font(bold=True);ws['A1'].alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[1].height=110
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=min(len(head),12))
    ws.append(head)
    for cc in ws[2]: cc.font=Font(bold=True);cc.fill=PatternFill('solid',fgColor='FFDDEBF7');cc.alignment=Alignment(wrap_text=True)
    for r in rows: ws.append(list(r))
    ws.auto_filter.ref=f'A2:{GL(len(head))}{len(rows)+2}';ws.freeze_panes='C3'
    for i,w in enumerate(widths,1): ws.column_dimensions[GL(i)].width=w
    for row in ws.iter_rows(min_row=3):
        for cc in row: cc.alignment=Alignment(wrap_text=True,vertical='top')
rows=[[o['sheet'],o['row'],o['label'],o['unit'],o['nform'],o['npat'],o['rep_cell'],str(o['rep_f'])[:300],round(o['rep_v'],6) if isinstance(o['rep_v'],float) else o['rep_v'],
       o['L1'],o['L2'],o['L3'],o['L4'],o['L5'],o['L6'],o['L7'],o['L8'],o['L9'],o['upY'],o['directY'],o['users'],o['gids'],o['sev']] for o in J]
sheet('逐列檢查','規則：第 0 段「C3_核心推導列」扣除「C7_m路徑」後屬 3A 工作表的 510 個頁＋列，逐列檢查 L1–L9。「—」＝沒有發現問題（不等於已證明正確）。L3 以 10_mech.py 機械檢查（世代×層級／成本頁之間同欄對位 0 處不符；Interface 層級常數 3*(世代−1)+k 對列標籤 0 處不符）；L7 以 20_l7_closure.py 在 4 個端點情境重算（文字或 0）；L8「直接」＝本列公式直接引用第 1 段 Y1–Y9 輸入格，「上游」＝依賴圖上游任一處含 Y 類輸入。G 編號對應報告第 3 節。嚴重度：高＝直接作用在下游契約名稱的口徑或閉合問題；中＝其他 L1–L9；低＝結構或顯示；資訊＝只有 L8。只標問題，不判定對錯、不提修法。',
 ['頁','列','列標籤','單位','公式格數','同列公式結構數（R1C1）','代表格','代表公式','代表值','L1 量綱','L2 口徑','L3 引用對位','L4 重複／漏算','L5 恆等式與閉合','L6 倒推','L7 條件分支','L8 與 Y1–Y9 的連結','L9 物理合理性','上游 Y 類','直接 Y 類','非治理頁下游格數','G 編號','嚴重度'],rows,
 [12,6,34,10,7,7,7,50,12,25,40,25,40,40,30,40,30,35,14,10,8,10,10])
# 傳遞強度
TR=[]
def fmt(x): return '—' if x is None else f'{x:+.1f}%'
TR.append(['G01','由上而下攤提閉合 0.82–0.88（Amortize 39）','IF_AmortTD_*、IF_AmortDefault_*、IF_AmortRev_*、IF_FullCostDefault_*','偏低',
  '；'.join(f"{x['gen'].split()[0]} 閉合 {x['ratio']:.3f}（閉合到訓練 GW × 持有成本時攤提 +{x['amort_up_pct']:.0f}%）" for x in A['G01']),'結構量（不擾動）；χ＝0 時閉合 1.00（Python 與 LibreOffice 各 1 次）'])
g2=[x for x in A['G02'] if x['gen'].startswith(('Hopper','VR200'))]
TR.append(['G02','自下而上攤提分母含免費 token（Amortize 13／22／32）','IF_AmortBU_*、IF_FullCost_*、IF_AmortDefault_*、IF_FullCostDefault_*','偏低（以付費 token 為分母時）',
  '；'.join(f"{x['gen'].split()[0]} {x['tier']}：分母比 {x['ratio']:.2f}，IF_FullCost {(x['FullCost_paid']/x['FullCost']-1)*100:+.0f}%、IF_FullCostDefault {(x['FullCostDefault_paid']/x['FullCostDefault']-1)*100:+.0f}%" for x in g2),'替代讀法（只重組工作簿數值；33_alt.py）'])
p=S3['Prod_over_Util']
TR.append(['G03','_Prod 層 token 線性 0.85、營收重解 SLO','IFW_TokGW_*_Prod、IFW_RevGW*_Prod','token 層偏高（相對營收層）',
  '；'.join(f"{nm}：{ '、'.join(f'{k.split()[0]} {v:.3f}' for k,v in d.items())}" for nm,d in p.items() if nm in('IFW_RevGWFleet','IFW_RevGW_Astra','IFW_RevGW_Sol','IFW_TokGW_Sol')),'結構量（_Prod ÷ _Util）'])
TR.append(['G04','兩組 _Life（IF_ 不含折減、IFW_ 含重解折減）','IF_RevGW_*_Life、IF_RevGWFleet_Life 對 IFW_*_Life','IF_ 組偏高',
  '；'.join(f"{nm.split('÷')[0]}：{'、'.join(f'{k.split()[0]} {v:.3f}' for k,v in d.items())}" for nm,d in S3['Life_IF_vs_IFW'].items()),'結構量（IFW ÷ IF）'])
TR.append(['G06','組合旗標使可達層級也變文字；單一層級列變 0','IF_AmortBU_*、IF_AmortDefault_*、IF_FullCost*、IFW_RevGWFleet*；IFW_TokGW_Astra_*、IFW_RevGW_Astra_* 等','文字／0',
  '；'.join(f"{sc}：變文字 {len(o['text'])} 個名稱、變 0 {len(o['zero'])} 個名稱" for sc,o in A['G06'].items()),'4 個 L7 情境（20_l7_closure.py）'])
TR.append(['G14','縮放倍數 2.2–12.0；預設攤提總額由 B4_Serve 決定','IF_AmortDefault_*、IF_FullCostDefault_*','視 B4_Serve',
  '見下方 A02 列（B4_Serve 0.3／0.6）','端點擾動'])
for r in S3['perturb']:
    sw=r['swing']
    TR.append(['端點',r['item'],'、'.join(k for k,s in sw.items() if any(v is not None and abs(v)>0.05 for v in s)) or '（無）','低端／高端',
      '；'.join(f"{k} {fmt(s[0])}／{fmt(s[1])}" for k,s in sw.items() if any(v is not None and abs(v)>0.05 for v in s)) or '所有列出的契約名稱變化 < 0.05%',f"單一輸入端點 {r['endpoints']}；20_eval.py"])
sheet('傳遞強度','規則：嚴重度「高」的問題（G01–G06、G14、G15），估計對下游契約名稱（docs/plan/Tokenomics_downstream_contract.md：IF_／IFW_／L1_）的影響方向與大小。結構量直接由工作簿數值計算；端點擾動以第 2 段 20_eval.py（工作簿公式原文求值；3A 5,628 格對快取值 0 不符；LibreOffice 1 次 27 格 0 不符）單一輸入取 Gov_Map 區間端點、其他不變重算。「替代讀法」只重組工作簿數值，不是活頁簿機制，只用來顯示大小。G05（η_p）、G15（B4_Serve、B4_Free）的大小見端點列。',
 ['G 編號','問題','受影響的契約名稱','方向','大小','方法'],TR,[8,36,50,14,110,30])
hi=[[o['sheet'],o['row'],o['label'],o['gids'],o['L2'],o['L4'],o['L5'],o['L6'],o['L7'],o['L9'],o['L8']] for o in J if o['sev']=='高']
sheet('高嚴重度','規則：「逐列檢查」頁嚴重度＝高的列（直接作用在下游契約名稱的口徑或閉合問題）。只標問題，不判定對錯、不提修法。',
 ['頁','列','列標籤','G 編號','L2','L4','L5','L6','L7','L9','L8'],hi,[12,6,40,12,45,45,45,30,40,30,30])
L=[[x['gen'].split('（')[0],x['tier'].split('（')[0],round(x['B_star'],2),round(x['decode_tok_s_gpu'],1),round(x['user_speed'],1),round(x['decode_MFU_incl_draft'],4),round(x['prefill_MFU'],3),round(x['decode_HBM_BW_util'],3),round(x['TTFT_s'],3),x['TTFT_ok'],round(x['tok_per_J'],2),round(x['J_per_decode_tok'],3),round(x['closure_ratio'],3),round(x['weights_share_HBM'],3)] for x in L9['L9']]
sheet('L9 物理量','規則：Perf 15 欄（世代 × 成本情境欄依層級）在基準下的中間物理量。decode 有效 MFU＝每 GPU decode tok/s × decode FLOPs/token ×（草稿數＋1）÷ 每步接受 token ÷ 計算用峰值（含草稿驗證）；prefill MFU＝prefill tok/s × prefill FLOPs/token ÷ 峰值；decode HBM 頻寬利用率＝每 GPU decode tok/s × 每接受 token HBM 讀取 ÷ HBM 頻寬。全部由工作簿數值計算（31_l9_closure_cause.py），不新增機制。只顯示，不判定。',
 ['世代','層級','B*（序列/GPU）','decode tok/s/GPU','每用戶速度','decode 有效 MFU','prefill MFU','decode HBM 頻寬利用率','TTFT 計算（s）','TTFT 符合','tok/J','J/decode token','能量閉合比','權重占可用 HBM'],L,[14,8,10,12,10,10,10,12,10,8,8,10,10,10])
wb.save('docs/reports/20261015_稽核3A_明細.xlsx');print('ok',len(rows),len(hi),len(TR))
