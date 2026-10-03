# Block 3 (v5.5): Tech_Registry, Train_In, Perf_Batch, Training, Sens_Train
from common import *
from perf import write_perf
from outputs import COLS15
from openpyxl.workbook.defined_name import DefinedName

# ---------------------------------------------------------------- Tech_Registry
HOOKS = [  # code, meaning, where it acts
    ("H_ETAD", "decode η_d 倍數", "Perf／Perf_Batch：etadm 列"),
    ("H_TL", "每層延遲倍數", "Perf／Perf_Batch：tlm 列"),
    ("H_FLOP", "每 token FLOPs 倍數（推論與訓練）", "Perf／Perf_Batch：hflop 列；Training：hflop 列"),
    ("H_KV", "KV bytes/token 倍數", "Perf／Perf_Batch：hkv 列"),
    ("H_WB", "權重 bytes/param 倍數", "Perf／Perf_Batch：hwb 列"),
    ("H_TPK", "訓練峰值倍數", "Training：htpk 列"),
    ("H_MFU", "訓練 MFU 倍數", "Training：hmfu 列"),
    ("H_TOK", "達同等預訓練品質所需 token 倍數", "Training：htok 列"),
    ("H_ROLL", "rollout 效率倍數", "Training：hroll 列"),
    ("H_CAP", "能力增量（Block 4 占位，尚無作用）", "Block 4"),
    ("H_HAR", "harness（v5.9 起不作倍數；T12 的開關 × 採用比例＝Har_In 混合權重 w）", "Har_In → Workload B 節、Harness"),
]

# id, tech, hook, acts-on, lo, base, hi, sel, status, labs, override, in-base, adopt, switch, tag, source/note, trigger
ENTRIES = [
 ("T01", "MTP 推測解碼", "—", "Serving!C6:C8（N、α）", 1, 1, 1, 2, "主流", 3, "", "是", 1, 1, "Verified",
  "DeepSeek V3／V4、GLM-4.5、MiMo 公開採用；已在 Serving 基準", "—"),
 ("T02", "壓縮／稀疏注意力（CSA/HCA 類）", "—", "Arch 列 16–19", 1, 1, 1, 2, "主流", 3, "", "是", 1, 1, "Verified",
  "DeepSeek V4 CSA/HCA、DeepSeek／GLM 稀疏注意力、Qwen3-Next 混合注意力；Luna／Sol 基準已含（S27、S35）", "—"),
 ("T03", "FP4 權重推論（NVFP4／MXFP4）", "—", "Spec_Rack 專家 bytes/param", 1, 1, 1, 2, "主流", 2, "", "是", 1, 1, "Verified",
  "DeepSeek V4 原生 FP4 專家權重、OpenAI gpt-oss MXFP4；已在基準", "—"),
 ("T04", "FP4 量化感知訓練（後訓練）", "—", "Spec_Rack 專家 bytes/param", 1, 1, 1, 2, "主流", 2, "", "是", 1, 1, "Interested-party",
  "DeepSeek V4：主權重量化到 FP4 再反量化到 FP8 計算，沿用 FP8 訓練框架（S35）", "—"),
 ("T05", "On-policy 蒸餾（多專家整合）", "—", "Training G 節", 1, 1, 1, 2, "主流", 2, "", "是", 1, 1, "Interested-party",
  "DeepSeek V4（專家 SFT＋GRPO 後以 on-policy 蒸餾整合，S35）、Qwen3；已在 Training 基準", "—"),
 ("T06", "Muon 優化器", "H_TOK", "預訓練所需 token", 0.7, 0.85, 1, 2, "早期採用", 2, "", "是", 1, 0, "Interested-party",
  "Moonshot Kimi K2、DeepSeek V4 採用；Luna／Sol 的 token 數為實際值，已隱含其效果，故不套倍數", "第三家前沿實驗室採用，或公開同品質 token 節省的對照實驗"),
 ("T07", "NVFP4 預訓練", "H_TPK", "訓練峰值", 1.5, 2, 3, 2, "早期採用", 1, "", "否", 1, 0, "Interested-party",
  "倍數＝NVFP4 訓練峰值 ÷ FP8（VR 35/17.5＝2；GB300 15/5＝3 [Assumed]）。NVIDIA 公開 12B 模型 10T token NVFP4 預訓練；前沿實驗室未見公開採用。MFU 可能同步下降，未計", "任一前沿實驗室公開以 FP4 完成主預訓練"),
 ("T08", "Looped transformer（權重共享迴圈）— 計算量", "H_FLOP", "每 token FLOPs", 1.5, 2, 4, 2, "研究", 0, "", "否", 1, 0, "Assumed",
  "同一組權重迴圈 k 次：FLOPs × k、權重 bytes × 1；能力增量於 Block 4（H_CAP）處理", "前沿模型卡或技術報告揭露迴圈深度"),
 ("T09", "Looped transformer — 每層延遲", "H_TL", "每步逐層延遲", 1.5, 2, 4, 2, "研究", 0, "", "否", 1, 0, "Assumed",
  "迴圈使有效層數 × k，逐層延遲同比增加；與 T08 同時開關", "同 T08"),
 ("T10", "非同步 RL（rollout 與訓練解耦、部分 rollout）", "H_ROLL", "rollout 效率", 1.2, 1.4, 1.6, 2, "主流", 4, "", "是", 1, 0, "Interested-party",
  "旗艦模型公開採用：DeepSeek V4.1（幾乎全部 RL 與 OPD，S45）、Zhipu GLM-5（slime，S46）、Moonshot Kimi-Researcher（完全非同步 rollout，S47）、Meta Llama 3（LlamaRL，S48）。閉源四家未找到披露。v5.6 起併入基準：Train_In rollout 效率 0.85（Andy 2026-10-01 決定 (a)）", "閉源實驗室披露同步做法，或前沿規模對照實驗顯示 off-policy 偏差抵銷吞吐增益"),
 ("T11", "KV 快取壓縮（FP4 KV 等）", "H_KV", "KV bytes/token", 0.5, 0.5, 0.75, 2, "早期採用", 1, "", "否", 1, 0, "Assumed",
  "FP8 → FP4 KV：bytes × 0.5；準確率損失待查", "第二家實驗室在生產服務公開採用"),
 ("T12", "Harness（代理框架）", "H_HAR", "Har_In 選定檔案（混合權重 w＝開關 × 採用比例）", 1, 1, 1, 2, "追蹤中", 0, "", "否", 1, 0, "Assumed",
  "v5.9：harness 以參數組表達（輪數、思考保留、思考量、歷史壓縮、快取命中、子代理、狀態保留）加成功率（L1、L2）；倍數欄不使用。基準不開（L3）。"
  "ARC-AGI-3（ARC Prize 自測）：GPT-6 Astra 標準 harness 最高推理 62.7%、$26,098 → OpenAI Provider Adapter 同檔 98.6%、$17,332（S58）。主流判定（J14）尚未評估", 
  "中立方（ARC Prize、METR、Artificial Analysis）在同條件下測得第二個 harness 對照"),
]

def tech_registry(wb):
    ws = wb.create_sheet("Tech_Registry")
    title(ws, "Tech_Registry — 新技術登錄與掛鉤（每列＝技術 × 作用物理量；基準值只在『已在基準＝否』且開關＝1 時套用）",
          "有效倍數＝1＋開關 × 採用比例 ×（所選倍數−1）；已在基準者恆為 1，避免重複計算。主流判定（J14）：至少兩家實驗室公開採用；L 欄可由 Andy 覆寫")
    hdr = ["ID", "技術", "掛鉤代碼", "作用物理量", "倍數 低", "倍數 基準", "倍數 高", "情境（1／2／3）", "狀態",
           "公開採用實驗室數", "主流判定（J14）", "Andy 覆寫", "已在基準", "採用比例", "開關（0／1）", "有效倍數",
           "證據標記", "來源／說明", "下次檢查觸發", "一致性檢查"]
    widths = [6, 30, 10, 22, 8, 8, 8, 9, 10, 9, 10, 9, 8, 8, 8, 9, 14, 70, 36, 26]
    for i, (h, w) in enumerate(zip(hdr, widths)):
        put(ws, f"{L(i+1)}4", h, F_BOLD, wrap=True); ws.column_dimensions[L(i+1)].width = w
    ws.row_dimensions[4].height = 30
    r0 = 5
    for i, e in enumerate(ENTRIES):
        r = r0 + i
        (eid, tech, hook, acts, lo, ba, hi, sel, st, labs, ovr, inb, adopt, sw, tag, src, trig) = e
        for c, v in zip("ABCD", (eid, tech, hook, acts)): put(ws, f"{c}{r}", v, wrap=(c == "B"))
        for c, v in zip("EFG", (lo, ba, hi)): put(ws, f"{c}{r}", v, fmt="0.00")
        put(ws, f"H{r}", sel, fmt="0"); put(ws, f"I{r}", st, F_IN); put(ws, f"J{r}", labs, fmt="0")
        put(ws, f"K{r}", f'=IF(J{r}>=CST_MainMin,"主流","非主流")')
        put(ws, f"L{r}", ovr if ovr else None, F_IN)
        put(ws, f"M{r}", inb, F_IN); put(ws, f"N{r}", adopt, fmt="0%"); put(ws, f"O{r}", sw, fmt="0")
        put(ws, f"P{r}", f'=IF(M{r}="是",1,1+O{r}*N{r}*(CHOOSE(H{r},E{r},F{r},G{r})-1))', fmt="0.00", fill=FILL_KEY)
        put(ws, f"Q{r}", tag, F_NOTE); put(ws, f"R{r}", src, F_NOTE, wrap=True); put(ws, f"S{r}", trig, F_NOTE, wrap=True)
        put(ws, f"T{r}", f'=IF(AND(IF(L{r}="",K{r},L{r})="主流",M{r}="否"),"主流但未入基準：須 Andy 判定",'
                         f'IF(AND(IF(L{r}="",K{r},L{r})="非主流",M{r}="是",I{r}<>"早期採用"),"非主流卻在基準","一致"))', wrap=True)
        ws.row_dimensions[r].height = 42
    r1 = r0 + len(ENTRIES) - 1
    # v5.12 (A): J14 threshold ("at least two labs") moved out of the K-column formula into a named input (value unchanged)
    put(ws, f"A{r1+1}", "門檻"); put(ws, f"B{r1+1}", "主流判定門檻：公開採用實驗室數 ≥ 本格（J14；K 欄共用）", wrap=True)
    put(ws, f"J{r1+1}", 2, fmt="0"); put(ws, f"Q{r1+1}", "Decision", F_NOTE)
    put(ws, f"R{r1+1}", "J14：至少兩家實驗室公開採用即為主流（CST_MainMin）", F_NOTE, wrap=True)
    wb.defined_names["CST_MainMin"] = DefinedName("CST_MainMin", attr_text=f"Tech_Registry!$J${r1+1}")
    r = r1 + 2
    section(ws, r, "掛鉤彙總（同一掛鉤多條目時取乘積；Perf、Perf_Batch、Training 連結本表 E 欄）", 20); r += 1
    for c, v in zip("ABCDE", ["代碼", "意義", "", "作用位置", "倍數"]): put(ws, f"{c}{r}", v, F_BOLD)
    r += 1
    TR = {}
    for code, meaning, where in HOOKS:
        put(ws, f"A{r}", code, F_BOLD); put(ws, f"B{r}", meaning); put(ws, f"D{r}", where, F_NOTE, wrap=True)
        put(ws, f"E{r}", f'=EXP(SUMPRODUCT(($C${r0}:$C${r1}="{code}")*LN($P${r0}:$P${r1})))', fmt="0.000", fill=FILL_KEY)
        TR[code] = f"Tech_Registry!$E${r}"; r += 1
    r += 1
    put(ws, f"A{r}", "注意", F_BOLD)
    put(ws, f"B{r}", "掛鉤作用於全部世代與層級。新條目：在表中插入一列（範圍內），填倍數、狀態與證據；開關預設 0。"
                     "Sens_Perf 的 η_d 與每層延遲情境為覆寫值，不受 Registry 影響。", F_NOTE, wrap=True)
    ws.merge_cells(f"B{r}:R{r}"); ws.row_dimensions[r].height = 30
    ws.freeze_panes = "C5"
    TR["_rows"] = (r0, r1)
    TR["_hooks"] = (r1 + 4, r1 + 3 + len(HOOKS))
    return TR

# ---------------------------------------------------------------- Train_In
ROUT_LUNA, ROUT_SOL, ROUT_ASTRA = 2.36, 2.36, 9.6   # v5.6: 由 v5.5 的 2, 2, 8 重校，維持 VR200 上 RL ÷ 預訓練 GPU 小時（J9 (a)）

def train_in(wb):
    ws = wb.create_sheet("Train_In")
    title(ws, "Train_In — 訓練輸入（藍字＝輸入；Analogy／Assumed 一律附區間）",
          "Block 3 命題：各層級代表模型從預訓練到可發布、以及含研發實驗的整個計畫，需要多少 GPU 小時、美元與 1 GW 年。決策 J7–J14 見 README")
    for c, w in zip("ABCDEFGH", [46, 12, 14, 14, 14, 14, 14, 90]): ws.column_dimensions[c].width = w
    T = {}
    put(ws, "A4", "全域設定", F_BOLD); put(ws, "B4", "單位", F_BOLD); put(ws, "C4", "值", F_BOLD); put(ws, "D4", "標記", F_BOLD); put(ws, "H4", "說明", F_BOLD)
    g = [
      ("prec", "訓練精度（MFU 分母）", "", "FP8 dense", None, "Decision", "J7：基準 FP8（DeepSeek V4 以 FP8 計算＋FP4 QAT，S35）；NVFP4 預訓練列於 Tech_Registry T07"),
      ("mfub", "預訓練 MFU 基準（FP8 分母；GB200／GB300）", "%", 0.25, "0%", "Analogy",
       "DeepSeek V3 推導 H800 FP8 口徑約 20%（Checks）；Llama 3 405B BF16 口徑 38–43%＝FP8 口徑約 19–22%（S37，待查）；NVL72 域大、EP 通訊較佳取 0.25。區間 0.15–0.40"),
      ("gp", "goodput（扣除故障、重啟、checkpoint 的有效時間比）", "%", 0.90, "0%", "Analogy", "Meta Llama 3：有效訓練時間 >90%（S37，待查）；區間 0.80–0.95"),
      ("sftk", "SFT MFU 係數（× 預訓練 MFU）", "x", 0.8, "0.00", "Assumed", "較短批次、較多變長序列；區間 0.6–1.0"),
      ("rlk", "RL trainer MFU 係數（× 預訓練 MFU）", "x", 0.6, "0.00", "Assumed", "長序列、小批次、與 rollout 交替；區間 0.4–0.9。蒸餾學生更新沿用"),
      ("ref", "參考模型前向（KL 懲罰；0＝無、1＝有）", "選擇", 0, "0", "Assumed", "GRPO 原版含 KL；DAPO 等後續做法移除。取 0，區間 0–1"),
      ("reff", "rollout 效率（長尾等待、權重同步、閒置）", "%", 0.85, "0%", "Analogy",
       "v5.6：非同步 RL 已屬主流（Tech_Registry T10）。各家報告非同步增益 1.5–2.7 倍（S45–S49）；同步約 0.4–0.5 → 非同步 0.85，區間 0.7–0.95。v5.5 為 0.6（混合）"),
      ("G", "GRPO 每題取樣數（共用 prompt，prefill 只算一次）", "個", 16, "0", "Analogy", "DeepSeekMath／R1 GRPO 群組 16；區間 8–64"),
      ("rdm", "研發倍數（研發總 GPU 小時 ÷ 最終訓練）", "x", 8, "0.0", "Analogy",
       "J10。Epoch：最終訓練占研發支出 OpenAI 9.6%、MiniMax 22.6%、Z.ai 12.3%（S38）＝4.4–10.4 倍；取 8。以支出為口徑，已含實驗的低利用率，故乘在 GPU 小時上"),
      ("rdmode", "研發歸屬", "", "家族合計乘一次，依最終訓練 GPU 小時比例分攤", None, "Decision",
       "J10。比例分攤下，各層級分得＝該層級最終訓練 × 倍數；Training J 節另列家族合計"),
      ("gtier", "合成資料生成層級（1 Luna／2 Sol／3 Astra）", "選擇", 3, "0", "Assumed", "以頂層模型生成；區間 2–3"),
      ("gdef", "下游預設訓練世代（世代索引）", "索引", 4, "0", "Decision", "J13：Andy 2026-09-30 決定 VR200（索引 4）；具名範圍 IF_TrainGenDefault"),
      ("galt", "並列訓練世代（世代索引）", "索引", 3, "0", "Decision", "J13：GB300（索引 3）並列；具名範圍 IF_TrainGenAlt"),
    ]
    r = 5
    for key, lab, unit, v, fmt, tag, note in g:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", unit); put(ws, f"C{r}", v, fmt=fmt, font=F_IN)
        put(ws, f"D{r}", tag, F_NOTE); put(ws, f"H{r}", note, F_NOTE, wrap=True)
        T[key] = r; r += 1
    for k in ("gdef", "galt"):
        put(ws, f"E{T[k]}", f"=INDEX(Spec_Rack!$C$4:$G$4,$C${T[k]})", F_HLINK)
    r += 1
    section(ws, r, "各世代", 8); r += 1
    put(ws, f"A{r}", "世代", F_BOLD)
    for c in "CDEFG": put(ws, f"{c}{r}", f"=Spec_Rack!{c}4", F_HLINK)
    r += 1
    T["mfug"] = r
    put(ws, f"A{r}", "預訓練 MFU 世代倍數"); put(ws, f"B{r}", "x")
    for c, v in zip("CDEFG", [0.8, 1, 1, 0.9, 0.9]): put(ws, f"{c}{r}", v, fmt="0.00")
    put(ws, f"H{r}", "Hopper 0.8 使 MFU＝20%，對應 V3 推導值 [Derived]；VR200 峰值 3.5 倍而 HBM 頻寬 2.75 倍，取 0.9 [Assumed，區間 0.7–1.1]；RU 同 VR [Assumed]", F_NOTE, wrap=True)
    r += 2
    section(ws, r, "各層級", 8); r += 1
    put(ws, f"A{r}", "項目", F_BOLD); put(ws, f"B{r}", "單位", F_BOLD)
    for c, t in zip("CDE", "CDE"): put(ws, f"{c}{r}", f"=Arch!{t}4", F_HLINK)
    put(ws, f"H{r}", "說明", F_BOLD); r += 1
    t = [
      ("ptok", "預訓練 token", "T", [32, 33, 60], "#,##0.0", "Verified：V4-Flash 32T、V4-Pro 33T（S35）；Astra [Assumed，區間 30–100T；J8 以 Arch 一致為基準，前沿錨點見 Checks]"),
      ("pseq", "預訓練序列長（dense 注意力）", "tok", [4096, 4096, 8192], "#,##0", "V4：4K → 16K 漸增，其間以 dense 注意力訓練（S35）；多數 token 在 4K [Assumed，區間 4K–16K]"),
      ("ltok", "長上下文延伸 token", "T", [1, 1, 2], "#,##0.0", "Assumed，區間 0.3–3T（V4 未揭露）"),
      ("lseq", "長上下文延伸序列長（稀疏注意力路徑）", "tok", [131072, 131072, 131072], "#,##0", "Assumed"),
      ("stok", "SFT token", "T", [0.05, 0.05, 0.1], "#,##0.00", "Assumed，區間 0.01–0.3T"),
      ("sseq", "SFT 序列長", "tok", [16384, 16384, 16384], "#,##0", "Assumed"),
      ("bs", "批次推論速度下限（rollout、合成、評測）", "tok/s", [20, 20, 20], "#,##0", "取代互動 SLO；過低時長尾拉長牆鐘時間 [Assumed，區間 10–40]"),
      ("risl", "rollout 參考輸入（prompt＋工具結果）", "tok", [4096, 8192, 16384], "#,##0", "Assumed"),
      ("rosl", "rollout 參考輸出（含思考）", "tok", [8192, 16384, 32768], "#,##0", "Assumed：推理與代理任務的長輸出"),
      ("rout", "RL rollout 輸出 token（含各領域專家）", "T", [ROUT_LUNA, ROUT_SOL, ROUT_ASTRA], "#,##0.00",
       "J9：由下而上輸入；基準值校到 VR200 上 RL ÷ 預訓練 GPU 小時約 Luna／Sol 0.3、Astra 1.0（結果見 Training rlH 列與 Checks）；v5.6 因 rollout 效率 0.6→0.85 由 2／2／8 上調為 2.36／2.36／9.6，GPU 小時不變（Andy 決定 (a)）。外部點：R1 約 5%、Grok 4 約 1 倍（S40、S42；GPU 小時或預算口徑）。區間：Luna／Sol 0.3–6T、Astra 2.5–25T"),
      ("dtok", "On-policy 蒸餾：學生 rollout 輸出 token", "T", [0.2, 0.2, 0.2], "#,##0.00", "Assumed，區間 0.05–1T"),
      ("dtea", "蒸餾教師層級（1／2／3）", "選擇", [1, 2, 3], "0", "V4 做法：同尺寸各領域專家為教師 [Interested-party，S35]"),
      ("syn", "合成資料生成 token", "T", [1, 2, 3], "#,##0.0", "Assumed，區間 0.5–10T"),
      ("ev", "評測 token", "T", [0.02, 0.05, 0.1], "#,##0.00", "Assumed"),
    ]
    for key, lab, unit, vals, fmt, note in t:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", unit)
        for c, v in zip("CDE", vals): put(ws, f"{c}{r}", v, fmt=fmt)
        put(ws, f"H{r}", note, F_NOTE, wrap=True)
        T[key] = r; r += 1
    for c in "CDE": ws[f"{c}{T['rout']}"].fill = FILL_KEY
    r += 1
    section(ws, r, "外部錨點（Checks 用；不進推導）", 8); r += 1
    a = [
      ("v3tok", "DeepSeek V3 預訓練 token", "T", 14.8, "#,##0.0", "Interested-party：DeepSeek V3 技術報告（S44）"),
      ("v3h", "DeepSeek V3 預訓練 H800 GPU 小時", "M", 2.664, "#,##0.000", "S44（全部 2.788M，含長上下文 0.119M、後訓練 0.005M）"),
      ("v3seq", "DeepSeek V3 預訓練序列長", "tok", 4096, "#,##0", "S44"),
      ("r1rl", "DeepSeek R1 RL GPU 小時 ÷ V3 預訓練", "x", 0.055, "0.0%", "147K ÷ 2.664M（S40，Lambert 轉述；GPU 小時口徑）"),
      ("grok", "Grok 4 RL ÷ 預訓練（宣稱）", "x", 1.0, "0.0", "Interested-party：xAI『預訓練規模』（S42）；口徑不明"),
      ("an1", "前沿預訓練錨點 低", "FLOP", 2e26, "0.0E+00", "Analogy：Epoch 估計（Grok-3 約 4.6e26，S39）；誤差 2–5 倍"),
      ("an2", "前沿預訓練錨點 中", "FLOP", 5e26, "0.0E+00", "同上"),
      ("an3", "前沿預訓練錨點 高", "FLOP", 2e27, "0.0E+00", "同上"),
    ]
    for key, lab, unit, v, fmt, note in a:
        put(ws, f"A{r}", lab); put(ws, f"B{r}", unit); put(ws, f"C{r}", v, fmt=fmt)
        put(ws, f"H{r}", note, F_NOTE, wrap=True); T[key] = r; r += 1
    ws.freeze_panes = "C5"
    return T

# ---------------------------------------------------------------- Perf_Batch
def perf_batch(wb, SP, AR, CAL, TR, TI):
    ws = wb.create_sheet("Perf_Batch")
    title(ws, "Perf_Batch — 批次推論引擎（RL rollout、蒸餾、合成資料、評測）：與 Perf 同一套公式，只換速度下限與參考任務",
          "C–Q 欄＝世代 × 層級基準；R、S 欄＝VR200 以 FP8 rollout 的情境（J12 替代）。綠字＝連結 Train_In（非情境改動）；黃底＝情境改動")
    put(ws, "A4", "世代", F_BOLD); put(ws, "A5", "層級", F_BOLD); put(ws, "A6", "世代索引", F_BOLD); put(ws, "A7", "層級索引", F_BOLD)
    put(ws, "A3", "欄位", F_BOLD)
    cols = COLS15 + ["R", "S"]
    for i, X in enumerate(cols):
        if i < 15: gi, t = i // 3 + 1, i % 3 + 1
        else: gi, t = 4, i - 13
        put(ws, f"{X}6", gi, fmt="0"); put(ws, f"{X}7", t, fmt="0")
        put(ws, f"{X}4", f"=INDEX(Spec_Rack!$C$4:$G$4,{X}6)", F_HLINK)
        put(ws, f"{X}5", f"=INDEX(Arch!$C$4:$E$4,{X}7)", F_HLINK)
        ws.column_dimensions[X].width = 13
    put(ws, "R3", "情境：FP8 rollout", F_BOLD); put(ws, "S3", "情境：FP8 rollout", F_BOLD)
    link = {"s": f"=INDEX(Train_In!$C${TI['bs']}:$E${TI['bs']},{{X}}$7)",
            "isl": f"=INDEX(Train_In!$C${TI['risl']}:$E${TI['risl']},{{X}}$7)",
            "osl": f"=INDEX(Train_In!$C${TI['rosl']}:$E${TI['rosl']},{{X}}$7)"}
    ov = {X: dict(link) for X in COLS15}
    R, r = write_perf(ws, COLS15, SP, AR, CAL, TR, overrides=ov, ov_link=True)
    # scenario columns (FP8 rollout on VR200): re-write the same rows into R, S with extra overrides
    fp8 = {"P": f"=Spec_Rack!$F${SP['tfp8']}", "be": 1, "bn": 1}
    for X in ("R", "S"):
        rr = 9
        from perf import perf_rows
        spec = perf_rows(SP, AR, CAL, TR)
        m = {}
        for item in spec:
            if item[0] == "§": rr += 1; continue
            key, lab, unit, fmt, tpl = item
            m[key] = rr
            if key in link:
                put(ws, f"{X}{rr}", link[key].replace("{X}", X), F_LINK, fmt=fmt)
            elif key in fp8:
                put(ws, f"{X}{rr}", fp8[key], F_LINK if isinstance(fp8[key], str) else F_IN, fmt=fmt, fill=PatternFill("solid", fgColor="FFFFFF00"))
            elif isinstance(tpl, (int, float)):
                put(ws, f"{X}{rr}", tpl, F_IN, fmt=fmt)
            else:
                mm = dict(m); mm["X"] = X
                put(ws, f"{X}{rr}", tpl.format(**mm), fmt=fmt)
            rr += 1
    section(ws, r, "L. 批次推論的有效 MFU（接受 token 的模型 FLOPs ÷ 峰值）", 19); r += 1
    R["mfu4"], R["mfu8"] = r, r + 1
    put(ws, f"A{r}", "decode 有效 MFU（對 Perf 計算用峰值）"); put(ws, f"B{r}", "%")
    put(ws, f"A{r+1}", "decode 有效 MFU（對 FP8 訓練峰值）"); put(ws, f"B{r+1}", "%")
    for X in cols:
        put(ws, f"{X}{r}", f"=IF({X}{R['D']}>0,{X}{R['D']}*{X}{R['Fd']}*1E9/({X}{R['P']}*1E15),0)", fmt="0.0%")
        put(ws, f"{X}{r+1}", f"=IF({X}{R['D']}>0,{X}{R['D']}*{X}{R['Fd']}*1E9/(INDEX(Spec_Rack!$C${SP['tfp8']}:$G${SP['tfp8']},{X}$6)*1E15),0)", fmt="0.0%", fill=FILL_KEY)
    ws.freeze_panes = "C8"
    return R

# ---------------------------------------------------------------- Training engine
def train_rows(SP, AR, TI, PB, TR):
    ti = lambda k: f"=INDEX(Train_In!$C${TI[k]}:$E${TI[k]},{{X}}$7)"
    tg = lambda k: f"=INDEX(Train_In!$C${TI[k]}:$G${TI[k]},{{X}}$6)"
    tgl = lambda k: f"=Train_In!$C${TI[k]}"
    sp = lambda row: f"=INDEX(Spec_Rack!$C${row}:$G${row},{{X}}$6)"
    ar = lambda row: f"=INDEX(Arch!$C${row}:$E${row},{{X}}$7)"
    dc = lambda row: f"=INDEX(DC_Cost!$C${row}:$Q${row},3*({{X}}$6-1)+2)"
    pb = lambda k: f"=INDEX(Perf_Batch!$C${PB[k]}:$Q${PB[k]},3*({{X}}$6-1)+{{X}}$7)"
    pbt = lambda k, tier: f"=INDEX(Perf_Batch!$C${PB[k]}:$Q${PB[k]},3*({{X}}$6-1)+{{X}}{{{tier}}})"
    hk = lambda code: f"={TR[code]}"
    sparse = lambda pos: ("={X}{ff}*IF({X}{cap}=0," + pos + ",MIN(" + pos + ",{X}{cap}))+(1-{X}{ff})*(" + pos + "/{X}{comp}+{X}{win})")
    gpuh = lambda C, k="1": "={X}{" + C + "}/({X}{pk}*1E15*{X}{mfu}*" + k + "*{X}{gp})/3600"
    return [
      ("§", "A. 規格、成本與 Tech_Registry 掛鉤"),
      ("pk8", "FP8 dense 峰值", "PF/GPU", "#,##0.000", sp(SP["tfp8"])),
      ("htpk", "訓練峰值倍數（H_TPK）", "x", "0.00", hk("H_TPK")),
      ("pk", "訓練用峰值（MFU 分母）", "PF/GPU", "#,##0.000", "={X}{pk8}*{X}{htpk}"),
      ("mfub", "預訓練 MFU 基準", "%", "0%", tgl("mfub")),
      ("mfug", "MFU 世代倍數", "x", "0.00", tg("mfug")),
      ("hmfu", "訓練 MFU 倍數（H_MFU）", "x", "0.00", hk("H_MFU")),
      ("mfu", "預訓練 MFU（採用）", "%", "0.0%", "={X}{mfub}*{X}{mfug}*{X}{hmfu}"),
      ("gp", "goodput", "%", "0%", tgl("gp")),
      ("gpus", "每 GW GPU 數", "顆", "#,##0", dc(14) + "/CTL_GW"),   # v5.11 F14: DC_Cost is the facility total
      ("ce", "每 GPU 小時持有成本 — 經濟（基準成本）", "$/GPU-hr", "#,##0.00", dc(58)),
      ("ca", "每 GPU 小時持有成本 — 會計（基準成本）", "$/GPU-hr", "#,##0.00", dc(57)),
      ("§", "B. 架構（連結 Arch）"),
      ("A", "啟用參數", "B", "#,##0", ar(AR["A"])),
      ("attc", "注意力 FLOPs／被注意 token", "FLOP", "#,##0", ar(AR["attc"])),
      ("ff", "全注意力層比例", "%", "0%", ar(AR["ff"])),
      ("cap", "全注意力跨度上限", "tok", "#,##0", ar(AR["cap"])),
      ("comp", "其他層壓縮比", "x", "#,##0", ar(AR["comp"])),
      ("win", "其他層滑動窗", "tok", "#,##0", ar(AR["win"])),
      ("hflop", "FLOPs/token 倍數（H_FLOP）", "x", "0.00", hk("H_FLOP")),
      ("§", "C. 預訓練（反向傳播型；每 token 訓練 FLOPs＝3 × 前向）"),
      ("htok", "所需 token 倍數（H_TOK）", "x", "0.00", hk("H_TOK")),
      ("Dp", "預訓練 token", "T", "#,##0.0", ti("ptok") + "*{X}{htok}"),
      ("Sp", "序列長", "tok", "#,##0", ti("pseq")),
      ("attp", "平均被注意 token（dense 注意力＝平均位置）", "tok", "#,##0", "={X}{Sp}/2"),
      ("Fpt", "訓練 FLOPs/token", "GFLOP", "#,##0.0", "=3*(2*{X}{A}*1E9+{X}{attc}*{X}{attp})*{X}{hflop}/1E9"),
      ("Cp", "預訓練 FLOPs", "FLOP", "0.00E+00", "={X}{Fpt}*1E9*{X}{Dp}*1E12"),
      ("Hp", "預訓練 GPU 小時", "GPU-hr", "#,##0", gpuh("Cp")),
      ("§", "D. 長上下文延伸（反向傳播型；稀疏注意力路徑）"),
      ("Dl", "延伸 token", "T", "#,##0.0", ti("ltok")),
      ("Sl", "序列長", "tok", "#,##0", ti("lseq")),
      ("attl", "平均被注意 token", "tok", "#,##0", sparse("{X}{Sl}/2")),
      ("Flt", "訓練 FLOPs/token", "GFLOP", "#,##0.0", "=3*(2*{X}{A}*1E9+{X}{attc}*{X}{attl})*{X}{hflop}/1E9"),
      ("Cl", "延伸 FLOPs", "FLOP", "0.00E+00", "={X}{Flt}*1E9*{X}{Dl}*1E12"),
      ("Hl", "延伸 GPU 小時", "GPU-hr", "#,##0", gpuh("Cl")),
      ("§", "E. SFT（反向傳播型）"),
      ("Ds", "SFT token", "T", "#,##0.00", ti("stok")),
      ("Ss", "序列長", "tok", "#,##0", ti("sseq")),
      ("atts", "平均被注意 token", "tok", "#,##0", sparse("{X}{Ss}/2")),
      ("Fst", "訓練 FLOPs/token", "GFLOP", "#,##0.0", "=3*(2*{X}{A}*1E9+{X}{attc}*{X}{atts})*{X}{hflop}/1E9"),
      ("Cs", "SFT FLOPs", "FLOP", "0.00E+00", "={X}{Fst}*1E9*{X}{Ds}*1E12"),
      ("sftk", "SFT MFU 係數", "x", "0.00", tgl("sftk")),
      ("Hs", "SFT GPU 小時", "GPU-hr", "#,##0", gpuh("Cs", "{X}{sftk}")),
      ("§", "F. RL（rollout＝推論型，以 Perf_Batch 計價；trainer＝反向傳播型）"),
      ("Rout", "rollout 輸出 token", "T", "#,##0.0", ti("rout")),
      ("isl", "每樣本輸入", "tok", "#,##0", pb("isl")),
      ("osl", "每樣本輸出", "tok", "#,##0", pb("osl")),
      ("ns", "樣本數", "個", "#,##0", "={X}{Rout}*1E12/{X}{osl}"),
      ("G", "GRPO 每題取樣數", "個", "0", tgl("G")),
      ("pref", "新鮮 prefill token（每題一次）", "T", "#,##0.00", "={X}{ns}*{X}{isl}/{X}{G}/1E12"),
      ("gsdr", "rollout decode GPU 秒／M（Perf_Batch 本層級）", "GPU-s/M", "#,##0.0", pb("gsd")),
      ("gsfr", "rollout prefill GPU 秒／M", "GPU-s/M", "#,##0.0", pb("gsf")),
      ("fdb", "rollout decode FLOPs/token", "GFLOP", "#,##0.0", pb("Fd")),
      ("fpb", "rollout prefill FLOPs/token", "GFLOP", "#,##0.0", pb("Fp")),
      ("hroll", "rollout 效率倍數（H_ROLL）", "x", "0.00", hk("H_ROLL")),
      ("reff", "rollout 效率（採用）", "%", "0%", tgl("reff") + "*{X}{hroll}"),
      ("Hro", "rollout GPU 小時（含閒置）", "GPU-hr", "#,##0", "=({X}{Rout}*1E6*{X}{gsdr}+{X}{pref}*1E6*{X}{gsfr})/3600/{X}{reff}"),
      ("Cro", "rollout FLOPs（接受 token）", "FLOP", "0.00E+00", "={X}{Rout}*1E12*{X}{fdb}*1E9+{X}{pref}*1E12*{X}{fpb}*1E9"),
      ("Tt", "trainer 處理 token（輸入＋輸出）", "T", "#,##0.0", "={X}{ns}*({X}{isl}+{X}{osl})/1E12"),
      ("attr", "trainer 平均被注意 token", "tok", "#,##0", sparse("({X}{isl}+{X}{osl})/2")),
      ("Frt", "trainer FLOPs/token", "GFLOP", "#,##0.0", "=3*(2*{X}{A}*1E9+{X}{attc}*{X}{attr})*{X}{hflop}/1E9"),
      ("Crt", "trainer FLOPs", "FLOP", "0.00E+00", "={X}{Frt}*1E9*{X}{Tt}*1E12"),
      ("rlk", "RL trainer MFU 係數", "x", "0.00", tgl("rlk")),
      ("Hrt", "trainer GPU 小時", "GPU-hr", "#,##0", gpuh("Crt", "{X}{rlk}")),
      ("refk", "參考模型前向（0／1）", "", "0", tgl("ref")),
      ("Hrf", "參考模型前向 GPU 小時（推論型）", "GPU-hr", "#,##0", "={X}{refk}*{X}{Hrt}/3"),
      ("Hrl", "RL GPU 小時合計", "GPU-hr", "#,##0", "={X}{Hro}+{X}{Hrt}+{X}{Hrf}"),
      ("Crl", "RL FLOPs 合計", "FLOP", "0.00E+00", "={X}{Cro}+{X}{Crt}*(1+{X}{refk}/3)"),
      ("rlmfu", "RL 有效 MFU（FLOPs ÷ GPU 小時 ÷ 訓練峰值）", "%", "0.0%", "={X}{Crl}/({X}{Hrl}*3600*{X}{pk}*1E15)"),
      ("rlH", "RL ÷ 預訓練（GPU 小時）", "x", "0.00", "={X}{Hrl}/{X}{Hp}"),
      ("rlF", "RL ÷ 預訓練（FLOPs）", "x", "0.00", "={X}{Crl}/{X}{Cp}"),
      ("§", "G. On-policy 蒸餾（學生 rollout＋教師評分＝推論型；學生更新＝反向傳播型）"),
      ("Dd", "學生 rollout 輸出 token", "T", "#,##0.00", ti("dtok")),
      ("dt", "教師層級", "", "0", ti("dtea")),
      ("Td", "教師評分 token（輸入＋輸出）", "T", "#,##0.00", "={X}{Dd}*1E12/{X}{osl}*({X}{isl}+{X}{osl})/1E12"),
      ("gsft", "教師 prefill GPU 秒／M", "GPU-s/M", "#,##0.0", pbt("gsf", "dt")),
      ("fpt", "教師 prefill FLOPs/token", "GFLOP", "#,##0.0", pbt("Fp", "dt")),
      ("Hds", "學生 rollout GPU 小時", "GPU-hr", "#,##0", "={X}{Dd}*1E6*{X}{gsdr}/3600/{X}{reff}"),
      ("Hdt", "教師評分 GPU 小時", "GPU-hr", "#,##0", "={X}{Td}*1E6*{X}{gsft}/3600"),
      ("Cdu", "學生更新 FLOPs", "FLOP", "0.00E+00", "={X}{Frt}*1E9*{X}{Td}*1E12"),
      ("Hdu", "學生更新 GPU 小時", "GPU-hr", "#,##0", gpuh("Cdu", "{X}{rlk}")),
      ("Hd", "蒸餾 GPU 小時合計", "GPU-hr", "#,##0", "={X}{Hds}+{X}{Hdt}+{X}{Hdu}"),
      ("Cd", "蒸餾 FLOPs 合計", "FLOP", "0.00E+00", "={X}{Dd}*1E12*{X}{fdb}*1E9+{X}{Td}*1E12*{X}{fpt}*1E9+{X}{Cdu}"),
      ("§", "H. 合成資料與評測（推論型）"),
      ("Dsy", "合成資料 token", "T", "#,##0.0", ti("syn")),
      ("gt", "生成層級", "", "0", tgl("gtier")),
      ("gsdg", "生成層級 decode GPU 秒／M", "GPU-s/M", "#,##0.0", pbt("gsd", "gt")),
      ("fdg", "生成層級 decode FLOPs/token", "GFLOP", "#,##0.0", pbt("Fd", "gt")),
      ("Hsy", "合成資料 GPU 小時", "GPU-hr", "#,##0", "={X}{Dsy}*1E6*{X}{gsdg}/3600"),
      ("Csy", "合成資料 FLOPs", "FLOP", "0.00E+00", "={X}{Dsy}*1E12*{X}{fdg}*1E9"),
      ("Dev", "評測 token", "T", "#,##0.00", ti("ev")),
      ("Hev", "評測 GPU 小時", "GPU-hr", "#,##0", "={X}{Dev}*1E6*{X}{gsdr}/3600"),
      ("Cev", "評測 FLOPs", "FLOP", "0.00E+00", "={X}{Dev}*1E12*{X}{fdb}*1E9"),
      ("§", "I. 最終訓練合計（單一模型）"),
      ("Hpre", "預訓練側 GPU 小時（預訓練＋長上下文）", "GPU-hr", "#,##0", "={X}{Hp}+{X}{Hl}"),
      ("Hpost", "後訓練 GPU 小時（SFT＋RL＋蒸餾）", "GPU-hr", "#,##0", "={X}{Hs}+{X}{Hrl}+{X}{Hd}"),
      ("Hoth", "合成資料與評測 GPU 小時", "GPU-hr", "#,##0", "={X}{Hsy}+{X}{Hev}"),
      ("Hfin", "最終訓練 GPU 小時", "GPU-hr", "#,##0", "={X}{Hpre}+{X}{Hpost}+{X}{Hoth}"),
      ("Cpre", "預訓練側 FLOPs", "FLOP", "0.00E+00", "={X}{Cp}+{X}{Cl}"),
      ("Cpost", "後訓練 FLOPs", "FLOP", "0.00E+00", "={X}{Cs}+{X}{Crl}+{X}{Cd}"),
      ("Cfin", "最終訓練 FLOPs", "FLOP", "0.00E+00", "={X}{Cpre}+{X}{Cpost}+{X}{Csy}+{X}{Cev}"),
      ("psF", "後訓練占比（FLOPs 口徑）", "%", "0.0%", "={X}{Cpost}/({X}{Cpre}+{X}{Cpost})"),
      ("psH", "後訓練占比（GPU 小時口徑）", "%", "0.0%", "={X}{Hpost}/({X}{Hpre}+{X}{Hpost})"),
      ("HI", "推論型 GPU 小時", "GPU-hr", "#,##0", "={X}{Hro}+{X}{Hrf}+{X}{Hds}+{X}{Hdt}+{X}{Hsy}+{X}{Hev}"),
      ("HB", "反向傳播型 GPU 小時", "GPU-hr", "#,##0", "={X}{Hp}+{X}{Hl}+{X}{Hs}+{X}{Hrt}+{X}{Hdu}"),
      ("Ish", "推論型占最終訓練", "%", "0.0%", "={X}{HI}/{X}{Hfin}"),
      ("Ufin", "最終訓練成本 — 經濟", "$M", "#,##0.0", "={X}{Hfin}*{X}{ce}/1E6"),
      ("Ufa", "最終訓練成本 — 會計", "$M", "#,##0.0", "={X}{Hfin}*{X}{ca}/1E6"),
      ("GWf", "最終訓練占 1 GW 一年", "%", "0.000%", "={X}{Hfin}/({X}{gpus}*8760)"),
      ("§", "J. 研發計畫（J10：研發倍數乘在 GPU 小時上；依最終訓練比例分攤＝各層級最終 × 倍數）"),
      ("rdm", "研發倍數", "x", "0.0", tgl("rdm")),
      ("Hprog", "研發計畫 GPU 小時", "GPU-hr", "#,##0", "={X}{Hfin}*{X}{rdm}"),
      ("Hexp", "其中：研發實驗（型態未拆分）", "GPU-hr", "#,##0", "={X}{Hprog}-{X}{Hfin}"),
      ("Uprog", "研發計畫成本 — 經濟", "$M", "#,##0", "={X}{Hprog}*{X}{ce}/1E6"),
      ("GWp", "研發計畫占 1 GW 一年", "%", "0.00%", "={X}{Hprog}/({X}{gpus}*8760)"),
      ("§", "K. 用途 × 型態（GPU 小時；單一模型計畫；對外服務不在第 0 層，J11）"),
      ("xPreB", "預訓練（含長上下文）｜反向傳播型", "GPU-hr", "#,##0", "={X}{Hpre}"),
      ("xPostI", "後訓練｜推論型（rollout、參考、蒸餾學生與教師）", "GPU-hr", "#,##0", "={X}{Hro}+{X}{Hrf}+{X}{Hds}+{X}{Hdt}"),
      ("xPostB", "後訓練｜反向傳播型（SFT、RL trainer、學生更新）", "GPU-hr", "#,##0", "={X}{Hs}+{X}{Hrt}+{X}{Hdu}"),
      ("xOthI", "合成資料與評測｜推論型", "GPU-hr", "#,##0", "={X}{Hoth}"),
      ("xExp", "研發實驗｜未拆分", "GPU-hr", "#,##0", "={X}{Hexp}"),
    ]

TKEY = {"Hp", "Hrl", "Hfin", "psF", "psH", "rlmfu", "Ufin", "GWf", "Hprog", "Uprog", "GWp", "rlH"}

def write_train(ws, cols, SP, AR, TI, PB, TR, start=9, overrides=None):
    overrides = overrides or {}
    R = {}; r = start
    for item in train_rows(SP, AR, TI, PB, TR):
        if item[0] == "§":
            section(ws, r, item[1], 2 + len(cols)); r += 1; continue
        key, lab, unit, fmt, tpl = item
        R[key] = r
        put(ws, f"A{r}", lab); put(ws, f"B{r}", unit)
        for X in cols:
            ov = overrides.get(X, {})
            if key in ov:
                v = ov[key].replace("{X}", X) if isinstance(ov[key], str) else ov[key]
                put(ws, f"{X}{r}", v, F_LINK if isinstance(v, str) and v.startswith("=") else F_IN, fmt=fmt, fill=PatternFill("solid", fgColor="FFFFFF00")); continue
            m = dict(R); m["X"] = X
            put(ws, f"{X}{r}", tpl.format(**m), fmt=fmt, fill=FILL_KEY if key in TKEY else None)
        r += 1
    ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 10
    return R, r

def hdr15(ws, cols, gens_tiers):
    put(ws, "A4", "世代", F_BOLD); put(ws, "A5", "層級", F_BOLD); put(ws, "A6", "世代索引", F_BOLD); put(ws, "A7", "層級索引", F_BOLD)
    for X, (g, t) in zip(cols, gens_tiers):
        put(ws, f"{X}6", g, fmt="0"); put(ws, f"{X}7", t, fmt="0")
        put(ws, f"{X}4", f"=INDEX(Spec_Rack!$C$4:$G$4,{X}6)", F_HLINK)
        put(ws, f"{X}5", f"=INDEX(Arch!$C$4:$E$4,{X}7)", F_HLINK)
        ws.column_dimensions[X].width = 13

def training_sheet(wb, SP, AR, TI, PB, TR):
    ws = wb.create_sheet("Training")
    title(ws, "Training — 訓練與研發計畫的 GPU 小時、成本、1 GW 年占比（世代 × 層級；下游預設 VR200，GB300 並列，J13）",
          "推論型運算以 Perf_Batch 計價；反向傳播型以訓練峰值 × MFU × goodput 計價。後訓練占比同時列 FLOPs 與 GPU 小時口徑（一級追蹤指標）")
    hdr15(ws, COLS15, [(i // 3 + 1, i % 3 + 1) for i in range(15)])
    R, r = write_train(ws, COLS15, SP, AR, TI, PB, TR)
    section(ws, r, "L. 模型家族合計（同世代三層級；J10 研發歸屬的總額）", 17); r += 1
    for key, lab, unit, fmt, src in [("fHfin", "家族最終訓練 GPU 小時", "GPU-hr", "#,##0", "Hfin"),
                                     ("fHprog", "家族研發計畫 GPU 小時", "GPU-hr", "#,##0", "Hprog"),
                                     ("fUprog", "家族研發計畫成本 — 經濟", "$M", "#,##0", "Uprog"),
                                     ("fGWp", "家族研發計畫占 1 GW 一年", "%", "0.00%", "GWp")]:
        R[key] = r; put(ws, f"A{r}", lab); put(ws, f"B{r}", unit)
        for X in COLS15:
            put(ws, f"{X}{r}", f"=SUMPRODUCT(($C$6:$Q$6={X}$6)*$C${R[src]}:$Q${R[src]})", fmt=fmt, fill=FILL_KEY)
        r += 1
    ws.freeze_panes = "C8"
    return R

SCEN_T = [  # label, overrides applied to both columns (Sol, Astra); special "fp8" flag
 ("基準", {}),
 ("預訓練 MFU 基準 0.15", {"mfub": 0.15}),
 ("預訓練 MFU 基準 0.40", {"mfub": 0.40}),
 ("NVFP4 訓練峰值（J7 替代）", {"pk8": "=INDEX(Spec_Rack!$C${tfp4}:$G${tfp4},{X}$6)"}),
 ("預訓練 token × 0.5", {"htok": 0.5}),
 ("預訓練 token × 1.67", {"htok": 1.67}),
 ("RL rollout token × 0.3", {"Rout": "=INDEX(Train_In!$C${rout}:$E${rout},{X}$7)*{X}$8", "_mult": 0.3}),
 ("RL rollout token × 3", {"Rout": "=INDEX(Train_In!$C${rout}:$E${rout},{X}$7)*{X}$8", "_mult": 3}),
 ("rollout 效率 0.6（v5.5 混合現況；rollout token 不變）", {"reff": 0.6}),
 ("rollout 效率 0.95（rollout token 不變）", {"reff": 0.95}),
 ("rollout 精度 FP8（J12 替代）", "fp8"),
 ("合成資料 token × 0", {"Dsy": "=INDEX(Train_In!$C${syn}:$E${syn},{X}$7)*{X}$8", "_mult": 0}),
 ("合成資料 token × 3", {"Dsy": "=INDEX(Train_In!$C${syn}:$E${syn},{X}$7)*{X}$8", "_mult": 3}),
 ("研發倍數 4.4（MiniMax）", {"rdm": 4.4}),
 ("研發倍數 10.4（OpenAI）", {"rdm": 10.4}),
 ("goodput 0.80", {"gp": 0.8}),
]

MULT_ROW = 8
MULT_LABEL = "情境倍數（×；黃底藍字＝作用中，乘在該情境改動的量上）"
# v5.11 had the "合成資料 token × 0" scenario written as a blue 0 on the synthetic-token row; v5.12 moves it to MULT_ROW.
# build.py remaps that snapshot key so an Excel-edited value is carried over (restore stays unmatched 0).
SNAP_MOVES = [(("Sens_Train", ("合成資料 token", 0), c), ("Sens_Train", (MULT_LABEL, 0), c)) for c in (25, 26)]   # Y, Z

def sens_train(wb, SP, AR, TI, PB, TR):
    ws = wb.create_sheet("Sens_Train")
    title(ws, "Sens_Train — 訓練與研發計畫的單變數敏感度（VR200；每組左 Sol、右 Astra；先看敏感度，再看基準）",
          "黃底藍字＝該情境改動的輸入。L 節為對基準欄（C、D）的比值")
    cols, ov, gt = [], {}, []
    put(ws, f"A{MULT_ROW}", MULT_LABEL, F_BOLD); put(ws, f"B{MULT_ROW}", "x")
    for i, (lab, o) in enumerate(SCEN_T):
        xs, xa = L(3 + 2 * i), L(4 + 2 * i)
        cols += [xs, xa]; gt += [(4, 2), (4, 3)]
        # v5.12 (A): scenario multipliers live in one input row (MULT_ROW); "—" where the scenario is not a multiplier
        m = o.get("_mult") if isinstance(o, dict) else None
        for X in (xs, xa):
            if m is None: put(ws, f"{X}{MULT_ROW}", "—", F_NOTE)
            else: put(ws, f"{X}{MULT_ROW}", m, fmt="0.00", fill=PatternFill("solid", fgColor="FFFFFF00"))
        if isinstance(o, dict): o = {k: v for k, v in o.items() if not k.startswith("_")}
        if o == "fp8":
            ov[xs] = {k: f"=Perf_Batch!R{PB[k2]}" for k, k2 in (("gsdr", "gsd"), ("gsfr", "gsf"), ("fdb", "Fd"), ("fpb", "Fp"))}
            ov[xa] = {k: f"=Perf_Batch!S{PB[k2]}" for k, k2 in (("gsdr", "gsd"), ("gsfr", "gsf"), ("fdb", "Fd"), ("fpb", "Fp"))}
        else:
            oo = {k: (v.replace("{tfp4}", str(SP["tfp4"])).replace("{rout}", str(TI["rout"])).replace("{syn}", str(TI["syn"])) if isinstance(v, str) else v)
                  for k, v in o.items()}
            ov[xs] = dict(oo); ov[xa] = dict(oo)
        put(ws, f"{xs}3", lab, F_BOLD, wrap=True); ws.merge_cells(f"{xs}3:{xa}3")
    ws.row_dimensions[3].height = 44; put(ws, "A3", "情境", F_BOLD)
    hdr15(ws, cols, gt)
    wb.defined_names["CST_STMult"] = DefinedName("CST_STMult", attr_text=f"Sens_Train!$C${MULT_ROW}:${cols[-1]}${MULT_ROW}")
    R, r = write_train(ws, cols, SP, AR, TI, PB, TR, overrides=ov)
    R["mult"] = MULT_ROW
    section(ws, r, "L. 對基準的比值（同層級）", 2 + len(cols)); r += 1
    for key, lab in [("Hfin", "最終訓練 GPU 小時 ÷ 基準"), ("Hprog", "研發計畫 GPU 小時 ÷ 基準"), ("psH", "後訓練占比（GPU 小時）÷ 基準")]:
        R["r_" + key] = r; put(ws, f"A{r}", lab); put(ws, f"B{r}", "x")
        for i, X in enumerate(cols):
            base = "C" if i % 2 == 0 else "D"
            put(ws, f"{X}{r}", f"={X}{R[key]}/${base}${R[key]}", fmt="0.00", fill=FILL_KEY)
        r += 1
    ws.freeze_panes = "C8"
    return R
