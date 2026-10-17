"""4-3 A 段步驟 1：在 v13.1 重跑死格分類。用 tests/parity/deadcells_lib（含 Sens_／_Prod 複本為根，與 test_slim 同口徑）。
輸出 pickle：dead（全部死格）、unmarked（無「用途」標記）、base（fixture 基準名單）。
用法：python3 20_dead_v131.py <out.pkl>   （工作目錄＝repo 根）"""
import sys, json, pickle
from pathlib import Path
REPO = Path.cwd()
sys.path.insert(0, str(REPO / "tests" / "parity"))
import deadcells_lib
import openpyxl
model = REPO / "model" / (REPO / "model" / "CURRENT").read_text().strip()
CONTRACT = REPO / "docs" / "plan" / "Tokenomics_downstream_contract.md"
dead, nform, per_page = deadcells_lib.dead_formula_cells(model, CONTRACT)
base = set(json.loads((REPO / "tests/parity/fixtures/dead_unmarked_baseline_v10.3.json").read_text(encoding="utf-8"))["unmarked"])
wb = openpyxl.load_workbook(model)
pcol = {}
for ws in wb.worksheets:
    for c in range(3, ws.max_column + 2):
        if ws.cell(3, c).value == "用途":
            pcol[ws.title] = c
unmarked, marked = [], []
for s, a in sorted(dead):
    ws = wb[s]
    col = pcol.get(s)
    m = bool(col and ws.cell(ws[a].row, col).value not in (None, ""))
    (marked if m else unmarked).append((s, a))
um = {f"{s}!{a}" for s, a in unmarked}
print("model", model.name, "formula cells", nform, "dead", len(dead), "marked", len(marked), "unmarked", len(unmarked))
print("unmarked in fixture", len(um & base), "unmarked not in fixture", len(um - base), "fixture not unmarked now", len(base - um))
print("rows unmarked", len({(s, wb[s][a].row) for s, a in unmarked}))
print(per_page.most_common(40))
pickle.dump(dict(model=model.name, nform=nform, dead=sorted(dead), unmarked=unmarked, marked=marked, base=sorted(base), per_page=dict(per_page), pcol=pcol), open(sys.argv[1], "wb"))
