"""鏡像與複本一致性檢查的共用函式（瘦身 4-1，永久測試用）。只讀公式文字、填色與字色，不需要重算。

由 4-0 盤點腳本 02_audit.py 的 C4 與 08_c4_extra.py 精簡而來。檢查對象：
- 8 個 *_Prod 鏡像頁 ↔ 核心頁（同格、R1C1 相對結構；*_Prod 頁名視同核心頁）
- Sens_Perf／Sens_Train ↔ Perf／Training（依世代、層級索引對應欄）
- Sensitivity／Sens_Har：同列各情境欄與首欄（C 欄）的差異
- Perf ↔ Perf_Batch（頁首聲明：同一套公式，只換速度下限與參考任務）
- 藍字情境輸入格是否上黃底
"""
from __future__ import annotations

import collections
import re

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L

BLUE = "FF0000FF"
YELLOWS = ("FFFFF2CC", "FFFFFF00")      # 淡黃（情境改動的公式格，沿用各敏感度頁既有寫法）、亮黃（情境輸入格）
# 明訂差異（白名單）：頁首或說明已聲明的差異列
PERF_BATCH_ROWS = {41, 43, 44, 53, 76, 93, 94, 120, 121}     # Perf_Batch ↔ Perf：頁首「與 Perf 同一套公式，只換速度下限與參考任務」
PROD_OWN_ROWS = {"Theory_Rev_Prod": set(range(90, 97))}       # 鏡像頁專屬的自我檢查列（頁首明訂，不鏡像核心頁同列）
SUMMARY_FROM = {"Sensitivity": 64}                            # Sensitivity 第 64 列起為摘要表（各欄意義不同，不比對）
NOT_COMPARED = {"Sens_Rev"}                                   # 封閉式重算頁，無逐格對應的核心頁


def load(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    cells = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                fill = None
                if c.fill is not None and c.fill.fill_type == "solid" and c.fill.fgColor.type == "rgb":
                    fill = str(c.fill.fgColor.rgb)
                col = None
                if c.font is not None and c.font.color is not None and c.font.color.type == "rgb":
                    col = str(c.font.color.rgb)
                cells[(ws.title, c.coordinate)] = dict(f=c.value, r=c.row, c=c.column, fill=fill, color=col)
    names = {n: dn.attr_text for n, dn in wb.defined_names.items()}
    return cells, names, [ws.title for ws in wb.worksheets]


def isf(x):
    return isinstance(x["f"], str) and x["f"].startswith("=")


def r1c1(formula, own, r0, c0, lname, lnames, shift=0):
    """公式轉成相對 R1C1 文字。*_Prod 頁名一律換成核心頁名，並去掉與本頁同名的前綴。"""
    base = own[:-5] if own.endswith("_Prod") else own
    try:
        toks = Tokenizer(formula).items
    except Exception:
        return formula
    out = []
    for t in toks:
        val = t.value
        if t.type == "OPERAND" and t.subtype == "RANGE" and val.lower() in lname:
            val = lname[val.lower()]
        if t.type == "OPERAND" and t.subtype == "RANGE" and not (val.lower() in lnames and val.lower() not in lname):
            sh = None
            ref = val
            if "!" in ref:
                sh, ref = ref.rsplit("!", 1)
                sh = sh.strip("'")
            is_local = sh is None or sh == own or sh.endswith("_Prod")
            if sh is not None and sh.endswith("_Prod"):
                sh = sh[:-5]
            if sh == base:
                sh = None
            parts = []
            for p in ref.split(":"):
                m = re.fullmatch(r"(\$?)([A-Z]{1,3})?(\$?)(\d+)?", p)
                if not m:
                    parts.append(p)
                    continue
                ca, col, ra, row = m.groups()
                s_ = ""
                if row:
                    s_ += f"R{row}" if ra else f"R[{int(row) - r0}]"
                if col:
                    cn = CI(col)
                    if shift and is_local and cn >= 19:
                        cn -= shift
                    s_ += f"C{cn}" if ca else f"C[{cn - c0}]"
                parts.append(s_)
            out.append((sh + "!" if sh else "") + ":".join(parts))
        elif t.type == "FUNC":
            out.append(t.value.upper())
        else:
            out.append(val)
    return "".join(out)


class Model:
    def __init__(self, path):
        self.cells, self.names, self.sheets = load(path)
        self.lname = {n.lower(): t for n, t in self.names.items() if "," not in t}
        self.lnames = {n.lower() for n in self.names}

    def rc(self, sheet, x):
        return r1c1(x["f"], sheet, x["r"], x["c"], self.lname, self.lnames)

    def label(self, sheet, r):
        x = self.cells.get((sheet, f"A{r}"))
        return str(x["f"]) if x and not isf(x) and x["f"] is not None else (str(x["f"]) if x else "")

    # ---------------------------------------------------------------- *_Prod 鏡像頁 ↔ 核心頁
    def prod_undocumented(self):
        """回傳未明訂差異清單（空＝通過）。直接讀核心頁同格、R1C1 結構相同、第 50 列 CTL_ProdDerate、專屬自我檢查列皆視為明訂。"""
        bad = []
        for mp in (s for s in self.sheets if s.endswith("_Prod")):
            core = mp[:-5]
            for (s, a), x in self.cells.items():
                if s != mp or not isf(x):
                    continue
                r, c = x["r"], x["c"]
                kind = "副本" if c >= 19 else "主體"
                cc = c - 16 if c >= 19 else c
                if r in PROD_OWN_ROWS.get(mp, ()):
                    continue
                mf = x["f"]
                if re.fullmatch(rf"='?{re.escape(core)}'?!\$?{L(cc)}\$?{r}", mf):
                    continue
                cx = self.cells.get((core, f"{L(cc)}{r}"))
                ml = self.cells.get((mp, f"A{r}"))
                if ml and not isf(ml) and ml["f"] and str(ml["f"]) != self.label(core, r):
                    bad.append((mp, a, "鏡像頁專屬列未列入白名單"))
                    continue
                if cx is None or not isf(cx):
                    bad.append((mp, a, "核心頁同位置非公式"))
                    continue
                shift = 16 if kind == "副本" else 0
                a1 = r1c1(mf, mp, r, cc, self.lname, self.lnames, shift)
                b1 = r1c1(cx["f"], core, r, cc, self.lname, self.lnames)
                if a1 == b1:
                    continue
                if r == 50 and "CTL_ProdDerate" in mf and kind == "主體":
                    continue
                if r == 50 and "Serving!$C$18" in mf and kind == "副本":
                    continue
                bad.append((mp, a, "R1C1 結構不同"))
        return bad

    # ---------------------------------------------------------------- Sens_Perf／Sens_Train ↔ Perf／Training
    def sens_core_undocumented(self):
        bad = []
        for sp, core in (("Sens_Perf", "Perf"), ("Sens_Train", "Training")):
            colmap = {}
            for c in range(3, 40):
                g, t = self.cells.get((core, f"{L(c)}6")), self.cells.get((core, f"{L(c)}7"))
                if g and t:
                    colmap[(g["f"], t["f"])] = c
            for (s, a), x in self.cells.items():
                if s != sp or x["c"] < 3 or x["r"] < 8:
                    continue
                g, t = self.cells.get((sp, f"{L(x['c'])}6")), self.cells.get((sp, f"{L(x['c'])}7"))
                cc = colmap.get((g["f"] if g else None, t["f"] if t else None))
                if cc is None or self.label(sp, x["r"]) != self.label(core, x["r"]):
                    continue
                if x["fill"] == "FFFFFF00" or (x["color"] == BLUE and not isf(x)):
                    continue          # 情境輸入格
                cx = self.cells.get((core, f"{L(cc)}{x['r']}"))
                if not isf(x):
                    if cx is not None and not isf(cx) and cx["f"] == x["f"]:
                        continue
                    bad.append((sp, a, "非情境常數與核心頁不同"))
                    continue
                if cx is None or not isf(cx):
                    bad.append((sp, a, "核心頁同位置非公式"))
                    continue
                if r1c1(x["f"], core, x["r"], x["c"], self.lname, self.lnames) != r1c1(cx["f"], core, x["r"], cc, self.lname, self.lnames):
                    bad.append((sp, a, "R1C1 結構不同"))
        return bad

    # ---------------------------------------------------------------- Sensitivity／Sens_Har：同列與首欄不同的格
    def row_diff_cells(self, sp):
        """回傳 (與同列首欄（C 欄）不同的公式格集合, 黃底公式格集合)；Sensitivity 第 64 列起（摘要表）不比對。"""
        rows = collections.defaultdict(list)
        for (s, a), x in self.cells.items():
            if s == sp and x["c"] >= 3 and isf(x) and x["r"] < SUMMARY_FROM.get(sp, 10 ** 6):
                rows[x["r"]].append(x)
        diff, yellow = set(), set()
        for r, xs in rows.items():
            xs.sort(key=lambda z: z["c"])
            base = xs[0]
            bf = self.rc(sp, base)
            for x in xs[1:]:
                if x["f"] != base["f"] and self.rc(sp, x) != bf:
                    diff.add(f"{L(x['c'])}{r}")
            if base["fill"] in YELLOWS:
                continue                      # 整列黃底＝關鍵輸出列，不是情境標記
            for x in xs[1:]:
                if x["fill"] in YELLOWS:
                    yellow.add(f"{L(x['c'])}{r}")
        return diff, yellow

    def blue_scenario_without_yellow(self):
        """五個敏感度頁中，藍字數值輸入格（第 6、7 列的世代／層級索引除外）沒有黃底者。"""
        bad = []
        for sp in ("Sens_Perf", "Sens_Train", "Sensitivity", "Sens_Rev", "Sens_Har"):
            for (s, a), x in self.cells.items():
                if s != sp or x["color"] != BLUE or isf(x) or not isinstance(x["f"], (int, float)):
                    continue
                if sp in ("Sens_Perf", "Sens_Train") and x["r"] in (6, 7):
                    continue
                if x["fill"] not in YELLOWS:
                    bad.append((sp, a))
        return bad

    # ---------------------------------------------------------------- Perf ↔ Perf_Batch
    def perf_batch_diff_rows(self):
        """Perf_Batch 與 Perf 在 C:Q 同格結構不同的列（集合）與格數。"""
        rows = collections.Counter()
        for (s, a), x in self.cells.items():
            if s != "Perf_Batch" or not isf(x) or x["c"] < 3 or x["c"] > 17 or x["r"] < 8:
                continue
            p = self.cells.get(("Perf", a))
            if p is None or not isf(p):
                rows[x["r"]] += 1
                continue
            if self.rc("Perf_Batch", x) != self.rc("Perf", p):
                rows[x["r"]] += 1
        return rows
