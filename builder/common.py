# Shared styling + helpers for Tokenomics v5.1 builder
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

BLUE, BLACK, GREEN = "FF0000FF", "FF000000", "FF008000"
F_IN   = Font(name="Arial", size=10, color=BLUE)
F_CALC = Font(name="Arial", size=10, color=BLACK)
F_LINK = Font(name="Arial", size=10, color=GREEN)
F_BOLD = Font(name="Arial", size=10, bold=True)
F_HLINK= Font(name="Arial", size=10, bold=True, color=GREEN)
F_TITLE= Font(name="Arial", size=13, bold=True)
F_NOTE = Font(name="Arial", size=9, color="FF595959")
FILL_SEC = PatternFill("solid", fgColor="FFD9E1F2")
FILL_KEY = PatternFill("solid", fgColor="FFFFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")

def put(ws, ref, v, font=None, fmt=None, fill=None, wrap=False):
    c = ws[ref]; c.value = v
    if font is None:
        if isinstance(v, str) and v.startswith("="):
            font = F_LINK if "!" in v else F_CALC
        elif isinstance(v, (int, float)):
            font = F_IN
        else:
            font = F_CALC
    c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if wrap: c.alignment = WRAP
    return c

def section(ws, row, text, ncols):
    for i in range(1, ncols + 1):
        ws.cell(row=row, column=i).fill = FILL_SEC
    put(ws, f"A{row}", text, F_BOLD, fill=FILL_SEC)

def title(ws, t1, t2):
    put(ws, "A1", t1, F_TITLE)
    put(ws, "A2", t2, F_NOTE)
