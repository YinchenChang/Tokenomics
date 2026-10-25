# 產生「NonNV_TPU_Trainium.xlsx」審閱檔（唯讀查證，不碰 model/）
# 用法：python3 build_xlsx.py <輸出.xlsx>，再以 LibreOffice headless 轉存寫入計算值
# 方法沿用 A 段（docs/evidence/20261023_NonNV昇騰Maia_A段/build_xlsx.py）：
#   產出比＝（候選每 kW 峰值算力 ÷ VR200）× 效率因子 f；持有比＝1 − s ×（1 − k）
# 推導頁藍字＝輸入，黑字＝公式；欄 D/E/F＝低端／基準／高端
import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, __file__.rsplit("/", 1)[0] if "/" in __file__ else ".")
from evidence_data import EVIDENCE, QUESTIONS  # noqa: E402

out = sys.argv[1]
wb = Workbook()
BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7")
SEC = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")

ws = wb.active
ws.title = "推導"
ws["A1"] = "推導 — Google TPU v7 Ironwood 與 AWS Trainium3 相對 VR200 NVL72 的每 GW 產出比與持有成本比（B 段；方法同 A 段）"
ws["A1"].font = BOLD
ws["A2"] = ("欄 D／E／F＝低端／基準／高端（產出列：悲觀→樂觀；持有列：成本低→成本高）。藍字＝輸入；黑字＝公式。"
            "標記沿用 Verified／Interested-party／Analogy／Assumed／Derived。VR200 分母取自 model/20261025_Tokenomics_v21.5.xlsx（基準成本欄 Interface!M）。"
            "證據編號 T01–T15 見『證據』頁。")
ws["A2"].alignment = WRAP
heads = ["編號", "物理量", "單位", "低端", "基準", "高端", "來源（證據編號／模型格）", "標記", "算式或說明"]
for i, h in enumerate(heads, 1):
    c = ws.cell(4, i, h)
    c.font = BOLD
    c.fill = HDR

R = {}
row = [5]


def sec(title):
    r = row[0]
    ws.cell(r, 1, title).font = BOLD
    for c in range(1, 10):
        ws.cell(r, c).fill = SEC
    row[0] += 1


def ref(key, col):
    return f"{col}{R[key]}"


def add(key, name, unit, vals, src, tag, note="", fmt=None):
    r = row[0]
    R[key] = r
    ws.cell(r, 1, key)
    ws.cell(r, 2, name)
    ws.cell(r, 3, unit)
    for j, v in enumerate(vals):
        col = "DEF"[j]
        cell = ws[f"{col}{r}"]
        if isinstance(v, str) and v.startswith("="):
            def rep(m):
                return ref(m.group(1), m.group(2) or col)
            cell.value = re.sub(r"\{([A-Za-z0-9_]+)(?:@([DEF]))?\}", rep, v)
        else:
            cell.value = v
            cell.font = BLUE
        if fmt:
            cell.number_format = fmt
    ws.cell(r, 7, src)
    ws.cell(r, 8, tag)
    ws.cell(r, 9, note)
    row[0] += 1


def same(v):
    return [v, v, v]


# ---- A. VR200 分母（與 A 段相同，值改讀 v21.5；數值與 v20.1 相同）
sec("A. VR200 NVL72 分母（模型 v21.5；同層級、同 SLO）")
add("A1", "VR200 每 GPU 計算用峰值（NVFP4 dense；模型服務精度）", "PF", same(35), "Spec_Rack!F35（＝F23；F22 服務精度 NVFP4）", "Interested-party（NVIDIA 規格）", "")
add("A2", "VR200 機架配電設計功率（基準）", "kW", same(227), "Spec_Rack!F9", "Interested-party（Schneider 轉述 NVIDIA）", "")
add("A3", "每機架 GPU 封裝數", "顆", same(72), "Spec_Rack!F6", "Verified", "")
add("A4", "VR200 每顆 GPU 所攤機架功率", "kW", ["={A2}/{A3}"] * 3, "Derived", "Derived", "227÷72")
add("A5", "VR200 每 kW 峰值算力", "PF/kW", ["={A1}/{A4}"] * 3, "Derived", "Derived", "")
add("A6", "VR200 每 GW 機架數", "架", same(4052), "Interface!M6", "Derived（模型）", "")
add("A7", "VR200 每 GW 機架 IT 功率（＝A6×A2）", "MW", ["={A6}*{A2}/1000"] * 3, "Derived", "Derived", "候選每 GW 顆數＝此值 ÷ 候選每顆所攤功率（與 A 段相同，含同一機房加成）")
add("A8", "VR200 機架價格（基準）", "$M", same(8.4), "Spec_Rack!F12", "Interested-party", "")
add("A9", "VR200 每 GW IT 設備資本支出（基準）", "$B", same(37.586352), "Interface!M9", "Derived（模型）", "")
add("A10", "IT 資本支出 ÷（機架數×機架價）＝網路等加成", "x", ["={A9}/({A6}*{A8}/1000)"] * 3, "Derived", "Derived", "約 1.10；候選沿用同一加成")
add("A11", "VR200 每 GW 廠房資本支出（基準）", "$B", same(12.67), "Interface!M10", "Derived（模型）", "")
add("A12", "VR200 每 GW 年持有成本 — 經濟口徑（基準）", "$B/年", same(12.2254710045841), "Interface!M13", "Derived（模型）", "分母")
add("A13", "VR200 每 GW 年持有成本 — 會計口徑（基準）", "$B/年", same(8.9923042149807), "Interface!M12", "Derived（模型）", "")
add("A14", "VR200 每 GW IT 折舊（基準）", "$B/年", same(6.264392), "Interface!M220", "Derived（模型）", "")
add("A15", "VR200 每 GW IT 維護（基準）", "$B/年", same(0.591045788314028), "Interface!M224", "Derived（模型）", "")
add("A16", "VR200 每 GW 財產稅與保險（基準）", "$B/年", same(0.25128176), "Interface!M227", "Derived（模型）", "")
add("A17", "資金成本（經濟 − 會計）", "$B/年", ["={A12}-{A13}"] * 3, "Derived", "Derived", "")
add("A18", "IT 設備占總資本支出比例", "—", ["={A9}/({A9}+{A11})"] * 3, "Derived", "Derived", "")
add("A19", "持有成本中隨 IT 設備資本支出成比例的部分", "$B/年", ["={A14}+{A15}+({A16}+{A17})*{A18}"] * 3, "Derived", "Derived", "")
add("A20", "上述部分占 VR200 持有成本比例 s", "—", ["={A19}/{A12}"] * 3, "Derived", "Derived", "候選持有比 ＝ 1 − s×(1 − k)")
add("A21", "VR200 每 GPU FP8 dense", "PF", same(17.5), "Spec_Rack!F40", "Interested-party", "只用於 E 節『FP8 對齊』旁證")
add("A22", "VR200 GPU 功率（能量用）", "W", same(1800), "Spec_Rack!F33", "Interested-party", "")
add("A23", "VR200 機架功率 ÷ 晶片功率總和", "x", ["={A2}*1000/({A3}*{A22})"] * 3, "Derived", "Derived", "1.75；A 段 Maia 旁證為 1.78。用於由 Trn3 晶片 TDP 推機架功率")

# ---- B. 校準（同 A 段）
sec("B. 還原既有世代：模型實際產出比 ÷ 『每 kW 峰值算力比』＝有效效率因子 f（同 A 段）")
for key, nm, pf, kw, g, src in [
    ("H", "Hopper H100", 1.979, 41, 32, "Spec_Rack!C35,C9,C6"),
    ("G2", "GB200 NVL72", 10, 132, 72, "Spec_Rack!D35,D9,D6"),
    ("G3", "GB300 NVL72", 15, 136, 72, "Spec_Rack!E35,E9,E6"),
]:
    add(f"B_{key}_pf", f"{nm} 每 GPU 計算用峰值", "PF", same(pf), src, "Interested-party", "Hopper 為 FP8（無 FP4），其餘 NVFP4")
    add(f"B_{key}_kw", f"{nm} 機架配電功率（基準）", "kW", same(kw), src, "Interested-party／Derived", "")
    add(f"B_{key}_n", f"{nm} 每機架 GPU", "顆", same(g), src, "Verified", "")
    add(f"B_{key}_kwg", f"{nm} 每顆所攤功率", "kW", [f"={{B_{key}_kw}}/{{B_{key}_n}}"] * 3, "Derived", "Derived", "")
    add(f"B_{key}_r", f"{nm} 每 kW 峰值算力 ÷ VR200", "—", [f"=({{B_{key}_pf}}/{{B_{key}_kwg}})/{{A5}}"] * 3, "Derived", "Derived", "")
vr = {"H": (0.107484660941882, 0.10544360110803, 0.0977954294066359),
      "G2": (0.534632077509924, 0.513240130160715, 0.483657996161255),
      "G3": (0.693216239743152, 0.665656345719338, 0.563845155935545)}
for key, nm in [("H", "Hopper"), ("G2", "GB200"), ("G3", "GB300")]:
    add(f"B_{key}_v", f"{nm} VR-eq（Luna／Sol／Astra）", "—", list(vr[key]), "Interface IF_VReq_Luna/Sol/Astra（列 22/32/42）", "Derived（模型）", "欄 D/E/F＝Luna／Sol／Astra（此行不是低/基準/高）")
    add(f"B_{key}_vm", f"{nm} VR-eq 三層平均", "—", [f"=AVERAGE(D{R[f'B_{key}_v']}:F{R[f'B_{key}_v']})"] * 3, "Derived", "Derived", "")
    add(f"B_{key}_f", f"{nm} 有效效率因子 f", "—", [f"={{B_{key}_vm}}/{{B_{key}_r}}"] * 3, "Derived", "Derived", "")

# ---- C. TPU v7
sec("C. Google TPU v7 Ironwood")
add("C1", "每顆 FP8 dense 峰值（無 FP4）", "PF", same(4.614), "T01", "Interested-party（Google）", "以自身最高原生精度計（同模型對 Hopper 的處理）；精度折減由 f 吸收，見待判斷 Q1")
add("C2", "9,216 顆 pod 功率", "MW", [12, 10, 9.5], "T02（Google『nearly 10 MW』）；低端 12＝設計功率 +20%（Assumed，比照 VR200 MaxP 227 對 MaxQ 188 約 1.2 倍）；高端 9.5（『nearly』）", "Interested-party（基準）／Assumed（區間）", "pod 內含主機、ICI、OCS；是否含設施冷卻未明（若含，IT 更低→產出比更高，見待判斷 Q2）")
add("C3", "每 pod 晶片數", "顆", same(9216), "T01", "Verified（官方）", "144 架 × 64 顆（T03）")
add("C4", "每顆所攤 pod 功率", "kW", ["={C2}*1000/{C3}"] * 3, "Derived", "Derived", "基準 1.085；旁證 T04（100 萬顆↔超過 1 GW）")
add("C5", "每 kW 峰值算力", "PF/kW", ["={C1}/{C4}"] * 3, "Derived", "Derived", "", fmt="0.000")
add("C6", "÷ VR200 每 kW 峰值算力", "—", ["={C5}/{A5}"] * 3, "Derived", "Derived", "", fmt="0.0000")
add("C7", "有效效率因子 f（TPU）", "—", ["={B_H_f@E}", f"=AVERAGE({ref('B_H_f','E')},AVERAGE({ref('B_G2_f','E')},{ref('B_G3_f','E')}))", "={B_G2_f@E}"],
    "低＝Hopper 還原值 0.744；基準＝Hopper 還原值與 GB200／GB300 還原值平均（0.97）的中間（r1，Project 決定方案 (a)）；高＝GB200 還原值 1.04，以 T05 InferenceX 佐證（單一來源 SemiAnalysis，依規則只作佐證，不單獨作依據）", "Assumed（區間）", "r1：原基準 0.97 實質只靠 InferenceX，已改")
add("C8", "方法一：產出比 ＝ C6 × C7", "—", ["={C6}*{C7}"] * 3, "Derived", "Derived", "", fmt="0.0000")
add("C9", "TPU v7 實測吞吐（20 tok/s/user，FP8）", "tok/s/chip", same(9364), "T05", "第三方實測（SemiAnalysis InferenceX）", "")
add("C10", "B200 同條件實測吞吐（FP8）", "tok/s/GPU", same(8903), "T05", "第三方實測", "HGX B200；模型 GB200 為 NVL72，以此近似")
add("C11", "TPU ÷ B200 每顆 FP8 實測比", "—", ["={C9}/{C10}"] * 3, "Derived", "Derived", "", fmt="0.000")
add("C12", "GB200 VR-eq 三層平均（模型，NVFP4）", "—", ["={B_G2_vm@E}"] * 3, "模型", "Derived", "")
add("C13", "GB200 每顆所攤功率（模型）", "kW", ["={B_G2_kwg@E}"] * 3, "模型", "Derived", "")
add("C14", "B200 由 FP8 改 NVFP4 的實際吞吐倍數 u", "x", [2.0, 1.75, 1.5], "低端產出（u＝2.0）＝TK 引擎的算力綁定假設（GB200 峰值 FP4 為 FP8 的 2 倍）；基準 1.75、高端 1.5＝實務上 FP4 增益常低於 2 倍（Assumed）", "Assumed（區間）", "u 越大，TPU（只能 FP8）相對越吃虧")
add("C15", "方法二（實測錨定）：產出比 ＝ C12 ×（C13 ÷ C4）× C11 ÷ C14", "—", ["={C12}*({C13}/{C4})*{C11}/{C14}"] * 3, "Derived", "對照（不進建議值）", "r1：效率全靠 T05 InferenceX（單一來源），依規則降為對照、不進平均；與方法一共用每顆功率 C4", fmt="0.0000")
add("C16", "TPU v7 每 GW 產出比（原始）", "—", ["={C8}"] * 3, "Derived", "Derived（含 Assumed）", "r1：只取方法一；方法二（C15）為對照，基準 0.52 高於本列，表示本列可能偏保守", fmt="0.0000")
add("C17", "每顆取得價（含機架）", "$K", [21, 25, 35], "低 T07（Broadcom $21B 訂單 ÷ 約 1 GW≈100 萬顆）；基準 T06（SemiAnalysis：$10B≈40 萬顆成品機架）；高 T08（$35B 融資 ÷ 1 GW，可能含網路與融資費用）", "Interested-party／第三方估計（區間）", "Google 自用也向 Broadcom 購買，Broadcom 毛利已含在內")
add("C18", "每 GW 顆數（以基準功率）", "顆", [f"={{A7}}*1000/$E${R['C4']}"] * 3, "Derived", "Derived", "VR200 為 291,744")
add("C19", "每 GW IT 設備資本支出", "$B", ["={C18}*{C17}/1000000*{A10}"] * 3, "Derived", "Derived", "乘 VR200 同一網路加成（高端可能重複計算，見 C17）")
add("C20", "k ＝ ÷ VR200 每 GW IT 資本支出", "—", ["={C19}/{A9}"] * 3, "Derived", "Derived", "", fmt="0.000")
add("C21", "TPU v7 每 GW 持有成本比（原始）", "—", ["=1-{A20}*(1-{C20})"] * 3, "Derived", "Derived（含 Assumed）", "", fmt="0.0000")
add("C22", "交叉：直接用每 GW 金額（$21B／—／$35B，不乘加成）→ 持有比", "—", [f"=1-{ref('A20','E')}*(1-21/{ref('A9','E')})", f"={ref('C21','E')}", f"=1-{ref('A20','E')}*(1-35/{ref('A9','E')})"], "T07、T08", "Derived", "不進建議值；GW 口徑未明（若為設施口徑，每 IT GW 金額更高）", fmt="0.0000")
add("C23", "交叉：SemiAnalysis TCO/chip-hr（內部 $1.03／—／外部 $1.60）× 每 GW 顆數 × 8,760 ÷ VR200 持有成本", "—", [f"=1.03*{ref('C18','E')}*8760/1E9/{ref('A12','E')}", f"={ref('C21','E')}", f"=1.60*{ref('C18','E')}*8760/1E9/{ref('A12','E')}"], "T05、T06", "第三方估計", "不進建議值；TCO 口徑（是否含電費、廠房）未公開，外部價含 Google 利潤", fmt="0.0000")

# ---- D. Trainium3
sec("D. AWS Trainium3（Teton3 Max 液冷機架，72 顆；UltraServer＝2 架 144 顆）")
add("D1", "每顆 FP8 dense 峰值（MXFP4 同速）", "PF", same(2.517), "T09", "Interested-party（AWS）", "FP4 無額外算力，故以 FP8 值為最高原生精度")
add("D2", "晶片 TDP", "kW", [1.0, 0.85, 0.70], "低端產出 1.0＝DCD 依 AWS 副總裁液冷說法（『晶片超過 1,000 W 才需液冷』）推論約 1 kW（T11；副總裁未給瓦數）；高端 0.70＝KGI 2026-07（T10）；基準 0.85＝中點（Assumed）", "Derived／二手（1.0）；賣方研究（0.70）", "旁證：AWS 另有氣冷 Teton3 PDS 32 顆版，晶片功率不會遠超 1 kW。Naddod 500 W（T12）品質低、與兩來源不符，不採用")
add("D3", "每機架晶片數", "顆", same(72), "T10（18 托盤×4）", "Verified（供應鏈）／Interested-party", "")
add("D3a", "TK 機架÷晶片倍數：GB200（132 kW ÷ 72×1.2 kW）", "x", same("=132/(72*1.2)"), "Spec_Rack!D9、D6、D33", "Derived（模型）", "")
add("D3b", "TK 機架÷晶片倍數：GB300（136 ÷ 72×1.4）", "x", same("=136/(72*1.4)"), "Spec_Rack!E9、E6、E33", "Derived（模型）", "")
add("D3c", "TK 機架÷晶片倍數：VR200 Max-Q（188 ÷ 72×1.8）", "x", same("=188/(72*1.8)"), "Spec_Rack!F8、F6、F33", "Derived（模型）", "")
add("D3r", "Trn3 機架÷晶片倍數（採用）", "x", ["={A23}", f"=AVERAGE({ref('D3a','E')},{ref('D3b','E')},{ref('D3c','E')},{ref('A23','E')})", "={D3b}"], "低端產出＝VR200 Max-P 1.75；基準＝TK 四值平均；高端產出＝GB300 1.35", "Analogy（區間）", "r1：原為單值 1.75。Trn3 機架另含 18 CPU、20 交換晶片（T10），倍數落在 TK 各世代區間內的假設")
add("D4", "機架功率 ＝ D2 × D3 × D3r", "kW", ["={D2}*{D3}*{D3r}"] * 3, "Derived（Analogy：TK 機架÷晶片倍數）", "Derived／Assumed", "上限旁證：電源 3 層×6×8.5＝153 kW（5+1 約 127.5），500 kW 電源櫃供兩架（T10）")
add("D5", "每顆所攤機架功率", "kW", ["={D4}/{D3}"] * 3, "Derived", "Derived", "")
add("D6", "每 kW 峰值算力", "PF/kW", ["={D1}/{D5}"] * 3, "Derived", "Derived", "", fmt="0.000")
add("D7", "÷ VR200 每 kW 峰值算力", "—", ["={D6}/{A5}"] * 3, "Derived", "Derived", "", fmt="0.0000")
add("D8", "有效效率因子 f（Trn3）", "—", [0.50, 0.80, "={B_G2_f@E}"], "無第三方實測（T14）；比照 A 段 Maia：低 0.50（Neuron 軟體成熟度）、基準 0.80（略高於 Hopper，MXFP4 權重省記憶體）、高＝GB200 還原值", "Assumed（區間）", "")
add("D9", "Trainium3 每 GW 產出比（原始）", "—", ["={D7}*{D8}"] * 3, "Derived", "Derived（含 Assumed）", "", fmt="0.0000")
add("D10", "每機架取得價", "$M", [1.0, 1.5, 2.2], "低 T12（$2M／144 顆）；基準 T10（KGI $1.4–1.6M 中點）；高＝KGI 上緣 1.6＋若晶片另計約 72×$8K（Assumed）", "賣方研究／Assumed（區間）", "")
add("D11", "每 GW 機架數（以基準功率）", "架", [f"={{A7}}*1000/$E${R['D4']}"] * 3, "Derived", "Derived", "")
add("D12", "每 GW IT 設備資本支出", "$B", ["={D11}*{D10}/1000*{A10}"] * 3, "Derived", "Derived", "")
add("D13", "k ＝ ÷ VR200 每 GW IT 資本支出", "—", ["={D12}/{A9}"] * 3, "Derived", "Derived", "", fmt="0.000")
add("D14", "Trainium3 每 GW 持有成本比（原始）", "—", ["=1-{A20}*(1-{D13})"] * 3, "Derived", "Derived（含 Assumed）", "", fmt="0.0000")
add("D15", "每 GW 顆數", "顆", ["={D11}*{D3}"] * 3, "Derived", "Derived", "")
add("D16", "GB300 每 FP8 PF 的 IT 資本支出（模型）", "$K/PF", same("=32.237224*1E6/(487008*5)"), "Interface!J9、J7；Spec_Rack!E40（FP8 5 PF）", "Derived（模型）", "")
add("D17", "交叉：Trn3 每 FP8 PF 資本支出 ÷ GB300", "—", [f"=({{D12}}*1E6/({ref('D15','E')}*{ref('D1','E')}))/{{D16}}"] * 3, "T13（SemiAnalysis：Trn3 每標稱 FP8 效能 TCO 比 GB300 好約 30%）", "Derived", "≈0.7 表示本表價格與 SemiAnalysis 說法相容；不進建議值", fmt="0.000")

# ---- E. 精度對齊旁證
sec("E. 精度對齊旁證（不進建議值）：若 VR200 也只以 FP8 服務")
add("E1", "TPU v7 方法一，VR200 改 FP8（比值 × A1÷A21）", "—", [f"={{C8}}*{ref('A1','E')}/{ref('A21','E')}"] * 3, "Derived", "Derived", "只用方法一（與 C16 一致）；f 仍取 C7；嚴格而言 f 須用 FP8 口徑重新還原（Hopper 會減半），此列只示意量級", fmt="0.0000")
add("E2", "Trainium3，VR200 改 FP8", "—", [f"={{D9}}*{ref('A1','E')}/{ref('A21','E')}"] * 3, "Derived", "Derived", "同上", fmt="0.0000")

# ---- F. 敏感度
sec("F. 敏感度（一次動一個參數，其餘取基準；D＝參數低值對應結果，F＝高值）")
add("F1", "TPU：方法一，f 取 低/基準/高", "產出比", [f"={ref('C6','E')}*{ref('C7','D')}", f"={ref('C8','E')}", f"={ref('C6','E')}*{ref('C7','F')}"], "Derived", "Derived", "", fmt="0.0000")
add("F2", "TPU：方法一，每顆功率 1.5625（冷卻上限 100 kW÷64）／1.085／0.99（若 10 MW 含 PUE 1.1）", "產出比", [f"={ref('C1','E')}/1.5625/{ref('A5','E')}*{ref('C7','E')}", f"={ref('C8','E')}", f"={ref('C1','E')}/(10000/1.1/9216)/{ref('A5','E')}*{ref('C7','E')}"], "T03／T02", "Derived", "", fmt="0.0000")
add("F3", "TPU：方法二（對照），u＝2.0／1.75／1.5", "產出比", [f"={ref('C12','E')}*({ref('C13','E')}/{ref('C4','E')})*{ref('C11','E')}/{ref('C14','D')}", f"={ref('C15','E')}", f"={ref('C12','E')}*({ref('C13','E')}/{ref('C4','E')})*{ref('C11','E')}/{ref('C14','F')}"], "Derived", "Derived", "", fmt="0.0000")
add("F4", "TPU：持有比，每顆功率 1.30／1.085／1.03 kW（價格取基準）", "持有比", [f"=1-{ref('A20','E')}*(1-({ref('A7','E')}*1000/{ref('C4','D')}*{ref('C17','E')}/1E6*{ref('A10','E')})/{ref('A9','E')})", f"={ref('C21','E')}", f"=1-{ref('A20','E')}*(1-({ref('A7','E')}*1000/{ref('C4','F')}*{ref('C17','E')}/1E6*{ref('A10','E')})/{ref('A9','E')})"], "Derived", "Derived", "功率↑→每 GW 顆數↓→資本支出↓", fmt="0.0000")
add("F5", "Trn3：f 0.50／0.80／1.04", "產出比", [f"={ref('D7','E')}*{ref('D8','D')}", f"={ref('D9','E')}", f"={ref('D7','E')}*{ref('D8','F')}"], "Derived", "Derived", "", fmt="0.0000")
add("F6", "Trn3：晶片 TDP 1.0／0.85／0.70 kW", "產出比", [f"={ref('D7','D')}*{ref('D8','E')}", f"={ref('D9','E')}", f"={ref('D7','F')}*{ref('D8','E')}"], "Derived", "Derived", "", fmt="0.0000")
add("F7", "Trn3：機架功率改用電源裝置 127.5 kW（5+1）", "產出比", [f"={ref('D1','E')}/(127.5/72)/{ref('A5','E')}*{ref('D8','E')}", f"={ref('D9','E')}", f"={ref('D9','E')}"], "T10", "Derived", "只有 D 欄有意義", fmt="0.0000")
add("F9", "Trn3：機架倍數 1.75／基準／1.35（TDP 取基準）", "產出比", [f"={ref('D1','E')}/({ref('D2','E')}*{ref('D3r','D')})/{ref('A5','E')}*{ref('D8','E')}", f"={ref('D9','E')}", f"={ref('D1','E')}/({ref('D2','E')}*{ref('D3r','F')})/{ref('A5','E')}*{ref('D8','E')}"], "Derived", "Derived", "r1 新增", fmt="0.0000")
add("F8", "Trn3：持有比，晶片 TDP 1.0／0.85／0.70 kW（價格取基準）", "持有比", [f"=1-{ref('A20','E')}*(1-({ref('A7','E')}*1000/{ref('D4','D')}*{ref('D10','E')}/1000*{ref('A10','E')})/{ref('A9','E')})", f"={ref('D14','E')}", f"=1-{ref('A20','E')}*(1-({ref('A7','E')}*1000/{ref('D4','F')}*{ref('D10','E')}/1000*{ref('A10','E')})/{ref('A9','E')})"], "Derived", "Derived", "", fmt="0.0000")

for col, w in zip("ABCDEFGHI", [8, 56, 10, 12, 12, 12, 50, 28, 60]):
    ws.column_dimensions[col].width = w
for r in range(5, row[0]):
    for c in (4, 5, 6):
        cell = ws.cell(r, c)
        if cell.number_format == "General" and ws.cell(r, 3).value in ("—", "PF/kW", "x", "kW"):
            cell.number_format = "0.0000"
    for c in (7, 8, 9):
        ws.cell(r, c).alignment = WRAP
ws.freeze_panes = "D5"

# ---------------------------------------------------------------- 建議值
sv = wb.create_sheet("建議值", 0)
sv["A1"] = "建議值 — NonNV 第 6、7 列（TPU v7、Trainium3）以每 kW 方法重推｜B 段唯讀查證，未寫入 model/；與 A 段 Maia／昇騰並排"
sv["A1"].font = BOLD
sv["A2"] = ("第 4–5 列＝推導頁公式的原始結果；第 8–9 列＝建議輸入值（四捨五入至 0.05，Trn3 產出比保留兩位小數）；第 12–13 列＝TK v21.5 現值。"
            "H＝F÷C（每 token 成本比）、I＝E÷D（最佳角落）、J＝G÷B（最差角落），同 NonNV。")
sv["A2"].alignment = WRAP
cols = ["候選", "產出比 低", "產出比 基準", "產出比 高", "持有比 低", "持有比 基準", "持有比 高", "每 token 成本比 基準", "最佳角落", "最差角落", "標記", "說明"]
for i, h in enumerate(cols, 1):
    c = sv.cell(3, i, h)
    c.font = BOLD
    c.fill = HDR


def ratio_cols(r):
    sv[f"H{r}"] = f"=F{r}/C{r}"
    sv[f"I{r}"] = f"=E{r}/D{r}"
    sv[f"J{r}"] = f"=G{r}/B{r}"


sv["A4"] = "Google TPU v7 Ironwood（原始公式結果）"
sv["A5"] = "AWS Trainium3（原始公式結果）"
for r, (o, h) in zip((4, 5), (("C16", "C21"), ("D9", "D14"))):
    for col, src in zip("BCD", "DEF"):
        sv[f"{col}{r}"] = f"=推導!{src}{R[o]}"
    for col, src in zip("EFG", "DEF"):
        sv[f"{col}{r}"] = f"=推導!{src}{R[h]}"
    ratio_cols(r)
    sv[f"L{r}"] = "公式（連結推導頁）"
sv["A7"] = "建議輸入值（若寫入 NonNV 第 6、7 列）"
sv["A7"].font = BOLD
sv["A8"] = "Google TPU v7 Ironwood"
sv["A9"] = "AWS Trainium3"
inp = {8: (0.25, 0.35, 0.40, 0.65, 0.70, 0.90), 9: (0.06, 0.14, 0.25, 0.45, 0.55, 0.70)}
note = {
    8: "每 kW 法：pod 9,216 顆約 10 MW（Google）→ 每顆 1.085 kW；FP8 4.61 PF、無 FP4，每 kW 峰值約 VR200 的 0.38；效率因子 Assumed（Hopper 與 Blackwell 還原值之間；InferenceX FP8 實測只作高端佐證）。持有比依每顆 $21–35K（Broadcom 訂單／SemiAnalysis／融資額）推導。證據：docs/evidence/20261025_NonNV_TPU_Trainium。",
    9: "每 kW 法：晶片 TDP 0.7–1.0 kW（KGI／DCD 依 AWS 副總裁液冷說法推論）× 72 顆 × TK 機架倍數 1.35–1.75 → 機架約 68–126 kW；FP8＝MXFP4 2.52 PF，每 kW 峰值約 VR200 的 0.18；無第三方實測，效率因子 Assumed。持有比依 T3 Max 機架 $1.0–2.2M（KGI）推導。",
}
for r in (8, 9):
    for j, v in enumerate(inp[r]):
        sv.cell(r, 2 + j, v).font = BLUE
    ratio_cols(r)
    sv[f"K{r}"] = "Assumed"
    sv[f"L{r}"] = note[r]
    sv[f"L{r}"].alignment = WRAP
sv["A11"] = "TK v21.5 現值（NonNV 第 6、7 列；慣例值、未推導）"
sv["A11"].font = BOLD
cur = {12: ("Google TPU v7 Ironwood（現值）", 0.6, 0.8, 1.1, 0.7, 0.8, 1.0), 13: ("AWS Trainium3（現值）", 0.3, 0.5, 0.8, 0.6, 0.7, 0.9)}
for r, vals in cur.items():
    sv.cell(r, 1, vals[0])
    for j, v in enumerate(vals[1:]):
        sv.cell(r, 2 + j, v)
    ratio_cols(r)
    sv[f"K{r}"] = "Assumed"
    sv[f"L{r}"] = "model/20261025_Tokenomics_v21.5.xlsx NonNV!B:G"

sv["A15"] = "與 A 段並排（每 kW 峰值算力比為各自最高原生精度；Maia、昇騰取 TK v21.5 NonNV 第 10–11 列現值）"
sv["A15"].font = BOLD
for i, h in enumerate(["候選", "每 kW 峰值 ÷ VR200（基準）", "產出比 基準", "持有比 基準", "每 token 成本比", "產出比 區間", "持有比 區間", "來源"], 1):
    c = sv.cell(16, i, h)
    c.font = BOLD
    c.fill = HDR
side = [
    ("Google TPU v7（本提案）", f"=推導!E{R['C6']}", "=C8", "=F8", 'TEXT(B8,"0.00")&"–"&TEXT(D8,"0.00")', 'TEXT(E8,"0.00")&"–"&TEXT(G8,"0.00")', "本表"),
    ("AWS Trainium3（本提案）", f"=推導!E{R['D7']}", "=C9", "=F9", 'TEXT(B9,"0.00")&"–"&TEXT(D9,"0.00")', 'TEXT(E9,"0.00")&"–"&TEXT(G9,"0.00")', "本表"),
    ("Microsoft Maia 200（A 段）", 0.676, 0.55, 0.50, "0.25–0.75", "0.40–0.60", "A 段推導 D6 基準 0.676；NonNV 第 11 列"),
    ("華為昇騰 910C（A 段）", 0.093, 0.08, 0.55, "0.05–0.15", "0.45–0.70", "A 段推導 C6 基準 0.093；NonNV 第 10 列"),
]
for k, (nm, pk, ob, hb, ol, hl, src) in enumerate(side):
    r = 17 + k
    sv.cell(r, 1, nm)
    sv.cell(r, 2, pk)
    sv.cell(r, 3, ob)
    sv.cell(r, 4, hb)
    sv.cell(r, 5, f"=D{r}/C{r}")
    sv.cell(r, 6, ("=" + ol) if ol.startswith("TEXT") else ol)
    sv.cell(r, 7, ("=" + hl) if hl.startswith("TEXT") else hl)
    sv.cell(r, 8, src)
    if not isinstance(pk, str):
        for cc in (2, 3, 4):
            sv.cell(r, cc).font = BLUE
for col, w in zip("ABCDEFGHIJKL", [36, 12, 11, 11, 11, 11, 11, 14, 10, 10, 10, 80]):
    sv.column_dimensions[col].width = w

# ---------------------------------------------------------------- 證據
ev = wb.create_sheet("證據")
eh = ["編號", "主張", "數值", "單位", "口徑", "出處網址", "來源等級", "立場", "標記", "取得方式與備註", "用於"]
for i, h in enumerate(eh, 1):
    c = ev.cell(1, i, h)
    c.font = BOLD
    c.fill = HDR
for rr, rowv in enumerate(EVIDENCE, 2):
    for cc, v in enumerate(rowv, 1):
        ev.cell(rr, cc, v).alignment = WRAP
for col, w in zip("ABCDEFGHIJK", [6, 52, 22, 12, 30, 56, 18, 10, 24, 52, 12]):
    ev.column_dimensions[col].width = w
ev.freeze_panes = "C2"

# ---------------------------------------------------------------- 待判斷
qd = wb.create_sheet("待判斷")
for i, h in enumerate(["編號", "題目", "預設（本提案採用）", "替代方案", "影響"], 1):
    c = qd.cell(1, i, h)
    c.font = BOLD
    c.fill = HDR
for rr, rowv in enumerate(QUESTIONS, 2):
    for cc, v in enumerate(rowv, 1):
        qd.cell(rr, cc, v).alignment = WRAP
for col, w in zip("ABCDE", [6, 50, 50, 50, 40]):
    qd.column_dimensions[col].width = w

wb.save(out)
print("saved", out, "rows", row[0])
