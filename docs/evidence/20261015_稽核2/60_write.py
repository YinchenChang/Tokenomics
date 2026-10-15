"""寫明細 docs/reports/20261015_稽核2_明細.xlsx（推導鏈、逐列檢查、傳遞強度、高嚴重度）。"""
import json,sys,pickle,openpyxl
from openpyxl.styles import Font,Alignment,PatternFill
S=sys.argv[1]
d=pickle.load(open(S+'/wb.pkl','rb'));c=d['cells']
v=lambda s,a:c[(s,a)]['v']
J=json.load(open('docs/evidence/20261015_稽核2/50_逐列判讀.json'))
S3=json.load(open('docs/evidence/20261015_稽核2/30_s3.json'))
wb=openpyxl.Workbook(); wb.remove(wb.active)
def sheet(name,rule,head,rows,widths):
    ws=wb.create_sheet(name); ws.append([rule]); ws['A1'].font=Font(bold=True); ws['A1'].alignment=Alignment(wrap_text=True,vertical='top'); ws.row_dimensions[1].height=90
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=min(len(head),12))
    ws.append(head)
    for cc in ws[2]: cc.font=Font(bold=True); cc.fill=PatternFill('solid',fgColor='FFDDEBF7'); cc.alignment=Alignment(wrap_text=True)
    for r in rows: ws.append(list(r))
    ws.auto_filter.ref=f'A2:{openpyxl.utils.get_column_letter(len(head))}{len(rows)+2}'
    ws.freeze_panes='C3'
    for i,w in enumerate(widths,1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    return ws
CH=[
('需求端','D1','API token（2025）','Alloc!C5','AL_DAPI','SRC_DEM_010（6 B tok/min，2025-10 時點）× 525,600 × AL_APIratio 0.75','M tok/年',v('Alloc','C5')),
('需求端','D2','ChatGPT token＝提示數 × 每則 token × 365','Alloc!C6','AL_DChat','SRC_DEM_012（2.5 B 則/日）× AL_TokPerPrompt 4,000 × 365','M tok/年',v('Alloc','C6')),
('需求端','D3','切分付費／免費（token 口徑）','Alloc!C7:C10','AL_DChatFree／AL_DChatPaid／AL_DPaid／AL_DFree','× AL_FreeShare 0.6（token 口徑）','M tok/年',v('Alloc','C9')),
('需求端','D4','付費混合有效單價（2026-09 快照）','Price_Frontier!C45:E45 → Cap_In!C90','B4_EffRef → IF_PriceRef_* → Cap_In C90','牌價 ×（1 − B4_Disc 0.2）× 參考請求 ISL／OSL、χ 0.55 混合；× 付費層級組合 B4_MixPaid（0.40／0.45／0.15）','$/M',v('Cap_In','C90')),
('需求端','D5','需求路線 m＝營收 ÷（付費 token × 單價）','Cap_In!C94 → L1!D76','L1_MImplied','SRC_DEM_007 13.07 $B × 1e9 ÷（AL_DPaid × C90）','x',v('L1','D76')),
('需求端','D5b','按收費模式分拆','Cap_In!C135:C146 → L1!D80:D82','L1_MAPI／L1_MSub／L1_MSeat','營收三分（訂閱 0.70、席次 0.10、API＝餘額）÷ token 三分 × 單價','x',f"{v('L1','D80'):.3f}／{v('L1','D81'):.3f}／{v('L1','D82'):.3f}"),
('機隊端','F1','每架產出（物理）','Perf!C98:Q98（Perf 第 10–98 列）','IF_TokGW_Luna／Sol／Astra（Interface 第 21／31／41 列）','規格（Spec_Rack）× 架構（Arch）× SLO 與校準（Serving、Calib）→ decode 批次 B* → 每請求 GPU 秒 → 每架 tok/s × 每 GW 機架數 × 8,760 h','M tok/GW/年（100%）',v('Interface','M31')),
('機隊端','F2','世代組合後每 GW 年產能','Alloc!C20:F25','AL_CapPaid／AL_BlendPaid／AL_BlendFree','IF_Util 0.6 ÷ Σ（層級組合 ÷ 各層級產出），再以 2025 世代組合（Hopper 0.6／GB200 0.4）加權','M tok/GW/年',v('Alloc','C24')),
('機隊端','F3','支出路線服務 GW','Alloc!C61','AL_ServeGWSpend','SRC_DEM_004 8.4 $B ÷ Σ（世代占比 × IF_HoldEcon 經濟口徑持有成本）','GW',v('Alloc','C61')),
('機隊端','F4','÷ 計價倍數 × κ','Cap_In!C109','AL_SpendBasis 1.3／AL_Kappa 1.0','服務 GW ÷ 1.3 × 1.0','GW',v('Cap_In','C109')),
('機隊端','F5','× 利用率修正','Cap_In!C114 → C177','AL_UtilActual 0.35 ÷ IF_Util 0.6','有效服務 GW','GW',v('Cap_In','C177')),
('機隊端','F6','機隊付費 token','Cap_In!C116','—','有效服務 GW ÷（1 ÷ 付費產能 ＋ r ÷ 免費產能），r＝AL_DFree ÷ AL_DPaid（取自需求端）','M tok/年',v('Cap_In','C116')),
('機隊端','F7','機隊路線 m＝營收 ÷（機隊付費 token × 單價）','Cap_In!C119 → L1!D78','L1_MFleet','同 D5 的分子與單價','x',v('L1','D78')),
('機隊端','F7b','m 下限（上限讀法）','Cap_In!C118, C121','L1_TokCeiling','以 CTL_ProdDerate 0.85 取代 0.35/0.6','x',v('Cap_In','C121')),
('相交','X1','兩路相交：同一分子（營收 13.07）、同一單價（C90），只差 token 分母','Cap_In!C171:C175','Checks M4','閉合比＝需求付費 token ÷ 機隊付費 token','x',v('Cap_In','C174')),
('相交','X2','缺口恆等式','Cap_In!C255:C256','L1_ScaleServe','ScaleServe ÷（計價倍數 × IF_Util ÷（實際利用率 × κ））＝1 ÷ 閉合比；驗算差恆為 0','x',v('Cap_In','C255')),
('相交','X3','落差分解（四項之積）','L1!D69:D74, D90','L1_GapPrompt／GapSpendBasis／GapISL／GapUtil／GapProduct','1.301 × 1.3 × 1.0 × 1.714＝2.900，對 ScaleServe2000 6.401','x',v('L1','D74')),
('相交','X4','閉合解（倒推）','Cap_In!C177:C181 → L1!D84','L1_TokPerPromptClose','令需求＝機隊有效產能，倒推每則 token τ*','tok',v('L1','D84')),
('回本','B1','損益兩平乘積（VR200）','L1!D52','L1_FleetBreakeven','IF_HoldEcon（VR200 基準）÷（IF_RevGWFleet ÷ IF_Util）','x',v('L1','D52')),
('回本','B2','回本所需 L（m＝m_fleet）','Cap_In!C240 → L1!D89','L1_LReqFleet','L1_FleetBreakeven ÷（IF_Util × CTL_ProdDerate × L1_MFleet）','x',v('L1','D89')),
('回本','B3','壽命期價格係數 L（自算）','Cap_In!C226:C234 → L1!D86:D88','L1_LPrice／L1_LDemote／L1_LTotal','g −10%／年、T1 4、ρ 0.5、T 6','x',v('L1','D88')),
('回本','B4','實收單價','L1!D91:D93','L1_PriceRealFleet／Demand／PriceReqBE','m × C90；回本所需＝FleetBreakeven ÷ 0.51 × C90 ÷ L_total','$/M',f"{v('L1','D91'):.3f}／{v('L1','D92'):.3f}／{v('L1','D93'):.3f}"),
('輸出','O1','× L × m 輸出（下游契約）','Interface!C236:Q271（_Life 列）','IFW_RevGW*_Life、IFW_RevGWFleet*_Life','Prod 列 × CTL_PriceLife（1）× CTL_Monetize（1）','$B/GW/年',v('Interface','M251')),
]
sheet('推導鏈','規則（S1）：從 m 相關的根往上游追的主幹。「需求端」＝提示數 × 每則 token × 單價 → 營收對照；「機隊端」＝推論支出 → GW → token 產能 → 可收費 token。兩路在「相交」列以同一分子與單價相交、以缺口恆等式閉合。值為 v5.36 檔內快取值（基準欄）。只描述，不判定。',
 ['側','步','內容','頁＋列','具名範圍','公式要旨','單位','值（基準）'],CH,[8,6,40,30,36,70,16,18])
rows=[]
for x in J:
    rows.append([x['sheet'],x['row'],x['label'],x['unit'],x['nform'],x['npat'],x['rep_cell'],str(x['rep_f'])[:300],x['rep_v'] if not isinstance(x['rep_v'],float) else round(x['rep_v'],6),
                 x['L1'],x['L2'],x['L3'],x['L4'],x['L5'],x['L6'],x['L7'],x['L8'],x['upY'],x['mode'],x['fids'],x['sev']])
sheet('逐列檢查','規則（S2）：第 0 段「C7_m路徑」561 個頁＋列逐列檢查 L1–L8。「—」＝沒有發現問題（不等於已證明正確）。L3 欄對位以 11_align.py 機械檢查（世代×層級／成本頁之間同欄對位 0 處不符；INDEX 常數欄號逐一對照列標籤）；L7 以 31_l7.py 在 4 個端點情境重算（文字或不可達被當成 0 的列）；L8「直接」＝本列公式直接引用第 1 段 Y1–Y9 所列輸入格，「上游」＝依賴圖上游任何一處含 Y 類輸入。F 編號對應報告第 3 節。嚴重度：高＝屬 F01–F08（直接作用在 m、缺口或回本判讀）；中＝其他 L1–L7；資訊＝只有 L8。只標問題，不判定對錯、不提修法。',
 ['頁','列','列標籤','單位','公式格數','同列公式結構數（R1C1）','代表格','代表公式','代表值','L1 量綱','L2 口徑','L3 引用對位','L4 重複／漏算','L5 閉合恆等式','L6 循環／倒推','L7 條件分支','L8 與 Y1–Y9 的連結（直接）','上游含 Y 類','傳遞方式','F 編號','嚴重度'],rows,
 [12,6,36,10,8,8,8,60,14,30,40,30,40,50,40,40,40,14,18,10,12])
METRICS=['L1_MImplied','L1_MFleet','L1_MAPI','L1_MSub','L1_MSeat','m 下限（Cap_In C121）','缺口＝機隊÷需求（1÷Cap_In C174）','L1_GapProduct','L1_LReqFleet','IFW_RevGWFleet_Life（VR200 基準）']
def _mx(r):
    sw=r.get('swing',{})
    return max([abs(v) for n in METRICS[:6] for v in (sw.get(n,[None,None])[:2]) if v is not None] or [0])
def _oth(r):
    sw=r.get('swing',{})
    return max([abs(v) for n in METRICS[6:] for v in (sw.get(n,[None,None])[:2]) if v is not None] or [0])
rank=sorted(S3['res'],key=lambda r:(-_mx(r),-_oth(r)))
rows=[]
for i,r in enumerate(rank,1):
    sw=r.get('swing',{})
    f=lambda n:'—' if n not in sw else '／'.join('—' if v is None else f'{v:+.1f}%' for v in sw[n][:2])
    mx=max([abs(v) for n in METRICS[:6] for v in (sw.get(n,[None,None])[:2]) if v is not None] or [0])
    rows.append([i,f"{r['sheet']}!{r['row']}",r['label'],r['cells'],r['Y'] or '（未歸 Y 類）',r.get('base',''),r.get('endpoint_lo',''),r.get('endpoint_hi',''),round(mx,1) if sw else '無法定量']+[f(n) for n in METRICS]+[r.get('note','')])
sheet('傳遞強度','規則（S3）：第 1 段嚴重度「高」且在 m 路徑上的 93 格，同一列合併為 46 列。單一輸入取 Gov_Map 區間端點（低／高）、其他不變，以 20_eval.py（以工作簿公式原文逐格求值；m 路徑 6,464 格對檔內快取值 0 不符；另以 LibreOffice 重算 3 個情境交叉驗證 0 不符）重算。數值為相對基準的變化 %（低端／高端）；「—」＝該情境下指標為文字（SLO 不可達）或無端點。世代組合三列的端點為「本世代 0%（其他依原比例放大）／100%」。排序依 m 指標（m_整體、m_fleet、m_API、m_訂閱、m_席次、m 下限）最大絕對變化；_Life 輸出受 CTL 基準 1 影響另列。不新增任何機制。',
 ['名次','頁＋列','列標籤','格','所屬 Y 類','基準值','低端','高端','m 指標最大 |變化|（%）']+[f'{n}（低／高）' for n in METRICS]+['備註'],rows,[6,14,30,14,10,16,20,20,12]+[16]*10+[40])
hi=[[x['sheet'],x['row'],x['label'],x['fids'],x['L1'],x['L2'],x['L3'],x['L4'],x['L5'],x['L6'],x['L7'],x['L8']] for x in J if x['sev']=='高']
sheet('高嚴重度','規則：「逐列檢查」頁嚴重度＝高的列（屬 F01–F08：直接作用在 m、缺口或回本判讀的邏輯問題）。只標問題，不判定對錯、不提修法。',
 ['頁','列','列標籤','F 編號','L1','L2','L3','L4','L5','L6','L7','L8'],hi,[12,6,36,10,30,50,30,50,50,40,40,40])
for ws in wb.worksheets:
    for row in ws.iter_rows(min_row=3):
        for cc in row: cc.alignment=Alignment(wrap_text=True,vertical='top')
wb.save('docs/reports/20261015_稽核2_明細.xlsx')
print('rows',len(J),'hi',len(hi),'s3',len(rows))
