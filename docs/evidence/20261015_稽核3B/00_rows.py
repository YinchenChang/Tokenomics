"""3B 列清單：第 0 段 C3_核心推導列 扣除 C7_m路徑（頁＋列），只取 3B 工作表。輸出 00_rows.json。"""
import openpyxl, json, collections
SH3B=['Training','Harness','L1','Alloc','Har_In','Cap_In','Train_In','Theory_Rev','Alloc_In']
wb=openpyxl.load_workbook('docs/reports/20261015_稽核0_明細.xlsx',read_only=True)
c3=[r for i,r in enumerate(wb['C3_核心推導列'].iter_rows(values_only=True)) if i>=2 and r[0]]
mp=set((r[1],r[2]) for i,r in enumerate(wb['C7_m路徑'].iter_rows(values_only=True)) if i>=2 and r[1])
rows=[dict(sh=r[0],r=r[1],label=r[2],nf=r[3],ni=r[4]) for r in c3 if r[0] in SH3B and (r[0],r[1]) not in mp]
cnt=collections.Counter(x['sh'] for x in rows)
print(len(rows),cnt)
json.dump(rows,open('docs/evidence/20261015_稽核3B/00_rows.json','w'),ensure_ascii=False,indent=0)
