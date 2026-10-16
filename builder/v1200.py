# v12.0 (work order docs/workorders/20261016_v11.1修補.md, branch claude/v11-1; H1–H7).
#   H1  Checks V1–V3 (MLPerf server soft caps) get a failure line: ratio <= 1 pass, 1 < ratio <= CHK_SoftCapFail warn, ratio > CHK_SoftCapFail fail. New blue input CHK_SoftCapFail = 2 (Calib M section; Assumed 1.5–3).
#   H2  CAL_BatchEtaD GB200 (Calib D156) 3.04 -> 2.36 = sqrt(1.83 * 3.04) (geometric mean of the MLPerf v6.0 offline back-solve 1.83 and the vLLM 10.1K back-solve 3.04); range 1.83–3.04. V20 becomes "校準相關（非獨立）" and leaves the verification count.
#   H3  Hopper CAL_BatchEtaD (Calib C156) stays 1 (now a blue input); marked "known underestimate, no solution at C68 = 1.5"; Gov_Map range 1–3.74.
#   H4  Calib C68 stays 1.5; correction note (no source; integer assumption since v5.2; the "倒推" wording of the audit-2 work order was wrong); Gov_Map range 1.0–1.5; Decisions correction record. Original wording is kept.
#   H5  IF_AllocRDGW / IF_AllocQ1 have no low/high columns: Alloc section N shows scenario values (constants written at build from LibreOffice scenario runs; the base row is live) + names AL_Rng*; tests guard the constants.
#   H6  texts: contract header, Calib H156 interval text, scenarios.yaml descriptions; IF_AmortDefault_Astra +0.9% traced (report).
#   H7  Checks V19 recomputes the verdict column T should show (from I, J, L, M and CST_MainMin) and compares it with T; V18 is marked an identity (not a verification) and leaves the count.
# Major version: new contract name CHK_SoftCapFail (CLAUDE.md section 4; work order H1). Every write to an Excel-owned cell is guarded (old value / old text).
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as L
from common import put, F_IN, F_NOTE, F_BOLD, F_CALC, F_LINK, FILL_KEY, section
import v1020
import v1030
import v1100

VERSION = "20261016_Tokenomics_v12.0"      # single source of the version string: README!B5 and README!A1 (finish.readme)
DATE = "2026-10-16"
V = "v12.0"
WO = "工作單 docs/workorders/20261016_v11.1修補.md"
DEC_BY = "Andy 2026-10-16 23:45「all ok」（H1–H5 預設）；H6、H7 為 chat 端預設，可撤回"
DASH = "—"
BATCH_GB200 = 2.36          # H2: sqrt(1.83 * 3.04) = 2.358, rounded to 3 significant digits
GB200_LO, GB200_HI = 1.83, 3.04       # the two back-solves (Calib L section V20 column C = 1.8297; Calib J section D column = 3.0373)
HOPPER_HI = 3.74            # Hopper multiplier at which the output saturates at C68 = 1.5 (0.941 of the target); upper end of the Gov_Map range
SOFT_FAIL = 2.0             # H1: CHK_SoftCapFail
SOFT_LO, SOFT_HI = 1.5, 3.0
C68_LO, C68_HI = 1.0, 1.5
SOFT_LABEL = "驗證軟上限的失敗門檻（CHK_SoftCapFail）"

# ================================================================== H5: scenario values (constants; filled from LibreOffice scenario runs, guarded by tests/parity)
# row: (label, GB200 multiplier, Hopper multiplier, C68, IF_AllocRDGW, IF_AllocQ1)
RANGE_ROWS = [      # values from LibreOffice scenario runs of the v12.0 workbook (tests/parity scenarios batch_gb200_lo／hi, batch_hopper_hi, batch_range_lo, c68_1)
    ("高端（GB200 1.83、Hopper 1）", 1.83, 1, 1.5, 0.211597, 0.190680),
    ("低端（GB200 3.04、Hopper 3.74）", 3.04, 3.74, 1.5, 0.162622, 0.153313),
    ("GB200 3.04、Hopper 1（v11.0 的基準）", 3.04, 1, 1.5, 0.197258, 0.180086),
    ("GB200 2.36、Hopper 3.74（只加 Hopper）", 2.36, 3.74, 1.5, 0.166368, 0.156293),
    ("C68＝1.0（GB200 2.36、Hopper 1；服務端 L1_MFleet 也會改變）", 2.36, 1, 1.0, 0.199873, 0.182039),
]


# ================================================================== H1: Calib M section (new blue input, created before restore())
def calib_m(wb):
    ws = wb["Calib"]
    r0 = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    section(ws, r0, "M. 驗證判準參數（v12.0 H1）：不是物理量，只決定 Checks V1–V3（MLPerf server 軟上限）何時由『警示』改判『失敗』", 10)
    r = r0 + 1
    put(ws, f"A{r}", SOFT_LABEL); put(ws, f"B{r}", "x")
    put(ws, f"C{r}", SOFT_FAIL, F_IN, fmt="0.0")
    put(ws, f"J{r}", "Assumed｜MLPerf server 的 TPOT 80 ms 為 p99 口徑，模型取平均，故模型略高於實測屬口徑差（V1–V3 比值 1 至本值之間判『警示』，不計失敗）；超過本值（基準 2 倍）表示模型高估，不再只是口徑差，判『失敗』。"
        "區間 1.5–3（Gov_Map；chat 端預設）。無來源；改本格只改 Checks V1–V3 的警示／失敗判定，不影響任何 IF_／L1_ 數值（CHK_VerifyPass／Fail／Warn 隨之改變）", F_NOTE, wrap=True)
    wb.defined_names["CHK_SoftCapFail"] = DefinedName("CHK_SoftCapFail", attr_text=f"Calib!$C${r}")
    return {"sec": r0, "row": r}


# ================================================================== H2: input write (after restore())
def inputs_update(wb):
    """Calib D156 (CAL_BatchEtaD GB200): 3.04 -> 2.36 (old-value guard)."""
    log = []
    ws = wb["Calib"]
    r = next(rr for rr in range(150, 170) if str(ws.cell(rr, 1).value or "").startswith("批次口徑 η_d 倍數（X10；CAL_BatchEtaD）"))
    assert r == 156, r
    if ws[f"D{r}"].value == 3.04:
        put(ws, f"D{r}", BATCH_GB200, F_IN, fmt="0.00"); log.append(f"Calib!D{r}: 3.04 -> {BATCH_GB200} (v12.0 H2)")
    return log


NOTE_156 = ("v12.0 H2：GB200 D156 由 3.04 改為 2.36＝√(1.83 × 3.04)（Derived；兩點幾何平均）。1.83＝MLPerf v6.0 offline GB200 6,752 tok/s/GPU 反推（Calib L 節 C 欄『達標所需批次口徑 η_d 倍數』，1.8297），"
            "3.04＝vLLM GB200 10.1K 反推（Calib J 節 D 欄，3.0373）；兩個點都是 NVIDIA／社群的利害關係方數據，無法判定何者代表真實批次口徑，取幾何平均並以 1.83–3.04 為區間（Gov_Map）。"
            "後果：V20（MLPerf offline GB200）與 vLLM 點同時被用於決定本格，V20 不再是獨立驗證點，Checks 標『校準相關（非獨立）』並移出驗證計數；Calib J 節 D 欄『模型 ÷ 目標』由 1.00 降為約 0.8（預期，本格已不再只對 vLLM 點）。"
            "Hopper C156 維持 1：已知低估、無解——在 C68＝1.5 下 η_d 再大，產出上限也只有 DeepSeek 剖析點目標的 0.941（C156＝3.74 即已飽和）；Gov_Map 區間 1–3.74（Decisions BB2-H3）")
NOTE_C68 = ("v12.0 H4 更正（Decisions BB2-H4）：上文『倒推（v5.5 建檔時與 C67 手算配對；Excel 內無公式關聯）』一語是稽核第 2 批工作單 J4（2026-10-16，d25bb75）誤寫，無證據："
            "v5.2–v5.12 各版工作簿的標記皆為 Assumed。正確標記：Assumed，無來源（v5.2 建檔以來的整數假設，非反推）；數值維持 1.5，區間 1.0–1.5（Gov_Map）；情境 c68_1 試算 C68＝1.0。原字樣保留，不刪改")
H156_OLD = "Derived（SRC_PERF_052／053／056／057；區間 0.70–1.0）"
H156_NEW = ("VR200（F156）：Derived（SRC_PERF_052／053／056／057；區間 0.70–1.0）｜GB200（D156）：Derived（兩點幾何平均 √(1.83 × 3.04)＝2.36；區間 1.83–3.04；SRC_PERF_074／079）｜"
            "Hopper（C156）：已知低估、無解（C68＝1.5 下 η_d 再大產出上限只有目標的 0.941；區間 1–3.74）｜GB300（E156）、Rubin Ultra（G156）：1（未校準）")


def notes_update(wb, append_text):
    log = {"j156": 0, "j68": 0, "h156": 0}
    ws = wb["Calib"]
    log["j156"] += append_text(ws["J156"], " ｜" + NOTE_156)
    log["j68"] += append_text(ws["J68"], " ｜" + NOTE_C68)
    if ws["H156"].value == H156_OLD:
        ws["H156"].value = H156_NEW; log["h156"] += 1
    return log


# ================================================================== Gov_Map rows for the new inputs (appended by gov.gm_append_c when absent)
def gov_map_rows(wb):
    ws = wb["Calib"]
    out = []
    r = next(rr for rr in range(150, ws.max_row + 1) if ws.cell(rr, 1).value == SOFT_LABEL)
    out.append({'scope': '切片四（v12.0 v11.1 修補）', 'sheet': 'Calib', 'cell': f'C{r}', 'label': SOFT_LABEL, 'check': SOFT_LABEL, 'cls': 'Assumed', 'role': '單值', 'src': '', 'rel': '',
                'dec': 'BB2-H1', 'lo': SOFT_LO, 'hi': SOFT_HI, 'rtext': f'{SOFT_LO}–{SOFT_HI}（chat 端預設）',
                'reason': ("Checks V1–V3（MLPerf server 軟上限）由『警示』改判『失敗』的比值門檻（基準 2）。MLPerf server 的 TPOT 80 ms 為 p99 口徑，模型取平均，略高於實測屬口徑差；超過 2 倍表示模型高估，不再只是口徑差。"
                           "非物理量，只決定檢查判定，不進任何 IF_／L1_ 數值。無來源；區間 1.5–3 為 chat 端預設。已搜尋：無可比對象（判定門檻不是可量測物理量）"), 'retag': '', 'seg': '低'})
    r = 156
    assert str(ws.cell(r, 1).value or "").startswith("批次口徑 η_d 倍數（X10；CAL_BatchEtaD）"), r
    out.append({'scope': '切片四（v12.0 v11.1 修補）', 'sheet': 'Calib', 'cell': f'C{r}', 'label': '批次口徑 η_d 倍數：Hopper（CAL_BatchEtaD）', 'check': ws.cell(r, 1).value, 'cls': 'Assumed', 'role': '單值',
                'src': '', 'rel': '', 'dec': 'BB2-H3', 'lo': 1, 'hi': HOPPER_HI, 'rtext': f'1–{HOPPER_HI}（下端＝未校準；上端＝C68＝1.5 下產出飽和處）',
                'reason': ("Hopper 批次口徑 η_d 倍數（只作用在 Perf_Batch）。數值維持 1（未校準）；已知低估、無解：在 C68＝1.5 下，η_d 再大，產出上限也只有 DeepSeek 剖析點目標 2,324 tok/s/GPU 的 0.941（超出 ±2%），"
                           "倍數 3.74 即已飽和；若 C68＝1.0 則需 4.32。MLPerf v6.0／TensorRT-LLM H200（V22）顯示 Hopper 批次口徑實際倍數約 3.3–3.7，故 1 是下界。"
                           "區間 1–3.74 用於呈現研發 GW 的方向（情境 batch_hopper_hi；Alloc N 節）：在 C68＝1.5 下這個值不會讓 Hopper 達標。無來源（Assumed）；待 Project 判斷是否改 C68"), 'retag': '', 'seg': '高'})
    return out


def gov_update(ws, append_text):
    n = 0
    r = v1020._gm_row(ws, "Calib", "D156")
    for col, old, new in (("K", 1.5, GB200_LO), ("L", 6.7, GB200_HI)):
        if ws[f"{col}{r}"].value == old:
            ws[f"{col}{r}"].value = new; n += 1
    old_rt = f'1.5–6.7（各批次口徑點反推值：V20 MLPerf offline、V8 LMSYS、校準點 3.04；見 Decisions BB-G2）'
    if ws[f"M{r}"].value == old_rt:
        ws[f"M{r}"].value = f'{GB200_LO}–{GB200_HI}（兩個批次口徑反推點：V20 MLPerf offline {GB200_LO}、vLLM 10.1K {GB200_HI}；基準＝幾何平均 {BATCH_GB200}；v12.0 H2）'; n += 1
    n += append_text(ws[f"N{r}"], " ｜v12.0 H2：基準由 3.04 改為 2.36＝√(1.83 × 3.04)（兩點幾何平均）；區間改為 1.83–3.04（原 1.5–6.7：上端 6.65 來自服務口徑的 V8，與批次倍數口徑不一致，下端 1.5 無對應資料點；查核 #89 觀察 6）。見 Decisions BB2-H2")
    r = v1020._gm_row(ws, "Calib", "C68")
    for col, old, new in (("K", 1, C68_LO), ("L", 3, C68_HI)):
        if ws[f"{col}{r}"].value == old:
            ws[f"{col}{r}"].value = new; n += 1
    if ws[f"M{r}"].value == "1–3（原註）":
        ws[f"M{r}"].value = f"{C68_LO}–{C68_HI}（v12.0 H4：下端＝無額外跨節點延遲；上端＝現值）"; n += 1
    if ws[f"O{r}"].value in (DASH, "", None):
        ws[f"O{r}"].value = "倒推 → Assumed（v12.0 H4 更正；類別欄原字樣『倒推』保留，見 Decisions BB2-H4）"; n += 1
    n += append_text(ws[f"N{r}"], " ｜v12.0 H4 更正：本格標記為 Assumed，無來源（v5.2 建檔以來的整數假設，非反推）；類別欄的『倒推』與 Calib J68 的『倒推（…手算配對）』是稽核第 2 批工作單 J4 誤寫，無證據，原字樣保留；"
                       "區間改為 1.0–1.5；情境 c68_1 試算。見 Decisions BB2-H4")
    return n


# ================================================================== H5: Alloc section N (scenario range of RD GW and Q1)
RANGE_TITLE = ("N. 區間（情境）：研發 GW（IF_AllocRDGW）與 Q1（IF_AllocQ1）對『批次口徑 η_d 倍數』與 C68 的區間（v12.0 H5）。第一列為即時值（公式）；其餘列為情境值＝建置時以 LibreOffice 重算寫入的常數"
               "（不是公式，改輸入後不會自動更新，由 tests/parity 以情境重算比對守護）。下游應把 IF_AllocRDGW、IF_AllocQ1 視為區間，不是點值")
RANGE_HDR = ["情境", "", "GB200 倍數（CAL_BatchEtaD[2]）", "Hopper 倍數（CAL_BatchEtaD[1]）", "C68（Hopper 每層延遲倍數）", "IF_AllocRDGW（GW·年）", "IF_AllocQ1（R1）"]


def alloc_range(wb, append_mark):
    """Alloc section N below the last used row. Returns the row map. Names: AL_RngRDGW_Lo／_Hi, AL_RngQ1_Lo／_Hi (low／high end of the range; display only)."""
    ws = wb["Alloc"]
    r0 = max(c.row for row in ws.iter_rows() for c in row if c.value is not None) + 2
    section(ws, r0, RANGE_TITLE, 7)
    r = r0 + 1
    for i, h in enumerate(RANGE_HDR):
        if h: put(ws, f"{L(i + 1)}{r}", h, F_BOLD, wrap=True)
    ws.row_dimensions[r].height = 44; r += 1
    R = {"sec": r0, "live": r}
    put(ws, f"A{r}", "基準（即時值；公式）")
    put(ws, f"C{r}", "=INDEX(CAL_BatchEtaD,1,2)", fmt="0.00"); put(ws, f"D{r}", "=INDEX(CAL_BatchEtaD,1,1)", fmt="0.00"); put(ws, f"E{r}", "=Calib!$C$68", fmt="0.00")
    put(ws, f"F{r}", "=AL_RDGW", fmt="0.0000", fill=FILL_KEY); put(ws, f"G{r}", "=AL_Q1", fmt="0.0000", fill=FILL_KEY)
    append_mark(ws, r, "顯示用：區間表基準列（即時值；與下方情境值並列）")
    r += 1
    for lab, gb, hp, c68, rd, q1 in RANGE_ROWS:
        put(ws, f"A{r}", lab, wrap=True)
        for col, v, fmt in (("C", gb, "0.00"), ("D", hp, "0.00"), ("E", c68, "0.00"), ("F", rd, "0.0000"), ("G", q1, "0.0000")):
            put(ws, f"{col}{r}", v, F_CALC, fmt=fmt)
        R[lab] = r; r += 1
    R["end"] = r - 1
    lo = next((R[l] for l, *_ in RANGE_ROWS if l.startswith("低端")), None)
    hi = next((R[l] for l, *_ in RANGE_ROWS if l.startswith("高端")), None)
    if lo and hi:
        for n, col, rr in (("AL_RngRDGW_Lo", "F", lo), ("AL_RngRDGW_Hi", "F", hi), ("AL_RngQ1_Lo", "G", lo), ("AL_RngQ1_Hi", "G", hi)):
            wb.defined_names[n] = DefinedName(n, attr_text=f"Alloc!${col}${rr}")
    return R


# ================================================================== end of build
def apply(wb, append_text):
    log = {"notes": notes_update(wb, append_text)}
    ws = wb["Calib"]
    sec = next(r for r in range(150, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("M. 驗證判準參數（v12.0 H1）"))
    log["range"] = alloc_range(wb, v1030._set_mark)
    log["decisions"] = decisions_append(wb)
    return log


# ================================================================== Decisions (10 columns; layout of v1100)
def _fmt(x):
    return f"{x:.4f}"


def _range_txt():
    if not RANGE_ROWS: return "（見 Alloc N 節）"
    return "；".join(f"{lab}：IF_AllocRDGW {_fmt(rd)}、IF_AllocQ1 {_fmt(q1)}" for lab, gb, hp, c68, rd, q1 in RANGE_ROWS)


_D = [
    ("H1", "MLPerf server 軟上限加失敗門檻 CHK_SoftCapFail（原則 4：檢查必須能失敗）",
     "查核 #89：V1–V3 的 F 欄公式在比值 > 1 時一律顯示『警示』，沒有第二道上限，比值為 10 也不會失敗。改為：比值 ≤ 1 通過；1 < 比值 ≤ CHK_SoftCapFail 警示；> CHK_SoftCapFail 失敗。"
     "新增具名輸入 CHK_SoftCapFail＝2（Calib M 節，藍字，Assumed，區間 1.5–3）：MLPerf server 的 TPOT 80 ms 為 p99 口徑，超過 2 倍表示模型高估，不再只是口徑差。恆等式 Pass＋Fail＋Warn＝Total 維持成立。"
     "更正舊描述：v11.0 報告、契約文件、CHANGELOG 寫『高於 1 但不超過上限＝警示』，實際上沒有上限，V1–V3 永遠不會失敗；現已有上限。新增專屬測試：把 V1 的比值以情境推到 2.5 判失敗、推到 1.5 判警示。"
     "新增契約名稱 CHK_SoftCapFail，故為主版 v12.0。",
     "Calib M 節；Checks V1–V3 的 F 欄公式與 V 節說明；下游契約文件", "v12.0 已建"),
    ("H2", "GB200 批次口徑倍數改用兩點幾何平均 2.36；V20 改為校準相關（非獨立）",
     f"v11.0 的 3.04 只由 vLLM 10.1K 一點校準，與獨立的 MLPerf v6.0 offline 不一致（V20 反推 1.83）。兩點各自來源：1.83＝MLPerf v6.0 offline GB200 6,752 tok/s/GPU 反推（Calib L 節 C 欄，1.8297）；3.04＝vLLM GB200 10.1K 反推（Calib J 節 D 欄，3.0373）；"
     f"兩個值都以 v11.0 的 Calib 重算確認。取幾何平均 √(1.83 × 3.04)＝2.358，寫入 D156＝2.36（Derived）；區間 1.83–3.04（Gov_Map）。新增情境 batch_gb200_lo（1.83）、batch_gb200_hi（3.04）。"
     f"V20 與 vLLM 點同時被用到，V20 不再是獨立驗證點：Checks 標『校準相關（非獨立）』，F 欄改顯示文字，移出驗證計數（若有更好的做法列入待 Project 判斷）。"
     f"影響：GB200 批次路徑 → Training → IF_AllocRDGW、IF_AllocQ1、AL_Q2、L1_ScaleRD 與 GB200 欄訓練鏈名稱；服務端（IF_TokGW_*、L1_MFleet、IF_AllocServeGW）、L1_LReqFleet_Id、L1_HoldEconMW_*、維護率不變。",
     "Calib D156、J156、H156；Gov_Map D156；Checks V20；Perf_Batch 第 53 列 → Training → Alloc", "v12.0 已建"),
    ("H3", "Hopper 批次口徑倍數維持 1：已知低估、無解；區間 1–3.74",
     "數值維持 1。標記改為『已知低估、無解』：在 C68＝1.5 下，η_d 再大，產出上限也只有 DeepSeek 剖析點目標的 0.941。Gov_Map 區間 1–3.74（新增列；C156 同時改為藍字輸入）。新增情境 batch_hopper_hi（3.74）："
     "在 C68＝1.5 下這個值不會讓 Hopper 達標，但能呈現研發 GW 的下限方向。獨立的 TensorRT-LLM H200 點（V22）在 Hopper 3.74 下由 0.50 升到 1.05，顯示 1 是下界。",
     "Calib C156、J156；Gov_Map 新列；情境 batch_hopper_hi", "v12.0 已建"),
    ("H4", "更正：C68 標記為 Assumed，無來源（原『倒推』一語無證據）",
     "數值維持 1.5。依 CLAUDE.md 第 6 節不得刪改既有的來源與標記欄，故原字樣保留：Gov_Map GM346 類別欄『倒推』、Calib J68 的『倒推（v5.5 建檔時與 C67 手算配對；Excel 內無公式關聯）』不動；"
     "更正做法＝在 Calib J68、Gov_Map 理由欄與『標記變更』欄加更正註記，並在本筆 Decisions 記錄：『倒推、與 C67 手算配對』一語出自稽核第 2 批工作單 J4（2026-10-16，d25bb75），v5.2–v5.12 各版工作簿標記皆為 Assumed，沒有證據；"
     "正確標記為 Assumed，說明『無來源（v5.2 建檔以來的整數假設）』。Gov_Map 區間 1.0–1.5（原 1–3）。新增情境 c68_1（C68＝1.0）。是否改 C68 仍待 Project 判斷（見 BB-G1）。",
     "Calib J68；Gov_Map GM346 的 K／L／M／N／O 欄；情境 c68_1", "v12.0 已建（更正註記）"),
    ("H5", "研發 GW 與 Q1 以區間呈現（Alloc N 節）",
     "IF_AllocRDGW、IF_AllocQ1 沒有低／高欄（L1_Ans1 的 Lo／Hi 只含 Alloc G 節一次動一個輸入的敏感度，不含批次口徑倍數），故在 Alloc 新增 N 節『區間（情境）』並建 AL_RngRDGW_Lo／_Hi、AL_RngQ1_Lo／_Hi（顯示用，下游不得連結）。"
     "低端＝GB200 3.04、Hopper 3.74；高端＝GB200 1.83、Hopper 1。情境值為建置時以 LibreOffice 重算寫入的常數（第一列基準為即時公式），tests/parity 以情境重算逐列比對。"
     "各情境值：" + _range_txt() + "。下游應視為區間，不是點值。",
     "Alloc N 節；AL_Rng*；情境 batch_gb200_lo／hi、batch_hopper_hi、batch_range_lo、c68_1", "v12.0 已建"),
    ("H6", "文字與解釋",
     "契約抬頭 v10.2 → v12.0；Calib H156 的區間字樣（原僅 VR200 的 0.70–1.0，與 D156 並列時自相矛盾）改為逐世代說明；scenarios.yaml 的過時說明更新。"
     "IF_AmortDefault_Astra 在 v11.0 的 GB200 欄上升 0.9% 的原因（追溯公式鏈，不改機制）：K6 預設 (c)＝由上而下總額（固定，3,904）× 該層級在自下而上總額中的占比（Amortize 第 47 列＝第 32 列 × 第 41 列）；"
     "GB200 批次倍數只壓低 RL rollout 成本，研發計畫成本 Luna −30.7%、Sol −23.7%、Astra −20.2%（Astra 的預訓練占比最高），自下而上總額 −20.9%，Astra 占比因此由 3.86% 升到 3.90%（+0.9%）。",
     "契約文件抬頭；Calib H156；tests/parity/scenarios.yaml；報告", "v12.0 已建"),
    ("H7", "V19 改為重算 T 欄應有判定再比對；V18 標恆等式（非驗證）",
     "查核 #87：V19 以 COUNTA 對 COUNTIF(…,\"一致\")，T 欄被清空或被手填『一致』時差值仍為 0。改為由公式自行重算每列應有的判定（I、J、L、M 欄與 CST_MainMin；不讀 T 公式的結果），與 T 欄逐列比對；"
     "結果＝不符列數＋實質不一致列數，為 0 才通過。V18：K 欄公式本身就是 J ≥ CST_MainMin，V18 重述同一條件，無法改成獨立於 K 的檢查，標『恆等式（非驗證）』並移出驗證計數。"
     "新增專屬測試：清空 T 欄一格，V19 必須判失敗。T06（狀態『早期採用』、J＝2 ≥ 門檻）的口徑差仍列待 Project 判斷。",
     "Checks V18、V19；CHK_Verify*", "v12.0 已建"),
]
DECISIONS_V1200 = [[f"BB2-{i}", V, f"v11.1 修補：{t}", txt, DATE, DEC_BY, st, where, f"{WO} {i}", "否"] for i, t, txt, where, st in _D]


def decisions_append(wb):
    ws = wb["Decisions"]
    have = {ws.cell(r, 1).value for r in range(5, ws.max_row + 1)}
    r = max([rr for rr in range(5, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, "")] or [4]) + 1
    n = 0
    for row in DECISIONS_V1200:
        if row[0] in have: continue
        for i, v in enumerate(row): put(ws, f"{L(i + 1)}{r}", v, F_CALC, wrap=i in (2, 3, 5, 7))
        r += 1; n += 1
    return n


# ================================================================== README
README_VERSION = (VERSION + "（主版：v11.1 修補，" + WO + "——H1 Checks V1–V3 軟上限加失敗門檻 CHK_SoftCapFail（新契約名稱）；H2 CAL_BatchEtaD GB200 2.36（兩點幾何平均，區間 1.83–3.04），V20 改校準相關；"
                  "H3 Hopper 維持 1（已知低估、無解，區間 1–3.74）；H4 C68 更正註記（Assumed，無來源，區間 1.0–1.5）；H5 研發 GW 與 Q1 區間（Alloc N 節）；H6 文字；H7 V19 重算比對、V18 標恆等式；Decisions BB2-*）。以下為 "
                  + v1100.README_VERSION.split("_Tokenomics_", 1)[1])
_T = v1100.README_TITLE
README_TITLE = V + _T[len(v1100.V):]
README_ROW = ("v11.1 修補（v12.0）",
              "驗證軟上限加失敗門檻 CHK_SoftCapFail；GB200 批次倍數 2.36（兩點幾何平均）、Hopper 維持 1（已知低估、無解）；C68 更正註記；研發 GW 與 Q1 區間（Alloc N 節）；V19 重算比對、V18 與 V20 移出驗證計數。細節見 Decisions BB2-*。")
