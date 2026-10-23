"""v19.1 W1 第二來源 B2（工作單 docs/workorders/20261020_W1第二來源_B2.md r1）專屬測試。讀 LibreOffice 重算後的值。

  1. Arch!C17＝512 × 4 ＋ 128＝2176；D17 仍 4224；SRC_MOD_106＝512、SRC_MOD_018＝1024（獨立重算）。
  2. 數值變動只出現在 Arch!C17 的下游：把 SRC_MOD_106 設回 1024（C17＝4224）時，所有既有具名範圍的值與封存的 v19.0 相同。
  3. 預期變動的關鍵名稱（新值）；L1_LReqFleet_IdNoK 不變。
  4. f_registry（Tech_Registry!O11:O13＝1）與 B4_CacheHit＝0 情境：整本活頁簿無錯誤值，C17 仍 2176。
  5. 記錄與治理：SRC_MOD_106、SRC_MOD_018 適用對象與 C1 免除、Gov_Map GM201、DB_Evidence E400、Decisions SL54-01、README。
  6. 逐格比較（對封存的 v19.0）：Arch 只有 C17、G17 改；SRC_Model 既有列只有 SRC_MOD_018 列與 X 欄範圍延伸。
"""
import re
from pathlib import Path

import openpyxl
import pytest
import yaml
from openpyxl.utils import range_boundaries

from conftest import model_path

REPO = Path(__file__).resolve().parents[2]
V190 = REPO / "model" / "archive" / "20261022_Tokenomics_v19.0.xlsx"
SC = {s["id"]: s["inputs"] for s in yaml.safe_load((Path(__file__).resolve().parent / "scenarios.yaml").read_text(encoding="utf-8"))["scenarios"]}
F_REGISTRY = SC["f_registry_t07_t09_on"]
CACHE_OFF = SC["d9_cache_cap_off"]
C17_OLD = {"SRC_MOD_106": 1024}      # 把新紀錄設回舊共用值 1024：C17＝4224，應回到 v19.0
PREWARM = {"test_v1910_stress_scenarios_no_error_values": [F_REGISTRY, CACHE_OFF], "test_v1910_changes_only_downstream_of_c17": [C17_OLD],
           "test_v1910_direction": [C17_OLD]}

NEW = {"SRC_MOD_106", "RNG_CHK_VerifyFail_Lo", "RNG_CHK_VerifyFail_Hi", "RNG_CHK_VerifyPass_Lo", "RNG_CHK_VerifyPass_Hi"}      # 新增名稱（區間端點檔多納入 2 個 CHK_ 名稱）
SKIP = ("RNG_", "IDX_", "GOV_", "IFC_")      # 區間端點（隨 Luna 值重算）、SRC_Index 範圍、治理計數另行檢查
VERSION_NAMES = {"RNG_Version", "RNG_Current"}


V20_CHANGED = {"L1_M_2026Q2_Lo", "L1_M_2026Q2_Hi", "L1_M_Lo", "L1_M_Hi"} | {f"IFW_{s}_LifeSelf" for s in ("RevGW_Luna", "RevGW_Sol", "RevGW_Astra", "RevGWFleet", "RevGWFront_Luna", "RevGWFront_Sol", "RevGWFront_Astra", "RevGWFleetFront")} | {"L1_PriceReqBE", "L1_PriceReqBE_Lo", "L1_PriceReqBE_Hi"}      # v20.0：H3 計數與 L1_M 的 2026 區間（test_v2000 看守）；v20.1：LifeSelf 與 L1_PriceReqBE 三格（test_v2010 看守）


def _flat(v):
    if isinstance(v, (list, tuple)):
        return [x for r in v for x in _flat(r)]
    return [v]


def _eq(a, b):
    if a in (None, "") and b in (None, ""):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= 1e-9 * max(abs(a), abs(b), 1e-300) + 1e-12
    return a == b


def test_v1910_c17_and_sources(lo):
    v = lo.view({})
    assert v.get_name("SRC_MOD_106") == 512 and v.get_name("SRC_MOD_018") == 1024
    assert v.get_name("SRC_MOD_019") == 4 and v.get_name("SRC_MOD_017") == 128
    assert v.get("Arch", "C17") == 512 * 4 + 128 == 2176
    assert v.get("Arch", "D17") == 1024 * 4 + 128 == 4224
    assert v.get("Arch", "E17") == 0
    assert v.get("Gov_Map", "Q205") == 2176      # GM201 直接引用 Arch!C17


# 預期變動的關鍵名稱：(名稱, 索引, v19.1 值)；v19.0 值在封存檔（下一項逐名稱比對）
KEY = [
    ("L1_MFleet", 0, 0.447574227389844),
    ("L1_RevGW_PQ_VR200", 0, 49.5137840363362),
    ("L1_LTotalFleetVDisc", 0, 0.758597367273996),
    ("L1_LTotalFleetVDisc_Lo", 0, 0.478526587147247),
    ("L1_TokLedger", 0, 5.61593255464987),
    ("IF_TokGWPaidMix", 0, 26017336543.4207),
    ("IF_LTotalFleetVDisc", 0, 0.758597367273996),
    ("L1_LReqFleet_IdNoK", 0, 0.756109635897205),      # 所需 L（不含 κ）與 Luna 跨度無關：不變（v19.0 同值）
]


@pytest.mark.parametrize("name,i,new", KEY)
def test_v1910_key_names(lo, name, i, new):
    got = _flat(lo.view({}).get_name(name))[i]
    assert abs(got / new - 1) <= 1e-9, (name, got, new)


def test_v1910_changes_only_downstream_of_c17(lo):
    """SRC_MOD_106＝1024（C17＝4224）時，所有既有數值型具名範圍等於封存的 v19.0：變動只來自 C17。"""
    pw = openpyxl.load_workbook(V190, data_only=True)
    v = lo.view(C17_OLD)
    assert v.get("Arch", "C17") == 4224
    bad, n = [], 0
    for name, d in pw.defined_names.items():
        if name.startswith(SKIP) or name in VERSION_NAMES or name in V20_CHANGED or "!" not in d.attr_text or "," in d.attr_text:
            continue
        sh, rg = d.attr_text.rsplit("!", 1)
        c1, r1, c2, r2 = range_boundaries(rg.replace("$", ""))
        want = [pw[sh.strip("'")].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        got = _flat(v.get_name(name))
        n += 1
        if not (len(want) == len(got) and all(_eq(a, b) for a, b in zip(want, got))):
            bad.append(name)
    assert n > 1400
    assert bad == [], bad[:20]


def test_v1910_direction(lo):
    """Luna 跨度上限縮短（4224 → 2176）：Luna 每 GW token 增加、m 下降。"""
    v, w = lo.view({}), lo.view(C17_OLD)
    assert _flat(v.get_name("IF_TokGW_Luna"))[0] > _flat(w.get_name("IF_TokGW_Luna"))[0]
    assert _flat(v.get_name("L1_MFleet"))[0] < _flat(w.get_name("L1_MFleet"))[0]
    assert _flat(v.get_name("L1_LReqFleet_IdNoK"))[0] == _flat(w.get_name("L1_LReqFleet_IdNoK"))[0]


def test_v1910_stress_scenarios_no_error_values(lo):
    for label, inp in (("f_registry", F_REGISTRY), ("cache_off", CACHE_OFF)):
        v = lo.view(inp)
        errs = [(s, c) for (s, c), x in v.evaluate_all().items() if isinstance(x, str) and x.startswith("#")]
        assert errs == [], (label, errs[:5])
        assert v.get("Arch", "C17") == 2176, label


def test_v1910_governance_counts(lo):
    v = lo.view({})
    assert (v.get_name("GOV_Errors"), v.get_name("GOV_Warnings"), v.get_name("GOV_Info")) == (0, 135, 182)      # v21.1: 133 → 135 (SRC_DEM_089/090 lack a second source, W1 +2)
    #
    assert v.get_name("RNG_StaleFlag") == "一致" and v.get_name("RNG_OutCount") == 0 and v.get_name("CHK_L1Order") == 0


# ---------------------------------------------------------------- static
@pytest.fixture(scope="module")
def wbf():
    return openpyxl.load_workbook(model_path())


def _rows(ws, col, val):
    return [r for r in range(5, ws.max_row + 1) if ws.cell(r, col).value == val]


def test_v1910_records(wbf):
    ws = wbf["SRC_Model"]
    r106, r18 = _rows(ws, 1, "SRC_MOD_106"), _rows(ws, 1, "SRC_MOD_018")
    assert len(r106) == 1 and len(r18) == 1
    a, b = r106[0], r18[0]
    g = lambda r, c: ws[f"{c}{r}"].value
    C1 = "免（C1 官方規格；Andy 2026-10-20 17:17）"
    assert g(a, "C") == 512 and g(a, "F") == "塊" and g(a, "H") == "DeepSeek V4-Flash" and g(a, "O") == "Active" and g(a, "K") == 2 and g(a, "R") == C1
    assert "index_topk" in g(a, "J") and "vLLM" in g(a, "J") and g(a, "Q") == "E400" and g(a, "V").startswith("Arch!C17")
    assert g(b, "C") == 1024 and g(b, "H") == "DeepSeek V4-Pro" and g(b, "R") == C1 and g(b, "V").startswith("Arch!D17")
    assert "『DeepSeek V4』" in g(b, "W") and "『Arch!C17｜換算』" in g(b, "W") and "SL54-01" in g(b, "W")      # 修訂九：原值逐字保留
    ar = wbf["Arch"]
    assert ar["C17"].value == "=SRC_MOD_106*SRC_MOD_019+SRC_MOD_017" and ar["D17"].value == "=SRC_MOD_018*SRC_MOD_019+SRC_MOD_017" and ar["E17"].value == 0
    assert "512" in ar["G17"].value and "1,024" in ar["G17"].value
    gm = wbf["Gov_Map"]
    rr = [r for r in range(5, gm.max_row + 1) if gm[f"C{r}"].value == "Arch" and gm[f"D{r}"].value == "C17"]
    assert len(rr) == 1 and gm[f"H{rr[0]}"].value == "SRC_MOD_106" and "SRC_MOD_018" in gm[f"N{rr[0]}"].value
    de = wbf["DB_Evidence"]
    re_ = _rows(de, 1, "E400")
    assert len(re_) == 1 and de[f"M{re_[0]}"].value == "SRC_MOD_106" and "512" in de[f"H{re_[0]}"].value and "4,224" in de[f"G{re_[0]}"].value
    assert len(_rows(de, 1, "E399")) == 1      # v18.0 的 E399 仍在，未被覆蓋
    d = wbf["Decisions"]
    rd = _rows(d, 1, "SL54-01")
    assert len(rd) == 1 and d.cell(rd[0], 2).value == "v19.1" and d.cell(rd[0], 7).value == "v19.1 已建" and d.cell(rd[0], 10).value == "否"
    assert "17:29" in d.cell(rd[0], 6).value and "SRC_MOD_106" in d.cell(rd[0], 4).value and "修訂九" in d.cell(rd[0], 4).value
    assert d.max_row < 1000
    rm = wbf["README"]
    assert "v21.0（主版：" in str(rm["B5"].value) and "v20.1（次版，底版 v20.0：" in str(rm["B5"].value) and "v20.0（主版：" in str(rm["B5"].value) and "v19.1（次版，底版 v19.0：" in str(rm["B5"].value) and "v19.0（主版：" in str(rm["B5"].value) and str(rm["A1"].value).startswith("Tokenomics v21.")
    assert len([x for x in range(4, rm.max_row + 1) if rm.cell(x, 1).value == "v19.1 W1 第二來源 B2"]) == 1


def test_v1910_diffs_only_expected_vs_v190():
    a, b = openpyxl.load_workbook(V190), openpyxl.load_workbook(V190.parent.parent / 'archive' / '20261022_Tokenomics_v19.1.xlsx')      # v20.0：現行檔已是 v20.0，本檢驗改比封存的 v19.1
    assert a.sheetnames == b.sheetnames
    d = {}
    for ws in a:
        wn = b[ws.title]
        for r in range(1, max(ws.max_row, wn.max_row) + 1):
            for c in range(1, max(ws.max_column, wn.max_column) + 1):
                x, y = ws.cell(r, c).value, wn.cell(r, c).value
                if x != y:
                    d.setdefault(ws.title, []).append((ws.cell(r, c).coordinate, x, y))
    # Arch：只有 C17（公式）與 G17（說明文字）
    assert sorted(t[0] for t in d["Arch"]) == ["C17", "G17"]
    # 其餘模型頁的公式／輸入不變（其餘為治理與範圍延伸）
    allowed = {"Arch", "SRC_Model", "SRC_Index", "Gov_Map", "Alloc", "Checks", "README", "Decisions", "DB_Evidence"}
    assert set(d) <= allowed, set(d) - allowed
    # SRC_Model 既有列：除 SRC_MOD_018 列（H、R、V、W）外，只有 X 欄（SUMPRODUCT 範圍 $158 → $159）與新增列
    sm = a["SRC_Model"]
    for coord, x, y in d["SRC_Model"]:
        r = int(re.sub(r"[A-Z]+", "", coord)); col = re.sub(r"\d+", "", coord)
        if x is None:
            continue      # 新增列（SRC_MOD_106）與其 X–AH 公式
        if col == "X":
            assert str(y) == str(x).replace("$158", "$159"), (coord, x, y)
        else:
            assert sm.cell(r, 1).value == "SRC_MOD_018" and col in ("H", "R", "V", "W"), (coord, x, y)
    # 名稱：新增 SRC_MOD_106 與 RNG_CHK_Verify*；其餘只有 IDX_／RNG_／GOV_ 範圍延伸
    na, nb = set(a.defined_names), set(b.defined_names)
    assert na <= nb and nb - na == NEW, (nb - na)
    changed = {k for k in na if a.defined_names[k].attr_text != b.defined_names[k].attr_text}
    assert all(k.startswith(("IDX_", "RNG_", "GOV_", "IFC_")) for k in changed), sorted(changed)[:10]      # v21.0：IFC_ 契約欄範圍延伸
